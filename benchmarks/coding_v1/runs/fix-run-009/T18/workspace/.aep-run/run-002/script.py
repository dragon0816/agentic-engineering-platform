from pathlib import Path
import subprocess
import sys
import csv
from collections import defaultdict
from openpyxl import load_workbook

ws = Path.cwd()
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)

fixed_code = r'''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals: dict[str, float] = {}
    counts: dict[str, int] = {}

    with open(args.input, newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            model = row["Model"]
            value_text = row["Value"]
            if value_text == "":
                continue
            value = float(value_text)
            totals[model] = totals.get(model, 0.0) + value
            counts[model] = counts.get(model, 0) + 1

    output_path = Path(args.output)
    if output_path.parent:
        output_path.parent.mkdir(parents=True, exist_ok=True)

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(output_path)


if __name__ == "__main__":
    main()
'''

export_path = out_dir / 'export.py'
export_path.write_text(fixed_code, encoding='utf-8', newline='\n')
print('WROTE', export_path)

cmd = [sys.executable, str(export_path), '--input', str(ws/'input'/'input.csv'), '--output', str(out_dir/'summary.xlsx')]
print('RUNNING', ' '.join(cmd))
completed = subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('exit_code:', completed.returncode)
if completed.stdout:
    print('stdout:\n' + completed.stdout)
if completed.stderr:
    print('stderr:\n' + completed.stderr)
if completed.returncode != 0:
    raise SystemExit(completed.returncode)

# Independent expected result from CSV, excluding blank Value.
expected_totals = defaultdict(float)
expected_counts = defaultdict(int)
with (ws/'input'/'input.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('SOURCE HEADERS:', reader.fieldnames)
    for row in reader:
        if row['Value'] == '':
            continue
        expected_totals[row['Model']] += float(row['Value'])
        expected_counts[row['Model']] += 1
expected_rows = [(m, expected_totals[m] / expected_counts[m], expected_counts[m]) for m in sorted(expected_totals)]
print('EXPECTED ROWS:', expected_rows)

# Reopen and validate the saved workbook.
xlsx_path = out_dir / 'summary.xlsx'
if not xlsx_path.exists():
    raise AssertionError('summary.xlsx was not created')
wb = load_workbook(xlsx_path, data_only=True)
if wb.sheetnames != ['summary']:
    raise AssertionError(f'Unexpected sheet names: {wb.sheetnames}')
wsheet = wb['summary']
actual_rows = list(wsheet.iter_rows(values_only=True))
print('WORKBOOK ROWS:', actual_rows)
header = actual_rows[0]
if header != ('Model', 'Average', 'ValidCount'):
    raise AssertionError(f'Unexpected header: {header}')
actual_data = actual_rows[1:]
if len(actual_data) != len(expected_rows):
    raise AssertionError(f'Unexpected data row count: {len(actual_data)} != {len(expected_rows)}')
for actual, expected in zip(actual_data, expected_rows):
    amodel, aavg, acount = actual
    emodel, eavg, ecount = expected
    if amodel != emodel:
        raise AssertionError(f'Model mismatch: {amodel!r} != {emodel!r}')
    if not isinstance(aavg, (int, float)):
        raise AssertionError(f'Average is not numeric for {amodel}: {aavg!r}')
    if abs(float(aavg) - eavg) > 1e-9:
        raise AssertionError(f'Average mismatch for {amodel}: {aavg} != {eavg}')
    if acount != ecount or not isinstance(acount, int):
        raise AssertionError(f'ValidCount mismatch/type for {amodel}: {acount!r} != {ecount}')

# Confirm input files were not changed enough? We cannot know prior hash here, but only wrote output.
print('VALIDATION PASSED: output/export.py exists; output/summary.xlsx has sheet summary, headers Model/Average/ValidCount, numeric averages, integer ValidCount, and expected row count/results.')