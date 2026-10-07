from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urljoin
import urllib.request
import urllib.error
import json
import re
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:50351/t13/list'
XLSX_PATH = OUTPUT / 'robust_products.xlsx'
WARNINGS_PATH = OUTPUT / 'warnings.json'
HEADERS = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']
NA = 'N/A'

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.links = []
        self._current_href = None
        self._current_text = []
        self._in_products_depth = 0

    def handle_starttag(self, tag, attrs):
        attrs_d = dict(attrs)
        if tag == 'ul' and attrs_d.get('id') == 'products':
            self._in_products_depth = 1
        elif self._in_products_depth:
            self._in_products_depth += 1
        if tag == 'a' and self._in_products_depth:
            self._current_href = attrs_d.get('href')
            self._current_text = []

    def handle_data(self, data):
        if self._current_href is not None:
            self._current_text.append(data)

    def handle_endtag(self, tag):
        if tag == 'a' and self._current_href is not None:
            text = ' '.join(' '.join(self._current_text).split())
            if self._current_href and text:
                self.links.append((text, self._current_href))
            self._current_href = None
            self._current_text = []
        if self._in_products_depth:
            self._in_products_depth -= 1

class DetailParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.stack = []  # each item: [tag, attrs, text_parts]
        self.elements = []
        self.all_text_parts = []
        self.rows = []
        self.current_row_cells = None
        self.current_cell = None

    def handle_starttag(self, tag, attrs):
        attrs_d = {k.lower(): (v or '') for k, v in attrs}
        self.stack.append([tag.lower(), attrs_d, []])
        if tag.lower() == 'tr':
            self.current_row_cells = []
        if tag.lower() in ('td', 'th') and self.current_row_cells is not None:
            self.current_cell = []

    def handle_data(self, data):
        if data and data.strip():
            self.all_text_parts.append(data)
        for item in self.stack:
            item[2].append(data)
        if self.current_cell is not None:
            self.current_cell.append(data)

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ('td', 'th') and self.current_cell is not None and self.current_row_cells is not None:
            cell = ' '.join(''.join(self.current_cell).split())
            self.current_row_cells.append(cell)
            self.current_cell = None
        if tag == 'tr' and self.current_row_cells is not None:
            if self.current_row_cells:
                self.rows.append(self.current_row_cells)
            self.current_row_cells = None
        if self.stack:
            item = self.stack.pop()
            elem_tag, attrs, texts = item
            text = ' '.join(''.join(texts).split())
            if text:
                self.elements.append({'tag': elem_tag, 'attrs': attrs, 'text': text})

def fetch_url(url, retry_500=False, max_attempts=3):
    attempts = 0
    last_status = None
    last_error = None
    while True:
        attempts += 1
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'AEP-product-collector/1.0'})
            with urllib.request.urlopen(req, timeout=10) as resp:
                body = resp.read().decode('utf-8', errors='replace')
                return {'ok': True, 'status': resp.status, 'body': body, 'attempts': attempts, 'error': None}
        except urllib.error.HTTPError as e:
            last_status = e.code
            try:
                e.read()
            except Exception:
                pass
            last_error = f'HTTP {e.code}'
            if retry_500 and e.code == 500 and attempts < max_attempts:
                continue
            return {'ok': False, 'status': last_status, 'body': '', 'attempts': attempts, 'error': last_error}
        except Exception as e:
            last_error = f'{type(e).__name__}: {e}'
            return {'ok': False, 'status': last_status, 'body': '', 'attempts': attempts, 'error': last_error}

def parse_list(html):
    p = ProductListParser()
    p.feed(html)
    return p.links[:5]

def first_non_na(*values):
    for value in values:
        if value is not None and str(value).strip() and str(value).strip() != NA:
            return str(value).strip()
    return NA

def clean_value(value):
    if value is None:
        return NA
    value = re.sub(r'\s+', ' ', str(value)).strip()
    value = re.sub(r'^[\s:：\-=]+|[\s;]+$', '', value)
    return value or NA

def parse_price(value):
    value = clean_value(value)
    if value == NA:
        return NA
    # Strip thousands separators and common currency decorations, keep numeric part.
    m = re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', value)
    if not m:
        return NA
    num = m.group(0).replace(',', '')
    if '.' in num:
        try:
            return float(num)
        except ValueError:
            return NA
    try:
        return int(num)
    except ValueError:
        return NA

def currency_from_text(value):
    value = clean_value(value)
    if value == NA:
        return NA
    # Prefer explicit ISO-style currency codes.
    for code in re.findall(r'\b[A-Z]{3}\b', value):
        if code not in {'CPU', 'RAM', 'ROM'}:
            return code
    symbols = {'$': 'USD', '€': 'EUR', '£': 'GBP', '¥': 'JPY', 'NT$': 'TWD'}
    for sym, code in symbols.items():
        if sym in value:
            return code
    return NA

def parse_detail(html, fallback_name):
    parser = DetailParser()
    parser.feed(html)

    by_key = {}
    attr_candidates = []
    for elem in parser.elements:
        attrs = elem['attrs']
        keys = ' '.join([attrs.get('id', ''), attrs.get('class', ''), attrs.get('name', ''), attrs.get('data-field', ''), attrs.get('data-key', '')]).lower()
        text = clean_value(elem['text'])
        if keys and text != NA:
            attr_candidates.append((keys, text))
        for attr_key, attr_val in attrs.items():
            low = attr_key.lower()
            if low.startswith('data-') or low in ('price', 'currency', 'cpu', 'name'):
                by_key[low.replace('data-', '')] = clean_value(attr_val)

    # Two-column tables/lists like Field | Value.
    for row in parser.rows:
        if len(row) >= 2:
            key = clean_value(row[0]).lower().rstrip(':：')
            val = clean_value(row[1])
            if key and val != NA:
                by_key[key] = val

    full_text = '\n'.join(' '.join(parser.all_text_parts).split('\n'))
    compact_text = re.sub(r'\s+', ' ', full_text)

    def find_by_label(labels):
        # attribute or table candidates
        for label in labels:
            if label in by_key:
                return by_key[label]
        for key, val in by_key.items():
            for label in labels:
                if label in key:
                    return val
        for keys, text in attr_candidates:
            for label in labels:
                if label in keys:
                    return text
        # textual labels, e.g. "CPU: Snapdragon"
        for label in labels:
            pattern = re.compile(rf'\b{re.escape(label)}\b\s*[:：\-]?\s*([^\n;|]+)', re.IGNORECASE)
            m = pattern.search(full_text)
            if m:
                val = clean_value(m.group(1))
                # trim if subsequent fields got concatenated
                val = re.split(r'\s+(?:Name|Price|Currency|CPU|Processor)\s*[:：\-]', val, flags=re.I)[0]
                return clean_value(val)
        return NA

    name = find_by_label(['name', 'product name', 'title'])
    if name == NA:
        # h1/title-like elements first, else list name.
        for elem in parser.elements:
            keys = ' '.join([elem['tag'], elem['attrs'].get('id', ''), elem['attrs'].get('class', '')]).lower()
            if elem['tag'] in ('h1', 'h2') or 'name' in keys or 'title' in keys:
                name = clean_value(elem['text'])
                break
    name = first_non_na(name, fallback_name)

    price_raw = find_by_label(['price'])
    price = parse_price(price_raw)
    currency = first_non_na(find_by_label(['currency']), currency_from_text(price_raw))
    cpu = find_by_label(['cpu', 'processor', 'chipset'])

    # If labels are absent but text has recognizable forms.
    if price == NA:
        price = parse_price(compact_text)
    if currency == NA:
        currency = currency_from_text(compact_text)

    return {'Name': name, 'Price': price, 'Currency': currency, 'CPU': clean_value(cpu)}

# Fetch and parse list.
list_result = fetch_url(LIST_URL, retry_500=False)
if not list_result['ok']:
    raise RuntimeError(f'Could not fetch product list: {list_result["error"]}')
products = parse_list(list_result['body'])
print(f'Found {len(products)} products in list: {products}')
if len(products) < 5:
    raise RuntimeError(f'Expected five products from list, found {len(products)}')

rows = []
warnings = []
for list_name, href in products[:5]:
    detail_url = urljoin(LIST_URL, href)
    result = fetch_url(detail_url, retry_500=True, max_attempts=3)
    if result['ok']:
        fields = parse_detail(result['body'], list_name)
        row = {
            'Name': fields['Name'],
            'Price': fields['Price'] if fields['Price'] != '' else NA,
            'Currency': fields['Currency'] if fields['Currency'] != '' else NA,
            'CPU': fields['CPU'] if fields['CPU'] != '' else NA,
            'URL': detail_url,
            'FetchStatus': 'ok',
        }
        warnings.append({'Name': row['Name'], 'URL': detail_url, 'FetchStatus': 'ok', 'attempts': result['attempts'], 'http_status': result['status']})
        print('OK', detail_url, row)
    else:
        row = {'Name': list_name, 'Price': NA, 'Currency': NA, 'CPU': NA, 'URL': detail_url, 'FetchStatus': 'error'}
        warnings.append({'Name': list_name, 'URL': detail_url, 'FetchStatus': 'error', 'attempts': result['attempts'], 'http_status': result['status'], 'error': result['error']})
        print('ERROR', detail_url, result)
    rows.append(row)

# Write workbook.
wb = Workbook()
ws = wb.active
ws.title = 'Products'
ws.append(HEADERS)
for row in rows:
    ws.append([row[h] for h in HEADERS])
wb.save(XLSX_PATH)

with WARNINGS_PATH.open('w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Reopen outputs and verify stated requirements.
check_wb = load_workbook(XLSX_PATH, data_only=True)
check_ws = check_wb['Products']
headers = [cell.value for cell in check_ws[1]]
assert headers == HEADERS, headers
actual_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
assert len(actual_rows) == 5, len(actual_rows)
for idx, values in enumerate(actual_rows, start=2):
    record = dict(zip(HEADERS, values))
    assert record['Name'] not in (None, ''), f'row {idx} missing Name'
    assert record['URL'] not in (None, ''), f'row {idx} missing URL'
    assert record['FetchStatus'] in ('ok', 'error'), f'row {idx} bad FetchStatus {record["FetchStatus"]}'
    for field in ['Price', 'Currency', 'CPU']:
        assert record[field] not in (None, ''), f'row {idx} missing placeholder/value for {field}'
    if record['FetchStatus'] == 'error':
        assert record['Price'] == NA and record['Currency'] == NA and record['CPU'] == NA, f'row {idx} error fields not N/A'
    else:
        # Price may be numeric or N/A; other unknowns use N/A.
        assert record['FetchStatus'] == 'ok'

with WARNINGS_PATH.open('r', encoding='utf-8') as f:
    warn_check = json.load(f)
assert isinstance(warn_check, list), 'warnings.json must be a list'
assert len(warn_check) == 5, len(warn_check)
for item in warn_check:
    assert item.get('FetchStatus') in ('ok', 'error'), item
    assert item.get('URL'), item
    if item.get('FetchStatus') == 'ok':
        assert item.get('attempts', 0) >= 1, item
    else:
        assert item.get('attempts', 0) <= 3, item

print('\nWrote:', XLSX_PATH)
print('Wrote:', WARNINGS_PATH)
print('Verified workbook headers, 5 rows, non-empty Name/URL/status, N/A placeholders for missing/error fields, and warnings.json statuses/attempts.')