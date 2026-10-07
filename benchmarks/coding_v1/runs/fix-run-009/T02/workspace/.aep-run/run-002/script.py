from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import TYPE_STRING

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read source CSV, preserving column names and row order.
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    if 'Status' not in headers:
        raise ValueError('Required column Status not found')
    failed_rows = [row for row in reader if (row.get('Status') or '').strip().casefold() == 'fail']

# Write workbook with requested sheet name.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

for col_idx, header in enumerate(headers, start=1):
    ws.cell(row=1, column=col_idx, value=header)

for row_idx, row in enumerate(failed_rows, start=2):
    for col_idx, header in enumerate(headers, start=1):
        value = row.get(header, '')
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        if header == 'Device ID':
            cell.value = '' if value is None else str(value)
            cell.data_type = 's'

# Also set text number format for Device ID column when present.
if 'Device ID' in headers:
    device_col = headers.index('Device ID') + 1
    for r in range(2, len(failed_rows) + 2):
        ws.cell(row=r, column=device_col).number_format = '@'

wb.save(output_path)
print('Wrote:', output_path)
print('Filtered rows:', len(failed_rows))

# Verification by reopening saved output.
check_wb = load_workbook(output_path, data_only=False)
if check_wb.sheetnames != ['failed_items']:
    raise AssertionError(f'Unexpected sheets: {check_wb.sheetnames}')
check_ws = check_wb['failed_items']
actual_headers = [check_ws.cell(row=1, column=c).value for c in range(1, check_ws.max_column + 1)]
if actual_headers != headers:
    raise AssertionError(f'Headers changed: {actual_headers} != {headers}')
if check_ws.max_row != len(failed_rows) + 1:
    raise AssertionError(f'Row count mismatch: {check_ws.max_row} != {len(failed_rows)+1}')

# Compare all values and confirm every output status matches rule.
for i, source_row in enumerate(failed_rows, start=2):
    output_row = {header: check_ws.cell(row=i, column=headers.index(header)+1).value for header in headers}
    expected_row = {header: ('' if source_row.get(header) is None else str(source_row.get(header, ''))) for header in headers}
    # openpyxl reads empty string cells back as None, normalize only blanks for comparison.
    normalized_output = {k: ('' if v is None else str(v)) for k, v in output_row.items()}
    if normalized_output != expected_row:
        raise AssertionError(f'Row {i} mismatch: {normalized_output} != {expected_row}')
    if normalized_output['Status'].strip().casefold() != 'fail':
        raise AssertionError(f'Row {i} Status does not match fail rule')

if 'Device ID' in headers:
    device_col = headers.index('Device ID') + 1
    device_values = []
    for r in range(2, check_ws.max_row + 1):
        cell = check_ws.cell(row=r, column=device_col)
        device_values.append((cell.value, cell.data_type, cell.number_format))
        expected = failed_rows[r-2]['Device ID']
        if cell.value != expected:
            raise AssertionError(f'Device ID leading zeros not preserved at row {r}: {cell.value!r} != {expected!r}')
        if cell.data_type != 's':
            raise AssertionError(f'Device ID is not text at row {r}: data_type={cell.data_type!r}')
    print('Device ID checks:', device_values)

print('Verification passed: sheet name, headers, row count, filter rule, original order/all values, and text Device ID with leading zeros.')
