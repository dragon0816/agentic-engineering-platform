from pathlib import Path
from html.parser import HTMLParser
from urllib.request import urlopen
from urllib.parse import urljoin
import re
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
URL = 'http://127.0.0.1:50351/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_catalog = False
        self.catalog_depth = 0
        self.current_card = None
        self.current_text_target = None
        self.products = []

    def handle_starttag(self, tag, attrs):
        attrs_d = dict(attrs)
        if tag == 'main' and attrs_d.get('id') == 'catalog':
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if self.in_catalog:
            self.catalog_depth += 1
            cls = attrs_d.get('class', '')
            classes = cls.split()
            if tag == 'div' and 'card' in classes:
                self.current_card = {'name': '', 'price_text': '', 'href': ''}
            elif self.current_card is not None and tag == 'a':
                self.current_card['href'] = attrs_d.get('href', '')
                self.current_text_target = 'name'
            elif self.current_card is not None and tag == 'span' and 'price' in classes:
                self.current_text_target = 'price_text'

    def handle_data(self, data):
        if self.in_catalog and self.current_card is not None and self.current_text_target:
            self.current_card[self.current_text_target] += data

    def handle_endtag(self, tag):
        if self.in_catalog:
            if self.current_card is not None:
                if tag == 'a' and self.current_text_target == 'name':
                    self.current_text_target = None
                elif tag == 'span' and self.current_text_target == 'price_text':
                    self.current_text_target = None
                elif tag == 'div':
                    if self.current_card.get('name') and self.current_card.get('price_text') and self.current_card.get('href'):
                        self.products.append(self.current_card)
                    self.current_card = None
                    self.current_text_target = None
            self.catalog_depth -= 1
            if self.catalog_depth <= 0:
                self.in_catalog = False
                self.catalog_depth = 0

def parse_price(price_text):
    s = price_text.strip()
    m = re.match(r'^([^\d\-+.,]*)([-+]?\d[\d,]*(?:\.\d+)?)', s)
    if not m:
        raise ValueError(f'Cannot parse price: {price_text!r}')
    currency = m.group(1).strip()
    number_text = m.group(2).replace(',', '')
    value = float(number_text) if '.' in number_text else int(number_text)
    return currency, value

with urlopen(URL, timeout=10) as resp:
    html = resp.read().decode('utf-8')

parser = CatalogParser()
parser.feed(html)

rows = []
for product in parser.products:
    currency, price = parse_price(product['price_text'])
    rows.append({
        '產品名稱': product['name'].strip(),
        '價格': price,
        '幣別': currency,
        '產品網址': urljoin(URL, product['href'].strip()),
    })

wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])

# keep URL as full text URL; price cells numeric
for cell in ws['D'][1:]:
    cell.style = 'Hyperlink'
wb.save(OUTFILE)
print(f'wrote {OUTFILE}')
print(f'rows written: {len(rows)}')
for r in rows:
    print(r)

# Verification by reopening the saved workbook
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
actual_headers = [ws2.cell(1, c).value for c in range(1, 5)]
assert actual_headers == headers, f'headers mismatch: {actual_headers}'
data_rows = []
for rr in range(2, ws2.max_row + 1):
    values = [ws2.cell(rr, c).value for c in range(1, 5)]
    if any(v is not None for v in values):
        data_rows.append(values)
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
assert [r[0] for r in data_rows] == expected_names, [r[0] for r in data_rows]
assert len(data_rows) == 5, len(data_rows)
for name, price, currency, url in data_rows:
    assert isinstance(name, str) and name
    assert isinstance(price, (int, float)) and not isinstance(price, bool), (name, price, type(price))
    assert currency == 'NT$', (name, currency)
    assert isinstance(url, str) and url.startswith('http://127.0.0.1:50351/'), (name, url)
    assert '/ad/' not in url and not name.startswith('推薦') and 'Sponsored' not in name, (name, url)
assert [r[1] for r in data_rows] == [1299, 2450, 99, 560, 12000]
print('verification ok: reopened products.xlsx; headers, 5 rows, page order, numeric prices, NT$ currency, full URLs, and no ad/recommendation rows confirmed')
