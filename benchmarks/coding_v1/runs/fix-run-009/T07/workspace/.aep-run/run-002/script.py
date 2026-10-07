from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

ws = Path.cwd()
input_path = ws / 'input' / 'specs.md'
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)
out_path = out_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split product sections beginning with level-2 headings. Only parse written specs.
sections = re.split(r'(?m)^##\s+', text)
product_sections = []
for sec in sections[1:]:
    if sec.strip():
        product_sections.append(sec)

unit_factor_to_mhz = {
    'mhz': 1.0,
    'ghz': 1000.0,
    'khz': 0.001,
    'hz': 0.000001,
}

def fmt_num(x):
    if x is None:
        return None
    # Keep numeric values numeric in workbook; avoid text IDs only applies Model.
    return float(x) if abs(x - round(x)) > 1e-9 else int(round(x))

def parse_number_with_unit(value):
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', value, flags=re.I)
    if not m:
        return None
    return float(m.group(1)) * unit_factor_to_mhz[m.group(2).lower()]

def parse_frequency_range(value):
    # Require a written numeric range and a written unit. Do not infer missing specs.
    v = value.replace('–', '-').replace('—', '-').replace('−', '-')
    # Common form: 2.4-2.5 GHz or 900-930 MHz
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*-\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', v, flags=re.I)
    if m:
        factor = unit_factor_to_mhz[m.group(3).lower()]
        return fmt_num(float(m.group(1)) * factor), fmt_num(float(m.group(2)) * factor)
    # More explicit form: 2.4 GHz - 2.5 GHz
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\s*-\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', v, flags=re.I)
    if m:
        return (
            fmt_num(float(m.group(1)) * unit_factor_to_mhz[m.group(2).lower()]),
            fmt_num(float(m.group(3)) * unit_factor_to_mhz[m.group(4).lower()]),
        )
    return None, None

def parse_interface(value):
    # Multiple interfaces separated by semicolon. Slash in source denotes multiple entries.
    parts = [p.strip() for p in re.split(r'[/,;]+|\band\b|\+|、|，', value) if p.strip()]
    return ';'.join(parts) if parts else ''

rows = []
for sec in product_sections:
    specs = {}
    for line in sec.splitlines():
        m = re.match(r'\s*-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            specs[m.group(1).strip()] = m.group(2).strip()
    freq_min, freq_max = (None, None)
    if 'Frequency' in specs and specs['Frequency']:
        freq_min, freq_max = parse_frequency_range(specs['Frequency'])
    bw = None
    if 'Bandwidth' in specs and specs['Bandwidth']:
        bw = fmt_num(parse_number_with_unit(specs['Bandwidth']))
    interface = ''
    if 'Interface' in specs and specs['Interface']:
        interface = parse_interface(specs['Interface'])
    rows.append({
        'Model': specs.get('Model', ''),
        'FrequencyMinMHz': freq_min if freq_min is not None else '',
        'FrequencyMaxMHz': freq_max if freq_max is not None else '',
        'BandwidthMHz': bw if bw is not None else '',
        'Interface': interface,
    })

headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
wb = Workbook()
wsheet = wb.active
wsheet.title = 'specs'
wsheet.append(headers)
for row in rows:
    wsheet.append([row[h] for h in headers])
wb.save(out_path)

# Verification by reopening saved workbook.
wb2 = load_workbook(out_path, data_only=True)
sh = wb2['specs']
actual_headers = [c.value for c in sh[1]]
actual_rows = list(sh.iter_rows(min_row=2, values_only=True))
print('WROTE', out_path)
print('HEADERS', actual_headers)
print('ROW_COUNT', len(actual_rows))
for r in actual_rows:
    print('ROW', r)

# Requirement checks
assert actual_headers == headers, f'headers mismatch: {actual_headers}'
assert len(actual_rows) == 3, f'expected 3 product rows, got {len(actual_rows)}'
expected = [
    ('RX-100', 2400, 2500, 80, 'USB;LAN'),
    ('RX-200', 900, 930, 20, 'LAN'),
    ('RX-300', 5000, 6000, None, 'USB'),
]
assert actual_rows == expected, f'rows mismatch: {actual_rows!r}'
# Value type checks: model/interface text, numeric converted MHz where present, missing blank reopens as None.
for row in actual_rows:
    model, fmin, fmax, bw, iface = row
    assert isinstance(model, str)
    assert isinstance(fmin, (int, float)) and isinstance(fmax, (int, float))
    assert bw is None or isinstance(bw, (int, float))
    assert isinstance(iface, str)
assert actual_rows[2][3] is None, 'missing Bandwidth should be blank/None'
print('CHECKS PASSED: specs.xlsx exists, sheet specs, expected headers, 3 rows, MHz conversions, semicolon interfaces, missing Bandwidth blank, numeric MHz values and text IDs verified.')