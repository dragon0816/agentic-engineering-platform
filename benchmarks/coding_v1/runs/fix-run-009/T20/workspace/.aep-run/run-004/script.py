from pathlib import Path
import csv
import os
import stat
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
THRESHOLD = 20

print(f'Workspace: {WORKSPACE}')
print(f'Input dir: {INPUT_DIR}')
print(f'Output dir: {OUTPUT_DIR}')

if not INPUT_DIR.exists():
    raise FileNotFoundError(f'Input folder not found: {INPUT_DIR}')

print('Input files:')
for p in sorted(INPUT_DIR.rglob('*')):
    kind = 'DIR' if p.is_dir() else f'FILE {p.stat().st_size} bytes'
    print(f' - {p.relative_to(WORKSPACE)} {kind}')

sop_path = INPUT_DIR / 'supplied_sop.md'
if not sop_path.exists():
    raise FileNotFoundError('supplied_sop.md not found under input/')
print('\nSOP preview:')
print(sop_path.read_text(encoding='utf-8', errors='replace')[:1200])

csv_files = sorted(INPUT_DIR.rglob('*.csv'))
if not csv_files:
    raise FileNotFoundError('No CSV file found under input/')
results_named = [p for p in csv_files if p.name.lower() == 'results.csv']
if results_named:
    csv_path = results_named[0]
elif len(csv_files) == 1:
    csv_path = csv_files[0]
else:
    raise RuntimeError('Multiple CSV files found and no results.csv; cannot choose one unambiguously: ' + ', '.join(str(p.relative_to(WORKSPACE)) for p in csv_files))
print(f'Using CSV: {csv_path.relative_to(WORKSPACE)}')

required_columns = ['Device ID', 'Status', 'Value']
rows = []
with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    print(f'CSV headers: {reader.fieldnames}')
    if reader.fieldnames is None:
        raise RuntimeError('CSV has no header row')
    missing = [c for c in required_columns if c not in reader.fieldnames]
    if missing:
        raise RuntimeError(f'CSV missing required columns: {missing}')
    for row in reader:
        rows.append(row)
print(f'Total rows read: {len(rows)}')

fail_rows = []
for row in rows:
    status = str(row['Status']).strip()
    value_raw = row['Value']
    try:
        value_num = float(value_raw)
        value_out = int(value_num) if value_num.is_integer() else value_num
    except Exception as exc:
        raise RuntimeError(f"Value is not numeric for Device ID {row.get('Device ID')}: {value_raw!r}") from exc
    if status == 'Fail':
        fail_rows.append({
            'Device ID': str(row['Device ID']),
            'Status': status,
            'Value': value_out,
            'Urgent': bool(value_num > THRESHOLD),
        })

summary = {
    'Total': len(rows),
    'Pass': sum(1 for r in rows if str(r['Status']).strip() == 'Pass'),
    'Fail': len(fail_rows),
    'Urgent': sum(1 for r in fail_rows if r['Urgent']),
}
print(f'Summary: {summary}')

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
failures_path = OUTPUT_DIR / 'failures.xlsx'
summary_path = OUTPUT_DIR / 'summary.md'
tmp_failures_path = OUTPUT_DIR / '~failures.tmp.xlsx'
tmp_summary_path = OUTPUT_DIR / '~summary.tmp.md'

# Clean leftovers from a previous failed run. If a target workbook is open/locked, report it clearly.
for p in [tmp_failures_path, tmp_summary_path, failures_path, summary_path]:
    if p.exists():
        try:
            os.chmod(p, stat.S_IWRITE | stat.S_IREAD)
            p.unlink()
            print(f'Removed old output: {p.relative_to(WORKSPACE)}')
        except PermissionError as exc:
            raise RuntimeError(f'Cannot replace {p}. Please close it if it is open in Excel/another app, then rerun the workflow.') from exc

wb = Workbook()
ws = wb.active
ws.title = 'Failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for r in fail_rows:
    ws.append([r['Device ID'], r['Status'], r['Value'], r['Urgent']])
wb.save(tmp_failures_path)
os.replace(tmp_failures_path, failures_path)

summary_lines = [
    '# Test Failure Summary',
    '',
    f'- Total: {summary["Total"]}',
    f'- Pass: {summary["Pass"]}',
    f'- Fail: {summary["Fail"]}',
    f'- Urgent: {summary["Urgent"]}',
    '',
    f'Urgent rule: Status is Fail and Value > {THRESHOLD}.',
]
tmp_summary_path.write_text('\n'.join(summary_lines) + '\n', encoding='utf-8')
os.replace(tmp_summary_path, summary_path)

print('\nVerifying saved outputs...')
if not failures_path.exists():
    raise RuntimeError('failures.xlsx was not created')
if not summary_path.exists():
    raise RuntimeError('summary.md was not created')

check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['Failures']
actual_headers = [cell.value for cell in next(check_ws.iter_rows(min_row=1, max_row=1))]
if actual_headers != headers:
    raise RuntimeError(f'Unexpected failures.xlsx headers: {actual_headers}')

saved_rows = []
for row in check_ws.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    saved_rows.append(row)
if len(saved_rows) != len(fail_rows):
    raise RuntimeError(f'Expected {len(fail_rows)} Fail rows in failures.xlsx, found {len(saved_rows)}')

for idx, (saved, expected) in enumerate(zip(saved_rows, fail_rows), start=2):
    dev, status, value, urgent = saved
    if dev != expected['Device ID'] or status != 'Fail':
        raise RuntimeError(f'Unexpected row {idx}: {saved}')
    if not isinstance(value, (int, float)):
        raise RuntimeError(f'Value in row {idx} is not numeric: {value!r}')
    if not isinstance(urgent, bool):
        raise RuntimeError(f'Urgent in row {idx} is not boolean: {urgent!r}')
    expected_urgent = bool(float(value) > THRESHOLD)
    if urgent != expected_urgent:
        raise RuntimeError(f'Urgent in row {idx} should be {expected_urgent}, got {urgent}')

summary_text = summary_path.read_text(encoding='utf-8')
for k, v in summary.items():
    token = f'- {k}: {v}'
    if token not in summary_text:
        raise RuntimeError(f'summary.md missing {token!r}')

print(f'Wrote: {failures_path}')
print(f'Wrote: {summary_path}')
print('Verified: output file names, workbook sheet/columns, Fail-only row count, numeric Value cells, boolean Urgent cells using Value > 20, and summary counts.')
