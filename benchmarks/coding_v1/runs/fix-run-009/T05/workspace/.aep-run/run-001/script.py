from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
input_dir = base / 'input'
print('Workspace:', base)
print('Input files:')
for p in sorted(input_dir.iterdir()):
    print(' -', p.name, 'size=', p.stat().st_size)

for filename in ['devices.xlsx', 'tests.xlsx']:
    path = input_dir / filename
    print('\n===', filename, '===')
    wb = load_workbook(path, data_only=True, read_only=False)
    print('Sheets:', wb.sheetnames)
    for ws in wb.worksheets:
        print('Sheet:', ws.title, 'max_row=', ws.max_row, 'max_col=', ws.max_column)
        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
        print('Headers:', headers)
        for r in range(2, min(ws.max_row, 6)+1):
            vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column+1)]
            print('Row', r, ':', vals)
