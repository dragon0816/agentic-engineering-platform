from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font

base = Path.cwd()
input_dir = base / 'input'
output_dir = base / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
out_path = output_dir / 'merged_devices_tests.xlsx'


def read_sheet(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    rows = []
    for r in range(2, ws.max_row + 1):
        rec = {}
        for c, h in enumerate(headers, start=1):
            value = ws.cell(r, c).value
            if h == 'Device ID' and value is not None:
                value = str(value)
            rec[h] = value
        if any(v is not None for v in rec.values()):
            rows.append(rec)
    return headers, rows


def write_rows(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for rec in rows:
        ws.append([rec.get(h) for h in headers])
    # format Device ID column(s) as text, preserving leading zeros already in values
    for col_idx, h in enumerate(headers, start=1):
        if h == 'Device ID':
            for row_idx in range(1, ws.max_row + 1):
                ws.cell(row_idx, col_idx).number_format = '@'
                if row_idx > 1 and ws.cell(row_idx, col_idx).value is not None:
                    ws.cell(row_idx, col_idx).value = str(ws.cell(row_idx, col_idx).value)
    for col_idx, h in enumerate(headers, start=1):
        max_len = len(str(h)) if h is not None else 0
        for row_idx in range(2, ws.max_row + 1):
            v = ws.cell(row_idx, col_idx).value
            max_len = max(max_len, len(str(v)) if v is not None else 0)
        ws.column_dimensions[ws.cell(1, col_idx).column_letter].width = min(max_len + 2, 40)


dev_headers, devices = read_sheet(input_dir / 'devices.xlsx')
test_headers, tests = read_sheet(input_dir / 'tests.xlsx')

if 'Device ID' not in dev_headers or 'Device ID' not in test_headers:
    raise ValueError('Both input files must contain a Device ID column')

# Validate the stated condition: each devices ID is unique.
seen = set()
duplicates = []
for rec in devices:
    did = rec.get('Device ID')
    if did in seen:
        duplicates.append(did)
    seen.add(did)
if duplicates:
    raise ValueError(f'Duplicate Device ID values found in devices.xlsx: {duplicates}')

dev_by_id = {rec.get('Device ID'): rec for rec in devices}
tests_by_id = {}
for rec in tests:
    tests_by_id.setdefault(rec.get('Device ID'), []).append(rec)

# matched: one row per matched test, with test columns followed by device-only columns.
matched_headers = list(test_headers) + [h for h in dev_headers if h not in test_headers]
matched_rows = []
for test in tests:
    did = test.get('Device ID')
    dev = dev_by_id.get(did)
    if dev is not None:
        merged = {}
        for h in test_headers:
            merged[h] = test.get(h)
        for h in dev_headers:
            if h not in test_headers:
                merged[h] = dev.get(h)
        matched_rows.append(merged)

# unmatched: devices with no tests, and tests with no device, in a union schema.
unmatched_headers = ['Source'] + list(dict.fromkeys(dev_headers + test_headers))
unmatched_rows = []
for dev in devices:
    did = dev.get('Device ID')
    if did not in tests_by_id:
        row = {'Source': 'devices'}
        row.update(dev)
        unmatched_rows.append(row)
for test in tests:
    did = test.get('Device ID')
    if did not in dev_by_id:
        row = {'Source': 'tests'}
        row.update(test)
        unmatched_rows.append(row)

wb_out = Workbook()
ws = wb_out.active
ws.title = 'matched'
write_rows(ws, matched_headers, matched_rows)
ws_un = wb_out.create_sheet('unmatched')
write_rows(ws_un, unmatched_headers, unmatched_rows)
wb_out.save(out_path)

print(f'Wrote: {out_path}')
print(f'Devices rows: {len(devices)}')
print(f'Tests rows: {len(tests)}')
print(f'Matched rows: {len(matched_rows)}')
print(f'Unmatched rows: {len(unmatched_rows)}')
print('Matched headers:', matched_headers)
print('Unmatched headers:', unmatched_headers)

# Verification by reopening the saved file.
check_wb = load_workbook(out_path, data_only=False)
assert check_wb.sheetnames == ['matched', 'unmatched'], check_wb.sheetnames
matched_ws = check_wb['matched']
unmatched_ws = check_wb['unmatched']

matched_read_headers = [matched_ws.cell(1, c).value for c in range(1, matched_ws.max_column + 1)]
unmatched_read_headers = [unmatched_ws.cell(1, c).value for c in range(1, unmatched_ws.max_column + 1)]
assert matched_read_headers == matched_headers, matched_read_headers
assert unmatched_read_headers == unmatched_headers, unmatched_read_headers
assert matched_ws.max_row - 1 == 3, matched_ws.max_row - 1
assert unmatched_ws.max_row - 1 == 2, unmatched_ws.max_row - 1
assert 'Source' in unmatched_read_headers

# Check Device ID values and text formatting.
for ws_check in [matched_ws, unmatched_ws]:
    id_cols = [c for c, h in enumerate([ws_check.cell(1, c).value for c in range(1, ws_check.max_column + 1)], start=1) if h == 'Device ID']
    assert id_cols, f'{ws_check.title} missing Device ID'
    for id_col in id_cols:
        assert ws_check.cell(1, id_col).number_format == '@'
        for r in range(2, ws_check.max_row + 1):
            v = ws_check.cell(r, id_col).value
            if v is not None:
                assert isinstance(v, str), (ws_check.title, r, v, type(v))
                assert ws_check.cell(r, id_col).number_format == '@'

matched_ids = [matched_ws.cell(r, matched_read_headers.index('Device ID') + 1).value for r in range(2, matched_ws.max_row + 1)]
unmatched_sources = [unmatched_ws.cell(r, unmatched_read_headers.index('Source') + 1).value for r in range(2, unmatched_ws.max_row + 1)]
unmatched_ids = [unmatched_ws.cell(r, unmatched_read_headers.index('Device ID') + 1).value for r in range(2, unmatched_ws.max_row + 1)]
assert matched_ids == ['001', '001', '002'], matched_ids
assert unmatched_sources == ['devices', 'tests'], unmatched_sources
assert unmatched_ids == ['003', '999'], unmatched_ids

print('Verification passed: sheets, headers, row counts, Source values, text Device ID formatting, and leading-zero IDs were checked from the saved workbook.')
