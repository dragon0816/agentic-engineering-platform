from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
print('workspace:', workspace)
print('input exists:', input_path.exists(), input_path)
print('output dir exists:', (workspace / 'output').exists())

wb = load_workbook(input_path, read_only=True, data_only=False)
print('sheetnames:', wb.sheetnames)
ws = wb['tasks']
print('tasks dimensions:', ws.max_row, ws.max_column)
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print('headers:', headers)
print('sample rows:')
for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 6), values_only=True):
    print(row)
wb.close()