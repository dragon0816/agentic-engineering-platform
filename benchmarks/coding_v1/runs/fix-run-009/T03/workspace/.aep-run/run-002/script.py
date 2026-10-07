from pathlib import Path
from collections import OrderedDict, Counter
from copy import copy
import re
from openpyxl import load_workbook, Workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
output_path = workspace / 'output' / 'owners.xlsx'

INVALID_SHEET_CHARS = r'[]:*?/\\'

def is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == '')

def clean_sheet_name(owner_key, used):
    # owner_key is display grouping key (blank already Unassigned); keep non-invalid chars and trim for Excel.
    base = str(owner_key)
    for ch in INVALID_SHEET_CHARS:
        base = base.replace(ch, '_')
    base = base.strip()
    if not base:
        base = 'Sheet'
    # Excel sheet names max 31 chars. Avoid case-insensitive duplicates as Excel does.
    def truncate_for_suffix(name, suffix):
        return name[:31 - len(suffix)] + suffix
    candidate = base[:31]
    n = 1
    used_lower = {name.lower() for name in used}
    while candidate.lower() in used_lower or candidate.lower() == 'summary':
        suffix = f'_{n}'
        candidate = truncate_for_suffix(base, suffix)
        n += 1
    used.add(candidate)
    return candidate

# Read input, preserving headers and source values.
src_wb = load_workbook(input_path, data_only=False)
if 'tasks' not in src_wb.sheetnames:
    raise ValueError('Input workbook does not contain a tasks worksheet')
src_ws = src_wb['tasks']
headers = [cell.value for cell in src_ws[1]]
if 'Owner' not in headers or 'Status' not in headers:
    raise ValueError(f'Required columns Owner and Status not found in headers: {headers}')
owner_idx = headers.index('Owner') + 1
status_idx = headers.index('Status') + 1

# Preserve grouping order as first seen in source rows; blanks grouped as Unassigned.
groups = OrderedDict()
for row in src_ws.iter_rows(min_row=2, max_row=src_ws.max_row):
    values = [cell.value for cell in row[:len(headers)]]
    # skip entirely empty trailing/input rows
    if all(v is None for v in values):
        continue
    owner_value = values[owner_idx - 1]
    group_key = 'Unassigned' if is_blank(owner_value) else owner_value
    groups.setdefault(group_key, []).append(values)

# Create output workbook.
out_wb = Workbook()
default = out_wb.active
out_wb.remove(default)

used_sheet_names = set()
summary_rows = []
sheet_map = {}

# Summary first, then owner sheets.
summary_ws = out_wb.create_sheet('Summary')
summary_headers = ['Owner', 'SheetName', 'Total', 'Pass', 'Fail']
summary_ws.append(summary_headers)
used_sheet_names.add('Summary')

for owner_key, rows in groups.items():
    sheet_name = clean_sheet_name(owner_key, used_sheet_names)
    sheet_map[owner_key] = sheet_name
    ws = out_wb.create_sheet(sheet_name)
    ws.append(headers)
    for values in rows:
        ws.append(values)
    statuses = Counter(row[status_idx - 1] for row in rows)
    summary_rows.append([owner_key, sheet_name, len(rows), statuses.get('Pass', 0), statuses.get('Fail', 0)])

for row in summary_rows:
    summary_ws.append(row)

# Basic widths for readability.
for ws in out_wb.worksheets:
    for col in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max(max_len + 2, 10), 50)

output_path.parent.mkdir(parents=True, exist_ok=True)
out_wb.save(output_path)
src_wb.close()
out_wb.close()
print('WROTE', output_path)

# Verification: reopen output and confirm requirements.
check_wb = load_workbook(output_path, data_only=False)
errors = []
if 'Summary' not in check_wb.sheetnames:
    errors.append('Missing Summary sheet')

summary = check_wb['Summary']
actual_summary_headers = [cell.value for cell in summary[1]]
if actual_summary_headers != summary_headers:
    errors.append(f'Summary headers mismatch: {actual_summary_headers} != {summary_headers}')

# Re-read source counts independently.
source_counts = OrderedDict()
source_owner_original_examples = {}
source_row_count = 0
src_wb2 = load_workbook(input_path, data_only=False, read_only=True)
src_ws2 = src_wb2['tasks']
for row in src_ws2.iter_rows(min_row=2, max_row=src_ws2.max_row, values_only=True):
    vals = list(row[:len(headers)])
    if all(v is None for v in vals):
        continue
    source_row_count += 1
    owner_value = vals[owner_idx - 1]
    group_key = 'Unassigned' if is_blank(owner_value) else owner_value
    source_counts.setdefault(group_key, Counter())
    source_counts[group_key]['Total'] += 1
    source_counts[group_key][vals[status_idx - 1]] += 1
src_wb2.close()

summary_data = []
for row in summary.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    summary_data.append(row[:5])

if len(summary_data) != len(source_counts):
    errors.append(f'Summary group count mismatch: {len(summary_data)} != {len(source_counts)}')

seen_sheet_names = set(['Summary'])
owner_sheet_row_total = 0
for owner, sheet_name, total, passed, failed in summary_data:
    if owner not in source_counts:
        errors.append(f'Summary owner not in source: {owner!r}')
        continue
    exp = source_counts[owner]
    if (total, passed, failed) != (exp['Total'], exp['Pass'], exp['Fail']):
        errors.append(f'Counts mismatch for {owner!r}: {(total, passed, failed)} != {(exp["Total"], exp["Pass"], exp["Fail"])}')
    if not sheet_name or sheet_name not in check_wb.sheetnames:
        errors.append(f'SheetName for {owner!r} missing/invalid: {sheet_name!r}')
        continue
    if len(sheet_name) > 31 or any(ch in sheet_name for ch in INVALID_SHEET_CHARS):
        errors.append(f'Illegal Excel sheet name for {owner!r}: {sheet_name!r}')
    if sheet_name in seen_sheet_names:
        errors.append(f'Duplicate sheet name: {sheet_name!r}')
    seen_sheet_names.add(sheet_name)
    ws = check_wb[sheet_name]
    actual_headers = [cell.value for cell in ws[1]]
    if actual_headers != headers:
        errors.append(f'Headers mismatch in {sheet_name}: {actual_headers} != {headers}')
    data_rows = [r[:len(headers)] for r in ws.iter_rows(min_row=2, values_only=True) if not all(v is None for v in r[:len(headers)])]
    owner_sheet_row_total += len(data_rows)
    if len(data_rows) != total:
        errors.append(f'Row count mismatch in {sheet_name}: {len(data_rows)} != {total}')
    for vals in data_rows:
        row_owner = vals[owner_idx - 1]
        row_group = 'Unassigned' if is_blank(row_owner) else row_owner
        if row_group != owner:
            errors.append(f'Owner grouping mismatch in {sheet_name}: row owner {row_owner!r} grouped under {owner!r}')
        if vals[status_idx - 1] not in ('Pass', 'Fail'):
            errors.append(f'Unexpected status in {sheet_name}: {vals[status_idx - 1]!r}')

if owner_sheet_row_total != source_row_count:
    errors.append(f'Total owner sheet rows mismatch: {owner_sheet_row_total} != {source_row_count}')

expected_sheets = {'Summary'} | {row[1] for row in summary_data}
if set(check_wb.sheetnames) != expected_sheets:
    errors.append(f'Workbook sheets mismatch: {check_wb.sheetnames} != {sorted(expected_sheets)}')

print('output sheets:', check_wb.sheetnames)
print('summary rows:')
for row in summary_data:
    print(row)
print('source row count:', source_row_count)
print('owner sheet row total:', owner_sheet_row_total)

check_wb.close()
if errors:
    print('VERIFICATION FAILED')
    for err in errors:
        print('ERROR:', err)
    raise SystemExit(1)
print('VERIFICATION PASSED: owners.xlsx exists; Summary has Owner/SheetName/Total/Pass/Fail; each owner group has one legal unique sheet; blank Owner rows are grouped under Unassigned while original Owner cell values are preserved; headers and row counts match source; statuses are Pass/Fail.')