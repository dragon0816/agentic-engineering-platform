from pathlib import Path
import subprocess
import sys
import csv
import math

from openpyxl import load_workbook

base = Path('.').resolve()
input_dir = base / 'input'
output_dir = base / 'output'
output_dir.mkdir(exist_ok=True)

fixed_code = '''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals: dict[str, float] = {}
    counts: dict[str, int] = {}

    with open(args.input, newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            model = row["Model"]
            value_text = (row.get("Value") or "").strip()
            if value_text == "":
                continue
            value = float(value_text)
            totals[model] = totals.get(model, 0.0) + value
            counts[model] = counts.get(model, 0) + 1

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(args.output)


if __name__ == "__main__":
    main()
'''

export_path = output_dir / 'export.py'
export_path.write_text(fixed_code, encoding='utf-8')
print('Wrote', export_path)

# Run the exact requested CLI shape from the workspace.
cmd = ['python', r'output\export.py', '--input', r'input\input.csv', '--output', r'output\summary.xlsx']
print('Running:', ' '.join(cmd))
proc = subprocess.run(cmd, cwd=base, text=True, capture_output=True)
print('Return code:', proc.returncode)
if proc.stdout:
    print('STDOUT:\n' + proc.stdout)
if proc.stderr:
    print('STDERR:\n' + proc.stderr)
if proc.returncode != 0:
    raise SystemExit(proc.returncode)

# Compute expected from input.csv independently.
expected_totals = {}
expected_counts = {}
with open(input_dir / 'input.csv', newline='', encoding='utf-8-sig') as handle:
    reader = csv.DictReader(handle)
    assert reader.fieldnames == ['Model', 'Value'], f'Unexpected input headers: {reader.fieldnames}'
    for row in reader:
        v = (row.get('Value') or '').strip()
        if not v:
            continue
        expected_totals[row['Model']] = expected_totals.get(row['Model'], 0.0) + float(v)
        expected_counts[row['Model']] = expected_counts.get(row['Model'], 0) + 1
expected_rows = [
    (model, expected_totals[model] / expected_counts[model], expected_counts[model])
    for model in sorted(expected_totals)
]

# Reopen and verify every output file.
assert export_path.exists(), 'output/export.py was not created'
export_text = export_path.read_text(encoding='utf-8')
assert 'parser.add_argument("--input", required=True)' in export_text, 'CLI --input changed/missing'
assert 'parser.add_argument("--output", required=True)' in export_text, 'CLI --output changed/missing'
assert 'row["Model"]' in export_text, 'Model column not used'
assert 'row.get("Value")' in export_text, 'Value column not used'

summary_path = output_dir / 'summary.xlsx'
assert summary_path.exists(), 'output/summary.xlsx was not created'
wb = load_workbook(summary_path, data_only=True)
assert wb.sheetnames == ['summary'], f'Unexpected sheet names: {wb.sheetnames}'
ws = wb['summary']
rows = list(ws.iter_rows(values_only=True))
print('Workbook rows:', rows)
assert rows[0] == ('Model', 'Average', 'ValidCount'), f'Unexpected output headers: {rows[0]}'
actual_rows = rows[1:]
assert len(actual_rows) == len(expected_rows), f'Expected {len(expected_rows)} data rows, got {len(actual_rows)}'
for actual, expected in zip(actual_rows, expected_rows):
    amodel, avg, cnt = actual
    emodel, eavg, ecnt = expected
    assert amodel == emodel, f'Model mismatch: {amodel} != {emodel}'
    assert isinstance(avg, (int, float)), f'Average for {amodel} is not numeric: {avg!r}'
    assert math.isclose(float(avg), eavg, rel_tol=0, abs_tol=1e-12), f'Average mismatch for {amodel}: {avg} != {eavg}'
    assert isinstance(cnt, int), f'ValidCount for {amodel} is not integer: {cnt!r}'
    assert cnt == ecnt, f'ValidCount mismatch for {amodel}: {cnt} != {ecnt}'
print('Expected rows:', expected_rows)
print('Verification passed: output/export.py and output/summary.xlsx meet the stated requirements.')
