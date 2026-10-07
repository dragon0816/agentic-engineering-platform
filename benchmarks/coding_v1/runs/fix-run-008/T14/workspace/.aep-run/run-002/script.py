from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

workspace = Path.cwd()
input_path = workspace / 'input' / 'large_results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

required_headers = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']

stats = defaultdict(lambda: {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0})
overall = {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0}

with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    print('Read headers:', headers)
    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise ValueError(f'Missing required columns: {missing}')
    for row_num, row in enumerate(reader, start=2):
        model = row['Device Model']
        status = row['Status']
        try:
            duration = float(row['DurationSeconds'])
        except Exception as e:
            raise ValueError(f'Invalid DurationSeconds at CSV row {row_num}: {row["DurationSeconds"]!r}') from e
        if status not in ('Pass', 'Fail'):
            raise ValueError(f'Unexpected Status at CSV row {row_num}: {status!r}')
        bucket = stats[model]
        bucket['Total'] += 1
        bucket['PassCount'] += 1 if status == 'Pass' else 0
        bucket['FailCount'] += 1 if status == 'Fail' else 0
        bucket['DurationSum'] += duration
        overall['Total'] += 1
        overall['PassCount'] += 1 if status == 'Pass' else 0
        overall['FailCount'] += 1 if status == 'Fail' else 0
        overall['DurationSum'] += duration

print('Models found:', len(stats))
print('Total rows processed:', overall['Total'])

wb = Workbook()
ws = wb.active
ws.title = 'Summary'
columns = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
ws.append(columns)
for cell in ws[1]:
    cell.font = Font(bold=True)

for model in sorted(stats):
    s = stats[model]
    total = s['Total']
    pass_rate = s['PassCount'] / total if total else 0
    avg_duration = s['DurationSum'] / total if total else 0
    ws.append([model, total, s['PassCount'], s['FailCount'], pass_rate, avg_duration])

# Formatting
for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
    for cell in row:
        cell.number_format = '0.00%'
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    for cell in row:
        cell.number_format = '0.00'
for col_idx, column_name in enumerate(columns, start=1):
    max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in ws[get_column_letter(col_idx)])
    ws.column_dimensions[get_column_letter(col_idx)].width = max(max_len + 2, len(column_name) + 2)

ows = wb.create_sheet('Overall')
ows.append(['Metric', 'Value'])
for cell in ows[1]:
    cell.font = Font(bold=True)
ows.append(['Total', overall['Total']])
ows.append(['PassCount', overall['PassCount']])
ows.append(['FailCount', overall['FailCount']])
ows.append(['PassRate', overall['PassCount'] / overall['Total'] if overall['Total'] else 0])
ows.append(['AvgDurationSeconds', overall['DurationSum'] / overall['Total'] if overall['Total'] else 0])
ows['B5'].number_format = '0.00%'
ows['B6'].number_format = '0.00'
ows.column_dimensions['A'].width = 22
ows.column_dimensions['B'].width = 18

wb.save(output_path)
print('Wrote:', output_path)

# Verification: reopen input and output, independently recompute and confirm requirements.
# Recompute independently from CSV.
expected = defaultdict(lambda: {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0})
expected_overall = {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0}
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        model = row['Device Model']
        status = row['Status']
        duration = float(row['DurationSeconds'])
        expected[model]['Total'] += 1
        expected[model]['PassCount'] += 1 if status == 'Pass' else 0
        expected[model]['FailCount'] += 1 if status == 'Fail' else 0
        expected[model]['DurationSum'] += duration
        expected_overall['Total'] += 1
        expected_overall['PassCount'] += 1 if status == 'Pass' else 0
        expected_overall['FailCount'] += 1 if status == 'Fail' else 0
        expected_overall['DurationSum'] += duration

check_wb = load_workbook(output_path, data_only=True)
assert set(['Summary', 'Overall']).issubset(set(check_wb.sheetnames)), check_wb.sheetnames
check_ws = check_wb['Summary']
actual_columns = [check_ws.cell(row=1, column=i).value for i in range(1, 7)]
assert actual_columns == columns, actual_columns
actual_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
assert len(actual_rows) == len(expected), (len(actual_rows), len(expected))
for row in actual_rows:
    model, total, pass_count, fail_count, pass_rate, avg_duration = row
    e = expected[model]
    assert isinstance(model, str), type(model)
    assert total == e['Total'], (model, total, e['Total'])
    assert pass_count == e['PassCount'], (model, pass_count, e['PassCount'])
    assert fail_count == e['FailCount'], (model, fail_count, e['FailCount'])
    assert abs(pass_rate - (e['PassCount'] / e['Total'])) < 1e-12, (model, pass_rate)
    assert abs(avg_duration - (e['DurationSum'] / e['Total'])) < 1e-12, (model, avg_duration)

check_ows = check_wb['Overall']
overall_values = {check_ows.cell(row=r, column=1).value: check_ows.cell(row=r, column=2).value for r in range(2, 7)}
assert overall_values['Total'] == expected_overall['Total']
assert overall_values['PassCount'] == expected_overall['PassCount']
assert overall_values['FailCount'] == expected_overall['FailCount']
assert abs(overall_values['PassRate'] - (expected_overall['PassCount'] / expected_overall['Total'])) < 1e-12
assert abs(overall_values['AvgDurationSeconds'] - (expected_overall['DurationSum'] / expected_overall['Total'])) < 1e-12
print('Verified output file: summary.xlsx')
print('Verified Summary columns:', actual_columns)
print('Verified model row count:', len(actual_rows))
print('Verified Overall metrics:', overall_values)
