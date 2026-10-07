from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV preserving original textual values, including leading zeroes.
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    fieldnames = reader.fieldnames
    if fieldnames is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in fieldnames:
        raise ValueError('Missing required column: Status')
    if 'Device ID' not in fieldnames:
        raise ValueError('Missing required column: Device ID')
    failed_rows = [row for row in reader if (row.get('Status') or '').strip().lower() == 'fail']

wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

# Write headers and rows. Keep all source columns and their original order.
for col_idx, name in enumerate(fieldnames, start=1):
    cell = ws.cell(row=1, column=col_idx, value=name)
    cell.font = Font(bold=True)

for row_idx, row in enumerate(failed_rows, start=2):
    for col_idx, name in enumerate(fieldnames, start=1):
        cell = ws.cell(row=row_idx, column=col_idx)
        value = row.get(name, '')
        # Ensure Device ID remains text, preserving leading zeroes.
        if name == 'Device ID':
            cell.number_format = '@'
            cell.value = '' if value is None else str(value)
        else:
            cell.value = value

# Also mark entire Device ID column as text.
device_col = fieldnames.index('Device ID') + 1
for r in range(1, len(failed_rows) + 2):
    ws.cell(row=r, column=device_col).number_format = '@'

wb.save(output_path)

# Verification: reopen output and confirm requirements.
check_wb = load_workbook(output_path, data_only=False)
if 'failed_items' not in check_wb.sheetnames:
    raise AssertionError('failed_items sheet missing')
check_ws = check_wb['failed_items']
headers = [check_ws.cell(row=1, column=c).value for c in range(1, len(fieldnames)+1)]
if headers != fieldnames:
    raise AssertionError(f'Headers changed: {headers!r} != {fieldnames!r}')
actual_rows = []
for r in range(2, check_ws.max_row + 1):
    actual_rows.append([check_ws.cell(row=r, column=c).value for c in range(1, len(fieldnames)+1)])
expected_rows = [[row.get(name, '') for name in fieldnames] for row in failed_rows]
if actual_rows != expected_rows:
    raise AssertionError(f'Rows/order mismatch: {actual_rows!r} != {expected_rows!r}')
for r in range(2, check_ws.max_row + 1):
    dev_cell = check_ws.cell(row=r, column=device_col)
    if not isinstance(dev_cell.value, str):
        raise AssertionError(f'Device ID at row {r} is not text: {dev_cell.value!r}')
    if dev_cell.number_format != '@':
        raise AssertionError(f'Device ID at row {r} not formatted as text')

print('Created:', output_path)
print('Source columns:', fieldnames)
print('Failed row count:', len(failed_rows))
print('Output sheet:', 'failed_items')
print('Verified: output file exists, worksheet name, all source columns/order, filtered Status strip/lower == fail, original row order, Device ID text with leading zeroes preserved.')