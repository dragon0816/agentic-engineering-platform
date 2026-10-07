import pathlib
import re
import urllib.request
from html.parser import HTMLParser
from urllib.parse import urljoin
from openpyxl import Workbook, load_workbook

WORKSPACE = pathlib.Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
CATALOG_URL = 'http://127.0.0.1:64650/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class Node:
    def __init__(self, tag=None, attrs=None, parent=None):
        self.tag = tag
        self.attrs = dict(attrs or [])
        self.parent = parent
        self.children = []
        self.text_parts = []
    def add_child(self, child):
        self.children.append(child)
    @property
    def text(self):
        parts = list(self.text_parts)
        for c in self.children:
            parts.append(c.text)
        return ''.join(parts)
    def has_class(self, cls):
        return cls in self.attrs.get('class', '').split()

class TreeParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('document')
        self.stack = [self.root]
    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.stack[-1])
        self.stack[-1].add_child(node)
        # HTML void tags should not be pushed
        if tag.lower() not in {'area','base','br','col','embed','hr','img','input','link','meta','param','source','track','wbr'}:
            self.stack.append(node)
    def handle_endtag(self, tag):
        for i in range(len(self.stack)-1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                break
    def handle_data(self, data):
        if self.stack:
            self.stack[-1].text_parts.append(data)

def walk(node):
    yield node
    for child in node.children:
        yield from walk(child)

def first_desc(node, predicate):
    for n in walk(node):
        if n is not node and predicate(n):
            return n
    return None

def parse_price(raw):
    raw = ' '.join((raw or '').split())
    m = re.match(r'^([^\d\-+.,]*)([\d,]+(?:\.\d+)?)', raw)
    if not m:
        raise ValueError(f'Cannot parse price: {raw!r}')
    currency = m.group(1).strip()
    value_text = m.group(2).replace(',', '')
    value = float(value_text) if '.' in value_text else int(value_text)
    return currency, value

# Fetch source page
with urllib.request.urlopen(CATALOG_URL, timeout=10) as response:
    html = response.read().decode('utf-8', errors='replace')

parser = TreeParser()
parser.feed(html)

main = None
for n in walk(parser.root):
    if n.tag == 'main' and n.attrs.get('id') == 'catalog':
        main = n
        break
if main is None:
    raise RuntimeError('Could not find <main id="catalog">')

products = []
for node in walk(main):
    if node.tag == 'div' and node.has_class('card'):
        link = first_desc(node, lambda x: x.tag == 'a' and x.attrs.get('href'))
        price_node = first_desc(node, lambda x: x.tag == 'span' and x.has_class('price'))
        if not link or not price_node:
            continue
        name = ' '.join(link.text.split())
        currency, price = parse_price(price_node.text)
        full_url = urljoin(CATALOG_URL, link.attrs['href'])
        products.append({
            '產品名稱': name,
            '價格': price,
            '幣別': currency,
            '產品網址': full_url,
        })

if not products:
    raise RuntimeError('No products found inside main catalog')

wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for p in products:
    ws.append([p[h] for h in headers])

# basic column widths
for idx, width in enumerate([20, 12, 10, 40], start=1):
    ws.column_dimensions[chr(64+idx)].width = width

wb.save(OUTFILE)
print(f'wrote {OUTFILE}')
print(f'product rows: {len(products)}')
for p in products:
    print(p)

# Verification: reopen output and assert requirements
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
read_headers = [ws2.cell(1, c).value for c in range(1, 5)]
assert read_headers == headers, f'headers mismatch: {read_headers}'
rows = list(ws2.iter_rows(min_row=2, max_col=4, values_only=True))
assert len(rows) == 5, f'expected 5 catalog rows, got {len(rows)}'
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
assert [r[0] for r in rows] == expected_names, [r[0] for r in rows]
for name, price, currency, product_url in rows:
    assert isinstance(name, str) and name, name
    assert isinstance(price, (int, float)) and not isinstance(price, bool), (name, price, type(price))
    assert currency == 'NT$', (name, currency)
    assert isinstance(product_url, str) and product_url.startswith('http://127.0.0.1:64650/'), product_url
    assert '/ad/' not in product_url, product_url
assert [r[1] for r in rows] == [1299, 2450, 99, 560, 12000], [r[1] for r in rows]
print('verification passed: reopened products.xlsx; checked sheet name, headers, 5 rows, page order, numeric prices, NT$ currency, full URLs, and no ad URLs.')
