import csv
import hashlib
import re
from pathlib import Path
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
THRESHOLD = 20

print('Workspace:', WORKSPACE)
print('Input directory:', INPUT_DIR)
print('Output directory:', OUTPUT_DIR)

if not INPUT_DIR.exists():
    raise FileNotFoundError(f'Input directory not found: {INPUT_DIR}')

input_files = sorted([p for p in INPUT_DIR.iterdir() if p.is_file()])
print('Input files:')
for p in input_files:
    print(' -', p.name, p.stat().st_size, 'bytes')

sop_path = INPUT_DIR / 'supplied_sop.md'
if not sop_path.exists():
    raise FileNotFoundError('Required SOP not found: input/supplied_sop.md')
sop_text = sop_path.read_text(encoding='utf-8-sig')
print('SOP first 800 chars:')
print(sop_text[:800])

csv_files = sorted(INPUT_DIR.glob('*.csv'))
if not csv_files:
    raise FileNotFoundError('No CSV file found under input/')
preferred = [p for p in csv_files if re.search(r'(test|result|data)', p.stem, re.I)]
if len(csv_files) == 1:
    csv_path = csv_files[0]
elif len(preferred) == 1:
    csv_path = preferred[0]
else:
    raise ValueError('Ambiguous CSV input files: ' + ', '.join(p.name for p in csv_files))
print('Selected CSV:', csv_path.name)

def file_hash(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(65536), b''):
            h.update(chunk)
    return h.hexdigest()

before_hashes = {str(p.relative_to(WORKSPACE)): file_hash(p) for p in input_files}

try:
    csv_text = csv_path.read_text(encoding='utf-8-sig')
    encoding_used = 'utf-8-sig'
except UnicodeDecodeError:
    csv_text = csv_path.read_text(encoding='cp950')
    encoding_used = 'cp950'
print('CSV encoding used:', encoding_used)

sample = csv_text[:4096]
dialect = csv.Sniffer().sniff(sample) if sample.strip() else csv.excel
reader = csv.DictReader(csv_text.splitlines(), dialect=dialect)
headers = reader.fieldnames
if not headers:
    raise ValueError(f'CSV has no header row: {csv_path.name}')
print('CSV headers:', headers)
rows = list(reader)
print('CSV row count:', len(rows))
print('CSV sample rows:', rows[:5])

def find_col(candidates):
    normalized = {h.strip().lower(): h for h in headers}
    for c in candidates:
        if c.lower() in normalized:
            return normalized[c.lower()]
    for h in headers:
        h_norm = h.strip().lower()
        if any(c.lower() in h_norm for c in candidates):
            return h
    return None

status_col = find_col(['Status', 'Result'])
value_col = find_col(['Value'])
if status_col is None:
    raise ValueError('Cannot find Status column in CSV headers: ' + ', '.join(headers))
if value_col is None:
    raise ValueError('Cannot find Value column in CSV headers: ' + ', '.join(headers))
print('Using Status column:', status_col)
print('Using Value column:', value_col)

fail_rows = []
for row in rows:
    if str(row.get(status_col, '')).strip().lower() == 'fail':
        raw_value = row.get(value_col, '')
        try:
            numeric_value = float(str(raw_value).strip())
        except Exception:
            numeric_value = None
        urgent = bool(numeric_value is not None and numeric_value > THRESHOLD)
        out_row = dict(row)
        out_row['Urgent'] = urgent
        fail_rows.append(out_row)

print('Fail rows:', len(fail_rows))
print('Urgent rows:', sum(1 for r in fail_rows if r['Urgent']))

out_xlsx = OUTPUT_DIR / 'failures.xlsx'
wb = Workbook()
ws = wb.active
ws.title = 'failures'
out_headers = list(headers) + ['Urgent']
ws.append(out_headers)
for row in fail_rows:
    ws.append([row.get(h, '') for h in headers] + [row['Urgent']])
for col_cells in ws.columns:
    max_len = 0
    letter = col_cells[0].column_letter
    for cell in col_cells:
        max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
    ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 60)
wb.save(out_xlsx)

urgent_count = sum(1 for r in fail_rows if r['Urgent'])
nonurgent_count = len(fail_rows) - urgent_count
summary_lines = [
    '# Test Failure Summary',
    '',
    f'- SOP: input/{sop_path.name}',
    f'- Source CSV: input/{csv_path.name}',
    f'- Total rows read: {len(rows)}',
    f'- Fail rows retained: {len(fail_rows)}',
    f'- Urgent threshold: {value_col} > {THRESHOLD}',
    f'- Urgent = true: {urgent_count}',
    f'- Urgent = false: {nonurgent_count}',
    '',
    '## Output files',
    '',
    '- output/failures.xlsx',
    '- output/summary.md',
]
(OUTPUT_DIR / 'summary.md').write_text('\n'.join(summary_lines) + '\n', encoding='utf-8')

check_wb = load_workbook(out_xlsx, data_only=True)
check_ws = check_wb['failures']
read_headers = [check_ws.cell(row=1, column=i).value for i in range(1, check_ws.max_column + 1)]
assert read_headers == out_headers, f'Workbook headers mismatch: {read_headers} != {out_headers}'
assert check_ws.max_row == len(fail_rows) + 1, f'Workbook row count mismatch: {check_ws.max_row} != {len(fail_rows)+1}'
urgent_col_idx = read_headers.index('Urgent') + 1
for r_idx in range(2, check_ws.max_row + 1):
    val = check_ws.cell(row=r_idx, column=urgent_col_idx).value
    assert isinstance(val, bool), f'Urgent value at row {r_idx} is not boolean: {val!r}'

summary_text = (OUTPUT_DIR / 'summary.md').read_text(encoding='utf-8')
for required in ['# Test Failure Summary', 'Fail rows retained:', 'Urgent threshold:', 'output/failures.xlsx']:
    assert required in summary_text, f'Missing summary content: {required}'

after_hashes = {str(p.relative_to(WORKSPACE)): file_hash(p) for p in input_files}
assert before_hashes == after_hashes, 'One or more input files changed during processing'

print('Created:', out_xlsx)
print('Created:', OUTPUT_DIR / 'summary.md')
print('Verification complete:')
print(' - input files unchanged')
print(' - failures.xlsx reopened; headers, row count, and boolean Urgent values checked')
print(' - summary.md reopened; required content checked')
