import os, csv, hashlib
from pathlib import Path
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

print(f'Workspace: {workspace}')
print(f'Input directory: {input_dir}')
print(f'Output directory: {output_dir}')
if not input_dir.exists():
    raise FileNotFoundError(f'Input directory not found: {input_dir}')

input_files = sorted([p for p in input_dir.rglob('*') if p.is_file()])
print('Input files:')
for p in input_files:
    print(f' - {p.relative_to(workspace)} {p.stat().st_size} bytes')

sop_path = input_dir / 'supplied_sop.md'
if not sop_path.exists():
    raise FileNotFoundError('Required SOP not found: input/supplied_sop.md')
sop_text = sop_path.read_text(encoding='utf-8-sig')
print('SOP first 800 chars:')
print(sop_text[:800])

manifest_path = input_dir / 'manifest.json'
if manifest_path.exists():
    print('manifest.json content:')
    print(manifest_path.read_text(encoding='utf-8-sig'))

def sha256(path):
    h = hashlib.sha256()
    with path.open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()
input_hashes_before = {str(p.relative_to(workspace)): sha256(p) for p in input_files}

results_path = input_dir / 'results.csv'
if not results_path.exists():
    csvs = sorted(input_dir.glob('*.csv'))
    if len(csvs) == 1:
        results_path = csvs[0]
        print(f'Using sole CSV under input/: {results_path.name}')
    elif not csvs:
        raise FileNotFoundError('Required test CSV not found. SOP names input/results.csv, but no CSV file exists under input/. Please add results.csv to input/ and rerun; input files were not modified.')
    else:
        names = ', '.join(p.name for p in csvs)
        raise ValueError(f'Multiple CSV files found under input/ ({names}) and SOP names results.csv; please provide input/results.csv or clarify which CSV to use.')
else:
    print('Using SOP-named CSV: input/results.csv')

with results_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    print(f'CSV headers: {headers}')
    sample_rows = []
    rows = []
    for row in reader:
        rows.append(row)
        if len(sample_rows) < 5:
            sample_rows.append(row)
print(f'CSV row count: {len(rows)}')
print(f'CSV sample rows: {sample_rows}')

required = ['Device ID', 'Status', 'Value']
missing = [c for c in required if c not in headers]
if missing:
    raise ValueError(f'CSV missing required columns from SOP: {missing}. Found headers: {headers}')

fail_rows = []
pass_count = 0
fail_count = 0
urgent_count = 0
for idx, row in enumerate(rows, start=2):
    status = (row.get('Status') or '').strip()
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        raw_value = row.get('Value')
        try:
            value_num = float(str(raw_value).strip())
        except Exception as exc:
            raise ValueError(f'Row {idx} has non-numeric Value for Fail row: {raw_value!r}') from exc
        urgent = value_num > 20
        if urgent:
            urgent_count += 1
        fail_rows.append({'Device ID': row.get('Device ID'), 'Status': row.get('Status'), 'Value': value_num, 'Urgent': urgent})

total_count = len(rows)

failures_path = output_dir / 'failures.xlsx'
wb = Workbook()
ws = wb.active
ws.title = 'failures'
out_headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(out_headers)
for r in fail_rows:
    ws.append([r['Device ID'], r['Status'], r['Value'], r['Urgent']])
wb.save(failures_path)

summary_path = output_dir / 'summary.md'
summary_text = '\n'.join(['# Summary', '', f'- Total: {total_count}', f'- Pass: {pass_count}', f'- Fail: {fail_count}', f'- Urgent: {urgent_count}', ''])
summary_path.write_text(summary_text, encoding='utf-8')

wb2 = load_workbook(failures_path, data_only=True)
ws2 = wb2['failures']
read_headers = [c.value for c in next(ws2.iter_rows(min_row=1, max_row=1))]
assert read_headers == out_headers, f'failures.xlsx headers mismatch: {read_headers}'
read_rows = list(ws2.iter_rows(min_row=2, values_only=True))
assert len(read_rows) == fail_count, f'failure row count mismatch: {len(read_rows)} != {fail_count}'
for i, row in enumerate(read_rows, start=2):
    device_id, status, value, urgent = row
    assert status == 'Fail', f'row {i} status is not Fail: {status!r}'
    assert isinstance(value, (int, float)), f'row {i} Value is not numeric: {value!r}'
    assert isinstance(urgent, bool), f'row {i} Urgent is not boolean: {urgent!r}'
    assert urgent == (float(value) > 20), f'row {i} Urgent mismatch for Value {value!r}'
summary_read = summary_path.read_text(encoding='utf-8')
for label, expected in [('Total', total_count), ('Pass', pass_count), ('Fail', fail_count), ('Urgent', urgent_count)]:
    token = f'- {label}: {expected}'
    assert token in summary_read, f'summary.md missing {token}'

input_files_after = sorted([p for p in input_dir.rglob('*') if p.is_file()])
input_hashes_after = {str(p.relative_to(workspace)): sha256(p) for p in input_files_after}
assert input_hashes_after == input_hashes_before, 'Input files changed during processing'

print('Created files:')
print(f' - {failures_path.relative_to(workspace)}')
print(f' - {summary_path.relative_to(workspace)}')
print('Verification passed: failures.xlsx columns, Fail-only rows, numeric Value, boolean Urgent Value > 20, summary counts, and unchanged inputs.')
