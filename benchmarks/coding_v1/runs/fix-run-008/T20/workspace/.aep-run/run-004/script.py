import csv
import json
import os
from pathlib import Path
from openpyxl import Workbook, load_workbook

WORKSPACE = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\coding_v1\runs\fix-run-008\T20\workspace')
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
SOP_PATH = INPUT_DIR / 'supplied_sop.md'
MANIFEST_PATH = INPUT_DIR / 'manifest.json'
THRESHOLD = 20

print(f'Workspace: {WORKSPACE}')
print(f'Input directory: {INPUT_DIR}')
print(f'Output directory: {OUTPUT_DIR}')
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

print('Input files:')
for p in sorted(INPUT_DIR.rglob('*')):
    if p.is_file():
        print(f' - {p.relative_to(WORKSPACE)} {p.stat().st_size} bytes')

if not SOP_PATH.exists():
    raise FileNotFoundError(f'Missing SOP file: {SOP_PATH}')
sop_text = SOP_PATH.read_text(encoding='utf-8')
print('SOP first 800 chars:')
print(sop_text[:800])

manifest = {}
if MANIFEST_PATH.exists():
    manifest = json.loads(MANIFEST_PATH.read_text(encoding='utf-8-sig'))
    print('manifest.json content:')
    print(json.dumps(manifest, ensure_ascii=False, indent=2))

# SOP names results.csv. The supplied manifest maps that logical input to the real file.
candidates = []
direct = INPUT_DIR / 'results.csv'
if direct.exists():
    candidates.append(direct)
manifest_value = manifest.get('results.csv')
if manifest_value:
    mp = Path(manifest_value)
    if not mp.is_absolute():
        mp = WORKSPACE / mp
    candidates.append(mp)
# Fallback only if exactly one CSV is supplied under input; this keeps behavior deterministic.
all_csvs = sorted(INPUT_DIR.rglob('*.csv'))
if len(all_csvs) == 1:
    candidates.append(all_csvs[0])

csv_path = None
seen = set()
for c in candidates:
    c = c.resolve()
    if c in seen:
        continue
    seen.add(c)
    try:
        c.relative_to(INPUT_DIR.resolve())
    except ValueError:
        raise ValueError(f'Manifest points outside input directory: {c}')
    if c.exists() and c.is_file():
        csv_path = c
        break

if csv_path is None:
    raise FileNotFoundError('Required test CSV not found. SOP names results.csv; no valid manifest mapping or single CSV under input was found. Input files were not modified.')

print(f'Using test CSV: {csv_path.relative_to(WORKSPACE)}')
print('CSV sample:')
with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    for i, line in enumerate(f):
        print(line.rstrip('\n'))
        if i >= 5:
            break

with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    source_headers = reader.fieldnames or []
    required_headers = ['Device ID', 'Status', 'Value']
    missing = [h for h in required_headers if h not in source_headers]
    if missing:
        raise ValueError(f'CSV missing required columns: {missing}; headers were: {source_headers}')
    rows = list(reader)

fail_rows = []
pass_count = 0
fail_count = 0
urgent_count = 0
for idx, row in enumerate(rows, start=2):
    status = (row.get('Status') or '').strip()
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        raw_value = row.get('Value')
        try:
            value_num = float(raw_value)
        except Exception as exc:
            raise ValueError(f'Row {idx} has non-numeric Value for Fail row: {raw_value!r}') from exc
        urgent = value_num > THRESHOLD
        if urgent:
            urgent_count += 1
        out_row = {
            'Device ID': '' if row.get('Device ID') is None else str(row.get('Device ID')),
            'Status': row.get('Status'),
            'Value': value_num,
            'Urgent': urgent,
        }
        fail_rows.append(out_row)

total_count = len(rows)
failures_path = OUTPUT_DIR / 'failures.xlsx'
summary_path = OUTPUT_DIR / 'summary.md'

wb = Workbook()
ws = wb.active
ws.title = 'Failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for out in fail_rows:
    ws.append([out['Device ID'], out['Status'], out['Value'], out['Urgent']])
# Keep IDs as text in workbook.
for cell in ws['A']:
    cell.number_format = '@'
wb.save(failures_path)

summary_text = '\n'.join([
    '# Test Failure Summary',
    '',
    f'Total: {total_count}',
    f'Pass: {pass_count}',
    f'Fail: {fail_count}',
    f'Urgent: {urgent_count}',
    '',
])
summary_path.write_text(summary_text, encoding='utf-8')

# Verification: reopen all outputs and confirm requirements.
print('Verification:')
if not failures_path.exists() or failures_path.stat().st_size <= 0:
    raise AssertionError('failures.xlsx was not created or is empty')
if not summary_path.exists() or summary_path.stat().st_size <= 0:
    raise AssertionError('summary.md was not created or is empty')

check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['Failures']
check_headers = [check_ws.cell(row=1, column=i).value for i in range(1, 5)]
if check_headers != headers:
    raise AssertionError(f'failures.xlsx headers mismatch: {check_headers}')
check_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
if len(check_rows) != fail_count:
    raise AssertionError(f'Expected {fail_count} Fail rows in failures.xlsx, found {len(check_rows)}')
calc_urgent = 0
for excel_idx, (device_id, status, value, urgent) in enumerate(check_rows, start=2):
    if not isinstance(device_id, str):
        raise AssertionError(f'Device ID at Excel row {excel_idx} is not text: {device_id!r}')
    if status != 'Fail':
        raise AssertionError(f'Non-Fail row found in failures.xlsx row {excel_idx}: {status!r}')
    if not isinstance(value, (int, float)):
        raise AssertionError(f'Value at Excel row {excel_idx} is not numeric: {value!r}')
    if not isinstance(urgent, bool):
        raise AssertionError(f'Urgent at Excel row {excel_idx} is not boolean: {urgent!r}')
    if urgent != (float(value) > THRESHOLD):
        raise AssertionError(f'Urgent mismatch at Excel row {excel_idx}: value={value}, urgent={urgent}')
    if urgent:
        calc_urgent += 1
if calc_urgent != urgent_count:
    raise AssertionError(f'Urgent count mismatch: {calc_urgent} vs {urgent_count}')

check_summary = summary_path.read_text(encoding='utf-8')
expected_lines = [f'Total: {total_count}', f'Pass: {pass_count}', f'Fail: {fail_count}', f'Urgent: {urgent_count}']
for line in expected_lines:
    if line not in check_summary:
        raise AssertionError(f'summary.md missing line: {line}')

print(f'Created {failures_path.relative_to(WORKSPACE)} with {len(check_rows)} Fail rows; headers={check_headers}')
print(f'Created {summary_path.relative_to(WORKSPACE)} with counts: Total={total_count}, Pass={pass_count}, Fail={fail_count}, Urgent={urgent_count}')
print('Input files were read only; no files under input were written by this script.')
