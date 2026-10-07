from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read CSV with UTF-8 BOM support and newline='' so quoted embedded newlines are parsed correctly.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.reader(f))

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
try:
    device_col_idx = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_col_idx = None

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'

for r_idx, row in enumerate(rows, start=1):
    # csv.reader preserves column order and duplicate rows. Write exactly the parsed fields.
    for c_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx)
        if value == '':
            # Keep CSV empty values as blank worksheet cells.
            cell.value = None
        else:
            # Assign text values from CSV; this preserves Chinese, commas, embedded newlines, and leading zeroes.
            cell.value = value
        if c_idx == device_col_idx:
            cell.number_format = '@'
            if value != '':
                cell.value = str(value)
        if isinstance(cell.value, str) and '\n' in cell.value:
            cell.alignment = Alignment(wrap_text=True)

wb.save(output_path)

# Verification: reopen source and output, then check stated requirements from saved xlsx.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    source_rows = list(csv.reader(f))

wb2 = load_workbook(output_path, data_only=False)
ws2 = wb2.active

max_row = ws2.max_row
max_col = ws2.max_column
xlsx_rows = []
for r in range(1, max_row + 1):
    row_vals = []
    for c in range(1, max_col + 1):
        v = ws2.cell(r, c).value
        row_vals.append('' if v is None else str(v))
    xlsx_rows.append(row_vals)

# Trim any all-empty trailing rows/cols defensively; should not be needed.
def trim_matrix(matrix):
    while matrix and all(v == '' for v in matrix[-1]):
        matrix.pop()
    if not matrix:
        return matrix
    last_col = 0
    for row in matrix:
        for i, v in enumerate(row, start=1):
            if v != '':
                last_col = max(last_col, i)
    return [row[:last_col] for row in matrix]

xlsx_rows_trimmed = trim_matrix([r[:] for r in xlsx_rows])

errors = []
if xlsx_rows_trimmed != source_rows:
    errors.append('xlsx content does not exactly match CSV rows after treating blank cells as empty strings')
if xlsx_rows_trimmed and xlsx_rows_trimmed[0] != headers:
    errors.append('headers/order changed')
if len(xlsx_rows_trimmed) != len(source_rows):
    errors.append(f'row count changed: source {len(source_rows)} xlsx {len(xlsx_rows_trimmed)}')
if source_rows.count(source_rows[1]) != xlsx_rows_trimmed.count(source_rows[1]):
    errors.append('duplicate row count changed')
if device_col_idx:
    for r in range(2, len(source_rows) + 1):
        src_val = source_rows[r-1][device_col_idx-1]
        cell = ws2.cell(r, device_col_idx)
        xlsx_val = '' if cell.value is None else str(cell.value)
        if xlsx_val != src_val:
            errors.append(f'Device ID mismatch row {r}: {xlsx_val!r} != {src_val!r}')
        if src_val != '' and (cell.data_type != 's' or cell.number_format != '@'):
            errors.append(f'Device ID row {r} is not stored/formatted as text: data_type={cell.data_type!r}, number_format={cell.number_format!r}')

if errors:
    print('Verification failed:')
    for e in errors:
        print('-', e)
    raise SystemExit(1)

print('Created:', output_path)
print('Verified output file exists:', output_path.exists())
print('Verified sheet:', ws2.title)
print('Verified columns preserved:', headers)
print('Verified row count including header:', len(xlsx_rows_trimmed))
print('Verified duplicate rows preserved for first data row count:', xlsx_rows_trimmed.count(source_rows[1]))
print('Verified Device ID cells are text with leading zeros preserved.')
print('Verified empty CSV fields reopened as blank cells.')
print('Verified quoted comma/newline and Chinese text preserved.')