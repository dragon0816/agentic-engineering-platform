from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_file = output_dir / 'merged_devices_tests.xlsx'

def id_as_text(cell):
    """Return a Device ID as text, preserving existing leading zeros and simple numeric zero formats."""
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, str):
        return value
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value)
    number_format = getattr(cell, 'number_format', '') or ''
    # If Excel stored the ID as a number with a format such as 000 or 0000, restore that width.
    if isinstance(value, int) and set(number_format) <= {'0'} and len(number_format) > len(text):
        text = text.zfill(len(number_format))
    return text

def read_sheet(path, expected_sheet=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[expected_sheet] if expected_sheet and expected_sheet in wb.sheetnames else wb.worksheets[0]
    headers = [cell.value for cell in ws[1]]
    if 'Device ID' not in headers:
        raise ValueError(f'{path.name} does not contain a Device ID column')
    id_idx = headers.index('Device ID') + 1
    rows = []
    for row_num in range(2, ws.max_row + 1):
        row_values = {}
        empty = True
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=row_num, column=col_idx)
            if header == 'Device ID':
                val = id_as_text(cell)
            else:
                val = cell.value
            row_values[header] = val
            if val not in (None, ''):
                empty = False
        if not empty:
            rows.append(row_values)
    wb.close()
    return headers, rows

dev_headers, dev_rows = read_sheet(input_dir / 'devices.xlsx', 'devices')
test_headers, test_rows = read_sheet(input_dir / 'tests.xlsx', 'tests')

# devices has unique IDs per requirement; retain first occurrence defensively.
dev_by_id = {}
for row in dev_rows:
    device_id = row.get('Device ID', '')
    if device_id not in dev_by_id:
        dev_by_id[device_id] = row

matched_rows = []
unmatched_device_ids = set(dev_by_id.keys())
unmatched_test_rows = []

for test in test_rows:
    device_id = test.get('Device ID', '')
    dev = dev_by_id.get(device_id)
    if dev is None:
        unmatched_test_rows.append(test)
    else:
        unmatched_device_ids.discard(device_id)
        combined = {}
        # Keep source column names: shared Device ID once, then test columns, then device columns not already present.
        combined['Device ID'] = device_id
        for h in test_headers:
            if h != 'Device ID':
                combined[h] = test.get(h)
        for h in dev_headers:
            if h != 'Device ID' and h not in combined:
                combined[h] = dev.get(h)
        matched_rows.append(combined)

matched_headers = ['Device ID'] + [h for h in test_headers if h != 'Device ID'] + [h for h in dev_headers if h != 'Device ID' and h not in test_headers]
unmatched_headers = ['Source'] + list(dict.fromkeys(dev_headers + [h for h in test_headers if h not in dev_headers]))

unmatched_rows = []
for device_id in sorted(unmatched_device_ids):
    row = {'Source': 'devices'}
    row.update(dev_by_id[device_id])
    unmatched_rows.append(row)
for test in unmatched_test_rows:
    row = {'Source': 'tests'}
    row.update(test)
    unmatched_rows.append(row)

wb_out = Workbook()
ws_matched = wb_out.active
ws_matched.title = 'matched'
ws_unmatched = wb_out.create_sheet('unmatched')

def write_table(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([row.get(h, '') if row.get(h, '') is not None else '' for h in headers])
    # Force Device ID column to text format in every output sheet.
    if 'Device ID' in headers:
        id_col = headers.index('Device ID') + 1
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(r, id_col)
            cell.value = '' if cell.value is None else str(cell.value)
            cell.number_format = '@'
    for col_cells in ws.columns:
        max_len = 0
        for cell in col_cells:
            max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 50)

write_table(ws_matched, matched_headers, matched_rows)
write_table(ws_unmatched, unmatched_headers, unmatched_rows)
wb_out.save(output_file)
print(f'wrote: {output_file}')
print(f'matched data rows: {len(matched_rows)}')
print(f'unmatched data rows: {len(unmatched_rows)}')

# Verification: reopen saved output and confirm sheets, headers, row counts, Source labels, and text IDs.
check_wb = load_workbook(output_file, data_only=False)
assert set(['matched', 'unmatched']).issubset(check_wb.sheetnames), check_wb.sheetnames
cm = check_wb['matched']
cu = check_wb['unmatched']
actual_matched_headers = [c.value for c in cm[1]]
actual_unmatched_headers = [c.value for c in cu[1]]
assert actual_matched_headers == matched_headers, (actual_matched_headers, matched_headers)
assert actual_unmatched_headers == unmatched_headers, (actual_unmatched_headers, unmatched_headers)
assert cm.max_row - 1 == len(matched_rows), (cm.max_row - 1, len(matched_rows))
assert cu.max_row - 1 == len(unmatched_rows), (cu.max_row - 1, len(unmatched_rows))
assert 'Source' in actual_unmatched_headers
source_col = actual_unmatched_headers.index('Source') + 1
sources = [cu.cell(r, source_col).value for r in range(2, cu.max_row + 1)]
assert all(s in {'devices', 'tests'} for s in sources), sources
for ws in [cm, cu]:
    headers = [c.value for c in ws[1]]
    id_col = headers.index('Device ID') + 1
    for r in range(2, ws.max_row + 1):
        val = ws.cell(r, id_col).value
        assert isinstance(val, str), (ws.title, r, val, type(val))
        assert ws.cell(r, id_col).number_format == '@', (ws.title, r, ws.cell(r, id_col).number_format)
# Confirm expected content from the observed inputs.
matched_ids = [cm.cell(r, 1).value for r in range(2, cm.max_row + 1)]
unmatched_records = []
for r in range(2, cu.max_row + 1):
    unmatched_records.append({actual_unmatched_headers[c-1]: cu.cell(r, c).value for c in range(1, cu.max_column + 1)})
assert matched_ids == ['001', '001', '002'], matched_ids
assert any(rec.get('Source') == 'devices' and rec.get('Device ID') == '003' for rec in unmatched_records), unmatched_records
assert any(rec.get('Source') == 'tests' and rec.get('Device ID') == '999' for rec in unmatched_records), unmatched_records
check_wb.close()
print('verification passed: workbook reopened; matched/unmatched sheets, headers, row counts, Source values, and text Device IDs were confirmed.')
