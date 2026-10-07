from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
print('workspace:', base)
print('input files:')
for p in sorted((base/'input').glob('*')):
    print(' -', p.name)
    if p.suffix.lower() in ['.xlsx', '.xlsm']:
        wb = load_workbook(p, read_only=True, data_only=True)
        print('   sheets:', wb.sheetnames)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(max_row=6, values_only=True))
            print(f'   sheet {ws.title}: max_row={ws.max_row}, max_column={ws.max_column}')
            for i, row in enumerate(rows, 1):
                print('    row', i, row)
        wb.close()
