import os, re, json, time, html
import urllib.request, urllib.error
from urllib.parse import urljoin
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import Workbook, load_workbook

BASE_LIST_URL = 'http://127.0.0.1:64650/t13/list'
OUTPUT_DIR = Path('output')
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
XLSX_PATH = OUTPUT_DIR / 'robust_products.xlsx'
WARN_PATH = OUTPUT_DIR / 'warnings.json'
COLUMNS = ['Name','Price','Currency','CPU','URL','FetchStatus']
UA = {'User-Agent': 'Mozilla/5.0 (compatible; AEP product collector)'}

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links=[]; self._in_a=False; self._href=None; self._text=[]
    def handle_starttag(self, tag, attrs):
        if tag.lower()=='a':
            self._in_a=True; self._href=dict(attrs).get('href'); self._text=[]
    def handle_data(self, data):
        if self._in_a: self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower()=='a' and self._in_a:
            text=' '.join(''.join(self._text).split())
            if self._href: self.links.append((text, self._href))
            self._in_a=False; self._href=None; self._text=[]

class NodeTextParser(HTMLParser):
    def __init__(self):
        super().__init__(); self.tag_stack=[]; self.records=[]; self.all_text=[]
    def handle_starttag(self, tag, attrs):
        self.tag_stack.append((tag.lower(), dict(attrs)))
    def handle_endtag(self, tag):
        tag=tag.lower()
        for i in range(len(self.tag_stack)-1, -1, -1):
            if self.tag_stack[i][0]==tag:
                del self.tag_stack[i:]; break
    def handle_data(self, data):
        text=' '.join(data.split())
        if not text: return
        current_tag=self.tag_stack[-1][0] if self.tag_stack else ''
        current_attrs=self.tag_stack[-1][1] if self.tag_stack else {}
        path='/'.join(t for t,a in self.tag_stack)
        self.records.append((path, current_tag, current_attrs, text)); self.all_text.append(text)

def fetch_url(url, detail=False, max_500_attempts=3):
    attempts=0
    while True:
        attempts += 1
        try:
            req=urllib.request.Request(url, headers=UA)
            with urllib.request.urlopen(req, timeout=10) as resp:
                data=resp.read(); charset=resp.headers.get_content_charset() or 'utf-8'
                return {'ok': True, 'status': resp.status, 'text': data.decode(charset, errors='replace'), 'attempts': attempts, 'error': None}
        except urllib.error.HTTPError as e:
            if detail and e.code == 500 and attempts < max_500_attempts:
                time.sleep(0.15); continue
            return {'ok': False, 'status': e.code, 'text': '', 'attempts': attempts, 'error': f'HTTP {e.code}'}
        except Exception as e:
            return {'ok': False, 'status': None, 'text': '', 'attempts': attempts, 'error': f'{type(e).__name__}: {e}'}

def normalize(v):
    if v is None: return 'N/A'
    v=html.unescape(str(v)).strip(); v=re.sub(r'\s+', ' ', v)
    return v if v else 'N/A'

def text_after_label(joined, label):
    pat=re.compile(r'\b'+re.escape(label)+r'\b\s*[:\-]?\s*([^|\n;]+)', re.I)
    m=pat.search(joined)
    if not m: return None
    val=m.group(1).strip()
    val=re.split(r'\b(?:Name|Product|Price|Currency|CPU|Processor|Spec|Specs|Specification|Specifications)\b\s*[:\-]?', val, maxsplit=1, flags=re.I)[0].strip()
    return val or None

def get_label_value(records, labels):
    canonical=[x.lower() for x in labels]
    for i, (_path, _tag, _attrs, text) in enumerate(records):
        for lab in labels:
            m=re.match(r'^\s*'+re.escape(lab)+r'\s*[:\-]\s*(.+)$', text, re.I)
            if m: return m.group(1).strip()
            if re.match(r'^\s*'+re.escape(lab)+r'\s*[:\-]?\s*$', text, re.I):
                for j in range(i+1, min(i+6, len(records))):
                    cand=records[j][3].strip()
                    if not cand: continue
                    if cand.lower().rstrip(':') in canonical: continue
                    return cand
    return None

def attr_value(records, names):
    for _path, _tag, attrs, text in records:
        blob=(' '.join(str(v) for v in attrs.values())+' '+text).strip()
        for name in names:
            if re.search(name, blob, re.I) and text and not re.fullmatch(r'(price|currency|cpu|processor|specs?|specifications?)\s*[:\-]?', text, re.I):
                return text
    return None

def currency_from_symbol(symbol_text):
    if symbol_text is None: return 'N/A'
    s=symbol_text.upper()
    if 'NT$' in s or 'NTD' in s or 'TWD' in s: return 'TWD'
    if '$' in s: return 'USD'
    if '€' in s: return 'EUR'
    if '¥' in s: return 'JPY'
    m=re.search(r'\b(USD|EUR|TWD|JPY|CNY|GBP|NTD)\b', s)
    if m: return 'TWD' if m.group(1)=='NTD' else m.group(1)
    return 'N/A'

def split_price_currency(value):
    value=normalize(value)
    if value=='N/A': return ('N/A','N/A')
    # code before amount, e.g. USD 799, TWD 12,900
    m=re.search(r'\b(USD|EUR|TWD|JPY|CNY|GBP|NTD)\b\s*([0-9][0-9,]*(?:\.\d+)?)', value, re.I)
    if m: return (m.group(2), 'TWD' if m.group(1).upper()=='NTD' else m.group(1).upper())
    # NT$9,900, $799, €399, ¥50000
    m=re.search(r'(NT\$|[$€¥])\s*([0-9][0-9,]*(?:\.\d+)?)', value, re.I)
    if m: return (m.group(2), currency_from_symbol(m.group(1)))
    # amount before code
    m=re.search(r'([0-9][0-9,]*(?:\.\d+)?)\s*\b(USD|EUR|TWD|JPY|CNY|GBP|NTD)\b', value, re.I)
    if m: return (m.group(1), 'TWD' if m.group(2).upper()=='NTD' else m.group(2).upper())
    m=re.search(r'([0-9][0-9,]*(?:\.\d+)?)', value)
    if m: return (m.group(1), 'N/A')
    return ('N/A','N/A')

def parse_detail(text, fallback_name):
    p=NodeTextParser(); p.feed(text)
    records=p.records; all_joined=' | '.join(p.all_text)
    name=get_label_value(records, ['Name','Product','Product Name'])
    if not name:
        for _path, tag, attrs, t in records:
            if tag in ('h1','h2'):
                name=t; break
    name=normalize(name if name else fallback_name)
    price_text=get_label_value(records, ['Price','Cost']) or attr_value(records, [r'price']) or text_after_label(all_joined, 'Price')
    currency_text=get_label_value(records, ['Currency']) or attr_value(records, [r'currency']) or text_after_label(all_joined, 'Currency')
    price, parsed_currency=split_price_currency(price_text)
    currency=normalize(currency_text)
    if currency=='N/A':
        currency=parsed_currency
    else:
        c=currency_from_symbol(currency)
        currency = c if c!='N/A' else currency
    cpu=get_label_value(records, ['CPU','Processor']) or attr_value(records, [r'cpu', r'processor']) or text_after_label(all_joined, 'CPU') or text_after_label(all_joined, 'Processor')
    return {'Name': name, 'Price': normalize(price), 'Currency': normalize(currency), 'CPU': normalize(cpu)}

list_result=fetch_url(BASE_LIST_URL)
if not list_result['ok']:
    raise RuntimeError(f'Could not fetch list URL {BASE_LIST_URL}: {list_result["error"]}')
list_parser=ProductListParser(); list_parser.feed(list_result['text'])
products=[(name, urljoin(BASE_LIST_URL, href)) for name, href in list_parser.links][:5]
print('discovered_products', json.dumps(products, ensure_ascii=False))
if len(products)!=5: raise RuntimeError(f'Expected five products from list, found {len(products)}')

rows=[]; warnings=[]
for idx,(list_name, detail_url) in enumerate(products, start=1):
    result=fetch_url(detail_url, detail=True, max_500_attempts=3)
    print('detail_fetch', idx, list_name, detail_url, 'ok', result['ok'], 'status', result['status'], 'attempts', result['attempts'], 'error', result['error'])
    if result['ok']:
        print('detail_sample', idx, result['text'][:500].replace('\n',' '))
        fields=parse_detail(result['text'], list_name)
        row={'Name': normalize(fields.get('Name') or list_name), 'Price': normalize(fields.get('Price')), 'Currency': normalize(fields.get('Currency')), 'CPU': normalize(fields.get('CPU')), 'URL': detail_url, 'FetchStatus': 'ok'}
        warnings.append({'Name': row['Name'], 'URL': detail_url, 'FetchStatus': 'ok', 'attempts': result['attempts'], 'warning': 'ok'})
    else:
        row={'Name': normalize(list_name), 'Price': 'N/A', 'Currency': 'N/A', 'CPU': 'N/A', 'URL': detail_url, 'FetchStatus': 'error'}
        warnings.append({'Name': row['Name'], 'URL': detail_url, 'FetchStatus': 'error', 'attempts': result['attempts'], 'warning': result['error'] or 'fetch failed'})
    rows.append(row)

wb=Workbook(); ws=wb.active; ws.title='Products'; ws.append(COLUMNS)
for row in rows: ws.append([row[col] for col in COLUMNS])
wb.save(XLSX_PATH)
with open(WARN_PATH, 'w', encoding='utf-8') as f:
    json.dump({'products': warnings}, f, ensure_ascii=False, indent=2)

# Verification: reopen every output and confirm requirements
check_wb=load_workbook(XLSX_PATH, data_only=True); check_ws=check_wb.active
headers=[check_ws.cell(row=1, column=i).value for i in range(1, len(COLUMNS)+1)]
assert headers==COLUMNS, f'headers mismatch: {headers}'
read_rows=[]
for r in range(2, check_ws.max_row+1):
    vals=[check_ws.cell(r,c).value for c in range(1,len(COLUMNS)+1)]
    read_rows.append(dict(zip(COLUMNS, vals)))
assert len(read_rows)==5, f'row count mismatch: {len(read_rows)}'
for i,rr in enumerate(read_rows, start=1):
    for col in COLUMNS:
        assert rr[col] is not None and str(rr[col])!='', f'blank {col} row {i}'
    assert rr['FetchStatus'] in ('ok','error'), f'bad FetchStatus row {i}'
    if rr['FetchStatus']=='error':
        assert rr['Price']=='N/A' and rr['Currency']=='N/A' and rr['CPU']=='N/A', f'failed row unknown fields not N/A row {i}'
with open(WARN_PATH,'r',encoding='utf-8') as f:
    warn_data=json.load(f)
assert isinstance(warn_data, dict) and 'products' in warn_data and len(warn_data['products'])==5
for item in warn_data['products']:
    assert item.get('FetchStatus') in ('ok','error')
    if item.get('FetchStatus')=='ok': assert item.get('warning')=='ok'
print('wrote', str(XLSX_PATH), str(WARN_PATH))
print('verified_headers', headers)
print('verified_row_count', len(read_rows))
print('verified_rows', json.dumps(read_rows, ensure_ascii=False))
print('verified_warnings', json.dumps(warn_data, ensure_ascii=False))
