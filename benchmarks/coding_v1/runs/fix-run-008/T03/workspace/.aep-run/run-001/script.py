from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
print('Workspace:', workspace)
print('Input files:')
for p in sorted(input_dir.iterdir()):
    print(' -', p.name)

src = input_dir / 'tasks.xlsx'
print('tasks.xlsx exists:', src.exists())
wb = load_workbook(src, read_only=True, data_only=False)
print('Workbook sheets:', wb.sheetnames)
if 'tasks' in wb.sheetnames:
    ws = wb['tasks']
    rows = list(ws.iter_rows(max_row=min(ws.max_row, 6), values_only=True))
    print('tasks dimensions:', ws.max_row, 'rows x', ws.max_column, 'columns')
    if rows:
        print('Headers:', rows[0])
        print('Sample rows:')
        for row in rows[1:]:
            print(row)
else:
    print('Sheet tasks not found')
