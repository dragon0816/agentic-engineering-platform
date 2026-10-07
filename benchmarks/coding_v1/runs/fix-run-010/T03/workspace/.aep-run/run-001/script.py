from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
print('Workspace:', workspace)
print('Input files:', [p.name for p in input_dir.iterdir()])
path = input_dir / 'tasks.xlsx'
print('tasks.xlsx exists:', path.exists())
wb = load_workbook(path, read_only=True, data_only=False)
print('Sheets:', wb.sheetnames)
ws = wb['tasks']
headers = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print('Headers:', headers)
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=min(ws.max_row, 6), values_only=True), start=2):
    print(f'Row {i}:', row)
print('Max row:', ws.max_row, 'Max col:', ws.max_column)
wb.close()
