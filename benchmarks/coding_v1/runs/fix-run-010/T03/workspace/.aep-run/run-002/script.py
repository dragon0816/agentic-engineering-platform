from pathlib import Path
from collections import OrderedDict, defaultdict
from copy import copy
import re
from openpyxl import load_workbook, Workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
out_dir = workspace / 'output'
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / 'owners.xlsx'

INVALID_SHEET_CHARS = re.compile(r'[\\/*?:\[\]]')

def is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == '')

def legal_sheet_name(owner_key, used):
    # Excel sheet names cannot contain: \ / * ? : [ ], cannot be blank, and max length is 31.
    base = INVALID_SHEET_CHARS.sub('_', str(owner_key))
    base = base.strip()
    if not base:
        base = 'Unassigned'
    base = base[:31]
    if base not in used:
        used.add(base)
        return base
    # Avoid duplicates while respecting 31-char Excel limit.
    n = 2
    while True:
        suffix = f'_{n}'
        candidate = base[:31 - len(suffix)] + suffix
        if candidate not in used:
            used.add(candidate)
            return candidate
        n += 1

# Read source workbook.
src_wb = load_workbook(input_path, data_only=False)
if 'tasks' not in src_wb.sheetnames:
    raise ValueError('Input workbook does not contain a tasks worksheet')
src_ws = src_wb['tasks']
headers = [cell.value for cell in src_ws[1]]
try:
    owner_col = headers.index('Owner') + 1
    status_col = headers.index('Status') + 1
except ValueError as exc:
    raise ValueError('tasks worksheet must contain Owner and Status columns') from exc

# Group rows by Owner key; blank Owner becomes Unassigned but row's original blank value is preserved.
groups = OrderedDict()
for row in src_ws.iter_rows(min_row=2, max_row=src_ws.max_row):
    values = [cell.value for cell in row[:len(headers)]]
    if all(v is None for v in values):
        continue
    owner_val = values[owner_col - 1]
    owner_key = 'Unassigned' if is_blank(owner_val) else str(owner_val)
    groups.setdefault(owner_key, []).append(values)

# Create output workbook.
out_wb = Workbook()
default_ws = out_wb.active
out_wb.remove(default_ws)

used_sheet_names = set()
summary_records = []

# Summary first.
summary_ws = out_wb.create_sheet('Summary')
summary_headers = ['Owner', 'SheetName', 'Total', 'Pass', 'Fail']
summary_ws.append(summary_headers)
used_sheet_names.add('Summary')

# Per-owner sheets.
for owner_key, rows in groups.items():
    sheet_name = legal_sheet_name(owner_key, used_sheet_names)
    ws = out_wb.create_sheet(sheet_name)
    ws.append(headers)
    total = 0
    pass_count = 0
    fail_count = 0
    for vals in rows:
        ws.append(vals)
        total += 1
        status = vals[status_col - 1]
        if status == 'Pass':
            pass_count += 1
        elif status == 'Fail':
            fail_count += 1
        else:
            raise ValueError(f'Unexpected status {status!r}; expected Pass or Fail')
    summary_records.append([owner_key, sheet_name, total, pass_count, fail_count])

for record in summary_records:
    summary_ws.append(record)

# Basic widths for readability; does not change data.
for ws in out_wb.worksheets:
    for col_cells in ws.columns:
        max_len = 0
        col_letter = col_cells[0].column_letter
        for cell in col_cells:
            max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[col_letter].width = min(max_len + 2, 40)

out_wb.save(out_path)
src_wb.close()
out_wb.close()
print('Wrote:', out_path)
print('Groups:', summary_records)

# Verification: reopen every output file and confirm requirements.
check_wb = load_workbook(out_path, data_only=False)
print('Output sheets:', check_wb.sheetnames)
assert 'Summary' in check_wb.sheetnames, 'Summary sheet missing'
summary = check_wb['Summary']
summary_header_values = [cell.value for cell in summary[1]]
assert summary_header_values == summary_headers, f'Summary headers differ: {summary_header_values}'

summary_data = []
for row in summary.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    owner, sheet_name, total, pass_count, fail_count = row
    assert isinstance(sheet_name, str) and sheet_name in check_wb.sheetnames, f'Bad SheetName for {owner}: {sheet_name}'
    assert isinstance(total, int), f'Total not numeric integer for {owner}'
    assert isinstance(pass_count, int), f'Pass not numeric integer for {owner}'
    assert isinstance(fail_count, int), f'Fail not numeric integer for {owner}'
    assert total == pass_count + fail_count, f'Summary counts do not add up for {owner}'
    summary_data.append((owner, sheet_name, total, pass_count, fail_count))

# Confirm there is exactly one Summary row per group and one per-owner sheet.
assert len(summary_data) == len(groups), f'Summary row count {len(summary_data)} != group count {len(groups)}'
assert len(check_wb.sheetnames) == len(groups) + 1, 'Unexpected number of sheets'

for owner_key, sheet_name, total, pass_count, fail_count in summary_data:
    ws = check_wb[sheet_name]
    per_headers = [cell.value for cell in ws[1]]
    assert per_headers == headers, f'Headers changed in {sheet_name}: {per_headers}'
    rows = [list(r) for r in ws.iter_rows(min_row=2, values_only=True) if not all(v is None for v in r)]
    assert len(rows) == total, f'Row count mismatch for {owner_key}'
    # Confirm owner original values preserved: blank owners are still blank/None in Unassigned rows.
    expected_rows = groups[owner_key]
    assert rows == expected_rows, f'Rows differ for {owner_key}'
    actual_pass = sum(1 for r in rows if r[status_col - 1] == 'Pass')
    actual_fail = sum(1 for r in rows if r[status_col - 1] == 'Fail')
    assert actual_pass == pass_count and actual_fail == fail_count, f'Count mismatch for {owner_key}'
    assert len(sheet_name) <= 31 and not INVALID_SHEET_CHARS.search(sheet_name), f'Illegal sheet name: {sheet_name}'

print('Verification passed: owners.xlsx exists; Summary columns Owner, SheetName, Total, Pass, Fail; one sheet per Owner group; blank Owner rows grouped as Unassigned with original Owner value preserved; legal unique sheet names; source headers and row data preserved; summary counts are numeric and match sheet contents.')
check_wb.close()
