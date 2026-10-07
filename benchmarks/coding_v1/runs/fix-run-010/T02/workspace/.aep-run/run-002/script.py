from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV preserving text values exactly as strings.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames
    if headers is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in headers:
        raise ValueError('results.csv is missing required Status column')
    if 'Device ID' not in headers:
        raise ValueError('results.csv is missing required Device ID column')
    source_rows = list(reader)

failed_rows = [row for row in source_rows if (row.get('Status') or '').strip().lower() == 'fail']

# Create workbook and worksheet with all original columns in original order.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = Font(bold=True)

for row_idx, row in enumerate(failed_rows, start=2):
    for col_idx, header in enumerate(headers, start=1):
        value = row.get(header, '')
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        # Preserve Device ID as text so leading zeros remain.
        if header == 'Device ID':
            cell.number_format = '@'
            cell.data_type = 's'

# Also mark entire Device ID column as text style.
device_col = headers.index('Device ID') + 1
for cell in ws.iter_cols(min_col=device_col, max_col=device_col, min_row=1, max_row=max(1, len(failed_rows) + 1)):
    for c in cell:
        c.number_format = '@'

wb.save(output_path)
print(f'Wrote: {output_path}')
print(f'Source data rows: {len(source_rows)}')
print(f'Failed rows written: {len(failed_rows)}')
print('Failed Device IDs in order:', [row['Device ID'] for row in failed_rows])

# Verification: reopen output and confirm requirements.
wb2 = load_workbook(output_path, data_only=False)
if 'failed_items' not in wb2.sheetnames:
    raise AssertionError('Missing failed_items worksheet')
ws2 = wb2['failed_items']
actual_headers = [ws2.cell(row=1, column=i).value for i in range(1, len(headers) + 1)]
if actual_headers != headers:
    raise AssertionError(f'Headers changed: {actual_headers!r} != {headers!r}')

actual_rows = []
for r in range(2, ws2.max_row + 1):
    actual_rows.append([ws2.cell(row=r, column=c).value for c in range(1, len(headers) + 1)])
expected_rows = [[row.get(h, '') for h in headers] for row in failed_rows]
if actual_rows != expected_rows:
    raise AssertionError(f'Rows differ: {actual_rows!r} != {expected_rows!r}')

# Confirm all output Status values match trimmed case-insensitive Fail and row order corresponds to source order.
status_col = headers.index('Status') + 1
for r in range(2, ws2.max_row + 1):
    status_value = ws2.cell(row=r, column=status_col).value
    if (status_value or '').strip().lower() != 'fail':
        raise AssertionError(f'Non-fail status in output row {r}: {status_value!r}')

# Confirm Device ID values are stored as strings and have text formatting.
for r in range(2, ws2.max_row + 1):
    cell = ws2.cell(row=r, column=device_col)
    if not isinstance(cell.value, str):
        raise AssertionError(f'Device ID at row {r} is not text: {cell.value!r} ({type(cell.value).__name__})')
    if cell.number_format != '@':
        raise AssertionError(f'Device ID at row {r} is not formatted as text: {cell.number_format!r}')

print('Verification passed:')
print('- output file exists')
print('- worksheet name is failed_items')
print('- all original columns preserved in order')
print('- rows match Status stripped/lower == fail, preserving source order')
print('- Device ID values are strings with text format, preserving leading zeros')
