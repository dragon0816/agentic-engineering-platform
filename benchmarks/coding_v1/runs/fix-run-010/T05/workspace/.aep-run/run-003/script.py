from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
out_path = base / 'output' / 'merged_results.xlsx'
if not out_path.exists():
    raise FileNotFoundError(out_path)
wb = load_workbook(out_path, data_only=False)
print('Output file:', out_path)
print('Sheets:', wb.sheetnames)
assert 'matched' in wb.sheetnames, 'missing matched sheet'
assert 'unmatched' in wb.sheetnames, 'missing unmatched sheet'

ws_m = wb['matched']
ws_u = wb['unmatched']
matched_headers = [c.value for c in ws_m[1]]
unmatched_headers = [c.value for c in ws_u[1]]
print('matched headers:', matched_headers)
print('unmatched headers:', unmatched_headers)
assert matched_headers == ['Device ID', 'Test', 'Result', 'Name', 'Location']
assert unmatched_headers == ['Source', 'Device ID', 'Name', 'Location', 'Test', 'Result']

matched = list(ws_m.iter_rows(min_row=2, values_only=True))
unmatched = list(ws_u.iter_rows(min_row=2, values_only=True))
print('matched data:', matched)
print('unmatched data:', unmatched)
assert len(matched) == 3
assert len(unmatched) == 2
assert matched == [
    ('001', 'RF-1', 'Pass', 'Alpha', 'Lab A'),
    ('001', 'RF-2', 'Fail', 'Alpha', 'Lab A'),
    ('002', 'RF-1', 'Pass', 'Beta', 'Lab B'),
]
assert unmatched == [
    ('devices', '003', 'Gamma', 'Lab C', None, None),
    ('tests', '999', None, None, 'RF-9', 'Fail'),
]

for ws in [ws_m, ws_u]:
    headers = [c.value for c in ws[1]]
    did_col = headers.index('Device ID') + 1
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=did_col)
        assert isinstance(cell.value, str), f'{ws.title} row {row} Device ID not text'
        assert cell.number_format == '@', f'{ws.title} row {row} Device ID number_format not text'
print('Checks passed: output workbook exists; required sheets, headers, row counts, expected matched/unmatched rows, Source values, and text Device ID formatting verified.')
wb.close()
