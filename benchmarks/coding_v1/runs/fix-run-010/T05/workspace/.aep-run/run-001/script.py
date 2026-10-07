from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
input_dir = base / 'input'
print('Workspace:', base)
print('Input files:')
for p in sorted(input_dir.glob('*')):
    print(' -', p.name)

for fname in ['devices.xlsx', 'tests.xlsx']:
    path = input_dir / fname
    print('\n===', fname, '===')
    wb = load_workbook(path, data_only=False, read_only=True)
    print('Sheets:', wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True))
    print('Dimensions:', ws.max_row, 'rows x', ws.max_column, 'cols')
    for i, row in enumerate(rows, 1):
        print(i, row)
    wb.close()
