from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font

base = Path.cwd()
input_dir = base / 'input'
output_dir = base / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
out_path = output_dir / 'merged_results.xlsx'

def read_sheet(path):
    wb = load_workbook(path, data_only=False)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        raise ValueError(f'{path.name} is empty')
    headers = [str(c) if c is not None else '' for c in rows[0]]
    data = []
    for row in rows[1:]:
        # pad row to headers length
        vals = list(row[:len(headers)]) + [None] * max(0, len(headers) - len(row))
        rec = {h: vals[i] for i, h in enumerate(headers)}
        data.append(rec)
    return headers, data

def id_text(value):
    if value is None:
        return ''
    return str(value)

dev_headers, devices = read_sheet(input_dir / 'devices.xlsx')
test_headers, tests = read_sheet(input_dir / 'tests.xlsx')
if 'Device ID' not in dev_headers:
    raise ValueError('devices.xlsx lacks Device ID column')
if 'Device ID' not in test_headers:
    raise ValueError('tests.xlsx lacks Device ID column')

# devices has unique IDs by requirement; still guard against duplicate to avoid silent overwrite.
dev_by_id = {}
for rec in devices:
    did = id_text(rec.get('Device ID'))
    if did in dev_by_id:
        raise ValueError(f'Duplicate Device ID in devices.xlsx: {did!r}')
    dev_by_id[did] = rec

# Matched: one row per test with a matching device. Use one Device ID column, then test-only columns, then device-only columns.
test_non_id = [h for h in test_headers if h != 'Device ID']
dev_non_id = [h for h in dev_headers if h != 'Device ID']
matched_headers = ['Device ID'] + test_non_id + dev_non_id

matched_rows = []
matched_device_ids = set()
unmatched_tests = []
for test in tests:
    did = id_text(test.get('Device ID'))
    dev = dev_by_id.get(did)
    if dev is None:
        unmatched_tests.append(test)
    else:
        matched_device_ids.add(did)
        row = {'Device ID': did}
        for h in test_non_id:
            row[h] = test.get(h)
        for h in dev_non_id:
            row[h] = dev.get(h)
        matched_rows.append(row)

unmatched_devices = [rec for rec in devices if id_text(rec.get('Device ID')) not in matched_device_ids]

# Unmatched: union of Source + original source column names. Use blanks where a source lacks a column.
union_headers = ['Source']
for h in dev_headers + test_headers:
    if h not in union_headers:
        union_headers.append(h)

unmatched_rows = []
for rec in unmatched_devices:
    row = {h: None for h in union_headers}
    row['Source'] = 'devices'
    for h in dev_headers:
        row[h] = id_text(rec[h]) if h == 'Device ID' else rec[h]
    unmatched_rows.append(row)
for rec in unmatched_tests:
    row = {h: None for h in union_headers}
    row['Source'] = 'tests'
    for h in test_headers:
        row[h] = id_text(rec[h]) if h == 'Device ID' else rec[h]
    unmatched_rows.append(row)

wb = Workbook()
ws_m = wb.active
ws_m.title = 'matched'
ws_u = wb.create_sheet('unmatched')

def write_table(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for rec in rows:
        ws.append([id_text(rec.get(h)) if h == 'Device ID' and rec.get(h) is not None else rec.get(h) for h in headers])
    # Make Device ID text-formatted to preserve leading zeroes.
    for idx, h in enumerate(headers, start=1):
        if h == 'Device ID':
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row):
                for c in cell:
                    c.number_format = '@'
                    if c.value is not None:
                        c.value = str(c.value)
    # Simple widths
    for idx, h in enumerate(headers, start=1):
        max_len = len(str(h))
        for c in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row):
            for cell in c:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max_len + 2, 40)

write_table(ws_m, matched_headers, matched_rows)
write_table(ws_u, union_headers, unmatched_rows)

wb.save(out_path)
print('Wrote:', out_path)
print('matched rows:', len(matched_rows))
print('unmatched devices:', len(unmatched_devices))
print('unmatched tests:', len(unmatched_tests))
print('unmatched rows:', len(unmatched_rows))
