from pathlib import Path
import re
from decimal import Decimal
from openpyxl import Workbook, load_workbook

ws = Path.cwd()
input_path = ws / 'input' / 'specs.md'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Parse exactly the product sections introduced by ## headings.
section_pattern = re.compile(r'^##\s+(.+?)\s*$([\s\S]*?)(?=^##\s+|\Z)', re.MULTILINE)
sections = section_pattern.findall(text)

unit_factor_to_mhz = {
    'ghz': Decimal('1000'),
    'mhz': Decimal('1'),
    'khz': Decimal('0.001'),
    'hz': Decimal('0.000001'),
}

def dec_to_cell(d: Decimal):
    """Return an int when exact, otherwise a float suitable for Excel numeric cells."""
    d = d.normalize()
    if d == d.to_integral():
        return int(d)
    return float(d)

def convert_to_mhz(value_text: str, unit_text: str):
    factor = unit_factor_to_mhz[unit_text.lower()]
    return dec_to_cell(Decimal(value_text) * factor)

def parse_frequency(value: str):
    # Accept ranges such as 2.4–2.5 GHz or 900-930 MHz. Do not infer a unit if absent.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*[–—-]\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', value, re.IGNORECASE)
    if not m:
        return '', ''
    return convert_to_mhz(m.group(1), m.group(3)), convert_to_mhz(m.group(2), m.group(3))

def parse_bandwidth(value: str):
    # Only convert explicitly written bandwidth values with units.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', value, re.IGNORECASE)
    if not m:
        return ''
    return convert_to_mhz(m.group(1), m.group(2))

def parse_interface(value: str):
    # Multiple interfaces in the source are slash-separated; output with semicolon separators.
    parts = [p.strip() for p in re.split(r'\s*/\s*', value.strip()) if p.strip()]
    return ';'.join(parts)

rows = []
for heading, body in sections:
    fields = {}
    for line in body.splitlines():
        m = re.match(r'\s*-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            fields[m.group(1).strip()] = m.group(2).strip()
    if not fields:
        continue
    freq_min, freq_max = ('', '')
    if fields.get('Frequency'):
        freq_min, freq_max = parse_frequency(fields['Frequency'])
    bandwidth_mhz = ''
    if fields.get('Bandwidth'):
        bandwidth_mhz = parse_bandwidth(fields['Bandwidth'])
    interface = ''
    if fields.get('Interface'):
        interface = parse_interface(fields['Interface'])
    rows.append({
        'Model': fields.get('Model', ''),
        'FrequencyMinMHz': freq_min,
        'FrequencyMaxMHz': freq_max,
        'BandwidthMHz': bandwidth_mhz,
        'Interface': interface,
    })

headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
wb = Workbook()
wsheet = wb.active
wsheet.title = 'specs'
wsheet.append(headers)
for row in rows:
    wsheet.append([row[h] for h in headers])
wb.save(output_path)

print('Created:', output_path)
print('Parsed rows:', len(rows))
for row in rows:
    print(row)

# Verification: reopen the saved workbook and confirm requirements.
check_wb = load_workbook(output_path, data_only=True)
check_ws = check_wb['specs']
read_headers = [cell.value for cell in check_ws[1]]
read_rows = [dict(zip(read_headers, values)) for values in check_ws.iter_rows(min_row=2, values_only=True)]

assert read_headers == headers, f'headers mismatch: {read_headers}'
assert len(read_rows) == 3, f'expected 3 product rows, got {len(read_rows)}'
expected = [
    {'Model': 'RX-100', 'FrequencyMinMHz': 2400, 'FrequencyMaxMHz': 2500, 'BandwidthMHz': 80, 'Interface': 'USB;LAN'},
    {'Model': 'RX-200', 'FrequencyMinMHz': 900, 'FrequencyMaxMHz': 930, 'BandwidthMHz': 20, 'Interface': 'LAN'},
    {'Model': 'RX-300', 'FrequencyMinMHz': 5000, 'FrequencyMaxMHz': 6000, 'BandwidthMHz': None, 'Interface': 'USB'},
]
assert read_rows == expected, f'content mismatch: {read_rows}'
for r in read_rows:
    assert isinstance(r['Model'], str) and r['Model'].startswith('RX-')
    assert isinstance(r['FrequencyMinMHz'], (int, float))
    assert isinstance(r['FrequencyMaxMHz'], (int, float))
    if r['BandwidthMHz'] is not None:
        assert isinstance(r['BandwidthMHz'], (int, float))
    assert isinstance(r['Interface'], str)
print('Verification passed: specs.xlsx reopened; headers, 3 rows, MHz conversions, semicolon interfaces, and blank missing bandwidth were confirmed.')