from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read CSV with utf-8-sig to support UTF-8 BOM and newline='' so quoted newlines are parsed correctly.
with open(input_path, 'r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.reader(f))

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
try:
    device_id_col = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_id_col = None

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'

for r_idx, row in enumerate(rows, start=1):
    # Preserve all columns and blanks. Pad short rows if any; fail on extra columns to avoid silent schema changes.
    if len(row) > len(headers):
        raise ValueError(f'Row {r_idx} has {len(row)} columns; header has {len(headers)}')
    row_to_write = row + [''] * (len(headers) - len(row))
    for c_idx, value in enumerate(row_to_write, start=1):
        cell = ws.cell(row=r_idx, column=c_idx)
        if value == '':
            cell.value = None  # blank Excel cell
        else:
            cell.value = str(value)
        if c_idx == device_id_col:
            cell.number_format = '@'  # force text display/storage for Device ID, preserving leading zeros

# A light bold header is presentation-only and does not add/change data.
for cell in ws[1]:
    cell.font = Font(bold=True)

wb.save(output_path)
print('Wrote:', output_path)
print('CSV rows including header:', len(rows))
print('CSV columns:', len(headers), headers)

# Verification: reopen output and compare every cell against CSV data, treating blank Excel cells as empty strings.
wb2 = load_workbook(output_path, data_only=False)
ws2 = wb2.active
actual_rows = []
for r in range(1, ws2.max_row + 1):
    actual_row = []
    for c in range(1, len(headers) + 1):
        v = ws2.cell(row=r, column=c).value
        actual_row.append('' if v is None else str(v))
    actual_rows.append(actual_row)

expected_rows = [row + [''] * (len(headers) - len(row)) for row in rows]

checks = []
checks.append(('output file exists', output_path.exists()))
checks.append(('no added index column / same column count', ws2.max_column == len(headers)))
checks.append(('same row count including duplicate rows', ws2.max_row == len(expected_rows)))
checks.append(('headers and column order preserved', actual_rows[0] == headers))
checks.append(('all cell values preserved including Chinese, commas, quoted newlines and blanks', actual_rows == expected_rows))

if device_id_col:
    data_ids = [ws2.cell(row=r, column=device_id_col).value for r in range(2, ws2.max_row + 1)]
    id_formats = [ws2.cell(row=r, column=device_id_col).number_format for r in range(1, ws2.max_row + 1)]
    expected_ids = [row[device_id_col - 1] for row in expected_rows[1:]]
    checks.append(('Device ID values are strings and preserve leading zeros', data_ids == expected_ids and all(isinstance(x, str) for x in data_ids if x is not None)))
    checks.append(('Device ID column formatted as text', all(fmt == '@' for fmt in id_formats)))

failed = [name for name, ok in checks if not ok]
for name, ok in checks:
    print(('PASS' if ok else 'FAIL') + ': ' + name)

print('Reopened workbook dimensions:', ws2.max_row, 'rows x', ws2.max_column, 'columns')
print('Reopened headers:', actual_rows[0])
print('Reopened first data rows:', actual_rows[1:4])
print('Reopened duplicate row check row2 vs row5:', actual_rows[1], actual_rows[4] if len(actual_rows) > 4 else None)

if failed:
    raise AssertionError('Verification failed: ' + '; '.join(failed))
