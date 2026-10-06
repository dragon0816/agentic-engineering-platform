import csv
import os
import re
import urllib.request
from html.parser import HTMLParser
from pathlib import Path

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)

URLS = [
    'http://127.0.0.1:56996/t12/a',
    'http://127.0.0.1:56996/t12/b',
    'http://127.0.0.1:56996/t12/c',
]

class SimpleHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.events = []
        self.stack = []
    def handle_starttag(self, tag, attrs):
        d = dict(attrs)
        self.events.append(('start', tag, d, ''))
        self.stack.append((tag, d))
    def handle_endtag(self, tag):
        self.events.append(('end', tag, {}, ''))
        # pop up to tag if present
        for i in range(len(self.stack)-1, -1, -1):
            if self.stack[i][0] == tag:
                self.stack = self.stack[:i]
                break
    def handle_data(self, data):
        if data.strip():
            self.events.append(('data', self.stack[-1][0] if self.stack else '', self.stack[-1][1] if self.stack else {}, data.strip()))

def text_from_html(html):
    parser = SimpleHTML()
    parser.feed(html)
    return parser.events

def clean_price(price_text):
    s = price_text.strip()
    # currency: any leading non-digit/sign chars or leading ASCII letters before the amount
    m = re.search(r'[-+]?\d[\d.,]*', s)
    if not m:
        raise ValueError(f'No numeric price in {price_text!r}')
    num = m.group(0)
    currency = (s[:m.start()] + s[m.end():]).strip()
    # remove prose/spacing left after extracting amount
    currency = currency.replace('Price:', '').strip()
    # If both separators exist, rightmost separator is decimal separator.
    if ',' in num and '.' in num:
        if num.rfind(',') > num.rfind('.'):
            normalized = num.replace('.', '').replace(',', '.')
        else:
            normalized = num.replace(',', '')
    elif ',' in num:
        # Treat comma as thousands if exactly three digits follow and no decimal-like ending; otherwise decimal.
        parts = num.split(',')
        if len(parts[-1]) == 3 and all(len(p) == 3 for p in parts[1:]):
            normalized = num.replace(',', '')
        else:
            normalized = num.replace(',', '.')
    else:
        normalized = num
    return float(normalized), currency

def clean_memory(mem_text):
    s = mem_text.strip()
    m = re.search(r'(\d+(?:\.\d+)?)\s*(GB|MB)\b', s, flags=re.I)
    if not m:
        raise ValueError(f'No memory in {mem_text!r}')
    value = float(m.group(1))
    unit = m.group(2).upper()
    gb = value if unit == 'GB' else value / 1024.0
    return int(gb) if gb.is_integer() else gb

def fetch(url):
    with urllib.request.urlopen(url, timeout=10) as response:
        return response.read().decode('utf-8', errors='replace')

def parse_a(html, url):
    # Table rows: skip header, collect td data in groups of 3.
    rows = re.findall(r'<tr>(.*?)</tr>', html, flags=re.S|re.I)
    out = []
    for row in rows[1:]:
        cells = [re.sub(r'<.*?>', '', c, flags=re.S).strip() for c in re.findall(r'<td>(.*?)</td>', row, flags=re.S|re.I)]
        if len(cells) == 3:
            price, currency = clean_price(cells[1])
            out.append({
                'Model': cells[0],
                'Price': price,
                'Currency': currency,
                'MemoryGB': clean_memory(cells[2]),
                'SourceURL': url,
            })
    return out

def parse_b(html, url):
    items = re.findall(r'<div class="item">(.*?)</div>', html, flags=re.S|re.I)
    out = []
    for item in items:
        model = re.search(r'<h2>(.*?)</h2>', item, flags=re.S|re.I).group(1).strip()
        price_text = re.search(r'<p class="price">(.*?)</p>', item, flags=re.S|re.I).group(1).strip()
        mem_text = re.search(r'<p class="mem">(.*?)</p>', item, flags=re.S|re.I).group(1).strip()
        price, currency = clean_price(price_text)
        out.append({'Model': model, 'Price': price, 'Currency': currency, 'MemoryGB': clean_memory(mem_text), 'SourceURL': url})
    return out

def parse_c(html, url):
    lis = re.findall(r'<li\b([^>]*)>(.*?)</li>', html, flags=re.S|re.I)
    out = []
    for attrs, body in lis:
        attr_model = re.search(r'data-model="([^"]+)"', attrs)
        text = re.sub(r'<.*?>', '', body, flags=re.S).strip()
        parts = [p.strip() for p in text.split('—')]
        model = attr_model.group(1) if attr_model else parts[0]
        price, currency = clean_price(parts[1])
        out.append({'Model': model, 'Price': price, 'Currency': currency, 'MemoryGB': clean_memory(parts[2]), 'SourceURL': url})
    return out

all_rows = []
for url in URLS:
    html = fetch(url)
    print(f'Fetched {url}: {len(html)} characters')
    if url.endswith('/a'):
        rows = parse_a(html, url)
    elif url.endswith('/b'):
        rows = parse_b(html, url)
    elif url.endswith('/c'):
        rows = parse_c(html, url)
    else:
        rows = []
    print(f'Parsed {len(rows)} rows from {url}')
    all_rows.extend(rows)

columns = ['Model', 'Price', 'Currency', 'MemoryGB', 'SourceURL']
csv_path = OUTPUT / 'merged_catalog.csv'
with csv_path.open('w', encoding='utf-8-sig', newline='') as f:
    writer = csv.DictWriter(f, fieldnames=columns)
    writer.writeheader()
    for row in all_rows:
        writer.writerow(row)

# Also write an XLSX copy for spreadsheet use.
try:
    from openpyxl import Workbook, load_workbook
    xlsx_path = OUTPUT / 'merged_catalog.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(columns)
    for row in all_rows:
        ws.append([row[c] for c in columns])
    wb.save(xlsx_path)
except Exception as e:
    xlsx_path = None
    print('XLSX write skipped:', repr(e))

# Verification: reopen saved outputs and check stated requirements.
with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    read_rows = list(csv.DictReader(f))
assert list(read_rows[0].keys()) == columns, f'CSV columns mismatch: {read_rows[0].keys()}'
assert len(read_rows) == 6, f'Expected 6 rows, got {len(read_rows)}'
expected = [
    ('NX-1', 12000.0, 'NT$', 16, URLS[0]),
    ('NX-2', 8500.0, 'NT$', 8, URLS[0]),
    ('NX-1', 399.5, 'USD', 16, URLS[1]),
    ('NX-3', 1250.0, 'USD', 32, URLS[1]),
    ('NX-4', 499.0, '€', 4, URLS[2]),
    ('NX-5', 1299.0, '€', 64, URLS[2]),
]
for row, exp in zip(read_rows, expected):
    assert row['Model'] == exp[0]
    assert abs(float(row['Price']) - exp[1]) < 1e-9, (row, exp)
    assert row['Currency'] == exp[2]
    assert abs(float(row['MemoryGB']) - float(exp[3])) < 1e-9, (row, exp)
    assert row['SourceURL'] == exp[4]

if xlsx_path:
    wb2 = load_workbook(xlsx_path, data_only=True)
    ws2 = wb2['Products']
    assert [cell.value for cell in ws2[1]] == columns
    assert ws2.max_row == 7 and ws2.max_column == 5
    for i, exp in enumerate(expected, start=2):
        vals = [ws2.cell(i, j).value for j in range(1, 6)]
        assert vals[0] == exp[0]
        assert abs(float(vals[1]) - exp[1]) < 1e-9
        assert vals[2] == exp[2]
        assert abs(float(vals[3]) - float(exp[3])) < 1e-9
        assert vals[4] == exp[4]

print('WROTE', csv_path)
if xlsx_path:
    print('WROTE', xlsx_path)
print('VERIFIED columns, 6 rows, duplicate model retained as separate rows, numeric Price, per-source Currency, MemoryGB conversion, and SourceURL values from saved output files.')
for row in read_rows:
    print(row)
