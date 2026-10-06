from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
path = ws/'input'/'tasks.xlsx'
print('tasks.xlsx exists:', path.exists(), 'size:', path.stat().st_size if path.exists() else None)
wb = load_workbook(path, read_only=True, data_only=False)
print('sheetnames:', wb.sheetnames)
sh = wb['tasks']
headers = [c.value for c in next(sh.iter_rows(min_row=1, max_row=1))]
print('headers:', headers)
for i,row in enumerate(sh.iter_rows(min_row=2, max_row=6, values_only=True), start=2):
    print('row', i, row)
print('max_row', sh.max_row, 'max_col', sh.max_column)
wb.close()