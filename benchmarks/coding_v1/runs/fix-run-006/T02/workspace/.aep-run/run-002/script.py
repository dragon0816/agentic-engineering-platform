from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV preserving original values as strings.
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    if 'Status' not in headers:
        raise ValueError('Required column Status not found')
    filtered_rows = []
    total_rows = 0
    for row in reader:
        total_rows += 1
        if (row.get('Status') or '').strip().lower() == 'fail':
            # Keep original column order and original string values.
            filtered_rows.append([row.get(h, '') for h in headers])

# Write workbook.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    cell.font = Font(bold=True)
    cell.number_format = '@'

# Ensure Device ID is written as text and all source values remain as strings.
device_col_idx = headers.index('Device ID') + 1 if 'Device ID' in headers else None
for row_idx, row_values in enumerate(filtered_rows, start=2):
    for col_idx, value in enumerate(row_values, start=1):
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        if col_idx == device_col_idx:
            cell.number_format = '@'
            cell.data_type = 's'

# Set text format for whole Device ID column including blanks/data.
if device_col_idx:
    col_letter = ws.cell(row=1, column=device_col_idx).column_letter
    for cell in ws[col_letter]:
        cell.number_format = '@'

wb.save(output_path)

# Verification by reopening saved output.
check_wb = load_workbook(output_path, data_only=False)
if 'failed_items' not in check_wb.sheetnames:
    raise AssertionError('Worksheet failed_items not found')
check_ws = check_wb['failed_items']
read_headers = [check_ws.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
if read_headers != headers:
    raise AssertionError(f'Headers mismatch: {read_headers} != {headers}')

read_rows = []
for r in range(2, check_ws.max_row + 1):
    read_rows.append([check_ws.cell(row=r, column=c).value for c in range(1, len(headers)+1)])
expected_rows = filtered_rows
if read_rows != expected_rows:
    raise AssertionError(f'Rows mismatch: {read_rows} != {expected_rows}')

# Confirm Status filter condition and original order via Device ID values.
status_idx = headers.index('Status')
for row in read_rows:
    if (row[status_idx] or '').strip().lower() != 'fail':
        raise AssertionError(f'Non-fail status present: {row}')

if device_col_idx:
    device_values = [check_ws.cell(row=r, column=device_col_idx).value for r in range(2, check_ws.max_row + 1)]
    device_types = [check_ws.cell(row=r, column=device_col_idx).data_type for r in range(2, check_ws.max_row + 1)]
    expected_device_values = [row[device_col_idx-1] for row in expected_rows]
    if device_values != expected_device_values:
        raise AssertionError(f'Device ID values changed: {device_values} != {expected_device_values}')
    if not all(isinstance(v, str) for v in device_values):
        raise AssertionError(f'Device ID values are not all text: {device_values}')
    if not all(t == 's' for t in device_types):
        raise AssertionError(f'Device ID cell types are not all string: {device_types}')

print('Created:', output_path)
print('Source rows:', total_rows)
print('Filtered fail rows:', len(filtered_rows))
print('Worksheet:', check_ws.title)
print('Headers checked:', read_headers)
print('Rows checked:')
for row in read_rows:
    print(row)
if device_col_idx:
    print('Device ID text values checked:', device_values)
    print('Device ID data types checked:', device_types)
print('Input file was not modified.')