from pathlib import Path
from html.parser import HTMLParser
from urllib.request import urlopen
from urllib.parse import urljoin
import re
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
CATALOG_URL = 'http://127.0.0.1:56996/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self, base_url):
        super().__init__(convert_charrefs=True)
        self.base_url = base_url
        self.in_catalog = False
        self.catalog_depth = 0
        self.card_stack = []
        self.current = None
        self.in_link = False
        self.in_price = False
        self.products = []

    def _attrs(self, attrs):
        return dict(attrs)

    def handle_starttag(self, tag, attrs):
        attrs_d = self._attrs(attrs)
        if self.in_catalog:
            self.catalog_depth += 1
        if tag == 'main' and attrs_d.get('id') == 'catalog' and not self.in_catalog:
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if not self.in_catalog:
            return
        classes = set((attrs_d.get('class') or '').split())
        if tag == 'div' and 'card' in classes:
            self.card_stack.append('card')
            self.current = {'name_parts': [], 'price_parts': [], 'url': ''}
        elif self.current is not None:
            self.card_stack.append(None)
        if self.current is not None and tag == 'a' and not self.current.get('url'):
            self.in_link = True
            self.current['url'] = urljoin(self.base_url, attrs_d.get('href', ''))
        if self.current is not None and 'price' in classes:
            self.in_price = True

    def handle_endtag(self, tag):
        if self.current is not None and tag == 'a':
            self.in_link = False
        # Only close price spans/classes on any endtag while in_price; enough for simple price element.
        if self.current is not None and self.in_price and tag in ('span', 'div', 'p'):
            self.in_price = False
        if self.in_catalog:
            if tag == 'div' and self.current is not None and self.card_stack:
                marker = self.card_stack.pop()
                if marker == 'card':
                    product = {
                        'name': ' '.join(''.join(self.current['name_parts']).split()),
                        'price_text': ' '.join(''.join(self.current['price_parts']).split()),
                        'url': self.current['url'],
                    }
                    self.products.append(product)
                    self.current = None
                    self.in_link = False
                    self.in_price = False
            elif self.current is not None and self.card_stack:
                self.card_stack.pop()
            self.catalog_depth -= 1
            if tag == 'main' or self.catalog_depth <= 0:
                self.in_catalog = False
                self.catalog_depth = 0
                self.current = None
                self.card_stack.clear()
                self.in_link = False
                self.in_price = False

    def handle_data(self, data):
        if self.current is not None:
            if self.in_link:
                self.current['name_parts'].append(data)
            if self.in_price:
                self.current['price_parts'].append(data)

def parse_price(text):
    s = text.strip()
    m = re.match(r'^([^\d\-+.,]*)([\d,]+(?:\.\d+)?)', s)
    if not m:
        raise ValueError(f'Cannot parse price: {text!r}')
    currency = m.group(1).strip()
    number_text = m.group(2).replace(',', '')
    price = float(number_text) if '.' in number_text else int(number_text)
    return currency, price

# Fetch and parse catalog HTML.
with urlopen(CATALOG_URL, timeout=10) as response:
    html = response.read().decode(response.headers.get_content_charset() or 'utf-8', errors='replace')

parser = CatalogParser(CATALOG_URL)
parser.feed(html)

rows = []
for item in parser.products:
    currency, price = parse_price(item['price_text'])
    rows.append({
        '產品名稱': item['name'],
        '價格': price,
        '幣別': currency,
        '產品網址': item['url'],
    })

# Write workbook.
wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(OUTFILE)

# Reopen and verify stated requirements from saved output.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
saved_headers = [cell.value for cell in ws2[1]]
saved_rows = list(ws2.iter_rows(min_row=2, values_only=True))
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
expected_prices = [1299, 2450, 99, 560, 12000]
expected_currency = ['NT$'] * 5
expected_urls = [urljoin(CATALOG_URL, path) for path in ['/p/router-a','/p/switch-b','/p/cable-c','/p/adapter-d','/p/antenna-e']]
assert OUTFILE.exists(), 'products.xlsx was not created'
assert saved_headers == headers, f'headers mismatch: {saved_headers}'
assert len(saved_rows) == 5, f'row count mismatch: {len(saved_rows)}'
for i, row in enumerate(saved_rows):
    name, price, currency, url = row
    assert name == expected_names[i], f'name/order mismatch row {i+2}: {name}'
    assert price == expected_prices[i] and isinstance(price, (int, float)), f'price mismatch/type row {i+2}: {price!r}'
    assert currency == expected_currency[i], f'currency mismatch row {i+2}: {currency}'
    assert url == expected_urls[i] and str(url).startswith('http://127.0.0.1:56996/'), f'url mismatch row {i+2}: {url}'
assert 'Sponsored Router' not in [r[0] for r in saved_rows], 'ad item included'
assert '推薦：高速網卡' not in [r[0] for r in saved_rows], 'recommended/ad item included'

print(f'Wrote: {OUTFILE}')
print('Saved headers:', saved_headers)
print('Saved row count:', len(saved_rows))
print('Saved rows:')
for r in saved_rows:
    print(r)
print('Checks passed: output file exists; headers are correct; 5 catalog rows in page order; prices are numeric; currency is NT$; URLs are complete; ads/recommendations excluded.')