from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
    if p.suffix.lower() in ['.xlsx', '.xlsm']:
        wb = load_workbook(p, read_only=True, data_only=False)
        print('  sheets:', wb.sheetnames)
        for s in wb.sheetnames:
            sh = wb[s]
            rows = list(sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True))
            print('  sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
            for r in rows:
                print('   ', r)
        wb.close()