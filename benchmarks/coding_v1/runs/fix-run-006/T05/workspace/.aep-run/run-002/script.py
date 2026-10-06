from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
outfile = OUTPUT / 'device_test_merge.xlsx'

def text_from_cell(cell):
    """Return cell value as text, preserving existing strings and simple zero-padded numeric formats."""
    v = cell.value
    if v is None:
        return ''
    if isinstance(v, str):
        return v
    # Preserve leading zeros if the Excel number format is like 000 or 0000.
    fmt = str(cell.number_format or '')
    if isinstance(v, (int, float)) and float(v).is_integer():
        stripped = fmt.replace('""', '')
        if stripped and set(stripped) <= {'0'}:
            return str(int(v)).zfill(len(stripped))
        return str(int(v))
    return str(v)

def read_sheet(path, expected_sheet=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[expected_sheet] if expected_sheet and expected_sheet in wb.sheetnames else wb[wb.sheetnames[0]]
    headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    if 'Device ID' not in headers:
        raise ValueError(f'{path.name} 缺少 Device ID 欄位')
    id_idx = headers.index('Device ID') + 1
    rows = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        if all(c.value is None for c in row):
            continue
        rec = {}
        for i, h in enumerate(headers, start=1):
            if not h:
                continue
            rec[h] = text_from_cell(row[i-1]) if i == id_idx else row[i-1].value
        rows.append(rec)
    wb.close()
    return headers, rows

dev_headers, dev_rows = read_sheet(INPUT / 'devices.xlsx', 'devices')
test_headers, test_rows = read_sheet(INPUT / 'tests.xlsx', 'tests')

# devices 中每個 ID 唯一；若輸入違反，停止以避免任意覆蓋。
seen = set()
for rec in dev_rows:
    did = rec['Device ID']
    if did in seen:
        raise ValueError(f'devices.xlsx 中 Device ID 不唯一: {did!r}')
    seen.add(did)

dev_by_id = {rec['Device ID']: rec for rec in dev_rows}
test_ids = {rec['Device ID'] for rec in test_rows}

# matched: one row per successfully matched test, with test columns first, then device columns excluding duplicate Device ID.
matched_headers = list(test_headers) + [h for h in dev_headers if h != 'Device ID']
matched_rows = []
for test in test_rows:
    did = test['Device ID']
    if did in dev_by_id:
        dev = dev_by_id[did]
        matched_rows.append([test.get(h) for h in test_headers] + [dev.get(h) for h in dev_headers if h != 'Device ID'])

# unmatched: union of device and test columns, Source added first.
unmatched_headers = ['Source'] + list(dev_headers) + [h for h in test_headers if h not in dev_headers]
unmatched_rows = []
for dev in dev_rows:
    if dev['Device ID'] not in test_ids:
        unmatched_rows.append(['devices'] + [dev.get(h) for h in dev_headers] + [None for h in test_headers if h not in dev_headers])
for test in test_rows:
    if test['Device ID'] not in dev_by_id:
        unmatched_rows.append(['tests'] + [test.get(h) if h in test_headers else None for h in dev_headers] + [test.get(h) for h in test_headers if h not in dev_headers])

wb = Workbook()
# Remove default and create exact sheet names.
default = wb.active
wb.remove(default)
for sheet_name, headers, rows in [
    ('matched', matched_headers, matched_rows),
    ('unmatched', unmatched_headers, unmatched_rows),
]:
    ws = wb.create_sheet(sheet_name)
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for c in ws[1]:
        c.font = Font(bold=True)
    # Format any Device ID column as text.
    for idx, header in enumerate(headers, start=1):
        if header == 'Device ID':
            for cell in ws.iter_cols(min_col=idx, max_col=idx, min_row=2, max_row=ws.max_row).__next__():
                cell.number_format = '@'
                if cell.value is not None:
                    cell.value = str(cell.value)
    # Basic column widths.
    for idx, header in enumerate(headers, start=1):
        values = [header] + [ws.cell(row=r, column=idx).value for r in range(2, ws.max_row + 1)]
        width = max(len(str(v)) if v is not None else 0 for v in values) + 2
        ws.column_dimensions[get_column_letter(idx)].width = min(max(width, 10), 40)
wb.save(outfile)

# Verification: reopen saved output and confirm requirements from the saved file.
check_wb = load_workbook(outfile, data_only=True)
assert check_wb.sheetnames == ['matched', 'unmatched'], check_wb.sheetnames
matched = check_wb['matched']
unmatched = check_wb['unmatched']
matched_saved_headers = [c.value for c in matched[1]]
unmatched_saved_headers = [c.value for c in unmatched[1]]
assert matched_saved_headers == matched_headers, matched_saved_headers
assert unmatched_saved_headers == unmatched_headers, unmatched_saved_headers
assert matched.max_row - 1 == len(matched_rows), (matched.max_row - 1, len(matched_rows))
assert unmatched.max_row - 1 == len(unmatched_rows), (unmatched.max_row - 1, len(unmatched_rows))
# IDs are text after saving/reopening, and leading zeros in sample are preserved.
for ws in [matched, unmatched]:
    id_col = [c.value for c in ws[1]].index('Device ID') + 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, id_col).value
        assert v is None or isinstance(v, str), (ws.title, r, v, type(v))
assert matched['A2'].value == '001', matched['A2'].value
# expected sample row membership
saved_matched = [[matched.cell(r, c).value for c in range(1, matched.max_column + 1)] for r in range(2, matched.max_row + 1)]
saved_unmatched = [[unmatched.cell(r, c).value for c in range(1, unmatched.max_column + 1)] for r in range(2, unmatched.max_row + 1)]
assert ['001', 'RF-1', 'Pass', 'Alpha', 'Lab A'] in saved_matched
assert ['001', 'RF-2', 'Fail', 'Alpha', 'Lab A'] in saved_matched
assert ['002', 'RF-1', 'Pass', 'Beta', 'Lab B'] in saved_matched
assert ['devices', '003', 'Gamma', 'Lab C', None, None] in saved_unmatched
assert ['tests', '999', None, None, 'RF-9', 'Fail'] in saved_unmatched
check_wb.close()

print(f'created: {outfile}')
print('matched headers:', matched_headers)
print('matched data rows:', len(matched_rows))
print('unmatched headers:', unmatched_headers)
print('unmatched data rows:', len(unmatched_rows))
print('verification: reopened output workbook; checked sheet names, headers, row counts, Device ID text type/leading zeros, and expected matched/unmatched rows.')