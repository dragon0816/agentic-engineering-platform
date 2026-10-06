from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

workspace = Path.cwd()
input_path = workspace / 'input' / 'large_results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

required_headers = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']

stats = defaultdict(lambda: {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0})
overall = {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0}
row_count = 0

with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('source headers:', reader.fieldnames)
    missing = [h for h in required_headers if h not in (reader.fieldnames or [])]
    if missing:
        raise ValueError(f'Missing required columns: {missing}')
    for row in reader:
        row_count += 1
        model = row['Device Model']
        status = row['Status']
        try:
            duration = float(row['DurationSeconds'])
        except Exception as exc:
            raise ValueError(f'Invalid DurationSeconds on row {row_count + 1}: {row["DurationSeconds"]!r}') from exc
        for bucket in (stats[model], overall):
            bucket['Total'] += 1
            if status == 'Pass':
                bucket['PassCount'] += 1
            elif status == 'Fail':
                bucket['FailCount'] += 1
            bucket['DurationSum'] += duration

headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']

wb = Workbook()
ws = wb.active
ws.title = 'Summary'
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True)

# Stable model ordering: natural numeric suffix when possible, otherwise lexical.
def model_key(m):
    import re
    match = re.match(r'^(.*?)(\d+)$', str(m))
    if match:
        return (match.group(1), int(match.group(2)))
    return (str(m), -1)

for model in sorted(stats.keys(), key=model_key):
    s = stats[model]
    total = s['Total']
    pass_rate = s['PassCount'] / total if total else 0
    avg_duration = s['DurationSum'] / total if total else 0
    ws.append([model, total, s['PassCount'], s['FailCount'], pass_rate, avg_duration])

ws_overall = wb.create_sheet('Overall')
ws_overall.append(headers)
for cell in ws_overall[1]:
    cell.font = Font(bold=True)
total = overall['Total']
ws_overall.append([
    'Overall',
    total,
    overall['PassCount'],
    overall['FailCount'],
    overall['PassCount'] / total if total else 0,
    overall['DurationSum'] / total if total else 0,
])

for sheet in (ws, ws_overall):
    for row in sheet.iter_rows(min_row=2, min_col=5, max_col=5):
        for cell in row:
            cell.number_format = '0.00%'
    for row in sheet.iter_rows(min_row=2, min_col=6, max_col=6):
        for cell in row:
            cell.number_format = '0.00'
    for col_idx in range(1, sheet.max_column + 1):
        max_len = max(len(str(sheet.cell(row=r, column=col_idx).value)) for r in range(1, sheet.max_row + 1))
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 30)

wb.save(output_path)
print('wrote:', output_path)
print('input rows processed:', row_count)
print('model count:', len(stats))

# Reopen and validate output requirements from saved file.
wb2 = load_workbook(output_path, data_only=True)
print('sheets:', wb2.sheetnames)
if 'Summary' not in wb2.sheetnames or 'Overall' not in wb2.sheetnames:
    raise AssertionError('summary.xlsx must contain Summary and Overall sheets')

summary = wb2['Summary']
overall_sheet = wb2['Overall']
summary_headers = [c.value for c in summary[1]]
overall_headers = [c.value for c in overall_sheet[1]]
print('Summary headers:', summary_headers)
print('Overall headers:', overall_headers)
if summary_headers != headers:
    raise AssertionError(f'Unexpected Summary headers: {summary_headers}')
if overall_headers != headers:
    raise AssertionError(f'Unexpected Overall headers: {overall_headers}')

summary_rows = list(summary.iter_rows(min_row=2, values_only=True))
overall_rows = list(overall_sheet.iter_rows(min_row=2, values_only=True))
print('Summary data rows:', len(summary_rows))
print('Overall data rows:', len(overall_rows))
if len(summary_rows) != len(stats):
    raise AssertionError('Summary row count does not match number of Device Model values')
if len(overall_rows) != 1:
    raise AssertionError('Overall sheet should have exactly one data row')

# Validate values, types, and formulas already materialized as numeric values.
summary_total = 0
summary_pass = 0
summary_fail = 0
for model, total, pass_count, fail_count, pass_rate, avg_duration in summary_rows:
    if not isinstance(model, str):
        raise AssertionError(f'Device Model should be text, got {model!r}')
    if not all(isinstance(v, int) for v in [total, pass_count, fail_count]):
        raise AssertionError(f'Counts should be integers for {model}')
    if not isinstance(pass_rate, (int, float)) or not isinstance(avg_duration, (int, float)):
        raise AssertionError(f'Rate and average should be numeric for {model}')
    expected = stats[model]
    expected_rate = expected['PassCount'] / expected['Total']
    expected_avg = expected['DurationSum'] / expected['Total']
    if (total, pass_count, fail_count) != (expected['Total'], expected['PassCount'], expected['FailCount']):
        raise AssertionError(f'Count mismatch for {model}')
    if abs(pass_rate - expected_rate) > 1e-12:
        raise AssertionError(f'PassRate mismatch for {model}: {pass_rate} vs {expected_rate}')
    if abs(avg_duration - expected_avg) > 1e-12:
        raise AssertionError(f'AvgDurationSeconds mismatch for {model}: {avg_duration} vs {expected_avg}')
    summary_total += total
    summary_pass += pass_count
    summary_fail += fail_count

omodel, ototal, opass, ofail, orate, oavg = overall_rows[0]
if omodel != 'Overall':
    raise AssertionError('Overall row label should be Overall')
if (ototal, opass, ofail) != (overall['Total'], overall['PassCount'], overall['FailCount']):
    raise AssertionError('Overall counts mismatch')
if abs(orate - (overall['PassCount'] / overall['Total'])) > 1e-12:
    raise AssertionError('Overall PassRate mismatch')
if abs(oavg - (overall['DurationSum'] / overall['Total'])) > 1e-12:
    raise AssertionError('Overall AvgDurationSeconds mismatch')
if (summary_total, summary_pass, summary_fail) != (overall['Total'], overall['PassCount'], overall['FailCount']):
    raise AssertionError('Summary totals do not match Overall')

print('validation: OK')
print('Overall:', overall_rows[0])
print('First 5 Summary rows:', summary_rows[:5])
