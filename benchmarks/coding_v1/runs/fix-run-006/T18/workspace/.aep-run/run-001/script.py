from pathlib import Path
import csv, os, sys, subprocess, statistics, json, textwrap
from openpyxl import load_workbook

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)
print('WORKSPACE:', ws)
print('INPUT FILES:')
for p in sorted(input_dir.iterdir()):
    print(' -', p.name, 'size=', p.stat().st_size)

export_in = input_dir / 'export.py'
csv_in = input_dir / 'input.csv'
deps = input_dir / 'DEPENDENCIES.txt'
print('\nDEPENDENCIES.txt:')
print(deps.read_text(encoding='utf-8', errors='replace'))
print('\nOriginal export.py first 200 lines:')
orig = export_in.read_text(encoding='utf-8', errors='replace')
for i, line in enumerate(orig.splitlines()[:200], 1):
    print(f'{i:03d}: {line}')

print('\ninput.csv headers/sample:')
with csv_in.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('headers=', reader.fieldnames)
    for idx, row in zip(range(10), reader):
        print(idx+1, row)

# Compute expected result from the source CSV, treating blank Value as invalid and not included.
from collections import defaultdict
sums = defaultdict(float)
counts = defaultdict(int)
order = []
with csv_in.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    if reader.fieldnames != ['Model', 'Value']:
        print('NOTE: source headers are', reader.fieldnames)
    for row in reader:
        model = row.get('Model', '')
        if model not in sums:
            order.append(model)
        raw = row.get('Value', '')
        if raw is None or str(raw).strip() == '':
            continue
        val = float(str(raw).strip())
        sums[model] += val
        counts[model] += 1
expected = []
for model in sorted(order):
    avg = sums[model] / counts[model] if counts[model] else None
    expected.append((model, avg, counts[model]))
print('\nExpected rows (sorted by Model):')
for row in expected:
    print(row)

# Write repaired script to output/export.py. Keep CLI: --input and --output.
fixed = r'''#!/usr/bin/env python3
"""Summarize average Value by Model from a CSV into an Excel workbook."""

from __future__ import annotations

import argparse
import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Summarize average Value by Model.")
    parser.add_argument("--input", required=True, help="Path to input CSV")
    parser.add_argument("--output", required=True, help="Path to output XLSX")
    return parser.parse_args()


def read_summary(input_path: str | Path) -> list[tuple[str, float | None, int]]:
    """Return (Model, Average, ValidCount) rows.

    Blank Value cells are ignored. Models are emitted in sorted order for stable
    output. If a model has no valid numeric Value cells, Average is left blank
    in the workbook and represented here as None.
    """
    totals: dict[str, float] = defaultdict(float)
    counts: dict[str, int] = defaultdict(int)
    models: set[str] = set()

    with open(input_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        required = {"Model", "Value"}
        missing = required.difference(reader.fieldnames or [])
        if missing:
            raise ValueError(f"Input CSV missing required column(s): {', '.join(sorted(missing))}")

        for line_number, row in enumerate(reader, start=2):
            model = row.get("Model", "")
            models.add(model)
            raw_value = row.get("Value", "")
            if raw_value is None or str(raw_value).strip() == "":
                continue
            try:
                value = float(str(raw_value).strip())
            except ValueError as exc:
                raise ValueError(f"Invalid numeric Value at CSV line {line_number}: {raw_value!r}") from exc
            totals[model] += value
            counts[model] += 1

    rows: list[tuple[str, float | None, int]] = []
    for model in sorted(models):
        valid_count = counts[model]
        average = totals[model] / valid_count if valid_count else None
        rows.append((model, average, valid_count))
    return rows


def write_workbook(rows: list[tuple[str, float | None, int]], output_path: str | Path) -> None:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Model", "Average", "ValidCount"])
    for model, average, valid_count in rows:
        ws.append([model, average, valid_count])
    wb.save(output_path)


def main() -> None:
    args = parse_args()
    rows = read_summary(args.input)
    write_workbook(rows, args.output)


if __name__ == "__main__":
    main()
'''
(export_out := output_dir / 'export.py').write_text(fixed, encoding='utf-8', newline='\n')
print('\nWrote repaired script:', export_out)

# Execute exactly as requested relative to workspace (using current Python executable for reliability).
summary_out = output_dir / 'summary.xlsx'
if summary_out.exists():
    summary_out.unlink()
cmd_display = r'python output\export.py --input input\input.csv --output output\summary.xlsx'
cmd = [sys.executable, str(output_dir / 'export.py'), '--input', str(input_dir / 'input.csv'), '--output', str(summary_out)]
print('\nRunning requested command equivalent:')
print(cmd_display)
completed = subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('returncode=', completed.returncode)
if completed.stdout:
    print('stdout:\n' + completed.stdout)
if completed.stderr:
    print('stderr:\n' + completed.stderr)
if completed.returncode != 0:
    raise SystemExit('repaired export.py execution failed')

# Verify saved outputs by reopening them.
print('\nVerification:')
print('output/export.py exists:', export_out.exists(), 'size=', export_out.stat().st_size)
print('output/summary.xlsx exists:', summary_out.exists(), 'size=', summary_out.stat().st_size)
wb = load_workbook(summary_out, data_only=True)
wsheet = wb.active
headers = [cell.value for cell in next(wsheet.iter_rows(min_row=1, max_row=1))]
print('workbook sheet:', wsheet.title)
print('headers:', headers)
actual = []
for row in wsheet.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    actual.append(tuple(row[:3]))
print('actual rows:')
for row in actual:
    print(row)

# strict checks: required names, counts, numeric averages/counts, equality against recomputed expected.
assert headers == ['Model', 'Average', 'ValidCount'], f'bad headers: {headers!r}'
assert len(actual) == len(expected), f'row count mismatch {len(actual)} != {len(expected)}'
for got, exp in zip(actual, expected):
    gm, ga, gc = got
    em, ea, ec = exp
    assert gm == em, f'model mismatch {gm!r} != {em!r}'
    assert gc == ec and isinstance(gc, int), f'ValidCount mismatch/type for {gm}: {gc!r} != {ec!r}'
    if ea is None:
        assert ga is None, f'Average expected blank for {gm}, got {ga!r}'
    else:
        assert isinstance(ga, (int, float)), f'Average not numeric for {gm}: {ga!r}'
        assert abs(float(ga) - float(ea)) < 1e-9, f'Average mismatch for {gm}: {ga!r} != {ea!r}'
print('All checks passed: CLI ran, output files reopen, headers match, blank Values excluded, averages/counts match CSV.')
