import csv
import json
import os
import stat
from pathlib import Path
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = json.load(open(params_path, encoding='utf-8')) if params_path else {}
threshold = params.get('urgent_threshold')
if threshold is None:
    raise SystemExit('Missing required parameter: urgent_threshold')
threshold = float(threshold)

input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

print(f'Workspace: {workspace}')
print('Input files:')
for p in sorted(input_dir.rglob('*')):
    if p.is_file():
        print(f'- {p.relative_to(workspace)} {p.stat().st_size}')

sop_path = input_dir / 'supplied_sop.md'
manifest_path = input_dir / 'manifest.json'
print('SOP first 500 chars:')
print(sop_path.read_text(encoding='utf-8')[:500])

with manifest_path.open('r', encoding='utf-8') as f:
    manifest = json.load(f)
csv_rel = manifest.get('csv') or manifest.get('input_csv') or manifest.get('results_csv') or 'data/test_results.csv'
csv_path = input_dir / csv_rel
if not csv_path.exists():
    alt = input_dir / 'data' / 'test_results.csv'
    if alt.exists():
        csv_path = alt
    else:
        raise FileNotFoundError(f'CSV not found from manifest path {csv_rel!r}; also checked {alt}')

rows = []
with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    print(f'CSV: {csv_path.relative_to(workspace)}')
    print(f'CSV headers: {headers}')
    required = ['Device ID', 'Status', 'Value']
    missing = [h for h in required if h not in headers]
    if missing:
        raise SystemExit(f'Missing required CSV columns: {missing}')
    for row in reader:
        rows.append(row)
print('Sample rows:')
for row in rows[:5]:
    print(row)

def parse_value(v):
    if v is None or str(v).strip() == '':
        return None
    return float(str(v).strip())

fail_rows = []
for row in rows:
    if str(row.get('Status', '')).strip() == 'Fail':
        value = parse_value(row.get('Value'))
        urgent = bool(value is not None and value > threshold)
        fail_rows.append({'Device ID': row.get('Device ID', ''), 'Status': row.get('Status', ''), 'Value': value if value is not None else '', 'Urgent': urgent})

summary = {
    'Total': len(rows),
    'Pass': sum(1 for r in rows if str(r.get('Status', '')).strip() == 'Pass'),
    'Fail': len(fail_rows),
    'Urgent': sum(1 for r in fail_rows if r['Urgent']),
}

failures_path = output_dir / 'failures.xlsx'
summary_path = output_dir / 'summary.md'
for target in [failures_path, summary_path]:
    if target.exists():
        try:
            target.chmod(stat.S_IWRITE | stat.S_IREAD)
        except Exception as exc:
            print(f'Warning: could not chmod {target}: {exc}')
        try:
            target.unlink()
        except PermissionError as exc:
            raise SystemExit(f'Cannot replace {target}; it may be open in another program. Close it and rerun. Details: {exc}') from exc

wb = Workbook()
ws = wb.active
ws.title = 'Failures'
columns = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(columns)
for r in fail_rows:
    ws.append([r['Device ID'], r['Status'], r['Value'], r['Urgent']])
tmp_xlsx = output_dir / 'failures.tmp.xlsx'
if tmp_xlsx.exists():
    tmp_xlsx.unlink()
wb.save(tmp_xlsx)
os.replace(tmp_xlsx, failures_path)

summary_lines = ['# Test Results Summary', '', f'- Total: {summary["Total"]}', f'- Pass: {summary["Pass"]}', f'- Fail: {summary["Fail"]}', f'- Urgent: {summary["Urgent"]}', '', f'Urgent rule: Status is Fail and Value > {threshold:g}.']
summary_path.write_text('\n'.join(summary_lines) + '\n', encoding='utf-8')

check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['Failures']
check_headers = [cell.value for cell in next(check_ws.iter_rows(min_row=1, max_row=1))]
if check_headers != columns:
    raise SystemExit(f'failures.xlsx headers mismatch: {check_headers}')
check_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
if len(check_rows) != len(fail_rows):
    raise SystemExit(f'failures.xlsx row count mismatch: got {len(check_rows)}, expected {len(fail_rows)}')
for idx, row in enumerate(check_rows, start=2):
    device_id, status, value, urgent = row
    if status != 'Fail':
        raise SystemExit(f'Row {idx} status is not Fail: {status!r}')
    if not isinstance(device_id, str):
        raise SystemExit(f'Row {idx} Device ID is not text: {device_id!r}')
    if value is not None and not isinstance(value, (int, float)):
        raise SystemExit(f'Row {idx} Value is not numeric/blank: {value!r}')
    expected_urgent = bool(value is not None and float(value) > threshold)
    if urgent is not expected_urgent:
        raise SystemExit(f'Row {idx} Urgent mismatch: got {urgent!r}, expected {expected_urgent!r}')
summary_text = summary_path.read_text(encoding='utf-8')
for key, val in summary.items():
    token = f'- {key}: {val}'
    if token not in summary_text:
        raise SystemExit(f'summary.md missing {token!r}')
print('Verified outputs from disk:')
print(f'- {failures_path.relative_to(workspace)}: {len(check_rows)} Fail rows, headers {check_headers}, Urgent booleans verified with Value > {threshold:g}')
print(f'- {summary_path.relative_to(workspace)}: Total={summary["Total"]}, Pass={summary["Pass"]}, Fail={summary["Fail"]}, Urgent={summary["Urgent"]}')
