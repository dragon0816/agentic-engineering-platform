import csv, json, os
from pathlib import Path
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = json.load(open(params_path, encoding='utf-8')) if params_path and Path(params_path).exists() else {}
threshold = int(params['urgent_threshold'])
print('Workspace:', WORKSPACE)
print('Input files:')
for p in sorted(INPUT_DIR.rglob('*')):
    if p.is_file(): print('-', p.relative_to(WORKSPACE), p.stat().st_size)
sop_path = INPUT_DIR / 'supplied_sop.md'
print('SOP first 500 chars:')
print(sop_path.read_text(encoding='utf-8')[:500])
manifest_path = INPUT_DIR / 'manifest.json'
csv_rel = json.loads(manifest_path.read_text(encoding='utf-8')).get('results.csv', 'input/data/test_results.csv') if manifest_path.exists() else 'input/data/test_results.csv'
csv_path = WORKSPACE / csv_rel
rows = []
with csv_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('CSV:', csv_path.relative_to(WORKSPACE))
    print('CSV headers:', reader.fieldnames)
    required = ['Device ID', 'Status', 'Value']
    missing = [c for c in required if c not in (reader.fieldnames or [])]
    if missing: raise ValueError(f'CSV is missing required columns: {missing}')
    rows = list(reader)
print('Sample rows:')
for row in rows[:5]: print(row)

def parse_number(v):
    num = float(str(v).strip())
    return int(num) if num.is_integer() else num

total_count = len(rows)
pass_count = sum(1 for r in rows if str(r.get('Status','')).strip() == 'Pass')
fail_rows = [r for r in rows if str(r.get('Status','')).strip() == 'Fail']
fail_count = len(fail_rows)
output_rows=[]; urgent_count=0
for r in fail_rows:
    value_num = parse_number(r['Value'])
    urgent = value_num > threshold
    urgent_count += int(urgent)
    output_rows.append([str(r['Device ID']), str(r['Status']), value_num, bool(urgent)])
failures_path = OUTPUT_DIR / 'failures.xlsx'
summary_path = OUTPUT_DIR / 'summary.md'
wb = Workbook(); ws = wb.active; ws.title = 'Failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for row in output_rows: ws.append(row)
wb.save(failures_path)
summary_path.write_text('\n'.join(['# Test Failure Summary','',f'- Total: {total_count}',f'- Pass: {pass_count}',f'- Fail: {fail_count}',f'- Urgent: {urgent_count}',''])+'\n', encoding='utf-8')
check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['Failures']
actual_headers = [c.value for c in next(check_ws.iter_rows(min_row=1, max_row=1))]
assert actual_headers == headers, actual_headers
actual_data = list(check_ws.iter_rows(min_row=2, values_only=True))
assert len(actual_data) == fail_count, (len(actual_data), fail_count)
for row in actual_data:
    device_id, status, value, urgent = row
    assert isinstance(device_id, str), row
    assert status == 'Fail', row
    assert isinstance(value, (int, float)), row
    assert isinstance(urgent, bool), row
    assert urgent == (value > threshold), row
saved_summary = summary_path.read_text(encoding='utf-8')
for line in [f'- Total: {total_count}', f'- Pass: {pass_count}', f'- Fail: {fail_count}', f'- Urgent: {urgent_count}']:
    assert line in saved_summary, line
print('Wrote:', failures_path.relative_to(WORKSPACE))
print('Wrote:', summary_path.relative_to(WORKSPACE))
print('Checked failures.xlsx columns, row count, Status=Fail, Device ID text, Value numeric, Urgent boolean/rule.')
print('Checked summary.md counts:', {'Total': total_count, 'Pass': pass_count, 'Fail': fail_count, 'Urgent': urgent_count})
