from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

ws = Path.cwd()
input_path = ws / 'input' / 'specs.md'
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)
out_path = out_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split product sections by level-2 headings.
sections = []
current_title = None
current_lines = []
for line in text.splitlines():
    m = re.match(r'^##\s+(.+?)\s*$', line)
    if m:
        if current_title is not None:
            sections.append((current_title, '\n'.join(current_lines)))
        current_title = m.group(1).strip()
        current_lines = []
    elif current_title is not None:
        current_lines.append(line)
if current_title is not None:
    sections.append((current_title, '\n'.join(current_lines)))

unit_to_mhz = {
    'hz': 1/1_000_000,
    'khz': 1/1_000,
    'mhz': 1,
    'ghz': 1000,
}

def fmt_num(x):
    # Keep numeric cells numeric in Excel; return int when exact, otherwise float.
    if x is None:
        return None
    if abs(x - round(x)) < 1e-9:
        return int(round(x))
    return float(x)

def parse_frequency(value):
    if not value:
        return None, None
    s = value.strip().replace('—', '-').replace('–', '-').replace('−', '-')
    # Examples: 2.4-2.5 GHz, 900-930 MHz. If one unit is written after the range,
    # apply only that explicitly written shared unit.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(?:hz|khz|mhz|ghz)?\s*-\s*([0-9]+(?:\.[0-9]+)?)\s*(hz|khz|mhz|ghz)\b', s, re.I)
    if not m:
        return None, None
    unit = m.group(3).lower()
    factor = unit_to_mhz[unit]
    return fmt_num(float(m.group(1)) * factor), fmt_num(float(m.group(2)) * factor)

def parse_bandwidth(value):
    if not value:
        return None
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(hz|khz|mhz|ghz)\b', value, re.I)
    if not m:
        return None
    return fmt_num(float(m.group(1)) * unit_to_mhz[m.group(2).lower()])

def parse_interface(value):
    if not value:
        return ''
    parts = [p.strip() for p in re.split(r'\s*(?:/|,|;|、|，)\s*', value.strip()) if p.strip()]
    return ';'.join(parts)

rows = []
for title, body in sections:
    fields = {}
    for line in body.splitlines():
        m = re.match(r'^\s*-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            fields[m.group(1).strip().lower()] = m.group(2).strip()
    fmin, fmax = parse_frequency(fields.get('frequency', ''))
    rows.append({
        'Model': fields.get('model', ''),
        'FrequencyMinMHz': fmin,
        'FrequencyMaxMHz': fmax,
        'Bandwidth': parse_bandwidth(fields.get('bandwidth', '')),
        'Interface': parse_interface(fields.get('interface', '')),
    })

headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'Bandwidth', 'Interface']
wb = Workbook()
wsheet = wb.active
wsheet.title = 'specs'
wsheet.append(headers)
for row in rows:
    wsheet.append([row[h] if row[h] is not None else '' for h in headers])
wb.save(out_path)

# Reopen and verify stated requirements from saved output.
wb2 = load_workbook(out_path, data_only=True)
sh = wb2['specs']
actual_headers = [c.value for c in sh[1]]
data = list(sh.iter_rows(min_row=2, values_only=True))
print('wrote', out_path)
print('headers', actual_headers)
print('row_count', len(data))
for r in data:
    print(r)

assert actual_headers == headers, f'headers mismatch: {actual_headers}'
assert len(data) == 3, f'expected 3 product rows, got {len(data)}'
expected = [
    ('RX-100', 2400, 2500, 80, 'USB;LAN'),
    ('RX-200', 900, 930, 20, 'LAN'),
    ('RX-300', 5000, 6000, None, 'USB'),
]
assert data == expected, f'data mismatch: {data!r}'
# Confirm type expectations: text IDs/interfaces, numeric frequencies/bandwidth where present, missing bandwidth blank/None.
for model, fmin, fmax, bw, interface in data:
    assert isinstance(model, str)
    assert isinstance(fmin, (int, float)) and isinstance(fmax, (int, float))
    assert bw is None or isinstance(bw, (int, float))
    assert isinstance(interface, str)
print('validation passed: output file, sheet name, headers, 3 rows, MHz conversions, semicolon interfaces, and blank missing bandwidth confirmed')
