from pathlib import Path
import csv
import subprocess
import sys
from openpyxl import load_workbook

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)

# Record input metadata so we can confirm we did not modify inputs.
before = {p.name: (p.stat().st_size, p.stat().st_mtime_ns) for p in input_dir.iterdir() if p.is_file()}
print('INPUT_METADATA_BEFORE', before)

# Read the real source CSV and compute expected results from its existing column names.
csv_path = input_dir / 'input.csv'
with csv_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('CSV_FIELDNAMES', reader.fieldnames)
    source_rows = list(reader)
print('CSV_ROW_COUNT', len(source_rows))
print('CSV_ROWS', source_rows)

expected_totals = {}
expected_counts = {}
for row in source_rows:
    model = row['Model']
    raw_value = row['Value']
    if raw_value is None or raw_value.strip() == '':
        continue
    value = float(raw_value)
    expected_totals[model] = expected_totals.get(model, 0.0) + value
    expected_counts[model] = expected_counts.get(model, 0) + 1
expected = {
    model: (expected_totals[model] / expected_counts[model], expected_counts[model])
    for model in sorted(expected_totals)
}
print('EXPECTED', expected)

fixed_code = r'''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals = {}
    counts = {}

    with open(args.input, newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            model = row["Model"]
            value_text = row["Value"]
            if value_text is None or value_text.strip() == "":
                continue
            value = float(value_text)
            totals[model] = totals.get(model, 0.0) + value
            counts[model] = counts.get(model, 0) + 1

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(args.output)


if __name__ == "__main__":
    main()
'''

export_out = output_dir / 'export.py'
summary_out = output_dir / 'summary.xlsx'
export_out.write_text(fixed_code, encoding='utf-8')
print('WROTE', export_out, 'size=', export_out.stat().st_size)

cmd = [sys.executable, str(export_out), '--input', str(input_dir / 'input.csv'), '--output', str(summary_out)]
print('RUNNING', ' '.join(cmd))
completed = subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('CLI_EXIT_CODE', completed.returncode)
print('CLI_STDOUT', completed.stdout)
print('CLI_STDERR', completed.stderr)
if completed.returncode != 0:
    raise SystemExit(completed.returncode)

# Reopen and validate output workbook.
if not summary_out.exists() or summary_out.stat().st_size == 0:
    raise AssertionError('summary.xlsx was not created or is empty')
wb = load_workbook(summary_out, data_only=True)
if 'summary' not in wb.sheetnames:
    raise AssertionError(f'missing summary sheet; sheets={wb.sheetnames}')
wsheet = wb['summary']
rows = list(wsheet.iter_rows(values_only=True))
print('SUMMARY_ROWS', rows)
header = rows[0]
if header != ('Model', 'Average', 'ValidCount'):
    raise AssertionError(f'bad header: {header!r}')
actual = {}
for row in rows[1:]:
    model, avg, count = row
    if not isinstance(model, str):
        raise AssertionError(f'Model is not text: {row!r}')
    if not isinstance(avg, (int, float)):
        raise AssertionError(f'Average is not numeric: {row!r}')
    if not isinstance(count, int):
        raise AssertionError(f'ValidCount is not int: {row!r}')
    actual[model] = (float(avg), count)
if set(actual) != set(expected):
    raise AssertionError(f'model mismatch: actual={actual}, expected={expected}')
for model, (exp_avg, exp_count) in expected.items():
    act_avg, act_count = actual[model]
    if abs(act_avg - exp_avg) > 1e-9 or act_count != exp_count:
        raise AssertionError(f'value mismatch for {model}: actual={(act_avg, act_count)}, expected={(exp_avg, exp_count)}')
if len(rows) != len(expected) + 1:
    raise AssertionError(f'row count mismatch: rows={len(rows)}, expected={len(expected)+1}')

# Confirm input files were not modified.
after = {p.name: (p.stat().st_size, p.stat().st_mtime_ns) for p in input_dir.iterdir() if p.is_file()}
print('INPUT_METADATA_AFTER', after)
if before != after:
    raise AssertionError('input directory files changed')

print('VERIFICATION_OK', {
    'export_py': str(export_out),
    'summary_xlsx': str(summary_out),
    'sheet': 'summary',
    'header': header,
    'data_rows': len(rows) - 1,
    'actual': actual,
})
