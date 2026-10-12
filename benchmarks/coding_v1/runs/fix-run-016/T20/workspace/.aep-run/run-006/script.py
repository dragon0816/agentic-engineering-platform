from pathlib import Path
from openpyxl import load_workbook

ws = Path.cwd()
print('Workspace:', ws)
print('Input listing:')
for p in sorted((ws/'input').rglob('*')):
    print('-', p.relative_to(ws), 'dir' if p.is_dir() else f'file size={p.stat().st_size}')
print('Output listing:')
for p in sorted((ws/'output').rglob('*')):
    print('-', p.relative_to(ws), 'dir' if p.is_dir() else f'file size={p.stat().st_size}')

xlsx = ws/'output'/'failures.xlsx'
summary = ws/'output'/'summary.md'
assert xlsx.exists() and xlsx.stat().st_size > 0, 'failures.xlsx missing or empty'
assert summary.exists() and summary.stat().st_size > 0, 'summary.md missing or empty'

wb = load_workbook(xlsx, data_only=True)
sh = wb.active
rows = list(sh.iter_rows(values_only=True))
print('XLSX rows:')
for r in rows:
    print(r)
headers = list(rows[0])
expected_headers = ['Device ID', 'Status', 'Value', 'Urgent']
assert headers == expected_headers, f'headers mismatch: {headers}'
data = rows[1:]
assert len(data) == 3, f'expected 3 Fail rows, got {len(data)}'
for device_id, status, value, urgent in data:
    assert isinstance(device_id, str), f'Device ID not text: {device_id!r}'
    assert status == 'Fail', f'non-Fail row found: {status!r}'
    assert isinstance(value, (int, float)), f'Value not numeric: {value!r}'
    assert isinstance(urgent, bool), f'Urgent not boolean: {urgent!r}'
    assert urgent == (value > 20), f'Urgent rule mismatch for {device_id}: {value}, {urgent}'

text = summary.read_text(encoding='utf-8')
print('summary.md content:')
print(text)
for expected in ['Total: 5', 'Pass: 2', 'Fail: 3', 'Urgent: 1']:
    assert expected in text, f'missing summary count: {expected}'
print('VERIFICATION PASSED: outputs exist, failures.xlsx content/rules are correct, summary.md counts are correct.')