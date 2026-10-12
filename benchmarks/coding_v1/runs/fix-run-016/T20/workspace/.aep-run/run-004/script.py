from pathlib import Path
import csv
import json
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
THRESHOLD = Decimal('20')

print(f'Workspace: {WORKSPACE}')
print('Input files:')
for p in sorted(INPUT_DIR.rglob('*')):
    kind = 'dir' if p.is_dir() else f'file size={p.stat().st_size}'
    print(f'- {p.relative_to(WORKSPACE)} {kind}')

manifest_path = INPUT_DIR / 'manifest.json'
if manifest_path.exists():
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    mapped = manifest.get('results.csv')
    if not mapped:
        raise FileNotFoundError('manifest.json does not contain a results.csv entry')
    csv_path = WORKSPACE / mapped
else:
    csv_path = INPUT_DIR / 'results.csv'

if not csv_path.exists():
    raise FileNotFoundError(f'Results CSV not found: {csv_path}')
if not csv_path.resolve().is_relative_to(INPUT_DIR.resolve()):
    raise ValueError(f'Results CSV must be under input/: {csv_path}')

required_columns = ['Device ID', 'Status', 'Value']
rows = []
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print(f'CSV path: {csv_path.relative_to(WORKSPACE)}')
    print(f'CSV headers: {reader.fieldnames}')
    missing = [c for c in required_columns if c not in (reader.fieldnames or [])]
    if missing:
        raise ValueError(f'Missing required CSV columns: {missing}')
    for row in reader:
        rows.append(row)

fail_rows = []
pass_count = 0
fail_count = 0
urgent_count = 0
for row in rows:
    status = (row.get('Status') or '').strip()
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        raw_value = row.get('Value')
        try:
            numeric_value = Decimal(str(raw_value).strip())
        except (InvalidOperation, AttributeError):
            raise ValueError(f"Value is not numeric for Device ID {row.get('Device ID')}: {raw_value!r}")
        urgent = numeric_value > THRESHOLD
        if urgent:
            urgent_count += 1
        value_for_excel = int(numeric_value) if numeric_value == numeric_value.to_integral_value() else float(numeric_value)
        fail_rows.append({'Device ID': row.get('Device ID'), 'Status': row.get('Status'), 'Value': value_for_excel, 'Urgent': bool(urgent)})

total_count = len(rows)

xlsx_path = OUTPUT_DIR / 'failures.xlsx'
wb = Workbook()
ws = wb.active
ws.title = 'Failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for row in fail_rows:
    ws.append([row['Device ID'], row['Status'], row['Value'], row['Urgent']])
wb.save(xlsx_path)

summary_path = OUTPUT_DIR / 'summary.md'
summary = '\n'.join(['# Summary', '', f'- Total: {total_count}', f'- Pass: {pass_count}', f'- Fail: {fail_count}', f'- Urgent: {urgent_count}', ''])
summary_path.write_text(summary, encoding='utf-8')

check_wb = load_workbook(xlsx_path, data_only=True)
check_ws = check_wb['Failures']
actual_headers = [cell.value for cell in next(check_ws.iter_rows(min_row=1, max_row=1))]
if actual_headers != headers:
    raise AssertionError(f'failures.xlsx headers mismatch: {actual_headers} != {headers}')
data_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
if len(data_rows) != fail_count:
    raise AssertionError(f'failures.xlsx row count mismatch: {len(data_rows)} != {fail_count}')
for row in data_rows:
    device_id, status, value, urgent = row
    if not isinstance(device_id, str):
        raise AssertionError(f'Device ID is not text for row: {row!r}')
    if status != 'Fail':
        raise AssertionError(f'Non-Fail row found in failures.xlsx: {row!r}')
    if not isinstance(value, (int, float)):
        raise AssertionError(f'Value is not numeric for row: {row!r}')
    if not isinstance(urgent, bool):
        raise AssertionError(f'Urgent is not boolean for row: {row!r}')
    if urgent != (Decimal(str(value)) > THRESHOLD):
        raise AssertionError(f'Urgent value mismatch for row: {row!r}')

summary_text = summary_path.read_text(encoding='utf-8')
expected_fragments = [f'- Total: {total_count}', f'- Pass: {pass_count}', f'- Fail: {fail_count}', f'- Urgent: {urgent_count}']
for fragment in expected_fragments:
    if fragment not in summary_text:
        raise AssertionError(f'summary.md missing {fragment!r}')

print(f'Wrote {xlsx_path.relative_to(WORKSPACE)} with {len(data_rows)} Fail rows')
print(f'Wrote {summary_path.relative_to(WORKSPACE)}')
print('Counts:', {'Total': total_count, 'Pass': pass_count, 'Fail': fail_count, 'Urgent': urgent_count})
print('Checked output files: failures.xlsx headers/row count/types/Urgent rule; summary.md counts')
