from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
import re
from openpyxl import Workbook, load_workbook

BASE_URL = 'http://127.0.0.1:64493/t11/list'
WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
OUTFILE = OUTPUT / 'details.xlsx'

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._in_a = False
        self._href = None
        self._text = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attr = dict(attrs)
            self._in_a = True
            self._href = attr.get('href')
            self._text = []
    def handle_data(self, data):
        if self._in_a:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._in_a:
            self.links.append((self._href, ''.join(self._text).strip()))
            self._in_a = False
            self._href = None
            self._text = []

class TableParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self._in_cell = False
        self._cell = []
        self._row = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'tr':
            self._row = []
        if tag.lower() in ('td','th'):
            self._in_cell = True
            self._cell = []
    def handle_data(self, data):
        if self._in_cell:
            self._cell.append(data)
    def handle_endtag(self, tag):
        if tag.lower() in ('td','th') and self._in_cell:
            txt = ' '.join(''.join(self._cell).split())
            self._row.append(txt)
            self._in_cell = False
            self._cell = []
        elif tag.lower() == 'tr':
            if self._row:
                self.rows.append(self._row)
                self._row = []

def fetch(url):
    with urlopen(url, timeout=10) as r:
        return r.read().decode('utf-8', errors='replace')

def html_text(html):
    text = re.sub(r'<script.*?</script>', ' ', html, flags=re.I|re.S)
    text = re.sub(r'<style.*?</style>', ' ', text, flags=re.I|re.S)
    text = re.sub(r'<[^>]+>', ' ', text)
    text = re.sub(r'\s+', ' ', text)
    return text.strip()

def normalize_label(label):
    return re.sub(r'[^a-z0-9]', '', label.lower())

def parse_price_currency(value):
    if value is None:
        return ('', '')
    s = str(value).strip()
    if not s:
        return ('', '')
    # Currency may be before or after the number, e.g. USD 999, $999, 999 EUR
    m = re.search(r'([A-Z]{3}|[$€£¥])?\s*([-+]?\d+(?:[,.]\d+)*)\s*([A-Z]{3}|[$€£¥])?', s)
    if not m:
        return (s, '')
    cur = (m.group(1) or m.group(3) or '').strip()
    num = m.group(2).replace(',', '')
    try:
        price = float(num)
        if price.is_integer():
            price = int(price)
    except Exception:
        price = num
    symbol_map = {'$':'USD','€':'EUR','£':'GBP','¥':'JPY'}
    cur = symbol_map.get(cur, cur)
    return (price, cur)

def memory_to_gb(value):
    if value is None:
        return ''
    s = str(value).strip()
    if not s:
        return ''
    m = re.search(r'([-+]?\d+(?:\.\d+)?)\s*([KMGT]?I?B|[KMGT]B|[KMGT])?', s, flags=re.I)
    if not m:
        return s
    num = float(m.group(1))
    unit = (m.group(2) or 'GB').upper().replace('IB','B')
    if unit in ('G','GB'):
        gb = num
    elif unit in ('M','MB'):
        gb = num / 1024
    elif unit in ('K','KB'):
        gb = num / (1024*1024)
    elif unit in ('T','TB'):
        gb = num * 1024
    else:
        gb = num
    if abs(gb - round(gb)) < 1e-9:
        return int(round(gb))
    return round(gb, 4)

def parse_detail(html, url):
    tp = TableParser()
    tp.feed(html)
    fields = {}
    for row in tp.rows:
        if len(row) >= 2:
            fields[normalize_label(row[0])] = row[1]
    # Fallbacks for common non-table patterns like <dt>Name</dt><dd>...</dd> or text labels.
    if not fields:
        pairs = re.findall(r'<(?:dt|th|span|b)[^>]*>\s*([^<:]+):?\s*</(?:dt|th|span|b)>\s*<(?:dd|td|span)[^>]*>\s*([^<]+)\s*</(?:dd|td|span)>', html, flags=re.I|re.S)
        for k, v in pairs:
            fields[normalize_label(k)] = ' '.join(v.split())
    text = html_text(html)
    def find_field(*names):
        for n in names:
            key = normalize_label(n)
            if key in fields:
                return fields[key]
        # Last-resort text search: Label: value until next known label-like separator.
        for n in names:
            pat = re.compile(r'\b' + re.escape(n) + r'\b\s*[:：]\s*([^|;]+?)(?=\s+(?:Name|Price|Currency|CPU|Memory|MemoryGB)\s*[:：]|$)', re.I)
            mm = pat.search(text)
            if mm:
                return mm.group(1).strip()
        return ''
    name = find_field('Name', 'Product Name')
    price_raw = find_field('Price')
    currency = find_field('Currency')
    price, price_cur = parse_price_currency(price_raw)
    if not currency:
        currency = price_cur
    cpu = find_field('CPU', 'Processor')
    mem_raw = find_field('MemoryGB', 'Memory', 'RAM')
    mem_gb = memory_to_gb(mem_raw)
    return {'Name': name, 'Price': price, 'Currency': currency, 'CPU': cpu, 'MemoryGB': mem_gb, 'URL': url}

list_html = fetch(BASE_URL)
print('Fetched product list:', BASE_URL)
print('List markup sample:', list_html[:500])
link_parser = LinkParser()
link_parser.feed(list_html)
links = []
seen = set()
for href, label in link_parser.links:
    if href:
        full = urljoin(BASE_URL, href)
        if full not in seen:
            links.append((full, label))
            seen.add(full)
print('Detail links found:', len(links), links)

records = []
for idx, (url, label) in enumerate(links, start=1):
    html = fetch(url)
    print(f'Detail {idx} URL: {url}')
    print('Detail markup sample:', html[:800])
    rec = parse_detail(html, url)
    print('Parsed record:', rec)
    records.append(rec)

headers = ['Name', 'Price', 'Currency', 'CPU', 'MemoryGB', 'URL']
wb = Workbook()
ws = wb.active
ws.title = 'details'
ws.append(headers)
for rec in records:
    ws.append([rec.get(h, '') for h in headers])
wb.save(OUTFILE)
print('Wrote:', OUTFILE)

# Verification by reopening saved output.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['details']
out_headers = [c.value for c in ws2[1]]
rows = list(ws2.iter_rows(min_row=2, values_only=True))
print('CHECK file exists:', OUTFILE.exists(), 'size:', OUTFILE.stat().st_size if OUTFILE.exists() else 0)
print('CHECK sheet:', ws2.title)
print('CHECK headers:', out_headers)
print('CHECK row_count:', len(rows), 'expected:', len(links))
for i, row in enumerate(rows, start=2):
    rowdict = dict(zip(out_headers, row))
    price = rowdict.get('Price')
    mem = rowdict.get('MemoryGB')
    print(f'CHECK row {i}: Price type={type(price).__name__}, MemoryGB={mem} type={type(mem).__name__}, URL={rowdict.get("URL")}')
assert out_headers == headers, f'headers mismatch: {out_headers}'
assert len(rows) == len(links), f'row count mismatch {len(rows)} != {len(links)}'
for row in rows:
    assert len(row) == len(headers)
    assert row[5] and str(row[5]).startswith('http://127.0.0.1:64493/t11/item/'), row[5]
print('All checks passed.')