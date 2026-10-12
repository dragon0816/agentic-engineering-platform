from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
import re
from openpyxl import Workbook, load_workbook

BASE_URL = 'http://127.0.0.1:64493/t11/list'
WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
OUTFILE = OUTPUT / 'details.xlsx'

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._in_a = False
        self._href = None
        self._text = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            self._in_a = True
            self._href = dict(attrs).get('href')
            self._text = []
    def handle_data(self, data):
        if self._in_a:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._in_a:
            self.links.append((self._href, ''.join(self._text).strip()))
            self._in_a = False
            self._href = None
            self._text = []

class DetailParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows = []
        self.name = ''
        self.title = ''
        self._in_cell = False
        self._cell = []
        self._row = []
        self._capture_name = False
        self._name_text = []
        self._capture_title = False
        self._title_text = []
    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        attr = dict(attrs)
        classes = set((attr.get('class') or '').split())
        if tag == 'tr':
            self._row = []
        if tag in ('td','th'):
            self._in_cell = True
            self._cell = []
        if tag == 'h1' and 'name' in classes:
            self._capture_name = True
            self._name_text = []
        if tag == 'title':
            self._capture_title = True
            self._title_text = []
    def handle_data(self, data):
        if self._in_cell:
            self._cell.append(data)
        if self._capture_name:
            self._name_text.append(data)
        if self._capture_title:
            self._title_text.append(data)
    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ('td','th') and self._in_cell:
            txt = ' '.join(''.join(self._cell).split())
            self._row.append(txt)
            self._in_cell = False
            self._cell = []
        elif tag == 'tr':
            if self._row:
                self.rows.append(self._row)
                self._row = []
        elif tag == 'h1' and self._capture_name:
            self.name = ' '.join(''.join(self._name_text).split())
            self._capture_name = False
        elif tag == 'title' and self._capture_title:
            self.title = ' '.join(''.join(self._title_text).split())
            self._capture_title = False

def fetch(url):
    with urlopen(url, timeout=10) as r:
        return r.read().decode('utf-8', errors='replace')

def normalize_label(label):
    return re.sub(r'[^a-z0-9]', '', label.lower())

def parse_price_currency(value):
    if value is None:
        return ('', '')
    s = str(value).strip()
    if not s:
        return ('', '')
    currency = ''
    # Preserve explicit textual/symbol currency. NT$ is common in source pages.
    currency_patterns = [
        (r'NT\$', 'NT$'), (r'TWD', 'TWD'), (r'USD', 'USD'), (r'EUR', 'EUR'),
        (r'JPY', 'JPY'), (r'GBP', 'GBP'), (r'\$', '$'), (r'€', '€'), (r'£', '£'), (r'¥', '¥'),
    ]
    for pat, cur in currency_patterns:
        if re.search(pat, s, flags=re.I):
            currency = cur
            break
    m = re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', s)
    if not m:
        return (s, currency)
    num = m.group(0).replace(',', '')
    try:
        price = float(num)
        if price.is_integer():
            price = int(price)
    except Exception:
        price = num
    return (price, currency)

def memory_to_gb(value):
    if value is None:
        return ''
    s = str(value).strip()
    if not s:
        return ''
    m = re.search(r'([-+]?\d+(?:\.\d+)?)\s*([KMGT]?I?B|[KMGT]B|[KMGT])?', s, flags=re.I)
    if not m:
        return s
    num = float(m.group(1))
    unit = (m.group(2) or 'GB').upper().replace('IB','B')
    factors = {'KB': 1/(1024*1024), 'K': 1/(1024*1024), 'MB': 1/1024, 'M': 1/1024, 'GB': 1, 'G': 1, 'TB': 1024, 'T': 1024}
    gb = num * factors.get(unit, 1)
    return int(round(gb)) if abs(gb - round(gb)) < 1e-9 else round(gb, 4)

def parse_detail(html, url):
    dp = DetailParser()
    dp.feed(html)
    fields = {}
    for row in dp.rows:
        if len(row) >= 2:
            fields[normalize_label(row[0])] = row[1]
    def find_field(*names):
        for n in names:
            key = normalize_label(n)
            if key in fields:
                return fields[key]
        return ''
    price, price_currency = parse_price_currency(find_field('Price'))
    currency = find_field('Currency') or price_currency
    return {
        'Name': find_field('Name', 'Product Name') or dp.name or dp.title,
        'Price': price,
        'Currency': currency,
        'CPU': find_field('CPU', 'Processor'),
        'MemoryGB': memory_to_gb(find_field('MemoryGB', 'Memory', 'RAM')),
        'URL': url,
    }

list_html = fetch(BASE_URL)
print('Fetched product list:', BASE_URL)
print('List markup sample:', list_html[:500])
link_parser = LinkParser()
link_parser.feed(list_html)
links = []
seen = set()
for href, label in link_parser.links:
    if href:
        full = urljoin(BASE_URL, href)
        if full not in seen:
            links.append((full, label))
            seen.add(full)
print('Detail links found:', len(links), links)

records = []
for idx, (url, label) in enumerate(links, start=1):
    html = fetch(url)
    print(f'Detail {idx} URL: {url}')
    print('Detail markup sample:', html[:800])
    rec = parse_detail(html, url)
    print('Parsed record:', rec)
    records.append(rec)

headers = ['Name', 'Price', 'Currency', 'CPU', 'MemoryGB', 'URL']
wb = Workbook()
ws = wb.active
ws.title = 'details'
ws.append(headers)
for rec in records:
    ws.append([rec.get(h, '') for h in headers])
wb.save(OUTFILE)
print('Wrote:', OUTFILE)

# Verification by reopening saved output.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['details']
out_headers = [c.value for c in ws2[1]]
rows = list(ws2.iter_rows(min_row=2, values_only=True))
print('CHECK file exists:', OUTFILE.exists(), 'size:', OUTFILE.stat().st_size if OUTFILE.exists() else 0)
print('CHECK sheet:', ws2.title)
print('CHECK headers:', out_headers)
print('CHECK row_count:', len(rows), 'expected:', len(links))
assert out_headers == headers, f'headers mismatch: {out_headers}'
assert len(rows) == len(links), f'row count mismatch {len(rows)} != {len(links)}'
for i, row in enumerate(rows, start=2):
    rowdict = dict(zip(out_headers, row))
    print(f'CHECK row {i}: {rowdict}')
    assert rowdict['Name'], f'missing Name row {i}'
    assert isinstance(rowdict['Price'], (int, float)), f'Price not numeric row {i}: {rowdict["Price"]!r}'
    assert rowdict['Currency'] in ('NT$', 'TWD', 'USD', 'EUR', 'JPY', 'GBP', '$', '€', '£', '¥'), f'unexpected Currency row {i}: {rowdict["Currency"]!r}'
    assert rowdict['MemoryGB'] is None or isinstance(rowdict['MemoryGB'], (int, float)), f'MemoryGB not numeric/blank row {i}: {rowdict["MemoryGB"]!r}'
    assert rowdict['URL'] and str(rowdict['URL']).startswith('http://127.0.0.1:64493/t11/item/'), rowdict['URL']
# Specific check based on page markup just fetched.
assert [r[0] for r in rows] == ['Laptop 13', 'Laptop 15', 'Mini PC', 'Tablet 11']
assert [r[2] for r in rows] == ['NT$', 'NT$', 'NT$', 'NT$']
assert [r[4] for r in rows] == [16, 32, 16, 8]
print('All checks passed.')