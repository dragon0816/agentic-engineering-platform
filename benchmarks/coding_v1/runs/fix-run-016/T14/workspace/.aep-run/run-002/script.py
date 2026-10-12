from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

ws = Path.cwd()
input_path = ws / 'input' / 'large_results.csv'
output_dir = ws / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

required_headers = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']
rows = []
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('Input headers:', reader.fieldnames)
    missing = [h for h in required_headers if h not in (reader.fieldnames or [])]
    if missing:
        raise ValueError(f'Missing required headers: {missing}')
    for line_no, row in enumerate(reader, start=2):
        device = row['Device Model']
        status = row['Status']
        try:
            duration = float(row['DurationSeconds'])
        except Exception as e:
            raise ValueError(f'Invalid DurationSeconds at CSV line {line_no}: {row["DurationSeconds"]!r}') from e
        rows.append({
            'Test ID': row['Test ID'],
            'Device Model': device,
            'Status': status,
            'DurationSeconds': duration,
        })

print('Input row count:', len(rows))
print('First 5 parsed rows:', rows[:5])

stats = defaultdict(lambda: {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0})
overall = {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0}

for row in rows:
    device = row['Device Model']
    status = row['Status']
    duration = row['DurationSeconds']
    for bucket in (stats[device], overall):
        bucket['Total'] += 1
        bucket['DurationSum'] += duration
        if status == 'Pass':
            bucket['PassCount'] += 1
        elif status == 'Fail':
            bucket['FailCount'] += 1

headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
overall_headers = ['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']

wb = Workbook()
summary_ws = wb.active
summary_ws.title = 'Summary'
summary_ws.append(headers)
for cell in summary_ws[1]:
    cell.font = Font(bold=True)

for device in sorted(stats.keys()):
    bucket = stats[device]
    total = bucket['Total']
    pass_count = bucket['PassCount']
    fail_count = bucket['FailCount']
    pass_rate = pass_count / total if total else 0
    avg_duration = bucket['DurationSum'] / total if total else 0
    summary_ws.append([device, total, pass_count, fail_count, pass_rate, avg_duration])

for row in summary_ws.iter_rows(min_row=2, min_col=5, max_col=5):
    row[0].number_format = '0.00%'
for row in summary_ws.iter_rows(min_row=2, min_col=6, max_col=6):
    row[0].number_format = '0.00'
summary_ws.freeze_panes = 'A2'
for col in range(1, len(headers) + 1):
    summary_ws.column_dimensions[chr(64 + col)].width = 20

overall_ws = wb.create_sheet('Overall')
overall_ws.append(overall_headers)
for cell in overall_ws[1]:
    cell.font = Font(bold=True)
total = overall['Total']
pass_count = overall['PassCount']
fail_count = overall['FailCount']
pass_rate = pass_count / total if total else 0
avg_duration = overall['DurationSum'] / total if total else 0
overall_ws.append([total, pass_count, fail_count, pass_rate, avg_duration])
overall_ws['D2'].number_format = '0.00%'
overall_ws['E2'].number_format = '0.00'
for col in range(1, len(overall_headers) + 1):
    overall_ws.column_dimensions[chr(64 + col)].width = 20

wb.save(output_path)
print('Wrote:', output_path)

# Reopen and verify every stated requirement from saved output.
wb2 = load_workbook(output_path, data_only=True)
print('Workbook sheets:', wb2.sheetnames)
if 'Summary' not in wb2.sheetnames:
    raise AssertionError('Missing Summary sheet')
if 'Overall' not in wb2.sheetnames:
    raise AssertionError('Missing Overall sheet')

sws = wb2['Summary']
ows = wb2['Overall']
actual_headers = [sws.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
actual_overall_headers = [ows.cell(row=1, column=i).value for i in range(1, len(overall_headers)+1)]
print('Summary headers:', actual_headers)
print('Overall headers:', actual_overall_headers)
if actual_headers != headers:
    raise AssertionError(f'Summary headers mismatch: {actual_headers}')
if actual_overall_headers != overall_headers:
    raise AssertionError(f'Overall headers mismatch: {actual_overall_headers}')

summary_rows = list(sws.iter_rows(min_row=2, values_only=True))
# Remove completely blank rows, if any.
summary_rows = [r for r in summary_rows if any(v is not None for v in r)]
expected_device_count = len(stats)
print('Summary data row count:', len(summary_rows), 'expected:', expected_device_count)
if len(summary_rows) != expected_device_count:
    raise AssertionError('Summary row count does not match number of Device Model groups')

saved_stats = {}
for r in summary_rows:
    device, total, pc, fc, pr, avg = r
    if not isinstance(device, str):
        raise AssertionError(f'Device Model is not text: {device!r}')
    for name, value in [('Total', total), ('PassCount', pc), ('FailCount', fc)]:
        if not isinstance(value, int):
            raise AssertionError(f'{name} for {device} is not integer: {value!r}')
    for name, value in [('PassRate', pr), ('AvgDurationSeconds', avg)]:
        if not isinstance(value, (int, float)):
            raise AssertionError(f'{name} for {device} is not numeric: {value!r}')
    b = stats[device]
    exp_total = b['Total']
    exp_pc = b['PassCount']
    exp_fc = b['FailCount']
    exp_pr = exp_pc / exp_total if exp_total else 0
    exp_avg = b['DurationSum'] / exp_total if exp_total else 0
    if total != exp_total or pc != exp_pc or fc != exp_fc or abs(pr - exp_pr) > 1e-12 or abs(avg - exp_avg) > 1e-12:
        raise AssertionError(f'Stats mismatch for {device}: saved={r}, expected={(device, exp_total, exp_pc, exp_fc, exp_pr, exp_avg)}')
    saved_stats[device] = r

overall_data_rows = [r for r in ows.iter_rows(min_row=2, values_only=True) if any(v is not None for v in r)]
print('Overall data row count:', len(overall_data_rows), 'expected: 1')
if len(overall_data_rows) != 1:
    raise AssertionError('Overall sheet should contain exactly one data row')
orow = overall_data_rows[0]
for name, value in zip(overall_headers[:3], orow[:3]):
    if not isinstance(value, int):
        raise AssertionError(f'Overall {name} is not integer: {value!r}')
for name, value in zip(overall_headers[3:], orow[3:]):
    if not isinstance(value, (int, float)):
        raise AssertionError(f'Overall {name} is not numeric: {value!r}')
exp_overall = (
    overall['Total'],
    overall['PassCount'],
    overall['FailCount'],
    overall['PassCount'] / overall['Total'] if overall['Total'] else 0,
    overall['DurationSum'] / overall['Total'] if overall['Total'] else 0,
)
if any(abs(a-b) > 1e-12 if isinstance(a, float) or isinstance(b, float) else a != b for a, b in zip(orow, exp_overall)):
    raise AssertionError(f'Overall mismatch: saved={orow}, expected={exp_overall}')

print('Verification passed.')
print('Overall saved:', orow)
print('First 5 summary rows:', summary_rows[:5])
print('Output files:', [p.name for p in sorted(output_dir.iterdir())])
