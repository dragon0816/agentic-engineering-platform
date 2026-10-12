from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'specs.md'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split product sections by Markdown H2 headings.
sections = []
current_title = None
current_lines = []
for line in text.splitlines():
    m = re.match(r'^##\s+(.+?)\s*$', line)
    if m:
        if current_title is not None:
            sections.append((current_title, '\n'.join(current_lines)))
        current_title = m.group(1).strip()
        current_lines = []
    elif current_title is not None:
        current_lines.append(line)
if current_title is not None:
    sections.append((current_title, '\n'.join(current_lines)))

print('Detected product sections:', [title for title, _ in sections])

unit_to_mhz = {
    'mhz': 1.0,
    'ghz': 1000.0,
    'khz': 0.001,
    'hz': 0.000001,
}

def fmt_num(value):
    if value is None:
        return None
    if abs(value - round(value)) < 1e-9:
        return int(round(value))
    return float(value)

def extract_field(body, name):
    # Extract only explicitly written bullet fields. Do not infer anything.
    pattern = re.compile(r'^\s*-\s*' + re.escape(name) + r'\s*:\s*(.*?)\s*$', re.IGNORECASE | re.MULTILINE)
    m = pattern.search(body)
    return m.group(1).strip() if m else None

def parse_frequency(value):
    if not value:
        return None, None
    s = value.strip().replace('−', '-').replace('–', '-').replace('—', '-')
    # Match explicit numeric range with one unit at end or units on both ends.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)?\s*-\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', s, re.IGNORECASE)
    if not m:
        return None, None
    a = float(m.group(1))
    unit_a = (m.group(2) or m.group(4)).lower()
    b = float(m.group(3))
    unit_b = m.group(4).lower()
    return fmt_num(a * unit_to_mhz[unit_a]), fmt_num(b * unit_to_mhz[unit_b])

def parse_bandwidth(value):
    if not value:
        return None
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', value, re.IGNORECASE)
    if not m:
        return None
    return fmt_num(float(m.group(1)) * unit_to_mhz[m.group(2).lower()])

def normalize_interface(value):
    if not value:
        return None
    # Multiple interfaces separated by slash/comma/semicolon are written with semicolons.
    parts = [p.strip() for p in re.split(r'\s*(?:/|,|;|、)\s*', value.strip()) if p.strip()]
    return ';'.join(parts) if parts else None

rows = []
for title, body in sections:
    model = extract_field(body, 'Model')
    freq = extract_field(body, 'Frequency')
    bandwidth = extract_field(body, 'Bandwidth')
    interface = extract_field(body, 'Interface')
    fmin, fmax = parse_frequency(freq)
    bw_mhz = parse_bandwidth(bandwidth)
    iface = normalize_interface(interface)
    row = {
        'Model': model,
        'FrequencyMinMHz': fmin,
        'FrequencyMaxMHz': fmax,
        'BandwidthMHz': bw_mhz,
        'Interface': iface,
    }
    rows.append(row)
    print('Parsed', title, row)

wb = Workbook()
ws = wb.active
ws.title = 'specs'
headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(output_path)
print('Wrote:', output_path)

# Verification: reopen the saved workbook and confirm stated requirements.
check_wb = load_workbook(output_path, data_only=True)
check_ws = check_wb['specs']
read_headers = [cell.value for cell in check_ws[1]]
read_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
print('Reopened headers:', read_headers)
print('Reopened row count:', len(read_rows))
for r in read_rows:
    print('Reopened row:', r)

assert read_headers == headers, f'Unexpected headers: {read_headers}'
assert len(read_rows) == 3, f'Expected 3 product rows, got {len(read_rows)}'
expected = [
    ('RX-100', 2400, 2500, 80, 'USB;LAN'),
    ('RX-200', 900, 930, 20, 'LAN'),
    ('RX-300', 5000, 6000, None, 'USB'),
]
assert read_rows == expected, f'Rows differ. Expected {expected}, got {read_rows}'
# Confirm numeric columns are numeric where present and missing bandwidth remains blank/None.
for row in read_rows:
    for idx in [1, 2, 3]:
        val = row[idx]
        if val is not None:
            assert isinstance(val, (int, float)), f'Column {headers[idx]} should be numeric, got {val!r}'
assert read_rows[2][3] is None, 'Missing Bandwidth for RX-300 should remain blank'
print('Verification passed: output file, sheet name, headers, 3 rows, MHz conversions, semicolon interfaces, numeric values, and blank missing Bandwidth were confirmed.')