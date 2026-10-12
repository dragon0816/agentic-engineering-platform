from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV without modifying input; keep all source column names and row order.
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    rows = list(reader)

if 'Status' not in headers:
    raise ValueError('Required column Status not found')
if 'Device ID' not in headers:
    raise ValueError('Required column Device ID not found')

failed_rows = [row for row in rows if (row.get('Status') or '').strip().lower() == 'fail']

wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

# Write headers
for col_idx, header in enumerate(headers, start=1):
    ws.cell(row=1, column=col_idx, value=header)

# Write rows, keeping all fields. Device ID as text to preserve leading zeros.
device_col_idx = headers.index('Device ID') + 1
for row_idx, row in enumerate(failed_rows, start=2):
    for col_idx, header in enumerate(headers, start=1):
        value = row.get(header, '')
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        if col_idx == device_col_idx:
            cell.value = '' if value is None else str(value)
            cell.number_format = '@'

# Set Device ID column text format, including header column cells for future clarity.
for cell in ws.iter_cols(min_col=device_col_idx, max_col=device_col_idx, min_row=1, max_row=max(1, len(failed_rows)+1))[0]:
    cell.number_format = '@'

wb.save(output_path)

# Verification by reopening saved file.
wb2 = load_workbook(output_path, data_only=False)
if wb2.sheetnames != ['failed_items']:
    raise AssertionError(f'Unexpected sheet names: {wb2.sheetnames}')
ws2 = wb2['failed_items']
actual_headers = [ws2.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
if actual_headers != headers:
    raise AssertionError(f'Headers mismatch: {actual_headers} != {headers}')

actual_rows = []
for r in range(2, ws2.max_row + 1):
    actual_rows.append([ws2.cell(row=r, column=c).value for c in range(1, len(headers)+1)])
expected_rows = [[row.get(h, '') for h in headers] for row in failed_rows]
# openpyxl may return None for empty cells; expected failed rows here should not depend on empty cells, but normalize just in case.
actual_norm = [['' if v is None else str(v) for v in row] for row in actual_rows]
expected_norm = [['' if v is None else str(v) for v in row] for row in expected_rows]
if actual_norm != expected_norm:
    raise AssertionError(f'Rows mismatch: {actual_norm} != {expected_norm}')

# Check every output row still satisfies the requested Status rule and order matches CSV filtering.
status_idx = headers.index('Status')
for row in actual_norm:
    if row[status_idx].strip().lower() != 'fail':
        raise AssertionError(f'Non-fail row found: {row}')

# Check Device ID cells are stored as strings with text format and leading zeros preserved.
for excel_row, source_row in enumerate(failed_rows, start=2):
    cell = ws2.cell(row=excel_row, column=device_col_idx)
    expected = str(source_row['Device ID'])
    if cell.value != expected:
        raise AssertionError(f'Device ID mismatch at row {excel_row}: {cell.value!r} != {expected!r}')
    if not isinstance(cell.value, str):
        raise AssertionError(f'Device ID not text at row {excel_row}: {type(cell.value)}')
    if cell.number_format != '@':
        raise AssertionError(f'Device ID cell not text formatted at row {excel_row}: {cell.number_format!r}')

print('created:', output_path)
print('input row count excluding header:', len(rows))
print('failed row count:', len(failed_rows))
print('headers preserved:', headers)
print('failed Device IDs:', [row['Device ID'] for row in failed_rows])
print('checks passed: sheet name failed_items; all columns preserved; original order preserved; Status stripped/case-insensitive == Fail; Device ID text with leading zeros preserved')
print('output files:', [p.name for p in output_dir.iterdir() if p.is_file()])