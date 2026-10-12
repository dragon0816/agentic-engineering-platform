from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
input_dir = base / 'input'
print('Workspace:', base)
print('Input files:')
for p in sorted(input_dir.glob('*')):
    print(' -', p.name, 'size=', p.stat().st_size)
    if p.suffix.lower() in ['.xlsx', '.xlsm']:
        wb = load_workbook(p, read_only=True, data_only=False)
        print('   sheets:', wb.sheetnames)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5), values_only=True))
            print(f'   sheet {ws.title}: rows={ws.max_row}, cols={ws.max_column}')
            for i, row in enumerate(rows, start=1):
                print('    row', i, row)
            break
        wb.close()
