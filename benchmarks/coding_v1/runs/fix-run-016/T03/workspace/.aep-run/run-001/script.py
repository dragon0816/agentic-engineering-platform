from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
print('workspace:', workspace)
print('input files:', [p.name for p in input_dir.iterdir()])
path = input_dir / 'tasks.xlsx'
print('tasks.xlsx exists:', path.exists())
wb = load_workbook(path, read_only=True, data_only=False)
print('sheetnames:', wb.sheetnames)
ws = wb['tasks']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print('headers:', headers)
print('max_row:', ws.max_row, 'max_column:', ws.max_column)
print('sample rows:')
for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 6), values_only=True):
    print(row)
wb.close()
print('output dir exists:', output_dir.exists())