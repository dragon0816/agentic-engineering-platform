from pathlib import Path
from collections import OrderedDict, Counter
from openpyxl import load_workbook, Workbook
from openpyxl.utils import get_column_letter

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'owners.xlsx'

INVALID_SHEET_CHARS = set('[]:*?/\\')
RESERVED_SHEET_NAMES = {'Summary'}

def owner_key(value):
    """Group key: blank/whitespace/None Owner becomes Unassigned; otherwise preserve exact value as string."""
    if value is None:
        return 'Unassigned'
    s = str(value)
    if s.strip() == '':
        return 'Unassigned'
    return s

def make_legal_sheet_name(owner, used):
    """Convert Owner to a legal, unique Excel worksheet name (<=31 chars, no invalid chars)."""
    base = ''.join('_' if ch in INVALID_SHEET_CHARS else ch for ch in str(owner))
    base = base.strip()
    if not base:
        base = 'Unassigned'
    # Excel sheet names cannot exceed 31 characters.
    base = base[:31]
    if base in RESERVED_SHEET_NAMES:
        base = base[:24] + '_Owner'
    candidate = base
    i = 1
    while candidate in used or candidate in RESERVED_SHEET_NAMES:
        suffix = f'_{i}'
        candidate = base[:31-len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate

# Read source workbook without modifying it.
src_wb = load_workbook(input_path, data_only=False)
if 'tasks' not in src_wb.sheetnames:
    raise ValueError('Input workbook does not contain a tasks worksheet')
src_ws = src_wb['tasks']
headers = [cell.value for cell in src_ws[1]]
if 'Owner' not in headers:
    raise ValueError('tasks worksheet does not contain Owner column')
if 'Status' not in headers:
    raise ValueError('tasks worksheet does not contain Status column')
owner_idx = headers.index('Owner')
status_idx = headers.index('Status')

# Preserve first-seen Owner grouping order and row values.
groups = OrderedDict()
for row in src_ws.iter_rows(min_row=2, values_only=True):
    # Skip completely empty trailing rows, if any.
    if all(v is None for v in row):
        continue
    key = owner_key(row[owner_idx])
    groups.setdefault(key, []).append(tuple(row))
src_wb.close()

# Create output workbook.
out_wb = Workbook()
# Rename default sheet to Summary and populate after owner sheets are prepared.
summary_ws = out_wb.active
summary_ws.title = 'Summary'
summary_headers = ['Owner', 'Total', 'Pass', 'Fail', 'SheetName']
summary_ws.append(summary_headers)

used_sheet_names = {'Summary'}
summary_rows = []
for owner, rows in groups.items():
    sheet_name = make_legal_sheet_name(owner, used_sheet_names)
    ws = out_wb.create_sheet(title=sheet_name)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    counts = Counter(str(row[status_idx]) if row[status_idx] is not None else '' for row in rows)
    total = len(rows)
    passed = counts.get('Pass', 0)
    failed = counts.get('Fail', 0)
    summary_rows.append([owner, total, passed, failed, sheet_name])

for row in summary_rows:
    summary_ws.append(row)

# Basic readable column widths.
for ws in out_wb.worksheets:
    for col_idx, header in enumerate(ws[1], start=1):
        max_len = len(str(header.value)) if header.value is not None else 0
        for cell in ws.iter_cols(min_col=col_idx, max_col=col_idx, min_row=2, values_only=True):
            for value in cell:
                max_len = max(max_len, len(str(value)) if value is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

out_wb.save(output_path)
out_wb.close()

# Reopen and verify all stated requirements from the saved output.
check_wb = load_workbook(output_path, data_only=True)
errors = []
if 'Summary' not in check_wb.sheetnames:
    errors.append('Missing Summary worksheet')
else:
    s = check_wb['Summary']
    got_headers = [c.value for c in s[1]]
    if got_headers != summary_headers:
        errors.append(f'Summary headers mismatch: {got_headers}')
    summary_data = list(s.iter_rows(min_row=2, values_only=True))
    if len(summary_data) != len(groups):
        errors.append(f'Summary row count {len(summary_data)} != group count {len(groups)}')
    for owner, total, passed, failed, sheet_name in summary_data:
        if sheet_name not in check_wb.sheetnames:
            errors.append(f'Summary references missing sheet: {sheet_name}')
            continue
        if any(ch in sheet_name for ch in INVALID_SHEET_CHARS) or len(sheet_name) > 31 or not sheet_name:
            errors.append(f'Illegal worksheet name: {sheet_name}')
        ws = check_wb[sheet_name]
        sheet_headers = [c.value for c in ws[1]]
        if sheet_headers != headers:
            errors.append(f'{sheet_name} headers changed: {sheet_headers}')
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        if len(data_rows) != total:
            errors.append(f'{sheet_name} row count {len(data_rows)} != Summary Total {total}')
        pass_count = sum(1 for r in data_rows if r[status_idx] == 'Pass')
        fail_count = sum(1 for r in data_rows if r[status_idx] == 'Fail')
        if pass_count != passed or fail_count != failed:
            errors.append(f'{sheet_name} Pass/Fail count mismatch: {pass_count}/{fail_count} vs {passed}/{failed}')
        for r in data_rows:
            expected_owner = owner_key(r[owner_idx])
            if expected_owner != owner:
                errors.append(f'{sheet_name} contains row for owner group {expected_owner}, expected {owner}')

# Verify duplicate worksheet names are absent by nature and expected output file exists/nonempty.
if len(check_wb.sheetnames) != len(set(check_wb.sheetnames)):
    errors.append('Duplicate worksheet names found')
check_wb.close()
if not output_path.exists() or output_path.stat().st_size == 0:
    errors.append('owners.xlsx missing or empty')

if errors:
    print('VALIDATION FAILED')
    for e in errors:
        print('-', e)
    raise SystemExit(1)

print('Created:', output_path)
print('Workbook sheets:', load_workbook(output_path, read_only=True).sheetnames)
print('Summary rows:')
check_wb = load_workbook(output_path, data_only=True, read_only=True)
for row in check_wb['Summary'].iter_rows(values_only=True):
    print(row)
check_wb.close()
print('Validation passed: output file exists/nonempty; Summary headers Owner, Total, Pass, Fail, SheetName; one sheet per Owner group; blank Owner grouped as Unassigned while row Owner values are preserved; worksheet names are legal and unique; row counts and Pass/Fail counts match saved sheets.')