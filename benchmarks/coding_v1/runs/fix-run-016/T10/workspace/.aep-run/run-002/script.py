from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
import re
from openpyxl import Workbook, load_workbook

CATALOG_URL = 'http://127.0.0.1:64493/t10/catalog'
ws = Path.cwd()
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)
out_file = out_dir / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_catalog = False
        self.catalog_depth = 0
        self.in_card = False
        self.card_depth = 0
        self.current = None
        self.capture = None
        self.products = []

    def handle_starttag(self, tag, attrs):
        attrs_d = dict(attrs)
        if tag == 'main' and attrs_d.get('id') == 'catalog':
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if self.in_catalog:
            self.catalog_depth += 1
            classes = attrs_d.get('class', '').split()
            if tag == 'div' and 'card' in classes and not self.in_card:
                self.in_card = True
                self.card_depth = 1
                self.current = {'產品名稱': '', '價格': None, '幣別': '', '產品網址': ''}
                return
            if self.in_card:
                self.card_depth += 1
                if tag == 'a' and not self.current['產品網址']:
                    self.current['產品網址'] = urljoin(CATALOG_URL, attrs_d.get('href', ''))
                    self.capture = 'name'
                elif tag == 'span' and 'price' in classes:
                    self.capture = 'price'

    def handle_endtag(self, tag):
        if self.in_card:
            self.card_depth -= 1
            if self.card_depth == 0:
                self.products.append(self.current)
                self.current = None
                self.in_card = False
                self.capture = None
                return
        if self.in_catalog:
            self.catalog_depth -= 1
            if self.catalog_depth == 0:
                self.in_catalog = False
        self.capture = None

    def handle_data(self, data):
        if not (self.in_card and self.current is not None and self.capture):
            return
        text = data.strip()
        if not text:
            return
        if self.capture == 'name':
            self.current['產品名稱'] += text
        elif self.capture == 'price':
            raw = text
            m = re.match(r'\s*([^\d\-.,]+)?\s*([\d,]+(?:\.\d+)?)', raw)
            if not m:
                raise ValueError(f'Cannot parse price: {raw!r}')
            currency = (m.group(1) or '').strip()
            number = m.group(2).replace(',', '')
            value = float(number) if '.' in number else int(number)
            self.current['幣別'] = currency
            self.current['價格'] = value

with urlopen(CATALOG_URL, timeout=10) as r:
    html = r.read().decode('utf-8', errors='replace')

parser = CatalogParser()
parser.feed(html)
products = parser.products
print('Parsed products:', len(products))
for i, p in enumerate(products, 1):
    print(i, p)

wb = Workbook()
wsheet = wb.active
wsheet.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
wsheet.append(headers)
for p in products:
    wsheet.append([p['產品名稱'], p['價格'], p['幣別'], p['產品網址']])
wb.save(out_file)
print('Wrote:', out_file)

# Verification: reopen saved output and confirm requirements.
check_wb = load_workbook(out_file, data_only=True)
check_ws = check_wb['products']
read_headers = [check_ws.cell(1, c).value for c in range(1, 5)]
rows = []
for r in range(2, check_ws.max_row + 1):
    rows.append([check_ws.cell(r, c).value for c in range(1, 5)])

assert read_headers == headers, f'Headers mismatch: {read_headers}'
assert len(rows) == 5, f'Expected 5 products from main catalog, got {len(rows)}'
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
expected_prices = [1299, 2450, 99, 560, 12000]
expected_currency = ['NT$'] * 5
expected_urls = [urljoin(CATALOG_URL, path) for path in ['/p/router-a','/p/switch-b','/p/cable-c','/p/adapter-d','/p/antenna-e']]
for idx, row in enumerate(rows):
    name, price, currency, product_url = row
    assert name == expected_names[idx], f'Name/order mismatch row {idx+2}: {name}'
    assert price == expected_prices[idx] and isinstance(price, (int, float)), f'Price mismatch/type row {idx+2}: {price} {type(price)}'
    assert currency == expected_currency[idx], f'Currency mismatch row {idx+2}: {currency}'
    assert product_url == expected_urls[idx], f'URL mismatch row {idx+2}: {product_url}'
    assert '/ad/' not in product_url and not str(name).startswith('Sponsored') and '推薦' not in str(name), f'Ad/recommendation included row {idx+2}'
print('Verification passed: products.xlsx exists, headers correct, 5 main-catalog rows in page order, prices numeric, URLs absolute, ads/recommendations excluded.')
print('Output files:', [p.name for p in out_dir.iterdir()])