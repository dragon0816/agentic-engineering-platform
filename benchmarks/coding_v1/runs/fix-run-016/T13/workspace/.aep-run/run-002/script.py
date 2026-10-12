from pathlib import Path
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from urllib.parse import urljoin
from html.parser import HTMLParser
from html import unescape
import json
import re
import time
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:64493/t13/list'
MAX_500_ATTEMPTS = 3
COLUMNS = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._current_href = None
        self._text_parts = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs = dict(attrs)
            self._current_href = attrs.get('href')
            self._text_parts = []
    def handle_data(self, data):
        if self._current_href is not None:
            self._text_parts.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._current_href is not None:
            name = ' '.join(''.join(self._text_parts).split())
            self.links.append({'name': name, 'href': self._current_href})
            self._current_href = None
            self._text_parts = []

def fetch_url(url, detail=False):
    """Return dict with ok, status, text, attempts. For detail HTTP 500, retry up to MAX_500_ATTEMPTS total attempts."""
    attempts = 0
    max_attempts = MAX_500_ATTEMPTS if detail else 1
    last = {'ok': False, 'status': None, 'text': '', 'attempts': 0, 'error': None}
    while attempts < max_attempts:
        attempts += 1
        try:
            req = Request(url, headers={'User-Agent': 'AEP-product-fetch/1.0'})
            with urlopen(req, timeout=10) as resp:
                body = resp.read().decode('utf-8', errors='replace')
                return {'ok': 200 <= resp.status < 400, 'status': resp.status, 'text': body, 'attempts': attempts, 'error': None}
        except HTTPError as e:
            body = e.read().decode('utf-8', errors='replace') if e.fp else ''
            last = {'ok': False, 'status': e.code, 'text': body, 'attempts': attempts, 'error': f'HTTP {e.code} {e.reason}'}
            if not (detail and e.code == 500 and attempts < max_attempts):
                break
            time.sleep(0.15)
        except URLError as e:
            last = {'ok': False, 'status': None, 'text': '', 'attempts': attempts, 'error': f'URL error: {e.reason}'}
            break
    return last

def strip_tags_to_lines(html):
    s = re.sub(r'(?is)<script.*?</script>|<style.*?</style>', ' ', html)
    s = re.sub(r'(?i)<br\s*/?>', '\n', s)
    s = re.sub(r'(?i)</?(tr|td|th|dt|dd|li|p|div|section|article|h[1-6]|span|label|strong|b)[^>]*>', '\n', s)
    s = re.sub(r'<[^>]+>', ' ', s)
    s = unescape(s)
    lines = [' '.join(line.split()) for line in s.splitlines()]
    return [line for line in lines if line]

def clean_html_fragment(fragment):
    return ' '.join(unescape(re.sub(r'<[^>]+>', ' ', fragment)).split())

def extract_class_value(html, key):
    # Values in an element whose id/class/data-field contains the key.
    pat = re.compile(r'(?is)<([a-z0-9]+)\b[^>]*(?:class|id|data-field|data-key)\s*=\s*["\'][^"\']*' + re.escape(key) + r'[^"\']*["\'][^>]*>(.*?)</\1>')
    m = pat.search(html)
    if m:
        val = clean_html_fragment(m.group(2))
        if val:
            return val
    return None

def extract_label_value(html, lines, label):
    # Table/dl pattern: <td>Price</td><td>123</td> or <dt>CPU</dt><dd>...</dd>
    label_re = re.escape(label)
    pat = re.compile(r'(?is)>\s*' + label_re + r'\s*</(?:td|th|dt|span|div|label|strong|b)[^>]*>\s*<[^>]+>\s*(.*?)\s*</[^>]+>')
    m = pat.search(html)
    if m:
        val = clean_html_fragment(m.group(1))
        if val and val.lower() != label.lower():
            return val
    # Inline text pattern.
    text = '\n'.join(lines)
    m = re.search(r'(?im)\b' + label_re + r'\b\s*[:：]\s*([^\n\r]+)', text)
    if m:
        val = m.group(1).strip()
        if val:
            return val
    # Adjacent line pattern.
    for i, line in enumerate(lines):
        if line.strip().lower() == label.lower() and i + 1 < len(lines):
            return lines[i + 1]
        if line.lower().startswith(label.lower() + ' '):
            val = line[len(label):].strip(' :-：')
            if val:
                return val
    return None

def extract_h1(html):
    m = re.search(r'(?is)<h1[^>]*>(.*?)</h1>', html)
    if m:
        val = clean_html_fragment(m.group(1))
        if val:
            return val
    return None

def split_price_currency(price_text, currency_text):
    if not price_text:
        return 'N/A', (currency_text or 'N/A')
    raw = price_text.strip()
    if not raw or raw.upper() == 'N/A':
        return 'N/A', (currency_text or 'N/A')
    currency = currency_text.strip() if currency_text else None
    upper = raw.upper()
    if not currency:
        for code in ['USD', 'EUR', 'TWD', 'NTD', 'JPY', 'CNY', 'GBP', 'AUD', 'CAD']:
            if re.search(r'\b' + code + r'\b', upper):
                currency = code
                break
    if not currency:
        symbols = {'$': 'USD', 'US$': 'USD', '€': 'EUR', '£': 'GBP', '¥': 'JPY', 'NT$': 'TWD'}
        for sym, code in sorted(symbols.items(), key=lambda kv: -len(kv[0])):
            if sym in raw:
                currency = code
                break
    m = re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', raw)
    if not m:
        return raw, (currency or 'N/A')
    num_text = m.group(0).replace(',', '')
    try:
        num = float(num_text)
        price = int(num) if num.is_integer() else num
    except ValueError:
        price = num_text
    return price, (currency or 'N/A')

def parse_detail(html, list_name):
    lines = strip_tags_to_lines(html)
    name = extract_class_value(html, 'name') or extract_class_value(html, 'title') or extract_label_value(html, lines, 'Name') or extract_h1(html) or list_name or 'N/A'
    price_text = extract_class_value(html, 'price') or extract_label_value(html, lines, 'Price')
    currency_text = extract_class_value(html, 'currency') or extract_label_value(html, lines, 'Currency')
    cpu = extract_class_value(html, 'cpu') or extract_label_value(html, lines, 'CPU') or 'N/A'
    price, currency = split_price_currency(price_text, currency_text)
    return {
        'Name': name or list_name or 'N/A',
        'Price': price if price not in (None, '') else 'N/A',
        'Currency': currency if currency not in (None, '') else 'N/A',
        'CPU': cpu if cpu not in (None, '') else 'N/A',
    }

# Fetch and parse list.
list_result = fetch_url(LIST_URL, detail=False)
if not list_result['ok']:
    raise RuntimeError(f"Cannot fetch product list {LIST_URL}: {list_result['error'] or list_result['status']}")
print('List status:', list_result['status'])
print('List markup sample:', list_result['text'][:500])
parser = LinkParser()
parser.feed(list_result['text'])
products = parser.links[:5]
print('Products found:', len(parser.links), 'using:', len(products))
for p in products:
    p['url'] = urljoin(LIST_URL, p['href'])
    print(' -', p['name'], p['url'])
if len(products) != 5:
    raise RuntimeError(f'Expected to collect five products from list, found {len(products)}')

rows = []
warnings = {'source_url': LIST_URL, 'products': []}
for p in products:
    result = fetch_url(p['url'], detail=True)
    if result['ok']:
        print(f"Detail OK: {p['name']} status={result['status']} attempts={result['attempts']} sample={result['text'][:250]!r}")
        parsed = parse_detail(result['text'], p['name'])
        row = {
            'Name': parsed['Name'],
            'Price': parsed['Price'],
            'Currency': parsed['Currency'],
            'CPU': parsed['CPU'],
            'URL': p['url'],
            'FetchStatus': 'ok',
        }
        warnings['products'].append({
            'Name': p['name'], 'URL': p['url'], 'FetchStatus': 'ok',
            'attempts': result['attempts'], 'http_status': result['status'], 'warning': None
        })
    else:
        print(f"Detail ERROR: {p['name']} status={result['status']} attempts={result['attempts']} error={result['error']}")
        row = {
            'Name': p['name'] or 'N/A',
            'Price': 'N/A',
            'Currency': 'N/A',
            'CPU': 'N/A',
            'URL': p['url'],
            'FetchStatus': 'error',
        }
        warnings['products'].append({
            'Name': p['name'], 'URL': p['url'], 'FetchStatus': 'error',
            'attempts': result['attempts'], 'http_status': result['status'],
            'warning': result['error'] or 'detail fetch failed'
        })
    rows.append(row)

# Write Excel.
xlsx_path = OUTPUT / 'robust_products.xlsx'
wb = Workbook()
ws = wb.active
ws.title = 'Products'
ws.append(COLUMNS)
for row in rows:
    ws.append([row[c] for c in COLUMNS])
wb.save(xlsx_path)

# Write warnings JSON.
json_path = OUTPUT / 'warnings.json'
with json_path.open('w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Verification: reopen outputs and check requirements.
wb2 = load_workbook(xlsx_path, data_only=True)
ws2 = wb2['Products']
headers = [cell.value for cell in ws2[1]]
data_rows = list(ws2.iter_rows(min_row=2, values_only=True))
assert headers == COLUMNS, f'Bad headers: {headers}'
assert len(data_rows) == 5, f'Expected 5 rows, got {len(data_rows)}'
for idx, values in enumerate(data_rows, start=2):
    record = dict(zip(COLUMNS, values))
    assert record['Name'] not in (None, ''), f'Missing Name row {idx}'
    assert record['URL'] and str(record['URL']).startswith('http://127.0.0.1:64493/t13/item/'), f'Bad URL row {idx}: {record["URL"]}'
    assert record['FetchStatus'] in ('ok', 'error'), f'Bad FetchStatus row {idx}: {record["FetchStatus"]}'
    if record['FetchStatus'] == 'error':
        assert record['Price'] == 'N/A' and record['Currency'] == 'N/A' and record['CPU'] == 'N/A', f'Error row {idx} unknown fields not N/A: {record}'
    else:
        # Successful rows may have N/A for missing price/spec by requirement; if present, price is numeric or text.
        assert record['Price'] not in (None, ''), f'Blank Price row {idx}'
        assert record['Currency'] not in (None, ''), f'Blank Currency row {idx}'
        assert record['CPU'] not in (None, ''), f'Blank CPU row {idx}'
with json_path.open('r', encoding='utf-8') as f:
    warnings2 = json.load(f)
assert warnings2['source_url'] == LIST_URL
assert len(warnings2['products']) == 5
for item in warnings2['products']:
    assert item['FetchStatus'] in ('ok', 'error')
    if item['FetchStatus'] == 'error' and item.get('http_status') == 500:
        assert item['attempts'] == 3, f"HTTP 500 error should have 3 attempts: {item}"
print('WROTE', xlsx_path)
print('WROTE', json_path)
print('VERIFIED robust_products.xlsx headers, 5 rows, URLs, FetchStatus values, N/A fields for errors, and warnings.json statuses/attempts.')
print('FINAL ROWS:')
for r in rows:
    print(r)
print('FINAL WARNINGS:')
print(json.dumps(warnings, ensure_ascii=False, indent=2))
