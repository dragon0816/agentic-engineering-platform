from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

ws = Path.cwd()
input_path = ws / 'input' / 'mixed.csv'
out_dir = ws / 'output'
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / 'converted.xlsx'

# Read CSV safely: UTF-8 BOM, Chinese, quoted commas/newlines, blank cells.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    rows = [row for row in reader]

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
try:
    device_col_idx = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_col_idx = None

# Write workbook without adding any index column.
wb = Workbook()
sh = wb.active
sh.title = 'converted'

for r_idx, row in enumerate(rows, start=1):
    # Preserve column count/order from the CSV header. csv.reader preserves blank fields.
    for c_idx, value in enumerate(row, start=1):
        cell = sh.cell(row=r_idx, column=c_idx, value=value)
        # Keep Device ID as text, including header and data cells, to preserve leading zeros.
        if device_col_idx is not None and c_idx == device_col_idx:
            cell.number_format = '@'

# Also set entire Device ID column style as text for future Excel display/editing.
if device_col_idx is not None:
    col_letter = sh.cell(row=1, column=device_col_idx).column_letter
    sh.column_dimensions[col_letter].number_format = '@'

wb.save(out_path)
print('WROTE:', out_path)
print('CSV rows including header:', len(rows))
print('CSV columns:', len(headers))
print('Headers:', headers)
print('Output files:', [p.name for p in sorted(out_dir.iterdir())])

# Verification: reopen source and output, compare exactly.
wb2 = load_workbook(out_path, data_only=False)
sh2 = wb2.active
xlsx_rows = []
for row in sh2.iter_rows(min_row=1, max_row=sh2.max_row, min_col=1, max_col=len(headers)):
    xlsx_rows.append(['' if cell.value is None else cell.value for cell in row])

# Check dimensions and exact cell values.
assert sh2.max_row == len(rows), f'Row count mismatch: xlsx={sh2.max_row}, csv={len(rows)}'
assert sh2.max_column == len(headers), f'Column count mismatch: xlsx={sh2.max_column}, csv={len(headers)}'
assert xlsx_rows == rows, f'Cell data mismatch: xlsx={xlsx_rows!r}, csv={rows!r}'

# Specific checks required by prompt.
assert xlsx_rows[0] == headers, 'Headers/order changed'
assert headers == ['Device ID', '名稱', '備註', 'Value'], 'Unexpected headers after conversion'
assert xlsx_rows[1] == xlsx_rows[4] == ['0001', '測試機甲', '正常', '10'], 'Duplicate row not preserved'
assert xlsx_rows[2][2] == '含, 逗號', 'Quoted comma field not preserved'
assert xlsx_rows[3][2] == '第一行\n第二行', 'Quoted newline field not preserved'
assert xlsx_rows[3][3] == '', 'Empty Value cell not kept blank'
assert xlsx_rows[5][1] == '' and xlsx_rows[5][2] == '', 'Empty text cells not kept blank'

if device_col_idx is not None:
    for r in range(2, sh2.max_row + 1):
        cell = sh2.cell(row=r, column=device_col_idx)
        assert isinstance(cell.value, str), f'Device ID row {r} is not text: {cell.value!r}'
        assert cell.number_format == '@', f'Device ID row {r} number format is not text: {cell.number_format!r}'
    assert sh2['A2'].value == '0001' and sh2['A3'].value == '0002' and sh2['A4'].value == '0003', 'Leading zeros not preserved'

print('CHECK PASSED: reopened converted.xlsx')
print(' - file exists and is non-empty:', out_path.exists(), out_path.stat().st_size)
print(' - no index column added; column count/order:', sh2.max_column, xlsx_rows[0])
print(' - row count including header:', sh2.max_row)
print(' - duplicate rows preserved:', xlsx_rows[1], xlsx_rows[4])
print(' - blank cells remain blank:', repr(xlsx_rows[3][3]), repr(xlsx_rows[5][1]), repr(xlsx_rows[5][2]))
print(' - Device ID values/types/formats:', [(sh2.cell(r, device_col_idx).value, type(sh2.cell(r, device_col_idx).value).__name__, sh2.cell(r, device_col_idx).number_format) for r in range(2, sh2.max_row + 1)])