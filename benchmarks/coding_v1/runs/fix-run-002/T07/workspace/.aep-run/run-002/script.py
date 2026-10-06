from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

input_path = Path('input/specs.md')
output_dir = Path('output')
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
rows = []

# Split into product sections beginning with level-2 headings.
sections = re.split(r'(?m)^##\s+', text)
for sec in sections[1:]:
    data = {}
    for line in sec.splitlines():
        m = re.match(r'^-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            data[m.group(1).strip()] = m.group(2).strip()
    if not data:
        continue

    model = data.get('Model', '')

    fmin = fmax = ''
    freq = data.get('Frequency', '')
    if freq:
        m = re.match(r'^\s*([0-9]+(?:\.[0-9]+)?)\s*[–-]\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\s*$', freq, re.I)
        if m:
            v1 = float(m.group(1))
            v2 = float(m.group(2))
            unit = m.group(3).lower()
            factor = 1000 if unit == 'ghz' else 1
            fmin = v1 * factor
            fmax = v2 * factor

    bw_mhz = ''
    bw = data.get('Bandwidth', '')
    if bw:
        m = re.match(r'^\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\s*$', bw, re.I)
        if m:
            val = float(m.group(1))
            unit = m.group(2).lower()
            bw_mhz = val * (1000 if unit == 'ghz' else 1)

    interface = data.get('Interface', '')
    if interface:
        # Multiple interfaces are written as slash-separated in the source.
        interface = ';'.join(part.strip() for part in interface.split('/') if part.strip())

    rows.append([model, fmin, fmax, bw_mhz, interface])

# Write workbook
wb = Workbook()
ws = wb.active
ws.title = 'specs'
ws.append(headers)
for row in rows:
    # Write whole-number floats as ints for cleaner cells.
    clean = []
    for x in row:
        if isinstance(x, float) and x.is_integer():
            clean.append(int(x))
        else:
            clean.append(x)
    ws.append(clean)
wb.save(output_path)

# Verify by reading back
wb2 = load_workbook(output_path, data_only=True)
ws2 = wb2.active
for row in ws2.iter_rows(values_only=True):
    print(row)
print(f'Wrote {output_path}')
