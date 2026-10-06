from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook

ws = Path.cwd()
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)
export_path = out_dir / 'export.py'
summary_path = out_dir / 'summary.xlsx'

fixed_code = '''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals = {}
    counts = {}
    with open(args.input, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            model = row["Model"]
            value_text = row["Value"].strip()
            if value_text == "":
                continue
            value = float(value_text)
            totals[model] = totals.get(model, 0.0) + value
            counts[model] = counts.get(model, 0) + 1

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(args.output)


if __name__ == "__main__":
    main()
'''

export_path.write_text(fixed_code, encoding='utf-8')
print(f'Wrote {export_path.relative_to(ws)}')

# Execute the requested CLI from the workspace. Use platform-neutral paths, equivalent to the requested command.
cmd = [sys.executable, str(export_path), '--input', str(ws / 'input' / 'input.csv'), '--output', str(summary_path)]
print('Running:', ' '.join(cmd))
completed = subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('exit_code:', completed.returncode)
if completed.stdout:
    print('stdout:', completed.stdout)
if completed.stderr:
    print('stderr:', completed.stderr)
completed.check_returncode()

wb = load_workbook(summary_path, data_only=True)
wsheet = wb['summary']
rows = list(wsheet.iter_rows(values_only=True))
print('Workbook rows:')
for row in rows:
    print(row)

expected = [('Model', 'Average', 'ValidCount'), ('M1', 15, 2), ('M2', 10, 2)]
assert rows == expected, f'Unexpected workbook contents: {rows!r}'
print('Verification passed: summary.xlsx contains expected Model, Average, ValidCount rows.')
