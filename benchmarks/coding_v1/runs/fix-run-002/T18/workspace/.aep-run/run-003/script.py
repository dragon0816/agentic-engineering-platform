from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook

ws = Path.cwd()
summary_path = ws / 'output' / 'summary.xlsx'
if summary_path.exists():
    summary_path.unlink()

cmd = [sys.executable, 'output\\export.py', '--input', 'input\\input.csv', '--output', 'output\\summary.xlsx']
print('Running from workspace:', ' '.join(cmd))
try:
    completed = subprocess.run(cmd, cwd=ws, text=True)
    print('exit_code:', completed.returncode)
    completed.check_returncode()
except OSError as exc:
    print('Subprocess failed with OSError:', repr(exc))
    raise

wb = load_workbook(summary_path, data_only=True)
rows = list(wb['summary'].iter_rows(values_only=True))
print('Workbook rows:')
for row in rows:
    print(row)
expected = [('Model', 'Average', 'ValidCount'), ('M1', 15, 2), ('M2', 10, 2)]
assert rows == expected, f'Unexpected workbook contents: {rows!r}'
print('Verification passed.')
