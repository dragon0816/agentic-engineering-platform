from pathlib import Path
import runpy
import sys
from openpyxl import load_workbook

ws = Path.cwd()
script_path = ws / 'output' / 'export.py'
summary_path = ws / 'output' / 'summary.xlsx'
if summary_path.exists():
    summary_path.unlink()

print('Executing script file with CLI argv equivalent:')
print('python output\\export.py --input input\\input.csv --output output\\summary.xlsx')
old_argv = sys.argv[:]
try:
    sys.argv = ['output\\export.py', '--input', 'input\\input.csv', '--output', 'output\\summary.xlsx']
    runpy.run_path(str(script_path), run_name='__main__')
finally:
    sys.argv = old_argv

print('summary exists:', summary_path.exists())
wb = load_workbook(summary_path, data_only=True)
rows = list(wb['summary'].iter_rows(values_only=True))
print('Workbook rows:')
for row in rows:
    print(row)
expected = [('Model', 'Average', 'ValidCount'), ('M1', 15, 2), ('M2', 10, 2)]
assert rows == expected, f'Unexpected workbook contents: {rows!r}'
print('Verification passed: output/export.py produced output/summary.xlsx with expected contents.')
