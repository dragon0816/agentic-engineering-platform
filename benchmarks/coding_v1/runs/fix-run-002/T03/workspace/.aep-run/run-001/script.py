from pathlib import Path
import re
from collections import OrderedDict, defaultdict
from copy import copy
from openpyxl import load_workbook, Workbook

workspace = Path.cwd()
input_path = workspace / "input" / "tasks.xlsx"
output_dir = workspace / "output"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / "owners.xlsx"

if not input_path.exists():
    raise FileNotFoundError(f"Input file not found: {input_path}")

src_wb = load_workbook(input_path)
if "tasks" not in src_wb.sheetnames:
    raise ValueError(f"Worksheet 'tasks' not found. Available sheets: {src_wb.sheetnames}")
ws = src_wb["tasks"]

# Read header row and locate required columns.
headers = [cell.value for cell in ws[1]]
try:
    owner_col_idx = next(i + 1 for i, v in enumerate(headers) if v == "Owner")
except StopIteration:
    raise ValueError("Column 'Owner' not found in tasks worksheet header row")
try:
    status_col_idx = next(i + 1 for i, v in enumerate(headers) if v == "Status")
except StopIteration:
    raise ValueError("Column 'Status' not found in tasks worksheet header row")

INVALID_TITLE_CHARS = re.compile(r"[\\/*?:\[\]]")

def display_text(value):
    if value is None:
        return ""
    return str(value)

def is_blank(value):
    return value is None or (isinstance(value, str) and value.strip() == "")

def owner_key(value):
    return "Unassigned" if is_blank(value) else str(value)

def sanitize_sheet_name(raw):
    """Make a legal Excel worksheet title while preserving as much text as possible."""
    name = str(raw) if raw is not None else "Unassigned"
    name = INVALID_TITLE_CHARS.sub("_", name)
    # Excel does not allow blank titles; avoid problematic surrounding apostrophes.
    name = name.strip("'")
    if name == "":
        name = "Sheet"
    return name[:31]

def unique_sheet_name(raw, used_lower):
    base = sanitize_sheet_name(raw)
    # openpyxl/Excel treat sheet names case-insensitively for duplicates.
    candidate = base
    n = 1
    while candidate.lower() in used_lower:
        suffix = f"_{n}"
        candidate = base[:31 - len(suffix)] + suffix
        n += 1
    used_lower.add(candidate.lower())
    return candidate

# Group source row numbers by owner, preserving first-seen owner-group order.
groups = OrderedDict()
counts = defaultdict(lambda: {"Total": 0, "Pass": 0, "Fail": 0})

for row_idx in range(2, ws.max_row + 1):
    owner_value = ws.cell(row_idx, owner_col_idx).value
    key = owner_key(owner_value)
    if key not in groups:
        groups[key] = []
    groups[key].append(row_idx)

    status = ws.cell(row_idx, status_col_idx).value
    counts[key]["Total"] += 1
    if status == "Pass":
        counts[key]["Pass"] += 1
    elif status == "Fail":
        counts[key]["Fail"] += 1

# Create target workbook.
out_wb = Workbook()
# Remove default sheet; create Summary explicitly first.
default_ws = out_wb.active
out_wb.remove(default_ws)
summary_ws = out_wb.create_sheet("Summary")

used_names = {"summary"}
sheet_for_owner = {}
for key in groups.keys():
    sheet_for_owner[key] = unique_sheet_name(key, used_names)

# Summary: Owner, Total, Pass, Fail, SheetName.
summary_headers = ["Owner", "Total", "Pass", "Fail", "SheetName"]
summary_ws.append(summary_headers)
for key, row_indices in groups.items():
    summary_ws.append([
        key,
        counts[key]["Total"],
        counts[key]["Pass"],
        counts[key]["Fail"],
        sheet_for_owner[key],
    ])

# Basic header formatting in summary.
for src_cell, dst_cell in zip(ws[1], summary_ws[1]):
    if src_cell.has_style:
        dst_cell.font = copy(src_cell.font)
        dst_cell.fill = copy(src_cell.fill)
        dst_cell.border = copy(src_cell.border)
        dst_cell.alignment = copy(src_cell.alignment)
        dst_cell.number_format = src_cell.number_format
        dst_cell.protection = copy(src_cell.protection)

# Owner sheets with original columns and original cell values/styles.
for key, row_indices in groups.items():
    dst_ws = out_wb.create_sheet(sheet_for_owner[key])

    # Copy column widths.
    for col_letter, dim in ws.column_dimensions.items():
        dst_ws.column_dimensions[col_letter].width = dim.width

    # Copy header row.
    for src_cell in ws[1]:
        dst_cell = dst_ws.cell(row=1, column=src_cell.column, value=src_cell.value)
        if src_cell.has_style:
            dst_cell.font = copy(src_cell.font)
            dst_cell.fill = copy(src_cell.fill)
            dst_cell.border = copy(src_cell.border)
            dst_cell.alignment = copy(src_cell.alignment)
            dst_cell.number_format = src_cell.number_format
            dst_cell.protection = copy(src_cell.protection)

    # Copy grouped rows. Owner cell original value is copied unchanged.
    dst_row = 2
    for src_row in row_indices:
        for src_cell in ws[src_row]:
            dst_cell = dst_ws.cell(row=dst_row, column=src_cell.column, value=src_cell.value)
            if src_cell.has_style:
                dst_cell.font = copy(src_cell.font)
                dst_cell.fill = copy(src_cell.fill)
                dst_cell.border = copy(src_cell.border)
                dst_cell.alignment = copy(src_cell.alignment)
                dst_cell.number_format = src_cell.number_format
                dst_cell.protection = copy(src_cell.protection)
        dst_row += 1

# Make summary columns readable.
for col in summary_ws.columns:
    max_len = 0
    col_letter = col[0].column_letter
    for cell in col:
        max_len = max(max_len, len(display_text(cell.value)))
    summary_ws.column_dimensions[col_letter].width = min(max_len + 2, 50)

out_wb.save(output_path)

print(f"Created: {output_path}")
print(f"Owner groups: {len(groups)}")
for key in groups:
    print(f"- {key!r}: total={counts[key]['Total']} pass={counts[key]['Pass']} fail={counts[key]['Fail']} sheet={sheet_for_owner[key]!r}")
