from pathlib import Path
import csv
from openpyxl import Workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / "input" / "results.csv"
output_dir = workspace / "output"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / "failed_items.xlsx"

if not input_path.exists():
    raise FileNotFoundError(f"Input file not found: {input_path}")

with input_path.open("r", encoding="utf-8-sig", newline="") as f:
    reader = csv.DictReader(f)
    if reader.fieldnames is None:
        raise ValueError("results.csv has no header row")
    fieldnames = reader.fieldnames
    status_col = next((name for name in fieldnames if name.strip().lower() == "status"), None)
    if status_col is None:
        raise ValueError("No Status column found in results.csv")

    failed_rows = []
    for row in reader:
        status_value = row.get(status_col, "")
        if status_value.strip().lower() == "fail":
            failed_rows.append(row)

wb = Workbook()
ws = wb.active
ws.title = "failed_items"

# Header
for col_idx, name in enumerate(fieldnames, start=1):
    cell = ws.cell(row=1, column=col_idx, value=name)
    cell.font = Font(bold=True)

# Data: write as strings to preserve original CSV text exactly, including Device ID leading zeros.
device_id_cols = {i + 1 for i, name in enumerate(fieldnames) if name.strip().lower() == "device id"}
for row_idx, row in enumerate(failed_rows, start=2):
    for col_idx, name in enumerate(fieldnames, start=1):
        value = row.get(name, "")
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        # Explicit text format for Device ID; strings also preserve leading zeros.
        if col_idx in device_id_cols:
            cell.number_format = "@"

# Keep Device ID header/data explicitly formatted as text.
for col_idx in device_id_cols:
    ws.cell(row=1, column=col_idx).number_format = "@"

wb.save(output_path)
print(f"Read: {input_path}")
print(f"Status column: {status_col}")
print(f"Failed rows written: {len(failed_rows)}")
print(f"Output: {output_path}")