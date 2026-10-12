from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV preserving values as strings so Device ID leading zeros remain intact.
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames
    if headers is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in headers:
        raise ValueError('Missing required column: Status')
    if 'Device ID' not in headers:
        raise ValueError('Missing required column: Device ID')
    failed_rows = [row for row in reader if (row.get('Status') or '').strip().lower() == 'fail']

wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

# Write header and data rows, keeping source column names/order.
ws.append(headers)
for row in failed_rows:
    values = [row.get(h, '') for h in headers]
    ws.append(values)

# Force Device ID column to text format and type.
device_col = headers.index('Device ID') + 1
for cell in ws.iter_cols(min_col=device_col, max_col=device_col, min_row=1, max_row=ws.max_row):
    for c in cell:
        c.number_format = '@'
        if c.value is not None:
            c.value = str(c.value)

wb.save(output_path)
print('wrote:', output_path)
print('failed row count:', len(failed_rows))
print('headers:', headers)
print('failed Device IDs:', [r['Device ID'] for r in failed_rows])

# Verification: reopen output and validate requirements.
wb2 = load_workbook(output_path, data_only=False)
if 'failed_items' not in wb2.sheetnames:
    raise AssertionError('failed_items worksheet missing')
ws2 = wb2['failed_items']
read_headers = [ws2.cell(row=1, column=i).value for i in range(1, ws2.max_column + 1)]
if read_headers != headers:
    raise AssertionError(f'Header mismatch: {read_headers!r} != {headers!r}')

read_rows = []
for r in range(2, ws2.max_row + 1):
    read_rows.append({headers[c-1]: ws2.cell(row=r, column=c).value for c in range(1, ws2.max_column + 1)})

if len(read_rows) != len(failed_rows):
    raise AssertionError(f'Row count mismatch: {len(read_rows)} != {len(failed_rows)}')

for idx, (expected, actual) in enumerate(zip(failed_rows, read_rows), start=2):
    for h in headers:
        if str(actual[h] if actual[h] is not None else '') != str(expected.get(h, '')):
            raise AssertionError(f'Value mismatch at row {idx}, column {h}: {actual[h]!r} != {expected.get(h, "")!r}')
    if str(actual['Status']).strip().lower() != 'fail':
        raise AssertionError(f'Non-fail status found at output row {idx}: {actual["Status"]!r}')
    dev_cell = ws2.cell(row=idx, column=device_col)
    if dev_cell.data_type != 's':
        raise AssertionError(f'Device ID is not text at row {idx}: data_type={dev_cell.data_type!r}')
    if str(actual['Device ID']) != expected['Device ID']:
        raise AssertionError(f'Device ID leading zeros not preserved at row {idx}: {actual["Device ID"]!r} != {expected["Device ID"]!r}')

# Check ordering by Device ID sequence against filtered input.
actual_device_ids = [str(row['Device ID']) for row in read_rows]
expected_device_ids = [row['Device ID'] for row in failed_rows]
if actual_device_ids != expected_device_ids:
    raise AssertionError(f'Order mismatch: {actual_device_ids!r} != {expected_device_ids!r}')

print('verification passed:')
print('- workbook exists:', output_path.exists())
print('- worksheet name: failed_items')
print('- columns preserved:', read_headers)
print('- output rows:', len(read_rows))
print('- all Status values trim/casefold to fail')
print('- Device ID cells are text and preserve values:', actual_device_ids)
print('- original filtered order preserved')