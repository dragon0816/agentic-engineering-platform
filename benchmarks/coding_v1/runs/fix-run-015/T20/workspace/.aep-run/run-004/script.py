import csv
import json
import os
from pathlib import Path
from openpyxl import Workbook, load_workbook

params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = json.load(open(params_path, encoding='utf-8')) if params_path else {}
threshold = params.get('threshold')
if threshold is None:
    raise SystemExit('Missing required runtime parameter: threshold')
try:
    threshold = float(threshold)
except Exception as exc:
    raise SystemExit(f'threshold must be numeric, got {threshold!r}') from exc

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

sop_path = input_dir / 'supplied_sop.md'
if not sop_path.exists():
    raise FileNotFoundError(f'SOP not found: {sop_path}')
sop_text = sop_path.read_text(encoding='utf-8')
print('Read SOP:', sop_path.relative_to(workspace), 'chars=', len(sop_text))

preferred = [input_dir / 'results.csv', input_dir / 'data' / 'results.csv', input_dir / 'data' / 'test_results.csv']
candidates = [p for p in preferred if p.exists()]
if not candidates:
    candidates = sorted(input_dir.rglob('*.csv'))
if not candidates:
    raise FileNotFoundError('No CSV file found under input/')
if len(candidates) > 1:
    # Prefer the explicit observed file if present, otherwise fail to avoid guessing between multiple CSVs.
    observed = input_dir / 'data' / 'test_results.csv'
    if observed in candidates:
        csv_path = observed
    else:
        raise SystemExit('Multiple CSV files found and no unambiguous results CSV: ' + ', '.join(str(p.relative_to(workspace)) for p in candidates))
else:
    csv_path = candidates[0]
print('Read CSV:', csv_path.relative_to(workspace))

required_cols = ['Device ID', 'Status', 'Value']
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    print('CSV headers:', headers)
    missing = [c for c in required_cols if c not in headers]
    if missing:
        raise SystemExit('CSV missing required columns: ' + ', '.join(missing))
    rows = list(reader)

fail_rows = []
total = len(rows)
pass_count = 0
fail_count = 0
urgent_count = 0

for row in rows:
    status = (row.get('Status') or '').strip()
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        raw_value = row.get('Value')
        try:
            value_num = float(raw_value)
        except Exception as exc:
            raise SystemExit(f"Value is not numeric for Device ID {row.get('Device ID')!r}: {raw_value!r}") from exc
        urgent = value_num > threshold
        if urgent:
            urgent_count += 1
        fail_rows.append({
            'Device ID': row.get('Device ID', ''),
            'Status': row.get('Status', ''),
            'Value': value_num,
            'Urgent': bool(urgent),
        })

failures_path = output_dir / 'failures.xlsx'
summary_path = output_dir / 'summary.md'

wb = Workbook()
ws = wb.active
ws.title = 'Failures'
headers_out = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers_out)
for item in fail_rows:
    ws.append([item['Device ID'], item['Status'], item['Value'], item['Urgent']])
wb.save(failures_path)

summary = (
    '# Test Failure Summary\n\n'
    f'- Total: {total}\n'
    f'- Pass: {pass_count}\n'
    f'- Fail: {fail_count}\n'
    f'- Urgent: {urgent_count}\n'
)
summary_path.write_text(summary, encoding='utf-8')

# Reopen and verify every stated output requirement.
check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['Failures']
actual_headers = [cell.value for cell in check_ws[1]]
assert actual_headers == headers_out, f'Unexpected failures.xlsx headers: {actual_headers}'
actual_data = list(check_ws.iter_rows(min_row=2, values_only=True))
assert len(actual_data) == fail_count == len(fail_rows), f'Fail row count mismatch: workbook={len(actual_data)} expected={fail_count}'
for row_values, source in zip(actual_data, fail_rows):
    assert isinstance(row_values[0], str), 'Device ID must remain text'
    assert row_values[1] == 'Fail', 'All failures.xlsx rows must be Fail rows'
    assert isinstance(row_values[2], (int, float)), 'Value must be numeric in failures.xlsx'
    assert isinstance(row_values[3], bool), 'Urgent must be boolean in failures.xlsx'
    assert row_values[3] == (float(row_values[2]) > threshold), 'Urgent flag does not match Value > threshold'

summary_read = summary_path.read_text(encoding='utf-8')
for expected in [f'- Total: {total}', f'- Pass: {pass_count}', f'- Fail: {fail_count}', f'- Urgent: {urgent_count}']:
    assert expected in summary_read, f'Missing summary line: {expected}'

print('Wrote:', failures_path.relative_to(workspace))
print('Wrote:', summary_path.relative_to(workspace))
print('Checked failures.xlsx columns, row count, text Device ID, numeric Value, boolean Urgent, and summary counts.')
print(json.dumps({'total': total, 'pass': pass_count, 'fail': fail_count, 'urgent': urgent_count, 'threshold': threshold}, ensure_ascii=False))
