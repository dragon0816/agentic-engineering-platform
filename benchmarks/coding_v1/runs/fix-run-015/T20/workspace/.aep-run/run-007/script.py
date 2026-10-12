from pathlib import Path
from openpyxl import load_workbook
import csv, json, re

root = Path.cwd()
input_dir = root / 'input'
output_dir = root / 'output'
sop_path = input_dir / 'supplied_sop.md'
csv_path = input_dir / 'data' / 'test_results.csv'
failures_path = output_dir / 'failures.xlsx'
summary_path = output_dir / 'summary.md'

print('Workspace:', root)
print('SOP path:', sop_path.relative_to(root), 'bytes=', sop_path.stat().st_size)
print('SOP contents:\n' + sop_path.read_text(encoding='utf-8'))
print('CSV path:', csv_path.relative_to(root), 'bytes=', csv_path.stat().st_size)
print('Output files:')
for p in sorted(output_dir.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(root), p.stat().st_size, 'bytes')

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
print('CSV headers:', list(rows[0].keys()) if rows else [])
print('CSV sample rows:', rows[:5])

fail_rows = [r for r in rows if str(r.get('Status','')).strip() == 'Fail']
pass_count = sum(1 for r in rows if str(r.get('Status','')).strip() == 'Pass')
urgent_count = sum(float(r['Value']) > 20 for r in fail_rows)

assert failures_path.exists() and failures_path.stat().st_size > 0, 'failures.xlsx missing or empty'
wb = load_workbook(failures_path, data_only=True)
ws = wb.active
headers = [cell.value for cell in ws[1]]
data = list(ws.iter_rows(min_row=2, values_only=True))
print('Workbook title:', ws.title)
print('Workbook headers:', headers)
print('Workbook rows:', data)

assert headers == ['Device ID', 'Status', 'Value', 'Urgent'], f'Unexpected workbook headers: {headers}'
assert len(data) == len(fail_rows), f'Expected all Fail rows ({len(fail_rows)}), got {len(data)}'
for row in data:
    device_id, status, value, urgent = row
    assert isinstance(device_id, str), f'Device ID not preserved as text: {row!r}'
    assert status == 'Fail', f'Non-Fail row included: {row!r}'
    assert isinstance(value, (int, float)), f'Value not numeric: {row!r}'
    assert urgent in (True, False), f'Urgent not boolean: {row!r}'
    assert urgent == (float(value) > 20), f'Urgent should be Value > 20: {row!r}'

assert summary_path.exists() and summary_path.stat().st_size > 0, 'summary.md missing or empty'
summary = summary_path.read_text(encoding='utf-8')
print('Summary.md contents:\n' + summary)

def find_count(label):
    m = re.search(rf'^-\s*{re.escape(label)}\s*:\s*(\d+)\s*$', summary, re.MULTILINE)
    assert m, f'Missing summary line for {label}'
    return int(m.group(1))

assert find_count('Total') == len(rows), 'Summary Total mismatch'
assert find_count('Pass') == pass_count, 'Summary Pass mismatch'
assert find_count('Fail') == len(fail_rows), 'Summary Fail mismatch'
assert find_count('Urgent') == urgent_count, 'Summary Urgent mismatch'

print('CHECK PASSED:', json.dumps({
    'threshold_rule': 'Value > 20',
    'total': len(rows),
    'pass': pass_count,
    'fail': len(fail_rows),
    'urgent': urgent_count,
    'outputs': {
        'failures.xlsx': failures_path.stat().st_size,
        'summary.md': summary_path.stat().st_size
    }
}, ensure_ascii=False))
