from pathlib import Path
from openpyxl import load_workbook
import csv, json, os

root = Path.cwd()
input_dir = root / 'input'
output_dir = root / 'output'
print('Workspace:', root)
print('Input files:')
for p in sorted(input_dir.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(root), p.stat().st_size, 'bytes')
print('Output files:')
for p in sorted(output_dir.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(root), p.stat().st_size, 'bytes')

csv_path = input_dir / 'data' / 'test_results.csv'
failures_path = output_dir / 'failures.xlsx'
summary_path = output_dir / 'summary.md'

assert csv_path.exists(), f'Missing input CSV: {csv_path}'
assert failures_path.exists() and failures_path.stat().st_size > 0, 'failures.xlsx missing or empty'
assert summary_path.exists() and summary_path.stat().st_size > 0, 'summary.md missing or empty'

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
print('CSV headers:', list(rows[0].keys()) if rows else [])
fail_rows = [r for r in rows if str(r.get('Status','')).strip().lower() == 'fail']
expected_fail_count = len(fail_rows)
expected_urgent_count = sum(float(r['Value']) > 20 for r in fail_rows)

wb = load_workbook(failures_path, data_only=False)
ws = wb.active
headers = [cell.value for cell in ws[1]]
data = list(ws.iter_rows(min_row=2, values_only=True))
print('Workbook sheet:', ws.title)
print('Workbook headers:', headers)
print('Workbook rows:', len(data))
print('Workbook data:', data)

assert headers == ['Device ID', 'Status', 'Value', 'Urgent'], f'Unexpected headers: {headers}'
assert len(data) == expected_fail_count, f'Expected {expected_fail_count} fail rows, got {len(data)}'
for row in data:
    device_id, status, value, urgent = row
    assert isinstance(device_id, str), f'Device ID is not text: {device_id!r}'
    assert status == 'Fail', f'Non-Fail row included: {row!r}'
    assert isinstance(value, (int, float)), f'Value is not numeric: {row!r}'
    assert isinstance(urgent, bool), f'Urgent is not boolean: {row!r}'
    assert urgent == (float(value) > 20), f'Urgent mismatch for row: {row!r}'

summary = summary_path.read_text(encoding='utf-8')
print('Summary.md contents:\n' + summary)
required_fragments = [
    f'Total records: {len(rows)}',
    f'Pass records: {sum(str(r.get("Status","")).strip().lower() == "pass" for r in rows)}',
    f'Fail records: {expected_fail_count}',
    f'Urgent failures: {expected_urgent_count}',
]
for frag in required_fragments:
    assert frag in summary, f'Missing summary fragment: {frag}'

print('CHECK PASSED:', json.dumps({
    'total': len(rows),
    'failures_rows': len(data),
    'urgent_failures': expected_urgent_count,
    'outputs': [str(failures_path.relative_to(root)), str(summary_path.relative_to(root))]
}, ensure_ascii=False))
