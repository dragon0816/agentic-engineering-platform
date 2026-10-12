from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'specs.md'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split into product sections headed by level-2 markdown headings.
sections = []
for m in re.finditer(r'(?m)^##\s+(.+?)\s*$', text):
    start = m.end()
    next_m = re.search(r'(?m)^##\s+.+?\s*$', text[start:])
    end = start + next_m.start() if next_m else len(text)
    sections.append((m.group(1).strip(), text[start:end]))

unit_to_mhz = {
    'mhz': 1.0,
    'ghz': 1000.0,
    'khz': 0.001,
    'hz': 0.000001,
}

def clean_number(s):
    return float(s.replace(',', '').strip())

def maybe_int(x):
    if x is None:
        return None
    return int(x) if abs(x - int(x)) < 1e-9 else x

def parse_frequency(value):
    if not value:
        return None, None
    v = value.strip().replace('—', '–').replace('-', '–')
    # Unit may be written once after the range: 2.4–2.5 GHz
    m = re.search(r'([0-9][0-9,]*(?:\.\d+)?)\s*–\s*([0-9][0-9,]*(?:\.\d+)?)\s*([GMK]?Hz)', v, re.I)
    if m:
        unit = m.group(3).lower()
        factor = unit_to_mhz[unit]
        return maybe_int(clean_number(m.group(1)) * factor), maybe_int(clean_number(m.group(2)) * factor)
    # Or unit on both endpoints.
    m = re.search(r'([0-9][0-9,]*(?:\.\d+)?)\s*([GMK]?Hz)\s*–\s*([0-9][0-9,]*(?:\.\d+)?)\s*([GMK]?Hz)', v, re.I)
    if m:
        f1 = unit_to_mhz[m.group(2).lower()]
        f2 = unit_to_mhz[m.group(4).lower()]
        return maybe_int(clean_number(m.group(1)) * f1), maybe_int(clean_number(m.group(3)) * f2)
    return None, None

def parse_bandwidth(value):
    if not value:
        return None
    m = re.search(r'([0-9][0-9,]*(?:\.\d+)?)\s*([GMK]?Hz)', value, re.I)
    if not m:
        return None
    return maybe_int(clean_number(m.group(1)) * unit_to_mhz[m.group(2).lower()])

def parse_interface(value):
    if not value:
        return None
    # The source uses slash to list multiple interfaces. Preserve names, separate multiple with semicolon.
    parts = [p.strip() for p in re.split(r'\s*(?:/|,|;|、)\s*', value.strip()) if p.strip()]
    return ';'.join(parts) if parts else None

def field_value(section_text, field):
    m = re.search(rf'(?mi)^\s*-\s*{re.escape(field)}\s*:\s*(.*?)\s*$', section_text)
    return m.group(1).strip() if m else None

rows = []
for heading, body in sections:
    model = field_value(body, 'Model')
    freq = field_value(body, 'Frequency')
    bw = field_value(body, 'Bandwidth')
    iface = field_value(body, 'Interface')
    fmin, fmax = parse_frequency(freq)
    rows.append({
        'Model': model,
        'FrequencyMinMHz': fmin,
        'FrequencyMaxMHz': fmax,
        'Bandwidth': parse_bandwidth(bw),
        'Interface': parse_interface(iface),
    })

wb = Workbook()
ws = wb.active
ws.title = 'specs'
headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'Bandwidth', 'Interface']
ws.append(headers)
for row in rows:
    ws.append([row[h] if row[h] is not None else '' for h in headers])
wb.save(output_path)

# Reopen and verify output file, columns, row count, value types, conversions, blanks, and interface separator.
wb2 = load_workbook(output_path, data_only=True)
ws2 = wb2['specs']
actual_headers = [cell.value for cell in ws2[1]]
assert actual_headers == headers, actual_headers
actual_rows = list(ws2.iter_rows(min_row=2, values_only=True))
assert len(actual_rows) == 3, actual_rows
expected = [
    ('RX-100', 2400, 2500, 80, 'USB;LAN'),
    ('RX-200', 900, 930, 20, 'LAN'),
    ('RX-300', 5000, 6000, None, 'USB'),
]
assert actual_rows == expected, actual_rows
for r in actual_rows:
    assert isinstance(r[0], str)
    assert isinstance(r[1], (int, float)) and isinstance(r[2], (int, float))
    if r[3] is not None:
        assert isinstance(r[3], (int, float))
    assert isinstance(r[4], str)

print('wrote:', output_path)
print('headers:', actual_headers)
print('rows:')
for row in actual_rows:
    print(row)
print('checks: file reopened; sheet specs exists; columns match; 3 data rows; frequency ranges converted to MHz; bandwidth converted to MHz with missing value blank; interfaces separated by semicolon; input file unchanged by script.')