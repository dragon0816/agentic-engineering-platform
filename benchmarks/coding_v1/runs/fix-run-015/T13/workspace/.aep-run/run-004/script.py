import re, json
from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urljoin
import urllib.request
import urllib.error
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:50447/t13/list'
XLSX_PATH = OUTPUT / 'robust_products.xlsx'
WARN_PATH = OUTPUT / 'warnings.json'
HEADERS = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__(); self.links=[]; self._href=None; self._text=[]
    def handle_starttag(self, tag, attrs):
        if tag.lower()=='a':
            self._href=dict(attrs).get('href'); self._text=[]
    def handle_data(self, data):
        if self._href is not None: self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower()=='a' and self._href is not None:
            self.links.append((''.join(self._text).strip(), self._href)); self._href=None; self._text=[]

class DetailParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_h1=False; self.h1=[]
        self.current_cell=None; self.cell_text=[]; self.row=[]; self.rows=[]
        self.text_parts=[]
    def handle_starttag(self, tag, attrs):
        tag=tag.lower()
        if tag=='h1': self.in_h1=True
        if tag in ('th','td'):
            self.current_cell=tag; self.cell_text=[]
        if tag=='tr':
            self.row=[]
    def handle_data(self, data):
        s=data.strip()
        if s: self.text_parts.append(s)
        if self.in_h1: self.h1.append(data)
        if self.current_cell is not None: self.cell_text.append(data)
    def handle_endtag(self, tag):
        tag=tag.lower()
        if tag=='h1': self.in_h1=False
        if tag in ('th','td') and self.current_cell is not None:
            self.row.append(''.join(self.cell_text).strip())
            self.current_cell=None; self.cell_text=[]
        if tag=='tr':
            if self.row: self.rows.append(self.row)
            self.row=[]
    def table_dict(self):
        d={}
        for row in self.rows:
            if len(row)>=2 and row[0]:
                d[row[0].strip().lower()] = row[1].strip()
        return d
    def text(self):
        return '\n'.join(self.text_parts)

def fetch_once(url):
    req=urllib.request.Request(url, headers={'User-Agent':'aep-product-fetcher/1.0'})
    with urllib.request.urlopen(req, timeout=10) as resp:
        return resp.status, resp.read().decode('utf-8', errors='replace')

def fetch_detail(url, max_500_attempts=3):
    attempts=0; last_error=None
    while attempts < max_500_attempts:
        attempts += 1
        try:
            status, body = fetch_once(url)
            return {'ok': True, 'status': status, 'body': body, 'attempts': attempts, 'error': None}
        except urllib.error.HTTPError as e:
            last_error=f'HTTP {e.code}'
            if e.code == 500 and attempts < max_500_attempts:
                continue
            return {'ok': False, 'status': e.code, 'body': '', 'attempts': attempts, 'error': last_error}
        except Exception as e:
            last_error=f'{type(e).__name__}: {e}'
            return {'ok': False, 'status': None, 'body': '', 'attempts': attempts, 'error': last_error}
    return {'ok': False, 'status': 500, 'body': '', 'attempts': attempts, 'error': last_error or 'HTTP 500'}

def clean(v):
    if v is None: return 'N/A'
    v=re.sub(r'\s+', ' ', str(v)).strip()
    return v if v else 'N/A'

def parse_price_currency(raw):
    raw=clean(raw)
    if raw=='N/A': return 'N/A','N/A'
    currency='N/A'
    # Keep common currency strings/symbols; NT$ maps naturally as currency text from page.
    cm=re.search(r'(NT\$|US\$|HK\$|USD|TWD|NTD|EUR|JPY|CNY|GBP|\$|€|¥)', raw, re.I)
    if cm:
        currency=cm.group(1)
    nums=re.findall(r'\d[\d,]*(?:\.\d+)?', raw)
    price=nums[0].replace(',', '') if nums else 'N/A'
    return price, currency

def parse_detail(html, fallback_name):
    p=DetailParser(); p.feed(html)
    table=p.table_dict()
    name=clean(''.join(p.h1)) if ''.join(p.h1).strip() else clean(fallback_name)
    raw_price=table.get('price')
    price,currency=parse_price_currency(raw_price)
    cpu=clean(table.get('cpu'))
    return {'Name': name, 'Price': price, 'Currency': currency, 'CPU': cpu}

# List
status, list_html=fetch_once(LIST_URL)
lp=LinkParser(); lp.feed(list_html)
products=lp.links[:5]
print('List status:', status)
print('Products found:', len(lp.links), 'using:', len(products))
for i,(name,href) in enumerate(products,1): print(f'{i}. {name} -> {href}')

rows=[]; warnings=[]
for list_name, href in products:
    detail_url=urljoin(LIST_URL, href)
    result=fetch_detail(detail_url, 3)
    warn={'name': list_name, 'url': detail_url, 'status': 'ok' if result['ok'] else 'error', 'http_status': result['status'], 'attempts': result['attempts']}
    if result['ok']:
        parsed=parse_detail(result['body'], list_name)
        row={**parsed, 'URL': detail_url, 'FetchStatus': 'ok'}
    else:
        warn['error']=result['error']
        row={'Name': clean(list_name), 'Price':'N/A', 'Currency':'N/A', 'CPU':'N/A', 'URL': detail_url, 'FetchStatus':'error'}
    rows.append(row); warnings.append(warn)

wb=Workbook(); ws=wb.active; ws.title='Products'; ws.append(HEADERS)
for row in rows: ws.append([row[h] for h in HEADERS])
wb.save(XLSX_PATH)
with open(WARN_PATH,'w',encoding='utf-8') as f: json.dump(warnings,f,ensure_ascii=False,indent=2)

print('Wrote:', XLSX_PATH)
print('Wrote:', WARN_PATH)
for r in rows: print(r)
print(json.dumps(warnings, ensure_ascii=False, indent=2))

# Verification from saved files
print('\nVerification:')
wb2=load_workbook(XLSX_PATH, data_only=True)
ws2=wb2.active
headers=[c.value for c in ws2[1]]
saved=[dict(zip(headers,row)) for row in ws2.iter_rows(min_row=2, values_only=True)]
print('xlsx headers:', headers)
print('xlsx row count:', len(saved))
print('saved rows:', saved)
assert headers == HEADERS
assert len(saved) == 5
for idx,row in enumerate(saved,2):
    for h in HEADERS:
        assert row.get(h) is not None and row.get(h) != '', f'blank {h} row {idx}'
    assert row['FetchStatus'] in ('ok','error')
    assert isinstance(row['URL'], str) and row['URL'].startswith('http://127.0.0.1:50447/t13/item/')
    if row['FetchStatus']=='error':
        assert row['Price']=='N/A' and row['Currency']=='N/A' and row['CPU']=='N/A'
    else:
        # Price may be N/A when missing; numeric prices should be stored as int or numeric-like strings by openpyxl.
        if row['Price'] != 'N/A':
            assert str(row['Price']).replace('.0','').isdigit(), f'price not numeric-like row {idx}: {row["Price"]}'
with open(WARN_PATH,encoding='utf-8') as f: w=json.load(f)
print('warnings count:', len(w))
assert isinstance(w, list) and len(w)==5
for item in w:
    assert item['status'] in ('ok','error')
    if item['status']=='ok': assert item['attempts']==1
    if item['status']=='error' and item.get('http_status')==500: assert item['attempts']==3
print('verification passed')