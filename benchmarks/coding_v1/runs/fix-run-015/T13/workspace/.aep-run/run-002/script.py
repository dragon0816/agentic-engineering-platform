import os, re, json, time
from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urljoin
import urllib.request
import urllib.error
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:50447/t13/list'
XLSX_PATH = OUTPUT / 'robust_products.xlsx'
WARN_PATH = OUTPUT / 'warnings.json'
HEADERS = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links=[]
        self._href=None
        self._text=[]
    def handle_starttag(self, tag, attrs):
        if tag.lower()=='a':
            attrs=dict(attrs)
            self._href=attrs.get('href')
            self._text=[]
    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower()=='a' and self._href is not None:
            text=''.join(self._text).strip()
            self.links.append((text, self._href))
            self._href=None
            self._text=[]

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts=[]
    def handle_data(self, data):
        s=data.strip()
        if s:
            self.parts.append(s)
    def text(self):
        return '\n'.join(self.parts)

def fetch_once(url):
    req = urllib.request.Request(url, headers={'User-Agent':'aep-product-fetcher/1.0'})
    with urllib.request.urlopen(req, timeout=10) as resp:
        body = resp.read().decode('utf-8', errors='replace')
        return resp.status, body

def fetch_detail(url, max_500_attempts=3):
    attempts = 0
    last_error = None
    while attempts < max_500_attempts:
        attempts += 1
        try:
            status, body = fetch_once(url)
            return {'ok': True, 'status': status, 'body': body, 'attempts': attempts, 'error': None}
        except urllib.error.HTTPError as e:
            last_error = f'HTTP {e.code}'
            if e.code == 500 and attempts < max_500_attempts:
                continue
            return {'ok': False, 'status': e.code, 'body': '', 'attempts': attempts, 'error': last_error}
        except Exception as e:
            last_error = f'{type(e).__name__}: {e}'
            return {'ok': False, 'status': None, 'body': '', 'attempts': attempts, 'error': last_error}
    return {'ok': False, 'status': 500, 'body': '', 'attempts': attempts, 'error': last_error or 'HTTP 500'}

def regex_first(patterns, text, flags=re.I|re.S):
    for pat in patterns:
        m=re.search(pat, text, flags)
        if m:
            return m.group(1).strip()
    return None

def clean_value(v):
    if v is None:
        return 'N/A'
    v = re.sub(r'\s+', ' ', str(v)).strip()
    return v if v else 'N/A'

def strip_tags(html):
    p=TextParser(); p.feed(html); return p.text()

def parse_detail(html, fallback_name):
    text=strip_tags(html)
    # Name: prefer visible h1/title-ish labels, then fallback list name
    name=regex_first([
        r'Name\s*[:\-]\s*([^\n]+)',
        r'Product\s*[:\-]\s*([^\n]+)',
        r'^\s*([^\n]+)'
    ], text) or fallback_name
    # If first line is a generic word, fallback to list name only when no explicit name found.
    explicit_name=regex_first([r'Name\s*[:\-]\s*([^\n]+)', r'Product\s*[:\-]\s*([^\n]+)'], text)
    if explicit_name:
        name=explicit_name
    elif fallback_name:
        # Keep the list name unless page clearly starts with a non-label product name not containing price/cpu.
        first_line = text.split('\n')[0].strip() if text else ''
        if first_line and not re.search(r'price|currency|cpu|spec', first_line, re.I):
            name=first_line
        else:
            name=fallback_name

    price=regex_first([
        r'Price\s*[:\-]\s*([0-9]+(?:[.,][0-9]+)?)',
        r'([0-9]+(?:[.,][0-9]+)?)\s*(?:USD|EUR|TWD|NTD|JPY|CNY|GBP|\$|€|¥)'
    ], text)
    currency=regex_first([
        r'Currency\s*[:\-]\s*([A-Z]{3}|\$|€|¥)',
        r'(USD|EUR|TWD|NTD|JPY|CNY|GBP)'
    ], text)
    # If currency symbol appears before price
    if not currency:
        symbol = regex_first([r'(\$|€|¥)\s*[0-9]'], text)
        currency = symbol
    cpu=regex_first([
        r'CPU\s*[:\-]\s*([^\n]+)',
        r'Processor\s*[:\-]\s*([^\n]+)',
        r'Chip(?:set)?\s*[:\-]\s*([^\n]+)'
    ], text)
    return {
        'Name': clean_value(name),
        'Price': clean_value(price),
        'Currency': clean_value(currency),
        'CPU': clean_value(cpu),
    }

# Fetch list
status, list_html = fetch_once(LIST_URL)
parser=LinkParser(); parser.feed(list_html)
products = parser.links[:5]
print('List status:', status)
print('Products found:', len(parser.links), 'using:', len(products))
for i,(name,href) in enumerate(products,1):
    print(f'  {i}. {name} -> {href}')

rows=[]
warnings=[]
for list_name, href in products:
    detail_url=urljoin(LIST_URL, href)
    result=fetch_detail(detail_url, max_500_attempts=3)
    warn={'name': list_name, 'url': detail_url, 'status': 'ok' if result['ok'] else 'error', 'http_status': result['status'], 'attempts': result['attempts']}
    if result['ok']:
        parsed=parse_detail(result['body'], list_name)
        row={**parsed, 'URL': detail_url, 'FetchStatus': 'ok'}
    else:
        warn['error']=result['error']
        row={'Name': clean_value(list_name), 'Price':'N/A', 'Currency':'N/A', 'CPU':'N/A', 'URL':detail_url, 'FetchStatus':'error'}
    rows.append(row)
    warnings.append(warn)

# Write workbook
wb=Workbook()
ws=wb.active
ws.title='Products'
ws.append(HEADERS)
for row in rows:
    ws.append([row[h] for h in HEADERS])
wb.save(XLSX_PATH)

with open(WARN_PATH, 'w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

print('Wrote:', XLSX_PATH)
print('Wrote:', WARN_PATH)
print('Rows:')
for r in rows:
    print(r)
print('Warnings:')
print(json.dumps(warnings, ensure_ascii=False, indent=2))

# Verification: reopen output files and check requirements
print('\nVerification:')
wb2=load_workbook(XLSX_PATH, data_only=True)
ws2=wb2.active
headers=[c.value for c in ws2[1]]
data=[]
for row in ws2.iter_rows(min_row=2, values_only=True):
    data.append(dict(zip(headers,row)))
print('xlsx headers:', headers)
print('xlsx row count:', len(data))
assert headers == HEADERS, f'headers mismatch: {headers}'
assert len(data) == 5, f'expected 5 rows got {len(data)}'
for idx,row in enumerate(data,2):
    for h in HEADERS:
        assert h in row, f'missing {h}'
        assert row[h] is not None and row[h] != '', f'blank {h} at row {idx}'
    assert row['FetchStatus'] in ('ok','error'), f'bad FetchStatus at row {idx}'
    if row['FetchStatus']=='error':
        assert row['Price']=='N/A' and row['Currency']=='N/A' and row['CPU']=='N/A', f'error row unknown fields not N/A at row {idx}'
    else:
        # Missing price/spec may be N/A by design; just ensure status ok.
        pass
with open(WARN_PATH, encoding='utf-8') as f:
    w=json.load(f)
print('warnings count:', len(w))
assert isinstance(w, list), 'warnings.json must be a list'
assert len(w)==5, 'warnings.json should include one entry per product'
for item in w:
    assert item.get('status') in ('ok','error'), 'warning status must be ok/error'
    if item.get('status')=='error' and item.get('http_status')==500:
        assert item.get('attempts')==3, 'HTTP 500 must be attempted 3 times'
print('verification passed')