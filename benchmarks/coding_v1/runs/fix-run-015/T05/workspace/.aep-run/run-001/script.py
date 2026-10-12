from pathlib import Path
from openpyxl import load_workbook

ws = Path.cwd()
print('Workspace:', ws)
input_dir = ws / 'input'
output_dir = ws / 'output'
print('Input files:', [p.name for p in sorted(input_dir.glob('*'))])
for filename in ['devices.xlsx', 'tests.xlsx']:
    path = input_dir / filename
    print('\nFILE', filename, 'exists=', path.exists())
    wb = load_workbook(path, data_only=False, read_only=True)
    print('Sheets:', wb.sheetnames)
    sh = wb[wb.sheetnames[0]]
    rows = list(sh.iter_rows(values_only=True, max_row=min(sh.max_row, 6)))
    print('Dimensions:', sh.max_row, 'rows x', sh.max_column, 'cols')
    for i, row in enumerate(rows, start=1):
        print(f'Row {i}:', row)
    wb.close()