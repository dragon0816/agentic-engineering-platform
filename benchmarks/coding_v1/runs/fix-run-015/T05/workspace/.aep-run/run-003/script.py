from pathlib import Path
from openpyxl import load_workbook

out_path = Path.cwd() / 'output' / 'merged_devices_tests.xlsx'
wb = load_workbook(out_path, data_only=True)
assert set(wb.sheetnames) == {'matched', 'unmatched'}, wb.sheetnames

m = wb['matched']
u = wb['unmatched']
matched_headers = [c.value for c in m[1]]
unmatched_headers = [c.value for c in u[1]]
assert matched_headers == ['Device ID', 'Test', 'Result', 'Name', 'Location'], matched_headers
assert unmatched_headers == ['Source', 'Device ID', 'Name', 'Location', 'Test', 'Result'], unmatched_headers

matched = [[c.value for c in row] for row in m.iter_rows(min_row=2)]
unmatched = [[c.value for c in row] for row in u.iter_rows(min_row=2)]
assert len(matched) == 3, matched
assert len(unmatched) == 2, unmatched
assert matched == [
    ['001', 'RF-1', 'Pass', 'Alpha', 'Lab A'],
    ['001', 'RF-2', 'Fail', 'Alpha', 'Lab A'],
    ['002', 'RF-1', 'Pass', 'Beta', 'Lab B'],
], matched
assert unmatched == [
    ['devices', '003', 'Gamma', 'Lab C', None, None],
    ['tests', '999', None, None, 'RF-9', 'Fail'],
], unmatched

for ws in [m, u]:
    headers = [c.value for c in ws[1]]
    did_col = headers.index('Device ID') + 1
    for row in range(2, ws.max_row + 1):
        v = ws.cell(row=row, column=did_col).value
        assert isinstance(v, str), (ws.title, row, v, type(v))
        assert ws.cell(row=row, column=did_col).number_format == '@'

wb.close()
print('Checked output file exists and reopens successfully.')
print('Checked sheet names: matched, unmatched.')
print('Checked matched headers, unmatched headers including Source.')
print('Checked row counts: matched=3, unmatched=2.')
print('Checked Device ID values are strings with leading zeroes preserved and text format applied.')
print('Checked matched and unmatched row contents.')