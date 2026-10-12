from pathlib import Path
from copy import copy
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
out_path = OUTPUT / 'merged_devices_tests.xlsx'

def read_sheet(path):
    wb = load_workbook(path, data_only=False, read_only=False)
    ws = wb[wb.sheetnames[0]]
    rows = []
    for row in ws.iter_rows():
        rows.append([cell.value for cell in row])
    wb.close()
    if not rows:
        return [], []
    headers = [str(v) if v is not None else '' for v in rows[0]]
    records = []
    for vals in rows[1:]:
        if all(v is None for v in vals):
            continue
        rec = {}
        for h, v in zip(headers, vals):
            if h == 'Device ID' and v is not None:
                rec[h] = str(v)
            else:
                rec[h] = v
        # pad if short rows were returned
        for h in headers[len(vals):]:
            rec[h] = None
        records.append(rec)
    return headers, records

dev_headers, devices = read_sheet(INPUT / 'devices.xlsx')
test_headers, tests = read_sheet(INPUT / 'tests.xlsx')

if 'Device ID' not in dev_headers:
    raise ValueError('devices.xlsx missing Device ID column')
if 'Device ID' not in test_headers:
    raise ValueError('tests.xlsx missing Device ID column')

# Treat IDs as exact text values and preserve leading zeroes.
device_by_id = {}
duplicate_device_ids = []
for rec in devices:
    did = rec.get('Device ID')
    if did in device_by_id:
        duplicate_device_ids.append(did)
    device_by_id[did] = rec
if duplicate_device_ids:
    raise ValueError(f'devices contains duplicate Device ID values: {duplicate_device_ids}')

test_ids = {rec.get('Device ID') for rec in tests}
matched_rows = []
matched_device_ids = set()
unmatched_test_rows = []
for test in tests:
    did = test.get('Device ID')
    dev = device_by_id.get(did)
    if dev is not None:
        matched_device_ids.add(did)
        row = {'Device ID': did}
        for h in test_headers:
            if h != 'Device ID':
                row[h] = test.get(h)
        for h in dev_headers:
            if h != 'Device ID':
                row[h] = dev.get(h)
        matched_rows.append(row)
    else:
        unmatched_test_rows.append(test)

unmatched_device_rows = [dev for dev in devices if dev.get('Device ID') not in test_ids]

matched_headers = ['Device ID'] + [h for h in test_headers if h != 'Device ID'] + [h for h in dev_headers if h != 'Device ID']
unmatched_headers = ['Source']
for h in dev_headers + test_headers:
    if h not in unmatched_headers:
        unmatched_headers.append(h)

wb = Workbook()
ws_m = wb.active
ws_m.title = 'matched'
ws_u = wb.create_sheet('unmatched')

def write_table(ws, headers, records):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
    for rec in records:
        ws.append([rec.get(h) for h in headers])
    # Force Device ID cells to text format.
    if 'Device ID' in headers:
        col_idx = headers.index('Device ID') + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            if cell.value is not None:
                cell.value = str(cell.value)
            cell.number_format = '@'
    for col_cells in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 10), 40)

write_table(ws_m, matched_headers, matched_rows)

unmatched_records = []
for rec in unmatched_device_rows:
    out = {'Source': 'devices'}
    out.update(rec)
    unmatched_records.append(out)
for rec in unmatched_test_rows:
    out = {'Source': 'tests'}
    out.update(rec)
    unmatched_records.append(out)
write_table(ws_u, unmatched_headers, unmatched_records)

wb.save(out_path)
print(f'Wrote {out_path}')
print(f'matched rows: {len(matched_rows)}')
print(f'unmatched rows: {len(unmatched_records)}')
print('matched headers:', matched_headers)
print('unmatched headers:', unmatched_headers)