from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
import re
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
URL = 'http://127.0.0.1:50447/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self, base_url):
        super().__init__(convert_charrefs=True)
        self.base_url = base_url
        self.in_catalog = False
        self.catalog_depth = 0
        self.current_card = None
        self.capture = None
        self.products = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'main' and attrs.get('id') == 'catalog':
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if self.in_catalog:
            self.catalog_depth += 1
            cls = attrs.get('class', '')
            classes = set(cls.split())
            if tag == 'div' and 'card' in classes:
                self.current_card = {'name': '', 'price_text': '', 'url': ''}
            elif self.current_card is not None and tag == 'a':
                self.current_card['url'] = urljoin(self.base_url, attrs.get('href', ''))
                self.capture = 'name'
            elif self.current_card is not None and tag == 'span' and 'price' in classes:
                self.capture = 'price_text'

    def handle_endtag(self, tag):
        if self.in_catalog:
            if tag == 'a' and self.capture == 'name':
                self.capture = None
            elif tag == 'span' and self.capture == 'price_text':
                self.capture = None
            elif tag == 'div' and self.current_card is not None:
                # End of product card in main catalog.
                if self.current_card.get('name') and self.current_card.get('price_text') and self.current_card.get('url'):
                    self.products.append(self.current_card)
                self.current_card = None
                self.capture = None
            self.catalog_depth -= 1
            if tag == 'main' or self.catalog_depth <= 0:
                self.in_catalog = False
                self.catalog_depth = 0
                self.capture = None
                self.current_card = None

    def handle_data(self, data):
        if self.in_catalog and self.current_card is not None and self.capture:
            self.current_card[self.capture] += data.strip()

def parse_price(text):
    m = re.match(r'\s*([A-Za-z$€£¥₩₫₱₹₽₺₪฿₴₦₡₲₵₸₼₾₿]+)\s*([\d,]+(?:\.\d+)?)\s*$', text)
    if not m:
        raise ValueError(f'Unrecognized price format: {text!r}')
    currency = m.group(1)
    number = float(m.group(2).replace(',', ''))
    if number.is_integer():
        number = int(number)
    return currency, number

with urlopen(URL, timeout=10) as r:
    html = r.read().decode('utf-8', errors='replace')

parser = CatalogParser(URL)
parser.feed(html)
products = []
for item in parser.products:
    currency, price = parse_price(item['price_text'])
    products.append({
        '產品名稱': item['name'],
        '價格': price,
        '幣別': currency,
        '產品網址': item['url'],
    })

print('Parsed products:', len(products))
for p in products:
    print(p)

wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for p in products:
    ws.append([p[h] for h in headers])
wb.save(OUTFILE)
print('Wrote', OUTFILE)

# Verification: reopen output and check all stated requirements.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
rows = list(ws2.iter_rows(values_only=True))
assert rows[0] == tuple(headers), f'Unexpected headers: {rows[0]!r}'
data = rows[1:]
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
expected_prices = [1299, 2450, 99, 560, 12000]
expected_urls = [urljoin(URL, path) for path in ['/p/router-a','/p/switch-b','/p/cable-c','/p/adapter-d','/p/antenna-e']]
assert len(data) == 5, f'Expected 5 product rows, got {len(data)}'
for idx, row in enumerate(data):
    name, price, currency, product_url = row
    assert name == expected_names[idx], f'Row {idx+2} name/order mismatch: {name!r}'
    assert isinstance(price, (int, float)) and not isinstance(price, bool), f'Row {idx+2} price is not numeric: {price!r}'
    assert price == expected_prices[idx], f'Row {idx+2} price mismatch: {price!r}'
    assert currency == 'NT$', f'Row {idx+2} currency mismatch: {currency!r}'
    assert product_url == expected_urls[idx], f'Row {idx+2} URL mismatch: {product_url!r}'
    assert '/ad/' not in product_url and 'Sponsored' not in str(name) and '推薦' not in str(name), f'Ad/recommendation included at row {idx+2}'
print('CHECK PASSED: products.xlsx reopened; headers, row count, page order, numeric prices, currencies, full URLs, and exclusion of ads/recommendations verified.')