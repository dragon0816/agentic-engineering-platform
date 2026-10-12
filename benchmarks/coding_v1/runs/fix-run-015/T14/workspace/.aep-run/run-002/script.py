from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import Workbook, load_workbook

ws = Path.cwd()
input_path = ws / 'input' / 'large_results.csv'
output_dir = ws / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

required_headers = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']

# First pass: read and aggregate
stats = defaultdict(lambda: {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0})
overall = {'Total': 0, 'PassCount': 0, 'FailCount': 0, 'DurationSum': 0.0}

with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('Input headers:', reader.fieldnames)
    missing = [h for h in required_headers if h not in (reader.fieldnames or [])]
    if missing:
        raise ValueError(f'Missing required headers: {missing}')
    
    for row_num, row in enumerate(reader, start=2):
        model = row['Device Model']
        status = row['Status']
        try:
            duration = float(row['DurationSeconds'])
        except Exception as e:
            raise ValueError(f'Invalid DurationSeconds at CSV row {row_num}: {row["DurationSeconds"]!r}') from e
        
        for bucket in (stats[model], overall):
            bucket['Total'] += 1
            if status == 'Pass':
                bucket['PassCount'] += 1
            elif status == 'Fail':
                bucket['FailCount'] += 1
            bucket['DurationSum'] += duration

# Create workbook
wb = Workbook()
ws_by_model = wb.active
ws_by_model.title = 'By Device Model'
headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
ws_by_model.append(headers)

for model in sorted(stats.keys()):
    s = stats[model]
    total = s['Total']
    pass_rate = s['PassCount'] / total if total else 0
    avg_duration = s['DurationSum'] / total if total else 0
    ws_by_model.append([model, total, s['PassCount'], s['FailCount'], pass_rate, avg_duration])

ws_overall = wb.create_sheet('Overall')
ws_overall.append(['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds'])
total = overall['Total']
ws_overall.append([
    total,
    overall['PassCount'],
    overall['FailCount'],
    overall['PassCount'] / total if total else 0,
    overall['DurationSum'] / total if total else 0,
])

# Formatting
for sheet in (ws_by_model, ws_overall):
    for cell in sheet[1]:
        cell.style = 'Headline 4'
    for col in sheet.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col)
        sheet.column_dimensions[col[0].column_letter].width = max(max_len + 2, 12)

# Save
wb.save(output_path)
print('Wrote:', output_path)

# Verification: reopen output and compare every aggregate to recomputed stats
check_wb = load_workbook(output_path, data_only=True)
expected_sheets = ['By Device Model', 'Overall']
actual_sheets = check_wb.sheetnames
print('Workbook sheets:', actual_sheets)
if actual_sheets != expected_sheets:
    raise AssertionError(f'Expected sheets {expected_sheets}, got {actual_sheets}')

sh = check_wb['By Device Model']
actual_headers = [c.value for c in sh[1]]
if actual_headers != headers:
    raise AssertionError(f'By Device Model headers mismatch: {actual_headers}')
expected_model_rows = len(stats)
actual_model_rows = sh.max_row - 1
print('By Device Model rows:', actual_model_rows)
if actual_model_rows != expected_model_rows:
    raise AssertionError(f'Expected {expected_model_rows} model rows, got {actual_model_rows}')

for row in sh.iter_rows(min_row=2, values_only=True):
    model, total_v, pass_v, fail_v, pass_rate_v, avg_v = row
    if model not in stats:
        raise AssertionError(f'Unexpected model in output: {model}')
    s = stats[model]
    exp = [s['Total'], s['PassCount'], s['FailCount'], s['PassCount']/s['Total'], s['DurationSum']/s['Total']]
    got = [total_v, pass_v, fail_v, pass_rate_v, avg_v]
    if got[:3] != exp[:3]:
        raise AssertionError(f'Counts mismatch for {model}: got {got[:3]}, expected {exp[:3]}')
    if abs(got[3] - exp[3]) > 1e-12 or abs(got[4] - exp[4]) > 1e-12:
        raise AssertionError(f'Rate/avg mismatch for {model}: got {got[3:]}, expected {exp[3:]}')

osh = check_wb['Overall']
overall_headers = [c.value for c in osh[1]]
if overall_headers != ['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']:
    raise AssertionError(f'Overall headers mismatch: {overall_headers}')
if osh.max_row != 2:
    raise AssertionError(f'Overall should have exactly one data row, got {osh.max_row-1}')
orow = [c.value for c in osh[2]]
expected_overall = [
    overall['Total'],
    overall['PassCount'],
    overall['FailCount'],
    overall['PassCount']/overall['Total'],
    overall['DurationSum']/overall['Total'],
]
print('Overall row:', orow)
print('Expected overall:', expected_overall)
if orow[:3] != expected_overall[:3]:
    raise AssertionError(f'Overall counts mismatch: got {orow[:3]}, expected {expected_overall[:3]}')
if abs(orow[3] - expected_overall[3]) > 1e-12 or abs(orow[4] - expected_overall[4]) > 1e-12:
    raise AssertionError(f'Overall rate/avg mismatch: got {orow[3:]}, expected {expected_overall[3:]}')

# Type checks
for row in sh.iter_rows(min_row=2, values_only=True):
    if not isinstance(row[0], str):
        raise AssertionError('Device Model must be text')
    for value in row[1:4]:
        if not isinstance(value, int):
            raise AssertionError(f'Count value must be int, got {value!r} ({type(value)})')
    for value in row[4:6]:
        if not isinstance(value, (int, float)):
            raise AssertionError(f'Numeric metric must be numeric, got {value!r} ({type(value)})')

print('Verification passed: output file, sheets, column names, row counts, counts, pass rates, average durations, and value types checked.')
print('Output file:', output_path)
