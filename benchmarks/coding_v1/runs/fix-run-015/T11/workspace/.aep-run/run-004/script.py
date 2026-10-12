from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
import re
import openpyxl

BASE_LIST_URL = 'http://127.0.0.1:50447/t11/list'
workspace = Path.cwd()
out_dir = workspace / 'output'
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / 'details.xlsx'

class ListLinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            href = dict(attrs).get('href')
            if href:
                self.links.append(href)

class DetailParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_h1_name = False
        self.name_parts = []
        self.in_th = False
        self.in_td = False
        self.cur_th = []
        self.cur_td = []
        self.pending_th = None
        self.fields = {}
    def handle_starttag(self, tag, attrs):
        attrs_d = dict(attrs)
        if tag.lower() == 'h1' and 'name' in attrs_d.get('class', '').split():
            self.in_h1_name = True
        elif tag.lower() == 'th':
            self.in_th = True
            self.cur_th = []
        elif tag.lower() == 'td':
            self.in_td = True
            self.cur_td = []
    def handle_data(self, data):
        if self.in_h1_name:
            self.name_parts.append(data)
        if self.in_th:
            self.cur_th.append(data)
        if self.in_td:
            self.cur_td.append(data)
    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == 'h1':
            self.in_h1_name = False
        elif tag == 'th':
            self.in_th = False
            self.pending_th = ''.join(self.cur_th).strip()
        elif tag == 'td':
            self.in_td = False
            value = ''.join(self.cur_td).strip()
            if self.pending_th:
                self.fields[self.pending_th.strip().lower()] = value
                self.pending_th = None
    @property
    def name(self):
        return ''.join(self.name_parts).strip()

def fetch(url):
    with urlopen(url, timeout=10) as resp:
        raw = resp.read()
        enc = resp.headers.get_content_charset() or 'utf-8'
        return raw.decode(enc, errors='replace')

def parse_price_currency(raw):
    if raw is None or str(raw).strip() == '':
        return '', ''
    s = str(raw).strip()
    m = re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', s)
    if not m:
        return s, ''
    num_text = m.group(0).replace(',', '')
    f = float(num_text)
    price = int(f) if abs(f - round(f)) < 1e-9 else f
    before = s[:m.start()].strip()
    after = s[m.end():].strip()
    currency = (before or after).strip()
    return price, currency

def memory_to_gb(value):
    if value is None or str(value).strip() == '':
        return ''
    s = str(value).strip()
    m = re.search(r'([-+]?\d+(?:\.\d+)?)\s*([KMGT]?I?B|[KMGT])?', s, re.I)
    if not m:
        return ''
    num = float(m.group(1))
    unit = (m.group(2) or 'GB').upper().replace('IB', 'B')
    if unit in ('GB', 'G'):
        gb = num
    elif unit in ('MB', 'M'):
        gb = num / 1024
    elif unit in ('KB', 'K'):
        gb = num / (1024 * 1024)
    elif unit in ('TB', 'T'):
        gb = num * 1024
    else:
        gb = num
    return int(round(gb)) if abs(gb - round(gb)) < 1e-9 else round(gb, 6)

print('WORKSPACE:', workspace)
print('OUTPUT:', out_path)
list_html = fetch(BASE_LIST_URL)
lp = ListLinkParser()
lp.feed(list_html)
seen = set()
detail_urls = []
for href in lp.links:
    full = urljoin(BASE_LIST_URL, href)
    if full not in seen:
        seen.add(full)
        detail_urls.append(full)
print('FOUND DETAIL URLS:', len(detail_urls))
for u in detail_urls:
    print(' -', u)

headers = ['Name', 'Price', 'Currency', 'CPU', 'MemoryGB', 'URL']
rows = []
for url in detail_urls:
    html = fetch(url)
    dp = DetailParser()
    dp.feed(html)
    price, currency = parse_price_currency(dp.fields.get('price', ''))
    row = {
        'Name': dp.name,
        'Price': price,
        'Currency': currency,
        'CPU': dp.fields.get('cpu', ''),
        'MemoryGB': memory_to_gb(dp.fields.get('memorygb', dp.fields.get('memory', ''))),
        'URL': url,
    }
    print('PARSED:', row)
    rows.append(row)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'details'
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(out_path)
print('SAVED:', out_path)

# Reopen and validate all stated requirements.
wb2 = openpyxl.load_workbook(out_path, data_only=True)
assert 'details' in wb2.sheetnames
ws2 = wb2['details']
read_headers = [cell.value for cell in ws2[1]]
assert read_headers == headers, f'headers mismatch: {read_headers}'
data_rows = list(ws2.iter_rows(min_row=2, values_only=True))
assert len(data_rows) == len(detail_urls), f'row count mismatch {len(data_rows)} != {len(detail_urls)}'
expected_urls = set(detail_urls)
actual_urls = set()
for idx, row in enumerate(data_rows, start=2):
    assert len(row) == len(headers)
    name, price, currency, cpu, mem, url = row
    assert isinstance(name, str) and name.strip(), f'missing Name row {idx}'
    assert price in (None, '') or isinstance(price, (int, float)), f'Price not numeric/blank row {idx}: {price!r}'
    assert currency in (None, '') or isinstance(currency, str), f'Currency not text/blank row {idx}: {currency!r}'
    assert cpu in (None, '') or isinstance(cpu, str), f'CPU not text/blank row {idx}: {cpu!r}'
    assert mem in (None, '') or isinstance(mem, (int, float)), f'MemoryGB not numeric/blank row {idx}: {mem!r}'
    assert isinstance(url, str) and url in expected_urls, f'bad URL row {idx}: {url!r}'
    actual_urls.add(url)
assert actual_urls == expected_urls, f'URL set mismatch: {actual_urls} != {expected_urls}'
print('VALIDATION PASSED')
print(' - File:', out_path)
print(' - Headers:', read_headers)
print(' - Data rows:', len(data_rows))
print(' - URLs match all detail pages from list')
print(' - Name populated from detail page h1.name')
print(' - Price numeric, Currency text/blank, CPU text/blank, MemoryGB numeric GB/blank')
