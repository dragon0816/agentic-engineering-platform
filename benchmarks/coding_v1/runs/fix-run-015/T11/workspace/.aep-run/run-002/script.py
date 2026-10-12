from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
from html import unescape
import re
import openpyxl

BASE_LIST_URL = 'http://127.0.0.1:50447/t11/list'
workspace = Path.cwd()
out_dir = workspace / 'output'
out_dir.mkdir(parents=True, exist_ok=True)
out_path = out_dir / 'details.xlsx'

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attr = dict(attrs)
            href = attr.get('href')
            if href:
                self.links.append(href)

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, data):
        if data and data.strip():
            self.parts.append(data.strip())
    def text(self):
        return '\n'.join(self.parts)

def fetch(url):
    with urlopen(url, timeout=10) as resp:
        raw = resp.read()
        enc = resp.headers.get_content_charset() or 'utf-8'
        return raw.decode(enc, errors='replace')

def page_text(html):
    p = TextParser()
    p.feed(html)
    return p.text()

def norm_label(label):
    return re.sub(r'[^a-z0-9]', '', label.lower())

def parse_fields(text):
    fields = {}
    # Lines can be "Name: ..." or table/list text where label/value are adjacent.
    for line in text.splitlines():
        line = line.strip()
        if not line:
            continue
        m = re.match(r'^\s*([^:：]+)\s*[:：]\s*(.*?)\s*$', line)
        if m:
            fields[norm_label(m.group(1))] = m.group(2).strip()
    # Fallback: scan whole text for common labels if not colon-separated line by line.
    labels = ['Name', 'Price', 'Currency', 'CPU', 'Memory', 'MemoryGB']
    for label in labels:
        key = norm_label(label)
        if key in fields:
            continue
        pattern = re.compile(rf'{re.escape(label)}\s*[:：]?\s*([^\n]+)', re.I)
        m = pattern.search(text)
        if m:
            val = m.group(1).strip()
            # Avoid swallowing following labels on same line.
            val = re.split(r'\s+(?:Name|Price|Currency|CPU|MemoryGB|Memory)\s*[:：]', val, maxsplit=1, flags=re.I)[0].strip()
            fields[key] = val
    return fields

def memory_to_gb(value):
    if value is None:
        return ''
    s = str(value).strip()
    if not s:
        return ''
    # Prefer explicit numeric + unit.
    m = re.search(r'([-+]?\d+(?:\.\d+)?)\s*([KMGT]?I?B|[KMGT])?', s, re.I)
    if not m:
        return s
    num = float(m.group(1))
    unit = (m.group(2) or 'GB').upper()
    unit = unit.replace('IB', 'B')
    if unit in ('GB', 'G'):
        gb = num
    elif unit in ('MB', 'M'):
        gb = num / 1024
    elif unit in ('KB', 'K'):
        gb = num / (1024 * 1024)
    elif unit in ('TB', 'T'):
        gb = num * 1024
    else:
        gb = num
    if abs(gb - round(gb)) < 1e-9:
        return int(round(gb))
    return round(gb, 6)

print('WORKSPACE:', workspace)
print('OUTPUT:', out_path)

list_html = fetch(BASE_LIST_URL)
link_parser = LinkParser()
link_parser.feed(list_html)
detail_urls = []
seen = set()
for href in link_parser.links:
    full = urljoin(BASE_LIST_URL, href)
    if full not in seen:
        seen.add(full)
        detail_urls.append(full)
print('FOUND DETAIL URLS:', len(detail_urls))
for u in detail_urls:
    print(' -', u)

rows = []
for url in detail_urls:
    html = fetch(url)
    text = page_text(html)
    print('\nDETAIL PAGE:', url)
    print(text[:1000])
    fields = parse_fields(text)
    name = fields.get('name', '')
    price_raw = fields.get('price', '')
    currency = fields.get('currency', '')
    cpu = fields.get('cpu', '')
    mem_raw = fields.get('memorygb', fields.get('memory', ''))
    mem_gb = memory_to_gb(mem_raw)
    # Convert price to numeric if possible while preserving blank for missing.
    price = ''
    if price_raw != '':
        cleaned = re.sub(r'[^0-9.\-]', '', price_raw)
        if cleaned not in ('', '.', '-', '-.'):
            f = float(cleaned)
            price = int(f) if abs(f - round(f)) < 1e-9 else f
        else:
            price = price_raw
    row = {'Name': name, 'Price': price, 'Currency': currency, 'CPU': cpu, 'MemoryGB': mem_gb, 'URL': url}
    print('PARSED:', row)
    rows.append(row)

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'details'
headers = ['Name', 'Price', 'Currency', 'CPU', 'MemoryGB', 'URL']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(out_path)
print('\nSAVED:', out_path)

# Validation: reopen output and confirm requested requirements.
wb2 = openpyxl.load_workbook(out_path, data_only=True)
ws2 = wb2['details']
read_headers = [c.value for c in ws2[1]]
assert read_headers == headers, f'headers mismatch: {read_headers}'
data_rows = list(ws2.iter_rows(min_row=2, values_only=True))
assert len(data_rows) == len(detail_urls), f'row count mismatch {len(data_rows)} != {len(detail_urls)}'
for idx, row in enumerate(data_rows, start=2):
    assert len(row) == len(headers), f'row {idx} column count mismatch'
    # URL should be populated and detail page URL.
    assert isinstance(row[5], str) and row[5].startswith('http://127.0.0.1:50447/t11/item/'), f'bad URL row {idx}: {row[5]}'
    # Missing fields are allowed as blank/None in Excel; populated MemoryGB must be numeric.
    if row[4] not in (None, ''):
        assert isinstance(row[4], (int, float)), f'MemoryGB not numeric row {idx}: {row[4]!r}'
    if row[1] not in (None, ''):
        assert isinstance(row[1], (int, float)), f'Price not numeric row {idx}: {row[1]!r}'
print('VALIDATION PASSED')
print(' - File exists:', out_path.exists())
print(' - Sheet:', ws2.title)
print(' - Headers:', read_headers)
print(' - Data rows:', len(data_rows))
print(' - MemoryGB populated values are numeric; blanks represent missing fields')
print(' - Price populated values are numeric; blanks represent missing fields')
