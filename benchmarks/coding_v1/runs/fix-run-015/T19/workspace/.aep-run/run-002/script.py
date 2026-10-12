from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read CSV with utf-8-sig so UTF-8 BOM is handled; newline='' preserves quoted newlines for csv parser.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    rows = list(reader)

if not rows:
    raise SystemExit('mixed.csv is empty; no headers to write')

headers = rows[0]
try:
    device_col_idx = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_col_idx = None

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'

for r_idx, row in enumerate(rows, start=1):
    # csv.reader keeps empty fields as ''. Keep all columns from the file; if a malformed short row exists,
    # write provided fields only and leave absent cells blank.
    for c_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        # Store Device ID as text, including header and all data values, preserving leading zeros.
        if device_col_idx is not None and c_idx == device_col_idx:
            cell.number_format = numbers.FORMAT_TEXT
            cell.value = '' if value is None else str(value)

# Also format the whole Device ID column as text for clarity in Excel.
if device_col_idx is not None:
    col_letter = ws.cell(row=1, column=device_col_idx).column_letter
    ws.column_dimensions[col_letter].number_format = numbers.FORMAT_TEXT

wb.save(output_path)

# Verification by reopening both source and output.
wb2 = load_workbook(output_path, data_only=False)
ws2 = wb2.active
xlsx_rows = []
for r in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=len(headers)):
    xlsx_rows.append(['' if cell.value is None else str(cell.value) for cell in r])

# Normalize source rows to header width for comparison; absent cells are blank.
source_norm = []
for row in rows:
    padded = list(row) + [''] * (len(headers) - len(row))
    source_norm.append(['' if v is None else str(v) for v in padded[:len(headers)]])

assert xlsx_rows == source_norm, 'Output rows do not exactly match source rows/order/duplicates/blank values'
assert ws2.max_column == len(headers), f'Expected {len(headers)} columns, got {ws2.max_column}'
assert [ws2.cell(row=1, column=i).value for i in range(1, len(headers)+1)] == headers, 'Headers changed'
if device_col_idx is not None:
    for r_idx in range(1, ws2.max_row + 1):
        cell = ws2.cell(row=r_idx, column=device_col_idx)
        assert isinstance(cell.value, str), f'Device ID cell {r_idx} is not text/string'
        assert cell.number_format == numbers.FORMAT_TEXT or cell.number_format == '@', f'Device ID cell {r_idx} not formatted as text'
    # Ensure no leading zeros were lost for known leading-zero values.
    for src, out in zip(source_norm[1:], xlsx_rows[1:]):
        assert out[device_col_idx-1] == src[device_col_idx-1], 'Device ID value changed'

print('Wrote:', output_path)
print('Rows including header:', len(rows))
print('Columns:', len(headers))
print('Headers:', headers)
print('Device ID column:', device_col_idx if device_col_idx else 'not present')
print('Verification passed: reopened converted.xlsx; preserved all columns, order, duplicate rows, blanks, quoted commas/newlines, Chinese text; no index column; Device ID stored/formatted as text.')
