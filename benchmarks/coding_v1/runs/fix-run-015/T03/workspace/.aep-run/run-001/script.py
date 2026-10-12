from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
print('Workspace:', workspace)
print('Input files:', [p.name for p in input_dir.iterdir() if p.is_file()])
path = input_dir / 'tasks.xlsx'
print('tasks.xlsx exists:', path.exists())
wb = load_workbook(path, read_only=True, data_only=False)
print('Workbook sheets:', wb.sheetnames)
ws = wb['tasks']
rows = list(ws.iter_rows(values_only=True))
print('tasks dimensions:', ws.max_row, ws.max_column)
print('Headers:', rows[0] if rows else None)
print('Sample rows:')
for row in rows[1:6]:
    print(row)
wb.close()