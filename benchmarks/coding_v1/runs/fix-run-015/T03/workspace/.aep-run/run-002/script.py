from pathlib import Path
from collections import OrderedDict, Counter
from openpyxl import load_workbook, Workbook
from copy import copy
import re

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'owners.xlsx'

INVALID_CHARS = r'[:\\/?*\[\]]'
MAX_SHEET_LEN = 31

def owner_key(value):
    if value is None:
        return 'Unassigned'
    if isinstance(value, str) and value.strip() == '':
        return 'Unassigned'
    return value

def safe_sheet_name(owner, used):
    base = 'Unassigned' if owner == 'Unassigned' else str(owner)
    base = re.sub(INVALID_CHARS, '_', base)
    base = base.strip()
    if not base:
        base = 'Sheet'
    # Excel worksheet title cannot exceed 31 chars.
    base = base[:MAX_SHEET_LEN]
    candidate = base
    idx = 1
    while candidate in used:
        suffix = f'_{idx}'
        candidate = base[:MAX_SHEET_LEN - len(suffix)] + suffix
        idx += 1
    used.add(candidate)
    return candidate

src_wb = load_workbook(input_path)
if 'tasks' not in src_wb.sheetnames:
    raise ValueError('Input workbook does not contain a tasks worksheet')
src_ws = src_wb['tasks']
headers = [cell.value for cell in src_ws[1]]
try:
    owner_idx = headers.index('Owner') + 1
    status_idx = headers.index('Status') + 1
except ValueError as e:
    raise ValueError('tasks worksheet must contain Owner and Status columns') from e

groups = OrderedDict()
for row in src_ws.iter_rows(min_row=2):
    values = [cell.value for cell in row]
    # Keep any row that exists in the worksheet; skip only fully empty trailing-style rows if present.
    if all(v is None for v in values):
        continue
    owner = owner_key(row[owner_idx - 1].value)
    groups.setdefault(owner, []).append(row)

out_wb = Workbook()
default = out_wb.active
out_wb.remove(default)
used_names = set()
summary_records = []

for owner, rows in groups.items():
    sheet_name = safe_sheet_name(owner, used_names)
    ws = out_wb.create_sheet(sheet_name)
    # Write header row exactly as source column names.
    for col_idx, src_cell in enumerate(src_ws[1], start=1):
        dst = ws.cell(row=1, column=col_idx, value=src_cell.value)
        if src_cell.has_style:
            dst.font = copy(src_cell.font)
            dst.fill = copy(src_cell.fill)
            dst.border = copy(src_cell.border)
            dst.alignment = copy(src_cell.alignment)
            dst.number_format = src_cell.number_format
            dst.protection = copy(src_cell.protection)
    # Write grouped rows, preserving original Owner values (including blanks/None).
    pass_count = fail_count = 0
    for out_row_idx, src_row in enumerate(rows, start=2):
        status = src_row[status_idx - 1].value
        if status == 'Pass':
            pass_count += 1
        elif status == 'Fail':
            fail_count += 1
        else:
            raise ValueError(f'Unexpected Status {status!r}; expected Pass or Fail')
        for col_idx, src_cell in enumerate(src_row, start=1):
            dst = ws.cell(row=out_row_idx, column=col_idx, value=src_cell.value)
            if src_cell.has_style:
                dst.font = copy(src_cell.font)
                dst.fill = copy(src_cell.fill)
                dst.border = copy(src_cell.border)
                dst.alignment = copy(src_cell.alignment)
                dst.number_format = src_cell.number_format
                dst.protection = copy(src_cell.protection)
    summary_records.append({
        'Owner': owner,
        'SheetName': sheet_name,
        'Total': len(rows),
        'Pass': pass_count,
        'Fail': fail_count,
    })

summary_ws = out_wb.create_sheet('Summary', 0)
summary_headers = ['Owner', 'SheetName', 'Total', 'Pass', 'Fail']
for col_idx, header in enumerate(summary_headers, start=1):
    summary_ws.cell(row=1, column=col_idx, value=header)
for row_idx, rec in enumerate(summary_records, start=2):
    for col_idx, header in enumerate(summary_headers, start=1):
        summary_ws.cell(row=row_idx, column=col_idx, value=rec[header])

out_wb.save(output_path)
src_wb.close()
out_wb.close()

print('Wrote:', output_path)
print('Groups:')
for rec in summary_records:
    print(rec)

# Verification: reopen output and confirm stated requirements.
out = load_workbook(output_path, data_only=False)
summary = out['Summary']
actual_summary_headers = [c.value for c in summary[1]]
assert actual_summary_headers == summary_headers, actual_summary_headers
assert 'Summary' in out.sheetnames

# Reopen input to compare source rows.
src = load_workbook(input_path, data_only=False)
src_ws = src['tasks']
source_data_rows = []
expected_groups = OrderedDict()
for row in src_ws.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    source_data_rows.append(row)
    expected_groups.setdefault(owner_key(row[owner_idx - 1]), []).append(row)

summary_rows = list(summary.iter_rows(min_row=2, values_only=True))
assert len(summary_rows) == len(expected_groups), (len(summary_rows), len(expected_groups))
seen_sheet_names = set()
reconstructed = []
for owner, sheet_name, total, p, f in summary_rows:
    assert sheet_name in out.sheetnames, sheet_name
    assert sheet_name != 'Summary'
    assert sheet_name not in seen_sheet_names
    seen_sheet_names.add(sheet_name)
    assert len(sheet_name) <= 31, sheet_name
    assert not re.search(INVALID_CHARS, sheet_name), sheet_name
    ws = out[sheet_name]
    assert [c.value for c in ws[1]] == headers
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert rows == expected_groups[owner], (owner, rows, expected_groups[owner])
    assert total == len(rows), (owner, total, len(rows))
    assert p == sum(1 for r in rows if r[status_idx - 1] == 'Pass')
    assert f == sum(1 for r in rows if r[status_idx - 1] == 'Fail')
    reconstructed.extend(rows)
assert reconstructed == source_data_rows
assert set(out.sheetnames) == {'Summary'} | seen_sheet_names

print('Verification passed.')
print('Output sheets:', out.sheetnames)
print('Summary headers:', actual_summary_headers)
print('Summary rows:', summary_rows)
print('Total source data rows:', len(source_data_rows))
out.close()
src.close()