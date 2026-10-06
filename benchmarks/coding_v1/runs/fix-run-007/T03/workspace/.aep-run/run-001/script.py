from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
print('Workspace:', workspace)
print('Input files:')
for p in sorted((workspace/'input').glob('*')):
    print(' -', p.name)

src = workspace/'input'/'tasks.xlsx'
wb = load_workbook(src, read_only=True, data_only=False)
print('Workbook sheets:', wb.sheetnames)
ws = wb['tasks']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print('tasks headers:', headers)
print('Sample rows:')
for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 6), values_only=True):
    print(row)
print('Max row:', ws.max_row, 'Max column:', ws.max_column)
wb.close()
