from pathlib import Path
from collections import OrderedDict
import re
from openpyxl import load_workbook, Workbook

workspace = Path.cwd()
input_file = workspace / 'input' / 'tasks.xlsx'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)
output_file = output_dir / 'owners.xlsx'

INVALID_CHARS = r'[]:*?/\\'
INVALID_RE = re.compile(r'[\[\]\:\*\?\/\\]')
MAX_SHEET_LEN = 31
RESERVED = {'Summary'}

def display_owner(value):
    if value is None:
        return 'Unassigned'
    if isinstance(value, str) and value.strip() == '':
        return 'Unassigned'
    return value

def safe_sheet_name(owner_value, used):
    # Use grouping/display owner as the candidate sheet name.
    name = str(owner_value) if owner_value is not None else 'Unassigned'
    name = INVALID_RE.sub('_', name)
    # Excel also dislikes sheet names that are empty or all apostrophes in practice.
    name = name.strip()
    if not name:
        name = 'Sheet'
    # Avoid reserved Summary name for group sheets.
    if name in RESERVED:
        name = name + '_Group'
    base = name[:MAX_SHEET_LEN]
    candidate = base
    i = 1
    while candidate in used:
        suffix = f'_{i}'
        candidate = base[:MAX_SHEET_LEN - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate

src_wb = load_workbook(input_file, data_only=False)
if 'tasks' not in src_wb.sheetnames:
    raise ValueError('tasks worksheet not found')
src_ws = src_wb['tasks']
rows = list(src_ws.iter_rows(values_only=True))
if not rows:
    raise ValueError('tasks worksheet is empty')
headers = list(rows[0])
try:
    owner_idx = headers.index('Owner')
    status_idx = headers.index('Status')
except ValueError as exc:
    raise ValueError('Required columns Owner and Status were not found') from exc

groups = OrderedDict()
for row in rows[1:]:
    values = list(row) + [None] * (len(headers) - len(row))
    owner_key = display_owner(values[owner_idx])
    groups.setdefault(owner_key, []).append(values[:len(headers)])

out_wb = Workbook()
# Remove default sheet.
out_wb.remove(out_wb.active)
used_names = set()
summary_records = []

for owner_key, group_rows in groups.items():
    sheet_name = safe_sheet_name(owner_key, used_names)
    ws = out_wb.create_sheet(title=sheet_name)
    ws.append(headers)
    for row in group_rows:
        ws.append(row)
    total = len(group_rows)
    pass_count = sum(1 for row in group_rows if row[status_idx] == 'Pass')
    fail_count = sum(1 for row in group_rows if row[status_idx] == 'Fail')
    summary_records.append({
        'Owner': owner_key,
        'SheetName': sheet_name,
        'Total': total,
        'Pass': pass_count,
        'Fail': fail_count,
    })

summary_ws = out_wb.create_sheet(title='Summary', index=0)
summary_headers = ['Owner', 'SheetName', 'Total', 'Pass', 'Fail']
summary_ws.append(summary_headers)
for record in summary_records:
    summary_ws.append([record[h] for h in summary_headers])

out_wb.save(output_file)
src_wb.close()
out_wb.close()

# Verification: reopen output and confirm requirements.
check_wb = load_workbook(output_file, data_only=True)
print('Output file:', output_file)
print('Output sheets:', check_wb.sheetnames)
assert 'Summary' in check_wb.sheetnames, 'Summary sheet missing'
summary = check_wb['Summary']
actual_summary_headers = [c.value for c in summary[1]]
print('Summary headers:', actual_summary_headers)
assert actual_summary_headers == summary_headers, f'Summary headers mismatch: {actual_summary_headers}'

# Build expected row counts and counts from source groups.
expected = OrderedDict()
for owner_key, group_rows in groups.items():
    expected[owner_key] = {
        'Total': len(group_rows),
        'Pass': sum(1 for row in group_rows if row[status_idx] == 'Pass'),
        'Fail': sum(1 for row in group_rows if row[status_idx] == 'Fail'),
    }

seen_owners = []
for row in summary.iter_rows(min_row=2, values_only=True):
    owner, sheet_name, total, p, f = row
    seen_owners.append(owner)
    assert owner in expected, f'Unexpected owner in Summary: {owner}'
    assert sheet_name in check_wb.sheetnames, f'SheetName not found: {sheet_name}'
    assert len(sheet_name) <= 31 and not any(ch in sheet_name for ch in INVALID_CHARS), f'Invalid sheet name: {sheet_name}'
    assert isinstance(total, int) and isinstance(p, int) and isinstance(f, int), 'Counts must be numeric integers'
    assert (total, p, f) == (expected[owner]['Total'], expected[owner]['Pass'], expected[owner]['Fail']), f'Counts mismatch for {owner}'
    data_ws = check_wb[sheet_name]
    data_headers = [c.value for c in data_ws[1]]
    assert data_headers == headers, f'Headers changed in {sheet_name}'
    assert data_ws.max_row - 1 == total, f'Row count mismatch in {sheet_name}'

assert list(expected.keys()) == seen_owners, 'Owner order or membership mismatch'
assert len(check_wb.sheetnames) == len(expected) + 1, 'Unexpected number of sheets'
print('Verified Summary rows:', summary.max_row - 1)
print('Verified group sheets:', len(expected))
for row in summary.iter_rows(min_row=2, values_only=True):
    print(row)
check_wb.close()
