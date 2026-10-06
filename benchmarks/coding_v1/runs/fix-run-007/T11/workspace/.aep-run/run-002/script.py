import os, re, html, urllib.request, urllib.parse
from html.parser import HTMLParser
from openpyxl import Workbook, load_workbook

LIST_URL = 'http://127.0.0.1:57997/t11/list'
OUT_PATH = os.path.join('output', 'details.xlsx')
HEADERS = ['Name', 'Price', 'Currency', 'CPU', 'MemoryGB', 'URL']

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            d = dict(attrs)
            href = d.get('href')
            if href:
                self.links.append(href)

class Node:
    def __init__(self, tag='document', attrs=None, parent=None):
        self.tag = tag
        self.attrs = dict(attrs or [])
        self.parent = parent
        self.children = []
        self.text = ''

class DOMParser(HTMLParser):
    SELF_CLOSING = {'br','hr','img','input','meta','link'}
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node()
        self.stack = [self.root]
    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        node = Node(tag, attrs, self.stack[-1])
        self.stack[-1].children.append(node)
        if tag not in self.SELF_CLOSING:
            self.stack.append(node)
    def handle_endtag(self, tag):
        tag = tag.lower()
        for i in range(len(self.stack)-1, 0, -1):
            if self.stack[i].tag == tag:
                del self.stack[i:]
                break
    def handle_data(self, data):
        if data:
            self.stack[-1].text += data

def fetch(url):
    with urllib.request.urlopen(url, timeout=10) as r:
        raw = r.read()
        charset = r.headers.get_content_charset() or 'utf-8'
        return raw.decode(charset, errors='replace')

def text_content(node):
    parts = []
    if node.text:
        parts.append(node.text)
    for c in node.children:
        parts.append(text_content(c))
    return html.unescape(' '.join(p for p in parts if p)).strip()

def norm_space(s):
    return re.sub(r'\s+', ' ', html.unescape(str(s or ''))).strip()

def find_all(node, tag=None):
    out = []
    for c in node.children:
        if tag is None or c.tag == tag:
            out.append(c)
        out.extend(find_all(c, tag))
    return out

def child_tags(node, tags):
    return [c for c in node.children if c.tag in tags]

def label_key(s):
    k = re.sub(r'[^a-z0-9]', '', (s or '').lower())
    mapping = {
        'name':'Name', 'productname':'Name', 'title':'Name',
        'price':'Price', 'currency':'Currency',
        'cpu':'CPU', 'processor':'CPU',
        'memory':'Memory', 'ram':'Memory', 'memorygb':'Memory',
    }
    return mapping.get(k)

def parse_key_values(root, raw_html):
    kv = {}
    # Tables: first cell as label, following cells as value.
    for tr in find_all(root, 'tr'):
        cells = child_tags(tr, {'th','td'})
        if len(cells) >= 2:
            key = label_key(text_content(cells[0]))
            if key and key not in kv:
                kv[key] = norm_space(' '.join(text_content(c) for c in cells[1:]))
    # Definition lists: dt followed by dd.
    for dl in find_all(root, 'dl'):
        last_key = None
        for c in dl.children:
            if c.tag == 'dt':
                last_key = label_key(text_content(c))
            elif c.tag == 'dd' and last_key and last_key not in kv:
                kv[last_key] = norm_space(text_content(c))
                last_key = None
    # Elements with useful class/id names.
    for n in find_all(root):
        attrs_joined = ' '.join(str(n.attrs.get(a,'')) for a in ('id','class','name','data-field','data-label'))
        key = label_key(attrs_joined)
        val = norm_space(text_content(n))
        if key and val and key not in kv:
            kv[key] = val
        # currency in attributes such as data-currency="USD"
        for attr_name, attr_val in n.attrs.items():
            if 'currency' in attr_name.lower() and attr_val and 'Currency' not in kv:
                kv['Currency'] = norm_space(attr_val)
    # Plain visible text patterns.
    visible = norm_space(re.sub(r'<[^>]+>', ' ', raw_html))
    for label in ['Name', 'Price', 'Currency', 'CPU', 'Memory', 'MemoryGB', 'RAM', 'Processor']:
        canonical = label_key(label)
        if canonical and canonical not in kv:
            m = re.search(r'\b' + re.escape(label) + r'\b\s*[:=]\s*([^|;\n\r]+?)(?=\s+(?:Name|Price|Currency|CPU|MemoryGB|Memory|RAM|Processor)\b\s*[:=]|$)', visible, flags=re.I)
            if m:
                kv[canonical] = norm_space(m.group(1))
    return kv

def parse_price_currency(price_text, currency_text=''):
    ptxt = norm_space(price_text)
    ctxt = norm_space(currency_text).upper()
    # Infer currency from price text if separate field absent.
    if not ctxt:
        m = re.search(r'\b(USD|EUR|GBP|JPY|TWD|NTD|CNY|RMB)\b', ptxt, re.I)
        if m:
            ctxt = m.group(1).upper()
            if ctxt == 'NTD':
                ctxt = 'TWD'
        elif 'NT$' in ptxt or 'NTD' in ptxt.upper():
            ctxt = 'TWD'
        elif '$' in ptxt:
            ctxt = 'USD'
        elif '€' in ptxt:
            ctxt = 'EUR'
        elif '£' in ptxt:
            ctxt = 'GBP'
        elif '¥' in ptxt:
            ctxt = 'JPY'
    else:
        ctxt = {'NTD':'TWD', 'RMB':'CNY'}.get(ctxt, ctxt)
    # First number in the price text.
    m = re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', ptxt)
    price = None
    if m:
        num = m.group(0).replace(',', '')
        try:
            f = float(num)
            price = int(f) if f.is_integer() else f
        except ValueError:
            price = None
    return price, ctxt or None

def parse_memory_gb(mem_text):
    s = norm_space(mem_text)
    if not s:
        return None
    # Handle common multiplier forms such as 2 x 8 GB.
    mx = re.search(r'(\d+(?:\.\d+)?)\s*[x×]\s*(\d+(?:\.\d+)?)\s*(TB|TIB|GB|GIB|MB|MIB)?', s, re.I)
    if mx:
        count = float(mx.group(1)); size = float(mx.group(2)); unit = (mx.group(3) or 'GB').upper()
        gb = count * size * unit_factor(unit)
        return int(gb) if float(gb).is_integer() else round(gb, 4)
    m = re.search(r'(\d+(?:\.\d+)?)\s*(TB|TIB|GB|GIB|MB|MIB)?', s, re.I)
    if not m:
        return None
    val = float(m.group(1)); unit = (m.group(2) or 'GB').upper()
    gb = val * unit_factor(unit)
    return int(gb) if float(gb).is_integer() else round(gb, 4)

def unit_factor(unit):
    unit = (unit or 'GB').upper()
    if unit in ('TB','TIB'):
        return 1024.0
    if unit in ('MB','MIB'):
        return 1.0/1024.0
    return 1.0

def parse_detail(url, page_html):
    parser = DOMParser(); parser.feed(page_html)
    root = parser.root
    kv = parse_key_values(root, page_html)
    # Name fallback: h1, then title.
    name = kv.get('Name')
    if not name:
        h1s = [norm_space(text_content(n)) for n in find_all(root, 'h1') if norm_space(text_content(n))]
        if h1s:
            name = h1s[0]
    if not name:
        titles = [norm_space(text_content(n)) for n in find_all(root, 'title') if norm_space(text_content(n))]
        if titles:
            name = titles[0]
    price, currency = parse_price_currency(kv.get('Price',''), kv.get('Currency',''))
    memory = parse_memory_gb(kv.get('Memory',''))
    return {
        'Name': name or None,
        'Price': price,
        'Currency': currency,
        'CPU': kv.get('CPU') or None,
        'MemoryGB': memory,
        'URL': url,
        '_kv': kv,
    }

# Inspect list and detail pages.
list_html = fetch(LIST_URL)
link_parser = LinkParser(); link_parser.feed(list_html)
links = []
seen = set()
for href in link_parser.links:
    full = urllib.parse.urljoin(LIST_URL, href)
    if full not in seen:
        links.append(full); seen.add(full)
print('List URL:', LIST_URL)
print('Found detail links:', len(links), links)
if not links:
    raise RuntimeError('No product detail links found on list page')

rows = []
for i, url in enumerate(links, start=1):
    page = fetch(url)
    print(f'Detail {i} URL: {url}; bytes={len(page)}; snippet={norm_space(page)[:500]}')
    rec = parse_detail(url, page)
    print('Parsed:', {k: rec[k] for k in HEADERS}, 'raw_kv=', rec['_kv'])
    rows.append([rec[h] for h in HEADERS])

os.makedirs('output', exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = 'details'
ws.append(HEADERS)
for row in rows:
    ws.append(row)
# basic column sizing
for col_idx, header in enumerate(HEADERS, start=1):
    ws.column_dimensions[chr(64+col_idx)].width = max(12, len(header) + 2)
wb.save(OUT_PATH)
print('Saved:', OUT_PATH)

# Reopen and verify stated requirements.
wb2 = load_workbook(OUT_PATH, data_only=True)
ws2 = wb2['details']
headers = [ws2.cell(row=1, column=i).value for i in range(1, len(HEADERS)+1)]
assert headers == HEADERS, f'Header mismatch: {headers}'
assert ws2.max_row == len(rows) + 1, f'Row count mismatch: {ws2.max_row} vs {len(rows)+1}'
assert ws2.max_column == len(HEADERS), f'Column count mismatch: {ws2.max_column}'
for r in range(2, ws2.max_row+1):
    values = {HEADERS[c-1]: ws2.cell(row=r, column=c).value for c in range(1, len(HEADERS)+1)}
    assert values['URL'] in links, f'Unexpected URL at row {r}: {values["URL"]}'
    assert values['MemoryGB'] is None or isinstance(values['MemoryGB'], (int, float)), f'MemoryGB not numeric/blank at row {r}: {values["MemoryGB"]!r}'
    assert values['Price'] is None or isinstance(values['Price'], (int, float)), f'Price not numeric/blank at row {r}: {values["Price"]!r}'
    for text_col in ['Name','Currency','CPU','URL']:
        assert values[text_col] is None or isinstance(values[text_col], str), f'{text_col} not text/blank at row {r}: {values[text_col]!r}'
print('Verified output: file exists, sheet details, exact headers,', len(rows), 'data rows, URL set, numeric/blank Price and MemoryGB, text/blank Name Currency CPU URL.')
print('Output absolute path:', os.path.abspath(OUT_PATH))
