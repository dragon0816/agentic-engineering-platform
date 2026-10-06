import os, re, json, time
from pathlib import Path
from urllib.parse import urljoin
import urllib.request, urllib.error
from html.parser import HTMLParser
from openpyxl import Workbook, load_workbook

BASE_LIST_URL = 'http://127.0.0.1:57997/t13/list'
OUT_DIR = Path('output')
OUT_DIR.mkdir(exist_ok=True)
XLSX_PATH = OUT_DIR / 'robust_products.xlsx'
WARN_PATH = OUT_DIR / 'warnings.json'
COLUMNS = ['Name','Price','Currency','CPU','URL','FetchStatus']
NA = 'N/A'

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_a = False
        self.current_href = None
        self.current_text = []
        self.links = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs = dict(attrs)
            self.in_a = True
            self.current_href = attrs.get('href')
            self.current_text = []
    def handle_data(self, data):
        if self.in_a:
            self.current_text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self.in_a:
            text = ' '.join(''.join(self.current_text).split())
            if self.current_href:
                self.links.append((text, self.current_href))
            self.in_a = False
            self.current_href = None
            self.current_text = []

def fetch_url(url, retry_500=False, max_attempts=3):
    attempts = 0
    last_err = None
    while True:
        attempts += 1
        try:
            req = urllib.request.Request(url, headers={'User-Agent':'Agent/1.0'})
            with urllib.request.urlopen(req, timeout=10) as resp:
                data = resp.read()
                return {'ok': True, 'status': resp.status, 'text': data.decode('utf-8', errors='replace'), 'attempts': attempts, 'error': None}
        except urllib.error.HTTPError as e:
            body = ''
            try:
                body = e.read().decode('utf-8', errors='replace')
            except Exception:
                pass
            last_err = f'HTTP {e.code} {e.reason}'.strip()
            if retry_500 and e.code == 500 and attempts < max_attempts:
                time.sleep(0.15)
                continue
            return {'ok': False, 'status': e.code, 'text': body, 'attempts': attempts, 'error': last_err}
        except Exception as e:
            last_err = repr(e)
            return {'ok': False, 'status': None, 'text': '', 'attempts': attempts, 'error': last_err}

def strip_tags(html):
    html = re.sub(r'(?is)<script.*?</script>|<style.*?</style>', ' ', html)
    html = re.sub(r'(?is)<br\s*/?>', '\n', html)
    html = re.sub(r'(?is)</(p|div|li|tr|h[1-6]|dt|dd|span)>', '\n', html)
    text = re.sub(r'(?is)<[^>]+>', ' ', html)
    text = text.replace('&nbsp;', ' ').replace('&amp;', '&').replace('&lt;', '<').replace('&gt;', '>')
    return '\n'.join(line.strip() for line in text.splitlines() if line.strip())

def extract_label(text, labels):
    for lab in labels:
        # label: value on same line, permissive around colon symbols
        pat = re.compile(r'(?im)^\s*' + re.escape(lab) + r'\s*[:：-]\s*(.+?)\s*$')
        m = pat.search(text)
        if m:
            return m.group(1).strip()
        # label value if HTML flattening lost colon line break
        pat2 = re.compile(r'(?i)\b' + re.escape(lab) + r'\b\s*[:：-]?\s*([^\n]+)')
        m2 = pat2.search(text)
        if m2:
            val = m2.group(1).strip()
            if val and val.lower() != lab.lower():
                return val
    return None

def parse_name(html, text, fallback):
    # Prefer common headings/title/explicit labels, fall back to list name.
    for pat in [r'(?is)<h1[^>]*>(.*?)</h1>', r'(?is)<h2[^>]*>(.*?)</h2>', r'(?is)<title[^>]*>(.*?)</title>']:
        m = re.search(pat, html)
        if m:
            val = strip_tags(m.group(1)).strip()
            if val:
                return val
    val = extract_label(text, ['Name', 'Product', 'Product Name'])
    return val or fallback or NA

def parse_price_currency(html, text):
    val = extract_label(text, ['Price', '價錢', '價格'])
    if not val:
        # Currency symbol followed by number anywhere.
        m = re.search(r'(?i)(USD|EUR|TWD|NT\$|\$|€|£|¥)\s*([0-9][0-9,]*(?:\.\d+)?)', text)
        if m:
            val = m.group(0)
    if not val:
        return NA, NA
    currency = NA
    symbol_map = {'$':'USD', 'NT$':'TWD', '€':'EUR', '£':'GBP', '¥':'JPY'}
    mcur = re.search(r'(?i)\b(USD|EUR|TWD|GBP|JPY|CNY|NTD)\b|NT\$|\$|€|£|¥', val)
    if mcur:
        token = mcur.group(0).upper()
        currency = symbol_map.get(mcur.group(0), token)
        if currency == 'NTD':
            currency = 'TWD'
    mnum = re.search(r'([0-9][0-9,]*(?:\.\d+)?)', val)
    if not mnum:
        return NA, currency
    num_s = mnum.group(1).replace(',', '')
    try:
        num = float(num_s)
        price = int(num) if num.is_integer() else num
    except Exception:
        price = num_s
    return price, currency

def parse_cpu(text):
    val = extract_label(text, ['CPU', 'Processor', 'Chipset', 'SoC'])
    if val:
        # Avoid swallowing following labels when markup is compact.
        val = re.split(r'\s+(?:Price|Currency|Memory|RAM|Storage|URL)\s*[:：-]', val)[0].strip()
        return val or NA
    return NA

# Fetch list.
list_res = fetch_url(BASE_LIST_URL)
print('List fetch:', {'ok': list_res['ok'], 'status': list_res['status'], 'attempts': list_res['attempts'], 'error': list_res['error']})
if not list_res['ok']:
    raise RuntimeError(f"Cannot fetch product list: {list_res['error']}")
parser = LinkParser()
parser.feed(list_res['text'])
links = [(name, urljoin(BASE_LIST_URL, href)) for name, href in parser.links][:5]
print('Parsed list links:', links)
if len(links) != 5:
    raise RuntimeError(f'Expected five product links, found {len(links)}')

rows = []
warnings = []
page_samples = {}
for list_name, url in links:
    res = fetch_url(url, retry_500=True, max_attempts=3)
    info = {'name': list_name, 'url': url, 'http_status': res['status'], 'attempts': res['attempts']}
    if res['ok']:
        text = strip_tags(res['text'])
        page_samples[list_name] = text[:600]
        name = parse_name(res['text'], text, list_name)
        price, currency = parse_price_currency(res['text'], text)
        cpu = parse_cpu(text)
        rows.append({'Name': name, 'Price': price if price != '' else NA, 'Currency': currency if currency != '' else NA, 'CPU': cpu if cpu != '' else NA, 'URL': url, 'FetchStatus': 'ok'})
        missing = [k for k in ['Price','Currency','CPU'] if rows[-1][k] in (None, '', NA)]
        warnings.append({**info, 'fetchStatus': 'ok', 'missing': missing})
    else:
        page_samples[list_name] = (res['text'] or res['error'] or '')[:600]
        rows.append({'Name': list_name or NA, 'Price': NA, 'Currency': NA, 'CPU': NA, 'URL': url, 'FetchStatus': 'error'})
        warnings.append({**info, 'fetchStatus': 'error', 'error': res['error'] or f"HTTP {res['status']}", 'missing': ['Price','Currency','CPU']})

print('Detail page text samples:')
for k,v in page_samples.items():
    print('---', k, '---')
    print(v)
print('Parsed rows:')
for r in rows:
    print(r)

# Write workbook.
wb = Workbook()
ws = wb.active
ws.title = 'Products'
ws.append(COLUMNS)
for r in rows:
    ws.append([r[c] for c in COLUMNS])
for col in range(1, len(COLUMNS)+1):
    ws.cell(row=1, column=col).style = 'Headline 1'
wb.save(XLSX_PATH)

with open(WARN_PATH, 'w', encoding='utf-8') as f:
    json.dump({'source': BASE_LIST_URL, 'count': len(rows), 'warnings': warnings}, f, ensure_ascii=False, indent=2)

# Reopen and validate stated requirements.
wb2 = load_workbook(XLSX_PATH, data_only=True)
ws2 = wb2['Products']
headers = [ws2.cell(1, c).value for c in range(1, len(COLUMNS)+1)]
if headers != COLUMNS:
    raise AssertionError(f'Headers mismatch: {headers}')
read_rows = []
for rr in range(2, ws2.max_row+1):
    read_rows.append({COLUMNS[c-1]: ws2.cell(rr, c).value for c in range(1, len(COLUMNS)+1)})
if len(read_rows) != 5:
    raise AssertionError(f'Row count mismatch: {len(read_rows)}')
for i, r in enumerate(read_rows, start=2):
    if r['FetchStatus'] not in ('ok','error'):
        raise AssertionError(f'Bad FetchStatus at row {i}: {r["FetchStatus"]}')
    if not r['Name'] or not r['URL']:
        raise AssertionError(f'Missing Name/URL at row {i}: {r}')
    if r['FetchStatus'] == 'error':
        for c in ['Price','Currency','CPU']:
            if r[c] != NA:
                raise AssertionError(f'Error row {i} {c} should be N/A: {r}')
    for c in ['Price','Currency','CPU']:
        if r[c] is None or r[c] == '':
            raise AssertionError(f'Blank {c} at row {i}: {r}')
with open(WARN_PATH, 'r', encoding='utf-8') as f:
    warn_read = json.load(f)
if warn_read.get('count') != 5 or len(warn_read.get('warnings', [])) != 5:
    raise AssertionError('warnings.json count mismatch')
for w in warn_read['warnings']:
    if w.get('fetchStatus') not in ('ok','error'):
        raise AssertionError(f'Bad warnings status: {w}')
print('VALIDATION OK')
print('Created:', str(XLSX_PATH.resolve()))
print('Created:', str(WARN_PATH.resolve()))
print('Validated headers:', headers)
print('Validated row count:', len(read_rows))
print('Validated warnings count:', len(warn_read['warnings']))