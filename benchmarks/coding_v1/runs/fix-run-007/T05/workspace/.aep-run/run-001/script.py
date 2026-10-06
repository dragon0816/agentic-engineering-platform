from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
input_dir = base / 'input'
print('Workspace:', base)
print('Input files:', sorted(p.name for p in input_dir.iterdir()))
for fname in ['devices.xlsx', 'tests.xlsx']:
    path = input_dir / fname
    print('\nFILE', fname, 'exists=', path.exists())
    wb = load_workbook(path, data_only=False, read_only=True)
    print('Sheets:', wb.sheetnames)
    for ws in wb.worksheets:
        print(' Sheet:', ws.title, 'max_row=', ws.max_row, 'max_column=', ws.max_column)
        rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 5), values_only=True))
        for i, row in enumerate(rows, 1):
            print('  row', i, row)
        break
    wb.close()