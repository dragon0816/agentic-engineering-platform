import csv
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read CSV with utf-8-sig to support UTF-8 BOM, and csv.reader to correctly handle quoted commas/newlines.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.reader(f))

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
try:
    device_col = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_col = None

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'

for r_idx, row in enumerate(rows, start=1):
    # Preserve source columns only: no added index column. Missing CSV fields, if any, remain blank.
    for c_idx in range(1, len(headers) + 1):
        value = row[c_idx - 1] if c_idx <= len(row) else ''
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.value = value if value != '' else None  # blank cell for empty values
        if c_idx == device_col:
            cell.number_format = '@'
            if value != '':
                cell.value = str(value)  # explicit text, preserving leading zeros
        if isinstance(value, str) and '\n' in value:
            cell.alignment = Alignment(wrap_text=True)

# Also set text format for entire Device ID column, including blanks/future edits.
if device_col:
    for row_cells in ws.iter_rows(min_col=device_col, max_col=device_col, min_row=1, max_row=max(len(rows), 1)):
        row_cells[0].number_format = '@'

wb.save(output_path)

# Verification: reopen every output file and confirm stated requirements.
wb2 = load_workbook(output_path, data_only=False)
ws2 = wb2.active
read_rows = []
for r in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=1, max_col=len(headers)):
    read_rows.append(['' if cell.value is None else cell.value for cell in r])

errors = []
if output_path.name != 'converted.xlsx':
    errors.append('wrong output filename')
if read_rows[0] != headers:
    errors.append(f'headers/order mismatch: {read_rows[0]!r} != {headers!r}')
if len(read_rows) != len(rows):
    errors.append(f'row count mismatch: {len(read_rows)} != {len(rows)}')
for i, original in enumerate(rows, start=1):
    expected = original[:len(headers)] + [''] * max(0, len(headers) - len(original))
    actual = read_rows[i-1]
    if actual != expected:
        errors.append(f'row {i} mismatch: {actual!r} != {expected!r}')

# Check no extra non-empty/index column after the original headers.
if ws2.max_column != len(headers):
    extra_nonempty = []
    for row in ws2.iter_rows(min_row=1, max_row=ws2.max_row, min_col=len(headers)+1, max_col=ws2.max_column):
        for cell in row:
            if cell.value not in (None, ''):
                extra_nonempty.append(cell.coordinate)
    if extra_nonempty:
        errors.append(f'extra non-empty cells beyond source columns: {extra_nonempty}')

# Check Device ID values are strings where nonblank, preserve leading zeros, and cells are text formatted.
if device_col:
    for r_idx in range(2, ws2.max_row + 1):
        cell = ws2.cell(r_idx, device_col)
        original_value = rows[r_idx - 1][device_col - 1] if device_col - 1 < len(rows[r_idx - 1]) else ''
        if original_value != '':
            if not isinstance(cell.value, str):
                errors.append(f'Device ID row {r_idx} is not text/string: {cell.value!r} ({type(cell.value)})')
            if cell.value != original_value:
                errors.append(f'Device ID row {r_idx} changed: {cell.value!r} != {original_value!r}')
        if cell.number_format != '@':
            errors.append(f'Device ID row {r_idx} number format is not text: {cell.number_format!r}')

# Check duplicate rows preserved exactly by comparing row sequence and counts.
original_data_rows = [tuple(row[:len(headers)] + [''] * max(0, len(headers) - len(row))) for row in rows[1:]]
actual_data_rows = [tuple(row) for row in read_rows[1:]]
if original_data_rows != actual_data_rows:
    errors.append('data row sequence differs; duplicates/order may not be preserved')

if errors:
    raise AssertionError('\n'.join(errors))

print('created:', output_path)
print('source rows including header:', len(rows))
print('source columns:', len(headers), headers)
print('output worksheet:', ws2.title)
print('output max_row:', ws2.max_row)
print('output max_column:', ws2.max_column)
print('verified headers/order preserved:', read_rows[0])
print('verified no index column added')
print('verified duplicate/order preserved')
print('verified blank CSV values are blank Excel cells')
print('verified quoted comma/newline fields preserved')
print('verified Device ID is text-formatted and leading zeros preserved')