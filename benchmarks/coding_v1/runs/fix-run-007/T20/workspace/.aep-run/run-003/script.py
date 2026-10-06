import csv
import json
import os
import sys
from pathlib import Path
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = json.load(open(params_path, encoding='utf-8')) if params_path else {}
threshold = params.get('urgent_threshold')
if threshold is None:
    raise SystemExit('Missing required parameter: urgent_threshold')
try:
    threshold = float(threshold)
except Exception as exc:
    raise SystemExit(f'urgent_threshold must be numeric, got {threshold!r}') from exc

# Discover and display the real inputs without modifying them.
print('Workspace:', workspace)
print('Input files:', sorted(p.name for p in input_dir.iterdir()))
sop_path = input_dir / 'supplied_sop.md'
if not sop_path.exists():
    raise FileNotFoundError(f'Missing SOP: {sop_path}')
sop_text = sop_path.read_text(encoding='utf-8')
print('SOP first lines:')
for line in sop_text.splitlines()[:20]:
    print(line)

# Prefer the manifest if it names the results CSV; otherwise fall back to the known CSV location/name.
manifest_candidates = [input_dir / 'manifest.json', input_dir / 'manifest.csv', input_dir / 'manifest.md', input_dir / 'manifest.txt']
csv_path = None
for manifest_path in manifest_candidates:
    if manifest_path.exists():
        text = manifest_path.read_text(encoding='utf-8', errors='replace')
        print(f'Manifest found: {manifest_path.name}')
        print(text[:1000])
        for rel in ['data/test_results.csv', 'test_results.csv', 'results.csv']:
            if rel in text:
                candidate = input_dir / rel
                if candidate.exists():
                    csv_path = candidate
                    break
        if csv_path:
            break
if csv_path is None:
    candidates = list(input_dir.rglob('*.csv'))
    print('CSV candidates:', [str(p.relative_to(input_dir)) for p in candidates])
    if not candidates:
        raise FileNotFoundError('No CSV file found under input/')
    preferred = [p for p in candidates if p.name == 'test_results.csv']
    csv_path = preferred[0] if preferred else candidates[0]
print('Using CSV:', csv_path.relative_to(workspace))

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    fieldnames = reader.fieldnames or []
    rows = list(reader)
print('CSV headers:', fieldnames)
print('CSV sample rows:', rows[:5])
required = ['Device ID', 'Status', 'Value']
missing = [c for c in required if c not in fieldnames]
if missing:
    raise ValueError(f'Missing required CSV columns: {missing}; found {fieldnames}')

fail_rows = []
for row in rows:
    if str(row.get('Status', '')).strip().lower() == 'fail':
        value_text = str(row.get('Value', '')).strip()
        try:
            value_num = float(value_text)
        except Exception as exc:
            raise ValueError(f"Fail row has non-numeric Value: {row}") from exc
        out_row = dict(row)
        out_row['Urgent'] = value_num > threshold
        fail_rows.append(out_row)

# Keep source column names and append Urgent.
out_headers = list(fieldnames)
if 'Urgent' not in out_headers:
    out_headers.append('Urgent')

xlsx_path = output_dir / 'failures.xlsx'
wb = Workbook()
ws = wb.active
ws.title = 'failures'
ws.append(out_headers)
for row in fail_rows:
    values = []
    for header in out_headers:
        if header == 'Urgent':
            values.append(bool(row['Urgent']))
        else:
            values.append(row.get(header, ''))
    ws.append(values)
wb.save(xlsx_path)

urgent_count = sum(1 for r in fail_rows if r['Urgent'])
summary_path = output_dir / 'summary.md'
summary = (
    '# Failure Summary\n\n'
    f'- Source CSV: `{csv_path.relative_to(workspace).as_posix()}`\n'
    f'- Urgent rule: `Value > {threshold:g}`\n'
    f'- Total records: {len(rows)}\n'
    f'- Fail records kept: {len(fail_rows)}\n'
    f'- Urgent failures: {urgent_count}\n'
)
summary_path.write_text(summary, encoding='utf-8')

# Reopen and verify all stated requirements from saved outputs.
wb2 = load_workbook(xlsx_path, data_only=True)
ws2 = wb2['failures']
saved_headers = [cell.value for cell in ws2[1]]
saved_rows = list(ws2.iter_rows(min_row=2, values_only=True))
if saved_headers != out_headers:
    raise AssertionError(f'Unexpected headers in failures.xlsx: {saved_headers} != {out_headers}')
if len(saved_rows) != len(fail_rows):
    raise AssertionError(f'Unexpected saved fail row count: {len(saved_rows)} != {len(fail_rows)}')
status_idx = saved_headers.index('Status')
value_idx = saved_headers.index('Value')
urgent_idx = saved_headers.index('Urgent')
for saved in saved_rows:
    if str(saved[status_idx]).strip().lower() != 'fail':
        raise AssertionError(f'Non-Fail row found in output: {saved}')
    value_num = float(saved[value_idx])
    urgent = saved[urgent_idx]
    if not isinstance(urgent, bool):
        raise AssertionError(f'Urgent is not boolean in output row: {saved}')
    if urgent != (value_num > threshold):
        raise AssertionError(f'Urgent rule mismatch in output row: {saved}')
summary_saved = summary_path.read_text(encoding='utf-8')
for expected in [f'Fail records kept: {len(fail_rows)}', f'Urgent failures: {urgent_count}', f'Value > {threshold:g}']:
    if expected not in summary_saved:
        raise AssertionError(f'Missing expected summary text: {expected}')

print('Wrote:', xlsx_path)
print('Wrote:', summary_path)
print('Verification passed: output files exist, headers preserved with appended Urgent, all output rows are Fail, Urgent booleans follow Value > threshold, and summary counts match.')
print(json.dumps({'failures_xlsx': str(xlsx_path), 'summary_md': str(summary_path), 'total_records': len(rows), 'fail_records': len(fail_rows), 'urgent_failures': urgent_count}, ensure_ascii=False))