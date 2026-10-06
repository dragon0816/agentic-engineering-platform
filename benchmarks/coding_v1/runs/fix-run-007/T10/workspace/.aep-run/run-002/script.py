from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
import re
from openpyxl import Workbook, load_workbook

CATALOG_URL = 'http://127.0.0.1:57997/t10/catalog'
OUT = Path('output') / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_main = False
        self.main_depth = 0
        self.current_card = None
        self.in_a = False
        self.in_price = False
        self.products = []

    def handle_starttag(self, tag, attrs):
        attrs_dict = dict(attrs)
        if tag == 'main' and attrs_dict.get('id') == 'catalog' and not self.in_main:
            self.in_main = True
            self.main_depth = 1
            return
        if self.in_main:
            self.main_depth += 1
            if tag == 'div' and 'card' in attrs_dict.get('class', '').split():
                self.current_card = {'name_parts': [], 'price_parts': [], 'href': None}
            elif self.current_card is not None and tag == 'a':
                self.in_a = True
                self.current_card['href'] = attrs_dict.get('href')
            elif self.current_card is not None and tag == 'span' and 'price' in attrs_dict.get('class', '').split():
                self.in_price = True

    def handle_endtag(self, tag):
        if self.in_main:
            if self.current_card is not None and tag == 'a':
                self.in_a = False
            elif self.current_card is not None and tag == 'span':
                self.in_price = False
            elif self.current_card is not None and tag == 'div':
                name = ''.join(self.current_card['name_parts']).strip()
                price_text = ''.join(self.current_card['price_parts']).strip()
                href = self.current_card['href']
                if name and price_text and href:
                    self.products.append({'name': name, 'price_text': price_text, 'href': href})
                self.current_card = None
            if tag == 'main':
                self.in_main = False
                self.main_depth = 0
            else:
                self.main_depth -= 1

    def handle_data(self, data):
        if self.current_card is not None:
            if self.in_a:
                self.current_card['name_parts'].append(data)
            elif self.in_price:
                self.current_card['price_parts'].append(data)


def parse_price(text):
    s = text.strip()
    m = re.match(r'^([^\d\-+.,]*)([\d,]+(?:\.\d+)?)', s)
    if not m:
        raise ValueError(f'Cannot parse price: {text!r}')
    currency = m.group(1).strip()
    num_text = m.group(2).replace(',', '')
    price = float(num_text) if '.' in num_text else int(num_text)
    return currency, price

# Fetch and parse the catalog page.
with urlopen(CATALOG_URL, timeout=10) as resp:
    html = resp.read().decode('utf-8', errors='replace')

parser = CatalogParser()
parser.feed(html)
rows = []
for item in parser.products:
    currency, price = parse_price(item['price_text'])
    rows.append({
        '產品名稱': item['name'],
        '價格': price,
        '幣別': currency,
        '產品網址': urljoin(CATALOG_URL, item['href']),
    })

if not rows:
    raise RuntimeError('No products found in <main id="catalog">')

# Write output workbook.
OUT.parent.mkdir(parents=True, exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(OUT)
print(f'WROTE {OUT} with {len(rows)} products')
for row in rows:
    print(row)

# Verification: reopen workbook and check all stated requirements.
wb2 = load_workbook(OUT, data_only=True)
ws2 = wb2['products']
read_headers = [cell.value for cell in ws2[1]]
assert read_headers == headers, f'Headers mismatch: {read_headers}'
read_rows = list(ws2.iter_rows(min_row=2, values_only=True))
assert len(read_rows) == 5, f'Expected 5 catalog products, got {len(read_rows)}'
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
expected_prices = [1299, 2450, 99, 560, 12000]
expected_urls = [urljoin(CATALOG_URL, p) for p in ['/p/router-a','/p/switch-b','/p/cable-c','/p/adapter-d','/p/antenna-e']]
for i, (name, price, currency, product_url) in enumerate(read_rows):
    assert name == expected_names[i], f'Row {i+2} name/order mismatch: {name}'
    assert isinstance(price, (int, float)) and not isinstance(price, bool), f'Row {i+2} price is not numeric: {price!r}'
    assert price == expected_prices[i], f'Row {i+2} price mismatch: {price}'
    assert currency == 'NT$', f'Row {i+2} currency mismatch: {currency}'
    assert product_url == expected_urls[i], f'Row {i+2} URL mismatch: {product_url}'
    assert product_url.startswith('http://127.0.0.1:57997/'), f'Row {i+2} URL is not full: {product_url}'
assert 'Sponsored Router' not in expected_names and '推薦：高速網卡' not in expected_names
print('CHECK OK: products.xlsx reopened; headers, 5-row page order, numeric prices, NT$ currency, full URLs, and ad/recommendation exclusion verified.')