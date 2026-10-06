from pathlib import Path
import csv
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'large_results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

required_headers = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']
summary_headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
overall_headers = ['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']

stats = defaultdict(lambda: {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': Decimal('0')})
overall = {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': Decimal('0')}
source_row_count = 0

with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    print('Input headers:', headers)
    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise ValueError(f'Missing required headers: {missing}')

    for line_no, row in enumerate(reader, start=2):
        source_row_count += 1
        model = row['Device Model']
        status = row['Status']
        try:
            duration = Decimal(str(row['DurationSeconds']))
        except (InvalidOperation, TypeError) as e:
            raise ValueError(f'Invalid DurationSeconds at CSV line {line_no}: {row["DurationSeconds"]!r}') from e

        stats[model]['total'] += 1
        stats[model]['duration_sum'] += duration
        overall['total'] += 1
        overall['duration_sum'] += duration

        if status == 'Pass':
            stats[model]['pass'] += 1
            overall['pass'] += 1
        elif status == 'Fail':
            stats[model]['fail'] += 1
            overall['fail'] += 1
        else:
            raise ValueError(f'Unexpected Status at CSV line {line_no}: {status!r}')

print('Source rows:', source_row_count)
print('Device model count:', len(stats))

wb = Workbook()
ws = wb.active
ws.title = 'Summary'
ws.append(summary_headers)
for cell in ws[1]:
    cell.font = Font(bold=True)

# Sort model names naturally when they have trailing numbers (M0, M1, ...), else lexicographically.
def sort_key(model):
    s = str(model)
    prefix = s.rstrip('0123456789')
    suffix = s[len(prefix):]
    if suffix:
        return (prefix, int(suffix))
    return (s, -1)

for model in sorted(stats.keys(), key=sort_key):
    s = stats[model]
    total = s['total']
    pass_count = s['pass']
    fail_count = s['fail']
    pass_rate = pass_count / total if total else 0
    avg_duration = float(s['duration_sum'] / total) if total else 0
    ws.append([model, total, pass_count, fail_count, pass_rate, avg_duration])

# Format rate and average columns
for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
    row[0].number_format = '0.00%'
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    row[0].number_format = '0.00'

ws_overall = wb.create_sheet('Overall')
ws_overall.append(overall_headers)
for cell in ws_overall[1]:
    cell.font = Font(bold=True)

total = overall['total']
pass_count = overall['pass']
fail_count = overall['fail']
pass_rate = pass_count / total if total else 0
avg_duration = float(overall['duration_sum'] / total) if total else 0
ws_overall.append([total, pass_count, fail_count, pass_rate, avg_duration])
ws_overall['D2'].number_format = '0.00%'
ws_overall['E2'].number_format = '0.00'

# Simple column widths
for sheet in (ws, ws_overall):
    for column_cells in sheet.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        sheet.column_dimensions[column_cells[0].column_letter].width = max(max_len + 2, 12)

wb.save(output_path)
print('Wrote:', output_path)

# Reopen output and verify every stated requirement.
check_wb = load_workbook(output_path, data_only=True)
print('Output sheets:', check_wb.sheetnames)
if 'Summary' not in check_wb.sheetnames:
    raise AssertionError('Missing Summary sheet')
if 'Overall' not in check_wb.sheetnames:
    raise AssertionError('Missing Overall sheet')

check_summary = check_wb['Summary']
check_overall = check_wb['Overall']
actual_summary_headers = [check_summary.cell(row=1, column=i).value for i in range(1, len(summary_headers)+1)]
actual_overall_headers = [check_overall.cell(row=1, column=i).value for i in range(1, len(overall_headers)+1)]
print('Summary headers:', actual_summary_headers)
print('Overall headers:', actual_overall_headers)
if actual_summary_headers != summary_headers:
    raise AssertionError(f'Summary headers mismatch: {actual_summary_headers}')
if actual_overall_headers != overall_headers:
    raise AssertionError(f'Overall headers mismatch: {actual_overall_headers}')

summary_data_rows = check_summary.max_row - 1
if summary_data_rows != len(stats):
    raise AssertionError(f'Summary data row count {summary_data_rows} != model count {len(stats)}')

# Build expected values and compare cell values/types.
expected_by_model = {}
for model, s in stats.items():
    total = s['total']
    expected_by_model[model] = {
        'Total': total,
        'PassCount': s['pass'],
        'FailCount': s['fail'],
        'PassRate': s['pass'] / total if total else 0,
        'AvgDurationSeconds': float(s['duration_sum'] / total) if total else 0,
    }

summary_total_sum = 0
summary_pass_sum = 0
summary_fail_sum = 0
for r in range(2, check_summary.max_row + 1):
    model = check_summary.cell(r, 1).value
    vals = {
        'Total': check_summary.cell(r, 2).value,
        'PassCount': check_summary.cell(r, 3).value,
        'FailCount': check_summary.cell(r, 4).value,
        'PassRate': check_summary.cell(r, 5).value,
        'AvgDurationSeconds': check_summary.cell(r, 6).value,
    }
    if model not in expected_by_model:
        raise AssertionError(f'Unexpected model in Summary: {model!r}')
    exp = expected_by_model[model]
    for key in ['Total', 'PassCount', 'FailCount']:
        if not isinstance(vals[key], int):
            raise AssertionError(f'{model} {key} is not int: {vals[key]!r}')
        if vals[key] != exp[key]:
            raise AssertionError(f'{model} {key} {vals[key]} != {exp[key]}')
    for key in ['PassRate', 'AvgDurationSeconds']:
        if not isinstance(vals[key], (int, float)):
            raise AssertionError(f'{model} {key} is not numeric: {vals[key]!r}')
        if abs(vals[key] - exp[key]) > 1e-12:
            raise AssertionError(f'{model} {key} {vals[key]} != {exp[key]}')
    summary_total_sum += vals['Total']
    summary_pass_sum += vals['PassCount']
    summary_fail_sum += vals['FailCount']

if summary_total_sum != source_row_count:
    raise AssertionError(f'Summary Total sum {summary_total_sum} != source rows {source_row_count}')
if summary_pass_sum != overall['pass'] or summary_fail_sum != overall['fail']:
    raise AssertionError('Summary pass/fail sums do not match overall counts')

overall_vals = {
    'Total': check_overall.cell(2, 1).value,
    'PassCount': check_overall.cell(2, 2).value,
    'FailCount': check_overall.cell(2, 3).value,
    'PassRate': check_overall.cell(2, 4).value,
    'AvgDurationSeconds': check_overall.cell(2, 5).value,
}
expected_overall = {
    'Total': overall['total'],
    'PassCount': overall['pass'],
    'FailCount': overall['fail'],
    'PassRate': overall['pass'] / overall['total'] if overall['total'] else 0,
    'AvgDurationSeconds': float(overall['duration_sum'] / overall['total']) if overall['total'] else 0,
}
for key in ['Total', 'PassCount', 'FailCount']:
    if not isinstance(overall_vals[key], int):
        raise AssertionError(f'Overall {key} is not int: {overall_vals[key]!r}')
    if overall_vals[key] != expected_overall[key]:
        raise AssertionError(f'Overall {key} {overall_vals[key]} != {expected_overall[key]}')
for key in ['PassRate', 'AvgDurationSeconds']:
    if not isinstance(overall_vals[key], (int, float)):
        raise AssertionError(f'Overall {key} is not numeric: {overall_vals[key]!r}')
    if abs(overall_vals[key] - expected_overall[key]) > 1e-12:
        raise AssertionError(f'Overall {key} {overall_vals[key]} != {expected_overall[key]}')

print('Verified output file:', output_path)
print('Verified Summary sheet columns:', summary_headers)
print('Verified Overall sheet columns:', overall_headers)
print('Verified Summary model rows:', summary_data_rows)
print('Verified source rows represented:', source_row_count)
print('Verified Overall values:', expected_overall)
