from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV without modifying input.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames
    if headers is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in headers:
        raise ValueError('Required column Status not found')
    if 'Device ID' not in headers:
        raise ValueError('Required column Device ID not found')
    failed_rows = [row for row in reader if (row.get('Status') or '').strip().lower() == 'fail']

# Create workbook with exact requested sheet name.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

# Headers: keep source column names and order.
for col_idx, header in enumerate(headers, start=1):
    cell = ws.cell(row=1, column=col_idx, value=header)
    if header == 'Device ID':
        cell.number_format = '@'

# Data: keep all columns, original order; Device ID as text to preserve leading zeros.
for row_idx, row in enumerate(failed_rows, start=2):
    for col_idx, header in enumerate(headers, start=1):
        value = row.get(header, '')
        cell = ws.cell(row=row_idx, column=col_idx)
        if header == 'Device ID':
            cell.value = '' if value is None else str(value)
            cell.number_format = '@'
        else:
            cell.value = value

wb.save(output_path)
print(f'Wrote {output_path}')
print(f'Failed row count: {len(failed_rows)}')

# Verification: reopen the saved workbook and confirm requirements.
wb2 = load_workbook(output_path, data_only=False)
if 'failed_items' not in wb2.sheetnames:
    raise AssertionError('Sheet failed_items missing')
ws2 = wb2['failed_items']

saved_headers = [ws2.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
if saved_headers != headers:
    raise AssertionError(f'Headers mismatch: {saved_headers} != {headers}')

saved_row_count = ws2.max_row - 1
if saved_row_count != len(failed_rows):
    raise AssertionError(f'Row count mismatch: {saved_row_count} != {len(failed_rows)}')

# Expected filtered rows from source order.
expected_device_ids = [str(row['Device ID']) for row in failed_rows]
expected_statuses = [row['Status'] for row in failed_rows]
saved_device_ids = []
saved_statuses = []
device_col = headers.index('Device ID') + 1
status_col = headers.index('Status') + 1
for r in range(2, ws2.max_row + 1):
    dev_cell = ws2.cell(row=r, column=device_col)
    status_cell = ws2.cell(row=r, column=status_col)
    saved_device_ids.append(dev_cell.value)
    saved_statuses.append(status_cell.value)
    if dev_cell.data_type != 's':
        raise AssertionError(f'Device ID at row {r} is not stored as text: data_type={dev_cell.data_type}, value={dev_cell.value!r}')
    if dev_cell.number_format != '@':
        raise AssertionError(f'Device ID at row {r} does not have text number format: {dev_cell.number_format!r}')

if saved_device_ids != expected_device_ids:
    raise AssertionError(f'Device ID/order mismatch: {saved_device_ids} != {expected_device_ids}')
if saved_statuses != expected_statuses:
    raise AssertionError(f'Status/order mismatch: {saved_statuses} != {expected_statuses}')
for s in saved_statuses:
    if (s or '').strip().lower() != 'fail':
        raise AssertionError(f'Non-fail status included: {s!r}')

print('Verification passed:')
print(' - output file exists:', output_path.exists())
print(' - sheet name: failed_items')
print(' - headers:', saved_headers)
print(' - rows:', saved_row_count)
print(' - Device ID values/text:', saved_device_ids)
print(' - Status values:', saved_statuses)
print(' - all output rows satisfy Status.strip().lower() == fail')