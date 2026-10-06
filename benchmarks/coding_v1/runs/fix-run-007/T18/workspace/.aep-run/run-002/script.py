import os, csv, json, subprocess, sys
from pathlib import Path

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

print('WORKSPACE:', workspace)
print('INPUT FILES:')
for p in sorted(input_dir.iterdir()):
    print(' -', p.name, 'size=', p.stat().st_size)

# Inspect dependencies
_dep_path = input_dir / 'DEPENDENCIES.txt'
if _dep_path.exists():
    print('\nDEPENDENCIES.txt:')
    print(_dep_path.read_text(encoding='utf-8', errors='replace'))
else:
    print('\nDEPENDENCIES.txt: MISSING')

# Inspect CSV headers and sample rows
csv_path = input_dir / 'input.csv'
print('\nCSV INSPECTION:')
with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    print('headers:', reader.fieldnames)
    sample = []
    for i, row in enumerate(reader):
        if i < 8:
            sample.append(row)
        else:
            break
    print('sample_rows:', json.dumps(sample, ensure_ascii=False, indent=2))

# Inspect original script enough to identify the likely failure
orig_path = input_dir / 'export.py'
print('\nORIGINAL export.py:')
orig_text = orig_path.read_text(encoding='utf-8', errors='replace')
print(orig_text[:4000])

# Compute expected result from input.csv: blank Value ignored; group by Model in first-seen order.
expected = {}
order = []
invalid_rows = []
with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    if reader.fieldnames is None:
        raise RuntimeError('input.csv has no header row')
    required = {'Model', 'Value'}
    missing = required - set(reader.fieldnames)
    if missing:
        raise RuntimeError(f'input.csv missing required columns: {sorted(missing)}')
    for line_no, row in enumerate(reader, start=2):
        model = row.get('Model')
        value_raw = row.get('Value')
        if model not in expected:
            expected[model] = []
            order.append(model)
        if value_raw is None or str(value_raw).strip() == '':
            continue
        try:
            value = float(str(value_raw).strip())
        except ValueError:
            invalid_rows.append((line_no, model, value_raw))
            continue
        expected[model].append(value)

if invalid_rows:
    print('WARNING invalid nonblank Value rows ignored by validation expectation:', invalid_rows)

expected_rows = []
for model in order:
    vals = expected[model]
    avg = (sum(vals) / len(vals)) if vals else None
    expected_rows.append({'Model': model, 'Average': avg, 'ValidCount': len(vals)})
print('\nEXPECTED SUMMARY:', json.dumps(expected_rows, ensure_ascii=False, indent=2))

# Write fixed export.py to output/export.py. CLI unchanged: --input and --output.
fixed_code = '''#!/usr/bin/env python3
"""Read input.csv, calculate average Value by Model, and write summary.xlsx.

CLI intentionally remains compatible with the original interface:
    python export.py --input input.csv --output summary.xlsx
"""

from __future__ import annotations

import argparse
import csv
from collections import OrderedDict
from pathlib import Path
from typing import Optional

from openpyxl import Workbook


def _parse_value(value: object) -> Optional[float]:
    """Return a numeric Value, or None when the source cell is blank.

    Blank Value cells are excluded from both Average and ValidCount. Nonblank
    nonnumeric values are rejected because silently treating them as zero would
    corrupt the average.
    """
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return None
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"Value must be numeric or blank, got {value!r}") from exc


def build_summary(input_csv: str | Path) -> list[tuple[str, Optional[float], int]]:
    groups: "OrderedDict[str, list[float]]" = OrderedDict()
    input_csv = Path(input_csv)

    with input_csv.open("r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None:
            raise ValueError("Input CSV is missing a header row")
        missing = {"Model", "Value"} - set(reader.fieldnames)
        if missing:
            raise ValueError(f"Input CSV missing required column(s): {', '.join(sorted(missing))}")

        for row in reader:
            model = row["Model"]
            groups.setdefault(model, [])
            parsed = _parse_value(row.get("Value"))
            if parsed is not None:
                groups[model].append(parsed)

    summary: list[tuple[str, Optional[float], int]] = []
    for model, values in groups.items():
        average = (sum(values) / len(values)) if values else None
        summary.append((model, average, len(values)))
    return summary


def write_summary_xlsx(rows: list[tuple[str, Optional[float], int]], output_xlsx: str | Path) -> None:
    output_xlsx = Path(output_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "summary"
    ws.append(["Model", "Average", "ValidCount"])
    for model, average, valid_count in rows:
        ws.append([model, average, valid_count])
    wb.save(output_xlsx)


def main() -> None:
    parser = argparse.ArgumentParser(description="Summarize average Value by Model")
    parser.add_argument("--input", required=True, help="Path to input CSV")
    parser.add_argument("--output", required=True, help="Path to output XLSX")
    args = parser.parse_args()

    rows = build_summary(args.input)
    write_summary_xlsx(rows, args.output)


if __name__ == "__main__":
    main()
'''

fixed_path = output_dir / 'export.py'
fixed_path.write_text(fixed_code, encoding='utf-8', newline='')
print('\nWROTE:', fixed_path, 'size=', fixed_path.stat().st_size)

# Run the requested command from workspace.
summary_path = output_dir / 'summary.xlsx'
if summary_path.exists():
    summary_path.unlink()
cmd = [sys.executable, 'output\\export.py', '--input', 'input\\input.csv', '--output', 'output\\summary.xlsx']
print('\nRUNNING:', ' '.join(cmd))
proc = subprocess.run(cmd, cwd=workspace, text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
print('returncode:', proc.returncode)
print('stdout:', proc.stdout)
print('stderr:', proc.stderr)
if proc.returncode != 0:
    raise RuntimeError('fixed export.py command failed')

# Reopen output file and verify stated requirements.
from openpyxl import load_workbook
if not summary_path.exists():
    raise RuntimeError('summary.xlsx was not created')
wb = load_workbook(summary_path, data_only=True)
ws = wb.active
headers = [cell.value for cell in ws[1]]
print('\nOUTPUT SHEET:', ws.title)
print('OUTPUT HEADERS:', headers)
if headers != ['Model', 'Average', 'ValidCount']:
    raise AssertionError(f'Unexpected headers: {headers!r}')

actual_rows = []
for row in ws.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    actual_rows.append({'Model': row[0], 'Average': row[1], 'ValidCount': row[2]})
print('OUTPUT ROWS:', json.dumps(actual_rows, ensure_ascii=False, indent=2))

if len(actual_rows) != len(expected_rows):
    raise AssertionError(f'Row count mismatch: actual {len(actual_rows)} expected {len(expected_rows)}')

for idx, (actual, exp) in enumerate(zip(actual_rows, expected_rows), start=2):
    if actual['Model'] != exp['Model']:
        raise AssertionError(f'Row {idx} Model mismatch: {actual["Model"]!r} != {exp["Model"]!r}')
    if not isinstance(actual['ValidCount'], int):
        raise AssertionError(f'Row {idx} ValidCount is not int: {actual["ValidCount"]!r}')
    if actual['ValidCount'] != exp['ValidCount']:
        raise AssertionError(f'Row {idx} ValidCount mismatch: {actual["ValidCount"]!r} != {exp["ValidCount"]!r}')
    if exp['Average'] is None:
        if actual['Average'] is not None:
            raise AssertionError(f'Row {idx} Average expected blank/None, got {actual["Average"]!r}')
    else:
        if not isinstance(actual['Average'], (int, float)):
            raise AssertionError(f'Row {idx} Average is not numeric: {actual["Average"]!r}')
        if abs(float(actual['Average']) - float(exp['Average'])) > 1e-9:
            raise AssertionError(f'Row {idx} Average mismatch: {actual["Average"]!r} != {exp["Average"]!r}')

print('\nVERIFICATION PASSED:')
print(' - output/export.py exists and preserves --input/--output CLI')
print(' - command executed from workspace: python output\\export.py --input input\\input.csv --output output\\summary.xlsx')
print(' - summary.xlsx reopened successfully')
print(' - headers exactly Model, Average, ValidCount')
print(' - row count and per-Model averages/ValidCount match input.csv with blank Value excluded')
print(' - Average cells are numeric when count > 0; ValidCount cells are integers')
