from pathlib import Path
import re
from decimal import Decimal
from openpyxl import Workbook, load_workbook

ws = Path.cwd()
input_path = ws / 'input' / 'specs.md'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split markdown into product sections beginning with level-2 headings.
sections = []
current_title = None
current_lines = []
for line in text.splitlines():
    m = re.match(r'^##\s+(.+?)\s*$', line)
    if m:
        if current_title is not None:
            sections.append((current_title, '\n'.join(current_lines)))
        current_title = m.group(1).strip()
        current_lines = []
    elif current_title is not None:
        current_lines.append(line)
if current_title is not None:
    sections.append((current_title, '\n'.join(current_lines)))

unit_factor = {'mhz': Decimal('1'), 'ghz': Decimal('1000')}

def dec_to_number(d: Decimal):
    """Return int when exact integer, otherwise float for Excel numeric cell."""
    d = d.normalize()
    if d == d.to_integral():
        return int(d)
    return float(d)

def parse_kv(section_text):
    data = {}
    for line in section_text.splitlines():
        m = re.match(r'^\s*-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            data[m.group(1).strip()] = m.group(2).strip()
    return data

def parse_frequency(value):
    if not value:
        return None, None
    # Supports ranges such as "2.4–2.5 GHz", "900-930 MHz". Uses only written units.
    s = value.strip().replace('—', '–').replace('−', '-').replace(' to ', '–')
    m = re.match(r'^\s*([0-9]+(?:\.[0-9]+)?)\s*(?:-|–)\s*([0-9]+(?:\.[0-9]+)?)\s*([A-Za-z]+)\s*$', s)
    if not m:
        return None, None
    lo = Decimal(m.group(1))
    hi = Decimal(m.group(2))
    unit = m.group(3).lower()
    if unit not in unit_factor:
        return None, None
    factor = unit_factor[unit]
    return dec_to_number(lo * factor), dec_to_number(hi * factor)

def parse_bandwidth(value):
    if not value:
        return None
    m = re.match(r'^\s*([0-9]+(?:\.[0-9]+)?)\s*([A-Za-z]+)\s*$', value.strip())
    if not m:
        return None
    unit = m.group(2).lower()
    if unit not in unit_factor:
        return None
    return dec_to_number(Decimal(m.group(1)) * unit_factor[unit])

def parse_interface(value):
    if not value:
        return ''
    # Multiple interfaces in the source are separated by slash; output uses semicolon.
    parts = [p.strip() for p in re.split(r'\s*/\s*', value.strip()) if p.strip()]
    return ';'.join(parts)

rows = []
for title, body in sections:
    kv = parse_kv(body)
    fmin, fmax = parse_frequency(kv.get('Frequency', ''))
    rows.append({
        'Model': kv.get('Model', ''),
        'FrequencyMinMHz': fmin,
        'FrequencyMaxMHz': fmax,
        'BandwidthMHz': parse_bandwidth(kv.get('Bandwidth', '')),
        'Interface': parse_interface(kv.get('Interface', '')),
    })

# Produce workbook.
wb = Workbook()
sh = wb.active
sh.title = 'specs'
headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
sh.append(headers)
for row in rows:
    sh.append([row[h] if row[h] is not None else '' for h in headers])
wb.save(output_path)

# Verification: reopen the saved file and check stated requirements.
check_wb = load_workbook(output_path, data_only=True)
check_sh = check_wb['specs']
read_headers = [check_sh.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
assert read_headers == headers, f'Headers mismatch: {read_headers}'
assert check_sh.max_row == 4, f'Expected 3 data rows plus header, got {check_sh.max_row}'
expected = [
    ['RX-100', 2400, 2500, 80, 'USB;LAN'],
    ['RX-200', 900, 930, 20, 'LAN'],
    ['RX-300', 5000, 6000, None, 'USB'],
]
actual = []
for r in range(2, check_sh.max_row + 1):
    actual.append([check_sh.cell(row=r, column=c).value for c in range(1, len(headers)+1)])
assert actual == expected, f'Rows mismatch: {actual}'
for r in range(2, check_sh.max_row + 1):
    assert isinstance(check_sh.cell(r, 1).value, str), 'Model must be text'
    assert isinstance(check_sh.cell(r, 2).value, (int, float)), 'FrequencyMinMHz must be numeric'
    assert isinstance(check_sh.cell(r, 3).value, (int, float)), 'FrequencyMaxMHz must be numeric'
    bw = check_sh.cell(r, 4).value
    assert bw is None or isinstance(bw, (int, float)), 'BandwidthMHz must be numeric or blank'
    assert isinstance(check_sh.cell(r, 5).value, str), 'Interface must be text'

print('Created:', output_path)
print('Rows written:', len(rows))
print('Headers verified:', headers)
print('Data verified:')
for row in actual:
    print(row)
print('Confirmed missing Bandwidth for RX-300 is blank (empty Excel cell).')