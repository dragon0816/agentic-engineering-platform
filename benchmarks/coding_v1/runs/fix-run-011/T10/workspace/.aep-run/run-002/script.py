from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
from html import unescape
import re
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
CATALOG_URL = 'http://127.0.0.1:54763/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_main_catalog = False
        self.main_depth = 0
        self.in_card = False
        self.card_depth = 0
        self.current = None
        self.current_field = None
        self.products = []

    def _has_class(self, attrs, cls):
        classes = ''
        for k, v in attrs:
            if k == 'class' and v:
                classes = v
        return cls in classes.split()

    def _attr(self, attrs, name):
        for k, v in attrs:
            if k == name:
                return v
        return None

    def handle_starttag(self, tag, attrs):
        attrs = list(attrs)
        if self.in_main_catalog:
            self.main_depth += 1
            if tag == 'div' and self._has_class(attrs, 'card') and not self.in_card:
                self.in_card = True
                self.card_depth = 1
                self.current = {'name': '', 'href': '', 'price_text': ''}
                self.current_field = None
                return
        else:
            if tag == 'main' and self._attr(attrs, 'id') == 'catalog':
                self.in_main_catalog = True
                self.main_depth = 1
                return

        if self.in_card:
            if tag == 'a':
                href = self._attr(attrs, 'href')
                if href:
                    self.current['href'] = href
                self.current_field = 'name'
            elif tag == 'span' and self._has_class(attrs, 'price'):
                self.current_field = 'price_text'
            elif not (tag == 'div' and self._has_class(attrs, 'card')):
                pass
            # Count nested tags inside card except the div.card start handled above.
            if not (tag == 'div' and self._has_class(attrs, 'card')):
                self.card_depth += 1

    def handle_endtag(self, tag):
        if self.in_card:
            if tag in ('a', 'span'):
                self.current_field = None
            self.card_depth -= 1
            if self.card_depth <= 0:
                name = ' '.join(unescape(self.current.get('name', '')).split())
                href = self.current.get('href', '').strip()
                price_text = ' '.join(unescape(self.current.get('price_text', '')).split())
                if name and href and price_text:
                    self.products.append({'name': name, 'href': href, 'price_text': price_text})
                self.in_card = False
                self.current = None
                self.current_field = None
                return

        if self.in_main_catalog:
            self.main_depth -= 1
            if self.main_depth <= 0:
                self.in_main_catalog = False
                self.main_depth = 0

    def handle_data(self, data):
        if self.in_card and self.current_field and self.current is not None:
            self.current[self.current_field] += data

def parse_price(price_text):
    text = price_text.strip()
    m = re.match(r'^([^\d\-+.,]+)\s*([\d,]+(?:\.\d+)?)$', text)
    if not m:
        raise ValueError(f'Cannot parse price: {price_text!r}')
    currency = m.group(1).strip()
    number = float(m.group(2).replace(',', ''))
    if number.is_integer():
        number = int(number)
    return currency, number

with urlopen(CATALOG_URL, timeout=10) as response:
    html = response.read().decode('utf-8', errors='replace')
    base_url = response.geturl()

parser = CatalogParser()
parser.feed(html)

rows = []
for item in parser.products:
    currency, price = parse_price(item['price_text'])
    rows.append({
        '產品名稱': item['name'],
        '價格': price,
        '幣別': currency,
        '產品網址': urljoin(base_url, item['href']),
    })

wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(OUTFILE)

# Reopen and verify saved workbook.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
saved_headers = [ws2.cell(1, c).value for c in range(1, len(headers)+1)]
saved_rows = []
for r in range(2, ws2.max_row + 1):
    saved_rows.append([ws2.cell(r, c).value for c in range(1, len(headers)+1)])

expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
expected_prices = [1299, 2450, 99, 560, 12000]
expected_urls = [urljoin(base_url, href) for href in ['/p/router-a','/p/switch-b','/p/cable-c','/p/adapter-d','/p/antenna-e']]
assert saved_headers == headers, f'headers mismatch: {saved_headers}'
assert len(saved_rows) == 5, f'row count mismatch: {len(saved_rows)}'
assert [r[0] for r in saved_rows] == expected_names, f'order/name mismatch: {saved_rows}'
assert [r[1] for r in saved_rows] == expected_prices, f'price mismatch: {saved_rows}'
assert all(isinstance(r[1], (int, float)) for r in saved_rows), 'prices are not numeric'
assert [r[2] for r in saved_rows] == ['NT$'] * 5, f'currency mismatch: {saved_rows}'
assert [r[3] for r in saved_rows] == expected_urls, f'url mismatch: {saved_rows}'
assert 'Sponsored Router' not in [r[0] for r in saved_rows], 'ad included'
assert '推薦：高速網卡' not in [r[0] for r in saved_rows], 'recommendation included'

print('Created:', OUTFILE)
print('Rows written:', len(saved_rows))
print('Headers:', saved_headers)
print('Verified page order, numeric prices, currency, full URLs, and exclusion of ads/recommendations.')
for row in saved_rows:
    print(row)
