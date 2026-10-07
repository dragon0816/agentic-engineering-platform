from pathlib import Path
from urllib.request import urlopen, Request
from urllib.parse import urljoin
from html.parser import HTMLParser
from html import unescape
import re
from openpyxl import Workbook, load_workbook

BASE_URL = 'http://127.0.0.1:54763/t11/list'
OUTPUT = Path('output') / 'details.xlsx'
HEADERS = ['Name', 'Price', 'Currency', 'CPU', 'MemoryGB', 'URL']

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._a = None
        self._in_products = False
        self._products_depth = 0

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        tag = tag.lower()
        if tag == 'ul' and attrs.get('id') == 'products':
            self._in_products = True
            self._products_depth = 1
        elif self._in_products:
            self._products_depth += 1
        if tag == 'a' and self._in_products:
            self._a = {'href': attrs.get('href', ''), 'text': ''}

    def handle_data(self, data):
        if self._a is not None:
            self._a['text'] += data

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == 'a' and self._a is not None:
            if self._a.get('href'):
                self.links.append(self._a)
            self._a = None
        if self._in_products:
            self._products_depth -= 1
            if self._products_depth <= 0:
                self._in_products = False

class DetailParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.name = ''
        self.spec = {}
        self._capture_name = False
        self._in_table = False
        self._table_depth = 0
        self._cell = None
        self._cell_text = ''
        self._row = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        tag = tag.lower()
        classes = attrs.get('class', '').split()
        if tag == 'h1' and 'name' in classes:
            self._capture_name = True
            self._cell_text = ''
        if tag == 'table' and 'spec' in classes:
            self._in_table = True
            self._table_depth = 1
        elif self._in_table:
            self._table_depth += 1
        if self._in_table and tag == 'tr':
            self._row = []
        if self._in_table and tag in ('th', 'td'):
            self._cell = tag
            self._cell_text = ''

    def handle_data(self, data):
        if self._capture_name:
            self._cell_text += data
        elif self._cell is not None:
            self._cell_text += data

    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag == 'h1' and self._capture_name:
            self.name = clean_text(self._cell_text)
            self._capture_name = False
            self._cell_text = ''
        if self._in_table and tag in ('th', 'td') and self._cell is not None:
            self._row.append((self._cell, clean_text(self._cell_text)))
            self._cell = None
            self._cell_text = ''
        if self._in_table and tag == 'tr':
            ths = [v for t, v in self._row if t == 'th']
            tds = [v for t, v in self._row if t == 'td']
            if ths and tds:
                self.spec[ths[0]] = tds[0]
            self._row = []
        if self._in_table:
            self._table_depth -= 1
            if self._table_depth <= 0:
                self._in_table = False

def clean_text(s):
    return re.sub(r'\s+', ' ', unescape(s or '')).strip()

def fetch(url):
    return urlopen(Request(url, headers={'User-Agent': 'Mozilla/5.0'}), timeout=10).read().decode('utf-8', errors='replace')

def normalize_label(s):
    return re.sub(r'[^a-z0-9]', '', (s or '').lower())

def parse_price_currency(value):
    value = clean_text(value)
    if not value:
        return None, None
    # Currency is everything before the first digit, after stripping spaces.
    m = re.search(r'[-+]?\d', value)
    currency = None
    number_part = value
    if m:
        prefix = value[:m.start()].strip()
        suffix = value[m.end():]
        if prefix:
            currency = prefix
        else:
            # Fall back to trailing alphabetic currency codes, e.g. 999 USD.
            tail = re.search(r'\b([A-Z]{3})\b\s*$', value)
            if tail:
                currency = tail.group(1)
        number_part = value[m.start():]
    num_match = re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', number_part)
    price = None
    if num_match:
        n = num_match.group(0).replace(',', '')
        price = float(n) if '.' in n else int(n)
    return price, currency

def parse_memory_gb(value):
    value = clean_text(value)
    if not value:
        return None
    m = re.search(r'(\d+(?:\.\d+)?)\s*(TB|GB|MB|G|M|T)\b', value, flags=re.I)
    if not m:
        m = re.search(r'(\d+(?:\.\d+)?)', value)
        if not m:
            return None
        gb = float(m.group(1))
    else:
        amount = float(m.group(1))
        unit = m.group(2).upper()
        if unit in ('TB', 'T'):
            gb = amount * 1024
        elif unit in ('MB', 'M'):
            gb = amount / 1024
        else:
            gb = amount
    return int(gb) if gb.is_integer() else gb

def get_first(spec, labels):
    normalized = {normalize_label(k): v for k, v in spec.items()}
    for label in labels:
        key = normalize_label(label)
        if key in normalized:
            return normalized[key]
    return ''

# Read list and every detail page.
list_html = fetch(BASE_URL)
lp = ProductListParser(); lp.feed(list_html)
links = [(clean_text(l['text']), urljoin(BASE_URL, l['href'])) for l in lp.links]
print(f'Found {len(links)} product links')
for text, url in links:
    print(' -', text, url)

rows = []
for list_name, url in links:
    html = fetch(url)
    dp = DetailParser(); dp.feed(html)
    raw_price = get_first(dp.spec, ['Price'])
    price, currency = parse_price_currency(raw_price)
    cpu = get_first(dp.spec, ['CPU', 'Processor'])
    raw_memory = get_first(dp.spec, ['Memory', 'MemoryGB', 'RAM'])
    memory_gb = parse_memory_gb(raw_memory)
    row = {
        'Name': dp.name or list_name or None,
        'Price': price,
        'Currency': currency,
        'CPU': cpu or None,
        'MemoryGB': memory_gb,
        'URL': url,
        '_raw_spec': dp.spec,
        '_raw_price': raw_price,
        '_raw_memory': raw_memory,
    }
    rows.append(row)
    print('Parsed:', {k: row[k] for k in HEADERS}, 'raw_spec=', dp.spec)

# Write workbook.
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = 'details'
ws.append(HEADERS)
for row in rows:
    ws.append([row[h] for h in HEADERS])
for col in range(1, len(HEADERS) + 1):
    ws.cell(row=1, column=col).style = 'Headline 1'
wb.save(OUTPUT)
print('Wrote', OUTPUT)

# Verification by reopening saved workbook.
wb2 = load_workbook(OUTPUT, data_only=True)
ws2 = wb2['details']
actual_headers = [ws2.cell(row=1, column=i).value for i in range(1, len(HEADERS)+1)]
assert actual_headers == HEADERS, f'headers mismatch: {actual_headers}'
data_rows = list(ws2.iter_rows(min_row=2, values_only=True))
assert len(data_rows) == len(links), f'row count mismatch {len(data_rows)} != {len(links)}'
for idx, r in enumerate(data_rows, start=2):
    values = dict(zip(HEADERS, r))
    assert isinstance(values['Name'], str) and values['Name'], f'row {idx} missing Name'
    assert isinstance(values['URL'], str) and values['URL'].startswith('http://127.0.0.1:54763/t11/item/'), f'row {idx} bad URL'
    assert values['MemoryGB'] is None or isinstance(values['MemoryGB'], (int, float)), f'row {idx} MemoryGB not numeric/blank'
    assert values['Price'] is None or isinstance(values['Price'], (int, float)), f'row {idx} Price not numeric/blank'
print('Verified workbook: sheet=details, headers=', actual_headers)
print('Verified rows:', len(data_rows))
print('Verified Price and MemoryGB are numeric or blank; missing fields are blank cells.')
print('Saved absolute path:', (Path.cwd()/OUTPUT).resolve())
