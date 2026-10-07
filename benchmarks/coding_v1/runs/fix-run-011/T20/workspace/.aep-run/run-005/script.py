from pathlib import Path
import csv
import json
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
MANIFEST = INPUT_DIR / 'manifest.json'
SOP = INPUT_DIR / 'supplied_sop.md'
THRESHOLD = Decimal('20')

print(f'Workspace: {WORKSPACE}')
print(f'Input directory: {INPUT_DIR}')
print(f'Output directory: {OUTPUT_DIR}')

if not INPUT_DIR.exists():
    raise FileNotFoundError(f'Missing input directory: {INPUT_DIR}')
if not MANIFEST.exists():
    raise FileNotFoundError(f'Missing manifest: {MANIFEST}')
if not SOP.exists():
    raise FileNotFoundError(f'Missing SOP: {SOP}')

manifest = json.loads(MANIFEST.read_text(encoding='utf-8'))
print('SOP file found:', SOP.as_posix())
print('Manifest:', manifest)

csv_rel = manifest.get('results.csv')
if not csv_rel:
    raise KeyError("manifest.json must contain key 'results.csv'")
csv_path = (WORKSPACE / csv_rel).resolve()
input_root = INPUT_DIR.resolve()
if input_root not in csv_path.parents and csv_path != input_root:
    raise ValueError(f'Manifest CSV path is not under input/: {csv_path}')
if not csv_path.exists():
    raise FileNotFoundError(f'Missing CSV from manifest: {csv_path}')

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    fieldnames = reader.fieldnames or []
    print('CSV path:', csv_path)
    print('CSV headers:', fieldnames)
    required = ['Device ID', 'Status', 'Value']
    missing = [name for name in required if name not in fieldnames]
    if missing:
        raise ValueError(f'Missing required CSV columns: {missing}')
    rows = list(reader)

print('CSV data rows:', len(rows))
print('CSV sample rows:', rows[:5])

fail_rows = []
status_counts = {'Pass': 0, 'Fail': 0}
urgent_count = 0
for idx, row in enumerate(rows, start=2):
    status = row.get('Status', '')
    if status == 'Pass':
        status_counts['Pass'] += 1
    if status == 'Fail':
        status_counts['Fail'] += 1
        raw_value = row.get('Value', '')
        try:
            value_dec = Decimal(str(raw_value).strip())
        except InvalidOperation as exc:
            raise ValueError(f'Row {idx} has non-numeric Value for failed result: {raw_value!r}') from exc
        urgent = value_dec > THRESHOLD
        if urgent:
            urgent_count += 1
        fail_rows.append({'Device ID': row.get('Device ID', ''), 'Status': status, 'Value': raw_value, 'Urgent': urgent})

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
failures_path = OUTPUT_DIR / 'failures.xlsx'
summary_path = OUTPUT_DIR / 'summary.md'

wb = Workbook()
ws = wb.active
ws.title = 'Failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for row in fail_rows:
    value_dec = Decimal(str(row['Value']).strip())
    workbook_value = int(value_dec) if value_dec == value_dec.to_integral_value() else float(value_dec)
    ws.append([row['Device ID'], row['Status'], workbook_value, row['Urgent']])
wb.save(failures_path)

summary = {'Total': len(rows), 'Pass': status_counts['Pass'], 'Fail': status_counts['Fail'], 'Urgent': urgent_count}
summary_path.write_text('\n'.join(['# Test Results Summary', '', f'- Total: {summary["Total"]}', f'- Pass: {summary["Pass"]}', f'- Fail: {summary["Fail"]}', f'- Urgent: {summary["Urgent"]}', '']) + '\n', encoding='utf-8')

check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['Failures']
actual_headers = [cell.value for cell in check_ws[1]]
if actual_headers != headers:
    raise AssertionError(f'failures.xlsx headers mismatch: {actual_headers} != {headers}')
actual_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
if len(actual_rows) != len(fail_rows):
    raise AssertionError(f'failures.xlsx row count mismatch: {len(actual_rows)} != {len(fail_rows)}')
for i, (actual, expected) in enumerate(zip(actual_rows, fail_rows), start=2):
    if actual[0] != expected['Device ID'] or actual[1] != expected['Status']:
        raise AssertionError(f'failures.xlsx row {i} ID/Status mismatch: {actual} vs {expected}')
    if Decimal(str(actual[2])) != Decimal(str(expected['Value']).strip()):
        raise AssertionError(f'failures.xlsx row {i} Value mismatch: {actual[2]} vs {expected["Value"]}')
    if actual[3] is not expected['Urgent']:
        raise AssertionError(f'failures.xlsx row {i} Urgent mismatch: {actual[3]} vs {expected["Urgent"]}')

summary_text = summary_path.read_text(encoding='utf-8')
for key, value in summary.items():
    token = f'- {key}: {value}'
    if token not in summary_text:
        raise AssertionError(f'summary.md missing {token!r}')
if not csv_path.exists() or not MANIFEST.exists() or not SOP.exists():
    raise AssertionError('One or more input files are missing after processing')

print('Created:', failures_path.as_posix())
print('Created:', summary_path.as_posix())
print('Verified failures.xlsx columns:', actual_headers)
print('Verified failures.xlsx failed rows:', len(actual_rows))
print('Verified Urgent rule: Status == Fail and Value > 20 is true, otherwise false for Fail rows')
print('Verified summary counts:', summary)
print('Input files were not modified by this workflow script.')