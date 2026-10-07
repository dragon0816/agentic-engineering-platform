from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

ws = Path.cwd()
input_path = ws / 'input' / 'specs.md'
output_dir = ws / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split into product sections by level-2 headings.
sections = []
current = None
for line in text.splitlines():
    m = re.match(r'^##\s+(.+?)\s*$', line)
    if m:
        if current is not None:
            sections.append(current)
        current = {'heading': m.group(1), 'lines': []}
    elif current is not None:
        current['lines'].append(line)
if current is not None:
    sections.append(current)

print('detected product sections:', [s['heading'] for s in sections])

def parse_kv(lines):
    data = {}
    for line in lines:
        m = re.match(r'^\s*-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            data[m.group(1).strip()] = m.group(2).strip()
    return data

def unit_to_mhz(value, unit):
    unit_norm = unit.strip().lower()
    if unit_norm == 'ghz':
        return value * 1000
    if unit_norm == 'mhz':
        return value
    # Unknown unit: do not guess.
    return None

def clean_number(n):
    # Keep integers as int, non-integers as float.
    if n is None:
        return None
    if abs(n - round(n)) < 1e-9:
        return int(round(n))
    return float(n)

def parse_frequency(freq):
    if not freq:
        return None, None
    # Accept ranges like 2.4–2.5 GHz or 900-930 MHz. Unit applies to both ends.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*[–—-]\s*([0-9]+(?:\.[0-9]+)?)\s*([GM]Hz)\b', freq, flags=re.I)
    if not m:
        return None, None
    lo = unit_to_mhz(float(m.group(1)), m.group(3))
    hi = unit_to_mhz(float(m.group(2)), m.group(3))
    return clean_number(lo), clean_number(hi)

def parse_bandwidth(bw):
    if not bw:
        return None
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*([GM]Hz)\b', bw, flags=re.I)
    if not m:
        return None
    return clean_number(unit_to_mhz(float(m.group(1)), m.group(2)))

def parse_interface(interface):
    if not interface:
        return None
    parts = [p.strip() for p in re.split(r'\s*/\s*|\s*,\s*|\s*;\s*', interface) if p.strip()]
    return ';'.join(parts) if parts else None

rows = []
for section in sections:
    kv = parse_kv(section['lines'])
    fmin, fmax = parse_frequency(kv.get('Frequency', ''))
    rows.append({
        'Model': kv.get('Model') or None,
        'FrequencyMinMHz': fmin,
        'FrequencyMaxMHz': fmax,
        'BandwidthMHz': parse_bandwidth(kv.get('Bandwidth', '')),
        'Interface': parse_interface(kv.get('Interface', '')),
    })

print('parsed rows:')
for row in rows:
    print(row)

wb = Workbook()
wsheet = wb.active
wsheet.title = 'specs'
headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
wsheet.append(headers)
for row in rows:
    wsheet.append([row[h] for h in headers])
wb.save(output_path)

# Verification: reopen output file and confirm requirements.
wb2 = load_workbook(output_path, data_only=True)
sh = wb2['specs']
read_headers = [cell.value for cell in sh[1]]
read_rows = []
for r in sh.iter_rows(min_row=2, values_only=True):
    read_rows.append(dict(zip(read_headers, r)))

expected_rows = [
    {'Model': 'RX-100', 'FrequencyMinMHz': 2400, 'FrequencyMaxMHz': 2500, 'BandwidthMHz': 80, 'Interface': 'USB;LAN'},
    {'Model': 'RX-200', 'FrequencyMinMHz': 900, 'FrequencyMaxMHz': 930, 'BandwidthMHz': 20, 'Interface': 'LAN'},
    {'Model': 'RX-300', 'FrequencyMinMHz': 5000, 'FrequencyMaxMHz': 6000, 'BandwidthMHz': None, 'Interface': 'USB'},
]

assert read_headers == headers, f'headers mismatch: {read_headers}'
assert len(read_rows) == 3, f'row count mismatch: {len(read_rows)}'
assert read_rows == expected_rows, f'rows mismatch: {read_rows}'
for row in read_rows:
    assert isinstance(row['Model'], str) and row['Model'].startswith('RX-')
    assert isinstance(row['FrequencyMinMHz'], (int, float))
    assert isinstance(row['FrequencyMaxMHz'], (int, float))
    if row['BandwidthMHz'] is not None:
        assert isinstance(row['BandwidthMHz'], (int, float))
    if row['Interface'] is not None:
        assert isinstance(row['Interface'], str)

print('created:', output_path)
print('verification passed: specs.xlsx exists, sheet specs, headers correct, 3 data rows, MHz conversions correct, missing bandwidth left blank, interfaces separated by semicolon.')