from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_file = workspace / 'input' / 'tasks.xlsx'
print('workspace:', workspace)
print('input exists:', input_file.exists(), input_file)
wb = load_workbook(input_file, read_only=True, data_only=False)
print('sheetnames:', wb.sheetnames)
ws = wb['tasks']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print('headers:', headers)
print('max_row:', ws.max_row, 'max_column:', ws.max_column)
for row in ws.iter_rows(min_row=2, max_row=min(ws.max_row, 6), values_only=True):
    print('sample:', row)
wb.close()
print('output dir exists:', (workspace / 'output').exists())