from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read CSV with UTF-8 BOM support and proper CSV quoting/newline handling.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    rows = [row for row in reader]

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
try:
    device_col_idx = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_col_idx = None

# Write XLSX without adding any index column, preserving all rows/order/duplicates.
wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'

for r_idx, row in enumerate(rows, start=1):
    # Preserve empty cells as blanks. Preserve values as text as read from CSV.
    for c_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx)
        if value == '':
            cell.value = None
        else:
            cell.value = value
        # Ensure Device ID column is text and leading zeroes are preserved.
        if device_col_idx is not None and c_idx == device_col_idx:
            cell.number_format = '@'
            if r_idx > 1 and value != '':
                cell.value = str(value)

# Apply text format to the whole Device ID column, including blanks and header.
if device_col_idx is not None:
    for r_idx in range(1, len(rows) + 1):
        ws.cell(row=r_idx, column=device_col_idx).number_format = '@'

wb.save(output_path)
print('Wrote:', output_path)
print('CSV rows including header:', len(rows))
print('CSV columns:', len(headers))

# Verification: reopen output and compare against source requirements.
wb2 = load_workbook(output_path, data_only=False)
ws2 = wb2.active

if ws2.max_row != len(rows):
    raise AssertionError(f'Row count mismatch: xlsx={ws2.max_row}, csv={len(rows)}')
if ws2.max_column != len(headers):
    raise AssertionError(f'Column count mismatch: xlsx={ws2.max_column}, csv={len(headers)}')

xlsx_headers = [ws2.cell(row=1, column=c).value for c in range(1, ws2.max_column + 1)]
if xlsx_headers != headers:
    raise AssertionError(f'Header mismatch: {xlsx_headers!r} != {headers!r}')

for r_idx, row in enumerate(rows, start=1):
    if len(row) != len(headers):
        raise AssertionError(f'CSV row {r_idx} has {len(row)} fields, expected {len(headers)}')
    for c_idx, csv_value in enumerate(row, start=1):
        xval = ws2.cell(row=r_idx, column=c_idx).value
        expected = None if csv_value == '' else csv_value
        if xval != expected:
            raise AssertionError(
                f'Value mismatch at row {r_idx}, col {c_idx}: xlsx={xval!r}, expected={expected!r}'
            )

if device_col_idx is not None:
    for r_idx in range(1, ws2.max_row + 1):
        cell = ws2.cell(row=r_idx, column=device_col_idx)
        if cell.number_format != '@':
            raise AssertionError(f'Device ID cell {cell.coordinate} is not text format: {cell.number_format!r}')
        if r_idx > 1 and cell.value is not None and not isinstance(cell.value, str):
            raise AssertionError(f'Device ID cell {cell.coordinate} is not stored as text: {type(cell.value)}')

# Confirm no added index column: output headers exactly match input headers and column count already checked.
# Confirm duplicate rows preserved by full row-by-row comparison above.
print('Verification passed:')
print(' - output file exists and reopens')
print(' - all original columns preserved with same names/order')
print(' - row count/order and duplicate rows preserved')
print(' - quoted comma/newline fields preserved')
print(' - blanks remain blank')
print(' - Device ID stored/formatted as text with leading zeroes preserved')
print(' - no index column added')
print('Output file:', output_path.name)
