from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook

ws = Path.cwd()
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)
export_path = out_dir / 'export.py'
summary_path = out_dir / 'summary.xlsx'

fixed_code = r'''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals = {}
    counts = {}
    with open(args.input, newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        for row_number, row in enumerate(reader, start=2):
            model = row["Model"]
            raw_value = row["Value"]
            if raw_value is None or raw_value.strip() == "":
                continue
            try:
                value = float(raw_value)
            except ValueError as exc:
                raise ValueError(f"Invalid numeric Value on row {row_number}: {raw_value!r}") from exc
            totals[model] = totals.get(model, 0.0) + value
            counts[model] = counts.get(model, 0) + 1

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(args.output)


if __name__ == "__main__":
    main()
'''
export_path.write_text(fixed_code, encoding='utf-8')
print('WROTE', export_path)

if summary_path.exists():
    summary_path.unlink()

cmd = ['python', r'output\export.py', '--input', r'input\input.csv', '--output', r'output\summary.xlsx']
print('RUNNING:', ' '.join(cmd))
result = subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('EXIT_CODE:', result.returncode)
if result.stdout:
    print('STDOUT:\n' + result.stdout)
if result.stderr:
    print('STDERR:\n' + result.stderr)
if result.returncode != 0:
    raise SystemExit(result.returncode)

print('\nVERIFYING OUTPUT FILES')
assert export_path.exists(), 'output/export.py was not created'
assert summary_path.exists(), 'output/summary.xlsx was not created'

book = load_workbook(summary_path, data_only=True)
print('Workbook sheets:', book.sheetnames)
assert book.sheetnames == ['summary'], f'unexpected sheets: {book.sheetnames}'
sheet = book['summary']
rows = list(sheet.iter_rows(values_only=True))
print('Rows read from output/summary.xlsx:')
for row in rows:
    print(row)

expected_rows = [
    ('Model', 'Average', 'ValidCount'),
    ('M1', 15.0, 2),
    ('M2', 10.0, 2),
]
assert len(rows) == len(expected_rows), f'expected {len(expected_rows)} rows, got {len(rows)}'
assert rows[0] == expected_rows[0], f'headers mismatch: {rows[0]}'
for got, expected in zip(rows[1:], expected_rows[1:]):
    assert got[0] == expected[0], f'Model mismatch: got {got[0]} expected {expected[0]}'
    assert isinstance(got[1], (int, float)), f'Average is not numeric for {got[0]}: {got[1]!r}'
    assert abs(float(got[1]) - expected[1]) < 1e-9, f'Average mismatch for {got[0]}: got {got[1]} expected {expected[1]}'
    assert isinstance(got[2], int), f'ValidCount is not int for {got[0]}: {got[2]!r}'
    assert got[2] == expected[2], f'ValidCount mismatch for {got[0]}: got {got[2]} expected {expected[2]}'
print('CHECKS PASSED: CLI ran successfully; workbook has sheet summary, headers Model/Average/ValidCount, numeric averages, integer counts, blank Value excluded, expected row count and values.')
