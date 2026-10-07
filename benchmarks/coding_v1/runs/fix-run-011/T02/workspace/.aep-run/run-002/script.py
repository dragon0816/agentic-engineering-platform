from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read source CSV without modifying it.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames
    if headers is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in headers:
        raise ValueError('results.csv is missing required Status column')
    if 'Device ID' not in headers:
        raise ValueError('results.csv is missing required Device ID column')
    failed_rows = [row for row in reader if (row.get('Status') or '').strip().lower() == 'fail']

# Write workbook/sheet, preserving source header order and selected row order.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'
ws.append(headers)

device_col_idx = headers.index('Device ID') + 1
for row in failed_rows:
    values = []
    for h in headers:
        value = row.get(h, '')
        if h == 'Device ID' and value is not None:
            value = str(value)
        values.append(value)
    ws.append(values)

# Mark Device ID column as text so leading zeroes remain visible as text IDs.
for cell in ws.iter_cols(min_col=device_col_idx, max_col=device_col_idx, min_row=1, max_row=ws.max_row):
    for c in cell:
        c.number_format = '@'
        if c.value is not None:
            c.value = str(c.value)

wb.save(output_path)

# Verification: reopen saved file and confirm all requested requirements.
verify_wb = load_workbook(output_path, data_only=False)
if 'failed_items' not in verify_wb.sheetnames:
    raise AssertionError('failed_items sheet is missing')
verify_ws = verify_wb['failed_items']
actual_headers = [verify_ws.cell(row=1, column=i).value for i in range(1, verify_ws.max_column + 1)]
if actual_headers != headers:
    raise AssertionError(f'Header mismatch: {actual_headers!r} != {headers!r}')
actual_rows = []
for r in range(2, verify_ws.max_row + 1):
    actual_rows.append([verify_ws.cell(row=r, column=c).value for c in range(1, verify_ws.max_column + 1)])
expected_rows = [[str(row[h]) if h == 'Device ID' and row.get(h) is not None else row.get(h, '') for h in headers] for row in failed_rows]
if actual_rows != expected_rows:
    raise AssertionError(f'Rows/order mismatch: {actual_rows!r} != {expected_rows!r}')
for r in range(2, verify_ws.max_row + 1):
    cell = verify_ws.cell(row=r, column=device_col_idx)
    expected = expected_rows[r-2][device_col_idx-1]
    if cell.data_type != 's':
        raise AssertionError(f'Device ID at row {r} is not stored as text: data_type={cell.data_type!r}')
    if cell.value != expected:
        raise AssertionError(f'Device ID value changed at row {r}: {cell.value!r} != {expected!r}')

print('Created:', output_path)
print('Worksheet:', verify_ws.title)
print('Headers:', actual_headers)
print('Failed row count:', len(actual_rows))
print('Rows:')
for row in actual_rows:
    print(row)
print('Device ID text check:', [(verify_ws.cell(row=r, column=device_col_idx).value, verify_ws.cell(row=r, column=device_col_idx).data_type, verify_ws.cell(row=r, column=device_col_idx).number_format) for r in range(2, verify_ws.max_row + 1)])