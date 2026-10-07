from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font

WORKSPACE = Path('.')
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
OUTFILE = OUTPUT / 'merged_devices_tests.xlsx'


def id_cell_to_text(cell):
    """Return Device ID as text, preserving existing text and simple zero-padded numeric formats."""
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, str):
        return value
    # If Excel stores numeric ID with a zero-only display format, honor that width.
    fmt = (cell.number_format or '').split(';')[0]
    if isinstance(value, (int, float)) and float(value).is_integer():
        int_text = str(int(value))
        if fmt and set(fmt) <= {'0'} and len(fmt) > len(int_text):
            return int_text.zfill(len(fmt))
        return int_text
    return str(value)


def read_sheet_records(path, sheet_name=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    headers = [cell.value for cell in ws[1]]
    if 'Device ID' not in headers:
        raise ValueError(f'{path.name}: missing Device ID column')
    id_idx = headers.index('Device ID') + 1
    records = []
    for r in range(2, ws.max_row + 1):
        # Skip fully blank rows.
        if all(ws.cell(r, c).value is None for c in range(1, ws.max_column + 1)):
            continue
        record = {}
        for c, header in enumerate(headers, 1):
            cell = ws.cell(r, c)
            if header == 'Device ID':
                record[header] = id_cell_to_text(cell)
            else:
                record[header] = cell.value
        records.append(record)
    wb.close()
    return headers, records


dev_headers, devices = read_sheet_records(INPUT / 'devices.xlsx', 'devices')
test_headers, tests = read_sheet_records(INPUT / 'tests.xlsx', 'tests')

# Validate the stated rule: each device ID is unique.
device_by_id = {}
duplicate_device_ids = []
for rec in devices:
    did = rec['Device ID']
    if did in device_by_id:
        duplicate_device_ids.append(did)
    device_by_id[did] = rec
if duplicate_device_ids:
    raise ValueError('devices contains duplicate Device ID values: ' + ', '.join(sorted(set(duplicate_device_ids))))

matched_rows = []
matched_device_ids = set()
matched_test_row_indices = set()
for idx, test in enumerate(tests):
    did = test['Device ID']
    device = device_by_id.get(did)
    if device is not None:
        matched_device_ids.add(did)
        matched_test_row_indices.add(idx)
        row = []
        # Matched: one Device ID, test data, then device data.
        for h in test_headers:
            row.append(test.get(h))
        for h in dev_headers:
            if h != 'Device ID':
                row.append(device.get(h))
        matched_rows.append(row)

matched_headers = list(test_headers) + [h for h in dev_headers if h != 'Device ID']

unmatched_headers = ['Source']
for h in dev_headers + test_headers:
    if h not in unmatched_headers:
        unmatched_headers.append(h)

unmatched_rows = []
# Devices with no successful test correspondence.
for device in devices:
    if device['Device ID'] not in matched_device_ids:
        row = ['devices']
        for h in unmatched_headers[1:]:
            row.append(device.get(h))
        unmatched_rows.append(row)
# Tests whose Device ID has no device record.
for idx, test in enumerate(tests):
    if idx not in matched_test_row_indices:
        row = ['tests']
        for h in unmatched_headers[1:]:
            row.append(test.get(h))
        unmatched_rows.append(row)

wb_out = Workbook()
ws_m = wb_out.active
ws_m.title = 'matched'
ws_u = wb_out.create_sheet('unmatched')

for ws, headers, rows in [(ws_m, matched_headers, matched_rows), (ws_u, unmatched_headers, unmatched_rows)]:
    ws.append(headers)
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    # Set Device ID column cells to text format.
    for col_idx, header in enumerate(headers, 1):
        if header == 'Device ID':
            col_letter = ws.cell(1, col_idx).column_letter
            for cell in ws[col_letter]:
                cell.number_format = '@'
            break
    ws.freeze_panes = 'A2'

wb_out.save(OUTFILE)

# Verification by reopening output file.
wb_check = load_workbook(OUTFILE, data_only=True)
print('Created:', OUTFILE)
print('Sheets:', wb_check.sheetnames)
for s in ['matched', 'unmatched']:
    ws = wb_check[s]
    headers = [cell.value for cell in ws[1]]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    print(f'{s} headers:', headers)
    print(f'{s} data row count:', len(rows))
    id_col = headers.index('Device ID') + 1
    ids = [ws.cell(r, id_col).value for r in range(2, ws.max_row + 1)]
    id_types = [type(v).__name__ for v in ids]
    print(f'{s} Device ID values:', ids)
    print(f'{s} Device ID types:', id_types)
    print(f'{s} Device ID number formats:', [ws.cell(r, id_col).number_format for r in range(1, ws.max_row + 1)])
print('matched rows:')
for row in wb_check['matched'].iter_rows(min_row=2, values_only=True):
    print(row)
print('unmatched rows:')
for row in wb_check['unmatched'].iter_rows(min_row=2, values_only=True):
    print(row)
wb_check.close()
