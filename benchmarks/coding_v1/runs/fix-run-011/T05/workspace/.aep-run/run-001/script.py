from pathlib import Path
from openpyxl import load_workbook

base = Path('.')
input_dir = base / 'input'
print('Input files:')
for p in sorted(input_dir.iterdir()):
    print(' -', p.name)

for fname in ['devices.xlsx', 'tests.xlsx']:
    path = input_dir / fname
    print(f'\nWorkbook: {fname} exists={path.exists()}')
    wb = load_workbook(path, read_only=True, data_only=True)
    print('Sheets:', wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True))
    print('Max row/col:', ws.max_row, ws.max_column)
    for i, row in enumerate(rows, 1):
        print(f'Row {i}:', row)
    wb.close()
