import json
import re
import time
from html.parser import HTMLParser
from urllib.parse import urljoin
from urllib.request import urlopen, Request
from urllib.error import HTTPError, URLError
from pathlib import Path
from openpyxl import Workbook, load_workbook

WORKSPACE = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\coding_v1\runs\fix-run-011\T13\workspace')
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:54763/t13/list'
XLSX_PATH = OUTPUT / 'robust_products.xlsx'
WARN_PATH = OUTPUT / 'warnings.json'
UA = {'User-Agent': 'Mozilla/5.0'}

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.items = []
        self._in_a = False
        self._href = None
        self._text = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs = dict(attrs)
            href = attrs.get('href')
            if href:
                self._in_a = True
                self._href = href
                self._text = []
    def handle_data(self, data):
        if self._in_a:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._in_a:
            name = ' '.join(''.join(self._text).split())
            if name and self._href:
                self.items.append({'Name': name, 'URL': urljoin(LIST_URL, self._href)})
            self._in_a = False
            self._href = None
            self._text = []

class TextHTMLParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, data):
        txt = data.strip()
        if txt:
            self.parts.append(txt)

def fetch(url, retry_500=False, max_attempts=3):
    attempts = 0
    last_error = None
    while True:
        attempts += 1
        try:
            req = Request(url, headers=UA)
            with urlopen(req, timeout=10) as r:
                body = r.read().decode('utf-8', errors='replace')
                return {'ok': True, 'status': getattr(r, 'status', 200), 'body': body, 'attempts': attempts, 'error': None}
        except HTTPError as e:
            body = ''
            try:
                body = e.read().decode('utf-8', errors='replace')
            except Exception:
                pass
            last_error = f'HTTP {e.code}'
            if retry_500 and e.code == 500 and attempts < max_attempts:
                time.sleep(0.05)
                continue
            return {'ok': False, 'status': e.code, 'body': body, 'attempts': attempts, 'error': last_error}
        except URLError as e:
            last_error = f'URL error: {e.reason}'
            return {'ok': False, 'status': None, 'body': '', 'attempts': attempts, 'error': last_error}
        except Exception as e:
            last_error = f'{type(e).__name__}: {e}'
            return {'ok': False, 'status': None, 'body': '', 'attempts': attempts, 'error': last_error}

def html_text(html):
    p = TextHTMLParser()
    p.feed(html)
    return '\n'.join(p.parts)

def clean_value(v):
    if v is None:
        return 'N/A'
    v = re.sub(r'\s+', ' ', str(v)).strip()
    return v if v else 'N/A'

def parse_detail(html):
    text = html_text(html)
    flat = re.sub(r'\s+', ' ', text).strip()
    result = {'Name': 'N/A', 'Price': 'N/A', 'Currency': 'N/A', 'CPU': 'N/A'}

    # Generic label extraction: supports "Price: 123", "Price 123", "CPU: A16", etc.
    def label_value(label, following_labels):
        labels_alt = '|'.join(re.escape(x) for x in following_labels)
        pattern = rf'(?is)\b{re.escape(label)}\b\s*[:：]?\s*(.*?)(?=\s*(?:{labels_alt})\b\s*[:：]?|$)'
        m = re.search(pattern, flat)
        if m:
            return clean_value(m.group(1))
        # Line-oriented fallback
        for line in text.splitlines():
            m2 = re.match(rf'(?is)^\s*{re.escape(label)}\s*[:：]?\s*(.+?)\s*$', line)
            if m2:
                return clean_value(m2.group(1))
        return 'N/A'

    result['Name'] = label_value('Name', ['Price', 'Currency', 'CPU'])
    price_raw = label_value('Price', ['Currency', 'CPU', 'Name'])
    currency_raw = label_value('Currency', ['CPU', 'Name', 'Price'])
    result['CPU'] = label_value('CPU', ['Name', 'Price', 'Currency'])

    # If price and currency appear together, split common formats.
    price = price_raw
    currency = currency_raw
    if price != 'N/A':
        m = re.search(r'([A-Z]{3}|[$€£¥])?\s*([0-9][0-9,]*(?:\.\d+)?)\s*([A-Z]{3})?', price)
        if m:
            if currency == 'N/A':
                currency = clean_value(m.group(1) or m.group(3))
            price = m.group(2).replace(',', '')
    else:
        m = re.search(r'(?is)\b(USD|TWD|EUR|GBP|JPY|CNY|NTD|[$€£¥])\s*([0-9][0-9,]*(?:\.\d+)?)', flat)
        if m:
            currency = m.group(1)
            price = m.group(2).replace(',', '')
        else:
            m = re.search(r'(?is)\bPrice\b\s*[:：]?\s*([0-9][0-9,]*(?:\.\d+)?)', flat)
            if m:
                price = m.group(1).replace(',', '')
    result['Price'] = clean_value(price)
    result['Currency'] = clean_value(currency)

    # Fallback for title/headline if no Name label exists.
    if result['Name'] == 'N/A':
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        if lines:
            first = re.sub(r'\s+', ' ', lines[0]).strip()
            if not re.match(r'(?i)^(price|currency|cpu)\b', first):
                result['Name'] = first
    return result

# Read list
list_fetch = fetch(LIST_URL)
if not list_fetch['ok']:
    raise RuntimeError(f'Unable to fetch product list: {list_fetch["error"]}')
parser = ProductListParser()
parser.feed(list_fetch['body'])
listed = parser.items[:5]
print('Fetched list URL:', LIST_URL)
print('List item count:', len(parser.items))
print('Using first five list items:', listed)
if len(listed) < 5:
    raise RuntimeError(f'Expected at least five products in list, found {len(listed)}')

rows = []
warnings = []
for item in listed:
    detail = fetch(item['URL'], retry_500=True, max_attempts=3)
    warning = {
        'Name': item['Name'],
        'URL': item['URL'],
        'FetchStatus': 'ok' if detail['ok'] else 'error',
        'HTTPStatus': detail['status'],
        'Attempts': detail['attempts'],
        'Warning': 'ok' if detail['ok'] else clean_value(detail['error']),
    }
    warnings.append(warning)
    if detail['ok']:
        parsed = parse_detail(detail['body'])
        # Preserve list name if detail lacks name.
        if parsed['Name'] == 'N/A':
            parsed['Name'] = item['Name']
        row = {
            'Name': clean_value(parsed.get('Name')),
            'Price': clean_value(parsed.get('Price')),
            'Currency': clean_value(parsed.get('Currency')),
            'CPU': clean_value(parsed.get('CPU')),
            'URL': item['URL'],
            'FetchStatus': 'ok',
        }
    else:
        row = {
            'Name': item['Name'],
            'Price': 'N/A',
            'Currency': 'N/A',
            'CPU': 'N/A',
            'URL': item['URL'],
            'FetchStatus': 'error',
        }
    rows.append(row)

# Write workbook
headers = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']
wb = Workbook()
ws = wb.active
ws.title = 'Products'
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(XLSX_PATH)

with WARN_PATH.open('w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Verification: reopen both outputs and check stated requirements.
check_wb = load_workbook(XLSX_PATH, data_only=True)
check_ws = check_wb['Products']
actual_headers = [check_ws.cell(1, c).value for c in range(1, 7)]
assert actual_headers == headers, f'Header mismatch: {actual_headers}'
assert check_ws.max_row == 6, f'Expected 5 product rows + header, found {check_ws.max_row}'
for r in range(2, 7):
    values = {headers[c-1]: check_ws.cell(r, c).value for c in range(1, 7)}
    assert values['Name'] not in (None, ''), f'Missing Name row {r}'
    assert values['URL'] not in (None, ''), f'Missing URL row {r}'
    assert values['FetchStatus'] in ('ok', 'error'), f'Bad FetchStatus row {r}: {values["FetchStatus"]}'
    for k in ['Price', 'Currency', 'CPU']:
        assert values[k] not in (None, ''), f'Missing {k} row {r}'
    if values['FetchStatus'] == 'error':
        assert values['Price'] == 'N/A' and values['Currency'] == 'N/A' and values['CPU'] == 'N/A', f'Error row not N/A row {r}: {values}'
with WARN_PATH.open('r', encoding='utf-8') as f:
    check_warn = json.load(f)
assert isinstance(check_warn, list) and len(check_warn) == 5, f'warnings.json should have 5 entries, found {len(check_warn) if isinstance(check_warn, list) else type(check_warn)}'
for w in check_warn:
    assert w.get('FetchStatus') in ('ok', 'error'), f'Bad warning FetchStatus: {w}'
    if w.get('FetchStatus') == 'ok':
        assert w.get('Warning') == 'ok', f'Success page warning should be ok: {w}'
    if w.get('FetchStatus') == 'error' and w.get('HTTPStatus') == 500:
        assert w.get('Attempts') == 3, f'HTTP 500 should be attempted 3 times: {w}'

print('Wrote:', XLSX_PATH)
print('Wrote:', WARN_PATH)
print('Rows:')
for row in rows:
    print(row)
print('Warnings:')
print(json.dumps(warnings, ensure_ascii=False, indent=2))
print('Verification passed: workbook headers, five rows, nonblank fields/N/A placeholders, error rows, warnings.json statuses, and HTTP 500 retry count were checked from saved outputs.')
