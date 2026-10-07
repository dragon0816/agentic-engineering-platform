from pathlib import Path
from urllib.request import urlopen
from urllib.parse import urljoin
from html.parser import HTMLParser
from openpyxl import Workbook, load_workbook
import re

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
CATALOG_URL = 'http://127.0.0.1:63668/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_main_catalog = False
        self.main_depth = 0
        self.in_card = False
        self.card_depth = 0
        self.current = None
        self.capture = None
        self.products = []

    def _attrs(self, attrs):
        return dict(attrs)

    def handle_starttag(self, tag, attrs):
        attrs_d = self._attrs(attrs)
        if tag == 'main' and attrs_d.get('id') == 'catalog' and not self.in_main_catalog:
            self.in_main_catalog = True
            self.main_depth = 1
            return

        if self.in_main_catalog:
            self.main_depth += 1
            classes = attrs_d.get('class', '').split()
            if tag == 'div' and 'card' in classes and not self.in_card:
                self.in_card = True
                self.card_depth = 1
                self.current = {'name_parts': [], 'price_parts': [], 'href': None}
                return

            if self.in_card:
                self.card_depth += 1
                if tag == 'a' and self.current is not None and self.current.get('href') is None:
                    self.current['href'] = attrs_d.get('href', '')
                    self.capture = 'name'
                elif tag == 'span' and 'price' in classes:
                    self.capture = 'price'

    def handle_endtag(self, tag):
        if self.in_card:
            if (self.capture == 'name' and tag == 'a') or (self.capture == 'price' and tag == 'span'):
                self.capture = None
            self.card_depth -= 1
            if self.card_depth == 0:
                name = ''.join(self.current['name_parts']).strip() if self.current else ''
                price_text = ''.join(self.current['price_parts']).strip() if self.current else ''
                href = self.current.get('href') if self.current else ''
                if name and price_text and href:
                    self.products.append({'name': name, 'price_text': price_text, 'href': href})
                self.in_card = False
                self.current = None
                self.capture = None
                return

        if self.in_main_catalog:
            self.main_depth -= 1
            if self.main_depth == 0:
                self.in_main_catalog = False

    def handle_data(self, data):
        if self.in_card and self.current is not None and self.capture == 'name':
            self.current['name_parts'].append(data)
        elif self.in_card and self.current is not None and self.capture == 'price':
            self.current['price_parts'].append(data)

def parse_price(price_text):
    text = price_text.strip()
    m = re.match(r'^\s*([A-Za-z$€£¥NT]+\$?)\s*([0-9][0-9,]*(?:\.\d+)?)\s*$', text)
    if not m:
        # More general fallback: currency is leading non-numeric text; amount is first number.
        m2 = re.search(r'([0-9][0-9,]*(?:\.\d+)?)', text)
        if not m2:
            raise ValueError(f'No numeric price found in {price_text!r}')
        currency = text[:m2.start()].strip()
        amount = m2.group(1)
    else:
        currency = m.group(1).strip()
        amount = m.group(2)
    price_value = float(amount.replace(',', '')) if '.' in amount else int(amount.replace(',', ''))
    return currency, price_value

# Fetch and parse only <main id="catalog">; this excludes <aside class="ads"> recommendations/ads.
with urlopen(CATALOG_URL, timeout=10) as resp:
    html = resp.read().decode('utf-8', errors='replace')
parser = CatalogParser()
parser.feed(html)

rows = []
for product in parser.products:
    currency, price_value = parse_price(product['price_text'])
    rows.append({
        '產品名稱': product['name'],
        '價格': price_value,
        '幣別': currency,
        '產品網址': urljoin(CATALOG_URL, product['href']),
    })

# Write workbook.
wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
for cell in ws['B'][1:]:
    cell.number_format = '#,##0.00' if isinstance(cell.value, float) else '#,##0'
for col in ['A', 'C', 'D']:
    for cell in ws[col]:
        cell.number_format = '@'
# Set readable widths.
widths = {'A': 20, 'B': 12, 'C': 10, 'D': 45}
for col, width in widths.items():
    ws.column_dimensions[col].width = width
wb.save(OUTFILE)
print(f'Wrote {OUTFILE}')
print('Parsed products in page order:')
for r in rows:
    print(r)

# Verification: reopen output file and confirm stated requirements.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
read_headers = [ws2.cell(row=1, column=i).value for i in range(1, 5)]
assert read_headers == headers, f'headers mismatch: {read_headers!r}'
read_rows = []
for row_idx in range(2, ws2.max_row + 1):
    vals = [ws2.cell(row=row_idx, column=i).value for i in range(1, 5)]
    read_rows.append(dict(zip(headers, vals)))
assert len(read_rows) == len(rows), f'row count mismatch {len(read_rows)} != {len(rows)}'
assert len(read_rows) == 5, f'expected 5 catalog products, got {len(read_rows)}'
assert [r['產品名稱'] for r in read_rows] == [r['產品名稱'] for r in rows], 'page order/name mismatch'
assert all(isinstance(r['價格'], (int, float)) and not isinstance(r['價格'], bool) for r in read_rows), 'price is not numeric'
assert all(r['幣別'] == 'NT$' for r in read_rows), 'currency mismatch'
assert all(isinstance(r['產品網址'], str) and r['產品網址'].startswith('http://127.0.0.1:63668/') for r in read_rows), 'URL not complete'
excluded_names = {'Sponsored Router', '推薦：高速網卡'}
assert excluded_names.isdisjoint({r['產品名稱'] for r in read_rows}), 'ad/recommendation product included'
expected_order = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
assert [r['產品名稱'] for r in read_rows] == expected_order, 'unexpected product order/content'
print('\nVerification passed:')
print(' - products.xlsx exists in output')
print(' - headers:', read_headers)
print(' - row count:', len(read_rows))
print(' - prices are numeric')
print(' - URLs are complete')
print(' - page order preserved')
print(' - ads/recommendations excluded')