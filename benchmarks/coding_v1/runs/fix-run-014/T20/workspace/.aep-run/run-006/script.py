from pathlib import Path
import csv, json, os, stat
from openpyxl import Workbook, load_workbook

WORKSPACE = Path('.').resolve()
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
THRESHOLD = 20

print('Running test failure tracking workflow')
print('Workspace:', WORKSPACE)
print('Input dir:', INPUT_DIR)
print('Output dir:', OUTPUT_DIR)

if not INPUT_DIR.exists():
    raise FileNotFoundError(f'Input directory not found: {INPUT_DIR}')
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Resolve results.csv through manifest.json when present, because this benchmark maps logical SOP names to supplied files.
manifest_path = INPUT_DIR / 'manifest.json'
if manifest_path.exists():
    manifest = json.loads(manifest_path.read_text(encoding='utf-8'))
    rel = manifest.get('results.csv')
    if not rel:
        raise KeyError('manifest.json does not define results.csv')
    csv_path = (WORKSPACE / rel).resolve()
else:
    csv_path = (INPUT_DIR / 'results.csv').resolve()

# Keep processing read-only: only read under input, only write under output.
try:
    csv_path.relative_to(INPUT_DIR.resolve())
except ValueError:
    raise ValueError(f'Resolved CSV is not under input/: {csv_path}')
if not csv_path.exists():
    raise FileNotFoundError(f'Results CSV not found: {csv_path}')

print('Reading CSV:', csv_path)
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    print('CSV headers:', headers)
    required = ['Device ID', 'Status', 'Value']
    missing = [h for h in required if h not in headers]
    if missing:
        raise ValueError(f'Missing required column(s): {missing}; headers={headers}')
    rows = list(reader)

print('Total input rows:', len(rows))

def parse_number(value, row_number):
    text = str(value).strip()
    if text == '':
        raise ValueError(f'Blank Value at data row {row_number}')
    try:
        num = float(text)
    except ValueError as exc:
        raise ValueError(f'Non-numeric Value {value!r} at data row {row_number}') from exc
    if num.is_integer():
        return int(num)
    return num

fail_rows = []
pass_count = 0
fail_count = 0
urgent_count = 0
for idx, row in enumerate(rows, start=2):
    status = row.get('Status')
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        value_num = parse_number(row.get('Value'), idx)
        urgent = value_num > THRESHOLD
        if urgent:
            urgent_count += 1
        fail_rows.append({
            'Device ID': row.get('Device ID'),
            'Status': status,
            'Value': value_num,
            'Urgent': urgent,
        })

failures_path = OUTPUT_DIR / 'failures.xlsx'
summary_path = OUTPUT_DIR / 'summary.md'

def remove_existing_output(path):
    if path.exists():
        try:
            path.chmod(stat.S_IREAD | stat.S_IWRITE)
        except OSError:
            pass
        try:
            path.unlink()
        except PermissionError as exc:
            raise PermissionError(
                f'Cannot replace {path}. It may be open in another application or locked; close it and rerun the workflow.'
            ) from exc

remove_existing_output(failures_path)
remove_existing_output(summary_path)

wb = Workbook()
ws = wb.active
ws.title = 'Failures'
output_headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(output_headers)
for r in fail_rows:
    ws.append([r['Device ID'], r['Status'], r['Value'], r['Urgent']])

tmp_failures_path = OUTPUT_DIR / f'.failures.{os.getpid()}.tmp.xlsx'
if tmp_failures_path.exists():
    tmp_failures_path.unlink()
wb.save(tmp_failures_path)
os.replace(tmp_failures_path, failures_path)

summary_text = '\n'.join([
    '# Test Results Summary',
    '',
    f'- Total: {len(rows)}',
    f'- Pass: {pass_count}',
    f'- Fail: {fail_count}',
    f'- Urgent: {urgent_count}',
    '',
])
tmp_summary_path = OUTPUT_DIR / f'.summary.{os.getpid()}.tmp.md'
tmp_summary_path.write_text(summary_text, encoding='utf-8')
os.replace(tmp_summary_path, summary_path)

# Verification: reopen outputs and confirm names, headers, counts, types, and urgent rule.
if not failures_path.exists():
    raise AssertionError('failures.xlsx was not created')
if not summary_path.exists():
    raise AssertionError('summary.md was not created')

check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['Failures']
actual_headers = [check_ws.cell(row=1, column=i).value for i in range(1, 5)]
if actual_headers != output_headers:
    raise AssertionError(f'Unexpected failures.xlsx headers: {actual_headers}')
actual_data = list(check_ws.iter_rows(min_row=2, values_only=True))
if len(actual_data) != fail_count:
    raise AssertionError(f'Expected {fail_count} failure rows, found {len(actual_data)}')
for row_num, (device_id, status, value, urgent) in enumerate(actual_data, start=2):
    if status != 'Fail':
        raise AssertionError(f'Non-Fail row in failures.xlsx row {row_num}: {status!r}')
    if not isinstance(device_id, str):
        raise AssertionError(f'Device ID is not text in row {row_num}: {device_id!r}')
    if not isinstance(value, (int, float)):
        raise AssertionError(f'Value is not numeric in row {row_num}: {value!r}')
    if urgent not in (True, False):
        raise AssertionError(f'Urgent is not boolean in row {row_num}: {urgent!r}')
    if urgent != (value > THRESHOLD):
        raise AssertionError(f'Urgent rule mismatch in row {row_num}: value={value}, urgent={urgent}')

summary_check = summary_path.read_text(encoding='utf-8')
expected_lines = [
    f'- Total: {len(rows)}',
    f'- Pass: {pass_count}',
    f'- Fail: {fail_count}',
    f'- Urgent: {urgent_count}',
]
for line in expected_lines:
    if line not in summary_check:
        raise AssertionError(f'Missing summary line: {line}')

print('Wrote:', failures_path)
print('Wrote:', summary_path)
print('Verified failures.xlsx headers, failure row count, Device ID text, Value numeric, Urgent boolean and Value > 20 rule.')
print('Verified summary.md counts:', {'Total': len(rows), 'Pass': pass_count, 'Fail': fail_count, 'Urgent': urgent_count})
