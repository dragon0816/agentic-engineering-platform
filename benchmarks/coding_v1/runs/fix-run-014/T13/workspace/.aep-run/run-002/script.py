import json
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import Workbook, load_workbook

LIST_URL = 'http://127.0.0.1:63668/t13/list'
OUTPUT_DIR = Path('output')
XLSX_PATH = OUTPUT_DIR / 'robust_products.xlsx'
WARNINGS_PATH = OUTPUT_DIR / 'warnings.json'
COLUMNS = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._href = None
        self._text_parts = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs = dict(attrs)
            self._href = attrs.get('href')
            self._text_parts = []
    def handle_data(self, data):
        if self._href is not None:
            self._text_parts.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._href is not None:
            text = ' '.join(''.join(self._text_parts).split())
            self.links.append((text, self._href))
            self._href = None
            self._text_parts = []

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, data):
        data = data.strip()
        if data:
            self.parts.append(data)
    def text(self):
        return '\n'.join(self.parts)

def fetch_text(url, attempts=1, retry_http_500_only=False):
    last_error = None
    for attempt in range(1, attempts + 1):
        try:
            req = urllib.request.Request(url, headers={'User-Agent': 'AEP robust product collector'})
            with urllib.request.urlopen(req, timeout=10) as resp:
                raw = resp.read()
                charset = resp.headers.get_content_charset() or 'utf-8'
                return raw.decode(charset, errors='replace'), attempt, None
        except urllib.error.HTTPError as e:
            body = ''
            try:
                body = e.read().decode('utf-8', errors='replace')
            except Exception:
                pass
            last_error = f'HTTP {e.code}' + (f': {body[:120]}' if body else '')
            if not (retry_http_500_only and e.code == 500 and attempt < attempts):
                break
            time.sleep(0.15)
        except Exception as e:
            last_error = f'{type(e).__name__}: {e}'
            break
    return None, attempts if retry_http_500_only else 1, last_error

def extract_labeled_value(text, labels):
    # Matches common detail forms: "Price: 999 USD", "CPU - A16", table text, etc.
    for label in labels:
        patterns = [
            rf'(?im)^\s*{re.escape(label)}\s*[:：\-]\s*(.+?)\s*$',
            rf'(?is)\b{re.escape(label)}\b\s*[:：\-]?\s*\n\s*(.+?)(?:\n|$)',
        ]
        for pat in patterns:
            m = re.search(pat, text)
            if m:
                val = ' '.join(m.group(1).strip().split())
                if val:
                    return val
    return None

def parse_detail(html, fallback_name):
    parser = TextParser()
    parser.feed(html)
    text = parser.text()
    compact = ' '.join(text.split())

    name = extract_labeled_value(text, ['Name', 'Product', '產品', '名稱'])
    if not name:
        title_match = re.search(r'(?is)<h1[^>]*>(.*?)</h1>', html)
        if title_match:
            tp = TextParser(); tp.feed(title_match.group(1)); name = ' '.join(tp.text().split())
    if not name:
        name = fallback_name

    cpu = extract_labeled_value(text, ['CPU', 'Processor', '處理器', '規格']) or 'N/A'

    price = 'N/A'
    currency = 'N/A'
    price_value = extract_labeled_value(text, ['Price', '價格'])
    search_spaces = []
    if price_value:
        search_spaces.append(price_value)
    search_spaces.append(compact)
    # Handles "USD 699", "$699", "699 USD", "NT$ 123", "TWD 123".
    for space in search_spaces:
        m = re.search(r'\b(USD|EUR|TWD|NTD|JPY|CNY|RMB|GBP)\b\s*([$€¥NT]*\s*)?([0-9][0-9,]*(?:\.\d+)?)', space, re.I)
        if m:
            currency = m.group(1).upper()
            price = m.group(3).replace(',', '')
            break
        m = re.search(r'([$€¥]|NT\$)\s*([0-9][0-9,]*(?:\.\d+)?)\s*\b(USD|EUR|TWD|NTD|JPY|CNY|RMB|GBP)?\b', space, re.I)
        if m:
            sym = m.group(1)
            price = m.group(2).replace(',', '')
            currency = (m.group(3).upper() if m.group(3) else {'$':'USD','€':'EUR','¥':'JPY','NT$':'TWD'}.get(sym, 'N/A'))
            break
        m = re.search(r'([0-9][0-9,]*(?:\.\d+)?)\s*\b(USD|EUR|TWD|NTD|JPY|CNY|RMB|GBP)\b', space, re.I)
        if m:
            price = m.group(1).replace(',', '')
            currency = m.group(2).upper()
            break
    if currency == 'NTD':
        currency = 'TWD'

    return {
        'Name': name or fallback_name,
        'Price': price or 'N/A',
        'Currency': currency or 'N/A',
        'CPU': cpu or 'N/A',
    }

def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    list_html, _, err = fetch_text(LIST_URL)
    if err:
        raise RuntimeError(f'Could not fetch product list: {err}')
    lp = LinkParser()
    lp.feed(list_html)
    products = []
    for name, href in lp.links:
        if href:
            products.append((name or 'N/A', urllib.parse.urljoin(LIST_URL, href)))
    products = products[:5]
    print(f'Found {len(products)} list products:')
    for name, url in products:
        print(' -', name, url)
    if len(products) != 5:
        raise RuntimeError(f'Expected five product links, found {len(products)}')

    rows = []
    warnings = []
    for list_name, url in products:
        html, attempts, err = fetch_text(url, attempts=3, retry_http_500_only=True)
        if err:
            row = {
                'Name': list_name or 'N/A',
                'Price': 'N/A',
                'Currency': 'N/A',
                'CPU': 'N/A',
                'URL': url,
                'FetchStatus': 'error',
            }
            warnings.append({'Name': row['Name'], 'URL': url, 'FetchStatus': 'error', 'attempts': attempts, 'warning': err})
        else:
            parsed = parse_detail(html, list_name)
            row = {
                'Name': parsed.get('Name') or list_name or 'N/A',
                'Price': parsed.get('Price') or 'N/A',
                'Currency': parsed.get('Currency') or 'N/A',
                'CPU': parsed.get('CPU') or 'N/A',
                'URL': url,
                'FetchStatus': 'ok',
            }
            warnings.append({'Name': row['Name'], 'URL': url, 'FetchStatus': 'ok', 'attempts': attempts, 'warning': ''})
        rows.append(row)
        print('ROW', row)

    wb = Workbook()
    ws = wb.active
    ws.title = 'Products'
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row[col] for col in COLUMNS])
    wb.save(XLSX_PATH)

    with WARNINGS_PATH.open('w', encoding='utf-8') as f:
        json.dump(warnings, f, ensure_ascii=False, indent=2)

    # Verification: reopen both outputs and check stated requirements.
    wb2 = load_workbook(XLSX_PATH, data_only=True)
    ws2 = wb2.active
    headers = [cell.value for cell in ws2[1]]
    data_rows = list(ws2.iter_rows(min_row=2, values_only=True))
    assert headers == COLUMNS, f'headers mismatch: {headers}'
    assert len(data_rows) == 5, f'row count mismatch: {len(data_rows)}'
    for idx, vals in enumerate(data_rows, start=2):
        rec = dict(zip(COLUMNS, vals))
        for col in COLUMNS:
            assert rec[col] is not None and rec[col] != '', f'blank {col} at row {idx}'
        assert rec['FetchStatus'] in ('ok', 'error'), f'bad FetchStatus row {idx}'
        if rec['FetchStatus'] == 'error':
            assert rec['Price'] == 'N/A' and rec['Currency'] == 'N/A' and rec['CPU'] == 'N/A', f'error placeholders row {idx}'
    with WARNINGS_PATH.open('r', encoding='utf-8') as f:
        w = json.load(f)
    assert isinstance(w, list) and len(w) == 5, 'warnings.json should be a five-item list'
    assert all(item.get('FetchStatus') in ('ok', 'error') for item in w), 'warnings statuses invalid'
    print('VERIFIED robust_products.xlsx headers, 5 rows, nonblank fields, N/A placeholders for errors, FetchStatus values.')
    print('VERIFIED warnings.json is a 5-item list with ok/error FetchStatus values.')
    print('OUTPUTS:', XLSX_PATH.resolve(), WARNINGS_PATH.resolve())

if __name__ == '__main__':
    main()
