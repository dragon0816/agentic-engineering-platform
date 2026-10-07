from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
out_file = base / 'output' / 'merged_device_tests.xlsx'
wb = load_workbook(out_file, data_only=True)
assert wb.sheetnames == ['matched', 'unmatched'], wb.sheetnames

ws_m = wb['matched']
ws_u = wb['unmatched']
matched_headers = [c.value for c in ws_m[1]]
unmatched_headers = [c.value for c in ws_u[1]]
print('matched headers:', matched_headers)
print('unmatched headers:', unmatched_headers)
assert matched_headers == ['Device ID', 'Test', 'Result', 'Name', 'Location']
assert unmatched_headers == ['Source', 'Device ID', 'Name', 'Location', 'Test', 'Result']

matched_rows = list(ws_m.iter_rows(min_row=2, values_only=True))
unmatched_rows = list(ws_u.iter_rows(min_row=2, values_only=True))
print('matched row count:', len(matched_rows))
print('unmatched row count:', len(unmatched_rows))
print('matched rows:', matched_rows)
print('unmatched rows:', unmatched_rows)
assert len(matched_rows) == 3
assert len(unmatched_rows) == 2

# Verify IDs are text strings with leading zeros preserved where present.
for row in matched_rows:
    assert isinstance(row[0], str), (row[0], type(row[0]))
assert [r[0] for r in matched_rows] == ['001', '001', '002']
for row in unmatched_rows:
    assert isinstance(row[1], str), (row[1], type(row[1]))
assert set(r[0] for r in unmatched_rows) == {'devices', 'tests'}
assert ('devices', '003', 'Gamma', 'Lab C', None, None) in unmatched_rows
assert ('tests', '999', None, None, 'RF-9', 'Fail') in unmatched_rows

# Check Device ID cells are formatted as text.
for ws in [ws_m, ws_u]:
    headers = [c.value for c in ws[1]]
    col = headers.index('Device ID') + 1
    fmts = [ws.cell(row=r, column=col).number_format for r in range(2, ws.max_row + 1)]
    assert all(fmt == '@' for fmt in fmts), (ws.title, fmts)
    print(ws.title, 'Device ID formats:', fmts)

wb.close()
print('Validation passed for file, sheets, headers, row counts, text IDs, leading zeros, Source values, and unmatched rows.')