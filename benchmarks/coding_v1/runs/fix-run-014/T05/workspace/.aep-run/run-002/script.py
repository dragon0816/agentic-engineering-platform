from pathlib import Path
from openpyxl import load_workbook, Workbook

base = Path.cwd()
input_dir = base / 'input'
out_dir = base / 'output'
out_dir.mkdir(exist_ok=True)
out_file = out_dir / 'merged_device_tests.xlsx'

def get_headers(ws):
    return [cell.value for cell in ws[1]]

def value_as_text(v):
    if v is None:
        return ''
    return str(v)

def load_rows(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = get_headers(ws)
    if 'Device ID' not in headers:
        raise ValueError(f'{path.name} does not contain Device ID column')
    rows = []
    id_idx = headers.index('Device ID')
    for row in ws.iter_rows(min_row=2, values_only=True):
        values = list(row[:len(headers)])
        while len(values) < len(headers):
            values.append(None)
        values[id_idx] = value_as_text(values[id_idx])
        rows.append(dict(zip(headers, values)))
    wb.close()
    return headers, rows

dev_headers, dev_rows = load_rows(input_dir / 'devices.xlsx')
test_headers, test_rows = load_rows(input_dir / 'tests.xlsx')

dev_by_id = {}
duplicate_devices = []
for r in dev_rows:
    did = r['Device ID']
    if did in dev_by_id:
        duplicate_devices.append(did)
    dev_by_id[did] = r
if duplicate_devices:
    raise ValueError('devices contains duplicate Device ID(s): ' + ', '.join(duplicate_devices))

test_ids = {r['Device ID'] for r in test_rows}
matched_tests = [t for t in test_rows if t['Device ID'] in dev_by_id]
unmatched_tests = [t for t in test_rows if t['Device ID'] not in dev_by_id]
unmatched_devices = [d for d in dev_rows if d['Device ID'] not in test_ids]

wb_out = Workbook()
ws_m = wb_out.active
ws_m.title = 'matched'
matched_headers = list(test_headers) + [h for h in dev_headers if h != 'Device ID']
ws_m.append(matched_headers)
for t in matched_tests:
    d = dev_by_id[t['Device ID']]
    ws_m.append([t.get(h) for h in test_headers] + [d.get(h) for h in dev_headers if h != 'Device ID'])

ws_u = wb_out.create_sheet('unmatched')
unmatched_headers = ['Source'] + dev_headers + [h for h in test_headers if h not in dev_headers]
ws_u.append(unmatched_headers)
for d in unmatched_devices:
    row = {'Source': 'devices'}
    row.update(d)
    ws_u.append([row.get(h) for h in unmatched_headers])
for t in unmatched_tests:
    row = {'Source': 'tests'}
    row.update(t)
    ws_u.append([row.get(h) for h in unmatched_headers])

# Force Device ID cells to text format in both sheets.
for ws in [ws_m, ws_u]:
    headers = [c.value for c in ws[1]]
    if 'Device ID' in headers:
        col = headers.index('Device ID') + 1
        for cell in ws.iter_cols(min_col=col, max_col=col, min_row=2, max_row=ws.max_row):
            for c in cell:
                c.value = value_as_text(c.value)
                c.number_format = '@'

wb_out.save(out_file)
print('Wrote:', out_file)
print('Matched rows:', len(matched_tests))
print('Unmatched devices:', len(unmatched_devices))
print('Unmatched tests:', len(unmatched_tests))
print('Output sheets: matched, unmatched')