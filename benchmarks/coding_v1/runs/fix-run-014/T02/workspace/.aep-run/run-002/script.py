from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV as text so Device ID leading zeroes are preserved.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames
    if headers is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in headers:
        raise ValueError('results.csv does not contain a Status column')
    if 'Device ID' not in headers:
        raise ValueError('results.csv does not contain a Device ID column')
    failed_rows = []
    for row in reader:
        if (row.get('Status') or '').strip().lower() == 'fail':
            # Keep all original columns and original order.
            failed_rows.append([row.get(h, '') for h in headers])

# Write workbook.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'
ws.append(headers)
for row_values in failed_rows:
    ws.append(row_values)

# Ensure Device ID column is text in Excel.
device_col_idx = headers.index('Device ID') + 1
for cell in ws.iter_cols(min_col=device_col_idx, max_col=device_col_idx, min_row=1, max_row=ws.max_row):
    for c in cell:
        c.number_format = '@'

wb.save(output_path)

# Verification: reopen output and confirm requirements.
check_wb = load_workbook(output_path, data_only=False)
if 'failed_items' not in check_wb.sheetnames:
    raise AssertionError('Workbook is missing failed_items worksheet')
check_ws = check_wb['failed_items']
actual_headers = [check_ws.cell(row=1, column=i).value for i in range(1, len(headers) + 1)]
if actual_headers != headers:
    raise AssertionError(f'Headers changed: expected {headers}, got {actual_headers}')
actual_rows = []
for r in range(2, check_ws.max_row + 1):
    actual_rows.append([check_ws.cell(row=r, column=c).value for c in range(1, len(headers) + 1)])
if actual_rows != failed_rows:
    raise AssertionError(f'Rows changed or order changed: expected {failed_rows}, got {actual_rows}')
for r in range(2, check_ws.max_row + 1):
    status_value = check_ws.cell(row=r, column=headers.index('Status') + 1).value
    if (status_value or '').strip().lower() != 'fail':
        raise AssertionError(f'Non-fail row found at Excel row {r}: {status_value!r}')
    device_cell = check_ws.cell(row=r, column=device_col_idx)
    if not isinstance(device_cell.value, str):
        raise AssertionError(f'Device ID is not text at Excel row {r}: {device_cell.value!r}')
    if device_cell.number_format != '@':
        raise AssertionError(f'Device ID cell is not formatted as text at Excel row {r}')

print('Created:', output_path)
print('Worksheet:', 'failed_items')
print('Headers:', headers)
print('Filtered row count:', len(failed_rows))
print('Filtered rows:', failed_rows)
print('Verified: output file reopens, worksheet name is failed_items, all original columns are present in original order, row order is preserved, Status matches strip/lower == fail, and Device ID values are text with text number format.')
