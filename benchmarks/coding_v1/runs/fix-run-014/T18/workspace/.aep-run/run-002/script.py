from pathlib import Path
import subprocess
import sys
from openpyxl import load_workbook
import csv

base = Path('.').resolve()
input_dir = base / 'input'
output_dir = base / 'output'
output_dir.mkdir(exist_ok=True)

# Write the repaired program to output/export.py without changing input/export.py.
fixed_code = '''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv
from pathlib import Path

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals = {}
    counts = {}

    with open(args.input, newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            model = row["Model"]
            value_text = row["Value"]
            if value_text is None or value_text.strip() == "":
                continue
            value = float(value_text)
            totals[model] = totals.get(model, 0.0) + value
            counts[model] = counts.get(model, 0) + 1

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])

    output_path = Path(args.output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)


if __name__ == "__main__":
    main()
'''

repaired = output_dir / 'export.py'
repaired.write_text(fixed_code, encoding='utf-8')
print('wrote', repaired)

# Execute exactly in the requested working directory with the requested relative paths.
cmd = [sys.executable, 'output\\export.py', '--input', 'input\\input.csv', '--output', 'output\\summary.xlsx']
print('running:', ' '.join(cmd))
proc = subprocess.run(cmd, cwd=base, text=True, capture_output=True)
print('returncode:', proc.returncode)
if proc.stdout:
    print('stdout:\n' + proc.stdout)
if proc.stderr:
    print('stderr:\n' + proc.stderr)
if proc.returncode != 0:
    raise SystemExit(proc.returncode)

# Independently compute expected results from input.csv for verification.
expected_totals = {}
expected_counts = {}
with (input_dir / 'input.csv').open(newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        val = row['Value']
        if val is None or val.strip() == '':
            continue
        model = row['Model']
        expected_totals[model] = expected_totals.get(model, 0.0) + float(val)
        expected_counts[model] = expected_counts.get(model, 0) + 1
expected_rows = [(m, expected_totals[m] / expected_counts[m], expected_counts[m]) for m in sorted(expected_totals)]

# Reopen output/export.py and output/summary.xlsx to confirm saved outputs.
assert repaired.exists(), 'output/export.py missing'
reopened_code = repaired.read_text(encoding='utf-8')
assert 'parser.add_argument("--input", required=True)' in reopened_code
assert 'parser.add_argument("--output", required=True)' in reopened_code
assert 'row["Model"]' in reopened_code
assert 'row["Value"]' in reopened_code

summary = output_dir / 'summary.xlsx'
assert summary.exists(), 'output/summary.xlsx missing'
wb = load_workbook(summary, data_only=True)
ws = wb.active
headers = [ws.cell(1, c).value for c in range(1, 4)]
actual_rows = []
for r in range(2, ws.max_row + 1):
    actual_rows.append((ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value))

print('headers:', headers)
print('actual_rows:', actual_rows)
print('expected_rows:', expected_rows)
assert ws.title == 'summary', f'unexpected sheet title {ws.title!r}'
assert headers == ['Model', 'Average', 'ValidCount'], headers
assert len(actual_rows) == len(expected_rows), (actual_rows, expected_rows)
for actual, expected in zip(actual_rows, expected_rows):
    am, aa, ac = actual
    em, ea, ec = expected
    assert am == em, (actual, expected)
    assert isinstance(aa, (int, float)), f'Average not numeric: {actual}'
    assert abs(float(aa) - float(ea)) < 1e-12, (actual, expected)
    assert isinstance(ac, int), f'ValidCount not integer: {actual}'
    assert ac == ec, (actual, expected)

print('verification passed: output/export.py preserves --input/--output CLI; output/summary.xlsx has sheet summary, columns Model/Average/ValidCount, numeric averages, integer counts, and blank Value rows excluded.')