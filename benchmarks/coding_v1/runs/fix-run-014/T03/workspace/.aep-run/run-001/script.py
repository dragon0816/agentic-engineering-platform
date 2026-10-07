from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
print('Workspace:', workspace)
print('Input files:')
for p in sorted(input_dir.iterdir()):
    print(' -', p.name, 'size=', p.stat().st_size)

src = input_dir / 'tasks.xlsx'
print('tasks.xlsx exists:', src.exists())
wb = load_workbook(src, read_only=True, data_only=False)
print('Workbook sheets:', wb.sheetnames)
ws = wb['tasks']
print('tasks dimensions:', ws.max_row, ws.max_column)
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
print('Headers:', headers)
print('Sample rows:')
for r in range(2, min(ws.max_row, 6) + 1):
    print([ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)])
wb.close()
