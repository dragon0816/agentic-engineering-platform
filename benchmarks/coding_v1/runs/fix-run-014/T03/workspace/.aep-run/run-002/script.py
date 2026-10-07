from pathlib import Path
from collections import OrderedDict, Counter
from openpyxl import load_workbook, Workbook
import re

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'owners.xlsx'

INVALID_SHEET_CHARS = r'[]:*?/\\'
invalid_re = re.compile(r'[\[\]\:\*\?\/\\]')

def owner_key(value):
    """Grouping key: blank Owner goes to Unassigned; nonblank uses original value."""
    if value is None:
        return 'Unassigned'
    if isinstance(value, str) and value.strip() == '':
        return 'Unassigned'
    return value

def owner_display(key):
    return key

def make_legal_sheet_name(owner, used):
    raw = 'Unassigned' if owner is None or (isinstance(owner, str) and owner.strip() == '') else str(owner)
    # Excel sheet names cannot contain []:*?/\\ and cannot be empty; max 31 chars.
    cleaned = invalid_re.sub('_', raw).strip()
    if cleaned == '':
        cleaned = 'Sheet'
    # Avoid reserved/conflicting Summary name too.
    base = cleaned[:31]
    if base.lower() == 'summary':
        base = (base[:28] + '_01') if len(base) > 28 else base + '_01'
    candidate = base
    n = 1
    while candidate.lower() in used:
        suffix = f'_{n}'
        candidate = base[:31-len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate

# Read source workbook and group rows by Owner.
src_wb = load_workbook(input_path, data_only=False)
if 'tasks' not in src_wb.sheetnames:
    raise ValueError('Input workbook does not contain a tasks worksheet')
src_ws = src_wb['tasks']
headers = [src_ws.cell(row=1, column=c).value for c in range(1, src_ws.max_column + 1)]
try:
    owner_idx = headers.index('Owner') + 1
    status_idx = headers.index('Status') + 1
except ValueError as e:
    raise ValueError('Required columns Owner and Status were not found') from e

groups = OrderedDict()
for row in src_ws.iter_rows(min_row=2, max_row=src_ws.max_row, max_col=src_ws.max_column, values_only=True):
    # Keep all rows, including possible fully blank data rows if present in worksheet range.
    key = owner_key(row[owner_idx - 1])
    groups.setdefault(key, []).append(row)

# Create output workbook.
out_wb = Workbook()
# Remove default sheet; create Summary explicitly.
default = out_wb.active
out_wb.remove(default)
summary_ws = out_wb.create_sheet('Summary')
summary_headers = ['Owner', 'SheetName', 'Total', 'Pass', 'Fail']
summary_ws.append(summary_headers)

used_names = {'summary'}
owner_to_sheet = OrderedDict()
for key, rows in groups.items():
    sheet_name = make_legal_sheet_name(key, used_names)
    owner_to_sheet[key] = sheet_name
    ws = out_wb.create_sheet(sheet_name)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    statuses = Counter(row[status_idx - 1] for row in rows)
    summary_ws.append([owner_display(key), sheet_name, len(rows), statuses.get('Pass', 0), statuses.get('Fail', 0)])

# Simple width adjustment.
for ws in out_wb.worksheets:
    for col_cells in ws.columns:
        max_len = max(len('' if cell.value is None else str(cell.value)) for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 60)

out_wb.save(output_path)
src_wb.close()
out_wb.close()
print('Wrote:', output_path)
print('Groups and sheet names:')
for owner, sheet in owner_to_sheet.items():
    print(f' - {owner!r} -> {sheet}')

# Verification by reopening output file.
check_wb = load_workbook(output_path, data_only=False)
errors = []
if 'Summary' not in check_wb.sheetnames:
    errors.append('Missing Summary worksheet')
summary = check_wb['Summary']
actual_summary_headers = [summary.cell(row=1, column=c).value for c in range(1, summary.max_column + 1)]
if actual_summary_headers != summary_headers:
    errors.append(f'Summary headers mismatch: {actual_summary_headers!r}')

# Validate legal and unique sheet names.
if len(check_wb.sheetnames) != len(set(check_wb.sheetnames)):
    errors.append('Duplicate sheet names found')
for name in check_wb.sheetnames:
    if len(name) > 31 or any(ch in name for ch in INVALID_SHEET_CHARS) or name == '':
        errors.append(f'Illegal sheet name: {name!r}')

# Validate summary and each owner sheet contents/counts.
summary_rows = list(summary.iter_rows(min_row=2, values_only=True))
if len(summary_rows) != len(groups):
    errors.append(f'Summary row count {len(summary_rows)} != group count {len(groups)}')

for idx, (key, rows) in enumerate(groups.items(), start=0):
    if idx >= len(summary_rows):
        continue
    owner_val, sheet_name, total, pass_count, fail_count = summary_rows[idx]
    expected_sheet = owner_to_sheet[key]
    expected_counts = Counter(row[status_idx - 1] for row in rows)
    if owner_val != owner_display(key):
        errors.append(f'Summary owner mismatch for {key!r}: {owner_val!r}')
    if sheet_name != expected_sheet or sheet_name not in check_wb.sheetnames:
        errors.append(f'Summary sheet mapping mismatch for {key!r}: {sheet_name!r} expected {expected_sheet!r}')
        continue
    if total != len(rows) or not isinstance(total, int):
        errors.append(f'Total mismatch/type for {key!r}: {total!r}')
    if pass_count != expected_counts.get('Pass', 0) or not isinstance(pass_count, int):
        errors.append(f'Pass mismatch/type for {key!r}: {pass_count!r}')
    if fail_count != expected_counts.get('Fail', 0) or not isinstance(fail_count, int):
        errors.append(f'Fail mismatch/type for {key!r}: {fail_count!r}')
    ws = check_wb[sheet_name]
    sheet_headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
    if sheet_headers != headers:
        errors.append(f'Headers mismatch in sheet {sheet_name}: {sheet_headers!r}')
    data_rows = list(ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers), values_only=True))
    if data_rows != rows:
        errors.append(f'Data rows mismatch in sheet {sheet_name}: {data_rows!r} != {rows!r}')

# Validate total row count across owner sheets equals source data rows.
source_data_count = sum(len(rows) for rows in groups.values())
output_data_count = sum(check_wb[owner_to_sheet[key]].max_row - 1 for key in groups)
if output_data_count != source_data_count:
    errors.append(f'Output data row count {output_data_count} != source data row count {source_data_count}')

check_wb.close()

if errors:
    print('Verification failed:')
    for err in errors:
        print(' -', err)
    raise SystemExit(1)

print('Verification succeeded.')
print('Checked output file exists, Summary headers Owner/SheetName/Total/Pass/Fail, legal unique sheet names, original headers, original Owner cell values in group sheets, integer counts, status totals, and row counts.')
