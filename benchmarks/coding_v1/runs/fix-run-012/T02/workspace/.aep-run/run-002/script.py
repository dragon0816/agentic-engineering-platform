from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV without modifying input. Keep original headers and row order.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames
    if headers is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in headers:
        raise ValueError('Required column Status not found')
    if 'Device ID' not in headers:
        raise ValueError('Required column Device ID not found')
    failed_rows = []
    total_rows = 0
    for row in reader:
        total_rows += 1
        if (row.get('Status') or '').strip().lower() == 'fail':
            failed_rows.append(row)

wb = Workbook()
ws = wb.active
ws.title = 'failed_items'
ws.append(headers)

for row in failed_rows:
    values = []
    for h in headers:
        value = row.get(h, '')
        if h == 'Device ID' and value is not None:
            value = str(value)
        values.append(value)
    ws.append(values)

# Explicitly format Device ID column as text to preserve leading zeroes.
device_col_idx = headers.index('Device ID') + 1
for cell in ws.iter_cols(min_col=device_col_idx, max_col=device_col_idx, min_row=1, max_row=ws.max_row):
    for c in cell:
        c.number_format = '@'

wb.save(output_path)

# Verification: reopen output and confirm requirements.
wb2 = load_workbook(output_path, data_only=False)
if 'failed_items' not in wb2.sheetnames:
    raise AssertionError('failed_items worksheet missing')
ws2 = wb2['failed_items']
actual_headers = [ws2.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
if actual_headers != headers:
    raise AssertionError(f'Headers changed: {actual_headers} != {headers}')

# Re-read source to compute expected rows and compare exact cell values as strings, preserving order.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    expected = [row for row in reader if (row.get('Status') or '').strip().lower() == 'fail']

actual = []
for r in range(2, ws2.max_row + 1):
    actual.append({h: ws2.cell(row=r, column=i+1).value for i, h in enumerate(headers)})

if len(actual) != len(expected):
    raise AssertionError(f'Row count mismatch: actual {len(actual)} expected {len(expected)}')

for idx, (a, e) in enumerate(zip(actual, expected), start=1):
    for h in headers:
        expected_value = e.get(h, '')
        actual_value = a[h]
        if actual_value is None:
            actual_value = ''
        if str(actual_value) != str(expected_value):
            raise AssertionError(f'Value mismatch row {idx}, column {h}: {actual_value!r} != {expected_value!r}')
    device_cell = ws2.cell(row=idx+1, column=device_col_idx)
    if not isinstance(device_cell.value, str):
        raise AssertionError(f'Device ID row {idx} is not text: {device_cell.value!r} ({type(device_cell.value)})')
    if device_cell.number_format != '@':
        raise AssertionError(f'Device ID row {idx} number_format not text: {device_cell.number_format!r}')

print('Created:', output_path)
print('Input rows:', total_rows)
print('Failed rows written:', len(failed_rows))
print('Worksheet:', ws2.title)
print('Headers:', actual_headers)
print('Device ID values:', [ws2.cell(row=r, column=device_col_idx).value for r in range(2, ws2.max_row + 1)])
print('Verification passed: filtered Status strip/lower == fail, original order/all columns preserved, Device ID is text with leading zeroes retained.')