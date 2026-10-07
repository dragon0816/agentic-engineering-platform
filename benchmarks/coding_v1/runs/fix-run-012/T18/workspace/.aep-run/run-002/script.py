from pathlib import Path
import csv
import subprocess
import sys
import textwrap
from openpyxl import load_workbook

base = Path('.').resolve()
input_dir = base / 'input'
output_dir = base / 'output'
output_dir.mkdir(exist_ok=True)

fixed_code = r'''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals = {}
    counts = {}

    with open(args.input, newline="", encoding="utf-8-sig") as handle:
        for row in csv.DictReader(handle):
            model = row["Model"]
            raw_value = row["Value"]
            if raw_value is None or raw_value.strip() == "":
                continue
            value = float(raw_value)
            totals[model] = totals.get(model, 0.0) + value
            counts[model] = counts.get(model, 0) + 1

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(args.output)


if __name__ == "__main__":
    main()
'''

export_out = output_dir / 'export.py'
export_out.write_text(fixed_code, encoding='utf-8')
print('Wrote fixed script:', export_out)

cmd_display = r'python output\export.py --input input\input.csv --output output\summary.xlsx'
print('Running required command:', cmd_display)
# Run with forward-slash paths for portability while matching the same relative targets.
cmd = [sys.executable, 'output/export.py', '--input', 'input/input.csv', '--output', 'output/summary.xlsx']
result = subprocess.run(cmd, cwd=base, text=True, capture_output=True)
print('Return code:', result.returncode)
if result.stdout:
    print('STDOUT:\n' + result.stdout)
if result.stderr:
    print('STDERR:\n' + result.stderr)
if result.returncode != 0:
    raise SystemExit(result.returncode)

# Compute expected values directly from input.csv.
expected_totals = {}
expected_counts = {}
with (input_dir / 'input.csv').open(newline='', encoding='utf-8-sig') as handle:
    reader = csv.DictReader(handle)
    print('Source columns:', reader.fieldnames)
    for row in reader:
        model = row['Model']
        raw = row['Value']
        if raw is None or raw.strip() == '':
            continue
        expected_totals[model] = expected_totals.get(model, 0.0) + float(raw)
        expected_counts[model] = expected_counts.get(model, 0) + 1
expected_rows = [(m, expected_totals[m] / expected_counts[m], expected_counts[m]) for m in sorted(expected_totals)]
print('Expected rows:', expected_rows)

# Reopen and validate the saved workbook.
summary = output_dir / 'summary.xlsx'
if not summary.exists():
    raise AssertionError('summary.xlsx was not created')
wb = load_workbook(summary, data_only=True)
if 'summary' not in wb.sheetnames:
    raise AssertionError(f'Missing summary sheet; sheets={wb.sheetnames}')
ws = wb['summary']
rows = list(ws.iter_rows(values_only=True))
print('Workbook rows:', rows)

assert rows[0] == ('Model', 'Average', 'ValidCount'), f'Unexpected header: {rows[0]!r}'
actual_rows = rows[1:]
assert len(actual_rows) == len(expected_rows), f'Expected {len(expected_rows)} data rows, got {len(actual_rows)}'
for actual, expected in zip(actual_rows, expected_rows):
    amodel, aavg, acount = actual
    emodel, eavg, ecount = expected
    assert amodel == emodel, (actual, expected)
    assert isinstance(aavg, (int, float)), f'Average is not numeric for {amodel}: {aavg!r}'
    assert abs(float(aavg) - eavg) < 1e-9, (actual, expected)
    assert isinstance(acount, int), f'ValidCount is not int for {amodel}: {acount!r}'
    assert acount == ecount, (actual, expected)

print('Validation passed: output/export.py exists, summary.xlsx sheet/header/row count/numeric averages/integer ValidCount all match input.csv with blank Values excluded.')
