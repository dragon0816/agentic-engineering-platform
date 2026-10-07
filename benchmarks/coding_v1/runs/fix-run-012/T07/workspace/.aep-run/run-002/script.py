from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'specs.md'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Parse product sections beginning with markdown H2 headings.
sections = re.split(r'(?m)^##\s+', text)
rows = []
for section in sections[1:]:
    lines = [line.strip() for line in section.strip().splitlines() if line.strip()]
    if not lines:
        continue
    fields = {}
    for line in lines[1:]:  # skip heading such as P1/P2/P3
        m = re.match(r'^-\s*([^:]+):\s*(.*)$', line)
        if m:
            fields[m.group(1).strip()] = m.group(2).strip()

    def convert_range_to_mhz(freq_text):
        if not freq_text:
            return ('', '')
        # Accept explicit ranges like 2.4–2.5 GHz or 900–930 MHz only.
        m = re.match(r'^\s*([0-9]+(?:\.[0-9]+)?)\s*[–-]\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\s*$', freq_text, re.IGNORECASE)
        if not m:
            return ('', '')
        lo = float(m.group(1))
        hi = float(m.group(2))
        unit = m.group(3).lower()
        factor = 1000 if unit == 'ghz' else 1
        return (lo * factor, hi * factor)

    def bandwidth_to_mhz(bw_text):
        if not bw_text:
            return ''
        m = re.match(r'^\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\s*$', bw_text, re.IGNORECASE)
        if not m:
            return ''
        value = float(m.group(1))
        unit = m.group(2).lower()
        factor = 1000 if unit == 'ghz' else 1
        return value * factor

    def normalize_number(value):
        if value == '':
            return ''
        # Store numeric values as numbers; avoid unnecessary decimals for whole MHz values.
        if abs(value - round(value)) < 1e-9:
            return int(round(value))
        return value

    fmin, fmax = convert_range_to_mhz(fields.get('Frequency', ''))
    bandwidth = bandwidth_to_mhz(fields.get('Bandwidth', ''))
    interface = fields.get('Interface', '')
    if interface:
        # Multiple interfaces are written using slash in source; output semicolon-separated.
        interface = ';'.join(part.strip() for part in re.split(r'\s*/\s*', interface) if part.strip())

    rows.append({
        'Model': fields.get('Model', ''),
        'FrequencyMinMHz': normalize_number(fmin),
        'FrequencyMaxMHz': normalize_number(fmax),
        'Bandwidth': normalize_number(bandwidth),
        'Interface': interface,
    })

wb = Workbook()
ws = wb.active
ws.title = 'specs'
headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'Bandwidth', 'Interface']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(output_path)
print('WROTE', output_path)
print('ROWS_PARSED', len(rows))
for row in rows:
    print(row)

# Verification: reopen workbook and check requirements.
wb2 = load_workbook(output_path, data_only=True)
ws2 = wb2['specs']
read_headers = [cell.value for cell in ws2[1]]
read_rows = []
for r in ws2.iter_rows(min_row=2, values_only=True):
    read_rows.append(dict(zip(read_headers, r)))
print('\nVERIFY headers:', read_headers)
print('VERIFY data rows:', len(read_rows))
for row in read_rows:
    print('VERIFY row:', row)

expected = [
    {'Model': 'RX-100', 'FrequencyMinMHz': 2400, 'FrequencyMaxMHz': 2500, 'Bandwidth': 80, 'Interface': 'USB;LAN'},
    {'Model': 'RX-200', 'FrequencyMinMHz': 900, 'FrequencyMaxMHz': 930, 'Bandwidth': 20, 'Interface': 'LAN'},
    {'Model': 'RX-300', 'FrequencyMinMHz': 5000, 'FrequencyMaxMHz': 6000, 'Bandwidth': None, 'Interface': 'USB'},
]
assert read_headers == headers, f'headers mismatch: {read_headers}'
assert len(read_rows) == 3, f'row count mismatch: {len(read_rows)}'
for got, exp in zip(read_rows, expected):
    for key, val in exp.items():
        assert got[key] == val, f'{key} mismatch: got {got[key]!r}, expected {val!r}'
for row in read_rows:
    assert isinstance(row['Model'], str) and row['Model'], 'Model should be text'
    assert isinstance(row['FrequencyMinMHz'], (int, float)), 'FrequencyMinMHz should be numeric'
    assert isinstance(row['FrequencyMaxMHz'], (int, float)), 'FrequencyMaxMHz should be numeric'
    if row['Bandwidth'] is not None:
        assert isinstance(row['Bandwidth'], (int, float)), 'Bandwidth should be numeric when present'
    assert isinstance(row['Interface'], str) and row['Interface'], 'Interface should be text'
print('VERIFY OK: output file, columns, 3 rows, MHz conversions, semicolon interface separator, and blank missing bandwidth were confirmed.')