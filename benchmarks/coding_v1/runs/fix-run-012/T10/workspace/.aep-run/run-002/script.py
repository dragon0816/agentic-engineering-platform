from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urljoin
import urllib.request
import re
from openpyxl import Workbook, load_workbook

BASE_URL = 'http://127.0.0.1:64212/t10/catalog'
ws = Path.cwd()
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)
out_file = out_dir / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self, base_url):
        super().__init__(convert_charrefs=True)
        self.base_url = base_url
        self.stack = []
        self.in_catalog = False
        self.catalog_depth = None
        self.in_card = False
        self.card_depth = None
        self.current = None
        self.capture = None
        self.products = []

    def _attrs(self, attrs):
        return dict(attrs)

    def handle_starttag(self, tag, attrs):
        attrs_d = self._attrs(attrs)
        self.stack.append((tag, attrs_d))
        depth = len(self.stack)

        if tag == 'main' and attrs_d.get('id') == 'catalog':
            self.in_catalog = True
            self.catalog_depth = depth

        if self.in_catalog and tag == 'div':
            classes = attrs_d.get('class', '').split()
            if 'card' in classes and not self.in_card:
                self.in_card = True
                self.card_depth = depth
                self.current = {'name': '', 'price_text': '', 'url': ''}

        if self.in_catalog and self.in_card and self.current is not None:
            if tag == 'a' and not self.current['url']:
                self.current['url'] = urljoin(self.base_url, attrs_d.get('href', ''))
                self.capture = 'name'
            elif tag == 'span' and 'price' in attrs_d.get('class', '').split():
                self.capture = 'price_text'

    def handle_data(self, data):
        if self.capture and self.current is not None:
            self.current[self.capture] += data

    def handle_endtag(self, tag):
        depth = len(self.stack)
        if self.capture:
            if (self.capture == 'name' and tag == 'a') or (self.capture == 'price_text' and tag == 'span'):
                self.capture = None

        if self.in_card and depth == self.card_depth and tag == 'div':
            if self.current:
                self.current = {k: v.strip() for k, v in self.current.items()}
                self.products.append(self.current)
            self.in_card = False
            self.card_depth = None
            self.current = None
            self.capture = None

        if self.in_catalog and depth == self.catalog_depth and tag == 'main':
            self.in_catalog = False
            self.catalog_depth = None

        if self.stack:
            self.stack.pop()

with urllib.request.urlopen(BASE_URL, timeout=10) as r:
    html = r.read().decode('utf-8', errors='replace')

parser = CatalogParser(BASE_URL)
parser.feed(html)

rows = []
for item in parser.products:
    price_text = item['price_text']
    m = re.match(r'\s*([^\d\-+.,]*?)\s*([\d,]+(?:\.\d+)?)\s*$', price_text)
    if not m:
        raise ValueError(f'Cannot parse price: {price_text!r}')
    currency = m.group(1).strip()
    price_num = float(m.group(2).replace(',', ''))
    if price_num.is_integer():
        price_num = int(price_num)
    rows.append([item['name'], price_num, currency, item['url']])

wb = Workbook()
wsheet = wb.active
wsheet.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
wsheet.append(headers)
for row in rows:
    wsheet.append(row)
wb.save(out_file)
print('WROTE', out_file)
print('ROWS', len(rows))
for row in rows:
    print(row)

# Verification: reopen output and confirm file/columns/order/types/full URLs/no ads.
wb2 = load_workbook(out_file, data_only=True)
sh = wb2['products']
read_headers = [sh.cell(1, c).value for c in range(1, 5)]
assert read_headers == headers, f'headers mismatch: {read_headers}'
read_rows = []
for r in range(2, sh.max_row + 1):
    read_rows.append([sh.cell(r, c).value for c in range(1, 5)])
assert len(read_rows) == 5, f'expected 5 catalog products, got {len(read_rows)}'
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
assert [r[0] for r in read_rows] == expected_names, [r[0] for r in read_rows]
assert all(isinstance(r[1], (int, float)) and not isinstance(r[1], bool) for r in read_rows), 'price must be numeric'
assert [r[2] for r in read_rows] == ['NT$'] * 5, 'currency mismatch'
assert all(isinstance(r[3], str) and r[3].startswith('http://127.0.0.1:64212/') for r in read_rows), 'URLs not complete'
assert not any(r[0] in ('Sponsored Router', '推薦：高速網卡') for r in read_rows), 'ad/recommendation item included'
print('VERIFIED headers, 5 rows, page order, numeric prices, currency, full URLs, and excluded ads/recommendations.')