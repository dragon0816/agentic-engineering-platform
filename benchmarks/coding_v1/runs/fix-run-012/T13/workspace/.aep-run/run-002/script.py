import urllib.request, urllib.error
from urllib.parse import urljoin
from html.parser import HTMLParser
from pathlib import Path
import re, json, time
from openpyxl import Workbook, load_workbook

BASE_URL = 'http://127.0.0.1:64212/t13/list'
workspace = Path.cwd()
outdir = workspace / 'output'
outdir.mkdir(exist_ok=True)
xlsx_path = outdir / 'robust_products.xlsx'
json_path = outdir / 'warnings.json'

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links=[]
        self._href=None
        self._text=[]
    def handle_starttag(self, tag, attrs):
        if tag.lower()=='a':
            d=dict(attrs)
            self._href=d.get('href')
            self._text=[]
    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower()=='a' and self._href is not None:
            name=' '.join(''.join(self._text).split())
            self.links.append((name, self._href))
            self._href=None
            self._text=[]

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts=[]
    def handle_data(self, data):
        s=' '.join(data.split())
        if s:
            self.parts.append(s)
    def text(self):
        return '\n'.join(self.parts)

def fetch_once(url):
    req = urllib.request.Request(url, headers={'User-Agent':'agent-products-fetcher/1.0'})
    with urllib.request.urlopen(req, timeout=10) as resp:
        return resp.status, resp.read().decode('utf-8','replace'), dict(resp.headers)

def fetch_detail_with_retries(url, max_attempts=3):
    attempts=[]
    for i in range(1, max_attempts+1):
        try:
            status, html, headers = fetch_once(url)
            attempts.append({'attempt': i, 'status': status})
            return True, status, html, attempts
        except urllib.error.HTTPError as e:
            status=e.code
            attempts.append({'attempt': i, 'status': status})
            # Requirement: detail page HTTP 500 max 3 tries. Other HTTP failures need not be retried.
            if status != 500 or i == max_attempts:
                try:
                    body=e.read().decode('utf-8','replace')
                except Exception:
                    body=''
                return False, status, body, attempts
            time.sleep(0.2)
        except Exception as e:
            attempts.append({'attempt': i, 'status': type(e).__name__})
            # transient connection error: retry up to same attempt budget
            if i == max_attempts:
                return False, type(e).__name__, '', attempts
            time.sleep(0.2)
    return False, 'unknown', '', attempts

def strip_html(html):
    p=TextParser(); p.feed(html); return p.text()

def extract_labeled(text, labels):
    lines=[ln.strip() for ln in text.splitlines() if ln.strip()]
    for idx, ln in enumerate(lines):
        # Normalize variants like "Price: $123", "Price $123", "CPU - A16"
        for label in labels:
            m=re.match(rf'^{re.escape(label)}\s*[:\-]?\s*(.+)$', ln, re.I)
            if m and m.group(1).strip():
                return m.group(1).strip()
            if ln.lower()==label.lower() and idx+1 < len(lines):
                return lines[idx+1].strip()
    # Search across all text where parser emitted table cells on separate lines.
    joined='\n'.join(lines)
    for label in labels:
        m=re.search(rf'(?im)^{re.escape(label)}\s*$\n\s*(.+)$', joined)
        if m:
            return m.group(1).strip()
    return None

def parse_price_currency(raw):
    if not raw:
        return 'N/A', 'N/A'
    raw=' '.join(raw.split())
    # Capture currency symbols/codes and amount. Supports USD 799, $799, 799 USD, NT$ 1000, etc.
    code_map={'$':'USD','US$':'USD','USD':'USD','€':'EUR','EUR':'EUR','£':'GBP','GBP':'GBP','NT$':'TWD','TWD':'TWD','JPY':'JPY','¥':'JPY'}
    cur=None
    for token in sorted(code_map, key=len, reverse=True):
        if re.search(re.escape(token), raw, re.I):
            cur=code_map[token]
            break
    m=re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', raw)
    price=m.group(0).replace(',','') if m else 'N/A'
    return price, (cur if cur else 'N/A')

def parse_detail(html, fallback_name):
    text=strip_html(html)
    name=extract_labeled(text, ['Name','Product','Product Name']) or fallback_name or 'N/A'
    price_raw=extract_labeled(text, ['Price','Cost'])
    cpu=extract_labeled(text, ['CPU','Processor','Chipset']) or 'N/A'
    price, currency = parse_price_currency(price_raw)
    # If no labeled price, try line containing currency or number near price keyword.
    if price == 'N/A':
        m=re.search(r'(?i)price\s*[:\-]?\s*([^\n]+)', text)
        if m:
            price, currency = parse_price_currency(m.group(1))
    return {'Name': name, 'Price': price, 'Currency': currency, 'CPU': cpu}

# Fetch list.
status, list_html, headers = fetch_once(BASE_URL)
print('list status:', status)
print('list headers content-type:', headers.get('Content-Type') or headers.get('content-type'))
print('list sample:', list_html[:500].replace('\n',' '))
parser=LinkParser(); parser.feed(list_html)
items=[]
for name, href in parser.links:
    if href and name:
        items.append((name, urljoin(BASE_URL, href)))
items=items[:5]
print('product links:', items)

rows=[]
warnings=[]
for list_name, url in items:
    ok, status, html, attempts = fetch_detail_with_retries(url, 3)
    print(f'detail {list_name} {url} attempts={attempts} ok={ok}')
    if ok:
        print('detail sample:', html[:300].replace('\n',' '))
        data=parse_detail(html, list_name)
        row={
            'Name': data.get('Name') or list_name or 'N/A',
            'Price': data.get('Price') or 'N/A',
            'Currency': data.get('Currency') or 'N/A',
            'CPU': data.get('CPU') or 'N/A',
            'URL': url,
            'FetchStatus': 'ok'
        }
        warnings.append({'Name': row['Name'], 'URL': url, 'FetchStatus': 'ok', 'attempts': attempts})
    else:
        row={'Name': list_name or 'N/A', 'Price':'N/A', 'Currency':'N/A', 'CPU':'N/A', 'URL':url, 'FetchStatus':'error'}
        warnings.append({'Name': row['Name'], 'URL': url, 'FetchStatus': 'error', 'attempts': attempts, 'warning': f'detail fetch failed after {len(attempts)} attempt(s), last_status={status}'})
    # Ensure missing price/spec fields are N/A.
    for k in ['Name','Price','Currency','CPU','URL','FetchStatus']:
        if row.get(k) in (None, ''):
            row[k]='N/A'
    rows.append(row)

headers_out=['Name','Price','Currency','CPU','URL','FetchStatus']
wb=Workbook()
ws=wb.active
ws.title='Products'
ws.append(headers_out)
for row in rows:
    ws.append([row[h] for h in headers_out])
wb.save(xlsx_path)
with open(json_path, 'w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Verification: reopen files and check stated requirements.
wb2=load_workbook(xlsx_path, data_only=True)
ws2=wb2['Products']
read_headers=[c.value for c in ws2[1]]
read_rows=list(ws2.iter_rows(min_row=2, values_only=True))
with open(json_path, encoding='utf-8') as f:
    warnings2=json.load(f)
print('VERIFY xlsx exists:', xlsx_path.exists(), 'path:', xlsx_path)
print('VERIFY json exists:', json_path.exists(), 'path:', json_path)
print('VERIFY headers:', read_headers)
print('VERIFY row_count:', len(read_rows))
print('VERIFY warnings_count:', len(warnings2))
print('VERIFY statuses:', [r[5] for r in read_rows])
print('VERIFY no blank cells:', all(all(v not in (None, '') for v in r) for r in read_rows))
print('VERIFY rows:')
for r in read_rows:
    print(r)
assert read_headers == headers_out, read_headers
assert len(read_rows) == 5, len(read_rows)
assert len(warnings2) == 5, len(warnings2)
assert all(r[5] in ('ok','error') for r in read_rows)
assert all(all(v not in (None, '') for v in r) for r in read_rows)
assert all(w.get('FetchStatus') in ('ok','error') for w in warnings2)
print('DONE')