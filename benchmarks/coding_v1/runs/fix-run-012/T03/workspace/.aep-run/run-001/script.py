from pathlib import Path
from openpyxl import load_workbook

ws_path = Path('input') / 'tasks.xlsx'
print('Input exists:', ws_path.exists(), ws_path)
wb = load_workbook(ws_path, read_only=True, data_only=False)
print('Sheet names:', wb.sheetnames)
if 'tasks' in wb.sheetnames:
    ws = wb['tasks']
    rows = list(ws.iter_rows(values_only=True))
    print('Max row/col:', ws.max_row, ws.max_column)
    if rows:
        print('Headers:', rows[0])
        print('Sample rows:')
        for r in rows[1:6]:
            print(r)
wb.close()