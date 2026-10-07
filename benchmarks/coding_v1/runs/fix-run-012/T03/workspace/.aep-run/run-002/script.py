from pathlib import Path
from collections import OrderedDict
from openpyxl import load_workbook, Workbook

INPUT = Path('input') / 'tasks.xlsx'
OUTPUT_DIR = Path('output')
OUTPUT = OUTPUT_DIR / 'owners.xlsx'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

INVALID_CHARS = set('[]:*?/\\')

def owner_key(value):
    if value is None:
        return 'Unassigned'
    if isinstance(value, str) and value.strip() == '':
        return 'Unassigned'
    return value

def owner_label_for_sheet(key):
    return str(key) if key != 'Unassigned' else 'Unassigned'

def sanitize_sheet_name(name, used):
    # Excel worksheet names: <=31 chars and cannot contain []:*?/\\.
    base = ''.join('_' if ch in INVALID_CHARS else ch for ch in str(name))
    base = base.strip()
    if not base:
        base = 'Sheet'
    # Avoid leading/trailing apostrophe, which Excel UI rejects.
    base = base.strip("'") or 'Sheet'
    # Avoid the reserved Summary sheet and duplicate names, case-insensitive.
    max_len = 31
    base = base[:max_len]
    candidate = base
    i = 1
    while candidate.lower() in used:
        suffix = f'_{i}'
        candidate = base[:max_len - len(suffix)] + suffix
        i += 1
    used.add(candidate.lower())
    return candidate

# Read input workbook and collect rows from tasks sheet.
in_wb = load_workbook(INPUT, data_only=False)
if 'tasks' not in in_wb.sheetnames:
    raise ValueError('Input workbook does not contain a tasks worksheet')
in_ws = in_wb['tasks']
rows = list(in_ws.iter_rows(values_only=True))
if not rows:
    raise ValueError('tasks worksheet is empty')
headers = list(rows[0])
required = ['Owner', 'Status']
missing = [c for c in required if c not in headers]
if missing:
    raise ValueError(f'Missing required columns: {missing}; headers={headers}')
owner_idx = headers.index('Owner')
status_idx = headers.index('Status')

# Preserve first-seen order of owner groups.
groups = OrderedDict()
for row in rows[1:]:
    # Normalize row width to header width while preserving cell values.
    normalized = tuple(row[:len(headers)]) + tuple([None] * max(0, len(headers) - len(row)))
    key = owner_key(normalized[owner_idx])
    groups.setdefault(key, []).append(normalized)

# Create output workbook.
out_wb = Workbook()
default_ws = out_wb.active
default_ws.title = 'Summary'
summary_ws = default_ws
summary_headers = ['Owner', 'Total', 'Pass', 'Fail', 'SheetName']
summary_ws.append(summary_headers)

used_sheet_names = {'summary'}
summary_rows = []
for key, group_rows in groups.items():
    sheet_name = sanitize_sheet_name(owner_label_for_sheet(key), used_sheet_names)
    ws = out_wb.create_sheet(title=sheet_name)
    ws.append(headers)
    for row in group_rows:
        ws.append(list(row))
    total = len(group_rows)
    pass_count = sum(1 for row in group_rows if row[status_idx] == 'Pass')
    fail_count = sum(1 for row in group_rows if row[status_idx] == 'Fail')
    summary_rows.append([key, total, pass_count, fail_count, sheet_name])

for r in summary_rows:
    summary_ws.append(r)

out_wb.save(OUTPUT)
out_wb.close()
in_wb.close()

# Verification: reopen output and confirm stated requirements.
check_wb = load_workbook(OUTPUT, data_only=False)
print('Created:', OUTPUT)
print('Output sheets:', check_wb.sheetnames)
assert 'Summary' in check_wb.sheetnames, 'Summary sheet missing'
check_summary = check_wb['Summary']
actual_summary_headers = [cell.value for cell in check_summary[1]]
print('Summary headers:', actual_summary_headers)
assert actual_summary_headers == summary_headers, f'Unexpected Summary headers: {actual_summary_headers}'

# Validate every group sheet and summary row.
summary_by_owner = {}
for r in check_summary.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in r):
        continue
    owner, total, pass_count, fail_count, sheet_name = r
    summary_by_owner[owner] = {
        'Total': total,
        'Pass': pass_count,
        'Fail': fail_count,
        'SheetName': sheet_name,
    }
    assert isinstance(total, int), f'Total for {owner} is not numeric int'
    assert isinstance(pass_count, int), f'Pass for {owner} is not numeric int'
    assert isinstance(fail_count, int), f'Fail for {owner} is not numeric int'
    assert sheet_name in check_wb.sheetnames, f'SheetName {sheet_name} not found'
    assert len(sheet_name) <= 31, f'SheetName too long: {sheet_name}'
    assert not any(ch in sheet_name for ch in INVALID_CHARS), f'Illegal char in SheetName: {sheet_name}'

assert len(summary_by_owner) == len(groups), 'Summary row count does not match owner group count'
for key, group_rows in groups.items():
    assert key in summary_by_owner, f'Owner group {key!r} missing from Summary'
    info = summary_by_owner[key]
    sheet_name = info['SheetName']
    ws = check_wb[sheet_name]
    sheet_headers = [cell.value for cell in ws[1]]
    assert sheet_headers == headers, f'Headers changed on {sheet_name}: {sheet_headers}'
    data = [tuple(row) for row in ws.iter_rows(min_row=2, values_only=True)]
    assert data == list(group_rows), f'Rows not preserved for owner {key!r}'
    assert info['Total'] == len(group_rows), f'Total mismatch for {key!r}'
    assert info['Pass'] == sum(1 for row in group_rows if row[status_idx] == 'Pass'), f'Pass mismatch for {key!r}'
    assert info['Fail'] == sum(1 for row in group_rows if row[status_idx] == 'Fail'), f'Fail mismatch for {key!r}'

print('Verified owner groups:', list(groups.keys()))
print('Verified Summary rows:')
for row in check_summary.iter_rows(min_row=2, values_only=True):
    print(row)
check_wb.close()
print('All checks passed: output file, Summary columns, SheetName mapping, legal unique sheet names, preserved source headers/rows, numeric counts, and row counts.')