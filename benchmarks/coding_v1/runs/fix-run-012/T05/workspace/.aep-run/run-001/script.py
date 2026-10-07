from pathlib import Path
from openpyxl import load_workbook

base = Path('input')
print('Input files:')
for p in sorted(base.glob('*')):
    print(' -', p.name)

for fname in ['devices.xlsx', 'tests.xlsx']:
    path = base / fname
    print(f'\n=== {fname} ===')
    wb = load_workbook(path, read_only=True, data_only=False)
    print('Sheets:', wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=False))
    print('Dimensions:', ws.max_row, 'rows x', ws.max_column, 'cols')
    if rows:
        print('Headers:', [c.value for c in rows[0]])
        for r in rows[1:]:
            vals = []
            for c in r:
                vals.append({'value': c.value, 'type': c.data_type, 'number_format': c.number_format})
            print('Sample:', vals)
    wb.close()