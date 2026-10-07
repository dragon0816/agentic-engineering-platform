from pathlib import Path
from openpyxl import load_workbook, Workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
OUTFILE = OUTPUT / 'merged_devices_tests.xlsx'

DEVICES_FILE = INPUT / 'devices.xlsx'
TESTS_FILE = INPUT / 'tests.xlsx'
ID_COL = 'Device ID'


def read_sheet(path):
    wb = load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [cell.value for cell in ws[1]]
    records = []
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if all(v is None for v in row):
            continue
        rec = {headers[i]: row[i] if i < len(row) else None for i in range(len(headers))}
        if ID_COL in rec and rec[ID_COL] is not None:
            rec[ID_COL] = str(rec[ID_COL])
        records.append(rec)
    wb.close()
    return headers, records

dev_headers, devices = read_sheet(DEVICES_FILE)
test_headers, tests = read_sheet(TESTS_FILE)

if ID_COL not in dev_headers:
    raise ValueError(f'{DEVICES_FILE.name} missing {ID_COL}')
if ID_COL not in test_headers:
    raise ValueError(f'{TESTS_FILE.name} missing {ID_COL}')

# devices 中每個 ID 唯一；若輸入不符，明確報錯，避免合併結果不確定。
device_by_id = {}
for rec in devices:
    did = rec.get(ID_COL)
    if did in device_by_id:
        raise ValueError(f'Duplicate Device ID in devices.xlsx: {did}')
    device_by_id[did] = rec

test_ids = {rec.get(ID_COL) for rec in tests}
matched_tests = [t for t in tests if t.get(ID_COL) in device_by_id]
unmatched_tests = [t for t in tests if t.get(ID_COL) not in device_by_id]
unmatched_devices = [d for d in devices if d.get(ID_COL) not in test_ids]

# Matched: each matched test, followed by matching device fields. Keep source column names;
# use the shared Device ID once to avoid duplicate ID columns.
matched_headers = list(test_headers) + [h for h in dev_headers if h != ID_COL]
# Unmatched: union of original fields, plus Source marking devices/tests.
unmatched_headers = ['Source'] + list(dev_headers) + [h for h in test_headers if h not in dev_headers]

wb_out = Workbook()
ws_m = wb_out.active
ws_m.title = 'matched'
ws_u = wb_out.create_sheet('unmatched')

def write_header(ws, headers):
    for c, h in enumerate(headers, 1):
        ws.cell(row=1, column=c, value=h)

def write_value(ws, row, col, header, value):
    cell = ws.cell(row=row, column=col, value=value)
    if header == ID_COL:
        cell.number_format = '@'
        if value is not None:
            cell.value = str(value)
    return cell

write_header(ws_m, matched_headers)
for r_idx, test in enumerate(matched_tests, 2):
    dev = device_by_id[test[ID_COL]]
    row_values = []
    for h in test_headers:
        row_values.append(test.get(h))
    for h in dev_headers:
        if h != ID_COL:
            row_values.append(dev.get(h))
    for c_idx, (h, v) in enumerate(zip(matched_headers, row_values), 1):
        write_value(ws_m, r_idx, c_idx, h, v)

write_header(ws_u, unmatched_headers)
r_idx = 2
for dev in unmatched_devices:
    row = {'Source': 'devices'}
    row.update(dev)
    for c_idx, h in enumerate(unmatched_headers, 1):
        write_value(ws_u, r_idx, c_idx, h, row.get(h))
    r_idx += 1
for test in unmatched_tests:
    row = {'Source': 'tests'}
    row.update(test)
    for c_idx, h in enumerate(unmatched_headers, 1):
        write_value(ws_u, r_idx, c_idx, h, row.get(h))
    r_idx += 1

# Simple formatting only: autosize-ish widths, freeze header rows.
for ws in (ws_m, ws_u):
    ws.freeze_panes = 'A2'
    for col_cells in ws.columns:
        letter = col_cells[0].column_letter
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 40)

wb_out.save(OUTFILE)

print('Created:', OUTFILE)
print('Devices:', len(devices), 'Tests:', len(tests))
print('Matched rows:', len(matched_tests))
print('Unmatched devices:', len(unmatched_devices), 'Unmatched tests:', len(unmatched_tests))
print('Matched headers:', matched_headers)
print('Unmatched headers:', unmatched_headers)
