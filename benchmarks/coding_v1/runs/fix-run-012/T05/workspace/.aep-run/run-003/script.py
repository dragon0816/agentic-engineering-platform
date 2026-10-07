from pathlib import Path
from openpyxl import load_workbook

out = Path('output') / 'merged_devices_tests.xlsx'
if not out.exists():
    raise AssertionError('Output file missing')

wb = load_workbook(out, data_only=False)
required_sheets = ['matched', 'unmatched']
if wb.sheetnames[:2] != required_sheets:
    raise AssertionError(f'Expected first sheets {required_sheets}, got {wb.sheetnames}')

ws_m = wb['matched']
ws_u = wb['unmatched']
matched_headers = [c.value for c in ws_m[1]]
unmatched_headers = [c.value for c in ws_u[1]]
if matched_headers != ['Device ID', 'Test', 'Result', 'Name', 'Location']:
    raise AssertionError(f'Unexpected matched headers: {matched_headers}')
if unmatched_headers != ['Source', 'Device ID', 'Name', 'Location', 'Test', 'Result']:
    raise AssertionError(f'Unexpected unmatched headers: {unmatched_headers}')

matched_rows = list(ws_m.iter_rows(min_row=2, values_only=True))
unmatched_rows = list(ws_u.iter_rows(min_row=2, values_only=True))
if len(matched_rows) != 3:
    raise AssertionError(f'Expected 3 matched rows, got {len(matched_rows)}')
if len(unmatched_rows) != 2:
    raise AssertionError(f'Expected 2 unmatched rows, got {len(unmatched_rows)}')

expected_matched = [
    ('001', 'RF-1', 'Pass', 'Alpha', 'Lab A'),
    ('001', 'RF-2', 'Fail', 'Alpha', 'Lab A'),
    ('002', 'RF-1', 'Pass', 'Beta', 'Lab B'),
]
if matched_rows != expected_matched:
    raise AssertionError(f'Matched rows mismatch: {matched_rows}')

expected_unmatched = [
    ('devices', '003', 'Gamma', 'Lab C', None, None),
    ('tests', '999', None, None, 'RF-9', 'Fail'),
]
if unmatched_rows != expected_unmatched:
    raise AssertionError(f'Unmatched rows mismatch: {unmatched_rows}')

# Check Device ID cells are stored as strings with text format and keep leading zero where applicable.
for ws, id_col in [(ws_m, 1), (ws_u, 2)]:
    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=id_col)
        if cell.value is not None and not isinstance(cell.value, str):
            raise AssertionError(f'{ws.title}!{cell.coordinate} Device ID is not text: {cell.value!r}')
        if cell.value is not None and cell.number_format != '@':
            raise AssertionError(f'{ws.title}!{cell.coordinate} Device ID number_format not text: {cell.number_format}')

if ws_m['A2'].value != '001' or ws_m['A3'].value != '001' or ws_u['B2'].value != '003':
    raise AssertionError('Leading zeros were not preserved')

sources = [r[0] for r in unmatched_rows]
if sources != ['devices', 'tests']:
    raise AssertionError(f'Unexpected Source values: {sources}')

wb.close()
print('CHECK PASSED')
print('Output:', out.resolve())
print('Sheets:', required_sheets)
print('Matched rows:', len(matched_rows))
print('Unmatched rows:', len(unmatched_rows))
print('Device ID text/leading-zero checks passed')
print('Source values:', sources)
