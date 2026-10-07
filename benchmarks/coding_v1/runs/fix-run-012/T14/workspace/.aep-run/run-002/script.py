from pathlib import Path
import csv
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'large_results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

required_headers = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']
summary_headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
overall_headers = ['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']

stats = defaultdict(lambda: {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': Decimal('0')})
overall = {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': Decimal('0')}
status_counts = defaultdict(int)
row_count = 0

with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise ValueError(f'Missing required headers: {missing}; found {headers}')
    for line_no, row in enumerate(reader, start=2):
        row_count += 1
        model = row['Device Model']
        status = row['Status']
        status_counts[status] += 1
        try:
            duration = Decimal(str(row['DurationSeconds']).strip())
        except (InvalidOperation, AttributeError) as exc:
            raise ValueError(f'Invalid DurationSeconds at CSV line {line_no}: {row["DurationSeconds"]!r}') from exc

        stats[model]['total'] += 1
        stats[model]['duration_sum'] += duration
        overall['total'] += 1
        overall['duration_sum'] += duration

        if status == 'Pass':
            stats[model]['pass'] += 1
            overall['pass'] += 1
        elif status == 'Fail':
            stats[model]['fail'] += 1
            overall['fail'] += 1
        else:
            # Total and average still include all rows; PassCount/FailCount only count exact Pass/Fail.
            pass

wb = Workbook()
ws = wb.active
ws.title = 'Summary'
ws.append(summary_headers)
for cell in ws[1]:
    cell.font = Font(bold=True)

for model in sorted(stats.keys()):
    s = stats[model]
    total = s['total']
    pass_count = s['pass']
    fail_count = s['fail']
    pass_rate = pass_count / total if total else 0
    avg_duration = float(s['duration_sum'] / total) if total else 0
    ws.append([model, total, pass_count, fail_count, pass_rate, avg_duration])

for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
    row[0].number_format = '0.00%'
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    row[0].number_format = '0.00'
for col in range(1, 7):
    ws.column_dimensions[chr(64 + col)].width = 20

ows = wb.create_sheet('Overall')
ows.append(overall_headers)
for cell in ows[1]:
    cell.font = Font(bold=True)
total = overall['total']
pass_count = overall['pass']
fail_count = overall['fail']
pass_rate = pass_count / total if total else 0
avg_duration = float(overall['duration_sum'] / total) if total else 0
ows.append([total, pass_count, fail_count, pass_rate, avg_duration])
ows['D2'].number_format = '0.00%'
ows['E2'].number_format = '0.00'
for col in range(1, 6):
    ows.column_dimensions[chr(64 + col)].width = 20

wb.save(output_path)
print(f'Created: {output_path}')
print(f'Input rows: {row_count}')
print(f'Device models: {len(stats)}')
print('Status counts:', dict(status_counts))

# Verification: reopen the saved workbook and independently confirm requirements.
check_wb = load_workbook(output_path, data_only=True)
expected_sheets = ['Summary', 'Overall']
actual_sheets = check_wb.sheetnames
assert actual_sheets == expected_sheets, f'Expected sheets {expected_sheets}, got {actual_sheets}'

check_ws = check_wb['Summary']
read_summary_headers = [cell.value for cell in check_ws[1]]
assert read_summary_headers == summary_headers, f'Summary headers mismatch: {read_summary_headers}'
assert check_ws.max_row == len(stats) + 1, f'Summary row count mismatch: {check_ws.max_row} vs {len(stats)+1}'
assert check_ws.max_column == len(summary_headers), f'Summary column count mismatch: {check_ws.max_column}'

# Build expected float values and compare all rows.
for row in check_ws.iter_rows(min_row=2, values_only=True):
    model, total_v, pass_v, fail_v, rate_v, avg_v = row
    s = stats[model]
    exp_total = s['total']
    exp_pass = s['pass']
    exp_fail = s['fail']
    exp_rate = exp_pass / exp_total if exp_total else 0
    exp_avg = float(s['duration_sum'] / exp_total) if exp_total else 0
    assert isinstance(model, str), f'Device Model not text: {model!r}'
    assert total_v == exp_total, f'Total mismatch for {model}: {total_v} vs {exp_total}'
    assert pass_v == exp_pass, f'PassCount mismatch for {model}: {pass_v} vs {exp_pass}'
    assert fail_v == exp_fail, f'FailCount mismatch for {model}: {fail_v} vs {exp_fail}'
    assert abs(rate_v - exp_rate) < 1e-12, f'PassRate mismatch for {model}: {rate_v} vs {exp_rate}'
    assert abs(avg_v - exp_avg) < 1e-12, f'AvgDurationSeconds mismatch for {model}: {avg_v} vs {exp_avg}'

check_ows = check_wb['Overall']
read_overall_headers = [cell.value for cell in check_ows[1]]
assert read_overall_headers == overall_headers, f'Overall headers mismatch: {read_overall_headers}'
assert check_ows.max_row == 2, f'Overall row count mismatch: {check_ows.max_row}'
assert check_ows.max_column == len(overall_headers), f'Overall column count mismatch: {check_ows.max_column}'
ovalues = [check_ows.cell(2, col).value for col in range(1, 6)]
exp_overall = [
    overall['total'],
    overall['pass'],
    overall['fail'],
    overall['pass'] / overall['total'] if overall['total'] else 0,
    float(overall['duration_sum'] / overall['total']) if overall['total'] else 0,
]
for got, exp, name in zip(ovalues, exp_overall, overall_headers):
    if isinstance(exp, float):
        assert abs(got - exp) < 1e-12, f'Overall {name} mismatch: {got} vs {exp}'
    else:
        assert got == exp, f'Overall {name} mismatch: {got} vs {exp}'

print('Verification passed:')
print(' - output file exists and reopens')
print(' - sheets:', actual_sheets)
print(' - Summary headers and row count:', read_summary_headers, check_ws.max_row - 1)
print(' - Overall headers and one aggregate row:', read_overall_headers)
print(' - numeric totals/pass/fail, PassRate, and AvgDurationSeconds match source data')
print('Overall values:', dict(zip(overall_headers, exp_overall)))
