from pathlib import Path
import csv
from collections import Counter
from openpyxl import load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_path = workspace / 'output' / 'converted.xlsx'

with open(input_path, 'r', encoding='utf-8-sig', newline='') as f:
    csv_rows = list(csv.reader(f))

wb = load_workbook(output_path, data_only=False)
ws = wb.active
xlsx_rows = []
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
    xlsx_rows.append(['' if cell.value is None else str(cell.value) for cell in row])

headers_ok = xlsx_rows[0] == csv_rows[0]
row_count_ok = len(xlsx_rows) == len(csv_rows)
col_count_ok = all(len(r) == len(csv_rows[0]) for r in xlsx_rows) and ws.max_column == len(csv_rows[0])
order_and_values_ok = xlsx_rows == csv_rows
counter_ok = Counter(tuple(r) for r in xlsx_rows[1:]) == Counter(tuple(r) for r in csv_rows[1:])
no_index_ok = ws.max_column == len(csv_rows[0]) and xlsx_rows[0] == ['Device ID', '名稱', '備註', 'Value']

device_idx = csv_rows[0].index('Device ID') + 1
device_values = [ws.cell(row=r, column=device_idx).value for r in range(2, ws.max_row + 1)]
device_formats = [ws.cell(row=r, column=device_idx).number_format for r in range(1, ws.max_row + 1)]
device_text_ok = all((v is None or isinstance(v, str)) for v in device_values) and all(fmt == '@' for fmt in device_formats)
leading_zero_ok = str(ws.cell(row=2, column=device_idx).value) == '0001'
quoted_comma_ok = ws.cell(row=3, column=3).value == '含, 逗號'
quoted_newline_ok = ws.cell(row=4, column=3).value == '第一行\n第二行'
blank_ok = ws.cell(row=4, column=4).value is None and ws.cell(row=6, column=2).value is None and ws.cell(row=6, column=3).value is None
duplicate_ok = xlsx_rows[1] == xlsx_rows[4] == ['0001', '測試機甲', '正常', '10']

checks = {
    'output_file_exists': output_path.exists(),
    'headers_preserved_in_order': headers_ok,
    'row_count_preserved_including_duplicates': row_count_ok,
    'column_count_preserved_no_extra_index': col_count_ok and no_index_ok,
    'all_values_and_row_order_preserved': order_and_values_ok,
    'duplicate_multiset_preserved': counter_ok and duplicate_ok,
    'device_id_text_format_and_type': device_text_ok,
    'device_id_leading_zero_preserved': leading_zero_ok,
    'quoted_comma_field_preserved': quoted_comma_ok,
    'quoted_newline_field_preserved': quoted_newline_ok,
    'empty_values_are_blank_cells': blank_ok,
}
for name, ok in checks.items():
    print(f'{name}: {ok}')
print('xlsx_dimensions:', ws.max_row, 'rows x', ws.max_column, 'cols')
print('headers:', xlsx_rows[0])
print('device_values:', device_values)
print('device_formats:', device_formats)
if not all(checks.values()):
    raise AssertionError('One or more output checks failed')
