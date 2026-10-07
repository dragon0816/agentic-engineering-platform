from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read CSV with utf-8-sig to support UTF-8 BOM and newline='' to correctly handle quoted newlines.
with open(input_path, 'r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    rows = list(reader)

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
try:
    device_col_idx = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_col_idx = None

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'

# Write every field as-is. Empty strings remain blank cells. No index column is added.
for r_idx, row in enumerate(rows, start=1):
    for c_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx)
        if value == '':
            cell.value = None
        else:
            cell.value = str(value)
        # Device ID must be stored/displayed as text to preserve leading zeroes.
        if device_col_idx is not None and c_idx == device_col_idx:
            cell.number_format = '@'
            if r_idx > 1 and value != '':
                cell.value = str(value)

# Lightly emphasize header without changing content/order.
for cell in ws[1]:
    cell.font = Font(bold=True)

wb.save(output_path)
print('created:', output_path)
print('source_rows_including_header:', len(rows))
print('source_columns:', len(headers))
print('headers:', headers)
print('data_rows:', len(rows) - 1)
