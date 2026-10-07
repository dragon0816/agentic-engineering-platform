import csv
import json
import os
import uuid
from pathlib import Path
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
INPUT_DIR = WORKSPACE / 'input'
OUTPUT_DIR = WORKSPACE / 'output'
MANIFEST = INPUT_DIR / 'manifest.json'
THRESHOLD = 20

print(f'workspace={WORKSPACE}')
print(f'input_dir={INPUT_DIR} exists={INPUT_DIR.exists()}')
print(f'output_dir={OUTPUT_DIR} exists={OUTPUT_DIR.exists()}')
print('input entries:')
for p in sorted(INPUT_DIR.rglob('*')):
    if p.is_file():
        print(f'  {p.relative_to(WORKSPACE)} size={p.stat().st_size}')

if not MANIFEST.exists():
    raise FileNotFoundError(f'Missing manifest: {MANIFEST}')
manifest = json.loads(MANIFEST.read_text(encoding='utf-8'))
if 'results.csv' not in manifest:
    raise KeyError('manifest.json must map results.csv to the actual input CSV path')

csv_path = WORKSPACE / manifest['results.csv']
if not csv_path.exists():
    raise FileNotFoundError(f'Mapped results.csv does not exist: {csv_path}')

with csv_path.open('r', encoding='utf-8-sig', newline='') as fh:
    reader = csv.DictReader(fh)
    fieldnames = reader.fieldnames
    print(f'csv_path={csv_path.relative_to(WORKSPACE)}')
    print(f'csv_headers={fieldnames}')
    required = ['Device ID', 'Status', 'Value']
    missing = [c for c in required if c not in (fieldnames or [])]
    if missing:
        raise ValueError(f'CSV missing required columns: {missing}')
    rows = list(reader)

print(f'total_rows={len(rows)}')
for i, row in enumerate(rows[:5], 1):
    print(f'sample_row_{i}={row}')

fail_rows = []
pass_count = 0
fail_count = 0
urgent_count = 0
for row_num, row in enumerate(rows, start=2):
    status = row['Status']
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        try:
            value_num = float(row['Value'])
        except Exception as exc:
            raise ValueError(f'Row {row_num} has non-numeric Value: {row["Value"]!r}') from exc
        urgent = value_num > THRESHOLD
        if urgent:
            urgent_count += 1
        fail_rows.append({
            'Device ID': row['Device ID'],
            'Status': row['Status'],
            'Value': value_num if not value_num.is_integer() else int(value_num),
            'Urgent': urgent,
        })

OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
failures_path = OUTPUT_DIR / 'failures.xlsx'
summary_path = OUTPUT_DIR / 'summary.md'

# Write to temporary files first, then atomically replace the required outputs.
# This avoids leaving a corrupt final workbook if a prior run failed midway.
run_token = uuid.uuid4().hex
failures_tmp = OUTPUT_DIR / f'.failures.{run_token}.tmp.xlsx'
summary_tmp = OUTPUT_DIR / f'.summary.{run_token}.tmp.md'

try:
    wb = Workbook()
    ws = wb.active
    ws.title = 'failures'
    headers = ['Device ID', 'Status', 'Value', 'Urgent']
    ws.append(headers)
    for row in fail_rows:
        ws.append([row['Device ID'], row['Status'], row['Value'], row['Urgent']])
    wb.save(failures_tmp)

    summary_text = '\n'.join([
        '# Summary',
        '',
        f'- Total: {len(rows)}',
        f'- Pass: {pass_count}',
        f'- Fail: {fail_count}',
        f'- Urgent: {urgent_count}',
        '',
    ])
    summary_tmp.write_text(summary_text, encoding='utf-8')

    # Replace any stale output from a previous run. If the destination is open/locked,
    # report a clear blocker instead of silently creating a differently named file.
    try:
        os.replace(failures_tmp, failures_path)
        os.replace(summary_tmp, summary_path)
    except PermissionError as exc:
        raise PermissionError(
            f'Cannot replace required output file. Close any program using {failures_path} or {summary_path}, then rerun.'
        ) from exc
finally:
    for tmp in (failures_tmp, summary_tmp):
        try:
            if tmp.exists():
                tmp.unlink()
        except Exception:
            pass

# Reopen and verify every stated requirement from saved outputs.
check_wb = load_workbook(failures_path, data_only=True)
check_ws = check_wb['failures']
check_headers = [cell.value for cell in next(check_ws.iter_rows(min_row=1, max_row=1))]
if check_headers != headers:
    raise AssertionError(f'failures.xlsx headers mismatch: {check_headers}')
check_rows = list(check_ws.iter_rows(min_row=2, values_only=True))
if len(check_rows) != fail_count:
    raise AssertionError(f'failures.xlsx row count {len(check_rows)} != Fail count {fail_count}')
for idx, excel_row in enumerate(check_rows, start=2):
    device_id, status, value, urgent = excel_row
    if not isinstance(device_id, str):
        raise AssertionError(f'Row {idx} Device ID is not text: {device_id!r}')
    if status != 'Fail':
        raise AssertionError(f'Row {idx} Status is not Fail: {status!r}')
    if not isinstance(value, (int, float)):
        raise AssertionError(f'Row {idx} Value is not numeric: {value!r}')
    if not isinstance(urgent, bool):
        raise AssertionError(f'Row {idx} Urgent is not boolean: {urgent!r}')
    if urgent != (float(value) > THRESHOLD):
        raise AssertionError(f'Row {idx} Urgent {urgent!r} does not match Value > {THRESHOLD}')

saved_summary = summary_path.read_text(encoding='utf-8')
expected_lines = [
    f'- Total: {len(rows)}',
    f'- Pass: {pass_count}',
    f'- Fail: {fail_count}',
    f'- Urgent: {urgent_count}',
]
for line in expected_lines:
    if line not in saved_summary:
        raise AssertionError(f'summary.md missing line: {line}')

print(f'wrote={failures_path.relative_to(WORKSPACE)} bytes={failures_path.stat().st_size}')
print(f'wrote={summary_path.relative_to(WORKSPACE)} bytes={summary_path.stat().st_size}')
print('checked failures.xlsx: file exists, sheet failures, columns Device ID/Status/Value/Urgent, all rows Status=Fail, Value numeric, Urgent boolean Value>20, row count matches Fail count')
print('checked summary.md: Total, Pass, Fail, Urgent counts match CSV')
print(f'counts total={len(rows)} pass={pass_count} fail={fail_count} urgent={urgent_count}')
