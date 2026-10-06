import os, re, urllib.request, urllib.parse, html
from html.parser import HTMLParser
from pathlib import Path
from openpyxl import Workbook, load_workbook

BASE_LIST_URL = 'http://127.0.0.1:49758/t11/list'
WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
OUTFILE = OUTPUT / 'details.xlsx'

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._current_href = None
        self._text = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            d = dict(attrs)
            self._current_href = d.get('href')
            self._text = []
    def handle_data(self, data):
        if self._current_href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._current_href is not None:
            self.links.append((self._current_href, ''.join(self._text).strip()))
            self._current_href = None
            self._text = []

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, data):
        s = data.strip()
        if s:
            self.parts.append(s)
    def text(self):
        return '\n'.join(self.parts)

def fetch(url):
    with urllib.request.urlopen(url, timeout=10) as r:
        return r.read().decode('utf-8', errors='replace')

def page_text(markup):
    p = TextParser(); p.feed(markup); return p.text()

def clean_value(v):
    if v is None:
        return ''
    v = html.unescape(str(v)).strip()
    v = re.sub(r'\s+', ' ', v)
    return v.strip(' :-\t\r\n')

def parse_labeled_values(markup):
    # Turn common separators/tags into lines, then parse Label: value / Label value pairs.
    s = markup
    s = re.sub(r'(?i)<\s*br\s*/?\s*>', '\n', s)
    s = re.sub(r'(?i)</\s*(li|p|div|tr|h[1-6]|dt|dd)\s*>', '\n', s)
    s = re.sub(r'(?i)<\s*/\s*t[dh]\s*>', ': ', s)
    txt = page_text(s)
    lines = [clean_value(x) for x in txt.splitlines() if clean_value(x)]
    data = {}
    labels = {'name':'Name','price':'Price','currency':'Currency','cpu':'CPU','memory':'Memory','memorygb':'MemoryGB','memory gb':'MemoryGB'}
    for line in lines:
        m = re.match(r'^(Name|Price|Currency|CPU|MemoryGB|Memory\s*GB|Memory)\s*[:：\-]?\s*(.*)$', line, re.I)
        if m:
            key_raw = m.group(1).lower().replace('\xa0',' ')
            key_raw = re.sub(r'\s+', ' ', key_raw)
            key = labels.get(key_raw, labels.get(key_raw.replace(' ', ''), m.group(1)))
            val = clean_value(m.group(2))
            if val:
                data[key] = val
    # Also parse table rows directly to catch markup where text parser inserts weird separators.
    for m in re.finditer(r'(?is)<tr[^>]*>\s*<t[dh][^>]*>(.*?)</t[dh]>\s*<td[^>]*>(.*?)</td>\s*</tr>', markup):
        k = clean_value(re.sub(r'<[^>]+>', ' ', m.group(1)))
        v = clean_value(re.sub(r'<[^>]+>', ' ', m.group(2)))
        kl = k.lower()
        if kl in ('name','price','currency','cpu','memory','memorygb','memory gb'):
            data[labels.get(kl, labels.get(kl.replace(' ', ''), k))] = v
    return data, txt

def normalize_memory_gb(raw):
    raw = clean_value(raw)
    if not raw:
        return ''
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)', raw.replace(',', ''))
    if not m:
        return ''
    val = float(m.group(1))
    lower = raw.lower()
    if re.search(r'\btb\b|terabyte', lower):
        val *= 1024
    elif re.search(r'\bmb\b|megabyte', lower):
        val /= 1024
    # default/no unit is GB
    if abs(val - round(val)) < 1e-9:
        return int(round(val))
    return val

def normalize_price(raw):
    raw = clean_value(raw)
    if not raw:
        return ''
    m = re.search(r'-?\d+(?:,\d{3})*(?:\.\d+)?|-?\d+(?:\.\d+)?', raw)
    if not m:
        return raw
    n = m.group(0).replace(',', '')
    try:
        f = float(n)
        return int(f) if f.is_integer() else f
    except ValueError:
        return n

def split_price_currency(price_raw, currency_raw):
    price = normalize_price(price_raw)
    currency = clean_value(currency_raw)
    if not currency and price_raw:
        up = price_raw.upper()
        m = re.search(r'\b(USD|EUR|TWD|JPY|GBP|CNY|NTD)\b|[$€£¥]', up)
        if m:
            token = m.group(0)
            currency = {'$':'USD','€':'EUR','£':'GBP','¥':'JPY'}.get(token, token)
    return price, currency

print('Fetching list:', BASE_LIST_URL)
list_html = fetch(BASE_LIST_URL)
link_parser = LinkParser(); link_parser.feed(list_html)
item_urls = []
for href, label in link_parser.links:
    full = urllib.parse.urljoin(BASE_LIST_URL, href)
    if '/t11/item/' in full and full not in item_urls:
        item_urls.append(full)
print('Found detail URLs:', item_urls)

rows = []
for url in item_urls:
    markup = fetch(url)
    print('\nDETAIL', url, 'first 1000 chars:\n', markup[:1000])
    data, txt = parse_labeled_values(markup)
    print('Parsed text:', txt.replace('\n', ' | '))
    print('Parsed data:', data)
    price, currency = split_price_currency(data.get('Price',''), data.get('Currency',''))
    row = {
        'Name': clean_value(data.get('Name','')),
        'Price': price,
        'Currency': currency,
        'CPU': clean_value(data.get('CPU','')),
        'MemoryGB': normalize_memory_gb(data.get('MemoryGB') or data.get('Memory','')),
        'URL': url,
    }
    rows.append(row)

wb = Workbook()
ws = wb.active
ws.title = 'details'
headers = ['Name','Price','Currency','CPU','MemoryGB','URL']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(OUTFILE)
print('\nWrote', OUTFILE)

# Verification: reopen output and confirm stated requirements.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['details']
out_headers = [c.value for c in ws2[1]]
out_rows = list(ws2.iter_rows(min_row=2, values_only=True))
print('CHECK headers:', out_headers)
print('CHECK row_count:', len(out_rows), 'expected:', len(item_urls))
print('CHECK rows:')
for r in out_rows:
    print(r)
assert out_headers == headers, f'headers mismatch: {out_headers}'
assert len(out_rows) == len(item_urls), 'row count mismatch'
for r in out_rows:
    assert len(r) == len(headers)
    # openpyxl returns None for intentionally blank cells; that is the Excel blank placeholder.
    assert isinstance(r[5], str) and r[5].startswith('http://127.0.0.1:49758/t11/item/')
    if r[1] is not None:
        assert isinstance(r[1], (int, float)), f'Price not numeric: {r[1]!r}'
    if r[4] is not None:
        assert isinstance(r[4], (int, float)), f'MemoryGB not numeric: {r[4]!r}'
print('CHECK passed: details.xlsx has requested columns, one row per detail URL, numeric Price/MemoryGB where present, and blanks for missing values.')