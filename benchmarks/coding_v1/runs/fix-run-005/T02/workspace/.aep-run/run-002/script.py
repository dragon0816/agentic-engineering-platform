from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV, keep original column names and row order.
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    fieldnames = reader.fieldnames
    if not fieldnames:
        raise ValueError('results.csv has no header row')
    if 'Status' not in fieldnames:
        raise ValueError('results.csv is missing required Status column')
    failed_rows = [row for row in reader if (row.get('Status') or '').strip().lower() == 'fail']

# Write workbook.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'

# Header
for col_idx, name in enumerate(fieldnames, start=1):
    cell = ws.cell(row=1, column=col_idx, value=name)
    cell.font = Font(bold=True)

# Data; preserve Device ID as text to retain leading zeros.
device_id_col_idx = fieldnames.index('Device ID') + 1 if 'Device ID' in fieldnames else None
for row_idx, row in enumerate(failed_rows, start=2):
    for col_idx, name in enumerate(fieldnames, start=1):
        value = row.get(name, '')
        cell = ws.cell(row=row_idx, column=col_idx, value=value)
        if col_idx == device_id_col_idx:
            cell.number_format = '@'
            cell.value = '' if value is None else str(value)

if device_id_col_idx is not None:
    # Format whole Device ID column as text, including header/data range.
    col_letter = ws.cell(row=1, column=device_id_col_idx).column_letter
    for cell in ws[col_letter]:
        cell.number_format = '@'

wb.save(output_path)

# Verification from saved output.
wb2 = load_workbook(output_path, data_only=False)
if 'failed_items' not in wb2.sheetnames:
    raise AssertionError('missing failed_items worksheet')
ws2 = wb2['failed_items']
headers = [ws2.cell(row=1, column=i).value for i in range(1, ws2.max_column + 1)]
if headers != fieldnames:
    raise AssertionError(f'headers differ: expected {fieldnames}, got {headers}')

saved_rows = []
for r in range(2, ws2.max_row + 1):
    saved_rows.append({headers[c-1]: ws2.cell(row=r, column=c).value for c in range(1, ws2.max_column + 1)})

if len(saved_rows) != len(failed_rows):
    raise AssertionError(f'row count differs: expected {len(failed_rows)}, got {len(saved_rows)}')

for idx, (expected, actual) in enumerate(zip(failed_rows, saved_rows), start=2):
    for name in fieldnames:
        expected_value = expected.get(name, '')
        actual_value = actual.get(name)
        if actual_value is None:
            actual_value = ''
        if str(actual_value) != str(expected_value):
            raise AssertionError(f'row {idx} col {name}: expected {expected_value!r}, got {actual_value!r}')
    if device_id_col_idx is not None:
        cell = ws2.cell(row=idx, column=device_id_col_idx)
        if cell.data_type != 's':
            raise AssertionError(f'Device ID at row {idx} is not stored as text: data_type={cell.data_type}')
        if not str(cell.value).startswith('0'):
            raise AssertionError(f'Device ID leading zero not preserved at row {idx}: {cell.value!r}')

# Verify Status filter condition and order explicitly.
statuses = [row['Status'] for row in saved_rows]
device_ids = [row['Device ID'] for row in saved_rows] if 'Device ID' in headers else []
if any((s or '').strip().lower() != 'fail' for s in statuses):
    raise AssertionError(f'non-fail status found: {statuses}')

print('Wrote:', output_path)
print('Worksheet:', ws2.title)
print('Headers:', headers)
print('Failed row count:', len(saved_rows))
print('Device IDs in saved order:', device_ids)
print('Statuses:', statuses)
print('Device ID cell types:', [ws2.cell(row=r, column=device_id_col_idx).data_type for r in range(2, ws2.max_row + 1)] if device_id_col_idx else 'no Device ID column')