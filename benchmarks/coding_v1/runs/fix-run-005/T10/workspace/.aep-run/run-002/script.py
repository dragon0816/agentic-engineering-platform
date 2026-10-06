from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urljoin
import urllib.request
import re
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
CATALOG_URL = 'http://127.0.0.1:49758/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_catalog = False
        self.catalog_depth = 0
        self.in_card = False
        self.card_depth = 0
        self.current = None
        self.capture = None
        self.products = []

    def _attrs(self, attrs):
        return dict(attrs)

    def handle_starttag(self, tag, attrs):
        attrs = self._attrs(attrs)
        if tag == 'main' and attrs.get('id') == 'catalog':
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if self.in_catalog:
            self.catalog_depth += 1
            classes = set(attrs.get('class', '').split())
            if tag == 'div' and 'card' in classes:
                self.in_card = True
                self.card_depth = 1
                self.current = {'name': '', 'href': '', 'price_text': ''}
                return
            if self.in_card:
                self.card_depth += 1
                if tag == 'a' and not self.current.get('href'):
                    self.current['href'] = attrs.get('href', '')
                    self.capture = 'name'
                elif tag == 'span' and 'price' in classes:
                    self.capture = 'price_text'

    def handle_endtag(self, tag):
        if self.in_card:
            if self.capture == 'name' and tag == 'a':
                self.capture = None
            elif self.capture == 'price_text' and tag == 'span':
                self.capture = None
            self.card_depth -= 1
            if self.card_depth == 0:
                if self.current:
                    self.products.append(self.current)
                self.current = None
                self.in_card = False
                return
        if self.in_catalog:
            self.catalog_depth -= 1
            if self.catalog_depth == 0:
                self.in_catalog = False

    def handle_data(self, data):
        if self.in_card and self.current is not None and self.capture:
            self.current[self.capture] += data.strip()

def parse_price(price_text):
    s = price_text.strip()
    m = re.match(r'^([^\d\-+]*)([\d,]+(?:\.\d+)?)', s)
    if not m:
        raise ValueError(f'Cannot parse price: {price_text!r}')
    currency = m.group(1).strip()
    value_s = m.group(2).replace(',', '')
    value = float(value_s) if '.' in value_s else int(value_s)
    return currency, value

with urllib.request.urlopen(CATALOG_URL, timeout=10) as response:
    html = response.read().decode('utf-8', errors='replace')

parser = CatalogParser()
parser.feed(html)

rows = []
for item in parser.products:
    currency, price = parse_price(item['price_text'])
    rows.append({
        '產品名稱': item['name'],
        '價格': price,
        '幣別': currency,
        '產品網址': urljoin(CATALOG_URL, item['href']),
    })

wb = Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(OUTFILE)

# Verification: reopen saved workbook and check every stated requirement from output file.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
saved_headers = [ws2.cell(1, c).value for c in range(1, ws2.max_column + 1)]
saved_rows = []
for r in range(2, ws2.max_row + 1):
    saved_rows.append([ws2.cell(r, c).value for c in range(1, ws2.max_column + 1)])

expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
expected_prices = [1299, 2450, 99, 560, 12000]
expected_currency = ['NT$'] * 5
expected_urls = [urljoin(CATALOG_URL, href) for href in ['/p/router-a','/p/switch-b','/p/cable-c','/p/adapter-d','/p/antenna-e']]

assert OUTFILE.exists(), 'products.xlsx was not created'
assert saved_headers == headers, f'headers mismatch: {saved_headers}'
assert ws2.max_row == 6, f'row count mismatch: {ws2.max_row - 1}'
assert ws2.max_column == 4, f'column count mismatch: {ws2.max_column}'
assert [r[0] for r in saved_rows] == expected_names, 'product order/names mismatch'
assert [r[1] for r in saved_rows] == expected_prices, 'prices mismatch'
assert all(isinstance(r[1], (int, float)) for r in saved_rows), 'price column contains non-numeric values'
assert [r[2] for r in saved_rows] == expected_currency, 'currency mismatch'
assert [r[3] for r in saved_rows] == expected_urls, 'URL mismatch'
assert all(str(r[3]).startswith('http://127.0.0.1:49758/') for r in saved_rows), 'not all URLs are full URLs'
assert all('Sponsored' not in r[0] and '推薦' not in r[0] for r in saved_rows), 'ad/recommendation content included'

print(f'Wrote {OUTFILE}')
print('Verified workbook:')
print('- sheet:', ws2.title)
print('- headers:', saved_headers)
print('- product rows:', len(saved_rows))
print('- prices numeric:', [type(r[1]).__name__ for r in saved_rows])
print('- URLs:', [r[3] for r in saved_rows])
print('- names in page order:', [r[0] for r in saved_rows])
