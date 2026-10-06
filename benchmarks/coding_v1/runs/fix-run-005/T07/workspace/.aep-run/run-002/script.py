from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'specs.md'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split into product sections headed by Markdown H2 (## P1, etc.).
sections = re.split(r'(?m)^##\s+', text)
products = []
for sec in sections[1:]:
    lines = [line.strip() for line in sec.splitlines() if line.strip()]
    fields = {}
    for line in lines:
        m = re.match(r'^-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            fields[m.group(1).strip().lower()] = m.group(2).strip()
    if fields:
        products.append(fields)

unit_to_mhz = {
    'hz': 1e-6,
    'khz': 1e-3,
    'mhz': 1.0,
    'ghz': 1000.0,
}

def parse_frequency_range(value):
    """Return (min_mhz, max_mhz), or blanks where unavailable. No guessing."""
    if not value:
        return '', ''
    s = value.strip().replace('–', '-').replace('—', '-')
    # Supports ranges where a single unit appears after the max value, e.g. 2.4-2.5 GHz.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*-\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', s, re.I)
    if not m:
        return '', ''
    lo = float(m.group(1))
    hi = float(m.group(2))
    unit = m.group(3).lower()
    factor = unit_to_mhz[unit]
    return lo * factor, hi * factor

def parse_bandwidth_mhz(value):
    if not value:
        return ''
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz|kHz|Hz)\b', value, re.I)
    if not m:
        return ''
    return float(m.group(1)) * unit_to_mhz[m.group(2).lower()]

def format_interface(value):
    if not value:
        return ''
    # Multiple interfaces may be separated by slash or comma in the source; preserve names, join with semicolons.
    parts = [p.strip() for p in re.split(r'\s*(?:/|,|;)\s*', value.strip()) if p.strip()]
    return ';'.join(parts)

rows = []
for p in products:
    fmin, fmax = parse_frequency_range(p.get('frequency', ''))
    rows.append({
        'Model': p.get('model', ''),
        'FrequencyMinMHz': fmin,
        'FrequencyMaxMHz': fmax,
        'BandwidthMHz': parse_bandwidth_mhz(p.get('bandwidth', '')),
        'Interface': format_interface(p.get('interface', '')),
    })

wb = Workbook()
ws = wb.active
ws.title = 'specs'
headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])

# Basic formatting: keep text IDs as text and numeric MHz as numbers; blanks remain blank.
for cell in ws['A']:
    cell.number_format = '@'
for row in ws.iter_rows(min_row=2, min_col=2, max_col=4):
    for cell in row:
        if cell.value != '':
            cell.number_format = '0.###'
wb.save(output_path)

# Re-open and validate saved output against the request.
wb2 = load_workbook(output_path, data_only=True)
ws2 = wb2['specs']
actual_headers = [c.value for c in ws2[1]]
records = []
for row in ws2.iter_rows(min_row=2, values_only=True):
    records.append(dict(zip(actual_headers, row)))

print('Wrote:', output_path)
print('Sheet:', ws2.title)
print('Headers:', actual_headers)
print('Row count:', len(records))
for rec in records:
    print(rec)

# Assertions verify: required file/sheet/columns, 3 rows, model text IDs, numeric frequency/bandwidth when present,
# missing bandwidth remains blank, interfaces joined by semicolon.
assert output_path.exists(), 'specs.xlsx was not created'
assert actual_headers == headers, f'headers mismatch: {actual_headers}'
assert len(records) == 3, f'expected 3 products, got {len(records)}'
expected = [
    {'Model': 'RX-100', 'FrequencyMinMHz': 2400, 'FrequencyMaxMHz': 2500, 'BandwidthMHz': 80, 'Interface': 'USB;LAN'},
    {'Model': 'RX-200', 'FrequencyMinMHz': 900, 'FrequencyMaxMHz': 930, 'BandwidthMHz': 20, 'Interface': 'LAN'},
    {'Model': 'RX-300', 'FrequencyMinMHz': 5000, 'FrequencyMaxMHz': 6000, 'BandwidthMHz': None, 'Interface': 'USB'},
]
for rec, exp in zip(records, expected):
    assert rec['Model'] == exp['Model'] and isinstance(rec['Model'], str)
    assert rec['Interface'] == exp['Interface']
    for col in ('FrequencyMinMHz', 'FrequencyMaxMHz'):
        assert isinstance(rec[col], (int, float)), f'{col} is not numeric for {rec}'
        assert abs(rec[col] - exp[col]) < 1e-9
    if exp['BandwidthMHz'] is None:
        assert rec['BandwidthMHz'] is None, f'missing bandwidth should be blank for {rec}'
    else:
        assert isinstance(rec['BandwidthMHz'], (int, float)), f'BandwidthMHz is not numeric for {rec}'
        assert abs(rec['BandwidthMHz'] - exp['BandwidthMHz']) < 1e-9
print('Validation passed: specs.xlsx has 3 rows, required columns, numeric MHz values, semicolon-separated interfaces, and blank missing bandwidth.')