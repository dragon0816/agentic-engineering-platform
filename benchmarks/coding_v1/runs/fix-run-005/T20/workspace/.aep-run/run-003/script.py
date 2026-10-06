import csv
import json
import os
from pathlib import Path
from openpyxl import Workbook, load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
workspace = Path(params.get('workspace') or r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\coding_v1\runs\fix-run-005\T20\workspace')
threshold = float(params['urgent_threshold'])
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
failures_path = output_dir / 'failures.xlsx'
summary_path = output_dir / 'summary.md'

print('INSPECTING INPUTS')
print(f'workspace={workspace}')
print(f'input_dir_exists={input_dir.exists()} output_dir_exists={output_dir.exists()}')
input_files = sorted([p.name for p in input_dir.iterdir()]) if input_dir.exists() else []
print('input_files=' + ', '.join(input_files))

sop_path = input_dir / 'supplied_sop.md'
if not sop_path.exists():
    raise FileNotFoundError(f'Missing SOP: {sop_path}')
sop_text = sop_path.read_text(encoding='utf-8-sig')
print('sop_first_lines:')
print('\n'.join(sop_text.splitlines()[:10]))

manifest_path = input_dir / 'manifest.json'
manifest_text = ''
if manifest_path.exists():
    manifest_text = manifest_path.read_text(encoding='utf-8-sig')
    print('manifest.json:')
    print(manifest_text[:4000])

def materialize_from_manifest(text):
    if not text.strip():
        return None
    target = output_dir / '_materialized_results.csv'
    obj = json.loads(text)
    if isinstance(obj, dict):
        for k, v in obj.items():
            if isinstance(k, str) and k.lower().endswith('.csv') and isinstance(v, str):
                target.write_text(v, encoding='utf-8', newline='')
                return target
            if isinstance(v, str) and 'Device ID' in v and 'Status' in v and 'Value' in v and '\n' in v:
                target.write_text(v, encoding='utf-8', newline='')
                return target
        for k in ['results.csv','results','csv','data','rows','records']:
            v = obj.get(k)
            if isinstance(v, list) and all(isinstance(r, dict) for r in v):
                with target.open('w', encoding='utf-8', newline='') as f:
                    writer = csv.DictWriter(f, fieldnames=['Device ID','Status','Value'])
                    writer.writeheader()
                    for r in v:
                        writer.writerow({c: r.get(c, '') for c in ['Device ID','Status','Value']})
                return target
            if isinstance(v, str) and 'Device ID' in v and 'Status' in v and 'Value' in v:
                target.write_text(v, encoding='utf-8', newline='')
                return target
    return None

materialized = None
try:
    materialized = materialize_from_manifest(manifest_text)
    if materialized:
        print(f'materialized_csv={materialized.relative_to(workspace)}')
except Exception as exc:
    print(f'manifest_parse_or_materialize_warning={type(exc).__name__}: {exc}')

def is_excluded(p: Path) -> bool:
    parts = {part.lower() for part in p.parts}
    return 'output' in parts or '.aep-run' in parts

candidates = []
preferred = input_dir / 'results.csv'
if preferred.exists():
    candidates.append(preferred)
for p in input_dir.rglob('*.csv') if input_dir.exists() else []:
    if p not in candidates:
        candidates.append(p)
if materialized and materialized.exists() and materialized not in candidates:
    candidates.append(materialized)
for p in workspace.rglob('*.csv'):
    if p not in candidates and not is_excluded(p):
        candidates.append(p)

print('csv_candidates=' + ', '.join(str(p.relative_to(workspace)) for p in candidates))
if not candidates:
    raise FileNotFoundError('No CSV files were found under input/ or the workspace, and manifest.json did not contain CSV data. The SOP requires input/results.csv with columns Device ID, Status, Value. Current input files are: ' + ', '.join(input_files))

required = ['Device ID', 'Status', 'Value']
selected = None
for csv_path in candidates:
    with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f)
        fn = reader.fieldnames or []
        sample = []
        for i, row in enumerate(reader):
            if i < 3:
                sample.append(row)
            else:
                break
    print(f'candidate {csv_path.relative_to(workspace)} headers={fn} sample_rows={sample}')
    if all(col in fn for col in required):
        selected = csv_path
        break
if selected is None:
    raise ValueError('No CSV with required columns Device ID, Status, Value was found.')

with selected.open('r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.DictReader(f))
print(f'selected_csv={selected.relative_to(workspace)} rows={len(rows)}')

total = len(rows)
pass_count = 0
fail_count = 0
urgent_count = 0
failure_rows = []
for idx, row in enumerate(rows, start=2):
    status = (row.get('Status') or '').strip()
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        raw_value = row.get('Value')
        try:
            numeric_value = float(str(raw_value).strip())
        except Exception as exc:
            raise ValueError(f'Value is not numeric at CSV row {idx}: {raw_value!r}') from exc
        urgent = numeric_value > threshold
        if urgent:
            urgent_count += 1
        failure_rows.append({'Device ID': row.get('Device ID', ''), 'Status': row.get('Status', ''), 'Value': numeric_value, 'Urgent': bool(urgent)})

wb = Workbook()
ws = wb.active
ws.title = 'failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for r in failure_rows:
    ws.append([r['Device ID'], r['Status'], r['Value'], r['Urgent']])
wb.save(failures_path)

summary_text = '\n'.join(['# Summary','',f'- Total: {total}',f'- Pass: {pass_count}',f'- Fail: {fail_count}',f'- Urgent: {urgent_count}','',f'Source CSV: `{selected.relative_to(workspace)}`',f'Urgent rule: `Value > {int(threshold) if threshold.is_integer() else threshold}`']) + '\n'
summary_path.write_text(summary_text, encoding='utf-8')

print('VERIFYING OUTPUTS')
wb2 = load_workbook(failures_path, data_only=True)
ws2 = wb2['failures']
actual_headers = [cell.value for cell in ws2[1]]
print(f'xlsx_headers={actual_headers}')
if actual_headers != headers:
    raise AssertionError(f'Wrong failures.xlsx headers: {actual_headers}')
actual_rows = list(ws2.iter_rows(min_row=2, values_only=True))
print(f'xlsx_failure_rows={len(actual_rows)}')
if len(actual_rows) != fail_count:
    raise AssertionError(f'Expected {fail_count} failure rows, got {len(actual_rows)}')
for i, row in enumerate(actual_rows, start=2):
    _, status, value, urgent = row
    if status != 'Fail':
        raise AssertionError(f'Non-Fail row in failures.xlsx row {i}: {row}')
    if not isinstance(value, (int, float)):
        raise AssertionError(f'Value is not numeric in failures.xlsx row {i}: {value!r}')
    if not isinstance(urgent, bool):
        raise AssertionError(f'Urgent is not boolean in failures.xlsx row {i}: {urgent!r}')
    if urgent != (float(value) > threshold):
        raise AssertionError(f'Urgent rule mismatch in failures.xlsx row {i}: {row}')
summary_saved = summary_path.read_text(encoding='utf-8')
print('summary.md:')
print(summary_saved)
for label, expected in [('Total', total), ('Pass', pass_count), ('Fail', fail_count), ('Urgent', urgent_count)]:
    token = f'- {label}: {expected}'
    if token not in summary_saved:
        raise AssertionError(f'Missing summary token: {token}')
print('DONE')
print(f'created={failures_path.relative_to(workspace)}, {summary_path.relative_to(workspace)}')
