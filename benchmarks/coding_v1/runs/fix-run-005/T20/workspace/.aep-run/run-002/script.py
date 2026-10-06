import csv
import json
import os
import re
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook, load_workbook

params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = json.load(open(params_path, 'r', encoding='utf-8')) if params_path and os.path.exists(params_path) else {}
workspace = Path(params.get('workspace') or r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\coding_v1\runs\fix-run-005\T20\workspace')
threshold = float(params.get('urgent_threshold', 20))
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

print('INSPECTING INPUTS')
print(f'workspace={workspace}')
print(f'input_dir_exists={input_dir.exists()} output_dir_exists={output_dir.exists()}')
if not input_dir.exists():
    raise FileNotFoundError(f'Missing input directory: {input_dir}')

all_input_files = sorted([p for p in input_dir.iterdir() if p.is_file()])
print('input_files=' + ', '.join(p.name for p in all_input_files))

sop_path = input_dir / 'supplied_sop.md'
if not sop_path.exists():
    alt = workspace / 'supplied_sop.md'
    if alt.exists():
        sop_path = alt
    else:
        raise FileNotFoundError('supplied_sop.md not found in input/ or workspace root')
sop_text = sop_path.read_text(encoding='utf-8-sig')
print('sop_first_lines:')
for line in sop_text.splitlines()[:20]:
    print(line)

csv_paths = sorted(input_dir.glob('*.csv'))
if not csv_paths:
    raise FileNotFoundError(f'No CSV files found in {input_dir}')
print('csv_files=' + ', '.join(p.name for p in csv_paths))

# Read CSV files without modifying inputs. Preserve source column names/order.
all_rows = []
source_headers = None
for csv_path in csv_paths:
    with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
        sample = f.read(4096)
        f.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(f, dialect=dialect)
        if not reader.fieldnames:
            raise ValueError(f'CSV has no header: {csv_path.name}')
        headers = [h if h is not None else '' for h in reader.fieldnames]
        print(f'{csv_path.name} headers={headers}')
        preview = []
        for idx, row in enumerate(reader, start=1):
            clean = {k: (v if v is not None else '') for k, v in row.items() if k is not None}
            if idx <= 3:
                preview.append(clean)
            clean['_source_file'] = csv_path.name
            all_rows.append(clean)
        print(f'{csv_path.name} sample_rows={preview}')
        if source_headers is None:
            source_headers = headers
        elif headers != source_headers:
            raise ValueError('All CSV files must have the same header columns to preserve source column names')

if source_headers is None:
    raise ValueError('No readable CSV headers found')

# Detect required columns. Prefer common names and SOP wording.
def find_col(candidates):
    lower_map = {c.strip().lower(): c for c in source_headers}
    for cand in candidates:
        if cand.lower() in lower_map:
            return lower_map[cand.lower()]
    for col in source_headers:
        normalized = re.sub(r'[^a-z0-9]+', '', col.lower())
        for cand in candidates:
            if normalized == re.sub(r'[^a-z0-9]+', '', cand.lower()):
                return col
    return None

status_col = find_col(['Status', 'Result', 'Test Result', 'Outcome'])
value_col = find_col(['Value', 'Measured Value', 'Measurement', 'Reading'])
if status_col is None:
    raise ValueError(f'Could not find a status/result column in CSV headers: {source_headers}')
if value_col is None:
    raise ValueError(f'Could not find a Value column in CSV headers: {source_headers}')
print(f'detected status_col={status_col} value_col={value_col}')

# Keep all Fail rows. Urgent true iff Value > threshold; otherwise false.
def is_fail(value):
    return str(value).strip().lower() == 'fail'

def to_float(value):
    text = str(value).strip()
    if text == '':
        return None
    m = re.search(r'-?\d+(?:\.\d+)?', text.replace(',', ''))
    if not m:
        return None
    return float(m.group(0))

fail_rows = []
for row in all_rows:
    if is_fail(row.get(status_col, '')):
        numeric_value = to_float(row.get(value_col, ''))
        urgent = bool(numeric_value is not None and numeric_value > threshold)
        out = {h: row.get(h, '') for h in source_headers}
        out['Urgent'] = urgent
        fail_rows.append(out)

# Write failures.xlsx.
xlsx_path = output_dir / 'failures.xlsx'
headers_out = list(source_headers) + ['Urgent']
wb = Workbook()
ws = wb.active
ws.title = 'failures'
ws.append(headers_out)
for row in fail_rows:
    ws.append([row.get(h, '') for h in source_headers] + [bool(row['Urgent'])])
for col_cells in ws.columns:
    max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
    ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 10), 60)
wb.save(xlsx_path)

# Write summary.md.
total_rows = len(all_rows)
fail_count = len(fail_rows)
urgent_count = sum(1 for r in fail_rows if r['Urgent'])
nonurgent_count = fail_count - urgent_count
summary_path = output_dir / 'summary.md'
summary_lines = [
    '# Failure Summary',
    '',
    f'- Source CSV files: {len(csv_paths)}',
    f'- Total input rows: {total_rows}',
    f'- Fail rows: {fail_count}',
    f'- Urgent threshold: Value > {threshold:g}',
    f'- Urgent failures: {urgent_count}',
    f'- Non-urgent failures: {nonurgent_count}',
    '',
    '## CSV files processed',
]
summary_lines.extend([f'- {p.name}' for p in csv_paths])
summary_lines.extend(['', '## Output files', '- failures.xlsx', '- summary.md', ''])
summary_path.write_text('\n'.join(summary_lines), encoding='utf-8')

# Verification from saved outputs.
print('VERIFYING OUTPUTS')
if not xlsx_path.exists():
    raise AssertionError('failures.xlsx was not created')
if not summary_path.exists():
    raise AssertionError('summary.md was not created')
wb2 = load_workbook(xlsx_path, data_only=True)
ws2 = wb2['failures'] if 'failures' in wb2.sheetnames else wb2.active
saved_headers = [cell.value for cell in next(ws2.iter_rows(min_row=1, max_row=1))]
if saved_headers != headers_out:
    raise AssertionError(f'Unexpected XLSX headers: {saved_headers} != {headers_out}')
saved_rows = list(ws2.iter_rows(min_row=2, values_only=True))
if len(saved_rows) != fail_count:
    raise AssertionError(f'XLSX row count {len(saved_rows)} != expected fail count {fail_count}')
urgent_idx = saved_headers.index('Urgent')
status_idx = saved_headers.index(status_col)
value_idx = saved_headers.index(value_col)
for i, saved in enumerate(saved_rows, start=2):
    if str(saved[status_idx]).strip().lower() != 'fail':
        raise AssertionError(f'Row {i} is not Fail: {saved[status_idx]}')
    v = to_float(saved[value_idx])
    expected = bool(v is not None and v > threshold)
    if saved[urgent_idx] is not expected:
        raise AssertionError(f'Row {i} Urgent {saved[urgent_idx]} != expected {expected}')
summary_text = summary_path.read_text(encoding='utf-8')
for required in ['# Failure Summary', f'- Fail rows: {fail_count}', f'- Urgent failures: {urgent_count}']:
    if required not in summary_text:
        raise AssertionError(f'Missing summary text: {required}')
print(f'created={xlsx_path}')
print(f'created={summary_path}')
print(f'checked headers={headers_out}')
print(f'checked total_rows={total_rows} fail_rows={fail_count} urgent={urgent_count} nonurgent={nonurgent_count}')
print('DONE')
