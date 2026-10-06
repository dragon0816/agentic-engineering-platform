from pathlib import Path
from openpyxl import load_workbook

ws_dir = Path.cwd()
print('Workspace:', ws_dir)
print('Input files:', [p.name for p in (ws_dir/'input').iterdir()])
path = ws_dir/'input'/'tasks.xlsx'
print('tasks.xlsx exists:', path.exists())
wb = load_workbook(path, data_only=False, read_only=True)
print('Sheet names:', wb.sheetnames)
ws = wb['tasks']
headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
print('Headers:', headers)
print('Max rows:', ws.max_row, 'Max cols:', ws.max_column)
for i, row in enumerate(ws.iter_rows(min_row=2, max_row=min(ws.max_row, 8), values_only=True), start=2):
    print('Row', i, row)
wb.close()