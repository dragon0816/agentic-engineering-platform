from pathlib import Path
from collections import OrderedDict, Counter
import re
from openpyxl import load_workbook, Workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'owners.xlsx'

INVALID_SHEET_CHARS = r'[]:*?/\\'
INVALID_RE = re.compile(r'[\[\]\:\*\?\/\\]')


def is_blank_owner(value):
    return value is None or (isinstance(value, str) and value.strip() == '')


def owner_display(value):
    return 'Unassigned' if is_blank_owner(value) else value


def base_sheet_name(owner_value):
    text = 'Unassigned' if is_blank_owner(owner_value) else str(owner_value)
    text = INVALID_RE.sub('_', text)
    text = text.strip()
    if not text:
        text = 'Unassigned'
    # Excel also behaves badly with leading/trailing apostrophes in sheet names.
    text = text.strip("'") or 'Unassigned'
    # Avoid the special reserved History sheet name in Excel.
    if text.lower() == 'history':
        text = 'History_'
    return text[:31]


def unique_sheet_name(owner_value, used):
    base = base_sheet_name(owner_value)
    candidate = base
    n = 1
    while candidate in used:
        suffix = f'_{n}'
        candidate = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(candidate)
    return candidate

# Read source workbook.
src_wb = load_workbook(input_path, data_only=False, read_only=True)
if 'tasks' not in src_wb.sheetnames:
    raise SystemExit('Input workbook does not contain a tasks worksheet')
src_ws = src_wb['tasks']
headers = [cell.value for cell in next(src_ws.iter_rows(min_row=1, max_row=1))]
try:
    owner_idx = headers.index('Owner')
    status_idx = headers.index('Status')
except ValueError as exc:
    raise SystemExit('tasks worksheet must include Owner and Status columns') from exc

# Group rows in first-seen owner order. Blank owners share the Unassigned group.
groups = OrderedDict()
for row in src_ws.iter_rows(min_row=2, values_only=True):
    # Skip completely blank trailing rows, if any.
    if all(v is None for v in row):
        continue
    owner_value = row[owner_idx]
    group_key = '__UNASSIGNED__' if is_blank_owner(owner_value) else owner_value
    if group_key not in groups:
        groups[group_key] = {'owner_value': owner_value, 'rows': []}
    groups[group_key]['rows'].append(tuple(row))
src_wb.close()

# Create output workbook.
out_wb = Workbook()
# Remove default sheet; create Summary explicitly first.
default_ws = out_wb.active
out_wb.remove(default_ws)
summary_ws = out_wb.create_sheet('Summary')
summary_headers = ['Owner', 'Total', 'Pass', 'Fail', 'SheetName']
summary_ws.append(summary_headers)

used_sheet_names = {'Summary'}
summary_rows = []

for group in groups.values():
    owner_value = group['owner_value']
    rows = group['rows']
    sheet_name = unique_sheet_name(owner_value, used_sheet_names)
    ws = out_wb.create_sheet(sheet_name)
    ws.append(headers)
    for row in rows:
        ws.append(list(row))
    counts = Counter(row[status_idx] for row in rows)
    summary_owner = owner_display(owner_value)
    summary_row = [summary_owner, len(rows), counts.get('Pass', 0), counts.get('Fail', 0), sheet_name]
    summary_rows.append(summary_row)
    summary_ws.append(summary_row)

# Light usability formatting: freeze header row and set simple widths.
for ws in out_wb.worksheets:
    ws.freeze_panes = 'A2'
    for col_cells in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(max_len + 2, 10), 40)

out_wb.save(output_path)
out_wb.close()
print('Wrote:', output_path)
print('Groups written:', len(groups))
print('Summary rows:', summary_rows)

# Reopen and verify output requirements from the saved file.
check_wb = load_workbook(output_path, data_only=False, read_only=True)
errors = []
if 'Summary' not in check_wb.sheetnames:
    errors.append('Missing Summary worksheet')

# Legal and unique sheet names.
if len(check_wb.sheetnames) != len(set(check_wb.sheetnames)):
    errors.append('Worksheet names are not unique')
for name in check_wb.sheetnames:
    if len(name) > 31 or INVALID_RE.search(name) or not name:
        errors.append(f'Illegal worksheet name: {name!r}')

if 'Summary' in check_wb.sheetnames:
    sws = check_wb['Summary']
    out_summary_headers = [c.value for c in next(sws.iter_rows(min_row=1, max_row=1))]
    if out_summary_headers != summary_headers:
        errors.append(f'Summary headers mismatch: {out_summary_headers!r}')
    saved_summary = [tuple(r) for r in sws.iter_rows(min_row=2, values_only=True)]
    if len(saved_summary) != len(groups):
        errors.append(f'Summary row count {len(saved_summary)} != group count {len(groups)}')

    for owner, total, passed, failed, sheet_name in saved_summary:
        if not isinstance(total, int) or not isinstance(passed, int) or not isinstance(failed, int):
            errors.append(f'Non-integer count in summary row: {(owner, total, passed, failed, sheet_name)!r}')
        if sheet_name not in check_wb.sheetnames:
            errors.append(f'Summary SheetName does not exist: {sheet_name!r}')
            continue
        gws = check_wb[sheet_name]
        group_headers = [c.value for c in next(gws.iter_rows(min_row=1, max_row=1))]
        if group_headers != headers:
            errors.append(f'Headers mismatch in {sheet_name!r}: {group_headers!r}')
        data_rows = list(gws.iter_rows(min_row=2, values_only=True))
        if len(data_rows) != total:
            errors.append(f'Row count mismatch in {sheet_name!r}: {len(data_rows)} != {total}')
        actual_pass = sum(1 for r in data_rows if r[status_idx] == 'Pass')
        actual_fail = sum(1 for r in data_rows if r[status_idx] == 'Fail')
        if actual_pass != passed or actual_fail != failed:
            errors.append(f'Pass/Fail count mismatch in {sheet_name!r}: {(actual_pass, actual_fail)} != {(passed, failed)}')
        if owner == 'Unassigned':
            if not all(is_blank_owner(r[owner_idx]) for r in data_rows):
                errors.append('Unassigned sheet contains non-blank Owner value')
        else:
            # Owner original values in data rows should remain exactly as in source for the group.
            if not all(r[owner_idx] == owner for r in data_rows):
                errors.append(f'Owner values not preserved in {sheet_name!r}')

# Compare all nonblank source rows against output grouped rows to ensure no loss or modification.
src_wb2 = load_workbook(input_path, data_only=False, read_only=True)
src_rows = []
for r in src_wb2['tasks'].iter_rows(min_row=2, values_only=True):
    if not all(v is None for v in r):
        src_rows.append(tuple(r))
src_wb2.close()
out_rows = []
for name in check_wb.sheetnames:
    if name == 'Summary':
        continue
    out_rows.extend(tuple(r) for r in check_wb[name].iter_rows(min_row=2, values_only=True))
if Counter(src_rows) != Counter(out_rows):
    errors.append('Output task rows do not exactly match source task rows')

check_wb.close()
if errors:
    print('CHECK FAILED')
    for error in errors:
        print('-', error)
    raise SystemExit(1)
print('CHECK PASSED')
print('Verified output file:', output_path.name)
print('Verified source headers:', headers)
print('Verified Summary headers:', summary_headers)
print('Verified worksheet names:', load_workbook(output_path, read_only=True).sheetnames)
print('Verified source/output task row count:', len(src_rows))