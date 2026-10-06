from pathlib import Path
from urllib.parse import urljoin
import urllib.request, urllib.error
from html.parser import HTMLParser
import re, json, time
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:49758/t13/list'
XLSX_PATH = OUTPUT / 'robust_products.xlsx'
WARN_PATH = OUTPUT / 'warnings.json'

print('workspace:', WORKSPACE)
print('input files:')
for p in sorted((WORKSPACE / 'input').glob('*')):
    print(f' - {p.name} size={p.stat().st_size}')

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_a = False
        self.current_href = None
        self.current_text = []
        self.items = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs = dict(attrs)
            self.in_a = True
            self.current_href = attrs.get('href')
            self.current_text = []
    def handle_data(self, data):
        if self.in_a:
            self.current_text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self.in_a:
            name = ' '.join(''.join(self.current_text).split())
            if self.current_href and name:
                self.items.append({'Name': name, 'URL': urljoin(LIST_URL, self.current_href)})
            self.in_a = False
            self.current_href = None
            self.current_text = []

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, data):
        data = ' '.join(data.split())
        if data:
            self.parts.append(data)

def fetch_once(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Agent/1.0'})
    with urllib.request.urlopen(req, timeout=10) as response:
        return response.status, response.read().decode('utf-8', errors='replace')

def fetch_detail_with_retry(url, max_500_attempts=3):
    attempts = 0
    while attempts < max_500_attempts:
        attempts += 1
        try:
            status, body = fetch_once(url)
            return {'ok': True, 'status_code': status, 'body': body, 'attempts': attempts, 'error': None}
        except urllib.error.HTTPError as e:
            if e.code == 500 and attempts < max_500_attempts:
                time.sleep(0.1)
                continue
            return {'ok': False, 'status_code': e.code, 'body': '', 'attempts': attempts, 'error': f'HTTP {e.code}'}
        except Exception as e:
            return {'ok': False, 'status_code': None, 'body': '', 'attempts': attempts, 'error': repr(e)}
    return {'ok': False, 'status_code': 500, 'body': '', 'attempts': attempts, 'error': 'HTTP 500'}

def parse_price_currency(raw):
    if not raw:
        return 'N/A', 'N/A'
    s = raw.strip()
    currency = 'N/A'
    # Order matters: NT$ contains "$" but means TWD.
    if re.search(r'\b(?:TWD|NTD)\b|NT\$', s, re.I):
        currency = 'TWD'
    elif re.search(r'\bUSD\b|US\$', s, re.I):
        currency = 'USD'
    elif '$' in s:
        currency = 'USD'
    elif '€' in s or re.search(r'\bEUR\b', s, re.I):
        currency = 'EUR'
    elif '¥' in s or re.search(r'\bJPY\b', s, re.I):
        currency = 'JPY'
    else:
        mcur = re.search(r'\b([A-Z]{3})\b', s)
        if mcur:
            currency = mcur.group(1)
    mnum = re.search(r'([0-9][0-9,]*(?:\.\d+)?)', s)
    if not mnum:
        return 'N/A', currency
    num = mnum.group(1).replace(',', '')
    price = float(num) if '.' in num else int(num)
    return price, currency

def parse_detail(html, fallback_name, url):
    parser = TextParser()
    parser.feed(html)
    parts = parser.parts
    print('detail text parts for', fallback_name, ':', parts)
    rec = {'Name': fallback_name, 'Price': 'N/A', 'Currency': 'N/A', 'CPU': 'N/A', 'URL': url, 'FetchStatus': 'ok'}
    if parts:
        # First non-label text is the product title/name on these pages.
        if parts[0].lower() not in ('name', 'price', 'currency', 'cpu', 'processor'):
            rec['Name'] = parts[0]
    labels = {'name': 'Name', 'price': 'Price', 'currency': 'Currency', 'cpu': 'CPU', 'processor': 'CPU'}
    i = 0
    while i < len(parts):
        key = labels.get(parts[i].strip().lower())
        if key and i + 1 < len(parts):
            val = parts[i+1].strip()
            if key == 'Price':
                price, cur = parse_price_currency(val)
                rec['Price'] = price
                if cur != 'N/A':
                    rec['Currency'] = cur
            elif key == 'Currency':
                rec['Currency'] = val.upper() if val else 'N/A'
            elif key == 'CPU':
                rec['CPU'] = val if val else 'N/A'
            elif key == 'Name':
                rec['Name'] = val if val else fallback_name
            i += 2
        else:
            i += 1
    # Fallback regex if page is not label/value split.
    flat = ' '.join(parts)
    if rec['Price'] == 'N/A':
        m = re.search(r'Price\s*[:：]?\s*((?:NT|US)?[$€¥]?\s*[0-9][0-9,]*(?:\.\d+)?\s*[A-Z]{0,3})', flat, re.I)
        if m:
            rec['Price'], rec['Currency'] = parse_price_currency(m.group(1))
    if rec['CPU'] == 'N/A':
        m = re.search(r'(?:CPU|Processor)\s*[:：]?\s*(.+?)(?=\s+(?:Name|Price|Currency)\s*[:：]?|$)', flat, re.I)
        if m:
            cpu = m.group(1).strip()
            if cpu:
                rec['CPU'] = cpu
    return rec

# Read list
status, list_html = fetch_once(LIST_URL)
print('list status:', status)
print('list markup sample:', list_html[:500].replace('\n', ' '))
parser = ProductListParser()
parser.feed(list_html)
list_items = parser.items[:5]
if len(list_items) < 5:
    raise RuntimeError(f'Expected at least 5 products from list, found {len(list_items)}')
print('list items:', list_items)

rows = []
warnings = {'source_list_url': LIST_URL, 'products': []}
for item in list_items:
    result = fetch_detail_with_retry(item['URL'], max_500_attempts=3)
    warnings['products'].append({
        'Name': item['Name'],
        'URL': item['URL'],
        'FetchStatus': 'ok' if result['ok'] else 'error',
        'attempts': result['attempts'],
        'status_code': result['status_code'],
        'warning': None if result['ok'] else (result['error'] or 'detail fetch failed'),
    })
    if result['ok']:
        rows.append(parse_detail(result['body'], item['Name'], item['URL']))
    else:
        rows.append({'Name': item['Name'], 'Price': 'N/A', 'Currency': 'N/A', 'CPU': 'N/A', 'URL': item['URL'], 'FetchStatus': 'error'})

headers = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']
wb = Workbook()
ws = wb.active
ws.title = 'Products'
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(XLSX_PATH)
with WARN_PATH.open('w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Reopen/check saved outputs only: file names, headers, values, placeholders, statuses, row counts.
print('checking outputs from saved files...')
assert XLSX_PATH.exists(), 'robust_products.xlsx missing'
assert WARN_PATH.exists(), 'warnings.json missing'
check_wb = load_workbook(XLSX_PATH, data_only=True)
check_ws = check_wb.active
saved_headers = [check_ws.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
assert saved_headers == headers, f'headers mismatch: {saved_headers}'
saved_rows = list(check_ws.iter_rows(min_row=2, max_col=len(headers), values_only=True))
assert len(saved_rows) == 5, f'expected 5 data rows, got {len(saved_rows)}'
for idx, saved in enumerate(saved_rows, start=1):
    rec = dict(zip(headers, saved))
    assert rec['Name'], f'row {idx} Name blank'
    assert rec['URL'] and str(rec['URL']).startswith('http://127.0.0.1:49758/'), f'row {idx} URL invalid'
    assert rec['FetchStatus'] in ('ok', 'error'), f'row {idx} FetchStatus invalid'
    if rec['FetchStatus'] == 'error':
        assert rec['Price'] == 'N/A' and rec['Currency'] == 'N/A' and rec['CPU'] == 'N/A', f'row {idx} error placeholders invalid'
    else:
        assert rec['Price'] is not None, f'row {idx} Price missing'
        assert rec['Currency'] is not None, f'row {idx} Currency missing'
        assert rec['CPU'] is not None, f'row {idx} CPU missing'
        assert isinstance(rec['Price'], (int, float, str)), f'row {idx} Price bad type'
check_warnings = json.loads(WARN_PATH.read_text(encoding='utf-8'))
assert 'products' in check_warnings and len(check_warnings['products']) == 5, 'warnings product count invalid'
for idx, w in enumerate(check_warnings['products'], start=1):
    assert w['FetchStatus'] in ('ok', 'error'), f'warning {idx} status invalid'
    if w['FetchStatus'] == 'ok':
        assert w['warning'] is None, f'warning {idx} ok should have no warning'
    if w['status_code'] == 500:
        assert w['attempts'] <= 3, f'warning {idx} exceeded retry limit'

print('wrote:', XLSX_PATH)
print('wrote:', WARN_PATH)
print('checked: robust_products.xlsx exists, exact headers, 5 rows, URLs/statuses/placeholders/types; warnings.json exists, 5 product entries, ok/error statuses, HTTP 500 attempts <= 3')
print('final rows:')
for r in saved_rows:
    print(r)
print('warnings:', json.dumps(check_warnings, ensure_ascii=False))