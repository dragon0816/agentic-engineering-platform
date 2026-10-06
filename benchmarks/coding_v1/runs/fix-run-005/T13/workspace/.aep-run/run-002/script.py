from pathlib import Path
from urllib.parse import urljoin
import urllib.request, urllib.error
from html.parser import HTMLParser
import re, json, time
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:49758/t13/list'
XLSX_PATH = OUTPUT / 'robust_products.xlsx'
WARN_PATH = OUTPUT / 'warnings.json'

print('workspace:', WORKSPACE)
print('input files:')
for p in sorted((WORKSPACE / 'input').glob('*')):
    print(f' - {p.name} size={p.stat().st_size}')

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_a = False
        self.current_href = None
        self.current_text = []
        self.items = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs = dict(attrs)
            self.in_a = True
            self.current_href = attrs.get('href')
            self.current_text = []
    def handle_data(self, data):
        if self.in_a:
            self.current_text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self.in_a:
            name = ' '.join(''.join(self.current_text).split())
            if self.current_href and name:
                self.items.append({'Name': name, 'URL': urljoin(LIST_URL, self.current_href)})
            self.in_a = False
            self.current_href = None
            self.current_text = []

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []
    def handle_data(self, data):
        data = data.strip()
        if data:
            self.parts.append(data)
    def text(self):
        return '\n'.join(self.parts)

def fetch_once(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Agent/1.0'})
    with urllib.request.urlopen(req, timeout=10) as response:
        body = response.read().decode('utf-8', errors='replace')
        return response.status, body

def fetch_detail_with_retry(url, max_500_attempts=3):
    attempts = 0
    last_error = None
    while attempts < max_500_attempts:
        attempts += 1
        try:
            status, body = fetch_once(url)
            return {'ok': True, 'status_code': status, 'body': body, 'attempts': attempts, 'error': None}
        except urllib.error.HTTPError as e:
            last_error = f'HTTP {e.code}'
            # Per requirement, retry HTTP 500 up to three attempts total.
            if e.code == 500 and attempts < max_500_attempts:
                time.sleep(0.1)
                continue
            return {'ok': False, 'status_code': e.code, 'body': '', 'attempts': attempts, 'error': last_error}
        except Exception as e:
            last_error = repr(e)
            return {'ok': False, 'status_code': None, 'body': '', 'attempts': attempts, 'error': last_error}
    return {'ok': False, 'status_code': 500, 'body': '', 'attempts': attempts, 'error': last_error or 'HTTP 500'}

def parse_detail(html, fallback_name, url):
    parser = TextParser()
    parser.feed(html)
    text = parser.text()
    print('detail text sample for', fallback_name, ':', text[:300].replace('\n', ' | '))
    flat = ' '.join(text.split())

    name = fallback_name
    # Prefer obvious Name field/title if present, otherwise list name.
    m = re.search(r'\bName\s*[:：]\s*([^\n|]+?)(?=\s+(?:Price|Currency|CPU|Processor)\s*[:：]|$)', flat, flags=re.I)
    if m:
        name = m.group(1).strip() or fallback_name

    price = 'N/A'
    currency = 'N/A'
    cpu = 'N/A'

    # Field-form price/currency, e.g. Price: 699 Currency: USD
    m = re.search(r'\bPrice\s*[:：]\s*([^\s<|]+)', flat, flags=re.I)
    if m:
        price = m.group(1).strip()
    # Inline money format, if no explicit field.
    if price == 'N/A':
        m = re.search(r'\b(USD|EUR|TWD|NTD|JPY|CNY|GBP)\s*([0-9][0-9,]*(?:\.\d+)?)\b', flat, flags=re.I)
        if m:
            currency = m.group(1).upper()
            price = m.group(2).replace(',', '')
        else:
            m = re.search(r'([$€¥])\s*([0-9][0-9,]*(?:\.\d+)?)', flat)
            if m:
                symbol_map = {'$': 'USD', '€': 'EUR', '¥': 'JPY'}
                currency = symbol_map.get(m.group(1), 'N/A')
                price = m.group(2).replace(',', '')

    m = re.search(r'\bCurrency\s*[:：]\s*([A-Z]{3}|NTD|TWD|USD|EUR|JPY|CNY|GBP)\b', flat, flags=re.I)
    if m:
        currency = m.group(1).upper()
    # If explicit price included currency after amount, pick it up.
    if currency == 'N/A':
        m = re.search(r'\bPrice\s*[:：]\s*(?:[0-9][0-9,]*(?:\.\d+)?)\s*([A-Z]{3}|NTD|TWD|USD|EUR|JPY|CNY|GBP)\b', flat, flags=re.I)
        if m:
            currency = m.group(1).upper()

    m = re.search(r'\b(?:CPU|Processor)\s*[:：]\s*(.+?)(?=\s+(?:Name|Price|Currency|URL|FetchStatus)\s*[:：]|$)', flat, flags=re.I)
    if m:
        val = m.group(1).strip()
        if val:
            cpu = val

    # Normalize price: keep numeric-looking price as number where possible; otherwise text.
    if isinstance(price, str) and price != 'N/A':
        cleaned = price.replace(',', '').strip()
        if re.fullmatch(r'[0-9]+(?:\.\d+)?', cleaned):
            price = float(cleaned) if '.' in cleaned else int(cleaned)
        else:
            price = cleaned
    return {'Name': name, 'Price': price, 'Currency': currency, 'CPU': cpu, 'URL': url, 'FetchStatus': 'ok'}

# Read list
status, list_html = fetch_once(LIST_URL)
print('list status:', status)
print('list markup sample:', list_html[:500].replace('\n', ' '))
parser = ProductListParser()
parser.feed(list_html)
list_items = parser.items[:5]
if len(list_items) < 5:
    raise RuntimeError(f'Expected at least 5 products from list, found {len(list_items)}')
print('list items:', list_items)

rows = []
warnings = {'source_list_url': LIST_URL, 'products': []}
for item in list_items:
    result = fetch_detail_with_retry(item['URL'], max_500_attempts=3)
    warning_entry = {
        'Name': item['Name'],
        'URL': item['URL'],
        'FetchStatus': 'ok' if result['ok'] else 'error',
        'attempts': result['attempts'],
        'status_code': result['status_code'],
        'warning': None if result['ok'] else (result['error'] or 'detail fetch failed'),
    }
    warnings['products'].append(warning_entry)
    if result['ok']:
        rows.append(parse_detail(result['body'], item['Name'], item['URL']))
    else:
        rows.append({'Name': item['Name'], 'Price': 'N/A', 'Currency': 'N/A', 'CPU': 'N/A', 'URL': item['URL'], 'FetchStatus': 'error'})

headers = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']
wb = Workbook()
ws = wb.active
ws.title = 'Products'
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(XLSX_PATH)
with WARN_PATH.open('w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Reopen/check saved outputs only: file names, headers, values, placeholders, statuses, row counts.
print('checking outputs from saved files...')
assert XLSX_PATH.exists(), 'robust_products.xlsx missing'
assert WARN_PATH.exists(), 'warnings.json missing'
check_wb = load_workbook(XLSX_PATH, data_only=True)
check_ws = check_wb.active
saved_headers = [check_ws.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
assert saved_headers == headers, f'headers mismatch: {saved_headers}'
saved_rows = list(check_ws.iter_rows(min_row=2, max_col=len(headers), values_only=True))
assert len(saved_rows) == 5, f'expected 5 data rows, got {len(saved_rows)}'
for idx, saved in enumerate(saved_rows, start=1):
    rec = dict(zip(headers, saved))
    assert rec['Name'], f'row {idx} Name blank'
    assert rec['URL'] and str(rec['URL']).startswith('http://127.0.0.1:49758/'), f'row {idx} URL invalid'
    assert rec['FetchStatus'] in ('ok', 'error'), f'row {idx} FetchStatus invalid'
    if rec['FetchStatus'] == 'error':
        assert rec['Price'] == 'N/A' and rec['Currency'] == 'N/A' and rec['CPU'] == 'N/A', f'row {idx} error placeholders invalid'
    else:
        assert rec['Price'] is not None, f'row {idx} Price missing'
        assert rec['Currency'] is not None, f'row {idx} Currency missing'
        assert rec['CPU'] is not None, f'row {idx} CPU missing'
        assert isinstance(rec['Price'], (int, float, str)), f'row {idx} Price bad type'
check_warnings = json.loads(WARN_PATH.read_text(encoding='utf-8'))
assert 'products' in check_warnings and len(check_warnings['products']) == 5, 'warnings product count invalid'
for idx, w in enumerate(check_warnings['products'], start=1):
    assert w['FetchStatus'] in ('ok', 'error'), f'warning {idx} status invalid'
    if w['FetchStatus'] == 'ok':
        assert w['warning'] is None, f'warning {idx} ok should have no warning'
    if w['status_code'] == 500:
        assert w['attempts'] <= 3, f'warning {idx} exceeded retry limit'

print('wrote:', XLSX_PATH)
print('wrote:', WARN_PATH)
print('checked: robust_products.xlsx exists, exact headers, 5 rows, URLs/statuses/placeholders/types; warnings.json exists, 5 product entries, ok/error statuses, HTTP 500 attempts <= 3')
print('final rows:')
for r in saved_rows:
    print(r)
print('warnings:', json.dumps(check_warnings, ensure_ascii=False))