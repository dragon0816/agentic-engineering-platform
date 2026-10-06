from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('Workspace:', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name)
    if p.suffix.lower() in ['.xlsx', '.xlsm']:
        wb = load_workbook(p, read_only=True, data_only=False)
        print('  sheets:', wb.sheetnames)
        for s in wb.sheetnames:
            sh = wb[s]
            rows = list(sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True))
            print(f'  Sheet {s}: {sh.max_row} rows x {sh.max_column} cols')
            for r in rows:
                print('   ', r)
        wb.close()