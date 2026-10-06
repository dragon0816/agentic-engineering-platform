from pathlib import Path
from collections import OrderedDict
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
input_path = workspace / 'input' / 'large_results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

# Read and aggregate without modifying input files.
stats = OrderedDict()
overall = {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0}

with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    required = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']
    missing = [c for c in required if c not in (reader.fieldnames or [])]
    if missing:
        raise ValueError(f'Missing required columns: {missing}')
    for row_num, row in enumerate(reader, start=2):
        model = row['Device Model']
        status = row['Status'].strip()
        try:
            duration = float(row['DurationSeconds'])
        except Exception as exc:
            raise ValueError(f'Invalid DurationSeconds at CSV row {row_num}: {row["DurationSeconds"]!r}') from exc
        if model not in stats:
            stats[model] = {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0}
        stats[model]['total'] += 1
        stats[model]['duration_sum'] += duration
        overall['total'] += 1
        overall['duration_sum'] += duration
        if status == 'Pass':
            stats[model]['pass'] += 1
            overall['pass'] += 1
        elif status == 'Fail':
            stats[model]['fail'] += 1
            overall['fail'] += 1
        else:
            # Unknown statuses count toward Total and average duration but not Pass/Fail.
            pass

# Create workbook.
wb = Workbook()
ws = wb.active
ws.title = 'Summary'
headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
ws.append(headers)
for cell in ws[1]:
    cell.font = Font(bold=True)

for model, s in stats.items():
    total = s['total']
    pass_count = s['pass']
    fail_count = s['fail']
    pass_rate = pass_count / total if total else 0
    avg_duration = s['duration_sum'] / total if total else 0
    ws.append([model, total, pass_count, fail_count, pass_rate, avg_duration])

for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
    row[0].number_format = '0.00%'
for row in ws.iter_rows(min_row=2, min_col=6, max_col=6):
    row[0].number_format = '0.00'
for col in range(1, len(headers) + 1):
    ws.column_dimensions[chr(64 + col)].width = max(14, len(headers[col-1]) + 2)

ow = wb.create_sheet('Overall')
overall_headers = ['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
ow.append(overall_headers)
for cell in ow[1]:
    cell.font = Font(bold=True)
overall_total = overall['total']
ow.append([
    overall_total,
    overall['pass'],
    overall['fail'],
    overall['pass'] / overall_total if overall_total else 0,
    overall['duration_sum'] / overall_total if overall_total else 0,
])
ow['D2'].number_format = '0.00%'
ow['E2'].number_format = '0.00'
for col in range(1, len(overall_headers) + 1):
    ow.column_dimensions[chr(64 + col)].width = max(14, len(overall_headers[col-1]) + 2)

wb.save(output_path)
print(f'Wrote {output_path}')
print(f'Input rows: {overall_total}; Device models: {len(stats)}')

# Verification: reopen CSV and output workbook; check required sheets, columns, counts, types and formulas-by-values.
expected = OrderedDict()
expected_overall = {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0}
with input_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    for row in reader:
        model = row['Device Model']
        status = row['Status'].strip()
        duration = float(row['DurationSeconds'])
        expected.setdefault(model, {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0})
        expected[model]['total'] += 1
        expected[model]['duration_sum'] += duration
        expected_overall['total'] += 1
        expected_overall['duration_sum'] += duration
        if status == 'Pass':
            expected[model]['pass'] += 1
            expected_overall['pass'] += 1
        elif status == 'Fail':
            expected[model]['fail'] += 1
            expected_overall['fail'] += 1

wb2 = load_workbook(output_path, data_only=True)
assert 'Summary' in wb2.sheetnames, 'Missing Summary sheet'
assert 'Overall' in wb2.sheetnames, 'Missing Overall sheet'
ws2 = wb2['Summary']
ow2 = wb2['Overall']
assert [c.value for c in ws2[1]] == headers, f'Bad Summary headers: {[c.value for c in ws2[1]]}'
assert [c.value for c in ow2[1]] == overall_headers, f'Bad Overall headers: {[c.value for c in ow2[1]]}'
assert ws2.max_row == len(expected) + 1, f'Summary row count {ws2.max_row} != {len(expected)+1}'
assert ow2.max_row == 2, f'Overall row count {ow2.max_row} != 2'

for r in range(2, ws2.max_row + 1):
    model = ws2.cell(r, 1).value
    if model not in expected:
        raise AssertionError(f'Unexpected model in Summary: {model!r}')
    e = expected[model]
    values = [ws2.cell(r, c).value for c in range(2, 7)]
    exp_values = [
        e['total'],
        e['pass'],
        e['fail'],
        e['pass'] / e['total'],
        e['duration_sum'] / e['total'],
    ]
    assert isinstance(model, str), f'Device Model not text for row {r}'
    for idx, (got, exp) in enumerate(zip(values, exp_values), start=2):
        assert isinstance(got, (int, float)), f'Non-numeric value at Summary row {r}, col {idx}: {got!r}'
        assert abs(got - exp) < 1e-9, f'Mismatch row {r}, col {idx}: got {got}, expected {exp}'

ovals = [ow2.cell(2, c).value for c in range(1, 6)]
oexp = [
    expected_overall['total'],
    expected_overall['pass'],
    expected_overall['fail'],
    expected_overall['pass'] / expected_overall['total'],
    expected_overall['duration_sum'] / expected_overall['total'],
]
for idx, (got, exp) in enumerate(zip(ovals, oexp), start=1):
    assert isinstance(got, (int, float)), f'Non-numeric Overall col {idx}: {got!r}'
    assert abs(got - exp) < 1e-9, f'Overall mismatch col {idx}: got {got}, expected {exp}'

print('Verification passed:')
print('- summary.xlsx exists in output folder')
print('- sheets: Summary, Overall')
print('- Summary headers:', headers)
print('- Overall headers:', overall_headers)
print('- Summary has one row per Device Model:', len(expected))
print('- Overall Total equals CSV data rows:', expected_overall['total'])
print('- Total/PassCount/FailCount/PassRate/AvgDurationSeconds values are numeric and match the CSV')
