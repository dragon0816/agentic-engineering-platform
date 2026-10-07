from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
print('workspace:', workspace)
print('input exists:', input_path.exists(), input_path)
if input_path.exists():
    wb = load_workbook(input_path, read_only=True, data_only=False)
    print('sheetnames:', wb.sheetnames)
    if 'tasks' in wb.sheetnames:
        ws = wb['tasks']
        rows = list(ws.iter_rows(values_only=True))
        print('max_row:', ws.max_row, 'max_column:', ws.max_column)
        if rows:
            print('headers:', rows[0])
            print('sample rows:')
            for row in rows[1:6]:
                print(row)
    wb.close()
