from pathlib import Path
from collections import OrderedDict, Counter
import re
from openpyxl import load_workbook, Workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'owners.xlsx'

# Load input without modifying it
in_wb = load_workbook(input_path, data_only=False)
if 'tasks' not in in_wb.sheetnames:
    raise ValueError('Input workbook does not contain a tasks worksheet')
ws = in_wb['tasks']
rows = list(ws.iter_rows(values_only=True))
if not rows:
    raise ValueError('tasks worksheet is empty')
headers = list(rows[0])
if 'Owner' not in headers or 'Status' not in headers:
    raise ValueError(f'Required columns Owner and Status not found in headers: {headers}')
owner_idx = headers.index('Owner')
status_idx = headers.index('Status')

# Group rows preserving original Owner values in data rows; blank Owner grouped as Unassigned
# Use OrderedDict to preserve first appearance order.
groups = OrderedDict()
owner_display = OrderedDict()
for row in rows[1:]:
    # keep row width aligned to headers
    row = tuple(row[:len(headers)]) + tuple([None] * max(0, len(headers) - len(row)))
    owner_val = row[owner_idx]
    key = 'Unassigned' if owner_val is None or str(owner_val).strip() == '' else owner_val
    if key not in groups:
        groups[key] = []
        owner_display[key] = 'Unassigned' if key == 'Unassigned' else owner_val
    groups[key].append(row)

invalid_chars = r'[\\/*?:\[\]]'
def make_valid_sheet_name(owner_label, used):
    base = str(owner_label) if owner_label is not None else 'Unassigned'
    base = re.sub(invalid_chars, '_', base)
    base = base.strip()
    if not base:
        base = 'Unassigned'
    # Excel sheet names max length 31 and cannot be Summary reserved duplicate here.
    # Avoid duplicates case-insensitively, including Summary.
    def clip_for_suffix(name, suffix):
        return name[:31-len(suffix)] + suffix
    candidate = base[:31]
    if candidate.lower() == 'summary':
        candidate = clip_for_suffix(base[:28] or 'Sheet', '_1')
    i = 1
    while candidate.lower() in used:
        suffix = f'_{i}'
        truncated_base = base[:31-len(suffix)] if len(base) > 31-len(suffix) else base
        candidate = truncated_base + suffix
        if candidate.lower() == 'summary':
            i += 1
            continue
        i += 1
    used.add(candidate.lower())
    return candidate

used_sheet_names = {'summary'}
sheet_map = OrderedDict()
for key, display in owner_display.items():
    sheet_map[key] = make_valid_sheet_name(display, used_sheet_names)

# Create output workbook
out_wb = Workbook()
default_ws = out_wb.active
default_ws.title = 'Summary'
summary_ws = default_ws
summary_headers = ['Owner', 'SheetName', 'Total', 'Pass', 'Fail']
summary_ws.append(summary_headers)

# Summary plus one worksheet per group
for key, task_rows in groups.items():
    display_owner = owner_display[key]
    sheet_name = sheet_map[key]
    task_ws = out_wb.create_sheet(title=sheet_name)
    task_ws.append(headers)
    for row in task_rows:
        task_ws.append(list(row))
    counts = Counter((r[status_idx] for r in task_rows))
    summary_ws.append([display_owner, sheet_name, len(task_rows), counts.get('Pass', 0), counts.get('Fail', 0)])

out_wb.save(output_path)
in_wb.close()
out_wb.close()
print('created:', output_path)
print('groups:', [(owner_display[k], sheet_map[k], len(v)) for k, v in groups.items()])

# Verification: reopen output and check requirements
verify_wb = load_workbook(output_path, data_only=True)
print('output sheets:', verify_wb.sheetnames)
assert 'Summary' in verify_wb.sheetnames, 'Missing Summary worksheet'
summary = verify_wb['Summary']
actual_summary_headers = [cell.value for cell in summary[1]]
assert actual_summary_headers == summary_headers, f'Summary headers mismatch: {actual_summary_headers}'

# Check every summary entry and sheet contents
summary_rows = list(summary.iter_rows(min_row=2, values_only=True))
assert len(summary_rows) == len(groups), f'Summary row count {len(summary_rows)} != group count {len(groups)}'
for (key, task_rows), srow in zip(groups.items(), summary_rows):
    display_owner = owner_display[key]
    expected_sheet = sheet_map[key]
    owner, sheet_name, total, pass_count, fail_count = srow
    assert owner == display_owner, f'Owner mismatch for {key}: {owner} != {display_owner}'
    assert sheet_name == expected_sheet, f'SheetName mismatch for {key}: {sheet_name} != {expected_sheet}'
    assert sheet_name in verify_wb.sheetnames, f'Missing owner sheet {sheet_name}'
    assert total == len(task_rows), f'Total mismatch for {key}'
    counts = Counter((r[status_idx] for r in task_rows))
    assert pass_count == counts.get('Pass', 0), f'Pass count mismatch for {key}'
    assert fail_count == counts.get('Fail', 0), f'Fail count mismatch for {key}'
    owner_ws = verify_wb[sheet_name]
    sheet_headers = [cell.value for cell in owner_ws[1]]
    assert sheet_headers == headers, f'Headers not preserved in {sheet_name}: {sheet_headers}'
    data_rows = list(owner_ws.iter_rows(min_row=2nub_values_only=True)) if False else list(owner_ws.iter_rows(min_row=2, values_only=True))
    assert len(data_rows) == len(task_rows), f'Row count mismatch in {sheet_name}'
    # Compare values, preserving blank Owner as blank in row but group name Unassigned in summary.
    for actual, expected in zip(data_rows, task_rows):
        actual = tuple(actual[:len(headers)])
        expected = tuple(expected[:len(headers)])
        assert actual == expected, f'Data row mismatch in {sheet_name}: {actual} != {expected}'

# Check no illegal worksheet characters and unique sheet names by Excel/openpyxl constraints
illegal = set('\\/*?:[]')
for name in verify_wb.sheetnames:
    assert len(name) <= 31, f'Sheet name too long: {name}'
    assert not any(ch in illegal for ch in name), f'Illegal char in sheet name: {name}'
assert len({name.lower() for name in verify_wb.sheetnames}) == len(verify_wb.sheetnames), 'Duplicate sheet names case-insensitively'
verify_wb.close()
print('verification passed: Summary headers/rows/counts, SheetName mapping, owner sheet names, preserved task headers/data, row counts, and legal unique sheet names checked.')
