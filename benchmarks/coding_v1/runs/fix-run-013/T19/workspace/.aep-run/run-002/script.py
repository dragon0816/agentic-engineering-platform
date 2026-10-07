import csv
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read CSV robustly: utf-8-sig handles UTF-8 BOM; csv handles quoted commas/newlines.
with open(input_path, 'r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    rows = [row for row in reader]

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
device_col_index = headers.index('Device ID') + 1 if 'Device ID' in headers else None

wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'

for r_idx, row in enumerate(rows, start=1):
    # Preserve all columns in the header order. If a data row is short, keep missing trailing cells blank.
    for c_idx in range(1, len(headers) + 1):
        value = row[c_idx - 1] if c_idx - 1 < len(row) else ''
        cell = ws.cell(row=r_idx, column=c_idx)
        if value == '':
            cell.value = None  # blank cell in Excel
        else:
            cell.value = str(value)
        if device_col_index is not None and c_idx == device_col_index:
            cell.number_format = '@'
            if value != '':
                cell.data_type = 's'

wb.save(output_path)
print('Wrote', output_path)
print('source_rows_including_header=', len(rows))
print('source_headers=', headers)
print('device_col_index=', device_col_index)

# Validation by reopening output file and comparing to source semantics.
wb2 = load_workbook(output_path, data_only=False)
ws2 = wb2.active
xlsx_rows = []
for r_idx in range(1, ws2.max_row + 1):
    xlsx_rows.append([ws2.cell(r_idx, c_idx).value for c_idx in range(1, len(headers) + 1)])

# Normalize xlsx blank cells to empty strings for comparison to CSV values.
normalized_xlsx = [[('' if v is None else str(v)) for v in row] for row in xlsx_rows]
normalized_source = []
for row in rows:
    normalized_source.append([(row[i] if i < len(row) else '') for i in range(len(headers))])

checks = []
checks.append(('output_file_exists', output_path.exists()))
checks.append(('headers_preserved', normalized_xlsx[0] == headers))
checks.append(('row_count_including_header_preserved', len(normalized_xlsx) == len(normalized_source)))
checks.append(('all_values_preserved_in_order', normalized_xlsx == normalized_source))
checks.append(('no_extra_index_column', ws2.max_column == len(headers)))

if device_col_index is not None:
    device_checks = []
    for r_idx in range(1, ws2.max_row + 1):
        cell = ws2.cell(r_idx, device_col_index)
        expected = normalized_source[r_idx - 1][device_col_index - 1]
        ok_value = ('' if cell.value is None else str(cell.value)) == expected
        ok_type = (cell.value is None) or (cell.data_type == 's')
        ok_format = cell.number_format == '@'
        device_checks.append(ok_value and ok_type and ok_format)
    checks.append(('device_id_text_preserves_leading_zero_and_text_format', all(device_checks)))

# Explicit duplicate preservation: compare full row occurrence counts.
from collections import Counter
checks.append(('duplicate_rows_preserved', Counter(map(tuple, normalized_xlsx[1:])) == Counter(map(tuple, normalized_source[1:]))))

for name, ok in checks:
    print(f'{name}: {ok}')

if not all(ok for _, ok in checks):
    raise AssertionError('One or more validation checks failed')

print('validated_output=', str(output_path))
print('xlsx_max_row=', ws2.max_row, 'xlsx_max_column=', ws2.max_column)
print('first_data_rows=', normalized_xlsx[:6])