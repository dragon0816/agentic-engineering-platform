from pathlib import Path
from openpyxl import load_workbook, Workbook

base = Path.cwd()
input_dir = base / 'input'
output_dir = base / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

devices_path = input_dir / 'devices.xlsx'
tests_path = input_dir / 'tests.xlsx'
out_path = output_dir / 'merged_device_tests.xlsx'

def load_rows(path):
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    headers = [str(c.value).strip() if c.value is not None else '' for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None for v in row):
            continue
        rec = {}
        for h, v in zip(headers, row):
            if h == 'Device ID':
                rec[h] = '' if v is None else str(v)
            else:
                rec[h] = v
        rows.append(rec)
    wb.close()
    return headers, rows

dev_headers, dev_rows = load_rows(devices_path)
test_headers, test_rows = load_rows(tests_path)

if 'Device ID' not in dev_headers:
    raise ValueError('devices.xlsx missing Device ID column')
if 'Device ID' not in test_headers:
    raise ValueError('tests.xlsx missing Device ID column')

# Ensure devices IDs are unique as required by prompt; fail loudly if input violates it.
seen = {}
for idx, rec in enumerate(dev_rows, start=2):
    did = rec.get('Device ID', '')
    if did in seen:
        raise ValueError(f'Duplicate Device ID in devices.xlsx: {did!r} at rows {seen[did]} and {idx}')
    seen[did] = idx

device_by_id = {rec['Device ID']: rec for rec in dev_rows}
tests_by_id = {}
for rec in test_rows:
    tests_by_id.setdefault(rec['Device ID'], []).append(rec)

# Matched columns: tests columns first, then non-ID device columns to avoid duplicate Device ID.
device_extra_headers = [h for h in dev_headers if h != 'Device ID']
matched_headers = list(test_headers) + device_extra_headers

unmatched_headers = ['Source'] + list(dev_headers) + [h for h in test_headers if h not in dev_headers]

matched_rows = []
unmatched_rows = []

for test in test_rows:
    did = test['Device ID']
    dev = device_by_id.get(did)
    if dev is not None:
        out = {h: test.get(h) for h in test_headers}
        for h in device_extra_headers:
            out[h] = dev.get(h)
        matched_rows.append(out)
    else:
        out = {h: None for h in unmatched_headers}
        out['Source'] = 'tests'
        for h in test_headers:
            out[h] = test.get(h)
        unmatched_rows.append(out)

for dev in dev_rows:
    did = dev['Device ID']
    if did not in tests_by_id:
        out = {h: None for h in unmatched_headers}
        out['Source'] = 'devices'
        for h in dev_headers:
            out[h] = dev.get(h)
        unmatched_rows.append(out)

wb_out = Workbook()
ws_m = wb_out.active
ws_m.title = 'matched'
ws_u = wb_out.create_sheet('unmatched')

def write_sheet(ws, headers, rows):
    ws.append(headers)
    for rec in rows:
        ws.append([rec.get(h) for h in headers])
    # Keep Device ID as text and preserve leading zeroes.
    if 'Device ID' in headers:
        col_idx = headers.index('Device ID') + 1
        for row in range(2, ws.max_row + 1):
            cell = ws.cell(row=row, column=col_idx)
            if cell.value is not None:
                cell.value = str(cell.value)
            cell.number_format = '@'
    for cell in ws[1]:
        cell.number_format = '@'

write_sheet(ws_m, matched_headers, matched_rows)
write_sheet(ws_u, unmatched_headers, unmatched_rows)
wb_out.save(out_path)

# Verification: reopen the saved output and confirm requirements.
check_wb = load_workbook(out_path, data_only=False)
assert set(check_wb.sheetnames) >= {'matched', 'unmatched'}
cm = check_wb['matched']
cu = check_wb['unmatched']
matched_headers_saved = [c.value for c in cm[1]]
unmatched_headers_saved = [c.value for c in cu[1]]
assert matched_headers_saved == matched_headers, (matched_headers_saved, matched_headers)
assert unmatched_headers_saved == unmatched_headers, (unmatched_headers_saved, unmatched_headers)
assert 'Source' in unmatched_headers_saved
assert cm.max_row - 1 == len(matched_rows)
assert cu.max_row - 1 == len(unmatched_rows)
# Device IDs remain strings; known leading zero samples should stay intact if present.
for ws in (cm, cu):
    headers = [c.value for c in ws[1]]
    if 'Device ID' in headers:
        idx = headers.index('Device ID') + 1
        for r in range(2, ws.max_row + 1):
            val = ws.cell(r, idx).value
            assert val is None or isinstance(val, str), (ws.title, r, val, type(val))

# Cross-check matched count equals tests whose ID exists in devices, unmatched includes all nonmatched tests and devices without tests.
expected_matched = sum(1 for t in test_rows if t['Device ID'] in device_by_id)
expected_unmatched_tests = sum(1 for t in test_rows if t['Device ID'] not in device_by_id)
expected_unmatched_devices = sum(1 for d in dev_rows if d['Device ID'] not in tests_by_id)
assert len(matched_rows) == expected_matched
assert len(unmatched_rows) == expected_unmatched_tests + expected_unmatched_devices
source_idx = unmatched_headers_saved.index('Source') + 1
sources = [cu.cell(r, source_idx).value for r in range(2, cu.max_row + 1)]
assert all(s in ('devices', 'tests') for s in sources)
check_wb.close()

print('Created:', out_path)
print('devices rows:', len(dev_rows))
print('tests rows:', len(test_rows))
print('matched rows:', len(matched_rows))
print('unmatched rows:', len(unmatched_rows))
print('matched headers:', matched_headers)
print('unmatched headers:', unmatched_headers)
print('unmatched sources:', {s: sources.count(s) for s in sorted(set(sources))})
print('Verification passed: output reopened; sheet names, headers, row counts, Source values, and text Device ID cells checked.')
