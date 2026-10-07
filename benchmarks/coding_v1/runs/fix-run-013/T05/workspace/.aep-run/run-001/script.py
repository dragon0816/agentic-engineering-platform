from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
input_dir = base / 'input'
print('Workspace:', base)
print('Input files:')
for p in sorted(input_dir.iterdir()):
    print(' -', p.name, p.stat().st_size, 'bytes')
    if p.suffix.lower() in ('.xlsx', '.xlsm'):
        wb = load_workbook(p, read_only=True, data_only=False)
        print('   sheets:', wb.sheetnames)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=True))
            print(f'   sheet {ws.title}: max_row={ws.max_row}, max_col={ws.max_column}')
            if rows:
                print('   headers:', rows[0])
                for r in rows[1:]:
                    print('   sample:', r)
            break
        wb.close()
