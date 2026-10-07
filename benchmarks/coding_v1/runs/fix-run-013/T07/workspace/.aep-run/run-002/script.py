from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'specs.md'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')

# Split product sections headed by markdown level-2 headings.
sections = []
for match in re.finditer(r'(?ms)^##\s+(.+?)\s*$\n(.*?)(?=^##\s+|\Z)', text):
    heading = match.group(1).strip()
    body = match.group(2)
    sections.append((heading, body))

headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'Bandwidth', 'Interface']
rows = []

def convert_to_mhz(value: str, unit: str):
    num = float(value)
    unit_l = unit.lower()
    if unit_l == 'ghz':
        return num * 1000
    if unit_l == 'mhz':
        return num
    return None

def fmt_num(n):
    if n is None:
        return ''
    if abs(n - round(n)) < 1e-9:
        return int(round(n))
    return n

for heading, body in sections:
    fields = {}
    for line in body.splitlines():
        m = re.match(r'^\s*-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            fields[m.group(1).strip().lower()] = m.group(2).strip()

    model = fields.get('model', '')

    fmin = fmax = ''
    freq = fields.get('frequency', '')
    if freq:
        # Supports ranges like 2.4–2.5 GHz, 900-930 MHz. Unit must be written.
        fm = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*[–—-]\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\b', freq, flags=re.I)
        if fm:
            fmin = fmt_num(convert_to_mhz(fm.group(1), fm.group(3)))
            fmax = fmt_num(convert_to_mhz(fm.group(2), fm.group(3)))

    bandwidth = ''
    bw = fields.get('bandwidth', '')
    if bw:
        bm = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\b', bw, flags=re.I)
        if bm:
            bandwidth = fmt_num(convert_to_mhz(bm.group(1), bm.group(2)))

    interface = fields.get('interface', '')
    if interface:
        parts = [p.strip() for p in re.split(r'[/,;、]+|\band\b', interface) if p.strip()]
        interface = ';'.join(parts)

    rows.append([model, fmin, fmax, bandwidth, interface])

wb = Workbook()
ws = wb.active
ws.title = 'specs'
ws.append(headers)
for row in rows:
    ws.append(row)
wb.save(output_path)

# Verification: reopen output and confirm stated requirements.
wb2 = load_workbook(output_path, data_only=True)
ws2 = wb2['specs']
actual_headers = [ws2.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
actual_rows = []
for r in range(2, ws2.max_row + 1):
    actual_rows.append([ws2.cell(row=r, column=c).value for c in range(1, len(headers)+1)])

print('created:', output_path)
print('sheet:', ws2.title)
print('headers:', actual_headers)
print('row_count:', len(actual_rows))
for row in actual_rows:
    print(row)

assert actual_headers == headers, f'headers mismatch: {actual_headers}'
assert len(actual_rows) == 3, f'expected 3 product rows, got {len(actual_rows)}'
# Check blanks stay blank/None in Excel for missing specs.
assert actual_rows[2][3] is None, 'P3 missing bandwidth should be blank'
# Check numeric MHz conversions and interface separator.
expected = [
    ['RX-100', 2400, 2500, 80, 'USB;LAN'],
    ['RX-200', 900, 930, 20, 'LAN'],
    ['RX-300', 5000, 6000, None, 'USB'],
]
assert actual_rows == expected, f'rows mismatch: {actual_rows}'
print('verification passed: specs.xlsx reopened; headers, 3 rows, MHz conversions, semicolon interfaces, and blank missing bandwidth checked.')
