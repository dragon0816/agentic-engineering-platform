from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'large_results.csv'
output_path = workspace / 'output' / 'summary.xlsx'

# Recompute expected values from source
stats = defaultdict(lambda: {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0})
overall = {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0}
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        model = row['Device Model']
        status = row['Status']
        duration = float(row['DurationSeconds'])
        stats[model]['total'] += 1
        overall['total'] += 1
        if status == 'Pass':
            stats[model]['pass'] += 1
            overall['pass'] += 1
        elif status == 'Fail':
            stats[model]['fail'] += 1
            overall['fail'] += 1
        stats[model]['duration_sum'] += duration
        overall['duration_sum'] += duration

wb = load_workbook(output_path, data_only=True)
expected_sheets = ['By Device Model', 'Overall']
assert wb.sheetnames == expected_sheets, f'Unexpected sheets: {wb.sheetnames}'

ws = wb['By Device Model']
expected_headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
actual_headers = [ws.cell(1, c).value for c in range(1, 7)]
assert actual_headers == expected_headers, f'Bad headers: {actual_headers}'
assert ws.max_row == len(stats) + 1, f'Bad row count: {ws.max_row}'

for r in range(2, ws.max_row + 1):
    model, total, pass_count, fail_count, pass_rate, avg_duration = [ws.cell(r, c).value for c in range(1, 7)]
    assert isinstance(model, str), f'Model not text at row {r}'
    s = stats[model]
    assert total == s['total'], (model, total, s['total'])
    assert pass_count == s['pass'], (model, pass_count, s['pass'])
    assert fail_count == s['fail'], (model, fail_count, s['fail'])
    assert isinstance(total, int), f'Total not int at row {r}'
    assert isinstance(pass_count, int), f'PassCount not int at row {r}'
    assert isinstance(fail_count, int), f'FailCount not int at row {r}'
    exp_rate = s['pass'] / s['total']
    exp_avg = s['duration_sum'] / s['total']
    assert abs(pass_rate - exp_rate) < 1e-12, (model, pass_rate, exp_rate)
    assert abs(avg_duration - exp_avg) < 1e-12, (model, avg_duration, exp_avg)
    assert isinstance(pass_rate, (int, float)), f'PassRate not numeric at row {r}'
    assert isinstance(avg_duration, (int, float)), f'AvgDurationSeconds not numeric at row {r}'

wso = wb['Overall']
overall_headers = [wso.cell(1, c).value for c in range(1, 6)]
assert overall_headers == ['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds'], overall_headers
assert wso.max_row == 2, f'Unexpected Overall row count: {wso.max_row}'
vals = [wso.cell(2, c).value for c in range(1, 6)]
exp_vals = [
    overall['total'],
    overall['pass'],
    overall['fail'],
    overall['pass'] / overall['total'],
    overall['duration_sum'] / overall['total'],
]
for actual, expected in zip(vals, exp_vals):
    if isinstance(expected, float):
        assert abs(actual - expected) < 1e-12, (actual, expected)
    else:
        assert actual == expected, (actual, expected)

print('verified output:', output_path)
print('sheets:', wb.sheetnames)
print('By Device Model rows including header:', ws.max_row)
print('Overall values:', vals)
print('source rows checked:', overall['total'])
