from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'large_results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'summary.xlsx'

required_cols = ['Test ID', 'Device Model', 'Status', 'DurationSeconds']
summary_headers = ['Device Model', 'Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']
overall_headers = ['Total', 'PassCount', 'FailCount', 'PassRate', 'AvgDurationSeconds']

stats = defaultdict(lambda: {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0})
overall = {'total': 0, 'pass': 0, 'fail': 0, 'duration_sum': 0.0}

with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    missing = [c for c in required_cols if c not in headers]
    if missing:
        raise ValueError(f'Missing required columns: {missing}; found {headers}')
    for row_num, row in enumerate(reader, start=2):
        model = row['Device Model']
        status = row['Status']
        try:
            duration = float(row['DurationSeconds'])
        except Exception as e:
            raise ValueError(f'Invalid DurationSeconds at CSV row {row_num}: {row["DurationSeconds"]!r}') from e
        bucket = stats[model]
        bucket['total'] += 1
        overall['total'] += 1
        if status == 'Pass':
            bucket['pass'] += 1
            overall['pass'] += 1
        elif status == 'Fail':
            bucket['fail'] += 1
            overall['fail'] += 1
        bucket['duration_sum'] += duration
        overall['duration_sum'] += duration

wb = Workbook()
ws = wb.active
ws.title = 'By Device Model'
ws.append(summary_headers)
for model in sorted(stats):
    s = stats[model]
    total = s['total']
    pass_count = s['pass']
    fail_count = s['fail']
    pass_rate = pass_count / total if total else 0
    avg_duration = s['duration_sum'] / total if total else 0
    ws.append([model, total, pass_count, fail_count, pass_rate, avg_duration])

# Basic number formats
for cell in ws['E'][1:]:
    cell.number_format = '0.00%'
for cell in ws['F'][1:]:
    cell.number_format = '0.00'

ws_overall = wb.create_sheet('Overall')
ws_overall.append(overall_headers)
ototal = overall['total']
ws_overall.append([
    ototal,
    overall['pass'],
    overall['fail'],
    overall['pass'] / ototal if ototal else 0,
    overall['duration_sum'] / ototal if ototal else 0,
])
ws_overall['D2'].number_format = '0.00%'
ws_overall['E2'].number_format = '0.00'

for sheet in [ws, ws_overall]:
    for col_cells in sheet.columns:
        max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col_cells)
        sheet.column_dimensions[col_cells[0].column_letter].width = max_len + 2

wb.save(output_path)
print('wrote:', output_path)
print('models:', len(stats))
print('total_rows:', overall['total'])
