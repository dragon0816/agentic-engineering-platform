from pathlib import Path
import subprocess
import sys
import csv
import math
from openpyxl import load_workbook

ws = Path.cwd()
out = ws / 'output'
out.mkdir(exist_ok=True)

fixed_code = r'''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv
from collections import defaultdict

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    totals = defaultdict(float)
    counts = defaultdict(int)

    with open(args.input, newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        for row in reader:
            model = row["Model"]
            value_text = row["Value"]
            if value_text is None or value_text.strip() == "":
                continue
            totals[model] += float(value_text)
            counts[model] += 1

    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(counts):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(args.output)


if __name__ == "__main__":
    main()
'''

export_path = out / 'export.py'
export_path.write_text(fixed_code, encoding='utf-8')
print('Wrote', export_path)

cmd = [sys.executable, str(export_path), '--input', str(ws / 'input' / 'input.csv'), '--output', str(out / 'summary.xlsx')]
print('Running:', ' '.join(cmd))
completed = subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('Return code:', completed.returncode)
if completed.stdout:
    print('STDOUT:', completed.stdout)
if completed.stderr:
    print('STDERR:', completed.stderr)
if completed.returncode != 0:
    raise SystemExit(completed.returncode)

# Independently compute expected results from input.csv for validation.
expected = {}
with (ws / 'input' / 'input.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('Input headers:', reader.fieldnames)
    for row in reader:
        value = row['Value']
        if value is None or value.strip() == '':
            continue
        model = row['Model']
        total, count = expected.get(model, (0.0, 0))
        expected[model] = (total + float(value), count + 1)
expected_rows = [(model, total / count, count) for model, (total, count) in sorted(expected.items())]
print('Expected rows:', expected_rows)

summary_path = out / 'summary.xlsx'
if not summary_path.exists():
    raise AssertionError('summary.xlsx was not created')

wb = load_workbook(summary_path, data_only=True)
if wb.sheetnames != ['summary']:
    raise AssertionError(f'Unexpected sheet names: {wb.sheetnames}')
wsheet = wb['summary']
actual_rows = list(wsheet.iter_rows(values_only=True))
print('Actual rows:', actual_rows)

headers = actual_rows[0]
if headers != ('Model', 'Average', 'ValidCount'):
    raise AssertionError(f'Unexpected headers: {headers}')
if len(actual_rows) != len(expected_rows) + 1:
    raise AssertionError(f'Unexpected row count: {len(actual_rows)} vs expected {len(expected_rows)+1}')

for actual, expected_row in zip(actual_rows[1:], expected_rows):
    model, avg, count = actual
    exp_model, exp_avg, exp_count = expected_row
    if model != exp_model:
        raise AssertionError(f'Model mismatch: {model} vs {exp_model}')
    if not isinstance(avg, (int, float)) or not math.isclose(float(avg), exp_avg, rel_tol=0, abs_tol=1e-12):
        raise AssertionError(f'Average mismatch for {model}: {avg} vs {exp_avg}')
    if not isinstance(count, int) or count != exp_count:
        raise AssertionError(f'ValidCount mismatch for {model}: {count} vs {exp_count}')

print('Validation passed: output/export.py exists, CLI command ran, summary.xlsx sheet/headers/rows/types/values verified.')