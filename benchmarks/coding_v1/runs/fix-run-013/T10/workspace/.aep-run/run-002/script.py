from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urljoin
import urllib.request
import re
import openpyxl

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
URL = 'http://127.0.0.1:63795/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_catalog = False
        self.catalog_depth = 0
        self.in_card = False
        self.card_depth = 0
        self.current = None
        self.capture = None
        self.products = []

    def _attrs(self, attrs):
        return dict(attrs)

    def handle_starttag(self, tag, attrs):
        a = self._attrs(attrs)
        if not self.in_catalog and tag == 'main' and a.get('id') == 'catalog':
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if self.in_catalog:
            self.catalog_depth += 1
            classes = set(a.get('class', '').split())
            if not self.in_card and tag == 'div' and 'card' in classes:
                self.in_card = True
                self.card_depth = 1
                self.current = {'name_parts': [], 'href': '', 'price_parts': []}
                return
            if self.in_card:
                self.card_depth += 1
                if tag == 'a':
                    self.current['href'] = a.get('href', '')
                    self.capture = 'name'
                elif tag == 'span' and 'price' in classes:
                    self.capture = 'price'

    def handle_endtag(self, tag):
        if self.in_card:
            if (tag == 'a' and self.capture == 'name') or (tag == 'span' and self.capture == 'price'):
                self.capture = None
            self.card_depth -= 1
            if self.card_depth == 0:
                name = ''.join(self.current['name_parts']).strip()
                price_text = ''.join(self.current['price_parts']).strip()
                href = self.current['href'].strip()
                if name and price_text and href:
                    self.products.append({'name': name, 'price_text': price_text, 'href': href})
                self.in_card = False
                self.current = None
                self.capture = None
                return
        if self.in_catalog:
            self.catalog_depth -= 1
            if self.catalog_depth == 0:
                self.in_catalog = False

    def handle_data(self, data):
        if self.in_card and self.capture == 'name':
            self.current['name_parts'].append(data)
        elif self.in_card and self.capture == 'price':
            self.current['price_parts'].append(data)

def parse_price(price_text):
    s = price_text.strip()
    m = re.match(r'^([^\d\-+.,]*)([\d,]+(?:\.\d+)?)', s)
    if not m:
        raise ValueError(f'Cannot parse price: {price_text!r}')
    currency = m.group(1).strip()
    number = m.group(2).replace(',', '')
    value = float(number) if '.' in number else int(number)
    return currency, value

with urllib.request.urlopen(URL, timeout=10) as r:
    html = r.read().decode('utf-8', errors='replace')
    base_url = r.geturl()

parser = CatalogParser()
parser.feed(html)

rows = []
for p in parser.products:
    currency, price = parse_price(p['price_text'])
    rows.append({
        '產品名稱': p['name'],
        '價格': price,
        '幣別': currency,
        '產品網址': urljoin(base_url, p['href']),
    })

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'products'
headers = ['產品名稱', '價格', '幣別', '產品網址']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
for cell in ws[1]:
    cell.style = 'Headline 4'
for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
    row[0].number_format = '#,##0.00' if isinstance(row[0].value, float) else '#,##0'
for col in ws.columns:
    max_len = max(len(str(c.value)) if c.value is not None else 0 for c in col)
    ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 80)
wb.save(OUTFILE)

# Verification: reopen output and confirm all stated requirements.
wb2 = openpyxl.load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
read_headers = [ws2.cell(1, c).value for c in range(1, 5)]
read_rows = []
for r in range(2, ws2.max_row + 1):
    read_rows.append([ws2.cell(r, c).value for c in range(1, 5)])

assert read_headers == headers, f'headers mismatch: {read_headers}'
assert len(read_rows) == 5, f'expected 5 catalog products, got {len(read_rows)}'
expected_names = ['Router A', 'Switch B', 'Cable C', 'Adapter D', 'Antenna E']
assert [r[0] for r in read_rows] == expected_names, 'page order or product names mismatch'
assert all(isinstance(r[1], (int, float)) and not isinstance(r[1], bool) for r in read_rows), 'prices are not numeric'
assert [r[2] for r in read_rows] == ['NT$'] * 5, 'currency mismatch'
assert all(str(r[3]).startswith('http://127.0.0.1:63795/') for r in read_rows), 'URLs are not complete'
assert not any('Sponsored' in str(r[0]) or '推薦' in str(r[0]) for r in read_rows), 'ad/recommendation item included'

print(f'Wrote: {OUTFILE}')
print('Rows:', len(read_rows))
print('Headers:', read_headers)
print('Data:')
for r in read_rows:
    print(r)
print('Verified: products.xlsx exists; sheet products has required columns; 5 main catalog rows in page order; prices are numeric; currency is NT$; product URLs are absolute; ads/recommendations excluded.')
