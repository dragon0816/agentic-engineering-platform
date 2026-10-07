import csv
import json
import os
from pathlib import Path
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = json.load(open(params_path, 'r', encoding='utf-8')) if params_path and Path(params_path).exists() else {}
threshold = int(params.get('urgent_threshold', 20))

print(f'workspace={workspace}')
print(f'urgent_threshold={threshold}')
print('input files:')
for p in sorted(input_dir.rglob('*')):
    print(' -', p.relative_to(workspace), 'dir' if p.is_dir() else f'file {p.stat().st_size} bytes')

sop_path = input_dir / 'supplied_sop.md'
if sop_path.exists():
    sop_text = sop_path.read_text(encoding='utf-8')
    print(f'supplied_sop.md bytes={len(sop_text.encode("utf-8"))}')
else:
    raise FileNotFoundError(f'Missing SOP: {sop_path}')

candidates = [input_dir / 'results.csv', input_dir / 'data' / 'results.csv', input_dir / 'data' / 'test_results.csv']
csv_path = next((p for p in candidates if p.exists()), None)
if csv_path is None:
    all_csvs = sorted(input_dir.rglob('*.csv'))
    if len(all_csvs) == 1:
        csv_path = all_csvs[0]
    elif not all_csvs:
        raise FileNotFoundError('No CSV file found under input/')
    else:
        raise RuntimeError('Multiple CSV files found and none matched expected names: ' + ', '.join(str(p.relative_to(workspace)) for p in all_csvs))
print(f'csv_path={csv_path.relative_to(workspace)}')

required_cols = ['Device ID', 'Status', 'Value']
rows = []
with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    print(f'csv_headers={reader.fieldnames}')
    if reader.fieldnames is None:
        raise RuntimeError('CSV has no header row')
    missing = [c for c in required_cols if c not in reader.fieldnames]
    if missing:
        raise RuntimeError(f'CSV missing required columns: {missing}')
    for row in reader:
        rows.append(row)

failures = []
pass_count = 0
fail_count = 0
urgent_count = 0
for row in rows:
    status = str(row['Status']).strip()
    if status == 'Pass':
        pass_count += 1
    if status == 'Fail':
        fail_count += 1
        value_text = row['Value']
        try:
            value = float(value_text)
        except Exception as exc:
            raise RuntimeError(f"Value is not numeric for Device ID {row.get('Device ID')}: {value_text!r}") from exc
        urgent = value > threshold
        if urgent:
            urgent_count += 1
        failures.append({'Device ID': row['Device ID'], 'Status': row['Status'], 'Value': value_text, 'Urgent': urgent})

total_count = len(rows)

xlsx_path = output_dir / 'failures.xlsx'
md_path = output_dir / 'summary.md'

wb = Workbook()
ws = wb.active
ws.title = 'Failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for item in failures:
    value_text = item['Value']
    try:
        n = float(value_text)
        cell_value = int(n) if n.is_integer() else n
    except Exception:
        cell_value = value_text
    ws.append([item['Device ID'], item['Status'], cell_value, item['Urgent']])
wb.save(xlsx_path)

summary = (
    '# Test Failure Summary\n\n'
    f'- Total: {total_count}\n'
    f'- Pass: {pass_count}\n'
    f'- Fail: {fail_count}\n'
    f'- Urgent: {urgent_count}\n'
)
md_path.write_text(summary, encoding='utf-8')

check_wb = load_workbook(xlsx_path, data_only=True)
check_ws = check_wb['Failures']
read_headers = [check_ws.cell(row=1, column=i).value for i in range(1, 5)]
assert read_headers == headers, f'Unexpected XLSX headers: {read_headers}'
assert check_ws.max_row == len(failures) + 1, f'Unexpected XLSX row count: {check_ws.max_row}'
urgent_values = []
for r in range(2, check_ws.max_row + 1):
    assert check_ws.cell(r, 2).value == 'Fail', f'Non-Fail row in failures.xlsx row {r}'
    device = check_ws.cell(r, 1).value
    val = check_ws.cell(r, 3).value
    urgent = check_ws.cell(r, 4).value
    assert isinstance(device, str) and device, f'Device ID not text at row {r}'
    assert isinstance(val, (int, float)), f'Value not numeric at row {r}: {val!r}'
    assert isinstance(urgent, bool), f'Urgent not boolean at row {r}: {urgent!r}'
    assert urgent == (float(val) > threshold), f'Urgent mismatch at row {r}'
    urgent_values.append(urgent)
check_summary = md_path.read_text(encoding='utf-8')
expected_lines = [f'- Total: {total_count}', f'- Pass: {pass_count}', f'- Fail: {fail_count}', f'- Urgent: {urgent_count}']
for line in expected_lines:
    assert line in check_summary, f'Missing summary line: {line}'

print(f'wrote={xlsx_path.relative_to(workspace)}')
print(f'wrote={md_path.relative_to(workspace)}')
print(f'checked failures.xlsx headers={headers}, fail_rows={len(failures)}, urgent_bools={urgent_values}')
print(f'checked summary counts total={total_count} pass={pass_count} fail={fail_count} urgent={urgent_count}')
