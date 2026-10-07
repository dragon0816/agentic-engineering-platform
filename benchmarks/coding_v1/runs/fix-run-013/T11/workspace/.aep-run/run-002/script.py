import os, re, urllib.request
from urllib.parse import urljoin
from html.parser import HTMLParser
from openpyxl import Workbook, load_workbook

BASE_LIST_URL = 'http://127.0.0.1:63795/t11/list'
OUT_PATH = os.path.join('output', 'details.xlsx')
COLUMNS = ['Name', 'Price', 'Currency', 'CPU', 'MemoryGB', 'URL']

class HTMLTextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self.text_parts = []
        self.tag_stack = []
        self.rows = []
        self.current_row = None
        self.current_cell = None
        self.current_cell_text = None
    def handle_starttag(self, tag, attrs):
        tag = tag.lower()
        attrs_d = dict(attrs)
        self.tag_stack.append(tag)
        if tag == 'a':
            self.links.append(attrs_d.get('href', ''))
        if tag == 'tr':
            self.current_row = []
        if tag in ('td','th') and self.current_row is not None:
            self.current_cell = tag
            self.current_cell_text = []
    def handle_endtag(self, tag):
        tag = tag.lower()
        if tag in ('td','th') and self.current_row is not None and self.current_cell_text is not None:
            txt = ' '.join(''.join(self.current_cell_text).split())
            self.current_row.append(txt)
            self.current_cell = None
            self.current_cell_text = None
        if tag == 'tr' and self.current_row is not None:
            self.rows.append(self.current_row)
            self.current_row = None
        if self.tag_stack:
            try:
                idx = len(self.tag_stack) - 1 - self.tag_stack[::-1].index(tag)
                self.tag_stack.pop(idx)
            except ValueError:
                pass
    def handle_data(self, data):
        if data and data.strip():
            self.text_parts.append(data.strip())
        if self.current_cell_text is not None:
            self.current_cell_text.append(data)


def fetch(url):
    with urllib.request.urlopen(url, timeout=10) as r:
        data = r.read()
    return data.decode('utf-8', errors='replace')


def parser_for(html):
    p = HTMLTextParser()
    p.feed(html)
    return p


def normalize_key(k):
    return re.sub(r'[^a-z0-9]', '', k.lower())


def parse_money(value):
    if value is None:
        return ('', '')
    v = str(value).strip()
    if not v:
        return ('', '')
    # Currency can be a symbol or code before/after the number.
    m = re.search(r'(?P<before>[A-Z]{3}|[$€£¥])?\s*(?P<num>[+-]?(?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d+)?)\s*(?P<after>[A-Z]{3}|[$€£¥])?', v)
    if not m:
        return (v, '')
    num_s = m.group('num').replace(',', '')
    try:
        price = float(num_s) if '.' in num_s else int(num_s)
    except ValueError:
        price = num_s
    cur = m.group('before') or m.group('after') or ''
    symbol_map = {'$':'USD', '€':'EUR', '£':'GBP', '¥':'JPY'}
    cur = symbol_map.get(cur, cur)
    return (price, cur)


def parse_memory_gb(value):
    if value is None:
        return ''
    v = str(value).strip()
    if not v:
        return ''
    # Accept forms like 8192 MB, 16GB, 1 TB, 0.5TB. If unit omitted, leave numeric as GB.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*([KMGT]?B)?', v, re.I)
    if not m:
        return v
    num = float(m.group(1))
    unit = (m.group(2) or 'GB').upper()
    if unit == 'KB':
        gb = num / (1024*1024)
    elif unit == 'MB':
        gb = num / 1024
    elif unit == 'GB':
        gb = num
    elif unit == 'TB':
        gb = num * 1024
    else:
        gb = num
    return int(gb) if abs(gb - int(gb)) < 1e-9 else gb


def extract_detail(html):
    p = parser_for(html)
    text = '\n'.join(p.text_parts)
    data = {}

    # Two-column detail tables / definition-like rows.
    for row in p.rows:
        if len(row) >= 2:
            key = normalize_key(row[0].rstrip(':'))
            val = row[1].strip()
            if key:
                data[key] = val

    # Label: value from plain text.
    for label in ['Name', 'Price', 'Currency', 'CPU', 'Memory', 'MemoryGB']:
        pattern = re.compile(r'\b' + re.escape(label) + r'\b\s*[:：]\s*([^\n\r<]+)', re.I)
        m = pattern.search(text)
        if m:
            data[normalize_key(label)] = m.group(1).strip()

    # Common JSON-ish or data attributes fallback.
    for label in ['name', 'price', 'currency', 'cpu', 'memory', 'memoryGB']:
        pattern = re.compile(r'["\']' + re.escape(label) + r'["\']\s*[:=]\s*["\']?([^,"\'<>}\n]+)', re.I)
        m = pattern.search(html)
        if m:
            data[normalize_key(label)] = m.group(1).strip()

    name = data.get('name', '')
    price = ''
    currency = data.get('currency', '')
    if 'price' in data:
        price, cur2 = parse_money(data.get('price'))
        if not currency:
            currency = cur2
    cpu = data.get('cpu', '')
    mem_source = data.get('memorygb', data.get('memory', ''))
    mem_gb = parse_memory_gb(mem_source)
    return {'Name': name, 'Price': price, 'Currency': currency, 'CPU': cpu, 'MemoryGB': mem_gb}

# Inspect/fetch list
list_html = fetch(BASE_LIST_URL)
list_parser = parser_for(list_html)
detail_urls = []
for href in list_parser.links:
    full = urljoin(BASE_LIST_URL, href)
    if full not in detail_urls:
        detail_urls.append(full)
print('detail URL count:', len(detail_urls))
for u in detail_urls:
    print('detail URL:', u)

records = []
print('\nDetail page samples and extracted records:')
for url in detail_urls:
    html = fetch(url)
    print('\n---', url, 'length', len(html), '---')
    print(html[:1000])
    rec = extract_detail(html)
    rec['URL'] = url
    records.append(rec)
    print('extracted:', rec)

# Write workbook
os.makedirs('output', exist_ok=True)
wb = Workbook()
ws = wb.active
ws.title = 'details'
ws.append(COLUMNS)
for rec in records:
    ws.append([rec.get(c, '') for c in COLUMNS])
wb.save(OUT_PATH)
print('\nwrote:', OUT_PATH)

# Verification: reopen output and confirm requirements.
wb2 = load_workbook(OUT_PATH, data_only=True)
ws2 = wb2.active
headers = [cell.value for cell in ws2[1]]
rows = list(ws2.iter_rows(min_row=2, values_only=True))
print('verification headers:', headers)
print('verification row count:', len(rows))
assert headers == COLUMNS, f'headers mismatch: {headers}'
assert len(rows) == len(detail_urls), f'row count mismatch {len(rows)} vs {len(detail_urls)}'
for i, row in enumerate(rows, start=2):
    row_d = dict(zip(COLUMNS, row))
    assert row_d['URL'] in detail_urls, f'bad URL in row {i}: {row_d["URL"]}'
    mem = row_d['MemoryGB']
    assert mem is None or isinstance(mem, (int, float)), f'MemoryGB not numeric/blank in row {i}: {mem!r}'
    price = row_d['Price']
    assert price is None or isinstance(price, (int, float)), f'Price not numeric/blank in row {i}: {price!r}'
print('verification rows:')
for row in rows:
    print(row)
print('all checks passed')
