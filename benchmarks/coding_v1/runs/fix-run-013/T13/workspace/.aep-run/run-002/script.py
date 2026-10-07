import json
import re
import time
import urllib.request
import urllib.error
from html.parser import HTMLParser
from urllib.parse import urljoin
from pathlib import Path
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
LIST_URL = 'http://127.0.0.1:63795/t13/list'
XLSX_PATH = OUTPUT / 'robust_products.xlsx'
WARNINGS_PATH = OUTPUT / 'warnings.json'
COLUMNS = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.items = []
        self._href = None
        self._text_parts = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs_d = dict(attrs)
            if 'href' in attrs_d:
                self._href = attrs_d['href']
                self._text_parts = []
    def handle_data(self, data):
        if self._href is not None:
            self._text_parts.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._href is not None:
            name = ' '.join(''.join(self._text_parts).split())
            self.items.append({'Name': name or 'N/A', 'URL': urljoin(LIST_URL, self._href)})
            self._href = None
            self._text_parts = []

class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.texts = []
    def handle_data(self, data):
        s = data.strip()
        if s:
            self.texts.append(s)

def fetch(url):
    req = urllib.request.Request(url, headers={'User-Agent': 'Agent/1.0'})
    with urllib.request.urlopen(req, timeout=10) as r:
        return r.status, r.read().decode('utf-8', errors='replace')

def fetch_detail_with_retry(url, max_attempts=3):
    last_error = None
    for attempt in range(1, max_attempts + 1):
        try:
            status, body = fetch(url)
            return {'ok': True, 'status': status, 'body': body, 'attempts': attempt, 'error': None}
        except urllib.error.HTTPError as e:
            last_error = f'HTTP {e.code}'
            if e.code == 500 and attempt < max_attempts:
                time.sleep(0.1)
                continue
            return {'ok': False, 'status': e.code, 'body': '', 'attempts': attempt, 'error': last_error}
        except Exception as e:
            last_error = type(e).__name__ + ': ' + str(e)
            return {'ok': False, 'status': None, 'body': '', 'attempts': attempt, 'error': last_error}
    return {'ok': False, 'status': None, 'body': '', 'attempts': max_attempts, 'error': last_error or 'unknown'}

def normalize_key(s):
    return re.sub(r'[^a-z0-9]+', '', s.lower())

def parse_detail(html, fallback_name):
    parser = TextParser()
    parser.feed(html)
    texts = parser.texts
    result = {'Name': fallback_name or 'N/A', 'Price': 'N/A', 'Currency': 'N/A', 'CPU': 'N/A'}

    # Parse common label/value text patterns: Price: 123 USD, CPU: Snapdragon, Name: Phone A
    joined = '\n'.join(texts)
    for raw_line in texts + joined.split('\n'):
        line = ' '.join(raw_line.split())
        if not line:
            continue
        m = re.match(r'^(Name|Product Name)\s*[:：-]\s*(.+)$', line, re.I)
        if m and m.group(2).strip():
            result['Name'] = m.group(2).strip()
        m = re.match(r'^(CPU|Processor|Chipset|Spec|Specs|Specification|Specifications)\s*[:：-]\s*(.+)$', line, re.I)
        if m and m.group(2).strip():
            result['CPU'] = m.group(2).strip()
        m = re.match(r'^(Price)\s*[:：-]\s*(.+)$', line, re.I)
        if m:
            price_text = m.group(2).strip()
            parse_price_into(result, price_text)

    # Parse adjacent labels and values if the page is table-like and parser split them.
    for i, t in enumerate(texts[:-1]):
        key = normalize_key(t)
        val = texts[i+1].strip()
        if key in {'name', 'productname'} and val:
            result['Name'] = val
        elif key in {'cpu', 'processor', 'chipset', 'spec', 'specs', 'specification', 'specifications'} and val:
            result['CPU'] = val
        elif key == 'price' and val:
            parse_price_into(result, val)

    # Last resort: any currency-looking amount in whole page.
    if result['Price'] == 'N/A':
        parse_price_into(result, joined)
    return {k: (str(v).strip() if str(v).strip() else 'N/A') for k, v in result.items()}

def parse_price_into(result, price_text):
    text = ' '.join(str(price_text).split())
    # currency before or after amount, allowing symbols and 3-letter codes.
    m = re.search(r'(?:(USD|TWD|NTD|EUR|JPY|CNY|GBP|\$|NT\$|€|¥)\s*)?([0-9]+(?:[,][0-9]{3})*(?:\.[0-9]+)?|[0-9]+(?:\.[0-9]+)?)(?:\s*(USD|TWD|NTD|EUR|JPY|CNY|GBP|dollars?))?', text, re.I)
    if not m:
        return
    currency = m.group(1) or m.group(3) or 'N/A'
    currency_map = {'$': 'USD', 'NT$': 'TWD', '€': 'EUR', '¥': 'JPY'}
    currency = currency_map.get(currency, currency.upper() if currency != 'N/A' else 'N/A')
    if currency == 'DOLLARS' or currency == 'DOLLAR':
        currency = 'USD'
    result['Price'] = m.group(2).replace(',', '')
    result['Currency'] = currency

# Fetch product list
status, list_html = fetch(LIST_URL)
list_parser = ProductListParser()
list_parser.feed(list_html)
products = list_parser.items[:5]
print('discovered products:', products)
if len(products) != 5:
    raise RuntimeError(f'Expected 5 products from list, found {len(products)}')

rows = []
warnings = []
for item in products:
    detail = fetch_detail_with_retry(item['URL'], max_attempts=3)
    if detail['ok']:
        parsed = parse_detail(detail['body'], item['Name'])
        row = {
            'Name': parsed.get('Name') or item['Name'] or 'N/A',
            'Price': parsed.get('Price') or 'N/A',
            'Currency': parsed.get('Currency') or 'N/A',
            'CPU': parsed.get('CPU') or 'N/A',
            'URL': item['URL'],
            'FetchStatus': 'ok',
        }
        warn = {'Name': row['Name'], 'URL': item['URL'], 'FetchStatus': 'ok', 'status': 'ok', 'attempts': detail['attempts'], 'http_status': detail['status']}
    else:
        row = {
            'Name': item['Name'] or 'N/A',
            'Price': 'N/A',
            'Currency': 'N/A',
            'CPU': 'N/A',
            'URL': item['URL'],
            'FetchStatus': 'error',
        }
        warn = {'Name': item['Name'] or 'N/A', 'URL': item['URL'], 'FetchStatus': 'error', 'status': 'error', 'attempts': detail['attempts'], 'http_status': detail['status'], 'warning': detail['error'] or 'detail fetch failed'}
    rows.append(row)
    warnings.append(warn)
    print('row:', row, 'warning:', warn)

# Write workbook
wb = Workbook()
ws = wb.active
ws.title = 'Products'
ws.append(COLUMNS)
for row in rows:
    ws.append([row[col] for col in COLUMNS])
for cell in ws['A'] + ws['B'] + ws['C'] + ws['D'] + ws['E'] + ws['F']:
    cell.number_format = '@'
wb.save(XLSX_PATH)

with open(WARNINGS_PATH, 'w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Reopen and validate stated requirements
wb2 = load_workbook(XLSX_PATH, data_only=True)
ws2 = wb2['Products']
headers = [ws2.cell(1, c).value for c in range(1, len(COLUMNS)+1)]
assert headers == COLUMNS, headers
loaded_rows = []
for r in range(2, ws2.max_row + 1):
    loaded_rows.append({headers[c-1]: ws2.cell(r, c).value for c in range(1, len(COLUMNS)+1)})
assert len(loaded_rows) == 5, len(loaded_rows)
for row in loaded_rows:
    assert set(row) == set(COLUMNS)
    assert row['Name'] not in (None, '')
    assert row['URL'] not in (None, '') and str(row['URL']).startswith('http://127.0.0.1:63795/t13/item/')
    assert row['FetchStatus'] in ('ok', 'error')
    for field in ['Price', 'Currency', 'CPU']:
        assert row[field] not in (None, ''), (field, row)
    if row['FetchStatus'] == 'error':
        assert row['Price'] == 'N/A' and row['Currency'] == 'N/A' and row['CPU'] == 'N/A'
with open(WARNINGS_PATH, 'r', encoding='utf-8') as f:
    warnings2 = json.load(f)
assert isinstance(warnings2, list) and len(warnings2) == 5
for w, row in zip(warnings2, loaded_rows):
    assert w['URL'] == row['URL']
    assert w['FetchStatus'] == row['FetchStatus']
    assert w['status'] == row['FetchStatus']
    if row['FetchStatus'] == 'ok':
        assert w['status'] == 'ok'
print('validated xlsx:', XLSX_PATH)
print('validated warnings:', WARNINGS_PATH)
print('validated headers:', headers)
print('validated row count:', len(loaded_rows))
print('validated rows:', loaded_rows)
print('validated warnings json:', warnings2)
