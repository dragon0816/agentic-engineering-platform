from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read source rows, preserving source column names and row order.
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames
    if headers is None:
        raise ValueError('results.csv has no header row')
    if 'Status' not in headers:
        raise ValueError('Missing required column: Status')
    if 'Device ID' not in headers:
        raise ValueError('Missing required column: Device ID')
    source_rows = list(reader)

failed_rows = [row for row in source_rows if (row.get('Status') or '').strip().lower() == 'fail']

wb = Workbook()
ws = wb.active
ws.title = 'failed_items'
ws.append(headers)

for row in failed_rows:
    values = []
    for h in headers:
        value = row.get(h, '')
        # Keep CSV values as text where appropriate; Device ID explicitly text.
        values.append(value)
    ws.append(values)

# Explicitly mark Device ID data cells as text to preserve leading zeroes in Excel.
device_col_idx = headers.index('Device ID') + 1
for cell in ws.iter_cols(min_col=device_col_idx, max_col=device_col_idx, min_row=2, max_row=ws.max_row):
    for c in cell:
        c.value = '' if c.value is None else str(c.value)
        c.number_format = '@'

wb.save(output_path)

# Verification by reopening saved output.
check_wb = load_workbook(output_path, data_only=False)
assert 'failed_items' in check_wb.sheetnames, 'failed_items sheet missing'
check_ws = check_wb['failed_items']
actual_headers = [check_ws.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
assert actual_headers == headers, f'Headers changed: {actual_headers} != {headers}'
actual_data = []
for r in range(2, check_ws.max_row + 1):
    actual_data.append({h: check_ws.cell(row=r, column=i+1).value for i, h in enumerate(headers)})
assert len(actual_data) == len(failed_rows), f'Row count mismatch: {len(actual_data)} != {len(failed_rows)}'
for i, (actual, expected) in enumerate(zip(actual_data, failed_rows), start=1):
    for h in headers:
        av = actual[h]
        ev = expected.get(h, '')
        if av is None:
            av = ''
        assert str(av) == str(ev), f'Row {i} column {h}: {av!r} != {ev!r}'
    assert expected['Status'].strip().lower() == 'fail', f'Non-fail row included at output row {i+1}'

# Device ID text format and leading zero checks.
for out_idx, expected in enumerate(failed_rows, start=2):
    cell = check_ws.cell(row=out_idx, column=device_col_idx)
    assert isinstance(cell.value, str), f'Device ID at row {out_idx} is not text: {cell.value!r}'
    assert cell.value == expected['Device ID'], f'Device ID leading zero/value mismatch at row {out_idx}'
    assert cell.number_format == '@', f'Device ID row {out_idx} not formatted as text'

print('Created:', output_path)
print('Worksheet:', check_ws.title)
print('Headers:', actual_headers)
print('Failed row count:', len(failed_rows))
print('Output Device IDs:', [check_ws.cell(row=r, column=device_col_idx).value for r in range(2, check_ws.max_row + 1)])
print('Checks passed: output file exists, sheet name, all original columns, original order, fail filter, row count, and Device ID text/leading zero preservation.')
