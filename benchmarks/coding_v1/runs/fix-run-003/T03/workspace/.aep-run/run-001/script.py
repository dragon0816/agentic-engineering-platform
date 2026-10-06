from pathlib import Path
from openpyxl import load_workbook, Workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'tasks.xlsx'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'owners.xlsx'

# Load source workbook without modifying it
src_wb = load_workbook(input_path, data_only=False)
if 'tasks' not in src_wb.sheetnames:
    raise ValueError("Input workbook does not contain a 'tasks' worksheet")
ws = src_wb['tasks']

# Read header and locate required columns
headers = [cell.value for cell in ws[1]]

def find_col(name):
    for idx, val in enumerate(headers, start=1):
        if val == name:
            return idx
    raise ValueError(f"Required column {name!r} not found in tasks worksheet")

owner_col = find_col('Owner')
status_col = find_col('Status')

# Collect rows by display/group owner. Preserve original row values in detail sheets.
groups = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    # Skip completely empty trailing rows
    if all(v is None for v in row):
        continue
    owner_value = row[owner_col - 1] if owner_col - 1 < len(row) else None
    if owner_value is None or str(owner_value).strip() == '':
        group_owner = 'Unassigned'
    else:
        group_owner = owner_value
    groups.setdefault(group_owner, []).append(row)

invalid_chars = set('[]:*?/\\')

def legal_sheet_base(owner):
    # Convert to text for sheet naming, replacing illegal Excel sheet-name characters.
    text = str(owner) if owner is not None else 'Unassigned'
    text = ''.join('_' if ch in invalid_chars else ch for ch in text)
    text = text.strip()
    if not text:
        text = 'Unassigned'
    # Excel worksheet names are limited to 31 chars. Reserve suffix room later.
    return text[:31]

used_names = {'Summary'}
owner_to_sheet = {}

for owner in groups.keys():
    base = legal_sheet_base(owner)
    # Avoid collisions with Summary or with other sanitized names, case-insensitively like Excel.
    candidate = base[:31]
    counter = 1
    used_lower = {n.lower() for n in used_names}
    while candidate.lower() in used_lower:
        suffix = f"_{counter}"
        candidate = base[:31 - len(suffix)] + suffix
        counter += 1
        used_lower = {n.lower() for n in used_names}
    used_names.add(candidate)
    owner_to_sheet[owner] = candidate

# Build output workbook
out_wb = Workbook()
summary_ws = out_wb.active
summary_ws.title = 'Summary'
summary_ws.append(['Owner', 'Total', 'Pass', 'Fail', 'SheetName'])

for owner, rows in groups.items():
    sheet_name = owner_to_sheet[owner]
    detail_ws = out_wb.create_sheet(title=sheet_name)
    detail_ws.append(headers)
    for row in rows:
        # Pad rows to header length for a consistent table shape.
        padded = list(row) + [None] * max(0, len(headers) - len(row))
        detail_ws.append(padded[:len(headers)])

    total = len(rows)
    pass_count = 0
    fail_count = 0
    for row in rows:
        status = row[status_col - 1] if status_col - 1 < len(row) else None
        if status == 'Pass':
            pass_count += 1
        elif status == 'Fail':
            fail_count += 1

    summary_owner = owner if not (owner is None or str(owner).strip() == '') else 'Unassigned'
    summary_ws.append([summary_owner, total, pass_count, fail_count, sheet_name])

# Optional: give readable widths
for ws_out in out_wb.worksheets:
    for col_cells in ws_out.columns:
        letter = col_cells[0].column_letter
        max_len = 0
        for cell in col_cells:
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws_out.column_dimensions[letter].width = min(max_len + 2, 60)

out_wb.save(output_path)
print(f"Created {output_path}")
print(f"Groups written: {len(groups)}")
for owner, sheet in owner_to_sheet.items():
    print(f"{owner!r} -> {sheet}")