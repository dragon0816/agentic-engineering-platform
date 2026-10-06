from pathlib import Path
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import re

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

devices_path = input_dir / 'devices.xlsx'
tests_path = input_dir / 'tests.xlsx'
out_path = output_dir / 'merged_device_tests.xlsx'


def is_zero_format(fmt):
    """Return a likely zero-padding width from an Excel number format, or None."""
    if not fmt:
        return None
    fmt = str(fmt).split(';')[0]
    # Ignore quoted literals and bracket sections for this simple ID use case.
    fmt_clean = re.sub(r'"[^"]*"', '', fmt)
    fmt_clean = re.sub(r'\[[^\]]*\]', '', fmt_clean)
    m = re.fullmatch(r'0+', fmt_clean.strip())
    if m:
        return len(m.group(0))
    return None


def cell_to_text(cell, is_id=False):
    v = cell.value
    if v is None:
        return ''
    if is_id:
        if isinstance(v, str):
            return v
        width = is_zero_format(cell.number_format)
        if isinstance(v, bool):
            s = str(v)
        elif isinstance(v, int):
            s = str(v)
        elif isinstance(v, float):
            s = str(int(v)) if v.is_integer() else str(v)
        else:
            s = str(v)
        if width and s.isdigit():
            s = s.zfill(width)
        return s
    return v


def read_sheet(path):
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows())
    if not rows:
        raise ValueError(f'{path.name} is empty')
    headers = []
    seen = {}
    for idx, cell in enumerate(rows[0], start=1):
        name = '' if cell.value is None else str(cell.value)
        name = name.strip()
        if not name:
            name = f'Column{idx}'
        # make headers unique in case source workbook has duplicates
        if name in seen:
            seen[name] += 1
            name = f'{name}_{seen[name]}'
        else:
            seen[name] = 1
        headers.append(name)
    # Find ID column, preferring exact Device ID (case-insensitive) then common variants.
    norm = {h.lower().strip(): i for i, h in enumerate(headers)}
    id_idx = None
    for candidate in ('device id', 'device_id', 'id'):
        if candidate in norm:
            id_idx = norm[candidate]
            break
    if id_idx is None:
        raise ValueError(f'Could not find a Device ID column in {path.name}; headers={headers}')
    records = []
    for row_cells in rows[1:]:
        # pad row to header length
        padded = list(row_cells) + [None] * max(0, len(headers) - len(row_cells))
        rec = {}
        has_any = False
        for i, h in enumerate(headers):
            cell = padded[i]
            if cell is None:
                val = ''
            else:
                val = cell_to_text(cell, is_id=(i == id_idx))
            if val not in ('', None):
                has_any = True
            rec[h] = val
        if not has_any:
            continue
        rec['_DeviceID'] = rec[headers[id_idx]]
        records.append(rec)
    return headers, headers[id_idx], records


def prefixed_headers(prefix, headers, id_header):
    return [f'{prefix}{h}' for h in headers if h != id_header]


dev_headers, dev_id_header, devices = read_sheet(devices_path)
test_headers, test_id_header, tests = read_sheet(tests_path)

# devices has unique IDs by requirement; if there are accidental duplicates, keep first and report.
device_by_id = {}
duplicate_device_ids = []
for rec in devices:
    did = rec['_DeviceID']
    if did in device_by_id:
        duplicate_device_ids.append(did)
        continue
    device_by_id[did] = rec

matched_device_ids = set()
matched_rows = []
unmatched_tests = []
for test in tests:
    did = test['_DeviceID']
    dev = device_by_id.get(did)
    if dev is None:
        unmatched_tests.append(test)
    else:
        matched_device_ids.add(did)
        matched_rows.append((test, dev))

unmatched_devices = [dev for dev in devices if dev['_DeviceID'] not in matched_device_ids]

# Create output workbook.
wb_out = Workbook()
ws_m = wb_out.active
ws_m.title = 'matched'
ws_u = wb_out.create_sheet('unmatched')

matched_headers = ['Device ID'] + prefixed_headers('tests.', test_headers, test_id_header) + prefixed_headers('devices.', dev_headers, dev_id_header)
ws_m.append(matched_headers)
for test, dev in matched_rows:
    row = [test['_DeviceID']]
    row += [test[h] for h in test_headers if h != test_id_header]
    row += [dev[h] for h in dev_headers if h != dev_id_header]
    ws_m.append(row)

unmatched_headers = ['Source', 'Device ID'] + prefixed_headers('tests.', test_headers, test_id_header) + prefixed_headers('devices.', dev_headers, dev_id_header)
ws_u.append(unmatched_headers)
for test in unmatched_tests:
    row = ['tests', test['_DeviceID']]
    row += [test[h] for h in test_headers if h != test_id_header]
    row += ['' for h in dev_headers if h != dev_id_header]
    ws_u.append(row)
for dev in unmatched_devices:
    row = ['devices', dev['_DeviceID']]
    row += ['' for h in test_headers if h != test_id_header]
    row += [dev[h] for h in dev_headers if h != dev_id_header]
    ws_u.append(row)

# Format header and Device ID columns as text, preserving leading zero strings.
for ws in (ws_m, ws_u):
    for cell in ws[1]:
        cell.font = Font(bold=True)
    # Device ID is column 1 in matched, column 2 in unmatched.
    id_col = 1 if ws.title == 'matched' else 2
    for row in ws.iter_rows(min_row=2, min_col=id_col, max_col=id_col):
        row[0].number_format = '@'
    # Auto-ish width
    for col_idx, col in enumerate(ws.columns, start=1):
        max_len = 0
        for cell in col:
            text = '' if cell.value is None else str(cell.value)
            if len(text) > max_len:
                max_len = len(text)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 10), 45)

wb_out.save(out_path)

print(f'Devices rows read: {len(devices)}')
print(f'Tests rows read: {len(tests)}')
print(f'Matched rows written: {len(matched_rows)}')
print(f'Unmatched tests written: {len(unmatched_tests)}')
print(f'Unmatched devices written: {len(unmatched_devices)}')
if duplicate_device_ids:
    print(f'Warning: duplicate device IDs in devices ignored after first: {duplicate_device_ids}')
print(f'Output: {out_path}')
print(f'Device ID header in devices: {dev_id_header}')
print(f'Device ID header in tests: {test_id_header}')
print(f'Matched headers: {matched_headers}')
print(f'Unmatched headers: {unmatched_headers}')