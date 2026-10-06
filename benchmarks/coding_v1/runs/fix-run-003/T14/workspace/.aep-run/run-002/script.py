from pathlib import Path
import csv
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

workspace = Path.cwd()
input_path = workspace / "input" / "large_results.csv"
output_dir = workspace / "output"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / "summary.xlsx"

if not input_path.exists():
    raise FileNotFoundError(f"Input file not found: {input_path}")

def normalize(name: str) -> str:
    return ''.join(ch.lower() for ch in (name or '') if ch.isalnum())

with input_path.open("r", newline="", encoding="utf-8-sig") as f:
    reader = csv.DictReader(f)
    if not reader.fieldnames:
        raise ValueError("CSV has no header row")
    fieldnames = reader.fieldnames
    norm_to_name = {normalize(n): n for n in fieldnames}

    def pick(*candidates):
        for c in candidates:
            key = normalize(c)
            if key in norm_to_name:
                return norm_to_name[key]
        return None

    model_col = pick("Device Model", "DeviceModel", "Model")
    result_col = pick("Result", "Status", "Test Result", "Outcome")
    duration_col = pick("DurationSeconds", "Duration Seconds", "Duration", "ElapsedSeconds", "Elapsed Seconds")

    missing = [label for label, col in [("Device Model", model_col), ("Result/Status", result_col), ("DurationSeconds", duration_col)] if col is None]
    if missing:
        raise ValueError(f"Missing required column(s): {missing}. Found columns: {fieldnames}")

    stats = defaultdict(lambda: {"total": 0, "pass": 0, "fail": 0, "duration_sum": 0.0})
    overall = {"total": 0, "pass": 0, "fail": 0, "duration_sum": 0.0}

    for row_num, row in enumerate(reader, start=2):
        model = (row.get(model_col) or "").strip() or "(blank)"
        result = (row.get(result_col) or "").strip().lower()
        duration_raw = (row.get(duration_col) or "").strip()
        try:
            duration = float(duration_raw) if duration_raw != "" else 0.0
        except ValueError:
            raise ValueError(f"Invalid duration at CSV row {row_num}: {duration_raw!r}")

        is_pass = result in {"pass", "passed", "true", "1", "ok", "success"}
        is_fail = result in {"fail", "failed", "false", "0", "ng", "error"}
        if not is_pass and not is_fail:
            is_fail = True

        for bucket in (stats[model], overall):
            bucket["total"] += 1
            bucket["pass"] += 1 if is_pass else 0
            bucket["fail"] += 1 if is_fail else 0
            bucket["duration_sum"] += duration

wb = Workbook()
ws = wb.active
ws.title = "By Device Model"
ws.append(["Device Model", "Total", "PassCount", "FailCount", "PassRate", "AvgDurationSeconds"])

for model in sorted(stats):
    s = stats[model]
    total = s["total"]
    ws.append([
        model,
        total,
        s["pass"],
        s["fail"],
        s["pass"] / total if total else 0,
        s["duration_sum"] / total if total else 0,
    ])

ws_overall = wb.create_sheet("Overall")
ws_overall.append(["Total", "PassCount", "FailCount", "PassRate", "AvgDurationSeconds"])
total = overall["total"]
ws_overall.append([
    total,
    overall["pass"],
    overall["fail"],
    overall["pass"] / total if total else 0,
    overall["duration_sum"] / total if total else 0,
])

for sheet in (ws, ws_overall):
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    headers = [cell.value for cell in sheet[1]]
    for row in sheet.iter_rows(min_row=2):
        for idx, header in enumerate(headers):
            if header == "PassRate":
                row[idx].number_format = "0.00%"
            elif header == "AvgDurationSeconds":
                row[idx].number_format = "0.00"
    for col_idx, column_cells in enumerate(sheet.columns, start=1):
        max_len = 0
        for cell in column_cells:
            value = cell.value
            max_len = max(max_len, len(str(value)) if value is not None else 0)
        sheet.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)

wb.save(output_path)

# Verify workbook can be opened and print a concise summary.
check = load_workbook(output_path, read_only=True, data_only=True)
print(f"Read: {input_path}")
print(f"Columns used: model={model_col!r}, result={result_col!r}, duration={duration_col!r}")
print(f"Workbook sheets: {check.sheetnames}")
print(f"Device models summarized: {len(stats)}")
print(f"Overall rows: {overall['total']}, pass={overall['pass']}, fail={overall['fail']}")
print(f"Wrote: {output_path}")