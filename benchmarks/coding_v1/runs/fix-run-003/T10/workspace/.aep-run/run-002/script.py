from pathlib import Path
from urllib.request import urlopen, Request
from urllib.parse import urljoin
from html.parser import HTMLParser
import re
from openpyxl import Workbook, load_workbook

CATALOG_URL = 'http://127.0.0.1:51979/t10/catalog'
OUT = Path('output/products.xlsx')
OUT.parent.mkdir(parents=True, exist_ok=True)

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_catalog = False
        self.catalog_depth = 0
        self.in_card = False
        self.card_depth = 0
        self.current = None
        self.capture = None
        self.items = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if self.in_catalog:
            self.catalog_depth += 1
        if tag == 'main' and attrs.get('id') == 'catalog' and not self.in_catalog:
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if not self.in_catalog:
            return
        classes = attrs.get('class', '').split()
        if tag == 'div' and 'card' in classes and not self.in_card:
            self.in_card = True
            self.card_depth = 1
            self.current = {'name_parts': [], 'price_parts': [], 'href': ''}
            return
        if self.in_card:
            if tag == 'a' and not self.current.get('href'):
                self.current['href'] = attrs.get('href', '')
                self.capture = 'name'
            elif tag == 'span' and 'price' in classes:
                self.capture = 'price'

    def handle_endtag(self, tag):
        if self.in_card:
            if tag in ('a', 'span'):
                self.capture = None
            self.card_depth -= 1
            if self.card_depth == 0:
                name = ' '.join(''.join(self.current['name_parts']).split())
                price_text = ''.join(self.current['price_parts']).strip()
                href = self.current.get('href', '')
                if name and price_text and href:
                    self.items.append((name, price_text, href))
                self.in_card = False
                self.current = None
                self.capture = None
                return
        if self.in_catalog:
            self.catalog_depth -= 1
            if self.catalog_depth == 0:
                self.in_catalog = False

    def handle_data(self, data):
        if self.in_card and self.capture == 'name':
            self.current['name_parts'].append(data)
        elif self.in_card and self.capture == 'price':
            self.current['price_parts'].append(data)

# Fix depth handling for nested tags inside card: increment separately after card starts.
class CatalogParser2(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_catalog=False
        self.catalog_stack=[]
        self.in_card=False
        self.card_depth=0
        self.current=None
        self.capture=None
        self.items=[]
    def handle_starttag(self, tag, attrs):
        attrs=dict(attrs)
        if tag=='main' and attrs.get('id')=='catalog' and not self.in_catalog:
            self.in_catalog=True
            self.catalog_stack=[tag]
            return
        if self.in_catalog:
            self.catalog_stack.append(tag)
            classes=attrs.get('class','').split()
            if tag=='div' and 'card' in classes and not self.in_card:
                self.in_card=True
                self.card_depth=1
                self.current={'name_parts':[], 'price_parts':[], 'href':''}
            elif self.in_card:
                self.card_depth += 1
                if tag=='a' and not self.current['href']:
                    self.current['href']=attrs.get('href','')
                    self.capture='name'
                elif tag=='span' and 'price' in classes:
                    self.capture='price'
    def handle_endtag(self, tag):
        if not self.in_catalog:
            return
        if self.in_card:
            if tag in ('a','span'):
                self.capture=None
            self.card_depth -= 1
            if self.card_depth == 0:
                name=' '.join(''.join(self.current['name_parts']).split())
                price=''.join(self.current['price_parts']).strip()
                href=self.current['href']
                if name and price and href:
                    self.items.append((name,price,href))
                self.in_card=False
                self.current=None
                self.capture=None
        if self.catalog_stack:
            self.catalog_stack.pop()
        if tag=='main' or not self.catalog_stack:
            self.in_catalog=False
            self.catalog_stack=[]
    def handle_data(self,data):
        if self.in_card and self.capture=='name':
            self.current['name_parts'].append(data)
        elif self.in_card and self.capture=='price':
            self.current['price_parts'].append(data)

req = Request(CATALOG_URL, headers={'User-Agent':'Mozilla/5.0'})
with urlopen(req, timeout=10) as r:
    html = r.read().decode('utf-8', errors='replace')

parser = CatalogParser2()
parser.feed(html)

wb = Workbook()
ws = wb.active
ws.title = 'products'
ws.append(['產品名稱', '價格', '幣別', '產品網址'])
for name, price_text, href in parser.items:
    m = re.match(r'\s*([^\d\s,\.]+)\s*([\d,]+(?:\.\d+)?)', price_text)
    if not m:
        raise ValueError(f'Cannot parse price: {price_text!r}')
    currency = m.group(1)
    price_num = float(m.group(2).replace(',', '')) if '.' in m.group(2) else int(m.group(2).replace(',', ''))
    ws.append([name, price_num, currency, urljoin(CATALOG_URL, href)])

wb.save(OUT)
print(f'wrote {OUT} with {len(parser.items)} products')
# Verify workbook values and types
wb2 = load_workbook(OUT, data_only=True)
for row in wb2.active.iter_rows(values_only=True):
    print(row)
print('price cell types:', [type(wb2.active.cell(row=i, column=2).value).__name__ for i in range(2, wb2.active.max_row+1)])
