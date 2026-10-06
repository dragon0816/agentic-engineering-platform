from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_path = workspace / "input" / "mixed.csv"
output_dir = workspace / "output"
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / "converted.xlsx"

if not input_path.exists():
    raise FileNotFoundError(f"Input file not found: {input_path}")

# Read with utf-8-sig to transparently support UTF-8 BOM. newline='' lets csv
# correctly handle quoted fields containing commas and embedded newlines.
with input_path.open("r", encoding="utf-8-sig", newline="") as f:
    rows = list(csv.reader(f, delimiter=",", quotechar='"', doublequote=True))

wb = Workbook()
ws = wb.active
ws.title = "converted"

# Write exactly the CSV rows/columns in order, without adding an index column.
# Empty CSV fields are written as blank Excel cells. Non-empty fields are kept as
# Python strings so leading zeros and Chinese/text/newlines are preserved.
for r_idx, row in enumerate(rows, start=1):
    for c_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx)
        cell.value = None if value == "" else value

# Ensure Device ID column is formatted as text and values remain strings.
device_col = None
if rows:
    for idx, header in enumerate(rows[0], start=1):
        if header == "Device ID":
            device_col = idx
            break

if device_col is not None:
    for row_idx in range(1, len(rows) + 1):
        cell = ws.cell(row=row_idx, column=device_col)
        cell.number_format = "@"
        if cell.value is not None:
            cell.value = str(cell.value)

wb.save(output_path)

# Quick verification readback.
check_wb = load_workbook(output_path, data_only=False)
check_ws = check_wb.active
max_row = check_ws.max_row
max_col = check_ws.max_column
headers = [check_ws.cell(row=1, column=c).value for c in range(1, max_col + 1)] if max_row else []
print(f"Wrote: {output_path}")
print(f"Rows read/written: {len(rows)} / {max_row}")
print(f"Columns: {max_col}")
print(f"Headers: {headers}")
if device_col is not None:
    sample = []
    for r in range(2, min(max_row, 6) + 1):
        cell = check_ws.cell(row=r, column=device_col)
        sample.append((cell.value, cell.data_type, cell.number_format))
    print(f"Device ID column: {device_col}; sample value/type/format: {sample}")
else:
    print("Device ID column not found.")
