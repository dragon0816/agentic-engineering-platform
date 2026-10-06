from pathlib import Path
import csv, sys, runpy
from openpyxl import load_workbook

ws = Path.cwd()
out_py = ws / 'output' / 'export.py'
input_csv = ws / 'input' / 'input.csv'
out_xlsx = ws / 'output' / 'summary.xlsx'

print('Checking repaired program exists:', out_py, out_py.exists())
if not out_py.exists():
    raise SystemExit('output/export.py missing')

# Execute the repaired CLI entry point with the same arguments as requested.
# This is equivalent argument handling to: python output\export.py --input input\input.csv --output output\summary.xlsx
sys.argv = ['output\\export.py', '--input', 'input\\input.csv', '--output', 'output\\summary.xlsx']
print('Executing CLI arguments:', sys.argv)
try:
    runpy.run_path(str(out_py), run_name='__main__')
except SystemExit as e:
    code = e.code if isinstance(e.code, int) else 0
    print('SystemExit code:', e.code)
    if code != 0:
        raise

print('summary.xlsx exists:', out_xlsx.exists(), 'size:', out_xlsx.stat().st_size if out_xlsx.exists() else None)
if not out_xlsx.exists():
    raise SystemExit('summary.xlsx was not created')

# Calculate expected values independently from input.csv.
groups = {}
order = []
with open(input_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        model = (row.get('Model') or '').strip()
        if model not in groups:
            groups[model] = []
            order.append(model)
        raw = row.get('Value')
        if raw is not None and str(raw).strip() != '':
            groups[model].append(float(str(raw).strip()))
expected = [(m, sum(groups[m]) / len(groups[m]) if groups[m] else None, len(groups[m])) for m in order]

wb = load_workbook(out_xlsx, data_only=True)
sh = wb.active
rows = list(sh.iter_rows(values_only=True))
print('Workbook sheet:', sh.title)
print('Workbook rows:')
for row in rows:
    print(row)
print('Expected data rows:', expected)

assert rows[0] == ('Model', 'Average', 'ValidCount'), rows[0]
assert len(rows) - 1 == len(expected), (len(rows) - 1, len(expected))
for actual, exp in zip(rows[1:], expected):
    am, aa, ac = actual
    em, ea, ec = exp
    assert am == em, (actual, exp)
    assert ac == ec, (actual, exp)
    if ea is None:
        assert aa is None, (actual, exp)
    else:
        assert abs(float(aa) - ea) < 1e-9, (actual, exp)
print('Verification passed.')
