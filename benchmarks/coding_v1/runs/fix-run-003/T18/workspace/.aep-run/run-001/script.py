from pathlib import Path
import csv, subprocess, sys, json, textwrap
from openpyxl import load_workbook

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)

orig = input_dir / 'export.py'
deps = input_dir / 'DEPENDENCIES.txt'
input_csv = input_dir / 'input.csv'
out_py = output_dir / 'export.py'
out_xlsx = output_dir / 'summary.xlsx'

print('Workspace:', ws)
print('Input files:', sorted(p.name for p in input_dir.iterdir()))
print('DEPENDENCIES.txt:')
print(deps.read_text(encoding='utf-8', errors='replace') if deps.exists() else '(missing)')
print('Original export.py:')
print(orig.read_text(encoding='utf-8', errors='replace'))
print('input.csv preview:')
print('\n'.join(input_csv.read_text(encoding='utf-8-sig', errors='replace').splitlines()[:10]))

# Write repaired program to output/export.py. It preserves the CLI flags --input and --output.
fixed = r'''#!/usr/bin/env python3
"""Create an Excel summary by Model from an input CSV.

CLI is intentionally kept compatible with the original script:
    python export.py --input input.csv --output summary.xlsx
"""

from __future__ import annotations

import argparse
import csv
from collections import OrderedDict
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook


def _parse_value(raw: object) -> float | None:
    """Return a numeric value, or None for blank/missing cells.

    Blank Value cells are ignored by the average and ValidCount as required.
    Non-blank, non-numeric values are treated as invalid data and raise a clear
    error instead of being silently counted incorrectly.
    """
    if raw is None:
        return None
    text = str(raw).strip()
    if text == "":
        return None
    try:
        return float(text)
    except ValueError as exc:
        raise ValueError(f"Value must be numeric or blank, got {raw!r}") from exc


def build_summary(input_csv: str | Path) -> list[tuple[str, float | None, int]]:
    """Read CSV and return rows: (Model, Average, ValidCount)."""
    groups: "OrderedDict[str, list[float]]" = OrderedDict()
    with open(input_csv, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        if reader.fieldnames is None:
            raise ValueError("Input CSV is missing a header row")
        required = {"Model", "Value"}
        missing = required.difference(reader.fieldnames)
        if missing:
            raise ValueError(f"Input CSV missing required column(s): {', '.join(sorted(missing))}")

        for row_num, row in enumerate(reader, start=2):
            model = (row.get("Model") or "").strip()
            if model == "":
                raise ValueError(f"Row {row_num}: Model is blank")
            groups.setdefault(model, [])
            value = _parse_value(row.get("Value"))
            if value is not None:
                groups[model].append(value)

    summary: list[tuple[str, float | None, int]] = []
    for model, values in groups.items():
        count = len(values)
        average = (sum(values) / count) if count else None
        summary.append((model, average, count))
    return summary


def write_workbook(rows: Iterable[tuple[str, float | None, int]], output_xlsx: str | Path) -> None:
    output_path = Path(output_xlsx)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Model", "Average", "ValidCount"])
    for model, average, count in rows:
        ws.append([model, average, count])
    wb.save(output_path)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Summarize average Value by Model")
    parser.add_argument("--input", required=True, help="Input CSV path")
    parser.add_argument("--output", required=True, help="Output XLSX path")
    args = parser.parse_args(argv)

    rows = build_summary(args.input)
    write_workbook(rows, args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
'''
out_py.write_text(fixed, encoding='utf-8')
print('Wrote repaired program:', out_py)

# Execute exactly the requested command shape from the workspace.
cmd = [sys.executable, str(out_py), '--input', str(input_csv), '--output', str(out_xlsx)]
print('Running:', ' '.join(cmd))
completed = subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('Return code:', completed.returncode)
print('STDOUT:', completed.stdout)
print('STDERR:', completed.stderr)
if completed.returncode != 0:
    raise SystemExit('Repaired export.py failed')
if not out_xlsx.exists():
    raise SystemExit('summary.xlsx was not created')

# Independent expected calculation from CSV.
expected_groups = {}
order = []
with open(input_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        model = (row.get('Model') or '').strip()
        if model not in expected_groups:
            expected_groups[model] = []
            order.append(model)
        raw = row.get('Value')
        if raw is not None and str(raw).strip() != '':
            expected_groups[model].append(float(str(raw).strip()))
expected = []
for model in order:
    vals = expected_groups[model]
    expected.append((model, (sum(vals)/len(vals) if vals else None), len(vals)))
print('Expected:', expected)

wb = load_workbook(out_xlsx, data_only=True)
wsheet = wb.active
actual = []
for r in wsheet.iter_rows(values_only=True):
    actual.append(r)
print('Workbook rows:', actual)

if actual[0] != ('Model', 'Average', 'ValidCount'):
    raise SystemExit(f'Unexpected header: {actual[0]!r}')
actual_data = actual[1:]
if len(actual_data) != len(expected):
    raise SystemExit(f'Expected {len(expected)} data rows, got {len(actual_data)}')
for got, exp in zip(actual_data, expected):
    gm, ga, gc = got
    em, ea, ec = exp
    if gm != em or gc != ec or (ea is None and ga is not None) or (ea is not None and abs(float(ga) - ea) > 1e-9):
        raise SystemExit(f'Mismatch: got {got!r}, expected {exp!r}')
print('Verification passed: output/summary.xlsx has Model, Average, ValidCount and ignores blank Value cells.')
