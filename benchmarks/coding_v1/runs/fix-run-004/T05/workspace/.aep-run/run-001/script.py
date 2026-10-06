from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
input_dir = base / 'input'
print('Workspace:', base)
print('Input files:', [p.name for p in input_dir.iterdir()])
for name in ['devices.xlsx', 'tests.xlsx']:
    path = input_dir / name
    print('\n---', name, '---')
    wb = load_workbook(path, read_only=True, data_only=False)
    print('Sheets:', wb.sheetnames)
    ws = wb[wb.sheetnames[0]]
    print('Dimensions:', ws.max_row, 'rows x', ws.max_column, 'cols')
    rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row, 6), values_only=False))
    for r_idx, row in enumerate(rows, 1):
        vals = []
        for cell in row:
            vals.append({'value': cell.value, 'type': cell.data_type, 'number_format': cell.number_format})
        print('Row', r_idx, vals)
    wb.close()