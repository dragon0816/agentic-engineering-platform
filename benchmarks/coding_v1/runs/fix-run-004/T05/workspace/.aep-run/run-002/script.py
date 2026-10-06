from pathlib import Path
from openpyxl import load_workbook, Workbook

base = Path.cwd()
input_dir = base / 'input'
output_dir = base / 'output'
output_dir.mkdir(exist_ok=True)
out_path = output_dir / 'merged_devices_tests.xlsx'


def load_table(path, sheet_name=None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(v) if v is not None else '' for v in rows[0]]
    records = []
    for row in rows[1:]:
        # keep rows that have any content
        if not any(v is not None for v in row):
            continue
        rec = {}
        for i, header in enumerate(headers):
            val = row[i] if i < len(row) else None
            if header == 'Device ID' and val is not None:
                rec[header] = str(val)
            else:
                rec[header] = val
        records.append(rec)
    wb.close()
    return headers, records

dev_headers, devices = load_table(input_dir / 'devices.xlsx')
test_headers, tests = load_table(input_dir / 'tests.xlsx')

if 'Device ID' not in dev_headers or 'Device ID' not in test_headers:
    raise ValueError('Both input files must contain a Device ID column')

devices_by_id = {}
for rec in devices:
    did = rec.get('Device ID')
    if did in devices_by_id:
        raise ValueError(f'Duplicate Device ID in devices.xlsx: {did}')
    devices_by_id[did] = rec

test_ids_with_match = set()
matched_rows = []
# Matched columns: all test columns, then device columns except duplicate Device ID.
matched_headers = list(test_headers) + [h for h in dev_headers if h != 'Device ID']
for test in tests:
    did = test.get('Device ID')
    device = devices_by_id.get(did)
    if device is not None:
        test_ids_with_match.add(did)
        row = [test.get(h) for h in test_headers] + [device.get(h) for h in dev_headers if h != 'Device ID']
        matched_rows.append(row)

# Unmatched sheet: Source + union of original columns, preserving source column names/order.
unmatched_headers = ['Source'] + list(dev_headers) + [h for h in test_headers if h not in dev_headers]
unmatched_rows = []
matched_device_ids = {test.get('Device ID') for test in tests if test.get('Device ID') in devices_by_id}

for device in devices:
    did = device.get('Device ID')
    if did not in matched_device_ids:
        unmatched_rows.append(['devices'] + [device.get(h) for h in dev_headers] + [None for h in test_headers if h not in dev_headers])

for test in tests:
    did = test.get('Device ID')
    if did not in devices_by_id:
        values = []
        for h in unmatched_headers[1:]:
            if h in test_headers:
                values.append(test.get(h))
            else:
                values.append(None)
        unmatched_rows.append(['tests'] + values)

wb_out = Workbook()
ws_m = wb_out.active
ws_m.title = 'matched'
ws_m.append(matched_headers)
for row in matched_rows:
    ws_m.append(row)
ws_u = wb_out.create_sheet('unmatched')
ws_u.append(unmatched_headers)
for row in unmatched_rows:
    ws_u.append(row)

# Force Device ID cells to text format in both sheets.
for ws in [ws_m, ws_u]:
    for cell in ws[1]:
        if cell.value == 'Device ID':
            col = cell.column
            for row in range(1, ws.max_row + 1):
                c = ws.cell(row=row, column=col)
                c.number_format = '@'
                if row > 1 and c.value is not None:
                    c.value = str(c.value)

wb_out.save(out_path)
print(f'Wrote {out_path}')
print(f'Devices rows: {len(devices)}')
print(f'Tests rows: {len(tests)}')
print(f'Matched data rows: {len(matched_rows)}')
print(f'Unmatched data rows: {len(unmatched_rows)}')
print('Matched headers:', matched_headers)
print('Unmatched headers:', unmatched_headers)
