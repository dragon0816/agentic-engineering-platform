from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
out_path = base / 'output' / 'merged_devices_tests.xlsx'
wb = load_workbook(out_path, data_only=False)
print('Output exists:', out_path.exists())
print('Sheets:', wb.sheetnames)
assert wb.sheetnames == ['matched', 'unmatched'], wb.sheetnames

ws_m = wb['matched']
ws_u = wb['unmatched']
matched_headers = [c.value for c in ws_m[1]]
unmatched_headers = [c.value for c in ws_u[1]]
print('Matched headers:', matched_headers)
print('Unmatched headers:', unmatched_headers)
assert matched_headers == ['Device ID', 'Test', 'Result', 'Name', 'Location']
assert unmatched_headers == ['Source', 'Device ID', 'Name', 'Location', 'Test', 'Result']

matched = [tuple(c.value for c in row) for row in ws_m.iter_rows(min_row=2)]
unmatched = [tuple(c.value for c in row) for row in ws_u.iter_rows(min_row=2)]
print('Matched rows:', matched)
print('Unmatched rows:', unmatched)
expected_matched = [
    ('001', 'RF-1', 'Pass', 'Alpha', 'Lab A'),
    ('001', 'RF-2', 'Fail', 'Alpha', 'Lab A'),
    ('002', 'RF-1', 'Pass', 'Beta', 'Lab B'),
]
expected_unmatched = [
    ('devices', '003', 'Gamma', 'Lab C', None, None),
    ('tests', '999', None, None, 'RF-9', 'Fail'),
]
assert matched == expected_matched, matched
assert unmatched == expected_unmatched, unmatched

# Confirm Device ID is text/string and text-formatted in both sheets; leading zeros preserved.
for ws in [ws_m, ws_u]:
    headers = [c.value for c in ws[1]]
    did_col = headers.index('Device ID') + 1
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, did_col)
        assert isinstance(cell.value, str), (ws.title, r, cell.value, type(cell.value))
        assert cell.number_format == '@', (ws.title, r, cell.number_format)
print('Checked: workbook, sheet names, headers, row counts/content, unmatched Source values, and Device ID text formatting/leading zeros.')
wb.close()
