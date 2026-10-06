from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook

workspace = Path.cwd()
input_path = workspace / 'input' / 'results.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'failed_items.xlsx'

# Read CSV as text values so IDs such as Device ID keep leading zeroes.
with input_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    headers = reader.fieldnames or []
    if 'Status' not in headers:
        raise ValueError('Status column not found')
    failed_rows = [row for row in reader if (row.get('Status') or '').strip().lower() == 'fail']

# Write failed_items workbook/sheet, preserving original column names and row order.
wb = Workbook()
ws = wb.active
ws.title = 'failed_items'
ws.append(headers)
for row in failed_rows:
    values = []
    for h in headers:
        v = row.get(h, '')
        # Ensure all source CSV fields are written as strings (or blank), preserving leading zeroes.
        values.append('' if v is None else str(v))
    ws.append(values)

# Explicitly set Device ID cells to text format if present.
if 'Device ID' in headers:
    device_col = headers.index('Device ID') + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(r, device_col).number_format = '@'
        ws.cell(r, device_col).value = str(ws.cell(r, device_col).value) if ws.cell(r, device_col).value is not None else ''

wb.save(output_path)

# Verification: reopen output and confirm stated requirements.
check_wb = load_workbook(output_path, data_only=False)
if 'failed_items' not in check_wb.sheetnames:
    raise AssertionError('failed_items worksheet missing')
check_ws = check_wb['failed_items']
out_headers = [check_ws.cell(1, c).value for c in range(1, check_ws.max_column + 1)]
if out_headers != headers:
    raise AssertionError(f'Header mismatch: {out_headers!r} != {headers!r}')

out_rows = []
for r in range(2, check_ws.max_row + 1):
    out_rows.append([check_ws.cell(r, c).value for c in range(1, check_ws.max_column + 1)])

expected_rows = [[row.get(h, '') for h in headers] for row in failed_rows]
# openpyxl may read blank cells as None; normalize only for comparison of blanks.
normalized_out = [['' if v is None else str(v) for v in row] for row in out_rows]
normalized_expected = [['' if v is None else str(v) for v in row] for row in expected_rows]
if normalized_out != normalized_expected:
    raise AssertionError(f'Row data/order mismatch: {normalized_out!r} != {normalized_expected!r}')

if 'Device ID' in headers:
    device_col = headers.index('Device ID') + 1
    for r in range(2, check_ws.max_row + 1):
        cell = check_ws.cell(r, device_col)
        if cell.value is not None and not isinstance(cell.value, str):
            raise AssertionError(f'Device ID row {r} is not text: {cell.value!r} ({type(cell.value)})')
        if cell.number_format != '@':
            raise AssertionError(f'Device ID row {r} not formatted as text')

# Confirm every included row matches trimmed/case-insensitive fail and count equals source filter.
status_col = headers.index('Status') + 1
for r in range(2, check_ws.max_row + 1):
    status_value = check_ws.cell(r, status_col).value or ''
    if str(status_value).strip().lower() != 'fail':
        raise AssertionError(f'Non-fail row included at output row {r}: {status_value!r}')

print('created:', output_path)
print('worksheet:', check_ws.title)
print('headers:', out_headers)
print('failed row count:', len(out_rows))
print('output rows:', normalized_out)
print('checks passed: output file exists, worksheet name is failed_items, all original columns preserved, row order preserved, Status filter is trim/case-insensitive fail, Device ID values are text with leading zeroes preserved.')