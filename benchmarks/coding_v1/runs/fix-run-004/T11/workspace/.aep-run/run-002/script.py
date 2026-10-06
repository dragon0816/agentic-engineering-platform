import os, re, html, json, urllib.request, urllib.parse
from pathlib import Path
from html.parser import HTMLParser
from openpyxl import Workbook, load_workbook

BASE_LIST_URL = 'http://127.0.0.1:58649/t11/list'
WORKSPACE = Path.cwd()
OUT = WORKSPACE / 'output' / 'details.xlsx'
OUT.parent.mkdir(parents=True, exist_ok=True)

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._in_a = False
        self._href = None
        self._text = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == 'a':
            attrs = dict(attrs)
            self._in_a = True
            self._href = attrs.get('href')
            self._text = []
    def handle_data(self, data):
        if self._in_a:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == 'a' and self._in_a:
            if self._href:
                self.links.append((self._href, ' '.join(''.join(self._text).split())))
            self._in_a = False
            self._href = None
            self._text = []

def fetch(url):
    with urllib.request.urlopen(url, timeout=10) as r:
        return r.read().decode('utf-8', errors='replace')

def clean_text(s):
    s = html.unescape(re.sub(r'<[^>]+>', ' ', s or ''))
    return ' '.join(s.split())

def parse_cells(row_html):
    return [clean_text(c) for c in re.findall(r'<(?:th|td)\b[^>]*>(.*?)</(?:th|td)>', row_html, flags=re.I|re.S)]

def parse_keyvals(doc):
    kv = {}
    # table rows
    for row in re.findall(r'<tr\b[^>]*>(.*?)</tr>', doc, flags=re.I|re.S):
        cells = parse_cells(row)
        if len(cells) >= 2:
            key = cells[0].strip().rstrip(':：').lower()
            val = cells[1].strip()
            kv[key] = val
    # definition lists
    for m in re.finditer(r'<dt\b[^>]*>(.*?)</dt>\s*<dd\b[^>]*>(.*?)</dd>', doc, flags=re.I|re.S):
        key = clean_text(m.group(1)).rstrip(':：').lower()
        kv[key] = clean_text(m.group(2))
    # generic text labels
    text_doc = re.sub(r'</(?:p|div|li|tr|h\d|dt|dd)>', '\n', doc, flags=re.I)
    text = html.unescape(re.sub(r'<[^>]+>', ' ', text_doc))
    lines = [' '.join(line.split()) for line in text.splitlines() if line.strip()]
    labels = ['name', 'price', 'currency', 'cpu', 'memory', 'ram']
    for line in lines:
        m = re.match(r'^(Name|Price|Currency|CPU|Memory|RAM)\s*[:：]\s*(.+)$', line, flags=re.I)
        if m:
            kv[m.group(1).lower()] = m.group(2).strip()
    return kv, lines

def first_match(pattern, doc, flags=re.I|re.S):
    m = re.search(pattern, doc, flags)
    return clean_text(m.group(1)) if m else ''

def json_ld_values(doc):
    vals = {}
    for script in re.findall(r'<script\b[^>]*type=["\']application/ld\+json["\'][^>]*>(.*?)</script>', doc, flags=re.I|re.S):
        try:
            obj = json.loads(html.unescape(script.strip()))
        except Exception:
            continue
        stack = obj if isinstance(obj, list) else [obj]
        while stack:
            item = stack.pop(0)
            if isinstance(item, dict):
                if 'name' in item and not vals.get('name'):
                    vals['name'] = str(item['name'])
                offers = item.get('offers')
                if isinstance(offers, dict):
                    if 'price' in offers: vals.setdefault('price', str(offers['price']))
                    if 'priceCurrency' in offers: vals.setdefault('currency', str(offers['priceCurrency']))
                for prop in item.get('additionalProperty', []) if isinstance(item.get('additionalProperty'), list) else []:
                    if isinstance(prop, dict):
                        n = str(prop.get('name','')).lower()
                        v = str(prop.get('value',''))
                        if 'cpu' in n: vals.setdefault('cpu', v)
                        if 'memory' in n or 'ram' in n: vals.setdefault('memory', v)
                for v in item.values():
                    if isinstance(v, (dict, list)):
                        stack.append(v)
            elif isinstance(item, list):
                stack.extend(item)
    return vals

def normalize_price_currency(price_text, currency_text):
    price_text = (price_text or '').strip()
    currency_text = (currency_text or '').strip()
    currency_symbols = {'$':'USD', '€':'EUR', '£':'GBP', '¥':'JPY', 'NT$':'TWD'}
    if not currency_text:
        for sym, cur in currency_symbols.items():
            if sym in price_text:
                currency_text = cur
                break
    mcur = re.search(r'\b(USD|EUR|GBP|JPY|TWD|NTD|CNY|RMB)\b', price_text, flags=re.I)
    if mcur and not currency_text:
        currency_text = mcur.group(1).upper()
    num = ''
    m = re.search(r'[-+]?\d[\d,]*(?:\.\d+)?', price_text)
    if m:
        raw = m.group(0).replace(',', '')
        try:
            val = float(raw)
            num = int(val) if val.is_integer() else val
        except Exception:
            num = raw
    return num, currency_text.upper() if currency_text else ''

def memory_to_gb(mem):
    mem = (mem or '').strip()
    if not mem:
        return ''
    m = re.search(r'(\d+(?:\.\d+)?)\s*(TB|T|GB|G|MB|M)\b', mem, flags=re.I)
    if not m:
        m = re.search(r'(\d+(?:\.\d+)?)', mem)
        if not m:
            return ''
        return float(m.group(1)) if '.' in m.group(1) else int(m.group(1))
    val = float(m.group(1))
    unit = m.group(2).upper()
    if unit in ('TB','T'):
        gb = val * 1024
    elif unit in ('MB','M'):
        gb = val / 1024
    else:
        gb = val
    return int(gb) if abs(gb - round(gb)) < 1e-9 else gb

def parse_detail(doc, url, link_text=''):
    kv, lines = parse_keyvals(doc)
    jv = json_ld_values(doc)
    name = jv.get('name') or kv.get('name') or first_match(r'<h1\b[^>]*>(.*?)</h1>', doc) or link_text
    price_text = jv.get('price') or kv.get('price') or first_match(r'(?:class|id)=["\'][^"\']*price[^"\']*["\'][^>]*>(.*?)<', doc)
    currency = jv.get('currency') or kv.get('currency')
    cpu = jv.get('cpu') or kv.get('cpu') or ''
    mem_text = jv.get('memory') or kv.get('memory') or kv.get('ram') or ''
    price, currency = normalize_price_currency(price_text, currency)
    return {
        'Name': name,
        'Price': price,
        'Currency': currency,
        'CPU': cpu,
        'MemoryGB': memory_to_gb(mem_text),
        'URL': url,
        '_debug_kv': kv,
        '_debug_lines': lines[:8],
        '_debug_mem': mem_text,
        '_debug_price': price_text,
    }

print('Workspace:', WORKSPACE)
print('Input files:', [p.name for p in (WORKSPACE/'input').glob('*')])
list_html = fetch(BASE_LIST_URL)
print('List page snippet:', list_html[:500].replace('\n',' '))
parser = LinkParser(); parser.feed(list_html)
links = []
seen = set()
for href, text in parser.links:
    abs_url = urllib.parse.urljoin(BASE_LIST_URL, href)
    if abs_url not in seen:
        links.append((abs_url, text)); seen.add(abs_url)
print('Found detail links:', links)

rows = []
for url, link_text in links:
    doc = fetch(url)
    print('\nDETAIL', url, 'snippet:', doc[:800].replace('\n',' '))
    rec = parse_detail(doc, url, link_text)
    print('Parsed:', {k:v for k,v in rec.items() if not k.startswith('_')}, 'debug_kv=', rec['_debug_kv'])
    rows.append({k: rec[k] for k in ['Name','Price','Currency','CPU','MemoryGB','URL']})

wb = Workbook()
ws = wb.active
ws.title = 'details'
headers = ['Name','Price','Currency','CPU','MemoryGB','URL']
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(OUT)
print('Wrote', OUT)

# Verification: reopen output and check requirements.
wb2 = load_workbook(OUT, data_only=True)
ws2 = wb2.active
actual_headers = [ws2.cell(row=1, column=i).value for i in range(1, len(headers)+1)]
assert actual_headers == headers, f'headers mismatch: {actual_headers}'
assert ws2.max_row == len(rows) + 1, f'row count mismatch: {ws2.max_row} vs {len(rows)+1}'
for r in range(2, ws2.max_row + 1):
    vals = {headers[c-1]: ws2.cell(row=r, column=c).value for c in range(1, len(headers)+1)}
    assert vals['URL'] and str(vals['URL']).startswith('http://127.0.0.1:58649/t11/item/'), f'bad URL row {r}: {vals}'
    assert vals['Name'] is None or isinstance(vals['Name'], str), f'Name not text row {r}'
    assert vals['Currency'] is None or isinstance(vals['Currency'], str), f'Currency not text row {r}'
    assert vals['CPU'] is None or isinstance(vals['CPU'], str), f'CPU not text row {r}'
    assert vals['Price'] is None or isinstance(vals['Price'], (int, float)), f'Price not numeric/blank row {r}: {vals["Price"]}'
    assert vals['MemoryGB'] is None or isinstance(vals['MemoryGB'], (int, float)), f'MemoryGB not numeric/blank row {r}: {vals["MemoryGB"]}'
print('Verified output: details.xlsx exists, headers match, row count =', ws2.max_row-1, ', URL/name/text columns and numeric Price/MemoryGB or blanks are valid.')