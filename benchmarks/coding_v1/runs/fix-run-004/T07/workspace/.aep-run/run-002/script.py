from pathlib import Path
import re
from openpyxl import Workbook, load_workbook

ws = Path.cwd()
input_path = ws / 'input' / 'specs.md'
out_dir = ws / 'output'
out_dir.mkdir(exist_ok=True)
out_path = out_dir / 'specs.xlsx'

text = input_path.read_text(encoding='utf-8')
print('Read', input_path, 'chars', len(text))
print('Input files:', [p.name for p in (ws/'input').iterdir()])
print('Sample content:', text[:300].replace('\n', ' | '))

# Split product sections by markdown H2 headings.
sections = []
current_title = None
current_lines = []
for line in text.splitlines():
    m = re.match(r'^##\s+(.+?)\s*$', line)
    if m:
        if current_title is not None:
            sections.append((current_title, '\n'.join(current_lines)))
        current_title = m.group(1).strip()
        current_lines = []
    elif current_title is not None:
        current_lines.append(line)
if current_title is not None:
    sections.append((current_title, '\n'.join(current_lines)))
print('Sections found:', [t for t, _ in sections])

def parse_fields(body):
    fields = {}
    for line in body.splitlines():
        m = re.match(r'^\s*-\s*([^:]+):\s*(.*?)\s*$', line)
        if m:
            fields[m.group(1).strip().lower()] = m.group(2).strip()
    return fields

def to_mhz(value, unit):
    if value is None or unit is None:
        return None
    v = float(value)
    u = unit.lower()
    if u == 'ghz':
        return v * 1000.0
    if u == 'mhz':
        return v
    return None

def clean_number(n):
    if n is None:
        return None
    if abs(n - round(n)) < 1e-9:
        return int(round(n))
    return n

def parse_frequency(s):
    if not s:
        return None, None
    # Accept ranges written with hyphen/en dash/em dash and a single trailing unit.
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*[\-–—]\s*([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\b', s, re.I)
    if not m:
        return None, None
    mn = to_mhz(m.group(1), m.group(3))
    mx = to_mhz(m.group(2), m.group(3))
    return clean_number(mn), clean_number(mx)

def parse_bandwidth(s):
    if not s:
        return None
    m = re.search(r'([0-9]+(?:\.[0-9]+)?)\s*(GHz|MHz)\b', s, re.I)
    if not m:
        return None
    return clean_number(to_mhz(m.group(1), m.group(2)))

def parse_interface(s):
    if not s:
        return None
    # Multiple interfaces in source are slash/comma separated; output semicolon separated.
    parts = [p.strip() for p in re.split(r'[/,;]+', s) if p.strip()]
    return ';'.join(parts) if parts else None

rows = []
for title, body in sections:
    fields = parse_fields(body)
    fmin, fmax = parse_frequency(fields.get('frequency'))
    rows.append({
        'Model': fields.get('model') or None,
        'FrequencyMinMHz': fmin,
        'FrequencyMaxMHz': fmax,
        'BandwidthMHz': parse_bandwidth(fields.get('bandwidth')),
        'Interface': parse_interface(fields.get('interface')),
    })
print('Parsed rows:', rows)

headers = ['Model', 'FrequencyMinMHz', 'FrequencyMaxMHz', 'BandwidthMHz', 'Interface']
wb = Workbook()
wsheet = wb.active
wsheet.title = 'specs'
wsheet.append(headers)
for row in rows:
    wsheet.append([row[h] for h in headers])
wb.save(out_path)
print('Wrote', out_path)

# Verification: reopen the generated workbook and check every stated requirement.
wb2 = load_workbook(out_path, data_only=True)
assert 'specs' in wb2.sheetnames, wb2.sheetnames
sh = wb2['specs']
read_headers = [cell.value for cell in sh[1]]
assert read_headers == headers, read_headers
read_rows = []
for r in sh.iter_rows(min_row=2, values_only=True):
    read_rows.append(dict(zip(headers, r)))
assert len(read_rows) == 3, len(read_rows)
expected = [
    {'Model': 'RX-100', 'FrequencyMinMHz': 2400, 'FrequencyMaxMHz': 2500, 'BandwidthMHz': 80, 'Interface': 'USB;LAN'},
    {'Model': 'RX-200', 'FrequencyMinMHz': 900, 'FrequencyMaxMHz': 930, 'BandwidthMHz': 20, 'Interface': 'LAN'},
    {'Model': 'RX-300', 'FrequencyMinMHz': 5000, 'FrequencyMaxMHz': 6000, 'BandwidthMHz': None, 'Interface': 'USB'},
]
assert read_rows == expected, read_rows
# Check numeric columns are numeric when present and blanks are None.
for row in read_rows:
    assert isinstance(row['Model'], str)
    assert isinstance(row['FrequencyMinMHz'], (int, float))
    assert isinstance(row['FrequencyMaxMHz'], (int, float))
    assert row['BandwidthMHz'] is None or isinstance(row['BandwidthMHz'], (int, float))
    assert row['Interface'] is None or isinstance(row['Interface'], str)
assert read_rows[2]['BandwidthMHz'] is None
print('Verification passed: specs.xlsx exists, sheet specs, headers match, 3 rows, MHz conversions correct, interfaces semicolon-separated, missing bandwidth is blank, numeric fields are numeric.')
