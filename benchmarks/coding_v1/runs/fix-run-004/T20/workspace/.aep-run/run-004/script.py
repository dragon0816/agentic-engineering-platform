from pathlib import Path
import csv
import json
import os
from openpyxl import Workbook, load_workbook

workspace = Path('.').resolve()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = {}
if params_path and Path(params_path).exists():
    with open(params_path, 'r', encoding='utf-8') as f:
        params = json.load(f)
threshold = int(params.get('threshold', 20))

manifest_path = input_dir / 'manifest.json'
if not manifest_path.exists():
    raise FileNotFoundError(f'Missing manifest: {manifest_path}')
with manifest_path.open('r', encoding='utf-8') as f:
    manifest = json.load(f)
rel_csv = manifest.get('results.csv')
if not rel_csv:
    raise ValueError('manifest.json must map "results.csv" to the CSV path')
csv_path = workspace / rel_csv
if not csv_path.exists():
    raise FileNotFoundError(f'Missing CSV: {csv_path}')
try:
    csv_path.relative_to(input_dir)
except ValueError:
    raise ValueError(f'CSV must be under input/: {csv_path}')

print('Workspace:', workspace)
print('SOP:', input_dir / 'supplied_sop.md')
print('CSV:', csv_path)
print('Threshold rule: Urgent is Value >', threshold)

rows = []
with csv_path.open('r', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    expected = ['Device ID', 'Status', 'Value']
    if reader.fieldnames != expected:
        raise ValueError(f'CSV columns must be {expected}, got {reader.fieldnames}')
    for line_no, row in enumerate(reader, start=2):
        try:
            value = float(row['Value'])
        except Exception as exc:
            raise ValueError(f'Line {line_no}: Value is not numeric: {row["Value"]!r}') from exc
        if value.is_integer():
            value = int(value)
        normalized = {
            'Device ID': str(row['Device ID']),
            'Status': str(row['Status']),
            'Value': value,
        }
        rows.append(normalized)

failures = []
for row in rows:
    if row['Status'] == 'Fail':
        out = dict(row)
        out['Urgent'] = bool(row['Value'] > threshold)
        failures.append(out)

summary = {
    'Total': len(rows),
    'Pass': sum(1 for r in rows if r['Status'] == 'Pass'),
    'Fail': sum(1 for r in rows if r['Status'] == 'Fail'),
    'Urgent': sum(1 for r in failures if r['Urgent']),
}

xlsx_path = output_dir / 'failures.xlsx'
wb = Workbook()
ws = wb.active
ws.title = 'Failures'
headers = ['Device ID', 'Status', 'Value', 'Urgent']
ws.append(headers)
for row in failures:
    ws.append([row['Device ID'], row['Status'], row['Value'], row['Urgent']])
wb.save(xlsx_path)

summary_path = output_dir / 'summary.md'
summary_text = '\n'.join([
    '# Test Failure Summary',
    '',
    f'- Total: {summary["Total"]}',
    f'- Pass: {summary["Pass"]}',
    f'- Fail: {summary["Fail"]}',
    f'- Urgent: {summary["Urgent"]}',
    '',
])
summary_path.write_text(summary_text, encoding='utf-8')

# Validation: reopen all outputs and confirm stated requirements.
if not xlsx_path.exists():
    raise AssertionError('failures.xlsx was not created')
if not summary_path.exists():
    raise AssertionError('summary.md was not created')

check_wb = load_workbook(xlsx_path, data_only=True)
check_ws = check_wb['Failures']
actual_headers = [cell.value for cell in check_ws[1]]
if actual_headers != headers:
    raise AssertionError(f'failure.xlsx headers mismatch: {actual_headers}')
actual_rows = []
for row in check_ws.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    actual_rows.append(row)
if len(actual_rows) != summary['Fail']:
    raise AssertionError(f'Expected {summary["Fail"]} failure rows, got {len(actual_rows)}')
for device_id, status, value, urgent in actual_rows:
    if not isinstance(device_id, str):
        raise AssertionError(f'Device ID must be text, got {device_id!r}')
    if status != 'Fail':
        raise AssertionError(f'Non-Fail row in failures.xlsx: {status!r}')
    if not isinstance(value, (int, float)):
        raise AssertionError(f'Value must be numeric, got {value!r}')
    if not isinstance(urgent, bool):
        raise AssertionError(f'Urgent must be boolean, got {urgent!r}')
    if urgent != (value > threshold):
        raise AssertionError(f'Urgent rule mismatch for {device_id}: value={value}, urgent={urgent}')

summary_check = summary_path.read_text(encoding='utf-8')
for key, val in summary.items():
    needle = f'- {key}: {val}'
    if needle not in summary_check:
        raise AssertionError(f'summary.md missing {needle!r}')

print('Created:', xlsx_path)
print('Created:', summary_path)
print('Validation passed:')
print(' - failures.xlsx exists with columns Device ID, Status, Value, Urgent')
print(' - all failure rows retained:', len(actual_rows))
print(' - Status is Fail for every output row')
print(' - Value cells are numeric')
print(' - Urgent cells are boolean and follow Value > threshold')
print(' - summary.md counts match:', summary)