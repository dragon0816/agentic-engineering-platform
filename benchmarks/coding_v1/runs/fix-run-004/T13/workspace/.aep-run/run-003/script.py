from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urljoin
import urllib.request, urllib.error
import json, re, time
from openpyxl import Workbook, load_workbook

BASE_LIST_URL = 'http://127.0.0.1:58649/t13/list'
OUT_DIR = Path('output')
OUT_DIR.mkdir(exist_ok=True)
XLSX_PATH = OUT_DIR / 'robust_products.xlsx'
WARN_PATH = OUT_DIR / 'warnings.json'

class ProductListParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_products = False
        self.in_a = False
        self.current_href = None
        self.current_text = []
        self.items = []
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'ul' and attrs.get('id') == 'products':
            self.in_products = True
        elif self.in_products and tag == 'a':
            self.in_a = True
            self.current_href = attrs.get('href')
            self.current_text = []
    def handle_data(self, data):
        if self.in_a:
            self.current_text.append(data)
    def handle_endtag(self, tag):
        if tag == 'a' and self.in_a:
            name = ''.join(self.current_text).strip()
            if self.current_href:
                self.items.append({'name': name, 'url': urljoin(BASE_LIST_URL, self.current_href)})
            self.in_a = False
            self.current_href = None
            self.current_text = []
        elif tag == 'ul' and self.in_products:
            self.in_products = False

class DetailParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_name = False
        self.name_parts = []
        self.in_th = False
        self.in_td = False
        self.current_th = []
        self.current_td = []
        self.pending_key = None
        self.spec = {}
    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'h1' and 'name' in attrs.get('class', '').split():
            self.in_name = True
            self.name_parts = []
        elif tag == 'th':
            self.in_th = True
            self.current_th = []
        elif tag == 'td':
            self.in_td = True
            self.current_td = []
    def handle_data(self, data):
        if self.in_name:
            self.name_parts.append(data)
        if self.in_th:
            self.current_th.append(data)
        if self.in_td:
            self.current_td.append(data)
    def handle_endtag(self, tag):
        if tag == 'h1' and self.in_name:
            self.in_name = False
        elif tag == 'th' and self.in_th:
            self.pending_key = ''.join(self.current_th).strip()
            self.in_th = False
        elif tag == 'td' and self.in_td:
            val = ''.join(self.current_td).strip()
            if self.pending_key:
                self.spec[self.pending_key] = val
            self.pending_key = None
            self.in_td = False
    @property
    def name(self):
        return ''.join(self.name_parts).strip()

def fetch_url(url, max_attempts=1):
    last_error = None
    for attempt in range(1, max_attempts + 1):
        try:
            with urllib.request.urlopen(url, timeout=10) as r:
                return {'ok': True, 'status': getattr(r, 'status', 200), 'text': r.read().decode('utf-8', 'replace'), 'attempts': attempt, 'error': None}
        except urllib.error.HTTPError as e:
            body = e.read().decode('utf-8', 'replace')[:300]
            last_error = f'HTTP {e.code}: {body}'
            if e.code == 500 and attempt < max_attempts:
                time.sleep(0.1)
                continue
            return {'ok': False, 'status': e.code, 'text': '', 'attempts': attempt, 'error': last_error}
        except Exception as e:
            last_error = f'{type(e).__name__}: {e}'
            return {'ok': False, 'status': None, 'text': '', 'attempts': attempt, 'error': last_error}
    return {'ok': False, 'status': None, 'text': '', 'attempts': max_attempts, 'error': last_error or 'unknown error'}

def parse_price(raw):
    if not raw:
        return 'N/A', 'N/A'
    text = raw.strip()
    m = re.match(r'^([^\d\-+.,]*?)\s*([\d,]+(?:\.\d+)?)', text)
    if not m:
        return text, 'N/A'
    currency = m.group(1).strip() or 'N/A'
    num_text = m.group(2).replace(',', '')
    try:
        price = float(num_text) if '.' in num_text else int(num_text)
    except ValueError:
        price = num_text
    return price, currency

# Collect list
list_result = fetch_url(BASE_LIST_URL, max_attempts=1)
if not list_result['ok']:
    raise RuntimeError(f'Could not fetch product list: {list_result["error"]}')
list_parser = ProductListParser()
list_parser.feed(list_result['text'])
products = list_parser.items[:5]
print('Collected product links:', products)
if len(products) != 5:
    raise RuntimeError(f'Expected 5 products, found {len(products)}')

rows = []
warnings = []
for product in products:
    detail = fetch_url(product['url'], max_attempts=3)
    entry = {
        'name': product['name'],
        'url': product['url'],
        'status': 'ok' if detail['ok'] else 'error',
        'http_status': detail['status'],
        'attempts': detail['attempts'],
        'warnings': []
    }
    if detail['ok']:
        parser = DetailParser()
        parser.feed(detail['text'])
        name = parser.name or product['name'] or 'N/A'
        price, currency = parse_price(parser.spec.get('Price', ''))
        cpu = parser.spec.get('CPU', '').strip() or 'N/A'
        if price == 'N/A':
            entry['warnings'].append('missing_price')
        if cpu == 'N/A':
            entry['warnings'].append('missing_cpu')
        rows.append([name, price, currency, cpu, product['url'], 'ok'])
    else:
        entry['warnings'].append(detail['error'])
        rows.append([product['name'] or 'N/A', 'N/A', 'N/A', 'N/A', product['url'], 'error'])
    warnings.append(entry)

# Write workbook
wb = Workbook()
ws = wb.active
ws.title = 'Products'
headers = ['Name', 'Price', 'Currency', 'CPU', 'URL', 'FetchStatus']
ws.append(headers)
for row in rows:
    ws.append(row)
wb.save(XLSX_PATH)

# Write warnings JSON
with WARN_PATH.open('w', encoding='utf-8') as f:
    json.dump(warnings, f, ensure_ascii=False, indent=2)

# Verification: reopen outputs and confirm requirements
wb2 = load_workbook(XLSX_PATH, data_only=True)
ws2 = wb2['Products']
read_headers = [c.value for c in ws2[1]]
read_rows = [list(r) for r in ws2.iter_rows(min_row=2, values_only=True)]
with WARN_PATH.open('r', encoding='utf-8') as f:
    warn_data = json.load(f)

checks = []
checks.append(('xlsx exists', XLSX_PATH.exists()))
checks.append(('warnings exists', WARN_PATH.exists()))
checks.append(('headers exact', read_headers == headers))
checks.append(('five data rows', len(read_rows) == 5))
checks.append(('all names and urls present', all(r[0] not in (None, '') and r[4] not in (None, '') for r in read_rows)))
checks.append(('missing fields use N/A', all(all(v is not None for v in r) for r in read_rows)))
checks.append(('fetch status only ok/error', all(r[5] in ('ok', 'error') for r in read_rows)))
checks.append(('successful pages status ok', all(w['status'] == 'ok' for w in warn_data if w.get('http_status') == 200)))
checks.append(('error page retried 3', any(w['status'] == 'error' and w['attempts'] == 3 for w in warn_data)))
checks.append(('warnings has five entries', isinstance(warn_data, list) and len(warn_data) == 5))
# Expected observed details from source inspection
checks.append(('Phone C missing price/currency N/A', any(r[0] == 'Phone C' and r[1] == 'N/A' and r[2] == 'N/A' and r[3] != 'N/A' and r[5] == 'ok' for r in read_rows)))
checks.append(('Phone D missing cpu N/A', any(r[0] == 'Phone D' and r[1] != 'N/A' and r[3] == 'N/A' and r[5] == 'ok' for r in read_rows)))
checks.append(('Phone E error preserved', any(r[0] == 'Phone E' and r[1] == 'N/A' and r[2] == 'N/A' and r[3] == 'N/A' and r[5] == 'error' for r in read_rows)))

failed = [name for name, ok in checks if not ok]
print('Rows written:')
for r in read_rows:
    print(r)
print('Warnings JSON:')
print(json.dumps(warn_data, ensure_ascii=False, indent=2))
print('Checks:')
for name, ok in checks:
    print(f'- {name}: {ok}')
if failed:
    raise AssertionError('Failed checks: ' + ', '.join(failed))
print('Output files:', XLSX_PATH, WARN_PATH)