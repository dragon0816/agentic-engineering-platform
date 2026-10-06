from pathlib import Path
from html.parser import HTMLParser
from urllib.request import urlopen
from urllib.parse import urljoin
import re
from openpyxl import Workbook, load_workbook

WORKSPACE = Path.cwd()
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
CATALOG_URL = 'http://127.0.0.1:58649/t10/catalog'
OUTFILE = OUTPUT / 'products.xlsx'

class CatalogParser(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.in_catalog = False
        self.catalog_depth = 0
        self.in_card = False
        self.card_depth = 0
        self.current = None
        self.capture = None
        self.rows = []

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'main' and attrs.get('id') == 'catalog' and not self.in_catalog:
            self.in_catalog = True
            self.catalog_depth = 1
            return
        if not self.in_catalog:
            return
        self.catalog_depth += 1
        classes = attrs.get('class', '').split()
        if tag == 'div' and 'card' in classes and not self.in_card:
            self.in_card = True
            self.card_depth = 1
            self.current = {'產品名稱': '', '價格': None, '幣別': '', '產品網址': ''}
            return
        if self.in_card:
            self.card_depth += 1
            if tag == 'a' and not self.current.get('產品網址'):
                self.current['產品網址'] = urljoin(CATALOG_URL, attrs.get('href', ''))
                self.capture = '產品名稱'
            elif tag == 'span' and 'price' in classes:
                self.capture = 'price_text'

    def handle_endtag(self, tag):
        if self.in_card:
            if self.capture == '產品名稱' and tag == 'a':
                self.capture = None
            elif self.capture == 'price_text' and tag == 'span':
                self.capture = None
            self.card_depth -= 1
            if self.card_depth == 0:
                if self.current and self.current.get('產品名稱') and self.current.get('價格') is not None:
                    self.rows.append(self.current)
                self.current = None
                self.in_card = False
                return
        if self.in_catalog:
            self.catalog_depth -= 1
            if self.catalog_depth == 0:
                self.in_catalog = False

    def handle_data(self, data):
        if not (self.in_catalog and self.in_card and self.capture and self.current is not None):
            return
        text = data.strip()
        if not text:
            return
        if self.capture == '產品名稱':
            self.current['產品名稱'] = (self.current['產品名稱'] + ' ' + text).strip()
        elif self.capture == 'price_text':
            # Examples: NT$1,299. Keep currency separate and price numeric.
            m = re.match(r'([^0-9\-\.]+)?\s*([0-9,]+(?:\.\d+)?)', text)
            if not m:
                raise ValueError(f'Cannot parse price: {text!r}')
            self.current['幣別'] = (m.group(1) or '').strip()
            num_text = m.group(2).replace(',', '')
            value = float(num_text) if '.' in num_text else int(num_text)
            self.current['價格'] = value

with urlopen(CATALOG_URL, timeout=10) as resp:
    html = resp.read().decode('utf-8', errors='replace')

parser = CatalogParser()
parser.feed(html)
rows = parser.rows

headers = ['產品名稱', '價格', '幣別', '產品網址']
wb = Workbook()
ws = wb.active
ws.title = 'products'
ws.append(headers)
for row in rows:
    ws.append([row[h] for h in headers])
wb.save(OUTFILE)

# Verification: reopen products.xlsx and confirm requirements.
wb2 = load_workbook(OUTFILE, data_only=True)
ws2 = wb2['products']
read_headers = [cell.value for cell in ws2[1]]
read_rows = list(ws2.iter_rows(min_row=2, values_only=True))
expected = [
    ('Router A', 1299, 'NT$', 'http://127.0.0.1:58649/p/router-a'),
    ('Switch B', 2450, 'NT$', 'http://127.0.0.1:58649/p/switch-b'),
    ('Cable C', 99, 'NT$', 'http://127.0.0.1:58649/p/cable-c'),
    ('Adapter D', 560, 'NT$', 'http://127.0.0.1:58649/p/adapter-d'),
    ('Antenna E', 12000, 'NT$', 'http://127.0.0.1:58649/p/antenna-e'),
]
assert read_headers == headers, f'headers mismatch: {read_headers}'
assert read_rows == expected, f'rows mismatch: {read_rows}'
for r in read_rows:
    assert isinstance(r[1], (int, float)), f'price not numeric: {r}'
    assert str(r[3]).startswith('http://127.0.0.1:58649/'), f'not full URL: {r[3]}'
assert all('Sponsored' not in r[0] and '推薦' not in r[0] for r in read_rows), 'ad/recommendation row found'
print(f'Wrote {OUTFILE}')
print('Checked: file exists, sheet name products, headers, 5 catalog rows in page order, numeric price values, NT$ currency, full URLs, and no ad/recommendation rows.')
for r in read_rows:
    print(r)
