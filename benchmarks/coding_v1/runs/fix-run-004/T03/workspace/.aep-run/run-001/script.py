from pathlib import Path
from openpyxl import load_workbook

base = Path('.')
print('Workspace:', base.resolve())
print('Input files:', [p.name for p in (base/'input').iterdir()])
path = base/'input'/'tasks.xlsx'
print('tasks.xlsx exists:', path.exists())
wb = load_workbook(path, data_only=False)
print('Sheets:', wb.sheetnames)
ws = wb['tasks']
print('Max rows/cols:', ws.max_row, ws.max_column)
headers = [ws.cell(1,c).value for c in range(1, ws.max_column+1)]
print('Headers:', headers)
for r in range(2, min(ws.max_row, 7)+1):
    print('Row', r, [ws.cell(r,c).value for c in range(1, ws.max_column+1)])