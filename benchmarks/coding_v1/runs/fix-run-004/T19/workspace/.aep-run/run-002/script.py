from pathlib import Path
import csv
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment

workspace = Path.cwd()
input_path = workspace / 'input' / 'mixed.csv'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)
output_path = output_dir / 'converted.xlsx'

# Read with utf-8-sig to support UTF-8 BOM. newline='' lets csv correctly handle quoted newlines.
with open(input_path, 'r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.reader(f))

if not rows:
    raise ValueError('mixed.csv is empty')

headers = rows[0]
try:
    device_id_col = headers.index('Device ID') + 1  # openpyxl is 1-based
except ValueError:
    device_id_col = None

wb = Workbook()
ws = wb.active
ws.title = 'converted'

for r_idx, row in enumerate(rows, start=1):
    # csv.reader preserves order and duplicates; write exactly the parsed fields.
    for c_idx, value in enumerate(row, start=1):
        cell = ws.cell(row=r_idx, column=c_idx)
        # Empty CSV fields are left as truly blank Excel cells.
        if value == '':
            cell.value = None
        else:
            cell.value = value
        # Device ID must always be text and preserve leading zeroes.
        if device_id_col is not None and c_idx == device_id_col:
            cell.number_format = '@'
            if value != '':
                cell.value = str(value)
        # Keep embedded newlines visible when opened in Excel.
        if isinstance(cell.value, str) and '\n' in cell.value:
            cell.alignment = Alignment(wrap_text=True)

wb.save(output_path)
print('wrote:', output_path)
print('source rows:', len(rows))
print('source columns:', len(headers), headers)

# Verification: reopen source and output, confirm requirements.
with open(input_path, 'r', encoding='utf-8-sig', newline='') as f:
    source_rows = list(csv.reader(f))

check_wb = load_workbook(output_path, data_only=False)
check_ws = check_wb.active

# Confirm no extra index column: dimensions must match exact CSV row/column counts.
expected_rows = len(source_rows)
expected_cols = max(len(r) for r in source_rows) if source_rows else 0
actual_rows = check_ws.max_row
actual_cols = check_ws.max_column
assert actual_rows == expected_rows, f'row count mismatch: {actual_rows} != {expected_rows}'
assert actual_cols == expected_cols, f'column count mismatch: {actual_cols} != {expected_cols}'

# Confirm headers/order and every value, with empty CSV fields represented as blank cells.
for r_idx, src_row in enumerate(source_rows, start=1):
    for c_idx in range(1, expected_cols + 1):
        src_val = src_row[c_idx - 1] if c_idx <= len(src_row) else ''
        got = check_ws.cell(r_idx, c_idx).value
        if src_val == '':
            assert got is None, f'expected blank at row {r_idx} col {c_idx}, got {got!r}'
        else:
            assert got == src_val, f'value mismatch at row {r_idx} col {c_idx}: {got!r} != {src_val!r}'

# Confirm Device ID is text, including preserved leading zeroes.
if device_id_col is not None:
    for r_idx in range(2, expected_rows + 1):
        src_val = source_rows[r_idx - 1][device_id_col - 1]
        cell = check_ws.cell(r_idx, device_id_col)
        assert cell.number_format == '@', f'Device ID number format not text at row {r_idx}: {cell.number_format!r}'
        if src_val != '':
            assert isinstance(cell.value, str), f'Device ID not string at row {r_idx}: {type(cell.value)}'
            assert cell.value == src_val, f'Device ID changed at row {r_idx}: {cell.value!r} != {src_val!r}'
            if src_val.startswith('0'):
                assert cell.value.startswith('0'), f'Device ID leading zero lost at row {r_idx}: {cell.value!r}'

# Confirm duplicate rows retained by comparing full data sequence.
workbook_rows_as_csv = []
for r_idx in range(1, expected_rows + 1):
    out_row = []
    for c_idx in range(1, expected_cols + 1):
        v = check_ws.cell(r_idx, c_idx).value
        out_row.append('' if v is None else v)
    workbook_rows_as_csv.append(out_row)
assert workbook_rows_as_csv == source_rows, 'Workbook row sequence does not exactly match CSV rows'

print('verification passed:')
print(' - output file exists:', output_path.exists())
print(' - headers/order:', headers)
print(' - rows preserved including header:', expected_rows)
print(' - columns preserved/no index column:', expected_cols)
print(' - duplicate rows preserved by exact row-sequence comparison')
print(' - UTF-8 BOM/Chinese/quoted comma/quoted newline parsed via csv module')
print(' - empty fields are blank cells')
print(' - Device ID cells are text (@) and leading zeroes preserved')
