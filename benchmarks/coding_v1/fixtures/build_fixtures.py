"""Build the frozen inputs of coding_v1 deterministically.

    python benchmarks/coding_v1/fixtures/build_fixtures.py

Writes `fixtures/T??/input/...` (what the Agent under test may read), the web
pages served by `servers.py` under `fixtures/web/`, and `MANIFEST.json` with
the SHA-256 of every file. Answers live only in `evaluator/`, never here.
"""

from __future__ import annotations

import csv
import hashlib
import io
import json
import shutil
from pathlib import Path

HERE = Path(__file__).resolve().parent


def _write(path: Path, data: str | bytes, *, encoding: str = "utf-8") -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if isinstance(data, str):
        path.write_bytes(data.encode(encoding))
    else:
        path.write_bytes(data)


def _xlsx(
    path: Path, sheet: str, rows: list[list[object]], *, text_columns: tuple[int, ...] = ()
) -> None:
    from openpyxl import Workbook

    book = Workbook()
    ws = book.active
    ws.title = sheet
    for r, row in enumerate(rows, start=1):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if (c - 1) in text_columns:
                cell.number_format = "@"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)
    _normalize_xlsx(path)


def _normalize_xlsx(path: Path) -> None:
    """Fixed timestamps inside the zip so the hash is reproducible."""
    import zipfile

    data = io.BytesIO(path.read_bytes())
    out = io.BytesIO()
    with zipfile.ZipFile(data) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in sorted(src.infolist(), key=lambda i: i.filename):
            content = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                import re

                content = re.sub(
                    rb"<dcterms:(created|modified)[^>]*>[^<]*</dcterms:\1>",
                    lambda m: (
                        m.group(0).split(b">")[0]
                        + b">2026-01-01T00:00:00Z</dcterms:"
                        + m.group(1)
                        + b">"
                    ),
                    content,
                )
            info = zipfile.ZipInfo(item.filename, date_time=(2026, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            dst.writestr(info, content)
    path.write_bytes(out.getvalue())


def t01(root: Path) -> None:
    d = root / "input" / "texts"
    _write(d / "a.txt", "第一行中文內容\n第二行：測試資料")
    _write(d / "b.txt", "Hello world\n\nThe line above is blank.\n")
    _write(d / "empty.txt", "")
    _write(d / "ignore.csv", "x,y\n1,2\n")
    _write(d / "nested" / "c.txt", "nested file must not be merged\n")


def t02(root: Path) -> None:
    _write(
        root / "input" / "results.csv",
        "Device ID,Status,Value\n001,Pass,10\n002,Fail,20\n003, fail ,30\n004,FAIL,40\n005,,50\n",
    )


def t03(root: Path) -> None:
    rows = [
        ["Task", "Owner", "Status"],
        ["T-1", "Alice", "Pass"],
        ["T-2", "Alice", "Fail"],
        ["T-3", "Bob", "Pass"],
        ["T-4", None, "Fail"],
        ["T-5", "A/B", "Pass"],
        ["T-6", "A:B", "Fail"],
        ["T-7", "X" * 40, "Pass"],
    ]
    _xlsx(root / "input" / "tasks.xlsx", "tasks", rows)


def t04(root: Path) -> None:
    d = root / "input" / "files"
    files = {
        "dir1/report.PDF": b"%PDF-1.4 dir1 report",
        "dir2/report.PDF": b"%PDF-1.4 dir2 report different",
        "budget.xls": b"legacy xls bytes",
        "plan.xlsx": b"xlsx bytes (name decides the class)",
        "logo.png": b"\x89PNG fake",
        "photo.JPG": b"\xff\xd8\xff fake jpeg",
        "notes.txt": b"plain notes",
        "README": b"no extension",
    }
    for name, data in files.items():
        _write(d / name, data)


def t05(root: Path) -> None:
    _xlsx(
        root / "input" / "devices.xlsx",
        "devices",
        [
            ["Device ID", "Name", "Location"],
            ["001", "Alpha", "Lab A"],
            ["002", "Beta", "Lab B"],
            ["003", "Gamma", "Lab C"],
        ],
        text_columns=(0,),
    )
    _xlsx(
        root / "input" / "tests.xlsx",
        "tests",
        [
            ["Device ID", "Test", "Result"],
            ["001", "RF-1", "Pass"],
            ["001", "RF-2", "Fail"],
            ["002", "RF-1", "Pass"],
            ["999", "RF-9", "Fail"],
        ],
        text_columns=(0,),
    )


def t06(root: Path) -> None:
    _write(
        root / "input" / "sop.md",
        """# 量測準備 SOP

本文件說明每日量測前的準備工作，適用於所有實驗室人員。

1. Action: 開啟量測儀器電源
   Input: 主電源開關
   Expected result: 儀器開機畫面出現

2. Action: 執行自我校正
   Expected result: 校正狀態顯示 OK

3. Action: 記錄環境溫度
   Input: 溫度計讀值

## 注意事項

若儀器顯示錯誤代碼，請停止作業並通知主管。
""",
    )


def t07(root: Path) -> None:
    _write(
        root / "input" / "specs.md",
        """# 產品規格

## P1

- Model: RX-100
- Frequency: 2.4–2.5 GHz
- Bandwidth: 80 MHz
- Interface: USB/LAN

## P2

- Model: RX-200
- Frequency: 900–930 MHz
- Bandwidth: 0.02 GHz
- Interface: LAN

## P3

- Model: RX-300
- Frequency: 5–6 GHz
- Interface: USB
""",
    )


def t08(root: Path) -> None:
    _write(
        root / "input" / "report_a.csv",
        "DUT,Item,Result,Measurement,Unit\nD001,TxPower,OK,12.5,dBm\nD002,TxPower,NG,,dBm\n",
    )
    _write(
        root / "input" / "report_b.json",
        json.dumps(
            [
                {"device": "D003", "test": "EVM", "outcome": "PASS", "value": -35.2, "unit": "dB"},
                {"device": "D004", "test": "EVM", "outcome": "FAIL", "value": -20.1},
            ],
            ensure_ascii=False,
            indent=2,
        ),
    )
    _write(
        root / "input" / "report_c.md",
        """| Device | Test | Status | Value | Unit |
|---|---|---|---|---|
| D005 | Sensitivity | OK | -90 | dBm |
| D006 | Sensitivity | FAIL | -70 | dBm |
""",
    )


def t09(root: Path) -> None:
    _write(
        root / "input" / "login_sop.md",
        """# 資料下載 SOP

1. 以帳號 bench / 密碼 bench-pass 呼叫服務的 POST /login（JSON body: {"user": ..., "password": ...}）。
   回應 JSON 的 ok 為 true 時，token 欄位是本次的存取 token；ok 為 false 表示登入失敗。
2. 登入失敗時等待 0.1 秒後重試，最多嘗試 3 次。
3. 登入成功後，以 header Authorization: Bearer <token> 呼叫 GET /data。
4. 將 /data 的 JSON 回應存成 data.json。
5. 3 次都登入失敗時停止，不可呼叫 /data，並回報失敗原因。
""",
    )


def t10(root: Path) -> None:
    _write(
        root / "input" / "catalog_notes.md",
        '主 catalog 位於頁面中 <main id="catalog"> 的區域；<aside class="ads"> 是廣告與推薦區。\n',
    )


def t11(root: Path) -> None:
    pass


def t12(root: Path) -> None:
    pass


def t13(root: Path) -> None:
    pass


def t14(root: Path) -> None:
    buffer = io.StringIO()
    writer = csv.writer(buffer, lineterminator="\n")
    writer.writerow(["Test ID", "Device Model", "Status", "DurationSeconds"])
    for i in range(10000):
        writer.writerow([f"R{i:05d}", f"M{i % 5}", "Fail" if i % 4 == 0 else "Pass", 1 + (i % 10)])
    _write(root / "input" / "large_results.csv", buffer.getvalue())


def t15(root: Path) -> None:
    lines = [
        "[2026-01-05T08:00:00][DUT-01][TxPower][PASS]",
        "[2026-01-05T08:00:05][DUT-01][EVM][FAIL]",
        "",
        "[2026-01-05T08:01:00][DUT-02][TxPower][PASS]",
        "[2026-01-05T08:01:05][DUT-02][EVM]PASS]",
        "[2026-01-05T08:02:00][DUT-03][TxPower][PASS]",
        "[2026-01-05T08:02:05][DUT-03][PASS]",
        "[2026-01-05T08:03:00][DUT-04][TxPower][FAIL]",
        "[2026-01-05T08:03:05][DUT-04][EVM][SKIP]",
        "[2026-01-05T08:04:00][DUT-05][EVM][PASS]",
    ]
    _write(root / "input" / "test.log", "\n".join(lines) + "\n")


def t16(root: Path) -> None:
    _write(
        root / "input" / "logs" / "run_a.log",
        "[2026-02-01T10:00:00][DUT-11][TxPower][PASS]\n"
        "[2026-02-01T10:00:05][DUT-11][EVM][FAIL]\n"
        "[2026-02-01T10:01:00][DUT-12][TxPower][PASS]\n",
    )
    _write(
        root / "input" / "logs" / "run_b.log",
        "\n".join(
            json.dumps(item)
            for item in [
                {
                    "ts": "2026-02-02T09:00:00",
                    "dut": "DUT-21",
                    "case": "Sensitivity",
                    "result": "PASS",
                },
                {"ts": "2026-02-02T09:00:10", "dut": "DUT-21", "case": "EVM", "result": "FAIL"},
                {"ts": "2026-02-02T09:01:00", "dut": "DUT-22", "case": "TxPower", "result": "PASS"},
            ]
        )
        + "\n",
    )
    _write(
        root / "input" / "logs" / "run_c.log",
        "2026/02/03 11:00 DUT-31 TxPower => PASS\n2026/02/03 11:01 DUT-31 EVM => FAIL\n",
    )


def t17(root: Path) -> None:
    plan = {
        "name": "RF plan",
        "description": "frequency values in GHz or MHz; keep this text",
        "center_frequency": "2.4GHz",
        "bands": [
            {"label": "WLAN 2G", "frequency": "2.4GHz", "bandwidth": "20MHz", "power": 18},
            {"label": "WLAN 5G", "frequency": "5 GHz", "bandwidth": "80MHz", "power": 17},
            {"label": "ISM", "frequency": "2400MHz", "timeout": 30},
        ],
        "sub": {"nested": [{"frequency": 915, "power": 10}, {"frequency": "100kHz", "note": "LF"}]},
        "timeout": 120,
    }
    _write(root / "input" / "testplan.json", json.dumps(plan, ensure_ascii=False, indent=2))


def t18(root: Path) -> None:
    _write(root / "input" / "input.csv", "Model,Value\nM1,10\nM1,20\nM1,\nM2,5\nM2,15\n")
    _write(
        root / "input" / "export.py",
        '''"""Average Value per Model from input.csv into summary.xlsx.

Usage: python export.py --input input.csv --output summary.xlsx
Requires: openpyxl
"""

import argparse
import csv

from openpyxl import Workbook


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()
    totals = {}
    counts = {}
    with open(args.input, newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            model = row["DeviceModel"]
            value = row["Value"]
            if value == "":
                continue
            totals[model] = totals.get(model, "") + value
            counts[model] = counts.get(model, 0) + 1
    book = Workbook()
    sheet = book.active
    sheet.title = "summary"
    sheet.append(["Model", "Average", "ValidCount"])
    for model in sorted(totals):
        sheet.append([model, totals[model] / counts[model], counts[model]])
    book.save(args.output)


if __name__ == "__main__":
    main()
''',
    )
    _write(root / "input" / "DEPENDENCIES.txt", "Python 3.12, openpyxl\n")


def t19(root: Path) -> None:
    rows = [
        ["Device ID", "名稱", "備註", "Value"],
        ["0001", "測試機甲", "正常", "10"],
        ["0002", "測試機乙", "含, 逗號", "20"],
        ["0003", "測試機丙", "第一行\n第二行", ""],
        ["0001", "測試機甲", "正常", "10"],
        ["0005", "", "", "50"],
        ["0006", "測試機己", "說明", "60"],
    ]
    buffer = io.StringIO()
    csv.writer(buffer, lineterminator="\r\n").writerows(rows)
    _write(root / "input" / "mixed.csv", "﻿" + buffer.getvalue())


def t20(root: Path) -> None:
    _write(
        root / "input" / "supplied_sop.md",
        """# 測試失敗追蹤 SOP

1. 讀取測試結果 results.csv（欄位：Device ID、Status、Value）。
2. 挑出 Status 為 Fail 的資料，輸出 failures.xlsx，保留 Device ID、Status、Value，並加上 Urgent 欄位。
3. 超過門檻的 Fail 列標記需立即處理（Urgent）。
4. 產生 summary.md，列出 Total、Pass、Fail、Urgent 的數量。
5. 原始資料不可修改。
""",
    )
    _write(
        root / "input" / "manifest.json",
        json.dumps({"results.csv": "input/data/test_results.csv"}, indent=2),
    )
    _write(
        root / "input" / "data" / "test_results.csv",
        "Device ID,Status,Value\nD01,Fail,10\nD02,Pass,5\nD03,Fail,20\nD04,Pass,7\nD05,Fail,30\n",
    )


# -- web pages (served by servers.py; not visible as files to the Agent) ----


def web(root: Path) -> None:
    ads = (
        '<aside class="ads"><div class="card"><a href="/ad/1">Sponsored Router</a>'
        '<span class="price">NT$99</span></div><div class="card"><a href="/ad/2">推薦：高速網卡</a>'
        '<span class="price">NT$199</span></div></aside>'
    )
    items = [
        ("Router A", "NT$1,299", "/p/router-a"),
        ("Switch B", "NT$2,450", "/p/switch-b"),
        ("Cable C", "NT$99", "/p/cable-c"),
        ("Adapter D", "NT$560", "/p/adapter-d"),
        ("Antenna E", "NT$12,000", "/p/antenna-e"),
    ]
    cards = "".join(
        f'<div class="card"><a href="{u}">{n}</a><span class="price">{p}</span></div>'
        for n, p, u in items
    )
    _write(
        root / "t10" / "catalog.html",
        f'<!doctype html><html><head><meta charset="utf-8"><title>Catalog</title></head><body>'
        f'{ads}<main id="catalog"><h1>產品目錄</h1>{cards}</main></body></html>',
    )

    details = [
        ("Laptop 13", "NT$32,900", "Intel i5-1340P", "16 GB", "/t11/item/1"),
        ("Laptop 15", "NT$45,500", "Intel i7-1360P", "32GB", "/t11/item/2"),
        ("Mini PC", "NT$18,800", None, "16 GB", "/t11/item/3"),
        ("Tablet 11", "NT$12,300", "ARM A78", "8 GB", "/t11/item/4"),
    ]
    listing = "".join(f'<li><a href="{u}">{n}</a></li>' for n, _p, _c, _m, u in details)
    _write(
        root / "t11" / "list.html",
        f'<!doctype html><html><head><meta charset="utf-8"><title>List</title></head>'
        f'<body><ul id="products">{listing}</ul></body></html>',
    )
    for n, p, c, m, u in details:
        rows = f"<tr><th>Price</th><td>{p}</td></tr>"
        if c:
            rows += f"<tr><th>CPU</th><td>{c}</td></tr>"
        rows += f"<tr><th>Memory</th><td>{m}</td></tr>"
        _write(
            root / "t11" / f"item{u.rsplit('/', 1)[1]}.html",
            f'<!doctype html><html><head><meta charset="utf-8"><title>{n}</title></head><body>'
            f'<h1 class="name">{n}</h1><table class="spec">{rows}</table></body></html>',
        )

    _write(
        root / "t12" / "a.html",
        '<!doctype html><html><head><meta charset="utf-8"></head><body>'
        '<table id="p"><tr><th>型號</th><th>價格</th><th>記憶體</th></tr>'
        "<tr><td>NX-1</td><td>NT$12,000</td><td>16 GB</td></tr>"
        "<tr><td>NX-2</td><td>NT$8,500</td><td>8192 MB</td></tr></table></body></html>",
    )
    _write(
        root / "t12" / "b.html",
        '<!doctype html><html><head><meta charset="utf-8"></head><body>'
        '<div class="item"><h2>NX-1</h2><p class="price">USD 399.50</p><p class="mem">Memory: 16384 MB</p></div>'
        '<div class="item"><h2>NX-3</h2><p class="price">USD 1,250.00</p><p class="mem">Memory: 32 GB</p></div>'
        "</body></html>",
    )
    _write(
        root / "t12" / "c.html",
        '<!doctype html><html><head><meta charset="utf-8"></head><body>'
        '<ul><li data-model="NX-4">NX-4 — €499.00 — 4096 MB</li>'
        '<li data-model="NX-5">NX-5 — €1.299,00 — 64 GB</li></ul></body></html>',
    )

    t13 = [
        ("Phone A", "NT$9,900", "Snap 7", "/t13/item/1"),
        ("Phone B", "NT$12,900", "Snap 8", "/t13/item/2"),
        ("Phone C", None, "Dim 9", "/t13/item/3"),
        ("Phone D", "NT$7,500", None, "/t13/item/4"),
        ("Phone E", "NT$15,000", "Snap 9", "/t13/item/5"),
    ]
    listing = "".join(f'<li><a href="{u}">{n}</a></li>' for n, _p, _c, u in t13)
    _write(
        root / "t13" / "list.html",
        f'<!doctype html><html><head><meta charset="utf-8"></head>'
        f'<body><ul id="products">{listing}</ul></body></html>',
    )
    for n, p, c, u in t13:
        rows = (f"<tr><th>Price</th><td>{p}</td></tr>" if p else "") + (
            f"<tr><th>CPU</th><td>{c}</td></tr>" if c else ""
        )
        _write(
            root / "t13" / f"item{u.rsplit('/', 1)[1]}.html",
            f'<!doctype html><html><head><meta charset="utf-8"></head><body><h1 class="name">{n}</h1>'
            f'<table class="spec">{rows}</table></body></html>',
        )


BUILDERS = {f"T{n:02d}": globals()[f"t{n:02d}"] for n in range(1, 21)}


def manifest(root: Path) -> dict[str, str]:
    return {
        path.relative_to(root).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(root.rglob("*"))
        if path.is_file()
        and path.name != "MANIFEST.json"
        and path.suffix != ".py"
        and "__pycache__" not in path.parts
    }


def main() -> None:
    for task, build in BUILDERS.items():
        target = HERE / task
        if target.exists():
            shutil.rmtree(target)
        (target / "input").mkdir(parents=True)
        build(target)
    if (HERE / "web").exists():
        shutil.rmtree(HERE / "web")
    web(HERE / "web")
    files = manifest(HERE)
    (HERE / "MANIFEST.json").write_text(
        json.dumps(files, indent=2, ensure_ascii=False) + "\n", encoding="utf-8"
    )
    print(f"wrote {len(files)} fixture files")


if __name__ == "__main__":
    main()
