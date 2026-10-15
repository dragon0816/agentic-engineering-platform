"""Self-test of the grader: a reference fill must pass, a damaged one must fail.

The reference fill is written by the operator from the expected answers; it
exists only to prove the checks are neither too strict nor too loose and is
never given to the Agent or counted as an Agent result.

  python operator/selftest.py
"""

from __future__ import annotations

import shutil
import sys
import tempfile
from decimal import Decimal
from pathlib import Path
from typing import Any

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from expected import expected, source, twd  # noqa: E402
from grade import TEMPLATE, grade_workbook, verdict  # noqa: E402
from openpyxl import load_workbook  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402

LAB = HERE.parent
CASES = [
    ("W01", "main"),
    ("W01", "rerun"),
    ("W02", "main"),
    ("W03", "main"),
    ("W04", "main"),
    ("D01", "main"),
    ("D02", "main"),
    ("D03", "main"),
    ("D04", "main"),
    ("C01", "main"),
    ("C02", "main"),
    ("C03", "main"),
    ("C04", "main"),
    ("X01", "main"),
    ("X02", "main"),
    ("X03", "main"),
]


def put_rows(sheet: Any, rows: list[list[Any]], money_columns: set[int]) -> None:
    for offset, values in enumerate(rows):
        for index, value in enumerate(values, 1):
            cell = sheet.cell(row=5 + offset, column=index, value=value)
            if index in money_columns and isinstance(value, Decimal):
                cell.value = float(value)
                cell.number_format = "0.00"


def group_rows(summary: dict[str, dict[str, Any]], width: int) -> list[list[Any]]:
    out = []
    for key, value in summary.items():
        parts = key.split("|")
        keys = (parts[:-1] + ["", ""])[:2] if width else ["", ""]
        out.append(
            [*keys, parts[-1], Decimal(value["amount"]), Decimal(value["twd"]), value["count"]]
        )
    return out


def fill(task: str, variant: str, target: Path) -> None:
    exp = expected(task, variant)
    data = source()
    kind = TEMPLATE[task]
    shutil.copy(LAB / "templates" / f"{kind}_template.xlsx", target)
    book = load_workbook(target)
    book["Report"]["B9"] = "missing: " + ", ".join(exp.get("missing_detail", []))
    if kind == "sales":
        orders = {o["OrderID"]: o for o in data["sales_orders"]}
        rows = []
        for ident in exp["ids"]:
            o = orders[ident]
            rows.append(
                [
                    ident,
                    o["CustomerID"],
                    o["CustomerName"],
                    o["ProjectID"],
                    o["OrderDate"],
                    o["Status"],
                    o["Currency"],
                    Decimal(o["NetAmount"]),
                    twd(o["NetAmount"], o["Currency"]),
                    o["CustomerContact"],
                    o["PaymentTerms"],
                ]
            )
        put_rows(book["Details"], rows, {8, 9})
        put_rows(book["Summary"], group_rows(exp["summary_by_currency"], 0), {4, 5})
        book["Summary"]["B2"] = f"=SUM(Details!I5:I{max(5, 4 + len(rows))})"
    elif kind == "expense":
        expenses = {e["ExpenseID"]: e for e in data["expenses"]}
        rows = []
        for ident in exp["ids"]:
            e = expenses[ident]
            amount_twd = twd(e["Amount"], e["Currency"])
            rows.append(
                [
                    ident,
                    e["EmployeeID"],
                    e["EmployeeName"],
                    e["ProjectID"],
                    e["ExpenseDate"],
                    e["Category"],
                    e["Currency"],
                    Decimal(e["Amount"]),
                    amount_twd,
                    e["Status"],
                    e["ReceiptID"],
                    e["Description"],
                    str(e["MissingReceipt"]).lower(),
                    str(amount_twd > 5000).lower(),
                ]
            )
        put_rows(book["Details"], rows, {8, 9})
        put_rows(book["Summary"], group_rows(exp["summary"], len(exp["summary_group"])), {4, 5})
        book["Summary"]["B2"] = f"=SUM(Details!I5:I{max(5, 4 + len(rows))})"
        put_rows(
            book["Exceptions"],
            [
                [i, expenses[i]["EmployeeID"], expenses[i]["ExpenseDate"], "MissingReceipt"]
                for i in exp.get("exceptions", [])
            ],
            set(),
        )
    elif kind == "quote":
        quotes = {q["QuoteID"]: q for q in data["quotes"]}
        rows, lines = [], []
        for ident in exp["ids"]:
            q = quotes[ident]
            rows.append(
                [
                    ident,
                    q["CustomerID"],
                    q["ProjectID"],
                    q["QuoteDate"],
                    q["ValidUntil"],
                    q["Status"],
                    q["Currency"],
                    Decimal(q["Subtotal"]),
                    int(q["DiscountPercent"]),
                    Decimal(q["NetAmount"]),
                    twd(q["NetAmount"], q["Currency"]),
                ]
            )
            lines += [
                [
                    ident,
                    x["ProductID"],
                    x["ProductName"],
                    x["Quantity"],
                    Decimal(x["UnitPrice"]),
                    Decimal(x["LineAmount"]),
                ]
                for x in q["Lines"]
            ]
        put_rows(book["Quotes"], rows, {8, 10, 11})
        put_rows(book["Lines"], lines, {5, 6})
        if task == "C03":
            put_rows(
                book["Summary"],
                [[p, "", "TWD", "", Decimal(v), ""] for p, v in exp["by_project_twd"].items()],
                {5},
            )
        else:
            put_rows(book["Summary"], group_rows(exp["summary"], len(exp["summary_group"])), {4, 5})
        book["Summary"]["B2"] = f"=SUM(Quotes!K5:K{max(5, 4 + len(rows))})"
    else:
        sources = {
            "Orders": {o["OrderID"]: o for o in data["sales_orders"]},
            "Invoices": {i["InvoiceID"]: i for i in data["invoices"]},
            "Expenses": {e["ExpenseID"]: e for e in data["expenses"]},
            "Quotes": {q["QuoteID"]: q for q in data["quotes"]},
        }
        for name, sheet in exp["sheets"].items():
            rows = []
            for ident in sheet["ids"]:
                r = sources[name][ident]
                if name == "Orders":
                    rows.append(
                        [
                            ident,
                            r["ProjectID"],
                            r["CustomerID"],
                            r["OrderDate"],
                            r["Status"],
                            r["Currency"],
                            Decimal(r["NetAmount"]),
                            Decimal(sheet["twd"][ident]),
                        ]
                    )
                elif name == "Invoices":
                    rows.append(
                        [
                            ident,
                            r["OrderID"],
                            r["ProjectID"],
                            r["CustomerID"],
                            r["InvoiceDate"],
                            r["DueDate"],
                            r["Currency"],
                            Decimal(r["Amount"]),
                            Decimal(r["PaidAmount"]),
                            Decimal(r["Balance"]),
                            Decimal(sheet["twd"][ident]),
                        ]
                    )
                elif name == "Expenses":
                    rows.append(
                        [
                            ident,
                            r["ProjectID"],
                            r["EmployeeID"],
                            r["ExpenseDate"],
                            r["Status"],
                            r["Currency"],
                            Decimal(r["Amount"]),
                            Decimal(sheet["twd"][ident]),
                        ]
                    )
                else:
                    rows.append(
                        [
                            ident,
                            r["ProjectID"],
                            r["CustomerID"],
                            r["QuoteDate"],
                            r["ValidUntil"],
                            r["Status"],
                            r["Currency"],
                            Decimal(r["NetAmount"]),
                            Decimal(sheet["twd"][ident]),
                        ]
                    )
            put_rows(book[name], rows, set(range(7, 12)))
        summary = [
            [
                p,
                *(
                    Decimal(v.get(c, "0"))
                    for c in (
                        "OrderTWD",
                        "InvoiceBalanceTWD",
                        "ApprovedExpenseTWD",
                        "AcceptedQuoteTWD",
                    )
                ),
            ]
            for p, v in exp["summary_projects"].items()
        ]
        put_rows(book["Summary"], summary, {2, 3, 4, 5})
        for column in range(2, 6):
            letter = get_column_letter(column)
            book["Summary"][f"{letter}2"] = f"=SUM({letter}5:{letter}{4 + len(summary)})"
        attention: list[list[Any]] = []
        for ident in exp.get("overdue", []):
            attention.append(
                ["Invoices", ident, sources["Invoices"][ident]["ProjectID"], "Overdue"]
            )
        for item in exp.get("attention", []):
            name, ident = item.split(":")
            attention.append([name, ident, sources[name][ident]["ProjectID"], "Attention"])
        put_rows(book["Exceptions"], attention, set())
    book.save(target)
    if task == "C02":
        (target.parent / "findings.md").write_text(
            "".join(
                f"- {q}: ValidUntil before 2026-09-30, not Accepted\n"
                for q in exp["expired_not_accepted"]
            ),
            encoding="utf-8",
        )


def damage(path: Path, task: str) -> None:
    book = load_workbook(path)
    sheet = {"sales": "Details", "expense": "Details", "quote": "Quotes"}.get(TEMPLATE[task])
    if sheet is None:
        sheet = "Summary"
    if book[sheet].max_row >= 5 and book[sheet]["A5"].value is not None:
        book[sheet].delete_rows(5)
    else:
        book[sheet]["A5"] = "BOGUS-0001"
    book.save(path)


def rerun_rule_checks() -> int:
    """Grader 1.3: a revision after the rerun request is counted; a reply without one is not."""
    from run_lab import drafted

    def exchange(*names: str) -> dict[str, Any]:
        tools = [{"name": n, "host": {"status": "succeeded", "data": {}}} for n in names]
        return {"reply": {"answer": {"outcome": {"conversation": {"tools": tools}}}}}

    cases = [
        (drafted([exchange("authoring__run_workflow")]), 0),
        (
            drafted(
                [
                    exchange("authoring__draft_workflow", "authoring__run_workflow"),
                    exchange("authoring__draft_workflow"),
                ]
            ),
            2,
        ),
        (drafted([{"reply": None}]), 0),
    ]
    wrong = sum(1 for got, want in cases if got != want)
    print(f"{'PASS' if not wrong else 'FAIL'}  rerun rule: revisions counted after the request")
    return wrong


def main() -> int:
    failures = rerun_rule_checks()
    with tempfile.TemporaryDirectory() as folder:
        for task, variant in CASES:
            target = Path(folder) / f"{task}-{variant}.xlsx"
            if task == "C04" and variant == "force_exit2":
                continue
            fill(task, variant, target)
            checks = grade_workbook(task, variant, target, "", Path(folder))
            good = verdict(checks)
            bad_checks = [
                c["name"] + ": " + str(c["detail"])[:160] for c in checks if c["ok"] is False
            ]
            damage(target, task)
            broken = verdict(grade_workbook(task, variant, target, "", Path(folder)))
            ok = good is True and broken is False
            failures += 0 if ok else 1
            print(
                f"{'PASS' if ok else 'FAIL'}  {task}:{variant}  reference={good} damaged={broken}"
            )
            for line in bad_checks:
                print("      ", line)
    print("selftest:", "all good" if not failures else f"{failures} case(s) wrong")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
