"""Check a produced workbook against the template and the expected answer.

Every check is a dict {"name", "ok", "detail"}; `ok` is None when the check
cannot be decided from the file (for example a formula with no cached value),
which is reported as such and never counted as a pass.
"""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

LAB = Path(__file__).resolve().parents[1]
TEMPLATES = LAB / "templates"
REPORT_LABELS = ("A1", "A3", "C3", "A4", "A5", "A7", "C7", "A9", "A20")
DETAIL_SHEETS = {"Details", "Quotes", "Lines", "Orders", "Invoices", "Expenses", "Exceptions"}


def check(name: str, ok: bool | None, detail: Any = "") -> dict[str, Any]:
    return {"name": name, "ok": ok, "detail": detail}


def dec(value: Any) -> Decimal | None:
    if value is None or isinstance(value, bool) or value == "":
        return None
    try:
        return Decimal(str(value).replace(",", "").strip()).quantize(Decimal("0.01"))
    except (InvalidOperation, ValueError):
        return None


def text(value: Any) -> str:
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def truthy(value: Any) -> bool | None:
    if isinstance(value, bool):
        return value
    lowered = text(value).lower()
    if lowered in {"true", "yes", "y", "1", "是", "true."}:
        return True
    if lowered in {"false", "no", "n", "0", "否", ""}:
        return False
    return None


def table(sheet: Worksheet) -> tuple[list[str], list[dict[str, Any]], list[int]]:
    """Column names from row 4 and data rows from row 5 down to the last non-empty row."""
    headers = [text(c.value) for c in sheet[4]]
    rows, numbers = [], []
    for index in range(5, sheet.max_row + 1):
        values = [sheet.cell(row=index, column=i + 1).value for i in range(len(headers))]
        if all(v is None or text(v) == "" for v in values):
            continue
        rows.append({h: v for h, v in zip(headers, values, strict=True) if h})
        numbers.append(index)
    return headers, rows, numbers


def template_checks(path: Path, template: str) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    produced = load_workbook(path)
    pristine = load_workbook(TEMPLATES / template)
    out.append(check("template: sheet names and order", produced.sheetnames == pristine.sheetnames,
                     produced.sheetnames))
    if "Report" in produced.sheetnames:
        report, original = produced["Report"], pristine["Report"]
        changed = [c for c in REPORT_LABELS if text(report[c].value) != text(original[c].value)]
        out.append(check("template: Report fixed labels and footer cell", not changed, changed))
        merged = {str(r) for r in report.merged_cells.ranges}
        out.append(check("template: Report A1:F1 merged", "A1:F1" in merged, sorted(merged)))
    footers = [s.title for s in produced if s.title in pristine.sheetnames
               and text(s.oddFooter.center.text) != text(pristine[s.title].oddFooter.center.text)]
    out.append(check("template: page footers kept", not footers, footers))
    if "Instructions" in produced.sheetnames:
        a, b = produced["Instructions"], pristine["Instructions"]
        diffs = [c.coordinate for row in b.iter_rows() for c in row
                 if text(a[c.coordinate].value) != text(c.value)]
        out.append(check("template: Instructions unchanged", not diffs, diffs[:10]))
    header_issues, layout_issues, width_issues = [], [], []
    for sheet in pristine:
        if sheet.title not in produced.sheetnames or sheet.title in {"Report", "Instructions"}:
            continue
        mine = produced[sheet.title]
        if [text(c.value) for c in mine[4]][: sheet.max_column] != [text(c.value) for c in sheet[4]]:
            header_issues.append(sheet.title)
        if sheet.title in DETAIL_SHEETS:
            if mine.freeze_panes != "A5":
                layout_issues.append(f"{sheet.title}: freeze {mine.freeze_panes}")
            if not (mine.auto_filter.ref or "").startswith("A4"):
                layout_issues.append(f"{sheet.title}: auto-filter {mine.auto_filter.ref}")
        for letter, dim in sheet.column_dimensions.items():
            if dim.width and abs((mine.column_dimensions[letter].width or 0) - dim.width) > 0.6:
                width_issues.append(f"{sheet.title}!{letter}")
    out.append(check("template: row-4 column names kept", not header_issues, header_issues))
    out.append(check("template: freeze panes and auto-filter kept", not layout_issues, layout_issues))
    out.append(check("template: column widths kept", not width_issues, width_issues[:10]))
    return out


def formula_checks(path: Path, total_cell: tuple[str, str, str, str] | None, template: str) -> list[dict[str, Any]]:
    """Every template formula is still a formula over the same column, extended to the data."""
    out = []
    produced = load_workbook(path)
    cached = load_workbook(path, data_only=True)
    pristine = load_workbook(TEMPLATES / template)
    for sheet in pristine:
        for row in sheet.iter_rows():
            for cell in row:
                if not (isinstance(cell.value, str) and cell.value.startswith("=")):
                    continue
                where = f"{sheet.title}!{cell.coordinate}"
                if sheet.title not in produced.sheetnames:
                    out.append(check(f"formula kept: {where}", False, "sheet missing"))
                    continue
                value = produced[sheet.title][cell.coordinate].value
                found = re.match(r"=SUM\(((?:\w+!)?)([A-Z]+)5:([A-Z]+)(\d+)\)", str(value or "").replace("$", ""))
                expected = re.match(r"=SUM\(((?:\w+!)?)([A-Z]+)5", cell.value)
                same = bool(found and expected and found.group(1) == expected.group(1)
                            and found.group(2) == expected.group(2))
                out.append(check(f"formula kept: {where}", same, value))
                if same and found:
                    target = found.group(1).rstrip("!") or sheet.title
                    if target in produced.sheetnames:
                        _, _, rows = table(produced[target])
                        last = rows[-1] if rows else 5
                        out.append(check(f"formula range covers the data: {where}", int(found.group(4)) >= last,
                                         f"{value}; last data row {last}"))
                shown = cached[sheet.title][cell.coordinate].value if sheet.title in cached.sheetnames else None
                out.append(check(f"formula cached value present: {where}", None if shown is None else True,
                                 "no cached value (the writer did not calculate formulas); "
                                 "totals are verified independently" if shown is None else shown))
    del total_cell
    return out


def detail_checks(path: Path, exp: dict[str, Any], sheet_name: str, id_column: str,
                  fields: tuple[str, ...]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    book = load_workbook(path)
    if sheet_name not in book.sheetnames:
        return [check(f"{sheet_name}: sheet exists", False)]
    _, rows, _ = table(book[sheet_name])
    raw_ids = [r.get(id_column) for r in rows]
    numeric_ids = [v for v in raw_ids if isinstance(v, (int, float))]
    ids = [text(v) for v in raw_ids]
    wanted = list(exp["ids"])
    out.append(check(f"{sheet_name}: {id_column} set matches", sorted(set(ids)) == sorted(wanted),
                     {"missing": sorted(set(wanted) - set(ids)), "extra": sorted(set(ids) - set(wanted))}))
    out.append(check(f"{sheet_name}: no duplicate {id_column}", len(ids) == len(set(ids)),
                     sorted({i for i in ids if ids.count(i) > 1})))
    out.append(check(f"{sheet_name}: row count", len(ids) == len(wanted), f"{len(ids)} vs {len(wanted)}"))
    out.append(check(f"{sheet_name}: IDs stored as text", not numeric_ids, numeric_ids[:5]))
    wrong: list[str] = []
    for row in rows:
        ident = text(row.get(id_column))
        want = exp["values"].get(ident)
        if want is None:
            continue
        for field in fields:
            if field not in want:
                continue
            got, expected_value = row.get(field), want[field]
            if isinstance(expected_value, bool):
                good = truthy(got) == expected_value
            elif dec(expected_value) is not None and field not in {"ProjectID", "ReceiptID", "CustomerID"}:
                good = dec(got) == dec(expected_value)
            else:
                good = " ".join(text(got).split()) == " ".join(text(expected_value).split())
            if not good:
                wrong.append(f"{ident}.{field}: {text(got)!r} != {text(expected_value)!r}")
    out.append(check(f"{sheet_name}: field values", not wrong, wrong[:12]))
    return out


def column_total(path: Path, sheet_name: str, column: str) -> Decimal | None:
    book = load_workbook(path)
    if sheet_name not in book.sheetnames:
        return None
    _, rows, _ = table(book[sheet_name])
    values = [dec(r.get(column)) for r in rows]
    return sum((v for v in values if v is not None), Decimal(0)) if rows else Decimal(0)


def summary_rows(path: Path) -> list[tuple[set[str], set[Decimal]]]:
    book = load_workbook(path)
    if "Summary" not in book.sheetnames:
        return []
    out = []
    for index in range(5, book["Summary"].max_row + 1):
        cells = [c.value for c in book["Summary"][index]]
        strings = {text(c) for c in cells if c is not None and dec(c) is None}
        numbers = {dec(c) for c in cells if dec(c) is not None}
        if strings or numbers:
            out.append((strings, numbers))
    return out


def grouped_summary_check(path: Path, summary: dict[str, dict[str, Any]], need_amount: bool) -> dict[str, Any]:
    rows = summary_rows(path)
    missing = []
    for key, value in summary.items():
        parts = {p for p in key.split("|") if p}
        hit = any(parts <= strings and dec(value["twd"]) in numbers
                  and (not need_amount or dec(value["amount"]) in numbers) for strings, numbers in rows)
        if not hit:
            missing.append(f"{key}: amount {value['amount']}, TWD {value['twd']}")
    return check("Summary: every expected group row (keys, amounts)", not missing, missing[:10])


def project_summary_check(path: Path, wanted: dict[str, dict[str, str]], columns: list[str]) -> dict[str, Any]:
    book = load_workbook(path)
    if "Summary" not in book.sheetnames:
        return check("Summary: project rows", False, "no Summary sheet")
    _, rows, _ = table(book["Summary"])
    by_project = {text(r.get("ProjectID")): r for r in rows}
    wrong = []
    for project, values in wanted.items():
        row = by_project.get(project)
        if row is None:
            wrong.append(f"{project}: row missing")
            continue
        for column in columns:
            if dec(row.get(column)) != dec(values[column]) and not (
                    dec(values[column]) == 0 and row.get(column) in (None, "")):
                wrong.append(f"{project}.{column}: {text(row.get(column))} != {values[column]}")
    return check("Summary: one row per project with the expected TWD totals", not wrong, wrong[:12])


def ids_in(path: Path, sheet_name: str, column: str) -> list[str]:
    book = load_workbook(path)
    if sheet_name not in book.sheetnames:
        return []
    _, rows, _ = table(book[sheet_name])
    return [text(r.get(column)) for r in rows if text(r.get(column))]


def empty_sheet_check(path: Path, names: list[str]) -> dict[str, Any]:
    book = load_workbook(path)
    filled = [n for n in names if n in book.sheetnames and table(book[n])[1]]
    return check("source sheets that must stay empty are empty", not filled, filled)
