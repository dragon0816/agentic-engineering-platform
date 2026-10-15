"""Build the lab's four Excel templates and the fixed exchange-rate file.

Every template follows one layout so a workflow can fill it from the
Instructions sheet alone:
- Report: A1:F1 merged title; B3 / D3 period start / end; B4 condition text;
  B5 generated-at; row 7 source name (B7) and fetched-at (D7); B9 notes;
  a fixed footer line in A20 and in the page footer.
- Detail sheets: title in A1, column names in row 4, data from row 5; row 4
  frozen, auto-filter on row 4, fixed column widths; IDs stored as text,
  dates as YYYY-MM-DD, money with two decimals.
- Summary: SUM formulas in row 2 over the detail column named in
  Instructions; the summary table's column names in row 4, rows from row 5.

Run:  python benchmarks/business_workflow_lab/templates/build_templates.py
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

HERE = Path(__file__).resolve().parent
FOOTER = "合成測試資料 Synthetic lab data · template v1 · 請勿修改本行"
TITLE_FONT = Font(bold=True, size=14, color="FFFFFF")
TITLE_FILL = PatternFill("solid", fgColor="0A3D62")
HEAD_FONT = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="D9E2EC")
THIN = Side(style="thin", color="9AA5B1")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
ID_COLUMNS = {
    "OrderID",
    "CustomerID",
    "ProjectID",
    "InvoiceID",
    "ExpenseID",
    "EmployeeID",
    "QuoteID",
    "ProductID",
    "ReceiptID",
    "DocumentID",
}
DATE_COLUMNS = {"OrderDate", "InvoiceDate", "DueDate", "ExpenseDate", "QuoteDate", "ValidUntil"}
MONEY_HINT = ("Amount", "TWD", "Subtotal", "UnitPrice", "Balance")


def kind(column: str) -> str:
    if column in ID_COLUMNS:
        return "text (文字，保留前導零)"
    if column in DATE_COLUMNS:
        return "date YYYY-MM-DD"
    if column == "DiscountPercent" or column == "Quantity" or column == "Count":
        return "number"
    if any(h in column for h in MONEY_HINT):
        return "number, 2 decimals (0.00)"
    return "text"


def report(book: Workbook, title: str) -> Worksheet:
    sheet = book.active
    sheet.title = "Report"
    sheet.merge_cells("A1:F1")
    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    sheet["A1"].fill = TITLE_FILL
    sheet["A1"].alignment = Alignment(horizontal="center")
    for cell, label in {
        "A3": "起始日期 From",
        "C3": "結束日期 To",
        "A4": "條件 Condition",
        "A5": "產出時間 Generated at",
        "A7": "資料來源 Source",
        "C7": "取得時間 Fetched at",
        "A9": "備註 Notes",
    }.items():
        sheet[cell] = label
        sheet[cell].font = HEAD_FONT
    sheet["A20"] = FOOTER
    sheet["A20"].font = Font(italic=True, color="7B8794")
    for column, width in zip("ABCDEF", (22, 26, 18, 26, 14, 14), strict=True):
        sheet.column_dimensions[column].width = width
    sheet.oddFooter.center.text = FOOTER
    return sheet


def detail(book: Workbook, name: str, title: str, columns: list[str]) -> Worksheet:
    sheet = book.create_sheet(name)
    sheet["A1"] = title
    sheet["A1"].font = Font(bold=True, size=12)
    sheet["A2"] = "第 4 列為欄名，第 5 列起為資料；欄位格式見 Instructions。"
    for index, column in enumerate(columns, 1):
        cell = sheet.cell(row=4, column=index, value=column)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.border = BOX
        width = 12 if column in ID_COLUMNS or column in DATE_COLUMNS else 16
        if column in {
            "CustomerName",
            "EmployeeName",
            "Description",
            "ProductName",
            "CustomerContact",
            "Reason",
        }:
            width = 32
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A5"
    sheet.auto_filter.ref = f"A4:{get_column_letter(len(columns))}4"
    sheet.oddFooter.center.text = FOOTER
    return sheet


def summary(book: Workbook, columns: list[str], totals: dict[str, str], title: str) -> Worksheet:
    sheet = detail(book, "Summary", title, columns)
    sheet["A2"] = None
    sheet["A2"] = "總計 TWD Totals"
    sheet["A2"].font = HEAD_FONT
    for cell, formula in totals.items():
        sheet[cell] = formula
        sheet[cell].number_format = "#,##0.00"
        sheet[cell].font = Font(bold=True)
    sheet.freeze_panes = None
    sheet.auto_filter.ref = None
    return sheet


def instructions(book: Workbook, rows: list[tuple[str, str, str]]) -> None:
    sheet = book.create_sheet("Instructions")
    sheet["A1"] = "填寫說明 Instructions（請勿修改本工作表）"
    sheet["A1"].font = Font(bold=True, size=12)
    for index, head in enumerate(("Sheet", "Cell / Range", "Content and rule"), 1):
        cell = sheet.cell(row=3, column=index, value=head)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
    common = [
        ("Report", "A1:F1", "合併儲存格標題，保留不改。"),
        ("Report", "B3 / D3", "查詢期間起 / 迄，格式 YYYY-MM-DD（含起訖兩日）。"),
        ("Report", "B4", "本次查詢條件的文字說明。"),
        ("Report", "B5", "產出時間（ISO 8601）。"),
        ("Report", "B7 / D7", "資料來源系統名稱 / 取得資料的時間；多個來源以分號分隔。"),
        ("Report", "B9", "備註：缺少的資料、例外或限制。"),
        ("Report", "A20 與頁尾", "固定 footer，保留不改。"),
        ("All detail sheets", "Row 4", "欄名，保留不改；資料自第 5 列起，一筆一列，不留空列。"),
        (
            "All detail sheets",
            "Formats",
            "ID 一律存為文字（保留前導零）；日期 YYYY-MM-DD；金額為數字並使用兩位小數格式 0.00。",
        ),
        (
            "All sheets",
            "-",
            "不可刪除或改名工作表、不可取消合併儲存格、不可移除公式；可新增資料列並延伸公式範圍。",
        ),
        (
            "All sheets",
            "Layout",
            "欄寬、凍結窗格、自動篩選、字型與填色保留模板設定，不要調整或自動調整欄寬。",
        ),
        (
            "All detail sheets",
            "Missing values",
            "來源沒有的值留空（空白儲存格），不要填 N/A、0 或自行推測的值。",
        ),
        (
            "All sheets",
            "AmountTWD",
            "原幣金額 × 匯率（fx_rates.csv），四捨五入 ROUND_HALF_UP 至兩位小數；不同幣別的原幣金額不可直接相加。",
        ),
    ]
    for offset, (sheet_name, cell_ref, rule) in enumerate(common + rows, 4):
        sheet.cell(row=offset, column=1, value=sheet_name)
        sheet.cell(row=offset, column=2, value=cell_ref)
        sheet.cell(row=offset, column=3, value=rule).alignment = Alignment(wrap_text=True)
    sheet.column_dimensions["A"].width = 20
    sheet.column_dimensions["B"].width = 18
    sheet.column_dimensions["C"].width = 100
    sheet.oddFooter.center.text = FOOTER


def column_rules(sheet_name: str, columns: list[str]) -> list[tuple[str, str, str]]:
    return [
        (sheet_name, f"{get_column_letter(i)}4", f"{c}: {kind(c)}")
        for i, c in enumerate(columns, 1)
    ]


SALES_DETAILS = [
    "OrderID",
    "CustomerID",
    "CustomerName",
    "ProjectID",
    "OrderDate",
    "Status",
    "Currency",
    "NetAmount",
    "NetAmountTWD",
    "CustomerContact",
    "PaymentTerms",
]
EXPENSE_DETAILS = [
    "ExpenseID",
    "EmployeeID",
    "EmployeeName",
    "ProjectID",
    "ExpenseDate",
    "Category",
    "Currency",
    "Amount",
    "AmountTWD",
    "Status",
    "ReceiptID",
    "Description",
    "MissingReceipt",
    "NeedsReview",
]
QUOTE_HEADERS = [
    "QuoteID",
    "CustomerID",
    "ProjectID",
    "QuoteDate",
    "ValidUntil",
    "Status",
    "Currency",
    "Subtotal",
    "DiscountPercent",
    "NetAmount",
    "NetAmountTWD",
]
QUOTE_LINES = ["QuoteID", "ProductID", "ProductName", "Quantity", "UnitPrice", "LineAmount"]
GROUP_SUMMARY = ["Group", "SubGroup", "Currency", "Amount", "AmountTWD", "Count"]
PROJECT_ORDERS = [
    "OrderID",
    "ProjectID",
    "CustomerID",
    "OrderDate",
    "Status",
    "Currency",
    "NetAmount",
    "AmountTWD",
]
PROJECT_INVOICES = [
    "InvoiceID",
    "OrderID",
    "ProjectID",
    "CustomerID",
    "InvoiceDate",
    "DueDate",
    "Currency",
    "Amount",
    "PaidAmount",
    "Balance",
    "BalanceTWD",
]
PROJECT_EXPENSES = [
    "ExpenseID",
    "ProjectID",
    "EmployeeID",
    "ExpenseDate",
    "Status",
    "Currency",
    "Amount",
    "AmountTWD",
]
PROJECT_QUOTES = [
    "QuoteID",
    "ProjectID",
    "CustomerID",
    "QuoteDate",
    "ValidUntil",
    "Status",
    "Currency",
    "NetAmount",
    "AmountTWD",
]
PROJECT_SUMMARY = [
    "ProjectID",
    "OrderTWD",
    "InvoiceBalanceTWD",
    "ApprovedExpenseTWD",
    "AcceptedQuoteTWD",
]
EXCEPTIONS = ["Source", "DocumentID", "ProjectID", "Reason"]
GROUP_RULE = (
    "Summary",
    "A4:F4 / row 5+",
    "彙總表：Group、SubGroup 為分組鍵（依題目要求，例如 Status、Currency、"
    "EmployeeID、Category；不用時留空），Currency 為原幣，Amount 為該組原幣合計，AmountTWD 為該組 TWD 合計，"
    "Count 為筆數；全部 TWD 總額不得混加不同原幣。",
)


def sales() -> Workbook:
    book = Workbook()
    report(book, "Sales Order Report 銷售訂單報表")
    detail(book, "Details", "Sales order details", SALES_DETAILS)
    summary(book, GROUP_SUMMARY, {"B2": "=SUM(Details!I5:I5)"}, "Sales summary")
    instructions(
        book,
        column_rules("Details", SALES_DETAILS)
        + [
            (
                "Summary",
                "B2",
                "公式 =SUM(Details!I5:I{最後資料列})：Details 的 NetAmountTWD 合計；填入資料後把範圍延伸到最後一列，零筆時保留 I5:I5（結果為 0）。",
            ),
            GROUP_RULE,
        ],
    )
    return book


def expense() -> Workbook:
    book = Workbook()
    report(book, "Expense Report 費用報表")
    detail(book, "Details", "Expense details", EXPENSE_DETAILS)
    summary(book, GROUP_SUMMARY, {"B2": "=SUM(Details!I5:I5)"}, "Expense summary")
    detail(
        book, "Exceptions", "Exceptions 例外", ["ExpenseID", "EmployeeID", "ExpenseDate", "Reason"]
    )
    instructions(
        book,
        column_rules("Details", EXPENSE_DETAILS)
        + [
            (
                "Details",
                "M / N",
                "MissingReceipt、NeedsReview 填 true 或 false（布林值或文字 true/false）。",
            ),
            (
                "Summary",
                "B2",
                "公式 =SUM(Details!I5:I{最後資料列})：Details 的 AmountTWD 合計；填入資料後把範圍延伸到最後一列，零筆時保留 I5:I5。",
            ),
            GROUP_RULE,
            (
                "Exceptions",
                "A4:D4 / row 5+",
                "例外清單：原始 ExpenseID、EmployeeID、ExpenseDate 與原因；列入例外的費用仍須保留在 Details。",
            ),
        ],
    )
    return book


def quote() -> Workbook:
    book = Workbook()
    report(book, "Quotation Report 報價報表")
    detail(book, "Quotes", "Quote headers", QUOTE_HEADERS)
    detail(book, "Lines", "Quote line items", QUOTE_LINES)
    summary(book, GROUP_SUMMARY, {"B2": "=SUM(Quotes!K5:K5)"}, "Quote summary")
    instructions(
        book,
        column_rules("Quotes", QUOTE_HEADERS)
        + column_rules("Lines", QUOTE_LINES)
        + [
            ("Quotes", "I", "DiscountPercent 為百分比數字，例如 5 代表 5%。"),
            ("Lines", "A", "每個明細列都要有對應的 QuoteID（與 Quotes 工作表 join）。"),
            (
                "Summary",
                "B2",
                "公式 =SUM(Quotes!K5:K{最後資料列})：Quotes 的 NetAmountTWD 合計；填入資料後把範圍延伸到最後一列，零筆時保留 K5:K5。",
            ),
            GROUP_RULE,
        ],
    )
    return book


def project() -> Workbook:
    book = Workbook()
    report(book, "Project Integration Report 專案整合報表")
    detail(book, "Orders", "Sales orders by project", PROJECT_ORDERS)
    detail(book, "Invoices", "Invoices by project", PROJECT_INVOICES)
    detail(book, "Expenses", "Expenses by project", PROJECT_EXPENSES)
    detail(book, "Quotes", "Quotes by project", PROJECT_QUOTES)
    summary(
        book,
        PROJECT_SUMMARY,
        {"B2": "=SUM(B5:B5)", "C2": "=SUM(C5:C5)", "D2": "=SUM(D5:D5)", "E2": "=SUM(E5:E5)"},
        "Project summary (TWD)",
    )
    detail(book, "Exceptions", "Exceptions 例外與關注項目", EXCEPTIONS)
    instructions(
        book,
        column_rules("Orders", PROJECT_ORDERS)
        + column_rules("Invoices", PROJECT_INVOICES)
        + column_rules("Expenses", PROJECT_EXPENSES)
        + column_rules("Quotes", PROJECT_QUOTES)
        + [
            (
                "Invoices",
                "J / K",
                "Balance = Amount − PaidAmount；BalanceTWD 為 Balance 換算 TWD。",
            ),
            (
                "Summary",
                "A4:E4 / row 5+",
                "每個專案一列：OrderTWD 訂單金額、InvoiceBalanceTWD 發票未收金額、"
                "ApprovedExpenseTWD 核准支出、AcceptedQuoteTWD 已接受報價，皆為 TWD；各來源分開彙總，不互相加總。",
            ),
            (
                "Summary",
                "B2:E2",
                "公式 =SUM(B5:B{最後資料列}) 等，填入資料後延伸到最後一列；沒有資料時保留 B5:B5。",
            ),
            (
                "Exceptions",
                "A4:D4 / row 5+",
                "Source 為 Orders / Invoices / Expenses / Quotes，DocumentID 為來源文件 ID，Reason 為原因（例如 Overdue）。",
            ),
        ],
    )
    return book


FX = (
    "Currency,RateToTWD,Note\r\n"
    "TWD,1,測試固定匯率，非真實行情 (fixed test rate, not a market quote)\r\n"
    "USD,32,測試固定匯率，非真實行情 (fixed test rate, not a market quote)\r\n"
    "EUR,35,測試固定匯率，非真實行情 (fixed test rate, not a market quote)\r\n"
)


def main() -> None:
    for name, build in {
        "sales_template.xlsx": sales,
        "expense_template.xlsx": expense,
        "quote_template.xlsx": quote,
        "project_template.xlsx": project,
    }.items():
        build().save(HERE / name)
        print(f"wrote {HERE / name}")
    (HERE / "fx_rates.csv").write_bytes(FX.encode("utf-8-sig"))
    print(f"wrote {HERE / 'fx_rates.csv'}")


if __name__ == "__main__":
    main()
