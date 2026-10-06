from pathlib import Path
from openpyxl import load_workbook
import csv, json

root = Path.cwd()
out = root / 'output'
print('Workspace:', root)

expected_files = [
    'project_report_2026-09_workflow.xlsx',
    'summary_2026-09.csv',
    'sap_orders_2026-09_all.csv',
    'sap_orders_2026-09_non_cancelled.csv',
    'sap_invoices_2026-09.csv',
    'expenses_2026-09_approved.csv',
    'expenses_2026-09_approved_normalized.csv',
    'quotes_2026-09_accepted.csv',
]
for name in expected_files:
    p = out / name
    assert p.exists(), f'missing {name}'
    print(f'FOUND {name}: {p.stat().st_size} bytes')

def read_csv(name):
    with (out/name).open('r', encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

orders_all = read_csv('sap_orders_2026-09_all.csv')
orders_nc = read_csv('sap_orders_2026-09_non_cancelled.csv')
invoices = read_csv('sap_invoices_2026-09.csv')
expenses = read_csv('expenses_2026-09_approved.csv')
expenses_norm = read_csv('expenses_2026-09_approved_normalized.csv')
quotes = read_csv('quotes_2026-09_accepted.csv')
summary = read_csv('summary_2026-09.csv')

# CSV checks
assert len(orders_all) == 20
assert len(orders_nc) == 18
assert len(invoices) == 13
assert len(expenses) == 15
assert len(expenses_norm) == 15
assert len(quotes) == 4
assert len(summary) == 5
assert sum(1 for r in orders_all if r['Status'].lower() == 'cancelled') == 2
assert all(r['Status'].lower() != 'cancelled' for r in orders_nc)
assert all(r['Status'] == 'Approved' for r in expenses)
assert all(r['Status'] == 'Accepted' for r in quotes)

# ID text and leading zero checks in CSV
assert all(isinstance(r['OrderID'], str) for r in orders_all)
assert any(r['OrderID'].startswith('00') for r in orders_all)
assert all(isinstance(r['InvoiceID'], str) for r in invoices)
assert any(r['InvoiceID'].startswith('00') for r in invoices)

expected_summary = {
    'P001': {'OrderTWD': 693570.00, 'InvoiceBalanceTWD': 386445.00, 'ApprovedExpenseTWD': 18305.00, 'AcceptedQuoteTWD': 331200.00},
    'P002': {'OrderTWD': 516084.50, 'InvoiceBalanceTWD': 156256.00, 'ApprovedExpenseTWD': 9100.00, 'AcceptedQuoteTWD': 97450.85},
    'P003': {'OrderTWD': 150000.00, 'InvoiceBalanceTWD': 0.00, 'ApprovedExpenseTWD': 13037.50, 'AcceptedQuoteTWD': 0.00},
    'P004': {'OrderTWD': 794876.25, 'InvoiceBalanceTWD': 78750.00, 'ApprovedExpenseTWD': 8840.00, 'AcceptedQuoteTWD': 282150.00},
    'P005': {'OrderTWD': 770874.20, 'InvoiceBalanceTWD': 236371.20, 'ApprovedExpenseTWD': 6424.00, 'AcceptedQuoteTWD': 0.00},
}
assert [r['ProjectID'] for r in summary] == list(expected_summary.keys())
for r in summary:
    for col, val in expected_summary[r['ProjectID']].items():
        assert abs(float(r[col]) - val) < 0.01, (r['ProjectID'], col, r[col], val)

# Workbook checks, template has titles in rows 1-3 and headers in row 4.
report = out / 'project_report_2026-09_workflow.xlsx'
wb = load_workbook(report, data_only=True)
expected_sheets = ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']
assert wb.sheetnames == expected_sheets, wb.sheetnames
expected_headers = {
    'Summary': ['ProjectID', 'OrderTWD', 'InvoiceBalanceTWD', 'ApprovedExpenseTWD', 'AcceptedQuoteTWD'],
    'Orders': ['OrderID', 'ProjectID', 'CustomerID', 'OrderDate', 'Status', 'Currency', 'NetAmount', 'AmountTWD'],
    'Invoices': ['InvoiceID', 'OrderID', 'ProjectID', 'CustomerID', 'InvoiceDate', 'DueDate', 'Currency', 'Amount', 'PaidAmount', 'Balance', 'BalanceTWD'],
    'Expenses': ['ExpenseID', 'ProjectID', 'EmployeeID', 'ExpenseDate', 'Status', 'Currency', 'Amount', 'AmountTWD'],
    'Quotes': ['QuoteID', 'ProjectID', 'CustomerID', 'QuoteDate', 'ValidUntil', 'Status', 'Currency', 'NetAmount', 'AmountTWD'],
    'Exceptions': ['Source', 'DocumentID', 'ProjectID', 'Reason'],
}
expected_counts = {'Summary': 5, 'Orders': 18, 'Invoices': 13, 'Expenses': 15, 'Quotes': 4, 'Exceptions': 21}
for sheet, headers in expected_headers.items():
    ws = wb[sheet]
    got_headers = [ws.cell(4, c).value for c in range(1, len(headers)+1)]
    assert got_headers == headers, f'{sheet} headers {got_headers} != {headers}'
    rows = ws.max_row - 4
    assert rows == expected_counts[sheet], f'{sheet} row count {rows} != {expected_counts[sheet]}'
    print(f'{sheet}: headers OK, data rows {rows}')

# workbook value checks and ID stored as text
ws = wb['Summary']
for i in range(5, ws.max_row + 1):
    pid = ws.cell(i, 1).value
    if pid in expected_summary:
        vals = {
            'OrderTWD': ws.cell(i, 2).value,
            'InvoiceBalanceTWD': ws.cell(i, 3).value,
            'ApprovedExpenseTWD': ws.cell(i, 4).value,
            'AcceptedQuoteTWD': ws.cell(i, 5).value,
        }
        for col, val in expected_summary[pid].items():
            assert abs(float(vals[col]) - val) < 0.01, (pid, col, vals[col], val)
for sheet, col in [('Orders', 1), ('Invoices', 1)]:
    ws = wb[sheet]
    vals = [ws.cell(r, col).value for r in range(5, ws.max_row + 1)]
    assert all(isinstance(v, str) for v in vals if v is not None), f'{sheet} IDs not text'
    assert any(v.startswith('00') for v in vals if isinstance(v, str)), f'{sheet} no leading-zero sample'

# verify separate revenue-like columns (orders and accepted quotes not combined)
for r in summary:
    pid = r['ProjectID']
    assert float(r['OrderTWD']) == expected_summary[pid]['OrderTWD']
    assert float(r['AcceptedQuoteTWD']) == expected_summary[pid]['AcceptedQuoteTWD']

print('\nVALIDATION PASSED')
print(json.dumps({
    'report': str(report),
    'files_checked': expected_files,
    'counts': {
        'sap_orders_all': len(orders_all),
        'sap_orders_non_cancelled': len(orders_nc),
        'cancelled_excluded': 2,
        'sap_invoices': len(invoices),
        'approved_expenses': len(expenses),
        'accepted_quotes': len(quotes),
        'summary_projects': len(summary),
        'workbook_exceptions': expected_counts['Exceptions']
    },
    'summary': expected_summary
}, ensure_ascii=False, indent=2))
