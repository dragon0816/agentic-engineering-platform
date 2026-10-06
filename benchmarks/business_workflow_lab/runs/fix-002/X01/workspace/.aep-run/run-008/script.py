from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from urllib.request import build_opener, HTTPCookieProcessor, Request
from urllib.parse import urlencode
from http.cookiejar import CookieJar
import csv, io, json, subprocess, datetime, os, sys
import openpyxl

params=json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
WORKSPACE=Path('.').resolve(); INPUT=WORKSPACE/'input'; OUTPUT=WORKSPACE/'output'; OUTPUT.mkdir(exist_ok=True)
START=params['start_date']; END=params['end_date']; AS_OF=datetime.date.fromisoformat(params['as_of_date'])
month_tag=START[:7]
BASE_URL='http://127.0.0.1:8810'
QUOTE_CLI=r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
EXPENSE_CSV=Path(params.get('expense_csv') or (OUTPUT/f'expenses_{month_tag}_approved.csv'))
if not EXPENSE_CSV.is_absolute(): EXPENSE_CSV=WORKSPACE/EXPENSE_CSV
REPORT_PATH=OUTPUT/(params.get('output_name') or f'project_report_{month_tag}.xlsx')
TWOPLACES=Decimal('0.01')
def money(x): return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def to_twd(amount,currency,fx):
    if currency not in fx: raise KeyError(f'Missing FX rate for {currency}')
    return (Decimal(str(amount))*fx[currency]).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def read_fx(path):
    fx={}
    with path.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f); next(r)
        for row in r:
            if row and row[0]: fx[row[0]]=Decimal(row[1])
    return fx
def read_csv_dict(path):
    with path.open(newline='', encoding='utf-8-sig') as f: return list(csv.DictReader(f))
def write_csv_dict(path, rows, fieldnames):
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        w=csv.DictWriter(f, fieldnames=fieldnames); w.writeheader(); w.writerows(rows)
def sap_opener():
    op=build_opener(HTTPCookieProcessor(CookieJar()))
    op.open(Request(BASE_URL+'/login', data=urlencode({'next':'/orders','sso':'company'}).encode(), method='POST'), timeout=20).read()
    return op
def sap_export(op, endpoint, params, out_path):
    raw=op.open(BASE_URL+endpoint+'?'+urlencode(params), timeout=30).read().decode('utf-8-sig')
    rows=list(csv.DictReader(io.StringIO(raw)))
    write_csv_dict(out_path, rows, list(rows[0].keys()) if rows else [])
    return rows
def quote_list():
    cp=subprocess.run([QUOTE_CLI,'quotes','list','--from',START,'--to',END,'--status','Accepted','--format','json'], capture_output=True, text=True, timeout=60)
    if cp.returncode!=0: raise RuntimeError(f'quote_cli failed {cp.returncode}: {cp.stderr}')
    return json.loads(cp.stdout)['items']
def clear_data(ws):
    if ws.max_row>=5: ws.delete_rows(5, ws.max_row-4)
def set_text(cell,val): cell.value='' if val is None else str(val); cell.number_format='@'
def set_money(cell,val): cell.value=float(money(val)); cell.number_format='0.00'
def set_date_text(cell,val): cell.value='' if val is None else str(val); cell.number_format='yyyy-mm-dd'

fx=read_fx(INPUT/'fx_rates.csv')
op=sap_opener()
orders_all=sap_export(op,'/orders/export.csv',{'from':START,'to':END},OUTPUT/f'sap_orders_{month_tag}_all.csv')
invoices=sap_export(op,'/invoices/export.csv',{'from':START,'to':END},OUTPUT/f'sap_invoices_{month_tag}.csv')
orders=[r for r in orders_all if r.get('Status')!='Cancelled']
write_csv_dict(OUTPUT/f'sap_orders_{month_tag}_non_cancelled.csv', orders, ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount'])
if not EXPENSE_CSV.exists():
    raise FileNotFoundError(f'Expense Viewer export not found: {EXPENSE_CSV}. Export Approved expenses for {START}..{END} from the registered desktop app before running this workflow.')
expenses=[r for r in read_csv_dict(EXPENSE_CSV) if START<=r.get('ExpenseDate','')<=END and r.get('Status')=='Approved']
write_csv_dict(OUTPUT/f'expenses_{month_tag}_approved_normalized.csv', expenses, ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt'])
quotes=quote_list()
write_csv_dict(OUTPUT/f'quotes_{month_tag}_accepted.csv', quotes, ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount'])

wb=openpyxl.load_workbook(INPUT/'project_template.xlsx')
for s in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']: clear_data(wb[s])
rep=wb['Report']; rep['B3']=START; rep['D3']=END
rep['B4']='SAP Sales Orders 非 Cancelled；SAP Invoices；Expense Approved；Quote Accepted；依 ProjectID 彙總，訂單與報價分開不相加。'
rep['B5']=datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat(timespec='seconds')
rep['B7']='Lab ERP SAP-like web; Expense Viewer desktop export; quote_cli; fx_rates.csv'
exceptions=[]; summary={}
def add(project,key,amount):
    if not project: return
    summary.setdefault(project, {'OrderTWD':Decimal('0'),'InvoiceBalanceTWD':Decimal('0'),'ApprovedExpenseTWD':Decimal('0'),'AcceptedQuoteTWD':Decimal('0')})
    summary[project][key]+=Decimal(amount)
ws=wb['Orders']; row=5
for r in orders:
    amt=money(r['NetAmount']); twd=to_twd(amt,r['Currency'],fx); vals=[r['OrderID'],r['ProjectID'],r['CustomerID'],r['OrderDate'],r['Status'],r['Currency'],amt,twd]
    for col,val in enumerate(vals,1):
        c=ws.cell(row,col); (set_text if col in [1,2,3,5,6] else set_date_text if col==4 else set_money)(c,val)
    if r['ProjectID']: add(r['ProjectID'],'OrderTWD',twd)
    else: exceptions.append(['Orders',r['OrderID'],'','Missing ProjectID; not included in project summary'])
    row+=1
ws=wb['Invoices']; row=5
for r in invoices:
    bal=money(r['Balance']); btwd=to_twd(bal,r['Currency'],fx); vals=[r['InvoiceID'],r['OrderID'],r['ProjectID'],r['CustomerID'],r['InvoiceDate'],r['DueDate'],r['Currency'],money(r['Amount']),money(r['PaidAmount']),bal,btwd]
    for col,val in enumerate(vals,1):
        c=ws.cell(row,col); (set_text if col in [1,2,3,4,7] else set_date_text if col in [5,6] else set_money)(c,val)
    if r['ProjectID']: add(r['ProjectID'],'InvoiceBalanceTWD',btwd)
    else: exceptions.append(['Invoices',r['InvoiceID'],'','Missing ProjectID; not included in project summary'])
    if bal>0 and datetime.date.fromisoformat(r['DueDate'])<AS_OF: exceptions.append(['Invoices',r['InvoiceID'],r['ProjectID'],f'Overdue unpaid balance as of {AS_OF.isoformat()}'])
    row+=1
ws=wb['Expenses']; row=5
for r in expenses:
    amt=money(r['Amount']); twd=to_twd(amt,r['Currency'],fx); vals=[r['ExpenseID'],r['ProjectID'],r['EmployeeID'],r['ExpenseDate'],r['Status'],r['Currency'],amt,twd]
    for col,val in enumerate(vals,1):
        c=ws.cell(row,col); (set_text if col in [1,2,3,5,6] else set_date_text if col==4 else set_money)(c,val)
    if r['ProjectID']: add(r['ProjectID'],'ApprovedExpenseTWD',twd)
    else: exceptions.append(['Expenses',r['ExpenseID'],'','Missing ProjectID; not included in project summary'])
    if str(r.get('MissingReceipt','')).lower()=='true': exceptions.append(['Expenses',r['ExpenseID'],r['ProjectID'],'Missing receipt'])
    row+=1
ws=wb['Quotes']; row=5
for r in quotes:
    amt=money(r['NetAmount']); twd=to_twd(amt,r['Currency'],fx); vals=[r['QuoteID'],r['ProjectID'],r['CustomerID'],r['QuoteDate'],r['ValidUntil'],r['Status'],r['Currency'],amt,twd]
    for col,val in enumerate(vals,1):
        c=ws.cell(row,col); (set_text if col in [1,2,3,6,7] else set_date_text if col in [4,5] else set_money)(c,val)
    if r['ProjectID']: add(r['ProjectID'],'AcceptedQuoteTWD',twd)
    else: exceptions.append(['Quotes',r['QuoteID'],'','Missing ProjectID; not included in project summary'])
    row+=1
ws=wb['Summary']; row=5
for project in sorted(summary):
    vals=[project, summary[project]['OrderTWD'], summary[project]['InvoiceBalanceTWD'], summary[project]['ApprovedExpenseTWD'], summary[project]['AcceptedQuoteTWD']]
    for col,val in enumerate(vals,1): (set_text if col==1 else set_money)(ws.cell(row,col),val)
    row+=1
last=max(5,row-1)
for col in 'BCDE': ws[f'{col}2']=f'=SUM({col}5:{col}{last})'; ws[f'{col}2'].number_format='0.00'
ws=wb['Exceptions']; row=5
for ex in exceptions:
    for col,val in enumerate(ex,1): set_text(ws.cell(row,col),val)
    row+=1
rep['B9']=f'Rows used: orders {len(orders)} (excluded Cancelled {len(orders_all)-len(orders)}), invoices {len(invoices)}, expenses {len(expenses)}, quotes {len(quotes)}. Exceptions {len(exceptions)}. Blank ProjectID rows are listed in Exceptions and excluded from Summary.'
wb.save(REPORT_PATH)
with (OUTPUT/f'summary_{month_tag}.csv').open('w', newline='', encoding='utf-8-sig') as f:
    w=csv.writer(f); w.writerow(['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'])
    for p in sorted(summary): w.writerow([p]+[str(summary[p][k].quantize(TWOPLACES)) for k in ['OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']])
# Reopen checks
wb2=openpyxl.load_workbook(REPORT_PATH, data_only=False)
headers={'Orders':['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],'Invoices':['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],'Expenses':['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],'Quotes':['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],'Summary':['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],'Exceptions':['Source','DocumentID','ProjectID','Reason']}
for s,h in headers.items():
    assert [wb2[s].cell(4,c).value for c in range(1,len(h)+1)]==h, s
counts={s:max(0, wb2[s].max_row-4) for s in headers}
assert counts['Orders']==len(orders) and all(wb2['Orders'].cell(r,5).value!='Cancelled' for r in range(5,wb2['Orders'].max_row+1))
assert counts['Invoices']==len(invoices)
assert counts['Expenses']==len(expenses) and all(wb2['Expenses'].cell(r,5).value=='Approved' for r in range(5,wb2['Expenses'].max_row+1))
assert counts['Quotes']==len(quotes) and all(wb2['Quotes'].cell(r,6).value=='Accepted' for r in range(5,wb2['Quotes'].max_row+1))
assert isinstance(wb2['Orders']['A5'].value,str) and isinstance(wb2['Invoices']['A5'].value,str)
print(json.dumps({'report':str(REPORT_PATH),'counts':counts,'cancelled_excluded':len(orders_all)-len(orders),'projects':sorted(summary),'summary_csv':str(OUTPUT/f'summary_{month_tag}.csv')}, ensure_ascii=False, indent=2))