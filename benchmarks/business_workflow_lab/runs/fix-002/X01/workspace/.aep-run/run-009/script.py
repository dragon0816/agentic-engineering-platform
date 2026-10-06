from pathlib import Path
import csv, json, math
from openpyxl import load_workbook

root = Path.cwd()
out = root / 'output'
print('Workspace:', root)
print('Output files present:')
for p in sorted(out.glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

expected_files = [
    'project_report_2026-09_workflow.xlsx',
    'summary_2026-09.csv',
    'sap_orders_2026-09_all.csv',
    'sap_orders_2026-09_non_cancelled.csv',
    'sap_invoices_2026-09.csv',
    'expenses_2026-09_approved.csv',
    'expenses_2026-09_approved_normalized.csv',
    'quotes_2026-09_accepted.csv',
]
missing = [f for f in expected_files if not (out/f).exists()]
assert not missing, f'Missing expected files: {missing}'

def read_csv(name):
    path = out / name
    with path.open('r', encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    print('\nCSV', name)
    print(' headers:', rows[0].keys() if rows else 'NO ROWS')
    print(' rows:', len(rows))
    if rows:
        print(' sample:', json.dumps(rows[0], ensure_ascii=False))
    return rows

orders_all = read_csv('sap_orders_2026-09_all.csv')
orders_nc = read_csv('sap_orders_2026-09_non_cancelled.csv')
invoices = read_csv('sap_invoices_2026-09.csv')
expenses = read_csv('expenses_2026-09_approved.csv')
expenses_norm = read_csv('expenses_2026-09_approved_normalized.csv')
quotes = read_csv('quotes_2026-09_accepted.csv')
summary = read_csv('summary_2026-09.csv')

# Requirement checks from the request and run result
assert len(orders_all) == 20, f'orders_all row count expected 20, got {len(orders_all)}'
assert len(orders_nc) == 18, f'non-cancelled orders expected 18, got {len(orders_nc)}'
assert len(invoices) == 13, f'invoices expected 13, got {len(invoices)}'
assert len(expenses) == 15, f'approved expenses expected 15, got {len(expenses)}'
assert len(expenses_norm) == 15, f'normalized approved expenses expected 15, got {len(expenses_norm)}'
assert len(quotes) == 4, f'accepted quotes expected 4, got {len(quotes)}'
assert len(summary) == 5, f'summary projects expected 5, got {len(summary)}'

# Status filters
cancelled = [r for r in orders_nc if r.get('Status','').strip().lower() == 'cancelled']
assert not cancelled, f'Cancelled orders present in non-cancelled output: {cancelled[:3]}'
assert sum(1 for r in orders_all if r.get('Status','').strip().lower() == 'cancelled') == 2, 'cancelled excluded count is not 2'
assert all(r.get('Status','').strip().lower() == 'approved' for r in expenses), 'non-Approved expense found'
assert all(r.get('Status','').strip().lower() == 'accepted' for r in quotes), 'non-Accepted quote found'

# ProjectID join/summarization checks: every transactional output has ProjectID and summary only uses ProjectID keys.
for name, rows in [('orders_nc', orders_nc), ('invoices', invoices), ('expenses_norm', expenses_norm), ('quotes', quotes)]:
    assert rows, f'{name} empty'
    assert 'ProjectID' in rows[0], f'{name} missing ProjectID column'

expected_projects = ['P001','P002','P003','P004','P005']
projects = [r['ProjectID'] for r in summary]
assert projects == expected_projects, f'summary projects mismatch: {projects}'

# Summary values from prior successful run; also proves orders and quotes are separate columns.
expected_summary = {
    'P001': (693570.00, 386445.00, 18305.00, 331200.00),
    'P002': (516084.50, 156256.00, 9100.00, 97450.85),
    'P003': (150000.00, 0.00, 13037.50, 0.00),
    'P004': (794876.25, 78750.00, 8840.00, 282150.00),
    'P005': (770874.20, 236371.20, 6424.00, 0.00),
}
for r in summary:
    vals = (float(r['OrderTWD']), float(r['InvoiceBalanceTWD']), float(r['ApprovedExpenseTWD']), float(r['AcceptedQuoteTWD']))
    exp = expected_summary[r['ProjectID']]
    assert all(abs(a-b) < 0.01 for a,b in zip(vals, exp)), f'summary mismatch for {r["ProjectID"]}: {vals} != {exp}'

# Workbook check
report = out / 'project_report_2026-09_workflow.xlsx'
wb = load_workbook(report, data_only=True)
print('\nWorkbook', report.name)
print(' sheets:', wb.sheetnames)
expected_sheets = ['Summary','Orders','Invoices','Expenses','Quotes','Exceptions']
for sheet in expected_sheets:
    assert sheet in wb.sheetnames, f'missing workbook sheet: {sheet}'
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    data_rows = ws.max_row - 1
    print(f' {sheet}: headers={headers}; data_rows={data_rows}')

assert wb['Summary'].max_row - 1 == 5, 'workbook Summary row count mismatch'
assert wb['Orders'].max_row - 1 == 18, 'workbook Orders row count mismatch'
assert wb['Invoices'].max_row - 1 == 13, 'workbook Invoices row count mismatch'
assert wb['Expenses'].max_row - 1 == 15, 'workbook Expenses row count mismatch'
assert wb['Quotes'].max_row - 1 == 4, 'workbook Quotes row count mismatch'
assert wb['Exceptions'].max_row - 1 == 21, 'workbook Exceptions row count mismatch'

# ID text preservation in workbook where columns exist.
for sheet, id_cols in [('Orders', ['OrderID']), ('Invoices', ['InvoiceID'])]:
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    for col in id_cols:
        if col in headers:
            idx = headers.index(col) + 1
            vals = [ws.cell(row=i, column=idx).value for i in range(2, ws.max_row+1)]
            assert all(isinstance(v, str) for v in vals if v is not None), f'{sheet}.{col} not all text'
            print(f' {sheet}.{col}: text ids confirmed; sample={vals[:3]}')

print('\nVALIDATION PASSED')
print(json.dumps({
    'files_checked': expected_files,
    'counts': {
        'orders_all': len(orders_all),
        'orders_non_cancelled': len(orders_nc),
        'invoices': len(invoices),
        'expenses_approved': len(expenses),
        'quotes_accepted': len(quotes),
        'summary_projects': len(summary),
        'workbook_exceptions': wb['Exceptions'].max_row - 1,
    },
    'report': str(report),
}, ensure_ascii=False, indent=2))
