from pathlib import Path
from openpyxl import load_workbook
import json
root=Path.cwd(); out=root/'output'; report=out/'project_report_2026-09_workflow.xlsx'
print('Inspecting', report)
wb=load_workbook(report, data_only=True)
for sheet in ['Report','Summary','Orders','Invoices','Expenses','Quotes','Exceptions','Instructions']:
    ws=wb[sheet]
    print('\nSHEET', sheet, 'max_row', ws.max_row, 'max_col', ws.max_column)
    for r in range(1, min(ws.max_row, 12)+1):
        vals=[ws.cell(r,c).value for c in range(1, min(ws.max_column, 12)+1)]
        print(r, vals)
