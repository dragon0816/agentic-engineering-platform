from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
for sh in ['Instructions','Report','Summary']:
    ws=wb[sh]
    print('---',sh,'---')
    for row in ws.iter_rows(values_only=True): print(row)
