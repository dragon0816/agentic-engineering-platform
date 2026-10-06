from pathlib import Path
from urllib.request import build_opener, HTTPCookieProcessor, Request
from urllib.parse import urlencode, urljoin
from http.cookiejar import CookieJar
from html.parser import HTMLParser
import csv, io, re
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

WORKSPACE = Path.cwd(); INPUT = WORKSPACE/'input'; OUTPUT = WORKSPACE/'output'; OUTPUT.mkdir(exist_ok=True)
BASE='http://127.0.0.1:8810/'; FROM_DATE='2026-09-01'; TO_DATE='2026-09-30'; STATUS='Open'; NOW='2026-10-06T00:00:00+08:00'
PLACEHOLDER='N/A'

class ERPClient:
    def __init__(self, base):
        self.base=base; self.cj=CookieJar(); self.opener=build_opener(HTTPCookieProcessor(self.cj)); self.relogin_count=0; self.max_relogin=3
        self.login('/orders')
    def _open(self, path, data=None):
        url=urljoin(self.base,path)
        req=Request(url, data=urlencode(data).encode('utf-8') if data is not None else None, method='POST' if data is not None else 'GET')
        resp=self.opener.open(req, timeout=10); raw=resp.read(); text=raw.decode('utf-8', errors='replace')
        return text, resp.geturl(), resp.headers.get('content-type','')
    def is_login(self, text, final_url):
        return '<h1>Sign in</h1>' in text or "id='login-form'" in text or '/login' in final_url
    def login(self, next_path):
        text, final, ctype = self._open('/login', {'next':next_path, 'sso':'company'})
        if self.is_login(text, final): raise RuntimeError('SSO login did not leave login page')
    def get(self, path):
        text, final, ctype = self._open(path)
        if self.is_login(text, final):
            if self.relogin_count >= self.max_relogin:
                raise RuntimeError(f'恢復三次仍失敗：仍回到登入頁；last_url={final}')
            self.relogin_count += 1
            print(f'Re-login recovery {self.relogin_count}/{self.max_relogin} for {path}')
            self.login(path)
            text, final, ctype = self._open(path)
            if self.is_login(text, final):
                raise RuntimeError(f'重新登入後仍回到登入頁；last_url={final}')
        return text, final, ctype

class HeaderParser(HTMLParser):
    def __init__(self):
        super().__init__(); self.in_table=False; self.in_tr=False; self.cell=None; self.buf=[]; self.cur=[]; self.rows=[]
    def handle_starttag(self, tag, attrs):
        attrs=dict(attrs)
        if tag=='table' and attrs.get('id')=='header': self.in_table=True
        elif self.in_table and tag=='tr': self.in_tr=True; self.cur=[]
        elif self.in_table and self.in_tr and tag in ('th','td'): self.cell=tag; self.buf=[]
    def handle_data(self, data):
        if self.in_table and self.in_tr and self.cell: self.buf.append(data)
    def handle_endtag(self, tag):
        if self.in_table and self.in_tr and self.cell==tag:
            self.cur.append(''.join(self.buf).strip()); self.cell=None
        elif self.in_table and tag=='tr':
            if len(self.cur)>=2: self.rows.append(self.cur[:2])
            self.in_tr=False
        elif self.in_table and tag=='table': self.in_table=False
class ResultCountParser(HTMLParser):
    def __init__(self): super().__init__(); self.on=False; self.text=[]
    def handle_starttag(self, tag, attrs):
        if dict(attrs).get('id')=='result-count': self.on=True
    def handle_data(self, data):
        if self.on: self.text.append(data)
    def handle_endtag(self, tag):
        if self.on and tag=='p': self.on=False

def parse_header(html):
    p=HeaderParser(); p.feed(html); return {k:v for k,v in p.rows}
def parse_total(html):
    p=ResultCountParser(); p.feed(html); s=' '.join(''.join(p.text).split()); m=re.search(r'Total:\s*(\d+)\s*rows', s); return int(m.group(1)) if m else None
def txt(v):
    return v if v not in (None,'') else PLACEHOLDER

# Source/input inspection printout for traceability
print('Inputs:', [p.name for p in sorted(INPUT.glob('*'))])
wb0=load_workbook(INPUT/'sales_template.xlsx', data_only=False)
print('Template sheets:', wb0.sheetnames)
print('Details headers:', [wb0['Details'].cell(4,c).value for c in range(1,12)])
with (INPUT/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    print('FX CSV sample:', [row for _,row in zip(range(4), csv.reader(f))])

client=ERPClient(BASE)
query=urlencode({'from':FROM_DATE,'to':TO_DATE,'status':STATUS,'sort':'OrderDate','dir':'asc'})
list_html,_,_=client.get('/orders?'+query)
total=parse_total(list_html); print('Filtered list total from page:', total)
csv_text,_,ctype=client.get('/orders/export.csv?'+query)
if not ctype.lower().startswith('text/csv'): raise RuntimeError('Export did not return CSV: '+ctype)
orders=[]
for row in csv.DictReader(io.StringIO(csv_text.lstrip('\ufeff'))):
    if any((v or '').strip() for v in row.values()): orders.append({k.lstrip('\ufeff'):(v or '') for k,v in row.items()})
print('CSV exported rows:', len(orders)); print('First CSV rows:', orders[:3])
if total is not None and len(orders)!=total: raise RuntimeError(f'CSV row count {len(orders)} != list total {total}')
seen=set()
for row in orders:
    oid=row['OrderID']
    if oid in seen: raise RuntimeError('Duplicate OrderID from export: '+oid)
    seen.add(oid)
    if not (FROM_DATE <= row['OrderDate'] <= TO_DATE and row['Status']==STATUS): raise RuntimeError('Out-of-filter row: '+repr(row))

details={}
for row in orders:
    oid=row['OrderID']
    html,_,_=client.get('/orders/'+oid)
    hdr=parse_header(html)
    if hdr.get('OrderID') != oid: raise RuntimeError(f'Detail mismatch {oid}: {hdr.get("OrderID")}')
    details[oid]=hdr
print('Detail pages fetched:', len(details), 're-login recoveries:', client.relogin_count)

rates={}
with (INPUT/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f); next(r)
    for row in r:
        if len(row)>=2 and row[0]: rates[row[0]]=Decimal(row[1])
print('FX rates used:', rates)

wb=load_workbook(INPUT/'sales_template.xlsx')
wsr=wb['Report']; wsd=wb['Details']; wss=wb['Summary']
if wsd.max_row>=5: wsd.delete_rows(5, wsd.max_row-4)
if wss.max_row>=5: wss.delete_rows(5, wss.max_row-4)
wsr['B3']=FROM_DATE; wsr['D3']=TO_DATE
wsr['B4']='Sales Orders: Status=Open; OrderDate 2026-09-01..2026-09-30; sorted by OrderDate asc; detail CustomerContact/PaymentTerms fetched per OrderID'
wsr['B5']=NOW; wsr['B7']='Lab ERP (synthetic) Sales Orders'; wsr['D7']=NOW
missing=[]
summary={}
for idx,row in enumerate(orders, start=5):
    oid=row['OrderID']; hdr=details[oid]; cur=row['Currency']; amount=Decimal(row['NetAmount'])
    twd=(amount*rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    proj=txt(row.get('ProjectID','')); contact=txt(hdr.get('CustomerContact','')); terms=txt(hdr.get('PaymentTerms',''))
    if proj==PLACEHOLDER: missing.append(f'{oid}: ProjectID')
    if contact==PLACEHOLDER: missing.append(f'{oid}: CustomerContact')
    if terms==PLACEHOLDER: missing.append(f'{oid}: PaymentTerms')
    vals=[oid, row['CustomerID'], row['CustomerName'], proj, row['OrderDate'], row['Status'], cur, float(amount), float(twd), contact, terms]
    for col,val in enumerate(vals,1):
        c=wsd.cell(idx,col,val)
        c.number_format = '0.00' if col in (8,9) else '@'
    key=(row['Status'], cur)
    summary.setdefault(key, {'amount':Decimal('0'), 'twd':Decimal('0'), 'count':0})
    summary[key]['amount'] += amount; summary[key]['twd'] += twd; summary[key]['count'] += 1
for r_idx, ((status,cur), agg) in enumerate(sorted(summary.items()), start=5):
    vals=['Status', status, cur, float(agg['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), float(agg['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), agg['count']]
    for col,val in enumerate(vals,1):
        c=wss.cell(r_idx,col,val)
        c.number_format='0.00' if col in (4,5) else ('0' if col==6 else '@')
last_detail=4+len(orders)
wss['B2']=f'=SUM(Details!I5:I{last_detail})' if orders else '=SUM(Details!I5:I5)'; wss['B2'].number_format='0.00'
wsr['B9']=f'{len(orders)} orders exported; {len(details)} detail pages fetched; session re-login recoveries: {client.relogin_count}; no duplicate OrderID. Placeholder "{PLACEHOLDER}" used for missing text fields: {len(missing)} ({"; ".join(missing) if missing else "none"}).'

outfile=OUTPUT/'sales_2026_09_open_orders.xlsx'
wb.save(outfile)
print('Wrote', outfile)

# Verification from saved file
wb2=load_workbook(outfile, data_only=False); wsd2=wb2['Details']; wss2=wb2['Summary']; wsr2=wb2['Report']
expected_headers=['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert [wsd2.cell(4,c).value for c in range(1,12)] == expected_headers
rows=list(wsd2.iter_rows(min_row=5, max_row=wsd2.max_row, max_col=11, values_only=False))
assert len(rows)==len(orders), (len(rows), len(orders))
ids=[r[0].value for r in rows]
assert ids==[o['OrderID'] for o in orders] and len(ids)==len(set(ids))
for r in rows:
    vals=[c.value for c in r]
    assert isinstance(vals[0], str) and vals[0].startswith('00')
    assert all(vals[i] not in (None,'') for i in [0,1,2,3,4,5,6,9,10]), vals
    assert FROM_DATE <= vals[4] <= TO_DATE and vals[5]==STATUS, vals
    assert isinstance(vals[7], (int,float)) and isinstance(vals[8], (int,float)), vals
    expected=(Decimal(str(vals[7]))*rates[vals[6]]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert Decimal(str(vals[8])).quantize(Decimal('0.01')) == expected, (vals, expected)
assert wsr2['B3'].value==FROM_DATE and wsr2['D3'].value==TO_DATE
assert wss2['B2'].value==f'=SUM(Details!I5:I{last_detail})'
summary_rows=list(wss2.iter_rows(min_row=5, max_row=wss2.max_row, max_col=6, values_only=True))
assert len(summary_rows)==len(summary) and sum(r[5] for r in summary_rows)==len(orders)
assert all(isinstance(r[3], (int,float)) and isinstance(r[4], (int,float)) for r in summary_rows)
print('VERIFIED saved workbook requirements')
print('Details row count:', len(rows))
print('OrderIDs:', ids)
print('Missing placeholders:', missing)
print('Summary rows:', summary_rows)
print('Summary!B2 formula:', wss2['B2'].value)
print('Output file:', outfile)
