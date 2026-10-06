import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path('.')
print('cwd', Path.cwd())
for p in sorted(Path('input').glob('*')):
    print('INPUT', p, 'size', p.stat().st_size)
    if p.suffix.lower()=='.csv':
        with p.open(newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f)
            for i,row in zip(range(5), r): print('CSVROW', row)
    elif p.suffix.lower()=='.md':
        txt=p.read_text(encoding='utf-8')
        print('MD HEAD:\n', txt[:2000])
    elif p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print('XLS sheets', wb.sheetnames)
        for sh in wb.sheetnames:
            ws2=wb[sh]
            print('SHEET', sh, ws2.max_row, ws2.max_column)
            for row in ws2.iter_rows(min_row=1, max_row=min(ws2.max_row,5), values_only=True):
                print(row)
