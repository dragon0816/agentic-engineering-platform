from pathlib import Path
import subprocess, sys, json, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'quote_template.xlsx'
OUT_MD = OUTPUT / 'findings.md'
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')

print('Inputs really present:')
for p in sorted(INPUT.glob('*')):
    print(f'- {p.name} ({p.stat().st_size} bytes)')
print('Task period:', FROM_DATE, 'to', TO_DATE, 'base date:', BASE_DATE.isoformat())

# Locate and use the synthetic quote CLI described by input/quote_cli_manual.md.
def find_cli():
    for parent in [WORKSPACE] + list(WORKSPACE.parents):
        cand = parent / 'quote_cli_mock' / 'run_quote_cli.py'
        if cand.exists():
            return cand
    raise FileNotFoundError('Could not locate quote_cli_mock/run_quote_cli.py in workspace parents')
CLI = find_cli()
print('Using quote CLI:', CLI)

def run_cli(args, retries=2):
    cmd = [sys.executable, str(CLI)] + args
    for attempt in range(retries + 1):
        cp = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
        if cp.returncode == 0:
            return cp.stdout
        if cp.returncode == 4 and attempt < retries:
            continue
        raise RuntimeError(f'CLI failed rc={cp.returncode} args={args}\nSTDERR={cp.stderr}\nSTDOUT={cp.stdout}')

# FX rates; only Currency and RateToTWD columns are required.
rates = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', {k: str(v) for k, v in rates.items()})

# Fetch all September 2026 quote headers, then full details for each quote.
listed = json.loads(run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json']))
headers = listed.get('items', [])
print('List schema_version:', listed.get('schema_version'), 'count:', listed.get('count'), 'items:', len(headers))
quotes = []
all_lines = []
for h in headers:
    qid = h['QuoteID']
    doc = json.loads(run_cli(['quotes', 'get', qid, '--format', 'json']))
    header = doc['header']
    line_items = doc.get('line_items', [])
    quotes.append(header)
    for ln in line_items:
        rec = dict(ln)
        rec['QuoteID'] = qid
        all_lines.append(rec)
print('Fetched quote documents:', len(quotes), 'line items:', len(all_lines))
print('First quote header sample:', quotes[0] if quotes else None)
print('First line sample:', all_lines[0] if all_lines else None)

TWOPLACES = Decimal('0.01')
def D(x): return Decimal(str(x))
def money(x): return D(x).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def calc_net(subtotal, discount_percent):
    return (D(subtotal) * (Decimal('1') - D(discount_percent) / Decimal('100'))).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def parse_date(s): return date.fromisoformat(str(s))
def text_or_blank(x): return '' if x is None else str(x)

def fmt_money(d): return f'{money(d):,.2f}'

# Check calculations.
lines_by_q = defaultdict(list)
line_amount_issues = []
for ln in all_lines:
    lines_by_q[ln['QuoteID']].append(ln)
    expected_line = money(D(ln['Quantity']) * D(ln['UnitPrice']))
    actual_line = money(ln['LineAmount'])
    if expected_line != actual_line:
        line_amount_issues.append({'QuoteID': ln['QuoteID'], 'ProductID': ln.get('ProductID'), 'expected': expected_line, 'actual': actual_line})
subtotal_issues = []
net_issues = []
expired_unaccepted = []
for q in quotes:
    qid = q['QuoteID']
    line_sum = sum((money(ln['LineAmount']) for ln in lines_by_q[qid]), Decimal('0.00')).quantize(TWOPLACES)
    subtotal = money(q['Subtotal'])
    if line_sum != subtotal:
        subtotal_issues.append({'QuoteID': qid, 'line_sum': line_sum, 'subtotal': subtotal, 'diff': line_sum - subtotal})
    expected_net = calc_net(q['Subtotal'], q['DiscountPercent'])
    actual_net = money(q['NetAmount'])
    if expected_net != actual_net:
        net_issues.append({'QuoteID': qid, 'expected': expected_net, 'actual': actual_net, 'diff': actual_net - expected_net})
    if parse_date(q['ValidUntil']) < BASE_DATE and q['Status'] != 'Accepted':
        expired_unaccepted.append(q)

# Fill workbook based on template.
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)
now_iso = datetime.now().astimezone().isoformat(timespec='seconds')
report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = 'QuoteDate between 2026-09-01 and 2026-09-30 inclusive; all quotes retained regardless of expiry.'
report['B5'] = now_iso
report['B7'] = 'quote_cli synthetic quote system; fx_rates.csv'
report['D7'] = now_iso
report['B9'] = '; '.join([
    f'{len(line_amount_issues)} line amount mismatch(es)',
    f'{len(subtotal_issues)} subtotal mismatch(es)',
    f'{len(net_issues)} net amount mismatch(es)',
    f'{len(expired_unaccepted)} quote(s) ValidUntil < 2026-09-30 and not Accepted',
    'blank source IDs are kept blank'
])

def clear_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)
def set_money_cell(cell, val):
    cell.value = float(money(val))
    cell.number_format = '0.00'
def set_number_cell(cell, val):
    dec = D(val)
    cell.value = int(dec) if dec == dec.to_integral_value() else float(dec)

qws = wb['Quotes']
clear_data(qws, 5)
quote_rows = []
for q in sorted(quotes, key=lambda r: r['QuoteID']):
    rate = rates.get(q['Currency'])
    if rate is None:
        raise KeyError(f'Missing FX rate for {q["Currency"]}')
    net_twd = (money(q['NetAmount']) * rate).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    quote_rows.append([
        str(q['QuoteID']), text_or_blank(q.get('CustomerID')), text_or_blank(q.get('ProjectID')),
        str(q['QuoteDate']), str(q['ValidUntil']), str(q['Status']), str(q['Currency']),
        money(q['Subtotal']), D(q['DiscountPercent']), money(q['NetAmount']), net_twd
    ])
for r_idx, row in enumerate(quote_rows, 5):
    for c_idx, value in enumerate(row, 1):
        cell = qws.cell(r_idx, c_idx)
        if c_idx in (1,2,3,4,5,6,7):
            cell.value = value
            if c_idx in (1,2,3): cell.number_format = '@'
            if c_idx in (4,5): cell.number_format = 'yyyy-mm-dd'
        elif c_idx in (8,10,11):
            set_money_cell(cell, value)
        elif c_idx == 9:
            set_number_cell(cell, value)

lws = wb['Lines']
clear_data(lws, 5)
line_rows = []
for ln in sorted(all_lines, key=lambda r: (r['QuoteID'], r['ProductID'], r.get('ProductName',''))):
    line_rows.append([str(ln['QuoteID']), str(ln['ProductID']), str(ln['ProductName']), D(ln['Quantity']), money(ln['UnitPrice']), money(ln['LineAmount'])])
for r_idx, row in enumerate(line_rows, 5):
    for c_idx, value in enumerate(row, 1):
        cell = lws.cell(r_idx, c_idx)
        if c_idx in (1,2,3):
            cell.value = value
            if c_idx in (1,2): cell.number_format = '@'
        elif c_idx == 4:
            set_number_cell(cell, value)
        else:
            set_money_cell(cell, value)

sws = wb['Summary']
clear_data(sws, 5)
last_quote_row = 4 + len(quote_rows)
sws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})' if quote_rows else '=SUM(Quotes!K5:K5)'
sws['B2'].number_format = '0.00'
summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0})
for qrow in quote_rows:
    key = (qrow[5], qrow[6])
    summary[key]['amount'] += money(qrow[9])
    summary[key]['amount_twd'] += money(qrow[10])
    summary[key]['count'] += 1
for r_idx, ((status, curr), vals) in enumerate(sorted(summary.items()), 5):
    sws.cell(r_idx, 1).value = 'Status'
    sws.cell(r_idx, 2).value = status
    sws.cell(r_idx, 3).value = curr
    set_money_cell(sws.cell(r_idx, 4), vals['amount'])
    set_money_cell(sws.cell(r_idx, 5), vals['amount_twd'])
    sws.cell(r_idx, 6).value = vals['count']

for wsname in ['Quotes', 'Lines', 'Summary']:
    ws = wb[wsname]
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = max(len(str(ws.cell(row, col).value or '')) for row in range(1, min(ws.max_row, 1000) + 1))
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 45)
wb.save(OUT_XLSX)

# findings.md
md = []
md.append('# Quote Findings — September 2026')
md.append('')
md.append(f'- Period fetched: `{FROM_DATE}` to `{TO_DATE}` inclusive')
md.append(f'- Baseline date for expiry check: `{BASE_DATE.isoformat()}`')
md.append(f'- Source: quote_cli synthetic quote system; FX: `input/fx_rates.csv`')
md.append(f'- Quotes retained in workbook: **{len(quotes)}** (no quote removed because of expiry)')
md.append(f'- Line items retained in workbook: **{len(all_lines)}**')
missing_project = sum(1 for q in quotes if q.get('ProjectID') in (None, ''))
if missing_project:
    md.append(f'- Source ProjectID blank/missing values kept blank in workbook: **{missing_project}**')
md.append('')
md.append('## Consistency checks')
md.append('')
md.append(f'- LineAmount = Quantity × UnitPrice mismatches: **{len(line_amount_issues)}**')
if line_amount_issues:
    md += ['| QuoteID | ProductID | Expected LineAmount | Actual LineAmount |', '|---|---:|---:|---:|']
    for i in line_amount_issues:
        md.append(f"| {i['QuoteID']} | {i['ProductID']} | {fmt_money(i['expected'])} | {fmt_money(i['actual'])} |")
    md.append('')
md.append(f'- Subtotal vs. sum(LineAmount) mismatches: **{len(subtotal_issues)}**')
if subtotal_issues:
    md += ['| QuoteID | Sum of LineAmount | Header Subtotal | Difference |', '|---|---:|---:|---:|']
    for i in subtotal_issues:
        md.append(f"| {i['QuoteID']} | {fmt_money(i['line_sum'])} | {fmt_money(i['subtotal'])} | {fmt_money(i['diff'])} |")
    md.append('')
md.append(f'- NetAmount = Subtotal × (1 − DiscountPercent / 100), ROUND_HALF_UP 2 decimals mismatches: **{len(net_issues)}**')
if net_issues:
    md += ['| QuoteID | Expected NetAmount | Header NetAmount | Difference |', '|---|---:|---:|---:|']
    for i in net_issues:
        md.append(f"| {i['QuoteID']} | {fmt_money(i['expected'])} | {fmt_money(i['actual'])} | {fmt_money(i['diff'])} |")
    md.append('')
md.append('## ValidUntil < 2026-09-30 and not Accepted')
md.append('')
md.append(f'Count: **{len(expired_unaccepted)}**')
md.append('')
if expired_unaccepted:
    md += ['| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |', '|---|---|---|---|---|---|---|---:|']
    for q in sorted(expired_unaccepted, key=lambda r: (r['ValidUntil'], r['QuoteID'])):
        md.append(f"| {q['QuoteID']} | {q.get('CustomerID','') or ''} | {q.get('ProjectID','') or ''} | {q['QuoteDate']} | {q['ValidUntil']} | {q['Status']} | {q['Currency']} | {fmt_money(q['NetAmount'])} |")
else:
    md.append('None.')
md.append('')
md.append('## Totals by Status and Currency')
md.append('')
md += ['| Status | Currency | NetAmount | NetAmountTWD | Count |', '|---|---|---:|---:|---:|']
for (status, curr), vals in sorted(summary.items()):
    md.append(f"| {status} | {curr} | {fmt_money(vals['amount'])} | {fmt_money(vals['amount_twd'])} | {vals['count']} |")
OUT_MD.write_text('\n'.join(md) + '\n', encoding='utf-8')

# Reopen saved outputs and verify stated requirements.
print('\nVerification from saved outputs:')
assert OUT_XLSX.exists(), 'Workbook output missing'
assert OUT_MD.exists(), 'findings.md missing'
wb2 = load_workbook(OUT_XLSX, data_only=False)
for sheet in ['Report','Quotes','Lines','Summary','Instructions']:
    assert sheet in wb2.sheetnames, f'Missing sheet {sheet}'
q2, l2, s2 = wb2['Quotes'], wb2['Lines'], wb2['Summary']
assert [q2.cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [l2.cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [s2.cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
def count_rows(ws, key_col=1, start=5):
    return sum(1 for r in range(start, ws.max_row+1) if ws.cell(r,key_col).value not in (None,''))
q_count_saved = count_rows(q2)
l_count_saved = count_rows(l2)
s_count_saved = count_rows(s2)
assert q_count_saved == len(quotes), (q_count_saved, len(quotes))
assert l_count_saved == len(all_lines), (l_count_saved, len(all_lines))
assert s_count_saved == len(summary), (s_count_saved, len(summary))
# Check data types from saved workbook: IDs with source values are text; blank ProjectID remains blank; money and quantities are numeric.
for r in range(5, 5+q_count_saved):
    assert isinstance(q2.cell(r,1).value, str) and q2.cell(r,1).value.startswith('Q')
    assert isinstance(q2.cell(r,2).value, str) and q2.cell(r,2).value.startswith('C')
    assert q2.cell(r,3).value is None or isinstance(q2.cell(r,3).value, str)
    for c in (8,9,10,11):
        assert isinstance(q2.cell(r,c).value, (int,float)), f'Quote numeric col {c} row {r} not numeric'
for r in range(5, 5+l_count_saved):
    for c in (1,2,3):
        assert isinstance(l2.cell(r,c).value, str), f'Line text col {c} row {r} not text'
    for c in (4,5,6):
        assert isinstance(l2.cell(r,c).value, (int,float)), f'Line numeric col {c} row {r} not numeric'
expected_formula = f'=SUM(Quotes!K5:K{last_quote_row})' if quote_rows else '=SUM(Quotes!K5:K5)'
assert s2['B2'].value == expected_formula, (s2['B2'].value, expected_formula)
# Recalculate the saved rows' quote-vs-line and net checks to confirm output contents match source checks.
saved_line_sum = defaultdict(Decimal)
for r in range(5, 5+l_count_saved):
    saved_line_sum[l2.cell(r,1).value] += money(l2.cell(r,6).value)
saved_subtotal_issues = []
saved_net_issues = []
saved_expired_unaccepted = []
for r in range(5, 5+q_count_saved):
    qid = q2.cell(r,1).value
    subtotal = money(q2.cell(r,8).value)
    net = money(q2.cell(r,10).value)
    discount = D(q2.cell(r,9).value)
    if saved_line_sum[qid] != subtotal:
        saved_subtotal_issues.append(qid)
    if calc_net(subtotal, discount) != net:
        saved_net_issues.append(qid)
    if parse_date(q2.cell(r,5).value) < BASE_DATE and q2.cell(r,6).value != 'Accepted':
        saved_expired_unaccepted.append(qid)
assert len(saved_subtotal_issues) == len(subtotal_issues)
assert len(saved_net_issues) == len(net_issues)
assert set(saved_expired_unaccepted) == {q['QuoteID'] for q in expired_unaccepted}
md_text = OUT_MD.read_text(encoding='utf-8')
for phrase in ['Quote Findings', 'Consistency checks', 'ValidUntil < 2026-09-30 and not Accepted', f'Quotes retained in workbook: **{len(quotes)}**']:
    assert phrase in md_text, f'Missing markdown phrase: {phrase}'
for q in expired_unaccepted:
    assert q['QuoteID'] in md_text, f'Missing expired-unaccepted quote {q["QuoteID"]} in findings.md'
print('Saved workbook:', OUT_XLSX)
print('Saved markdown:', OUT_MD)
print('Saved quote rows:', q_count_saved, 'line rows:', l_count_saved, 'summary rows:', s_count_saved)
print('Mismatches - line:', len(line_amount_issues), 'subtotal:', len(subtotal_issues), 'net:', len(net_issues))
print('Expired & unaccepted:', len(expired_unaccepted), [q['QuoteID'] for q in sorted(expired_unaccepted, key=lambda r: (r['ValidUntil'], r['QuoteID']))])
print('Summary formula:', s2['B2'].value)
print('All saved-output checks passed.')
