from pathlib import Path
import csv, os
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
# inspect csv
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv sample:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(10), r): print(row)
# inspect manual first lines
man = ws/'input'/'quote_cli_manual.md'
if man.exists():
    print('\nquote_cli_manual.md head:')
    print('\n'.join(man.read_text(encoding='utf-8', errors='replace').splitlines()[:80]))
# inspect workbook sheets/cells used range
xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    wb=load_workbook(xlsx)
    print('\nworkbook sheets:', wb.sheetnames)
    for sh in wb.worksheets:
        print('sheet', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
            print(row)
        print('---')