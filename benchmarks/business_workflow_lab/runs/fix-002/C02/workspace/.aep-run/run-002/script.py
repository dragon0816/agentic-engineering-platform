from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
