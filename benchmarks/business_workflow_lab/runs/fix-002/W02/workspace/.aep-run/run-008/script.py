from pathlib import Path
from decimal import Decimal
from openpyxl import load_workbook
import os
OUT=Path('output')
X=OUT/'sales_template_filled.xlsx'
S=OUT/'sales_order_summary.txt'
CUSTOMER_ID='C001'; DATE_FROM='2026-08-01'; DATE_TO='2026-10-31'
expected_oids=['0045000101','0045000107','0045000109','0045000114','0045000119','0045000125','0045000129','0045000132','0045000136']
allowed_missing={('0045000132','PaymentTerms')}
print('Output files:')
for p in sorted(OUT.iterdir()): print('-', p, p.stat().st_size)
wb=load_workbook(X, data_only=False)
errors=[]
# workbook/sheets/headers
for sheet in ['Report','Details','Summary','Instructions']:
    if sheet not in wb.sheetnames: errors.append(f'Missing sheet {sheet}')
wd=wb['Details']; wr=wb['Report']; ws=wb['Summary']
headers=[wd.cell(4,c).value for c in range(1,12)]
expected_headers=['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
print('Details headers:', headers)
if headers != expected_headers: errors.append('Details headers changed')
if wr['B3'].value != DATE_FROM or wr['D3'].value != DATE_TO: errors.append('Date range incorrect')
if CUSTOMER_ID not in str(wr['B4'].value): errors.append('Condition does not mention CustomerID')
if '0045000132: PaymentTerms' not in str(wr['B9'].value): errors.append('Missing PaymentTerms not noted in Report B9')
# rows
if wd.max_row != 13: errors.append(f'Expected 9 data rows ending row 13, got max_row {wd.max_row}')
seen=[]; total=Decimal('0.00'); blanks=[]
for r in range(5, wd.max_row+1):
    oid=wd.cell(r,1).value; seen.append(oid)
    cust=wd.cell(r,2).value; dt=str(wd.cell(r,5).value)
    if cust != CUSTOMER_ID: errors.append(f'Non-target customer at row {r}: {cust}')
    if not (DATE_FROM <= dt <= DATE_TO): errors.append(f'Date out of range row {r}: {dt}')
    if not isinstance(oid, str) or not oid.startswith('00'): errors.append(f'OrderID not text/leading zeros row {r}: {oid!r}')
    for c, name in [(8,'NetAmount'),(9,'NetAmountTWD')]:
        if not isinstance(wd.cell(r,c).value, (int,float)): errors.append(f'{name} not numeric row {r}: {wd.cell(r,c).value!r}')
    if wd.cell(r,10).value in (None,''):
        if (oid,'CustomerContact') not in allowed_missing: errors.append(f'Unexpected blank CustomerContact for {oid}')
        blanks.append((oid,'CustomerContact'))
    if wd.cell(r,11).value in (None,''):
        if (oid,'PaymentTerms') not in allowed_missing: errors.append(f'Unexpected blank PaymentTerms for {oid}')
        blanks.append((oid,'PaymentTerms'))
    total += Decimal(str(wd.cell(r,9).value)).quantize(Decimal('0.01'))
print('Details OrderIDs:', seen)
print('Blank allowed fields:', blanks)
if seen != expected_oids: errors.append(f'OrderIDs mismatch: {seen}')
# Summary formula and rows
print('Summary B2 formula:', ws['B2'].value)
if ws['B2'].value != '=SUM(Details!I5:I13)': errors.append('Summary formula not extended to last detail row')
summary_count=0; summary_twd=Decimal('0.00')
summary_rows=[]
for r in range(5, ws.max_row+1):
    if ws.cell(r,1).value is None: continue
    summary_rows.append([ws.cell(r,c).value for c in range(1,7)])
    if ws.cell(r,1).value != 'Status': errors.append(f'Unexpected Summary Group at row {r}')
    if not isinstance(ws.cell(r,4).value, (int,float)) or not isinstance(ws.cell(r,5).value, (int,float)): errors.append(f'Summary amounts not numeric row {r}')
    summary_count += int(ws.cell(r,6).value)
    summary_twd += Decimal(str(ws.cell(r,5).value)).quantize(Decimal('0.01'))
print('Summary rows:', summary_rows)
print('Details TWD total:', total, 'Summary TWD total:', summary_twd, 'Summary count:', summary_count)
if summary_count != 9: errors.append('Summary Count total is not 9')
if summary_twd != total: errors.append('Summary TWD total mismatch')
# external summary
text=S.read_text(encoding='utf-8')
print('Summary text:\n'+text)
if 'CustomerID: C001 (exact match only)' not in text: errors.append('Summary text does not confirm exact match')
if '0045000132: PaymentTerms' not in text: errors.append('Summary text does not describe missing data')
if errors:
    print('ERRORS:')
    for e in errors: print('-', e)
    raise SystemExit(1)
print('All saved-output checks passed.')