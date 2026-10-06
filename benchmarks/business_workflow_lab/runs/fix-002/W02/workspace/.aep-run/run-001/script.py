import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
base=Path('.')
print('CWD', Path.cwd())
print('Input files:')
for p in sorted((base/'input').iterdir()):
    print('-', p, p.stat().st_size)
print('\nfx_rates.csv preview:')
with open(base/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r):
        print(i, row)
print('\nsap_manual.md preview:')
text=(base/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[:4000])
print('\nsales_template.xlsx sheets/headers/sample:')
wb=load_workbook(base/'input'/'sales_template.xlsx', data_only=False)
for ws in wb.worksheets:
    print('Sheet', ws.title, 'max', ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,8), values_only=True):
        print(row)
    print('---')