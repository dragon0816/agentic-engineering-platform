from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
for wsname in ['Report','Details','Summary']:
    ws=wb[wsname]
    print('\n',wsname)
    for r in range(1, min(ws.max_row,25)+1):
        vals=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        print(r, vals)
print('\nInstructions rows')
for r in range(1, wb['Instructions'].max_row+1):
    vals=[wb['Instructions'].cell(r,c).value for c in range(1,4)]
    print(r, vals)