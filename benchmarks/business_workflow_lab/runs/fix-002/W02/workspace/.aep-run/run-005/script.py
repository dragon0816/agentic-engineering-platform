import urllib.request, urllib.parse, http.cookiejar, re, csv, io
from pathlib import Path
from openpyxl import load_workbook
print('Workbook instructions full:')
wb=load_workbook('input/sales_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True): print(row)
base='http://127.0.0.1:8810'
cj=http.cookiejar.CookieJar(); opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
def get(path, data=None):
    if data is not None: data=urllib.parse.urlencode(data).encode()
    r=opener.open(base+path, data=data, timeout=10)
    return r.geturl(), r.read().decode('utf-8-sig', errors='replace')
get('/'); get('/login', {'next':'/orders','sso':'company'})
q=urllib.parse.urlencode({'customer':'C001','from':'2026-08-01','to':'2026-10-31'})
_, csvbody=get('/orders/export.csv?'+q)
rows=list(csv.DictReader(io.StringIO(csvbody)))
print('\nExport headers:', rows[0].keys() if rows else [])
print('Export rows:', len(rows))
for r in rows: print(r)
for oid in [r['OrderID'] for r in rows[:2]]:
    url, det=get('/orders/'+oid)
    print('\nDETAIL', oid, url)
    print(det[:6000])
    print('tables?', re.findall(r'<table.*?</table>', det, flags=re.S)[:2])