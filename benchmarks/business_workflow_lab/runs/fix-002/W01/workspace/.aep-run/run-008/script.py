from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
print('Workspace:', workspace)
print('Input files:', [p.name for p in (workspace/'input').iterdir()])
out = workspace/'output'/'september_open_orders.xlsx'
print('Checking output:', out)
if not out.exists():
    raise FileNotFoundError(out)
wb = load_workbook(out, data_only=False)
print('Sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print(f'Sheet {ws.title}: max_row={ws.max_row}, max_column={ws.max_column}')

details = wb['Details']
headers = [details.cell(1, c).value for c in range(1, details.max_column+1)]
print('Details headers:', headers)
rows = []
for r in range(2, details.max_row+1):
    vals = [details.cell(r, c).value for c in range(1, len(headers)+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('Details data rows:', len(rows))
print('First row:', rows[0] if rows else None)
print('Last row:', rows[-1] if rows else None)
order_ids = [str(r['OrderID']) for r in rows]
statuses = sorted({r['Status'] for r in rows})
order_dates = [r['OrderDate'] for r in rows]
print('Order IDs sorted:', order_ids == sorted(order_ids))
print('Statuses:', statuses)
print('Date min/max:', min(order_dates), max(order_dates))

summary = wb['Summary']
summary_rows = []
for r in range(1, summary.max_row+1):
    vals = [summary.cell(r, c).value for c in range(1, summary.max_column+1)]
    if any(v is not None for v in vals):
        summary_rows.append(vals)
print('Summary rows:')
for vals in summary_rows:
    print(vals)

# Verify core requirements independently from saved workbook.
assert len(rows) == 14, f'expected 14 detail rows, got {len(rows)}'
assert statuses == ['Open'], statuses
assert order_ids == sorted(order_ids), 'OrderID not sorted ascending'
for d in order_dates:
    assert str(d) >= '2026-09-01' and str(d) <= '2026-09-30', d
assert any('TWD' in [str(x) for x in vals] or 'Total' in [str(x) for x in vals] for vals in summary_rows), 'summary appears missing totals'
print('Independent re-open check OK: september_open_orders.xlsx meets row count/status/date/sort and has summary rows.')