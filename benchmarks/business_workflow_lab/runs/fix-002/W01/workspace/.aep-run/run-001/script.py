from pathlib import Path
import csv
from openpyxl import load_workbook
base=Path('.')
print('Workspace:', base.resolve())
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print('-', p, p.stat().st_size)
md=base/'input'/'sap_manual.md'
print('\nSAP manual first 2000 chars:')
print(md.read_text(encoding='utf-8')[:2000])
fx=base/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r): print(row)
wb_path=base/'input'/'sales_template.xlsx'
wb=load_workbook(wb_path)
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('Sheet:', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,8), values_only=True):
        print(row)
    print('---')