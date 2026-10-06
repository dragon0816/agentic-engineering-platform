import csv
import html
import json
import os
import re
import shutil
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from urllib import parse, request
import http.cookiejar

from openpyxl import load_workbook

WORKSPACE = Path('.').resolve()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
BASE_URL = 'http://127.0.0.1:8810'

params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
run_params = {}
if params_path and Path(params_path).exists():
    run_params = json.load(open(params_path, encoding='utf-8'))

start_date = run_params.get('start_date') or '2026-09-01'
end_date = run_params.get('end_date') or '2026-09-30'
status = run_params.get('status') or 'Open'
output_filename = run_params.get('output_filename') or 'sales_orders_report.xlsx'
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'

# Print actual inputs for auditability.
print('Workspace:', WORKSPACE)
print('Run parameters:', {'start_date': start_date, 'end_date': end_date, 'status': status, 'output_filename': output_filename})
print('Input files:')
for p in sorted(INPUT.glob('*')):
    print(f'- {p.name} ({p.stat().st_size} bytes)')

# Validate date inputs.
start_dt = datetime.strptime(start_date, '%Y-%m-%d').date()
end_dt = datetime.strptime(end_date, '%Y-%m-%d').date()
if start_dt > end_dt:
    raise ValueError('start_date must be on or before end_date')

# Read FX rates. The Note column in the lab CSV may contain an unquoted comma; only the first two columns are needed.
fx_rates = {}
with (INPUT / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header:', header)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            fx_rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates loaded:', {k: str(v) for k, v in sorted(fx_rates.items())})

# Minimal HTML helpers for the known Lab ERP markup.
def clean_cell(s):
    s = re.sub(r'<[^>]+>', '', s, flags=re.S)
    return html.unescape(s).strip()

def extract_table_rows(page_html):
    table_m = re.search(r"<table id='results'.*?<thead>.*?<tr>(.*?)</tr>.*?</thead>.*?<tbody>(.*?)</tbody>", page_html, re.S)
    if not table_m:
        return [], []
    headers = [clean_cell(x) for x in re.findall(r'<th[^>]*>(.*?)</th>', table_m.group(1), re.S)]
    rows = []
    for tr in re.findall(r'<tr>(.*?)</tr>', table_m.group(2), re.S):
        vals = [clean_cell(x) for x in re.findall(r'<td[^>]*>(.*?)</td>', tr, re.S)]
        if vals:
            rows.append(dict(zip(headers, vals)))
    return headers, rows

def extract_result_count(page_html):
    m = re.search(r"<p id='result-count'>\s*Total:\s*<strong>(\d+)</strong>\s*rows\s*·\s*Page\s*(\d+)\s*of\s*(\d+)", page_html, re.S)
    if not m:
        return None, None, None
    return int(m.group(1)), int(m.group(2)), int(m.group(3))

def extract_detail_fields(page_html):
    fields = {}
    for tr in re.findall(r'<tr>(.*?)</tr>', page_html, re.S):
        th = re.search(r'<th[^>]*scope=["\']row["\'][^>]*>(.*?)</th>', tr, re.S)
        td = re.search(r'<td[^>]*data-field=["\']([^"\']+)["\'][^>]*>(.*?)</td>', tr, re.S)
        if th and td:
            fields[clean_cell(th.group(1))] = clean_cell(td.group(2))
            fields[td.group(1)] = clean_cell(td.group(2))
    return fields

# Log in using the documented SSO form; no password is used.
cj = http.cookiejar.CookieJar()
opener = request.build_opener(request.HTTPCookieProcessor(cj))
login_data = parse.urlencode({'next': '/orders', 'sso': 'company'}).encode('utf-8')
opener.open(request.Request(BASE_URL + '/login', data=login_data, method='POST'), timeout=20).read()

# Fetch every result page, sorted by OrderID asc.
query = {'from': start_date, 'to': end_date, 'sort': 'OrderID', 'dir': 'asc'}
if status:
    query['status'] = status
all_rows = []
headers_seen = None
expected_total = None
page_count = None
page = 1
while True:
    query['page'] = str(page)
    url = BASE_URL + '/orders?' + parse.urlencode(query)
    page_html = opener.open(url, timeout=20).read().decode('utf-8', 'replace')
    if 'Sign in' in page_html and 'login-form' in page_html:
        opener.open(request.Request(BASE_URL + '/login', data=login_data, method='POST'), timeout=20).read()
        page_html = opener.open(url, timeout=20).read().decode('utf-8', 'replace')
    total, current_page, total_pages = extract_result_count(page_html)
    if total is None:
        if '查無資料' in page_html or '0 筆' in page_html:
            expected_total, current_page, total_pages = 0, 1, 1
            rows = []
        else:
            raise RuntimeError('Could not parse result count from orders page')
    else:
        expected_total, page_count = total, total_pages
        headers, rows = extract_table_rows(page_html)
        headers_seen = headers_seen or headers
    print(f'Fetched page {page}: {len(rows)} rows')
    all_rows.extend(rows)
    if expected_total == 0 or page >= (page_count or 1):
        break
    page += 1

print('Result headers:', headers_seen)
print('Expected total rows:', expected_total)
print('Collected rows from pages:', len(all_rows))
if expected_total is not None and len(all_rows) != expected_total:
    raise RuntimeError(f'Collected {len(all_rows)} rows but page reported {expected_total}')

# Sort explicitly as text to preserve leading zeros.
all_rows.sort(key=lambda r: r.get('OrderID', ''))

# Fetch detail-only fields for each order.
for row in all_rows:
    order_id = row['OrderID']
    detail_html = opener.open(BASE_URL + '/orders/' + parse.quote(order_id), timeout=20).read().decode('utf-8', 'replace')
    if 'Sign in' in detail_html and 'login-form' in detail_html:
        opener.open(request.Request(BASE_URL + '/login', data=login_data, method='POST'), timeout=20).read()
        detail_html = opener.open(BASE_URL + '/orders/' + parse.quote(order_id), timeout=20).read().decode('utf-8', 'replace')
    fields = extract_detail_fields(detail_html)
    row['CustomerContact'] = fields.get('CustomerContact', '')
    row['PaymentTerms'] = fields.get('PaymentTerms', '')

# Build workbook from template.
template = INPUT / 'sales_template.xlsx'
out_path = OUTPUT / output_filename
wb = load_workbook(template)
ws_report = wb['Report']
ws_details = wb['Details']
ws_summary = wb['Summary']

# Clear old detail and summary data below headers, preserving sheets and header rows.
if ws_details.max_row >= 5:
    ws_details.delete_rows(5, ws_details.max_row - 4)
if ws_summary.max_row >= 5:
    ws_summary.delete_rows(5, ws_summary.max_row - 4)

now_iso = datetime.now().astimezone().isoformat(timespec='seconds')
condition = f'Sales Orders; OrderDate {start_date} to {end_date}; Status={status or "(all)"}; Sort=OrderID asc'
ws_report['B3'] = start_date
ws_report['D3'] = end_date
ws_report['B4'] = condition
ws_report['B5'] = now_iso
ws_report['B7'] = 'Lab ERP (synthetic); fx_rates.csv'
ws_report['D7'] = now_iso
ws_report['B9'] = f'Fetched all {len(all_rows)} rows across {page_count or 1} page(s). NetAmountTWD uses ROUND_HALF_UP and fx_rates.csv.'

money_q = Decimal('0.01')
detail_cols = ['OrderID', 'CustomerID', 'CustomerName', 'ProjectID', 'OrderDate', 'Status', 'Currency', 'NetAmount', 'NetAmountTWD', 'CustomerContact', 'PaymentTerms']
summary = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
for idx, row in enumerate(all_rows, start=5):
    currency = row['Currency']
    if currency not in fx_rates:
        raise KeyError(f'Missing FX rate for currency {currency}')
    net = Decimal(str(row['NetAmount']).replace(',', '')).quantize(money_q, rounding=ROUND_HALF_UP)
    twd = (net * fx_rates[currency]).quantize(money_q, rounding=ROUND_HALF_UP)
    values = [
        str(row.get('OrderID', '')),
        str(row.get('CustomerID', '')),
        str(row.get('CustomerName', '')),
        str(row.get('ProjectID', '')),
        str(row.get('OrderDate', '')),
        str(row.get('Status', '')),
        str(currency),
        float(net),
        float(twd),
        str(row.get('CustomerContact', '')),
        str(row.get('PaymentTerms', '')),
    ]
    for col_idx, value in enumerate(values, start=1):
        cell = ws_details.cell(row=idx, column=col_idx, value=value)
        if col_idx in (1, 2, 4):
            cell.number_format = '@'
        elif col_idx == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col_idx in (8, 9):
            cell.number_format = '0.00'
    summary[currency]['amount'] += net
    summary[currency]['twd'] += twd
    summary[currency]['count'] += 1

last_detail_row = max(5, 4 + len(all_rows))
ws_summary['B2'] = f'=SUM(Details!I5:I{last_detail_row})'
summary_row = 5
for currency in sorted(summary):
    item = summary[currency]
    vals = ['Currency', '', currency, float(item['amount'].quantize(money_q)), float(item['twd'].quantize(money_q)), item['count']]
    for col_idx, value in enumerate(vals, start=1):
        cell = ws_summary.cell(row=summary_row, column=col_idx, value=value)
        if col_idx in (4, 5):
            cell.number_format = '0.00'
    summary_row += 1
total_twd = sum((item['twd'] for item in summary.values()), Decimal('0')).quantize(money_q)
vals = ['Total', 'All currencies', 'TWD', None, float(total_twd), len(all_rows)]
for col_idx, value in enumerate(vals, start=1):
    cell = ws_summary.cell(row=summary_row, column=col_idx, value=value)
    if col_idx in (4, 5):
        cell.number_format = '0.00'

wb.save(out_path)
print('Saved workbook:', out_path)

# Re-open and verify all stated requirements from the saved workbook.
wb2 = load_workbook(out_path, data_only=False)
assert wb2.sheetnames == ['Report', 'Details', 'Summary', 'Instructions'], wb2.sheetnames
rd = wb2['Details']
rs = wb2['Summary']
headers = [rd.cell(4, c).value for c in range(1, 12)]
assert headers == detail_cols, headers
saved_rows = []
for r in range(5, 5 + len(all_rows)):
    saved_rows.append({detail_cols[c-1]: rd.cell(r, c).value for c in range(1, 12)})
assert len(saved_rows) == len(all_rows)
assert [r['OrderID'] for r in saved_rows] == sorted([r['OrderID'] for r in saved_rows])
for r in saved_rows:
    assert isinstance(r['OrderID'], str) and r['OrderID'].startswith('0')
    assert start_date <= r['OrderDate'] <= end_date, r
    if status:
        assert r['Status'] == status, r
    assert isinstance(r['NetAmount'], (int, float))
    assert isinstance(r['NetAmountTWD'], (int, float))
    assert r['Currency'] in fx_rates
assert rs['B2'].value == f'=SUM(Details!I5:I{last_detail_row})'
# Verify summary currency rows and total row.
saved_summary = []
for r in range(5, rs.max_row + 1):
    if rs.cell(r, 1).value:
        saved_summary.append([rs.cell(r, c).value for c in range(1, 7)])
expected_summary_rows = len(summary) + 1
assert len(saved_summary) == expected_summary_rows, saved_summary
for row in saved_summary:
    if row[0] == 'Currency':
        cur = row[2]
        assert round(row[3], 2) == float(summary[cur]['amount'])
        assert round(row[4], 2) == float(summary[cur]['twd'])
        assert row[5] == summary[cur]['count']
    elif row[0] == 'Total':
        assert row[2] == 'TWD'
        assert round(row[4], 2) == float(total_twd)
        assert row[5] == len(all_rows)
    else:
        raise AssertionError(row)
print('Verification OK: workbook exists, Details headers/row count/text IDs/date range/status/sort/numeric amounts verified; Summary currency totals, all-currency TWD total, and B2 formula verified.')
print('Output:', str(out_path))