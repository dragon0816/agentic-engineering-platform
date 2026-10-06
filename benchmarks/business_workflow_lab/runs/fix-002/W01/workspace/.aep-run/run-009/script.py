from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal

workspace = Path.cwd()
out = workspace/'output'/'september_open_orders.xlsx'
print('Workspace:', workspace)
print('Input files:', sorted(p.name for p in (workspace/'input').iterdir()))
print('Checking output:', out)
if not out.exists():
    raise FileNotFoundError(out)
wb = load_workbook(out, data_only=False)
print('Sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print(f'Sheet {ws.title}: max_row={ws.max_row}, max_column={ws.max_column}')

# Details sheet uses template title/instruction rows; row 4 is header, row 5+ data.
ws = wb['Details']
header_row = 4
headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column+1)]
print('Details header row:', header_row)
print('Details headers:', headers)
rows = []
for r in range(header_row+1, ws.max_row+1):
    vals = [ws.cell(r, c).value for c in range(1, len(headers)+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('Details data rows:', len(rows))
print('First detail row:', rows[0] if rows else None)
print('Last detail row:', rows[-1] if rows else None)

required_headers = ['OrderID', 'CustomerID', 'CustomerName', 'ProjectID', 'OrderDate', 'Status', 'Currency', 'NetAmount']
missing = [h for h in required_headers if h not in headers]
assert not missing, f'missing detail headers: {missing}'
order_ids = [str(r['OrderID']) for r in rows]
statuses = sorted({r['Status'] for r in rows})
order_dates = [str(r['OrderDate'])[:10] for r in rows]
currencies = sorted({r['Currency'] for r in rows})
net_amounts = [r['NetAmount'] for r in rows]
print('Order IDs:', order_ids)
print('Order IDs sorted:', order_ids == sorted(order_ids))
print('Statuses:', statuses)
print('Date min/max:', min(order_dates), max(order_dates))
print('Currencies:', currencies)
print('NetAmount types:', sorted({type(x).__name__ for x in net_amounts}))

assert len(rows) == 14, f'expected 14 detail rows, got {len(rows)}'
assert statuses == ['Open'], statuses
assert all('2026-09-01' <= d <= '2026-09-30' for d in order_dates), order_dates
assert order_ids == sorted(order_ids), 'OrderID not sorted ascending'
assert all(isinstance(r['OrderID'], str) for r in rows), 'OrderID must be text IDs'
assert all(isinstance(r['CustomerID'], str) for r in rows), 'CustomerID must be text IDs'
assert all(isinstance(x, (int, float)) for x in net_amounts), 'NetAmount must be numeric'

# Summary sheet: print non-empty rows and verify per-currency original totals + all TWD total.
sumws = wb['Summary']
summary_nonempty = []
for r in range(1, sumws.max_row+1):
    vals = [sumws.cell(r, c).value for c in range(1, sumws.max_column+1)]
    if any(v is not None for v in vals):
        summary_nonempty.append(vals)
print('Summary non-empty rows:')
for vals in summary_nonempty:
    print(vals)

# Find expected totals from Details workbook values and attached rates.
rates = {'TWD': Decimal('1'), 'USD': Decimal('32'), 'EUR': Decimal('35')}
expected_orig = {}
expected_twd = Decimal('0')
for r in rows:
    cur = r['Currency']
    amt = Decimal(str(r['NetAmount']))
    expected_orig[cur] = expected_orig.get(cur, Decimal('0')) + amt
    expected_twd += amt * rates[cur]
print('Expected original totals by currency:', {k: str(v) for k, v in sorted(expected_orig.items())})
print('Expected all-currency TWD total:', str(expected_twd))

# Extract numeric values from Summary rows by currency labels and total label.
summary_text = '\n'.join(' | '.join('' if v is None else str(v) for v in vals) for vals in summary_nonempty)
for cur, total in expected_orig.items():
    assert cur in summary_text, f'{cur} missing in Summary'
    # Check a numeric cell with that row equals original total.
    found = False
    for vals in summary_nonempty:
        if cur in [str(v) for v in vals]:
            nums = [Decimal(str(v)) for v in vals if isinstance(v, (int, float))]
            if total in nums:
                found = True
    assert found, f'original total for {cur}={total} not found in Summary'

found_twd_total = any(
    any(isinstance(v, (int, float)) and Decimal(str(v)) == expected_twd for v in vals)
    for vals in summary_nonempty
)
assert found_twd_total, f'TWD total {expected_twd} not found in Summary'
print('Independent re-open check OK: workbook meets Details and Summary requirements.')