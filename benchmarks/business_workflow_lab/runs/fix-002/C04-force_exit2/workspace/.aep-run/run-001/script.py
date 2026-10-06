from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size, 'bytes')
manual = ws/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md head:')
print(manual.read_text(encoding='utf-8', errors='replace')[:4000])
fx = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i,row in zip(range(8), reader):
        print(row)

xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nworkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('sheet', s, 'size', sh.max_row, sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
        print(r)
    print('---')