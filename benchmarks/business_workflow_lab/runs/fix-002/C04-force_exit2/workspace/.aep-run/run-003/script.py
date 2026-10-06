from pathlib import Path
import subprocess, sys, json, csv, shutil, os, re
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI_PY = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
OUT_XLSX = OUTPUT / 'sent_quotes_2026-09.xlsx'
RUN_LOG = OUTPUT / 'quote_cli_run_log.json'
ERROR_TXT = OUTPUT / 'quote_cli_error.txt'
CHECK_TXT = OUTPUT / 'validation_report.txt'

TW = timezone(timedelta(hours=8))
now = datetime.now(TW).replace(microsecond=0).isoformat()

class CliError(Exception):
    def __init__(self, cmd, exit_code, stderr, attempts):
        super().__init__(f'CLI failed exit {exit_code}: {stderr}')
        self.cmd = cmd
        self.exit_code = exit_code
        self.stderr = stderr
        self.attempts = attempts

def run_cli(args, parse_json=True):
    """Run quote CLI. Retry exit code 4 up to two retries. Never use stdout unless rc=0 and parsable."""
    cmd = [sys.executable, str(CLI_PY)] + args
    attempts = []
    max_attempts = 3  # first try + at most two retries for exit 4
    for attempt in range(1, max_attempts + 1):
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True, encoding='utf-8', errors='replace')
        rec = {
            'attempt': attempt,
            'cmd': cmd,
            'exit_code': proc.returncode,
            'stderr': proc.stderr.strip(),
            'stdout_bytes': len(proc.stdout.encode('utf-8', errors='replace')),
        }
        attempts.append(rec)
        if proc.returncode == 0:
            if parse_json:
                try:
                    return json.loads(proc.stdout), attempts
                except Exception as e:
                    raise CliError(cmd, 0, f'stdout was not parseable JSON: {e}', attempts)
            return proc.stdout, attempts
        if proc.returncode == 4:
            if attempt < max_attempts:
                continue
            raise CliError(cmd, 4, proc.stderr.strip(), attempts)
        if proc.returncode in (2, 3):
            # Do not retry. stderr only is authoritative.
            raise CliError(cmd, proc.returncode, proc.stderr.strip(), attempts)
        raise CliError(cmd, proc.returncode, proc.stderr.strip(), attempts)
    raise CliError(cmd, 4, 'temporary service error after retries', attempts)

def load_fx():
    rates = {}
    # File has unquoted commas in Note; only first two columns matter.
    with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        header = next(reader)
        for row in reader:
            if len(row) >= 2 and row[0]:
                rates[row[0]] = Decimal(row[1])
    return rates

def money(v):
    return Decimal(str(v)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def set_text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.number_format = '@'

def clear_data(sh, start_row, max_col):
    for row in sh.iter_rows(min_row=start_row, max_row=max(sh.max_row, start_row), max_col=max_col):
        for cell in row:
            cell.value = None

def fill_error_workbook(error: CliError, run_log):
    shutil.copyfile(TEMPLATE, OUT_XLSX)
    wb = load_workbook(OUT_XLSX)
    rep = wb['Report']
    rep['B3'] = FROM_DATE
    rep['D3'] = TO_DATE
    rep['B4'] = f'Status={STATUS}; QuoteDate {FROM_DATE}..{TO_DATE}'
    rep['B5'] = now
    rep['B7'] = 'quote_cli'
    rep['D7'] = now
    rep['B9'] = f'CLI 查詢失敗；exit code {error.exit_code}；未使用 stdout。stderr: {error.stderr}'
    wb['Summary']['B2'] = '=SUM(Quotes!K5:K5)'
    wb.save(OUT_XLSX)

def fill_success_workbook(list_data, details, run_log):
    rates = load_fx()
    shutil.copyfile(TEMPLATE, OUT_XLSX)
    wb = load_workbook(OUT_XLSX)
    rep = wb['Report']
    quotes = wb['Quotes']
    lines = wb['Lines']
    summ = wb['Summary']

    # Clear existing possible rows without touching headers/formula rows.
    clear_data(quotes, 5, 11)
    clear_data(lines, 5, 6)
    clear_data(summ, 5, 6)

    rep['B3'] = FROM_DATE
    rep['D3'] = TO_DATE
    rep['B4'] = f'Status={STATUS}; QuoteDate {FROM_DATE}..{TO_DATE}'
    rep['B5'] = now
    rep['B7'] = 'quote_cli; fx_rates.csv'
    rep['D7'] = now
    rep['B9'] = '成功取得並填入資料；金額以 Decimal ROUND_HALF_UP 計算。'

    quote_headers = []
    line_rows = []
    for item in list_data.get('items', []):
        qid = str(item['QuoteID'])
        detail = details[qid]
        header = detail['header']
        quote_headers.append(header)
        for li in detail.get('line_items', []):
            line_rows.append((qid, li))

    # sort stable by QuoteDate then QuoteID
    quote_headers.sort(key=lambda h: (str(h.get('QuoteDate','')), str(h.get('QuoteID',''))))
    line_rows.sort(key=lambda pair: (str(pair[0]), str(pair[1].get('ProductID',''))))

    for idx, h in enumerate(quote_headers, start=5):
        cur = h['Currency']
        rate = rates[cur]
        subtotal = money(h['Subtotal'])
        disc = Decimal(str(h['DiscountPercent']))
        net = money(h['NetAmount'])
        calc_net = money(subtotal * (Decimal('1') - disc / Decimal('100')))
        if net != calc_net:
            raise ValueError(f"NetAmount mismatch for {h['QuoteID']}: CLI {net}, calculated {calc_net}")
        net_twd = money(net * rate)
        vals = [h['QuoteID'], h['CustomerID'], h['ProjectID'], h['QuoteDate'], h['ValidUntil'], h['Status'], cur,
                float(subtotal), float(disc), float(net), float(net_twd)]
        for col, val in enumerate(vals, start=1):
            c = quotes.cell(idx, col, val)
            if col in (1,2,3):
                c.number_format = '@'
                c.value = str(val)
            elif col in (4,5):
                c.number_format = 'yyyy-mm-dd'
            elif col in (8,10,11):
                c.number_format = '0.00'
            elif col == 9:
                c.number_format = '0.00'

    for idx, (qid, li) in enumerate(line_rows, start=5):
        qty = Decimal(str(li['Quantity']))
        unit = money(li['UnitPrice'])
        line_amt = money(li['LineAmount'])
        calc_line = money(qty * unit)
        if line_amt != calc_line:
            raise ValueError(f"LineAmount mismatch for {qid}/{li['ProductID']}: CLI {line_amt}, calculated {calc_line}")
        vals = [qid, li['ProductID'], li['ProductName'], int(qty) if qty == qty.to_integral_value() else float(qty), float(unit), float(line_amt)]
        for col, val in enumerate(vals, start=1):
            c = lines.cell(idx, col, val)
            if col in (1,2):
                c.number_format = '@'
                c.value = str(val)
            elif col in (5,6):
                c.number_format = '0.00'

    last_q_row = 4 + len(quote_headers) if quote_headers else 5
    summ['B2'] = f'=SUM(Quotes!K5:K{last_q_row})' if quote_headers else '=SUM(Quotes!K5:K5)'

    groups = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    for h in quote_headers:
        key = (h['Status'], '', h['Currency'])
        net = money(h['NetAmount'])
        twd = money(net * rates[h['Currency']])
        groups[key]['amount'] += net
        groups[key]['twd'] += twd
        groups[key]['count'] += 1
    row = 5
    for (grp, subgrp, cur), agg in sorted(groups.items()):
        vals = [grp, subgrp, cur, float(money(agg['amount'])), float(money(agg['twd'])), agg['count']]
        for col, val in enumerate(vals, start=1):
            c = summ.cell(row, col, val)
            if col in (4,5): c.number_format = '0.00'
        row += 1

    wb.save(OUT_XLSX)

def validate_output(expected_success, expected_quote_count=None, expected_line_count=None, expected_error=None):
    wb = load_workbook(OUT_XLSX, data_only=False)
    problems = []
    if wb.sheetnames != ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']:
        problems.append(f'sheet names changed: {wb.sheetnames}')
    if str(wb['Report']['A20'].value or '') != '合成測試資料 Synthetic lab data · template v1 · 請勿修改本行':
        problems.append('Report A20 footer changed')
    if str(wb['Report']['B3'].value) != FROM_DATE or str(wb['Report']['D3'].value) != TO_DATE:
        problems.append('Report date range incorrect')
    if STATUS not in str(wb['Report']['B4'].value):
        problems.append('Report condition missing status')
    if expected_error and expected_error not in str(wb['Report']['B9'].value):
        problems.append('Report notes do not include expected CLI error')
    q_headers = [wb['Quotes'].cell(4,c).value for c in range(1,12)]
    if q_headers != ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']:
        problems.append(f'Quotes headers wrong: {q_headers}')
    l_headers = [wb['Lines'].cell(4,c).value for c in range(1,7)]
    if l_headers != ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']:
        problems.append(f'Lines headers wrong: {l_headers}')
    s_headers = [wb['Summary'].cell(4,c).value for c in range(1,7)]
    if s_headers != ['Group','SubGroup','Currency','Amount','AmountTWD','Count']:
        problems.append(f'Summary headers wrong: {s_headers}')

    def data_count(sh, key_col=1):
        count = 0
        for r in range(5, sh.max_row+1):
            if sh.cell(r,key_col).value not in (None, ''):
                count += 1
        return count
    q_count = data_count(wb['Quotes'])
    l_count = data_count(wb['Lines'])
    if expected_quote_count is not None and q_count != expected_quote_count:
        problems.append(f'Quotes row count {q_count} != {expected_quote_count}')
    if expected_line_count is not None and l_count != expected_line_count:
        problems.append(f'Lines row count {l_count} != {expected_line_count}')
    # IDs text and monetary numeric for any data rows
    for r in range(5, 5+q_count):
        for c in (1,2,3):
            if not isinstance(wb['Quotes'].cell(r,c).value, str): problems.append(f'Quotes {r},{c} ID is not text')
        for c in (8,9,10,11):
            if not isinstance(wb['Quotes'].cell(r,c).value, (int,float)): problems.append(f'Quotes {r},{c} amount is not numeric')
    for r in range(5, 5+l_count):
        for c in (1,2):
            if not isinstance(wb['Lines'].cell(r,c).value, str): problems.append(f'Lines {r},{c} ID is not text')
        for c in (4,5,6):
            if not isinstance(wb['Lines'].cell(r,c).value, (int,float)): problems.append(f'Lines {r},{c} numeric field is not numeric')
    status = 'OK' if not problems else 'FAIL'
    report = {
        'status': status,
        'output_file': str(OUT_XLSX),
        'quotes_rows': q_count,
        'lines_rows': l_count,
        'summary_formula': wb['Summary']['B2'].value,
        'problems': problems,
        'checked': ['file exists', 'sheet names', 'headers', 'Report dates/condition/notes', 'footer', 'row counts', 'ID text types', 'numeric amount types']
    }
    CHECK_TXT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    if problems:
        raise AssertionError(json.dumps(report, ensure_ascii=False, indent=2))
    return report

all_attempts = []
try:
    list_args = ['quotes','list','--from',FROM_DATE,'--to',TO_DATE,'--status',STATUS,'--format','json']
    list_data, attempts = run_cli(list_args, parse_json=True)
    all_attempts.extend(attempts)
    if list_data.get('schema_version') != 1 or 'items' not in list_data or not isinstance(list_data['items'], list):
        raise ValueError('quotes list JSON schema is not recognized')
    if list_data.get('count') != len(list_data.get('items', [])):
        raise ValueError('quotes list count does not match items length')
    details = {}
    total_lines = 0
    for item in list_data['items']:
        qid = str(item['QuoteID'])
        detail, attempts = run_cli(['quotes','get',qid,'--format','json'], parse_json=True)
        all_attempts.extend(attempts)
        if detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
            raise ValueError(f'quote detail JSON schema is not recognized for {qid}')
        details[qid] = detail
        total_lines += len(detail['line_items'])
    RUN_LOG.write_text(json.dumps({'status':'success','attempts':all_attempts}, ensure_ascii=False, indent=2), encoding='utf-8')
    fill_success_workbook(list_data, details, all_attempts)
    report = validate_output(True, expected_quote_count=len(list_data['items']), expected_line_count=total_lines)
    print('SUCCESS')
    print(json.dumps(report, ensure_ascii=False, indent=2))
except CliError as e:
    all_attempts.extend(e.attempts)
    RUN_LOG.write_text(json.dumps({'status':'cli_error','exit_code':e.exit_code,'stderr':e.stderr,'attempts':all_attempts}, ensure_ascii=False, indent=2), encoding='utf-8')
    ERROR_TXT.write_text(f'exit code {e.exit_code}\nstderr: {e.stderr}\n', encoding='utf-8')
    fill_error_workbook(e, all_attempts)
    report = validate_output(False, expected_quote_count=0, expected_line_count=0, expected_error=f'exit code {e.exit_code}')
    print('CLI_ERROR')
    print(json.dumps({'exit_code':e.exit_code,'stderr':e.stderr,'attempts':len(e.attempts),'output_file':str(OUT_XLSX),'validation':report}, ensure_ascii=False, indent=2))
except Exception as e:
    RUN_LOG.write_text(json.dumps({'status':'script_error','error':repr(e),'attempts':all_attempts}, ensure_ascii=False, indent=2), encoding='utf-8')
    raise