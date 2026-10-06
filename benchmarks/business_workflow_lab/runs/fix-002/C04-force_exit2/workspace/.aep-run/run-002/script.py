from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True):
    print(row)
print('Report merged:', list(wb['Report'].merged_cells.ranges))
for s in ['Report','Quotes','Lines','Summary']:
    sh=wb[s]
    print('\n',s,'merged',list(sh.merged_cells.ranges))
    print('dimensions',sh.max_row,sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=25, values_only=False):
        vals=[c.value for c in r]
        print(r[0].row, vals)