from pathlib import Path
import os, csv
from openpyxl import load_workbook
ws=Path.cwd()
print('CWD', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
    if p.suffix.lower()=='.csv':
        with p.open('r', encoding='utf-8-sig', newline='') as f:
            rdr=csv.reader(f)
            for i,row in zip(range(5), rdr):
                print('  CSV row', i+1, row)
    elif p.suffix.lower()=='.md':
        txt=p.read_text(encoding='utf-8', errors='replace')
        print('  MD head:', repr(txt[:1000]))
    elif p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print('  Sheets:', wb.sheetnames)
        for s in wb.sheetnames:
            sh=wb[s]
            print('   ', s, 'max', sh.max_row, sh.max_column)
            for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
                print('     ', r)
