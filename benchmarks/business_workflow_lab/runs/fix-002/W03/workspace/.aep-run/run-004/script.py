from pathlib import Path
import urllib.request, urllib.parse, http.cookiejar, csv, io, re, os
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from copy import copy

WORK=Path.cwd()
IN=WORK/'input'
OUT=WORK/'output'
OUT.mkdir(exist_ok=True)
BASE='http://127.0.0.1:8810'
FROM='2026-09-01'
TO='2026-09-30'
ASOF=date(2026,9,30)

def dparse(s):
    if isinstance(s, date): return s
    return datetime.strptime(str(s), '%Y-%m-%d').date()

def money(v):
    return Decimal(str(v).replace(',','')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# Load FX rates (first two fields are authoritative; Note may contain unquoted comma in synthetic file)
fx={}
with (IN/'fx_rates.csv').open('r', encoding='utf-8-sig', newline='') as f:
    rdr=csv.reader(f)
    header=next(rdr)
    print('FX header:', header)
    for row in rdr:
        if len(row)>=2 and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', fx)

# Query ERP invoice export for all September 2026 invoices
cj=http.cookiejar.CookieJar()
opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
login_data=urllib.parse.urlencode({'next':'/invoices','sso':'company'}).encode('utf-8')
with opener.open(urllib.request.Request(BASE+'/login', data=login_data, method='POST'), timeout=20) as r:
    print('Logged in:', r.status, r.geturl())
params=urllib.parse.urlencode({'date_from':FROM, 'date_to':TO})
export_url=BASE+'/invoices/export.csv?'+params
with opener.open(export_url, timeout=20) as r:
    raw=r.read()
    ctype=r.headers.get('content-type')
    final_url=r.geturl()
print('Invoice export:', final_url, ctype, len(raw), 'bytes')
text=raw.decode('utf-8-sig')
print('Export sample:')
for line in text.splitlines()[:5]:
    print(' ', line)
reader=csv.DictReader(io.StringIO(text))
rows=list(reader)
print('Invoice columns:', reader.fieldnames)
print('Invoice rows:', len(rows))
if not rows:
    raise SystemExit('No invoice rows returned; aborting')

expected_cols=['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance']
missing=[c for c in expected_cols if c not in (reader.fieldnames or [])]
if missing:
    raise SystemExit(f'Missing expected invoice columns: {missing}')

# Verify filter range and prepare records
records=[]
for row in rows:
    inv_date=dparse(row['InvoiceDate'])
    if not (dparse(FROM) <= inv_date <= dparse(TO)):
        raise SystemExit(f'Invoice outside requested range: {row}')
    cur=row['Currency'].strip()
    if cur not in fx:
        raise SystemExit(f'Missing FX rate for {cur}')
    amount=money(row['Amount'])
    paid=money(row['PaidAmount'])
    bal=money(row['Balance'])
    if bal != (amount-paid).quantize(Decimal('0.01')):
        raise SystemExit(f'Balance mismatch for {row["InvoiceID"]}: {bal} vs {amount-paid}')
    bal_twd=(bal*fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    records.append({
        'InvoiceID': str(row['InvoiceID']),
        'OrderID': str(row['OrderID']),
        'ProjectID': str(row['ProjectID']),
        'CustomerID': str(row['CustomerID']),
        'InvoiceDate': inv_date,
        'DueDate': dparse(row['DueDate']),
        'Currency': cur,
        'Amount': amount,
        'PaidAmount': paid,
        'Balance': bal,
        'BalanceTWD': bal_twd,
    })
records.sort(key=lambda r: (r['InvoiceDate'], r['InvoiceID']))

# Load template and clear data rows in source/summarization sheets
out_path=OUT/'project_invoice_report_2026-09.xlsx'
wb=load_workbook(IN/'project_template.xlsx')

# helpers for preserving row 4 style/number formats
header_row=4
start_row=5

def clear_sheet_data(ws):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row-start_row+1)

def style_like_header(ws, row_idx):
    # copy basic style from row 4 cells into new data row, if available
    for col in range(1, ws.max_column+1):
        src=ws.cell(header_row, col)
        dst=ws.cell(row_idx, col)
        if src.has_style:
            dst.font=copy(src.font)
            dst.fill=copy(src.fill)
            dst.border=copy(src.border)
            dst.alignment=copy(src.alignment)
            dst.number_format=src.number_format
            dst.protection=copy(src.protection)

for s in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    clear_sheet_data(wb[s])

# Fill Report metadata where applicable
rep=wb['Report']
rep['B3']=FROM
rep['D3']=TO
rep['B4']='Invoices only; as-of 2026-09-30; overdue when Balance > 0 and DueDate < 2026-09-30'
rep['B5']=datetime.now().strftime('%Y-%m-%d %H:%M:%S')

# Fill Invoices sheet
inv_ws=wb['Invoices']
headers=[inv_ws.cell(header_row,c).value for c in range(1, inv_ws.max_column+1)]
print('Template invoice headers:', headers)
for idx, rec in enumerate(records, start=start_row):
    style_like_header(inv_ws, idx)
    values=[rec[h] for h in headers]
    for col, val in enumerate(values, start=1):
        cell=inv_ws.cell(idx,col,value=val)
        if headers[col-1] in ['InvoiceID','OrderID','ProjectID','CustomerID','Currency']:
            cell.number_format='@'
        elif headers[col-1] in ['InvoiceDate','DueDate']:
            cell.number_format='yyyy-mm-dd'
        elif headers[col-1] in ['Amount','PaidAmount','Balance','BalanceTWD']:
            cell.number_format='#,##0.00'

# Fill Summary: project aggregated unpaid TWD only; other source totals blank/zero per source blank request
summary={}
for rec in records:
    summary.setdefault(rec['ProjectID'], Decimal('0.00'))
    summary[rec['ProjectID']] += rec['BalanceTWD']

sum_ws=wb['Summary']
sum_headers=[sum_ws.cell(header_row,c).value for c in range(1, sum_ws.max_column+1)]
for r_idx, project in enumerate(sorted(summary), start=start_row):
    style_like_header(sum_ws, r_idx)
    row_map={
        'ProjectID': project,
        'OrderTWD': 0,
        'InvoiceBalanceTWD': float(summary[project].quantize(Decimal('0.01'))),
        'ApprovedExpenseTWD': 0,
        'AcceptedQuoteTWD': 0,
    }
    for c_idx,h in enumerate(sum_headers, start=1):
        val=row_map.get(h)
        cell=sum_ws.cell(r_idx,c_idx,value=val)
        if h=='ProjectID':
            cell.number_format='@'
        elif h:
            cell.number_format='#,##0.00'
# update totals row formulas to exact data range
last_sum_row=start_row+len(summary)-1
if len(summary)>0:
    for col in range(2, sum_ws.max_column+1):
        letter=sum_ws.cell(header_row,col).column_letter
        sum_ws.cell(2,col).value=f'=SUM({letter}{start_row}:{letter}{last_sum_row})'
else:
    for col in range(2, sum_ws.max_column+1):
        sum_ws.cell(2,col).value=0

# Fill Exceptions: Balance > 0 and DueDate < 2026-09-30 only, not same-day due
exc_ws=wb['Exceptions']
exc_headers=[exc_ws.cell(header_row,c).value for c in range(1, exc_ws.max_column+1)]
exceptions=[rec for rec in records if rec['Balance'] > 0 and rec['DueDate'] < ASOF]
exceptions.sort(key=lambda r: (r['DueDate'], r['InvoiceID']))
for r_idx, rec in enumerate(exceptions, start=start_row):
    style_like_header(exc_ws, r_idx)
    row_map={'Source':'Invoices','DocumentID':rec['InvoiceID'],'ProjectID':rec['ProjectID'],'Reason':'Overdue'}
    for c_idx,h in enumerate(exc_headers, start=1):
        cell=exc_ws.cell(r_idx,c_idx,value=row_map.get(h))
        if h in ['Source','DocumentID','ProjectID','Reason']:
            cell.number_format='@'

wb.save(out_path)
print('Saved', out_path)

# Reopen and verify all stated requirements from saved file
vwb=load_workbook(out_path, data_only=False)
# Check other source sheets blank after row 4
for s in ['Orders','Expenses','Quotes']:
    ws=vwb[s]
    nonempty=[]
    for row in ws.iter_rows(min_row=start_row):
        for cell in row:
            if cell.value is not None:
                nonempty.append(cell.coordinate)
    print(f'{s} data cells after row 4:', len(nonempty))
    if nonempty:
        raise SystemExit(f'{s} is not blank: {nonempty[:10]}')
# Invoice checks
ws=vwb['Invoices']
saved_headers=[ws.cell(header_row,c).value for c in range(1, ws.max_column+1)]
if saved_headers != ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']:
    raise SystemExit(f'Invoice headers changed: {saved_headers}')
saved_inv=[]
for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, values_only=True):
    if any(v is not None for v in row):
        saved_inv.append(row)
print('Saved invoice rows:', len(saved_inv))
if len(saved_inv)!=len(records):
    raise SystemExit(f'Invoice row count mismatch {len(saved_inv)} != {len(records)}')
for row in saved_inv:
    inv_id, order_id, project_id, cust_id, inv_date, due_date, cur, amount, paid, bal, bal_twd=row
    if not isinstance(inv_id, str) or not isinstance(order_id, str):
        raise SystemExit('InvoiceID/OrderID not saved as text')
    if dparse(inv_date) < dparse(FROM) or dparse(inv_date) > dparse(TO):
        raise SystemExit('Saved invoice outside September')
    calc=(money(bal)*fx[cur]).quantize(Decimal('0.01'))
    if money(bal_twd) != calc:
        raise SystemExit(f'BalanceTWD mismatch for {inv_id}: {bal_twd} vs {calc}')
# Summary checks
ws=vwb['Summary']
saved_summary=[]
for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, values_only=True):
    if any(v is not None for v in row):
        saved_summary.append(row)
print('Saved summary rows:', len(saved_summary), saved_summary)
if len(saved_summary)!=len(summary):
    raise SystemExit('Summary project count mismatch')
for project, order_twd, inv_bal_twd, exp_twd, quote_twd in saved_summary:
    if order_twd != 0 or exp_twd != 0 or quote_twd != 0:
        raise SystemExit(f'Non-invoice source summary not zero for {project}')
    if money(inv_bal_twd) != summary[project].quantize(Decimal('0.01')):
        raise SystemExit(f'Summary mismatch for {project}')
# Exceptions checks
ws=vwb['Exceptions']
saved_exc=[]
for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, values_only=True):
    if any(v is not None for v in row):
        saved_exc.append(row)
print('Saved exception rows:', len(saved_exc), saved_exc)
if len(saved_exc)!=len(exceptions):
    raise SystemExit('Exception count mismatch')
exc_ids={r['InvoiceID'] for r in exceptions}
for source, doc_id, project_id, reason in saved_exc:
    if source!='Invoices' or reason!='Overdue' or doc_id not in exc_ids:
        raise SystemExit(f'Invalid exception row: {(source, doc_id, project_id, reason)}')
# Specifically ensure due date equal 2026-09-30 not flagged
same_day=[r['InvoiceID'] for r in records if r['Balance']>0 and r['DueDate']==ASOF]
flagged_same_day=[x for x in same_day if x in {row[1] for row in saved_exc}]
print('Same-day due open invoices not flagged:', same_day)
if flagged_same_day:
    raise SystemExit(f'Same-day due invoices incorrectly flagged: {flagged_same_day}')
print('Verification OK')
print('Output file:', out_path)
