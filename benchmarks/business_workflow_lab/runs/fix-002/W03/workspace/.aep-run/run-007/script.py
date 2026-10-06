from pathlib import Path
import urllib.request, urllib.parse, http.cookiejar, csv, io
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
from copy import copy

WORK=Path.cwd(); IN=WORK/'input'; OUT=WORK/'output'; OUT.mkdir(exist_ok=True)
BASE='http://127.0.0.1:8810'; FROM='2026-09-01'; TO='2026-09-30'; ASOF=date(2026,9,30)

def dparse(s):
    if isinstance(s, datetime): return s.date()
    if isinstance(s, date): return s
    return datetime.strptime(str(s), '%Y-%m-%d').date()

def money(v):
    return Decimal(str(v).replace(',','')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# Inspect and load FX rates
print('Workspace:', WORK)
print('Inputs:', [p.name for p in sorted(IN.glob('*'))])
fx={}
with (IN/'fx_rates.csv').open('r', encoding='utf-8-sig', newline='') as f:
    rdr=csv.reader(f); header=next(rdr)
    print('FX header:', header)
    for i,row in enumerate(rdr):
        if i < 3: print('FX sample row:', row)
        if len(row)>=2 and row[0].strip(): fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates loaded:', fx)
wb0=load_workbook(IN/'project_template.xlsx', data_only=False)
print('Template sheets:', wb0.sheetnames)
for s in ['Invoices','Summary','Exceptions','Orders','Expenses','Quotes']:
    sh=wb0[s]
    print(s, 'headers:', [sh.cell(4,c).value for c in range(1, sh.max_column+1)])

# Query ERP export using documented query parameter names from/to
cj=http.cookiejar.CookieJar(); opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
login_data=urllib.parse.urlencode({'next':'/invoices','sso':'company'}).encode('utf-8')
with opener.open(urllib.request.Request(BASE+'/login', data=login_data, method='POST'), timeout=20) as r:
    print('Logged in:', r.status, r.geturl())
params=urllib.parse.urlencode({'from':FROM, 'to':TO})
export_url=BASE+'/invoices/export.csv?'+params
with opener.open(export_url, timeout=20) as r:
    raw=r.read(); print('Invoice export:', r.geturl(), r.headers.get('content-type'), len(raw), 'bytes')
text=raw.decode('utf-8-sig')
print('Export sample lines:')
for line in text.splitlines()[:6]: print(' ', line)
reader=csv.DictReader(io.StringIO(text)); rows=list(reader)
print('Invoice columns:', reader.fieldnames); print('Invoice rows:', len(rows))
expected_cols=['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance']
missing=[c for c in expected_cols if c not in (reader.fieldnames or [])]
if missing: raise SystemExit(f'Missing expected invoice columns: {missing}')
if not rows: raise SystemExit('No invoices returned')

records=[]
for row in rows:
    inv_date=dparse(row['InvoiceDate'])
    if not (dparse(FROM) <= inv_date <= dparse(TO)):
        raise SystemExit(f'Invoice outside requested invoice-date range: {row}')
    cur=row['Currency'].strip()
    if cur not in fx: raise SystemExit(f'Missing FX rate for {cur}')
    amount=money(row['Amount']); paid=money(row['PaidAmount']); bal=money(row['Balance'])
    if bal != (amount-paid).quantize(Decimal('0.01')):
        raise SystemExit(f'Balance mismatch for {row["InvoiceID"]}')
    records.append({
        'InvoiceID': str(row['InvoiceID']), 'OrderID': str(row['OrderID']),
        'ProjectID': str(row['ProjectID']) if row['ProjectID'] is not None else '',
        'CustomerID': str(row['CustomerID']), 'InvoiceDate': inv_date, 'DueDate': dparse(row['DueDate']),
        'Currency': cur, 'Amount': amount, 'PaidAmount': paid, 'Balance': bal,
        'BalanceTWD': (bal*fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP),
    })
records.sort(key=lambda r: (r['InvoiceDate'], r['InvoiceID']))

# Write workbook
out_path=OUT/'project_invoice_report_2026-09.xlsx'
wb=load_workbook(IN/'project_template.xlsx')
header_row=4; start_row=5

def clear_sheet_data(ws):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row-start_row+1)

def style_like_header(ws, row_idx):
    for col in range(1, ws.max_column+1):
        src=ws.cell(header_row, col); dst=ws.cell(row_idx, col)
        if src.has_style:
            dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border)
            dst.alignment=copy(src.alignment); dst.number_format=src.number_format; dst.protection=copy(src.protection)

for s in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    clear_sheet_data(wb[s])

# Report metadata
rep=wb['Report']; rep['B3']=FROM; rep['D3']=TO
rep['B4']='Invoices only; as-of 2026-09-30; overdue when Balance > 0 and DueDate < 2026-09-30'
rep['B5']=datetime.now().strftime('%Y-%m-%d %H:%M:%S')

# Invoices
inv_ws=wb['Invoices']; headers=[inv_ws.cell(header_row,c).value for c in range(1, inv_ws.max_column+1)]
for idx, rec in enumerate(records, start=start_row):
    style_like_header(inv_ws, idx)
    for col,h in enumerate(headers, start=1):
        val=rec[h]
        if h in ['Amount','PaidAmount','Balance','BalanceTWD']:
            val=float(val)
        cell=inv_ws.cell(idx,col,value=val)
        if h in ['InvoiceID','OrderID','ProjectID','CustomerID','Currency']:
            cell.number_format='@'
        elif h in ['InvoiceDate','DueDate']:
            cell.number_format='yyyy-mm-dd'
        elif h in ['Amount','PaidAmount','Balance','BalanceTWD']:
            cell.number_format='#,##0.00'

# Summary by project unpaid TWD; non-invoice source columns kept zero, and source sheets blank
summary={}
for rec in records:
    project=rec['ProjectID']
    summary.setdefault(project, Decimal('0.00'))
    summary[project] += rec['BalanceTWD']
sum_ws=wb['Summary']; sum_headers=[sum_ws.cell(header_row,c).value for c in range(1, sum_ws.max_column+1)]
for r_idx, project in enumerate(sorted(summary), start=start_row):
    style_like_header(sum_ws, r_idx)
    row_map={'ProjectID':project,'OrderTWD':0,'InvoiceBalanceTWD':float(summary[project].quantize(Decimal('0.01'))),'ApprovedExpenseTWD':0,'AcceptedQuoteTWD':0}
    for c_idx,h in enumerate(sum_headers, start=1):
        cell=sum_ws.cell(r_idx,c_idx,value=row_map.get(h))
        if h=='ProjectID': cell.number_format='@'
        elif h: cell.number_format='#,##0.00'
last_sum_row=start_row+len(summary)-1
for col in range(2, sum_ws.max_column+1):
    letter=sum_ws.cell(header_row,col).column_letter
    sum_ws.cell(2,col).value=f'=SUM({letter}{start_row}:{letter}{last_sum_row})' if len(summary)>0 else 0

# Exceptions
exc_ws=wb['Exceptions']; exc_headers=[exc_ws.cell(header_row,c).value for c in range(1, exc_ws.max_column+1)]
exceptions=[rec for rec in records if rec['Balance'] > 0 and rec['DueDate'] < ASOF]
exceptions.sort(key=lambda r: (r['DueDate'], r['InvoiceID']))
for r_idx, rec in enumerate(exceptions, start=start_row):
    style_like_header(exc_ws, r_idx)
    row_map={'Source':'Invoices','DocumentID':rec['InvoiceID'],'ProjectID':rec['ProjectID'],'Reason':'Overdue'}
    for c_idx,h in enumerate(exc_headers, start=1):
        cell=exc_ws.cell(r_idx,c_idx,value=row_map.get(h)); cell.number_format='@'
wb.save(out_path)
print('Saved:', out_path)

# Verify saved output from file
vwb=load_workbook(out_path, data_only=False)
# source sheets blank
for s in ['Orders','Expenses','Quotes']:
    ws=vwb[s]
    nonempty=[cell.coordinate for row in ws.iter_rows(min_row=start_row) for cell in row if cell.value is not None]
    print(f'CHECK {s} blank data cells:', len(nonempty))
    if nonempty: raise SystemExit(f'{s} not blank after row 4: {nonempty[:10]}')
# Invoices
ws=vwb['Invoices']; saved_headers=[ws.cell(header_row,c).value for c in range(1, ws.max_column+1)]
required_invoice_headers=['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
print('CHECK invoice headers:', saved_headers)
if saved_headers != required_invoice_headers: raise SystemExit('Invoice headers incorrect')
saved_inv=[row for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, values_only=True) if any(v is not None for v in row)]
print('CHECK invoice row count:', len(saved_inv))
if len(saved_inv)!=len(records): raise SystemExit('Invoice row count mismatch')
for row in saved_inv:
    inv_id, order_id, project_id, cust_id, inv_date, due_date, cur, amount, paid, bal, bal_twd=row
    if not isinstance(inv_id, str) or not isinstance(order_id, str): raise SystemExit('IDs not text')
    if not (dparse(FROM) <= dparse(inv_date) <= dparse(TO)): raise SystemExit('Saved invoice outside date range')
    float(amount); float(bal_twd)
    calc=(money(bal)*fx[cur]).quantize(Decimal('0.01'))
    if money(bal_twd)!=calc: raise SystemExit(f'BalanceTWD mismatch for {inv_id}')
# Summary
ws=vwb['Summary']; saved_summary=[row for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, values_only=True) if any(v is not None for v in row)]
print('CHECK summary row count:', len(saved_summary)); print('CHECK summary rows:', saved_summary)
if len(saved_summary)!=len(summary): raise SystemExit('Summary row count mismatch')
for project, order_twd, inv_bal_twd, exp_twd, quote_twd in saved_summary:
    if order_twd != 0 or exp_twd != 0 or quote_twd != 0: raise SystemExit(f'Other source totals non-zero for {project}')
    if money(inv_bal_twd)!=summary[project].quantize(Decimal('0.01')): raise SystemExit(f'Summary amount mismatch for {project}')
# Exceptions
ws=vwb['Exceptions']; saved_exc=[row for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, values_only=True) if any(v is not None for v in row)]
print('CHECK exception row count:', len(saved_exc)); print('CHECK exception rows:', saved_exc)
if len(saved_exc)!=len(exceptions): raise SystemExit('Exception row count mismatch')
exc_ids={r['InvoiceID'] for r in exceptions}
for source, doc_id, project_id, reason in saved_exc:
    if source!='Invoices' or reason!='Overdue' or doc_id not in exc_ids: raise SystemExit(f'Bad exception row {source,doc_id,project_id,reason}')
same_day=[r['InvoiceID'] for r in records if r['Balance']>0 and r['DueDate']==ASOF]
flagged_same_day=[x for x in same_day if x in {row[1] for row in saved_exc}]
print('CHECK due-on-2026-09-30 open invoices not overdue:', same_day)
if flagged_same_day: raise SystemExit(f'Same-day due flagged: {flagged_same_day}')
print('Verification OK')
print('Projects summarized:', [(k, str(v)) for k,v in sorted(summary.items())])
print('Output file:', out_path)
