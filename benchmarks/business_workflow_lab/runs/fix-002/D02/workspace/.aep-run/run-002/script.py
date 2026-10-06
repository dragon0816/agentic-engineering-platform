from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path('.')
print('Output files:')
for p in sorted((ws/'output').glob('*')):
    print(p, p.stat().st_size)
    if p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print(' sheets', wb.sheetnames)
        for sh in wb.sheetnames:
            s=wb[sh]
            print(' ', sh, s.max_row, s.max_column)
            for r in s.iter_rows(min_row=1, max_row=min(5,s.max_row), values_only=True): print('   ', r)
print('\nInstructions full:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
s=wb['Instructions']
for r in s.iter_rows(values_only=True): print(r)
print('\nSubmitted preview:')
wb=load_workbook(ws/'output'/'submitted_export.xlsx', data_only=False)
for sh in wb.sheetnames:
    s=wb[sh]
    print(sh, s.max_row, s.max_column)
    for r in s.iter_rows(min_row=1, max_row=min(10,s.max_row), values_only=True): print(r)
print('\nApproved preview:')
wb=load_workbook(ws/'output'/'approved_export.xlsx', data_only=False)
for sh in wb.sheetnames:
    s=wb[sh]
    print(sh, s.max_row, s.max_column)
    for r in s.iter_rows(min_row=1, max_row=min(10,s.max_row), values_only=True): print(r)