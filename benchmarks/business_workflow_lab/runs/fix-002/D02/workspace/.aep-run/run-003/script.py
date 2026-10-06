from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv

workspace = Path('.')
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

template_path = input_dir / 'expense_template.xlsx'
submitted_path = output_dir / 'submitted_export.xlsx'
approved_path = output_dir / 'approved_export.xlsx'
fx_path = input_dir / 'fx_rates.csv'
out_path = output_dir / 'expense_report_2026-09_submitted_approved.xlsx'

# Read fx rates; the Note column may contain commas, but first two columns are stable.
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    headers = next(reader)
    for row in reader:
        if not row or len(row) < 2:
            continue
        fx[row[0]] = Decimal(row[1])
print('FX rates loaded:', fx)

# Read exports from the desktop app. These include detail fields ReceiptID, Description, MissingReceipt.
def read_expense_export(path):
    wb = load_workbook(path, data_only=True)
    ws = wb['Expenses']
    headers = [c.value for c in ws[1]]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in row):
            continue
        rec = dict(zip(headers, row))
        rows.append(rec)
    return headers, rows

submitted_headers, submitted_rows = read_expense_export(submitted_path)
approved_headers, approved_rows = read_expense_export(approved_path)
print('Submitted rows:', len(submitted_rows), 'headers:', submitted_headers)
print('Approved rows:', len(approved_rows), 'headers:', approved_headers)

all_rows = submitted_rows + approved_rows
# Filter defensively to the requested period/statuses and sort deterministically.
wanted_statuses = {'Submitted', 'Approved'}
all_rows = [r for r in all_rows if str(r.get('ExpenseDate')) >= '2026-09-01' and str(r.get('ExpenseDate')) <= '2026-09-30' and r.get('Status') in wanted_statuses]
all_rows.sort(key=lambda r: (str(r['ExpenseDate']), str(r['ExpenseID'])))

# Convert/export row values.
def money_decimal(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def to_bool_text(value):
    # Desktop export returns real booleans; keep output stable as text per template instruction.
    if isinstance(value, bool):
        return 'true' if value else 'false'
    s = '' if value is None else str(value).strip().lower()
    return 'true' if s in {'true', '1', 'yes', 'y'} else 'false'

def amount_twd(amount, currency):
    return (money_decimal(amount) * fx[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# Load template.
wb = load_workbook(template_path)

# Helper to clear data rows while preserving header and sheets.
def clear_below(ws, start_row, max_col):
    if ws.max_row >= start_row:
        for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

# Report sheet metadata.
rpt = wb['Report']
rpt['B3'] = '2026-09-01'
rpt['D3'] = '2026-09-30'
rpt['B4'] = 'ExpenseDate 2026-09-01..2026-09-30; Status in Submitted, Approved; all employees/projects'
rpt['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
rpt['B7'] = 'Expense Viewer (synthetic lab) desktop export: submitted_export.xlsx; approved_export.xlsx'
rpt['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')
missing_count = sum(1 for r in all_rows if to_bool_text(r.get('MissingReceipt')) == 'true')
rpt['B9'] = f'{len(all_rows)} expenses retained. {missing_count} expense(s) missing receipt are listed in Exceptions and retained in Details.'

# Details sheet.
details = wb['Details']
clear_below(details, 5, 14)
# Copy row 4 style into rows we write.
for idx, rec in enumerate(all_rows, start=5):
    values = []
    amount = money_decimal(rec['Amount'])
    twd = amount_twd(amount, rec['Currency'])
    missing = to_bool_text(rec.get('MissingReceipt'))
    needs_review = 'true' if missing == 'true' else 'false'
    values = [
        str(rec['ExpenseID']),
        str(rec['EmployeeID']),
        str(rec['EmployeeName']),
        str(rec['ProjectID']),
        str(rec['ExpenseDate']),
        str(rec['Category']),
        str(rec['Currency']),
        float(amount),
        float(twd),
        str(rec['Status']),
        '' if rec.get('ReceiptID') is None else str(rec.get('ReceiptID')),
        '' if rec.get('Description') is None else str(rec.get('Description')),
        missing,
        needs_review,
    ]
    for col, val in enumerate(values, start=1):
        c = details.cell(row=idx, column=col, value=val)
        # Keep header-ish styling and formats simple.
        if idx == 5:
            pass
        c._style = copy(details.cell(row=4, column=col)._style)
        if col in (8, 9):
            c.number_format = '0.00'
        elif col == 5:
            c.number_format = 'yyyy-mm-dd'
        else:
            c.number_format = '@'

# Summary sheet: status/category/currency summary.
summary = wb['Summary']
clear_below(summary, 5, 6)
last_detail_row = 4 + len(all_rows) if all_rows else 5
summary['B2'] = f'=SUM(Details!I5:I{last_detail_row})' if all_rows else '=SUM(Details!I5:I5)'
summary_map = {}
for rec in all_rows:
    key = (str(rec['Status']), str(rec['Category']), str(rec['Currency']))
    amount = money_decimal(rec['Amount'])
    twd = amount_twd(amount, rec['Currency'])
    if key not in summary_map:
        summary_map[key] = {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0}
    summary_map[key]['amount'] += amount
    summary_map[key]['twd'] += twd
    summary_map[key]['count'] += 1
for row_idx, (key, agg) in enumerate(sorted(summary_map.items()), start=5):
    vals = [key[0], key[1], key[2], float(agg['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), float(agg['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), int(agg['count'])]
    for col, val in enumerate(vals, start=1):
        c = summary.cell(row=row_idx, column=col, value=val)
        c._style = copy(summary.cell(row=4, column=col)._style)
        if col in (4,5): c.number_format = '0.00'

# Exceptions for missing receipt; do not remove them from Details.
exceptions = wb['Exceptions']
clear_below(exceptions, 5, 4)
exc_rows = [r for r in all_rows if to_bool_text(r.get('MissingReceipt')) == 'true']
for row_idx, rec in enumerate(exc_rows, start=5):
    vals = [str(rec['ExpenseID']), str(rec['EmployeeID']), str(rec['ExpenseDate']), 'Missing receipt']
    for col, val in enumerate(vals, start=1):
        c = exceptions.cell(row=row_idx, column=col, value=val)
        c._style = copy(exceptions.cell(row=4, column=col)._style)
        c.number_format = '@'

# Save output.
wb.save(out_path)
print('Wrote', out_path)

# Verification: reopen saved output and confirm all requirements.
wb2 = load_workbook(out_path, data_only=False)
errors = []
# Sheet names intact.
expected_sheets = ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
if wb2.sheetnames != expected_sheets:
    errors.append(f'Sheet names/order changed: {wb2.sheetnames}')

# Details checks.
d2 = wb2['Details']
headers2 = [d2.cell(4,c).value for c in range(1,15)]
expected_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
if headers2 != expected_headers:
    errors.append(f'Details headers mismatch: {headers2}')
actual_detail_rows = []
for r in range(5, 5+len(all_rows)):
    actual_detail_rows.append([d2.cell(r,c).value for c in range(1,15)])
if len(actual_detail_rows) != 22:
    errors.append(f'Expected 22 detail rows, got {len(actual_detail_rows)}')
ids = [row[0] for row in actual_detail_rows]
if len(ids) != len(set(ids)):
    errors.append('Duplicate ExpenseID in Details')
if set(ids) != {str(r['ExpenseID']) for r in all_rows}:
    errors.append('Details ExpenseID set differs from source rows')
for row in actual_detail_rows:
    if row[9] not in ('Submitted','Approved'):
        errors.append(f'Bad status {row[9]} for {row[0]}')
    if not ('2026-09-01' <= str(row[4]) <= '2026-09-30'):
        errors.append(f'Bad date {row[4]} for {row[0]}')
    if not isinstance(row[7], (int,float)) or not isinstance(row[8], (int,float)):
        errors.append(f'Amounts not numeric for {row[0]}: {row[7]}, {row[8]}')
    if row[12] not in ('true','false') or row[13] not in ('true','false'):
        errors.append(f'Boolean text invalid for {row[0]}: {row[12]}, {row[13]}')
    if row[12] == 'false' and not row[10]:
        errors.append(f'Non-missing receipt has blank ReceiptID: {row[0]}')

# Exceptions checks.
e2 = wb2['Exceptions']
exc_actual = []
for r in range(5, e2.max_row+1):
    vals = [e2.cell(r,c).value for c in range(1,5)]
    if any(v is not None for v in vals):
        exc_actual.append(vals)
expected_missing_ids = {str(r['ExpenseID']) for r in exc_rows}
actual_missing_ids = {row[0] for row in exc_actual}
if expected_missing_ids != actual_missing_ids:
    errors.append(f'Exception IDs mismatch: expected {expected_missing_ids}, got {actual_missing_ids}')
if not expected_missing_ids.issubset(set(ids)):
    errors.append('Missing receipt expenses were not retained in Details')

# Summary / formula checks.
s2 = wb2['Summary']
if s2['B2'].value != f'=SUM(Details!I5:I{4+len(all_rows)})':
    errors.append(f'Summary B2 formula wrong: {s2["B2"].value}')
summary_nonempty = 0
for r in range(5, s2.max_row+1):
    if any(s2.cell(r,c).value is not None for c in range(1,7)):
        summary_nonempty += 1
if summary_nonempty != len(summary_map):
    errors.append(f'Summary row count mismatch: expected {len(summary_map)}, got {summary_nonempty}')

if errors:
    print('VERIFICATION FAILED')
    for e in errors:
        print(' -', e)
    raise SystemExit(1)
else:
    print('VERIFICATION PASSED')
    print('Details rows:', len(actual_detail_rows))
    print('Exceptions rows:', len(exc_actual), 'IDs:', sorted(actual_missing_ids))
    print('Summary groups:', summary_nonempty)
    print('Output file:', out_path)
