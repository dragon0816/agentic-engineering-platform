import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path('.')
print('CWD', Path.cwd())
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(p, p.stat().st_size)
print('\nManual preview:')
manual=ws/'input'/'expense_manual.md'
print(manual.read_text(encoding='utf-8')[:4000])
print('\nFX rates preview:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    for i,row in enumerate(csv.reader(f)):
        print(row)
        if i>=10: break
print('\nWorkbook sheets/headers/sample:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    s=wb[sh]
    print('SHEET', sh, 'max_row', s.max_row, 'max_col', s.max_column)
    for r in s.iter_rows(min_row=1, max_row=min(5,s.max_row), values_only=True): print(r)
    print('---')