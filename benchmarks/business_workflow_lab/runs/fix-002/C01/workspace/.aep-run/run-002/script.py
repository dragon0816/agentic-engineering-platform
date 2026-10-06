from pathlib import Path
from openpyxl import load_workbook
import subprocess, json, os
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
print('Instructions:')
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
cli=r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
print('\nCLI exists:', Path(cli).exists(), cli)
cmd=[cli,'quotes','list','--customer','C001','--from','2026-09-01','--to','2026-09-30','--format','json']
print('running', cmd)
p=subprocess.run(cmd, capture_output=True, text=True, shell=False)
print('return', p.returncode)
print('stdout', p.stdout[:2000])
print('stderr', p.stderr[:1000])
if p.returncode==0:
    data=json.loads(p.stdout)
    print('count', data.get('count'))
    for item in data.get('items',[])[:3]: print(item)
    if data.get('items'):
        q=data['items'][0]['QuoteID']
        p2=subprocess.run([cli,'quotes','get',q,'--format','json'], capture_output=True, text=True, shell=False)
        print('get return', p2.returncode)
        print(p2.stdout[:2000])
        print('get err', p2.stderr[:1000])
