import os, sys, json, csv, subprocess, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from zoneinfo import ZoneInfo
from copy import copy
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

customer_id = str(params.get('customer_id', '')).strip()
from_date = str(params.get('from_date', '')).strip()
to_date = str(params.get('to_date', '')).strip()
cli_path = str(params.get('cli_path', '')).strip()
if not customer_id:
    raise ValueError('customer_id is required')
if not from_date or not to_date:
    raise ValueError('from_date and to_date are required')
# Validate date strings and order.
from_dt = date.fromisoformat(from_date)
to_dt = date.fromisoformat(to_date)
if from_dt > to_dt:
    raise ValueError('from_date must be on or before to_date')
cli = Path(cli_path)
if not cli.exists():
    raise FileNotFoundError(f'CLI path not found: {cli_path}')

template_path = input_dir / 'quote_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
if not template_path.exists():
    raise FileNotFoundError(f'Missing template: {template_path}')
if not fx_path.exists():
    raise FileNotFoundError(f'Missing FX rates: {fx_path}')

TWOPLACES = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def cli_cmd(*args):
    if cli.suffix.lower() == '.py':
        return [sys.executable, str(cli), *args]
    return [str(cli), *args]

def run_cli(*args):
    cmd = cli_cmd(*args)
    p = subprocess.run(cmd, capture_output=True, text=True, shell=False)
    if p.returncode != 0:
        raise RuntimeError(f'CLI failed exit={p.returncode}: {" ".join(cmd)}\nstderr={p.stderr}')
    try:
        return json.loads(p.stdout)
    except json.JSONDecodeError as e:
        raise RuntimeError(f'CLI returned non-JSON stdout for {" ".join(cmd)}: {p.stdout[:500]}') from e

# Read fx rates; the Note field may contain unquoted commas in the fixture, so use only first two columns.
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader, None)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
if not fx:
    raise ValueError('No FX rates loaded')

listed = run_cli('quotes', 'list', '--customer', customer_id, '--from', from_date, '--to', to_date, '--format', 'json')
if listed.get('schema_version') != 1:
    raise ValueError(f'Unexpected quotes list schema_version: {listed.get("schema_version")}')
items = listed.get('items', [])
if listed.get('count') != len(items):
    raise ValueError(f'CLI count mismatch: count={listed.get("count")} items={len(items)}')

quote_ids = [str(item['QuoteID']) for item in items]
headers = []
lines = []
for qid in quote_ids:
    detail = run_cli('quotes', 'get', qid, '--format', 'json')
    if detail.get('schema_version') != 1:
        raise ValueError(f'Unexpected quote detail schema_version for {qid}: {detail.get("schema_version")}')
    h = detail.get('header') or {}
    if str(h.get('QuoteID')) != qid:
        raise ValueError(f'QuoteID mismatch in detail: list={qid}, detail={h.get("QuoteID")}')
    currency = h.get('Currency')
    if currency not in fx:
        raise ValueError(f'Missing FX rate for currency {currency} quote {qid}')
    net = money(h['NetAmount'])
    rate = fx[currency]
    h['_NetAmountTWD'] = money(net * rate)
    headers.append(h)
    for line in detail.get('line_items', []):
        row = dict(line)
        row['QuoteID'] = qid
        lines.append(row)

# Sort for stable output.
headers.sort(key=lambda h: (h.get('QuoteDate',''), h.get('QuoteID','')))
quote_order = {h['QuoteID']: i for i, h in enumerate(headers)}
lines.sort(key=lambda r: (quote_order.get(r['QuoteID'], 999999), r.get('ProductID',''), r.get('ProductName','')))

wb = load_workbook(template_path)
required_sheets = ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
missing = [s for s in required_sheets if s not in wb.sheetnames]
if missing:
    raise ValueError(f'Template missing sheets: {missing}')

# Clear old data rows but preserve row 1-4 headers/notes and sheet names.
def clear_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

for sheet in ['Quotes', 'Lines', 'Summary']:
    clear_data(wb[sheet], 5)

report = wb['Report']
now = datetime.now(ZoneInfo('Asia/Taipei')).replace(microsecond=0).isoformat()
report['B3'] = from_date
report['D3'] = to_date
report['B4'] = f'CustomerID={customer_id}; QuoteDate {from_date}..{to_date}'
report['B5'] = now
report['B7'] = 'quote_cli'
report['D7'] = now
report['B9'] = f'取得 {len(headers)} 筆報價、{len(lines)} 筆明細；匯率來源 fx_rates.csv。'

quotes_ws = wb['Quotes']
quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for r_idx, h in enumerate(headers, start=5):
    values = [
        str(h['QuoteID']), str(h['CustomerID']), str(h['ProjectID']), h['QuoteDate'], h['ValidUntil'],
        str(h['Status']), str(h['Currency']), float(money(h['Subtotal'])), float(Decimal(str(h['DiscountPercent']))),
        float(money(h['NetAmount'])), float(h['_NetAmountTWD'])
    ]
    for c_idx, value in enumerate(values, start=1):
        cell = quotes_ws.cell(r_idx, c_idx, value)
        if c_idx in [1,2,3,6,7]:
            cell.number_format = '@'
        elif c_idx in [4,5]:
            cell.number_format = 'yyyy-mm-dd'
        elif c_idx in [8,10,11]:
            cell.number_format = '0.00'
        elif c_idx == 9:
            cell.number_format = '0.00'

lines_ws = wb['Lines']
line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for r_idx, line in enumerate(lines, start=5):
    values = [str(line['QuoteID']), str(line['ProductID']), str(line['ProductName']), int(line['Quantity']), float(money(line['UnitPrice'])), float(money(line['LineAmount']))]
    for c_idx, value in enumerate(values, start=1):
        cell = lines_ws.cell(r_idx, c_idx, value)
        if c_idx in [1,2,3]:
            cell.number_format = '@'
        elif c_idx == 4:
            cell.number_format = '0'
        elif c_idx in [5,6]:
            cell.number_format = '0.00'

summary_ws = wb['Summary']
summary = {}
for h in headers:
    key = (str(h['Status']), str(h['Currency']))
    if key not in summary:
        summary[key] = {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0}
    summary[key]['amount'] += money(h['NetAmount'])
    summary[key]['amount_twd'] += money(h['_NetAmountTWD'])
    summary[key]['count'] += 1
for r_idx, ((status, currency), agg) in enumerate(sorted(summary.items()), start=5):
    values = ['Status', status, currency, float(money(agg['amount'])), float(money(agg['amount_twd'])), int(agg['count'])]
    for c_idx, value in enumerate(values, start=1):
        cell = summary_ws.cell(r_idx, c_idx, value)
        if c_idx in [1,2,3]:
            cell.number_format = '@'
        elif c_idx in [4,5]:
            cell.number_format = '0.00'
        elif c_idx == 6:
            cell.number_format = '0'
last_quote_row = 4 + len(headers) if headers else 5
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
summary_ws['B2'].number_format = '0.00'

safe_customer = ''.join(ch if ch.isalnum() or ch in '-_' else '_' for ch in customer_id)
out_path = output_dir / f'quote_report_{safe_customer}_{from_date}_to_{to_date}.xlsx'
wb.save(out_path)

# Reopen and validate saved workbook from disk.
check_wb = load_workbook(out_path, data_only=False)
assert check_wb.sheetnames == wb.sheetnames, 'Sheet names/order changed unexpectedly'
qws = check_wb['Quotes']; lws = check_wb['Lines']; sws = check_wb['Summary']; rws = check_wb['Report']
expected_quote_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
expected_line_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
expected_summary_headers = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
assert [qws.cell(4,c).value for c in range(1,12)] == expected_quote_headers
assert [lws.cell(4,c).value for c in range(1,7)] == expected_line_headers
assert [sws.cell(4,c).value for c in range(1,7)] == expected_summary_headers
assert rws['B3'].value == from_date and rws['D3'].value == to_date
assert qws.max_row == 4 + len(headers), f'Quote row count mismatch: {qws.max_row} vs {4+len(headers)}'
assert lws.max_row == 4 + len(lines), f'Line row count mismatch: {lws.max_row} vs {4+len(lines)}'
assert sws.max_row == (4 + len(summary) if summary else 4), f'Summary row count mismatch: {sws.max_row}'
# Text IDs and numeric values.
for row in range(5, qws.max_row + 1):
    assert isinstance(qws.cell(row,1).value, str)
    assert isinstance(qws.cell(row,2).value, str)
    assert isinstance(qws.cell(row,3).value, str)
    assert isinstance(qws.cell(row,8).value, (int,float))
    assert isinstance(qws.cell(row,10).value, (int,float))
    assert isinstance(qws.cell(row,11).value, (int,float))
for row in range(5, lws.max_row + 1):
    assert isinstance(lws.cell(row,1).value, str)
    assert isinstance(lws.cell(row,2).value, str)
    assert isinstance(lws.cell(row,4).value, int)
    assert isinstance(lws.cell(row,5).value, (int,float))
    assert isinstance(lws.cell(row,6).value, (int,float))
# Validate no blank data rows and summary totals by status/currency.
quote_net_twd_sum = Decimal('0')
summary_from_quotes = {}
for row in range(5, qws.max_row + 1):
    assert all(qws.cell(row,c).value is not None for c in range(1,12)), f'Blank quote value in row {row}'
    status = qws.cell(row,6).value; currency = qws.cell(row,7).value
    amount = money(qws.cell(row,10).value); amount_twd = money(qws.cell(row,11).value)
    quote_net_twd_sum += amount_twd
    summary_from_quotes.setdefault((status,currency), {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
    summary_from_quotes[(status,currency)]['amount'] += amount
    summary_from_quotes[(status,currency)]['amount_twd'] += amount_twd
    summary_from_quotes[(status,currency)]['count'] += 1
summary_from_sheet = {}
for row in range(5, sws.max_row + 1):
    if sws.cell(row,1).value is None:
        continue
    assert sws.cell(row,1).value == 'Status'
    key = (sws.cell(row,2).value, sws.cell(row,3).value)
    summary_from_sheet[key] = {'amount': money(sws.cell(row,4).value), 'amount_twd': money(sws.cell(row,5).value), 'count': int(sws.cell(row,6).value)}
assert summary_from_sheet == {k: {'amount': money(v['amount']), 'amount_twd': money(v['amount_twd']), 'count': v['count']} for k,v in summary_from_quotes.items()}
expected_formula = f'=SUM(Quotes!K5:K{last_quote_row})'
assert sws['B2'].value == expected_formula, f'Summary B2 formula mismatch: {sws["B2"].value} vs {expected_formula}'

print(json.dumps({
    'output_file': str(out_path),
    'customer_id': customer_id,
    'from_date': from_date,
    'to_date': to_date,
    'quote_count': len(headers),
    'line_count': len(lines),
    'summary_group_count': len(summary),
    'summary_total_twd': str(money(quote_net_twd_sum)),
    'checked': [
        'reopened saved workbook',
        'sheet names and row-4 headers preserved',
        'Report date/condition/source fields filled',
        'Quotes row count and text ID/numeric amount types',
        'Lines row count and text ID/numeric amount types',
        'Summary grouped by Status and Currency',
        'Summary B2 formula range extended'
    ]
}, ensure_ascii=False, indent=2))
