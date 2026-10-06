import csv
import json
import os
import shutil
import subprocess
import time
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
customer_id = str(params['customer_id']).strip()
from_date = str(params['from_date']).strip()
to_date = str(params['to_date']).strip()
cli_path = str(params['cli_path']).strip()

if not customer_id:
    raise ValueError('customer_id is required')
if not Path(cli_path).exists():
    raise FileNotFoundError(f'CLI not found: {cli_path}')

# Fixed local timestamp without zoneinfo/tzdata dependency.
TW = timezone(timedelta(hours=8))
def now_tw_iso():
    return datetime.now(TW).replace(microsecond=0).isoformat()

def money(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [cli_path] + args
    last = None
    for attempt in range(1, 4):
        try:
            p = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True, encoding='utf-8', errors='replace')
        except OSError:
            # .cmd fallback if the platform shell requires command processor handling.
            quoted = ' '.join('"' + a.replace('"', '\\"') + '"' if (' ' in a or '\\' in a or ':' in a) else a for a in cmd)
            p = subprocess.run(quoted, cwd=str(WORKSPACE), shell=True, capture_output=True, text=True, encoding='utf-8', errors='replace')
        last = p
        if p.returncode == 0:
            return p.stdout
        if p.returncode == 4 and attempt < 3:
            time.sleep(1)
            continue
        raise RuntimeError(f'CLI failed rc={p.returncode} args={args}\nSTDERR:\n{p.stderr}\nSTDOUT:\n{p.stdout}')
    raise RuntimeError(f'CLI failed after retries: {last.stderr if last else "unknown"}')

print('Inputs really are:')
print(' workspace=', WORKSPACE)
print(' input files=', sorted(p.name for p in INPUT.iterdir()))
print(' customer_id=', customer_id, 'from_date=', from_date, 'to_date=', to_date)
print(' cli_path=', cli_path)

with open(INPUT / 'fx_rates.csv', encoding='utf-8-sig', newline='') as f:
    rates = {row['Currency']: Decimal(row['RateToTWD']) for row in csv.DictReader(f)}
print(' fx currencies=', sorted(rates))

list_json = run_cli(['quotes', 'list', '--customer', customer_id, '--from', from_date, '--to', to_date, '--format', 'json'])
listing = json.loads(list_json)
if listing.get('schema_version') != 1:
    raise ValueError(f'Unexpected list schema_version: {listing.get("schema_version")}')
items = listing.get('items', [])
print(' quotes listed=', len(items))

headers = []
lines = []
for item in items:
    qid = str(item['QuoteID'])
    detail_json = run_cli(['quotes', 'get', qid, '--format', 'json'])
    detail = json.loads(detail_json)
    if detail.get('schema_version') != 1:
        raise ValueError(f'Unexpected detail schema_version for {qid}: {detail.get("schema_version")}')
    h = detail['header']
    # Keep source field names, but normalize numeric values as Decimal for writing/calculation.
    currency = h['Currency']
    if currency not in rates:
        raise ValueError(f'Missing FX rate for currency {currency}')
    subtotal = money(h['Subtotal'])
    discount = Decimal(str(h['DiscountPercent']))
    net = money(h['NetAmount'])
    expected_net = money(subtotal * (Decimal('1') - discount / Decimal('100')))
    if net != expected_net:
        raise ValueError(f'NetAmount mismatch for {qid}: {net} != {expected_net}')
    h2 = dict(h)
    h2['QuoteID'] = str(h2['QuoteID'])
    h2['CustomerID'] = str(h2['CustomerID'])
    h2['ProjectID'] = str(h2['ProjectID'])
    h2['Subtotal'] = subtotal
    h2['DiscountPercent'] = discount
    h2['NetAmount'] = net
    h2['NetAmountTWD'] = money(net * rates[currency])
    headers.append(h2)
    for line in detail.get('line_items', []):
        qty = Decimal(str(line['Quantity']))
        unit = money(line['UnitPrice'])
        amt = money(line['LineAmount'])
        expected_amt = money(qty * unit)
        if amt != expected_amt:
            raise ValueError(f'LineAmount mismatch for {qid}/{line.get("ProductID")}: {amt} != {expected_amt}')
        lines.append({
            'QuoteID': str(qid),
            'ProductID': str(line['ProductID']),
            'ProductName': str(line['ProductName']),
            'Quantity': int(qty) if qty == qty.to_integral_value() else float(qty),
            'UnitPrice': unit,
            'LineAmount': amt,
        })

headers.sort(key=lambda h: (h['QuoteDate'], h['QuoteID']))
lines.sort(key=lambda l: (l['QuoteID'], l['ProductID']))

out_file = OUTPUT / f'quote_report_{customer_id}_{from_date}_to_{to_date}.xlsx'
shutil.copyfile(INPUT / 'quote_template.xlsx', out_file)
wb = load_workbook(out_file)

# Report
report = wb['Report']
generated_at = now_tw_iso()
report['B3'] = from_date
report['D3'] = to_date
report['B4'] = f'CustomerID={customer_id}; QuoteDate between {from_date} and {to_date}'
report['B5'] = generated_at
report['B7'] = 'quote_cli mock; fx_rates.csv'
report['D7'] = generated_at

# Clear/fill Quotes from row 5, preserving row 4 headers.
quotes_ws = wb['Quotes']
for row in quotes_ws.iter_rows(min_row=5, max_row=max(quotes_ws.max_row, 5), max_col=quotes_ws.max_column):
    for cell in row:
        cell.value = None
quote_cols = [quotes_ws.cell(4, c).value for c in range(1, quotes_ws.max_column + 1)]
for r, h in enumerate(headers, start=5):
    for c, col in enumerate(quote_cols, start=1):
        val = h.get(col)
        if isinstance(val, Decimal):
            val = float(val)
        quotes_ws.cell(r, c).value = val

# Clear/fill Lines.
lines_ws = wb['Lines']
for row in lines_ws.iter_rows(min_row=5, max_row=max(lines_ws.max_row, 5), max_col=lines_ws.max_column):
    for cell in row:
        cell.value = None
line_cols = [lines_ws.cell(4, c).value for c in range(1, lines_ws.max_column + 1)]
for r, line in enumerate(lines, start=5):
    for c, col in enumerate(line_cols, start=1):
        val = line.get(col)
        if isinstance(val, Decimal):
            val = float(val)
        lines_ws.cell(r, c).value = val

# Summary by status and currency.
summary_ws = wb['Summary']
for row in summary_ws.iter_rows(min_row=5, max_row=max(summary_ws.max_row, 5), max_col=summary_ws.max_column):
    for cell in row:
        cell.value = None
summary_ws['B2'] = f'=SUM(Quotes!K5:K{max(5, 4 + len(headers))})'
summary_cols = [summary_ws.cell(4, c).value for c in range(1, summary_ws.max_column + 1)]
groups = defaultdict(lambda: {'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0})
for h in headers:
    key = (h['Status'], h['Currency'])
    groups[key]['Amount'] += h['NetAmount']
    groups[key]['AmountTWD'] += h['NetAmountTWD']
    groups[key]['Count'] += 1
for r, ((status, currency), agg) in enumerate(sorted(groups.items()), start=5):
    row_data = {
        'Group': 'Status',
        'SubGroup': status,
        'Currency': currency,
        'Amount': money(agg['Amount']),
        'AmountTWD': money(agg['AmountTWD']),
        'Count': agg['Count'],
    }
    for c, col in enumerate(summary_cols, start=1):
        val = row_data.get(col)
        if isinstance(val, Decimal):
            val = float(val)
        summary_ws.cell(r, c).value = val

wb.save(out_file)

# Reopen saved output and verify stated requirements from the saved file.
check_wb = load_workbook(out_file, data_only=False)
for name in ['Report', 'Quotes', 'Lines', 'Summary']:
    if name not in check_wb.sheetnames:
        raise AssertionError(f'Missing sheet {name}')
qws = check_wb['Quotes']
lws = check_wb['Lines']
sws = check_wb['Summary']
rws = check_wb['Report']
if [qws.cell(4, c).value for c in range(1, 12)] != ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']:
    raise AssertionError('Quotes row 4 headers changed')
if [lws.cell(4, c).value for c in range(1, 7)] != ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']:
    raise AssertionError('Lines row 4 headers changed')
if [sws.cell(4, c).value for c in range(1, 7)] != ['Group','SubGroup','Currency','Amount','AmountTWD','Count']:
    raise AssertionError('Summary row 4 headers changed')
if rws['B3'].value != from_date or rws['D3'].value != to_date or customer_id not in str(rws['B4'].value):
    raise AssertionError('Report criteria not filled')
if rws['B7'].value != 'quote_cli mock; fx_rates.csv' or not rws['D7'].value:
    raise AssertionError('Report source/fetched time not filled')
# Count data rows.
def data_count(ws, col=1):
    n = 0
    row = 5
    while ws.cell(row, col).value not in (None, ''):
        n += 1
        row += 1
    return n
if data_count(qws) != len(headers):
    raise AssertionError(f'Quotes row count mismatch: {data_count(qws)} != {len(headers)}')
if data_count(lws) != len(lines):
    raise AssertionError(f'Lines row count mismatch: {data_count(lws)} != {len(lines)}')
if data_count(sws) != len(groups):
    raise AssertionError(f'Summary row count mismatch: {data_count(sws)} != {len(groups)}')
# Type/value checks.
for row in range(5, 5 + len(headers)):
    if not isinstance(qws.cell(row, 1).value, str) or not isinstance(qws.cell(row, 2).value, str):
        raise AssertionError(f'Quote text ID check failed at row {row}')
    for col in [8, 9, 10, 11]:
        if not isinstance(qws.cell(row, col).value, (int, float)):
            raise AssertionError(f'Quote numeric check failed at row {row} col {col}')
for row in range(5, 5 + len(lines)):
    if not isinstance(lws.cell(row, 1).value, str) or not isinstance(lws.cell(row, 2).value, str):
        raise AssertionError(f'Line text ID check failed at row {row}')
    for col in [4, 5, 6]:
        if not isinstance(lws.cell(row, col).value, (int, float)):
            raise AssertionError(f'Line numeric check failed at row {row} col {col}')
expected_formula = f'=SUM(Quotes!K5:K{max(5, 4 + len(headers))})'
if sws['B2'].value != expected_formula:
    raise AssertionError(f'Summary B2 formula mismatch: {sws["B2"].value} != {expected_formula}')
# Verify Summary numbers against saved Quotes.
saved_groups = defaultdict(lambda: {'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0})
for row in range(5, 5 + len(headers)):
    key = (qws.cell(row, 6).value, qws.cell(row, 7).value)
    saved_groups[key]['Amount'] += money(qws.cell(row, 10).value)
    saved_groups[key]['AmountTWD'] += money(qws.cell(row, 11).value)
    saved_groups[key]['Count'] += 1
for row in range(5, 5 + len(groups)):
    key = (sws.cell(row, 2).value, sws.cell(row, 3).value)
    if key not in saved_groups:
        raise AssertionError(f'Unexpected summary key {key}')
    if money(sws.cell(row, 4).value) != money(saved_groups[key]['Amount']):
        raise AssertionError(f'Summary amount mismatch for {key}')
    if money(sws.cell(row, 5).value) != money(saved_groups[key]['AmountTWD']):
        raise AssertionError(f'Summary TWD mismatch for {key}')
    if int(sws.cell(row, 6).value) != saved_groups[key]['Count']:
        raise AssertionError(f'Summary count mismatch for {key}')

print('Wrote', out_file)
print('Checked: workbook re-opened; sheet names; row-4 column names; Report date/condition/source; Quotes rows/text IDs/numeric amounts; Lines rows/text IDs/numeric values; Summary status+currency aggregation; Summary!B2 formula range.')
print(json.dumps({'output_file': str(out_file), 'quotes': len(headers), 'lines': len(lines), 'summary_groups': len(groups)}, ensure_ascii=False))
