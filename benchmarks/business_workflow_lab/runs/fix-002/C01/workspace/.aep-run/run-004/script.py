from pathlib import Path
from openpyxl import load_workbook
import json, os
ws=Path.cwd()
print('cwd', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
manual=ws/'input'/'quote_cli_manual.md'
print('\n--- quote_cli_manual.md ---')
print(manual.read_text(encoding='utf-8'))
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
print('\n--- workbook sheets ---', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in range(1, min(sh.max_row,8)+1):
        print(r, [sh.cell(r,c).value for c in range(1, min(sh.max_column,12)+1)])
    print()
print('\n--- fx_rates.csv head ---')
print((ws/'input'/'fx_rates.csv').read_text(encoding='utf-8').splitlines()[:10])