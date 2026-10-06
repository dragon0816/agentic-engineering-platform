import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
for root, dirs, files in os.walk(ws/'input'):
    level = Path(root).relative_to(ws).parts
    print('DIR', root)
    for f in files:
        p=Path(root)/f
        print(' FILE', f, 'size', p.stat().st_size)
print('\nquote_cli_manual.md head:')
manual=ws/'input'/'quote_cli_manual.md'
print(manual.read_text(encoding='utf-8', errors='replace')[:4000])
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r): print(row)
print('\nWorkbook sheets/headers/sample dimensions:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for sh in wb.sheetnames:
    s=wb[sh]
    print(sh, s.max_row, s.max_column)
    rows=list(s.iter_rows(min_row=1, max_row=min(s.max_row,5), values_only=True))
    for row in rows: print(' ', row)
