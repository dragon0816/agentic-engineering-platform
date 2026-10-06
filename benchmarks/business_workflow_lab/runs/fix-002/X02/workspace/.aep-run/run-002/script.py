from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in range(1, sh.max_row+1):
    vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
    print(r, vals)
