from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)

# workbook structure
wb_path=ws/'input'/'project_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print(f'[{s}] max_row={sh.max_row} max_col={sh.max_column}')
    for r in range(1, min(sh.max_row,8)+1):
        vals=[sh.cell(r,c).value for c in range(1,min(sh.max_column,8)+1)]
        print(' row', r, vals)

# csv sample
csv_path=ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(10),rdr):
        print(i,row)

for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=ws/'input'/name
    print('\n---', name, '---')
    txt=p.read_text(encoding='utf-8')
    print(txt[:3000])