from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in (ws/'input').iterdir():
    print('-', p.name, p.stat().st_size)

# Inspect template workbook
xlsx = ws/'input'/'expense_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
        print(' ', r)

# Inspect fx csv
csvp = ws/'input'/'fx_rates.csv'
print('fx_rates sample:')
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in enumerate(reader):
        print(row)
        if i>=10: break

# Inspect manual
mdp = ws/'input'/'expense_manual.md'
print('manual head:')
print(mdp.read_text(encoding='utf-8')[:4000])