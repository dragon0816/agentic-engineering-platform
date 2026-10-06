from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('output files currently:')
for p in (ws/'output').iterdir():
    print('-', p.name, p.stat().st_size)
    if p.suffix.lower()=='.csv':
        print('csv content repr:', repr(p.read_text(encoding='utf-8-sig')))

wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
print('Instructions all rows:')
for r in sh.iter_rows(values_only=True):
    print(r)
print('Report merged ranges:', list(wb['Report'].merged_cells.ranges))
for s in ['Report','Details','Summary','Exceptions']:
    sh=wb[s]
    print('Merged', s, list(sh.merged_cells.ranges))
    print('column widths', {k:v.width for k,v in sh.column_dimensions.items() if v.width})
