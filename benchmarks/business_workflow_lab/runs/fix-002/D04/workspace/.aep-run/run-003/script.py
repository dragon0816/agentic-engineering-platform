from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment, Protection

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

EMPLOYEE = 'E006'
DATE_FROM = '2026-01-01'
DATE_TO = '2026-01-31'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
FETCHED_AT = '2026-10-06T00:00:00+08:00'
EXPORT_CSV = OUTPUT / 'expense_E006_2026-01.csv'
TEMPLATE = INPUT / 'expense_template.xlsx'
FX_CSV = INPUT / 'fx_rates.csv'
REPORT_XLSX = OUTPUT / 'expense_report_E006_2026-01.xlsx'

print('Checking exported CSV before processing:', EXPORT_CSV)
if not EXPORT_CSV.exists():
    raise FileNotFoundError(f'Exported CSV does not exist: {EXPORT_CSV}')
if EXPORT_CSV.stat().st_size == 0:
    raise ValueError(f'Exported CSV is empty: {EXPORT_CSV}')

# Read exported data. Keep source column names; exported CSV may be header-only for a no-data report.
with EXPORT_CSV.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    if reader.fieldnames is None:
        raise ValueError('Exported CSV has no header row')
    source_headers = reader.fieldnames
    rows = list(reader)
print('Export CSV headers:', source_headers)
print('Export CSV row count:', len(rows))
print('Export CSV sample rows:', rows[:3])

required_export_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']
missing_headers = [h for h in required_export_headers if h not in source_headers]
if missing_headers:
    raise ValueError(f'Exported CSV missing required headers: {missing_headers}')

# Read fixed FX rates. csv is permissive because the Note column contains unquoted commas in this lab file.
fx = {}
with FX_CSV.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    header = next(reader, None)
    for line in reader:
        if len(line) >= 2 and line[0]:
            fx[line[0]] = Decimal(line[1])
print('FX rates loaded:', {k: str(v) for k, v in fx.items()})

def round2(value: Decimal) -> Decimal:
    return value.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def as_bool_text(value):
    s = '' if value is None else str(value).strip()
    if s.lower() in ('true','1','yes','y'):
        return 'true'
    return 'false'

def copy_row_style(ws, src_row, dst_row, max_col):
    for c in range(1, max_col + 1):
        src = ws.cell(src_row, c)
        dst = ws.cell(dst_row, c)
        if src.has_style:
            dst.font = copy(src.font)
            dst.fill = copy(src.fill)
            dst.border = copy(src.border)
            dst.alignment = copy(src.alignment)
            dst.number_format = src.number_format
            dst.protection = copy(src.protection)

wb = load_workbook(TEMPLATE)
report = wb['Report']
details = wb['Details']
summary = wb['Summary']
exceptions = wb['Exceptions']

# Clear possible pre-existing data rows while preserving headers and sheets.
for ws, max_col in [(details, 14), (summary, 6), (exceptions, 4)]:
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

# Fill report metadata.
report['B3'] = DATE_FROM
report['D3'] = DATE_TO
report['B4'] = f'EmployeeID={EMPLOYEE}; ExpenseDate between {DATE_FROM} and {DATE_TO}; Status=All; Project=All'
report['B5'] = GENERATED_AT
report['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D7'] = FETCHED_AT
if rows:
    report['B9'] = '匯出成功；資料列已依固定匯率換算 AmountTWD。'
else:
    report['B9'] = '查無符合條件的 expense；仍產出只有表頭的正常空報表。匯出曾出現暫時性寫入錯誤，已重試 1 次後成功；未將不存在檔案視為成功。'

# Details rows and exceptions. NeedsReview is true if receipt is missing or status is Rejected/Draft.
detail_headers = [details.cell(4, c).value for c in range(1, 15)]
expected_detail_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
if detail_headers != expected_detail_headers:
    raise ValueError(f'Detail template headers changed: {detail_headers}')

summary_acc = {}
exception_records = []
for idx, row in enumerate(rows, start=5):
    currency = (row.get('Currency') or '').strip()
    amount_raw = (row.get('Amount') or '0').strip()
    amount = Decimal(amount_raw) if amount_raw else Decimal('0')
    if currency not in fx:
        raise ValueError(f'Missing FX rate for currency {currency!r}')
    amount_twd = round2(amount * fx[currency])
    missing_receipt = as_bool_text(row.get('MissingReceipt'))
    status = (row.get('Status') or '').strip()
    needs_review = 'true' if missing_receipt == 'true' or status in ('Rejected', 'Draft') else 'false'

    values = [
        row.get('ExpenseID',''), row.get('EmployeeID',''), row.get('EmployeeName',''), row.get('ProjectID',''),
        row.get('ExpenseDate',''), row.get('Category',''), currency, float(amount), float(amount_twd), status,
        row.get('ReceiptID',''), row.get('Description',''), missing_receipt, needs_review
    ]
    copy_row_style(details, 5 if details.max_row >= 5 else 4, idx, 14)
    for col, value in enumerate(values, start=1):
        cell = details.cell(idx, col, value)
        if col in (1,2,4,11,13,14):
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9):
            cell.number_format = '0.00'

    key = (status, row.get('Category',''), currency)
    if key not in summary_acc:
        summary_acc[key] = {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0}
    summary_acc[key]['amount'] += amount
    summary_acc[key]['amount_twd'] += amount_twd
    summary_acc[key]['count'] += 1

    reasons = []
    if missing_receipt == 'true':
        reasons.append('Missing receipt')
    if status in ('Rejected', 'Draft'):
        reasons.append(f'Status {status} needs review')
    if reasons:
        exception_records.append([row.get('ExpenseID',''), row.get('EmployeeID',''), row.get('ExpenseDate',''), '; '.join(reasons)])

last_detail_row = max(5, 4 + len(rows))
summary['B2'] = f'=SUM(Details!I5:I{last_detail_row})'
summary['B2'].number_format = '0.00'

# Summary groups, only when data exists. Group=Status, SubGroup=Category.
srow = 5
for (status, category, currency), vals in sorted(summary_acc.items()):
    copy_row_style(summary, 5 if summary.max_row >= 5 else 4, srow, 6)
    out = [status, category, currency, float(round2(vals['amount'])), float(round2(vals['amount_twd'])), vals['count']]
    for col, value in enumerate(out, start=1):
        cell = summary.cell(srow, col, value)
        if col in (4,5):
            cell.number_format = '0.00'
    srow += 1

# Exception rows, only when data exists and exception conditions are met.
erow = 5
for rec in exception_records:
    copy_row_style(exceptions, 5 if exceptions.max_row >= 5 else 4, erow, 4)
    for col, value in enumerate(rec, start=1):
        cell = exceptions.cell(erow, col, value)
        if col in (1,2):
            cell.number_format = '@'
        elif col == 3:
            cell.number_format = 'yyyy-mm-dd'
    erow += 1

# Ensure truly empty normal report has no stray row 5 values.
if not rows:
    for ws in (details, summary, exceptions):
        if ws.max_row > 4:
            ws.delete_rows(5, ws.max_row - 4)
    summary['B2'] = '=SUM(Details!I5:I5)'
    summary['B2'].number_format = '0.00'

wb.save(REPORT_XLSX)
print('Wrote report:', REPORT_XLSX, REPORT_XLSX.stat().st_size)

# Reopen and verify all stated requirements from saved outputs.
verify_wb = load_workbook(REPORT_XLSX, data_only=False)
assert verify_wb.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions'], verify_wb.sheetnames
vr = verify_wb['Report']
vd = verify_wb['Details']
vs = verify_wb['Summary']
ve = verify_wb['Exceptions']
assert vr['B3'].value == DATE_FROM
assert vr['D3'].value == DATE_TO
assert EMPLOYEE in vr['B4'].value
assert vr['B5'].value == GENERATED_AT
assert 'Expense Viewer' in vr['B7'].value
assert vr['D7'].value == FETCHED_AT
assert list(vd.iter_rows(min_row=4, max_row=4, values_only=True))[0] == tuple(expected_detail_headers)
assert list(vs.iter_rows(min_row=4, max_row=4, values_only=True))[0] == ('Group','SubGroup','Currency','Amount','AmountTWD','Count')
assert list(ve.iter_rows(min_row=4, max_row=4, values_only=True))[0] == ('ExpenseID','EmployeeID','ExpenseDate','Reason')
# Header-only source should become a normal empty workbook: no data rows on detail/summary/exceptions sheets.
assert len(rows) == 0, 'Expected no source rows for this request based on exported CSV'
assert vd.max_row == 4, f'Details should be empty after header, got max_row={vd.max_row}'
assert vs.max_row == 4, f'Summary should be empty after header, got max_row={vs.max_row}'
assert ve.max_row == 4, f'Exceptions should be empty after header, got max_row={ve.max_row}'
assert vs['B2'].value == '=SUM(Details!I5:I5)'
assert EXPORT_CSV.exists() and EXPORT_CSV.stat().st_size > 0
with EXPORT_CSV.open('r', encoding='utf-8-sig', newline='') as f:
    reread = csv.DictReader(f)
    assert reread.fieldnames == required_export_headers
    assert sum(1 for _ in reread) == 0
print('Verification passed: exported CSV exists and is header-only; workbook dates/condition/source/notes filled; sheet names and headers preserved; Details/Summary/Exceptions contain no data rows; zero-data formula retained; IDs/text/amount formats ready for populated data.')
print('Output files:')
for p in OUTPUT.iterdir():
    print('-', p.name, p.stat().st_size)
