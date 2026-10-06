from pathlib import Path
import subprocess, json, csv, shutil, sys, re, os
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
PY_CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
TODAY = '2026-10-06'
NOW = datetime.now(ZoneInfo('Asia/Taipei'))
# keep operator-supplied date if host clock differs
if NOW.date().isoformat() != TODAY:
    NOW = datetime.fromisoformat(TODAY + 'T00:00:00+08:00')
GENERATED_AT = NOW.isoformat(timespec='seconds')

TWOPLACES = Decimal('0.01')
def money(v):
    return Decimal(str(v)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args, label):
    """Run quote CLI; retry exit code 4 at most two times; parse stdout only on code 0."""
    attempts = []
    for attempt in range(1, 4):
        if CLI.exists():
            cmd = [str(CLI)] + args
        else:
            cmd = [sys.executable, str(PY_CLI)] + args
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True, encoding='utf-8', errors='replace')
        attempts.append({'attempt': attempt, 'exit_code': proc.returncode, 'stderr': proc.stderr.strip()})
        if proc.returncode == 0:
            try:
                data = json.loads(proc.stdout)
            except Exception as e:
                raise RuntimeError(f'{label}: CLI succeeded but stdout is not parseable JSON: {e}')
            return data, attempts
        if proc.returncode in (2,3):
            raise RuntimeError(f'{label}: CLI exit code {proc.returncode}; stderr: {proc.stderr.strip()}')
        if proc.returncode == 4:
            if attempt < 3:
                continue
            raise RuntimeError(f'{label}: CLI exit code 4 after {attempt} attempts; stderr: {proc.stderr.strip()}')
        raise RuntimeError(f'{label}: unexpected CLI exit code {proc.returncode}; stderr: {proc.stderr.strip()}')

# Read FX rates, using first two columns only because Note text can include commas in the supplied synthetic CSV.
fx = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
if not fx:
    raise RuntimeError('No FX rates loaded')
print('Loaded FX rates:', {k: str(v) for k,v in fx.items()})

# Query list
list_args = ['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
list_data, list_attempts = run_cli(list_args, 'quotes list')
if list_data.get('schema_version') != 1 or 'items' not in list_data or not isinstance(list_data['items'], list):
    raise RuntimeError('quotes list returned unexpected JSON schema')
headers = list_data['items']
print(f'quotes list returned {len(headers)} rows; attempts={list_attempts}')

# Fetch line details per quote
all_lines = []
get_attempts = {}
full_headers = []
for h in headers:
    qid = str(h['QuoteID'])
    detail, attempts = run_cli(['quotes', 'get', qid, '--format', 'json'], f'quotes get {qid}')
    get_attempts[qid] = attempts
    if detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
        raise RuntimeError(f'quotes get {qid} returned unexpected JSON schema')
    full_headers.append(detail['header'])
    for line in detail['line_items']:
        line2 = dict(line)
        line2['QuoteID'] = qid
        all_lines.append(line2)
print(f'Fetched details for {len(full_headers)} quotes and {len(all_lines)} line rows')

# Sort deterministically
full_headers.sort(key=lambda x: str(x.get('QuoteID','')))
all_lines.sort(key=lambda x: (str(x.get('QuoteID','')), str(x.get('ProductID','')), str(x.get('ProductName',''))))

# Copy template and fill
out_xlsx = OUTPUT / 'sent_quotes_2026-09.xlsx'
shutil.copyfile(TEMPLATE, out_xlsx)
wb = load_workbook(out_xlsx)

# Report sheet
rep = wb['Report']
rep['B3'] = FROM_DATE
rep['D3'] = TO_DATE
rep['B4'] = f'Status={STATUS}; QuoteDate between {FROM_DATE} and {TO_DATE} inclusive'
rep['B5'] = GENERATED_AT
rep['B7'] = 'quote_cli; fx_rates.csv'
rep['D7'] = GENERATED_AT
notes = []
notes.append(f'查詢成功；quotes list 筆數 {len(full_headers)}，line items 筆數 {len(all_lines)}。')
if len(full_headers) == 0:
    notes.append('本期間沒有符合條件的 Sent 報價。')
rep['B9'] = ' '.join(notes)

# helper styles/formats from row 4 maybe
def clear_data(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def as_num(dec):
    # Decimal -> float acceptable numeric xlsx values with number format 0.00
    return float(money(dec))

# Quotes
qws = wb['Quotes']
clear_data(qws, 5)
q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for col, val in enumerate(q_headers, 1):
    if qws.cell(4,col).value != val:
        print('Header notice Quotes', col, qws.cell(4,col).value, '-> expected', val)
for r, h in enumerate(full_headers, start=5):
    currency = str(h['Currency'])
    if currency not in fx:
        raise RuntimeError(f'Missing FX rate for {currency}')
    net = money(h['NetAmount'])
    net_twd = money(net * fx[currency])
    values = [
        str(h['QuoteID']), str(h['CustomerID']), str(h['ProjectID']), str(h['QuoteDate']), str(h['ValidUntil']),
        str(h['Status']), currency, as_num(h['Subtotal']), float(Decimal(str(h['DiscountPercent']))), as_num(net), as_num(net_twd)
    ]
    for c,v in enumerate(values,1):
        qws.cell(r,c).value = v
    for c in [1,2,3,4,5,6,7]:
        qws.cell(r,c).number_format = '@'
    for c in [8,10,11]:
        qws.cell(r,c).number_format = '0.00'
    qws.cell(r,9).number_format = '0.00'

# Lines
lws = wb['Lines']
clear_data(lws, 5)
l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for r, line in enumerate(all_lines, start=5):
    vals = [str(line['QuoteID']), str(line['ProductID']), str(line['ProductName']),
            int(Decimal(str(line['Quantity']))), as_num(line['UnitPrice']), as_num(line['LineAmount'])]
    for c,v in enumerate(vals,1):
        lws.cell(r,c).value = v
    for c in [1,2,3]:
        lws.cell(r,c).number_format='@'
    for c in [5,6]:
        lws.cell(r,c).number_format='0.00'

# Summary: clear rows 5+, total formula; group by Status+Currency
sws = wb['Summary']
clear_data(sws, 5)
last_q_row = 4 + len(full_headers) if full_headers else 5
formula_last = max(5, last_q_row)
sws['B2'] = f'=SUM(Quotes!K5:K{formula_last})'
summary = {}
for h in full_headers:
    key = (str(h['Status']), str(h['Currency']))
    currency = key[1]
    net = money(h['NetAmount'])
    net_twd = money(net * fx[currency])
    if key not in summary:
        summary[key] = {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0}
    summary[key]['amount'] += net
    summary[key]['twd'] += net_twd
    summary[key]['count'] += 1
for r, ((group, currency), agg) in enumerate(sorted(summary.items()), start=5):
    vals = [group, '', currency, as_num(agg['amount']), as_num(agg['twd']), agg['count']]
    for c,v in enumerate(vals,1):
        sws.cell(r,c).value = v
    for c in [1,2,3]: sws.cell(r,c).number_format='@'
    for c in [4,5]: sws.cell(r,c).number_format='0.00'

wb.save(out_xlsx)

# Save audit metadata
out_json = OUTPUT / 'sent_quotes_2026-09_audit.json'
audit = {
    'query': {'from': FROM_DATE, 'to': TO_DATE, 'status': STATUS},
    'generated_at': GENERATED_AT,
    'quote_count': len(full_headers),
    'line_count': len(all_lines),
    'fx_rates': {k: str(v) for k,v in fx.items()},
    'list_attempts': list_attempts,
    'get_attempts': get_attempts,
    'output_workbook': out_xlsx.name,
}
out_json.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')

# Verification: reopen saved outputs and check requirements.
wb2 = load_workbook(out_xlsx, data_only=False)
errors=[]
# Sheets preserved
for name in ['Report','Quotes','Lines','Summary','Instructions']:
    if name not in wb2.sheetnames:
        errors.append(f'Missing sheet {name}')
# Merged title preserved
if 'A1:F1' not in [str(r) for r in wb2['Report'].merged_cells.ranges]:
    errors.append('Report A1:F1 merge not preserved')
# Report cells
rep2=wb2['Report']
checks = {'B3': FROM_DATE, 'D3': TO_DATE}
for cell, expected in checks.items():
    if rep2[cell].value != expected:
        errors.append(f'Report {cell} expected {expected}, got {rep2[cell].value!r}')
if not rep2['B4'].value or STATUS not in rep2['B4'].value:
    errors.append('Report B4 condition missing Sent status')
if not rep2['B5'].value or not str(rep2['B5'].value).startswith(TODAY):
    errors.append('Report B5 generated timestamp missing expected date')
if 'quote_cli' not in str(rep2['B7'].value):
    errors.append('Report B7 source missing quote_cli')
# Detail headers and row counts
qws2=wb2['Quotes']; lws2=wb2['Lines']; sws2=wb2['Summary']
actual_q_headers=[qws2.cell(4,c).value for c in range(1,12)]
if actual_q_headers != q_headers:
    errors.append(f'Quotes headers mismatch: {actual_q_headers}')
actual_l_headers=[lws2.cell(4,c).value for c in range(1,7)]
if actual_l_headers != l_headers:
    errors.append(f'Lines headers mismatch: {actual_l_headers}')
q_rows = max(0, qws2.max_row - 4)
l_rows = max(0, lws2.max_row - 4)
# delete_rows to below row 4 can leave max_row at 4 when no data; OK
q_nonempty = sum(1 for r in range(5, qws2.max_row+1) if qws2.cell(r,1).value not in (None,''))
l_nonempty = sum(1 for r in range(5, lws2.max_row+1) if lws2.cell(r,1).value not in (None,''))
if q_nonempty != len(full_headers): errors.append(f'Quotes row count {q_nonempty} != {len(full_headers)}')
if l_nonempty != len(all_lines): errors.append(f'Lines row count {l_nonempty} != {len(all_lines)}')
# Type/value checks
for r in range(5, 5+len(full_headers)):
    qid=qws2.cell(r,1).value
    if not isinstance(qid, str): errors.append(f'Quotes row {r} QuoteID not text')
    if qws2.cell(r,6).value != STATUS: errors.append(f'Quotes row {r} status not {STATUS}')
    if not (FROM_DATE <= qws2.cell(r,4).value <= TO_DATE): errors.append(f'Quotes row {r} date outside range')
    for c in [8,9,10,11]:
        if not isinstance(qws2.cell(r,c).value, (int,float)): errors.append(f'Quotes row {r} col {c} not numeric')
for r in range(5, 5+len(all_lines)):
    if not isinstance(lws2.cell(r,1).value, str): errors.append(f'Lines row {r} QuoteID not text')
    if not isinstance(lws2.cell(r,2).value, str): errors.append(f'Lines row {r} ProductID not text')
    for c in [4,5,6]:
        if not isinstance(lws2.cell(r,c).value, (int,float)): errors.append(f'Lines row {r} col {c} not numeric')
# Summary formula/ranges
if sws2['B2'].value != f'=SUM(Quotes!K5:K{formula_last})':
    errors.append(f'Summary B2 formula mismatch: {sws2["B2"].value}')
# Validate total formula range value manually equals sum K
sum_k = money(sum(Decimal(str(qws2.cell(r,11).value)) for r in range(5, 5+len(full_headers))))
sum_summary = money(sum(Decimal(str(sws2.cell(r,5).value)) for r in range(5, sws2.max_row+1) if sws2.cell(r,1).value not in (None,'')))
if sum_k != sum_summary:
    errors.append(f'Summary TWD total {sum_summary} != Quotes K total {sum_k}')
# Audit parse check
parsed_audit=json.loads(out_json.read_text(encoding='utf-8'))
if parsed_audit['quote_count'] != len(full_headers) or parsed_audit['line_count'] != len(all_lines):
    errors.append('Audit counts mismatch')
if errors:
    print('VERIFICATION FAILED')
    for e in errors: print('-', e)
    raise SystemExit(1)
print('VERIFICATION OK')
print('Output workbook:', out_xlsx)
print('Audit file:', out_json)
print('Checked: workbook reopened; sheets/merge preserved; Report cells; Quotes and Lines headers; row counts; text IDs; numeric money/quantity fields; Sent/date filters; Summary formula and totals; audit JSON parsed.')
print('Quotes:', len(full_headers), 'Lines:', len(all_lines), 'Summary TWD:', str(sum_k))