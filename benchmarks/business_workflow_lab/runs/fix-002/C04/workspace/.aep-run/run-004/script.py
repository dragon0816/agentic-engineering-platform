from pathlib import Path
import subprocess, json, csv, shutil, sys
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
PY_CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
TODAY = '2026-10-06'
TW = timezone(timedelta(hours=8))
NOW = datetime.now(TW)
# Use operator-provided current date if runtime clock is not the stated lab date.
if NOW.date().isoformat() != TODAY:
    NOW = datetime.fromisoformat(TODAY + 'T00:00:00+08:00')
GENERATED_AT = NOW.isoformat(timespec='seconds')

TWOPLACES = Decimal('0.01')
def money(v):
    return Decimal(str(v)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args, label):
    """Run quote CLI; retry exit code 4 at most two times; parse stdout only on code 0."""
    attempts = []
    for attempt in range(1, 4):
        if CLI.exists():
            cmd = [str(CLI)] + args
        else:
            cmd = [sys.executable, str(PY_CLI)] + args
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True, encoding='utf-8', errors='replace')
        attempts.append({'attempt': attempt, 'exit_code': proc.returncode, 'stderr': proc.stderr.strip()})
        if proc.returncode == 0:
            try:
                data = json.loads(proc.stdout)
            except Exception as e:
                raise RuntimeError(f'{label}: CLI succeeded but stdout is not parseable JSON: {e}')
            return data, attempts
        if proc.returncode in (2, 3):
            raise RuntimeError(f'{label}: CLI exit code {proc.returncode}; stderr: {proc.stderr.strip()}')
        if proc.returncode == 4:
            if attempt < 3:
                continue
            raise RuntimeError(f'{label}: CLI exit code 4 after {attempt} attempts; stderr: {proc.stderr.strip()}')
        raise RuntimeError(f'{label}: unexpected CLI exit code {proc.returncode}; stderr: {proc.stderr.strip()}')

# Read FX rates. Use first two CSV fields; the synthetic Note field contains unquoted commas.
fx = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader, None)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
print('Loaded FX rates:', {k: str(v) for k, v in fx.items()})

list_args = ['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
list_data, list_attempts = run_cli(list_args, 'quotes list')
if list_data.get('schema_version') != 1 or 'items' not in list_data or not isinstance(list_data['items'], list):
    raise RuntimeError('quotes list returned unexpected JSON schema')
headers = list_data['items']
print(f'quotes list returned {len(headers)} rows; attempts={list_attempts}')

all_lines = []
get_attempts = {}
full_headers = []
for h in headers:
    qid = str(h['QuoteID'])
    detail, attempts = run_cli(['quotes', 'get', qid, '--format', 'json'], f'quotes get {qid}')
    get_attempts[qid] = attempts
    if detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
        raise RuntimeError(f'quotes get {qid} returned unexpected JSON schema')
    full_headers.append(detail['header'])
    for line in detail['line_items']:
        line2 = dict(line)
        line2['QuoteID'] = qid
        all_lines.append(line2)
print(f'Fetched details for {len(full_headers)} quotes and {len(all_lines)} line rows')

full_headers.sort(key=lambda x: str(x.get('QuoteID', '')))
all_lines.sort(key=lambda x: (str(x.get('QuoteID', '')), str(x.get('ProductID', '')), str(x.get('ProductName', ''))))

out_xlsx = OUTPUT / 'sent_quotes_2026-09.xlsx'
shutil.copyfile(TEMPLATE, out_xlsx)
wb = load_workbook(out_xlsx)

rep = wb['Report']
rep['B3'] = FROM_DATE
rep['D3'] = TO_DATE
rep['B4'] = f'Status={STATUS}; QuoteDate between {FROM_DATE} and {TO_DATE} inclusive'
rep['B5'] = GENERATED_AT
rep['B7'] = 'quote_cli; fx_rates.csv'
rep['D7'] = GENERATED_AT
rep['B9'] = f'查詢成功；quotes list 筆數 {len(full_headers)}，line items 筆數 {len(all_lines)}。' + ('' if full_headers else ' 本期間沒有符合條件的 Sent 報價。')

def clear_data(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def as_num(dec):
    return float(money(dec))

qws = wb['Quotes']
clear_data(qws, 5)
q_headers = ['QuoteID', 'CustomerID', 'ProjectID', 'QuoteDate', 'ValidUntil', 'Status', 'Currency', 'Subtotal', 'DiscountPercent', 'NetAmount', 'NetAmountTWD']
for c, val in enumerate(q_headers, 1):
    qws.cell(4, c).value = val
for r, h in enumerate(full_headers, start=5):
    currency = str(h['Currency'])
    if currency not in fx:
        raise RuntimeError(f'Missing FX rate for {currency}')
    net = money(h['NetAmount'])
    net_twd = money(net * fx[currency])
    values = [str(h['QuoteID']), str(h['CustomerID']), str(h['ProjectID']), str(h['QuoteDate']), str(h['ValidUntil']), str(h['Status']), currency, as_num(h['Subtotal']), float(Decimal(str(h['DiscountPercent']))), as_num(net), as_num(net_twd)]
    for c, v in enumerate(values, 1):
        qws.cell(r, c).value = v
    for c in [1, 2, 3, 4, 5, 6, 7]:
        qws.cell(r, c).number_format = '@'
    for c in [8, 10, 11]:
        qws.cell(r, c).number_format = '0.00'
    qws.cell(r, 9).number_format = '0.00'

lws = wb['Lines']
clear_data(lws, 5)
l_headers = ['QuoteID', 'ProductID', 'ProductName', 'Quantity', 'UnitPrice', 'LineAmount']
for c, val in enumerate(l_headers, 1):
    lws.cell(4, c).value = val
for r, line in enumerate(all_lines, start=5):
    vals = [str(line['QuoteID']), str(line['ProductID']), str(line['ProductName']), int(Decimal(str(line['Quantity']))), as_num(line['UnitPrice']), as_num(line['LineAmount'])]
    for c, v in enumerate(vals, 1):
        lws.cell(r, c).value = v
    for c in [1, 2, 3]:
        lws.cell(r, c).number_format = '@'
    for c in [5, 6]:
        lws.cell(r, c).number_format = '0.00'

sws = wb['Summary']
clear_data(sws, 5)
formula_last = max(5, 4 + len(full_headers))
sws['B2'] = f'=SUM(Quotes!K5:K{formula_last})'
summary = {}
for h in full_headers:
    key = (str(h['Status']), str(h['Currency']))
    currency = key[1]
    net = money(h['NetAmount'])
    net_twd = money(net * fx[currency])
    summary.setdefault(key, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    summary[key]['amount'] += net
    summary[key]['twd'] += net_twd
    summary[key]['count'] += 1
for r, ((group, currency), agg) in enumerate(sorted(summary.items()), start=5):
    vals = [group, '', currency, as_num(agg['amount']), as_num(agg['twd']), agg['count']]
    for c, v in enumerate(vals, 1):
        sws.cell(r, c).value = v
    for c in [1, 2, 3]:
        sws.cell(r, c).number_format = '@'
    for c in [4, 5]:
        sws.cell(r, c).number_format = '0.00'

wb.save(out_xlsx)

out_json = OUTPUT / 'sent_quotes_2026-09_audit.json'
audit = {
    'query': {'from': FROM_DATE, 'to': TO_DATE, 'status': STATUS},
    'generated_at': GENERATED_AT,
    'quote_count': len(full_headers),
    'line_count': len(all_lines),
    'fx_rates': {k: str(v) for k, v in fx.items()},
    'list_attempts': list_attempts,
    'get_attempts': get_attempts,
    'output_workbook': out_xlsx.name,
}
out_json.write_text(json.dumps(audit, ensure_ascii=False, indent=2), encoding='utf-8')

# Reopen and verify saved outputs.
wb2 = load_workbook(out_xlsx, data_only=False)
errors = []
for name in ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']:
    if name not in wb2.sheetnames:
        errors.append(f'Missing sheet {name}')
if 'A1:F1' not in [str(r) for r in wb2['Report'].merged_cells.ranges]:
    errors.append('Report A1:F1 merge not preserved')
rep2 = wb2['Report']
if rep2['B3'].value != FROM_DATE: errors.append(f'Report B3 mismatch: {rep2["B3"].value!r}')
if rep2['D3'].value != TO_DATE: errors.append(f'Report D3 mismatch: {rep2["D3"].value!r}')
if STATUS not in str(rep2['B4'].value): errors.append('Report B4 condition missing Sent status')
if not str(rep2['B5'].value).startswith(TODAY): errors.append('Report B5 generated timestamp missing expected date')
if 'quote_cli' not in str(rep2['B7'].value): errors.append('Report B7 source missing quote_cli')
qws2 = wb2['Quotes']; lws2 = wb2['Lines']; sws2 = wb2['Summary']
if [qws2.cell(4, c).value for c in range(1, 12)] != q_headers:
    errors.append('Quotes headers mismatch')
if [lws2.cell(4, c).value for c in range(1, 7)] != l_headers:
    errors.append('Lines headers mismatch')
q_nonempty = sum(1 for r in range(5, qws2.max_row + 1) if qws2.cell(r, 1).value not in (None, ''))
l_nonempty = sum(1 for r in range(5, lws2.max_row + 1) if lws2.cell(r, 1).value not in (None, ''))
if q_nonempty != len(full_headers): errors.append(f'Quotes row count {q_nonempty} != {len(full_headers)}')
if l_nonempty != len(all_lines): errors.append(f'Lines row count {l_nonempty} != {len(all_lines)}')
for r in range(5, 5 + len(full_headers)):
    if not isinstance(qws2.cell(r, 1).value, str): errors.append(f'Quotes row {r} QuoteID not text')
    if not isinstance(qws2.cell(r, 2).value, str): errors.append(f'Quotes row {r} CustomerID not text')
    if not isinstance(qws2.cell(r, 3).value, str): errors.append(f'Quotes row {r} ProjectID not text')
    if qws2.cell(r, 6).value != STATUS: errors.append(f'Quotes row {r} status not {STATUS}')
    if not (FROM_DATE <= qws2.cell(r, 4).value <= TO_DATE): errors.append(f'Quotes row {r} date outside range')
    for c in [8, 9, 10, 11]:
        if not isinstance(qws2.cell(r, c).value, (int, float)): errors.append(f'Quotes row {r} col {c} not numeric')
for r in range(5, 5 + len(all_lines)):
    if not isinstance(lws2.cell(r, 1).value, str): errors.append(f'Lines row {r} QuoteID not text')
    if not isinstance(lws2.cell(r, 2).value, str): errors.append(f'Lines row {r} ProductID not text')
    if not isinstance(lws2.cell(r, 3).value, str): errors.append(f'Lines row {r} ProductName not text')
    for c in [4, 5, 6]:
        if not isinstance(lws2.cell(r, c).value, (int, float)): errors.append(f'Lines row {r} col {c} not numeric')
expected_formula = f'=SUM(Quotes!K5:K{formula_last})'
if sws2['B2'].value != expected_formula:
    errors.append(f'Summary B2 formula mismatch: {sws2["B2"].value}')
sum_k = money(sum(Decimal(str(qws2.cell(r, 11).value)) for r in range(5, 5 + len(full_headers))))
sum_summary = money(sum(Decimal(str(sws2.cell(r, 5).value)) for r in range(5, sws2.max_row + 1) if sws2.cell(r, 1).value not in (None, '')))
if sum_k != sum_summary:
    errors.append(f'Summary TWD total {sum_summary} != Quotes K total {sum_k}')
parsed_audit = json.loads(out_json.read_text(encoding='utf-8'))
if parsed_audit['quote_count'] != len(full_headers) or parsed_audit['line_count'] != len(all_lines):
    errors.append('Audit counts mismatch')
if errors:
    print('VERIFICATION FAILED')
    for e in errors:
        print('-', e)
    raise SystemExit(1)
print('VERIFICATION OK')
print('Output workbook:', out_xlsx)
print('Audit file:', out_json)
print('Checked: workbook reopened; sheets/merge preserved; Report cells; Quotes and Lines headers; row counts; text IDs; numeric money/quantity fields; Sent/date filters; Summary formula and totals; audit JSON parsed.')
print('Quotes:', len(full_headers), 'Lines:', len(all_lines), 'Summary TWD:', str(sum_k))