from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
manual = ws/'input'/'quote_cli_manual.md'
print('\nmanual head:')
print(manual.read_text(encoding='utf-8', errors='replace')[:4000])
fx = ws/'input'/'fx_rates.csv'
print('\nfx sample:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i,row in zip(range(8), r):
        print(row)
tpl = ws/'input'/'quote_template.xlsx'
wb = load_workbook(tpl, data_only=False)
print('\nworkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=False):
        print([c.value for c in row[:min(sh.max_column,10)]])
    print('---')