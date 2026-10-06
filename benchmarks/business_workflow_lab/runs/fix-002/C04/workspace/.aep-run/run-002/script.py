from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for s in ['Instructions','Report','Summary']:
    sh=wb[s]
    print('\nSHEET',s)
    for r in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True):
        print(r)
print('\nmerged ranges')
for s in wb.sheetnames:
    print(s, list(wb[s].merged_cells.ranges))