from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
for p in (ws/'input').iterdir():
    print('\nFILE', p.name, 'size', p.stat().st_size)
    if p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, read_only=True, data_only=False)
        print('sheets:', wb.sheetnames)
        for s in wb.sheetnames:
            sh=wb[s]
            print(' sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
            rows=[]
            for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True): rows.append(row)
            print(' sample', rows)
    elif p.suffix.lower()=='.csv':
        with p.open(newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f); rows=[next(r, None) for _ in range(5)]
        print(' sample', rows)
    elif p.suffix.lower()=='.md':
        txt=p.read_text(encoding='utf-8', errors='replace')
        print(' head:', txt[:2000])