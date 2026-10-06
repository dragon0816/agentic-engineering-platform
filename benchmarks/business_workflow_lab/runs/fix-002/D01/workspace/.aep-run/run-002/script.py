from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('output files:', [p.name for p in (ws/'output').iterdir()])
exp=ws/'output'/'E001_september_export.xlsx'
if exp.exists():
    wb=load_workbook(exp, read_only=True, data_only=True)
    print('export sheets', wb.sheetnames)
    for s in wb.sheetnames:
        sh=wb[s]
        print('sheet', s, sh.max_row, sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(5, sh.max_row), values_only=True): print(row)
print('\nInstructions:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', read_only=True, data_only=True)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True): print(row)
print('\nReport cells:')
sh=wb['Report']
for row in sh.iter_rows(min_row=1,max_row=20,min_col=1,max_col=6,values_only=True): print(row)