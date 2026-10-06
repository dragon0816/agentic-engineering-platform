from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
exp=ws/'output'/'E001_september_export.xlsx'
wb=load_workbook(exp, read_only=True, data_only=True)
sh=wb['Expenses']
print('headers', [c.value for c in sh[1]])
for row in sh.iter_rows(min_row=2, values_only=True): print(row)