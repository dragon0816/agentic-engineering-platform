from pathlib import Path
import csv, json, subprocess, sys
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUTFILE = OUTPUT / 'quote_report_P001_P005_2026-09_Accepted.xlsx'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
PROJECTS = [f'P{i:03d}' for i in range(1, 6)]
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Accepted'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
Q = Decimal('0.01')

def money(x):
    return Decimal(str(x)).quantize(Q, rounding=ROUND_HALF_UP)

def clear_data(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

# Read FX rates; Note text contains commas in this fixture, but Currency and RateToTWD are first two columns.
fx = {}
with open(FX_FILE, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        if row and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates loaded:', fx)

# Fetch quote headers and line items from the documented CLI.
headers = []
lines = []
fetch_log = []
for project in PROJECTS:
    list_cmd = [sys.executable, str(CLI), 'quotes', 'list', '--project', project, '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
    res = subprocess.run(list_cmd, text=True, capture_output=True)
    if res.returncode != 0:
        raise RuntimeError(f'quote_cli list failed for {project}: rc={res.returncode} stderr={res.stderr}')
    data = json.loads(res.stdout)
    fetch_log.append((project, data.get('count', len(data.get('items', [])))))
    for item in data.get('items', []):
        qid = item['QuoteID']
        get_cmd = [sys.executable, str(CLI), 'quotes', 'get', qid, '--format', 'json']
        gres = subprocess.run(get_cmd, text=True, capture_output=True)
        if gres.returncode != 0:
            raise RuntimeError(f'quote_cli get failed for {qid}: rc={gres.returncode} stderr={gres.stderr}')
        detail = json.loads(gres.stdout)
        h = detail['header']
        cur = h['Currency']
        if cur not in fx:
            raise RuntimeError(f'Missing FX rate for currency {cur}')
        h['_NetAmountTWD'] = money(money(h['NetAmount']) * fx[cur])
        headers.append(h)
        for li in detail.get('line_items', []):
            lines.append({'QuoteID': qid, **li})
headers.sort(key=lambda h: (h['ProjectID'], h['QuoteID']))
lines.sort(key=lambda r: (r['QuoteID'], r['ProductID']))
print('Fetched counts by project:', fetch_log)
print('Fetched quote IDs:', [h['QuoteID'] for h in headers])

wb = load_workbook(TEMPLATE)
report = wb['Report']
quotes_ws = wb['Quotes']
lines_ws = wb['Lines']
summary_ws = wb['Summary']
for sheet in (quotes_ws, lines_ws, summary_ws):
    clear_data(sheet, 5)

# Report metadata
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = 'ProjectID P001-P005; Status Accepted; QuoteDate 2026-09-01 through 2026-09-30; FX from input/fx_rates.csv'
report['B5'] = GENERATED_AT
report['B7'] = 'quote_cli_mock; input/fx_rates.csv'
report['D7'] = GENERATED_AT
zero_projects = [p for p, c in fetch_log if c == 0]
report['B9'] = 'Projects with no Accepted quote kept as zero rows in Summary: ' + ', '.join(zero_projects) if zero_projects else 'None'

# Quotes sheet
for r, h in enumerate(headers, start=5):
    values = [h['QuoteID'], h['CustomerID'], h['ProjectID'], h['QuoteDate'], h['ValidUntil'], h['Status'], h['Currency'],
              float(money(h['Subtotal'])), float(Decimal(str(h['DiscountPercent']))), float(money(h['NetAmount'])), float(h['_NetAmountTWD'])]
    for c, v in enumerate(values, start=1):
        cell = quotes_ws.cell(r, c, v)
        if c in (1, 2, 3, 4, 5, 6, 7):
            cell.number_format = '@'
        elif c in (8, 10, 11):
            cell.number_format = '0.00'
        elif c == 9:
            cell.number_format = '0'

# Lines sheet
for r, li in enumerate(lines, start=5):
    values = [li['QuoteID'], li['ProductID'], li['ProductName'], int(li['Quantity']), float(money(li['UnitPrice'])), float(money(li['LineAmount']))]
    for c, v in enumerate(values, start=1):
        cell = lines_ws.cell(r, c, v)
        if c in (1, 2, 3):
            cell.number_format = '@'
        elif c == 4:
            cell.number_format = '0'
        elif c in (5, 6):
            cell.number_format = '0.00'

# Summary by ProjectID, with Currency retained so original currencies are not mixed.
summary_rows = []
for project in PROJECTS:
    project_headers = [h for h in headers if h['ProjectID'] == project]
    if not project_headers:
        summary_rows.append((project, '', '', Decimal('0.00'), Decimal('0.00'), 0))
        continue
    by_currency = {}
    for h in project_headers:
        cur = h['Currency']
        by_currency.setdefault(cur, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
        by_currency[cur]['amount'] += money(h['NetAmount'])
        by_currency[cur]['twd'] += money(h['_NetAmountTWD'])
        by_currency[cur]['count'] += 1
    for cur in sorted(by_currency):
        d = by_currency[cur]
        summary_rows.append((project, '', cur, money(d['amount']), money(d['twd']), d['count']))

for r, row in enumerate(summary_rows, start=5):
    for c, v in enumerate(row, start=1):
        cell = summary_ws.cell(r, c, float(v) if isinstance(v, Decimal) else v)
        if c in (1, 2, 3):
            cell.number_format = '@'
        elif c in (4, 5):
            cell.number_format = '0.00'
        elif c == 6:
            cell.number_format = '0'

last_quote_row = max(5, 4 + len(headers))
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
summary_ws['B2'].number_format = '0.00'
wb.save(OUTFILE)
print('Saved:', OUTFILE)

# Verification from saved workbook.
wb2 = load_workbook(OUTFILE, data_only=False)
qws, lws, sws, rws = wb2['Quotes'], wb2['Lines'], wb2['Summary'], wb2['Report']
assert [qws.cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [lws.cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [sws.cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
quote_rows = list(qws.iter_rows(min_row=5, max_row=qws.max_row, values_only=True)) if qws.max_row >= 5 else []
line_rows = list(lws.iter_rows(min_row=5, max_row=lws.max_row, values_only=True)) if lws.max_row >= 5 else []
summary_saved = list(sws.iter_rows(min_row=5, max_row=sws.max_row, values_only=True)) if sws.max_row >= 5 else []
assert len(quote_rows) == 3, len(quote_rows)
assert len(line_rows) == 5, len(line_rows)
assert len(summary_saved) == 5, len(summary_saved)
assert {r[0] for r in quote_rows} == {'Q0011','Q0013','Q0015'}
for row in quote_rows:
    qid, customer, project, qdate, valid, status, cur, subtotal, discount, net, net_twd = row
    assert isinstance(qid, str) and isinstance(customer, str) and isinstance(project, str)
    assert project in PROJECTS and FROM_DATE <= qdate <= TO_DATE and status == STATUS
    assert isinstance(subtotal, (int, float)) and isinstance(net, (int, float)) and isinstance(net_twd, (int, float))
    assert money(net_twd) == money(Decimal(str(net)) * fx[cur]), (qid, net_twd, net, fx[cur])
summary_by_project = {r[0]: r for r in summary_saved}
assert list(summary_by_project.keys()) == PROJECTS, list(summary_by_project.keys())
assert summary_by_project['P003'][3] == 0 and summary_by_project['P003'][4] == 0 and summary_by_project['P003'][5] == 0
assert summary_by_project['P005'][3] == 0 and summary_by_project['P005'][4] == 0 and summary_by_project['P005'][5] == 0
assert summary_by_project['P001'][2] == 'USD' and money(summary_by_project['P001'][3]) == Decimal('10350.00') and money(summary_by_project['P001'][4]) == Decimal('331200.00')
assert summary_by_project['P002'][2] == 'EUR' and money(summary_by_project['P002'][3]) == Decimal('2784.31') and money(summary_by_project['P002'][4]) == Decimal('97450.85')
assert summary_by_project['P004'][2] == 'TWD' and money(summary_by_project['P004'][3]) == Decimal('282150.00') and money(summary_by_project['P004'][4]) == Decimal('282150.00')
assert sws['B2'].value == '=SUM(Quotes!K5:K7)', sws['B2'].value
quote_twd_total = sum(money(r[10]) for r in quote_rows)
summary_twd_total = sum(money(r[4]) for r in summary_saved)
assert quote_twd_total == summary_twd_total == Decimal('710800.85'), (quote_twd_total, summary_twd_total)
assert rws['B3'].value == FROM_DATE and rws['D3'].value == TO_DATE
print('Verification passed:')
print('- output workbook exists and was reopened:', OUTFILE)
print('- Quotes headers preserved; 3 Accepted rows Q0011/Q0013/Q0015; IDs are text; money fields are numeric')
print('- NetAmountTWD checked from saved Quotes using fx_rates.csv: USD 32, EUR 35, TWD 1')
print('- Lines headers preserved; 5 line-item rows')
print('- Summary has one row for each P001-P005; P003 and P005 are zero rows')
print('- Summary original Amount grouped with Currency (USD/EUR/TWD), so original currencies are not mixed')
print('- Summary B2 formula is =SUM(Quotes!K5:K7); saved TWD total is 710800.85')