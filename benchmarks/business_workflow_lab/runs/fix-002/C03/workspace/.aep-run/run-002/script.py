from pathlib import Path
import subprocess, json, sys, csv
from openpyxl import load_workbook
ws=Path.cwd()
print('Instructions rows:')
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
print('\nTest quote_cli list P001 accepted Sept JSON:')
cli=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
cmd=[sys.executable, str(cli), 'quotes', 'list', '--project','P001','--from','2026-09-01','--to','2026-09-30','--status','Accepted','--format','json']
print('cmd:', cmd)
res=subprocess.run(cmd, text=True, capture_output=True)
print('returncode', res.returncode)
print('stdout', res.stdout[:2000])
print('stderr', res.stderr[:1000])