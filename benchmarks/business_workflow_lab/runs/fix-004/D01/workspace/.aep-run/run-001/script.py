from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
# template structure
wb=load_workbook(ws/'input'/'expense_template.xlsx')
print('template sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)
# fx sample
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r): print('fx', row)
# manual first lines
print('manual sample:')
print('\n'.join((ws/'input'/'expense_manual.md').read_text(encoding='utf-8').splitlines()[:40]))