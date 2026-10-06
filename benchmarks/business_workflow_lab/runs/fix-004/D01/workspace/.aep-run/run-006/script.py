from pathlib import Path
from openpyxl import load_workbook
from collections import defaultdict
import json

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size)
print('OUTPUT FILES')
for p in sorted((ws/'output').glob('*')):
    print(' -', p.name, p.stat().st_size)

report = ws/'output'/'E001_september_expenses.xlsx'
export = ws/'output'/'expense_export_E001_2026-09.xlsx'
assert report.exists(), 'missing output/E001_september_expenses.xlsx'
assert export.exists(), 'missing output/expense_export_E001_2026-09.xlsx'

wb = load_workbook(report, data_only=False)
print('SHEETS', wb.sheetnames)
required_sheets = ['Report','Details','Summary','Exceptions','Instructions']
assert wb.sheetnames == required_sheets, f'unexpected sheets {wb.sheetnames}'

# Details
sh = wb['Details']
headers = [sh.cell(4,c).value for c in range(1, sh.max_column+1)]
print('DETAILS_HEADERS', headers)
required_cols = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert headers[:len(required_cols)] == required_cols, 'Details columns do not match template'
idx = {h:i+1 for i,h in enumerate(headers) if h}
rows=[]
for r in range(5, sh.max_row+1):
    if not sh.cell(r, idx['ExpenseID']).value:
        continue
    rows.append({c: sh.cell(r, idx[c]).value for c in required_cols})
print('DETAILS_ROW_COUNT', len(rows))
print('DETAILS_SAMPLE', json.dumps(rows[:3], ensure_ascii=False, default=str))
assert len(rows) == 13, f'expected 13 details, got {len(rows)}'
for row in rows:
    assert row['EmployeeID'] == 'E001', f"wrong employee {row}"
    ds = row['ExpenseDate'].strftime('%Y-%m-%d') if hasattr(row['ExpenseDate'], 'strftime') else str(row['ExpenseDate'])[:10]
    assert '2026-09-01' <= ds <= '2026-09-30', f'date outside range: {row}'
    assert isinstance(row['Amount'], (int,float)), f"Amount not numeric: {row}"
    assert isinstance(row['AmountTWD'], (int,float)), f"AmountTWD not numeric: {row}"

# Report metadata
rep = wb['Report']
print('REPORT_B3_D3_B4_B5_B7_D7', rep['B3'].value, rep['D3'].value, rep['B4'].value, rep['B5'].value, rep['B7'].value, rep['D7'].value)
assert str(rep['B3'].value)[:10] == '2026-09-01'
assert str(rep['D3'].value)[:10] == '2026-09-30'
assert 'E001' in str(rep['B4'].value)

# Summary validation using template columns Group/SubGroup/Currency/Amount/AmountTWD/Count
sumsh = wb['Summary']
print('SUMMARY_TOTAL_FORMULA_B2', sumsh['B2'].value)
assert sumsh['B2'].value == '=SUM(Details!I5:I17)', f"unexpected total formula {sumsh['B2'].value}"
summary_headers = [sumsh.cell(4,c).value for c in range(1, sumsh.max_column+1)]
print('SUMMARY_HEADERS', summary_headers)
assert summary_headers[:6] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
sidx = {h:i+1 for i,h in enumerate(summary_headers) if h}
summary_rows=[]
for r in range(5, sumsh.max_row+1):
    if not sumsh.cell(r, sidx['SubGroup']).value:
        continue
    row = {h: sumsh.cell(r,c).value for h,c in sidx.items()}
    summary_rows.append(row)
print('SUMMARY_ROWS', json.dumps(summary_rows, ensure_ascii=False, default=str))
calc = defaultdict(lambda: {'Amount':0.0,'AmountTWD':0.0,'Count':0})
for row in rows:
    key=(row['Status'], row['Currency'])
    calc[key]['Amount'] += float(row['Amount'])
    calc[key]['AmountTWD'] += float(row['AmountTWD'])
    calc[key]['Count'] += 1
summary_map={(r['SubGroup'], r['Currency']): r for r in summary_rows if r['Group']=='Status'}
assert set(summary_map) == set(calc), f'summary groups mismatch: {set(summary_map)} vs {set(calc)}'
for key, expected in calc.items():
    got=summary_map[key]
    assert abs(float(got['Amount'])-expected['Amount']) <= 0.01, f'{key} Amount mismatch'
    assert abs(float(got['AmountTWD'])-expected['AmountTWD']) <= 0.01, f'{key} AmountTWD mismatch'
    assert int(got['Count']) == expected['Count'], f'{key} Count mismatch'
twd_total = sum(float(row['AmountTWD']) for row in rows)
print('TWD_TOTAL', round(twd_total, 2))
assert abs(twd_total - 31877.90) <= 0.01, f'TWD total mismatch {twd_total}'

# Exceptions validation
ex = wb['Exceptions']
ex_headers = [ex.cell(4,c).value for c in range(1, ex.max_column+1)]
print('EXCEPTION_HEADERS', ex_headers)
assert ex_headers[:4] == ['ExpenseID','EmployeeID','ExpenseDate','Reason']
ex_idx = {h:i+1 for i,h in enumerate(ex_headers) if h}
ex_rows=[]
for r in range(5, ex.max_row+1):
    if ex.cell(r, ex_idx['ExpenseID']).value:
        ex_rows.append({h: ex.cell(r,c).value for h,c in ex_idx.items()})
print('EXCEPTION_ROWS', json.dumps(ex_rows, ensure_ascii=False, default=str))
expected_ex=[]
for row in rows:
    missing_receipt = str(row['MissingReceipt']).lower() == 'true'
    blank_receipt = row['ReceiptID'] in (None, '')
    if missing_receipt or blank_receipt:
        expected_ex.append(row['ExpenseID'])
assert [r['ExpenseID'] for r in ex_rows] == expected_ex, f'exceptions mismatch: {ex_rows} expected {expected_ex}'

# Export file inspection
exwb = load_workbook(export, data_only=True, read_only=True)
print('EXPORT_SHEETS', exwb.sheetnames)
exsh = exwb[exwb.sheetnames[0]]
print('EXPORT_DIM', exsh.max_row, exsh.max_column)
export_headers = [exsh.cell(1,c).value for c in range(1, exsh.max_column+1)]
print('EXPORT_HEADERS', export_headers)
assert exsh.max_row == 14, f'export should include 1 header + 13 rows, got max_row={exsh.max_row}'

print('VALIDATION_OK')