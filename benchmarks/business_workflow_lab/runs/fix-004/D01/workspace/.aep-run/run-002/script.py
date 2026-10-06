from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
out=ws/'output'
print('output files:', [p.name for p in out.iterdir()])
export=out/'expense_export_E001_2026-09.xlsx'
print('export exists', export.exists(), 'size', export.stat().st_size if export.exists() else None)
wb=load_workbook(export, data_only=False)
print('export sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(6, sh.max_row), values_only=True):
        print(row)