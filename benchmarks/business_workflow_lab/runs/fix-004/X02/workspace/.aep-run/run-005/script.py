from pathlib import Path
import csv, json, os
from openpyxl import load_workbook

ws = Path.cwd()
out = ws / 'output'
print('WORKSPACE', ws)
print('OUTPUT_EXISTS', out.exists())
print('OUTPUT_FILES')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print('-', p.relative_to(ws), p.stat().st_size)

required = [
    out / 'project_report_C001_2026-09.xlsx',
    out / 'downloads' / 'orders_C001_2026-09.csv',
    out / 'downloads' / 'invoices_C001_2026-09.csv',
    out / 'expenses_2026-09_all_projects.csv',
    out / 'quotes_C001_2026-09.json',
]
missing = [str(p.relative_to(ws)) for p in required if not p.exists()]
print('MISSING', missing)
if missing:
    raise SystemExit(1)

# count source/export rows
with open(out / 'downloads' / 'orders_C001_2026-09.csv', newline='', encoding='utf-8-sig') as f:
    orders = list(csv.DictReader(f))
with open(out / 'downloads' / 'invoices_C001_2026-09.csv', newline='', encoding='utf-8-sig') as f:
    invoices = list(csv.DictReader(f))
with open(out / 'expenses_2026-09_all_projects.csv', newline='', encoding='utf-8-sig') as f:
    expenses = list(csv.DictReader(f))
with open(out / 'quotes_C001_2026-09.json', encoding='utf-8') as f:
    quotes_obj = json.load(f)
quotes = quotes_obj if isinstance(quotes_obj, list) else quotes_obj.get('quotes', [])
print('ROW_COUNTS', {'orders_csv': len(orders), 'invoices_csv': len(invoices), 'expenses_csv': len(expenses), 'quotes_json': len(quotes)})

wb = load_workbook(out / 'project_report_C001_2026-09.xlsx', data_only=False)
print('WORKBOOK_SHEETS', wb.sheetnames)
checks = {}
for sheet in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    if sheet not in wb.sheetnames:
        checks[f'{sheet}_exists'] = False
        continue
    sh = wb[sheet]
    headers = [c.value for c in sh[1]]
    data_rows = []
    for row in sh.iter_rows(min_row=2, values_only=True):
        if any(v is not None for v in row):
            data_rows.append(row)
    checks[f'{sheet}_headers'] = headers
    checks[f'{sheet}_rows'] = len(data_rows)

# Look for obvious formulas/ranges in Summary
if 'Summary' in wb.sheetnames:
    sh = wb['Summary']
    formulas = []
    for row in sh.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith('='):
                formulas.append(cell.value)
    print('SUMMARY_FORMULAS', formulas[:20])
print('WORKBOOK_CHECKS', checks)

# Validate expected row counts from prior deliverable
expected_counts = {
    'Orders_rows': 4,
    'Invoices_rows': 3,
    'Expenses_rows': 18,
    'Quotes_rows': 4,
    'Summary_rows': 2,
    'Exceptions_rows': 18,
}
failures = []
for key, expected in expected_counts.items():
    actual = checks.get(key)
    if actual != expected:
        failures.append((key, actual, expected))
print('COUNT_CHECK_FAILURES', failures)
if failures:
    raise SystemExit(2)
print('CHECK_RESULT OK')