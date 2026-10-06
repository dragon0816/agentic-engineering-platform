from pathlib import Path
from openpyxl import load_workbook
import csv
base=Path('.')
print('INPUT FILES:')
for p in sorted((base/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(base/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(10), rdr):
        print(row)
print('\nManual headers/samples:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=base/'input'/name
    txt=p.read_text(encoding='utf-8')
    print('\n---', name, '---')
    print(txt[:3000])
print('\nWorkbook sheets and dimensions:')
wb=load_workbook(base/'input'/'project_template.xlsx', data_only=False)
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,12), values_only=True):
        print(r)
    print('---')