from pathlib import Path
import json, csv
from openpyxl import load_workbook
ws=Path.cwd(); out=ws/'output'
print('FILES')
for p in sorted(out.rglob('*')):
    if p.is_file(): print(p.relative_to(ws), p.stat().st_size)
qpath=out/'quotes_C001_2026-09.json'
obj=json.load(open(qpath,encoding='utf-8'))
print('QUOTE_JSON_TYPE', type(obj).__name__)
if isinstance(obj, dict):
    print('QUOTE_JSON_KEYS', list(obj.keys()))
    for k,v in obj.items():
        print('KEY', k, 'TYPE', type(v).__name__, 'LEN', len(v) if hasattr(v,'__len__') else None)
        if isinstance(v, list) and v: print('SAMPLE', v[0])
else:
    print('QUOTE_JSON_LEN', len(obj)); print('SAMPLE', obj[0] if obj else None)

wb=load_workbook(out/'project_report_C001_2026-09.xlsx', data_only=False)
for s in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    sh=wb[s]
    print('\nSHEET',s,'max',sh.max_row,sh.max_column)
    for r in range(1, min(sh.max_row,10)+1):
        vals=[sh.cell(r,c).value for c in range(1,min(sh.max_column,12)+1)]
        print(r, vals)

# infer row counts after real header rows by locating expected header names
header_markers = {
 'Orders':'OrderID', 'Invoices':'InvoiceID', 'Expenses':'ExpenseID', 'Quotes':'QuoteID', 'Summary':'ProjectID', 'Exceptions':'ProjectID'
}
for s,marker in header_markers.items():
    sh=wb[s]
    header_row=None
    for r in range(1, sh.max_row+1):
        vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
        if marker in vals:
            header_row=r; break
    rows=[]
    if header_row:
        for r in range(header_row+1, sh.max_row+1):
            vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
            if any(v is not None for v in vals): rows.append(vals)
    print('INFERRED', s, 'header_row', header_row, 'data_rows', len(rows))
