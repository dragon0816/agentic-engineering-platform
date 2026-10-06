from pathlib import Path
import os, csv
from openpyxl import load_workbook
ws = Path.cwd()
print('cwd=', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, 'size=', p.stat().st_size)
manual = ws/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md first 120 lines:')
print('\n'.join(manual.read_text(encoding='utf-8').splitlines()[:120]))
fx = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r): print(row)
wb_path = ws/'input'/'quote_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('\nworkbook sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('sheet', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,12), values_only=False):
        vals=[]
        for c in row[:min(wsht.max_column,8)]:
            vals.append(c.value)
        print(vals)
    print('---')
