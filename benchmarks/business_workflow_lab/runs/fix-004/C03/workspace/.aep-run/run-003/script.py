from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
ins=wb['Instructions']
print('Instructions:')
for row in ins.iter_rows(values_only=True):
    print(row)
print('\nReport merged:', list(wb['Report'].merged_cells.ranges))
for sh in ['Quotes','Lines','Summary']:
    s=wb[sh]
    print(sh, 'merged', list(s.merged_cells.ranges))
