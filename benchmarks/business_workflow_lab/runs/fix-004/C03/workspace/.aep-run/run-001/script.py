from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r):
        print(i, row)
print('\nquote_cli_manual.md first 80 lines:')
text=(ws/'input'/'quote_cli_manual.md').read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:80],1):
    print(f'{i:02d}: {line}')
print('\nquote_template.xlsx structure:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for sh in wb.sheetnames:
    s=wb[sh]
    print('sheet', sh, 'max_row', s.max_row, 'max_col', s.max_column)
    for row in s.iter_rows(min_row=1, max_row=min(s.max_row,8), values_only=True):
        print(row)
    print('---')
