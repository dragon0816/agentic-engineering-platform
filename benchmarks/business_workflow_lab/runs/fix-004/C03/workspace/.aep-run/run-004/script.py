from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, json, subprocess, sys, shutil
from datetime import datetime, timezone, timedelta, date
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
OUTFILE = OUTPUT / 'quote_template.xlsx'
PROJECTS = [f'P{i:03d}' for i in range(1, 6)]
DATE_FROM = '2026-09-01'
DATE_TO = '2026-09-30'
STATUS = 'Accepted'
TW = timezone(timedelta(hours=8))
GENERATED_AT = datetime.now(TW).replace(microsecond=0).isoformat()

TWOPLACES = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def to_number(d):
    # OpenPyXL writes Decimal as numeric in xlsx, but use float for broad Excel compatibility.
    return float(money(d))

# Read fixed exchange rates. Notes contain commas, but first two columns are stable.
fx = {}
with open(FX_FILE, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0].strip():
            continue
        fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates loaded:', fx)

# Fetch accepted September 2026 quotes for each requested project.
headers = []
lines = []
raw_counts = {}
for project in PROJECTS:
    cmd = [sys.executable, str(CLI), 'quotes', 'list', '--project', project,
           '--from', DATE_FROM, '--to', DATE_TO, '--status', STATUS, '--format', 'json']
    cp = subprocess.run(cmd, capture_output=True, text=True)
    if cp.returncode != 0:
        raise RuntimeError(f'quote list failed for {project}: exit {cp.returncode}; stderr={cp.stderr}')
    data = json.loads(cp.stdout)
    if data.get('schema_version') != 1:
        raise RuntimeError(f'unexpected schema_version for list {project}: {data.get("schema_version")}')
    items = data.get('items', [])
    raw_counts[project] = len(items)
    print(f'{project}: {len(items)} accepted quote(s)')
    for item in items:
        # Confirm full quote and collect line items.
        qid = item['QuoteID']
        gp = subprocess.run([sys.executable, str(CLI), 'quotes', 'get', qid, '--format', 'json'],
                            capture_output=True, text=True)
        if gp.returncode != 0:
            raise RuntimeError(f'quote get failed for {qid}: exit {gp.returncode}; stderr={gp.stderr}')
        qdoc = json.loads(gp.stdout)
        if qdoc.get('schema_version') != 1:
            raise RuntimeError(f'unexpected schema_version for get {qid}: {qdoc.get("schema_version")}')
        h = qdoc['header']
        # Defensively re-check filters; do not include anything outside requested scope.
        if not (h['ProjectID'] == project and h['Status'] == STATUS and DATE_FROM <= h['QuoteDate'] <= DATE_TO):
            raise RuntimeError(f'quote {qid} did not match requested filters: {h}')
        ccy = h['Currency']
        if ccy not in fx:
            raise RuntimeError(f'missing FX rate for {ccy}')
        net = money(h['NetAmount'])
        h = dict(h)
        h['NetAmountTWD'] = money(net * fx[ccy])
        headers.append(h)
        for li in qdoc.get('line_items', []):
            row = dict(li)
            row['QuoteID'] = qid
            lines.append(row)

headers.sort(key=lambda r: (r['ProjectID'], r['QuoteDate'], r['QuoteID']))
lines.sort(key=lambda r: (r['QuoteID'], r['ProductID'], r['ProductName']))
print('Fetched headers:', [h['QuoteID'] for h in headers])
print('Fetched line count:', len(lines))

# Load template and fill it; never modify input file.
wb = load_workbook(TEMPLATE)

# Helper to clear old data rows without disturbing headers/template areas.
def clear_rows(ws, start_row=5):
    if ws.max_row >= start_row:
        for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row):
            for cell in row:
                cell.value = None

def style_data_row(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)

# Report sheet
report = wb['Report']
report['B3'] = DATE_FROM
report['D3'] = DATE_TO
report['B4'] = 'ProjectID P001-P005; QuoteDate 2026-09-01..2026-09-30; Status Accepted'
report['B5'] = GENERATED_AT
report['B7'] = 'quote_cli; fx_rates.csv'
report['D7'] = GENERATED_AT
no_quote = [p for p in PROJECTS if raw_counts.get(p, 0) == 0]
report['B9'] = 'No Accepted quotes retained as zero rows for: ' + ', '.join(no_quote) if no_quote else ''

# Quotes sheet
qws = wb['Quotes']
clear_rows(qws, 5)
q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for idx, h in enumerate(headers, start=5):
    style_data_row(qws, 4, idx, len(q_headers))
    for col, key in enumerate(q_headers, start=1):
        cell = qws.cell(idx, col)
        if key in ('Subtotal','DiscountPercent','NetAmount','NetAmountTWD'):
            cell.value = to_number(h[key])
            cell.number_format = '0.00' if key != 'DiscountPercent' else '0'
        else:
            cell.value = str(h[key])
    # Force ID/text and date display formats.
    for col in [1,2,3,4,5,6,7]:
        qws.cell(idx, col).number_format = '@'

# Lines sheet
lws = wb['Lines']
clear_rows(lws, 5)
l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for idx, li in enumerate(lines, start=5):
    style_data_row(lws, 4, idx, len(l_headers))
    for col, key in enumerate(l_headers, start=1):
        cell = lws.cell(idx, col)
        if key == 'Quantity':
            cell.value = int(li[key])
            cell.number_format = '0'
        elif key in ('UnitPrice','LineAmount'):
            cell.value = to_number(li[key])
            cell.number_format = '0.00'
        else:
            cell.value = str(li[key])
            cell.number_format = '@'

# Summary by ProjectID. Keep one row for every P001..P005; zero rows use Amount=0, AmountTWD=0, Count=0.
sws = wb['Summary']
clear_rows(sws, 5)
summary_rows = []
for project in PROJECTS:
    hs = [h for h in headers if h['ProjectID'] == project]
    count = len(hs)
    twd_total = sum((money(h['NetAmountTWD']) for h in hs), Decimal('0.00'))
    currencies = sorted({h['Currency'] for h in hs})
    if count == 0:
        currency = ''
        amount = Decimal('0.00')
    elif len(currencies) == 1:
        currency = currencies[0]
        amount = sum((money(h['NetAmount']) for h in hs), Decimal('0.00'))
    else:
        # Do not mix original-currency amounts. This branch is not expected for the supplied data.
        currency = 'MIXED'
        amount = None
    summary_rows.append({
        'Group': project,
        'SubGroup': '',
        'Currency': currency,
        'Amount': amount,
        'AmountTWD': money(twd_total),
        'Count': count,
    })

for idx, r in enumerate(summary_rows, start=5):
    style_data_row(sws, 4, idx, 6)
    sws.cell(idx, 1).value = r['Group']; sws.cell(idx, 1).number_format = '@'
    sws.cell(idx, 2).value = r['SubGroup']; sws.cell(idx, 2).number_format = '@'
    sws.cell(idx, 3).value = r['Currency']; sws.cell(idx, 3).number_format = '@'
    sws.cell(idx, 4).value = None if r['Amount'] is None else to_number(r['Amount']); sws.cell(idx, 4).number_format = '0.00'
    sws.cell(idx, 5).value = to_number(r['AmountTWD']); sws.cell(idx, 5).number_format = '0.00'
    sws.cell(idx, 6).value = int(r['Count']); sws.cell(idx, 6).number_format = '0'

last_quote_row = max(5, 4 + len(headers))
sws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'

# Preserve workbook structure; save output.
wb.save(OUTFILE)
print('Saved', OUTFILE)

# Reopen saved workbook and verify requirements from the actual output file.
check_wb = load_workbook(OUTFILE, data_only=False)
errors = []
if set(check_wb.sheetnames) != set(load_workbook(TEMPLATE).sheetnames):
    errors.append('sheet names changed')
for sh in load_workbook(TEMPLATE).sheetnames:
    if sh not in check_wb.sheetnames:
        errors.append(f'missing sheet {sh}')
# Header checks
for sh, expected in [('Quotes', q_headers), ('Lines', l_headers), ('Summary', ['Group','SubGroup','Currency','Amount','AmountTWD','Count'])]:
    actual = [check_wb[sh].cell(4, c).value for c in range(1, len(expected)+1)]
    if actual != expected:
        errors.append(f'{sh} headers mismatch: {actual}')
# Quote checks
cq = check_wb['Quotes']
quote_rows = []
for row in range(5, 5+len(headers)):
    vals = [cq.cell(row, c).value for c in range(1, 12)]
    quote_rows.append(vals)
    for c in [1,2,3]:
        if not isinstance(cq.cell(row,c).value, str):
            errors.append(f'Quotes row {row} col {c} is not text')
    for c in [8,9,10,11]:
        if not isinstance(cq.cell(row,c).value, (int,float)):
            errors.append(f'Quotes row {row} col {c} is not numeric')
if len(quote_rows) != 3:
    errors.append(f'expected 3 quote rows, got {len(quote_rows)}')
if any(r[2] not in PROJECTS or r[5] != STATUS or not (DATE_FROM <= r[3] <= DATE_TO) for r in quote_rows):
    errors.append('a quote row is outside requested filters')
expected_twd_by_quote = {
    h['QuoteID']: float(money(Decimal(str(h['NetAmount'])) * fx[h['Currency']])) for h in headers
}
for vals in quote_rows:
    if abs(vals[10] - expected_twd_by_quote[vals[0]]) > 0.001:
        errors.append(f'NetAmountTWD mismatch for {vals[0]}: {vals[10]} vs {expected_twd_by_quote[vals[0]]}')
# Lines checks
cl = check_wb['Lines']
line_count = 0
quote_ids = {h['QuoteID'] for h in headers}
for row in range(5, 5+len(lines)):
    line_count += 1
    if cl.cell(row,1).value not in quote_ids:
        errors.append(f'line row {row} quote id does not join to Quotes')
    if not isinstance(cl.cell(row,1).value, str) or not isinstance(cl.cell(row,2).value, str):
        errors.append(f'Lines row {row} IDs are not text')
    for c in [4,5,6]:
        if not isinstance(cl.cell(row,c).value, (int,float)):
            errors.append(f'Lines row {row} col {c} is not numeric')
if line_count != len(lines):
    errors.append(f'line count mismatch: {line_count} vs {len(lines)}')
# Summary checks
cs = check_wb['Summary']
summary = []
for row in range(5, 10):
    summary.append([cs.cell(row,c).value for c in range(1,7)])
if [r[0] for r in summary] != PROJECTS:
    errors.append(f'summary projects mismatch: {[r[0] for r in summary]}')
for r in summary:
    if r[0] in no_quote and not (r[3] == 0 and r[4] == 0 and r[5] == 0):
        errors.append(f'zero project row incorrect: {r}')
    for cidx, val in [(4,r[3]),(5,r[4]),(6,r[5])]:
        if not isinstance(val, (int,float)):
            errors.append(f'Summary {r[0]} col {cidx} is not numeric')
if cs['B2'].value != '=SUM(Quotes!K5:K7)':
    errors.append(f'Summary B2 formula mismatch: {cs["B2"].value}')
# Confirm no mixed original-currency additions within each summary row.
for r in summary:
    project = r[0]
    ccys = {h['Currency'] for h in headers if h['ProjectID'] == project}
    if len(ccys) > 1 and r[3] is not None:
        errors.append(f'project {project} mixed original currencies but Amount is not blank')
if errors:
    print('CHECK FAILED')
    for e in errors:
        print('-', e)
    raise SystemExit(1)
print('CHECK PASSED')
print('Output file:', OUTFILE)
print('Quote rows:', len(quote_rows), 'Line rows:', line_count, 'Summary rows:', len(summary))
print('Summary rows:')
for r in summary:
    print(r)
print('Summary B2 formula:', cs['B2'].value)
