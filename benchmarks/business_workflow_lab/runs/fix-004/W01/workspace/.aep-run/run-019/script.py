from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
import csv, datetime

ws=Path.cwd()
out=ws/'output'/'september_open_orders.xlsx'
print('workspace:', ws)
print('output files:', sorted(str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()))

for rel in ['output/downloads/sales_orders_raw.csv','output/downloads/september_open_orders_raw.csv']:
    p=ws/rel
    print('\nCSV:', rel, 'exists', p.exists())
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as f:
            rdr=csv.DictReader(f); rows=list(rdr)
        print('headers:', rdr.fieldnames)
        print('row_count:', len(rows))
        print('all_statuses:', sorted(set(r['Status'] for r in rows)))
        print('date_min_max:', min(r['OrderDate'] for r in rows), max(r['OrderDate'] for r in rows))

wb=load_workbook(out, data_only=True)
print('\nWorkbook:', out, 'exists', out.exists())
print('sheets:', wb.sheetnames)

# Details parse row 4
sh=wb['Details']
headers=[sh.cell(4,c).value for c in range(1, sh.max_column+1)]
rows=[]
for rr in range(5, sh.max_row+1):
    vals=[sh.cell(rr,c).value for c in range(1, sh.max_column+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('\nDetails headers:', headers)
print('Details row_count:', len(rows))
print('first_detail:', rows[0] if rows else None)
print('last_detail:', rows[-1] if rows else None)

errors=[]
if len(rows)!=14: errors.append(f'Details row count expected 14 got {len(rows)}')
ids=[r['OrderID'] for r in rows]
if ids!=sorted(ids): errors.append('OrderID is not ascending')
if any(not isinstance(x,str) or not x.startswith('00') for x in ids): errors.append('OrderID text/leading zero issue')
if any(r['Status']!='Open' for r in rows): errors.append('non-Open status in Details')
for r in rows:
    d=r['OrderDate']
    if isinstance(d, datetime.datetime): d=d.date().isoformat()
    elif isinstance(d, datetime.date): d=d.isoformat()
    else: d=str(d)
    if d<'2026-09-01' or d>'2026-09-30': errors.append(f'out of range date {r["OrderID"]}: {d}')
blank_contacts=[r['OrderID'] for r in rows if not r.get('CustomerContact')]
blank_terms=[r['OrderID'] for r in rows if not r.get('PaymentTerms')]
print('blank CustomerContact IDs:', blank_contacts)
print('blank PaymentTerms IDs:', blank_terms)
if set(blank_contacts) - {'0045000113','0045000124'}: errors.append('CustomerContact is unexpectedly blank for most/all orders')
if blank_terms: errors.append('PaymentTerms has blanks')

rates={'TWD':Decimal('1'),'USD':Decimal('32'),'EUR':Decimal('35')}
total=Decimal('0')
for r in rows:
    expected=(Decimal(str(r['NetAmount']))*rates[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    got=Decimal(str(r['NetAmountTWD'])).quantize(Decimal('0.01'))
    if Decimal(str(r['RateToTWD']))!=rates[r['Currency']]: errors.append(f'rate mismatch {r["OrderID"]}')
    if got!=expected: errors.append(f'twd mismatch {r["OrderID"]}: {got} vs {expected}')
    total += got
print('Details total TWD:', total)

# Summary parse row 6 headers, row 7 onward
sumsh=wb['Summary']
summary_headers=[sumsh.cell(6,c).value for c in range(1,sumsh.max_column+1)]
summary_rows=[]
for rr in range(7,sumsh.max_row+1):
    vals=[sumsh.cell(rr,c).value for c in range(1,sumsh.max_column+1)]
    if any(v is not None for v in vals): summary_rows.append(dict(zip(summary_headers, vals)))
print('\nSummary headers:', summary_headers)
print('Summary rows:', summary_rows)
summary_total=sum(Decimal(str(r.get('AmountTWD') or 0)).quantize(Decimal('0.01')) for r in summary_rows)
print('Summary total TWD:', summary_total)
if summary_total!=total: errors.append(f'Summary total {summary_total} != Details total {total}')

report_total=wb['Report']['B3'].value
print('\nReport B3 grand total:', report_total)
if Decimal(str(report_total)).quantize(Decimal('0.01'))!=total: errors.append(f'Report total {report_total} != Details total {total}')

print('\nFINAL_CHECK_ERRORS:', errors)
if errors: raise SystemExit(4)
print('FINAL_CHECK_OK')