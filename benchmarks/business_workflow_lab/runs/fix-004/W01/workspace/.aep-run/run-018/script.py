import csv
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace:', ws)
print('input files:', sorted([p.name for p in (ws/'input').glob('*')]) if (ws/'input').exists() else [])
print('output files:', sorted([str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()]))

for rel in ['output/downloads/sales_orders_raw.csv', 'output/downloads/september_open_orders_raw.csv']:
    p = ws / rel
    print('\nCSV', rel, 'exists=', p.exists())
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        print(' headers:', reader.fieldnames)
        print(' row_count:', len(rows))
        print(' sample:', rows[:3])

out = ws / 'output' / 'september_open_orders.xlsx'
print('\nWorkbook exists=', out.exists(), out)
if not out.exists():
    raise SystemExit(2)
wb = load_workbook(out, data_only=True)
print('sheets:', wb.sheetnames)

# Inspect dimensions / first rows
for s in ['Report','Details','Summary']:
    sh = wb[s]
    print(f'\nSheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in range(1, min(sh.max_row, 6)+1):
        print(r, [sh.cell(r,c).value for c in range(1, min(sh.max_column, 14)+1)])

# Details: template states row 4 headers, data starts row 5
sh = wb['Details']
headers = [sh.cell(4,c).value for c in range(1, sh.max_column+1) if sh.cell(4,c).value is not None]
rows=[]
for r in range(5, sh.max_row+1):
    vals = [sh.cell(r,c).value for c in range(1, len(headers)+1)]
    if all(v is None for v in vals):
        continue
    rows.append(dict(zip(headers, vals)))
print('\nDetails parsed headers:', headers)
print('Details parsed row_count:', len(rows))
print('Details first:', rows[0] if rows else None)
print('Details last:', rows[-1] if rows else None)

errors=[]
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','CustomerContact','PaymentTerms','SalesRep','RateToTWD','NetAmountTWD']
for h in expected_headers:
    if h not in headers:
        errors.append(f'missing Details column {h}')
if len(rows) != 14:
    errors.append(f'expected 14 Details rows, got {len(rows)}')
ids = [str(r.get('OrderID')) for r in rows]
if ids != sorted(ids):
    errors.append('OrderID not sorted ascending')
if any(not isinstance(r.get('OrderID'), str) or not str(r.get('OrderID')).startswith('00') for r in rows):
    errors.append('some OrderID values are not text with leading zeroes')
if any(r.get('Status') != 'Open' for r in rows):
    errors.append('some Status values are not Open')
if any(str(r.get('OrderDate')) < '2026-09-01' or str(r.get('OrderDate')) > '2026-09-30' for r in rows):
    errors.append('some OrderDate values are outside September 2026')
allowed_blank_contacts = {'0045000113','0045000124'}
blank_contacts = {str(r.get('OrderID')) for r in rows if not r.get('CustomerContact')}
if not blank_contacts.issubset(allowed_blank_contacts):
    errors.append(f'unexpected blank CustomerContact IDs {sorted(blank_contacts)}')
if any(not r.get('PaymentTerms') for r in rows):
    errors.append('some PaymentTerms values are blank')

rates = {'TWD': Decimal('1'), 'USD': Decimal('32'), 'EUR': Decimal('35')}
detail_total = Decimal('0')
for r in rows:
    oid = str(r.get('OrderID'))
    cur = r.get('Currency')
    if cur not in rates:
        errors.append(f'unknown currency for {oid}: {cur}')
        continue
    rate = Decimal(str(r.get('RateToTWD')))
    if rate != rates[cur]:
        errors.append(f'rate mismatch for {oid}: {cur} {rate}')
    net = Decimal(str(r.get('NetAmount')))
    twd = Decimal(str(r.get('NetAmountTWD')))
    expected = (net * rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    if twd != expected:
        errors.append(f'NetAmountTWD mismatch for {oid}: got {twd}, expected {expected}')
    detail_total += twd
print('blank_contacts:', sorted(blank_contacts))
print('detail_total_twd:', detail_total)

# Summary parse: find header row containing Currency and NetAmountTWD
shs = wb['Summary']
summary_rows=[]
summary_headers=[]
for r in range(1, shs.max_row+1):
    vals=[shs.cell(r,c).value for c in range(1, shs.max_column+1)]
    if 'Currency' in vals and 'NetAmountTWD' in vals:
        summary_headers=[v for v in vals if v is not None]
        start=r+1
        for rr in range(start, shs.max_row+1):
            vv=[shs.cell(rr,c).value for c in range(1, len(summary_headers)+1)]
            if all(v is None for v in vv):
                continue
            summary_rows.append(dict(zip(summary_headers, vv)))
        break
print('Summary parsed headers:', summary_headers)
print('Summary parsed rows:', summary_rows)
summary_total = Decimal('0')
if not summary_rows:
    errors.append('could not parse Summary rows')
else:
    for r in summary_rows:
        if r.get('Currency') in (None, 'Grand Total'):
            continue
        val = r.get('NetAmountTWD')
        if val is not None:
            summary_total += Decimal(str(val))
print('summary_total_twd:', summary_total)
if summary_total != detail_total:
    errors.append(f'Summary total {summary_total} != Details total {detail_total}')

# Report: scan for detail total or grand total label
shr=wb['Report']
report_cells=[]
for row in shr.iter_rows(values_only=True):
    report_cells.extend([v for v in row if v is not None])
print('Report non-empty values:', report_cells)
found_total = any((isinstance(v,(int,float)) and Decimal(str(v)).quantize(Decimal('0.01')) == detail_total) or (isinstance(v,str) and str(detail_total) in v) for v in report_cells)
if not found_total:
    errors.append('Report does not contain checked grand total value')

print('\nCHECK_ERRORS:', errors)
if errors:
    raise SystemExit(3)
print('CHECK_OK: workbook and CSV outputs satisfy checked requirements')