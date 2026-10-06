from pathlib import Path
import csv, json, math
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

ws = Path.cwd()
in_dir = ws / 'input'
out_dir = ws / 'output'
download_dir = out_dir / 'downloads'
print('workspace:', ws)
print('input files:', sorted(p.name for p in in_dir.iterdir()) if in_dir.exists() else 'MISSING')
print('output files:', sorted(p.name for p in out_dir.iterdir()) if out_dir.exists() else 'MISSING')
print('download files:', sorted(p.name for p in download_dir.iterdir()) if download_dir.exists() else 'MISSING')

xlsx = out_dir / 'september_open_orders.xlsx'
raw_named = download_dir / 'september_open_orders_raw.csv'
raw_latest = download_dir / 'sales_orders_raw.csv'
fx_file = in_dir / 'fx_rates.csv'

errors = []
for p in [xlsx, raw_named, raw_latest, fx_file]:
    if not p.exists():
        errors.append(f'missing file: {p.relative_to(ws)}')

if raw_named.exists():
    with raw_named.open(newline='', encoding='utf-8-sig') as f:
        raw_rows = list(csv.DictReader(f))
    print('named raw headers:', list(raw_rows[0].keys()) if raw_rows else [])
    print('named raw row count:', len(raw_rows))
    print('named raw sample:', raw_rows[:3])
else:
    raw_rows = []

if raw_latest.exists():
    with raw_latest.open(newline='', encoding='utf-8-sig') as f:
        latest_rows = list(csv.DictReader(f))
    print('latest raw row count:', len(latest_rows))
    if raw_rows and latest_rows != raw_rows:
        errors.append('sales_orders_raw.csv and september_open_orders_raw.csv differ')

rates = {}
if fx_file.exists():
    with fx_file.open(newline='', encoding='utf-8-sig') as f:
        for r in csv.DictReader(f):
            rates[r['Currency']] = Decimal(str(r['RateToTWD']))
    print('fx rates:', {k: str(v) for k, v in rates.items()})
    for k, expected in {'TWD':'1', 'USD':'32', 'EUR':'35'}.items():
        if str(rates.get(k)) != expected:
            errors.append(f'fx rate {k} expected {expected}, got {rates.get(k)}')

def dec2(v):
    return Decimal(str(v)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

if xlsx.exists():
    wb = load_workbook(xlsx, data_only=True)
    print('workbook sheets:', wb.sheetnames)
    for sheet in ['Summary', 'Details']:
        if sheet not in wb.sheetnames:
            errors.append(f'missing sheet: {sheet}')
    if 'Details' in wb.sheetnames:
        wsdet = wb['Details']
        # find header row containing OrderID
        header_row = None
        headers = []
        for r in range(1, min(wsdet.max_row, 20) + 1):
            vals = [wsdet.cell(r, c).value for c in range(1, wsdet.max_column + 1)]
            if 'OrderID' in vals:
                header_row = r
                headers = vals
                break
        print('Details max row/col:', wsdet.max_row, wsdet.max_column)
        print('Details header row:', header_row)
        print('Details headers:', headers)
        if header_row is None:
            errors.append('Details header row with OrderID not found')
        else:
            rows = []
            for r in range(header_row + 1, wsdet.max_row + 1):
                vals = [wsdet.cell(r, c).value for c in range(1, len(headers) + 1)]
                if any(v is not None for v in vals):
                    rows.append(dict(zip(headers, vals)))
            print('Details row count:', len(rows))
            print('Details sample:', rows[:3])
            required_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','RateToTWD','NetAmountTWD','CustomerContact']
            for h in required_headers:
                if h not in headers:
                    errors.append(f'Details missing header: {h}')
            if raw_rows and len(rows) != len(raw_rows):
                errors.append(f'Details row count {len(rows)} != raw CSV row count {len(raw_rows)}')
            order_ids = [str(r.get('OrderID')) for r in rows]
            if order_ids != sorted(order_ids):
                errors.append('OrderID values are not sorted ascending')
            statuses = {r.get('Status') for r in rows}
            print('Statuses:', statuses)
            if statuses != {'Open'}:
                errors.append(f'Status values expected only Open, got {statuses}')
            start = date(2026,9,1); end = date(2026,9,30)
            for r in rows:
                od = r.get('OrderDate')
                if isinstance(od, datetime):
                    d = od.date()
                elif isinstance(od, date):
                    d = od
                else:
                    d = datetime.strptime(str(od), '%Y-%m-%d').date()
                if not (start <= d <= end):
                    errors.append(f'OrderDate out of range for {r.get("OrderID")}: {od}')
                oid = str(r.get('OrderID'))
                cur = r.get('Currency')
                if cur not in rates:
                    errors.append(f'No FX rate for {cur} in row {oid}')
                    continue
                try:
                    net = Decimal(str(r.get('NetAmount')))
                    rate = Decimal(str(r.get('RateToTWD')))
                    twd = Decimal(str(r.get('NetAmountTWD')))
                except Exception as e:
                    errors.append(f'Numeric conversion failed for {oid}: {e}')
                    continue
                if rate != rates[cur]:
                    errors.append(f'Rate mismatch for {oid}: {rate} != {rates[cur]}')
                expected_twd = dec2(net * rates[cur])
                if dec2(twd) != expected_twd:
                    errors.append(f'NetAmountTWD mismatch for {oid}: {twd} != {expected_twd}')
            blank_contacts = [str(r.get('OrderID')) for r in rows if r.get('CustomerContact') in (None, '')]
            print('Blank CustomerContact OrderIDs:', blank_contacts)
            if set(blank_contacts) != {'0045000113','0045000124'}:
                errors.append(f'Unexpected blank CustomerContact set: {blank_contacts}')
            # Summary checks: collect visible key/value-ish rows and totals by currency from Details
            by_currency = {}
            total_twd = Decimal('0')
            for r in rows:
                cur = r.get('Currency')
                by_currency.setdefault(cur, {'count': 0, 'net': Decimal('0'), 'twd': Decimal('0')})
                by_currency[cur]['count'] += 1
                by_currency[cur]['net'] += Decimal(str(r.get('NetAmount')))
                by_currency[cur]['twd'] += Decimal(str(r.get('NetAmountTWD')))
                total_twd += Decimal(str(r.get('NetAmountTWD')))
            print('Computed totals by currency:', {k: {'count': v['count'], 'net': str(dec2(v['net'])), 'twd': str(dec2(v['twd']))} for k,v in sorted(by_currency.items())})
            print('Computed total TWD:', str(dec2(total_twd)))
            if 'Summary' in wb.sheetnames:
                wss = wb['Summary']
                print('Summary max row/col:', wss.max_row, wss.max_column)
                summary_values = []
                for r in range(1, wss.max_row + 1):
                    vals = [wss.cell(r, c).value for c in range(1, min(wss.max_column, 8)+1)]
                    if any(v is not None for v in vals):
                        summary_values.append(vals)
                print('Summary non-empty rows:', summary_values[:20])
                # Search numeric total in sheet
                nums = []
                for row in wss.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, (int, float)):
                            nums.append(Decimal(str(cell.value)).quantize(Decimal('0.01')))
                if dec2(total_twd) not in nums:
                    errors.append('Summary does not contain computed total NetAmountTWD')

if errors:
    print('CHECK FAILED')
    for e in errors:
        print('ERROR:', e)
    raise SystemExit(1)
else:
    print('CHECK PASSED: output workbook and raw CSV satisfy the stated requirements.')
