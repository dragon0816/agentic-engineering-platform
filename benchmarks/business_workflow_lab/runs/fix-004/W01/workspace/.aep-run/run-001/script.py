from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(6), r): print(row)
print('\nsales_template.xlsx structure:')
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    sheet=wb[sh]
    print('sheet', sh, 'max_row', sheet.max_row, 'max_col', sheet.max_column)
    for row in sheet.iter_rows(min_row=1, max_row=min(8,sheet.max_row), values_only=True): print(row)
print('\nsap_manual.md:')
print((ws/'input'/'sap_manual.md').read_text(encoding='utf-8')[:4000])