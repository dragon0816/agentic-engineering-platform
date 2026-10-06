from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from openpyxl import load_workbook

ws=Path.cwd(); out=ws/'output'; dl=out/'downloads'
print('workspace:', ws)
print('output files:', sorted(str(p.relative_to(ws)) for p in out.rglob('*') if p.is_file()))

xlsx=out/'september_open_orders.xlsx'; raw=dl/'sales_orders_raw.csv'; named=dl/'september_open_orders_raw.csv'
with raw.open('r', encoding='utf-8-sig', newline='') as f:
    raw_rows=list(csv.DictReader(f))
print('RAW headers:', list(raw_rows[0].keys()))
print('RAW rows:', len(raw_rows))
print('RAW orderids:', [r['OrderID'] for r in raw_rows])
print('RAW rows full:')
for r in raw_rows: print(r)

rates={'TWD':Decimal('1'),'USD':Decimal('32'),'EUR':Decimal('35')}
def q2(x): return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
raw_total=sum((q2(r['NetAmount'])*rates[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP) for r in raw_rows)
print('RAW calc total TWD:', raw_total)
by={}
for r in raw_rows:
    cur=r['Currency']; by.setdefault(cur,[0,Decimal('0.00'),Decimal('0.00')]); by[cur][0]+=1; by[cur][1]+=q2(r['NetAmount']); by[cur][2]+=q2(q2(r['NetAmount'])*rates[cur])
print('RAW by currency:', {k:[v[0],str(v[1]),str(v[2])] for k,v in by.items()})

wb=load_workbook(xlsx, data_only=True)
print('sheets:', wb.sheetnames)
# find headers and rows
sh=wb['Details']; header_row=None
for r in range(1,20):
    vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
    if 'OrderID' in vals:
        header_row=r; headers=vals[:max(i for i,v in enumerate(vals,1) if v is not None)]; break
print('Details headers:', headers)
rows=[]
for r in range(header_row+1, sh.max_row+1):
    vals=[sh.cell(r,c).value for c in range(1,len(headers)+1)]
    if any(v is not None for v in vals): rows.append(dict(zip(headers, vals)))
print('Details row count:', len(rows))
for r in rows: print('DETAIL', r)

def date_ok(v):
    if isinstance(v, datetime): d=v.date().isoformat()
    elif isinstance(v, date): d=v.isoformat()
    else: d=str(v)[:10]
    return '2026-09-01' <= d <= '2026-09-30'
checks=[]
checks.append(('csv row count 14', len(raw_rows)==14))
checks.append(('details row count 14', len(rows)==14))
checks.append(('status Open', all(r['Status']=='Open' for r in rows)))
checks.append(('date range', all(date_ok(r['OrderDate']) for r in rows)))
checks.append(('order asc', [r['OrderID'] for r in rows]==sorted(r['OrderID'] for r in rows)))
checks.append(('OrderID text-like preserved', all(isinstance(r['OrderID'], str) and len(r['OrderID'])==10 and r['OrderID'].startswith('0045') for r in rows)))
checks.append(('headers exact', headers==['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','RateToTWD','NetAmountTWD','CustomerContact','PaymentTerms']))
# formula values
amount_ok=True
for r in rows:
    cur=r['Currency']; net=q2(r['NetAmount']); rate=Decimal(str(r['RateToTWD'])); twd=q2(r['NetAmountTWD'])
    if rate != rates[cur] or twd != q2(net*rate):
        print('amount mismatch', r['OrderID'], cur, net, rate, twd, 'expected', q2(net*rates[cur])); amount_ok=False
checks.append(('rates and NetAmountTWD', amount_ok))
details_total=sum(q2(r['NetAmountTWD']) for r in rows)
checks.append(('details total matches raw', details_total==raw_total))
# detail fields
blank_contacts=[r['OrderID'] for r in rows if r.get('CustomerContact') in (None,'')]
blank_terms=[r['OrderID'] for r in rows if r.get('PaymentTerms') in (None,'')]
print('blank_contacts:', blank_contacts)
print('blank_payment_terms:', blank_terms)
checks.append(('CustomerContact not all blank', len(blank_contacts)<len(rows)))
checks.append(('PaymentTerms not all blank', len(blank_terms)<len(rows)))
# summary report total
rep=wb['Report']; print('Report A1:F8:', [[rep.cell(r,c).value for c in range(1,7)] for r in range(1,9)])
sumsh=wb['Summary']; print('Summary all:', [[sumsh.cell(r,c).value for c in range(1,sumsh.max_column+1)] for r in range(1,sumsh.max_row+1)])
# summary values check
summary_total=None
for r in range(1,sumsh.max_row+1):
    vals=[sumsh.cell(r,c).value for c in range(1,sumsh.max_column+1)]
    if vals and vals[0]=='Grand Total TWD': summary_total=vals[3] if len(vals)>3 else None
print('summary_total:', summary_total, 'details_total:', details_total)
checks.append(('Summary grand total matches Details', q2(summary_total)==details_total if summary_total is not None else False))
checks.append(('Report grand total matches Details', q2(rep['B3'].value)==details_total if rep['B3'].value is not None else False))
for name, ok in checks: print(('PASS' if ok else 'FAIL'), name)
if not all(ok for _,ok in checks): raise SystemExit('VALIDATION FAILED')
print('VALIDATION PASSED')
