from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'output'/'september_open_orders.xlsx')
sh=wb['Details']
headers=[sh.cell(4,c).value for c in range(1, sh.max_column+1)]
print('headers=', headers)
for r in range(5, sh.max_row+1):
    row={h: sh.cell(r,i+1).value for i,h in enumerate(headers)}
    if any(v is not None for v in row.values()):
        print(r, row['OrderID'], 'contact=', row['CustomerContact'], 'terms=', row['PaymentTerms'])