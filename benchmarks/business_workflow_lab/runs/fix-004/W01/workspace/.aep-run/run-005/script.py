from pathlib import Path
from openpyxl import load_workbook
import csv, json

ws = Path.cwd()
print('workspace=', ws)
print('input files=', sorted(p.name for p in (ws/'input').glob('*')) if (ws/'input').exists() else 'NO input folder')
print('output files=', sorted(str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()) if (ws/'output').exists() else 'NO output folder')

out = ws/'output'/'september_open_orders.xlsx'
raw = ws/'output'/'downloads'/'september_open_orders_raw.csv'
print('workbook exists=', out.exists(), out)
print('raw csv exists=', raw.exists(), raw)

if not out.exists():
    raise SystemExit('missing output workbook')

wb = load_workbook(out, data_only=False)
print('sheets=', wb.sheetnames)
if 'Details' not in wb.sheetnames or 'Summary' not in wb.sheetnames:
    raise SystemExit('missing required sheets')

d = wb['Details']
s = wb['Summary']
headers = [d.cell(4, c).value for c in range(1, d.max_column+1)]
print('Details headers row 4=', headers)
rows = []
for r in range(5, d.max_row+1):
    vals = [d.cell(r, c).value for c in range(1, len(headers)+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('detail row count=', len(rows))
print('first 3 OrderID=', [str(row.get('OrderID')) for row in rows[:3]])
print('last 3 OrderID=', [str(row.get('OrderID')) for row in rows[-3:]])

# Locate relevant columns by known template names
required = ['OrderID','OrderDate','Status','Currency','Amount','FXRate','AmountTWD']
missing = [h for h in required if h not in headers]
print('missing required headers=', missing)
if missing:
    raise SystemExit('missing headers: '+','.join(missing))

order_ids = [str(row['OrderID']) for row in rows]
statuses = {row['Status'] for row in rows}
currencies = sorted({row['Currency'] for row in rows})
print('statuses=', sorted(statuses))
print('currencies=', currencies)
print('is sorted by OrderID=', order_ids == sorted(order_ids))
print('all dates in Sep 2026=', all(str(row['OrderDate'])[:10] >= '2026-09-01' and str(row['OrderDate'])[:10] <= '2026-09-30' for row in rows))
print('amount/FX/TWD numeric=', all(isinstance(row['Amount'], (int,float)) and isinstance(row['FXRate'], (int,float)) and isinstance(row['AmountTWD'], (int,float)) for row in rows))
print('Summary B2 formula/value=', s['B2'].value)

summary_rows=[]
for r in range(1, s.max_row+1):
    vals=[s.cell(r,c).value for c in range(1, min(s.max_column,5)+1)]
    if any(v is not None for v in vals):
        summary_rows.append(vals)
print('Summary non-empty rows=', summary_rows)

# Cross-check summary totals from details
from collections import defaultdict
orig=defaultdict(float); twd=defaultdict(float)
for row in rows:
    orig[row['Currency']] += float(row['Amount'])
    twd[row['Currency']] += float(row['AmountTWD'])
print('computed original by currency=', {k: round(v,2) for k,v in sorted(orig.items())})
print('computed TWD by currency=', {k: round(v,2) for k,v in sorted(twd.items())})
print('computed total TWD=', round(sum(twd.values()),2))

if raw.exists():
    with raw.open(newline='', encoding='utf-8-sig') as f:
        reader=csv.DictReader(f)
        raw_rows=list(reader)
    print('raw csv columns=', reader.fieldnames)
    print('raw csv row count=', len(raw_rows))

# Hard assertions for requested result
assert len(rows) > 0
assert statuses == {'Open'}
assert order_ids == sorted(order_ids)
assert all(str(row['OrderDate'])[:10] >= '2026-09-01' and str(row['OrderDate'])[:10] <= '2026-09-30' for row in rows)
assert all(isinstance(row['Amount'], (int,float)) and isinstance(row['FXRate'], (int,float)) and isinstance(row['AmountTWD'], (int,float)) for row in rows)
print('CHECK_RESULT=PASS')