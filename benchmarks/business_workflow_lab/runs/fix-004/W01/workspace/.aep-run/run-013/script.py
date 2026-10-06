import os, csv, json, re, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from urllib.request import urlopen
from html.parser import HTMLParser
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
date_from = params['date_from']
date_to = params['date_to']
status = params['status']
output_filename = params.get('output_filename') or 'sales_orders_report.xlsx'
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'

ws = Path.cwd()
in_dir = ws / 'input'
out_dir = ws / 'output'
download_dir = out_dir / 'downloads'
out_dir.mkdir(exist_ok=True)
download_dir.mkdir(exist_ok=True)
print('workspace:', ws)
print('input files:', sorted(p.name for p in in_dir.iterdir()))
print('download files:', sorted(p.name for p in download_dir.iterdir()))

raw_src = download_dir / 'sales_orders_raw.csv'
raw_named = download_dir / (Path(output_filename).stem + '_raw.csv')
if not raw_src.exists():
    raise FileNotFoundError(f'Missing downloaded CSV: {raw_src}')
shutil.copyfile(raw_src, raw_named)

with raw_src.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    raw_rows = list(reader)
    raw_headers = reader.fieldnames or []
print('raw headers:', raw_headers)
print('raw row count:', len(raw_rows))
print('raw sample:', raw_rows[:3])

fx_path = in_dir / 'fx_rates.csv'
rates = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    fx_reader = csv.DictReader(f)
    print('fx headers:', fx_reader.fieldnames)
    for r in fx_reader:
        rates[r['Currency']] = Decimal(str(r['RateToTWD']))
print('fx rates:', {k: str(v) for k, v in rates.items()})

def money(v):
    return Decimal(str(v)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

class DetailParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.in_cell = False
        self.current = ''
        self.cells = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() in ('td','th'):
            self.in_cell = True
            self.current = ''
    def handle_data(self, data):
        if self.in_cell:
            self.current += data
    def handle_endtag(self, tag):
        if tag.lower() in ('td','th') and self.in_cell:
            self.cells.append(' '.join(self.current.split()))
            self.in_cell = False
            self.current = ''

def fetch_detail(order_id):
    url = f'http://127.0.0.1:8810/orders/{order_id}'
    result = {'CustomerContact': '', 'PaymentTerms': ''}
    try:
        html = urlopen(url, timeout=10).read().decode('utf-8', errors='replace')
    except Exception as e:
        print(f'WARN: detail fetch failed for {order_id}: {e}')
        return result
    parser = DetailParser()
    parser.feed(html)
    cells = parser.cells
    for i in range(0, len(cells)-1):
        key = cells[i].strip().rstrip(':')
        val = cells[i+1].strip()
        if key in result:
            result[key] = val
    if not result['CustomerContact']:
        m = re.search(r'CustomerContact\s*</[^>]+>\s*<[^>]+>\s*([^<]*)', html, re.I)
        if m:
            result['CustomerContact'] = ' '.join(m.group(1).split())
    if not result['PaymentTerms']:
        m = re.search(r'PaymentTerms\s*</[^>]+>\s*<[^>]+>\s*([^<]*)', html, re.I)
        if m:
            result['PaymentTerms'] = ' '.join(m.group(1).split())
    return result

detail_cache = {r['OrderID']: fetch_detail(r['OrderID']) for r in raw_rows}
print('detail cache entries:', len(detail_cache))

start = datetime.strptime(date_from, '%Y-%m-%d').date()
end = datetime.strptime(date_to, '%Y-%m-%d').date()
rows = []
for r in raw_rows:
    order_date = datetime.strptime(r['OrderDate'], '%Y-%m-%d').date()
    if not (start <= order_date <= end):
        raise ValueError(f'OrderDate out of range for {r["OrderID"]}: {r["OrderDate"]}')
    if r['Status'] != status:
        raise ValueError(f'Status mismatch for {r["OrderID"]}: {r["Status"]}')
    cur = r['Currency']
    if cur not in rates:
        raise ValueError(f'Missing FX rate for {cur}')
    net = money(r['NetAmount'])
    rate = rates[cur]
    detail = detail_cache.get(r['OrderID'], {})
    rows.append({'OrderID': r['OrderID'], 'CustomerID': r['CustomerID'], 'CustomerName': r['CustomerName'], 'ProjectID': r.get('ProjectID') or '', 'OrderDate': order_date, 'Status': r['Status'], 'Currency': cur, 'NetAmount': float(net), 'RateToTWD': float(rate), 'NetAmountTWD': float(money(net * rate)), 'CustomerContact': detail.get('CustomerContact') or '', 'PaymentTerms': detail.get('PaymentTerms') or ''})
rows.sort(key=lambda x: x['OrderID'])

wb = load_workbook(in_dir / 'sales_template.xlsx')
if 'Details' not in wb.sheetnames:
    wb.create_sheet('Details')
wsd = wb['Details']
header_row = None
for rr in range(1, min(wsd.max_row, 25) + 1):
    vals = [wsd.cell(rr, cc).value for cc in range(1, wsd.max_column + 1)]
    if 'OrderID' in vals:
        header_row = rr
        break
if header_row is None:
    header_row = 4
headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','RateToTWD','NetAmountTWD','CustomerContact','PaymentTerms']
for rr in range(header_row, max(wsd.max_row, header_row + len(rows) + 5) + 1):
    for cc in range(1, max(wsd.max_column, len(headers)) + 1):
        wsd.cell(rr, cc).value = None
for cc, h in enumerate(headers, 1):
    wsd.cell(header_row, cc).value = h
for idx, row in enumerate(rows, header_row + 1):
    for cc, h in enumerate(headers, 1):
        wsd.cell(idx, cc).value = row[h]
for rr in range(header_row + 1, header_row + 1 + len(rows)):
    wsd.cell(rr, 5).number_format = 'yyyy-mm-dd'
    for cc in (8, 9, 10):
        wsd.cell(rr, cc).number_format = '#,##0.00'

if 'Summary' not in wb.sheetnames:
    wbs = wb.create_sheet('Summary')
else:
    wbs = wb['Summary']
for row in wbs.iter_rows():
    for cell in row:
        cell.value = None
wbs['A1'] = 'Sales summary'
wbs['A2'] = 'DateFrom'; wbs['B2'] = date_from
wbs['A3'] = 'DateTo'; wbs['B3'] = date_to
wbs['A4'] = 'Status'; wbs['B4'] = status
wbs['A6'] = 'Currency'; wbs['B6'] = 'Amount'; wbs['C6'] = 'RateToTWD'; wbs['D6'] = 'AmountTWD'; wbs['E6'] = 'Count'
by_currency = {}
for row in rows:
    cur = row['Currency']
    by_currency.setdefault(cur, {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0, 'rate': rates[cur]})
    by_currency[cur]['amount'] += Decimal(str(row['NetAmount']))
    by_currency[cur]['twd'] += Decimal(str(row['NetAmountTWD']))
    by_currency[cur]['count'] += 1
summary_row = 7
total_twd = Decimal('0')
for cur in sorted(by_currency):
    item = by_currency[cur]
    wbs.cell(summary_row, 1).value = cur
    wbs.cell(summary_row, 2).value = float(money(item['amount']))
    wbs.cell(summary_row, 3).value = float(item['rate'])
    wbs.cell(summary_row, 4).value = float(money(item['twd']))
    wbs.cell(summary_row, 5).value = item['count']
    total_twd += money(item['twd'])
    summary_row += 1
wbs.cell(summary_row, 1).value = 'Grand Total TWD'
wbs.cell(summary_row, 4).value = float(money(total_twd))
wbs.cell(summary_row, 5).value = len(rows)
for rr in range(7, summary_row + 1):
    for cc in (2,3,4):
        wbs.cell(rr, cc).number_format = '#,##0.00'

if 'Report' in wb.sheetnames:
    wbr = wb['Report']
    wbr['A1'] = 'Sales Orders Report'
    wbr['A2'] = f'{date_from} to {date_to} / {status}'
    wbr['A3'] = 'Grand Total TWD'
    wbr['B3'] = float(money(total_twd))

out_path = out_dir / output_filename
wb.save(out_path)
print('saved workbook:', out_path)
print('saved raw csv:', raw_named)

wb2 = load_workbook(out_path, data_only=True)
wsd2 = wb2['Details']
found = None
for rr in range(1, 25):
    vals = [wsd2.cell(rr, cc).value for cc in range(1, wsd2.max_column + 1)]
    if 'OrderID' in vals:
        found = rr
        saved_headers = vals[:len(headers)]
        break
if found is None:
    raise AssertionError('Details header not found after save')
print('saved Details header row:', found)
print('saved Details headers:', saved_headers)
if saved_headers != headers:
    raise AssertionError(f'Unexpected Details headers: {saved_headers}')
saved_rows = []
for rr in range(found + 1, wsd2.max_row + 1):
    vals = [wsd2.cell(rr, cc).value for cc in range(1, len(headers) + 1)]
    if any(v is not None for v in vals):
        saved_rows.append(dict(zip(headers, vals)))
if len(saved_rows) != len(raw_rows):
    raise AssertionError(f'Saved row count {len(saved_rows)} != raw row count {len(raw_rows)}')
if [str(r['OrderID']) for r in saved_rows] != sorted(str(r['OrderID']) for r in saved_rows):
    raise AssertionError('Saved rows not sorted by OrderID')
if {r['Status'] for r in saved_rows} != {status}:
    raise AssertionError('Saved statuses are not all requested status')
for r in saved_rows:
    cur = r['Currency']
    expected_rate = rates[cur]
    if Decimal(str(r['RateToTWD'])) != expected_rate:
        raise AssertionError(f'Rate mismatch for {r["OrderID"]}')
    expected_twd = money(Decimal(str(r['NetAmount'])) * expected_rate)
    if money(r['NetAmountTWD']) != expected_twd:
        raise AssertionError(f'TWD mismatch for {r["OrderID"]}')
blank_contacts = [str(r['OrderID']) for r in saved_rows if r['CustomerContact'] in (None, '')]
print('blank CustomerContact OrderIDs:', blank_contacts)
print('CHECK PASSED')
