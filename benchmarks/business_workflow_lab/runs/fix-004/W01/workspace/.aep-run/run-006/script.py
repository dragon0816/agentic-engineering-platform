from pathlib import Path
from openpyxl import load_workbook
import csv, datetime, json, math
from collections import defaultdict

ws = Path.cwd()
out = ws/'output'/'september_open_orders.xlsx'
raw = ws/'output'/'downloads'/'september_open_orders_raw.csv'
fxp = ws/'input'/'fx_rates.csv'

print('workspace=', ws)
print('input files=', sorted(p.name for p in (ws/'input').glob('*')))
print('output files=', sorted(str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()))

# Read FX rates
fx = {}
with fxp.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('fx headers=', reader.fieldnames)
    for row in reader:
        # accept common column names
        cur = row.get('Currency') or row.get('currency') or row.get('From') or row.get('Code')
        rate_val = row.get('TWD') or row.get('Rate') or row.get('FXRate') or row.get('rate_to_twd') or row.get('RateToTWD')
        if cur and rate_val:
            fx[cur.strip()] = float(rate_val)
print('fx=', fx)

# Read raw export
raw_rows = []
if raw.exists():
    with raw.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        raw_rows = list(reader)
        print('raw headers=', reader.fieldnames)
        print('raw row count=', len(raw_rows))
        print('raw first row=', raw_rows[0] if raw_rows else None)

# Read output workbook
wb = load_workbook(out, data_only=False)
print('sheets=', wb.sheetnames)
det = wb['Details']; summ = wb['Summary']
headers = [det.cell(4,c).value for c in range(1, det.max_column+1)]
print('details headers=', headers)
rows=[]
for r in range(5, det.max_row+1):
    vals=[det.cell(r,c).value for c in range(1,len(headers)+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('details row count=', len(rows))
print('order ids=', [str(r['OrderID']) for r in rows])

# Template actual names
required = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
missing=[h for h in required if h not in headers]
print('missing template headers=', missing)
assert not missing

order_ids=[str(r['OrderID']) for r in rows]
assert order_ids == sorted(order_ids), 'OrderID not sorted'
assert len(set(order_ids)) == len(order_ids), 'duplicate OrderID'
assert all(r['Status']=='Open' for r in rows), 'non-Open row present'

def date_ok(v):
    if isinstance(v, datetime.datetime):
        d=v.date()
    elif isinstance(v, datetime.date):
        d=v
    else:
        d=datetime.date.fromisoformat(str(v)[:10])
    return datetime.date(2026,9,1) <= d <= datetime.date(2026,9,30)
assert all(date_ok(r['OrderDate']) for r in rows), 'date outside September 2026'

# Verify all raw rows represented and detail pages fields present
if raw_rows:
    raw_ids = sorted(str(r.get('OrderID','')) for r in raw_rows)
    assert raw_ids == order_ids, f'raw IDs mismatch raw={raw_ids} details={order_ids}'
assert all(r['CustomerContact'] not in (None,'') for r in rows), 'blank CustomerContact'
assert all(r['PaymentTerms'] not in (None,'') for r in rows), 'blank PaymentTerms'

# Verify FX conversion from attached FX rates
calc_orig=defaultdict(float); calc_twd=defaultdict(float)
for r in rows:
    cur=str(r['Currency'])
    net=float(r['NetAmount'])
    twd=float(r['NetAmountTWD'])
    assert cur in fx, f'missing FX for {cur}'
    expected=round(net*fx[cur], 2)
    assert abs(twd-expected) <= 0.011, f'FX mismatch {r["OrderID"]}: {net}*{fx[cur]}={expected}, got {twd}'
    calc_orig[cur]+=net
    calc_twd[cur]+=twd
print('computed original by currency=', {k: round(v,2) for k,v in sorted(calc_orig.items())})
print('computed TWD by currency=', {k: round(v,2) for k,v in sorted(calc_twd.items())})
print('computed total TWD=', round(sum(calc_twd.values()),2))

# Inspect and verify Summary cells by reading non-empty rows
summary=[]
for rr in range(1, summ.max_row+1):
    vals=[summ.cell(rr,cc).value for cc in range(1, min(summ.max_column,6)+1)]
    if any(v is not None for v in vals):
        summary.append(vals)
print('summary rows=', summary)
# Check formulas/values include the total over Details NetAmountTWD range in B2 or elsewhere
formula_text=' '.join(str(v) for row in summary for v in row if v is not None)
assert 'Details!I5:I18' in formula_text or round(sum(calc_twd.values()),2) in [round(float(v),2) for row in summary for v in row if isinstance(v,(int,float))], 'summary does not contain expected total/formula'
# Check each currency appears in Summary with corresponding totals approximately somewhere in row
for cur in sorted(calc_orig):
    rows_for_cur=[row for row in summary if cur in [str(v) for v in row]]
    assert rows_for_cur, f'{cur} missing in Summary'
    nums=[]
    for row in rows_for_cur:
        for v in row:
            if isinstance(v,(int,float)):
                nums.append(round(float(v),2))
    assert round(calc_orig[cur],2) in nums, f'{cur} original total missing in Summary'
    assert round(calc_twd[cur],2) in nums, f'{cur} TWD total missing in Summary'

print('CHECK_RESULT=PASS')