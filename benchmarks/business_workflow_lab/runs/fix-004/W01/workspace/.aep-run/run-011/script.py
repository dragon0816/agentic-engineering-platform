import csv
import json
import os
import shutil
from collections import defaultdict
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
date_from = params['date_from']
date_to = params['date_to']
status = params['status']
output_filename = params.get('output_filename') or 'sales_orders_report.xlsx'
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
download_dir = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)
download_dir.mkdir(exist_ok=True)

print('workspace:', workspace)
print('input files:', sorted(p.name for p in input_dir.iterdir()))
print('download files:', sorted(p.name for p in download_dir.iterdir()))

raw_download = download_dir / 'sales_orders_raw.csv'
if not raw_download.exists():
    raise FileNotFoundError(f'Missing downloaded CSV: {raw_download}')
raw_named = download_dir / (Path(output_filename).stem + '_raw.csv')
if raw_named.resolve() != raw_download.resolve():
    shutil.copyfile(raw_download, raw_named)

with raw_download.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('raw headers:', reader.fieldnames)
    rows = list(reader)
print('raw row count:', len(rows))
print('raw sample:', rows[:3])

start_date = datetime.strptime(date_from, '%Y-%m-%d').date()
end_date = datetime.strptime(date_to, '%Y-%m-%d').date()
filtered = []
for row in rows:
    order_date = datetime.strptime(row['OrderDate'], '%Y-%m-%d').date()
    if start_date <= order_date <= end_date and row.get('Status') == status:
        filtered.append(row)
filtered.sort(key=lambda r: str(r.get('OrderID', '')))
if len(filtered) != len(rows):
    print('filtered row count after safety check:', len(filtered))

fx_path = input_dir / 'fx_rates.csv'
if not fx_path.exists():
    raise FileNotFoundError(f'Missing FX rates: {fx_path}')
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('fx headers:', reader.fieldnames)
    for r in reader:
        keys = list(r.keys())
        currency_key = next((k for k in keys if k.lower() in {'currency', 'ccy'}), keys[0])
        rate_key = next((k for k in keys if 'twd' in k.lower() or 'rate' in k.lower()), keys[1] if len(keys) > 1 else keys[0])
        if r.get(currency_key):
            fx[r[currency_key].strip()] = Decimal(str(r[rate_key]).strip())
print('fx rates:', {k: str(v) for k, v in fx.items()})

out_path = output_dir / output_filename
template_path = input_dir / 'sales_template.xlsx'
base_path = out_path if out_path.exists() else template_path
if not base_path.exists():
    raise FileNotFoundError('Missing both existing output workbook and input/sales_template.xlsx')

# Reuse existing output as a local cache for detail fields that came from earlier browser-operated detail pages.
detail_cache = {}
if out_path.exists():
    old_wb = load_workbook(out_path, data_only=False)
    if 'Details' in old_wb.sheetnames:
        old_ws = old_wb['Details']
        old_header_row = None
        for r in range(1, min(old_ws.max_row, 20) + 1):
            vals = [old_ws.cell(r, c).value for c in range(1, old_ws.max_column + 1)]
            if 'OrderID' in vals:
                old_header_row = r
                break
        if old_header_row:
            old_headers = {str(old_ws.cell(old_header_row, c).value): c for c in range(1, old_ws.max_column + 1) if old_ws.cell(old_header_row, c).value}
            for r in range(old_header_row + 1, old_ws.max_row + 1):
                oid = old_ws.cell(r, old_headers.get('OrderID', 1)).value
                if oid is None:
                    continue
                detail_cache[str(oid)] = {
                    'CustomerContact': old_ws.cell(r, old_headers['CustomerContact']).value if 'CustomerContact' in old_headers else '',
                    'PaymentTerms': old_ws.cell(r, old_headers['PaymentTerms']).value if 'PaymentTerms' in old_headers else '',
                }
        old_wb.close()
print('detail cache entries:', len(detail_cache))

wb = load_workbook(base_path)
if 'Details' not in wb.sheetnames:
    raise ValueError('Workbook has no Details sheet')
ws = wb['Details']

header_row = None
for r in range(1, min(ws.max_row, 30) + 1):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    if 'OrderID' in vals:
        header_row = r
        break
if header_row is None:
    raise ValueError('Could not find Details header row containing OrderID')
headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
header_map = {str(v): i + 1 for i, v in enumerate(headers) if v is not None}
print('Details header row:', header_row)
print('Details headers:', [h for h in headers if h is not None])

data_start = header_row + 1
for r in range(data_start, max(ws.max_row, data_start + len(filtered) + 5) + 1):
    for c in range(1, ws.max_column + 1):
        ws.cell(r, c).value = None

def money(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def is_twd_amount_header(h):
    s = str(h or '').lower().replace(' ', '')
    return ('twd' in s and 'amount' in s) or s in {'netamounttwd', 'amounttwd'}

for idx, row in enumerate(filtered, start=data_start):
    oid = str(row.get('OrderID', ''))
    net = money(row.get('NetAmount', '0'))
    curr = row.get('Currency', '')
    if curr not in fx:
        raise ValueError(f'Missing FX rate for {curr}')
    twd = money(net * fx[curr])
    cache = detail_cache.get(oid, {})
    for h, c in header_map.items():
        if h == 'OrderID':
            ws.cell(idx, c).value = oid
            ws.cell(idx, c).number_format = '@'
        elif h == 'NetAmount':
            ws.cell(idx, c).value = float(net)
            ws.cell(idx, c).number_format = '#,##0.00'
        elif is_twd_amount_header(h):
            ws.cell(idx, c).value = float(twd)
            ws.cell(idx, c).number_format = '#,##0.00'
        elif h == 'CustomerContact':
            ws.cell(idx, c).value = cache.get('CustomerContact') or ''
        elif h == 'PaymentTerms':
            ws.cell(idx, c).value = cache.get('PaymentTerms') or ''
        elif h in row:
            ws.cell(idx, c).value = row[h]
        else:
            ws.cell(idx, c).value = cache.get(h, '')

# Update Summary if present.
if 'Summary' in wb.sheetnames:
    sum_ws = wb['Summary']
    twd_col = None
    for h, c in header_map.items():
        if is_twd_amount_header(h):
            twd_col = c
            break
    if twd_col:
        col_letter = get_column_letter(twd_col)
        end_row = data_start + len(filtered) - 1
        sum_ws['B2'] = f'=SUM(Details!{col_letter}{data_start}:{col_letter}{end_row})'
    totals = defaultdict(lambda: {'orig': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
    for row in filtered:
        curr = row['Currency']
        net = money(row['NetAmount'])
        totals[curr]['orig'] += net
        totals[curr]['twd'] += money(net * fx[curr])
        totals[curr]['count'] += 1
    summary_header_row = None
    for r in range(1, min(sum_ws.max_row, 30) + 1):
        vals = [sum_ws.cell(r, c).value for c in range(1, min(sum_ws.max_column, 10) + 1)]
        if 'Currency' in vals:
            summary_header_row = r
            break
    if summary_header_row is None:
        summary_header_row = 4
        sum_ws.cell(summary_header_row, 1).value = 'Currency'
        sum_ws.cell(summary_header_row, 2).value = 'OriginalAmount'
        sum_ws.cell(summary_header_row, 3).value = 'TWD'
        sum_ws.cell(summary_header_row, 4).value = 'Count'
    sum_headers = {str(sum_ws.cell(summary_header_row, c).value): c for c in range(1, sum_ws.max_column + 1) if sum_ws.cell(summary_header_row, c).value}
    for r in range(summary_header_row + 1, summary_header_row + 20):
        for c in range(1, max(sum_ws.max_column, 4) + 1):
            sum_ws.cell(r, c).value = None
    for i, curr in enumerate(sorted(totals), start=summary_header_row + 1):
        sum_ws.cell(i, sum_headers.get('Currency', 1)).value = curr
        sum_ws.cell(i, sum_headers.get('OriginalAmount', 2)).value = float(money(totals[curr]['orig']))
        sum_ws.cell(i, sum_headers.get('TWD', 3)).value = float(money(totals[curr]['twd']))
        sum_ws.cell(i, sum_headers.get('Count', 4)).value = totals[curr]['count']

wb.save(out_path)
wb.close()
print('saved workbook:', out_path)
print('saved raw copy:', raw_named)

# Reopen and validate saved outputs.
check_wb = load_workbook(out_path, data_only=False)
check_ws = check_wb['Details']
check_headers = {str(check_ws.cell(header_row, c).value): c for c in range(1, check_ws.max_column + 1) if check_ws.cell(header_row, c).value}
actual_rows = []
for r in range(data_start, data_start + len(filtered)):
    rec = {h: check_ws.cell(r, c).value for h, c in check_headers.items()}
    actual_rows.append(rec)
if len(actual_rows) != len(filtered):
    raise AssertionError('Row count mismatch after save')
if [str(r['OrderID']) for r in actual_rows] != sorted(str(r['OrderID']) for r in actual_rows):
    raise AssertionError('OrderID is not sorted')
for rec in actual_rows:
    if not isinstance(rec.get('OrderID'), str):
        raise AssertionError('OrderID was not saved as text')
    if rec.get('Status') != status:
        raise AssertionError('Unexpected Status in saved workbook')
    od = datetime.strptime(str(rec.get('OrderDate')), '%Y-%m-%d').date()
    if not (start_date <= od <= end_date):
        raise AssertionError('OrderDate outside requested range')
    if not isinstance(rec.get('NetAmount'), (int, float)):
        raise AssertionError('NetAmount was not saved as numeric')
    twd_keys = [h for h in check_headers if is_twd_amount_header(h)]
    if twd_keys and not isinstance(rec.get(twd_keys[0]), (int, float)):
        raise AssertionError('TWD amount was not saved as numeric')
if 'PaymentTerms' in check_headers:
    missing_terms = [str(r['OrderID']) for r in actual_rows if not r.get('PaymentTerms')]
    if missing_terms:
        raise AssertionError('Missing PaymentTerms for: ' + ', '.join(missing_terms))
missing_contacts = []
if 'CustomerContact' in check_headers:
    missing_contacts = [str(r['OrderID']) for r in actual_rows if not r.get('CustomerContact')]
if 'Summary' in check_wb.sheetnames:
    b2 = check_wb['Summary']['B2'].value
    if not (isinstance(b2, str) and b2.startswith('=SUM(Details!')):
        raise AssertionError('Summary!B2 formula was not updated')
check_wb.close()

report = {
    'workbook': str(out_path),
    'raw_csv': str(raw_named),
    'download_csv': str(raw_download),
    'rows': len(actual_rows),
    'missing_customer_contact_order_ids': missing_contacts,
}
print('VALIDATION:', json.dumps(report, ensure_ascii=False, indent=2))
