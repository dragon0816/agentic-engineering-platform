import os, csv, json
from pathlib import Path
ws=Path.cwd()
print('workspace', ws)
print('input files', [p.name for p in (ws/'input').iterdir()])
for p in [ws/'output'/'downloads'/'september_open_orders_raw.csv', ws/'output'/'september_open_orders.xlsx']:
    print('exists', p, p.exists(), 'size', p.stat().st_size if p.exists() else None)
if (ws/'output'/'downloads'/'september_open_orders_raw.csv').exists():
    with open(ws/'output'/'downloads'/'september_open_orders_raw.csv', newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f); rows=list(r)
    print('raw header', rows[0] if rows else None)
    print('raw sample', rows[1:4])
# Inspect workbook sheets/header/sample only; no changes.
from openpyxl import load_workbook
out=ws/'output'/'september_open_orders.xlsx'
if out.exists():
    wb=load_workbook(out, data_only=False)
    print('sheets', wb.sheetnames)
    for s in wb.sheetnames:
        wsht=wb[s]
        print('sheet', s, 'max', wsht.max_row, wsht.max_column)
        for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
            print(row)
        print('---')