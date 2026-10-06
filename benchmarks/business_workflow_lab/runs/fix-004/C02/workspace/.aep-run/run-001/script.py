from pathlib import Path
import csv, os
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv preview:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(6), r):
        print(i,row)
print('\nmanual preview:')
print((ws/'input'/'quote_cli_manual.md').read_text(encoding='utf-8')[:2000])
print('\ntemplate sheets/dimensions/headers:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for sh in wb.worksheets:
    print(sh.title, sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(8,sh.max_row), values_only=True):
        print(row)
    print('---')