from pathlib import Path
import subprocess, json, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta, date
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
OUT_XLSX = OUTPUT / 'quote_template.xlsx'
OUT_MD = OUTPUT / 'findings.md'
CLI = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
TW = timezone(timedelta(hours=8))
GENERATED_AT = datetime.now(TW).replace(microsecond=0).isoformat()
Q = Decimal('0.01')

def money(v):
    return Decimal(str(v)).quantize(Q, rounding=ROUND_HALF_UP)

def parse_rates(path):
    rates = {}
    with open(path, encoding='utf-8-sig', newline='') as f:
        reader = csv.reader(f)
        header = next(reader)
        for row in reader:
            if len(row) >= 2 and row[0]:
                rates[row[0]] = Decimal(row[1])
    return rates

def run_cli(args):
    res = subprocess.run([CLI] + args, capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f"CLI failed rc={res.returncode} args={args}\nstderr={res.stderr}\nstdout={res.stdout}")
    return json.loads(res.stdout)

def clear_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def clone_row_style(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        s = ws.cell(source_row, col)
        t = ws.cell(target_row, col)
        if s.has_style:
            t._style = copy(s._style)
        if s.number_format:
            t.number_format = s.number_format
        if s.alignment:
            t.alignment = copy(s.alignment)
        if s.font:
            t.font = copy(s.font)
        if s.fill:
            t.fill = copy(s.fill)
        if s.border:
            t.border = copy(s.border)

rates = parse_rates(INPUT / 'fx_rates.csv')
listed = run_cli(['quotes','list','--from',FROM_DATE,'--to',TO_DATE,'--format','json'])
headers = listed['items']
# Stable order by QuoteDate then QuoteID, preserving all quotes returned for September.
headers.sort(key=lambda h: (h.get('QuoteDate',''), h.get('QuoteID','')))
quote_docs = []
for h in headers:
    doc = run_cli(['quotes','get',h['QuoteID'],'--format','json'])
    quote_docs.append(doc)

# Validate and prepare rows.
quote_rows = []
line_rows = []
findings = []
expired_unaccepted = []
summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
all_ok_line = True
all_ok_net = True

for doc in quote_docs:
    h = doc['header']
    qid = str(h['QuoteID'])
    ccy = h['Currency']
    if ccy not in rates:
        raise RuntimeError(f'Missing FX rate for {ccy}')
    subtotal = money(h['Subtotal'])
    disc = Decimal(str(h['DiscountPercent']))
    net = money(h['NetAmount'])
    line_sum = Decimal('0.00')
    line_issues = []
    for li in doc['line_items']:
        qty = Decimal(str(li['Quantity']))
        unit = Decimal(str(li['UnitPrice']))
        line_amount = money(li['LineAmount'])
        expected_line = money(qty * unit)
        if line_amount != expected_line:
            all_ok_line = False
            line_issues.append(f"{li['ProductID']}: LineAmount {line_amount} != Quantity×UnitPrice {expected_line}")
        line_sum += line_amount
        line_rows.append([
            qid,
            str(li.get('ProductID','')),
            li.get('ProductName',''),
            int(li['Quantity']) if Decimal(str(li['Quantity'])) == Decimal(str(int(Decimal(str(li['Quantity']))))) else float(li['Quantity']),
            float(money(li['UnitPrice'])),
            float(line_amount),
        ])
    line_sum = money(line_sum)
    expected_net = money(subtotal * (Decimal('1') - disc / Decimal('100')))
    detail_ok = (line_sum == subtotal)
    net_ok = (expected_net == net)
    if not detail_ok:
        all_ok_line = False
    if not net_ok:
        all_ok_net = False
    if line_issues or not detail_ok or not net_ok:
        issue_parts = []
        if line_issues:
            issue_parts.extend(line_issues)
        if not detail_ok:
            issue_parts.append(f"Subtotal {subtotal} != sum(LineAmount) {line_sum}")
        if not net_ok:
            issue_parts.append(f"NetAmount {net} != Subtotal×(1−DiscountPercent/100) {expected_net}")
        findings.append({'QuoteID': qid, 'issues': issue_parts})
    valid_until = date.fromisoformat(h['ValidUntil'])
    if valid_until < BASE_DATE and h['Status'] != 'Accepted':
        expired_unaccepted.append(h)
    net_twd = money(net * rates[ccy])
    quote_rows.append([
        qid,
        str(h.get('CustomerID','')),
        str(h.get('ProjectID','')) if h.get('ProjectID','') is not None else '',
        h.get('QuoteDate',''),
        h.get('ValidUntil',''),
        h.get('Status',''),
        ccy,
        float(subtotal),
        float(disc),
        float(net),
        float(net_twd),
    ])
    key = (h.get('Status',''), ccy)
    summary[key]['amount'] += net
    summary[key]['twd'] += net_twd
    summary[key]['count'] += 1

# Copy/fill workbook.
shutil.copyfile(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)
report = wb['Report']
quotes_ws = wb['Quotes']
lines_ws = wb['Lines']
summary_ws = wb['Summary']

# Report cells
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = 'QuoteDate 2026-09-01 through 2026-09-30; all September quotes retained; expiry reviewed against 2026-09-30; FX from fx_rates.csv.'
report['B5'] = GENERATED_AT
report['B7'] = 'quote_cli mock; fx_rates.csv'
report['D7'] = GENERATED_AT
report['B9'] = f"Fetched {len(quote_rows)} quotes and {len(line_rows)} line items. Expired before {BASE_DATE.isoformat()} and not Accepted: {len(expired_unaccepted)}. Detail/net issues: {len(findings)}."

# Clear and append data
clear_data(quotes_ws, 5)
clear_data(lines_ws, 5)
clear_data(summary_ws, 5)

for i, row in enumerate(quote_rows, start=5):
    clone_row_style(quotes_ws, 4, i, 11)
    for j, val in enumerate(row, start=1):
        c = quotes_ws.cell(i, j, val)
        if j in (1,2,3,6,7):
            c.number_format = '@'
        elif j in (4,5):
            c.number_format = 'yyyy-mm-dd'
        elif j in (8,10,11):
            c.number_format = '0.00'
        elif j == 9:
            c.number_format = '0.00'

for i, row in enumerate(line_rows, start=5):
    clone_row_style(lines_ws, 4, i, 6)
    for j, val in enumerate(row, start=1):
        c = lines_ws.cell(i, j, val)
        if j in (1,2,3):
            c.number_format = '@'
        elif j in (5,6):
            c.number_format = '0.00'

# Summary by Status and Currency so original-currency amounts are not mixed.
summary_rows = []
for (status, ccy), vals in sorted(summary.items(), key=lambda kv: (kv[0][0], kv[0][1])):
    summary_rows.append(['Status', status, ccy, float(money(vals['amount'])), float(money(vals['twd'])), vals['count']])
for i, row in enumerate(summary_rows, start=5):
    clone_row_style(summary_ws, 4, i, 6)
    for j, val in enumerate(row, start=1):
        c = summary_ws.cell(i, j, val)
        if j in (1,2,3):
            c.number_format = '@'
        elif j in (4,5):
            c.number_format = '0.00'
# Extend formula range
last_quote_row = 4 + len(quote_rows) if quote_rows else 5
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})' if quote_rows else '=SUM(Quotes!K5:K5)'
wb.save(OUT_XLSX)

# Findings markdown
md = []
md.append('# Quote findings — September 2026')
md.append('')
md.append(f'- Query period: {FROM_DATE} through {TO_DATE} inclusive')
md.append(f'- Expiry review baseline: {BASE_DATE.isoformat()}')
md.append(f'- Source: quote_cli mock; fx_rates.csv')
md.append(f'- Generated at: {GENERATED_AT}')
md.append(f'- Quotes retained in workbook: {len(quote_rows)} (no quote removed because of expiry)')
md.append(f'- Line items retained in workbook: {len(line_rows)}')
md.append('')
md.append('## Amount consistency checks')
md.append('')
if not findings:
    md.append('- All quotes passed: each quote Subtotal equals the sum of LineAmount values, each LineAmount equals Quantity × UnitPrice, and each NetAmount equals Subtotal × (1 − DiscountPercent / 100) rounded ROUND_HALF_UP to 2 decimals.')
else:
    md.append(f'- Issues found in {len(findings)} quote(s):')
    for f in findings:
        md.append(f"  - {f['QuoteID']}: " + '; '.join(f['issues']))
md.append('')
md.append('## ValidUntil before 2026-09-30 and Status not Accepted')
md.append('')
if not expired_unaccepted:
    md.append('- None.')
else:
    md.append('| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |')
    md.append('|---|---|---|---|---|---|---:|---:|')
    for h in sorted(expired_unaccepted, key=lambda x: (x['ValidUntil'], x['QuoteID'])):
        md.append(f"| {h['QuoteID']} | {h.get('CustomerID','')} | {h.get('ProjectID','')} | {h.get('QuoteDate','')} | {h.get('ValidUntil','')} | {h.get('Status','')} | {h.get('Currency','')} | {money(h.get('NetAmount','0'))} |")
md.append('')
md.append('## Summary by Status and Currency')
md.append('')
md.append('| Status | Currency | NetAmount | NetAmountTWD | Count |')
md.append('|---|---|---:|---:|---:|')
for (status, ccy), vals in sorted(summary.items(), key=lambda kv: (kv[0][0], kv[0][1])):
    md.append(f"| {status} | {ccy} | {money(vals['amount'])} | {money(vals['twd'])} | {vals['count']} |")
OUT_MD.write_text('\n'.join(md) + '\n', encoding='utf-8')

# Reopen and verify saved outputs against stated requirements.
wb2 = load_workbook(OUT_XLSX, data_only=False)
qws = wb2['Quotes']; lws = wb2['Lines']; sws = wb2['Summary']
q_headers = [qws.cell(4,c).value for c in range(1,12)]
l_headers = [lws.cell(4,c).value for c in range(1,7)]
s_headers = [sws.cell(4,c).value for c in range(1,7)]
assert q_headers == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD'], q_headers
assert l_headers == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount'], l_headers
assert s_headers == ['Group','SubGroup','Currency','Amount','AmountTWD','Count'], s_headers
saved_q_rows = list(qws.iter_rows(min_row=5, max_row=4+len(quote_rows), max_col=11, values_only=True))
saved_l_rows = list(lws.iter_rows(min_row=5, max_row=4+len(line_rows), max_col=6, values_only=True))
assert len(saved_q_rows) == len(headers) == 12, (len(saved_q_rows), len(headers))
assert len(saved_l_rows) == len(line_rows) and len(saved_l_rows) > 0
assert all(isinstance(r[0], str) and r[0].startswith('Q') for r in saved_q_rows)
assert all(isinstance(r[1], str) and (r[2] is None or isinstance(r[2], str)) for r in saved_q_rows)
assert all(isinstance(r[7], (int,float)) and isinstance(r[9], (int,float)) and isinstance(r[10], (int,float)) for r in saved_q_rows)
assert all(isinstance(r[0], str) and isinstance(r[1], str) and isinstance(r[3], (int,float)) and isinstance(r[4], (int,float)) and isinstance(r[5], (int,float)) for r in saved_l_rows)
assert sws['B2'].value == '=SUM(Quotes!K5:K16)', sws['B2'].value
md_text = OUT_MD.read_text(encoding='utf-8')
assert 'Q0010' in md_text and 'Q0014' in md_text and 'Quote findings' in md_text
# Validate expected expired unaccepted list from saved quote rows.
saved_expired = sorted([r[0] for r in saved_q_rows if date.fromisoformat(r[4]) < BASE_DATE and r[5] != 'Accepted'])
assert saved_expired == sorted([h['QuoteID'] for h in expired_unaccepted])
print('Wrote:', OUT_XLSX)
print('Wrote:', OUT_MD)
print('Checked workbook sheets/headers, quote rows:', len(saved_q_rows), 'line rows:', len(saved_l_rows), 'summary rows:', len(summary_rows))
print('Checked numeric amount cells and text ID cells from saved workbook.')
print('Checked Summary!B2 formula:', sws['B2'].value)
print('Expired before 2026-09-30 and not Accepted:', ', '.join(saved_expired) if saved_expired else 'none')
print('Amount consistency issue count:', len(findings))
print('Output files:', [p.name for p in OUTPUT.iterdir()])