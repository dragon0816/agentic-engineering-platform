from pathlib import Path
from openpyxl import load_workbook
import subprocess, json, sys, os
ws=Path.cwd()
print('Instructions rows:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
cli=r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
print('\nCLI list Sep 2026:')
res=subprocess.run([cli,'quotes','list','--from','2026-09-01','--to','2026-09-30','--format','json'], capture_output=True, text=True, shell=False)
print('rc',res.returncode)
print('stderr',res.stderr[:500])
print('stdout',res.stdout[:2000])
if res.returncode==0:
    data=json.loads(res.stdout)
    print('count', data.get('count'))
    if data.get('items'):
        qid=data['items'][0]['QuoteID']
        print('\nget first', qid)
        g=subprocess.run([cli,'quotes','get',qid,'--format','json'], capture_output=True, text=True)
        print('rc',g.returncode,'stderr',g.stderr[:200])
        print(g.stdout[:2000])