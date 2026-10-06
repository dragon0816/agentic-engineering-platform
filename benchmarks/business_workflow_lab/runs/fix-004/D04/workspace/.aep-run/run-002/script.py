from pathlib import Path
import csv, os
from openpyxl import load_workbook
ws=Path.cwd()
print('output files:', [p.name for p in (ws/'output').iterdir()])
p=ws/'output'/'expense_E006_2026-01.csv'
print('csv exists', p.exists(), 'size', p.stat().st_size if p.exists() else None)
if p.exists():
    with p.open(newline='', encoding='utf-8-sig') as f:
        rdr=csv.reader(f)
        for i,row in zip(range(10), rdr):
            print(i,row)
# print all instructions rows
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in sh.iter_rows(values_only=True):
    print(r)