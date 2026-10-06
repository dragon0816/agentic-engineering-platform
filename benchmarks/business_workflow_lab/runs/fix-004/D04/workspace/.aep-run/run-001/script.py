import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', sorted([p.name for p in (ws/'input').iterdir()]))
for p in sorted((ws/'input').iterdir()):
    print('\nFILE', p.name, 'size', p.stat().st_size)
    if p.suffix.lower()=='.csv':
        with p.open(newline='', encoding='utf-8-sig') as f:
            rdr=csv.reader(f)
            for i,row in zip(range(5), rdr):
                print('csv row', i, row)
    elif p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print('sheets', wb.sheetnames)
        for s in wb.sheetnames:
            sh=wb[s]
            print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
            for r in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
                print(r)
    elif p.suffix.lower()=='.md':
        txt=p.read_text(encoding='utf-8')
        print(txt[:2000])