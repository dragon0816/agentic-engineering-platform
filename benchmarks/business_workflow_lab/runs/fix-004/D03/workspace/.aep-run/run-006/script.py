from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime
from collections import defaultdict
from openpyxl import load_workbook

base=Path('.')
output_xlsx=base/'output'/'expense_report.xlsx'
export_csv=base/'output'/'expense_export.csv'
fx_csv=base/'input'/'fx_rates.csv'
assert output_xlsx.exists(), 'missing output workbook'

def blank_eq(a,b):
    return (a in (None,'') and b in (None,'')) or a==b

fx={}
with open(fx_csv, newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        fx[row['Currency']]=Decimal(row['RateToTWD'])
source=[]
with open(export_csv, newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        if row['Status']=='Approved' and '2026-09-01' <= row['ExpenseDate'] <= '2026-09-30':
            amt=Decimal(row['Amount'])
            row['_AmountTWD']=(amt*fx[row['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
            row['_NeedsReview']='true' if row['_AmountTWD'] > Decimal('5000') else 'false'
            source.append(row)
source.sort(key=lambda r:(r['ExpenseDate'], r['ExpenseID']))

wb=load_workbook(output_xlsx, data_only=False)
assert wb.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
ws=wb['Details']
headers=[ws.cell(4,c).value for c in range(1,15)]
expected_headers=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert headers==expected_headers, headers
last=ws.max_row
while last>=5 and all(ws.cell(last,c).value is None for c in range(1,15)):
    last-=1
detail_count=max(0,last-4)
assert detail_count==len(source)==15

for i, src in enumerate(source, start=5):
    vals=[ws.cell(i,c).value for c in range(1,15)]
    assert vals[0]==src['ExpenseID']
    assert vals[1]==src['EmployeeID']
    assert vals[2]==src['EmployeeName']
    assert vals[3]==src['ProjectID']
    d=vals[4].date() if hasattr(vals[4], 'date') else vals[4]
    assert d == datetime.strptime(src['ExpenseDate'], '%Y-%m-%d').date()
    assert vals[5]==src['Category']
    assert vals[6]==src['Currency']
    assert Decimal(str(vals[7])).quantize(Decimal('0.01')) == Decimal(src['Amount']).quantize(Decimal('0.01'))
    assert Decimal(str(vals[8])).quantize(Decimal('0.01')) == src['_AmountTWD']
    assert vals[9]=='Approved'
    assert blank_eq(vals[10], src['ReceiptID'])
    assert blank_eq(vals[11], src['Description'])
    assert str(vals[12]).lower() in ('true','false')
    assert str(vals[13]).lower()==src['_NeedsReview']
    assert isinstance(vals[0], str) and isinstance(vals[1], str) and isinstance(vals[3], str) and isinstance(vals[6], str)
    assert isinstance(vals[7], (int,float)) and isinstance(vals[8], (int,float))

summary_expected=defaultdict(lambda:{'Amount':Decimal('0.00'),'AmountTWD':Decimal('0.00'),'Count':0})
for src in source:
    key=(f"{src['EmployeeID']} {src['EmployeeName']}", src['Category'], src['Currency'])
    summary_expected[key]['Amount'] += Decimal(src['Amount'])
    summary_expected[key]['AmountTWD'] += src['_AmountTWD']
    summary_expected[key]['Count'] += 1
wss=wb['Summary']
assert wss['B2'].value == '=SUM(Details!I5:I19)', wss['B2'].value
summary_rows=[]
for r in range(5,wss.max_row+1):
    if wss.cell(r,1).value is not None:
        key=(wss.cell(r,1).value,wss.cell(r,2).value,wss.cell(r,3).value)
        assert key in summary_expected, key
        summary_rows.append(key)
        exp=summary_expected[key]
        assert Decimal(str(wss.cell(r,4).value)).quantize(Decimal('0.01')) == exp['Amount'].quantize(Decimal('0.01'))
        assert Decimal(str(wss.cell(r,5).value)).quantize(Decimal('0.01')) == exp['AmountTWD'].quantize(Decimal('0.01'))
        assert wss.cell(r,6).value == exp['Count']
assert len(summary_rows)==len(summary_expected)==14
assert set(summary_rows)==set(summary_expected.keys())

expected_ex=[src for src in source if str(src.get('MissingReceipt','')).lower()=='true']
wse=wb['Exceptions']
actual_ex=[]
for r in range(5,wse.max_row+1):
    if wse.cell(r,1).value is not None:
        actual_ex.append(wse.cell(r,1).value)
assert len(actual_ex)==len(expected_ex)==1
assert set(actual_ex)=={r['ExpenseID'] for r in expected_ex}

wr=wb['Report']
b3=wr['B3'].value.date() if hasattr(wr['B3'].value,'date') else wr['B3'].value
d3=wr['D3'].value.date() if hasattr(wr['D3'].value,'date') else wr['D3'].value
assert b3 == date(2026,9,1), b3
assert d3 == date(2026,9,30), d3
assert 'Approved' in wr['B4'].value

print('CHECK OK')
print('Saved output:', output_xlsx)
print('Details rows:', detail_count)
print('Summary rows:', len(summary_rows))
print('Exceptions rows:', len(actual_ex))
print('NeedsReview true:', sum(1 for s in source if s['_NeedsReview']=='true'))
print('NeedsReview false:', sum(1 for s in source if s['_NeedsReview']=='false'))
print('Total AmountTWD:', sum((s['_AmountTWD'] for s in source), Decimal('0.00')).quantize(Decimal('0.01')))