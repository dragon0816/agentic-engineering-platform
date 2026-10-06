from pathlib import Path
import csv, json
from openpyxl import load_workbook
base=Path('.')
print('INPUT files:', [str(p) for p in (base/'input').iterdir()])
print('OUTPUT files:', [str(p) for p in (base/'output').iterdir()])
for csv_path in [base/'input'/'fx_rates.csv', base/'output'/'expense_export.csv']:
    print('\nCSV', csv_path)
    with open(csv_path, newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        rows=[]
        for i,row in zip(range(6), r): rows.append(row)
    for row in rows: print(row)
md=base/'input'/'expense_manual.md'
print('\nManual sample:')
print(md.read_text(encoding='utf-8')[:1200])
wb=load_workbook(base/'input'/'expense_template.xlsx')
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('Sheet', ws.title, 'max', ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,8), values_only=True):
        print(row)
    print('---')