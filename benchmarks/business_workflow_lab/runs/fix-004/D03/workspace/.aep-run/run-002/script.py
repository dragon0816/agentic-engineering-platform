from pathlib import Path
from openpyxl import load_workbook
base=Path('.')
wb=load_workbook(base/'input'/'expense_template.xlsx')
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)