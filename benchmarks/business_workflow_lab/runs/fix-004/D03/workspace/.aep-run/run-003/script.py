from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

base = Path('.')
input_dir = base / 'input'
out_dir = base / 'output'
out_dir.mkdir(exist_ok=True)
export_csv = out_dir / 'expense_export.csv'
fx_csv = input_dir / 'fx_rates.csv'
template = input_dir / 'expense_template.xlsx'
output_xlsx = out_dir / 'expense_report.xlsx'

# Read fx rates
fx = {}
with open(fx_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        if row.get('Currency'):
            fx[row['Currency'].strip()] = Decimal(row['RateToTWD'].strip())

# Read exported expenses
rows = []
with open(export_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        rows.append(row)

# Filter defensively to requested scope, although desktop export was already set that way
filtered = []
for row in rows:
    if row.get('Status') != 'Approved':
        continue
    d = datetime.strptime(row['ExpenseDate'], '%Y-%m-%d').date()
    if not (date(2026, 9, 1) <= d <= date(2026, 9, 30)):
        continue
    cur = row['Currency']
    if cur not in fx:
        raise ValueError(f'Missing FX rate for {cur}')
    amt = Decimal(row['Amount'])
    amount_twd = (amt * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    row['_ExpenseDateObj'] = d
    row['_AmountDec'] = amt
    row['_AmountTWDDec'] = amount_twd
    row['_NeedsReview'] = 'true' if amount_twd > Decimal('5000') else 'false'
    # normalize MissingReceipt as requested true/false text
    mr = str(row.get('MissingReceipt','')).strip().lower()
    row['_MissingReceipt'] = 'true' if mr == 'true' else 'false'
    filtered.append(row)

# Sort by ExpenseDate then ExpenseID for stable output
filtered.sort(key=lambda r: (r['_ExpenseDateObj'], r['ExpenseID']))

wb = load_workbook(template)

# Helpers
def clear_data(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def copy_row_style(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

# Report
ws = wb['Report']
ws['B3'] = date(2026, 9, 1)
ws['B3'].number_format = 'yyyy-mm-dd'
ws['D3'] = date(2026, 9, 30)
ws['D3'].number_format = 'yyyy-mm-dd'
ws['B4'] = 'Status=Approved; ExpenseDate=2026-09-01..2026-09-30; AmountTWD=fixed fx_rates.csv; NeedsReview=true if AmountTWD > 5000'
now_iso = datetime.now().astimezone().isoformat(timespec='seconds')
ws['B5'] = now_iso
ws['B7'] = 'Expense Viewer desktop export; fx_rates.csv'
ws['D7'] = now_iso
ws['B9'] = 'All exported Approved September 2026 rows included; original Currency and Amount retained.'

# Details
wsd = wb['Details']
clear_data(wsd)
for idx, r in enumerate(filtered, start=5):
    if idx > 5:
        copy_row_style(wsd, 5, idx, 14)
    values = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['_ExpenseDateObj'],
        r['Category'], r['Currency'], float(r['_AmountDec']), float(r['_AmountTWDDec']), r['Status'],
        r['ReceiptID'], r['Description'], r['_MissingReceipt'], r['_NeedsReview']
    ]
    for col, val in enumerate(values, start=1):
        cell = wsd.cell(idx, col, val)
        if col in (1,2,3,4,6,7,10,11,12,13,14):
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9):
            cell.number_format = '0.00'

# Summary: by employee and Category, separated by original Currency so original Amount is meaningful
wss = wb['Summary']
clear_data(wss)
last_detail = 4 + len(filtered) if filtered else 5
wss['B2'] = f'=SUM(Details!I5:I{last_detail})' if filtered else '=SUM(Details!I5:I5)'
summary = defaultdict(lambda: {'Amount': Decimal('0'), 'AmountTWD': Decimal('0'), 'Count': 0})
for r in filtered:
    emp = f"{r['EmployeeID']} {r['EmployeeName']}"
    key = (emp, r['Category'], r['Currency'])
    summary[key]['Amount'] += r['_AmountDec']
    summary[key]['AmountTWD'] += r['_AmountTWDDec']
    summary[key]['Count'] += 1

for idx, ((group, sub, cur), vals) in enumerate(sorted(summary.items()), start=5):
    if idx > 5:
        copy_row_style(wss, 5, idx, 6)
    row_vals = [group, sub, cur, float(vals['Amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), float(vals['AmountTWD'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), vals['Count']]
    for col, val in enumerate(row_vals, start=1):
        cell = wss.cell(idx, col, val)
        if col in (1,2,3):
            cell.number_format = '@'
        elif col in (4,5):
            cell.number_format = '0.00'
        elif col == 6:
            cell.number_format = '0'

# Exceptions: missing receipts, if any; still kept in Details
wse = wb['Exceptions']
clear_data(wse)
ex_rows = [r for r in filtered if r['_MissingReceipt'] == 'true']
for idx, r in enumerate(ex_rows, start=5):
    if idx > 5:
        copy_row_style(wse, 5, idx, 4)
    vals = [r['ExpenseID'], r['EmployeeID'], r['_ExpenseDateObj'], 'MissingReceipt=true']
    for col, val in enumerate(vals, start=1):
        cell = wse.cell(idx, col, val)
        if col == 3:
            cell.number_format = 'yyyy-mm-dd'
        else:
            cell.number_format = '@'

wb.save(output_xlsx)
print(f'Wrote {output_xlsx}')
print(f'Detail rows: {len(filtered)}')
print(f'Summary rows: {len(summary)}')
print(f'Exception rows: {len(ex_rows)}')
print('FX rates:', {k: str(v) for k,v in fx.items()})
print('Total AmountTWD:', sum((r['_AmountTWDDec'] for r in filtered), Decimal('0.00')).quantize(Decimal('0.01')))
print('NeedsReview true:', sum(1 for r in filtered if r['_NeedsReview']=='true'))