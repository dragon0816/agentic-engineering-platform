from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)

TEMPLATE = input_dir / 'expense_template.xlsx'
EXPORTS = [output_dir / 'submitted_2026-09.xlsx', output_dir / 'approved_2026-09.xlsx']
FX = input_dir / 'fx_rates.csv'
OUT = output_dir / 'expense_report_2026-09_submitted_approved.xlsx'

print('Inputs to process:')
print(' template:', TEMPLATE.name, TEMPLATE.exists())
print(' fx:', FX.name, FX.exists())
for p in EXPORTS:
    print(' export:', p.name, p.exists(), p.stat().st_size if p.exists() else None)

# Load FX rates (first two columns only; note column may contain commas in synthetic file).
rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    headers = next(reader)
    print('FX headers:', headers)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

# Load exported expenses.
records = []
export_headers = None
for p in EXPORTS:
    wb_e = load_workbook(p, data_only=True)
    sh_e = wb_e.active
    headers = [c.value for c in sh_e[1]]
    print(f'{p.name} headers:', headers)
    if export_headers is None:
        export_headers = headers
    elif headers != export_headers:
        raise RuntimeError(f'Header mismatch in {p.name}')
    for row in sh_e.iter_rows(min_row=2, values_only=True):
        if not any(v is not None for v in row):
            continue
        rec = dict(zip(headers, row))
        records.append(rec)

# Filter defensively to exactly requested period/statuses.
requested_statuses = {'Submitted', 'Approved'}
filtered = []
for rec in records:
    d = rec['ExpenseDate']
    if hasattr(d, 'date'):
        dstr = d.date().isoformat()
    elif hasattr(d, 'isoformat') and not isinstance(d, str):
        dstr = d.isoformat()
    else:
        dstr = str(d)
    rec['ExpenseDate'] = dstr
    if '2026-09-01' <= dstr <= '2026-09-30' and rec['Status'] in requested_statuses:
        filtered.append(rec)
records = sorted(filtered, key=lambda r: (r['ExpenseDate'], r['ExpenseID']))
print('Filtered records:', len(records))
print('First records:', records[:3])

# Helper conversions.
def money_decimal(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def amount_twd(amount, currency):
    if currency not in rates:
        raise RuntimeError(f'Missing FX rate for {currency}')
    return (money_decimal(amount) * rates[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def bool_text(v):
    return 'true' if bool(v) else 'false'

# Load template and fill.
wb = load_workbook(TEMPLATE)

# Report sheet.
report = wb['Report']
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Status in Submitted, Approved; ExpenseDate 2026-09-01 through 2026-09-30 inclusive'
# Use supplied current date/time context, without inventing exact current clock time.
report['B5'] = '2026-10-06T00:00:00+08:00'
report['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D7'] = '2026-10-06T00:00:00+08:00'
missing_count = sum(1 for r in records if bool(r.get('MissingReceipt')))
report['B9'] = f'{len(records)} expenses included; {missing_count} missing receipt item(s) listed in Exceptions and retained in Details.'

# Clear utility function for data area.
def clear_rows(sh, min_row, max_col):
    if sh.max_row >= min_row:
        for row in sh.iter_rows(min_row=min_row, max_row=sh.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

# Details.
details = wb['Details']
clear_rows(details, 5, 14)
for idx, rec in enumerate(records, start=5):
    currency = rec['Currency']
    amt = money_decimal(rec['Amount'])
    twd = amount_twd(amt, currency)
    missing = bool(rec.get('MissingReceipt'))
    values = [
        str(rec['ExpenseID']),
        str(rec['EmployeeID']),
        str(rec['EmployeeName']),
        str(rec['ProjectID']),
        rec['ExpenseDate'],
        str(rec['Category']),
        str(currency),
        float(amt),
        float(twd),
        str(rec['Status']),
        None if rec.get('ReceiptID') in (None, '') else str(rec['ReceiptID']),
        None if rec.get('Description') in (None, '') else str(rec['Description']),
        bool_text(missing),
        bool_text(missing),
    ]
    for col, value in enumerate(values, start=1):
        cell = details.cell(row=idx, column=col, value=value)
        if col in [1,2,3,4,6,7,10,11,12,13,14]:
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in [8,9]:
            cell.number_format = '0.00'

last_detail_row = 4 + len(records) if records else 5
if records:
    details.auto_filter.ref = f'A4:N{last_detail_row}'

# Summary: group by Status and Currency so original-currency amounts are not mixed.
summary = wb['Summary']
clear_rows(summary, 5, 6)
last_for_formula = 4 + len(records) if records else 5
summary['B2'] = f'=SUM(Details!I5:I{last_for_formula})'
from collections import defaultdict
agg = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
for rec in records:
    key = (rec['Status'], rec['Currency'])
    amt = money_decimal(rec['Amount'])
    agg[key]['amount'] += amt
    agg[key]['twd'] += amount_twd(amt, rec['Currency'])
    agg[key]['count'] += 1
for row_idx, ((status, currency), vals) in enumerate(sorted(agg.items()), start=5):
    row_values = [status, currency, currency, float(vals['amount'].quantize(Decimal('0.01'))), float(vals['twd'].quantize(Decimal('0.01'))), vals['count']]
    for col, value in enumerate(row_values, start=1):
        c = summary.cell(row=row_idx, column=col, value=value)
        if col in [1,2,3]:
            c.number_format = '@'
        elif col in [4,5]:
            c.number_format = '0.00'
        elif col == 6:
            c.number_format = '0'
if agg:
    summary.auto_filter.ref = f'A4:F{4+len(agg)}'

# Exceptions.
exceptions = wb['Exceptions']
clear_rows(exceptions, 5, 4)
exception_records = [r for r in records if bool(r.get('MissingReceipt'))]
for row_idx, rec in enumerate(exception_records, start=5):
    vals = [str(rec['ExpenseID']), str(rec['EmployeeID']), rec['ExpenseDate'], 'MissingReceipt=true; receipt is missing']
    for col, value in enumerate(vals, start=1):
        c = exceptions.cell(row=row_idx, column=col, value=value)
        if col in [1,2,4]:
            c.number_format = '@'
        elif col == 3:
            c.number_format = 'yyyy-mm-dd'
if exception_records:
    exceptions.auto_filter.ref = f'A4:D{4+len(exception_records)}'

# Save.
wb.save(OUT)
print('Wrote:', OUT)

# Validation by reopening saved output.
wb2 = load_workbook(OUT, data_only=False)
errors = []
# Sheet names preserved.
expected_sheets = ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
if wb2.sheetnames != expected_sheets:
    errors.append(f'Sheet names/order changed: {wb2.sheetnames}')
# Details checks.
d = wb2['Details']
headers_actual = [d.cell(4,c).value for c in range(1,15)]
headers_expected = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
if headers_actual != headers_expected:
    errors.append(f'Details headers mismatch: {headers_actual}')
detail_rows = []
for r in range(5, d.max_row+1):
    vals = [d.cell(r,c).value for c in range(1,15)]
    if any(v is not None for v in vals):
        detail_rows.append(vals)
if len(detail_rows) != 22:
    errors.append(f'Details row count {len(detail_rows)} != 22')
seen_ids = set()
for vals in detail_rows:
    exp_id, emp_id, _, _, exp_date, _, cur, amt, twd, status, receipt, desc, missing, needs = vals
    seen_ids.add(exp_id)
    if not isinstance(exp_id, str) or not exp_id.startswith('EXP-'):
        errors.append(f'ExpenseID not text/id format: {exp_id!r}')
    if status not in requested_statuses:
        errors.append(f'Unexpected status {status} for {exp_id}')
    if not ('2026-09-01' <= str(exp_date) <= '2026-09-30'):
        errors.append(f'Date out of range {exp_date} for {exp_id}')
    if not isinstance(amt, (int,float)) or not isinstance(twd, (int,float)):
        errors.append(f'Amounts not numeric for {exp_id}: {amt!r}, {twd!r}')
    if missing not in ('true','false') or needs not in ('true','false'):
        errors.append(f'Missing/Needs not true/false text for {exp_id}: {missing!r}/{needs!r}')
    if missing == 'true' and receipt is not None:
        errors.append(f'Missing receipt row has ReceiptID for {exp_id}: {receipt!r}')
    if desc is None:
        errors.append(f'Description missing for {exp_id}')
# Summary formula and counts.
s = wb2['Summary']
if s['B2'].value != '=SUM(Details!I5:I26)':
    errors.append(f'Summary B2 formula wrong: {s["B2"].value}')
summary_count = sum(s.cell(r,6).value or 0 for r in range(5, s.max_row+1))
if summary_count != 22:
    errors.append(f'Summary count {summary_count} != 22')
# Exceptions checks.
e = wb2['Exceptions']
ex_rows = []
for r in range(5, e.max_row+1):
    vals = [e.cell(r,c).value for c in range(1,5)]
    if any(v is not None for v in vals):
        ex_rows.append(vals)
ex_ids = {row[0] for row in ex_rows}
expected_ex_ids = {vals[0] for vals in detail_rows if vals[12] == 'true'}
if ex_ids != expected_ex_ids:
    errors.append(f'Exception IDs {sorted(ex_ids)} != missing detail IDs {sorted(expected_ex_ids)}')
if ex_ids != {'EXP-00006', 'EXP-00015'}:
    errors.append(f'Expected two known missing receipt IDs; got {sorted(ex_ids)}')
if not ex_ids.issubset(seen_ids):
    errors.append('Exception IDs are not retained in Details')
# Report cells.
if wb2['Report']['B3'].value != '2026-09-01' or wb2['Report']['D3'].value != '2026-09-30':
    errors.append('Report period cells are wrong')
if wb2['Instructions']['A1'].value != '填寫說明 Instructions（請勿修改本工作表）':
    errors.append('Instructions sheet appears modified')

print('\nValidation results:')
print(' detail rows:', len(detail_rows))
print(' exception rows:', len(ex_rows), sorted(ex_ids))
print(' summary groups/count:', len(agg), summary_count)
print(' statuses:', sorted({row[9] for row in detail_rows}))
print(' currencies:', sorted({row[6] for row in detail_rows}))
print(' output:', OUT.name, OUT.stat().st_size, 'bytes')
if errors:
    print('ERRORS:')
    for err in errors:
        print(' -', err)
    raise SystemExit(1)
print('All checks passed.')