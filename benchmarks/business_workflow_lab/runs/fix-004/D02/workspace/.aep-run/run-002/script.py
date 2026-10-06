from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
for p in (ws/'output').glob('*'):
    print('output file', p.name, p.stat().st_size)
for fname in ['submitted_2026-09.xlsx','approved_2026-09.xlsx']:
    p=ws/'output'/fname
    wb=load_workbook(p, data_only=False)
    print('\nEXPORT', fname, wb.sheetnames)
    sh=wb[wb.sheetnames[0]]
    print('rows', sh.max_row, 'cols', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 6), values_only=True): print(row)

wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
ins=wb['Instructions']
print('\nInstructions:')
for row in ins.iter_rows(values_only=True): print(row)
print('\nReport sheet rows:')
for row in wb['Report'].iter_rows(values_only=False):
    print([c.coordinate+':'+repr(c.value) for c in row])