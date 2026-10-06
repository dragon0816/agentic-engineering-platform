from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
# workbook structure
wb=load_workbook(ws/'input'/'expense_template.xlsx')
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
        print(row)
# fx sample
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    rows=[]
    for i,row in zip(range(6),r): rows.append(row)
print('fx sample:')
for row in rows: print(row)
# manual sample
text=(ws/'input'/'expense_manual.md').read_text(encoding='utf-8')
print('manual first 2000 chars:\n', text[:2000])