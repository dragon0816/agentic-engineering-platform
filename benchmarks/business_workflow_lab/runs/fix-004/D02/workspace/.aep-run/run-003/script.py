from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
for fname in ['submitted_2026-09.xlsx','approved_2026-09.xlsx']:
    wb=load_workbook(ws/'output'/fname, data_only=True)
    sh=wb.active
    print('\n', fname)
    headers=[c.value for c in sh[1]]
    print(headers)
    for row in sh.iter_rows(min_row=2, values_only=True):
        print(row)