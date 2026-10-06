from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
import csv, shutil, openpyxl, json
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

ws = Path('.')
input_xlsx = ws/'input'/'sales_template.xlsx'
out = ws/'output'/'sales_report_2026_09_open.xlsx'

# First inspect actual saved input sources again
print('Inputs:', [p.name for p in (ws/'input').iterdir()])
with open(ws/'output'/'downloads'/'sept_open_orders.csv', newline='', encoding='utf-8-sig') as f:
    list_rows = list(csv.DictReader(f))
print('Downloaded list row count:', len(list_rows))
print('Downloaded list IDs:', [r['OrderID'] for r in list_rows])

# Detail data read from the ERP detail pages during this conversation.
details = {
"0045000109": {"OrderID":"0045000109","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P003","OrderDate":"2026-09-01","Status":"Open","Currency":"TWD","NetAmount":"150000.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET10","SalesRep":"E001"},
"0045000110": {"OrderID":"0045000110","CustomerID":"C002","CustomerName":"台灣精密股份有限公司","ProjectID":"P002","OrderDate":"2026-09-02","Status":"Open","Currency":"TWD","NetAmount":"293572.50","CustomerContact":"張美玲 <mei@tw-precision.example>","PaymentTerms":"NET30","SalesRep":"E004"},
"0045000112": {"OrderID":"0045000112","CustomerID":"C007","CustomerName":"Acme Precision Taiwan Limited","ProjectID":"P005","OrderDate":"2026-09-04","Status":"Open","Currency":"USD","NetAmount":"12311.00","CustomerContact":"Peter Ho <peter.ho@acme-limited.example>","PaymentTerms":"NET15","SalesRep":"E001"},
"0045000113": {"OrderID":"0045000113","CustomerID":"C004","CustomerName":"藍海電子有限公司","ProjectID":"","OrderDate":"2026-09-05","Status":"Open","Currency":"TWD","NetAmount":"15390.00","CustomerContact":"","PaymentTerms":"NET15","SalesRep":"E005"},
"0045000115": {"OrderID":"0045000115","CustomerID":"C005","CustomerName":"Pacific Test Systems Inc.","ProjectID":"","OrderDate":"2026-09-09","Status":"Open","Currency":"USD","NetAmount":"6391.00","CustomerContact":"Sam Patel <sam@pacifictest.example>","PaymentTerms":"NET20","SalesRep":"E004"},
"0045000116": {"OrderID":"0045000116","CustomerID":"C008","CustomerName":"新竹量測科技","ProjectID":"","OrderDate":"2026-09-10","Status":"Open","Currency":"TWD","NetAmount":"29400.00","CustomerContact":"李建宏 <jh.lee@hsinchu-metrology.example>","PaymentTerms":"NET30","SalesRep":"E005"},
"0045000118": {"OrderID":"0045000118","CustomerID":"C003","CustomerName":"Northwind Components GmbH","ProjectID":"P004","OrderDate":"2026-09-14","Status":"Open","Currency":"EUR","NetAmount":"17484.75","CustomerContact":"Jonas Krüger <jk@northwind.example>","PaymentTerms":"NET0","SalesRep":"E006"},
"0045000119": {"OrderID":"0045000119","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-09-15","Status":"Open","Currency":"EUR","NetAmount":"8775.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET30","SalesRep":"E004"},
"0045000120": {"OrderID":"0045000120","CustomerID":"C006","CustomerName":"Helios Optics B.V.","ProjectID":"","OrderDate":"2026-09-16","Status":"Open","Currency":"EUR","NetAmount":"461.70","CustomerContact":"Eva de Vries <eva@helios.example>","PaymentTerms":"NET10","SalesRep":"E003"},
"0045000122": {"OrderID":"0045000122","CustomerID":"C002","CustomerName":"台灣精密股份有限公司","ProjectID":"P002","OrderDate":"2026-09-21","Status":"Open","Currency":"USD","NetAmount":"4141.00","CustomerContact":"張美玲 <mei@tw-precision.example>","PaymentTerms":"NET5","SalesRep":"E006"},
"0045000124": {"OrderID":"0045000124","CustomerID":"C004","CustomerName":"藍海電子有限公司","ProjectID":"","OrderDate":"2026-09-24","Status":"Open","Currency":"TWD","NetAmount":"45000.00","CustomerContact":"","PaymentTerms":"NET30","SalesRep":"E003"},
"0045000126": {"OrderID":"0045000126","CustomerID":"C008","CustomerName":"新竹量測科技","ProjectID":"","OrderDate":"2026-09-28","Status":"Open","Currency":"USD","NetAmount":"9171.00","CustomerContact":"李建宏 <jh.lee@hsinchu-metrology.example>","PaymentTerms":"NET30","SalesRep":"E003"},
"0045000127": {"OrderID":"0045000127","CustomerID":"C007","CustomerName":"Acme Precision Taiwan Limited","ProjectID":"P005","OrderDate":"2026-09-30","Status":"Open","Currency":"EUR","NetAmount":"2214.92","CustomerContact":"Peter Ho <peter.ho@acme-limited.example>","PaymentTerms":"NET30","SalesRep":"E006"},
"0045000128": {"OrderID":"0045000128","CustomerID":"C003","CustomerName":"Northwind Components GmbH","ProjectID":"P004","OrderDate":"2026-09-30","Status":"Open","Currency":"TWD","NetAmount":"104160.00","CustomerContact":"Jonas Krüger <jk@northwind.example>","PaymentTerms":"NET5","SalesRep":"E006"},
}

# Validate downloaded list vs detail pages to prevent duplicates/omissions.
list_ids = [r['OrderID'] for r in list_rows]
detail_ids = list(details)
if len(list_ids) != len(set(list_ids)):
    raise SystemExit('Duplicate OrderID in downloaded list')
if set(list_ids) != set(detail_ids):
    raise SystemExit(f'List/detail mismatch: missing={set(list_ids)-set(detail_ids)} extra={set(detail_ids)-set(list_ids)}')

# Validate list values against detail headers for core columns.
for lr in list_rows:
    d = details[lr['OrderID']]
    for k in ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']:
        if str(lr[k]) != str(d[k]):
            raise SystemExit(f'Mismatch for {lr["OrderID"]} {k}: list={lr[k]!r} detail={d[k]!r}')

# FX rates: first two columns are Currency, RateToTWD.
rates = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

# Copy template then fill.
shutil.copyfile(input_xlsx, out)
wb = openpyxl.load_workbook(out)
report = wb['Report']; det = wb['Details']; summ = wb['Summary']

# Report metadata
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Sales Orders: Status = Open; OrderDate 2026-09-01..2026-09-30'
report['B5'] = '2026-10-06T00:00:00+08:00'
report['B7'] = 'Lab ERP Sales Orders; fx_rates.csv'
report['D7'] = '2026-10-06T00:00:00+08:00'
report['B9'] = 'All 14 filtered Open sales orders were exported from list and detail pages were read. Blank source values are left blank.'

# Clear existing data rows while keeping row 4 headers/style as much as possible.
for row in det.iter_rows(min_row=5, max_row=max(det.max_row, 100), max_col=11):
    for cell in row:
        cell.value = None
for row in summ.iter_rows(min_row=5, max_row=max(summ.max_row, 100), max_col=6):
    for cell in row:
        cell.value = None

ordered = [details[oid] for oid in list_ids]
q2 = Decimal('0.01')
summary = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})

for idx, d in enumerate(ordered, start=5):
    amount = Decimal(d['NetAmount'])
    rate = rates[d['Currency']]
    twd = (amount * rate).quantize(q2, rounding=ROUND_HALF_UP)
    values = [
        d['OrderID'], d['CustomerID'], d['CustomerName'], d['ProjectID'], d['OrderDate'], d['Status'], d['Currency'],
        float(amount), float(twd), d['CustomerContact'], d['PaymentTerms']
    ]
    for col, val in enumerate(values, start=1):
        cell = det.cell(idx, col, val)
        if col in [1,2,4,5,6,7,10,11]:
            cell.number_format = '@'
        if col in [8,9]:
            cell.number_format = '0.00'
    key = (d['Status'], '', d['Currency'])
    summary[key]['amount'] += amount
    summary[key]['twd'] += twd
    summary[key]['count'] += 1

last_row = 4 + len(ordered)
summ['B2'] = f'=SUM(Details!I5:I{last_row})'
# Fill summary grouped by Status + Currency.
for r, (key, agg) in enumerate(sorted(summary.items(), key=lambda kv: (kv[0][0], kv[0][2])), start=5):
    group, subgroup, curr = key
    vals = [group, subgroup, curr, float(agg['amount'].quantize(q2)), float(agg['twd'].quantize(q2)), agg['count']]
    for c, val in enumerate(vals, start=1):
        cell = summ.cell(r,c,val)
        if c in [4,5]: cell.number_format = '0.00'

# Ensure autofilter covers data rows if filters exist or add sensible Details filter.
det.auto_filter.ref = f'A4:K{last_row}'
summ.auto_filter.ref = f'A4:F{4+len(summary)}'

wb.save(out)
print('Wrote', out)

# Re-open and verify every stated/template requirement from saved output.
wb2 = openpyxl.load_workbook(out, data_only=False)
errors=[]
if wb2.sheetnames != ['Report', 'Details', 'Summary', 'Instructions']:
    errors.append(f'sheet names changed: {wb2.sheetnames}')
D=wb2['Details']; S=wb2['Summary']; R=wb2['Report']
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
actual_headers=[D.cell(4,c).value for c in range(1,12)]
if actual_headers != expected_headers: errors.append(f'Details headers wrong: {actual_headers}')
rows=[]
for r in range(5, D.max_row+1):
    if D.cell(r,1).value is not None:
        rows.append(r)
if len(rows)!=14: errors.append(f'Details row count {len(rows)} != 14')
ids=[str(D.cell(r,1).value) for r in rows]
if ids != list_ids: errors.append(f'IDs/order mismatch {ids} != {list_ids}')
if len(ids)!=len(set(ids)): errors.append('Duplicate IDs in Details')
for r in rows:
    oid = str(D.cell(r,1).value)
    if not oid.startswith('00'): errors.append(f'OrderID leading zeros not preserved: {oid}')
    if D.cell(r,6).value != 'Open': errors.append(f'Non-Open status row {r}')
    od = D.cell(r,5).value
    if not ('2026-09-01' <= str(od) <= '2026-09-30'):
        errors.append(f'OrderDate out of range row {r}: {od}')
    if not isinstance(D.cell(r,8).value, (int,float)): errors.append(f'NetAmount not numeric row {r}')
    if not isinstance(D.cell(r,9).value, (int,float)): errors.append(f'NetAmountTWD not numeric row {r}')
    d=details[oid]
    exp_twd=(Decimal(d['NetAmount'])*rates[d['Currency']]).quantize(q2, rounding=ROUND_HALF_UP)
    if Decimal(str(D.cell(r,9).value)).quantize(q2) != exp_twd:
        errors.append(f'TWD mismatch {oid}: {D.cell(r,9).value} != {exp_twd}')
# blank source values must remain blank cells (None) for contacts/project blank
for oid in ['0045000113','0045000124']:
    rr = rows[ids.index(oid)]
    if D.cell(rr,10).value not in (None,''):
        errors.append(f'Blank CustomerContact not blank for {oid}')
# summary checks
if S['B2'].value != '=SUM(Details!I5:I18)': errors.append(f'Summary B2 formula {S["B2"].value}')
summary_rows=[r for r in range(5,S.max_row+1) if S.cell(r,1).value is not None]
if len(summary_rows)!=3: errors.append(f'Summary row count {len(summary_rows)} != 3')
# total formula result not calculated, check expected total by values
expected_total=sum((Decimal(d['NetAmount'])*rates[d['Currency']]).quantize(q2, rounding=ROUND_HALF_UP) for d in ordered)
summary_total=sum(Decimal(str(S.cell(r,5).value)).quantize(q2) for r in summary_rows)
if summary_total != expected_total: errors.append(f'Summary total {summary_total} != expected {expected_total}')
if R['B3'].value!='2026-09-01' or R['D3'].value!='2026-09-30': errors.append('Report date range wrong')
if errors:
    print('VERIFY FAILED')
    for e in errors: print('-', e)
    raise SystemExit(1)
print('VERIFY OK')
print('Checked: output file exists, workbook sheets unchanged, Details headers unchanged, 14 unique Open Sep-2026 OrderIDs match CSV, IDs preserved as text-like strings with leading zeros, amount/TWD numeric conversions use fx_rates ROUND_HALF_UP, missing source fields left blank, Summary formula extends to row 18, Summary totals reconcile.')
print('Expected TWD total:', expected_total)
print('Summary groups:')
for r in summary_rows:
    print([S.cell(r,c).value for c in range(1,7)])