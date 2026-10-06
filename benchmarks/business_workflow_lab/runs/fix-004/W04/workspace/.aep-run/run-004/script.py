import os, csv, json
from pathlib import Path
ws=Path('.')
print('input files:', [p.name for p in (ws/'input').iterdir()])
print('output files:', [str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()])
dl=ws/'output'/'downloads'/'sept_open_orders.csv'
if dl.exists():
    print('download exists', dl)
    with open(dl, newline='', encoding='utf-8-sig') as f:
        rows=list(csv.reader(f))
    print('rows', len(rows)); print('header', rows[0] if rows else None); print('sample', rows[1:6])
# inspect manual/template/fx briefly
import openpyxl
wb=openpyxl.load_workbook(ws/'input'/'sales_template.xlsx')
print('template sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max', sh.max_row, sh.max_column, 'headers', [sh.cell(1,c).value for c in range(1, sh.max_column+1)])
print('fx sample')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    for i,row in zip(range(5), csv.reader(f)):
        print(row)
print('manual first 1200:')
print((ws/'input'/'sap_manual.md').read_text(encoding='utf-8')[:1200])