from pathlib import Path
import openpyxl, csv
ws=Path('.')
wb=openpyxl.load_workbook(ws/'input'/'sales_template.xlsx')
for s in wb.sheetnames:
    sh=wb[s]
    print('\nSHEET', s, 'max', sh.max_row, sh.max_column)
    for r in range(1, min(sh.max_row,30)+1):
        vals=[sh.cell(r,c).value for c in range(1, min(sh.max_column,15)+1)]
        print(r, vals)
print('\nmanual rest relevant:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[1200:3200])