from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
manual=ws/'input'/'sap_manual.md'
print('\n--- sap_manual.md head ---')
print(manual.read_text(encoding='utf-8')[:4000])
print('\n--- fx_rates.csv ---')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rows=list(csv.reader(f))
print('rows', len(rows))
for r in rows[:10]: print(r)
print('\n--- sales_template.xlsx ---')
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
print('sheets', wb.sheetnames)
for sh in wb.sheetnames:
    s=wb[sh]
    print('sheet', sh, 'max_row', s.max_row, 'max_col', s.max_column)
    for row in s.iter_rows(min_row=1, max_row=min(8,s.max_row), values_only=True):
        print(row)
    print('merged', list(s.merged_cells.ranges))