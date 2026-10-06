import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(6), r):
        print(row)
print('\nmanual excerpts:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    print(f'--- {name} ---')
    text=(ws/'input'/name).read_text(encoding='utf-8')
    print(text[:4000])
print('\nproject_template.xlsx:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
print('sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet',s,'max',sh.max_row,sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)
