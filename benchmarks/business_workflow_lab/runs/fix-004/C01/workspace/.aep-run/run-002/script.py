from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
print('\nReport cells:')
rep=wb['Report']
for r in range(1,21):
    print(r, [rep.cell(r,c).value for c in range(1,7)])
