from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('cwd', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
print('\nmanual head:')
print((ws/'input'/'quote_cli_manual.md').read_text(encoding='utf-8')[:4000])
print('\nfx_rates sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r): print(row)
print('\nworkbook sheets/headers:')
wb = load_workbook(ws/'input'/'quote_template.xlsx')
for sh in wb.sheetnames:
    wsht=wb[sh]
    print(sh, wsht.max_row, wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(5, wsht.max_row), values_only=True): print(' ', row)
