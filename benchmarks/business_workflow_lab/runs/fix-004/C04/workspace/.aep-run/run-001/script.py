from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('CWD', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    txt = manual.read_text(encoding='utf-8', errors='replace')
    print('\n--- quote_cli_manual.md first 2000 chars ---')
    print(txt[:2000])
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\n--- fx_rates.csv sample ---')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i,row in zip(range(6), r):
            print(row)
xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    print('\n--- workbook structure ---')
    wb = load_workbook(xlsx, data_only=False)
    print('sheets:', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
            print(row)
        print('---')