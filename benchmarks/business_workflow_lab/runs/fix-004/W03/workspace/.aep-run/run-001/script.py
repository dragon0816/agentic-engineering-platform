from pathlib import Path
from openpyxl import load_workbook
import csv
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r):
        print(row)
print('\nsap_manual.md head:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[:1500])
print('\nworkbook sheets/headers:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print(sh, wsht.max_row, wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(5, wsht.max_row), values_only=True):
        print(' ', row)
    print()