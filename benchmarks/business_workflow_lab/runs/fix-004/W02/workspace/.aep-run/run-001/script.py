from pathlib import Path
from openpyxl import load_workbook
import csv
ws = Path.cwd()
print('Workspace:', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
# Inspect sales_template.xlsx
xlsx = ws/'input'/'sales_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('Workbook sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print('Sheet:', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(5, sh.max_row), values_only=True):
        print('  ', row)
# Inspect fx_rates.csv
csvp = ws/'input'/'fx_rates.csv'
if csvp.exists():
    with csvp.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i,row in zip(range(5), reader):
            print('fx row', i+1, row)
# Inspect manual excerpt
md = ws/'input'/'sap_manual.md'
if md.exists():
    text = md.read_text(encoding='utf-8')
    print('Manual first 2000 chars:')
    print(text[:2000])