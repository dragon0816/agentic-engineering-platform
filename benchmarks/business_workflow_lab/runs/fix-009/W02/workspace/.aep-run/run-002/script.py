from pathlib import Path
from openpyxl import load_workbook
ws=Path('.')
print('Instructions full:')
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
print('\nReport merged ranges:', list(wb['Report'].merged_cells.ranges))
print('Details merged ranges:', list(wb['Details'].merged_cells.ranges))
print('Summary merged ranges:', list(wb['Summary'].merged_cells.ranges))
print('\nWorkbook formulas/styles check key rows:')
for title in ['Details','Summary','Report']:
    sh=wb[title]
    print('---', title)
    for r in range(1, min(sh.max_row,25)+1):
        print(r, [sh.cell(r,c).value for c in range(1, sh.max_column+1)])