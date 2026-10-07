from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime
import csv, shutil

workspace = Path('.')
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

source_template = input_dir / 'sales_template.xlsx'
orders_csv = output_dir / 'downloads' / 'orders_C001_2026-08-01_2026-10-31.csv'
fx_csv = input_dir / 'fx_rates.csv'
outfile = output_dir / 'sales_report_C001_2026-08-01_2026-10-31.xlsx'

print('Inputs for rerun:')
for p in [source_template, orders_csv, fx_csv]:
    print(p, 'exists=', p.exists(), 'bytes=', p.stat().st_size if p.exists() else None)

# Detail data read from ERP detail pages using browser in this conversation.
detail_headers = {
    '0045000101': {'OrderID':'0045000101','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-08-03','Status':'Closed','Currency':'TWD','NetAmount':'95895.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET15'},
    '0045000107': {'OrderID':'0045000107','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P003','OrderDate':'2026-08-27','Status':'Closed','Currency':'USD','NetAmount':'6250.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET30'},
    '0045000109': {'OrderID':'0045000109','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P003','OrderDate':'2026-09-01','Status':'Open','Currency':'TWD','NetAmount':'150000.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET10'},
    '0045000114': {'OrderID':'0045000114','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-09-08','Status':'Cancelled','Currency':'USD','NetAmount':'11956.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET5'},
    '0045000119': {'OrderID':'0045000119','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-09-15','Status':'Open','Currency':'EUR','NetAmount':'8775.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET30'},
    '0045000125': {'OrderID':'0045000125','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-09-25','Status':'Closed','Currency':'TWD','NetAmount':'386445.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET0'},
    '0045000129': {'OrderID':'0045000129','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-10-01','Status':'Open','Currency':'USD','NetAmount':'1602.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET30'},
    '0045000132': {'OrderID':'0045000132','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P003','OrderDate':'2026-10-09','Status':'Open','Currency':'TWD','NetAmount':'50400.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':''},
    '0045000136': {'OrderID':'0045000136','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-10-31','Status':'Open','Currency':'EUR','NetAmount':'2250.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET5'},
}

with open(orders_csv, newline='', encoding='utf-8-sig') as f:
    orders = list(csv.DictReader(f))
print('Downloaded order rows:', len(orders))
print('OrderIDs:', [r['OrderID'] for r in orders])
print('Sample row:', orders[0] if orders else None)

rates = {}
with open(fx_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates used:', rates)

for r in orders:
    assert r['CustomerID'] == 'C001', f"Non-C001 row found: {r}"
    assert '2026-08-01' <= r['OrderDate'] <= '2026-10-31', f"Out-of-period row found: {r}"
assert [r['OrderID'] for r in orders] == list(detail_headers.keys())
assert not any(r['CustomerID'] == 'C007' for r in orders)

TWOPLACES = Decimal('0.01')
def money_twd(amount, currency):
    return (Decimal(str(amount)) * rates[currency]).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def to_date(s):
    y,m,d = map(int, s.split('-'))
    return date(y,m,d)

shutil.copyfile(source_template, outfile)
wb = load_workbook(outfile)

report = wb['Report']
report['B3'] = '2026-08-01'
report['D3'] = '2026-10-31'
report['B4'] = 'Sales Orders; CustomerID exactly C001; OrderDate 2026-08-01 through 2026-10-31; all statuses; sort OrderID asc.'
report['B5'] = '2026-10-06T00:00:00+08:00'
report['B7'] = 'Lab ERP Sales Orders detail; fx_rates.csv'
report['D7'] = '2026-10-06T00:00:00+08:00'
report['B9'] = 'Missing data: PaymentTerms is blank for OrderID 0045000132. CustomerContact has no missing values. Exact CustomerID C001 filter used; similar-name C007 not included.'

details = wb['Details']
if details.max_row >= 5:
    details.delete_rows(5, details.max_row - 4)
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert [details.cell(4,c).value for c in range(1,12)] == expected_headers

summary_rows = []
for idx, order in enumerate(orders, start=5):
    d = detail_headers[order['OrderID']]
    net = Decimal(d['NetAmount'])
    twd = money_twd(net, d['Currency'])
    values = [
        d['OrderID'], d['CustomerID'], d['CustomerName'], d.get('ProjectID') or None,
        to_date(d['OrderDate']), d['Status'], d['Currency'], float(net), float(twd),
        d.get('CustomerContact') or None, d.get('PaymentTerms') or None,
    ]
    for col, value in enumerate(values, start=1):
        cell = details.cell(idx, col, value)
        if col in (1,2,4):
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9):
            cell.number_format = '0.00'
        elif col in (3,6,7,10,11):
            cell.number_format = '@'
    summary_rows.append((d['Status'], d['Currency'], net, twd))
details.auto_filter.ref = f'A4:K{4 + len(orders)}'

summary = wb['Summary']
if summary.max_row >= 5:
    summary.delete_rows(5, summary.max_row - 4)
last_detail_row = 4 + len(orders) if orders else 5
summary['B2'] = f'=SUM(Details!I5:I{last_detail_row if orders else 5})'
summary['B2'].number_format = '0.00'

groups = {}
order_keys = []
for status, currency, amount, twd in summary_rows:
    key = (status, currency)
    if key not in groups:
        groups[key] = [Decimal('0'), Decimal('0'), 0]
        order_keys.append(key)
    groups[key][0] += amount
    groups[key][1] += twd
    groups[key][2] += 1
for rix, key in enumerate(order_keys, start=5):
    status, currency = key
    amount, twd, count = groups[key]
    vals = [status, None, currency, float(amount.quantize(TWOPLACES)), float(twd.quantize(TWOPLACES)), count]
    for col, value in enumerate(vals, start=1):
        cell = summary.cell(rix, col, value)
        if col in (4,5):
            cell.number_format = '0.00'
        elif col == 6:
            cell.number_format = '0'

wb.save(outfile)
print('Wrote:', outfile)

# Reopen and verify saved output.
wb2 = load_workbook(outfile, data_only=False)
source_wb = load_workbook(source_template, read_only=True)
assert wb2.sheetnames == source_wb.sheetnames, wb2.sheetnames
assert str(list(wb2['Report'].merged_cells.ranges)) == str(list(load_workbook(source_template)['Report'].merged_cells.ranges))
rep = wb2['Report']
assert rep['B3'].value == '2026-08-01'
assert rep['D3'].value == '2026-10-31'
assert 'C001' in rep['B4'].value and '0045000132' in rep['B9'].value

det = wb2['Details']
assert [det.cell(4,c).value for c in range(1,12)] == expected_headers
assert det.max_row == 13, det.max_row
seen=[]
missing_payment=[]
for row in range(5,14):
    oid = det.cell(row,1).value
    seen.append(oid)
    assert isinstance(oid, str) and oid.startswith('0045')
    assert det.cell(row,2).value == 'C001'
    assert det.cell(row,3).value == 'Acme Precision Taiwan Ltd.'
    dt = det.cell(row,5).value
    if isinstance(dt, datetime):
        dt = dt.date()
    assert date(2026,8,1) <= dt <= date(2026,10,31), (oid, dt)
    assert isinstance(det.cell(row,8).value, (int,float))
    assert isinstance(det.cell(row,9).value, (int,float))
    if det.cell(row,11).value in (None, ''):
        missing_payment.append(oid)
assert seen == list(detail_headers.keys()), seen
assert missing_payment == ['0045000132'], missing_payment
assert all(det.cell(r,10).value for r in range(5,14))

summ = wb2['Summary']
assert summ['B2'].value == '=SUM(Details!I5:I13)', summ['B2'].value
assert summ.max_row == 10, summ.max_row
summary_total = sum(Decimal(str(summ.cell(r,5).value)) for r in range(5,11)).quantize(TWOPLACES)
details_total = sum(Decimal(str(det.cell(r,9).value)) for r in range(5,14)).quantize(TWOPLACES)
assert summary_total == details_total == Decimal('1702471.00'), (summary_total, details_total)
assert all(isinstance(summ.cell(r,4).value, (int,float)) for r in range(5,11))
assert all(isinstance(summ.cell(r,5).value, (int,float)) for r in range(5,11))
print('Checked: workbook sheets/merged range preserved; Report fields populated; 9 rows all exact CustomerID C001 and dates inclusive; no C007/similar-name customer rows; detail CustomerContact/PaymentTerms written with blank only where source blank; OrderIDs are text; amounts are numeric; Summary B2 formula and grouped totals equal Details total 1702471.00 TWD.')
print('Output file:', outfile.resolve())