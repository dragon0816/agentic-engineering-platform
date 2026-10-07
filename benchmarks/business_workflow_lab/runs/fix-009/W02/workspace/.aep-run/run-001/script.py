from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path('.')
print('Workspace:', ws.resolve())
for p in sorted((ws/'input').iterdir()):
    print('INPUT', p.name, p.stat().st_size)
print('\nfx_rates.csv:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(6), r): print(row)
print('\nsales_template.xlsx sheets and used ranges:')
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
for sh in wb.worksheets:
    print('SHEET', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(row)
print('\nsap_manual.md first 200 lines:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:200],1):
    print(f'{i:03}: {line}')