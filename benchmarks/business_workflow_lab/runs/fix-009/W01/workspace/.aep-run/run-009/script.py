from pathlib import Path
from openpyxl import load_workbook
import csv

ws = Path.cwd()
print('WORKSPACE', ws)
print('OUTPUT FILES:')
for p in sorted((ws/'output').rglob('*')):
    if p.is_file():
        print('-', p.relative_to(ws), p.stat().st_size)

for rel in ['output/august_closed_orders.xlsx', 'output/september_open_orders.xlsx', 'output/downloads/workflow_sales_orders_filtered.csv', 'output/downloads/workflow_all_sales_orders.csv']:
    p = ws / rel
    print('\nCHECK', rel, 'exists=', p.exists(), 'size=', p.stat().st_size if p.exists() else None)
    if p.suffix.lower() == '.xlsx' and p.exists():
        wb = load_workbook(p, data_only=False)
        print(' sheets=', wb.sheetnames)
        for s in ['Report','Details','Summary','Instructions']:
            if s in wb.sheetnames:
                sh = wb[s]
                print(f' {s}: max_row={sh.max_row}, max_col={sh.max_column}')
                if s == 'Details':
                    header = [sh.cell(4,c).value for c in range(1, sh.max_column+1)]
                    print('  header row 4=', header)
                    rows = []
                    r = 5
                    while r <= sh.max_row:
                        vals = [sh.cell(r,c).value for c in range(1, min(sh.max_column, 12)+1)]
                        if any(v is not None for v in vals):
                            rows.append(vals)
                        r += 1
                    print('  detail_rows=', len(rows))
                    print('  first rows=', rows[:5])
                if s == 'Summary':
                    for r in range(1, min(sh.max_row,10)+1):
                        vals=[sh.cell(r,c).value for c in range(1, min(sh.max_column,8)+1)]
                        if any(v is not None for v in vals):
                            print('  summary row', r, vals)
    elif p.suffix.lower() == '.csv' and p.exists():
        with p.open('r', encoding='utf-8-sig', newline='') as f:
            rdr = csv.DictReader(f)
            rows = list(rdr)
        print(' csv headers=', rdr.fieldnames)
        print(' csv rows=', len(rows))
        print(' csv first rows=', rows[:5])