from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
import csv
from datetime import datetime

ws = Path.cwd()
out = ws/'output'/'september_open_orders.xlsx'
filtered_csv = ws/'output'/'downloads'/'workflow_sales_orders_filtered.csv'
fx_csv = ws/'input'/'fx_rates.csv'
template = ws/'input'/'sales_template.xlsx'

print('CHECK START')
print('workspace', ws)
assert out.exists(), f'Missing {out}'
assert filtered_csv.exists(), f'Missing {filtered_csv}'
assert fx_csv.exists(), f'Missing {fx_csv}'
assert template.exists(), f'Missing {template}'

# Load FX rates
fx = {}
with fx_csv.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('fx headers:', reader.fieldnames)
    for r in reader:
        # tolerate common naming
        cur = (r.get('Currency') or r.get('currency') or r.get('CCY') or '').strip()
        rate = (r.get('RateToTWD') or r.get('TWD') or r.get('rate_to_twd') or r.get('Rate') or '').strip()
        if cur and rate:
            fx[cur] = Decimal(rate)
print('fx loaded:', fx)
assert fx == {'TWD': Decimal('1'), 'USD': Decimal('32'), 'EUR': Decimal('35')}, fx

# Load filtered CSV: all exported filtered rows
with filtered_csv.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
print('filtered CSV count:', len(rows))
print('filtered CSV headers:', rows[0].keys() if rows else [])
assert len(rows) == 14, len(rows)
ids = [r['OrderID'] for r in rows]
assert ids == sorted(ids), 'CSV OrderID not sorted asc'
for r in rows:
    assert r['Status'] == 'Open', r
    assert '2026-09-01' <= r['OrderDate'] <= '2026-09-30', r
    assert r['OrderID'].startswith('0045'), r['OrderID']
    assert r['Currency'] in fx, r['Currency']
print('CSV status/date/sort checks OK')

# Load workbooks
wb = load_workbook(out, data_only=False)
wbv = load_workbook(out, data_only=True)
twb = load_workbook(template, data_only=False)
assert wb.sheetnames == twb.sheetnames == ['Report', 'Details', 'Summary', 'Instructions'], wb.sheetnames
print('sheet names/template match OK')

# Report checks
rep = wb['Report']
assert rep['B3'].value == '2026-09-01', rep['B3'].value
assert rep['D3'].value == '2026-09-30', rep['D3'].value
assert 'Status=Open' in str(rep['B4'].value), rep['B4'].value
assert 'OrderID asc' in str(rep['B4'].value), rep['B4'].value
print('Report parameter cells OK')

# Details checks
sh = wb['Details']
headers = [sh.cell(4, c).value for c in range(1, sh.max_column+1)]
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
print('Details headers:', headers)
assert headers == expected_headers, headers
assert sh.max_row == 18, sh.max_row

detail_rows = []
for r in range(5, sh.max_row+1):
    rec = {headers[c-1]: sh.cell(r,c).value for c in range(1, len(headers)+1)}
    detail_rows.append(rec)
assert len(detail_rows) == 14, len(detail_rows)
assert [r['OrderID'] for r in detail_rows] == sorted([r['OrderID'] for r in detail_rows]), 'Workbook OrderID not sorted'
assert [r['OrderID'] for r in detail_rows] == ids, 'Workbook rows do not match filtered CSV'

# Compare each detail row to CSV and recalc TWD
csv_by_id = {r['OrderID']: r for r in rows}
summary_calc = {}
total_twd = Decimal('0.00')
for rec in detail_rows:
    oid = rec['OrderID']
    src = csv_by_id[oid]
    assert isinstance(oid, str), f'OrderID not text for {oid}'
    for col in ['CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency']:
        expected = src[col] if src[col] != '' else None
        assert rec[col] == expected, f'{oid} {col}: {rec[col]!r} != {expected!r}'
    assert rec['Status'] == 'Open'
    assert '2026-09-01' <= rec['OrderDate'] <= '2026-09-30'
    net = Decimal(str(src['NetAmount']))
    assert Decimal(str(rec['NetAmount'])) == net, f'{oid} NetAmount mismatch'
    twd = (net * fx[rec['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert Decimal(str(rec['NetAmountTWD'])).quantize(Decimal('0.01')) == twd, f'{oid} TWD {rec["NetAmountTWD"]} != {twd}'
    assert rec['PaymentTerms'] not in (None, ''), f'{oid} PaymentTerms missing'
    summary_calc.setdefault(rec['Currency'], {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    summary_calc[rec['Currency']]['amount'] += net
    summary_calc[rec['Currency']]['twd'] += twd
    summary_calc[rec['Currency']]['count'] += 1
    total_twd += twd
print('Details row checks OK')
print('calculated summary:', summary_calc, 'total_twd:', total_twd)

# Summary checks
sumsh = wb['Summary']
assert sumsh['B2'].value == '=SUM(Details!I5:I18)', sumsh['B2'].value
expected_summary = {
    'EUR': {'amount': Decimal('28936.37'), 'twd': Decimal('1012772.95'), 'count': 4},
    'TWD': {'amount': Decimal('637522.50'), 'twd': Decimal('637522.50'), 'count': 6},
    'USD': {'amount': Decimal('32014.00'), 'twd': Decimal('1024448.00'), 'count': 4},
}
for r in range(5,8):
    cur = sumsh.cell(r,3).value
    assert cur in expected_summary, cur
    exp = expected_summary[cur]
    calc = summary_calc[cur]
    assert calc['amount'].quantize(Decimal('0.01')) == exp['amount'], (cur, calc, exp)
    assert calc['twd'].quantize(Decimal('0.01')) == exp['twd'], (cur, calc, exp)
    assert calc['count'] == exp['count'], (cur, calc, exp)
    assert Decimal(str(sumsh.cell(r,4).value)).quantize(Decimal('0.01')) == exp['amount'], cur
    assert Decimal(str(sumsh.cell(r,5).value)).quantize(Decimal('0.01')) == exp['twd'], cur
    assert sumsh.cell(r,6).value == exp['count'], cur
assert sumsh.cell(8,1).value == 'Total'
assert Decimal(str(sumsh.cell(8,5).value)).quantize(Decimal('0.01')) == Decimal('2674743.45')
assert sumsh.cell(8,6).value == 14
assert total_twd.quantize(Decimal('0.01')) == Decimal('2674743.45')
print('Summary checks OK')

# Verify no unexpected formulas/value issue in calculated view for summary total may be stale if no engine; formula itself already checked.
print('generated_at:', rep['B5'].value)
print('OUTPUT FILE:', out, out.stat().st_size, 'bytes')
print('DOWNLOAD FILTERED:', filtered_csv, filtered_csv.stat().st_size, 'bytes')
print('CHECK PASSED')
