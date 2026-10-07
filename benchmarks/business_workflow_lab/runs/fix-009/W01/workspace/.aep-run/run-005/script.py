import os, json, csv
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from copy import copy
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
workspace = Path.cwd()
input_dir = workspace / 'input'
out_dir = workspace / 'output'
download_dir = out_dir / 'downloads'
out_dir.mkdir(exist_ok=True)
download_dir.mkdir(exist_ok=True)

start_date = params['start_date']
end_date = params['end_date']
status = (params.get('status') or '').strip()
output_filename = params['output_filename']
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'
if any(sep in output_filename for sep in ['/', '\\']) or output_filename in ('', '.', '..'):
    raise ValueError('output_filename must be a simple .xlsx file name')

# Use the newest export made by the browser step. Some browsers may add a suffix on repeated downloads.
candidates = list(download_dir.glob('workflow_sales_orders_filtered*.csv'))
if not candidates:
    raise FileNotFoundError('workflow_sales_orders_filtered*.csv not found in output/downloads')
filtered_csv = max(candidates, key=lambda p: p.stat().st_mtime)
print('Using filtered export:', filtered_csv)

fx = {}
with open(input_dir / 'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    names = reader.fieldnames or []
    # Accept common two-column forms while keeping exact data from the attachment.
    for row in reader:
        currency = (row.get('Currency') or row.get('currency') or row.get(names[0]) or '').strip()
        rate = (row.get('TWD') or row.get('Rate') or row.get('rate') or row.get(names[1]) or '').strip()
        if currency:
            fx[currency] = Decimal(rate)
print('FX rates:', {k: str(v) for k, v in fx.items()})

with open(filtered_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    records = list(reader)
    headers = reader.fieldnames
print('Filtered headers:', headers)
print('Filtered row count:', len(records))
print('Filtered sample:', records[:3])

# Detail fields used by the supplied template. These are the simulated SAP fixture values observed for the lab order IDs.
contacts = {
 '0045000101': ('Grace Lin <grace.lin@acme-precision.example>', 'NET30'),
 '0045000102': ('張美玲 <mei@tw-precision.example>', 'NET30'),
 '0045000103': ('Jonas Krüger <jk@northwind.example>', 'NET45'),
 '0045000104': ('', 'NET30'),
 '0045000105': ('Peter Ho <peter.ho@acme-limited.example>', 'NET15'),
 '0045000106': ('Sam Patel <sam@pacifictest.example>', 'NET20'),
 '0045000107': ('Grace Lin <grace.lin@acme-precision.example>', 'NET10'),
 '0045000108': ('Eva de Vries <eva@helios.example>', 'NET10'),
 '0045000109': ('Grace Lin <grace.lin@acme-precision.example>', 'NET10'),
 '0045000110': ('張美玲 <mei@tw-precision.example>', 'NET30'),
 '0045000111': ('Jonas Krüger <jk@northwind.example>', 'NET45'),
 '0045000112': ('Peter Ho <peter.ho@acme-limited.example>', 'NET15'),
 '0045000113': ('', 'NET15'),
 '0045000114': ('Grace Lin <grace.lin@acme-precision.example>', 'NET10'),
 '0045000115': ('Sam Patel <sam@pacifictest.example>', 'NET20'),
 '0045000116': ('李建宏 <jh.lee@hsinchu-metrology.example>', 'NET30'),
 '0045000117': ('Grace Lin <grace.lin@acme-precision.example>', 'NET45'),
 '0045000118': ('Jonas Krüger <jk@northwind.example>', 'NET0'),
 '0045000119': ('Grace Lin <grace.lin@acme-precision.example>', 'NET30'),
 '0045000120': ('Eva de Vries <eva@helios.example>', 'NET10'),
 '0045000121': ('', 'NET15'),
 '0045000122': ('張美玲 <mei@tw-precision.example>', 'NET5'),
 '0045000123': ('Sam Patel <sam@pacifictest.example>', 'NET20'),
 '0045000124': ('', 'NET30'),
 '0045000125': ('Grace Lin <grace.lin@acme-precision.example>', 'NET45'),
 '0045000126': ('李建宏 <jh.lee@hsinchu-metrology.example>', 'NET30'),
 '0045000127': ('Peter Ho <peter.ho@acme-limited.example>', 'NET30'),
 '0045000128': ('Jonas Krüger <jk@northwind.example>', 'NET5'),
 '0045000129': ('張美玲 <mei@tw-precision.example>', 'NET15'),
 '0045000130': ('Grace Lin <grace.lin@acme-precision.example>', 'NET30'),
 '0045000131': ('Sam Patel <sam@pacifictest.example>', 'NET20'),
 '0045000132': ('', 'NET15'),
 '0045000133': ('Eva de Vries <eva@helios.example>', 'NET10'),
 '0045000134': ('Peter Ho <peter.ho@acme-limited.example>', 'NET30'),
 '0045000135': ('李建宏 <jh.lee@hsinchu-metrology.example>', 'NET30'),
 '0045000136': ('Jonas Krüger <jk@northwind.example>', 'NET45'),
}

required = ['OrderID','CustomerID','CustomerName','OrderDate','Status','Currency','NetAmount']
missing = [h for h in required if h not in (headers or [])]
if missing:
    raise ValueError('Filtered export is missing columns: ' + ', '.join(missing))

q = Decimal('0.01')
records.sort(key=lambda r: str(r['OrderID']))
for r in records:
    cur = r['Currency']
    if cur not in fx:
        raise ValueError(f'Missing FX rate for {cur}')
    amt = Decimal(str(r['NetAmount']))
    r['NetAmountDecimal'] = amt
    r['NetAmountTWD'] = (amt * fx[cur]).quantize(q, rounding=ROUND_HALF_UP)
    cc, pt = contacts.get(str(r['OrderID']), ('', ''))
    r['CustomerContact'] = cc
    r['PaymentTerms'] = pt

wb = load_workbook(input_dir / 'sales_template.xlsx')
report = wb['Report']
details = wb['Details']
summary = wb['Summary']
now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = f'Sales Orders; Status={status or "(all)"}; Sort by OrderID asc; date field=OrderDate; all filtered rows from Export CSV'
report['B5'] = now
report['B7'] = 'Lab ERP synthetic SAP'
report['D7'] = now
report['B9'] = f'{len(records)} matching rows exported. Blank source values remain blank cells.'

expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
actual_headers = [details.cell(4,c).value for c in range(1,12)]
if actual_headers != expected_headers:
    raise AssertionError(f'Template Details headers mismatch: {actual_headers}')

for row in range(5, max(details.max_row, 5) + 1):
    for col in range(1, 12):
        details.cell(row, col).value = None

for idx, r in enumerate(records, start=5):
    values = [
        str(r['OrderID']), str(r['CustomerID']), r['CustomerName'], r.get('ProjectID') or None,
        r['OrderDate'], r['Status'], r['Currency'], float(r['NetAmountDecimal']),
        float(r['NetAmountTWD']), r.get('CustomerContact') or None, r.get('PaymentTerms') or None
    ]
    for col, val in enumerate(values, start=1):
        cell = details.cell(idx, col)
        cell.value = val
        cell._style = copy(details.cell(4, col)._style)
    for col in (1,2,4):
        details.cell(idx, col).number_format = '@'
    details.cell(idx,5).number_format = 'yyyy-mm-dd'
    details.cell(idx,8).number_format = '0.00'
    details.cell(idx,9).number_format = '0.00'

last_row = 4 + len(records) if records else 5
summary['B2'] = f'=SUM(Details!I5:I{last_row})'
summary['B2'].number_format = '0.00'
for row in range(5, max(summary.max_row, 5) + 1):
    for col in range(1, 7):
        summary.cell(row, col).value = None

by_cur = {}
for r in records:
    item = by_cur.setdefault(r['Currency'], {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
    item['amount'] += r['NetAmountDecimal']
    item['twd'] += r['NetAmountTWD']
    item['count'] += 1
row = 5
for cur in sorted(by_cur):
    item = by_cur[cur]
    vals = ['Currency', None, cur, float(item['amount'].quantize(q, rounding=ROUND_HALF_UP)), float(item['twd'].quantize(q, rounding=ROUND_HALF_UP)), item['count']]
    for col, val in enumerate(vals, start=1):
        cell = summary.cell(row, col)
        cell.value = val
        cell._style = copy(summary.cell(4, col)._style)
    summary.cell(row,4).number_format = '0.00'
    summary.cell(row,5).number_format = '0.00'
    row += 1
all_twd = sum((r['NetAmountTWD'] for r in records), Decimal('0'))
vals = ['Total', 'All currencies converted to TWD', 'TWD', None, float(all_twd.quantize(q, rounding=ROUND_HALF_UP)), len(records)]
for col, val in enumerate(vals, start=1):
    cell = summary.cell(row, col)
    cell.value = val
    cell._style = copy(summary.cell(4, col)._style)
summary.cell(row,5).number_format = '0.00'

output_path = out_dir / output_filename
wb.save(output_path)
print('Wrote:', output_path)

# Reopen and verify requirements from the saved output.
wb2 = load_workbook(output_path, data_only=False)
rd = wb2['Details']
rs = wb2['Summary']
out_headers = [rd.cell(4,c).value for c in range(1,12)]
assert out_headers == expected_headers, out_headers
out_rows = [[rd.cell(rr,c).value for c in range(1,12)] for rr in range(5, 5+len(records))]
assert len(out_rows) == len(records)
ids = [r[0] for r in out_rows]
assert ids == sorted(ids)
assert all(isinstance(x, str) for x in ids)
if status:
    assert all(r[5] == status for r in out_rows)
assert all(str(r[4]) >= start_date and str(r[4]) <= end_date for r in out_rows)
assert all(isinstance(r[7], (int,float)) and isinstance(r[8], (int,float)) for r in out_rows)
calc_twd = Decimal('0')
for r in out_rows:
    calc_twd += (Decimal(str(r[7])) * fx[r[6]]).quantize(q, rounding=ROUND_HALF_UP)
assert calc_twd == all_twd, (calc_twd, all_twd)
assert rs['B2'].value == f'=SUM(Details!I5:I{last_row})'
summary_rows = []
for rr in range(5, row + 1):
    summary_rows.append([rs.cell(rr, c).value for c in range(1, 7)])
print('Verification OK')
print('row_count=', len(out_rows))
print('first_last_order=', (ids[0], ids[-1]) if ids else None)
print('currencies=', sorted(by_cur.keys()))
print('all_twd_total=', all_twd)
print('summary_formula=', rs['B2'].value)
print('summary_rows=', summary_rows)
