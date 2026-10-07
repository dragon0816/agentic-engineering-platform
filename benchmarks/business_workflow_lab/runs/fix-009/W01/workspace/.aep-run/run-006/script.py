from pathlib import Path
from openpyxl import load_workbook
import csv

ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.relative_to(ws), p.stat().st_size)
print('output files:')
for p in sorted((ws/'output').glob('*')):
    if p.is_file():
        print(' -', p.relative_to(ws), p.stat().st_size)
print('download files:')
for p in sorted((ws/'output'/'downloads').glob('*')):
    print(' -', p.relative_to(ws), p.stat().st_size)

out = ws/'output'/'september_open_orders.xlsx'
if out.exists():
    wb = load_workbook(out, data_only=False)
    print('workbook sheets:', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print('sheet', s, 'max_row', sh.max_row, 'max_column', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
            print(' ', row)
else:
    print('missing output/september_open_orders.xlsx')

for name in ['workflow_sales_orders_filtered.csv','workflow_all_sales_orders.csv']:
    p = ws/'output'/'downloads'/name
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as f:
            rows = list(csv.reader(f))
        print(name, 'rows including header:', len(rows), 'header:', rows[0] if rows else None, 'first data:', rows[1] if len(rows)>1 else None)
