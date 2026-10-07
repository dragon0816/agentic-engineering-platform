from pathlib import Path
from openpyxl import load_workbook
import csv
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv preview:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(6), r):
        print(row)
print('\nsap_manual.md head:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[:3000])
print('\nsales_template.xlsx structure:')
wb=load_workbook(ws/'input'/'sales_template.xlsx')
print('sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)
    print('---')
