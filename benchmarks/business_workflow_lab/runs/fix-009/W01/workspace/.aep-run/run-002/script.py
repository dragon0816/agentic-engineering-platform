from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True): print(row)
