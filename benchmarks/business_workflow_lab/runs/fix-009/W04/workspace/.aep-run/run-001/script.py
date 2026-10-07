from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size)
# workbook structure
wbp=ws/'input'/'sales_template.xlsx'
wb=load_workbook(wbp, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print(f'[{s}] size {sh.max_row}x{sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(' ', r)
# fx rates
print('\nfx_rates sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(10), reader):
        print(row)
# manual
print('\nsap_manual.md:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[:4000])