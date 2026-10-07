from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
print('\nDownloaded CSV:')
p=ws/'output'/'downloads'/'orders_2026-09_open.csv'
print(p.exists(), p if p.exists() else '')
if p.exists():
    print(p.read_text(encoding='utf-8-sig')[:2000])