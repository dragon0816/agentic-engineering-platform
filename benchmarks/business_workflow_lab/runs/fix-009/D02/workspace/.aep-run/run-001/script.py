from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
# workbook sheets/headers/sample rows
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(r)
# csv headers/sample
csv_path=ws/'input'/'fx_rates.csv'
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader):
        print('FX row', i+1, row)
# manual preview
md_path=ws/'input'/'expense_manual.md'
print('manual preview:')
print(md_path.read_text(encoding='utf-8')[:4000])