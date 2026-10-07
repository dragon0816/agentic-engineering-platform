from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('output files:', [p.name for p in (ws/'output').iterdir()])
for name in ['submitted_2026_09.csv','approved_2026_09.csv']:
    print('\nCSV', name)
    with open(ws/'output'/name, newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(12), r):
            print(i+1, row)
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
print('\nInstructions full:')
for row in sh.iter_rows(values_only=True):
    print(row)
