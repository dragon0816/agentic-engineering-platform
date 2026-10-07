from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment, Protection
from openpyxl.utils import get_column_letter
from decimal import Decimal, ROUND_HALF_UP
import csv, shutil, datetime
from collections import defaultdict

workspace = Path.cwd()
in_dir = workspace / 'input'
out_dir = workspace / 'output'
out_dir.mkdir(exist_ok=True)

template = in_dir / 'expense_template.xlsx'
submitted_csv = out_dir / 'submitted_2026_09.csv'
approved_csv = out_dir / 'approved_2026_09.csv'
fx_csv = in_dir / 'fx_rates.csv'
outfile = out_dir / 'expense_2026_09_submitted_approved.xlsx'

print('Inputs really are:')
print(' template:', template, template.exists())
print(' submitted csv:', submitted_csv, submitted_csv.exists())
print(' approved csv:', approved_csv, approved_csv.exists())
print(' fx csv:', fx_csv, fx_csv.exists())
for p in [submitted_csv, approved_csv]:
    with open(p, newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        rows = list(reader)
    print(p.name, 'headers:', rows[0])
    print(p.name, 'row_count:', len(rows)-1)
    print(p.name, 'sample:', rows[1] if len(rows)>1 else None)

# Load FX rates; the Note column may contain unquoted commas in this synthetic file, so use first two fields only.
rates = {}
with open(fx_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx headers:', header)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0]] = Decimal(row[1])
print('fx rates:', rates)

# Read exported data.
records = []
for p in [submitted_csv, approved_csv]:
    with open(p, newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for row in reader:
            records.append(row)

# Enforce requested period/status and stable chronological ordering.
from_date = datetime.date(2026, 9, 1)
to_date = datetime.date(2026, 9, 30)
statuses = {'Submitted', 'Approved'}
for r in records:
    d = datetime.date.fromisoformat(r['ExpenseDate'])
    assert from_date <= d <= to_date, f'out of period: {r}'
    assert r['Status'] in statuses, f'unrequested status: {r}'
    assert 'ReceiptID' in r and 'Description' in r and 'MissingReceipt' in r
records.sort(key=lambda r: (r['ExpenseDate'], r['ExpenseID']))
print('combined record count:', len(records))

# Helpers
def copy_row_style(ws, src_row, dst_row, max_col):
    # If template row does not have data row styles, copy header style lightly is avoided.
    # We preserve workbook-level formatting and set only number formats/alignment needed by instructions.
    for col in range(1, max_col+1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = src._style.copy()
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = src.alignment.copy()

def bool_text(v):
    return str(v).strip().lower() == 'true'

def amount_twd(amount_s, currency):
    a = Decimal(str(amount_s))
    if currency not in rates:
        raise ValueError(f'Missing FX rate for {currency}')
    return (a * rates[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# Fill a copy of the provided template, not a rebuilt workbook.
shutil.copyfile(template, outfile)
wb = load_workbook(outfile)

# Report
report = wb['Report']
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'ExpenseDate 2026-09-01 through 2026-09-30; Status in Submitted, Approved'
now = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B5'] = now
report['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D7'] = now
missing_ids = [r['ExpenseID'] for r in records if bool_text(r['MissingReceipt'])]
report['B9'] = f"Missing receipt exceptions listed in Exceptions: {', '.join(missing_ids)}" if missing_ids else 'No missing receipt exceptions.'

# Details
headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
details = wb['Details']
actual_headers = [details.cell(4, c).value for c in range(1, len(headers)+1)]
assert actual_headers == headers, actual_headers
# Clear old rows if present after row 4
for row in details.iter_rows(min_row=5, max_row=max(details.max_row, 5), max_col=len(headers)):
    for cell in row:
        cell.value = None

for idx, r in enumerate(records, start=5):
    amt = Decimal(r['Amount']).quantize(Decimal('0.01'))
    twd = amount_twd(r['Amount'], r['Currency'])
    missing = bool_text(r['MissingReceipt'])
    vals = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['ExpenseDate'],
        r['Category'], r['Currency'], float(amt), float(twd), r['Status'],
        r['ReceiptID'] if r['ReceiptID'] != '' else None, r['Description'], str(missing).lower(), str(missing).lower()
    ]
    for c, val in enumerate(vals, start=1):
        cell = details.cell(idx, c, val)
        if c in [1,2,4,7,10,11,12,13,14]:
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in [8,9]:
            cell.number_format = '0.00'
    details.cell(idx, 12).alignment = Alignment(wrap_text=True, vertical='top')
# Autofilter if template has one? Keep/extend if present or set to data area.
details.auto_filter.ref = f"A4:N{4+len(records)}"

# Summary: by Status and Currency. Group=Status, SubGroup left blank, Currency=currency.
summary = wb['Summary']
summary['B2'] = f"=SUM(Details!I5:I{4+len(records)})" if records else '=SUM(Details!I5:I5)'
for row in summary.iter_rows(min_row=5, max_row=max(summary.max_row, 5), max_col=6):
    for cell in row:
        cell.value = None
summary_data = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
for r in records:
    key = (r['Status'], r['Currency'])
    summary_data[key]['amount'] += Decimal(r['Amount'])
    summary_data[key]['twd'] += amount_twd(r['Amount'], r['Currency'])
    summary_data[key]['count'] += 1
summary_rows = sorted(summary_data.items(), key=lambda kv: (kv[0][0], kv[0][1]))
for idx, ((status, curr), v) in enumerate(summary_rows, start=5):
    vals = [status, None, curr, float(v['amount'].quantize(Decimal('0.01'))), float(v['twd'].quantize(Decimal('0.01'))), v['count']]
    for c, val in enumerate(vals, start=1):
        cell = summary.cell(idx, c, val)
        if c in [1,2,3]: cell.number_format = '@'
        if c in [4,5]: cell.number_format = '0.00'
summary.auto_filter.ref = f"A4:F{4+len(summary_rows)}"

# Exceptions: missing receipts, preserving in Details.
exceptions = wb['Exceptions']
for row in exceptions.iter_rows(min_row=5, max_row=max(exceptions.max_row, 5), max_col=4):
    for cell in row:
        cell.value = None
exc_rows = [r for r in records if bool_text(r['MissingReceipt'])]
for idx, r in enumerate(exc_rows, start=5):
    vals = [r['ExpenseID'], r['EmployeeID'], r['ExpenseDate'], 'Missing receipt']
    for c, val in enumerate(vals, start=1):
        cell = exceptions.cell(idx, c, val)
        if c in [1,2,4]: cell.number_format = '@'
        elif c == 3: cell.number_format = 'yyyy-mm-dd'
exceptions.auto_filter.ref = f"A4:D{4+len(exc_rows)}" if exc_rows else 'A4:D4'

# Basic sheet calculation metadata
wb.calculation.fullCalcOnLoad = True
wb.calculation.forceFullCalc = True

wb.save(outfile)
print('wrote', outfile)

# Reopen and verify stated requirements from saved output.
wb2 = load_workbook(outfile, data_only=False)
assert wb2.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
d = wb2['Details']
# Count non-empty detail rows
out_records = []
for row in d.iter_rows(min_row=5, max_row=d.max_row, max_col=14, values_only=True):
    if row[0] is None:
        continue
    out_records.append(row)
print('verify detail rows:', len(out_records))
assert len(out_records) == 22, len(out_records)
assert [d.cell(4, c).value for c in range(1,15)] == headers
# IDs textual, amounts numeric, detail columns filled where available, missing retained
ids = [r[0] for r in out_records]
assert 'EXP-00015' in ids and 'EXP-00006' in ids
assert all(isinstance(r[0], str) for r in out_records)
assert all(isinstance(r[7], (int,float)) for r in out_records)
assert all(isinstance(r[8], (int,float)) for r in out_records)
assert all(r[9] in ['Submitted','Approved'] for r in out_records)
assert all(str(r[12]).lower() in ['true','false'] for r in out_records)
# Requested detail fields present in Details headers and populated per source (ReceiptID blank if missing)
for r in out_records:
    if str(r[12]).lower() == 'false':
        assert r[10] not in [None, ''], f'missing receipt id for non-missing {r[0]}'
    assert r[11] not in [None, ''], f'missing description for {r[0]}'
# Exceptions
ex = wb2['Exceptions']
exc = [row for row in ex.iter_rows(min_row=5, max_row=ex.max_row, max_col=4, values_only=True) if row[0] is not None]
print('verify exception rows:', len(exc), exc)
assert len(exc) == 2
assert sorted([r[0] for r in exc]) == ['EXP-00006','EXP-00015']
# Summary and formula
s = wb2['Summary']
print('summary B2 formula:', s['B2'].value)
assert s['B2'].value == '=SUM(Details!I5:I26)'
summary_nonempty = [row for row in s.iter_rows(min_row=5, max_row=s.max_row, max_col=6, values_only=True) if row[0] is not None]
print('summary rows:', summary_nonempty)
# Totals check
expected_total = sum(amount_twd(r['Amount'], r['Currency']) for r in records).quantize(Decimal('0.01'))
actual_total = sum(Decimal(str(r[8])) for r in out_records).quantize(Decimal('0.01'))
print('expected AmountTWD total:', expected_total, 'actual:', actual_total)
assert expected_total == actual_total
# Report period/source cells
rp = wb2['Report']
assert rp['B3'].value == '2026-09-01' and rp['D3'].value == '2026-09-30'
assert 'Submitted' in rp['B4'].value and 'Approved' in rp['B4'].value
print('verified output workbook:', outfile.name)
