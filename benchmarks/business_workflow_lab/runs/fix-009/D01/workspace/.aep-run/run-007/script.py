from pathlib import Path
from openpyxl import load_workbook
import csv, json, os

ws = Path.cwd()
print('workspace:', ws)
print('input files:', sorted([p.name for p in (ws/'input').glob('*')]))
print('output files:', sorted([p.name for p in (ws/'output').glob('*')]))

report = ws/'output'/'E001_september_expenses.xlsx'
export = ws/'output'/'E001_september_expenses_export.xlsx'
for p in [report, export]:
    print(p.name, 'exists=', p.exists(), 'size=', p.stat().st_size if p.exists() else None)

# Reopen and validate the final report workbook.
wb = load_workbook(report, data_only=False)
print('report sheets:', wb.sheetnames)
expected_sheets = ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
assert wb.sheetnames == expected_sheets, f'unexpected sheets {wb.sheetnames}'

details = wb['Details']
summary = wb['Summary']
exceptions = wb['Exceptions']

# Read Details assuming headers on row 4 and data from row 5, as used in the filled template.
headers = [details.cell(row=4, column=c).value for c in range(1, details.max_column+1)]
print('details headers:', headers)
rows = []
for r in range(5, details.max_row+1):
    vals = [details.cell(row=r, column=c).value for c in range(1, details.max_column+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('details row count:', len(rows))
assert len(rows) == 13, f'expected 13 detail rows, got {len(rows)}'

# Flexible column lookup.
def find_col(possible):
    for name in possible:
        if name in headers:
            return name
    return None
emp_col = find_col(['Employee ID','Employee','employee_id','員工編號'])
date_col = find_col(['Date','Expense Date','date','日期'])
status_col = find_col(['Status','status','狀態'])
currency_col = find_col(['Currency','currency','幣別'])
amount_col = find_col(['Amount','Original Amount','amount','金額'])
twd_col = find_col(['TWD Amount','Amount TWD','TWD','台幣金額'])
expense_id_col = find_col(['Expense ID','ID','ExpenseId','expense_id'])
print('detected columns:', {'employee': emp_col, 'date': date_col, 'status': status_col, 'currency': currency_col, 'amount': amount_col, 'twd': twd_col, 'expense_id': expense_id_col})
assert emp_col and date_col and status_col and currency_col and amount_col and twd_col

from datetime import datetime, date
for row in rows:
    assert str(row[emp_col]) == 'E001', row
    d = row[date_col]
    if isinstance(d, datetime):
        dd = d.date()
    elif isinstance(d, date):
        dd = d
    else:
        dd = datetime.fromisoformat(str(d)[:10]).date()
    assert date(2026,9,1) <= dd <= date(2026,9,30), row
    assert isinstance(row[amount_col], (int,float)), (amount_col, row[amount_col], type(row[amount_col]))
    assert isinstance(row[twd_col], (int,float)), (twd_col, row[twd_col], type(row[twd_col]))
    if expense_id_col:
        # openpyxl sees text IDs as str.
        assert isinstance(row[expense_id_col], str), (expense_id_col, row[expense_id_col], type(row[expense_id_col]))

# Summary checks.
summary_rows = []
summary_headers = [summary.cell(row=4, column=c).value for c in range(1, summary.max_column+1)]
for r in range(5, summary.max_row+1):
    vals = [summary.cell(row=r, column=c).value for c in range(1, summary.max_column+1)]
    if any(v is not None for v in vals):
        summary_rows.append(vals)
print('summary headers:', summary_headers)
print('summary data rows:', len(summary_rows))
print('summary B2:', summary['B2'].value)
assert len(summary_rows) == 7, f'expected 7 summary rows, got {len(summary_rows)}'
assert summary['B2'].value == '=SUM(Details!I5:I17)', summary['B2'].value

# Exceptions checks.
exception_rows = []
exc_headers = [exceptions.cell(row=4, column=c).value for c in range(1, exceptions.max_column+1)]
for r in range(5, exceptions.max_row+1):
    vals = [exceptions.cell(row=r, column=c).value for c in range(1, exceptions.max_column+1)]
    if any(v is not None for v in vals):
        exception_rows.append(vals)
print('exceptions headers:', exc_headers)
print('exceptions data rows:', len(exception_rows))
assert len(exception_rows) == 1, f'expected 1 exception row, got {len(exception_rows)}'

# Data-only formula value may not be cached by openpyxl, so compute the TWD total from Details.
twd_total = round(sum(float(row[twd_col]) for row in rows), 2)
print('computed TWD total:', f'{twd_total:.2f}')
assert abs(twd_total - 31877.90) < 0.01, twd_total

# Reopen export workbook and validate the exported rows.
exwb = load_workbook(export, data_only=True)
print('export sheets:', exwb.sheetnames)
exws = exwb[exwb.sheetnames[0]]
ex_headers = [exws.cell(row=1, column=c).value for c in range(1, exws.max_column+1)]
ex_rows = []
for r in range(2, exws.max_row+1):
    vals = [exws.cell(row=r, column=c).value for c in range(1, exws.max_column+1)]
    if any(v is not None for v in vals):
        ex_rows.append(dict(zip(ex_headers, vals)))
print('export headers:', ex_headers)
print('export row count:', len(ex_rows))
assert len(ex_rows) == 13, f'expected 13 exported rows, got {len(ex_rows)}'
print('CHECK_OK')