from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date
from collections import defaultdict

ws = Path.cwd()
print('workspace:', ws)
print('input files:', sorted([p.name for p in (ws/'input').glob('*')]))
print('output files:', sorted([p.name for p in (ws/'output').glob('*')]))

report = ws/'output'/'E001_september_expenses.xlsx'
export = ws/'output'/'E001_september_expenses_export.xlsx'
assert report.exists(), f'missing {report}'
assert export.exists(), f'missing {export}'
print('report size:', report.stat().st_size)
print('export size:', export.stat().st_size)

# ---------- Validate final workbook ----------
wb = load_workbook(report, data_only=False)
expected_sheets = ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
print('report sheets:', wb.sheetnames)
assert wb.sheetnames == expected_sheets, f'unexpected sheet names/order: {wb.sheetnames}'

details = wb['Details']
summary = wb['Summary']
exceptions = wb['Exceptions']

headers = [details.cell(row=4, column=c).value for c in range(1, details.max_column+1)]
print('details headers:', headers)
rows = []
for r in range(5, details.max_row+1):
    vals = [details.cell(row=r, column=c).value for c in range(1, details.max_column+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('details row count:', len(rows))
assert len(rows) == 13, f'expected 13 detail rows, got {len(rows)}'

required = ['ExpenseID', 'EmployeeID', 'EmployeeName', 'ProjectID', 'ExpenseDate', 'Category', 'Currency', 'Amount', 'AmountTWD', 'Status', 'ReceiptID', 'Description', 'MissingReceipt', 'NeedsReview']
for h in required:
    assert h in headers, f'missing detail column {h}'

for row in rows:
    assert row['EmployeeID'] == 'E001', row
    d = row['ExpenseDate']
    if isinstance(d, datetime):
        dd = d.date()
    elif isinstance(d, date):
        dd = d
    else:
        dd = datetime.fromisoformat(str(d)[:10]).date()
    assert date(2026, 9, 1) <= dd <= date(2026, 9, 30), row
    assert isinstance(row['ExpenseID'], str), (row['ExpenseID'], type(row['ExpenseID']))
    assert isinstance(row['EmployeeID'], str), (row['EmployeeID'], type(row['EmployeeID']))
    assert isinstance(row['ProjectID'], str), (row['ProjectID'], type(row['ProjectID']))
    assert isinstance(row['Amount'], (int, float)), (row['Amount'], type(row['Amount']))
    assert isinstance(row['AmountTWD'], (int, float)), (row['AmountTWD'], type(row['AmountTWD']))

for r in range(5, 5+len(rows)):
    assert details.cell(r, 1).data_type == 's', f'ExpenseID not text at row {r}'
    assert details.cell(r, 2).data_type == 's', f'EmployeeID not text at row {r}'
    assert details.cell(r, 4).data_type == 's', f'ProjectID not text at row {r}'
    assert details.cell(r, 8).data_type == 'n', f'Amount not numeric at row {r}'
    assert details.cell(r, 9).data_type == 'n', f'AmountTWD not numeric at row {r}'

# Summary by Status + Currency. Template uses Group/SubGroup; inspect actual values and accept Status in either column.
expected_groups = defaultdict(lambda: {'count': 0, 'amount': 0.0, 'twd': 0.0})
for row in rows:
    key = (row['Status'], row['Currency'])
    expected_groups[key]['count'] += 1
    expected_groups[key]['amount'] += float(row['Amount'])
    expected_groups[key]['twd'] += float(row['AmountTWD'])

summary_headers = [summary.cell(row=4, column=c).value for c in range(1, summary.max_column+1)]
summary_rows_raw = []
for r in range(5, summary.max_row+1):
    vals = [summary.cell(row=r, column=c).value for c in range(1, summary.max_column+1)]
    if any(v is not None for v in vals):
        summary_rows_raw.append(dict(zip(summary_headers, vals)))
print('summary headers:', summary_headers)
print('summary data rows:', len(summary_rows_raw))
print('summary rows:', summary_rows_raw)
print('summary B2:', summary['B2'].value)
assert len(summary_rows_raw) == 7, f'expected 7 summary rows, got {len(summary_rows_raw)}'
assert summary['B2'].value == '=SUM(Details!I5:I17)', f'unexpected TWD formula {summary["B2"].value}'
for h in ['Group', 'SubGroup', 'Currency', 'Amount', 'AmountTWD', 'Count']:
    assert h in summary_headers, f'missing summary column {h}'

status_values = set(row['Status'] for row in rows)
seen_groups = {}
for row in summary_rows_raw:
    candidates = [row.get('Group'), row.get('SubGroup')]
    matching_status = [v for v in candidates if v in status_values]
    assert len(matching_status) == 1, f'cannot identify status in summary row {row}'
    key = (matching_status[0], row['Currency'])
    seen_groups[key] = row
assert set(seen_groups) == set(expected_groups), f'summary groups differ: {set(seen_groups)} vs {set(expected_groups)}'
for key, exp in expected_groups.items():
    row = seen_groups[key]
    assert int(row['Count']) == exp['count'], (key, row, exp)
    assert abs(float(row['Amount']) - round(exp['amount'], 2)) < 0.01, (key, row, exp)
    assert abs(float(row['AmountTWD']) - round(exp['twd'], 2)) < 0.01, (key, row, exp)

# Exceptions.
exc_headers = [exceptions.cell(row=4, column=c).value for c in range(1, exceptions.max_column+1)]
exception_rows = []
for r in range(5, exceptions.max_row+1):
    vals = [exceptions.cell(row=r, column=c).value for c in range(1, exceptions.max_column+1)]
    if any(v is not None for v in vals):
        exception_rows.append(dict(zip(exc_headers, vals)))
print('exceptions headers:', exc_headers)
print('exceptions data rows:', len(exception_rows))
print('exceptions rows:', exception_rows)
assert len(exception_rows) == 1, f'expected 1 exception row, got {len(exception_rows)}'
assert any('Missing receipt' in str(v) for v in exception_rows[0].values()), exception_rows[0]

twd_total = round(sum(float(row['AmountTWD']) for row in rows), 2)
print('computed TWD total:', f'{twd_total:.2f}')
assert abs(twd_total - 31877.90) < 0.01, twd_total

# ---------- Validate exported workbook ----------
exwb = load_workbook(export, data_only=True)
print('export sheets:', exwb.sheetnames)
exws = exwb[exwb.sheetnames[0]]
ex_headers = [exws.cell(row=1, column=c).value for c in range(1, exws.max_column+1)]
ex_rows = []
for r in range(2, exws.max_row+1):
    vals = [exws.cell(row=r, column=c).value for c in range(1, exws.max_column+1)]
    if any(v is not None for v in vals):
        ex_rows.append(dict(zip(ex_headers, vals)))
print('export headers:', ex_headers)
print('export row count:', len(ex_rows))
assert len(ex_rows) == 13, f'expected 13 exported rows, got {len(ex_rows)}'
for row in ex_rows:
    assert row['EmployeeID'] == 'E001', row
    d = row['ExpenseDate']
    if isinstance(d, datetime):
        dd = d.date()
    elif isinstance(d, date):
        dd = d
    else:
        dd = datetime.fromisoformat(str(d)[:10]).date()
    assert date(2026, 9, 1) <= dd <= date(2026, 9, 30), row

print('CHECK_OK')