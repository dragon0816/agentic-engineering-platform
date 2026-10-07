from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
exp = ws/'output'/'E001_september_expenses_export.xlsx'
print('export exists', exp.exists(), exp)
if exp.exists():
    wb = load_workbook(exp, data_only=True)
    print('export sheets', wb.sheetnames)
    sh = wb[wb.sheetnames[0]]
    print('export max', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(6, sh.max_row), values_only=True): print(row)

manual = (ws/'input'/'expense_manual.md').read_text(encoding='utf-8')
print('\nmanual full:\n', manual)

wb = load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
inst = wb['Instructions']
print('\nInstructions all rows:')
for row in inst.iter_rows(values_only=True): print(row)