import os, json, csv, shutil
from pathlib import Path
from datetime import datetime, timezone, timedelta
from collections import defaultdict
from decimal import Decimal
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
employee_id = params['employee_id']
date_from = params['date_from']
date_to = params['date_to']
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)
export_path = output_dir / f'{employee_id}_{date_from}_{date_to}_expenses_export.xlsx'
report_path = output_dir / f'{employee_id}_{date_from}_{date_to}_expenses.xlsx'
template_path = input_dir / 'expense_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'

print('inputs:', employee_id, date_from, date_to)
print('expected export exists:', export_path.exists(), export_path)
if not export_path.exists():
    fallback_path = output_dir / 'tny.xlsx'
    print('fallback export exists:', fallback_path.exists(), fallback_path)
    if fallback_path.exists():
        shutil.copyfile(fallback_path, export_path)
        print('copied fallback export to expected path')
if not export_path.exists():
    raise FileNotFoundError(export_path)

fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        fx[row['Currency']] = Decimal(str(row['RateToTWD']))

wb_exp = load_workbook(export_path, data_only=True)
ws_exp = wb_exp[wb_exp.sheetnames[0]]
headers = [c.value for c in ws_exp[1]]
print('export headers:', headers)
rows = []
for r in ws_exp.iter_rows(min_row=2, values_only=True):
    if not any(v is not None for v in r):
        continue
    rec = dict(zip(headers, r))
    rows.append(rec)
print('export row count:', len(rows))

for rec in rows:
    if str(rec.get('EmployeeID')) != employee_id:
        raise AssertionError(f'unexpected EmployeeID {rec.get("EmployeeID")}')
    d = str(rec.get('ExpenseDate'))[:10]
    if not (date_from <= d <= date_to):
        raise AssertionError(f'out-of-range date {d}')

shutil.copyfile(template_path, report_path)
wb = load_workbook(report_path)
report = wb['Report']
details = wb['Details']
summary = wb['Summary']
exceptions = wb['Exceptions']

for sh in [details, summary, exceptions]:
    if sh.max_row > 4:
        sh.delete_rows(5, sh.max_row - 4)

now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B3'] = date_from
report['D3'] = date_to
report['B4'] = f'EmployeeID={employee_id}; ExpenseDate {date_from} to {date_to}; Status=All'
report['B5'] = now
report['B6'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D6'] = now

summary_map = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
exception_rows = []

for i, rec in enumerate(rows, start=5):
    cur = str(rec['Currency'])
    amount = Decimal(str(rec['Amount']))
    amount_twd = amount * fx[cur]
    status = str(rec['Status'])
    missing = rec.get('MissingReceipt') in (True, 'TRUE', 'True', 'true', 1, '1')
    values = [
        str(rec['ExpenseID']), str(rec['EmployeeID']), rec.get('EmployeeName'), rec.get('ProjectID'),
        str(rec['ExpenseDate'])[:10], rec.get('Category'), cur, float(amount), float(amount_twd),
        status, rec.get('ReceiptID'), rec.get('Description')
    ]
    for col, value in enumerate(values, start=1):
        details.cell(row=i, column=col, value=value)
    key = (status, cur)
    summary_map[key]['amount'] += amount
    summary_map[key]['twd'] += amount_twd
    summary_map[key]['count'] += 1
    if missing or not rec.get('ReceiptID'):
        exception_rows.append([str(rec['ExpenseID']), str(rec['EmployeeID']), str(rec['ExpenseDate'])[:10], 'Missing receipt'])

for row_num, ((status, cur), vals) in enumerate(sorted(summary_map.items()), start=5):
    summary.cell(row=row_num, column=1, value=status)
    summary.cell(row=row_num, column=2, value=None)
    summary.cell(row=row_num, column=3, value=cur)
    summary.cell(row=row_num, column=4, value=float(vals['amount']))
    summary.cell(row=row_num, column=5, value=float(vals['twd']))
    summary.cell(row=row_num, column=6, value=vals['count'])

last_detail = 4 + len(rows)
summary['B2'] = f'=SUM(Details!I5:I{last_detail})' if rows else '=0'
for row_num, vals in enumerate(exception_rows, start=5):
    for col, value in enumerate(vals, start=1):
        exceptions.cell(row=row_num, column=col, value=value)

wb.save(report_path)

check = load_workbook(report_path, data_only=False)
assert check.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
assert check['Details'].max_row == 4 + len(rows)
assert check['Summary']['B2'].value == (f'=SUM(Details!I5:I{last_detail})' if rows else '=0')
assert check['Report']['B3'].value == date_from
assert check['Report']['D3'].value == date_to
assert export_path.exists()
assert report_path.exists()
print('wrote export:', export_path)
print('wrote report:', report_path)
print('details rows:', len(rows))
print('summary groups:', len(summary_map))
print('exceptions:', len(exception_rows))
print('TWD total:', sum(Decimal(str(r['Amount'])) * fx[str(r['Currency'])] for r in rows))
