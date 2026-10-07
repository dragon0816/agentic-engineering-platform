from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
# workbook structure
wb_path = ws/'input'/'expense_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('template sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    wsht = wb[sh]
    print('\nSHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
        print(row)
# fx rates
fx = ws/'input'/'fx_rates.csv'
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    rows = list(reader)
print('\nfx rows sample:')
for r in rows[:8]: print(r)
# manual snippet
manual = (ws/'input'/'expense_manual.md').read_text(encoding='utf-8')
print('\nmanual first 1200 chars:\n', manual[:1200])