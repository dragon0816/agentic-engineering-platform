from pathlib import Path
from openpyxl import load_workbook
import csv

ws = Path.cwd()
print('WORKSPACE', ws)
for sub in ['input','output']:
    p = ws/sub
    print(f'[{sub}] exists={p.exists()}')
    if p.exists():
        for f in sorted(p.iterdir()):
            print(' -', f.name, f.stat().st_size, 'bytes')

for rel in ['input/expense_template.xlsx', 'output/E001_september_expenses.xlsx', 'output/E001_september_expenses_export.xlsx']:
    p = ws/rel
    print('\nFILE', rel, 'exists=', p.exists())
    if p.exists() and p.suffix.lower() == '.xlsx':
        wb = load_workbook(p, data_only=False)
        print(' sheets=', wb.sheetnames)
        for s in wb.sheetnames:
            sh = wb[s]
            print(f'  {s}: max_row={sh.max_row}, max_col={sh.max_column}')
            # print first few non-empty rows
            n = 0
            for row in sh.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    print('   ', row[:min(len(row), 12)])
                    n += 1
                    if n >= 5:
                        break

p = ws/'input/fx_rates.csv'
print('\nFILE input/fx_rates.csv exists=', p.exists())
if p.exists():
    with p.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i,row in zip(range(5), r):
            print(' ', row)
