from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime
from copy import copy
from openpyxl import load_workbook

ws = Path.cwd()
input_dir = ws/'input'
out_dir = ws/'output'
downloads = out_dir/'downloads'
out_dir.mkdir(exist_ok=True)

print('Inputs:')
for p in sorted(input_dir.glob('*')):
    print(' -', p.name)

template = input_dir/'project_template.xlsx'
invoice_csv = downloads/'invoices_2026-09.csv'
fx_csv = input_dir/'fx_rates.csv'
print('Invoice CSV:', invoice_csv, 'exists=', invoice_csv.exists())

rates = {}
with fx_csv.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if row and row[0].strip():
            rates[row[0].strip()] = Decimal(row[1].strip())
print('rates:', rates)

with invoice_csv.open(newline='', encoding='utf-8-sig') as f:
    rdr = csv.DictReader(f)
    print('invoice headers:', rdr.fieldnames)
    inv_rows = list(rdr)
print('invoice row count:', len(inv_rows))
print('invoice sample rows:', inv_rows[:3])

from_date = date(2026,9,1)
to_date = date(2026,9,30)
asof = date(2026,9,30)
q = Decimal('0.01')
processed = []
for r in inv_rows:
    inv_date = date.fromisoformat(r['InvoiceDate'])
    if not (from_date <= inv_date <= to_date):
        continue
    amt = Decimal(r['Amount'])
    paid = Decimal(r['PaidAmount'])
    balance = Decimal(r['Balance'])
    calc_bal = (amt - paid).quantize(q, ROUND_HALF_UP)
    if calc_bal != balance.quantize(q, ROUND_HALF_UP):
        raise ValueError(f"Balance mismatch for {r['InvoiceID']}: {balance} vs {calc_bal}")
    ccy = r['Currency']
    if ccy not in rates:
        raise ValueError(f'Missing FX rate for {ccy}')
    bal_twd = (balance * rates[ccy]).quantize(q, ROUND_HALF_UP)
    due = date.fromisoformat(r['DueDate'])
    out = dict(r)
    out['Amount_dec'] = amt.quantize(q, ROUND_HALF_UP)
    out['PaidAmount_dec'] = paid.quantize(q, ROUND_HALF_UP)
    out['Balance_dec'] = balance.quantize(q, ROUND_HALF_UP)
    out['BalanceTWD_dec'] = bal_twd
    out['InvoiceDate_date'] = inv_date
    out['DueDate_date'] = due
    out['Overdue'] = balance > 0 and due < asof
    processed.append(out)
print('processed in period:', len(processed))

def copy_row_style(sh, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = sh.cell(src_row, col)
        dst = sh.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.font = copy(src.font)

def clear_data(sh, max_col):
    if sh.max_row > 4:
        for row in sh.iter_rows(min_row=5, max_row=sh.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

wb = load_workbook(template)
print('Workbook sheets:', wb.sheetnames)
print('Invoices row 4 A:K:', [wb['Invoices'].cell(4,c).value for c in range(1,12)])
if wb['Invoices'].cell(4,11).value in (None, ''):
    wb['Invoices'].cell(4,11).value = 'BalanceTWD'
    src = wb['Invoices'].cell(4,10); dst = wb['Invoices'].cell(4,11)
    if src.has_style: dst._style = copy(src._style)

for sheet_name, max_col in [('Orders',8), ('Expenses',8), ('Quotes',9)]:
    clear_data(wb[sheet_name], max_col)

sh = wb['Invoices']
clear_data(sh, 11)
for idx, r in enumerate(processed, start=5):
    copy_row_style(sh, 4, idx, 11)
    vals = [r['InvoiceID'], r['OrderID'], r['ProjectID'] or None, r['CustomerID'], r['InvoiceDate_date'], r['DueDate_date'], r['Currency'], float(r['Amount_dec']), float(r['PaidAmount_dec']), float(r['Balance_dec']), float(r['BalanceTWD_dec'])]
    for col, val in enumerate(vals, start=1):
        cell = sh.cell(idx, col, val)
        if col in (1,2,3,4):
            cell.number_format = '@'
        elif col in (5,6):
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9,10,11):
            cell.number_format = '0.00'

summary = {}
for r in processed:
    pid = (r['ProjectID'] or '').strip()
    if not pid:
        continue
    summary.setdefault(pid, Decimal('0.00'))
    summary[pid] += r['BalanceTWD_dec']

sum_sh = wb['Summary']
clear_data(sum_sh, 5)
for idx, pid in enumerate(sorted(summary), start=5):
    copy_row_style(sum_sh, 4, idx, 5)
    vals = [pid, 0.00, float(summary[pid].quantize(q, ROUND_HALF_UP)), 0.00, 0.00]
    for col, val in enumerate(vals, 1):
        c = sum_sh.cell(idx, col, val)
        c.number_format = '@' if col == 1 else '0.00'
last_summary_row = max(5, 4 + len(summary))
for col_letter in ['B','C','D','E']:
    sum_sh[f'{col_letter}2'] = f'=SUM({col_letter}5:{col_letter}{last_summary_row})'
    sum_sh[f'{col_letter}2'].number_format = '0.00'

ex_sh = wb['Exceptions']
clear_data(ex_sh, 4)
overdue_rows = [r for r in processed if r['Overdue']]
for idx, r in enumerate(overdue_rows, start=5):
    copy_row_style(ex_sh, 4, idx, 4)
    vals = ['Invoices', r['InvoiceID'], r['ProjectID'] or None, 'Overdue']
    for col, val in enumerate(vals, 1):
        c = ex_sh.cell(idx, col, val)
        if col in (2,3): c.number_format = '@'

rep = wb['Report']
rep['B3'] = from_date
rep['D3'] = to_date
rep['B3'].number_format = 'yyyy-mm-dd'
rep['D3'].number_format = 'yyyy-mm-dd'
rep['B4'] = 'Invoices: InvoiceDate 2026-09-01 to 2026-09-30; as-of 2026-09-30; overdue if Balance > 0 and DueDate < 2026-09-30'
now = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B5'] = now
rep['B7'] = 'Lab ERP Invoices; fx_rates.csv'
rep['D7'] = now
blank_project_count = sum(1 for r in processed if not (r['ProjectID'] or '').strip())
rep['B9'] = f'Orders/Expenses/Quotes left blank per request. {blank_project_count} invoice row(s) have blank ProjectID and are not included in project Summary.'

out_file = out_dir/'project_invoice_report_2026-09.xlsx'
wb.save(out_file)
print('saved:', out_file)

# Verification
wb2 = load_workbook(out_file, data_only=False)
errors=[]
inv = wb2['Invoices']
headers = [inv.cell(4,c).value for c in range(1,12)]
expected_headers = ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
print('Invoices headers A:K:', headers)
if headers != expected_headers:
    errors.append(f'Invoice headers mismatch: {headers}')
inv_count = sum(1 for r in range(5, inv.max_row+1) if inv.cell(r,1).value not in (None,''))
if inv_count != len(processed) or inv_count != 13:
    errors.append(f'Expected 13 invoice rows, got {inv_count}')

def cell_date(v):
    if isinstance(v, datetime):
        return v.date()
    return v
for r in range(5, 5+inv_count):
    if not isinstance(inv.cell(r,1).value, str) or not inv.cell(r,1).value.startswith('009'):
        errors.append(f'InvoiceID not text at row {r}: {inv.cell(r,1).value!r}')
    if not isinstance(inv.cell(r,2).value, str) or not inv.cell(r,2).value.startswith('004'):
        errors.append(f'OrderID not text at row {r}: {inv.cell(r,2).value!r}')
    invd = cell_date(inv.cell(r,5).value)
    due = cell_date(inv.cell(r,6).value)
    if not (from_date <= invd <= to_date):
        errors.append(f'InvoiceDate out of range row {r}: {invd}')
    if not isinstance(due, date):
        errors.append(f'DueDate not date row {r}: {due!r}')
    for c in (8,9,10,11):
        if not isinstance(inv.cell(r,c).value, (int,float)):
            errors.append(f'Amount column {c} row {r} not numeric: {inv.cell(r,c).value!r}')

for sheet_name in ['Orders','Expenses','Quotes']:
    shx=wb2[sheet_name]
    filled=[]
    for row in shx.iter_rows(min_row=5, max_row=shx.max_row):
        if any(c.value not in (None,'') for c in row):
            filled.append(row[0].row)
    if filled:
        errors.append(f'{sheet_name} has data rows: {filled[:10]}')

ex=wb2['Exceptions']
ex_rows=[]
for r in range(5, ex.max_row+1):
    if ex.cell(r,1).value not in (None,''):
        ex_rows.append([ex.cell(r,c).value for c in range(1,5)])
expected_overdue_ids = [r['InvoiceID'] for r in overdue_rows]
actual_overdue_ids = [row[1] for row in ex_rows]
print('Expected overdue IDs:', expected_overdue_ids)
print('Actual exception rows:', ex_rows)
if actual_overdue_ids != expected_overdue_ids:
    errors.append(f'Overdue IDs mismatch expected {expected_overdue_ids} got {actual_overdue_ids}')
if any(row[3] != 'Overdue' or row[0] != 'Invoices' for row in ex_rows):
    errors.append('Exception source/reason mismatch')
# specifically confirm due date 2026-09-30 with balance positive is not in exceptions
not_overdue_930 = [r['InvoiceID'] for r in processed if r['Balance_dec'] > 0 and r['DueDate_date'] == asof]
print('Due on as-of date with balance, not overdue:', not_overdue_930)
if any(i in actual_overdue_ids for i in not_overdue_930):
    errors.append(f'Due on 2026-09-30 incorrectly marked overdue: {not_overdue_930}')

summ=wb2['Summary']
actual_summary={}
for r in range(5, summ.max_row+1):
    pid=summ.cell(r,1).value
    if pid not in (None,''):
        actual_summary[pid]=Decimal(str(summ.cell(r,3).value)).quantize(q)
expected_summary={pid: amt.quantize(q) for pid,amt in summary.items()}
print('Expected project invoice balance TWD:', expected_summary)
print('Actual project invoice balance TWD:', actual_summary)
if actual_summary != expected_summary:
    errors.append(f'Summary mismatch expected {expected_summary} got {actual_summary}')
for r in range(5, summ.max_row+1):
    if summ.cell(r,1).value not in (None,''):
        for c in [2,4,5]:
            if Decimal(str(summ.cell(r,c).value)).quantize(q) != Decimal('0.00'):
                errors.append(f'Summary non-invoice col row {r} col {c} not zero')
for col in ['B','C','D','E']:
    expected=f'=SUM({col}5:{col}{last_summary_row})'
    actual=summ[f'{col}2'].value
    print(f'Summary {col}2 formula:', actual)
    if actual != expected:
        errors.append(f'Formula {col}2 expected {expected}, got {actual}')

if errors:
    print('VERIFICATION FAILED')
    for e in errors:
        print(' -', e)
    raise SystemExit(1)
print('VERIFICATION PASSED')
print('Output file:', out_file.name)
print('Invoice rows:', inv_count)
print('Summary projects:', len(actual_summary))
print('Exceptions:', len(ex_rows))
