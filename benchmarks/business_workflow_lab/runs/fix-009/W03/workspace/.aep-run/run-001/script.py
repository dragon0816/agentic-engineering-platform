from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('cwd=', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# workbook structure
xlsx = ws/'input'/'project_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'\nSheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in range(1, min(sh.max_row, 8)+1):
        vals = [sh.cell(r,c).value for c in range(1, min(sh.max_column, 10)+1)]
        print(r, vals)

# fx_rates.csv
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv sample:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i,row in zip(range(10), reader):
            print(i+1, row)

# sap manual
md = ws/'input'/'sap_manual.md'
if md.exists():
    print('\nsap_manual.md first 120 lines:')
    text = md.read_text(encoding='utf-8-sig', errors='replace').splitlines()
    for i,line in enumerate(text[:120],1):
        print(f'{i}: {line}')
