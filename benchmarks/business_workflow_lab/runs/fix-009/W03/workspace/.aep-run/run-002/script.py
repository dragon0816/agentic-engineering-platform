from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
sh=wb['Instructions']
print('Instructions:')
for r in sh.iter_rows(values_only=True):
    print(r)
print('\nDownloaded CSV files:')
for p in sorted((ws/'output'/'downloads').glob('*')):
    print(p.name, p.stat().st_size)
    print(p.read_text(encoding='utf-8-sig')[:1000])
