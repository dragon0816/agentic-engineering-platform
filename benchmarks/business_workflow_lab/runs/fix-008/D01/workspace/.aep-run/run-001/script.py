from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
# workbook structure
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path)
print('template sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print('SHEET', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(' ', row)
# fx rates
fx=ws/'input'/'fx_rates.csv'
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader):
        print('fx row', i+1, row)
# manual preview
manual=ws/'input'/'expense_manual.md'
print('manual preview:')
print(manual.read_text(encoding='utf-8')[:2500])