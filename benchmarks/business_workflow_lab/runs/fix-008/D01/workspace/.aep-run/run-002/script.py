from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
export=ws/'output'/'E001_2026-09_expense_export.xlsx'
print('export exists', export.exists(), export.stat().st_size if export.exists() else None)
wb=load_workbook(export, data_only=False)
print('export sheets', wb.sheetnames)
for sh in wb.worksheets:
    print('SHEET', sh.title, sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(10,sh.max_row), values_only=True):
        print(row)
print('Instructions all rows:')
tpl=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
for row in tpl['Instructions'].iter_rows(values_only=True):
    print(row)