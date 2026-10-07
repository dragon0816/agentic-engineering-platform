import os
from pathlib import Path
from openpyxl import load_workbook
ws=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-008\D01\workspace')
out=ws/'output'
print('Output folder:', out)
for p in sorted(out.glob('*')):
    print(p.name, p.stat().st_size)
for name in ['E002_august_expenses.xlsx','E002_2026-08_expense_export.xlsx']:
    p=out/name
    print(name, 'exists=', p.exists())
    if p.exists():
        try:
            wb=load_workbook(p, data_only=False, read_only=True)
            print(' sheets=', wb.sheetnames)
            for s in wb.sheetnames:
                sh=wb[s]
                print(' ', s, 'max_row=', sh.max_row, 'max_col=', sh.max_column)
            wb.close()
        except Exception as e:
            print(' open_error=', repr(e))