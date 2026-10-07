from pathlib import Path
import subprocess, json, csv, shutil, sys, time
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
CLI_PY = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
GENERATED_AT = '2026-10-06'  # supplied environment date, Asia/Taipei
FETCHED_AT = GENERATED_AT
OUTFILE = OUTPUT / 'quote_report_2026-09_Sent.xlsx'
LOGFILE = OUTPUT / 'quote_cli_run_log.json'

Q2 = Decimal('0.01')
def d(value):
    if value is None or value == '':
        return None
    try:
        return Decimal(str(value))
    except InvalidOperation as e:
        raise ValueError(f'Invalid decimal value: {value!r}') from e

def money(value):
    dec = d(value)
    return None if dec is None else dec.quantize(Q2, rounding=ROUND_HALF_UP)

def to_number(dec):
    return None if dec is None else float(dec)

# Read fixed FX rates. The Note column contains unquoted commas in this test data;
# only the first two columns are part of the contract we need.
fx = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    if header[:2] != ['Currency', 'RateToTWD']:
        raise RuntimeError(f'Unexpected fx_rates.csv header: {header}')
    for row in reader:
        if not row or len(row) < 2:
            continue
        fx[row[0]] = d(row[1])
print('Loaded FX rates:', {k: str(v) for k, v in fx.items()})

if CLI.exists():
    base_cmd = [str(CLI)]
elif CLI_PY.exists():
    base_cmd = [sys.executable, str(CLI_PY)]
else:
    raise RuntimeError('quote_cli not found at documented paths')
print('CLI command base:', base_cmd)

run_log = []
def run_cli(args, parse_json=True):
    """Run quote CLI. Retry exit code 4 up to two retries; never retry 2/3."""
    attempts = 0
    while True:
        attempts += 1
        cmd = base_cmd + args
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
        entry = {
            'args': args,
            'attempt': attempts,
            'exit_code': proc.returncode,
            'stderr': proc.stderr.strip(),
            'stdout_len': len(proc.stdout),
        }
        run_log.append(entry)
        print(f"CLI attempt {attempts}: {' '.join(args)} -> exit {proc.returncode}")
        if proc.returncode == 4 and attempts <= 3:
            if attempts <= 2:
                print('Temporary service error; retrying because exit code 4:', proc.stderr.strip())
                time.sleep(0.2)
                continue
            # attempts == 3 means original + two retries exhausted
            raise RuntimeError(f'quote_cli temporary service error after 3 attempts. stderr: {proc.stderr.strip()}')
        if proc.returncode in (2, 3):
            raise RuntimeError(f'quote_cli returned exit code {proc.returncode}; not retried. stderr: {proc.stderr.strip()}')
        if proc.returncode != 0:
            raise RuntimeError(f'quote_cli failed with exit code {proc.returncode}. stderr: {proc.stderr.strip()}')
        if parse_json:
            try:
                return json.loads(proc.stdout)
            except json.JSONDecodeError as e:
                raise RuntimeError(f'quote_cli success stdout was not parseable JSON for args {args}. stdout was not used as data.') from e
        return proc.stdout

# Query September 2026 Sent quote headers.
listing = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'])
if listing.get('schema_version') != 1 or 'items' not in listing:
    raise RuntimeError(f'Unexpected list response schema: {listing}')
headers = listing['items']
if listing.get('count') != len(headers):
    raise RuntimeError(f"List count {listing.get('count')} does not match item length {len(headers)}")
print(f'List returned {len(headers)} quote headers')

# Fetch detail for each quote to obtain line items; use get header as authoritative complete detail.
quotes = []
lines = []
for h in headers:
    qid = h.get('QuoteID')
    if not qid:
        raise RuntimeError(f'Missing QuoteID in header: {h}')
    detail = run_cli(['quotes', 'get', str(qid), '--format', 'json'])
    if detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
        raise RuntimeError(f'Unexpected get response schema for {qid}: {detail}')
    qh = detail['header']
    if qh.get('QuoteID') != qid:
        raise RuntimeError(f'QuoteID mismatch between list {qid} and get {qh.get("QuoteID")}')
    quotes.append(qh)
    for li in detail['line_items']:
        row = dict(li)
        row['QuoteID'] = qid
        lines.append(row)
print(f'Fetched {len(quotes)} quote details and {len(lines)} line items')

# Sort for stable report output.
quotes.sort(key=lambda x: (x.get('QuoteDate') or '', x.get('QuoteID') or ''))
line_order = {q['QuoteID']: i for i, q in enumerate(quotes)}
lines.sort(key=lambda x: (line_order.get(x.get('QuoteID'), 999999), x.get('ProductID') or '', x.get('ProductName') or ''))

# Copy and fill the original template; do not rebuild workbook.
shutil.copy2(TEMPLATE, OUTFILE)
wb = load_workbook(OUTFILE)

# Helper: clear existing rows from row 5 down without touching headers/templates/layout.
def clear_rows(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def apply_cell_style_from_above(ws, row, max_col):
    # Prefer row 5 if it exists, else copy header row 4 style lightly.
    source_row = 5 if row != 5 and ws.max_row >= 5 else 4
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)

# Report metadata
wsr = wb['Report']
wsr['B3'] = FROM_DATE
wsr['D3'] = TO_DATE
wsr['B4'] = f'Status={STATUS}; QuoteDate between {FROM_DATE} and {TO_DATE} inclusive'
wsr['B5'] = GENERATED_AT
wsr['B7'] = 'quote_cli; fx_rates.csv'
wsr['D7'] = FETCHED_AT
wsr['B9'] = ''

# Quotes sheet
wsq = wb['Quotes']
clear_rows(wsq, 5)
quote_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
if [wsq.cell(4, c).value for c in range(1, 12)] != quote_headers:
    raise RuntimeError('Quotes row 4 headers do not match expected template')

for r_idx, q in enumerate(quotes, start=5):
    apply_cell_style_from_above(wsq, r_idx, 11)
    currency = q.get('Currency')
    if currency not in fx:
        raise RuntimeError(f'Missing FX rate for currency {currency!r}')
    subtotal = money(q.get('Subtotal'))
    discount = d(q.get('DiscountPercent'))
    net = money(q.get('NetAmount'))
    if subtotal is not None and discount is not None:
        expected_net = (subtotal * (Decimal('1') - discount / Decimal('100'))).quantize(Q2, rounding=ROUND_HALF_UP)
        if net != expected_net:
            raise RuntimeError(f"NetAmount mismatch for {q.get('QuoteID')}: CLI={net}, expected={expected_net}")
    net_twd = (net * fx[currency]).quantize(Q2, rounding=ROUND_HALF_UP) if net is not None else None
    values = [
        q.get('QuoteID'), q.get('CustomerID'), q.get('ProjectID'), q.get('QuoteDate'), q.get('ValidUntil'),
        q.get('Status'), currency, to_number(subtotal), to_number(discount), to_number(net), to_number(net_twd)
    ]
    for c, value in enumerate(values, start=1):
        cell = wsq.cell(r_idx, c, value)
        if c in (1,2,3,4,5,6,7):
            cell.number_format = '@' if c in (1,2,3) else 'yyyy-mm-dd' if c in (4,5) else 'General'
        if c in (8,10,11):
            cell.number_format = '0.00'
        if c == 9:
            cell.number_format = '0.00'

# Lines sheet
wsl = wb['Lines']
clear_rows(wsl, 5)
line_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if [wsl.cell(4, c).value for c in range(1, 7)] != line_headers:
    raise RuntimeError('Lines row 4 headers do not match expected template')
for r_idx, li in enumerate(lines, start=5):
    apply_cell_style_from_above(wsl, r_idx, 6)
    qty = d(li.get('Quantity'))
    unit = money(li.get('UnitPrice'))
    amount = money(li.get('LineAmount'))
    if qty is not None and unit is not None:
        expected_amount = (qty * unit).quantize(Q2, rounding=ROUND_HALF_UP)
        if amount != expected_amount:
            raise RuntimeError(f"LineAmount mismatch for {li.get('QuoteID')} {li.get('ProductID')}: CLI={amount}, expected={expected_amount}")
    values = [li.get('QuoteID'), li.get('ProductID'), li.get('ProductName'), to_number(qty), to_number(unit), to_number(amount)]
    for c, value in enumerate(values, start=1):
        cell = wsl.cell(r_idx, c, value)
        if c in (1,2):
            cell.number_format = '@'
        elif c in (5,6):
            cell.number_format = '0.00'

# Summary: by Status (Group) and Currency, because the query condition fixes Status and currency totals must not be mixed.
wss = wb['Summary']
clear_rows(wss, 5)
last_quote_row = 5 if not quotes else 4 + len(quotes)
wss['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
wss['B2'].number_format = '0.00'
summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0})
for q in quotes:
    cur = q.get('Currency')
    net = money(q.get('NetAmount')) or Decimal('0.00')
    amt_twd = (net * fx[cur]).quantize(Q2, rounding=ROUND_HALF_UP)
    key = (q.get('Status') or '', '', cur)
    summary[key]['amount'] += net
    summary[key]['amount_twd'] += amt_twd
    summary[key]['count'] += 1
for r_idx, (key, val) in enumerate(sorted(summary.items()), start=5):
    apply_cell_style_from_above(wss, r_idx, 6)
    group, subgroup, cur = key
    values = [group, subgroup, cur, to_number(val['amount'].quantize(Q2, rounding=ROUND_HALF_UP)), to_number(val['amount_twd'].quantize(Q2, rounding=ROUND_HALF_UP)), val['count']]
    for c, value in enumerate(values, start=1):
        cell = wss.cell(r_idx, c, value)
        if c in (4,5):
            cell.number_format = '0.00'
        if c == 6:
            cell.number_format = '0'

wb.save(OUTFILE)

# Save CLI run log after successful workbook creation.
LOGFILE.write_text(json.dumps(run_log, ensure_ascii=False, indent=2), encoding='utf-8')
print('Wrote workbook:', OUTFILE)
print('Wrote CLI log:', LOGFILE)

# Reopen output and validate stated/template requirements.
wbv = load_workbook(OUTFILE, data_only=False)
errors = []
# required files/sheets/headers
for sheet in ['Report','Quotes','Lines','Summary','Instructions']:
    if sheet not in wbv.sheetnames:
        errors.append(f'Missing sheet {sheet}')
if not errors:
    vq = wbv['Quotes']; vl = wbv['Lines']; vs = wbv['Summary']; vr = wbv['Report']
    if [vq.cell(4, c).value for c in range(1, 12)] != quote_headers:
        errors.append('Quotes headers changed')
    if [vl.cell(4, c).value for c in range(1, 7)] != line_headers:
        errors.append('Lines headers changed')
    if vr['B3'].value != FROM_DATE or vr['D3'].value != TO_DATE:
        errors.append('Report dates not filled correctly')
    if STATUS not in str(vr['B4'].value):
        errors.append('Report condition does not mention Sent status')
    q_rows = max(0, vq.max_row - 4) if vq.max_row >= 5 else 0
    l_rows = max(0, vl.max_row - 4) if vl.max_row >= 5 else 0
    s_rows = max(0, vs.max_row - 4) if vs.max_row >= 5 else 0
    if q_rows != len(quotes):
        errors.append(f'Quotes row count {q_rows} != {len(quotes)}')
    if l_rows != len(lines):
        errors.append(f'Lines row count {l_rows} != {len(lines)}')
    if vs['B2'].value != f'=SUM(Quotes!K5:K{last_quote_row})':
        errors.append(f'Summary B2 formula wrong: {vs["B2"].value}')
    # check types/formats for populated detail rows
    for r in range(5, 5 + len(quotes)):
        for c in (1,2,3):
            if not isinstance(vq.cell(r,c).value, str):
                errors.append(f'Quotes {vq.cell(r,c).coordinate} ID not text')
        if vq.cell(r,6).value != STATUS:
            errors.append(f'Quotes {r} status is not Sent')
        if not (FROM_DATE <= str(vq.cell(r,4).value) <= TO_DATE):
            errors.append(f'Quotes {r} QuoteDate outside range')
        for c in (8,9,10,11):
            if not isinstance(vq.cell(r,c).value, (int,float)):
                errors.append(f'Quotes {vq.cell(r,c).coordinate} not numeric')
    for r in range(5, 5 + len(lines)):
        for c in (1,2):
            if not isinstance(vl.cell(r,c).value, str):
                errors.append(f'Lines {vl.cell(r,c).coordinate} ID not text')
        for c in (4,5,6):
            if not isinstance(vl.cell(r,c).value, (int,float)):
                errors.append(f'Lines {vl.cell(r,c).coordinate} not numeric')
    print('Validation counts:', {'quotes': q_rows, 'lines': l_rows, 'summary': s_rows})
    print('Validation formula Summary!B2:', vs['B2'].value)

if errors:
    print('VALIDATION ERRORS:')
    for e in errors:
        print('-', e)
    raise RuntimeError('Output validation failed')
print('Validation passed: workbook reopened; required sheets/headers, dates, condition, row counts, ID text cells, numeric amount cells, Sent status/date range, and Summary formula were checked.')
