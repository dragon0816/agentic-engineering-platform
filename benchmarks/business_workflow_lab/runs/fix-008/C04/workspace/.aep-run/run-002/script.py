from pathlib import Path
from openpyxl import load_workbook
p=Path('input/quote_template.xlsx')
wb=load_workbook(p, data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
    print(row)
print('\nSummary sheet formulas/styles:')
for row in wb['Summary'].iter_rows(min_row=1, max_row=10, values_only=False):
    print([c.coordinate+':'+repr(c.value) for c in row])
print('\nReport merged ranges:', list(wb['Report'].merged_cells.ranges))
