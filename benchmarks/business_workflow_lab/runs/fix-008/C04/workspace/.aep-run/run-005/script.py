from pathlib import Path
import subprocess, json, csv, shutil, sys, time
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
CLI_PY = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
FROM_DATE, TO_DATE, STATUS = '2026-09-01', '2026-09-30', 'Sent'
GENERATED_AT = FETCHED_AT = '2026-10-06'
OUTFILE = OUTPUT / 'quote_report_2026-09_Sent.xlsx'
LOGFILE = OUTPUT / 'quote_cli_run_log.json'
Q2 = Decimal('0.01')

def text_or_blank(v):
    return None if v is None or v == '' else str(v)
def dec(v):
    return None if v is None or v == '' else Decimal(str(v))
def money(v):
    x = dec(v)
    return None if x is None else x.quantize(Q2, rounding=ROUND_HALF_UP)
def num(v):
    return None if v is None else float(v)

# Inputs actually used
print('Inputs:', sorted(p.name for p in INPUT.glob('*')))

fx = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    header = next(r)
    if header[:2] != ['Currency', 'RateToTWD']:
        raise RuntimeError(f'Unexpected fx header: {header}')
    for row in r:
        if len(row) >= 2:
            fx[row[0]] = dec(row[1])
print('FX rates:', {k: str(v) for k, v in fx.items()})

base_cmd = [str(CLI)] if CLI.exists() else [sys.executable, str(CLI_PY)]
if not CLI.exists() and not CLI_PY.exists():
    raise RuntimeError('quote_cli executable not found')
print('CLI base:', base_cmd)
log = []

def run_cli(args):
    attempts = 0
    while True:
        attempts += 1
        p = subprocess.run(base_cmd + args, cwd=str(WORKSPACE), text=True, capture_output=True)
        log.append({'args': args, 'attempt': attempts, 'exit_code': p.returncode, 'stderr': p.stderr.strip(), 'stdout_len': len(p.stdout)})
        print(f"CLI attempt {attempts}: {' '.join(args)} -> exit {p.returncode}")
        if p.returncode == 4:
            if attempts <= 2:  # original call + up to two retries total = max 3 attempts
                print('exit code 4 retry:', p.stderr.strip())
                time.sleep(0.2)
                continue
            raise RuntimeError('quote_cli exit code 4 after two retries; stderr: ' + p.stderr.strip())
        if p.returncode in (2, 3):
            raise RuntimeError(f'quote_cli exit code {p.returncode}; not retried; stderr: {p.stderr.strip()}')
        if p.returncode != 0:
            raise RuntimeError(f'quote_cli exit code {p.returncode}; stderr: {p.stderr.strip()}')
        try:
            return json.loads(p.stdout)
        except json.JSONDecodeError as e:
            raise RuntimeError('quote_cli stdout was successful but not parseable JSON; stdout not used as data') from e

listing = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'])
if listing.get('schema_version') != 1 or not isinstance(listing.get('items'), list):
    raise RuntimeError('Unexpected quotes list schema')
if listing.get('count') != len(listing['items']):
    raise RuntimeError('quotes list count does not match items length')
print('List count:', len(listing['items']))

quotes, lines = [], []
for h in listing['items']:
    qid = text_or_blank(h.get('QuoteID'))
    if not qid:
        raise RuntimeError(f'Missing QuoteID in list item: {h}')
    detail = run_cli(['quotes', 'get', qid, '--format', 'json'])
    if detail.get('schema_version') != 1 or not isinstance(detail.get('header'), dict) or not isinstance(detail.get('line_items'), list):
        raise RuntimeError(f'Unexpected quote detail schema for {qid}')
    if text_or_blank(detail['header'].get('QuoteID')) != qid:
        raise RuntimeError(f'QuoteID mismatch for {qid}')
    quotes.append(detail['header'])
    for li in detail['line_items']:
        row = dict(li)
        row['QuoteID'] = qid
        lines.append(row)
quotes.sort(key=lambda q: (text_or_blank(q.get('QuoteDate')) or '', text_or_blank(q.get('QuoteID')) or ''))
order = {text_or_blank(q.get('QuoteID')): i for i, q in enumerate(quotes)}
lines.sort(key=lambda li: (order.get(text_or_blank(li.get('QuoteID')), 999999), text_or_blank(li.get('ProductID')) or '', text_or_blank(li.get('ProductName')) or ''))
print('Fetched details:', {'quotes': len(quotes), 'lines': len(lines)})

shutil.copy2(TEMPLATE, OUTFILE)
wb = load_workbook(OUTFILE)

def clear_rows(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)
def copy_header_style(ws, row, max_col):
    for c in range(1, max_col + 1):
        src, dst = ws.cell(4, c), ws.cell(row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)

# Report metadata
rpt = wb['Report']
rpt['B3'], rpt['D3'] = FROM_DATE, TO_DATE
rpt['B4'] = f'Status={STATUS}; QuoteDate between {FROM_DATE} and {TO_DATE} inclusive'
rpt['B5'] = GENERATED_AT
rpt['B7'], rpt['D7'] = 'quote_cli; fx_rates.csv', FETCHED_AT
rpt['B9'] = ''

# Quotes
q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
wsq = wb['Quotes']; clear_rows(wsq, 5)
if [wsq.cell(4, c).value for c in range(1, 12)] != q_headers:
    raise RuntimeError('Quotes header row changed/unexpected')
for rr, q in enumerate(quotes, 5):
    copy_header_style(wsq, rr, 11)
    cur = text_or_blank(q.get('Currency'))
    if cur not in fx:
        raise RuntimeError(f'Missing FX for {cur}')
    subtotal, discount, net = money(q.get('Subtotal')), dec(q.get('DiscountPercent')), money(q.get('NetAmount'))
    if subtotal is not None and discount is not None:
        expected = (subtotal * (Decimal('1') - discount / Decimal('100'))).quantize(Q2, rounding=ROUND_HALF_UP)
        if net != expected:
            raise RuntimeError(f'NetAmount mismatch for {q.get("QuoteID")}: {net} != {expected}')
    net_twd = None if net is None else (net * fx[cur]).quantize(Q2, rounding=ROUND_HALF_UP)
    vals = [text_or_blank(q.get('QuoteID')), text_or_blank(q.get('CustomerID')), text_or_blank(q.get('ProjectID')), text_or_blank(q.get('QuoteDate')), text_or_blank(q.get('ValidUntil')), text_or_blank(q.get('Status')), cur, num(subtotal), num(discount), num(net), num(net_twd)]
    for c, v in enumerate(vals, 1):
        cell = wsq.cell(rr, c, v)
        if c in (1,2,3): cell.number_format = '@'
        elif c in (4,5): cell.number_format = 'yyyy-mm-dd'
        elif c in (8,9,10,11): cell.number_format = '0.00'

# Lines
l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
wsl = wb['Lines']; clear_rows(wsl, 5)
if [wsl.cell(4, c).value for c in range(1, 7)] != l_headers:
    raise RuntimeError('Lines header row changed/unexpected')
for rr, li in enumerate(lines, 5):
    copy_header_style(wsl, rr, 6)
    qty, unit, amount = dec(li.get('Quantity')), money(li.get('UnitPrice')), money(li.get('LineAmount'))
    if qty is not None and unit is not None:
        expected = (qty * unit).quantize(Q2, rounding=ROUND_HALF_UP)
        if amount != expected:
            raise RuntimeError(f'LineAmount mismatch for {li.get("QuoteID")}/{li.get("ProductID")}: {amount} != {expected}')
    vals = [text_or_blank(li.get('QuoteID')), text_or_blank(li.get('ProductID')), text_or_blank(li.get('ProductName')), num(qty), num(unit), num(amount)]
    for c, v in enumerate(vals, 1):
        cell = wsl.cell(rr, c, v)
        if c in (1,2): cell.number_format = '@'
        elif c in (5,6): cell.number_format = '0.00'

# Summary by Status + Currency; original currency totals are not mixed.
wss = wb['Summary']; clear_rows(wss, 5)
last_quote_row = 5 if not quotes else 4 + len(quotes)
wss['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
wss['B2'].number_format = '0.00'
summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0})
for q in quotes:
    cur = text_or_blank(q.get('Currency'))
    net = money(q.get('NetAmount')) or Decimal('0.00')
    key = (text_or_blank(q.get('Status')) or '', '', cur)
    summary[key]['amount'] += net
    summary[key]['amount_twd'] += (net * fx[cur]).quantize(Q2, rounding=ROUND_HALF_UP)
    summary[key]['count'] += 1
for rr, (key, val) in enumerate(sorted(summary.items()), 5):
    copy_header_style(wss, rr, 6)
    vals = [key[0], key[1], key[2], num(val['amount'].quantize(Q2, rounding=ROUND_HALF_UP)), num(val['amount_twd'].quantize(Q2, rounding=ROUND_HALF_UP)), val['count']]
    for c, v in enumerate(vals, 1):
        cell = wss.cell(rr, c, v)
        if c in (4,5): cell.number_format = '0.00'
        elif c == 6: cell.number_format = '0'

wb.save(OUTFILE)
LOGFILE.write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding='utf-8')
print('Wrote:', OUTFILE)
print('Wrote:', LOGFILE)

# Reopen and validate every stated/template requirement possible from saved outputs.
wbv = load_workbook(OUTFILE, data_only=False)
errors = []
for s in ['Report','Quotes','Lines','Summary','Instructions']:
    if s not in wbv.sheetnames: errors.append(f'Missing sheet {s}')
if not errors:
    vr, vq, vl, vs = wbv['Report'], wbv['Quotes'], wbv['Lines'], wbv['Summary']
    if [vq.cell(4, c).value for c in range(1, 12)] != q_headers: errors.append('Quotes headers not preserved')
    if [vl.cell(4, c).value for c in range(1, 7)] != l_headers: errors.append('Lines headers not preserved')
    if vr['B3'].value != FROM_DATE or vr['D3'].value != TO_DATE: errors.append('Report period wrong')
    if STATUS not in str(vr['B4'].value): errors.append('Report condition missing Sent')
    if vs['B2'].value != f'=SUM(Quotes!K5:K{last_quote_row})': errors.append(f'Summary B2 formula wrong: {vs["B2"].value}')
    if (vq.max_row - 4) != len(quotes): errors.append(f'Quotes rows {vq.max_row - 4} != {len(quotes)}')
    if (vl.max_row - 4) != len(lines): errors.append(f'Lines rows {vl.max_row - 4} != {len(lines)}')
    quote_ids = set()
    for rr in range(5, 5 + len(quotes)):
        for c in (1,2,3):
            v = vq.cell(rr,c).value
            if v is not None and not isinstance(v, str): errors.append(f'Quotes {vq.cell(rr,c).coordinate} nonblank ID not text')
        if not isinstance(vq.cell(rr,1).value, str): errors.append(f'Quotes {rr} QuoteID missing/not text')
        quote_ids.add(vq.cell(rr,1).value)
        if vq.cell(rr,6).value != STATUS: errors.append(f'Quotes row {rr} status not Sent')
        if not (FROM_DATE <= str(vq.cell(rr,4).value) <= TO_DATE): errors.append(f'Quotes row {rr} date outside period')
        for c in (8,9,10,11):
            if vq.cell(rr,c).value is not None and not isinstance(vq.cell(rr,c).value, (int,float)): errors.append(f'Quotes {vq.cell(rr,c).coordinate} amount not numeric')
    for rr in range(5, 5 + len(lines)):
        for c in (1,2):
            v = vl.cell(rr,c).value
            if v is not None and not isinstance(v, str): errors.append(f'Lines {vl.cell(rr,c).coordinate} nonblank ID not text')
        if vl.cell(rr,1).value not in quote_ids: errors.append(f'Lines row {rr} QuoteID not in Quotes')
        for c in (4,5,6):
            if vl.cell(rr,c).value is not None and not isinstance(vl.cell(rr,c).value, (int,float)): errors.append(f'Lines {vl.cell(rr,c).coordinate} numeric field not numeric')
    for rr in range(5, vs.max_row + 1):
        for c in (4,5,6):
            if vs.cell(rr,c).value is not None and not isinstance(vs.cell(rr,c).value, (int,float)): errors.append(f'Summary {vs.cell(rr,c).coordinate} numeric field not numeric')

print('Validation counts:', {'quotes': len(quotes), 'lines': len(lines), 'summary': max(0, wbv['Summary'].max_row - 4)})
print('Quote IDs:', sorted(quote_ids) if 'quote_ids' in locals() else [])
print('Summary B2:', wbv['Summary']['B2'].value)
if errors:
    print('VALIDATION ERRORS:')
    for e in errors: print('-', e)
    raise RuntimeError('Output validation failed')
print('Validation passed: saved workbook reopened; required sheets/headers, Report period/condition, row counts, nonblank text IDs with blank missing values allowed, numeric fields, Sent status, September date range, line joins, and Summary!B2 formula checked.')
