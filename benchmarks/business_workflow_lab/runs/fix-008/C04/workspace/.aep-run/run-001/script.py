from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nFX RATES SAMPLE:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i, row in enumerate(r):
            print(i, row)
            if i >= 5: break

tpl = ws/'input'/'quote_template.xlsx'
if tpl.exists():
    print('\nWORKBOOK STRUCTURE:')
    wb = load_workbook(tpl, data_only=False)
    for sh in wb.sheetnames:
        wsht = wb[sh]
        print('SHEET', sh, 'dim', wsht.max_row, wsht.max_column)
        for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 12), values_only=True):
            print(row)
        print('---')

manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\nMANUAL CONTENT:')
    text = manual.read_text(encoding='utf-8', errors='replace')
    print(text[:8000])
