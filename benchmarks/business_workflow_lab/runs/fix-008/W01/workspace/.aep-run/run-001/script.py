from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
# inspect fx
fx = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv exists', fx.exists())
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        rows = list(r)
    print('fx rows count', len(rows))
    print('fx header/sample', rows[:5])
# inspect manual
manual = ws/'input'/'sap_manual.md'
print('\nsap_manual.md exists', manual.exists())
if manual.exists():
    text = manual.read_text(encoding='utf-8')
    print('manual first 2000 chars:\n', text[:2000])
# inspect workbook
xlsx = ws/'input'/'sales_template.xlsx'
print('\nsales_template.xlsx exists', xlsx.exists())
if xlsx.exists():
    wb = load_workbook(xlsx)
    print('sheets', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print(f'-- sheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
            print(row)
        print()