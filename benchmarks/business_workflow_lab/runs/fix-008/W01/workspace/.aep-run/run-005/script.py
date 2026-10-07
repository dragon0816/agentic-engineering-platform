import os, json, csv, re
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from collections import defaultdict
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
start_date = params['start_date']
end_date = params['end_date']
status = (params.get('status') or '').strip()
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
download_dir = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)
download_dir.mkdir(exist_ok=True)

print('Input files:', sorted(p.name for p in input_dir.iterdir()))
print('Download files:', sorted(p.name for p in download_dir.glob('*.csv')))

safe_status = status if status else 'all'
expected_names = [
    f'sales_orders_{start_date}_{end_date}_{safe_status}.csv',
    f'sales_orders_{start_date}_{end_date}_{status}.csv',
    'sales_orders_export.csv',
]
existing = [download_dir / n for n in expected_names if (download_dir / n).exists()]
if existing:
    csv_path = max(existing, key=lambda p: p.stat().st_mtime)
else:
    matches = list(download_dir.glob(f'sales_orders_{start_date}_{end_date}_*.csv')) + list(download_dir.glob('sales_orders_export*.csv'))
    if not matches:
        raise FileNotFoundError(f'Downloaded CSV not found for {start_date}..{end_date}, status={status!r} under {download_dir}')
    csv_path = max(matches, key=lambda p: p.stat().st_mtime)
print('Using downloaded CSV:', csv_path.relative_to(workspace))

rates = {}
with (input_dir / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', {k: str(v) for k, v in rates.items()})

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
print('CSV headers:', list(rows[0].keys()) if rows else [])
print('CSV rows:', len(rows))

required = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']
missing_cols = [c for c in required if rows and c not in rows[0]]
if missing_cols:
    raise ValueError('CSV missing required columns: ' + ', '.join(missing_cols))

rows.sort(key=lambda r: str(r.get('OrderID','')))
q = Decimal('0.01')
for r in rows:
    if not (start_date <= r['OrderDate'] <= end_date):
        raise ValueError(f'Row outside requested date range: {r}')
    if status and r['Status'] != status:
        raise ValueError(f'Row status does not match requested status: {r}')
    cur = r['Currency']
    if cur not in rates:
        raise ValueError(f'Missing FX rate for {cur}')
    amt = Decimal(r['NetAmount'])
    r['_amount'] = amt
    r['_twd'] = (amt * rates[cur]).quantize(q, rounding=ROUND_HALF_UP)

wb = load_workbook(input_dir / 'sales_template.xlsx')
report = wb['Report']
details = wb['Details']
summary = wb['Summary']
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
actual_headers = [details.cell(4, c).value for c in range(1, 12)]
print('Template details headers:', actual_headers)
if actual_headers != expected_headers:
    raise ValueError('Unexpected Details header row in template')

generated_at = datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = f'Sales Orders; Status={status or "(all)"}; Sort by OrderID asc'
report['B5'] = generated_at
report['B7'] = 'Lab ERP (synthetic SAP-like)'
report['D7'] = generated_at
report['B9'] = 'Export CSV (all filtered rows) used for all pages; FX from input/fx_rates.csv. CustomerContact and PaymentTerms are blank because the export does not include detail-page-only fields.'

for row in details.iter_rows(min_row=5, max_row=max(details.max_row, 5), min_col=1, max_col=11):
    for cell in row:
        cell.value = None
for row in summary.iter_rows(min_row=5, max_row=max(summary.max_row, 5), min_col=1, max_col=6):
    for cell in row:
        cell.value = None

for idx, r in enumerate(rows, start=5):
    vals = [r['OrderID'], r['CustomerID'], r['CustomerName'], r.get('ProjectID') or None, r['OrderDate'], r['Status'], r['Currency'], float(r['_amount']), float(r['_twd']), None, None]
    for col, val in enumerate(vals, start=1):
        details.cell(idx, col).value = val
    for col in [1,2,4]:
        details.cell(idx, col).number_format = '@'
    details.cell(idx, 5).number_format = 'yyyy-mm-dd'
    for col in [8,9]:
        details.cell(idx, col).number_format = '0.00'

last_data_row = 4 + len(rows) if rows else 5
summary['B2'] = f'=SUM(Details!I5:I{last_data_row})'
by_currency = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
for r in rows:
    g = by_currency[r['Currency']]
    g['amount'] += r['_amount']
    g['twd'] += r['_twd']
    g['count'] += 1
for row_idx, cur in enumerate(sorted(by_currency.keys()), start=5):
    g = by_currency[cur]
    vals = ['Currency', cur, cur, float(g['amount'].quantize(q, rounding=ROUND_HALF_UP)), float(g['twd'].quantize(q, rounding=ROUND_HALF_UP)), g['count']]
    for col, val in enumerate(vals, start=1):
        summary.cell(row_idx, col).value = val
    for col in [4,5]:
        summary.cell(row_idx, col).number_format = '0.00'

if start_date == '2026-09-01' and end_date == '2026-09-30' and status == 'Open':
    out_name = 'september_open_orders.xlsx'
else:
    out_name = 'sales_orders_' + re.sub(r'[^0-9A-Za-z_-]+', '_', f'{start_date}_{end_date}_{status or "all"}') + '.xlsx'
out_path = output_dir / out_name
wb.save(out_path)
print('Saved:', out_path.relative_to(workspace))

wb2 = load_workbook(out_path, data_only=False)
if wb2.sheetnames != ['Report','Details','Summary','Instructions']:
    raise AssertionError('Workbook sheet names changed')
det2 = wb2['Details']; sum2 = wb2['Summary']; rep2 = wb2['Report']
if rep2['B3'].value != start_date or rep2['D3'].value != end_date:
    raise AssertionError('Report date cells mismatch')
read_ids = [det2.cell(r,1).value for r in range(5, 5+len(rows))]
if read_ids != sorted(read_ids):
    raise AssertionError('OrderID values are not sorted')
if not all(isinstance(x, str) for x in read_ids):
    raise AssertionError('OrderID values are not text')
if status and any(det2.cell(r,6).value != status for r in range(5,5+len(rows))):
    raise AssertionError('Status values mismatch')
if any(not isinstance(det2.cell(r,8).value, (int,float)) or not isinstance(det2.cell(r,9).value, (int,float)) for r in range(5,5+len(rows))):
    raise AssertionError('Amount fields are not numeric')
if sum2['B2'].value != f'=SUM(Details!I5:I{last_data_row})':
    raise AssertionError('Summary B2 formula range mismatch')
print('Verified: template sheets preserved, details rows=', len(rows), 'OrderID text sorted, amounts numeric, FX TWD conversion applied, summary by currency written.')