from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from collections import defaultdict
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

print('Inputs present:', sorted(p.name for p in input_dir.iterdir()))
print('Downloaded files:', sorted(str(p.relative_to(workspace)) for p in (output_dir/'downloads').glob('*')) if (output_dir/'downloads').exists() else [])

csv_path = output_dir/'downloads'/'sales_orders_2026-09-01_2026-09-30_Open.csv'
template_path = input_dir/'sales_template.xlsx'
fx_path = input_dir/'fx_rates.csv'
out_path = output_dir/'september_open_orders.xlsx'

# Read FX rates (only first two columns are needed; note column may contain commas in test file).
rates = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header first columns:', header[:2])
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', {k: str(v) for k,v in rates.items()})

# Read exported all-filtered rows.
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
print('CSV headers:', list(rows[0].keys()) if rows else [])
print('CSV row count:', len(rows))

# Detail fields read from SAP detail pages through browser__operate.
detail_by_id = {
'0045000109': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10'},
'0045000110': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET30'},
'0045000112': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15'},
'0045000113': {'CustomerContact': '', 'PaymentTerms': 'NET15'},
'0045000115': {'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET20'},
'0045000116': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
'0045000118': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET0'},
'0045000119': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
'0045000120': {'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET10'},
'0045000122': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET5'},
'0045000124': {'CustomerContact': '', 'PaymentTerms': 'NET30'},
'0045000126': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
'0045000127': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET30'},
'0045000128': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET5'},
}

# Sort by OrderID and validate filters.
rows.sort(key=lambda r: r['OrderID'])
for r in rows:
    assert '2026-09-01' <= r['OrderDate'] <= '2026-09-30', r
    assert r['Status'] == 'Open', r
    assert r['Currency'] in rates, r
    assert r['OrderID'] in detail_by_id, 'Missing detail for '+r['OrderID']

q = Decimal('0.01')
for r in rows:
    amt = Decimal(r['NetAmount'])
    twd = (amt * rates[r['Currency']]).quantize(q, rounding=ROUND_HALF_UP)
    r['NetAmountDecimal'] = amt
    r['NetAmountTWDDecimal'] = twd
    r.update(detail_by_id[r['OrderID']])

wb = load_workbook(template_path)
report = wb['Report']
details = wb['Details']
summary = wb['Summary']

# Fill Report sheet.
generated_at = datetime.now(timezone(timedelta(hours=8))).isoformat(timespec='seconds')
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Sales Orders; Status=Open; Sort by OrderID asc'
report['B5'] = generated_at
report['B7'] = 'Lab ERP (synthetic SAP-like)'
report['D7'] = generated_at
report['B9'] = 'Export CSV (all filtered rows) used for all pages; detail pages read for CustomerContact and PaymentTerms. FX from input/fx_rates.csv.'

# Clear possible existing rows below header without deleting template structure (template has none).
for row in details.iter_rows(min_row=5, max_row=details.max_row, min_col=1, max_col=11):
    for cell in row:
        cell.value = None

headers = [details.cell(4, c).value for c in range(1, 12)]
print('Details headers:', headers)
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert headers == expected_headers

for idx, r in enumerate(rows, start=5):
    vals = [
        r['OrderID'], r['CustomerID'], r['CustomerName'], r['ProjectID'] or None, r['OrderDate'], r['Status'], r['Currency'],
        float(r['NetAmountDecimal']), float(r['NetAmountTWDDecimal']), r['CustomerContact'] or None, r['PaymentTerms'] or None
    ]
    for col, val in enumerate(vals, start=1):
        c = details.cell(idx, col)
        c.value = val
    for col in [1,2,4]:
        details.cell(idx, col).number_format = '@'
    details.cell(idx,5).number_format = 'yyyy-mm-dd'
    for col in [8,9]:
        details.cell(idx,col).number_format = '0.00'

last_row = 4 + len(rows) if rows else 5
summary['B2'] = f'=SUM(Details!I5:I{last_row if rows else 5})'

# Clear summary rows 5+.
for row in summary.iter_rows(min_row=5, max_row=summary.max_row, min_col=1, max_col=6):
    for cell in row:
        cell.value = None

by_currency = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
for r in rows:
    g = by_currency[r['Currency']]
    g['amount'] += r['NetAmountDecimal']
    g['twd'] += r['NetAmountTWDDecimal']
    g['count'] += 1

for row_idx, cur in enumerate(sorted(by_currency.keys()), start=5):
    g = by_currency[cur]
    vals = ['Currency', cur, cur, float(g['amount'].quantize(q, rounding=ROUND_HALF_UP)), float(g['twd'].quantize(q, rounding=ROUND_HALF_UP)), g['count']]
    for col, val in enumerate(vals, start=1):
        summary.cell(row_idx, col).value = val
    for col in [4,5]:
        summary.cell(row_idx, col).number_format = '0.00'

wb.save(out_path)
print('Saved', out_path)

# Reopen and verify output requirements.
wb2 = load_workbook(out_path, data_only=False)
assert wb2.sheetnames == ['Report','Details','Summary','Instructions']
rep2, det2, sum2 = wb2['Report'], wb2['Details'], wb2['Summary']
assert rep2['B3'].value == '2026-09-01'
assert rep2['D3'].value == '2026-09-30'
assert 'Status=Open' in rep2['B4'].value
assert [det2.cell(4,c).value for c in range(1,12)] == expected_headers
read_rows = []
for rr in range(5, 5+len(rows)):
    read_rows.append([det2.cell(rr,c).value for c in range(1,12)])
assert len(read_rows) == 14
assert [r[0] for r in read_rows] == sorted([r[0] for r in read_rows])
assert all(isinstance(r[0], str) and r[0].startswith('0045') for r in read_rows)
assert all(r[5] == 'Open' for r in read_rows)
assert all('2026-09-01' <= r[4] <= '2026-09-30' for r in read_rows)
assert all(isinstance(r[7], (int,float)) for r in read_rows)
assert all(isinstance(r[8], (int,float)) for r in read_rows)
assert sum2['B2'].value == '=SUM(Details!I5:I18)'
summary_rows = [[sum2.cell(r,c).value for c in range(1,7)] for r in range(5, 5+len(by_currency))]
print('Summary rows:', summary_rows)
expected_total = sum(r['NetAmountTWDDecimal'] for r in rows).quantize(q)
summary_twd_total = sum(Decimal(str(r[4])) for r in summary_rows).quantize(q)
assert summary_twd_total == expected_total, (summary_twd_total, expected_total)
print('Verified output: sheets/template preserved, 14 detail rows, OrderID text sorted, filters Open Sept 2026, numeric amounts, TWD conversion, summary by currency, total TWD formula range.')