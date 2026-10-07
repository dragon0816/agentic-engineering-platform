from openpyxl import load_workbook
from pathlib import Path
wb=load_workbook('input/sales_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)