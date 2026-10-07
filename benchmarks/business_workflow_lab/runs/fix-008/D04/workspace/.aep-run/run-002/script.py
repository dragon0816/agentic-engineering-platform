from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print((ws/'input'/'expense_manual.md').read_text(encoding='utf-8'))
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    sht=wb[sh]
    print('\nSHEET', sh)
    for row in sht.iter_rows(min_row=1, max_row=sht.max_row, max_col=sht.max_column, values_only=False):
        vals=[]
        for c in row:
            if c.value is not None:
                vals.append(f'{c.coordinate}={c.value!r}')
        if vals:
            print('; '.join(vals))
    print('merged:', list(sht.merged_cells.ranges))