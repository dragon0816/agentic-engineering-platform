from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:', sorted([p.name for p in (ws/'input').iterdir()]))
# workbook sheets and dimensions
wb_path = ws/'input'/'expense_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('template sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    sht=wb[sh]
    print('sheet', sh, 'max_row', sht.max_row, 'max_col', sht.max_column)
    for r in sht.iter_rows(min_row=1, max_row=min(8,sht.max_row), values_only=True):
        print(' ', r)
# fx csv sample
fx=ws/'input'/'fx_rates.csv'
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(6), reader):
        print('fx row', i, row)
# manual sample
manual=ws/'input'/'expense_manual.md'
print('manual first 1200 chars:')
print(manual.read_text(encoding='utf-8')[:1200])