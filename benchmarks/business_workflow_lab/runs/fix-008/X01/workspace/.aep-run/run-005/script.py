from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, json, math
from openpyxl import load_workbook
from datetime import date

base = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-008\X01\workspace')
out = base/'output'
from_date='2026-09-01'; to_date='2026-09-30'
orders_csv = out/'downloads'/'sap_orders_export.csv'
invoices_csv = out/'downloads'/'sap_invoices_export.csv'
expenses_csv = out/'expense_approved_export.csv'
quotes_json = out/f'quotes_{from_date}_{to_date}_accepted.json'
report_xlsx = out/f'project_report_{from_date}_{to_date}.xlsx'
fx_csv = base/'input'/'fx_rates.csv'

required = [orders_csv, invoices_csv, expenses_csv, quotes_json, report_xlsx, fx_csv]
for p in required:
    assert p.exists(), f'missing {p}'
    assert p.stat().st_size > 0, f'empty {p}'
print('All required output/input files exist and are non-empty.')

# read fx
fx = {}
with fx_csv.open(encoding='utf-8-sig', newline='') as f:
    for row in csv.DictReader(f):
        cur = row.get('Currency') or row.get('currency')
        rate = row.get('RateToTWD') or row.get('rate_to_twd') or row.get('TWD') or row.get('Rate')
        if cur and rate:
            fx[cur] = float(rate)
print('FX currencies:', sorted(fx))

def in_range(s):
    return from_date <= s <= to_date

def add(d, project, field, amount):
    if not project:
        return
    d.setdefault(project, {'OrderTWD':0.0,'InvoiceBalanceTWD':0.0,'ApprovedExpenseTWD':0.0,'AcceptedQuoteTWD':0.0})
    d[project][field] += amount

def close(a,b):
    return abs(float(a)-float(b)) < 0.011

summary = {}
# orders
with orders_csv.open(encoding='utf-8-sig', newline='') as f:
    rows = list(csv.DictReader(f))
assert rows, 'orders no rows'
for r in rows:
    assert in_range(r['OrderDate']), f'order out of range {r}'
    assert r['Status'] != 'Cancelled', f'cancelled order included {r}'
    amt = float(r['NetAmount']) * fx[r['Currency']]
    add(summary, r.get('ProjectID',''), 'OrderTWD', amt)
print('Orders validated:', len(rows), 'rows; no Cancelled; all in period.')
# invoices
with invoices_csv.open(encoding='utf-8-sig', newline='') as f:
    inv_rows = list(csv.DictReader(f))
assert inv_rows, 'invoices no rows'
for r in inv_rows:
    assert in_range(r['InvoiceDate']), f'invoice out of range {r}'
    bal = float(r['Balance']) * fx[r['Currency']]
    add(summary, r.get('ProjectID',''), 'InvoiceBalanceTWD', bal)
print('Invoices validated:', len(inv_rows), 'rows; all in period; using Balance only.')
# expenses
with expenses_csv.open(encoding='utf-8-sig', newline='') as f:
    exp_rows = list(csv.DictReader(f))
assert exp_rows, 'expenses no rows'
for r in exp_rows:
    assert in_range(r['ExpenseDate']), f'expense out of range {r}'
    assert r['Status'] == 'Approved', f'non-approved expense included {r}'
    amt = float(r['Amount']) * fx[r['Currency']]
    add(summary, r.get('ProjectID',''), 'ApprovedExpenseTWD', amt)
print('Expenses validated:', len(exp_rows), 'rows; all Approved; all in period.')
# quotes
qdata = json.loads(quotes_json.read_text(encoding='utf-8'))
items = qdata['items']
assert qdata['count'] == len(items), 'quote count mismatch'
for r in items:
    qdate = r.get('QuoteDate') or r.get('quote_date')
    status = r.get('Status') or r.get('status')
    project = r.get('ProjectID') or r.get('project_id') or ''
    cur = r.get('Currency') or r.get('currency')
    amt_raw = r.get('NetAmount') if 'NetAmount' in r else r.get('net_amount')
    assert in_range(qdate), f'quote out of range {r}'
    assert status == 'Accepted', f'non-accepted quote included {r}'
    add(summary, project, 'AcceptedQuoteTWD', float(amt_raw) * fx[cur])
print('Quotes validated:', len(items), 'items; all Accepted; all in period.')

# workbook checks
wb = load_workbook(report_xlsx, data_only=False)
for s in ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']:
    assert s in wb.sheetnames, f'missing sheet {s}'
print('Workbook sheets present.')

# headers check row 4
expected_headers = {
    'Orders': ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
    'Invoices': ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
    'Expenses': ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
    'Quotes': ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
    'Summary': ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
}
for sheet, headers in expected_headers.items():
    ws = wb[sheet]
    got = [ws.cell(4, c).value for c in range(1, len(headers)+1)]
    assert got == headers, f'{sheet} headers mismatch {got}'
print('Workbook headers validated.')

# summary compare (rounded to cents tolerance)
wsum = wb['Summary']
wb_summary = {}
for row in range(5, wsum.max_row+1):
    project = wsum.cell(row,1).value
    if not project:
        continue
    wb_summary[str(project)] = {
        'OrderTWD': float(wsum.cell(row,2).value or 0),
        'InvoiceBalanceTWD': float(wsum.cell(row,3).value or 0),
        'ApprovedExpenseTWD': float(wsum.cell(row,4).value or 0),
        'AcceptedQuoteTWD': float(wsum.cell(row,5).value or 0),
    }
# remove zero-only generated? none expected
calc_projects = {p for p,v in summary.items() if any(abs(x)>0.0001 for x in v.values())}
assert set(wb_summary) == calc_projects, f'project set mismatch workbook={set(wb_summary)} calc={calc_projects}'
for p in sorted(calc_projects):
    for k in ['OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']:
        assert close(wb_summary[p][k], summary[p][k]), f'summary mismatch {p} {k}: wb={wb_summary[p][k]} calc={summary[p][k]}'
print('Summary totals match source exports by ProjectID and by separate amount columns.')
print('Projects:', sorted(wb_summary))
print('Final deliverable:', report_xlsx)
print('Supporting exports:', orders_csv, invoices_csv, expenses_csv, quotes_json, sep='\n- ')
