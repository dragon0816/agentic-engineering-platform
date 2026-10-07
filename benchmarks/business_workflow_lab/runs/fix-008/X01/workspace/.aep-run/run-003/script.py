import csv, json, os, sys, subprocess, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from collections import defaultdict
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
FROM = params['from_date']
TO = params['to_date']
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
downloads = output_dir / 'downloads'
output_dir.mkdir(parents=True, exist_ok=True)

orders_path = downloads / 'sap_orders_export.csv'
invoices_path = downloads / 'sap_invoices_export.csv'
expenses_path = output_dir / 'expense_approved_export.csv'
quote_json_path = output_dir / f'quotes_{FROM}_{TO}_accepted.json'
report_path = output_dir / f'project_report_{FROM}_{TO}.xlsx'

required = [input_dir/'project_template.xlsx', input_dir/'fx_rates.csv', orders_path, invoices_path, expenses_path]
print('Inputs expected:')
for p in required:
    print(f'- {p}: exists={p.exists()} size={p.stat().st_size if p.exists() else "missing"}')
missing = [str(p) for p in required if not p.exists()]
if missing:
    raise FileNotFoundError('Missing required input/export files: ' + '; '.join(missing))

TWOPLACES = Decimal('0.01')
def D(v):
    if v is None or str(v).strip() == '':
        return Decimal('0.00')
    return Decimal(str(v).strip())
def q2(v):
    return Decimal(v).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def money_cell(v):
    return float(q2(v))

def read_csv_dict(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    print(f'Read {len(rows)} rows from {path.name}; headers={list(rows[0].keys()) if rows else []}')
    return rows

fx = {}
with open(input_dir/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('FX headers:', reader.fieldnames)
    for row in reader:
        cur = (row.get('Currency') or row.get('currency') or '').strip()
        rate = row.get('TWD') or row.get('RateToTWD') or row.get('rate') or row.get('Rate')
        if not cur and row:
            vals = list(row.values())
            cur = str(vals[0]).strip() if vals else ''
            rate = vals[1] if len(vals) > 1 else rate
        if cur:
            fx[cur] = D(rate)
print('FX currencies:', sorted(fx))

def to_twd(amount, currency):
    currency = (currency or '').strip()
    if currency not in fx:
        raise ValueError(f'Missing FX rate for currency {currency!r}')
    return q2(D(amount) * fx[currency])

orders_raw = read_csv_dict(orders_path)
invoices = read_csv_dict(invoices_path)
expenses = read_csv_dict(expenses_path)
orders = [r for r in orders_raw if (r.get('Status') or '').strip() != 'Cancelled']
removed_cancelled = len(orders_raw) - len(orders)

cli_py = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
cmds = [
    [sys.executable, str(cli_py), 'quotes', 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json'],
    [sys.executable, str(cli_py), 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json'],
]
last = None
for cmd in cmds:
    for attempt in range(1, 4):
        proc = subprocess.run(cmd, text=True, capture_output=True)
        last = proc
        if proc.returncode == 0:
            break
        if proc.returncode == 4 and attempt < 3:
            continue
        break
    if last and last.returncode == 0:
        break
if not last or last.returncode != 0:
    raise RuntimeError(f'quote_cli failed rc={last.returncode if last else "none"}: {last.stderr if last else "no run"}')
quote_payload = json.loads(last.stdout)
quotes = quote_payload.get('items', quote_payload if isinstance(quote_payload, list) else [])
quote_json_path.write_text(json.dumps(quote_payload, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'Read {len(quotes)} accepted quotes from quote_cli; saved {quote_json_path.name}')

def in_range(date_s):
    return FROM <= (date_s or '') <= TO
for r in orders:
    assert in_range(r.get('OrderDate')), ('order out of range', r)
    assert (r.get('Status') or '').strip() != 'Cancelled', ('cancelled order included', r)
for r in invoices:
    assert in_range(r.get('InvoiceDate')), ('invoice out of range', r)
for r in expenses:
    assert in_range(r.get('ExpenseDate')), ('expense out of range', r)
    assert (r.get('Status') or '').strip() == 'Approved', ('non-approved expense included', r)
for r in quotes:
    assert in_range(r.get('QuoteDate')), ('quote out of range', r)
    assert (r.get('Status') or '').strip() == 'Accepted', ('non-accepted quote included', r)

shutil.copyfile(input_dir/'project_template.xlsx', report_path)
wb = load_workbook(report_path)
headers_expected = {
    'Orders': ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
    'Invoices': ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
    'Expenses': ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
    'Quotes': ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
    'Summary': ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
    'Exceptions': ['Source','DocumentID','ProjectID','Reason'],
}
for sheet, heads in headers_expected.items():
    actual = [wb[sheet].cell(4, c).value for c in range(1, len(heads)+1)]
    if actual != heads:
        raise AssertionError(f'Unexpected headers in {sheet}: {actual}')

def clear_sheet(sheet_name):
    ws = wb[sheet_name]
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)
    return ws

def set_text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.number_format = '@'
def set_date(cell, value):
    cell.value = value or None
    cell.number_format = 'yyyy-mm-dd'
def set_amount(cell, value):
    cell.value = money_cell(value)
    cell.number_format = '0.00'

summary = defaultdict(lambda: {'order': Decimal('0.00'), 'invoice': Decimal('0.00'), 'expense': Decimal('0.00'), 'quote': Decimal('0.00')})
exceptions = []

ws = clear_sheet('Orders')
for i, r in enumerate(orders, 5):
    amt_twd = to_twd(r.get('NetAmount'), r.get('Currency'))
    set_text(ws.cell(i,1), r.get('OrderID')); set_text(ws.cell(i,2), r.get('ProjectID')); set_text(ws.cell(i,3), r.get('CustomerID'))
    set_date(ws.cell(i,4), r.get('OrderDate')); set_text(ws.cell(i,5), r.get('Status')); set_text(ws.cell(i,6), r.get('Currency'))
    set_amount(ws.cell(i,7), D(r.get('NetAmount'))); set_amount(ws.cell(i,8), amt_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['order'] += amt_twd
    else:
        exceptions.append(['Orders', r.get('OrderID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Invoices')
for i, r in enumerate(invoices, 5):
    balance = D(r.get('Balance')) if r.get('Balance') not in (None, '') else D(r.get('Amount')) - D(r.get('PaidAmount'))
    bal_twd = to_twd(balance, r.get('Currency'))
    set_text(ws.cell(i,1), r.get('InvoiceID')); set_text(ws.cell(i,2), r.get('OrderID')); set_text(ws.cell(i,3), r.get('ProjectID')); set_text(ws.cell(i,4), r.get('CustomerID'))
    set_date(ws.cell(i,5), r.get('InvoiceDate')); set_date(ws.cell(i,6), r.get('DueDate')); set_text(ws.cell(i,7), r.get('Currency'))
    set_amount(ws.cell(i,8), D(r.get('Amount'))); set_amount(ws.cell(i,9), D(r.get('PaidAmount'))); set_amount(ws.cell(i,10), balance); set_amount(ws.cell(i,11), bal_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['invoice'] += bal_twd
    else:
        exceptions.append(['Invoices', r.get('InvoiceID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Expenses')
for i, r in enumerate(expenses, 5):
    amt_twd = to_twd(r.get('Amount'), r.get('Currency'))
    set_text(ws.cell(i,1), r.get('ExpenseID')); set_text(ws.cell(i,2), r.get('ProjectID')); set_text(ws.cell(i,3), r.get('EmployeeID'))
    set_date(ws.cell(i,4), r.get('ExpenseDate')); set_text(ws.cell(i,5), r.get('Status')); set_text(ws.cell(i,6), r.get('Currency'))
    set_amount(ws.cell(i,7), D(r.get('Amount'))); set_amount(ws.cell(i,8), amt_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['expense'] += amt_twd
    else:
        exceptions.append(['Expenses', r.get('ExpenseID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Quotes')
for i, r in enumerate(quotes, 5):
    amt_twd = to_twd(r.get('NetAmount'), r.get('Currency'))
    set_text(ws.cell(i,1), r.get('QuoteID')); set_text(ws.cell(i,2), r.get('ProjectID')); set_text(ws.cell(i,3), r.get('CustomerID'))
    set_date(ws.cell(i,4), r.get('QuoteDate')); set_date(ws.cell(i,5), r.get('ValidUntil')); set_text(ws.cell(i,6), r.get('Status')); set_text(ws.cell(i,7), r.get('Currency'))
    set_amount(ws.cell(i,8), D(r.get('NetAmount'))); set_amount(ws.cell(i,9), amt_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['quote'] += amt_twd
    else:
        exceptions.append(['Quotes', r.get('QuoteID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Summary')
projects = sorted(summary)
for i, pid in enumerate(projects, 5):
    vals = summary[pid]
    set_text(ws.cell(i,1), pid)
    set_amount(ws.cell(i,2), vals['order'])
    set_amount(ws.cell(i,3), vals['invoice'])
    set_amount(ws.cell(i,4), vals['expense'])
    set_amount(ws.cell(i,5), vals['quote'])
last_row = max(5, 4 + len(projects))
for col, cell in [('B','B2'),('C','C2'),('D','D2'),('E','E2')]:
    ws[cell] = f'=SUM({col}5:{col}{last_row})'
    ws[cell].number_format = '0.00'

ws = clear_sheet('Exceptions')
for i, row in enumerate(exceptions, 5):
    for c, v in enumerate(row, 1):
        set_text(ws.cell(i,c), v)

rep = wb['Report']
rep['B3'] = FROM
rep['D3'] = TO
rep['B4'] = 'SAP Orders: OrderDate in range and Status != Cancelled; SAP Invoices: InvoiceDate in range; Expense: ExpenseDate in range and Status = Approved; Quotes: QuoteDate in range and Status = Accepted; summary by ProjectID only.'
rep['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B7'] = 'SAP web exports; Expense Viewer export; quote_cli Accepted JSON; fx_rates.csv'
rep['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B9'] = f'Cancelled SAP orders removed: {removed_cancelled}. Orders, invoice balances, approved expenses, and accepted quotes are separate columns; orders and quotes are not added as revenue; CustomerName is not used as a join key.'
wb.save(report_path)

# Saved-output verification.
wb2 = load_workbook(report_path, data_only=False)
for sheet, heads in headers_expected.items():
    actual = [wb2[sheet].cell(4, c).value for c in range(1, len(heads)+1)]
    assert actual == heads, (sheet, actual)
assert wb2['Report']['B3'].value == FROM and wb2['Report']['D3'].value == TO
counts = {'Orders': len(orders), 'Invoices': len(invoices), 'Expenses': len(expenses), 'Quotes': len(quotes), 'Summary': len(projects), 'Exceptions': len(exceptions)}
for sheet, n in counts.items():
    ws = wb2[sheet]
    actual_rows = 0
    for row in range(5, ws.max_row + 1):
        if any(ws.cell(row, col).value not in (None, '') for col in range(1, ws.max_column + 1)):
            actual_rows += 1
    assert actual_rows == n, f'{sheet} row count {actual_rows} != {n}'
for sheet, id_cols in {'Orders':[1,2,3], 'Invoices':[1,2,3,4], 'Expenses':[1,2,3], 'Quotes':[1,2,3], 'Summary':[1], 'Exceptions':[1,2,3,4]}.items():
    ws = wb2[sheet]
    for row in range(5, ws.max_row + 1):
        if any(ws.cell(row, c).value not in (None, '') for c in range(1, ws.max_column + 1)):
            for c in id_cols:
                assert ws.cell(row,c).number_format == '@', f'{sheet}!{ws.cell(row,c).coordinate} not text formatted'
for sheet, cols in {'Orders':[7,8], 'Invoices':[8,9,10,11], 'Expenses':[7,8], 'Quotes':[8,9], 'Summary':[2,3,4,5]}.items():
    ws = wb2[sheet]
    for row in range(5, ws.max_row + 1):
        if any(ws.cell(row, c).value not in (None, '') for c in range(1, ws.max_column + 1)):
            for c in cols:
                assert isinstance(ws.cell(row,c).value, (int, float)), f'{sheet}!{ws.cell(row,c).coordinate} not numeric'
print('Verification OK from saved workbook:', report_path)
print('Record counts:', counts)
print('Output files:', report_path, quote_json_path, expenses_path, orders_path, invoices_path, sep='\n- ')
