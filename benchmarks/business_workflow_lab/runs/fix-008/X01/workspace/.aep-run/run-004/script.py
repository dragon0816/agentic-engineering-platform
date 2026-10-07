from pathlib import Path
from openpyxl import load_workbook
import csv, json, os

ws = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-008\X01\workspace')
out = ws / 'output'
print('Workspace:', ws)
print('Output exists:', out.exists())
print('\nOutput tree:')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print('-', p.relative_to(ws), 'size=', p.stat().st_size)

xlsx_files = sorted(out.glob('*.xlsx'))
print('\nXLSX files:', [p.name for p in xlsx_files])
for xlsx in xlsx_files:
    print('\nInspect workbook:', xlsx.name)
    wb = load_workbook(xlsx, data_only=False)
    print('Sheets:', wb.sheetnames)
    for sh in wb.sheetnames:
        wsht = wb[sh]
        print(f' Sheet {sh}: max_row={wsht.max_row}, max_column={wsht.max_column}')
        for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 6), values_only=True):
            print('  ', r)

csv_files = sorted(out.rglob('*.csv'))
print('\nCSV files:', [str(p.relative_to(ws)) for p in csv_files])
for csvp in csv_files:
    print('\nInspect CSV:', csvp.relative_to(ws))
    try:
        with csvp.open('r', encoding='utf-8-sig', newline='') as f:
            reader = csv.reader(f)
            rows = []
            for i, row in enumerate(reader):
                rows.append(row)
                if i >= 4:
                    break
        print(' sample rows:')
        for row in rows:
            print('  ', row)
        with csvp.open('r', encoding='utf-8-sig', newline='') as f:
            count = sum(1 for _ in f)
        print(' line_count:', count)
    except UnicodeDecodeError:
        print(' cannot decode as utf-8-sig')

json_files = sorted(out.rglob('*.json'))
print('\nJSON files:', [str(p.relative_to(ws)) for p in json_files])
for jp in json_files:
    print('\nInspect JSON:', jp.relative_to(ws))
    try:
        data = json.loads(jp.read_text(encoding='utf-8'))
        if isinstance(data, list):
            print(' list length:', len(data))
            print(' first item:', data[0] if data else None)
        elif isinstance(data, dict):
            print(' keys:', list(data.keys()))
            for k, v in list(data.items())[:3]:
                print(' ', k, ':', v if not isinstance(v, (list, dict)) else (type(v).__name__, len(v)))
        else:
            print(type(data), data)
    except Exception as e:
        print(' json read error:', repr(e))
