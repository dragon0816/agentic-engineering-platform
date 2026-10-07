from pathlib import Path
import csv, json
from openpyxl import load_workbook

base = Path.cwd()
print('Workspace:', base)
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print(f'- {p.relative_to(base)} size={p.stat().st_size}')
print('Output files:')
for p in sorted((base/'output').rglob('*')):
    if p.is_file():
        print(f'- {p.relative_to(base)} size={p.stat().st_size}')

print('\nCSV headers and sample rows:')
for rel in [
    'output/downloads/sap_orders_export.csv',
    'output/downloads/sap_orders_2026-09-01_2026-09-30.csv',
    'output/downloads/sap_invoices_export.csv',
    'output/downloads/sap_invoices_2026-09-01_2026-09-30.csv',
    'output/expense_approved_export.csv',
    'output/expense_2026-09-01_2026-09-30_approved.csv',
]:
    p = base / rel
    print(f'\n{rel}: exists={p.exists()}')
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as f:
            rows = list(csv.DictReader(f))
        print('headers:', rows[0].keys() if rows else 'NO ROWS')
        print('row_count:', len(rows))
        for r in rows[:3]:
            print('sample:', r)

print('\nJSON preview:')
p = base/'output/quotes_2026-09-01_2026-09-30_accepted.json'
print(f'{p.relative_to(base)}: exists={p.exists()}')
if p.exists():
    data = json.loads(p.read_text(encoding='utf-8'))
    print('type:', type(data).__name__, 'count:', len(data) if isinstance(data, list) else 'n/a')
    if isinstance(data, list):
        for r in data[:3]: print('sample:', r)

print('\nWorkbook sheets and dimensions:')
p = base/'output/project_report_2026-09-01_2026-09-30.xlsx'
print(f'{p.relative_to(base)}: exists={p.exists()}')
if p.exists():
    wb = load_workbook(p, data_only=False)
    print('sheets:', wb.sheetnames)
    for ws in wb.worksheets:
        print(ws.title, 'rows=', ws.max_row, 'cols=', ws.max_column)
        rows = list(ws.iter_rows(min_row=1, max_row=min(ws.max_row,4), values_only=True))
        for row in rows:
            print(' ', row)
