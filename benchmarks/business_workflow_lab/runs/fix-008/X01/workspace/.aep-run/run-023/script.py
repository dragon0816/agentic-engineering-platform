from pathlib import Path
import csv, json
from decimal import Decimal, InvalidOperation
from datetime import date
from collections import defaultdict
from openpyxl import load_workbook

ws = Path.cwd()
out = ws/'output'
start = date.fromisoformat('2026-08-01')
end = date.fromisoformat('2026-08-31')

paths = {
    'orders': out/'downloads'/'sap_orders_2026-08-01_2026-08-31.csv',
    'invoices': out/'downloads'/'sap_invoices_2026-08-01_2026-08-31.csv',
    'expenses': out/'expense_2026-08-01_2026-08-31_approved.csv',
    'quotes': out/'quotes_2026-08-01_2026-08-31_accepted.json',
    'workbook': out/'project_report_2026-08-01_2026-08-31.xlsx',
    'sep_orders': out/'downloads'/'sap_orders_2026-09-01_2026-09-30.csv',
    'sep_invoices': out/'downloads'/'sap_invoices_2026-09-01_2026-09-30.csv',
    'sep_expenses': out/'expense_2026-09-01_2026-09-30_approved.csv',
    'sep_quotes': out/'quotes_2026-09-01_2026-09-30_accepted.json',
    'sep_workbook': out/'project_report_2026-09-01_2026-09-30.xlsx',
}
print('Reopening outputs for final validation:')
for name, p in paths.items():
    print(f'- {name}: {p.relative_to(ws)} exists={p.exists()} size={p.stat().st_size if p.exists() else None}')
    assert p.exists(), f'Missing {name}: {p}'

def read_csv(p):
    with p.open('r', encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def dec(v, label):
    try:
        return Decimal(str(v))
    except InvalidOperation:
        raise AssertionError(f'{label} not numeric: {v!r}')

def d(v, label):
    try:
        return date.fromisoformat(str(v))
    except Exception:
        raise AssertionError(f'{label} not ISO date: {v!r}')

orders = read_csv(paths['orders'])
invoices = read_csv(paths['invoices'])
expenses = read_csv(paths['expenses'])
quotes_obj = json.loads(paths['quotes'].read_text(encoding='utf-8'))
quotes = quotes_obj['items']

# CSV/JSON checks
assert len(orders) == 7, len(orders)
assert len(invoices) == 6, len(invoices)
assert len(expenses) == 4, len(expenses)
assert len(quotes) == 4, len(quotes)
for r in orders:
    assert start <= d(r['OrderDate'], 'OrderDate') <= end, r
    assert r['Status'] != 'Cancelled', r
    assert isinstance(r['OrderID'], str) and r['OrderID'].startswith('0045') and len(r['OrderID']) == 10, r['OrderID']
    dec(r['NetAmount'], 'Order NetAmount')
for r in invoices:
    assert start <= d(r['InvoiceDate'], 'InvoiceDate') <= end, r
    assert isinstance(r['InvoiceID'], str) and r['InvoiceID'].startswith('0090') and len(r['InvoiceID']) == 10, r['InvoiceID']
    for c in ['Amount','PaidAmount','Balance']:
        dec(r[c], f'Invoice {c}')
for r in expenses:
    assert start <= d(r['ExpenseDate'], 'ExpenseDate') <= end, r
    assert r['Status'] == 'Approved', r
    assert isinstance(r['ExpenseID'], str) and r['ExpenseID'].startswith('EXP-'), r['ExpenseID']
    dec(r['Amount'], 'Expense Amount')
assert set(quotes_obj.keys()) == {'schema_version','items','count'}, quotes_obj.keys()
assert quotes_obj['count'] == len(quotes), 'Quote count mismatch'
for r in quotes:
    assert start <= d(r['QuoteDate'], 'QuoteDate') <= end, r
    assert r['Status'] == 'Accepted', r
    assert isinstance(r['QuoteID'], str) and r['QuoteID'].startswith('Q'), r['QuoteID']
    for c in ['Subtotal','NetAmount']:
        dec(r[c], f'Quote {c}')

# Workbook checks
wb = load_workbook(paths['workbook'], data_only=True, read_only=True)
expected_sheets = ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
assert wb.sheetnames == expected_sheets, wb.sheetnames
report = wb['Report']
assert report['B3'].value == '2026-08-01', report['B3'].value
assert report['D3'].value == '2026-08-31', report['D3'].value
assert 'Cancelled SAP orders removed: 1' in str(report['B9'].value), report['B9'].value

def records_from_sheet(shname):
    sh = wb[shname]
    headers = [sh.cell(4, c).value for c in range(1, sh.max_column+1)]
    rows = []
    for r in range(5, sh.max_row+1):
        vals = [sh.cell(r, c).value for c in range(1, sh.max_column+1)]
        if any(v is not None for v in vals):
            rows.append(dict(zip(headers, vals)))
    return headers, rows

headers_orders, x_orders = records_from_sheet('Orders')
headers_invoices, x_invoices = records_from_sheet('Invoices')
headers_expenses, x_expenses = records_from_sheet('Expenses')
headers_quotes, x_quotes = records_from_sheet('Quotes')
assert len(x_orders) == len(orders), (len(x_orders), len(orders))
assert len(x_invoices) == len(invoices), (len(x_invoices), len(invoices))
assert len(x_expenses) == len(expenses), (len(x_expenses), len(expenses))
assert len(x_quotes) == len(quotes), (len(x_quotes), len(quotes))

# ID and numeric type checks in workbook
for rows, idcol, amount_cols, name in [
    (x_orders, 'OrderID', ['NetAmount','AmountTWD'], 'Orders'),
    (x_invoices, 'InvoiceID', ['Amount','PaidAmount','Balance','BalanceTWD'], 'Invoices'),
    (x_expenses, 'ExpenseID', ['Amount','AmountTWD'], 'Expenses'),
    (x_quotes, 'QuoteID', ['NetAmount','AmountTWD'], 'Quotes'),
]:
    for r in rows:
        assert isinstance(r[idcol], str), f'{name} ID not text: {r[idcol]!r}'
        for c in amount_cols:
            if r[c] is not None:
                assert isinstance(r[c], (int,float)), f'{name} {c} not numeric: {r[c]!r}'

# Summary recalculation from workbook detail rows (ProjectID only, missing ProjectID excluded)
def add(sums, rows, project_col, amount_col):
    for r in rows:
        pid = r.get(project_col)
        if pid:
            sums[pid] += Decimal(str(r.get(amount_col) or 0))

order_sums = defaultdict(Decimal); invoice_sums = defaultdict(Decimal); expense_sums = defaultdict(Decimal); quote_sums = defaultdict(Decimal)
add(order_sums, x_orders, 'ProjectID', 'AmountTWD')
add(invoice_sums, x_invoices, 'ProjectID', 'BalanceTWD')
add(expense_sums, x_expenses, 'ProjectID', 'AmountTWD')
add(quote_sums, x_quotes, 'ProjectID', 'AmountTWD')
summary_headers, summary_rows = records_from_sheet('Summary')
summary_pids = {r['ProjectID'] for r in summary_rows}
expected_pids = set(order_sums) | set(invoice_sums) | set(expense_sums) | set(quote_sums)
assert summary_pids == expected_pids, (summary_pids, expected_pids)
for r in summary_rows:
    pid = r['ProjectID']
    checks = [
        ('OrderTWD', order_sums[pid]),
        ('InvoiceBalanceTWD', invoice_sums[pid]),
        ('ApprovedExpenseTWD', expense_sums[pid]),
        ('AcceptedQuoteTWD', quote_sums[pid]),
    ]
    for col, exp in checks:
        got = Decimal(str(r[col] or 0))
        assert abs(got-exp) < Decimal('0.001'), f'{pid} {col}: got {got}, expected {exp}'

# Exceptions: missing project IDs are present
exc_headers, exc_rows = records_from_sheet('Exceptions')
assert len(exc_rows) == 3, len(exc_rows)
assert {r['DocumentID'] for r in exc_rows} == {'0045000106','0045000108','0090000206'}, exc_rows

print('Final validation passed.')
print(f'Counts: orders={len(orders)}, invoices={len(invoices)}, expenses={len(expenses)}, quotes={len(quotes)}, exceptions={len(exc_rows)}')
print('August output files:')
for key in ['workbook','orders','invoices','expenses','quotes']:
    p = paths[key]
    print('-', p.relative_to(ws), p.stat().st_size, 'bytes')
print('Non-overwrite check: September dated outputs still present:')
for key in ['sep_workbook','sep_orders','sep_invoices','sep_expenses','sep_quotes']:
    p = paths[key]
    print('-', p.relative_to(ws), p.stat().st_size, 'bytes')
