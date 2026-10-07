from pathlib import Path
import csv, json, os
from openpyxl import load_workbook

root = Path.cwd()
print('WORKSPACE', root)
print('INPUT FILES:')
for p in sorted((root/'input').glob('*')):
    print('-', p.relative_to(root), 'exists', p.exists(), 'size', p.stat().st_size)
print('OUTPUT FILES:')
for p in sorted((root/'output').rglob('*')):
    if p.is_file():
        print('-', p.relative_to(root), 'size', p.stat().st_size)

def show_csv(rel, n=5):
    p = root/rel
    print('\nCSV', rel, 'exists', p.exists())
    if not p.exists(): return
    with p.open(newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    print('headers', rows[0].keys() if rows else [])
    print('row_count', len(rows))
    for r in rows[:n]: print('sample', r)
    for col in ['Status','OrderDate','InvoiceDate','Date','ExpenseDate']:
        vals = sorted({r.get(col,'') for r in rows if col in r})
        if vals: print(col, 'values/minmax', vals[:10], '...', vals[-3:])
    if 'Status' in (rows[0].keys() if rows else []):
        print('status_counts', {s: sum(1 for r in rows if r.get('Status')==s) for s in sorted({r.get('Status') for r in rows})})

for rel in ['output/downloads/sap_orders_export.csv','output/downloads/sap_orders_2026-09-01_2026-09-30.csv','output/downloads/sap_invoices_export.csv','output/downloads/sap_invoices_2026-09-01_2026-09-30.csv','output/expense_approved_export.csv','output/expense_2026-09-01_2026-09-30_approved.csv']:
    show_csv(rel)

q = root/'output/quotes_2026-09-01_2026-09-30_accepted.json'
print('\nJSON', q.relative_to(root), 'exists', q.exists())
if q.exists():
    data=json.load(q.open(encoding='utf-8'))
    print('type', type(data).__name__, 'count', len(data) if isinstance(data, list) else 'n/a')
    print('sample', data[:3] if isinstance(data, list) else data)

xlsx = root/'output/project_report_2026-09-01_2026-09-30.xlsx'
print('\nXLSX', xlsx.relative_to(root), 'exists', xlsx.exists())
if xlsx.exists():
    wb=load_workbook(xlsx, data_only=False)
    print('sheets', wb.sheetnames)
    for ws in wb.worksheets:
        print('sheet', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
        rows=[]
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,6), values_only=True):
            rows.append(row)
        for r in rows: print(' ', r)
    wb.close()
