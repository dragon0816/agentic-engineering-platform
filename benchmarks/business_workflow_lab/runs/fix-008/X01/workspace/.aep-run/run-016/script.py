from pathlib import Path
import csv, json, re
from decimal import Decimal
from openpyxl import load_workbook

base = Path.cwd()
FROM='2026-09-01'
TO='2026-09-30'
errors=[]
checks=[]

def fail(msg): errors.append(msg)
def ok(msg): checks.append(msg)
def between(s): return FROM <= str(s) <= TO

def read_csv(rel):
    p=base/rel
    if not p.exists():
        fail(f'missing {rel}')
        return [], []
    with p.open(newline='', encoding='utf-8-sig') as f:
        rdr=csv.DictReader(f)
        rows=list(rdr)
    return rdr.fieldnames or [], rows

def sheet_rows(wb, sheet):
    ws=wb[sheet]
    headers=[c.value for c in ws[4]]
    rows=[]
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(headers,row)))
    return headers, rows

def money_equal(a,b,tol=0.01): return abs(float(a)-float(b)) <= tol

required = [
 'output/downloads/sap_orders_export.csv',
 f'output/downloads/sap_orders_{FROM}_{TO}.csv',
 'output/downloads/sap_invoices_export.csv',
 f'output/downloads/sap_invoices_{FROM}_{TO}.csv',
 'output/expense_approved_export.csv',
 f'output/expense_{FROM}_{TO}_approved.csv',
 f'output/quotes_{FROM}_{TO}_accepted.json',
 f'output/project_report_{FROM}_{TO}.xlsx',
]
for rel in required:
    p=base/rel
    if p.exists() and p.stat().st_size>0: ok(f'{rel} exists size={p.stat().st_size}')
    else: fail(f'{rel} missing or empty')

orders_h, orders = read_csv('output/downloads/sap_orders_export.csv')
if orders_h == ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']: ok('orders CSV headers match source names')
else: fail(f'orders headers unexpected: {orders_h}')
if len(orders)==18: ok('orders CSV row count 18')
else: fail(f'orders row count {len(orders)} != 18')
for r in orders:
    if not between(r['OrderDate']): fail(f'order date out of range {r}')
    if r['Status']=='Cancelled': fail(f'cancelled order present {r["OrderID"]}')
    if not isinstance(r['OrderID'], str) or not re.fullmatch(r'\d+', r['OrderID']) or not r['OrderID'].startswith('00'):
        fail(f'OrderID text/leading zeros not preserved {r["OrderID"]}')
    Decimal(r['NetAmount'])
ok('orders are within date range, exclude Cancelled, preserve text IDs and numeric amounts')
_, orders_dated = read_csv(f'output/downloads/sap_orders_{FROM}_{TO}.csv')
if orders_dated == orders: ok('dated orders CSV matches export CSV')
else: fail('dated orders CSV differs from export CSV')

inv_h, inv = read_csv('output/downloads/sap_invoices_export.csv')
if inv_h == ['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance']: ok('invoice CSV headers match source names')
else: fail(f'invoice headers unexpected: {inv_h}')
if len(inv)==13: ok('invoice CSV row count 13')
else: fail(f'invoice row count {len(inv)} != 13')
for r in inv:
    if not between(r['InvoiceDate']): fail(f'invoice date out of range {r}')
    if not isinstance(r['InvoiceID'], str) or not r['InvoiceID'].startswith('00'): fail(f'InvoiceID text/leading zeros not preserved {r["InvoiceID"]}')
    for k in ['Amount','PaidAmount','Balance']: Decimal(r[k])
ok('invoices are within date range, preserve text IDs and numeric amounts')
_, inv_dated = read_csv(f'output/downloads/sap_invoices_{FROM}_{TO}.csv')
if inv_dated == inv: ok('dated invoices CSV matches export CSV')
else: fail('dated invoices CSV differs from export CSV')

exp_h, exp = read_csv('output/expense_approved_export.csv')
if exp_h == ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']: ok('expense CSV headers match source names')
else: fail(f'expense headers unexpected: {exp_h}')
if len(exp)==15: ok('expense CSV row count 15')
else: fail(f'expense row count {len(exp)} != 15')
for r in exp:
    if not between(r['ExpenseDate']): fail(f'expense date out of range {r}')
    if r['Status']!='Approved': fail(f'non-approved expense present {r["ExpenseID"]}: {r["Status"]}')
    if not isinstance(r['ExpenseID'], str): fail(f'ExpenseID not text {r["ExpenseID"]}')
    Decimal(r['Amount'])
ok('expenses are within date range, Approved only, preserve text IDs and numeric amounts')
_, exp_dated = read_csv(f'output/expense_{FROM}_{TO}_approved.csv')
if exp_dated == exp: ok('dated expense CSV matches export CSV')
else: fail('dated expense CSV differs from export CSV')

qobj=json.loads((base/f'output/quotes_{FROM}_{TO}_accepted.json').read_text(encoding='utf-8'))
if isinstance(qobj, dict):
    quotes = qobj.get('quotes') if isinstance(qobj.get('quotes'), list) else qobj.get('data') if isinstance(qobj.get('data'), list) else []
else:
    quotes = qobj if isinstance(qobj, list) else []
if len(quotes)==4: ok('quotes JSON accepted count 4')
else: fail(f'quotes count {len(quotes)} != 4')
for r in quotes:
    if not between(r.get('QuoteDate','')): fail(f'quote date out of range {r}')
    if r.get('Status')!='Accepted': fail(f'non-accepted quote present {r}')
    if not isinstance(r.get('QuoteID'), str) or not re.fullmatch(r'Q\d+', r.get('QuoteID','')):
        fail(f'QuoteID text/format not preserved {r}')
    Decimal(str(r.get('NetAmount')))
ok('quotes are within date range, Accepted only, preserve text IDs and numeric amounts')

wb=load_workbook(base/f'output/project_report_{FROM}_{TO}.xlsx', data_only=False)
expected_sheets=['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
if wb.sheetnames == expected_sheets: ok('workbook sheets preserved')
else: fail(f'workbook sheets {wb.sheetnames} != {expected_sheets}')
if wb['Report']['B3'].value == FROM and wb['Report']['D3'].value == TO: ok('Report sheet dates match run period')
else: fail(f'Report dates are {wb["Report"]["B3"].value}/{wb["Report"]["D3"].value}')

_, w_orders=sheet_rows(wb,'Orders')
_, w_inv=sheet_rows(wb,'Invoices')
_, w_exp=sheet_rows(wb,'Expenses')
_, w_quotes=sheet_rows(wb,'Quotes')
_, w_sum=sheet_rows(wb,'Summary')
if len(w_orders)==len(orders): ok('Orders sheet row count matches CSV')
else: fail(f'Orders sheet row count {len(w_orders)} != CSV {len(orders)}')
if len(w_inv)==len(inv): ok('Invoices sheet row count matches CSV')
else: fail(f'Invoices sheet row count {len(w_inv)} != CSV {len(inv)}')
if len(w_exp)==len(exp): ok('Expenses sheet row count matches CSV')
else: fail(f'Expenses sheet row count {len(w_exp)} != CSV {len(exp)}')
if len(w_quotes)==len(quotes): ok('Quotes sheet row count matches JSON')
else: fail(f'Quotes sheet row count {len(w_quotes)} != JSON {len(quotes)}')

for r in w_orders:
    if not isinstance(r['OrderID'], str) or not r['OrderID'].startswith('00'): fail(f'Workbook OrderID not text {r}')
    if r['Status']=='Cancelled' or not between(r['OrderDate']): fail(f'Workbook order invalid filter {r}')
    if not isinstance(r['NetAmount'], (int,float)) or not isinstance(r['AmountTWD'], (int,float)): fail(f'Workbook order amounts not numeric {r}')
for r in w_inv:
    if not isinstance(r['InvoiceID'], str) or not r['InvoiceID'].startswith('00'): fail(f'Workbook InvoiceID not text {r}')
    if not between(r['InvoiceDate']): fail(f'Workbook invoice invalid date {r}')
    for k in ['Amount','PaidAmount','Balance','BalanceTWD']:
        if not isinstance(r[k], (int,float)): fail(f'Workbook invoice {k} not numeric {r}')
for r in w_exp:
    if not isinstance(r['ExpenseID'], str): fail(f'Workbook ExpenseID not text {r}')
    if r['Status']!='Approved' or not between(r['ExpenseDate']): fail(f'Workbook expense invalid filter {r}')
    for k in ['Amount','AmountTWD']:
        if not isinstance(r[k], (int,float)): fail(f'Workbook expense {k} not numeric {r}')
for r in w_quotes:
    if not isinstance(r['QuoteID'], str) or not re.fullmatch(r'Q\d+', r['QuoteID']): fail(f'Workbook QuoteID not text/format {r}')
    if r['Status']!='Accepted' or not between(r['QuoteDate']): fail(f'Workbook quote invalid filter {r}')
    for k in ['NetAmount','AmountTWD']:
        if not isinstance(r[k], (int,float)): fail(f'Workbook quote {k} not numeric {r}')
ok('workbook filters, text IDs, and numeric amounts checked')

projects=sorted({r['ProjectID'] for rows in [w_orders,w_inv,w_exp,w_quotes] for r in rows if r.get('ProjectID')})
summary={r['ProjectID']:r for r in w_sum if r.get('ProjectID')}
if sorted(summary)==projects: ok('Summary project IDs match detail sheets')
else: fail(f'Summary project IDs {sorted(summary)} != detail projects {projects}')

def sum_by(rows, field, project):
    return sum(float(r[field] or 0) for r in rows if r.get('ProjectID')==project)
for p in projects:
    sr=summary[p]
    comps={
        'OrderTWD': sum_by(w_orders,'AmountTWD',p),
        'InvoiceBalanceTWD': sum_by(w_inv,'BalanceTWD',p),
        'ApprovedExpenseTWD': sum_by(w_exp,'AmountTWD',p),
        'AcceptedQuoteTWD': sum_by(w_quotes,'AmountTWD',p),
    }
    for k,v in comps.items():
        if not money_equal(sr[k], v): fail(f'Summary {p} {k}={sr[k]} != detail sum {v}')
ok('Summary amounts match detail sheet sums')

print('CHECKS PASSED:')
for c in checks:
    print('-', c)
if errors:
    print('\nERRORS:')
    for e in errors: print('-', e)
    raise SystemExit(1)
print('\nAll verification checks passed.')
print('\nFinal files:')
for rel in required:
    p=base/rel
    print(f'- {rel} ({p.stat().st_size} bytes)')
