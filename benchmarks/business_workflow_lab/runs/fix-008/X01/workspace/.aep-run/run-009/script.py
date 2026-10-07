from pathlib import Path
import csv, json
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from openpyxl import load_workbook

FROM='2026-09-01'; TO='2026-09-30'
root=Path.cwd()
out=root/'output'; downloads=out/'downloads'; inp=root/'input'
files={
 'report': out/f'project_report_{FROM}_{TO}.xlsx',
 'orders_export': downloads/'sap_orders_export.csv',
 'orders_dated': downloads/f'sap_orders_{FROM}_{TO}.csv',
 'invoices_export': downloads/'sap_invoices_export.csv',
 'invoices_dated': downloads/f'sap_invoices_{FROM}_{TO}.csv',
 'expense_export': out/'expense_approved_export.csv',
 'expense_dated': out/f'expense_{FROM}_{TO}_approved.csv',
 'quotes': out/f'quotes_{FROM}_{TO}_accepted.json',
 'template': inp/'project_template.xlsx',
 'fx': inp/'fx_rates.csv',
}
print('Checking files:')
for k,p in files.items():
    print(k, p.relative_to(root), 'exists=', p.exists(), 'size=', p.stat().st_size if p.exists() else None)
    assert p.exists(), f'missing {k}: {p}'

TWOPLACES=Decimal('0.01')
def D(x):
    if x is None or str(x).strip()=='' or str(x).strip()=='None': return Decimal('0.00')
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def q(x): return Decimal(x).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def in_range(s): return FROM <= (s or '') <= TO

def rows(path):
    with path.open(newline='', encoding='utf-8-sig') as f:
        r=list(csv.DictReader(f)); h=f.fieldnames
    print(path.relative_to(root), 'rows=', len(r), 'headers=', h)
    return r,h

fx={}
with files['fx'].open(newline='', encoding='utf-8-sig') as f:
    rd=csv.DictReader(f)
    for row in rd:
        fx[row['Currency']]=Decimal(row['RateToTWD'])
print('fx=', fx)
def twd(amount,currency): return q(Decimal(str(amount))*fx[currency])

orders, oh=rows(files['orders_export'])
orders2, oh2=rows(files['orders_dated'])
assert orders==orders2, 'orders export and dated copy differ'
assert all(in_range(r['OrderDate']) for r in orders), 'orders contain out-of-range dates'
assert all(r['Status']!='Cancelled' for r in orders), 'orders contain Cancelled rows'

invoices, ih=rows(files['invoices_export'])
invoices2, ih2=rows(files['invoices_dated'])
assert invoices==invoices2, 'invoices export and dated copy differ'
assert all(in_range(r['InvoiceDate']) for r in invoices), 'invoices contain out-of-range dates'

expenses, eh=rows(files['expense_export'])
expenses2, eh2=rows(files['expense_dated'])
assert expenses==expenses2, 'expense export and dated copy differ'
assert all(in_range(r['ExpenseDate']) and r['Status']=='Approved' for r in expenses), 'expenses contain invalid rows'

qdata=json.loads(files['quotes'].read_text(encoding='utf-8'))
quotes=qdata['items']
print(files['quotes'].relative_to(root), 'items=', len(quotes), 'count=', qdata.get('count'))
assert qdata.get('count')==len(quotes), 'quote count mismatch'
assert all(in_range(r['QuoteDate']) and r['Status']=='Accepted' for r in quotes), 'quotes contain invalid rows'

wb=load_workbook(files['report'], data_only=False)
print('workbook sheets=', wb.sheetnames)
expected_sheets=['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
assert wb.sheetnames==expected_sheets, f'sheets mismatch: {wb.sheetnames}'
expected_headers={
 'Orders':['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
 'Invoices':['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
 'Expenses':['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
 'Quotes':['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
 'Summary':['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
 'Exceptions':['Source','DocumentID','ProjectID','Reason'],
}
for s,heads in expected_headers.items():
    ws=wb[s]
    got=[ws.cell(4,c).value for c in range(1,len(heads)+1)]
    print(s, 'headers=', got, 'rows=', ws.max_row-4)
    assert got==heads, f'{s} headers mismatch'

# row counts vs sources
assert wb['Orders'].max_row-4==len(orders), 'Orders row count mismatch'
assert wb['Invoices'].max_row-4==len(invoices), 'Invoices row count mismatch'
assert wb['Expenses'].max_row-4==len(expenses), 'Expenses row count mismatch'
assert wb['Quotes'].max_row-4==len(quotes), 'Quotes row count mismatch'

# ID text and numeric amounts
for sheet,id_cols,amount_cols in [
 ('Orders',[1,2,3],[7,8]),('Invoices',[1,2,3,4],[8,9,10,11]),('Expenses',[1,2,3],[7,8]),('Quotes',[1,2,3],[8,9]),('Summary',[1],[2,3,4,5])]:
    ws=wb[sheet]
    for r in range(5,ws.max_row+1):
        if all(ws.cell(r,c).value in (None,'') for c in range(1, ws.max_column+1)): continue
        for c in id_cols:
            v=ws.cell(r,c).value
            if v not in (None,''):
                assert isinstance(v,str), f'{sheet} row {r} col {c} should be text ID, got {v!r}/{type(v)}'
        for c in amount_cols:
            v=ws.cell(r,c).value
            assert isinstance(v,(int,float)), f'{sheet} row {r} col {c} should be numeric, got {v!r}/{type(v)}'

# summary from workbook detail sheets
calc=defaultdict(lambda:[Decimal('0.00'),Decimal('0.00'),Decimal('0.00'),Decimal('0.00')])
for sheet,pcol,acol,idx in [('Orders',2,8,0),('Invoices',3,11,1),('Expenses',2,8,2),('Quotes',2,9,3)]:
    ws=wb[sheet]
    for r in range(5, ws.max_row+1):
        pid=ws.cell(r,pcol).value
        amt=ws.cell(r,acol).value
        if pid:
            calc[pid][idx]+=D(amt)
saved={}
ws=wb['Summary']
for r in range(5, ws.max_row+1):
    pid=ws.cell(r,1).value
    if pid:
        saved[pid]=[D(ws.cell(r,c).value) for c in range(2,6)]
print('summary projects saved=', sorted(saved), 'calc=', sorted(calc))
assert set(saved)==set(calc), 'summary project set mismatch'
for pid in saved:
    assert saved[pid]==[q(x) for x in calc[pid]], f'summary mismatch for {pid}: {saved[pid]} vs {calc[pid]}'

# ensure workbook was made from template by checking instructions/title persisted
assert wb['Report']['A1'].value == 'Project Integration Report 專案整合報表', 'template title missing'
assert wb['Instructions']['A1'].value == '填寫說明 Instructions（請勿修改本工作表）', 'template Instructions missing'

print('VALIDATION PASSED')
print('Counts: orders=%d invoices=%d expenses=%d quotes=%d summary_projects=%d exceptions=%d' % (len(orders),len(invoices),len(expenses),len(quotes),len(saved), wb['Exceptions'].max_row-4))
