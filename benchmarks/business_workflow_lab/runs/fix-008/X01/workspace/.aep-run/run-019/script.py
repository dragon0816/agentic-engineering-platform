import csv, json, os, sys, subprocess, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from collections import defaultdict
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
FROM = params['from_date']
TO = params['to_date']
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
downloads = output_dir / 'downloads'
output_dir.mkdir(parents=True, exist_ok=True)
downloads.mkdir(parents=True, exist_ok=True)

orders_raw_path = downloads / 'sap_orders_export.csv'
invoices_path = downloads / 'sap_invoices_export.csv'
expenses_path = output_dir / 'expense_approved_export.csv'
orders_dated_path = downloads / f'sap_orders_{FROM}_{TO}.csv'
invoices_dated_path = downloads / f'sap_invoices_{FROM}_{TO}.csv'
expenses_dated_path = output_dir / f'expense_{FROM}_{TO}_approved.csv'
quote_json_path = output_dir / f'quotes_{FROM}_{TO}_accepted.json'
report_path = output_dir / f'project_report_{FROM}_{TO}.xlsx'

required = [input_dir/'project_template.xlsx', input_dir/'fx_rates.csv', orders_raw_path, invoices_path, expenses_path]
print('Inputs expected:')
for p in required:
    print(f'- {p}: exists={p.exists()} size={p.stat().st_size if p.exists() else "missing"}')
missing = [str(p) for p in required if not p.exists()]
if missing:
    raise FileNotFoundError('Missing required input/export files: ' + '; '.join(missing))

TWOPLACES = Decimal('0.01')
def D(v):
    if v is None or str(v).strip() == '':
        return Decimal('0.00')
    return Decimal(str(v).strip())
def q2(v):
    return Decimal(v).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def money_cell(v):
    return float(q2(v))
def in_range(date_s):
    return FROM <= (date_s or '') <= TO

def read_csv_dict(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        headers = reader.fieldnames
    print(f'Read {len(rows)} rows from {path.name}; headers={headers}')
    return rows, headers

def write_csv_dict(path, rows, headers):
    with open(path, 'w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=headers)
        w.writeheader()
        w.writerows(rows)
    print(f'Wrote {len(rows)} rows to {path}')

fx = {}
with open(input_dir/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('FX headers:', reader.fieldnames)
    for row in reader:
        cur = (row.get('Currency') or row.get('currency') or '').strip()
        rate = row.get('RateToTWD') or row.get('TWD') or row.get('rate_to_twd') or row.get('rate') or row.get('Rate')
        if cur:
            fx[cur] = D(rate)
print('FX currencies:', sorted(fx))
def to_twd(amount, currency):
    currency = (currency or '').strip()
    if currency not in fx:
        raise ValueError(f'Missing FX rate for currency {currency!r}')
    return q2(D(amount) * fx[currency])

orders_raw, orders_headers = read_csv_dict(orders_raw_path)
invoices, invoices_headers = read_csv_dict(invoices_path)
expenses, expenses_headers = read_csv_dict(expenses_path)
orders = [r for r in orders_raw if in_range(r.get('OrderDate')) and (r.get('Status') or '').strip() != 'Cancelled']
removed_cancelled = len([r for r in orders_raw if in_range(r.get('OrderDate')) and (r.get('Status') or '').strip() == 'Cancelled'])
removed_other = len(orders_raw) - len(orders) - removed_cancelled
invoices = [r for r in invoices if in_range(r.get('InvoiceDate'))]
expenses = [r for r in expenses if in_range(r.get('ExpenseDate')) and (r.get('Status') or '').strip() == 'Approved']
write_csv_dict(orders_raw_path, orders, orders_headers)
write_csv_dict(orders_dated_path, orders, orders_headers)
write_csv_dict(invoices_path, invoices, invoices_headers)
write_csv_dict(invoices_dated_path, invoices, invoices_headers)
write_csv_dict(expenses_path, expenses, expenses_headers)
write_csv_dict(expenses_dated_path, expenses, expenses_headers)

cli_py = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
cmds = [
    [sys.executable, str(cli_py), 'quotes', 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json'],
    [sys.executable, str(cli_py), 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json'],
]
last = None
for cmd in cmds:
    for attempt in range(1, 4):
        proc = subprocess.run(cmd, text=True, capture_output=True)
        last = proc
        if proc.returncode == 0:
            break
        if proc.returncode == 4 and attempt < 3:
            continue
        break
    if last and last.returncode == 0:
        break
if not last or last.returncode != 0:
    raise RuntimeError(f'quote_cli failed rc={last.returncode if last else "none"}: {last.stderr if last else "no run"}')
quote_payload = json.loads(last.stdout)
quotes = quote_payload.get('items', quote_payload if isinstance(quote_payload, list) else [])
quotes = [r for r in quotes if in_range(r.get('QuoteDate')) and (r.get('Status') or '').strip() == 'Accepted']
if isinstance(quote_payload, dict):
    quote_payload['items'] = quotes
    quote_payload['count'] = len(quotes)
else:
    quote_payload = {'schema_version': 1, 'items': quotes, 'count': len(quotes)}
quote_json_path.write_text(json.dumps(quote_payload, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'Read/wrote {len(quotes)} accepted quotes from quote_cli; saved {quote_json_path.name}')

for r in orders:
    assert in_range(r.get('OrderDate')), ('order out of range', r)
    assert (r.get('Status') or '').strip() != 'Cancelled', ('cancelled order included', r)
for r in invoices:
    assert in_range(r.get('InvoiceDate')), ('invoice out of range', r)
for r in expenses:
    assert in_range(r.get('ExpenseDate')), ('expense out of range', r)
    assert (r.get('Status') or '').strip() == 'Approved', ('non-approved expense included', r)
for r in quotes:
    assert in_range(r.get('QuoteDate')), ('quote out of range', r)
    assert (r.get('Status') or '').strip() == 'Accepted', ('non-accepted quote included', r)

shutil.copyfile(input_dir/'project_template.xlsx', report_path)
wb = load_workbook(report_path)
headers_expected = {
    'Orders': ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
    'Invoices': ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
    'Expenses': ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
    'Quotes': ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
    'Summary': ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
    'Exceptions': ['Source','DocumentID','ProjectID','Reason'],
}
for sheet, heads in headers_expected.items():
    actual = [wb[sheet].cell(4, c).value for c in range(1, len(heads)+1)]
    if actual != heads:
        raise AssertionError(f'Unexpected headers in {sheet}: {actual}')

def clear_sheet(sheet_name):
    ws = wb[sheet_name]
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)
    return ws

def set_text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.number_format = '@'
def set_date(cell, value):
    cell.value = value or None
    cell.number_format = 'yyyy-mm-dd'
def set_amount(cell, value):
    cell.value = money_cell(value)
    cell.number_format = '0.00'

summary = defaultdict(lambda: {'order': Decimal('0.00'), 'invoice': Decimal('0.00'), 'expense': Decimal('0.00'), 'quote': Decimal('0.00')})
exceptions = []

ws = clear_sheet('Orders')
for i, r in enumerate(orders, 5):
    amt_twd = to_twd(r.get('NetAmount'), r.get('Currency'))
    set_text(ws.cell(i,1), r.get('OrderID')); set_text(ws.cell(i,2), r.get('ProjectID')); set_text(ws.cell(i,3), r.get('CustomerID'))
    set_date(ws.cell(i,4), r.get('OrderDate')); set_text(ws.cell(i,5), r.get('Status')); set_text(ws.cell(i,6), r.get('Currency'))
    set_amount(ws.cell(i,7), D(r.get('NetAmount'))); set_amount(ws.cell(i,8), amt_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['order'] += amt_twd
    else:
        exceptions.append(['Orders', r.get('OrderID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Invoices')
for i, r in enumerate(invoices, 5):
    balance = D(r.get('Balance')) if r.get('Balance') not in (None, '') else D(r.get('Amount')) - D(r.get('PaidAmount'))
    bal_twd = to_twd(balance, r.get('Currency'))
    set_text(ws.cell(i,1), r.get('InvoiceID')); set_text(ws.cell(i,2), r.get('OrderID')); set_text(ws.cell(i,3), r.get('ProjectID')); set_text(ws.cell(i,4), r.get('CustomerID'))
    set_date(ws.cell(i,5), r.get('InvoiceDate')); set_date(ws.cell(i,6), r.get('DueDate')); set_text(ws.cell(i,7), r.get('Currency'))
    set_amount(ws.cell(i,8), D(r.get('Amount'))); set_amount(ws.cell(i,9), D(r.get('PaidAmount'))); set_amount(ws.cell(i,10), balance); set_amount(ws.cell(i,11), bal_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['invoice'] += bal_twd
    else:
        exceptions.append(['Invoices', r.get('InvoiceID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Expenses')
for i, r in enumerate(expenses, 5):
    amt_twd = to_twd(r.get('Amount'), r.get('Currency'))
    set_text(ws.cell(i,1), r.get('ExpenseID')); set_text(ws.cell(i,2), r.get('ProjectID')); set_text(ws.cell(i,3), r.get('EmployeeID'))
    set_date(ws.cell(i,4), r.get('ExpenseDate')); set_text(ws.cell(i,5), r.get('Status')); set_text(ws.cell(i,6), r.get('Currency'))
    set_amount(ws.cell(i,7), D(r.get('Amount'))); set_amount(ws.cell(i,8), amt_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['expense'] += amt_twd
    else:
        exceptions.append(['Expenses', r.get('ExpenseID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Quotes')
for i, r in enumerate(quotes, 5):
    amt_twd = to_twd(r.get('NetAmount'), r.get('Currency'))
    set_text(ws.cell(i,1), r.get('QuoteID')); set_text(ws.cell(i,2), r.get('ProjectID')); set_text(ws.cell(i,3), r.get('CustomerID'))
    set_date(ws.cell(i,4), r.get('QuoteDate')); set_date(ws.cell(i,5), r.get('ValidUntil')); set_text(ws.cell(i,6), r.get('Status')); set_text(ws.cell(i,7), r.get('Currency'))
    set_amount(ws.cell(i,8), D(r.get('NetAmount'))); set_amount(ws.cell(i,9), amt_twd)
    if r.get('ProjectID'):
        summary[r['ProjectID']]['quote'] += amt_twd
    else:
        exceptions.append(['Quotes', r.get('QuoteID'), r.get('ProjectID'), 'Missing ProjectID; excluded from project summary'])

ws = clear_sheet('Summary')
projects = sorted(summary)
for i, pid in enumerate(projects, 5):
    vals = summary[pid]
    set_text(ws.cell(i,1), pid)
    set_amount(ws.cell(i,2), vals['order'])
    set_amount(ws.cell(i,3), vals['invoice'])
    set_amount(ws.cell(i,4), vals['expense'])
    set_amount(ws.cell(i,5), vals['quote'])
last_row = max(5, 4 + len(projects))
for col, cell in [('B','B2'),('C','C2'),('D','D2'),('E','E2')]:
    ws[cell] = f'=SUM({col}5:{col}{last_row})'
    ws[cell].number_format = '0.00'

ws = clear_sheet('Exceptions')
for i, row in enumerate(exceptions, 5):
    for c, v in enumerate(row, 1):
        set_text(ws.cell(i,c), v)

rep = wb['Report']
rep['B3'] = FROM
rep['D3'] = TO
rep['B4'] = 'SAP Orders: OrderDate in range and Status != Cancelled; SAP Invoices: InvoiceDate in range; Expense: ExpenseDate in range and Status = Approved; Quotes: QuoteDate in range and Status = Accepted; summary by ProjectID only.'
rep['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B7'] = 'SAP web exports; Expense Viewer export; quote_cli Accepted JSON; fx_rates.csv'
rep['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B9'] = f'Cancelled SAP orders removed: {removed_cancelled}; other raw order rows removed by date/status normalization: {removed_other}. Orders, invoice balances, approved expenses, and accepted quotes are separate columns; orders and quotes are not added as revenue; CustomerName is not used as a join key.'
wb.save(report_path)

# Saved-output verification.
wb2 = load_workbook(report_path, data_only=False)
for sheet, heads in headers_expected.items():
    actual = [wb2[sheet].cell(4, c).value for c in range(1, len(heads)+1)]
    assert actual == heads, f'headers after save mismatch {sheet}: {actual}'
for sheet, id_cols, amount_cols in [
    ('Orders',[1,2,3],[7,8]), ('Invoices',[1,2,3,4],[8,9,10,11]), ('Expenses',[1,2,3],[7,8]), ('Quotes',[1,2,3],[8,9]), ('Summary',[1],[2,3,4,5])]:
    ws = wb2[sheet]
    for row in range(5, ws.max_row+1):
        if all(ws.cell(row,c).value in (None,'') for c in range(1, ws.max_column+1)):
            continue
        for c in id_cols:
            v = ws.cell(row,c).value
            if v not in (None, ''):
                assert isinstance(v, str), f'{sheet} row {row} col {c} ID is not text: {v!r}'
        for c in amount_cols:
            v = ws.cell(row,c).value
            assert isinstance(v, (int,float)), f'{sheet} row {row} col {c} amount is not numeric: {v!r}'
assert wb2['Orders'].max_row == 4 + len(orders), f'Orders row count mismatch {wb2["Orders"].max_row} vs {4+len(orders)}'
assert wb2['Invoices'].max_row == 4 + len(invoices), 'Invoices row count mismatch'
assert wb2['Expenses'].max_row == 4 + len(expenses), 'Expenses row count mismatch'
assert wb2['Quotes'].max_row == 4 + len(quotes), 'Quotes row count mismatch'
assert wb2['Summary'].max_row == 4 + len(projects), 'Summary row count mismatch'
calc = defaultdict(lambda: [Decimal('0.00'), Decimal('0.00'), Decimal('0.00'), Decimal('0.00')])
for sheet, pcol, acol, idx in [('Orders',2,8,0), ('Invoices',3,11,1), ('Expenses',2,8,2), ('Quotes',2,9,3)]:
    ws = wb2[sheet]
    for row in range(5, ws.max_row+1):
        pid = ws.cell(row,pcol).value
        amt = ws.cell(row,acol).value
        if pid:
            calc[pid][idx] += D(amt)
ws = wb2['Summary']
saved_summary = {}
for row in range(5, ws.max_row+1):
    pid = ws.cell(row,1).value
    if pid:
        saved_summary[pid] = [D(ws.cell(row,c).value) for c in range(2,6)]
assert set(saved_summary) == set(calc), f'Summary project set mismatch {set(saved_summary)} vs {set(calc)}'
for pid in saved_summary:
    for a,b in zip(saved_summary[pid], calc[pid]):
        assert q2(a) == q2(b), f'Summary mismatch {pid}: {saved_summary[pid]} vs {calc[pid]}'
chk_orders, _ = read_csv_dict(orders_raw_path)
assert all((r.get('Status') or '').strip() != 'Cancelled' and in_range(r.get('OrderDate')) for r in chk_orders), 'clean orders export contains invalid rows'
chk_inv, _ = read_csv_dict(invoices_path)
assert all(in_range(r.get('InvoiceDate')) for r in chk_inv), 'clean invoices export contains invalid rows'
chk_exp, _ = read_csv_dict(expenses_path)
assert all((r.get('Status') or '').strip() == 'Approved' and in_range(r.get('ExpenseDate')) for r in chk_exp), 'clean expense export contains invalid rows'
qcheck = json.loads(quote_json_path.read_text(encoding='utf-8'))
assert qcheck.get('count') == len(qcheck.get('items', [])), 'quote count mismatch after save'
assert all((r.get('Status') or '').strip() == 'Accepted' and in_range(r.get('QuoteDate')) for r in qcheck.get('items', [])), 'quote export contains invalid rows'
print('Verification passed.')
print(f'Orders rows: {len(orders)}; cancelled removed: {removed_cancelled}; other removed: {removed_other}')
print(f'Invoices rows: {len(invoices)}; Expenses rows: {len(expenses)}; Quotes rows: {len(quotes)}; Projects: {len(projects)}; Exceptions: {len(exceptions)}')
print('Report:', report_path)
print('Supporting files:')
for p in [orders_raw_path, orders_dated_path, invoices_path, invoices_dated_path, expenses_path, expenses_dated_path, quote_json_path]:
    print('-', p)