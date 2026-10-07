from pathlib import Path
import csv, json, sys
from decimal import Decimal, InvalidOperation
from datetime import date
from openpyxl import load_workbook

ws = Path.cwd()
out = ws/'output'
start = date.fromisoformat('2026-08-01')
end = date.fromisoformat('2026-08-31')

paths = {
    'orders': out/'downloads'/'sap_orders_2026-08-01_2026-08-31.csv',
    'invoices': out/'downloads'/'sap_invoices_2026-08-01_2026-08-31.csv',
    'expenses': out/'expense_2026-08-01_2026-08-31_approved.csv',
    'quotes': out/'quotes_2026-08-01_2026-08-31_accepted.json',
    'workbook': out/'project_report_2026-08-01_2026-08-31.xlsx',
    'sep_orders': out/'downloads'/'sap_orders_2026-09-01_2026-09-30.csv',
    'sep_invoices': out/'downloads'/'sap_invoices_2026-09-01_2026-09-30.csv',
    'sep_expenses': out/'expense_2026-09-01_2026-09-30_approved.csv',
    'sep_quotes': out/'quotes_2026-09-01_2026-09-30_accepted.json',
    'sep_workbook': out/'project_report_2026-09-01_2026-09-30.xlsx',
}
print('Reopening outputs for validation:')
for name, p in paths.items():
    print(f'- {name}: {p.relative_to(ws)} exists={p.exists()} size={p.stat().st_size if p.exists() else None}')
    if not p.exists():
        raise AssertionError(f'Missing {name}: {p}')

def read_csv(p):
    with p.open('r', encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def dec(v, label):
    try:
        return Decimal(str(v))
    except InvalidOperation:
        raise AssertionError(f'{label} not numeric: {v!r}')

def d(v, label):
    try:
        return date.fromisoformat(str(v))
    except Exception:
        raise AssertionError(f'{label} not ISO date: {v!r}')

orders = read_csv(paths['orders'])
invoices = read_csv(paths['invoices'])
expenses = read_csv(paths['expenses'])
quotes_obj = json.loads(paths['quotes'].read_text(encoding='utf-8'))
quotes = quotes_obj['items']

expected_headers = {
    'orders': ['OrderID', 'CustomerID', 'CustomerName', 'ProjectID', 'OrderDate', 'Status', 'Currency', 'NetAmount'],
    'invoices': ['InvoiceID', 'OrderID', 'CustomerID', 'ProjectID', 'InvoiceDate', 'DueDate', 'Currency', 'Amount', 'PaidAmount', 'Balance'],
    'expenses': ['ExpenseID', 'EmployeeID', 'EmployeeName', 'ProjectID', 'ExpenseDate', 'Category', 'Currency', 'Amount', 'Status', 'ReceiptID', 'Description', 'MissingReceipt'],
}
for key, hdrs in expected_headers.items():
    actual = list((orders if key=='orders' else invoices if key=='invoices' else expenses)[0].keys()) if (orders if key=='orders' else invoices if key=='invoices' else expenses) else []
    assert actual == hdrs, f'{key} headers differ: {actual}'

# record-level checks
for r in orders:
    assert start <= d(r['OrderDate'], 'OrderDate') <= end, r
    assert r['Status'] != 'Cancelled', r
    assert isinstance(r['OrderID'], str) and r['OrderID'].startswith('0045') and len(r['OrderID']) == 10, r['OrderID']
    dec(r['NetAmount'], 'Order NetAmount')
for r in invoices:
    assert start <= d(r['InvoiceDate'], 'InvoiceDate') <= end, r
    assert isinstance(r['InvoiceID'], str) and r['InvoiceID'].startswith('0090') and len(r['InvoiceID']) == 10, r['InvoiceID']
    for c in ['Amount','PaidAmount','Balance']:
        dec(r[c], f'Invoice {c}')
for r in expenses:
    assert start <= d(r['ExpenseDate'], 'ExpenseDate') <= end, r
    assert r['Status'] == 'Approved', r
    assert isinstance(r['ExpenseID'], str) and r['ExpenseID'].startswith('EXP-'), r['ExpenseID']
    dec(r['Amount'], 'Expense Amount')
assert set(quotes_obj.keys()) == {'schema_version','items','count'}, quotes_obj.keys()
assert quotes_obj['count'] == len(quotes), 'Quote count mismatch'
for r in quotes:
    assert start <= d(r['QuoteDate'], 'QuoteDate') <= end, r
    assert r['Status'] == 'Accepted', r
    assert isinstance(r['QuoteID'], str) and r['QuoteID'].startswith('Q'), r['QuoteID']
    for c in ['Subtotal','NetAmount']:
        dec(r[c], f'Quote {c}')

# Workbook checks
wb = load_workbook(paths['workbook'], data_only=True, read_only=True)
expected_sheets = ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
assert wb.sheetnames == expected_sheets, wb.sheetnames
report = wb['Report']
# Check dates occur somewhere in Report sheet
report_vals = [cell.value for row in report.iter_rows() for cell in row]
assert any(str(v) == '2026-08-01' or v == start for v in report_vals), 'from date not found in Report'
assert any(str(v) == '2026-08-31' or v == end for v in report_vals), 'to date not found in Report'

# Count rows by reading non-empty data rows after title/header rows. Workflow template row1 title, row2 headers.
def sheet_data_rows(shname):
    sh = wb[shname]
    return [row for row in sh.iter_rows(min_row=3, values_only=True) if any(v is not None for v in row)]
assert len(sheet_data_rows('Orders')) == len(orders), (len(sheet_data_rows('Orders')), len(orders))
assert len(sheet_data_rows('Invoices')) == len(invoices), (len(sheet_data_rows('Invoices')), len(invoices))
assert len(sheet_data_rows('Expenses')) == len(expenses), (len(sheet_data_rows('Expenses')), len(expenses))
assert len(sheet_data_rows('Quotes')) == len(quotes), (len(sheet_data_rows('Quotes')), len(quotes))

# Check workbook ID cells are text-like and amount cells numeric in data rows.
for shname, id_col, amount_cols in [
    ('Orders', 1, [8]),
    ('Invoices', 1, [8,9,10]),
    ('Expenses', 1, [8]),
    ('Quotes', 1, [8,9]),
]:
    sh = wb[shname]
    for row in range(3, sh.max_row+1):
        if not any(sh.cell(row, c).value is not None for c in range(1, sh.max_column+1)):
            continue
        assert isinstance(sh.cell(row, id_col).value, str), f'{shname} row {row} ID not text: {sh.cell(row,id_col).value!r}'
        for c in amount_cols:
            v = sh.cell(row, c).value
            if v is not None:
                assert isinstance(v, (int, float)), f'{shname} row {row} col {c} amount not numeric: {v!r}'

print('Validation passed.')
print(f'Counts: orders={len(orders)}, invoices={len(invoices)}, expenses={len(expenses)}, quotes={len(quotes)}')
print('August output files:')
for key in ['workbook','orders','invoices','expenses','quotes']:
    p = paths[key]
    print('-', p.relative_to(ws), p.stat().st_size, 'bytes')
print('September dated outputs still present:')
for key in ['sep_workbook','sep_orders','sep_invoices','sep_expenses','sep_quotes']:
    p = paths[key]
    print('-', p.relative_to(ws), p.stat().st_size, 'bytes')
