from pathlib import Path
import csv, json, math
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

base = Path.cwd()
FROM='2026-09-01'
TO='2026-09-30'
errors=[]
checks=[]

def fail(msg):
    errors.append(msg)

def ok(msg):
    checks.append(msg)

def read_csv(rel):
    p=base/rel
    if not p.exists():
        fail(f'missing {rel}')
        return [], []
    with p.open(newline='', encoding='utf-8-sig') as f:
        rdr=csv.DictReader(f)
        rows=list(rdr)
        return rdr.fieldnames or [], rows

def between(s):
    return FROM <= s <= TO

def num(x):
    if x is None or x=='':
        return None
    return float(x)

def money_equal(a,b,tol=0.01):
    return abs(float(a)-float(b)) <= tol

def sheet_rows(wb, sheet):
    ws=wb[sheet]
    headers=[c.value for c in ws[4]]
    rows=[]
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(headers,row)))
    return headers, rows

# Required output files
required = [
 'output/downloads/sap_orders_export.csv',
 f'output/downloads/sap_orders_{FROM}_{TO}.csv',
 'output/downloads/sap_invoices_export.csv',
 f'output/downloads/sap_invoices_{FROM}_{TO}.csv',
 'output/expense_approved_export.csv',
 f'output/expense_{FROM}_{TO}_approved.csv',
 f'output/quotes_{FROM}_{TO}_accepted.json',
 f'output/project_report_{FROM}_{TO}.xlsx',
]
for rel in required:
    p=base/rel
    if p.exists() and p.stat().st_size>0:
        ok(f'{rel} exists size={p.stat().st_size}')
    else:
        fail(f'{rel} missing or empty')

# CSV checks
orders_h, orders = read_csv('output/downloads/sap_orders_export.csv')
expected_orders_h=['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']
if orders_h != expected_orders_h: fail(f'orders headers {orders_h} != {expected_orders_h}')
else: ok('orders CSV headers match source names')
if len(orders)!=18: fail(f'orders row count {len(orders)} != 18')
else: ok('orders CSV row count 18')
for r in orders:
    if not between(r['OrderDate']): fail(f'order date out of range {r}')
    if r['Status']=='Cancelled': fail(f'cancelled order present {r["OrderID"]}')
    if not isinstance(r['OrderID'], str) or not r['OrderID'].startswith('00'): fail(f'OrderID text/leading zeros not preserved {r["OrderID"]}')
    try: Decimal(r['NetAmount'])
    except Exception: fail(f'NetAmount not numeric string {r}')
if not any(e.startswith('order date') or e.startswith('cancelled') for e in errors): ok('orders are within date range and exclude Cancelled')

orders_h2, orders2 = read_csv(f'output/downloads/sap_orders_{FROM}_{TO}.csv')
if orders2 != orders: fail('dated orders CSV differs from export CSV')
else: ok('dated orders CSV matches export CSV')

inv_h, inv = read_csv('output/downloads/sap_invoices_export.csv')
expected_inv_h=['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance']
if inv_h != expected_inv_h: fail(f'invoice headers {inv_h} != {expected_inv_h}')
else: ok('invoice CSV headers match source names')
if len(inv)!=13: fail(f'invoice row count {len(inv)} != 13')
else: ok('invoice CSV row count 13')
for r in inv:
    if not between(r['InvoiceDate']): fail(f'invoice date out of range {r}')
    if not r['InvoiceID'].startswith('00'): fail(f'InvoiceID leading zeros not preserved {r["InvoiceID"]}')
    for k in ['Amount','PaidAmount','Balance']:
        try: Decimal(r[k])
        except Exception: fail(f'invoice {k} not numeric string {r}')
ok('invoice dates and numeric fields checked')
_, inv2=read_csv(f'output/downloads/sap_invoices_{FROM}_{TO}.csv')
if inv2 != inv: fail('dated invoices CSV differs from export CSV')
else: ok('dated invoices CSV matches export CSV')

exp_h, exp = read_csv('output/expense_approved_export.csv')
expected_exp_h=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']
if exp_h != expected_exp_h: fail(f'expense headers {exp_h} != {expected_exp_h}')
else: ok('expense CSV headers match source names')
if len(exp)!=15: fail(f'expense row count {len(exp)} != 15')
else: ok('expense CSV row count 15')
for r in exp:
    if not between(r['ExpenseDate']): fail(f'expense date out of range {r}')
    if r['Status']!='Approved': fail(f'non-approved expense present {r["ExpenseID"]}: {r["Status"]}')
    try: Decimal(r['Amount'])
    except Exception: fail(f'expense Amount not numeric string {r}')
ok('expense dates/status/numeric fields checked')
_, exp2=read_csv(f'output/expense_{FROM}_{TO}_approved.csv')
if exp2 != exp: fail('dated expense CSV differs from export CSV')
else: ok('dated expense CSV matches export CSV')

# JSON checks
jp=base/f'output/quotes_{FROM}_{TO}_accepted.json'
qobj=json.loads(jp.read_text(encoding='utf-8'))
if isinstance(qobj, dict):
    if 'quotes' in qobj and isinstance(qobj['quotes'], list): quotes=qobj['quotes']
    elif 'data' in qobj and isinstance(qobj['data'], list): quotes=qobj['data']
    else:
        # first list value
        lists=[v for v in qobj.values() if isinstance(v,list)]
        quotes=lists[0] if lists else []
else:
    quotes=qobj if isinstance(qobj,list) else []
if len(quotes)!=4: fail(f'quotes count {len(quotes)} != 4')
else: ok('quotes JSON accepted count 4')
for r in quotes:
    qdate=str(r.get('QuoteDate',''))
    status=str(r.get('Status',''))
    if not between(qdate): fail(f'quote date out of range {r}')
    if status!='Accepted': fail(f'non-accepted quote present {r}')
    if not str(r.get('QuoteID','')).startswith('Q-'): fail(f'QuoteID text unexpected {r}')
    try: Decimal(str(r.get('NetAmount')))
    except Exception: fail(f'quote NetAmount not numeric {r}')
ok('quote dates/status/numeric fields checked')

# Workbook checks
wb=load_workbook(base/f'output/project_report_{FROM}_{TO}.xlsx', data_only=False)
expected_sheets=['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
if wb.sheetnames != expected_sheets: fail(f'workbook sheets {wb.sheetnames} != {expected_sheets}')
else: ok('workbook sheets preserved')
# report date cells
ws=wb['Report']
if ws['B3'].value != FROM or ws['D3'].value != TO: fail(f'Report dates B3/D3 = {ws["B3"].value}/{ws["D3"].value}')
else: ok('Report sheet dates match run period')

_, w_orders=sheet_rows(wb,'Orders')
_, w_inv=sheet_rows(wb,'Invoices')
_, w_exp=sheet_rows(wb,'Expenses')
_, w_quotes=sheet_rows(wb,'Quotes')
_, w_sum=sheet_rows(wb,'Summary')
if len(w_orders)!=len(orders): fail(f'Orders sheet row count {len(w_orders)} != CSV {len(orders)}')
else: ok('Orders sheet row count matches CSV')
if len(w_inv)!=len(inv): fail(f'Invoices sheet row count {len(w_inv)} != CSV {len(inv)}')
else: ok('Invoices sheet row count matches CSV')
if len(w_exp)!=len(exp): fail(f'Expenses sheet row count {len(w_exp)} != CSV {len(exp)}')
else: ok('Expenses sheet row count matches CSV')
if len(w_quotes)!=len(quotes): fail(f'Quotes sheet row count {len(w_quotes)} != JSON {len(quotes)}')
else: ok('Quotes sheet row count matches JSON')

# workbook row constraints and ID text/numeric TWD fields
for r in w_orders:
    if not isinstance(r['OrderID'], str) or not r['OrderID'].startswith('00'): fail(f'Workbook OrderID not text {r}')
    if r['Status']=='Cancelled' or not between(str(r['OrderDate'])): fail(f'Workbook order invalid filter {r}')
    if not isinstance(r['NetAmount'], (int,float)) or not isinstance(r['AmountTWD'], (int,float)): fail(f'Workbook order amounts not numeric {r}')
for r in w_inv:
    if not isinstance(r['InvoiceID'], str) or not r['InvoiceID'].startswith('00'): fail(f'Workbook InvoiceID not text {r}')
    if not between(str(r['InvoiceDate'])): fail(f'Workbook invoice invalid date {r}')
    for k in ['Amount','PaidAmount','Balance','BalanceTWD']:
        if not isinstance(r[k], (int,float)): fail(f'Workbook invoice {k} not numeric {r}')
for r in w_exp:
    if not isinstance(r['ExpenseID'], str): fail(f'Workbook ExpenseID not text {r}')
    if r['Status']!='Approved' or not between(str(r['ExpenseDate'])): fail(f'Workbook expense invalid filter {r}')
    for k in ['Amount','AmountTWD']:
        if not isinstance(r[k], (int,float)): fail(f'Workbook expense {k} not numeric {r}')
for r in w_quotes:
    if not isinstance(r['QuoteID'], str): fail(f'Workbook QuoteID not text {r}')
    if r['Status']!='Accepted' or not between(str(r['QuoteDate'])): fail(f'Workbook quote invalid filter {r}')
    for k in ['NetAmount','AmountTWD']:
        if not isinstance(r[k], (int,float)): fail(f'Workbook quote {k} not numeric {r}')
ok('workbook filters, text IDs, and numeric amounts checked')

# Summary totals match sheet details by ProjectID, excluding blank project ids.
projects=sorted({r['ProjectID'] for rows in [w_orders,w_inv,w_exp,w_quotes] for r in rows if r.get('ProjectID')})
summary={r['ProjectID']:r for r in w_sum if r.get('ProjectID')}
if sorted(summary)!=projects: fail(f'Summary project IDs {sorted(summary)} != detail projects {projects}')
else: ok('Summary project IDs match detail sheets')

def sum_by(rows, field, project):
    return sum(float(r[field] or 0) for r in rows if r.get('ProjectID')==project)
for p in projects:
    sr=summary[p]
    comps={
        'OrderTWD': sum_by(w_orders,'AmountTWD',p),
        'InvoiceBalanceTWD': sum_by(w_inv,'BalanceTWD',p),
        'ApprovedExpenseTWD': sum_by(w_exp,'AmountTWD',p),
        'AcceptedQuoteTWD': sum_by(w_quotes,'AmountTWD',p),
    }
    for k,v in comps.items():
        if not money_equal(sr[k], v): fail(f'Summary {p} {k}={sr[k]} != detail sum {v}')
ok('Summary amounts match detail sheet sums')

print('CHECKS PASSED:')
for c in checks:
    print('-', c)
if errors:
    print('\nERRORS:')
    for e in errors:
        print('-', e)
    raise SystemExit(1)
print('\nAll verification checks passed.')
print('Final files:')
for rel in required:
    p=base/rel
    print(f'- {rel} ({p.stat().st_size} bytes)')
