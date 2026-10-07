from pathlib import Path
import csv, json, re
from decimal import Decimal
from openpyxl import load_workbook

base = Path.cwd()
FROM='2026-09-01'
TO='2026-09-30'
errors=[]
checks=[]

def fail(msg): errors.append(msg)
def ok(msg): checks.append(msg)
def between(s): return FROM <= str(s) <= TO

def read_csv(rel):
    p=base/rel
    if not p.exists():
        fail(f'missing {rel}')
        return [], []
    with p.open(newline='', encoding='utf-8-sig') as f:
        rdr=csv.DictReader(f)
        rows=list(rdr)
    return rdr.fieldnames or [], rows

def sheet_rows(wb, sheet):
    ws=wb[sheet]
    headers=[c.value for c in ws[4]]
    rows=[]
    for row in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        if all(v is None for v in row):
            continue
        rows.append(dict(zip(headers,row)))
    return headers, rows

def money_equal(a,b,tol=0.01): return abs(float(a)-float(b)) <= tol

required = [
 'output/downloads/sap_orders_export.csv',
 f'output/downloads/sap_orders_{FROM}_{TO}.csv',
 'output/downloads/sap_invoices_export.csv',
 f'output/downloads/sap_invoices_{FROM}_{TO}.csv',
 'output/expense_approved_export.csv',
 f'output/expense_{FROM}_{TO}_approved.csv',
 f'output/quotes_{FROM}_{TO}_accepted.json',
 f'output/project_report_{FROM}_{TO}.xlsx',
]
for rel in required:
    p=base/rel
    if p.exists() and p.stat().st_size>0: ok(f'{rel} exists size={p.stat().st_size}')
    else: fail(f'{rel} missing or empty')

orders_h, orders = read_csv('output/downloads/sap_orders_export.csv')
assert orders_h == ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount'], orders_h
ok('orders CSV headers match source names')
assert len(orders)==18, len(orders)
ok('orders CSV row count 18')
for r in orders:
    assert between(r['OrderDate']), r
    assert r['Status']!='Cancelled', r
    assert isinstance(r['OrderID'], str) and re.fullmatch(r'\d+', r['OrderID']) and r['OrderID'].startswith('00'), r
    Decimal(r['NetAmount'])
ok('orders are within date range, exclude Cancelled, preserve text IDs and numeric amounts')
_, orders_dated = read_csv(f'output/downloads/sap_orders_{FROM}_{TO}.csv')
assert orders_dated == orders
ok('dated orders CSV matches export CSV')

inv_h, inv = read_csv('output/downloads/sap_invoices_export.csv')
assert inv_h == ['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance'], inv_h
ok('invoice CSV headers match source names')
assert len(inv)==13, len(inv)
ok('invoice CSV row count 13')
for r in inv:
    assert between(r['InvoiceDate']), r
    assert isinstance(r['InvoiceID'], str) and r['InvoiceID'].startswith('00'), r
    for k in ['Amount','PaidAmount','Balance']: Decimal(r[k])
ok('invoices are within date range, preserve text IDs and numeric amounts')
_, inv_dated = read_csv(f'output/downloads/sap_invoices_{FROM}_{TO}.csv')
assert inv_dated == inv
ok('dated invoices CSV matches export CSV')

exp_h, exp = read_csv('output/expense_approved_export.csv')
assert exp_h == ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt'], exp_h
ok('expense CSV headers match source names')
assert len(exp)==15, len(exp)
ok('expense CSV row count 15')
for r in exp:
    assert between(r['ExpenseDate']), r
    assert r['Status']=='Approved', r
    assert isinstance(r['ExpenseID'], str), r
    Decimal(r['Amount'])
ok('expenses are within date range, Approved only, preserve text IDs and numeric amounts')
_, exp_dated = read_csv(f'output/expense_{FROM}_{TO}_approved.csv')
assert exp_dated == exp
ok('dated expense CSV matches export CSV')

qobj=json.loads((base/f'output/quotes_{FROM}_{TO}_accepted.json').read_text(encoding='utf-8'))
assert isinstance(qobj, dict), type(qobj)
assert qobj.get('schema_version') == 1, qobj
quotes=qobj.get('items')
assert isinstance(quotes, list), qobj.keys()
assert qobj.get('count') == len(quotes) == 4, (qobj.get('count'), len(quotes))
ok('quotes JSON schema_version/items/count valid; accepted count 4')
for r in quotes:
    assert between(r.get('QuoteDate','')), r
    assert r.get('Status')=='Accepted', r
    assert isinstance(r.get('QuoteID'), str) and re.fullmatch(r'Q\d+', r.get('QuoteID','')), r
    Decimal(str(r.get('NetAmount')))
ok('quotes are within date range, Accepted only, preserve text IDs and numeric amounts')

wb=load_workbook(base/f'output/project_report_{FROM}_{TO}.xlsx', data_only=False)
assert wb.sheetnames == ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions'], wb.sheetnames
ok('workbook sheets preserved')
assert wb['Report']['B3'].value == FROM and wb['Report']['D3'].value == TO, (wb['Report']['B3'].value, wb['Report']['D3'].value)
ok('Report sheet dates match run period')

_, w_orders=sheet_rows(wb,'Orders')
_, w_inv=sheet_rows(wb,'Invoices')
_, w_exp=sheet_rows(wb,'Expenses')
_, w_quotes=sheet_rows(wb,'Quotes')
_, w_sum=sheet_rows(wb,'Summary')
assert len(w_orders)==len(orders), (len(w_orders), len(orders))
ok('Orders sheet row count matches CSV')
assert len(w_inv)==len(inv), (len(w_inv), len(inv))
ok('Invoices sheet row count matches CSV')
assert len(w_exp)==len(exp), (len(w_exp), len(exp))
ok('Expenses sheet row count matches CSV')
assert len(w_quotes)==len(quotes), (len(w_quotes), len(quotes))
ok('Quotes sheet row count matches JSON')

for r in w_orders:
    assert isinstance(r['OrderID'], str) and r['OrderID'].startswith('00'), r
    assert r['Status']!='Cancelled' and between(r['OrderDate']), r
    assert isinstance(r['NetAmount'], (int,float)) and isinstance(r['AmountTWD'], (int,float)), r
for r in w_inv:
    assert isinstance(r['InvoiceID'], str) and r['InvoiceID'].startswith('00'), r
    assert between(r['InvoiceDate']), r
    for k in ['Amount','PaidAmount','Balance','BalanceTWD']: assert isinstance(r[k], (int,float)), (k,r)
for r in w_exp:
    assert isinstance(r['ExpenseID'], str), r
    assert r['Status']=='Approved' and between(r['ExpenseDate']), r
    for k in ['Amount','AmountTWD']: assert isinstance(r[k], (int,float)), (k,r)
for r in w_quotes:
    assert isinstance(r['QuoteID'], str) and re.fullmatch(r'Q\d+', r['QuoteID']), r
    assert r['Status']=='Accepted' and between(r['QuoteDate']), r
    for k in ['NetAmount','AmountTWD']: assert isinstance(r[k], (int,float)), (k,r)
ok('workbook filters, text IDs, and numeric amounts checked')

projects=sorted({r['ProjectID'] for rows in [w_orders,w_inv,w_exp,w_quotes] for r in rows if r.get('ProjectID')})
summary={r['ProjectID']:r for r in w_sum if r.get('ProjectID')}
assert sorted(summary)==projects, (sorted(summary), projects)
ok('Summary project IDs match detail sheets')

def sum_by(rows, field, project):
    return sum(float(r[field] or 0) for r in rows if r.get('ProjectID')==project)
for p in projects:
    sr=summary[p]
    comps={
        'OrderTWD': sum_by(w_orders,'AmountTWD',p),
        'InvoiceBalanceTWD': sum_by(w_inv,'BalanceTWD',p),
        'ApprovedExpenseTWD': sum_by(w_exp,'AmountTWD',p),
        'AcceptedQuoteTWD': sum_by(w_quotes,'AmountTWD',p),
    }
    for k,v in comps.items():
        assert money_equal(sr[k], v), (p,k,sr[k],v)
ok('Summary amounts match detail sheet sums')

print('CHECKS PASSED:')
for c in checks:
    print('-', c)
print('\nAll verification checks passed.')
print('\nFinal files:')
for rel in required:
    p=base/rel
    print(f'- {rel} ({p.stat().st_size} bytes)')
