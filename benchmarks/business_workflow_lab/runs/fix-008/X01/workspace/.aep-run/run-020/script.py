from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('Workspace:', ws)
print('Output files:')
for p in sorted((ws/'output').rglob('*')):
    if p.is_file():
        print('-', p.relative_to(ws), 'size=', p.stat().st_size)

# Print headers/sample for August dated outputs and key current exports
files = [
    ws/'output'/'downloads'/'sap_orders_2026-08-01_2026-08-31.csv',
    ws/'output'/'downloads'/'sap_invoices_2026-08-01_2026-08-31.csv',
    ws/'output'/'expense_2026-08-01_2026-08-31_approved.csv',
    ws/'output'/'quotes_2026-08-01_2026-08-31_accepted.json',
    ws/'output'/'project_report_2026-08-01_2026-08-31.xlsx',
]
for f in files:
    print('\nFILE', f.relative_to(ws), 'exists=', f.exists(), 'size=', f.stat().st_size if f.exists() else None)
    if not f.exists():
        continue
    if f.suffix.lower() == '.csv':
        with f.open('r', encoding='utf-8-sig', newline='') as fh:
            r = csv.DictReader(fh)
            print('headers:', r.fieldnames)
            for i, row in zip(range(3), r):
                print('sample', i+1, row)
    elif f.suffix.lower() == '.json':
        obj = json.loads(f.read_text(encoding='utf-8'))
        print('json keys:', list(obj.keys()))
        print('count:', obj.get('count'), 'items_len:', len(obj.get('items', [])))
        for i, item in enumerate(obj.get('items', [])[:3], 1):
            print('sample', i, item)
    elif f.suffix.lower() == '.xlsx':
        wb = load_workbook(f, data_only=False, read_only=True)
        print('sheets:', wb.sheetnames)
        for s in wb.sheetnames:
            sh = wb[s]
            print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
            if sh.max_row:
                vals = [sh.cell(1, c).value for c in range(1, min(sh.max_column, 10)+1)]
                print(' row1:', vals)
