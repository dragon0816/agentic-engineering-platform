from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
path = ws/'output'/'project_report_2026-08-01_2026-08-31.xlsx'
print('Inspecting workbook:', path.relative_to(ws), 'exists=', path.exists(), 'size=', path.stat().st_size if path.exists() else None)
wb = load_workbook(path, data_only=True, read_only=True)
for shname in ['Orders','Invoices','Expenses','Quotes','Summary','Report','Exceptions']:
    sh = wb[shname]
    print('\nSHEET', shname, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r_idx, row in enumerate(sh.iter_rows(min_row=1, max_row=min(sh.max_row, 15), values_only=True), 1):
        print(r_idx, row)
