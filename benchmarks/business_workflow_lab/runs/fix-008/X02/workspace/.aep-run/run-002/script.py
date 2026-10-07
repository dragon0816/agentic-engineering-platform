from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/project_template.xlsx', data_only=False)
s=wb['Instructions']
for row in s.iter_rows(values_only=True): print(row)