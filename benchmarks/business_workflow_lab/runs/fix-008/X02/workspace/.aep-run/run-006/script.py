from pathlib import Path
from openpyxl import load_workbook
import csv, json, os

ws = Path.cwd()
print('Workspace:', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.relative_to(ws), p.stat().st_size)
print('Output files:')
for p in sorted((ws/'output').rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(ws), p.stat().st_size)

report = ws/'output'/'project_report_C001_2026-09_TWD.xlsx'
if not report.exists():
    raise SystemExit('Missing expected report: ' + str(report))
wb = load_workbook(report, data_only=False)
print('Workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'Sheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    headers = [sh.cell(1,c).value for c in range(1, sh.max_column+1)]
    print('  headers:', headers)
    sample = []
    for r in range(2, min(sh.max_row, 4)+1):
        sample.append([sh.cell(r,c).value for c in range(1, sh.max_column+1)])
    print('  sample:', sample)

# Requirement checks
expected = {'Orders': 3, 'Invoices': 3, 'Expenses': 9, 'Quotes': 4, 'Summary': 2, 'Exceptions': 4}
actual = {}
for sheet, count in expected.items():
    sh = wb[sheet]
    actual[sheet] = max(0, sh.max_row - 1)
print('Row counts:', actual)
assert actual == expected, f'Row count mismatch: {actual} != {expected}'

# Check source filter semantics in saved workbook
orders = wb['Orders']
order_headers = [orders.cell(1,c).value for c in range(1, orders.max_column+1)]
status_col = order_headers.index('Status') + 1
assert all(orders.cell(r,status_col).value != 'Cancelled' for r in range(2, orders.max_row+1)), 'Cancelled order found'

expenses = wb['Expenses']
exp_headers = [expenses.cell(1,c).value for c in range(1, expenses.max_column+1)]
exp_status_col = exp_headers.index('Status') + 1
exp_proj_col = exp_headers.index('ProjectID') + 1
assert all(expenses.cell(r,exp_status_col).value == 'Approved' for r in range(2, expenses.max_row+1)), 'Non-Approved expense found'
assert set(expenses.cell(r,exp_proj_col).value for r in range(2, expenses.max_row+1)) <= {'P001','P003'}, 'Expense ProjectID outside P001/P003 found'

# Check exceptions content
exceptions = wb['Exceptions']
exc_headers = [exceptions.cell(1,c).value for c in range(1, exceptions.max_column+1)]
exc_rows = [[exceptions.cell(r,c).value for c in range(1, exceptions.max_column+1)] for r in range(2, exceptions.max_row+1)]
print('Exceptions rows:', json.dumps(exc_rows, ensure_ascii=False, default=str))
assert len(exc_rows) == 4, 'Expected 4 exceptions'
assert any(row[0] == 'Invoices' and row[1] == '0090000219' for row in exc_rows), 'Missing overdue unpaid invoice exception'
assert any(row[0] == 'Expenses' and row[1] == 'EXP-00006' for row in exc_rows), 'Missing missing receipt expense exception'
assert any(row[0] == 'Quotes' and row[1] == 'Q0010' for row in exc_rows), 'Missing Q0010 expired quote exception'
assert any(row[0] == 'Quotes' and row[1] == 'Q0014' for row in exc_rows), 'Missing Q0014 expired quote exception'

print('CHECK PASSED: output workbook exists, sheets/row counts/filters/exceptions match the stated requirements.')