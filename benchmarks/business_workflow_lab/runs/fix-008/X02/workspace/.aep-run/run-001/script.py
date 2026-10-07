from pathlib import Path
from openpyxl import load_workbook
import csv, json, os
ws = Path.cwd()
print('cwd', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nmanual previews:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=ws/'input'/name
    print(f'--- {name} ---')
    print(p.read_text(encoding='utf-8')[:4000])
print('\nfx_rates.csv:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r):
        print(row)
print('\nworkbook sheets and dimensions:')
wb = load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    s=wb[sh]
    print(sh, s.max_row, s.max_column)
    for row in s.iter_rows(min_row=1, max_row=min(s.max_row,8), values_only=True):
        print(row)
    print()