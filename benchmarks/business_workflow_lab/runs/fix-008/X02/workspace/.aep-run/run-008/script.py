from pathlib import Path
from openpyxl import load_workbook
import csv, json
from datetime import datetime, date

ws = Path.cwd()
print('Workspace:', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(f' - {p.relative_to(ws)} ({p.stat().st_size} bytes)')
print('Output files:')
for p in sorted((ws/'output').rglob('*')):
    if p.is_file():
        print(f' - {p.relative_to(ws)} ({p.stat().st_size} bytes)')

# Reopen and inspect CSV outputs.
expected_csvs = [
    ws/'output'/'downloads'/'orders_C001_2026-09.csv',
    ws/'output'/'downloads'/'invoices_C001_2026-09.csv',
    ws/'output'/'expenses_P001_Approved_2026-09.csv',
    ws/'output'/'expenses_P003_Approved_2026-09.csv',
]
for p in expected_csvs:
    assert p.exists(), f'Missing CSV: {p}'
    with p.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        rows = list(reader)
        print(f'{p.relative_to(ws)} headers={reader.fieldnames} rows={len(rows)} sample={rows[:2]}')

# Reopen workbook and validate template-filled content.
report = ws/'output'/'project_report_C001_2026-09_TWD.xlsx'
assert report.exists(), 'Missing workbook output'
wb = load_workbook(report, data_only=False)
print('Workbook sheets:', wb.sheetnames)
assert wb.sheetnames == ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']

# Template rule: row 4 is header, row 5+ is data on output sheets.
def sheet_headers(name):
    sh = wb[name]
    return [sh.cell(4,c).value for c in range(1, sh.max_column+1)]

def sheet_rows(name):
    sh = wb[name]
    out = []
    for r in range(5, sh.max_row+1):
        vals = [sh.cell(r,c).value for c in range(1, sh.max_column+1)]
        if any(v is not None for v in vals):
            out.append(vals)
    return out

def d(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.fromisoformat(str(v)[:10]).date()

start = date(2026, 9, 1)
end = date(2026, 9, 30)
expected_counts = {'Orders':3, 'Invoices':3, 'Expenses':9, 'Quotes':4, 'Summary':2, 'Exceptions':4}
actual_counts = {s: len(sheet_rows(s)) for s in expected_counts}
print('Workbook data row counts:', actual_counts)
assert actual_counts == expected_counts, f'Unexpected workbook row counts: {actual_counts}'

# Report metadata and source FX note.
report_sh = wb['Report']
print('Report B3/D3/B4:', report_sh['B3'].value, report_sh['D3'].value, report_sh['B4'].value)
assert d(report_sh['B3'].value) == start
assert d(report_sh['D3'].value) == end
condition = str(report_sh['B4'].value)
assert 'Customer C001' in condition
assert 'orders exclude Cancelled' in condition
assert 'expenses Approved' in condition
assert 'quotes all statuses' in condition
assert 'FX from input/fx_rates.csv' in condition

# Orders.
h = sheet_headers('Orders'); rows = sheet_rows('Orders'); idx = {name:i for i,name in enumerate(h) if name}
print('Orders rows:', rows)
assert h == ['OrderID', 'ProjectID', 'CustomerID', 'OrderDate', 'Status', 'Currency', 'NetAmount', 'AmountTWD']
assert all(isinstance(r[idx['OrderID']], str) for r in rows)
assert all(r[idx['CustomerID']] == 'C001' for r in rows)
assert all(r[idx['Status']] != 'Cancelled' for r in rows)
assert all(start <= d(r[idx['OrderDate']]) <= end for r in rows)
assert all(isinstance(r[idx['NetAmount']], (int, float)) and isinstance(r[idx['AmountTWD']], (int, float)) for r in rows)

# Invoices.
h = sheet_headers('Invoices'); rows = sheet_rows('Invoices'); idx = {name:i for i,name in enumerate(h) if name}
print('Invoices rows:', rows)
assert h == ['InvoiceID', 'OrderID', 'ProjectID', 'CustomerID', 'InvoiceDate', 'DueDate', 'Currency', 'Amount', 'PaidAmount', 'Balance', 'BalanceTWD']
assert all(isinstance(r[idx['InvoiceID']], str) for r in rows)
assert all(r[idx['CustomerID']] == 'C001' for r in rows)
assert all(start <= d(r[idx['InvoiceDate']]) <= end for r in rows)
assert all(isinstance(r[idx['Amount']], (int, float)) and isinstance(r[idx['PaidAmount']], (int, float)) and isinstance(r[idx['BalanceTWD']], (int, float)) for r in rows)

# Expenses.
h = sheet_headers('Expenses'); rows = sheet_rows('Expenses'); idx = {name:i for i,name in enumerate(h) if name}
print('Expenses rows:', rows)
assert h == ['ExpenseID', 'ProjectID', 'EmployeeID', 'ExpenseDate', 'Status', 'Currency', 'Amount', 'AmountTWD']
assert all(isinstance(r[idx['ExpenseID']], str) for r in rows)
assert set(r[idx['ProjectID']] for r in rows) <= {'P001', 'P003'}
assert all(r[idx['Status']] == 'Approved' for r in rows)
assert all(start <= d(r[idx['ExpenseDate']]) <= end for r in rows)
assert all(isinstance(r[idx['Amount']], (int, float)) and isinstance(r[idx['AmountTWD']], (int, float)) for r in rows)

# Quotes.
h = sheet_headers('Quotes'); rows = sheet_rows('Quotes'); idx = {name:i for i,name in enumerate(h) if name}
print('Quotes rows:', rows)
assert h == ['QuoteID', 'ProjectID', 'CustomerID', 'QuoteDate', 'ValidUntil', 'Status', 'Currency', 'NetAmount', 'AmountTWD']
assert all(isinstance(r[idx['QuoteID']], str) for r in rows)
assert all(r[idx['CustomerID']] == 'C001' for r in rows)
assert all(start <= d(r[idx['QuoteDate']]) <= end for r in rows)
assert {r[idx['Status']] for r in rows} == {'Sent','Accepted','Draft','Expired'}
assert all(isinstance(r[idx['NetAmount']], (int, float)) and isinstance(r[idx['AmountTWD']], (int, float)) for r in rows)

# Summary: two projects and formulas extended into data rows.
h = sheet_headers('Summary'); rows = sheet_rows('Summary'); idx = {name:i for i,name in enumerate(h) if name}
print('Summary rows:', rows)
assert h == ['ProjectID', 'OrderTWD', 'InvoiceBalanceTWD', 'ApprovedExpenseTWD', 'AcceptedQuoteTWD']
assert {r[idx['ProjectID']] for r in rows} == {'P001', 'P003'}
for rownum in range(5, 5+len(rows)):
    assert wb['Summary'].cell(rownum,2).data_type == 'f'
    assert wb['Summary'].cell(rownum,3).data_type == 'f'
    assert wb['Summary'].cell(rownum,4).data_type == 'f'
    assert wb['Summary'].cell(rownum,5).data_type == 'f'

# Exceptions: exactly the requested attention conditions.
h = sheet_headers('Exceptions'); rows = sheet_rows('Exceptions'); idx = {name:i for i,name in enumerate(h) if name}
print('Exceptions rows:', json.dumps(rows, ensure_ascii=False, default=str))
assert h == ['Source', 'DocumentID', 'ProjectID', 'Reason']
actual = {(r[idx['Source']], r[idx['DocumentID']], r[idx['ProjectID']]) for r in rows}
expected = {
    ('Invoices', '0090000219', 'P001'),
    ('Expenses', 'EXP-00006', 'P001'),
    ('Quotes', 'Q0010', 'P001'),
    ('Quotes', 'Q0014', 'P001'),
}
assert actual == expected, actual

print('CHECK PASSED: all output CSVs and workbook were reopened; row counts, filters, date ranges, source FX note, text IDs, numeric TWD values, summary formulas, and attention exceptions match the requirements.')