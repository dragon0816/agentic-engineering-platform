import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-008\D02\workspace')
print('workspace exists', ws.exists())
for sub in ['input','output']:
    p=ws/sub
    print(sub, 'exists', p.exists())
    if p.exists():
        for f in sorted(p.iterdir()):
            print(' -', f.name, f.stat().st_size)
# inspect workbook
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path)
print('workbook sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('\nSHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
        print(r)
# inspect csv
csv_path=ws/'input'/'fx_rates.csv'
if csv_path.exists():
    print('\nfx_rates.csv sample')
    with open(csv_path, newline='', encoding='utf-8-sig') as f:
        rdr=csv.reader(f)
        for i,row in zip(range(10), rdr):
            print(row)
# manual
manual=ws/'input'/'expense_manual.md'
if manual.exists():
    print('\nmanual first 200 lines')
    with open(manual, encoding='utf-8') as f:
        for i,line in zip(range(200), f):
            print(f'{i+1}: {line.rstrip()}')