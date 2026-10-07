from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-008\D02\workspace')
for fname in ['expense_submitted_2026-09.csv','expense_approved_2026-09.csv']:
    p=ws/'output'/fname
    print('\nFILE', fname, 'exists', p.exists(), 'size', p.stat().st_size if p.exists() else None)
    if p.exists():
        with open(p, newline='', encoding='utf-8-sig') as f:
            rdr=csv.DictReader(f)
            print('headers:', rdr.fieldnames)
            rows=list(rdr)
            print('rows:', len(rows))
            for r in rows[:5]: print(r)
wb=load_workbook(ws/'input'/'expense_template.xlsx')
inst=wb['Instructions']
print('\nInstructions all rows:')
for row in inst.iter_rows(values_only=True):
    print(row)