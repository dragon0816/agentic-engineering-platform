from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace=', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
print('output exists=', (ws/'output').exists())

# Inspect CSV headers/samples
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv sample:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i, row in zip(range(6), r):
            print(i, row)

manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\nquote_cli_manual.md head:')
    text = manual.read_text(encoding='utf-8', errors='replace')
    print(text[:3000])

qt = ws/'input'/'quote_template.xlsx'
if qt.exists():
    print('\nquote_template.xlsx sheets/cells:')
    wb = load_workbook(qt, data_only=False)
    for sh in wb.sheetnames:
        wsht = wb[sh]
        print('Sheet', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
        for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 10), values_only=True):
            print(row)
        print('---')
