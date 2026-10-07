from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from copy import copy
import subprocess, json, csv, shutil, sys
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
OUT_XLSX = OUTPUT / 'quote_template.xlsx'
OUT_MD = OUTPUT / 'findings.md'
CLI = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = date(2026, 9, 30)
GENERATED_AT = datetime.now().astimezone().isoformat(timespec='seconds')
TWOPLACES = Decimal('0.01')

def money(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [CLI] + args
    res = subprocess.run(cmd, capture_output=True, text=True, shell=False)
    if res.returncode != 0:
        raise RuntimeError(f'quote_cli failed ({res.returncode}) for {args}: {res.stderr}')
    return json.loads(res.stdout)

def parse_date(s):
    return datetime.strptime(s, '%Y-%m-%d').date()

def load_fx():
    rates = {}
    with (INPUT/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        header = next(reader)
        for row in reader:
            if len(row) >= 2 and row[0]:
                rates[row[0]] = Decimal(row[1])
    return rates

def copy_row_format(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        s = ws.cell(src_row, col)
        d = ws.cell(dst_row, col)
        if s.has_style:
            d._style = copy(s._style)
        if s.number_format:
            d.number_format = s.number_format
        if s.font:
            d.font = copy(s.font)
        if s.fill:
            d.fill = copy(s.fill)
        if s.border:
            d.border = copy(s.border)
        if s.alignment:
            d.alignment = copy(s.alignment)
        if s.protection:
            d.protection = copy(s.protection)

# Fetch all September quote headers, then full quote documents.
listed = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json'])
if listed.get('schema_version') != 1:
    raise RuntimeError(f'Unexpected list schema_version: {listed.get("schema_version")}')
headers = listed.get('items', [])
quote_ids = [h['QuoteID'] for h in headers]
full_quotes = []
for qid in quote_ids:
    doc = run_cli(['quotes', 'get', qid, '--format', 'json'])
    if doc.get('schema_version') != 1:
        raise RuntimeError(f'Unexpected get schema_version for {qid}: {doc.get("schema_version")}')
    full_quotes.append(doc)

# Load FX and validate data.
fx = load_fx()
calc_findings = []
expired_unaccepted = []
line_rows = []
quote_rows = []
summary = {}

for doc in full_quotes:
    h = doc['header']
    qid = h['QuoteID']
    currency = h['Currency']
    if currency not in fx:
        raise RuntimeError(f'Missing FX rate for {currency}')
    subtotal = money(h['Subtotal'])
    discount_percent = Decimal(str(h['DiscountPercent']))
    net = money(h['NetAmount'])
    line_sum = Decimal('0.00')
    line_issues = []
    for li in doc.get('line_items', []):
        qty = Decimal(str(li['Quantity']))
        unit = money(li['UnitPrice'])
        line_amount = money(li['LineAmount'])
        expected_line = money(qty * unit)
        if line_amount != expected_line:
            line_issues.append(f"line {li.get('ProductID','')} expected {expected_line} got {line_amount}")
        line_sum += line_amount
        line_rows.append({
            'QuoteID': qid,
            'ProductID': li.get('ProductID', ''),
            'ProductName': li.get('ProductName', ''),
            'Quantity': int(li['Quantity']) if Decimal(str(li['Quantity'])) == Decimal(str(int(li['Quantity']))) else float(li['Quantity']),
            'UnitPrice': unit,
            'LineAmount': line_amount,
        })
    line_sum = money(line_sum)
    expected_net = money(subtotal * (Decimal('1') - (discount_percent / Decimal('100'))))
    issues = []
    if line_issues:
        issues.extend(line_issues)
    if line_sum != subtotal:
        issues.append(f'line detail sum {line_sum} != Subtotal {subtotal}')
    if expected_net != net:
        issues.append(f'NetAmount expected {expected_net} from Subtotal {subtotal} and DiscountPercent {discount_percent} != reported {net}')
    calc_findings.append({
        'QuoteID': qid,
        'Subtotal': subtotal,
        'LineSum': line_sum,
        'DiscountPercent': discount_percent,
        'ExpectedNetAmount': expected_net,
        'NetAmount': net,
        'OK': not issues,
        'Issues': issues,
    })
    vu = parse_date(h['ValidUntil'])
    if vu < BASE_DATE and h['Status'] != 'Accepted':
        expired_unaccepted.append({
            'QuoteID': qid,
            'CustomerID': h.get('CustomerID',''),
            'ProjectID': h.get('ProjectID',''),
            'QuoteDate': h['QuoteDate'],
            'ValidUntil': h['ValidUntil'],
            'Status': h['Status'],
            'Currency': currency,
            'NetAmount': net,
        })
    net_twd = money(net * fx[currency])
    quote_rows.append({
        'QuoteID': qid,
        'CustomerID': h.get('CustomerID',''),
        'ProjectID': h.get('ProjectID',''),
        'QuoteDate': parse_date(h['QuoteDate']),
        'ValidUntil': vu,
        'Status': h['Status'],
        'Currency': currency,
        'Subtotal': subtotal,
        'DiscountPercent': discount_percent,
        'NetAmount': net,
        'NetAmountTWD': net_twd,
    })
    key = (h['Status'], currency)
    if key not in summary:
        summary[key] = {'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0}
    summary[key]['Amount'] += net
    summary[key]['AmountTWD'] += net_twd
    summary[key]['Count'] += 1

# Fill a copy of the template.
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

# Report sheet.
rep = wb['Report']
rep['B3'] = FROM_DATE
rep['D3'] = TO_DATE
rep['B4'] = 'QuoteDate 2026-09-01 至 2026-09-30（含起訖）；全部報價；未依到期狀態刪除'
rep['B5'] = GENERATED_AT
rep['B7'] = 'quote_cli mock; fx_rates.csv'
rep['D7'] = GENERATED_AT
rep['B9'] = f'以 {BASE_DATE.isoformat()} 為基準檢查 ValidUntil 小於基準日且 Status 非 Accepted；共 {len(expired_unaccepted)} 份。明細合計/折扣後 NetAmount 不一致共 {sum(1 for x in calc_findings if not x["OK"])} 份。'

# Helpers for writing sheets.
def clear_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def set_text(cell, value):
    cell.value = value if value != '' else None
    cell.number_format = '@'

def set_date(cell, value):
    cell.value = value
    cell.number_format = 'yyyy-mm-dd'

def set_money_cell(cell, value):
    cell.value = float(money(value))
    cell.number_format = '0.00'

quotes_ws = wb['Quotes']
clear_data(quotes_ws)
for idx, row in enumerate(quote_rows, start=5):
    copy_row_format(quotes_ws, 4, idx, 11)
    set_text(quotes_ws.cell(idx,1), row['QuoteID'])
    set_text(quotes_ws.cell(idx,2), row['CustomerID'])
    set_text(quotes_ws.cell(idx,3), row['ProjectID'])
    set_date(quotes_ws.cell(idx,4), row['QuoteDate'])
    set_date(quotes_ws.cell(idx,5), row['ValidUntil'])
    set_text(quotes_ws.cell(idx,6), row['Status'])
    set_text(quotes_ws.cell(idx,7), row['Currency'])
    set_money_cell(quotes_ws.cell(idx,8), row['Subtotal'])
    quotes_ws.cell(idx,9).value = float(row['DiscountPercent'])
    quotes_ws.cell(idx,9).number_format = '0.##'
    set_money_cell(quotes_ws.cell(idx,10), row['NetAmount'])
    set_money_cell(quotes_ws.cell(idx,11), row['NetAmountTWD'])

lines_ws = wb['Lines']
clear_data(lines_ws)
for idx, row in enumerate(line_rows, start=5):
    copy_row_format(lines_ws, 4, idx, 6)
    set_text(lines_ws.cell(idx,1), row['QuoteID'])
    set_text(lines_ws.cell(idx,2), row['ProductID'])
    lines_ws.cell(idx,3).value = row['ProductName']
    lines_ws.cell(idx,4).value = row['Quantity']
    set_money_cell(lines_ws.cell(idx,5), row['UnitPrice'])
    set_money_cell(lines_ws.cell(idx,6), row['LineAmount'])

summ_ws = wb['Summary']
clear_data(summ_ws)
last_quote_row = 4 + len(quote_rows)
if len(quote_rows) == 0:
    summ_ws['B2'] = '=SUM(Quotes!K5:K5)'
else:
    summ_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
srow = 5
for (status, currency), vals in sorted(summary.items()):
    copy_row_format(summ_ws, 4, srow, 6)
    summ_ws.cell(srow,1).value = 'Status'
    summ_ws.cell(srow,2).value = status
    summ_ws.cell(srow,3).value = currency
    set_money_cell(summ_ws.cell(srow,4), vals['Amount'])
    set_money_cell(summ_ws.cell(srow,5), vals['AmountTWD'])
    summ_ws.cell(srow,6).value = vals['Count']
    srow += 1

# Autofilter ranges if present can be expanded conservatively.
quotes_ws.auto_filter.ref = f'A4:K{max(5,last_quote_row)}'
lines_ws.auto_filter.ref = f'A4:F{max(5,4+len(line_rows))}'
summ_ws.auto_filter.ref = f'A4:F{max(5,4+len(summary))}'

wb.save(OUT_XLSX)

# Write findings markdown.
md = []
md.append('# Quote findings')
md.append('')
md.append(f'- Query period: {FROM_DATE} to {TO_DATE} inclusive')
md.append(f'- Baseline date for ValidUntil: {BASE_DATE.isoformat()}')
md.append(f'- Source: quote_cli mock; fx_rates.csv')
md.append(f'- Quotes retained in workbook: {len(quote_rows)} (no quote removed because of expiration)')
md.append(f'- Line rows retained in workbook: {len(line_rows)}')
md.append('')
md.append('## Calculation checks')
md.append('')
md.append('| QuoteID | Detail sum | Subtotal | DiscountPercent | Expected NetAmount | Reported NetAmount | Result |')
md.append('|---|---:|---:|---:|---:|---:|---|')
for f in calc_findings:
    result = 'OK' if f['OK'] else '; '.join(f['Issues'])
    md.append(f"| {f['QuoteID']} | {f['LineSum']} | {f['Subtotal']} | {f['DiscountPercent']} | {f['ExpectedNetAmount']} | {f['NetAmount']} | {result} |")
md.append('')
md.append('## ValidUntil < 2026-09-30 and Status not Accepted')
md.append('')
if expired_unaccepted:
    md.append('| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |')
    md.append('|---|---|---|---|---|---|---|---:|')
    for e in expired_unaccepted:
        md.append(f"| {e['QuoteID']} | {e['CustomerID']} | {e['ProjectID']} | {e['QuoteDate']} | {e['ValidUntil']} | {e['Status']} | {e['Currency']} | {e['NetAmount']} |")
else:
    md.append('None.')
md.append('')
md.append('## FX rates used')
md.append('')
md.append('| Currency | RateToTWD |')
md.append('|---|---:|')
for c in sorted(fx):
    md.append(f'| {c} | {fx[c]} |')
OUT_MD.write_text('\n'.join(md) + '\n', encoding='utf-8')

# Reopen and validate saved outputs against requirements.
wb2 = load_workbook(OUT_XLSX, data_only=False)
qws = wb2['Quotes']; lws = wb2['Lines']; sws = wb2['Summary']; rws = wb2['Report']
# Count non-empty data rows by QuoteID and Lines QuoteID.
saved_qids = [qws.cell(r,1).value for r in range(5, qws.max_row+1) if qws.cell(r,1).value]
saved_line_qids = [lws.cell(r,1).value for r in range(5, lws.max_row+1) if lws.cell(r,1).value]
assert saved_qids == quote_ids, (saved_qids, quote_ids)
assert len(saved_qids) == listed['count'] == 12
assert len(saved_line_qids) == len(line_rows)
assert rws['B3'].value == FROM_DATE and rws['D3'].value == TO_DATE
assert sws['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})'
# Check columns and types/formats.
assert [qws.cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [lws.cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for r in range(5, 5+len(quote_rows)):
    assert qws.cell(r,1).data_type == 's'
    assert qws.cell(r,2).data_type == 's'
    # ProjectID can be blank (None) for missing values, otherwise text.
    assert qws.cell(r,3).value is None or qws.cell(r,3).data_type == 's'
    assert qws.cell(r,8).data_type == 'n'
    assert qws.cell(r,10).data_type == 'n'
    assert qws.cell(r,11).data_type == 'n'
# Recalculate and compare saved amounts.
for i, row in enumerate(quote_rows, start=5):
    assert money(qws.cell(i,10).value) == row['NetAmount']
    assert money(qws.cell(i,11).value) == row['NetAmountTWD']
md_text = OUT_MD.read_text(encoding='utf-8')
for e in expired_unaccepted:
    assert e['QuoteID'] in md_text
assert 'Q0010' in md_text and 'Q0014' in md_text and 'Q0017' in md_text
# Ensure accepted but before baseline (Q0015) is not listed in expired-unaccepted table after the heading.
section = md_text.split('## ValidUntil < 2026-09-30 and Status not Accepted',1)[1].split('## FX rates used',1)[0]
assert 'Q0015' not in section

print('Wrote:', OUT_XLSX)
print('Wrote:', OUT_MD)
print('Quotes fetched/written:', len(quote_rows), 'IDs:', ', '.join(quote_ids))
print('Line rows written:', len(line_rows))
print('Calculation mismatches:', sum(1 for f in calc_findings if not f['OK']))
print('Expired and not Accepted as of 2026-09-30:', ', '.join(e['QuoteID'] for e in expired_unaccepted) or 'None')
print('Summary formula:', sws['B2'].value)
print('Validation complete: reopened workbook and markdown; checked sheet names/headers, query dates, row counts, text IDs, numeric amount cells, Summary B2 formula range, retained all quotes, and expired-not-accepted findings.')
