from pathlib import Path
import subprocess, json, sys, os
manual = Path('input/quote_cli_manual.md').read_text(encoding='utf-8')
cli = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
cmd=[cli,'quotes','list','--from','2026-09-01','--to','2026-09-30','--format','json']
print('Running:', cmd)
res=subprocess.run(cmd, capture_output=True, text=True, shell=False)
print('exit', res.returncode)
print('stderr', res.stderr[:1000])
print('stdout head', res.stdout[:4000])
if res.returncode==0:
    data=json.loads(res.stdout)
    print('count', data.get('count'), 'items sample')
    for item in data.get('items', [])[:3]: print(item)
    if data.get('items'):
        qid=data['items'][0]['QuoteID']
        cmd2=[cli,'quotes','get',qid,'--format','json']
        print('Running get:', cmd2)
        r2=subprocess.run(cmd2, capture_output=True, text=True, shell=False)
        print('get exit', r2.returncode, 'stderr', r2.stderr[:1000])
        print('get stdout', r2.stdout[:3000])
print('\nFull Instructions rows:')
from openpyxl import load_workbook
wb=load_workbook('input/quote_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
