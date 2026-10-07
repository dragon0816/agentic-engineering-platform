from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook(Path('input')/'quote_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True):
    print(r)
print('\nMerged ranges:')
for sh in wb.worksheets:
    print(sh.title, list(sh.merged_cells.ranges))
