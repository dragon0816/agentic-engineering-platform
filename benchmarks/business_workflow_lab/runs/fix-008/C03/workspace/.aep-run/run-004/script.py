from pathlib import Path
import subprocess, json, csv, shutil, os, sys
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta, date
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.cell.cell import MergedCell

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
OUTFILE = OUTPUT / 'quote_report_P001-P005_2026-09_Accepted.xlsx'
CLI = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
PROJECTS = [f'P{i:03d}' for i in range(1,6)]
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Accepted'
TW = timezone(timedelta(hours=8))
GENERATED_AT = datetime.now(TW).isoformat(timespec='seconds')

TWOPLACES = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [CLI] + args
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f'CLI failed rc={res.returncode}: {cmd}\nstderr={res.stderr}\nstdout={res.stdout}')
    return json.loads(res.stdout)

# Read fixed FX rates; note field may contain commas, only first two columns are required.
fx_rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    if header[:2] != ['Currency', 'RateToTWD']:
        raise ValueError(f'Unexpected FX header: {header}')
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx_rates[row[0]] = Decimal(row[1])
print('FX rates:', {k: str(v) for k,v in fx_rates.items()})

# Fetch quote headers and detail lines.
headers = []
lines = []
fetch_counts = {}
for project in PROJECTS:
    data = run_cli(['quotes','list','--project',project,'--from',FROM_DATE,'--to',TO_DATE,'--status',STATUS,'--format','json'])
    if data.get('schema_version') != 1:
        raise ValueError(f'Unexpected schema for {project}: {data.get("schema_version")}')
    items = data.get('items', [])
    fetch_counts[project] = len(items)
    for item in items:
        if item.get('ProjectID') != project or item.get('Status') != STATUS or not (FROM_DATE <= item.get('QuoteDate','') <= TO_DATE):
            raise ValueError(f'CLI returned out-of-filter quote: {item}')
        detail = run_cli(['quotes','get',item['QuoteID'],'--format','json'])
        header = detail['header']
        # Preserve list header values, but validate detail header matches.
        for k in ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount']:
            if str(header.get(k)) != str(item.get(k)):
                raise ValueError(f'Header mismatch {item["QuoteID"]} field {k}: list={item.get(k)} get={header.get(k)}')
        curr = item['Currency']
        if curr not in fx_rates:
            raise ValueError(f'Missing FX rate for currency {curr}')
        net_twd = money(Decimal(item['NetAmount']) * fx_rates[curr])
        item = dict(item)
        item['NetAmountTWD'] = net_twd
        headers.append(item)
        for line in detail.get('line_items', []):
            line = dict(line)
            line['QuoteID'] = item['QuoteID']
            lines.append(line)

headers.sort(key=lambda h: (h['ProjectID'], h['QuoteDate'], h['QuoteID']))
lines.sort(key=lambda l: (l['QuoteID'], l['ProductID']))
print('Fetched accepted quotes by project:', fetch_counts)
print('QuoteIDs:', [h['QuoteID'] for h in headers])
print('Line rows:', len(lines))

# Copy and fill template.
shutil.copy2(TEMPLATE, OUTFILE)
wb = load_workbook(OUTFILE)

def clear_data(ws, min_row, max_col):
    if ws.max_row >= min_row:
        for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row, max_col=max_col):
            for cell in row:
                if not isinstance(cell, MergedCell):
                    cell.value = None

def set_text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.number_format = '@'

def set_money(cell, value):
    cell.value = float(money(value))
    cell.number_format = '0.00'

def set_number(cell, value):
    cell.value = int(value) if Decimal(str(value)) == Decimal(str(value)).to_integral() else float(Decimal(str(value)))

# Report sheet.
rpt = wb['Report']
rpt['B3'] = FROM_DATE
rpt['D3'] = TO_DATE
rpt['B4'] = 'ProjectID P001-P005; QuoteDate 2026-09-01 through 2026-09-30 inclusive; Status Accepted'
rpt['B5'] = GENERATED_AT
rpt['B7'] = 'quote_cli mock; fx_rates.csv'
rpt['D7'] = GENERATED_AT
no_quote = [p for p,c in fetch_counts.items() if c == 0]
rpt['B9'] = ('No Accepted quotes for: ' + ', '.join(no_quote) + '. ' if no_quote else '') + 'NetAmountTWD calculated from NetAmount × fx_rates.csv using ROUND_HALF_UP; original currency amounts summarized by ProjectID and Currency to avoid mixed-currency addition.'

# Quotes sheet.
qws = wb['Quotes']
clear_data(qws, 5, 11)
quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for r, h in enumerate(headers, start=5):
    for c, key in enumerate(quote_cols, start=1):
        cell = qws.cell(r, c)
        if key in ['QuoteID','CustomerID','ProjectID','Status','Currency']:
            set_text(cell, h[key])
        elif key in ['QuoteDate','ValidUntil']:
            cell.value = h[key]
            cell.number_format = 'yyyy-mm-dd'
        elif key in ['Subtotal','NetAmount','NetAmountTWD']:
            set_money(cell, h[key])
        else:
            set_number(cell, h[key])

# Lines sheet.
lws = wb['Lines']
clear_data(lws, 5, 6)
line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for r, line in enumerate(lines, start=5):
    for c, key in enumerate(line_cols, start=1):
        cell = lws.cell(r, c)
        if key in ['QuoteID','ProductID','ProductName']:
            set_text(cell, line.get(key, ''))
        elif key in ['UnitPrice','LineAmount']:
            set_money(cell, line.get(key, '0'))
        else:
            set_number(cell, line.get(key, '0'))

# Summary: ProjectID + Currency to avoid adding original amounts across currencies. Projects with no quote retain one zero row.
sws = wb['Summary']
clear_data(sws, 5, 6)
# B2 total formula with extended Quotes range.
last_quote_row = max(5, 4 + len(headers))
sws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
summary_rows = []
for project in PROJECTS:
    ph = [h for h in headers if h['ProjectID'] == project]
    if not ph:
        summary_rows.append({'Group': project, 'SubGroup': '', 'Currency': '', 'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0})
    else:
        by_curr = {}
        for h in ph:
            curr = h['Currency']
            by_curr.setdefault(curr, {'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0})
            by_curr[curr]['Amount'] += Decimal(h['NetAmount'])
            by_curr[curr]['AmountTWD'] += Decimal(str(h['NetAmountTWD']))
            by_curr[curr]['Count'] += 1
        for curr in sorted(by_curr):
            summary_rows.append({'Group': project, 'SubGroup': '', 'Currency': curr, 'Amount': money(by_curr[curr]['Amount']), 'AmountTWD': money(by_curr[curr]['AmountTWD']), 'Count': by_curr[curr]['Count']})
for r, row in enumerate(summary_rows, start=5):
    set_text(sws.cell(r,1), row['Group'])
    set_text(sws.cell(r,2), row['SubGroup'])
    set_text(sws.cell(r,3), row['Currency'])
    set_money(sws.cell(r,4), row['Amount'])
    set_money(sws.cell(r,5), row['AmountTWD'])
    set_number(sws.cell(r,6), row['Count'])

# Keep footer/layout; save.
wb.save(OUTFILE)
print('Wrote:', OUTFILE)

# Reopen and validate stated requirements from the saved workbook.
wb2 = load_workbook(OUTFILE, data_only=False)
assert wb2.sheetnames == ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
q2 = wb2['Quotes']
l2 = wb2['Lines']
s2 = wb2['Summary']
assert [q2.cell(4,c).value for c in range(1,12)] == quote_cols
assert [l2.cell(4,c).value for c in range(1,7)] == line_cols
assert [s2.cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
# Quotes exact filters and TWD calculation.
quote_rows = []
for r in range(5, 5+len(headers)):
    vals = [q2.cell(r,c).value for c in range(1,12)]
    quote_rows.append(vals)
    assert vals[2] in PROJECTS
    assert FROM_DATE <= vals[3] <= TO_DATE
    assert vals[5] == STATUS
    assert isinstance(vals[0], str) and isinstance(vals[1], str) and isinstance(vals[2], str)
    expected_twd = money(Decimal(str(vals[9])) * fx_rates[vals[6]])
    assert money(vals[10]) == expected_twd, (vals, expected_twd)
# line quote IDs all correspond.
quote_ids = {h['QuoteID'] for h in headers}
for r in range(5, 5+len(lines)):
    assert l2.cell(r,1).value in quote_ids
    assert isinstance(l2.cell(r,1).value, str)
# Summary includes all projects and no mixed currency original amount.
summary_data = []
for r in range(5, 5+len(summary_rows)):
    row = [s2.cell(r,c).value for c in range(1,7)]
    summary_data.append(row)
projects_in_summary = [row[0] for row in summary_data]
for p in PROJECTS:
    assert p in projects_in_summary, (p, summary_data)
for p in no_quote:
    rows = [row for row in summary_data if row[0] == p]
    assert len(rows) == 1 and rows[0][3] == 0 and rows[0][4] == 0 and rows[0][5] == 0, rows
# Total formula extended correctly.
assert s2['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})'
# Formula evaluated value not available from openpyxl, so compare summary TWD total and quotes TWD total.
quote_twd_sum = money(sum(Decimal(str(row[10])) for row in quote_rows))
summary_twd_sum = money(sum(Decimal(str(row[4])) for row in summary_data))
assert quote_twd_sum == summary_twd_sum
print('VALIDATION OK')
print('Saved file:', OUTFILE.name)
print('Quotes rows:', len(headers), 'Lines rows:', len(lines), 'Summary rows:', len(summary_rows))
print('Quote TWD total:', quote_twd_sum)
print('Summary rows:')
for row in summary_data:
    print(row)
