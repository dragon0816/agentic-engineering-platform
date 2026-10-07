from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace=', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# workbook structure
wb_path = ws/'input'/'quote_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print(f'\nSheet {sh.title}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 10), values_only=True):
        print(r)

# fx csv
fx_path = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i, row in enumerate(reader):
        print(row)
        if i >= 10:
            break

# manual
manual = ws/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md first 120 lines:')
text = manual.read_text(encoding='utf-8-sig', errors='replace').splitlines()
for line in text[:120]:
    print(line)
