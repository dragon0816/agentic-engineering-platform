from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Alignment

WORKSPACE = Path.cwd()
IN = WORKSPACE / 'input'
OUT = WORKSPACE / 'output'
OUT.mkdir(exist_ok=True)
TEMPLATE = IN / 'expense_template.xlsx'
FX = IN / 'fx_rates.csv'
EXPORT = OUT / 'expense_2026-09_approved.csv'
OUTPUT = OUT / 'expense_2026-09_approved_report.xlsx'

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Approved'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
FETCHED_AT = '2026-10-06T00:00:00+08:00'

TWOPLACES = Decimal('0.01')
def money(d):
    return d.quantize(TWOPLACES, rounding=ROUND_HALF_UP)

# Read fixed FX rates (first two columns only; Note may contain unquoted commas in test file)
rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates loaded:', rates)

# Read exported Expense Viewer rows
rows = []
with EXPORT.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('Export headers:', reader.fieldnames)
    for r in reader:
        if r['ExpenseDate'] < FROM_DATE or r['ExpenseDate'] > TO_DATE or r['Status'] != STATUS:
            raise ValueError(f"Unexpected exported row outside requested filter: {r}")
        cur = r['Currency']
        if cur not in rates:
            raise ValueError(f'Missing FX rate for currency {cur}')
        amount = Decimal(r['Amount'])
        amount_twd = money(amount * rates[cur])
        r['_AmountDecimal'] = amount
        r['_AmountTWDDecimal'] = amount_twd
        r['_NeedsReview'] = 'true' if amount_twd > Decimal('5000.00') else 'false'
        # normalize MissingReceipt to requested true/false lower-case text
        mr = (r.get('MissingReceipt') or '').strip().lower()
        if mr in ('true', '1', 'yes'):
            r['_MissingReceipt'] = 'true'
        elif mr in ('false', '0', 'no'):
            r['_MissingReceipt'] = 'false'
        else:
            r['_MissingReceipt'] = ''
        rows.append(r)
print('Rows read:', len(rows))

# Load template copy and fill it
wb = load_workbook(TEMPLATE)
report = wb['Report']
details = wb['Details']
summary = wb['Summary']
exceptions = wb['Exceptions']

# Utility: clear data rows without touching headers/layout above row 5.
def clear_rows(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

clear_rows(details, 5)
clear_rows(summary, 5)
clear_rows(exceptions, 5)

# Report metadata
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = 'ExpenseDate 2026-09-01 to 2026-09-30; Status=Approved; Employee=all; Project=all'
report['B5'] = GENERATED_AT
report['B7'] = 'Expense Viewer desktop export; fx_rates.csv fixed rates'
report['D7'] = FETCHED_AT
report['B9'] = 'AmountTWD = original Amount × fixed RateToTWD, ROUND_HALF_UP to 2 decimals. NeedsReview=true when AmountTWD > 5000. Missing source values left blank.'

# Fill details
headers = [details.cell(4, c).value for c in range(1, details.max_column + 1)]
print('Details headers:', headers)
for idx, r in enumerate(rows, start=5):
    copy_row_style(details, 4, idx, details.max_column)
    vals = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['ExpenseDate'],
        r['Category'], r['Currency'], float(r['_AmountDecimal']), float(r['_AmountTWDDecimal']),
        r['Status'], r.get('ReceiptID') or None, r.get('Description') or None,
        r['_MissingReceipt'], r['_NeedsReview']
    ]
    for col, val in enumerate(vals, start=1):
        cell = details.cell(idx, col, val)
        if col in (1,2,3,4,6,7,10,11,12,13,14):
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9):
            cell.number_format = '0.00'
        if col == 12:
            cell.alignment = Alignment(wrap_text=True)

# Summary by Employee and Category, keeping each original Currency separate to avoid mixing original amounts.
summary_groups = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
emp_names = {}
for r in rows:
    emp_key = r['EmployeeID']
    emp_names[emp_key] = r['EmployeeName']
    key = (emp_key, r['Category'], r['Currency'])
    summary_groups[key]['amount'] += r['_AmountDecimal']
    summary_groups[key]['twd'] += r['_AmountTWDDecimal']
    summary_groups[key]['count'] += 1

for idx, ((emp, cat, cur), agg) in enumerate(sorted(summary_groups.items()), start=5):
    copy_row_style(summary, 4, idx, summary.max_column)
    # Group/SubGroup as requested: employee/category. Currency column preserves original currency.
    values = [f'{emp} {emp_names.get(emp, "")}', cat, cur, float(money(agg['amount'])), float(money(agg['twd'])), agg['count']]
    for col, val in enumerate(values, start=1):
        cell = summary.cell(idx, col, val)
        if col in (1,2,3):
            cell.number_format = '@'
        elif col in (4,5):
            cell.number_format = '0.00'
        elif col == 6:
            cell.number_format = '0'

last_detail = 4 + len(rows) if rows else 5
summary['B2'] = f'=SUM(Details!I5:I{last_detail})' if rows else '=SUM(Details!I5:I5)'

# No separate exceptions requested; keep Exceptions sheet header only.

wb.save(OUTPUT)
print('Saved:', OUTPUT)
print('Detail rows:', len(rows))
print('Summary rows:', len(summary_groups))
print('NeedsReview true count:', sum(1 for r in rows if r['_NeedsReview'] == 'true'))
print('Details AmountTWD total:', money(sum((r['_AmountTWDDecimal'] for r in rows), Decimal('0.00'))))
