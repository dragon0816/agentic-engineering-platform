from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('output files:')
for p in sorted((ws/'output').iterdir()):
    print(p.name, p.stat().st_size)
print('\nexport sample:')
with (ws/'output'/'expense_2026-09_approved.csv').open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r): print(row)
print('\nInstructions all:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True): print(row)
