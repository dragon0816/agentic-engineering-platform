from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace=', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size, 'bytes')

fx = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i,row in zip(range(10), r):
        print(row)

manual = ws/'input'/'expense_manual.md'
print('\nexpense_manual.md head:')
print(manual.read_text(encoding='utf-8')[:2000])

template = ws/'input'/'expense_template.xlsx'
wb = load_workbook(template)
print('\ntemplate sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('\nSheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        print(row)
