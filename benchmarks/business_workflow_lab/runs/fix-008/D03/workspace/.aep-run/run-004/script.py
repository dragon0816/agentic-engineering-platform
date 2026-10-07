from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from openpyxl import load_workbook

ws = Path.cwd()
out = ws/'output'/'expense_2026-09_approved_report.xlsx'
export = ws/'output'/'expense_2026-09_approved.csv'
fx = ws/'input'/'fx_rates.csv'
assert out.exists(), f'Missing output workbook: {out}'

# expected exported rows and FX
rates = {}
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        if row:
            rates[row[0]] = Decimal(row[1])
exp_rows = []
with export.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for r in reader:
        assert r['Status'] == 'Approved'
        assert '2026-09-01' <= r['ExpenseDate'] <= '2026-09-30'
        amount = Decimal(r['Amount'])
        twd = (amount * rates[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        exp_rows.append((r, twd, 'true' if twd > Decimal('5000.00') else 'false'))

wb = load_workbook(out, data_only=False)
assert wb.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions'], wb.sheetnames
rep, det, summ, exc = wb['Report'], wb['Details'], wb['Summary'], wb['Exceptions']

# report fields
assert rep['B3'].value == '2026-09-01'
assert rep['D3'].value == '2026-09-30'
assert 'Approved' in rep['B4'].value
assert 'Expense Viewer' in rep['B7'].value and 'fx_rates.csv' in rep['B7'].value

# details headers and rows
expected_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
headers = [det.cell(4,c).value for c in range(1,15)]
assert headers == expected_headers, headers
assert det.max_row == 4 + len(exp_rows), (det.max_row, len(exp_rows))
for i, (src, expected_twd, expected_review) in enumerate(exp_rows, start=5):
    vals = [det.cell(i,c).value for c in range(1,15)]
    assert vals[0] == src['ExpenseID']
    assert vals[1] == src['EmployeeID']
    assert vals[2] == src['EmployeeName']
    assert vals[3] == src['ProjectID']
    assert vals[4] == src['ExpenseDate']
    assert vals[5] == src['Category']
    assert vals[6] == src['Currency']
    assert Decimal(str(vals[7])).quantize(Decimal('0.01')) == Decimal(src['Amount']).quantize(Decimal('0.01'))
    assert Decimal(str(vals[8])).quantize(Decimal('0.01')) == expected_twd
    assert vals[9] == 'Approved'
    assert (vals[10] or '') == src['ReceiptID']
    assert (vals[11] or '') == src['Description']
    assert vals[12] in ('true','false','')
    assert vals[13] == expected_review

# summary: one row per EmployeeID/Category/Currency, with Group containing employee and SubGroup category
expected_groups = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
for src, twd, review in exp_rows:
    key = (src['EmployeeID'], src['Category'], src['Currency'])
    expected_groups[key]['amount'] += Decimal(src['Amount'])
    expected_groups[key]['twd'] += twd
    expected_groups[key]['count'] += 1
assert summ['B2'].value == f'=SUM(Details!I5:I{4+len(exp_rows)})'
summary_rows = {}
for i in range(5, summ.max_row+1):
    group = summ.cell(i,1).value
    emp = group.split(' ', 1)[0]
    key = (emp, summ.cell(i,2).value, summ.cell(i,3).value)
    summary_rows[key] = {
        'amount': Decimal(str(summ.cell(i,4).value)).quantize(Decimal('0.01')),
        'twd': Decimal(str(summ.cell(i,5).value)).quantize(Decimal('0.01')),
        'count': summ.cell(i,6).value,
    }
assert set(summary_rows) == set(expected_groups), (set(summary_rows), set(expected_groups))
for key, exp in expected_groups.items():
    got = summary_rows[key]
    assert got['amount'] == exp['amount'].quantize(Decimal('0.01'))
    assert got['twd'] == exp['twd'].quantize(Decimal('0.01'))
    assert got['count'] == exp['count']

assert exc.max_row == 4, f'Exceptions data rows unexpectedly present: max_row={exc.max_row}'
print('Validated output workbook:', out)
print('Validated details rows:', len(exp_rows))
print('Validated all rows are Approved and within 2026-09-01..2026-09-30')
print('Validated FX conversion ROUND_HALF_UP and NeedsReview > 5000 rule')
print('Validated original Currency and Amount are preserved')
print('Validated summary rows:', len(summary_rows), 'grouped by Employee/Category with Currency retained')
print('Validated Summary!B2 formula:', summ['B2'].value)
print('Validated Exceptions sheet has no data rows')
