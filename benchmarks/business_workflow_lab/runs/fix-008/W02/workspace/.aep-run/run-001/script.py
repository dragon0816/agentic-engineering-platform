import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('CWD', ws)
print('INPUT FILES')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
# Inspect workbook
wb_path=ws/'input'/'sales_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('WORKBOOK SHEETS', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(r)
# Inspect CSV
csv_path=ws/'input'/'fx_rates.csv'
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader):
        print('CSV', i, row)
# Inspect manual beginning
md_path=ws/'input'/'sap_manual.md'
print('MANUAL HEAD')
print(md_path.read_text(encoding='utf-8')[:4000])