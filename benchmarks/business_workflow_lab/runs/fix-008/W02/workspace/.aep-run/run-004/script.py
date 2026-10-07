from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
import csv
from collections import defaultdict

ws = Path.cwd()
infile = ws/'input'/'sales_template.xlsx'
outdir = ws/'output'
outdir.mkdir(exist_ok=True)
outfile = outdir/'sales_report_C001_2026-08-01_2026-10-31.xlsx'

orders = [
    {"OrderID":"0045000101","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-08-03","Status":"Closed","Currency":"TWD","NetAmount":"95895.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET15"},
    {"OrderID":"0045000107","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P003","OrderDate":"2026-08-27","Status":"Closed","Currency":"USD","NetAmount":"6250.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET30"},
    {"OrderID":"0045000109","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P003","OrderDate":"2026-09-01","Status":"Open","Currency":"TWD","NetAmount":"150000.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET10"},
    {"OrderID":"0045000114","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-09-08","Status":"Cancelled","Currency":"USD","NetAmount":"11956.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET5"},
    {"OrderID":"0045000119","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-09-15","Status":"Open","Currency":"EUR","NetAmount":"8775.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET30"},
    {"OrderID":"0045000125","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-09-25","Status":"Closed","Currency":"TWD","NetAmount":"386445.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET0"},
    {"OrderID":"0045000129","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-10-01","Status":"Open","Currency":"USD","NetAmount":"1602.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET30"},
    {"OrderID":"0045000132","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P003","OrderDate":"2026-10-09","Status":"Open","Currency":"TWD","NetAmount":"50400.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":""},
    {"OrderID":"0045000136","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-10-31","Status":"Open","Currency":"EUR","NetAmount":"2250.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET5"},
]

rates = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

q = Decimal('0.01')
for o in orders:
    o['NetAmountTWD'] = (Decimal(o['NetAmount']) * rates[o['Currency']]).quantize(q, rounding=ROUND_HALF_UP)

wb = load_workbook(infile)
report = wb['Report']
details = wb['Details']
summary = wb['Summary']

tz = timezone(timedelta(hours=8))
generated_at = datetime.now(tz).isoformat(timespec='seconds')
report['B3'] = '2026-08-01'
report['D3'] = '2026-10-31'
report['B4'] = 'CustomerID = C001 (exact match); Sales Orders by OrderDate, all statuses'
report['B5'] = generated_at
report['B7'] = 'Lab ERP Sales Orders; fx_rates.csv'
report['D7'] = generated_at
missing_payment = [o['OrderID'] for o in orders if not o.get('PaymentTerms')]
missing_contact = [o['OrderID'] for o in orders if not o.get('CustomerContact')]
notes = ['Filtered by exact CustomerID C001; similar customer C007/Acme Precision Taiwan Limited excluded.']
if missing_contact:
    notes.append('Missing CustomerContact: ' + ', '.join(missing_contact) + '.')
if missing_payment:
    notes.append('Missing PaymentTerms: ' + ', '.join(missing_payment) + '.')
if not missing_contact and not missing_payment:
    notes.append('No missing CustomerContact or PaymentTerms.')
report['B9'] = ' '.join(notes)

for row in details.iter_rows(min_row=5, max_row=max(details.max_row, 100), min_col=1, max_col=11):
    for cell in row:
        cell.value = None

for idx, o in enumerate(orders, start=5):
    values = [
        o['OrderID'], o['CustomerID'], o['CustomerName'], o['ProjectID'], o['OrderDate'],
        o['Status'], o['Currency'], float(Decimal(o['NetAmount'])), float(o['NetAmountTWD']),
        o['CustomerContact'], o['PaymentTerms'] if o['PaymentTerms'] != '' else None
    ]
    for col, val in enumerate(values, start=1):
        details.cell(idx, col).value = val
    for col in [1,2,4]:
        details.cell(idx, col).number_format = '@'
    for col in [8,9]:
        details.cell(idx, col).number_format = '0.00'

last_detail_row = 4 + len(orders) if orders else 5
summary['B2'] = f'=SUM(Details!I5:I{last_detail_row if orders else 5})'
for row in summary.iter_rows(min_row=5, max_row=max(summary.max_row, 100), min_col=1, max_col=6):
    for cell in row:
        cell.value = None

agg = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
for o in orders:
    key = (o['Status'], '', o['Currency'])
    agg[key]['amount'] += Decimal(o['NetAmount'])
    agg[key]['twd'] += o['NetAmountTWD']
    agg[key]['count'] += 1

status_order = {'Open':0, 'Closed':1, 'Cancelled':2}
currency_order = {'TWD':0, 'USD':1, 'EUR':2}
for row, ((status, subgroup, currency), a) in enumerate(sorted(agg.items(), key=lambda kv: (status_order.get(kv[0][0], 99), currency_order.get(kv[0][2], 99), kv[0][2])), start=5):
    summary.cell(row,1).value = status
    summary.cell(row,2).value = None
    summary.cell(row,3).value = currency
    summary.cell(row,4).value = float(a['amount'].quantize(q, rounding=ROUND_HALF_UP))
    summary.cell(row,5).value = float(a['twd'].quantize(q, rounding=ROUND_HALF_UP))
    summary.cell(row,6).value = a['count']
    summary.cell(row,4).number_format = '0.00'
    summary.cell(row,5).number_format = '0.00'

wb.save(outfile)
print('WROTE', outfile)

# Verification from saved output
wb2 = load_workbook(outfile, data_only=False)
rep2, det2, sum2 = wb2['Report'], wb2['Details'], wb2['Summary']
assert wb2.sheetnames == ['Report', 'Details', 'Summary', 'Instructions']
assert rep2['B3'].value == '2026-08-01'
assert rep2['D3'].value == '2026-10-31'
assert 'C001' in rep2['B4'].value
assert 'C007' in rep2['B9'].value and 'excluded' in rep2['B9'].value
rows=[]
for r in range(5, 5+len(orders)):
    rows.append([det2.cell(r,c).value for c in range(1,12)])
assert len(rows)==9
assert all(r[1]=='C001' for r in rows), rows
assert 'C007' not in [r[1] for r in rows]
assert [r[0] for r in rows] == [o['OrderID'] for o in orders]
assert det2['A5'].data_type == 's' and det2['A5'].value.startswith('00')
assert isinstance(det2['H5'].value, (int,float))
assert isinstance(det2['I5'].value, (int,float))
row_0132 = next(r for r in range(5,14) if det2.cell(r,1).value == '0045000132')
assert det2.cell(row_0132,11).value is None
assert '0045000132' in rep2['B9'].value
assert sum2['B2'].value == '=SUM(Details!I5:I13)'
summary_rows = []
for r in range(5, sum2.max_row+1):
    if sum2.cell(r,1).value is not None:
        summary_rows.append([sum2.cell(r,c).value for c in range(1,7)])
expected_twd = sum(o['NetAmountTWD'] for o in orders)
summary_twd = sum(Decimal(str(r[4])).quantize(q) for r in summary_rows)
assert summary_twd == expected_twd, (summary_twd, expected_twd)
print('CHECKS PASSED')
print('detail_rows', len(rows))
print('expected_total_twd', expected_twd)
print('summary_rows', len(summary_rows), summary_rows)
print('missing_payment_terms', missing_payment)
print('missing_customer_contact', missing_contact)
print('output_file', str(outfile))