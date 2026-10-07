from pathlib import Path
import csv, json, subprocess, sys, os, shutil
from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_CSV = INPUT / 'fx_rates.csv'
MANUAL = INPUT / 'quote_cli_manual.md'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
PYCLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
OUT_XLSX = OUTPUT / 'sent_quotes_2026-09.xlsx'
LOG = OUTPUT / 'quote_cli_run_log.json'
ERRTXT = OUTPUT / 'quote_cli_error.txt'

FROM = '2026-09-01'
TO = '2026-09-30'
STATUS = 'Sent'
TZ = timezone(timedelta(hours=8))

print('Workspace:', WORKSPACE)
print('Inputs exist:', {p.name: p.exists() for p in [TEMPLATE, FX_CSV, MANUAL]})
print('CLI exists:', CLI.exists(), 'Python CLI exists:', PYCLI.exists())
print('Template sheets/headers preview:')
wb_prev = load_workbook(TEMPLATE, data_only=False)
for ws in wb_prev.worksheets:
    print(' ', ws.title, ws.max_row, ws.max_column)
    if ws.title in ('Quotes','Lines','Summary'):
        print('   row4:', [ws.cell(4,c).value for c in range(1, ws.max_column+1)])
print('FX preview:')
with FX_CSV.open('r', encoding='utf-8-sig', newline='') as f:
    for i,row in zip(range(4), csv.reader(f)):
        print(' ', i, row)

run_log = []

def run_cli(args, label):
    # exit code 4: initial try + up to 2 retries = max 3 attempts total.
    attempts = []
    max_attempts = 3
    for attempt in range(1, max_attempts+1):
        if CLI.exists():
            cmd = [str(CLI)] + args
            shell = False
        else:
            cmd = [sys.executable, str(PYCLI)] + args
            shell = False
        cp = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, shell=shell)
        rec = {
            'label': label,
            'attempt': attempt,
            'args': args,
            'exit_code': cp.returncode,
            'stdout': cp.stdout,
            'stderr': cp.stderr,
        }
        attempts.append(rec)
        run_log.append(rec)
        print(f"CLI {label} attempt {attempt}: exit={cp.returncode}; stderr={cp.stderr.strip()!r}; stdout_len={len(cp.stdout)}")
        if cp.returncode == 4 and attempt < max_attempts:
            continue
        return cp.returncode, cp.stdout, cp.stderr, attempts
    return attempts[-1]['exit_code'], attempts[-1]['stdout'], attempts[-1]['stderr'], attempts

def parse_json_stdout(exit_code, stdout, label):
    if exit_code != 0:
        return None
    try:
        return json.loads(stdout)
    except Exception as e:
        raise RuntimeError(f'{label}: stdout was returned with exit 0 but is not parseable JSON: {e}')

def write_error(reason, code, stderr):
    payload = {
        'status': 'failed',
        'reason': reason,
        'exit_code': code,
        'stderr': stderr,
        'period': {'from': FROM, 'to': TO},
        'status_filter': STATUS,
        'note': 'Per instruction, stdout is ignored unless exit code is 0 and parseable; exit code 2/3 is not retried.',
        'run_log': run_log,
    }
    LOG.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding='utf-8')
    ERRTXT.write_text(f"quote_cli failed\nexit_code={code}\nstderr={stderr}\nreason={reason}\n", encoding='utf-8')
    print('Wrote error files:', LOG, ERRTXT)

# 1) Query headers.
args = ['quotes', 'list', '--from', FROM, '--to', TO, '--status', STATUS, '--format', 'json']
code, out, err, attempts = run_cli(args, 'quotes list Sent 2026-09')
if code in (2,3):
    write_error('Non-retryable quote_cli error while listing quotes', code, err.strip())
    print('STOP: non-retryable exit code; not filling workbook.')
    sys.exit(0)
if code == 4:
    write_error('Temporary service error remained after maximum retries while listing quotes', code, err.strip())
    print('STOP: retry limit reached; not filling workbook.')
    sys.exit(0)
headers_obj = parse_json_stdout(code, out, 'quotes list')
if not isinstance(headers_obj, dict) or headers_obj.get('schema_version') != 1 or not isinstance(headers_obj.get('items'), list):
    raise RuntimeError('quotes list returned JSON but not expected schema_version/items object')
headers = headers_obj['items']

# 2) Get line items for each quote.
details = []
for h in headers:
    qid = h.get('QuoteID')
    if not qid:
        raise RuntimeError('Quote header missing QuoteID')
    code, out, err, attempts = run_cli(['quotes','get', qid, '--format','json'], f'quotes get {qid}')
    if code in (2,3):
        write_error(f'Non-retryable quote_cli error while getting {qid}', code, err.strip())
        print('STOP: non-retryable exit code; not filling workbook.')
        sys.exit(0)
    if code == 4:
        write_error(f'Temporary service error remained after maximum retries while getting {qid}', code, err.strip())
        print('STOP: retry limit reached; not filling workbook.')
        sys.exit(0)
    obj = parse_json_stdout(code, out, f'quotes get {qid}')
    if not isinstance(obj, dict) or obj.get('schema_version') != 1 or 'header' not in obj or 'line_items' not in obj:
        raise RuntimeError(f'quotes get {qid} returned JSON but not expected detail schema')
    details.append(obj)

# 3) FX rates. The CSV Note field has commas in unquoted text; only first two columns matter.
rates = {}
with FX_CSV.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates parsed:', {k: str(v) for k,v in rates.items()})

# Decimal helpers.
def D(v):
    try:
        return Decimal(str(v))
    except InvalidOperation:
        raise RuntimeError(f'Cannot parse decimal: {v!r}')

def q2(v):
    return D(v).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# 4) Fill a copy of the template, preserving workbook structure.
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)
ws_report = wb['Report']
ws_quotes = wb['Quotes']
ws_lines = wb['Lines']
ws_summary = wb['Summary']

# Clear old data rows if any.
for ws in (ws_quotes, ws_lines, ws_summary):
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

now = datetime.now(TZ).replace(microsecond=0).isoformat()
source = 'quote_cli_mock quote_cli'
condition = f'QuoteDate {FROM} to {TO} inclusive; Status={STATUS}'
ws_report['B3'] = FROM
ws_report['D3'] = TO
ws_report['B4'] = condition
ws_report['B5'] = now
ws_report['B7'] = source
ws_report['D7'] = now
# Additional useful metadata if labels exist in first column.
for r in range(1, ws_report.max_row+1):
    v = ws_report.cell(r,1).value
    if isinstance(v, str):
        if '筆數' in v or 'Count' in v:
            ws_report.cell(r,2).value = len(headers)

quote_headers = [ws_quotes.cell(4,c).value for c in range(1, ws_quotes.max_column+1)]
line_headers = [ws_lines.cell(4,c).value for c in range(1, ws_lines.max_column+1)]
summary_headers = [ws_summary.cell(4,c).value for c in range(1, ws_summary.max_column+1)]
expected_q = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
expected_l = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if quote_headers != expected_q:
    raise RuntimeError(f'Unexpected Quotes headers: {quote_headers}')
if line_headers != expected_l:
    raise RuntimeError(f'Unexpected Lines headers: {line_headers}')

quote_rows = []
line_rows = []
summary = {}
for obj in details:
    h = obj['header']
    if h.get('Status') != STATUS or not (FROM <= h.get('QuoteDate','') <= TO):
        raise RuntimeError(f'CLI returned record outside filter: {h}')
    cur = h.get('Currency')
    if cur not in rates:
        raise RuntimeError(f'Missing FX rate for {cur}')
    subtotal = q2(h['Subtotal'])
    disc = D(h['DiscountPercent'])
    net = q2(h['NetAmount'])
    calc_net = q2(subtotal * (Decimal('1') - disc / Decimal('100')))
    if net != calc_net:
        raise RuntimeError(f'NetAmount check failed for {h.get("QuoteID")}: {net} != {calc_net}')
    net_twd = q2(net * rates[cur])
    row = [h.get(k) for k in expected_q[:-1]] + [net_twd]
    quote_rows.append(row)
    summary.setdefault(cur, {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0})
    summary[cur]['amount'] += net
    summary[cur]['amount_twd'] += net_twd
    summary[cur]['count'] += 1
    for li in obj['line_items']:
        qty = D(li['Quantity'])
        unit = q2(li['UnitPrice'])
        line_amount = q2(li['LineAmount'])
        if line_amount != q2(qty * unit):
            raise RuntimeError(f'LineAmount check failed for {h.get("QuoteID")} {li}')
        line_rows.append([h['QuoteID'], li.get('ProductID'), li.get('ProductName'), int(qty) if qty == qty.to_integral() else qty, unit, line_amount])

# Stable order from list result.
for r_idx, row in enumerate(quote_rows, start=5):
    for c_idx, val in enumerate(row, start=1):
        cell = ws_quotes.cell(r_idx, c_idx)
        cell.value = float(val) if isinstance(val, Decimal) else val
for r_idx, row in enumerate(line_rows, start=5):
    for c_idx, val in enumerate(row, start=1):
        cell = ws_lines.cell(r_idx, c_idx)
        cell.value = float(val) if isinstance(val, Decimal) else val

# Summary: one row per currency plus grand total row. Keep source column names.
r = 5
for cur in sorted(summary.keys()):
    item = summary[cur]
    vals = ['Status', STATUS, cur, q2(item['amount']), q2(item['amount_twd']), item['count']]
    for c, val in enumerate(vals, start=1):
        ws_summary.cell(r,c).value = float(val) if isinstance(val, Decimal) else val
    r += 1
total_twd = q2(sum((item['amount_twd'] for item in summary.values()), Decimal('0.00')))
vals = ['Total', STATUS, 'TWD', total_twd, total_twd, len(quote_rows)]
for c, val in enumerate(vals, start=1):
    ws_summary.cell(r,c).value = float(val) if isinstance(val, Decimal) else val
# Formula instruction row B2 should reflect all quote rows.
last_quote_row = 4 + len(quote_rows)
ws_summary['B2'] = f'=SUM(Quotes!K5:K{max(5,last_quote_row)})'

# Number formats.
for ws in (ws_quotes, ws_lines, ws_summary):
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            if isinstance(cell.value, (int,float)) and cell.column not in (4,6):
                cell.number_format = '#,##0.00'
# Keep quantity/count integer formats.
for row in ws_lines.iter_rows(min_row=5, min_col=4, max_col=4):
    row[0].number_format = '0'
for row in ws_summary.iter_rows(min_row=5, min_col=6, max_col=6):
    row[0].number_format = '0'

wb.save(OUT_XLSX)

# Save log after success.
LOG.write_text(json.dumps({
    'status': 'succeeded',
    'period': {'from': FROM, 'to': TO},
    'status_filter': STATUS,
    'quote_count': len(quote_rows),
    'line_count': len(line_rows),
    'output': str(OUT_XLSX),
    'run_log': run_log,
}, ensure_ascii=False, indent=2), encoding='utf-8')
if ERRTXT.exists():
    ERRTXT.unlink()

# 5) Reopen and verify stated requirements.
wb2 = load_workbook(OUT_XLSX, data_only=False)
checks = []
checks.append(('output workbook exists', OUT_XLSX.exists()))
checks.append(('source template sheets preserved', wb2.sheetnames == wb_prev.sheetnames))
checks.append(('report from date', wb2['Report']['B3'].value == FROM))
checks.append(('report to date', wb2['Report']['D3'].value == TO))
checks.append(('report condition mentions Sent', 'Sent' in str(wb2['Report']['B4'].value)))
q_rows = list(wb2['Quotes'].iter_rows(min_row=5, values_only=True))
l_rows = list(wb2['Lines'].iter_rows(min_row=5, values_only=True))
# Trim all-empty rows
q_rows = [row for row in q_rows if any(v is not None for v in row)]
l_rows = [row for row in l_rows if any(v is not None for v in row)]
checks.append(('quote row count matches CLI count', len(q_rows) == headers_obj.get('count') == len(headers)))
checks.append(('line row count matches details', len(l_rows) == len(line_rows)))
checks.append(('all QuoteID values are text', all(isinstance(row[0], str) for row in q_rows)))
checks.append(('all quote statuses Sent', all(row[5] == STATUS for row in q_rows)))
checks.append(('all quote dates in September 2026', all(FROM <= row[3] <= TO for row in q_rows)))
checks.append(('NetAmountTWD numeric', all(isinstance(row[10], (int,float)) for row in q_rows)))
checks.append(('summary B2 formula updated', wb2['Summary']['B2'].value == f'=SUM(Quotes!K5:K{max(5,last_quote_row)})'))
for name, ok in checks:
    print('CHECK', name, 'PASS' if ok else 'FAIL')
if not all(ok for _,ok in checks):
    raise RuntimeError('One or more output verification checks failed')
print('Wrote workbook:', OUT_XLSX)
print('Wrote run log:', LOG)
print('Quote rows:', len(q_rows), 'Line rows:', len(l_rows), 'Total TWD:', total_twd)
