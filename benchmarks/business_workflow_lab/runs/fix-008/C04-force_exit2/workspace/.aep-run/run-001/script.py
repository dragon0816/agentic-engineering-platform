from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

# CSV headers/sample
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nFX CSV:')
    with fx.open('r', encoding='utf-8-sig', newline='') as f:
        r = csv.reader(f)
        for i, row in zip(range(5), r):
            print(i, row)

# Workbook sheets and dimensions, first few rows values
xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    print('\nWORKBOOK:')
    wb = load_workbook(xlsx, data_only=False)
    print('sheets:', wb.sheetnames)
    for sh in wb.worksheets:
        print('sheet', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
            print(' ', row)

# Manual excerpt
md = ws/'input'/'quote_cli_manual.md'
if md.exists():
    print('\nMANUAL EXCERPT:')
    text = md.read_text(encoding='utf-8', errors='replace')
    print(text[:4000])