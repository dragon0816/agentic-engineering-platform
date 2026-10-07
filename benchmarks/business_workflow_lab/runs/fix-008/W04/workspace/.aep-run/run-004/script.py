from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook(Path('input/sales_template.xlsx'), data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
print('\nCSV full:')
print(Path('output/downloads/orders_2026-09_open.csv').read_text(encoding='utf-8-sig'))
