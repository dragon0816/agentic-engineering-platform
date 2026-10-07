from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:', [p.name for p in (ws/'input').glob('*')])
print('output files:', [str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()])
manual=ws/'input'/'sap_manual.md'
print('\nmanual head:')
print(manual.read_text(encoding='utf-8')[:2000])
for csvp in list((ws/'input').glob('*.csv'))+list((ws/'output'/'downloads').glob('*.csv')):
    print('\nCSV', csvp.relative_to(ws))
    with csvp.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(6),r): print(i,row)
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
print('\ntemplate sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet',s,'dim',sh.max_row,sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(8,sh.max_row), values_only=True): print(row)
