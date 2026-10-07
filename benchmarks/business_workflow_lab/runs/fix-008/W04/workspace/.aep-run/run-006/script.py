from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
import csv, shutil
ws=Path.cwd()
src=ws/'output'/'sales_template_filled_2026-09_open.xlsx'
dst=ws/'output'/'sales_template.xlsx'
shutil.copyfile(src, dst)

# Reopen and validate the saved workbook, not in-memory objects.
wb=load_workbook(dst, data_only=False)
assert wb.sheetnames == ['Report','Details','Summary','Instructions'], wb.sheetnames
report=wb['Report']; details=wb['Details']; summary=wb['Summary']
assert report['B3'].value == '2026-09-01'
assert report['D3'].value == '2026-09-30'
assert 'Status=Open' in report['B4'].value
assert 'Lab ERP' in report['B7'].value
assert details.max_row == 18, details.max_row
headers=[details.cell(4,c).value for c in range(1,12)]
expected=['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert headers==expected, headers
ids=[]
statuses=[]
for r in range(5, details.max_row+1):
    oid=details.cell(r,1).value
    ids.append(oid)
    assert isinstance(oid, str) and oid.startswith('00') and len(oid)==10, (r,oid,type(oid))
    assert details.cell(r,1).number_format == '@', (r, details.cell(r,1).number_format)
    assert isinstance(details.cell(r,2).value, str)
    # ProjectID can be blank/None or text
    pid=details.cell(r,4).value
    assert pid is None or isinstance(pid, str), (r,pid,type(pid))
    assert details.cell(r,5).value >= '2026-09-01' and details.cell(r,5).value <= '2026-09-30', (r, details.cell(r,5).value)
    statuses.append(details.cell(r,6).value)
    assert details.cell(r,6).value == 'Open', (r, details.cell(r,6).value)
    assert details.cell(r,7).value in {'TWD','USD','EUR'}
    assert isinstance(details.cell(r,8).value, (int,float)), (r, type(details.cell(r,8).value))
    assert isinstance(details.cell(r,9).value, (int,float)), (r, type(details.cell(r,9).value))
    assert details.cell(r,8).number_format == '0.00'
    assert details.cell(r,9).number_format == '0.00'
    assert details.cell(r,11).value and str(details.cell(r,11).value).startswith('NET')
assert len(ids)==14 and len(set(ids))==14, ids
expected_ids=[]
with (ws/'output'/'downloads'/'orders_2026-09_open.csv').open(newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f): expected_ids.append(row['OrderID'])
assert ids == expected_ids, (ids, expected_ids)
assert summary['B2'].value == '=SUM(Details!I5:I18)', summary['B2'].value
assert summary.max_row == 7, summary.max_row
# Validate AmountTWD with fixed rates
rates={'TWD':Decimal('1'),'USD':Decimal('32'),'EUR':Decimal('35')}
q=Decimal('0.01')
calc=Decimal('0')
for r in range(5,19):
    amt=Decimal(str(details.cell(r,8).value))
    cur=details.cell(r,7).value
    val=(amt*rates[cur]).quantize(q, rounding=ROUND_HALF_UP)
    assert Decimal(str(details.cell(r,9).value)).quantize(q) == val, (r, details.cell(r,9).value, val)
    calc += val
summary_total=Decimal('0')
for r in range(5,8):
    assert summary.cell(r,1).value=='Status'
    assert summary.cell(r,2).value=='Open'
    assert summary.cell(r,3).value in {'TWD','USD','EUR'}
    assert isinstance(summary.cell(r,4).value,(int,float))
    assert isinstance(summary.cell(r,5).value,(int,float))
    assert isinstance(summary.cell(r,6).value,int)
    summary_total += Decimal(str(summary.cell(r,5).value)).quantize(q)
assert summary_total == calc, (summary_total, calc)
print('VALIDATED output/sales_template.xlsx')
print('sheets:', wb.sheetnames)
print('details rows:', len(ids), 'unique order IDs:', len(set(ids)))
print('order IDs:', ', '.join(ids))
print('all statuses:', sorted(set(statuses)))
print('Summary B2:', summary['B2'].value)
print('TWD total from details:', calc)
print('output files:', [str(p.relative_to(ws)) for p in (ws/'output').glob('*.xlsx')])
