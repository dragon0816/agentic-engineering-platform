from pathlib import Path
import csv
p=Path('output/downloads/orders_2026-09_open.csv')
print('exists', p.exists(), p)
if p.exists():
    print(p.read_text(encoding='utf-8-sig')[:2000])
    with open(p, newline='', encoding='utf-8-sig') as f:
        rows=list(csv.DictReader(f))
    print('rows', len(rows))
    print('ids', [r['OrderID'] for r in rows])
# more instructions
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
for r in wb['Instructions'].iter_rows(values_only=True):
    print(r)