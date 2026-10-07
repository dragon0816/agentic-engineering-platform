from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, re, shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment
from openpyxl.utils import get_column_letter

ws = Path.cwd()
input_template = ws/'input'/'sales_template.xlsx'
output_file = ws/'output'/'sales_template_filled_2026-09_open.xlsx'

# Source order list exported from the ERP UI
orders = []
with (ws/'output'/'downloads'/'orders_2026-09_open.csv').open(newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        orders.append(row)

# Detail fields read from ERP order detail pages during browser operation.
detail_texts = {
'0045000109': '''OrderID\t0045000109\nCustomerID\tC001\nCustomerName\tAcme Precision Taiwan Ltd.\nProjectID\tP003\nOrderDate\t2026-09-01\nStatus\tOpen\nCurrency\tTWD\nNetAmount\t150000.00\nCustomerContact\tGrace Lin <grace.lin@acme-precision.example>\nPaymentTerms\tNET10\nSalesRep\tE001\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1001\t4\t37500.00\t150000.00''',
'0045000110': '''OrderID\t0045000110\nCustomerID\tC002\nCustomerName\t台灣精密股份有限公司\nProjectID\tP002\nOrderDate\t2026-09-02\nStatus\tOpen\nCurrency\tTWD\nNetAmount\t293572.50\nCustomerContact\t張美玲 <mei@tw-precision.example>\nPaymentTerms\tNET30\nSalesRep\tE004\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1006\t3\t29400.00\t88200.00\n20\tMAT-1005\t3\t18457.50\t55372.50\n30\tMAT-1001\t4\t37500.00\t150000.00''',
'0045000112': '''OrderID\t0045000112\nCustomerID\tC007\nCustomerName\tAcme Precision Taiwan Limited\nProjectID\tP005\nOrderDate\t2026-09-04\nStatus\tOpen\nCurrency\tUSD\nNetAmount\t12311.00\nCustomerContact\tPeter Ho <peter.ho@acme-limited.example>\nPaymentTerms\tNET15\nSalesRep\tE001\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1002\t2\t85.50\t171.00\n20\tMAT-1004\t4\t2300.00\t9200.00\n30\tMAT-1006\t3\t980.00\t2940.00''',
'0045000113': '''OrderID\t0045000113\nCustomerID\tC004\nCustomerName\t藍海電子有限公司\nProjectID\t\nOrderDate\t2026-09-05\nStatus\tOpen\nCurrency\tTWD\nNetAmount\t15390.00\nCustomerContact\t\nPaymentTerms\tNET15\nSalesRep\tE005\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1002\t6\t2565.00\t15390.00''',
'0045000115': '''OrderID\t0045000115\nCustomerID\tC005\nCustomerName\tPacific Test Systems Inc.\nProjectID\t\nOrderDate\t2026-09-09\nStatus\tOpen\nCurrency\tUSD\nNetAmount\t6391.00\nCustomerContact\tSam Patel <sam@pacifictest.example>\nPaymentTerms\tNET20\nSalesRep\tE004\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1002\t2\t85.50\t171.00\n20\tMAT-1004\t1\t2300.00\t2300.00\n30\tMAT-1006\t4\t980.00\t3920.00''',
'0045000116': '''OrderID\t0045000116\nCustomerID\tC008\nCustomerName\t新竹量測科技\nProjectID\t\nOrderDate\t2026-09-10\nStatus\tOpen\nCurrency\tTWD\nNetAmount\t29400.00\nCustomerContact\t李建宏 <jh.lee@hsinchu-metrology.example>\nPaymentTerms\tNET30\nSalesRep\tE005\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1006\t1\t29400.00\t29400.00''',
'0045000118': '''OrderID\t0045000118\nCustomerID\tC003\nCustomerName\tNorthwind Components GmbH\nProjectID\tP004\nOrderDate\t2026-09-14\nStatus\tOpen\nCurrency\tEUR\nNetAmount\t17484.75\nCustomerContact\tJonas Krüger <jk@northwind.example>\nPaymentTerms\tNET0\nSalesRep\tE006\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1002\t5\t76.95\t384.75\n20\tMAT-1004\t5\t2070.00\t10350.00\n30\tMAT-1007\t5\t1350.00\t6750.00''',
'0045000119': '''OrderID\t0045000119\nCustomerID\tC001\nCustomerName\tAcme Precision Taiwan Ltd.\nProjectID\tP001\nOrderDate\t2026-09-15\nStatus\tOpen\nCurrency\tEUR\nNetAmount\t8775.00\nCustomerContact\tGrace Lin <grace.lin@acme-precision.example>\nPaymentTerms\tNET30\nSalesRep\tE004\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1001\t3\t1125.00\t3375.00\n20\tMAT-1007\t4\t1350.00\t5400.00''',
'0045000120': '''OrderID\t0045000120\nCustomerID\tC006\nCustomerName\tHelios Optics B.V.\nProjectID\t\nOrderDate\t2026-09-16\nStatus\tOpen\nCurrency\tEUR\nNetAmount\t461.70\nCustomerContact\tEva de Vries <eva@helios.example>\nPaymentTerms\tNET10\nSalesRep\tE003\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1002\t6\t76.95\t461.70''',
'0045000122': '''OrderID\t0045000122\nCustomerID\tC002\nCustomerName\t台灣精密股份有限公司\nProjectID\tP002\nOrderDate\t2026-09-21\nStatus\tOpen\nCurrency\tUSD\nNetAmount\t4141.00\nCustomerContact\t張美玲 <mei@tw-precision.example>\nPaymentTerms\tNET5\nSalesRep\tE006\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1003\t4\t420.00\t1680.00\n20\tMAT-1005\t4\t615.25\t2461.00''',
'0045000124': '''OrderID\t0045000124\nCustomerID\tC004\nCustomerName\t藍海電子有限公司\nProjectID\t\nOrderDate\t2026-09-24\nStatus\tOpen\nCurrency\tTWD\nNetAmount\t45000.00\nCustomerContact\t\nPaymentTerms\tNET30\nSalesRep\tE003\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1007\t1\t45000.00\t45000.00''',
'0045000126': '''OrderID\t0045000126\nCustomerID\tC008\nCustomerName\t新竹量測科技\nProjectID\t\nOrderDate\t2026-09-28\nStatus\tOpen\nCurrency\tUSD\nNetAmount\t9171.00\nCustomerContact\t李建宏 <jh.lee@hsinchu-metrology.example>\nPaymentTerms\tNET30\nSalesRep\tE003\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1002\t2\t85.50\t171.00\n20\tMAT-1007\t6\t1500.00\t9000.00''',
'0045000127': '''OrderID\t0045000127\nCustomerID\tC007\nCustomerName\tAcme Precision Taiwan Limited\nProjectID\tP005\nOrderDate\t2026-09-30\nStatus\tOpen\nCurrency\tEUR\nNetAmount\t2214.92\nCustomerContact\tPeter Ho <peter.ho@acme-limited.example>\nPaymentTerms\tNET30\nSalesRep\tE006\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1005\t4\t553.73\t2214.92''',
'0045000128': '''OrderID\t0045000128\nCustomerID\tC003\nCustomerName\tNorthwind Components GmbH\nProjectID\tP004\nOrderDate\t2026-09-30\nStatus\tOpen\nCurrency\tTWD\nNetAmount\t104160.00\nCustomerContact\tJonas Krüger <jk@northwind.example>\nPaymentTerms\tNET5\nSalesRep\tE006\nLine items\nLineNo\tProductID\tQuantity\tUnitPrice\tLineAmount\n10\tMAT-1002\t2\t2565.00\t5130.00\n20\tMAT-1005\t4\t18457.50\t73830.00\n30\tMAT-1003\t2\t12600.00\t25200.00'''
}

def parse_detail(txt):
    d = {}
    for line in txt.splitlines():
        if '\t' in line:
            k,v = line.split('\t', 1)
            if k in ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','CustomerContact','PaymentTerms','SalesRep']:
                d[k] = v
    return d

details = {oid: parse_detail(txt) for oid,txt in detail_texts.items()}

rates = {}
with (ws/'input'/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        rates[row['Currency']] = Decimal(row['RateToTWD'])

q = Decimal('0.01')
def money(x): return Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP)
def twd(amount, currency): return (Decimal(str(amount))*rates[currency]).quantize(q, rounding=ROUND_HALF_UP)

wb = load_workbook(input_template)
report = wb['Report']; det = wb['Details']; summ = wb['Summary']

# Clear existing detail and summary data rows only
for sh, start_row in [(det,5),(summ,5)]:
    if sh.max_row >= start_row:
        sh.delete_rows(start_row, sh.max_row-start_row+1)

# Report metadata
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Sales Orders: Status=Open; OrderDate from 2026-09-01 to 2026-09-30 inclusive'
report['B5'] = '2026-10-06T00:00:00+08:00'
report['B7'] = 'Lab ERP Sales Orders; fx_rates.csv'
report['D7'] = '2026-10-06T00:00:00+08:00'
report['B9'] = '14 Open orders fetched. CustomerContact is blank where source detail page was blank. Session did not require more than three recovery attempts.'

# Details rows
headers = [det.cell(4,c).value for c in range(1, det.max_column+1)]
for idx, o in enumerate(orders, start=5):
    d = details[o['OrderID']]
    vals = {
        'OrderID': d['OrderID'],
        'CustomerID': d['CustomerID'],
        'CustomerName': d['CustomerName'],
        'ProjectID': d.get('ProjectID',''),
        'OrderDate': d['OrderDate'],
        'Status': d['Status'],
        'Currency': d['Currency'],
        'NetAmount': float(money(d['NetAmount'])),
        'NetAmountTWD': float(twd(d['NetAmount'], d['Currency'])),
        'CustomerContact': d.get('CustomerContact',''),
        'PaymentTerms': d.get('PaymentTerms',''),
    }
    for col, h in enumerate(headers, start=1):
        cell = det.cell(idx, col)
        cell.value = vals.get(h, '')
        if h in ['OrderID','CustomerID','ProjectID','OrderDate','Status','Currency','CustomerContact','PaymentTerms','CustomerName']:
            cell.number_format = '@' if h in ['OrderID','CustomerID','ProjectID'] else ('yyyy-mm-dd' if h=='OrderDate' else 'General')
        if h in ['NetAmount','NetAmountTWD']:
            cell.number_format = '0.00'

last_detail = 4 + len(orders)
summ['B2'] = f'=SUM(Details!I5:I{last_detail})'

# Summary by Status + Currency to avoid mixing original currencies
summary = {}
for o in orders:
    d = details[o['OrderID']]
    key = (d['Status'], d['Currency'])
    rec = summary.setdefault(key, {'amount': Decimal('0'), 'twd': Decimal('0'), 'count':0})
    rec['amount'] += Decimal(d['NetAmount'])
    rec['twd'] += twd(d['NetAmount'], d['Currency'])
    rec['count'] += 1

for r, ((status, currency), rec) in enumerate(sorted(summary.items(), key=lambda x: (x[0][0], x[0][1])), start=5):
    vals = ['Status', status, currency, float(rec['amount'].quantize(q, rounding=ROUND_HALF_UP)), float(rec['twd'].quantize(q, rounding=ROUND_HALF_UP)), rec['count']]
    for c, v in enumerate(vals, start=1):
        cell = summ.cell(r,c)
        cell.value = v
        if c in (4,5): cell.number_format = '0.00'

output_file.parent.mkdir(parents=True, exist_ok=True)
wb.save(output_file)
print('WROTE', output_file)
print('orders', len(orders), 'detail records', len(details), 'last_detail_row', last_detail)
print('summary rows', len(summary))
