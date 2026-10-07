from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
# workbook structure
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
print('workbook sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print('\nSHEET', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(r)
# fx rates
print('\nfx_rates sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader):
        print(row)
# manual excerpt
print('\nsap_manual excerpt:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[:4000])