from pathlib import Path
from openpyxl import load_workbook
import csv, json, os
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size, 'bytes')

manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    text = manual.read_text(encoding='utf-8', errors='replace')
    print('\nquote_cli_manual.md first 120 lines:')
    for i, line in enumerate(text.splitlines()[:120], 1):
        print(f'{i:03d}: {line}')

fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv sample:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i, row in zip(range(8), r):
            print(row)

xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    wb = load_workbook(xlsx, data_only=False)
    print('\nworkbook sheets:', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print(f'\nSheet {s}: max_row={sh.max_row}, max_column={sh.max_column}')
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(row)
        print('merged:', list(sh.merged_cells.ranges))
