from pathlib import Path
from openpyxl import load_workbook
p=Path('input/quote_template.xlsx')
wb=load_workbook(p, data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)
