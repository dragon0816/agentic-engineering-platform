from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('Downloaded files:')
for p in sorted((ws/'output'/'downloads').glob('*')):
    print(p.name, p.stat().st_size)
    with p.open(newline='', encoding='utf-8-sig') as f:
        for i,row in zip(range(20), csv.reader(f)):
            print(i+1, row)

wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
sh=wb['Instructions']
print('\nInstructions all rows:')
for r in range(1, sh.max_row+1):
    print(r, [sh.cell(r,c).value for c in range(1, sh.max_column+1)])
