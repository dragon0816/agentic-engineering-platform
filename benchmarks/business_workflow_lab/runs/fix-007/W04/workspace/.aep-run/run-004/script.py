from pathlib import Path
from openpyxl import load_workbook
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
import csv, json, datetime

ws = Path.cwd()
out = ws/'output'
out.mkdir(exist_ok=True)
# Data obtained from Lab ERP browser operations; CSV list was downloaded and detail pages read via browser.
details = [
 {'OrderID':'0045000109','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P003','OrderDate':'2026-09-01','Status':'Open','Currency':'TWD','NetAmount':'150000.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET10'},
 {'OrderID':'0045000110','CustomerID':'C002','CustomerName':'台灣精密股份有限公司','ProjectID':'P002','OrderDate':'2026-09-02','Status':'Open','Currency':'TWD','NetAmount':'293572.50','CustomerContact':'張美玲 <mei@tw-precision.example>','PaymentTerms':'NET30'},
 {'OrderID':'0045000112','CustomerID':'C007','CustomerName':'Acme Precision Taiwan Limited','ProjectID':'P005','OrderDate':'2026-09-04','Status':'Open','Currency':'USD','NetAmount':'12311.00','CustomerContact':'Peter Ho <peter.ho@acme-limited.example>','PaymentTerms':'NET15'},
 {'OrderID':'0045000113','CustomerID':'C004','CustomerName':'藍海電子有限公司','ProjectID':'','OrderDate':'2026-09-05','Status':'Open','Currency':'TWD','NetAmount':'15390.00','CustomerContact':'','PaymentTerms':'NET15'},
 {'OrderID':'0045000115','CustomerID':'C005','CustomerName':'Pacific Test Systems Inc.','ProjectID':'','OrderDate':'2026-09-09','Status':'Open','Currency':'USD','NetAmount':'6391.00','CustomerContact':'Sam Patel <sam@pacifictest.example>','PaymentTerms':'NET20'},
 {'OrderID':'0045000116','CustomerID':'C008','CustomerName':'新竹量測科技','ProjectID':'','OrderDate':'2026-09-10','Status':'Open','Currency':'TWD','NetAmount':'29400.00','CustomerContact':'李建宏 <jh.lee@hsinchu-metrology.example>','PaymentTerms':'NET30'},
 {'OrderID':'0045000118','CustomerID':'C003','CustomerName':'Northwind Components GmbH','ProjectID':'P004','OrderDate':'2026-09-14','Status':'Open','Currency':'EUR','NetAmount':'17484.75','CustomerContact':'Jonas Krüger <jk@northwind.example>','PaymentTerms':'NET0'},
 {'OrderID':'0045000119','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-09-15','Status':'Open','Currency':'EUR','NetAmount':'8775.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET30'},
 {'OrderID':'0045000120','CustomerID':'C006','CustomerName':'Helios Optics B.V.','ProjectID':'','OrderDate':'2026-09-16','Status':'Open','Currency':'EUR','NetAmount':'461.70','CustomerContact':'Eva de Vries <eva@helios.example>','PaymentTerms':'NET10'},
 {'OrderID':'0045000122','CustomerID':'C002','CustomerName':'台灣精密股份有限公司','ProjectID':'P002','OrderDate':'2026-09-21','Status':'Open','Currency':'USD','NetAmount':'4141.00','CustomerContact':'張美玲 <mei@tw-precision.example>','PaymentTerms':'NET5'},
 {'OrderID':'0045000124','CustomerID':'C004','CustomerName':'藍海電子有限公司','ProjectID':'','OrderDate':'2026-09-24','Status':'Open','Currency':'TWD','NetAmount':'45000.00','CustomerContact':'','PaymentTerms':'NET30'},
 {'OrderID':'0045000126','CustomerID':'C008','CustomerName':'新竹量測科技','ProjectID':'','OrderDate':'2026-09-28','Status':'Open','Currency':'USD','NetAmount':'9171.00','CustomerContact':'李建宏 <jh.lee@hsinchu-metrology.example>','PaymentTerms':'NET30'},
 {'OrderID':'0045000127','CustomerID':'C007','CustomerName':'Acme Precision Taiwan Limited','ProjectID':'P005','OrderDate':'2026-09-30','Status':'Open','Currency':'EUR','NetAmount':'2214.92','CustomerContact':'Peter Ho <peter.ho@acme-limited.example>','PaymentTerms':'NET30'},
 {'OrderID':'0045000128','CustomerID':'C003','CustomerName':'Northwind Components GmbH','ProjectID':'P004','OrderDate':'2026-09-30','Status':'Open','Currency':'TWD','NetAmount':'104160.00','CustomerContact':'Jonas Krüger <jk@northwind.example>','PaymentTerms':'NET5'},
]
# Read source CSV to verify record count/order and no missing/extra orders
list_path = ws/'output'/'downloads'/'orders_2026-09_open.csv'
with open(list_path, newline='', encoding='utf-8-sig') as f:
    list_rows = list(csv.DictReader(f))
print('downloaded filtered list rows:', len(list_rows))
print('downloaded headers:', list(list_rows[0].keys()))
print('downloaded order ids:', [r['OrderID'] for r in list_rows])
assert len(list_rows) == 14
assert [r['OrderID'] for r in list_rows] == [r['OrderID'] for r in details]
# Verify detail header fields match downloaded list
for lr, dr in zip(list_rows, details):
    for k in ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']:
        assert lr[k] == dr[k], (k, lr[k], dr[k])
# FX rates (CSV notes contain comma; only first two columns matter)
rates = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('fx rates:', rates)
q = Decimal('0.01')
for r in details:
    amt = Decimal(r['NetAmount'])
    r['NetAmountTWD'] = (amt * rates[r['Currency']]).quantize(q, rounding=ROUND_HALF_UP)
# Fill template copy
src = ws/'input'/'sales_template.xlsx'
output = out/'sales_2026-09_open_orders.xlsx'
wb = load_workbook(src)
rep = wb['Report']
det = wb['Details']
sumsh = wb['Summary']
# Report metadata
rep['B3'] = '2026-09-01'
rep['D3'] = '2026-09-30'
rep['B4'] = 'Sales Orders: Status=Open; OrderDate 2026-09-01 through 2026-09-30'
# Use supplied today date with timezone as generated timestamp
rep['B5'] = '2026-10-06T00:00:00+08:00'
rep['B7'] = 'Lab ERP (synthetic); fx_rates.csv'
rep['D7'] = '2026-10-06T00:00:00+08:00'
rep['B9'] = '取得 14 筆 Open 訂單；2 筆 CustomerContact 為來源空白；操作期間發生 1 次逾時並重新登入後續作。'
# Clear existing data rows if any; preserve row 5 style using row 4/header if needed
for sh in [det, sumsh]:
    if sh.max_row > 4:
        sh.delete_rows(5, sh.max_row-4)
# helper copy style from row 4 to row target? Better from row 4 header not good color. If existing no row5. Use row4? formats only.
def style_from(src_cell, dst_cell):
    if src_cell.has_style:
        dst_cell._style = copy(src_cell._style)
    if src_cell.number_format:
        dst_cell.number_format = src_cell.number_format
    if src_cell.alignment:
        dst_cell.alignment = copy(src_cell.alignment)
# Details rows
cols = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
for idx, r in enumerate(details, start=5):
    for c, key in enumerate(cols, start=1):
        cell = det.cell(idx, c)
        style_from(det.cell(4, c), cell)
        if key in ['NetAmount','NetAmountTWD']:
            cell.value = float(Decimal(str(r[key])))
            cell.number_format = '0.00'
        elif key == 'OrderDate':
            cell.value = r[key]
            cell.number_format = 'yyyy-mm-dd'
        else:
            val = r[key]
            cell.value = val if val != '' else None
            cell.number_format = '@'
# Auto-filter include data
last = 4 + len(details)
det.auto_filter.ref = f'A4:K{last}'
# Summary B2 formula
sumsh['B2'] = f'=SUM(Details!I5:I{last})'
# Summary by status/currency (all Open), grouped by currency
summary = {}
for r in details:
    key = r['Currency']
    d = summary.setdefault(key, {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
    d['amount'] += Decimal(r['NetAmount'])
    d['twd'] += r['NetAmountTWD']
    d['count'] += 1
for idx, cur in enumerate(sorted(summary.keys()), start=5):
    d = summary[cur]
    vals = ['Status', 'Open', cur, float(d['amount'].quantize(q, rounding=ROUND_HALF_UP)), float(d['twd'].quantize(q, rounding=ROUND_HALF_UP)), d['count']]
    for c, val in enumerate(vals, start=1):
        cell = sumsh.cell(idx, c)
        style_from(sumsh.cell(4, c), cell)
        cell.value = val
        if c in [4,5]: cell.number_format = '0.00'
        elif c == 6: cell.number_format = '0'
sumsh.auto_filter.ref = f'A4:F{4+len(summary)}'
# Save
wb.save(output)
print('wrote', output, output.stat().st_size)
print('summary:', {k:{'amount':str(v['amount']),'twd':str(v['twd']),'count':v['count']} for k,v in summary.items()})