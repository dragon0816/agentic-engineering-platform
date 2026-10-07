from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
import csv
ws=Path.cwd()
out=ws/'output'/'sales_2026-09_open_orders.xlsx'
print('checking', out, 'exists', out.exists(), 'bytes', out.stat().st_size if out.exists() else None)
wb=load_workbook(out, data_only=False)
print('sheets', wb.sheetnames)
assert wb.sheetnames == ['Report','Details','Summary','Instructions']
rep=wb['Report']; det=wb['Details']; summ=wb['Summary']
# report cells
checks = {
 'Report!B3': rep['B3'].value,
 'Report!D3': rep['D3'].value,
 'Report!B4': rep['B4'].value,
 'Report!B5': rep['B5'].value,
 'Report!B7': rep['B7'].value,
 'Report!D7': rep['D7'].value,
 'Report!B9': rep['B9'].value,
}
print('report cells:', checks)
assert rep['B3'].value == '2026-09-01'
assert rep['D3'].value == '2026-09-30'
assert 'Status=Open' in rep['B4'].value
# Details headers and rows
expected_headers=['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
headers=[det.cell(4,c).value for c in range(1,12)]
print('detail headers:', headers)
assert headers == expected_headers
# determine data rows
rows=[]
for r in range(5, det.max_row+1):
    vals=[det.cell(r,c).value for c in range(1,12)]
    if any(v is not None for v in vals): rows.append(vals)
print('detail data row count:', len(rows), 'max_row', det.max_row)
assert len(rows)==14
ids=[r[0] for r in rows]
print('ids:', ids)
assert len(ids)==len(set(ids))==14
expected_ids=[]
with open(ws/'output'/'downloads'/'orders_2026-09_open.csv', newline='', encoding='utf-8-sig') as f:
    csv_rows=list(csv.DictReader(f)); expected_ids=[r['OrderID'] for r in csv_rows]
assert ids==expected_ids
# Value type / business rules
rates={'TWD':Decimal('1'),'USD':Decimal('32'),'EUR':Decimal('35')}
missing_contacts=0
for vals in rows:
    oid,cid,cname,pid,od,status,cur,amt,twd,contact,terms=vals
    assert isinstance(oid,str) and oid.startswith('00')
    assert isinstance(cid,str)
    assert pid is None or isinstance(pid,str)
    assert od >= '2026-09-01' and od <= '2026-09-30'
    assert status == 'Open'
    assert cur in rates
    assert isinstance(amt,(int,float))
    assert isinstance(twd,(int,float))
    expected=(Decimal(str(amt))*rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert Decimal(str(twd)).quantize(Decimal('0.01')) == expected, (oid,twd,expected)
    if contact is None: missing_contacts += 1
    assert terms is not None and str(terms).startswith('NET')
print('missing contact blank cells:', missing_contacts)
assert missing_contacts == 2
print('Details auto_filter:', det.auto_filter.ref)
assert det.auto_filter.ref == 'A4:K18'
# Summary formula and rows
print('Summary!B2:', summ['B2'].value)
assert summ['B2'].value == '=SUM(Details!I5:I18)'
srows=[]
for r in range(5, summ.max_row+1):
    vals=[summ.cell(r,c).value for c in range(1,7)]
    if any(v is not None for v in vals): srows.append(vals)
print('summary rows:', srows)
assert len(srows)==3
counts=sum(int(r[5]) for r in srows)
assert counts==14
sum_twd=sum(Decimal(str(r[4])).quantize(Decimal('0.01')) for r in srows)
detail_twd=sum(Decimal(str(r[8])).quantize(Decimal('0.01')) for r in rows)
print('detail_twd_sum', detail_twd, 'summary_twd_sum', sum_twd)
assert sum_twd == detail_twd
# Verify source workbook not modified (load source headers/row count)
src=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
assert src['Details'].max_row == 4
assert [src['Details'].cell(4,c).value for c in range(1,12)] == expected_headers
print('All checks passed.')