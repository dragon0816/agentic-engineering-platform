from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
# print full instructions sheet and manual remaining
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
sh=wb['Instructions']
print('Instructions sheet:')
for row in sh.iter_rows(values_only=True): print(row)
print('\nManual full:')
print((ws/'input'/'sap_manual.md').read_text(encoding='utf-8'))