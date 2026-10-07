from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
# workbook structure
xlsx = ws/'input'/'sales_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'[{s}] size rows={sh.max_row} cols={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(r)
    print('---')
# fx sample
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader): print(row)
# manual first 200 lines
print('\nsap_manual.md first lines:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:200],1):
    print(f'{i:03d}: {line}')