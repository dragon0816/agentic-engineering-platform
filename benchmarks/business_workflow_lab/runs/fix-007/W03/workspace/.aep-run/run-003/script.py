from pathlib import Path
import csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
DOWNLOADS = OUTPUT / 'downloads'
OUTPUT.mkdir(exist_ok=True)

TEMPLATE = INPUT / 'project_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
INVOICE_CSV = DOWNLOADS / 'invoices_2026-09.csv'
OUT_XLSX = OUTPUT / 'project_invoices_2026-09_asof_2026-09-30.xlsx'

FROM_DATE = date(2026, 9, 1)
TO_DATE = date(2026, 9, 30)
AS_OF = date(2026, 9, 30)
GENERATED_AT = '2026-10-06T00:00:00+08:00'

print('Inputs:')
for p in [TEMPLATE, FX_FILE, INVOICE_CSV]:
    print(f'- {p}: exists={p.exists()} size={p.stat().st_size if p.exists() else "missing"}')

# Read FX rates. Notes may contain commas in this lab file; use first two fields only.
rates = {}
with open(FX_FILE, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header:', header)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

# Read invoices CSV.
with open(INVOICE_CSV, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    invoice_headers = reader.fieldnames
    rows = list(reader)
print('Invoice CSV headers:', invoice_headers)
print('Invoice row count:', len(rows))
print('Invoice sample rows:', rows[:3])

required_headers = ['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance']
if invoice_headers != required_headers:
    raise ValueError(f'Unexpected invoice headers: {invoice_headers}')
if not rows:
    raise ValueError('No invoice rows found')

TWOPLACES = Decimal('0.01')
def money(v):
    return Decimal(str(v)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def parse_date(v):
    return datetime.strptime(v, '%Y-%m-%d').date()
def to_float(d):
    return float(d)

def clear_data(sh, min_row=5):
    if sh.max_row >= min_row:
        sh.delete_rows(min_row, sh.max_row - min_row + 1)

def write_cell(cell, value, fmt=None, text=False):
    cell.value = value
    if fmt:
        cell.number_format = fmt
    if text:
        cell.number_format = '@'

# Validate query period and enrich rows.
invoices = []
for r in rows:
    inv_date = parse_date(r['InvoiceDate'])
    due_date = parse_date(r['DueDate'])
    if not (FROM_DATE <= inv_date <= TO_DATE):
        raise ValueError(f'Invoice outside September 2026: {r}')
    cur = r['Currency']
    if cur not in rates:
        raise ValueError(f'Missing FX rate for {cur}')
    amount = money(r['Amount'])
    paid = money(r['PaidAmount'])
    balance = money(r['Balance'])
    calc_balance = money(amount - paid)
    if balance != calc_balance:
        raise ValueError(f"Balance mismatch for {r['InvoiceID']}: csv={balance}, amount-paid={calc_balance}")
    balance_twd = money(balance * rates[cur])
    enriched = {
        **r,
        'InvoiceDate_obj': inv_date,
        'DueDate_obj': due_date,
        'Amount_dec': amount,
        'PaidAmount_dec': paid,
        'Balance_dec': balance,
        'BalanceTWD_dec': balance_twd,
        'Overdue': balance > 0 and due_date < AS_OF,
    }
    invoices.append(enriched)

# Copy and fill template.
shutil.copyfile(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

# Keep source sheets other than Invoices blank.
for sheet_name in ['Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions']:
    clear_data(wb[sheet_name], 5)

# Report metadata, preserving layout.
report = wb['Report']
report['B3'] = FROM_DATE
report['B3'].number_format = 'yyyy-mm-dd'
report['D3'] = TO_DATE
report['D3'].number_format = 'yyyy-mm-dd'
report['B4'] = 'Invoices only; InvoiceDate 2026-09-01 through 2026-09-30; unpaid balance as of 2026-09-30; overdue if Balance > 0 and DueDate < 2026-09-30.'
report['B5'] = GENERATED_AT
report['B7'] = 'Lab ERP Invoices export; fx_rates.csv'
report['D7'] = GENERATED_AT
report['B9'] = 'Orders, Expenses, and Quotes were intentionally left blank per request. Summary excludes invoices with blank ProjectID because it is grouped by project.'

# Fill Invoices sheet. Target headers: InvoiceID, OrderID, ProjectID, CustomerID, InvoiceDate, DueDate, Currency, Amount, PaidAmount, Balance, BalanceTWD
inv_sh = wb['Invoices']
for idx, r in enumerate(invoices, start=5):
    values = [
        r['InvoiceID'], r['OrderID'], r['ProjectID'] or None, r['CustomerID'],
        r['InvoiceDate_obj'], r['DueDate_obj'], r['Currency'],
        to_float(r['Amount_dec']), to_float(r['PaidAmount_dec']), to_float(r['Balance_dec']), to_float(r['BalanceTWD_dec'])
    ]
    for col, v in enumerate(values, start=1):
        c = inv_sh.cell(idx, col, v)
        if col in [1,2,3,4,7]:
            c.number_format = '@'
        elif col in [5,6]:
            c.number_format = 'yyyy-mm-dd'
        elif col in [8,9,10,11]:
            c.number_format = '0.00'

# Summary by nonblank ProjectID only; source columns other than InvoiceBalanceTWD remain 0.00 because no source data requested.
summary = {}
for r in invoices:
    project = (r['ProjectID'] or '').strip()
    if not project:
        continue
    summary.setdefault(project, Decimal('0.00'))
    summary[project] += r['BalanceTWD_dec']

sum_sh = wb['Summary']
for row_idx, project in enumerate(sorted(summary), start=5):
    sum_sh.cell(row_idx, 1, project).number_format = '@'
    vals = [Decimal('0.00'), summary[project].quantize(TWOPLACES), Decimal('0.00'), Decimal('0.00')]
    for j, d in enumerate(vals, start=2):
        cell = sum_sh.cell(row_idx, j, to_float(d))
        cell.number_format = '0.00'
last = max(5, 4 + len(summary))
for col in range(2,6):
    letter = sum_sh.cell(4, col).column_letter
    sum_sh.cell(2, col).value = f'=SUM({letter}5:{letter}{last})'
    sum_sh.cell(2, col).number_format = '0.00'

# Exceptions for overdue invoices.
exc_sh = wb['Exceptions']
overdue_fill = PatternFill(fill_type='solid', fgColor='FFF2CC')
overdue_rows = [r for r in invoices if r['Overdue']]
for row_idx, r in enumerate(overdue_rows, start=5):
    vals = ['Invoices', r['InvoiceID'], r['ProjectID'] or None, 'Overdue']
    for col, v in enumerate(vals, start=1):
        c = exc_sh.cell(row_idx, col, v)
        if col in [1,2,3,4]:
            c.number_format = '@'
        c.fill = overdue_fill

wb.save(OUT_XLSX)
print('Wrote:', OUT_XLSX)

# Reopen and verify output requirements.
wb2 = load_workbook(OUT_XLSX, data_only=False)
checks = []

def assert_equal(a,b,msg):
    if a != b:
        raise AssertionError(f'{msg}: {a!r} != {b!r}')
    checks.append(msg)

def assert_true(cond,msg):
    if not cond:
        raise AssertionError(msg)
    checks.append(msg)

# headers preserved.
assert_equal([wb2['Invoices'].cell(4,c).value for c in range(1,12)], ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], 'Invoices headers preserved')
assert_equal([wb2['Summary'].cell(4,c).value for c in range(1,6)], ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], 'Summary headers preserved')
assert_equal([wb2['Exceptions'].cell(4,c).value for c in range(1,5)], ['Source','DocumentID','ProjectID','Reason'], 'Exceptions headers preserved')

# Row counts.
inv_count = wb2['Invoices'].max_row - 4
summary_count = wb2['Summary'].max_row - 4
exc_count = wb2['Exceptions'].max_row - 4
assert_equal(inv_count, len(invoices), 'All 13 September invoices written')
assert_equal(summary_count, len(summary), 'Summary has one row per nonblank project')
assert_equal(exc_count, len(overdue_rows), 'Exceptions has one row per overdue invoice')
for s in ['Orders','Expenses','Quotes']:
    assert_equal(wb2[s].max_row, 4, f'{s} source sheet remains blank below headers')

# IDs text/date/number and date range.
for row_idx in range(5, 5 + len(invoices)):
    inv_id = wb2['Invoices'].cell(row_idx,1).value
    order_id = wb2['Invoices'].cell(row_idx,2).value
    inv_date = wb2['Invoices'].cell(row_idx,5).value
    due_date = wb2['Invoices'].cell(row_idx,6).value
    for col in [1,2,3,4,7]:
        assert_equal(wb2['Invoices'].cell(row_idx,col).number_format, '@', f'Invoices row {row_idx} col {col} text format')
    assert_true(isinstance(inv_id, str) and inv_id.startswith('009'), f'InvoiceID row {row_idx} is text with leading zeros')
    assert_true(isinstance(order_id, str) and order_id.startswith('004'), f'OrderID row {row_idx} is text with leading zeros')
    assert_true(FROM_DATE <= inv_date.date() <= TO_DATE if hasattr(inv_date, 'date') else FROM_DATE <= inv_date <= TO_DATE, f'InvoiceDate row {row_idx} in requested period')
    for col in [8,9,10,11]:
        assert_true(isinstance(wb2['Invoices'].cell(row_idx,col).value, (int,float)), f'Invoices row {row_idx} col {col} numeric')

# BalanceTWD calculations and summary totals.
by_inv = {r['InvoiceID']: r for r in invoices}
for row_idx in range(5, 5 + len(invoices)):
    inv_id = wb2['Invoices'].cell(row_idx,1).value
    expected = float(by_inv[inv_id]['BalanceTWD_dec'])
    assert_equal(round(wb2['Invoices'].cell(row_idx,11).value,2), expected, f'BalanceTWD for {inv_id}')
for row_idx in range(5, 5 + len(summary)):
    project = wb2['Summary'].cell(row_idx,1).value
    assert_equal(round(wb2['Summary'].cell(row_idx,3).value,2), float(summary[project]), f'Summary InvoiceBalanceTWD for {project}')
    assert_equal(round(wb2['Summary'].cell(row_idx,2).value,2), 0.0, f'Summary OrderTWD blank-source zero for {project}')
    assert_equal(round(wb2['Summary'].cell(row_idx,4).value,2), 0.0, f'Summary ApprovedExpenseTWD blank-source zero for {project}')
    assert_equal(round(wb2['Summary'].cell(row_idx,5).value,2), 0.0, f'Summary AcceptedQuoteTWD blank-source zero for {project}')

# Overdue rule: Balance > 0 and DueDate < 2026-09-30; due on 2026-09-30 is not overdue.
actual_exc = []
for row_idx in range(5, 5 + len(overdue_rows)):
    actual_exc.append((wb2['Exceptions'].cell(row_idx,1).value, wb2['Exceptions'].cell(row_idx,2).value, wb2['Exceptions'].cell(row_idx,3).value, wb2['Exceptions'].cell(row_idx,4).value))
expected_exc = [('Invoices', r['InvoiceID'], r['ProjectID'] or None, 'Overdue') for r in overdue_rows]
assert_equal(actual_exc, expected_exc, 'Exceptions match overdue rule only')
assert_true('0090000211' not in [x[1] for x in actual_exc], 'Invoice due on 2026-09-30 is not marked overdue')
assert_true('0090000213' not in [x[1] for x in actual_exc], 'Partly unpaid invoice due on 2026-09-30 is not marked overdue')

print('Checks passed:')
for c in checks:
    print('-', c)
print('Summary totals:')
for project in sorted(summary):
    print(project, summary[project])
print('Overdue invoice IDs:', [r['InvoiceID'] for r in overdue_rows])
print('Output file:', OUT_XLSX)