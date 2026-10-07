from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# workbook sheets/headers/sample
xlsx = ws/'input'/'project_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    rows = []
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
        rows.append(r)
    print('sample', rows)
# fx headers/sample
fx = ws/'input'/'fx_rates.csv'
with open(fx, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(5), reader):
        print('fx row', i+1, row)
# manual beginning
md = ws/'input'/'sap_manual.md'
print('sap_manual first 2000 chars:')
print(md.read_text(encoding='utf-8')[:2000])