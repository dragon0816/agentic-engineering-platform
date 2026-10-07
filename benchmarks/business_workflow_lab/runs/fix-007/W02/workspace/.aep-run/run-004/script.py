from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, shutil, copy
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

base = Path.cwd()
in_dir = base/'input'
out_dir = base/'output'
out_dir.mkdir(exist_ok=True)
template = in_dir/'sales_template.xlsx'
out_file = out_dir/'sales_C001_2026-08-01_2026-10-31.xlsx'

print('inputs:')
for p in sorted(in_dir.iterdir()):
    print(p.name, p.stat().st_size)

export_csv = out_dir/'downloads'/'c001_orders_2026-08-01_2026-10-31.csv'
print('export exists', export_csv, export_csv.exists())
with open(export_csv, newline='', encoding='utf-8-sig') as f:
    export_rows = list(csv.DictReader(f))
print('export headers', list(export_rows[0].keys()) if export_rows else [])
print('export row count', len(export_rows))
print('export sample', export_rows[:2])

detail_by_order = {
    '0045000101': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET15'},
    '0045000107': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000109': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10'},
    '0045000114': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET5'},
    '0045000119': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000125': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET0'},
    '0045000129': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000132': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': ''},
    '0045000136': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET5'},
}

rates = {}
with open(in_dir/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        if len(row) >= 2:
            rates[row[0]] = Decimal(row[1])
print('fx rates', rates)

expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']
if export_rows and list(export_rows[0].keys()) != expected_headers:
    raise AssertionError(f'Unexpected export headers: {list(export_rows[0].keys())}')

from_date = '2026-08-01'
to_date = '2026-10-31'
rows = []
for r in export_rows:
    if r['CustomerID'] != 'C001':
        print('excluded non-exact customer', r)
        continue
    if not (from_date <= r['OrderDate'] <= to_date):
        print('excluded out of date', r)
        continue
    oid = r['OrderID']
    d = detail_by_order.get(oid, {})
    net = Decimal(r['NetAmount'])
    twd = (net * rates[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    rows.append({**r, 'NetAmount': net, 'NetAmountTWD': twd,
                 'CustomerContact': d.get('CustomerContact',''), 'PaymentTerms': d.get('PaymentTerms','')})
rows.sort(key=lambda x: x['OrderID'])
print('filtered rows', len(rows))
for r in rows:
    print(r)

missing = [f"{r['OrderID']} {field}" for r in rows for field in ['CustomerContact','PaymentTerms'] if r.get(field,'') in (None,'')]

shutil.copyfile(template, out_file)
wb = load_workbook(out_file)
report, details, summary = wb['Report'], wb['Details'], wb['Summary']

# Clear rows in the copied file.
for ws, min_row, max_col in [(details, 5, 11), (summary, 5, 6)]:
    if ws.max_row >= min_row:
        for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

taipei = timezone(timedelta(hours=8))
now_iso = datetime.now(taipei).replace(microsecond=0).isoformat()
report['B3'] = from_date
report['D3'] = to_date
report['B4'] = 'Sales Orders: CustomerID = C001 (exact match), OrderDate 2026-08-01 to 2026-10-31 inclusive; all statuses'
report['B5'] = now_iso
report['B7'] = 'Lab ERP Sales Orders; sales order detail pages; fx_rates.csv'
report['D7'] = now_iso
report['B9'] = ('Missing source fields left blank: ' + '; '.join(missing) + '. Similar customer names excluded by exact CustomerID filter.') if missing else 'No missing CustomerContact or PaymentTerms. Similar customer names excluded by exact CustomerID filter.'

headers = [details.cell(4, c).value for c in range(1, 12)]
print('details template headers', headers)
style_src = [details.cell(4, c) for c in range(1, 12)]
for idx, r in enumerate(rows, start=5):
    values = [r['OrderID'], r['CustomerID'], r['CustomerName'], r['ProjectID'], r['OrderDate'], r['Status'], r['Currency'],
              float(r['NetAmount']), float(r['NetAmountTWD']), r['CustomerContact'], r['PaymentTerms']]
    for c, val in enumerate(values, start=1):
        cell = details.cell(idx, c)
        cell.value = val
        src = style_src[c-1]
        if src.has_style:
            cell._style = copy.copy(src._style)
        cell.font = copy.copy(src.font)
        cell.fill = copy.copy(src.fill)
        cell.border = copy.copy(src.border)
        cell.alignment = copy.copy(src.alignment)
    for c in [1,2,4]:
        details.cell(idx,c).number_format = '@'
    details.cell(idx,5).number_format = 'yyyy-mm-dd'
    for c in [8,9]:
        details.cell(idx,c).number_format = '0.00'

formula_last = 4 + len(rows) if rows else 5
summary['B2'] = f'=SUM(Details!I5:I{formula_last})'

groups = {}
for r in rows:
    key = (r['Status'], r['Currency'])
    groups.setdefault(key, {'Amount': Decimal('0'), 'AmountTWD': Decimal('0'), 'Count': 0})
    groups[key]['Amount'] += r['NetAmount']
    groups[key]['AmountTWD'] += r['NetAmountTWD']
    groups[key]['Count'] += 1
summary_style = [summary.cell(4, c) for c in range(1,7)]
rowno = 5
for (status, currency), agg in sorted(groups.items()):
    values = ['Status', status, currency, float(agg['Amount'].quantize(Decimal('0.01'))), float(agg['AmountTWD'].quantize(Decimal('0.01'))), agg['Count']]
    for c, val in enumerate(values, start=1):
        cell = summary.cell(rowno, c)
        cell.value = val
        src = summary_style[c-1]
        if src.has_style:
            cell._style = copy.copy(src._style)
        cell.font = copy.copy(src.font); cell.fill = copy.copy(src.fill); cell.border = copy.copy(src.border); cell.alignment = copy.copy(src.alignment)
    summary.cell(rowno,4).number_format = '0.00'
    summary.cell(rowno,5).number_format = '0.00'
    rowno += 1

wb.save(out_file)
print('wrote', out_file)
print('missing', missing)
