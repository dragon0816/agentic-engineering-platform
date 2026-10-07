from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
import csv

base=Path.cwd()
out_file=base/'output'/'sales_C001_2026-08-01_2026-10-31.xlsx'
print('checking', out_file, out_file.exists(), out_file.stat().st_size if out_file.exists() else None)
wb = load_workbook(out_file, data_only=False)
print('sheets', wb.sheetnames)
assert wb.sheetnames == ['Report','Details','Summary','Instructions']
assert str(wb['Report']['B3'].value) == '2026-08-01'
assert str(wb['Report']['D3'].value) == '2026-10-31'
assert 'CustomerID = C001' in wb['Report']['B4'].value
assert '0045000132 PaymentTerms' in wb['Report']['B9'].value
assert 'Similar customer names excluded' in wb['Report']['B9'].value

ws=wb['Details']
headers=[ws.cell(4,c).value for c in range(1,12)]
expected=['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
print('headers', headers)
assert headers == expected
rows=[]
for r in range(5, ws.max_row+1):
    vals=[ws.cell(r,c).value for c in range(1,12)]
    if any(v is not None for v in vals):
        rows.append(vals)
print('data rows', len(rows))
for row in rows:
    print(row)
assert len(rows)==9
expected_orderids=['0045000101','0045000107','0045000109','0045000114','0045000119','0045000125','0045000129','0045000132','0045000136']
assert [r[0] for r in rows] == expected_orderids
for i,row in enumerate(rows, start=5):
    assert row[1] == 'C001', row
    assert row[2] == 'Acme Precision Taiwan Ltd.', row
    assert 'Limited' not in row[2]
    assert '2026-08-01' <= str(row[4]) <= '2026-10-31'
    assert isinstance(row[0], str) and row[0].startswith('00')
    assert ws.cell(i,1).number_format == '@'
    assert ws.cell(i,2).number_format == '@'
    assert ws.cell(i,4).number_format == '@'
    assert isinstance(row[7], (int,float))
    assert isinstance(row[8], (int,float))
    assert ws.cell(i,8).number_format == '0.00'
    assert ws.cell(i,9).number_format == '0.00'
    assert row[9] == 'Grace Lin <grace.lin@acme-precision.example>'

# missing value check: only 0045000132 PaymentTerms blank/None.
missing=[]
for row in rows:
    if row[9] in (None,''): missing.append((row[0],'CustomerContact'))
    if row[10] in (None,''): missing.append((row[0],'PaymentTerms'))
print('missing', missing)
assert missing == [('0045000132','PaymentTerms')]

rates={'TWD':Decimal('1'),'USD':Decimal('32'),'EUR':Decimal('35')}
for row in rows:
    expected_twd=(Decimal(str(row[7]))*rates[row[6]]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert Decimal(str(row[8])).quantize(Decimal('0.01')) == expected_twd, (row, expected_twd)

summary=wb['Summary']
print('Summary B2', summary['B2'].value)
assert summary['B2'].value == '=SUM(Details!I5:I13)'
summary_rows=[]
for r in range(5, summary.max_row+1):
    vals=[summary.cell(r,c).value for c in range(1,7)]
    if any(v is not None for v in vals): summary_rows.append(vals)
print('summary rows', summary_rows)
assert summary_rows
# confirm total of summary TWD equals details TWD
details_total=sum(Decimal(str(row[8])) for row in rows)
summary_total=sum(Decimal(str(row[4])) for row in summary_rows)
print('details_total_twd', details_total, 'summary_total_twd', summary_total)
assert summary_total.quantize(Decimal('0.01')) == details_total.quantize(Decimal('0.01'))
print('ALL CHECKS PASSED')