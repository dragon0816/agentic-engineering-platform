from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook(Path('input')/'sales_template.xlsx', data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)
print('\nReport merged:', list(wb['Report'].merged_cells.ranges))
print('Details merged:', list(wb['Details'].merged_cells.ranges))
print('Summary merged:', list(wb['Summary'].merged_cells.ranges))