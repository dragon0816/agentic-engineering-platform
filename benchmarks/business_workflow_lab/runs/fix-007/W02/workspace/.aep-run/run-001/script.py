import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('cwd', ws)
for root, dirs, files in os.walk(ws/'input'):
    print('DIR', root)
    for f in files:
        p = Path(root)/f
        print(' FILE', f, 'size', p.stat().st_size)

# workbook structure
p = ws/'input'/'sales_template.xlsx'
wb = load_workbook(p, data_only=False)
print('workbook sheets', wb.sheetnames)
for sh in wb.sheetnames:
    wsht = wb[sh]
    print('SHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
        print(' ', r)

# fx headers/sample
p = ws/'input'/'fx_rates.csv'
with open(p, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(6), reader):
        print('fx row', i, row)

# manual head
p = ws/'input'/'sap_manual.md'
text = p.read_text(encoding='utf-8')
print('manual first 4000 chars:')
print(text[:4000])