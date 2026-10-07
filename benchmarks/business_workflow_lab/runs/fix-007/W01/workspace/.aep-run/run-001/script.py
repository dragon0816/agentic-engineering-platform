import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:', sorted([p.name for p in (ws/'input').iterdir()]))
# Inspect workbook
wb_path=ws/'input'/'sales_template.xlsx'
wb=load_workbook(wb_path)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(r)
# FX
fx_path=ws/'input'/'fx_rates.csv'
print('fx_rates.csv:')
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(10), reader):
        print(row)
# Manual excerpt
md_path=ws/'input'/'sap_manual.md'
print('sap_manual.md first 200 lines:')
text=md_path.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:200],1):
    print(f'{i}: {line}')