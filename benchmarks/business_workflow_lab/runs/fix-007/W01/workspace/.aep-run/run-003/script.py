from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx')
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True): print(row)