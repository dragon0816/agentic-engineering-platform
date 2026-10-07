from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print((ws/'input'/'quote_cli_manual.md').read_text(encoding='utf-8'))
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for s in wb.sheetnames:
    sh=wb[s]
    print('\n===', s, '===')
    print('dimensions', sh.max_row, sh.max_column)
    for r in range(1, sh.max_row+1):
        vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
        print(r, vals)
    print('merged', list(sh.merged_cells.ranges))