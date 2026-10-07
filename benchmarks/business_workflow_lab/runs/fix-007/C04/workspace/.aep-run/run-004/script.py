from pathlib import Path
import subprocess, sys, json, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import date
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI_PY = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
OUT_XLSX = OUTPUT / 'quote_report_2026-09_Sent.xlsx'
AUDIT_JSON = OUTPUT / 'quote_report_2026-09_Sent_audit.json'
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
SOURCE_NAME = 'quote_cli synthetic quote system; fx_rates.csv'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
FETCHED_AT = GENERATED_AT

print('Inputs confirmed:')
print(' template:', TEMPLATE, TEMPLATE.exists())
print(' fx rates:', FX_FILE, FX_FILE.exists())
print(' cli:', CLI_PY, CLI_PY.exists())
print(' output:', OUT_XLSX)

if not TEMPLATE.exists():
    raise FileNotFoundError(TEMPLATE)
if not FX_FILE.exists():
    raise FileNotFoundError(FX_FILE)
if not CLI_PY.exists():
    raise FileNotFoundError(CLI_PY)

fx = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx[row[0]] = Decimal(row[1])
print('fx currencies:', sorted(fx))

attempt_log = []

def run_cli_json(args, label):
    base = [sys.executable, str(CLI_PY)]
    max_retries = 2
    for attempt in range(1, max_retries + 2):
        cmd = base + args
        cp = subprocess.run(cmd, text=True, capture_output=True)
        record = {
            'label': label,
            'attempt': attempt,
            'exit_code': cp.returncode,
            'stderr': cp.stderr.strip(),
            'stdout_preview': cp.stdout[:200],
        }
        attempt_log.append(record)
        print(f"CLI {label} attempt {attempt}: exit={cp.returncode}; stderr={cp.stderr.strip()!r}")
        if cp.returncode == 0:
            try:
                return json.loads(cp.stdout)
            except json.JSONDecodeError as e:
                raise RuntimeError(f"{label}: stdout had exit code 0 but was not parseable JSON: {e}; preview={cp.stdout[:500]!r}")
        if cp.returncode in (2, 3):
            raise RuntimeError(f"{label}: quote CLI returned exit code {cp.returncode}; stderr: {cp.stderr.strip()}")
        if cp.returncode == 4:
            if attempt <= max_retries:
                continue
            raise RuntimeError(f"{label}: quote CLI temporary service error persisted after {max_retries} retries; stderr: {cp.stderr.strip()}")
        raise RuntimeError(f"{label}: unexpected quote CLI exit code {cp.returncode}; stderr: {cp.stderr.strip()}")

list_data = run_cli_json(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'], 'quotes list Sep 2026 Sent')
if list_data.get('schema_version') != 1 or not isinstance(list_data.get('items'), list):
    raise RuntimeError(f"quotes list returned unexpected schema: {list_data}")
headers = list_data['items']
print('quotes list count:', list_data.get('count'), 'items:', len(headers))
if list_data.get('count') != len(headers):
    raise RuntimeError(f"quotes list count mismatch: count={list_data.get('count')} items={len(headers)}")

quote_details = []
for h in headers:
    qid = h.get('QuoteID')
    if not qid:
        raise RuntimeError(f'Missing QuoteID in header item: {h}')
    detail = run_cli_json(['quotes', 'get', qid, '--format', 'json'], f'quotes get {qid}')
    if detail.get('schema_version') != 1 or not isinstance(detail.get('header'), dict) or not isinstance(detail.get('line_items'), list):
        raise RuntimeError(f"quotes get {qid} returned unexpected schema: {detail}")
    quote_details.append(detail)

from_dt = date.fromisoformat(FROM_DATE)
to_dt = date.fromisoformat(TO_DATE)
headers_by_id = {h['QuoteID']: h for h in headers}
for detail in quote_details:
    h = detail['header']
    qid = h.get('QuoteID')
    if qid not in headers_by_id:
        raise RuntimeError(f"Detail quote {qid} not in list results")
    qdate = date.fromisoformat(h['QuoteDate'])
    if h.get('Status') != STATUS or not (from_dt <= qdate <= to_dt):
        raise RuntimeError(f"Quote {qid} did not match requested filter: Status={h.get('Status')} QuoteDate={h.get('QuoteDate')}")
    for key in ['CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount']:
        if str(h.get(key)) != str(headers_by_id[qid].get(key)):
            raise RuntimeError(f"Header mismatch for {qid} field {key}: list={headers_by_id[qid].get(key)!r} detail={h.get(key)!r}")

shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)
expected_sheets = ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
if wb.sheetnames != expected_sheets:
    raise RuntimeError(f'Template sheets changed/unexpected: {wb.sheetnames}')
report = wb['Report']
quotes_ws = wb['Quotes']
lines_ws = wb['Lines']
summary_ws = wb['Summary']

def copy_row_style(ws, src_row, dst_row, max_col):
    for c in range(1, max_col + 1):
        src = ws.cell(src_row, c)
        dst = ws.cell(dst_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)

report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = f'QuoteDate {FROM_DATE} to {TO_DATE} inclusive; Status={STATUS}'
report['B5'] = GENERATED_AT
report['B7'] = SOURCE_NAME
report['D7'] = FETCHED_AT
report['B9'] = f'CLI retries applied per instruction: exit code 4 retried up to 2 times; exit code 2/3 not retried. Retrieved {len(headers)} quote headers and {sum(len(d["line_items"]) for d in quote_details)} line items.'

for ws, max_col in [(quotes_ws, 11), (lines_ws, 6), (summary_ws, 6)]:
    if ws.max_row > 4:
        ws.delete_rows(5, ws.max_row - 4)

money_fmt = '0.00'
date_fmt = 'yyyy-mm-dd'
text_fmt = '@'
quote_rows = []
for detail in sorted(quote_details, key=lambda d: (d['header']['QuoteDate'], d['header']['QuoteID'])):
    h = detail['header']
    cur = h['Currency']
    if cur not in fx:
        raise RuntimeError(f'Missing FX rate for {cur}')
    net = Decimal(str(h['NetAmount']))
    net_twd = (net * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    quote_rows.append((h, net_twd))

for idx, (h, net_twd) in enumerate(quote_rows, start=5):
    copy_row_style(quotes_ws, 4, idx, 11)
    values = [
        str(h['QuoteID']), str(h['CustomerID']), str(h['ProjectID']),
        date.fromisoformat(h['QuoteDate']), date.fromisoformat(h['ValidUntil']),
        str(h['Status']), str(h['Currency']),
        float(Decimal(str(h['Subtotal']))), float(Decimal(str(h['DiscountPercent']))),
        float(Decimal(str(h['NetAmount']))), float(net_twd),
    ]
    for col, val in enumerate(values, start=1):
        cell = quotes_ws.cell(idx, col, val)
        if col in (1,2,3,6,7):
            cell.number_format = text_fmt
        elif col in (4,5):
            cell.number_format = date_fmt
        elif col in (8,10,11):
            cell.number_format = money_fmt
        elif col == 9:
            cell.number_format = '0.##'

line_count = 0
for detail in sorted(quote_details, key=lambda d: (d['header']['QuoteDate'], d['header']['QuoteID'])):
    qid = str(detail['header']['QuoteID'])
    for li in detail['line_items']:
        line_count += 1
        r = 4 + line_count
        copy_row_style(lines_ws, 4, r, 6)
        quantity_raw = li['Quantity']
        try:
            quantity = int(quantity_raw)
        except Exception:
            quantity = float(Decimal(str(quantity_raw)))
        vals = [qid, str(li.get('ProductID','')), li.get('ProductName',''), quantity,
                float(Decimal(str(li['UnitPrice']))), float(Decimal(str(li['LineAmount'])))]
        for col, val in enumerate(vals, start=1):
            cell = lines_ws.cell(r, col, val)
            if col in (1,2):
                cell.number_format = text_fmt
            elif col in (5,6):
                cell.number_format = money_fmt
            elif col == 4:
                cell.number_format = '0'

summary = {}
for h, net_twd in quote_rows:
    key = (h['Status'], h['Currency'])
    amount = Decimal(str(h['NetAmount']))
    if key not in summary:
        summary[key] = {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0}
    summary[key]['amount'] += amount
    summary[key]['amount_twd'] += net_twd
    summary[key]['count'] += 1

for i, ((status, currency), vals) in enumerate(sorted(summary.items()), start=5):
    copy_row_style(summary_ws, 4, i, 6)
    row_vals = ['Status', status, currency,
                float(vals['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)),
                float(vals['amount_twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)),
                vals['count']]
    for col, val in enumerate(row_vals, start=1):
        cell = summary_ws.cell(i, col, val)
        if col in (1,2,3):
            cell.number_format = text_fmt
        elif col in (4,5):
            cell.number_format = money_fmt
        elif col == 6:
            cell.number_format = '0'

last_quote_row = 4 + len(quote_rows)
if len(quote_rows) == 0:
    summary_ws['B2'] = '=SUM(Quotes!K5:K5)'
else:
    summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
summary_ws['B2'].number_format = money_fmt

wb.save(OUT_XLSX)
AUDIT_JSON.write_text(json.dumps({
    'from': FROM_DATE,
    'to': TO_DATE,
    'status': STATUS,
    'quote_count': len(quote_rows),
    'line_count': line_count,
    'attempts': attempt_log,
}, ensure_ascii=False, indent=2), encoding='utf-8')

check_wb = load_workbook(OUT_XLSX, data_only=False)
checks = []
checks.append(('sheets unchanged', check_wb.sheetnames == expected_sheets))
checks.append(('Report B3/D3 period', check_wb['Report']['B3'].value == FROM_DATE and check_wb['Report']['D3'].value == TO_DATE))
checks.append(('Report condition mentions Sent', STATUS in str(check_wb['Report']['B4'].value)))
checks.append(('Report source/fetched populated', bool(check_wb['Report']['B7'].value) and bool(check_wb['Report']['D7'].value)))
checks.append(('Report footer kept', check_wb['Report']['A20'].value == '合成測試資料 Synthetic lab data · template v1 · 請勿修改本行'))
checks.append(('Report merge kept', 'A1:F1' in [str(r) for r in check_wb['Report'].merged_cells.ranges]))
checks.append(('Quotes headers kept', [check_wb['Quotes'].cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']))
checks.append(('Lines headers kept', [check_wb['Lines'].cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']))
checks.append(('Summary headers kept', [check_wb['Summary'].cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']))
qws = check_wb['Quotes']
lws = check_wb['Lines']
sws = check_wb['Summary']
actual_quote_rows = max(0, qws.max_row - 4)
actual_line_rows = max(0, lws.max_row - 4)
checks.append(('quote row count', actual_quote_rows == len(quote_rows)))
checks.append(('line row count', actual_line_rows == line_count))
expected_b2 = f'=SUM(Quotes!K5:K{last_quote_row})' if quote_rows else '=SUM(Quotes!K5:K5)'
checks.append(('summary B2 formula extended', sws['B2'].value == expected_b2))
quote_ids = set()
filter_ok = True
text_ids_ok = True
numeric_ok = True
fx_ok = True
for r in range(5, 5 + len(quote_rows)):
    qid = qws.cell(r,1).value
    quote_ids.add(qid)
    text_ids_ok = text_ids_ok and isinstance(qid, str) and isinstance(qws.cell(r,2).value, str) and isinstance(qws.cell(r,3).value, str)
    qdate_val = qws.cell(r,4).value
    qdate = qdate_val if isinstance(qdate_val, date) else date.fromisoformat(str(qdate_val))
    filter_ok = filter_ok and qws.cell(r,6).value == STATUS and from_dt <= qdate <= to_dt
    for c in (8,9,10,11):
        numeric_ok = numeric_ok and isinstance(qws.cell(r,c).value, (int,float))
    cur = qws.cell(r,7).value
    expected_twd = (Decimal(str(qws.cell(r,10).value)) * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    actual_twd = Decimal(str(qws.cell(r,11).value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    fx_ok = fx_ok and expected_twd == actual_twd
line_join_ok = True
for r in range(5, 5 + line_count):
    line_join_ok = line_join_ok and (lws.cell(r,1).value in quote_ids) and isinstance(lws.cell(r,1).value, str) and isinstance(lws.cell(r,2).value, str)
    for c in (4,5,6):
        numeric_ok = numeric_ok and isinstance(lws.cell(r,c).value, (int,float))
checks.append(('quotes filter Status Sent and Sep 2026', filter_ok))
checks.append(('ID columns are text values', text_ids_ok and line_join_ok))
checks.append(('amount/quantity cells are numeric', numeric_ok))
checks.append(('NetAmountTWD calculated with fx ROUND_HALF_UP', fx_ok))
summary_count = 0
summary_twd = Decimal('0')
for r in range(5, sws.max_row + 1):
    if sws.cell(r,1).value is None and sws.cell(r,2).value is None and sws.cell(r,3).value is None:
        continue
    summary_count += int(sws.cell(r,6).value or 0)
    if sws.cell(r,5).value is not None:
        summary_twd += Decimal(str(sws.cell(r,5).value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
quote_twd = Decimal('0')
for r in range(5, 5 + len(quote_rows)):
    quote_twd += Decimal(str(qws.cell(r,11).value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
checks.append(('summary count matches quotes', summary_count == len(quote_rows)))
checks.append(('summary TWD matches quote TWD total', summary_twd == quote_twd))
failed = [name for name, ok in checks if not ok]
print('\nVerification checks:')
for name, ok in checks:
    print(f" - {name}: {'OK' if ok else 'FAIL'}")
if failed:
    raise RuntimeError('Output verification failed: ' + ', '.join(failed))
print('\nCreated files:')
print(' -', OUT_XLSX)
print(' -', AUDIT_JSON)
print(f'Final counts: quotes={len(quote_rows)}, lines={line_count}, summary_rows={len(summary)}')
print('CLI attempts total:', len(attempt_log))