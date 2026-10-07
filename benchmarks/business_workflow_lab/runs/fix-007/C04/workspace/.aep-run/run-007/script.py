from pathlib import Path
import json, csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import date, datetime
from openpyxl import load_workbook

ws=Path.cwd()
out=ws/'output'/'quote_report_2026-09_Sent.xlsx'
audit_path=ws/'output'/'quote_report_2026-09_Sent_audit.json'
fx_path=ws/'input'/'fx_rates.csv'
tpl=ws/'input'/'quote_template.xlsx'
print('Rechecking saved outputs:')
print(' workbook:', out, out.exists())
print(' audit:', audit_path, audit_path.exists())

fx={}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f); next(rdr)
    for row in rdr:
        if len(row)>=2 and row[0]: fx[row[0]]=Decimal(row[1])

audit=json.loads(audit_path.read_text(encoding='utf-8'))
wb=load_workbook(out, data_only=False)
expected_sheets=['Report','Quotes','Lines','Summary','Instructions']
checks=[]
checks.append(('workbook exists', out.exists()))
checks.append(('audit exists', audit_path.exists()))
checks.append(('sheets unchanged', wb.sheetnames==expected_sheets))
checks.append(('Instructions sheet present and not removed', 'Instructions' in wb.sheetnames))
checks.append(('Report title merge kept', 'A1:F1' in [str(r) for r in wb['Report'].merged_cells.ranges]))
checks.append(('Report footer kept', wb['Report']['A20'].value=='合成測試資料 Synthetic lab data · template v1 · 請勿修改本行'))
checks.append(('period is Sep 2026 inclusive', wb['Report']['B3'].value=='2026-09-01' and wb['Report']['D3'].value=='2026-09-30'))
checks.append(('condition says Sent', 'Sent' in str(wb['Report']['B4'].value)))
checks.append(('generated/source/fetched populated', bool(wb['Report']['B5'].value) and bool(wb['Report']['B7'].value) and bool(wb['Report']['D7'].value)))
qws,lws,sws=wb['Quotes'],wb['Lines'],wb['Summary']
checks.append(('Quotes headers kept', [qws.cell(4,c).value for c in range(1,12)]==['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']))
checks.append(('Lines headers kept', [lws.cell(4,c).value for c in range(1,7)]==['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']))
checks.append(('Summary headers kept', [sws.cell(4,c).value for c in range(1,7)]==['Group','SubGroup','Currency','Amount','AmountTWD','Count']))
checks.append(('quote row count matches audit', max(0,qws.max_row-4)==audit['quote_count']))
checks.append(('line row count matches audit', max(0,lws.max_row-4)==audit['line_count']))
expected_b2=f"=SUM(Quotes!K5:K{4+audit['quote_count']})" if audit['quote_count'] else '=SUM(Quotes!K5:K5)'
checks.append(('Summary B2 formula range', sws['B2'].value==expected_b2))
from_dt=date.fromisoformat(audit['from']); to_dt=date.fromisoformat(audit['to'])
quote_ids=set(); filter_ok=True; id_text_or_blank=True; numeric_ok=True; fx_ok=True
quote_twd=Decimal('0')
for r in range(5,5+audit['quote_count']):
    qid=qws.cell(r,1).value; cid=qws.cell(r,2).value; pid=qws.cell(r,3).value
    quote_ids.add(qid)
    id_text_or_blank = id_text_or_blank and isinstance(qid,str) and isinstance(cid,str) and (pid is None or isinstance(pid,str))
    qdate_val=qws.cell(r,4).value
    qdate=qdate_val.date() if isinstance(qdate_val, datetime) else (qdate_val if isinstance(qdate_val,date) else date.fromisoformat(str(qdate_val)))
    filter_ok = filter_ok and qws.cell(r,6).value=='Sent' and from_dt<=qdate<=to_dt
    for c in (8,9,10,11): numeric_ok = numeric_ok and isinstance(qws.cell(r,c).value,(int,float))
    cur=qws.cell(r,7).value
    expected=(Decimal(str(qws.cell(r,10).value))*fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    actual=Decimal(str(qws.cell(r,11).value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    fx_ok = fx_ok and expected==actual
    quote_twd += actual
line_join_ok=True
for r in range(5,5+audit['line_count']):
    line_join_ok = line_join_ok and lws.cell(r,1).value in quote_ids and isinstance(lws.cell(r,1).value,str) and isinstance(lws.cell(r,2).value,str)
    for c in (4,5,6): numeric_ok = numeric_ok and isinstance(lws.cell(r,c).value,(int,float))
checks.append(('all Quotes rows are Sent in September 2026', filter_ok))
checks.append(('ID values are text, missing ProjectID stays blank', id_text_or_blank and line_join_ok))
checks.append(('amount and quantity cells are numeric', numeric_ok))
checks.append(('NetAmountTWD uses fx ROUND_HALF_UP', fx_ok))
summary_count=0; summary_twd=Decimal('0')
for r in range(5,sws.max_row+1):
    if sws.cell(r,1).value is None and sws.cell(r,2).value is None and sws.cell(r,3).value is None: continue
    summary_count += int(sws.cell(r,6).value or 0)
    if sws.cell(r,5).value is not None:
        summary_twd += Decimal(str(sws.cell(r,5).value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
checks.append(('summary count matches quotes', summary_count==audit['quote_count']))
checks.append(('summary TWD matches quote total', summary_twd==quote_twd))
# Retry policy from audit: stdout previews are informational only; data was only used after exit 0 and JSON parse by generating script.
by={}
for a in audit['attempts']:
    by.setdefault(a['label'],[]).append(a)
retry_ok=True; saw_temp_retry=False; no_retry_2_3=True
for label, attempts in by.items():
    retry_ok = retry_ok and len(attempts)<=3
    for idx,a in enumerate(attempts[:-1]):
        if a['exit_code']==4: saw_temp_retry=True
        retry_ok = retry_ok and a['exit_code']==4
        if a['exit_code'] in (2,3): no_retry_2_3=False
checks.append(('CLI retry policy: code 4 max two retries, code 2/3 not retried', retry_ok and no_retry_2_3))
checks.append(('CLI successful data only in audit/result', all(a['exit_code'] in (0,4) for a in audit['attempts'])))
failed=[n for n,ok in checks if not ok]
print('\nVerification checks:')
for n,ok in checks: print(f" - {n}: {'OK' if ok else 'FAIL'}")
if failed: raise RuntimeError('Verification failed: '+', '.join(failed))
print('\nVerified output files:')
print(' - output/quote_report_2026-09_Sent.xlsx')
print(' - output/quote_report_2026-09_Sent_audit.json')
print(f"Counts: quotes={audit['quote_count']}, lines={audit['line_count']}, summary rows={max(0,sws.max_row-4)}")
print('CLI attempts recorded:', len(audit['attempts']))