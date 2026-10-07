from pathlib import Path
from openpyxl import load_workbook
p=Path('output/quote_report_2026-09_Sent.xlsx')
wb=load_workbook(p, data_only=False)
for wsname in ['Quotes','Lines']:
    ws=wb[wsname]
    print(wsname, ws.max_row, ws.max_column)
    for r in range(5, ws.max_row+1):
        print(r, [(ws.cell(r,c).value, type(ws.cell(r,c).value).__name__, ws.cell(r,c).number_format) for c in range(1, min(ws.max_column,3)+1)])