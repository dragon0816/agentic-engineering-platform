from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'expense_template.xlsx'
FX = INPUT / 'fx_rates.csv'
EXPORT = OUTPUT / 'expense_2026-09_approved.csv'
OUTFILE = OUTPUT / 'expense_report_2026-09_approved.xlsx'

print('Inputs confirmed:')
for p in [TEMPLATE, FX, EXPORT]:
    print(f' - {p}: exists={p.exists()} bytes={p.stat().st_size if p.exists() else "NA"}')

# Helpers
TWOPLACES = Decimal('0.01')
def q2(x: Decimal) -> Decimal:
    return x.quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def clear_rows(ws, start_row, max_col):
    # Clear existing data area from start_row downward, preserving row 4 headers and sheet settings.
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def style_new_row_like(ws, src_row, dst_row, max_col):
    # Copy formatting/number formats from template/reference row to new row.
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.protection:
            dst.protection = copy(src.protection)

# Read FX rates
rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        cur = (row.get('Currency') or '').strip()
        rate = (row.get('RateToTWD') or '').strip()
        if cur:
            rates[cur] = Decimal(rate)
print('FX rates:', {k: str(v) for k,v in rates.items()})

# Read exported expense data
rows = []
with EXPORT.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('Export headers:', reader.fieldnames)
    for r in reader:
        amount = Decimal(r['Amount'])
        currency = r['Currency']
        if currency not in rates:
            raise ValueError(f'Missing FX rate for {currency}')
        amount_twd = q2(amount * rates[currency])
        missing_raw = str(r.get('MissingReceipt','')).strip().lower()
        missing_bool = (missing_raw == 'true')
        rec = {
            'ExpenseID': r['ExpenseID'],
            'EmployeeID': r['EmployeeID'],
            'EmployeeName': r['EmployeeName'],
            'ProjectID': r['ProjectID'],
            'ExpenseDate': r['ExpenseDate'],
            'Category': r['Category'],
            'Currency': currency,
            'Amount': q2(amount),
            'AmountTWD': amount_twd,
            'Status': r['Status'],
            'ReceiptID': r.get('ReceiptID',''),
            'Description': r.get('Description',''),
            'MissingReceipt': 'true' if missing_bool else 'false',
            'NeedsReview': 'true' if amount_twd > Decimal('5000') else 'false',
        }
        rows.append(rec)

rows.sort(key=lambda r: (r['ExpenseDate'], r['ExpenseID']))
print(f'Read {len(rows)} exported rows')
print('First rows:', rows[:3])

# Build summary by employee, category, currency (currency retained to avoid mixing original amounts)
summary = {}
for r in rows:
    key = (f"{r['EmployeeID']} {r['EmployeeName']}", r['Category'], r['Currency'])
    if key not in summary:
        summary[key] = {'Amount': Decimal('0'), 'AmountTWD': Decimal('0'), 'Count': 0}
    summary[key]['Amount'] += r['Amount']
    summary[key]['AmountTWD'] += r['AmountTWD']
    summary[key]['Count'] += 1
summary_rows = []
for key, val in sorted(summary.items(), key=lambda kv: (kv[0][0], kv[0][1], kv[0][2])):
    group, subgroup, currency = key
    summary_rows.append({
        'Group': group,
        'SubGroup': subgroup,
        'Currency': currency,
        'Amount': q2(val['Amount']),
        'AmountTWD': q2(val['AmountTWD']),
        'Count': val['Count'],
    })
print(f'Summary rows: {len(summary_rows)}')

exception_rows = []
for r in rows:
    reasons = []
    if r['MissingReceipt'] == 'true':
        reasons.append('MissingReceipt=true')
    if not r['ReceiptID']:
        reasons.append('ReceiptID blank')
    if reasons:
        exception_rows.append({
            'ExpenseID': r['ExpenseID'],
            'EmployeeID': r['EmployeeID'],
            'ExpenseDate': r['ExpenseDate'],
            'Reason': '; '.join(reasons),
        })
print(f'Exception rows: {len(exception_rows)}')

# Load and fill a copy of template
wb = load_workbook(TEMPLATE)
for required in ['Report','Details','Summary','Exceptions','Instructions']:
    if required not in wb.sheetnames:
        raise ValueError(f'Missing sheet {required}')

report = wb['Report']
details = wb['Details']
sumws = wb['Summary']
exws = wb['Exceptions']

# Report metadata
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Status=Approved; source=Expense Viewer export; AmountTWD uses fixed rates from fx_rates.csv; NeedsReview=true when AmountTWD > 5000'
report['B5'] = datetime.now().replace(microsecond=0).isoformat()
report['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D7'] = datetime.now().replace(microsecond=0).isoformat()
report['B9'] = f'{len(rows)} approved expenses exported. Summary grouped by Employee + Category + Currency to preserve original currency amounts. Exceptions list missing receipt data.'

# Details
clear_rows(details, 5, 14)
headers = [details.cell(4, c).value for c in range(1, 15)]
for i, r in enumerate(rows, start=5):
    style_new_row_like(details, 4, i, 14)
    values = [r[h] for h in headers]
    for c, v in enumerate(values, start=1):
        cell = details.cell(i, c, v)
        if c in [1,2,4,7,10,11,13,14]:
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in [8,9]:
            cell.number_format = '0.00'
            cell.value = float(v)

# Summary
clear_rows(sumws, 5, 6)
last_detail_row = 4 + len(rows) if rows else 5
sumws['B2'] = f'=SUM(Details!I5:I{last_detail_row if rows else 5})'
for i, r in enumerate(summary_rows, start=5):
    style_new_row_like(sumws, 4, i, 6)
    vals = [r['Group'], r['SubGroup'], r['Currency'], r['Amount'], r['AmountTWD'], r['Count']]
    for c, v in enumerate(vals, start=1):
        cell = sumws.cell(i, c, v)
        if c in [1,2,3]:
            cell.number_format = '@'
        elif c in [4,5]:
            cell.number_format = '0.00'
            cell.value = float(v)
        elif c == 6:
            cell.number_format = '0'

# Exceptions
clear_rows(exws, 5, 4)
for i, r in enumerate(exception_rows, start=5):
    style_new_row_like(exws, 4, i, 4)
    vals = [r['ExpenseID'], r['EmployeeID'], r['ExpenseDate'], r['Reason']]
    for c, v in enumerate(vals, start=1):
        cell = exws.cell(i, c, v)
        if c in [1,2,4]:
            cell.number_format = '@'
        elif c == 3:
            cell.number_format = 'yyyy-mm-dd'

# Preserve views / filters can remain. Save.
wb.save(OUTFILE)
print('Saved:', OUTFILE)

# Verification: reopen output and confirm requirements
wb2 = load_workbook(OUTFILE, data_only=False)
issues = []
# sheet names intact
expected_sheets = ['Report','Details','Summary','Exceptions','Instructions']
if wb2.sheetnames != expected_sheets:
    issues.append(f'Sheet names/order changed: {wb2.sheetnames}')

d2 = wb2['Details']
s2 = wb2['Summary']
e2 = wb2['Exceptions']
# headers intact
if [d2.cell(4,c).value for c in range(1,15)] != headers:
    issues.append('Details headers differ from template')
if [s2.cell(4,c).value for c in range(1,7)] != ['Group','SubGroup','Currency','Amount','AmountTWD','Count']:
    issues.append('Summary headers invalid')
if [e2.cell(4,c).value for c in range(1,5)] != ['ExpenseID','EmployeeID','ExpenseDate','Reason']:
    issues.append('Exceptions headers invalid')
# row counts
actual_detail_rows = max(0, d2.max_row - 4)
actual_summary_rows = max(0, s2.max_row - 4)
actual_exception_rows = max(0, e2.max_row - 4)
if actual_detail_rows != len(rows): issues.append(f'Details row count {actual_detail_rows} != {len(rows)}')
if actual_summary_rows != len(summary_rows): issues.append(f'Summary row count {actual_summary_rows} != {len(summary_rows)}')
if actual_exception_rows != len(exception_rows): issues.append(f'Exceptions row count {actual_exception_rows} != {len(exception_rows)}')
# Amount calculations, status, period, flags
export_by_id = {r['ExpenseID']: r for r in rows}
for rownum in range(5, 5+len(rows)):
    exp_id = d2.cell(rownum,1).value
    r = export_by_id.get(exp_id)
    if not r:
        issues.append(f'Unexpected ExpenseID {exp_id} at row {rownum}')
        continue
    if d2.cell(rownum,7).value != r['Currency']:
        issues.append(f'Currency mismatch {exp_id}')
    if round(float(d2.cell(rownum,8).value),2) != float(r['Amount']):
        issues.append(f'Amount mismatch {exp_id}')
    if round(float(d2.cell(rownum,9).value),2) != float(r['AmountTWD']):
        issues.append(f'AmountTWD mismatch {exp_id}')
    if d2.cell(rownum,10).value != 'Approved':
        issues.append(f'Status not Approved {exp_id}')
    if not ('2026-09-01' <= d2.cell(rownum,5).value <= '2026-09-30'):
        issues.append(f'Date out of range {exp_id}')
    expected_flag = 'true' if Decimal(str(d2.cell(rownum,9).value)) > Decimal('5000') else 'false'
    if d2.cell(rownum,14).value != expected_flag:
        issues.append(f'NeedsReview mismatch {exp_id}')
    if not isinstance(d2.cell(rownum,8).value, (int,float)) or not isinstance(d2.cell(rownum,9).value, (int,float)):
        issues.append(f'Amount cells not numeric {exp_id}')
# Summary totals vs Details
summary_twd = sum(Decimal(str(s2.cell(r,5).value)) for r in range(5,5+len(summary_rows)))
detail_twd = sum(Decimal(str(d2.cell(r,9).value)) for r in range(5,5+len(rows)))
if q2(summary_twd) != q2(detail_twd):
    issues.append(f'Summary TWD {summary_twd} != Details TWD {detail_twd}')
# Formula extended
expected_formula = f'=SUM(Details!I5:I{4+len(rows)})'
if s2['B2'].value != expected_formula:
    issues.append(f'Summary B2 formula {s2["B2"].value} != {expected_formula}')
# No blank rows in Details data range
for rr in range(5, 5+len(rows)):
    if all(d2.cell(rr, cc).value in (None,'') for cc in range(1,15)):
        issues.append(f'Blank detail row {rr}')

print('\nVerification:')
print(' - workbook reopened:', OUTFILE.name)
print(' - Details rows:', actual_detail_rows)
print(' - Summary rows:', actual_summary_rows)
print(' - Exceptions rows:', actual_exception_rows)
print(' - Details AmountTWD total:', q2(detail_twd))
print(' - Summary AmountTWD total:', q2(summary_twd))
print(' - Summary B2 formula:', s2['B2'].value)
print(' - NeedsReview true count:', sum(1 for r in range(5,5+len(rows)) if d2.cell(r,14).value == 'true'))
if issues:
    print('ISSUES:')
    for x in issues:
        print(' *', x)
    raise SystemExit(1)
print('All checks passed.')
