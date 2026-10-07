from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('Exported files:')
for p in sorted((ws/'output').glob('*')):
    print(p.name, p.stat().st_size)
print('\nExport CSV preview:')
with (ws/'output'/'expense_2026-09_approved.csv').open(newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(8), rdr):
        print(row)
print('\nInstructions sheet all rows:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
