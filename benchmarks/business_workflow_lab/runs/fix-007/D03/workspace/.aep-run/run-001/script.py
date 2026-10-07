from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('Workspace:', ws)
print('Input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size, 'bytes')

# CSV headers/sample
csv_path = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv:')
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.reader(f))
    for r in rows[:6]:
        print(r)

# manual first lines
manual = ws/'input'/'expense_manual.md'
print('\nexpense_manual.md first 80 lines:')
print('\n'.join(manual.read_text(encoding='utf-8').splitlines()[:80]))

# workbook sheets, dimensions, headers/sample values
xlsx = ws/'input'/'expense_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWorkbook sheets:')
for sh in wb.sheetnames:
    wsht = wb[sh]
    print(f' - {sh}: {wsht.max_row} rows x {wsht.max_column} cols')
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,5), values_only=True):
        print('   ', row)
