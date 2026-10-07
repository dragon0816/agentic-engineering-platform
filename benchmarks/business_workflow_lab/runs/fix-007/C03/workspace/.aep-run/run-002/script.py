from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
ins=wb['Instructions']
for row in ins.iter_rows(values_only=True):
    print(row)