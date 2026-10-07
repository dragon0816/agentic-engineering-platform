from pathlib import Path
import subprocess, json, sys, csv, re, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date, timezone, timedelta
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd(); INPUT = WORKSPACE/'input'; OUTPUT = WORKSPACE/'output'; OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT/'quote_template.xlsx'; FX_FILE = INPUT/'fx_rates.csv'; MANUAL = INPUT/'quote_cli_manual.md'
OUTFILE = OUTPUT/'quote_report_P001-P005_2026-09_Accepted.xlsx'
PROJECTS = [f'P{i:03d}' for i in range(1,6)]
FROM_DATE='2026-09-01'; TO_DATE='2026-09-30'; STATUS='Accepted'; TWOPLACES=Decimal('0.01')

def q2(x): return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def parse_date(s):
    y,m,d=map(int,s.split('-')); return date(y,m,d)
def as_date(v):
    if isinstance(v, datetime): return v.date()
    return v

def read_fx(path):
    rates={}
    with open(path, newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f); header=next(r); print('fx header sample:', header)
        for row in r:
            if row: rates[row[0]]=Decimal(row[1])
    return rates

def find_cli(text):
    m=re.search(r'python "([^"]*run_quote_cli\.py)" <command>', text)
    if not m: raise RuntimeError('Cannot find CLI path')
    return m.group(1)

def run_cli(cli,args):
    res=subprocess.run([sys.executable, cli]+args, capture_output=True, text=True)
    if res.returncode!=0: raise RuntimeError(f'CLI failed {res.returncode} {args}\n{res.stderr}\n{res.stdout}')
    return json.loads(res.stdout)

def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1,max_col+1):
        src=ws.cell(src_row,col); dst=ws.cell(dst_row,col)
        if src.has_style: dst._style=copy(src._style)
        dst.number_format=src.number_format; dst.alignment=copy(src.alignment); dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border)
    ws.row_dimensions[dst_row].height=ws.row_dimensions[src_row].height

def clear_data(ws,start_row=5):
    if ws.max_row>=start_row: ws.delete_rows(start_row, ws.max_row-start_row+1)

print('Workspace:', WORKSPACE)
print('Input files:')
for p in sorted(INPUT.iterdir()): print('-', p.name, p.stat().st_size)
fx=read_fx(FX_FILE); print('FX rates:', {k:str(v) for k,v in fx.items()})
cli=find_cli(MANUAL.read_text(encoding='utf-8')); print('CLI path:', cli)

headers=[]; line_rows=[]
for project in PROJECTS:
    listed=run_cli(cli, ['quotes','list','--project',project,'--from',FROM_DATE,'--to',TO_DATE,'--status',STATUS,'--format','json'])
    print(f'{project}: count={listed.get("count")}')
    for item in listed.get('items',[]):
        assert item['ProjectID']==project and item['Status']==STATUS and FROM_DATE<=item['QuoteDate']<=TO_DATE
        assert item['Currency'] in fx
        detail=run_cli(cli, ['quotes','get',item['QuoteID'],'--format','json'])
        hdr=detail['header']; assert hdr==item
        hdr=dict(hdr); hdr['NetAmountTWD']=q2(Decimal(hdr['NetAmount'])*fx[hdr['Currency']]); headers.append(hdr)
        for li in detail.get('line_items',[]):
            row=dict(li); row['QuoteID']=hdr['QuoteID']; line_rows.append(row)
headers.sort(key=lambda h:(h['ProjectID'],h['QuoteDate'],h['QuoteID']))
line_rows.sort(key=lambda r:(r['QuoteID'],r['ProductID']))
print('Fetched quote headers:', len(headers))
print('Fetched line items:', len(line_rows))
print('Quote headers sample:')
for h in headers: print(h)

by=defaultdict(lambda:{'Amount':Decimal('0.00'),'AmountTWD':Decimal('0.00'),'Count':0})
for h in headers:
    key=(h['ProjectID'],h['Currency']); by[key]['Amount']+=Decimal(h['NetAmount']); by[key]['AmountTWD']+=h['NetAmountTWD']; by[key]['Count']+=1
summary=[]
for project in PROJECTS:
    keys=sorted([k for k in by if k[0]==project], key=lambda k:k[1])
    if not keys: summary.append({'Group':project,'SubGroup':'','Currency':'','Amount':Decimal('0.00'),'AmountTWD':Decimal('0.00'),'Count':0})
    else:
        for proj,curr in keys:
            v=by[(proj,curr)]; summary.append({'Group':proj,'SubGroup':'','Currency':curr,'Amount':q2(v['Amount']),'AmountTWD':q2(v['AmountTWD']),'Count':v['Count']})
print('Summary to write:')
for r in summary: print(r)

shutil.copyfile(TEMPLATE, OUTFILE)
wb=load_workbook(OUTFILE); report=wb['Report']; qws=wb['Quotes']; lws=wb['Lines']; sws=wb['Summary']
for ws in (qws,lws,sws): clear_data(ws,5)
generated=datetime(2026,10,6,0,0,0,tzinfo=timezone(timedelta(hours=8))).isoformat()
report['B3']=FROM_DATE; report['D3']=TO_DATE
report['B4']='ProjectID P001-P005; Status Accepted; QuoteDate 2026-09-01..2026-09-30; Summary by ProjectID; original currency amounts are not mixed'
report['B5']=generated; report['B7']='quote_cli; fx_rates.csv'; report['D7']=generated
missing=[p for p in PROJECTS if not any(h['ProjectID']==p for h in headers)]
report['B9']='No Accepted quote in period for: '+', '.join(missing) if missing else ''

quote_cols=['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for i,h in enumerate(headers,5):
    copy_row_style(qws,4,i,11)
    for c,col in enumerate(quote_cols,1):
        cell=qws.cell(i,c)
        if col in ('QuoteID','CustomerID','ProjectID','Status','Currency'):
            cell.value=str(h[col]); cell.number_format='@'
        elif col in ('QuoteDate','ValidUntil'):
            cell.value=parse_date(h[col]); cell.number_format='yyyy-mm-dd'
        elif col=='DiscountPercent':
            cell.value=float(Decimal(str(h[col]))); cell.number_format='0'
        else:
            cell.value=float(q2(h[col])); cell.number_format='0.00'
line_cols=['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for i,r in enumerate(line_rows,5):
    copy_row_style(lws,4,i,6)
    for c,col in enumerate(line_cols,1):
        cell=lws.cell(i,c)
        if col in ('QuoteID','ProductID','ProductName'):
            cell.value=str(r[col]); cell.number_format='@'
        elif col=='Quantity':
            dec=Decimal(str(r[col])); cell.value=int(dec) if dec==dec.to_integral_value() else float(dec); cell.number_format='0'
        else:
            cell.value=float(q2(r[col])); cell.number_format='0.00'
for i,r in enumerate(summary,5):
    copy_row_style(sws,4,i,6)
    vals=[r['Group'],r['SubGroup'],r['Currency'],r['Amount'],r['AmountTWD'],r['Count']]
    for c,v in enumerate(vals,1):
        cell=sws.cell(i,c)
        if c<=3: cell.value=str(v) if v!='' else None; cell.number_format='@'
        elif c in (4,5): cell.value=float(q2(v)); cell.number_format='0.00'
        else: cell.value=int(v); cell.number_format='0'
last_quote_row=4+max(1,len(headers)); sws['B2']=f'=SUM(Quotes!K5:K{last_quote_row})'
if qws.auto_filter and qws.auto_filter.ref: qws.auto_filter.ref=f'A4:K{max(4,4+len(headers))}'
if lws.auto_filter and lws.auto_filter.ref: lws.auto_filter.ref=f'A4:F{max(4,4+len(line_rows))}'
if sws.auto_filter and sws.auto_filter.ref: sws.auto_filter.ref=f'A4:F{max(4,4+len(summary))}'
wb.save(OUTFILE); print('Wrote:', OUTFILE)

# Reopen checks
wbv=load_workbook(OUTFILE, data_only=False); qws=wbv['Quotes']; lws=wbv['Lines']; sws=wbv['Summary']; rws=wbv['Report']
assert wbv.sheetnames==['Report','Quotes','Lines','Summary','Instructions']
assert [qws.cell(4,c).value for c in range(1,12)]==quote_cols
assert [lws.cell(4,c).value for c in range(1,7)]==line_cols
assert [sws.cell(4,c).value for c in range(1,7)]==['Group','SubGroup','Currency','Amount','AmountTWD','Count']
assert rws['B3'].value==FROM_DATE and rws['D3'].value==TO_DATE
assert (qws.max_row-4)==len(headers); assert (lws.max_row-4)==len(line_rows); assert (sws.max_row-4)==len(summary)
for row in range(5,5+len(headers)):
    assert isinstance(qws.cell(row,1).value,str) and qws.cell(row,1).value.startswith('Q')
    assert isinstance(qws.cell(row,2).value,str) and qws.cell(row,2).value.startswith('C')
    assert isinstance(qws.cell(row,3).value,str) and qws.cell(row,3).value in PROJECTS
    assert qws.cell(row,6).value==STATUS
    d=as_date(qws.cell(row,4).value); assert isinstance(d,date) and date(2026,9,1)<=d<=date(2026,9,30)
    curr=qws.cell(row,7).value; net=Decimal(str(qws.cell(row,10).value)); twd=Decimal(str(qws.cell(row,11).value))
    assert twd==q2(net*fx[curr]), (row,curr,net,twd,q2(net*fx[curr]))
for row in range(5,5+len(line_rows)):
    assert isinstance(lws.cell(row,1).value,str) and lws.cell(row,1).value.startswith('Q')
    assert isinstance(lws.cell(row,2).value,str) and isinstance(lws.cell(row,3).value,str)
    assert isinstance(lws.cell(row,4).value,int) and isinstance(lws.cell(row,5).value,(int,float)) and isinstance(lws.cell(row,6).value,(int,float))
summary_projects=[sws.cell(row,1).value for row in range(5,5+len(summary))]
for project in PROJECTS: assert project in summary_projects
for row in range(5,5+len(summary)):
    proj=sws.cell(row,1).value; curr=sws.cell(row,3).value; amount=Decimal(str(sws.cell(row,4).value)); amount_twd=Decimal(str(sws.cell(row,5).value)); count=sws.cell(row,6).value
    assert proj in PROJECTS and isinstance(count,int)
    if count==0: assert amount==Decimal('0') and amount_twd==Decimal('0')
    else: assert curr in fx and amount_twd==q2(amount*fx[curr])
assert sws['B2'].value==f'=SUM(Quotes!K5:K{last_quote_row})'
print('Validation OK')
print('Checked output file:', OUTFILE.name)
print('Checked sheet names/header rows/template-based workbook')
print('Checked Quotes rows:', qws.max_row-4, 'IDs as text, date/status/project filters, numeric amounts, NetAmountTWD=NetAmount*FX ROUND_HALF_UP')
print('Checked Lines rows:', lws.max_row-4, 'IDs as text, numeric quantity and amounts')
print('Checked Summary rows:', sws.max_row-4, 'P001-P005 all present, zero rows for projects without Accepted quotes, original currency Amount not cross-mixed')
print('Checked Summary!B2:', sws['B2'].value)
print('Checked Report metadata and note:', rws['B9'].value)