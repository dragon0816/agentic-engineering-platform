import os, csv
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
    if p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print('  sheets:', wb.sheetnames)
        for s in wb.sheetnames:
            sh=wb[s]
            print('  sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
            for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
                print('   ', row)
    elif p.suffix.lower()=='.csv':
        with open(p, newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f)
            for i,row in zip(range(5), r): print('  ', row)
    elif p.suffix.lower()=='.md':
        print(p.read_text(encoding='utf-8')[:1000])
print('output exists', (ws/'output').exists())