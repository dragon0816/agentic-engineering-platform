import csv, shutil, os
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font

workspace=Path.cwd()
input_dir=workspace/'input'
output_dir=workspace/'output'
output_dir.mkdir(exist_ok=True)
export_csv=output_dir/'expense_export_E001_2026-09-01_2026-09-30.csv'
template=input_dir/'expense_template.xlsx'
fx_file=input_dir/'fx_rates.csv'
outfile=output_dir/'E001_september_expenses.xlsx'

print('Inputs:')
print(' export_csv', export_csv, export_csv.exists(), export_csv.stat().st_size if export_csv.exists() else None)
print(' template', template, template.exists(), template.stat().st_size if template.exists() else None)
print(' fx_file', fx_file, fx_file.exists(), fx_file.stat().st_size if fx_file.exists() else None)

# read rates
rates={}
with open(fx_file, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    header=next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row)>=2 and row[0]:
            rates[row[0]]=Decimal(row[1])
print('rates:', rates)

# read export rows
with open(export_csv, newline='', encoding='utf-8-sig') as f:
    reader=csv.DictReader(f)
    rows=list(reader)
print('export header:', reader.fieldnames)
print('export row count:', len(rows))
print('sample rows:', rows[:2])

q_from='2026-09-01'; q_to='2026-09-30'; employee='E001'
# validate exported rows match query
bad=[r for r in rows if r['EmployeeID']!=employee or not(q_from <= r['ExpenseDate'] <= q_to)]
if bad:
    raise ValueError(f'Export contains {len(bad)} rows outside query, e.g. {bad[:2]}')

# copy/open template (do not rebuild)
shutil.copyfile(template, outfile)
wb=load_workbook(outfile)

# helper clear values in data rows while preserving formatting where possible
for sheet_name in ['Details','Summary','Exceptions']:
    ws=wb[sheet_name]
    start=5
    if ws.max_row>=start:
        for row in ws.iter_rows(min_row=start, max_row=ws.max_row):
            for cell in row:
                cell.value=None

# Report fields
ws=wb['Report']
ws['B3']=q_from
ws['D3']=q_to
ws['B4']=f'EmployeeID={employee}; ExpenseDate {q_from} to {q_to}; Status=all'
ws['B5']='2026-10-06T00:00:00+08:00'
ws['B7']='Expense Viewer (synthetic lab) export CSV; fx_rates.csv'
ws['D7']='2026-10-06T00:00:00+08:00'
ws['B9']='All rows exported through the Expense Viewer export function. NeedsReview/Exceptions are based on MissingReceipt=true from source data.'

# Details
headers=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
ws=wb['Details']
# sort by date then ExpenseID for stable report
rows_sorted=sorted(rows, key=lambda r:(r['ExpenseDate'], r['ExpenseID']))
for idx,r in enumerate(rows_sorted, start=5):
    cur=r['Currency']
    if cur not in rates:
        raise ValueError(f'Missing FX rate for {cur}')
    amount=Decimal(r['Amount'])
    amount_twd=(amount*rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    missing=str(r.get('MissingReceipt','')).strip().lower() == 'true'
    needs_review=missing
    values=[
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['ExpenseDate'], r['Category'], cur,
        float(amount), float(amount_twd), r['Status'], r.get('ReceiptID',''), r.get('Description',''), str(missing).lower(), str(needs_review).lower()
    ]
    for col,val in enumerate(values, start=1):
        c=ws.cell(idx,col,val)
        if col in [1,2,4,11,13,14]:
            c.number_format='@'
        elif col==5:
            c.number_format='yyyy-mm-dd'
        elif col in [8,9]:
            c.number_format='0.00'

last_detail=4+len(rows_sorted) if rows_sorted else 5
# Summary
ws=wb['Summary']
last_formula_row=4+len(rows_sorted) if rows_sorted else 5
ws['B2']=f'=SUM(Details!I5:I{last_formula_row})'
summary=defaultdict(lambda: {'amount':Decimal('0'), 'twd':Decimal('0'), 'count':0})
for r in rows_sorted:
    amount=Decimal(r['Amount'])
    twd=(amount*rates[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    key=(r['Status'], r['Currency'])
    summary[key]['amount'] += amount
    summary[key]['twd'] += twd
    summary[key]['count'] += 1
for idx,((status,currency),vals) in enumerate(sorted(summary.items()), start=5):
    row_vals=[status, '', currency, float(vals['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), float(vals['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), vals['count']]
    for col,val in enumerate(row_vals, start=1):
        c=ws.cell(idx,col,val)
        if col in [4,5]: c.number_format='0.00'
# Add a visible TWD total row after grouped summary as requested
idx=5+len(summary)
ws.cell(idx,1,'TOTAL_TWD')
ws.cell(idx,2,'')
ws.cell(idx,3,'TWD')
ws.cell(idx,4,'')
ws.cell(idx,5,f'=B2')
ws.cell(idx,6,len(rows_sorted))
ws.cell(idx,5).number_format='0.00'
for c in ws[idx]:
    c.font = Font(bold=True)

# Exceptions: MissingReceipt=true only
ws=wb['Exceptions']
exc=[r for r in rows_sorted if str(r.get('MissingReceipt','')).strip().lower()=='true']
for idx,r in enumerate(exc, start=5):
    ws.cell(idx,1,r['ExpenseID']).number_format='@'
    ws.cell(idx,2,r['EmployeeID']).number_format='@'
    ws.cell(idx,3,r['ExpenseDate']).number_format='yyyy-mm-dd'
    ws.cell(idx,4,'MissingReceipt=true')

wb.save(outfile)
print('Wrote', outfile, outfile.stat().st_size)
print('Rows written Details:', len(rows_sorted), 'Summary groups:', len(summary), 'Exceptions:', len(exc))
print('TWD total formula Summary!B2:', wb['Summary']['B2'].value)
