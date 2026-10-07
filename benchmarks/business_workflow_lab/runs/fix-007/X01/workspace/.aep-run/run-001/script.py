from pathlib import Path
from openpyxl import load_workbook
import csv, json, os
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nManual previews:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=ws/'input'/name
    print(f'--- {name} ---')
    txt=p.read_text(encoding='utf-8', errors='replace')
    print(txt[:4000])
print('\nFX rates:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r): print(row)
print('\nTemplate workbook:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
print(wb.sheetnames)
for sh in wb.worksheets:
    print('sheet', sh.title, sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)
