from openpyxl import load_workbook
from pathlib import Path
wb=load_workbook('input/project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    print('\nSHEET', sh)
    ws=wb[sh]
    for r in range(1, min(ws.max_row,8)+1):
        print(r, [ws.cell(r,c).value for c in range(1, min(ws.max_column,12)+1)])