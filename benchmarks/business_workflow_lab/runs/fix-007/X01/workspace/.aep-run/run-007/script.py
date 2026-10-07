from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path('.')
print('Workspace:', ws.resolve())
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p, 'size=', p.stat().st_size)

print('\nfx_rates.csv:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rows = list(csv.reader(f))
    for r in rows[:10]:
        print(r)

print('\nWorkbook sheets and headers/sample dimensions:')
wb = load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    wsht = wb[sh]
    print(sh, 'max_row=', wsht.max_row, 'max_col=', wsht.max_column)
    print(' row1:', [wsht.cell(1,c).value for c in range(1, min(wsht.max_column,12)+1)])

for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    print('\n---', name, '---')
    text = (ws/'input'/name).read_text(encoding='utf-8', errors='replace')
    print(text[:4000])