import os, csv, json, subprocess, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
DOWNLOADS = OUTPUT / 'downloads'
OUTPUT.mkdir(exist_ok=True)
DOWNLOADS.mkdir(exist_ok=True)
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
PERIOD = '2026-09'

print('Inputs discovered:')
for p in sorted(INPUT.iterdir()): print(' input/', p.name, p.stat().st_size)
for p in sorted(DOWNLOADS.glob('*')): print(' download/', p.name, p.stat().st_size)

# Helpers
TWOPLACES = Decimal('0.01')
def D(v):
    if v is None: return Decimal('0')
    s = str(v).strip()
    return Decimal(s if s else '0')
def q2(v):
    return D(v).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def fmt_dec(v):
    return str(q2(v))
def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    print(path.name, 'headers=', rows[0].keys() if rows else 'NO ROWS', 'rows=', len(rows))
    if rows: print(' sample=', rows[0])
    return rows
def in_period(date_s):
    return FROM_DATE <= str(date_s) <= TO_DATE

# FX
fx_rates = {}
with open(INPUT/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for r in reader:
        fx_rates[r['Currency'].strip()] = D(r['RateToTWD'])
print('fx_rates=', fx_rates)
for c in ['TWD','USD','EUR']:
    assert c in fx_rates, f'Missing FX {c}'

def to_twd(amount, currency):
    if currency not in fx_rates:
        raise ValueError(f'Missing FX rate for {currency}')
    return q2(D(amount) * fx_rates[currency])

# Source files from browser/desktop
orders_raw = read_csv(DOWNLOADS/f'sap_orders_{PERIOD}.csv')
invoices_raw = read_csv(DOWNLOADS/f'sap_invoices_{PERIOD}.csv')
expenses_raw = read_csv(DOWNLOADS/f'expenses_{PERIOD}_approved.csv')

# quote CLI export to downloads
quote_cli_py = WORKSPACE.parent.parent.parent / 'quote_cli_mock' / 'run_quote_cli.py'
if not quote_cli_py.exists():
    # fallback from documented absolute sibling lab path
    quote_cli_py = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
quote_out = DOWNLOADS/f'quotes_{PERIOD}_accepted.csv'
cmd = ['python', str(quote_cli_py), 'quotes', 'export', '--from', FROM_DATE, '--to', TO_DATE, '--status', 'Accepted', '--output', str(quote_out)]
print('Running quote CLI:', cmd)
res = subprocess.run(cmd, text=True, capture_output=True)
print('quote_cli returncode=', res.returncode)
if res.stdout: print('quote_cli stdout=', res.stdout[:1000])
if res.stderr: print('quote_cli stderr=', res.stderr[:1000])
if res.returncode != 0:
    raise SystemExit('quote_cli failed')
quotes_raw = read_csv(quote_out)

# Filter and normalize
orders = []
for r in orders_raw:
    if not in_period(r.get('OrderDate','')): continue
    if r.get('Status') == 'Cancelled': continue
    currency = r['Currency']
    orders.append({
        'OrderID': str(r['OrderID']),
        'ProjectID': r.get('ProjectID') or None,
        'CustomerID': r.get('CustomerID') or None,
        'OrderDate': r['OrderDate'],
        'Status': r['Status'],
        'Currency': currency,
        'NetAmount': q2(r['NetAmount']),
        'AmountTWD': to_twd(r['NetAmount'], currency),
    })

invoices = []
for r in invoices_raw:
    if not in_period(r.get('InvoiceDate','')): continue
    currency = r['Currency']
    balance = q2(r.get('Balance') if r.get('Balance') not in [None,''] else D(r['Amount'])-D(r['PaidAmount']))
    invoices.append({
        'InvoiceID': str(r['InvoiceID']),
        'OrderID': str(r.get('OrderID') or ''),
        'ProjectID': r.get('ProjectID') or None,
        'CustomerID': r.get('CustomerID') or None,
        'InvoiceDate': r['InvoiceDate'],
        'DueDate': r['DueDate'],
        'Currency': currency,
        'Amount': q2(r['Amount']),
        'PaidAmount': q2(r['PaidAmount']),
        'Balance': balance,
        'BalanceTWD': to_twd(balance, currency),
    })

expenses = []
for r in expenses_raw:
    if not in_period(r.get('ExpenseDate','')): continue
    if r.get('Status') != 'Approved': continue
    currency = r['Currency']
    expenses.append({
        'ExpenseID': str(r['ExpenseID']),
        'ProjectID': r.get('ProjectID') or None,
        'EmployeeID': r.get('EmployeeID') or None,
        'ExpenseDate': r['ExpenseDate'],
        'Status': r['Status'],
        'Currency': currency,
        'Amount': q2(r['Amount']),
        'AmountTWD': to_twd(r['Amount'], currency),
    })

quotes = []
for r in quotes_raw:
    if not in_period(r.get('QuoteDate','')): continue
    if r.get('Status') != 'Accepted': continue
    currency = r['Currency']
    quotes.append({
        'QuoteID': str(r['QuoteID']),
        'ProjectID': r.get('ProjectID') or None,
        'CustomerID': r.get('CustomerID') or None,
        'QuoteDate': r['QuoteDate'],
        'ValidUntil': r['ValidUntil'],
        'Status': r['Status'],
        'Currency': currency,
        'NetAmount': q2(r['NetAmount']),
        'AmountTWD': to_twd(r['NetAmount'], currency),
    })

print('Filtered counts:', {'orders_non_cancelled':len(orders), 'invoices':len(invoices), 'approved_expenses':len(expenses), 'accepted_quotes':len(quotes)})

summary = {}
def add_summary(project, col, amount):
    if not project: return
    summary.setdefault(project, {'OrderTWD':Decimal('0'), 'InvoiceBalanceTWD':Decimal('0'), 'ApprovedExpenseTWD':Decimal('0'), 'AcceptedQuoteTWD':Decimal('0')})
    summary[project][col] += D(amount)
for r in orders: add_summary(r['ProjectID'], 'OrderTWD', r['AmountTWD'])
for r in invoices: add_summary(r['ProjectID'], 'InvoiceBalanceTWD', r['BalanceTWD'])
for r in expenses: add_summary(r['ProjectID'], 'ApprovedExpenseTWD', r['AmountTWD'])
for r in quotes: add_summary(r['ProjectID'], 'AcceptedQuoteTWD', r['AmountTWD'])

# Fill template copy
src_template = INPUT/'project_template.xlsx'
out_xlsx = OUTPUT/f'project_integration_{PERIOD}.xlsx'
shutil.copy2(src_template, out_xlsx)
wb = load_workbook(out_xlsx)

def clear_sheet_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row-start_row+1)

def append_dicts(ws, headers, rows):
    clear_sheet_data(ws)
    for row in rows:
        ws.append([row.get(h) for h in headers])
    # style IDs text and dates ISO strings as desired: keep dates as text? Template ok; preserve document IDs text.
    for c in range(1, len(headers)+1):
        ws.cell(4,c).font = Font(bold=True)

# Report
report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = 'SAP Sales Orders: OrderDate in period and Status != Cancelled; SAP Invoices: InvoiceDate in period; Expense Viewer: ExpenseDate in period and Status = Approved; quote_cli: QuoteDate in period and Status = Accepted; grouped by ProjectID; FX converted with input/fx_rates.csv; orders and quotes kept separate.'
now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B5'] = now
report['B7'] = 'Lab ERP SAP web; Expense Viewer desktop; quote_cli; fx_rates.csv'
report['D7'] = now

append_dicts(wb['Orders'], ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'], orders)
append_dicts(wb['Invoices'], ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], invoices)
append_dicts(wb['Expenses'], ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'], expenses)
append_dicts(wb['Quotes'], ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'], quotes)

# Summary
sum_ws = wb['Summary']
clear_sheet_data(sum_ws)
for project in sorted(summary):
    vals = summary[project]
    sum_ws.append([project, q2(vals['OrderTWD']), q2(vals['InvoiceBalanceTWD']), q2(vals['ApprovedExpenseTWD']), q2(vals['AcceptedQuoteTWD'])])
last = max(5, sum_ws.max_row)
sum_ws['B2'] = f'=SUM(B5:B{last})'
sum_ws['C2'] = f'=SUM(C5:C{last})'
sum_ws['D2'] = f'=SUM(D5:D{last})'
sum_ws['E2'] = f'=SUM(E5:E{last})'

# Exceptions
ex_ws = wb['Exceptions']
clear_sheet_data(ex_ws)
exceptions=[]
for source, idkey, rows in [('Orders','OrderID',orders), ('Invoices','InvoiceID',invoices), ('Expenses','ExpenseID',expenses), ('Quotes','QuoteID',quotes)]:
    for r in rows:
        if not r.get('ProjectID'):
            exceptions.append({'Source':source, 'DocumentID':str(r.get(idkey)), 'ProjectID':None, 'Reason':'Missing ProjectID; excluded from Summary'})
for e in exceptions:
    ex_ws.append([e['Source'], e['DocumentID'], e['ProjectID'], e['Reason']])

# formats: text IDs, numeric amounts
for sh, id_cols in {'Orders':[1,2,3], 'Invoices':[1,2,3,4], 'Expenses':[1,2,3], 'Quotes':[1,2,3], 'Summary':[1], 'Exceptions':[1,2,3]}.items():
    ws = wb[sh]
    for row in ws.iter_rows(min_row=5):
        for c in id_cols:
            row[c-1].number_format = '@'
for sh in ['Orders','Invoices','Expenses','Quotes','Summary']:
    ws=wb[sh]
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.value = float(cell.value)
            if isinstance(cell.value, (int,float)):
                cell.number_format = '#,##0.00'
wb.save(out_xlsx)
print('Saved', out_xlsx)

# Reopen and verify every stated requirement
wbv = load_workbook(out_xlsx, data_only=False)
assert set(['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']).issubset(wbv.sheetnames)
checks=[]
checks.append(('output workbook exists', out_xlsx.exists()))
checks.append(('source template copied with Instructions sheet', 'Instructions' in wbv.sheetnames))
checks.append(('Report from/to', wbv['Report']['B3'].value == FROM_DATE and wbv['Report']['D3'].value == TO_DATE))
# Headers unchanged
expected_headers={
 'Orders':['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
 'Invoices':['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
 'Expenses':['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
 'Quotes':['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
 'Summary':['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
 'Exceptions':['Source','DocumentID','ProjectID','Reason'],
}
for sh, heads in expected_headers.items():
    got=[wbv[sh].cell(4,i+1).value for i in range(len(heads))]
    checks.append((f'{sh} headers', got==heads))
# data row counts
counts={'Orders':len(orders),'Invoices':len(invoices),'Expenses':len(expenses),'Quotes':len(quotes),'Summary':len(summary),'Exceptions':len(exceptions)}
for sh,n in counts.items():
    actual=max(0, wbv[sh].max_row-4)
    checks.append((f'{sh} row count {actual}={n}', actual==n))
# status filters
orders_status=[wbv['Orders'].cell(r,5).value for r in range(5, wbv['Orders'].max_row+1)]
expenses_status=[wbv['Expenses'].cell(r,5).value for r in range(5, wbv['Expenses'].max_row+1)]
quotes_status=[wbv['Quotes'].cell(r,6).value for r in range(5, wbv['Quotes'].max_row+1)]
checks.append(('no Cancelled orders', all(s!='Cancelled' for s in orders_status)))
checks.append(('all expenses Approved', all(s=='Approved' for s in expenses_status)))
checks.append(('all quotes Accepted', all(s=='Accepted' for s in quotes_status)))
# Summary values match computed and no missing ProjectID
for idx, p in enumerate(sorted(summary), start=5):
    ws=wbv['Summary']
    row=(ws.cell(idx,1).value, Decimal(str(ws.cell(idx,2).value)), Decimal(str(ws.cell(idx,3).value)), Decimal(str(ws.cell(idx,4).value)), Decimal(str(ws.cell(idx,5).value)))
    vals=summary[p]
    checks.append((f'Summary {p}', row[0]==p and q2(row[1])==q2(vals['OrderTWD']) and q2(row[2])==q2(vals['InvoiceBalanceTWD']) and q2(row[3])==q2(vals['ApprovedExpenseTWD']) and q2(row[4])==q2(vals['AcceptedQuoteTWD'])))
checks.append(('separate summary columns', [wbv['Summary'].cell(4,i).value for i in range(2,6)] == ['OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']))
checks.append(('orders and quotes not combined as revenue', 'Revenue' not in [wbv['Summary'].cell(4,i).value for i in range(1, wbv['Summary'].max_column+1)]))
checks.append(('missing ProjectIDs excluded from Summary', all(wbv['Summary'].cell(r,1).value for r in range(5, wbv['Summary'].max_row+1))))
failed=[name for name, ok in checks if not ok]
for name,ok in checks: print('CHECK', 'OK' if ok else 'FAIL', '-', name)
if failed:
    raise AssertionError(f'Failed checks: {failed}')
print('Final files:')
for p in [DOWNLOADS/f'sap_orders_{PERIOD}.csv', DOWNLOADS/f'sap_invoices_{PERIOD}.csv', DOWNLOADS/f'expenses_{PERIOD}_approved.csv', quote_out, out_xlsx]:
    print(p.relative_to(WORKSPACE), p.stat().st_size)
print('Summary rows:')
for r in wbv['Summary'].iter_rows(min_row=5, values_only=True): print(r)
