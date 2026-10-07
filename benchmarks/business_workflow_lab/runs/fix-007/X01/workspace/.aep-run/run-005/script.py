import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace=', ws)
for sub in ['input','output','output/downloads']:
    p=ws/sub
    print('\nDIR', sub, 'exists=', p.exists())
    if p.exists():
        for f in sorted(p.iterdir()):
            print(' -', f.name, f.stat().st_size, 'bytes')
print('\nInput workbook sheets/headers:')
wbp=ws/'input'/'project_template.xlsx'
if wbp.exists():
    wb=load_workbook(wbp, data_only=False)
    print(wb.sheetnames)
    for sh in wb.sheetnames:
        s=wb[sh]
        print('SHEET', sh, 'max', s.max_row, s.max_column)
        for row in s.iter_rows(min_row=1, max_row=min(5,s.max_row), values_only=True):
            print(row)
print('\nFX sample:')
fx=ws/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as fh:
        for i,row in zip(range(5), csv.reader(fh)):
            print(row)
print('\nManual snippets:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=ws/'input'/name
    print('\n---', name, '---')
    if p.exists():
        txt=p.read_text(encoding='utf-8')
        print(txt[:2500])
print('\nExisting output workbook check:')
out=ws/'output'/'project_integration_2026-09.xlsx'
if out.exists():
    wb=load_workbook(out, data_only=True)
    print(wb.sheetnames)
    for sh in wb.sheetnames:
        s=wb[sh]
        print('SHEET', sh, 'max', s.max_row, s.max_column)
        for row in s.iter_rows(min_row=1, max_row=min(8,s.max_row), values_only=True):
            print(row)