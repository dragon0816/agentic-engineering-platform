from pathlib import Path
import csv, os
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, 'size', p.stat().st_size)
print('\nquote_cli_manual.md head:')
manual = ws/'input'/'quote_cli_manual.md'
print(manual.read_text(encoding='utf-8')[:2000])
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(6), r):
        print(row)
print('\nworkbook sheets / headers / sample dimensions:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for sh in wb.sheetnames:
    wsht=wb[sh]
    print(sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for rr in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,5), values_only=True):
        print(' ', rr)
print('\nworkspace candidates containing cli:')
for p in ws.rglob('*'):
    if 'cli' in p.name.lower() or 'quote' in p.name.lower():
        print(' ', p.relative_to(ws), 'dir' if p.is_dir() else 'file')