from pathlib import Path
from openpyxl import load_workbook
import subprocess, json, sys
ws=Path.cwd()
print('Instructions full rows:')
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
cli=r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
print('\nCLI list sample:')
res=subprocess.run([cli,'quotes','list','--customer','C001','--from','2026-09-01','--to','2026-09-30','--format','json'],capture_output=True,text=True)
print('returncode',res.returncode)
print('stdout',res.stdout[:2000])
print('stderr',res.stderr[:1000])
if res.returncode==0:
    data=json.loads(res.stdout)
    print('count',data.get('count'))
    if data.get('items'):
        qid=data['items'][0]['QuoteID']
        print('get first', qid)
        res2=subprocess.run([cli,'quotes','get',qid,'--format','json'],capture_output=True,text=True)
        print('returncode',res2.returncode)
        print('stdout',res2.stdout[:2000])
        print('stderr',res2.stderr[:500])