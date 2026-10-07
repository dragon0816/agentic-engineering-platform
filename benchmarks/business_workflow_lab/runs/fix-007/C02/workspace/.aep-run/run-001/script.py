from pathlib import Path
import csv, os
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
fx=ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv sample:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(6), r): print(row)
manual=ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\nquote_cli_manual.md first 4000 chars:')
    print(manual.read_text(encoding='utf-8', errors='replace')[:4000])
tpl=ws/'input'/'quote_template.xlsx'
if tpl.exists():
    wb=load_workbook(tpl, data_only=False)
    print('\nworkbook sheets:')
    for s in wb.sheetnames:
        sh=wb[s]
        print(s, sh.max_row, sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
            print(row)
        print('---')
