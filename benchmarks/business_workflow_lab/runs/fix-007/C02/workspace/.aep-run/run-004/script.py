from pathlib import Path
import subprocess, json, csv, sys, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from copy import copy
from collections import defaultdict
from openpyxl import load_workbook

WORKSPACE = Path.cwd(); INPUT = WORKSPACE/'input'; OUTPUT = WORKSPACE/'output'; OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT/'quote_template.xlsx'; FX_FILE = INPUT/'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
FROM_DATE='2026-09-01'; TO_DATE='2026-09-30'; BASE_DATE=date.fromisoformat('2026-09-30')
GENERATED_AT=datetime.now().astimezone().replace(microsecond=0).isoformat()
TWOPLACES=Decimal('0.01')
def D(x): return None if x is None or x=='' else Decimal(str(x))
def money(x): return D(x).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def run_cli(args):
    proc=subprocess.run([sys.executable, str(CLI)] + args, text=True, capture_output=True, cwd=str(WORKSPACE))
    if proc.returncode!=0: raise RuntimeError(f'CLI failed rc={proc.returncode} args={args}\nstdout={proc.stdout}\nstderr={proc.stderr}')
    return json.loads(proc.stdout)
def clear_data(ws,start,max_col):
    if ws.max_row>=start:
        for row in ws.iter_rows(min_row=start,max_row=ws.max_row,min_col=1,max_col=max_col):
            for cell in row: cell.value=None
def copy_row_style(ws,src_row,dst_row,max_col):
    for col in range(1,max_col+1):
        src=ws.cell(src_row,col); dst=ws.cell(dst_row,col)
        if src.has_style: dst._style=copy(src._style)
        dst.number_format=src.number_format
        dst.alignment=copy(src.alignment); dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border)

print('Workspace:', WORKSPACE)
print('Input files:', [p.name for p in sorted(INPUT.glob('*'))])
print('CLI query: quotes list --from', FROM_DATE, '--to', TO_DATE, '--format json')
fx={}
with FX_FILE.open(newline='',encoding='utf-8-sig') as f:
    reader=csv.reader(f); print('fx header:', next(reader))
    for row in reader:
        if len(row)>=2: fx[row[0]]=D(row[1])
print('fx rates:', {k:str(v) for k,v in fx.items()})
listing=run_cli(['quotes','list','--from',FROM_DATE,'--to',TO_DATE,'--format','json'])
print('listed count:', listing.get('count'), 'sample:', listing.get('items',[])[:3])
quotes=[]; all_lines=[]; findings=[]; expired_unaccepted=[]
for item in listing.get('items',[]):
    qid=item['QuoteID']; doc=run_cli(['quotes','get',qid,'--format','json']); h=doc['header']; lis=doc.get('line_items',[])
    quotes.append(h)
    for li in lis:
        r=dict(li); r['QuoteID']=qid; all_lines.append(r)
    line_sum=Decimal('0'); issue_parts=[]
    for i,li in enumerate(lis,1):
        expected=money(D(li['Quantity'])*D(li['UnitPrice'])); reported=money(li['LineAmount']); line_sum+=reported
        if reported!=expected: issue_parts.append(f'line {i} {li.get("ProductID")}: LineAmount {reported} != Quantity × UnitPrice {expected}')
    reported_subtotal=money(h['Subtotal']); expected_subtotal=money(line_sum)
    reported_net=money(h['NetAmount']); expected_net=money(reported_subtotal*(Decimal('1')-D(h.get('DiscountPercent') or '0')/Decimal('100')))
    if reported_subtotal!=expected_subtotal: issue_parts.append(f'Subtotal {reported_subtotal} != sum(LineAmount) {expected_subtotal}')
    if reported_net!=expected_net: issue_parts.append(f'NetAmount {reported_net} != Subtotal × (1 − DiscountPercent/100) {expected_net}')
    if issue_parts:
        findings.append({'QuoteID':qid,'message':'; '.join(issue_parts),'reported_subtotal':reported_subtotal,'expected_subtotal':expected_subtotal,'reported_net':reported_net,'expected_net':expected_net})
    if date.fromisoformat(h['ValidUntil']) < BASE_DATE and h.get('Status')!='Accepted': expired_unaccepted.append(h)
quotes.sort(key=lambda h:(h.get('QuoteDate',''),h.get('QuoteID',''))); all_lines.sort(key=lambda r:(r.get('QuoteID',''),r.get('ProductID','')))
print('retrieved detail quote rows:', len(quotes), 'line rows:', len(all_lines))
print('mismatch count:', len(findings), 'expired unaccepted:', [h['QuoteID'] for h in sorted(expired_unaccepted,key=lambda x:(x['ValidUntil'],x['QuoteID']))])

out_xlsx=OUTPUT/'quote_template.xlsx'; shutil.copy2(TEMPLATE,out_xlsx); wb=load_workbook(out_xlsx)
rep=wb['Report']; rep['B3']=FROM_DATE; rep['D3']=TO_DATE; rep['B4']='QuoteDate 2026-09-01..2026-09-30; all statuses; retain all quotes; expiry check base 2026-09-30'; rep['B5']=GENERATED_AT; rep['B7']='quote_cli; fx_rates.csv'; rep['D7']=GENERATED_AT; rep['B9']=f'Retrieved {len(quotes)} quotes and {len(all_lines)} line items. Amount mismatches: {len(findings)}. ValidUntil < 2026-09-30 and Status != Accepted: {len(expired_unaccepted)}.'
qws=wb['Quotes']; clear_data(qws,5,11)
for r,h in enumerate(quotes,5):
    copy_row_style(qws,5,r,11); ccy=h.get('Currency'); net=money(h['NetAmount']); rate=fx.get(ccy)
    vals=[h.get('QuoteID'), h.get('CustomerID'), h.get('ProjectID'), h.get('QuoteDate'), h.get('ValidUntil'), h.get('Status'), ccy, float(money(h['Subtotal'])), float(D(h['DiscountPercent'])), float(net), float(money(net*rate)) if rate is not None else None]
    for c,v in enumerate(vals,1):
        cell=qws.cell(r,c); cell.value=(str(v) if c in (1,2,3) and v is not None else v)
        if c in (1,2,3): cell.number_format='@'
        elif c in (8,10,11): cell.number_format='0.00'
        elif c==9: cell.number_format='0.##'
        elif c in (4,5): cell.number_format='yyyy-mm-dd'
if qws.auto_filter and qws.auto_filter.ref: qws.auto_filter.ref=f'A4:K{max(5,4+len(quotes))}'
lws=wb['Lines']; clear_data(lws,5,6)
for r,li in enumerate(all_lines,5):
    copy_row_style(lws,5,r,6); vals=[li.get('QuoteID'),li.get('ProductID'),li.get('ProductName'),float(D(li['Quantity'])),float(money(li['UnitPrice'])),float(money(li['LineAmount']))]
    for c,v in enumerate(vals,1):
        cell=lws.cell(r,c); cell.value=(str(v) if c in (1,2) and v is not None else v)
        if c in (1,2): cell.number_format='@'
        elif c in (5,6): cell.number_format='0.00'
if lws.auto_filter and lws.auto_filter.ref: lws.auto_filter.ref=f'A4:F{max(5,4+len(all_lines))}'
sws=wb['Summary']; clear_data(sws,5,6); last_q_row=max(5,4+len(quotes)); sws['B2']=f'=SUM(Quotes!K5:K{last_q_row})'
groups=defaultdict(lambda:{'amount':Decimal('0'),'amount_twd':Decimal('0'),'count':0})
for h in quotes:
    ccy=h.get('Currency'); net=money(h['NetAmount']); groups[(h.get('Status'),ccy)]['amount']+=net; groups[(h.get('Status'),ccy)]['amount_twd']+=money(net*fx[ccy]); groups[(h.get('Status'),ccy)]['count']+=1
for r,((status,ccy),data) in enumerate(sorted(groups.items()),5):
    copy_row_style(sws,5,r,6); vals=['Status',status,ccy,float(money(data['amount'])),float(money(data['amount_twd'])),data['count']]
    for c,v in enumerate(vals,1):
        cell=sws.cell(r,c); cell.value=v
        if c in (4,5): cell.number_format='0.00'
if sws.auto_filter and sws.auto_filter.ref: sws.auto_filter.ref=f'A4:F{max(5,4+len(groups))}'
wb.save(out_xlsx)

out_md=OUTPUT/'findings.md'; md=[]
md += ['# Findings for 2026-09 Quotes','',f'- Query period: `{FROM_DATE}` to `{TO_DATE}` inclusive (QuoteDate).',f'- Expiry baseline: `{BASE_DATE.isoformat()}`; flagged when `ValidUntil < {BASE_DATE.isoformat()}` and `Status != Accepted`.',f'- Quotes retained in workbook: `{len(quotes)}` (none removed due to expiry).',f'- Line items retained in workbook: `{len(all_lines)}`.','', '## Amount validation','']
if findings:
    md += ['| QuoteID | Issue | Reported Subtotal | Expected Subtotal | Reported NetAmount | Expected NetAmount |','|---|---|---:|---:|---:|---:|']
    for f in findings: md.append(f"| {f['QuoteID']} | {f['message'].replace('|','\\|')} | {f['reported_subtotal']} | {f['expected_subtotal']} | {f['reported_net']} | {f['expected_net']} |")
else:
    md.append('No amount mismatches found. For each quote, line amounts matched Quantity × UnitPrice, Subtotal matched sum(LineAmount), and NetAmount matched the discounted subtotal rounded HALF_UP to two decimals.')
md += ['', '## ValidUntil before 2026-09-30 and not Accepted','']
if expired_unaccepted:
    md += ['| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |','|---|---|---|---|---|---|---|---:|']
    for h in sorted(expired_unaccepted,key=lambda x:(x['ValidUntil'],x['QuoteID'])):
        md.append(f"| {h.get('QuoteID')} | {h.get('CustomerID') or ''} | {h.get('ProjectID') or ''} | {h.get('QuoteDate')} | {h.get('ValidUntil')} | {h.get('Status')} | {h.get('Currency')} | {money(h.get('NetAmount'))} |")
else: md.append('No quotes matched this condition.')
md += ['', '## Outputs','', '- `quote_template.xlsx` filled from the provided template.', '- `findings.md` generated by this run.']
out_md.write_text('\n'.join(md)+'\n',encoding='utf-8')

# Verification from saved outputs
wb2=load_workbook(out_xlsx,data_only=False); q2=wb2['Quotes']; l2=wb2['Lines']; s2=wb2['Summary']; r2=wb2['Report']
expected_q=['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']; expected_l=['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [q2.cell(4,c).value for c in range(1,12)]==expected_q
assert [l2.cell(4,c).value for c in range(1,7)]==expected_l
assert r2['B3'].value==FROM_DATE and r2['D3'].value==TO_DATE
q_rows=[row for row in q2.iter_rows(min_row=5,max_col=11,values_only=True) if row[0] is not None]
l_rows=[row for row in l2.iter_rows(min_row=5,max_col=6,values_only=True) if row[0] is not None]
assert len(q_rows)==len(quotes); assert len(l_rows)==len(all_lines)
assert all((row[0] is None or isinstance(row[0],str)) and (row[1] is None or isinstance(row[1],str)) and (row[2] is None or isinstance(row[2],str)) for row in q_rows)
assert all(isinstance(row[7],(int,float)) and isinstance(row[8],(int,float)) and isinstance(row[9],(int,float)) and isinstance(row[10],(int,float)) for row in q_rows)
assert all((row[0] is None or isinstance(row[0],str)) and (row[1] is None or isinstance(row[1],str)) for row in l_rows)
assert all(isinstance(row[3],(int,float)) and isinstance(row[4],(int,float)) and isinstance(row[5],(int,float)) for row in l_rows)
assert s2['B2'].value==f'=SUM(Quotes!K5:K{last_q_row})'
md_text=out_md.read_text(encoding='utf-8'); assert 'Amount validation' in md_text and 'ValidUntil before 2026-09-30' in md_text
# Verify no quote got removed: all listed IDs present in workbook.
assert sorted(row[0] for row in q_rows)==sorted(h['QuoteID'] for h in quotes)
print('\nVerification passed')
print('output workbook:', out_xlsx)
print('output findings:', out_md)
print('quote rows:', len(q_rows), 'line rows:', len(l_rows))
print('headers preserved:', expected_q, expected_l)
print('summary formula:', s2['B2'].value)
print('amount mismatches:', len(findings))
print('expired unaccepted count:', len(expired_unaccepted), 'ids:', [h['QuoteID'] for h in sorted(expired_unaccepted,key=lambda x:(x['ValidUntil'],x['QuoteID']))])
