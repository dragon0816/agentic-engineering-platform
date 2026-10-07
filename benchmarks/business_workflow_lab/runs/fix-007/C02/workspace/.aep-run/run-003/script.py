from pathlib import Path
import subprocess, json, csv, sys, os, shutil
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from datetime import datetime, date
from copy import copy
from collections import defaultdict
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
GENERATED_AT = datetime.now().astimezone().replace(microsecond=0).isoformat()

print('Workspace:', WORKSPACE)
print('Inputs present:', [p.name for p in sorted(INPUT.glob('*'))])
print('Using CLI:', CLI)
print('Query:', FROM_DATE, 'to', TO_DATE, 'all statuses')

TWOPLACES = Decimal('0.01')
def D(x):
    if x is None or x == '':
        return None
    return Decimal(str(x))

def money(x):
    return D(x).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [sys.executable, str(CLI)] + args
    proc = subprocess.run(cmd, text=True, capture_output=True, cwd=str(WORKSPACE))
    if proc.returncode != 0:
        raise RuntimeError(f'CLI failed rc={proc.returncode}\ncmd={cmd}\nstdout={proc.stdout}\nstderr={proc.stderr}')
    try:
        return json.loads(proc.stdout)
    except json.JSONDecodeError as e:
        raise RuntimeError(f'CLI returned non-JSON stdout for {cmd}: {proc.stdout[:500]}') from e

# Read fx rates; only first two columns are needed because Note contains commas in this synthetic file.
fx = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if not row or len(row) < 2:
            continue
        fx[row[0]] = D(row[1])
print('fx rates:', {k: str(v) for k,v in fx.items()})

# Fetch quotes
listing = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json'])
items = listing.get('items', [])
print('listed quote count:', len(items), 'declared:', listing.get('count'))
print('first listed rows:', items[:3])

quotes = []
lines = []
findings = []
expired_unaccepted = []
for item in items:
    qid = item['QuoteID']
    doc = run_cli(['quotes', 'get', qid, '--format', 'json'])
    header = doc['header']
    line_items = doc.get('line_items', [])
    quotes.append(header)
    for li in line_items:
        row = dict(li)
        row['QuoteID'] = qid
        lines.append(row)

    # Validate line amounts and subtotal/net.
    line_sum = Decimal('0')
    line_issues = []
    for idx, li in enumerate(line_items, start=1):
        qty = D(li.get('Quantity'))
        unit = D(li.get('UnitPrice'))
        reported_line = money(li.get('LineAmount'))
        expected_line = money(qty * unit)
        line_sum += reported_line
        if reported_line != expected_line:
            line_issues.append(f'line {idx} {li.get("ProductID")}: LineAmount {reported_line} != Quantity × UnitPrice {expected_line}')

    reported_subtotal = money(header.get('Subtotal'))
    expected_subtotal = money(line_sum)
    discount_pct = D(header.get('DiscountPercent') or '0')
    reported_net = money(header.get('NetAmount'))
    expected_net = money(reported_subtotal * (Decimal('1') - discount_pct / Decimal('100')))

    quote_issue_parts = []
    if line_issues:
        quote_issue_parts.extend(line_issues)
    if reported_subtotal != expected_subtotal:
        quote_issue_parts.append(f'Subtotal {reported_subtotal} != sum(LineAmount) {expected_subtotal}')
    if reported_net != expected_net:
        quote_issue_parts.append(f'NetAmount {reported_net} != Subtotal × (1 − DiscountPercent/100) {expected_net}')
    if quote_issue_parts:
        findings.append({
            'QuoteID': qid,
            'type': 'amount_mismatch',
            'message': '; '.join(quote_issue_parts),
            'reported_subtotal': reported_subtotal,
            'expected_subtotal': expected_subtotal,
            'reported_net': reported_net,
            'expected_net': expected_net,
        })

    valid_until = date.fromisoformat(header['ValidUntil'])
    if valid_until < BASE_DATE and header.get('Status') != 'Accepted':
        expired_unaccepted.append(header)

quotes.sort(key=lambda h: (h.get('QuoteDate',''), h.get('QuoteID','')))
lines.sort(key=lambda r: (r.get('QuoteID',''), r.get('ProductID','')))
print('retrieved quote details:', len(quotes), 'line rows:', len(lines))
print('amount mismatch count:', len(findings))
print('expired and unaccepted count:', len(expired_unaccepted))

# Prepare workbook copy
out_xlsx = OUTPUT / 'quote_template.xlsx'
shutil.copy2(TEMPLATE, out_xlsx)
wb = load_workbook(out_xlsx)

# Helper functions

def clear_data(ws, start_row, max_col):
    if ws.max_row >= start_row:
        for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, min_col=1, max_col=max_col):
            for cell in row:
                cell.value = None

def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

# Report
rep = wb['Report']
rep['B3'] = FROM_DATE
rep['D3'] = TO_DATE
rep['B4'] = 'QuoteDate 2026-09-01..2026-09-30; all statuses; retain all quotes; expiry check base 2026-09-30'
rep['B5'] = GENERATED_AT
rep['B7'] = 'quote_cli; fx_rates.csv'
rep['D7'] = GENERATED_AT
rep['B9'] = f'Retrieved {len(quotes)} quotes and {len(lines)} line items. Amount mismatches: {len(findings)}. ValidUntil < 2026-09-30 and Status != Accepted: {len(expired_unaccepted)}.'

# Quotes sheet
qws = wb['Quotes']
clear_data(qws, 5, 11)
q_headers = [qws.cell(4, c).value for c in range(1,12)]
for r_idx, h in enumerate(quotes, start=5):
    copy_row_style(qws, 5, r_idx, 11)
    ccy = h.get('Currency')
    rate = fx.get(ccy)
    net = money(h.get('NetAmount'))
    values = [
        str(h.get('QuoteID')) if h.get('QuoteID') is not None else None,
        str(h.get('CustomerID')) if h.get('CustomerID') is not None else None,
        str(h.get('ProjectID')) if h.get('ProjectID') is not None else None,
        h.get('QuoteDate'), h.get('ValidUntil'), h.get('Status'), ccy,
        float(money(h.get('Subtotal'))),
        float(D(h.get('DiscountPercent'))),
        float(net),
        float(money(net * rate)) if rate is not None else None,
    ]
    for c_idx, val in enumerate(values, start=1):
        cell = qws.cell(r_idx, c_idx)
        cell.value = val
        if c_idx in (1,2,3):
            cell.number_format = '@'
        elif c_idx in (8,10,11):
            cell.number_format = '0.00'
        elif c_idx == 9:
            cell.number_format = '0.##'
        elif c_idx in (4,5):
            cell.number_format = 'yyyy-mm-dd'
if qws.auto_filter and qws.auto_filter.ref:
    qws.auto_filter.ref = f'A4:K{max(5, 4+len(quotes))}'

# Lines sheet
lws = wb['Lines']
clear_data(lws, 5, 6)
for r_idx, li in enumerate(lines, start=5):
    copy_row_style(lws, 5, r_idx, 6)
    vals = [str(li.get('QuoteID')), str(li.get('ProductID')), li.get('ProductName'),
            float(D(li.get('Quantity'))), float(money(li.get('UnitPrice'))), float(money(li.get('LineAmount')))]
    for c_idx, val in enumerate(vals, start=1):
        cell = lws.cell(r_idx, c_idx)
        cell.value = val
        if c_idx in (1,2):
            cell.number_format = '@'
        elif c_idx in (5,6):
            cell.number_format = '0.00'
if lws.auto_filter and lws.auto_filter.ref:
    lws.auto_filter.ref = f'A4:F{max(5, 4+len(lines))}'

# Summary by Status and Currency, amounts are NetAmount per currency and converted TWD.
sws = wb['Summary']
clear_data(sws, 5, 6)
last_q_row = max(5, 4+len(quotes))
sws['B2'] = f'=SUM(Quotes!K5:K{last_q_row})'
groups = defaultdict(lambda: {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
for h in quotes:
    ccy = h.get('Currency')
    status = h.get('Status')
    net = money(h.get('NetAmount'))
    rate = fx.get(ccy)
    key = (status, ccy)
    groups[key]['amount'] += net
    groups[key]['amount_twd'] += money(net * rate) if rate is not None else Decimal('0')
    groups[key]['count'] += 1
for r_idx, ((status, ccy), data) in enumerate(sorted(groups.items()), start=5):
    copy_row_style(sws, 5, r_idx, 6)
    vals = ['Status', status, ccy, float(money(data['amount'])), float(money(data['amount_twd'])), data['count']]
    for c_idx, val in enumerate(vals, start=1):
        cell = sws.cell(r_idx, c_idx)
        cell.value = val
        if c_idx in (4,5):
            cell.number_format = '0.00'
if sws.auto_filter and sws.auto_filter.ref:
    sws.auto_filter.ref = f'A4:F{max(5, 4+len(groups))}'

wb.save(out_xlsx)

# Findings markdown
out_md = OUTPUT / 'findings.md'
lines_md = []
lines_md.append('# Findings for 2026-09 Quotes')
lines_md.append('')
lines_md.append(f'- Query period: `{FROM_DATE}` to `{TO_DATE}` inclusive (QuoteDate).')
lines_md.append(f'- Expiry baseline: `{BASE_DATE.isoformat()}`; flagged when `ValidUntil < {BASE_DATE.isoformat()}` and `Status != Accepted`.')
lines_md.append(f'- Quotes retained in workbook: `{len(quotes)}` (none removed due to expiry).')
lines_md.append(f'- Line items retained in workbook: `{len(lines)}`.')
lines_md.append('')
lines_md.append('## Amount validation')
lines_md.append('')
if findings:
    lines_md.append('| QuoteID | Issue | Reported Subtotal | Expected Subtotal | Reported NetAmount | Expected NetAmount |')
    lines_md.append('|---|---|---:|---:|---:|---:|')
    for f in findings:
        msg = str(f['message']).replace('|','\\|')
        lines_md.append(f"| {f['QuoteID']} | {msg} | {f['reported_subtotal']} | {f['expected_subtotal']} | {f['reported_net']} | {f['expected_net']} |")
else:
    lines_md.append('No amount mismatches found. For each quote, line amounts matched Quantity × UnitPrice, Subtotal matched sum(LineAmount), and NetAmount matched the discounted subtotal rounded HALF_UP to two decimals.')
lines_md.append('')
lines_md.append('## ValidUntil before 2026-09-30 and not Accepted')
lines_md.append('')
if expired_unaccepted:
    lines_md.append('| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |')
    lines_md.append('|---|---|---|---|---|---|---|---:|')
    for h in sorted(expired_unaccepted, key=lambda x: (x['ValidUntil'], x['QuoteID'])):
        lines_md.append(f"| {h.get('QuoteID')} | {h.get('CustomerID')} | {h.get('ProjectID')} | {h.get('QuoteDate')} | {h.get('ValidUntil')} | {h.get('Status')} | {h.get('Currency')} | {money(h.get('NetAmount'))} |")
else:
    lines_md.append('No quotes matched this condition.')
lines_md.append('')
lines_md.append('## Outputs')
lines_md.append('')
lines_md.append('- `quote_template.xlsx` filled from the provided template.')
lines_md.append('- `findings.md` generated by this run.')
out_md.write_text('\n'.join(lines_md) + '\n', encoding='utf-8')

# Reopen and verify requirements
print('\nVerification:')
wb2 = load_workbook(out_xlsx, data_only=False)
q2 = wb2['Quotes']; l2 = wb2['Lines']; s2 = wb2['Summary']; r2 = wb2['Report']
q_header2 = [q2.cell(4,c).value for c in range(1,12)]
l_header2 = [l2.cell(4,c).value for c in range(1,7)]
expected_q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
expected_l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert q_header2 == expected_q_headers, q_header2
assert l_header2 == expected_l_headers, l_header2
assert r2['B3'].value == FROM_DATE and r2['D3'].value == TO_DATE
# Count non-empty QuoteID rows
q_rows = [row for row in q2.iter_rows(min_row=5, max_col=11, values_only=True) if row[0] is not None]
l_rows = [row for row in l2.iter_rows(min_row=5, max_col=6, values_only=True) if row[0] is not None]
assert len(q_rows) == len(quotes), (len(q_rows), len(quotes))
assert len(l_rows) == len(lines), (len(l_rows), len(lines))
# IDs text-like strings, money numeric, no expired rows removed
assert all(isinstance(row[0], str) for row in q_rows)
assert all(isinstance(row[1], str) for row in q_rows)
assert all(isinstance(row[2], str) for row in q_rows)
assert all(isinstance(row[7], (int,float)) and isinstance(row[9], (int,float)) and isinstance(row[10], (int,float)) for row in q_rows)
assert all(isinstance(row[0], str) and isinstance(row[1], str) for row in l_rows)
assert all(isinstance(row[4], (int,float)) and isinstance(row[5], (int,float)) for row in l_rows)
assert s2['B2'].value == f'=SUM(Quotes!K5:K{last_q_row})'
md_text = out_md.read_text(encoding='utf-8')
assert 'Amount validation' in md_text and 'ValidUntil before 2026-09-30' in md_text
print('output workbook:', out_xlsx)
print('output findings:', out_md)
print('quote rows checked:', len(q_rows))
print('line rows checked:', len(l_rows))
print('summary formula checked:', s2['B2'].value)
print('quote headers checked:', q_header2)
print('line headers checked:', l_header2)
print('findings amount mismatches:', len(findings))
print('expired_unaccepted flagged:', len(expired_unaccepted), [h['QuoteID'] for h in sorted(expired_unaccepted, key=lambda x: (x['ValidUntil'], x['QuoteID']))])
