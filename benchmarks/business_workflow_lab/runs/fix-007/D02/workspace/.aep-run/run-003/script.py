from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in sh.iter_rows(values_only=True):
    print(r)
print('\nSubmitted all rows')
for name in ['submitted_2026_09.csv','approved_2026_09.csv']:
    print('---', name)
    with open(ws/'output'/name, newline='', encoding='utf-8-sig') as f:
        reader=csv.DictReader(f)
        print(reader.fieldnames)
        for row in reader:
            print(row)
