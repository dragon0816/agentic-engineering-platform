from pathlib import Path
import subprocess, sys, json, csv, shutil, os, traceback
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
OUT_XLSX = OUTPUT / 'sent_quotes_2026-09.xlsx'
ERR_TXT = OUTPUT / 'quote_cli_error.txt'
RUN_LOG = OUTPUT / 'quote_cli_run_log.json'

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'

attempt_log = []

def run_cli(args, parse_json=True):
    """Run CLI. Retry exit code 4 up to two retries. Do not retry exit 2/3. Only parse stdout on exit 0."""
    max_attempts = 3  # initial + two retries for exit 4
    for attempt in range(1, max_attempts + 1):
        cmd = [sys.executable, str(CLI)] + args
        p = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
        rec = {
            'args': args,
            'attempt': attempt,
            'exit_code': p.returncode,
            'stderr': p.stderr,
            'stdout_preview': p.stdout[:500] if p.returncode == 0 else ''
        }
        attempt_log.append(rec)
        if p.returncode == 0:
            if parse_json:
                try:
                    return json.loads(p.stdout)
                except Exception as e:
                    raise RuntimeError(f'CLI success but stdout was not parseable JSON for args {args}: {e}')
            return p.stdout
        if p.returncode in (2, 3):
            raise RuntimeError(f'CLI exit code {p.returncode}; not retried. stderr: {p.stderr.strip()}')
        if p.returncode == 4:
            if attempt < max_attempts:
                continue
            raise RuntimeError(f'CLI exit code 4 after {max_attempts} attempts. Last stderr: {p.stderr.strip()}')
        raise RuntimeError(f'CLI unexpected exit code {p.returncode}. stderr: {p.stderr.strip()}')

def dec(x):
    if x is None or x == '':
        return None
    return Decimal(str(x))

def money(x):
    return dec(x).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def read_fx():
    rates = {}
    with FX.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        header = next(r)
        for row in r:
            if not row:
                continue
            rates[row[0]] = Decimal(row[1])
    return rates

def set_text(cell):
    cell.number_format = '@'

def set_money(cell):
    cell.number_format = '0.00'

def set_date(cell):
    cell.number_format = 'yyyy-mm-dd'

def clear_data(ws, start_row):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def fill_success(items, details):
    rates = read_fx()
    shutil.copy2(TEMPLATE, OUT_XLSX)
    wb = load_workbook(OUT_XLSX)
    now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()

    report = wb['Report']
    report['B3'] = FROM_DATE
    report['D3'] = TO_DATE
    report['B4'] = f'Status={STATUS}; QuoteDate {FROM_DATE}..{TO_DATE} (inclusive)'
    report['B5'] = now
    report['B7'] = 'quote_cli; fx_rates.csv'
    report['D7'] = now
    report['B9'] = ''

    qws = wb['Quotes']
    lws = wb['Lines']
    sws = wb['Summary']
    clear_data(qws, 5)
    clear_data(lws, 5)
    clear_data(sws, 5)

    quote_rows = []
    for it in items:
        cur = it.get('Currency')
        if cur not in rates:
            raise RuntimeError(f'Missing FX rate for currency {cur}')
        net = money(it.get('NetAmount'))
        net_twd = money(net * rates[cur])
        quote_rows.append({**it, 'NetAmountTWD': net_twd})

    for idx, q in enumerate(quote_rows, start=5):
        vals = [q.get('QuoteID'), q.get('CustomerID'), q.get('ProjectID'), q.get('QuoteDate'), q.get('ValidUntil'), q.get('Status'), q.get('Currency'), dec(q.get('Subtotal')), dec(q.get('DiscountPercent')), dec(q.get('NetAmount')), q.get('NetAmountTWD')]
        for col, val in enumerate(vals, start=1):
            cell = qws.cell(idx, col, val)
            if col in (1,2,3): set_text(cell)
            if col in (4,5): set_date(cell)
            if col in (8,10,11): set_money(cell)

    line_count = 0
    for qid, detail in details.items():
        for line in detail.get('line_items', []):
            line_count += 1
            vals = [qid, line.get('ProductID'), line.get('ProductName'), dec(line.get('Quantity')), dec(line.get('UnitPrice')), dec(line.get('LineAmount'))]
            for col, val in enumerate(vals, start=1):
                cell = lws.cell(line_count + 4, col, val)
                if col in (1,2): set_text(cell)
                if col in (5,6): set_money(cell)

    last_q = 4 + len(quote_rows)
    if quote_rows:
        sws['B2'] = f'=SUM(Quotes!K5:K{last_q})'
    else:
        sws['B2'] = '=SUM(Quotes!K5:K5)'

    # Summary: by Status and Currency, preserving original-currency totals separately.
    groups = {}
    for q in quote_rows:
        key = (q.get('Status'), q.get('Currency'))
        groups.setdefault(key, {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
        groups[key]['amount'] += dec(q.get('NetAmount'))
        groups[key]['amount_twd'] += q.get('NetAmountTWD')
        groups[key]['count'] += 1
    for row_idx, ((status, cur), agg) in enumerate(sorted(groups.items()), start=5):
        vals = ['Status', status, cur, money(agg['amount']), money(agg['amount_twd']), agg['count']]
        for col, val in enumerate(vals, start=1):
            cell = sws.cell(row_idx, col, val)
            if col in (4,5): set_money(cell)

    wb.save(OUT_XLSX)
    return {'workbook': str(OUT_XLSX), 'quotes': len(quote_rows), 'lines': line_count, 'summary_rows': len(groups)}

try:
    list_data = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'])
    if list_data.get('schema_version') != 1 or 'items' not in list_data:
        raise RuntimeError('CLI list JSON schema was not the expected schema_version 1 with items')
    items = list_data['items']
    details = {}
    for item in items:
        qid = item.get('QuoteID')
        detail = run_cli(['quotes', 'get', qid, '--format', 'json'])
        if detail.get('schema_version') != 1 or 'line_items' not in detail:
            raise RuntimeError(f'CLI get JSON schema was not expected for {qid}')
        details[qid] = detail
    result = fill_success(items, details)
    status = 'success'
    error = None
except Exception as e:
    status = 'error'
    error = str(e)
    # Do not create a filled workbook when CLI failed or stdout is unusable.
    ERR_TXT.write_text(error + '\n', encoding='utf-8')
    result = {'error_file': str(ERR_TXT)}
finally:
    RUN_LOG.write_text(json.dumps({'status': status, 'error': error, 'attempts': attempt_log, 'result': result}, ensure_ascii=False, indent=2), encoding='utf-8')

print(json.dumps({'status': status, 'error': error, 'result': result, 'run_log': str(RUN_LOG)}, ensure_ascii=False, indent=2))

# Verification phase: reopen outputs and confirm requirements that can be met.
print('\nVERIFICATION')
log = json.loads(RUN_LOG.read_text(encoding='utf-8'))
print('run log exists:', RUN_LOG.exists())
for a in log['attempts']:
    print('attempt:', a['args'], 'try', a['attempt'], 'exit', a['exit_code'], 'stderr=', a['stderr'].strip())
if status == 'success':
    wb2 = load_workbook(OUT_XLSX, data_only=False)
    print('workbook exists:', OUT_XLSX.exists())
    print('sheets:', wb2.sheetnames)
    print('Report B3/D3/B4:', wb2['Report']['B3'].value, wb2['Report']['D3'].value, wb2['Report']['B4'].value)
    print('Quotes headers:', [wb2['Quotes'].cell(4,c).value for c in range(1,12)])
    print('Lines headers:', [wb2['Lines'].cell(4,c).value for c in range(1,7)])
    print('Summary formula B2:', wb2['Summary']['B2'].value)
    print('row counts quotes/lines/summary:', result['quotes'], result['lines'], result['summary_rows'])
    # Value type checks
    if result['quotes']:
        print('QuoteID number_format:', wb2['Quotes']['A5'].number_format)
        print('NetAmountTWD format:', wb2['Quotes']['K5'].number_format)
else:
    print('error file exists:', ERR_TXT.exists())
    print('error text:', ERR_TXT.read_text(encoding='utf-8').strip())
    if OUT_XLSX.exists():
        print('note: workbook exists from earlier run:', OUT_XLSX)
