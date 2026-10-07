from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
print('\nReport non-empty cells:')
for row in wb['Report'].iter_rows():
    vals=[]
    for c in row:
        if c.value is not None:
            vals.append(f'{c.coordinate}={c.value!r}')
    if vals:
        print('; '.join(vals))
print('\nSummary rows:')
for row in wb['Summary'].iter_rows(values_only=False):
    vals=[]
    for c in row:
        if c.value is not None:
            vals.append(f'{c.coordinate}={c.value!r}')
    if vals:
        print('; '.join(vals))