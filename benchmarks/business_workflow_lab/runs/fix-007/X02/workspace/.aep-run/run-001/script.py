from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)

# workbook structure
wb_path = ws/'input'/'project_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('WORKBOOK SHEETS:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'SHEET {s}: max_row={sh.max_row} max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(' ', r)

# fx headers/sample
fx = ws/'input'/'fx_rates.csv'
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    rows=[]
    for i,row in enumerate(reader):
        rows.append(row)
        if i>=5: break
print('FX SAMPLE:')
for r in rows: print(r)

for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=ws/'input'/name
    print('\nMANUAL', name)
    txt=p.read_text(encoding='utf-8', errors='replace')
    print(txt[:3000])