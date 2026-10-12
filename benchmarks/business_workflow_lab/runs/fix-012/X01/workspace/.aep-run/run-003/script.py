import os, json, csv, subprocess, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from zoneinfo import ZoneInfo
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
FROM = params['from_date']
TO = params['to_date']
YM = params.get('year_month') or FROM[:7]
TODAY = params.get('today') or datetime.now(ZoneInfo('Asia/Taipei')).date().isoformat()
ws = Path.cwd()
inp = ws/'input'
out = ws/'output'
dl = out/'downloads'
out.mkdir(exist_ok=True)
dl.mkdir(exist_ok=True)

TWOPLACES = Decimal('0.01')
def D(x):
    if x is None or str(x).strip() == '':
        return Decimal('0')
    return Decimal(str(x).strip())
def q2(x):
    return D(x).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def money(x):
    return float(q2(x))
def in_range(s):
    return FROM <= str(s) <= TO
def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

# Read FX rates. Currency and RateToTWD are the first two columns.
fx = {}
with open(inp/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    next(r, None)
    for row in r:
        if len(row) >= 2 and row[0]:
            fx[row[0].strip()] = D(row[1])
def to_twd(amount, cur):
    cur = (cur or '').strip()
    if cur not in fx:
        raise ValueError(f'Missing FX rate for currency {cur!r}')
    return q2(D(amount) * fx[cur])

# Browser/desktop steps use stable export names; this step copies them to dated deliverable names.
orders_src = dl/'sap_orders.csv'
invoices_src = dl/'sap_invoices.csv'
expenses_src = out/'expenses_approved.csv'
orders_path = out/f'sap_orders_{YM}.csv'
invoices_path = out/f'sap_invoices_{YM}.csv'
expenses_path = out/f'expenses_{YM}_approved.csv'
for src, dst in [(orders_src, orders_path), (invoices_src, invoices_path), (expenses_src, expenses_path)]:
    if not src.exists():
        raise FileNotFoundError(str(src))
    shutil.copyfile(src, dst)
for p in [orders_path, invoices_path, expenses_path, inp/'project_template.xlsx', inp/'fx_rates.csv']:
    if not p.exists():
        raise FileNotFoundError(str(p))

orders_raw = read_csv(orders_path)
invoices_raw = read_csv(invoices_path)
expenses_raw = read_csv(expenses_path)
orders = [r for r in orders_raw if in_range(r.get('OrderDate','')) and r.get('Status') != 'Cancelled']
invoices = [r for r in invoices_raw if in_range(r.get('InvoiceDate',''))]
expenses = [r for r in expenses_raw if in_range(r.get('ExpenseDate','')) and r.get('Status') == 'Approved']

# Pull accepted quotes from the quote CLI.
cli_candidates = [
    ws/'quote_cli.py',
    inp/'quote_cli.py',
    ws.parent/'quote_cli.py',
    Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py'),
]
cli = next((p for p in cli_candidates if p.exists()), None)
if cli is None:
    raise FileNotFoundError('quote CLI not found in expected locations')
cmd = ['python', str(cli), 'quotes', 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json']
proc = subprocess.run(cmd, text=True, capture_output=True)
if proc.returncode != 0:
    raise RuntimeError(f'quote_cli failed rc={proc.returncode}: {proc.stderr}')
quote_payload = json.loads(proc.stdout)
quotes = [r for r in quote_payload.get('items', []) if in_range(r.get('QuoteDate','')) and r.get('Status') == 'Accepted']
quotes_json = out/f'quotes_{YM}_accepted.json'
quotes_json.write_text(json.dumps(quote_payload, ensure_ascii=False, indent=2), encoding='utf-8')

out_xlsx = out/f'project_integration_{YM}.xlsx'
shutil.copyfile(inp/'project_template.xlsx', out_xlsx)
wb = load_workbook(out_xlsx)
def clear_data(sheet):
    wsht = wb[sheet]
    if wsht.max_row > 4:
        wsht.delete_rows(5, wsht.max_row - 4)
def set_text(c): c.number_format = '@'
def set_amt(c): c.number_format = '0.00'
def write_rows(sheet, rows):
    clear_data(sheet)
    wsht = wb[sheet]
    for ridx, vals in enumerate(rows, 5):
        for cidx, val in enumerate(vals, 1):
            cell = wsht.cell(ridx, cidx, val)
            header = wsht.cell(4, cidx).value
            if header and ('ID' in str(header) or header in ('Status','Currency')):
                set_text(cell)
            if header and any(k in str(header) for k in ['Amount','Balance','TWD','PaidAmount','NetAmount']):
                set_amt(cell)

summary = {}
exceptions = []
def addsum(project, key, amount):
    if not project:
        return
    summary.setdefault(project, {'OrderTWD':Decimal('0'), 'InvoiceBalanceTWD':Decimal('0'), 'ApprovedExpenseTWD':Decimal('0'), 'AcceptedQuoteTWD':Decimal('0')})
    summary[project][key] += q2(amount)

order_rows=[]
for r in orders:
    twd = to_twd(r['NetAmount'], r['Currency'])
    order_rows.append([r.get('OrderID'), r.get('ProjectID') or None, r.get('CustomerID') or None, r.get('OrderDate'), r.get('Status'), r.get('Currency'), money(r.get('NetAmount')), money(twd)])
    if r.get('ProjectID'): addsum(r['ProjectID'], 'OrderTWD', twd)
    else: exceptions.append(['Orders', r.get('OrderID'), None, 'Missing ProjectID; excluded from Summary'])

invoice_rows=[]
today_d = date.fromisoformat(TODAY)
for r in invoices:
    bal = q2(D(r.get('Balance')) if r.get('Balance') not in (None,'') else D(r.get('Amount')) - D(r.get('PaidAmount')))
    twd = to_twd(bal, r['Currency'])
    invoice_rows.append([r.get('InvoiceID'), r.get('OrderID'), r.get('ProjectID') or None, r.get('CustomerID') or None, r.get('InvoiceDate'), r.get('DueDate'), r.get('Currency'), money(r.get('Amount')), money(r.get('PaidAmount')), money(bal), money(twd)])
    if r.get('ProjectID'): addsum(r['ProjectID'], 'InvoiceBalanceTWD', twd)
    else: exceptions.append(['Invoices', r.get('InvoiceID'), None, 'Missing ProjectID; excluded from Summary'])
    try:
        if bal > 0 and date.fromisoformat(r.get('DueDate')) < today_d:
            exceptions.append(['Invoices', r.get('InvoiceID'), r.get('ProjectID') or None, 'Overdue unpaid balance'])
    except Exception:
        pass

expense_rows=[]
for r in expenses:
    twd = to_twd(r['Amount'], r['Currency'])
    expense_rows.append([r.get('ExpenseID'), r.get('ProjectID') or None, r.get('EmployeeID') or None, r.get('ExpenseDate'), r.get('Status'), r.get('Currency'), money(r.get('Amount')), money(twd)])
    if r.get('ProjectID'): addsum(r['ProjectID'], 'ApprovedExpenseTWD', twd)
    else: exceptions.append(['Expenses', r.get('ExpenseID'), None, 'Missing ProjectID; excluded from Summary'])

quote_rows=[]
for r in quotes:
    twd = to_twd(r['NetAmount'], r['Currency'])
    quote_rows.append([r.get('QuoteID'), r.get('ProjectID') or None, r.get('CustomerID') or None, r.get('QuoteDate'), r.get('ValidUntil'), r.get('Status'), r.get('Currency'), money(r.get('NetAmount')), money(twd)])
    if r.get('ProjectID'): addsum(r['ProjectID'], 'AcceptedQuoteTWD', twd)
    else: exceptions.append(['Quotes', r.get('QuoteID'), None, 'Missing ProjectID; excluded from Summary'])

write_rows('Orders', order_rows)
write_rows('Invoices', invoice_rows)
write_rows('Expenses', expense_rows)
write_rows('Quotes', quote_rows)
summary_rows=[]
for pid in sorted(summary):
    s = summary[pid]
    summary_rows.append([pid, money(s['OrderTWD']), money(s['InvoiceBalanceTWD']), money(s['ApprovedExpenseTWD']), money(s['AcceptedQuoteTWD'])])
write_rows('Summary', summary_rows)
last = max(5, 4 + len(summary_rows))
sum_ws = wb['Summary']
for col in range(2,6):
    letter = chr(ord('A') + col - 1)
    sum_ws.cell(2,col).value = f'=SUM({letter}5:{letter}{last})'
    set_amt(sum_ws.cell(2,col))
write_rows('Exceptions', exceptions)

rep = wb['Report']
now = datetime.now(ZoneInfo('Asia/Taipei')).replace(microsecond=0).isoformat()
rep['B3'] = FROM
rep['D3'] = TO
rep['B4'] = 'SAP orders: non-Cancelled by OrderDate; SAP invoices by InvoiceDate; Expense Viewer: Approved by ExpenseDate; quote_cli: Accepted by QuoteDate; grouped by ProjectID only.'
rep['B5'] = now
rep['B7'] = 'Lab ERP/SAP web; Expense Viewer desktop; quote_cli; fx_rates.csv'
rep['D7'] = now
rep['B9'] = 'Orders and accepted quotes are reported separately and not added as revenue. Customer name is not used as a join key. Blank ProjectID records are listed in Exceptions and excluded from Summary.'
wb.save(out_xlsx)

# Reopen and verify stated requirements.
wb2 = load_workbook(out_xlsx, data_only=False)
assert [c.value for c in wb2['Orders'][4]][:8] == ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD']
assert [c.value for c in wb2['Summary'][4]][:5] == ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']
assert all(wb2['Orders'].cell(r,5).value != 'Cancelled' for r in range(5, 5+len(order_rows)))
assert all(wb2['Expenses'].cell(r,5).value == 'Approved' for r in range(5, 5+len(expense_rows)))
assert all(wb2['Quotes'].cell(r,6).value == 'Accepted' for r in range(5, 5+len(quote_rows)))
for sheet, idcols, amtcols, n in [('Orders',[1,2,3],[7,8],len(order_rows)), ('Invoices',[1,2,3,4],[8,9,10,11],len(invoice_rows)), ('Expenses',[1,2,3],[7,8],len(expense_rows)), ('Quotes',[1,2,3],[8,9],len(quote_rows))]:
    wsht = wb2[sheet]
    for r in range(5, 5+n):
        for c in idcols:
            assert wsht.cell(r,c).number_format == '@'
        for c in amtcols:
            v = wsht.cell(r,c).value
            assert isinstance(v, (int,float)) or v is None
checks = {
  'orders_rows': len(order_rows),
  'invoices_rows': len(invoice_rows),
  'expenses_rows': len(expense_rows),
  'quotes_rows': len(quote_rows),
  'summary_rows': len(summary_rows),
  'exceptions_rows': len(exceptions),
  'outputs': [str(orders_path), str(invoices_path), str(expenses_path), str(quotes_json), str(out_xlsx)],
}
print(json.dumps(checks, ensure_ascii=False, indent=2))