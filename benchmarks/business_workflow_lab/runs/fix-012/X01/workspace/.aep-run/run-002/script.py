import csv, os, subprocess, json
from pathlib import Path
ws=Path.cwd()
for f in [ws/'output'/'downloads'/'sap_orders_2026-09.csv', ws/'output'/'downloads'/'sap_invoices_2026-09.csv', ws/'output'/'expenses_2026-09_approved.csv']:
    print('\nFILE', f, f.exists(), f.stat().st_size if f.exists() else None)
    if f.exists():
        with open(f, newline='', encoding='utf-8-sig') as fh:
            rdr=list(csv.reader(fh))
            for row in rdr[:6]: print(row)
            print('rows incl header', len(rdr))
print('\nInstructions full:')
print((ws/'input'/'project_template.xlsx').exists())
from openpyxl import load_workbook
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True): print(row)
print('\nQuote CLI test')
cmd=['python', r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py', 'quotes','list','--from','2026-09-01','--to','2026-09-30','--status','Accepted','--format','json']
p=subprocess.run(cmd, text=True, capture_output=True)
print('return', p.returncode)
print('stderr', p.stderr[:500])
print('stdout', p.stdout[:2000])