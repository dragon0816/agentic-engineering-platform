import os, csv, json
from pathlib import Path
from openpyxl import load_workbook

base = Path.cwd()
out = base / 'output'
print('CHECK_WORKSPACE', str(base))
print('OUTPUT_EXISTS', out.exists())
print('OUTPUT_FILES')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(str(p.relative_to(base)), p.stat().st_size)

expected = [
    'sap_orders_2026-09.csv',
    'sap_invoices_2026-09.csv',
    'expenses_2026-09_approved.csv',
    'quotes_2026-09_accepted.json',
    'project_integration_2026-09.xlsx',
]
checks = {}
for name in expected:
    p = out / name
    checks[name] = {'exists': p.exists(), 'size': p.stat().st_size if p.exists() else 0}

# CSV checks
for name in ['sap_orders_2026-09.csv','sap_invoices_2026-09.csv','expenses_2026-09_approved.csv']:
    p = out / name
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            rows = list(reader)
        checks[name]['header'] = rows[0] if rows else []
        checks[name]['data_rows'] = max(0, len(rows)-1)
        checks[name]['sample'] = rows[1][:5] if len(rows) > 1 else []

# JSON checks
jp = out / 'quotes_2026-09_accepted.json'
if jp.exists():
    data = json.loads(jp.read_text(encoding='utf-8'))
    checks[jp.name]['type'] = type(data).__name__
    checks[jp.name]['items'] = len(data) if isinstance(data, list) else len(data.keys()) if isinstance(data, dict) else None
    checks[jp.name]['sample_keys'] = list(data[0].keys()) if isinstance(data, list) and data else list(data.keys())[:10] if isinstance(data, dict) else []

# Workbook checks
wp = out / 'project_integration_2026-09.xlsx'
if wp.exists():
    wb = load_workbook(wp, data_only=False)
    checks[wp.name]['sheets'] = wb.sheetnames
    sheet_info = {}
    for ws in wb.worksheets:
        nonempty_rows = 0
        for row in ws.iter_rows(values_only=True):
            if any(v is not None for v in row):
                nonempty_rows += 1
        headers = [ws.cell(1, c).value for c in range(1, min(ws.max_column, 12)+1)]
        sheet_info[ws.title] = {'max_row': ws.max_row, 'max_column': ws.max_column, 'nonempty_rows': nonempty_rows, 'headers_first_row': headers}
    checks[wp.name]['sheet_info'] = sheet_info

# final assertions
assert all(v['exists'] and v['size'] > 0 for v in checks.values()), checks
print('CHECKS_JSON')
print(json.dumps(checks, ensure_ascii=False, indent=2, default=str))