import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=list(csv.reader(f))
    for row in r[:10]: print(row)
print('\nproject_template.xlsx sheets/headers:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    s=wb[sh]
    print('sheet', sh, 'max', s.max_row, s.max_column)
    for i,row in enumerate(s.iter_rows(min_row=1, max_row=min(s.max_row,8), values_only=True),1):
        print(i, row)
    print()
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    print('\n====',name,'====')
    txt=(ws/'input'/name).read_text(encoding='utf-8')
    print(txt[:4000])