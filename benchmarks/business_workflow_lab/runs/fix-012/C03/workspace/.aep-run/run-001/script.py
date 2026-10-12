from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('cwd=', ws)
print('input files=')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# fx sample
fx = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i, row in enumerate(r):
        print(i, row)
        if i >= 8: break

# workbook structure
xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nworkbook sheets=', wb.sheetnames)
for sh in wb.sheetnames:
    s = wb[sh]
    print('sheet', sh, 'max_row', s.max_row, 'max_col', s.max_column)
    for row in s.iter_rows(min_row=1, max_row=min(s.max_row, 8), values_only=True):
        print(' ', row)

# manual first lines
md = ws/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md:')
txt = md.read_text(encoding='utf-8', errors='replace')
print(txt[:4000])