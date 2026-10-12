from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook(Path('input/quote_template.xlsx'), data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)