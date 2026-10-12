from pathlib import Path
import subprocess, json, csv, shutil, sys, os
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

WORK = Path.cwd()
IN = WORK / 'input'
OUT = WORK / 'output'
OUT.mkdir(exist_ok=True)

PROJECTS = [f'P{i:03d}' for i in range(1, 6)]
DATE_FROM = '2026-09-01'
DATE_TO = '2026-09-30'
STATUS = 'Accepted'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
FETCHED_AT = datetime.now().astimezone().isoformat(timespec='seconds')
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

TWOPLACES = Decimal('0.01')
def money(x):
    if x is None or x == '':
        return None
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def twd(amount, currency, rates):
    return (Decimal(str(amount)) * rates[currency]).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [sys.executable, str(CLI)] + args
    p = subprocess.run(cmd, cwd=str(WORK), text=True, capture_output=True)
    if p.returncode != 0:
        raise RuntimeError(f'CLI failed {p.returncode}: {cmd}\nSTDERR={p.stderr}\nSTDOUT={p.stdout}')
    return p.stdout

# Read FX rates robustly; Note contains unquoted commas, only first two columns matter.
rates = {}
with (IN/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

# Fetch quote headers and detail lines.
headers = []
lines = []
queries = []
for project in PROJECTS:
    args = ['quotes', 'list', '--project', project, '--from', DATE_FROM, '--to', DATE_TO, '--status', STATUS, '--format', 'json']
    data = json.loads(run_cli(args))
    queries.append({'project': project, 'count': data.get('count')})
    for h in data.get('items', []):
        if h.get('ProjectID') != project or h.get('Status') != STATUS or not (DATE_FROM <= h.get('QuoteDate','') <= DATE_TO):
            raise AssertionError(f'Unexpected quote in list result: {h}')
        headers.append(h)

headers.sort(key=lambda h: (h.get('ProjectID',''), h.get('QuoteDate',''), h.get('QuoteID','')))
print('Query counts:', queries)
print('Fetched quote headers:', len(headers))
for h in headers:
    print(' HEADER', h)
    detail = json.loads(run_cli(['quotes', 'get', h['QuoteID'], '--format', 'json']))
    dh = detail['header']
    if dh['QuoteID'] != h['QuoteID']:
        raise AssertionError('Detail/header quote mismatch')
    for li in detail.get('line_items', []):
        li = dict(li)
        li['QuoteID'] = h['QuoteID']
        lines.append(li)
print('Fetched line rows:', len(lines))
for li in lines[:10]:
    print(' LINE', li)

# Copy template and fill it.
outfile = OUT / 'quote_report_P001-P005_2026-09_Accepted.xlsx'
shutil.copy2(IN/'quote_template.xlsx', outfile)
wb = load_workbook(outfile)

# Helper: copy row style/format from template header row to data rows where useful.
def style_like_row(ws, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

# Report sheet
ws = wb['Report']
ws['B3'] = DATE_FROM
ws['D3'] = DATE_TO
ws['B4'] = 'ProjectID P001-P005; QuoteDate 2026-09-01..2026-09-30 inclusive; Status Accepted; NetAmountTWD = NetAmount × fx_rates.csv RateToTWD'
ws['B5'] = GENERATED_AT
ws['B7'] = 'quote_cli; fx_rates.csv'
ws['D7'] = FETCHED_AT
ws['B9'] = 'Summary groups by ProjectID. Projects with no Accepted quote are retained with zero totals. Original-currency Amount is filled only when that project has a single currency; otherwise left blank to avoid mixing currencies.'

# Clear/fill Quotes from row 5
wsq = wb['Quotes']
for row in range(5, max(wsq.max_row, 5)+1):
    for col in range(1, 12):
        wsq.cell(row, col).value = None
for idx, h in enumerate(headers, start=5):
    style_like_row(wsq, 4, idx, 11)
    qid = h['QuoteID']; cur = h['Currency']
    if cur not in rates:
        raise AssertionError(f'Missing FX rate for {cur}')
    vals = [
        qid,
        h.get('CustomerID'),
        h.get('ProjectID'),
        h.get('QuoteDate'),
        h.get('ValidUntil'),
        h.get('Status'),
        cur,
        float(money(h.get('Subtotal'))),
        float(Decimal(str(h.get('DiscountPercent')))),
        float(money(h.get('NetAmount'))),
        float(twd(money(h.get('NetAmount')), cur, rates)),
    ]
    for col, val in enumerate(vals, 1):
        c = wsq.cell(idx, col, val)
        if col in [1,2,3,4,5,6,7]:
            c.number_format = '@' if col in [1,2,3] else 'yyyy-mm-dd' if col in [4,5] else 'General'
        if col in [8,10,11]: c.number_format = '0.00'
        if col == 9: c.number_format = '0.00'

# Lines
wsl = wb['Lines']
for row in range(5, max(wsl.max_row, 5)+1):
    for col in range(1, 7):
        wsl.cell(row, col).value = None
for idx, li in enumerate(lines, start=5):
    style_like_row(wsl, 4, idx, 6)
    vals = [li.get('QuoteID'), li.get('ProductID'), li.get('ProductName'), int(li.get('Quantity')), float(money(li.get('UnitPrice'))), float(money(li.get('LineAmount')))]
    for col, val in enumerate(vals, 1):
        c = wsl.cell(idx, col, val)
        if col in [1,2]: c.number_format = '@'
        if col in [5,6]: c.number_format = '0.00'

# Summary by ProjectID, preserving P001..P005.
wss = wb['Summary']
for row in range(5, max(wss.max_row, 5)+1):
    for col in range(1, 7):
        wss.cell(row, col).value = None

by_project = {p: [] for p in PROJECTS}
for h in headers:
    by_project[h['ProjectID']].append(h)

for r, project in enumerate(PROJECTS, start=5):
    style_like_row(wss, 4, r, 6)
    hs = by_project[project]
    currencies = sorted({h['Currency'] for h in hs})
    count = len(hs)
    amount_twd = sum((twd(money(h['NetAmount']), h['Currency'], rates) for h in hs), Decimal('0.00')).quantize(TWOPLACES)
    # Do not mix original currencies. If no rows, zero; if one currency, sum original amount; if multiple, leave blank and Currency='MULTI'.
    if count == 0:
        currency = ''
        amount = Decimal('0.00')
    elif len(currencies) == 1:
        currency = currencies[0]
        amount = sum((money(h['NetAmount']) for h in hs), Decimal('0.00')).quantize(TWOPLACES)
    else:
        currency = 'MULTI'
        amount = None
    vals = [project, '', currency, None if amount is None else float(amount), float(amount_twd), count]
    for col, val in enumerate(vals, 1):
        c = wss.cell(r, col, val)
        if col in [1,2,3]: c.number_format = '@'
        if col in [4,5]: c.number_format = '0.00'
        if col == 6: c.number_format = '0'

last_quote_row = 4 + max(len(headers), 1)
wss['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'

# Re-apply autofilter ranges if present enough.
if len(headers):
    wsq.auto_filter.ref = f'A4:K{4+len(headers)}'
if len(lines):
    wsl.auto_filter.ref = f'A4:F{4+len(lines)}'
wss.auto_filter.ref = 'A4:F9'

wb.save(outfile)
print('Wrote', outfile)

# Validation: reopen saved output and verify requirements.
wb2 = load_workbook(outfile, data_only=False)
assert wb2.sheetnames == ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions'], wb2.sheetnames
q = wb2['Quotes']; l = wb2['Lines']; s = wb2['Summary']; rep = wb2['Report']
expected_q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [q.cell(4,c).value for c in range(1,12)] == expected_q_headers
expected_l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [l.cell(4,c).value for c in range(1,7)] == expected_l_headers
expected_s_headers = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
assert [s.cell(4,c).value for c in range(1,7)] == expected_s_headers

# Check quote filters and TWD calculations.
for row in range(5, 5+len(headers)):
    qid = q.cell(row,1).value
    proj = q.cell(row,3).value
    qdate = q.cell(row,4).value
    stat = q.cell(row,6).value
    cur = q.cell(row,7).value
    net = Decimal(str(q.cell(row,10).value)).quantize(TWOPLACES)
    net_twd = Decimal(str(q.cell(row,11).value)).quantize(TWOPLACES)
    assert proj in PROJECTS, (row, proj)
    assert DATE_FROM <= qdate <= DATE_TO, (row, qdate)
    assert stat == STATUS, (row, stat)
    assert net_twd == twd(net, cur, rates), (qid, net, cur, net_twd)
    assert isinstance(qid, str) and isinstance(proj, str)

# Check line QuoteIDs all correspond.
quote_ids = {q.cell(row,1).value for row in range(5, 5+len(headers))}
for row in range(5, 5+len(lines)):
    assert l.cell(row,1).value in quote_ids
    assert isinstance(l.cell(row,1).value, str)
    assert isinstance(l.cell(row,2).value, str)

# Summary exactly one row per project and no mixed original sum.
summary_projects = [s.cell(row,1).value for row in range(5,10)]
assert summary_projects == PROJECTS, summary_projects
for row, project in zip(range(5,10), PROJECTS):
    hs = by_project[project]
    assert s.cell(row,6).value == len(hs), (project, s.cell(row,6).value, len(hs))
    expected_twd = sum((twd(money(h['NetAmount']), h['Currency'], rates) for h in hs), Decimal('0.00')).quantize(TWOPLACES)
    actual_twd = Decimal(str(s.cell(row,5).value)).quantize(TWOPLACES)
    assert actual_twd == expected_twd, (project, actual_twd, expected_twd)
    currencies = {h['Currency'] for h in hs}
    if len(hs) == 0:
        assert Decimal(str(s.cell(row,4).value)).quantize(TWOPLACES) == Decimal('0.00')
        assert actual_twd == Decimal('0.00')
    elif len(currencies) > 1:
        assert s.cell(row,4).value is None
        assert s.cell(row,3).value == 'MULTI'
    else:
        expected_amt = sum((money(h['NetAmount']) for h in hs), Decimal('0.00')).quantize(TWOPLACES)
        actual_amt = Decimal(str(s.cell(row,4).value)).quantize(TWOPLACES)
        assert actual_amt == expected_amt

assert s['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})'
print('VALIDATION OK')
print('Output file:', outfile.name)
print('Quote rows:', len(headers), 'Line rows:', len(lines), 'Summary rows:', len(PROJECTS))
for row in range(5,10):
    print('SUMMARY', [s.cell(row,c).value for c in range(1,7)])
print('Summary B2 formula:', s['B2'].value)
