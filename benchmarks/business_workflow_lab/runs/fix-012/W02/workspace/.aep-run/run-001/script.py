import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('CWD', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# workbook structure
path = ws/'input'/'sales_template.xlsx'
wb = load_workbook(path, data_only=False)
print('Workbook sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print(f'--- Sheet {sh.title} max_row={sh.max_row} max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
        print(r)
# fx rates sample
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('--- fx_rates.csv sample')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        rdr = csv.reader(f)
        for i,row in zip(range(10), rdr): print(row)
# manual
man = ws/'input'/'sap_manual.md'
if man.exists():
    print('--- sap_manual.md')
    print(man.read_text(encoding='utf-8')[:4000])