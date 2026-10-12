from pathlib import Path
import subprocess, sys, json, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import date
from collections import defaultdict
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
REPORT = OUTPUT / 'quote_report_2026-09.xlsx'
FINDINGS = OUTPUT / 'findings.md'
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
GENERATED_AT = '2026-10-06T00:00:00+08:00'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
PY_CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TWOPLACES = Decimal('0.01')

def D(v):
    if v is None or v == '':
        return Decimal('0')
    return Decimal(str(v))

def money(v):
    return D(v).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [str(CLI)] + args if CLI.exists() else [sys.executable, str(PY_CLI)] + args
    cp = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
    if cp.returncode != 0:
        raise RuntimeError(f'CLI failed code={cp.returncode}\nCMD={cmd}\nSTDERR={cp.stderr}\nSTDOUT={cp.stdout}')
    return cp.stdout

# read actual inputs again
print('Workspace:', WORKSPACE)
print('Inputs:', [p.name for p in sorted(INPUT.glob('*'))])
fx_rates = {}
with FX.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    print('FX header:', next(reader))
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            fx_rates[row[0].strip()] = D(row[1].strip())
print('Loaded FX rates:', fx_rates)
wb_probe = load_workbook(TEMPLATE, data_only=False)
print('Template sheets:', wb_probe.sheetnames)
print('Quotes headers:', [wb_probe['Quotes'].cell(row=4, column=c).value for c in range(1, wb_probe['Quotes'].max_column+1)])
print('Lines headers:', [wb_probe['Lines'].cell(row=4, column=c).value for c in range(1, wb_probe['Lines'].max_column+1)])

listed = json.loads(run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json']))
quotes_headers = listed.get('items', [])
print(f'Fetched {len(quotes_headers)} quote headers for {FROM_DATE}..{TO_DATE}')

quotes = []
all_lines = []
for h in quotes_headers:
    qid = h['QuoteID']
    detail = json.loads(run_cli(['quotes', 'get', qid, '--format', 'json']))
    header = detail['header']
    lines = detail.get('line_items', [])
    quotes.append(header)
    for line in lines:
        rec = dict(line)
        rec['QuoteID'] = qid
        all_lines.append(rec)
print(f'Fetched details for {len(quotes)} quotes and {len(all_lines)} line items')
if quotes:
    print('Sample quote header:', quotes[0])
if all_lines:
    print('Sample line:', all_lines[0])

line_findings = []
net_findings = []
expired_unaccepted = []
quote_lines = defaultdict(list)
for line in all_lines:
    quote_lines[line['QuoteID']].append(line)
    expected_line = money(D(line.get('Quantity')) * D(line.get('UnitPrice')))
    actual_line = money(line.get('LineAmount'))
    if actual_line != expected_line:
        line_findings.append({'QuoteID': line['QuoteID'], 'ProductID': line.get('ProductID', ''), 'Quantity': str(line.get('Quantity', '')), 'UnitPrice': str(line.get('UnitPrice', '')), 'LineAmount': str(actual_line), 'ExpectedLineAmount': str(expected_line), 'Diff': str(actual_line - expected_line)})

for q in quotes:
    qid = q['QuoteID']
    sum_lines = money(sum((money(l.get('LineAmount')) for l in quote_lines.get(qid, [])), Decimal('0')))
    subtotal = money(q.get('Subtotal'))
    if subtotal != sum_lines:
        line_findings.append({'QuoteID': qid, 'ProductID': '(subtotal)', 'Quantity': '', 'UnitPrice': '', 'LineAmount': str(subtotal), 'ExpectedLineAmount': str(sum_lines), 'Diff': str(subtotal - sum_lines)})
    expected_net = money(subtotal * (Decimal('1') - (D(q.get('DiscountPercent')) / Decimal('100'))))
    actual_net = money(q.get('NetAmount'))
    if actual_net != expected_net:
        net_findings.append({'QuoteID': qid, 'Subtotal': str(subtotal), 'DiscountPercent': str(q.get('DiscountPercent', '')), 'NetAmount': str(actual_net), 'ExpectedNetAmount': str(expected_net), 'Diff': str(actual_net - expected_net)})
    if date.fromisoformat(q['ValidUntil']) < BASE_DATE and q.get('Status') != 'Accepted':
        expired_unaccepted.append({'QuoteID': qid, 'CustomerID': q.get('CustomerID',''), 'ProjectID': q.get('ProjectID',''), 'QuoteDate': q.get('QuoteDate',''), 'ValidUntil': q.get('ValidUntil',''), 'Status': q.get('Status',''), 'Currency': q.get('Currency',''), 'NetAmount': str(actual_net)})

shutil.copyfile(TEMPLATE, REPORT)
wb = load_workbook(REPORT)
rep = wb['Report']
rep['B3'] = FROM_DATE
rep['D3'] = TO_DATE
rep['B4'] = '2026 年 9 月全部報價；檢查明細合計、折扣後 NetAmount；ValidUntil < 2026-09-30 且尚未 Accepted 者列入 findings；保留全部報價。'
rep['B5'] = GENERATED_AT
rep['B7'] = 'quote_cli_mock; fx_rates.csv'
rep['D7'] = GENERATED_AT
rep['B9'] = f'取得 {len(quotes)} 份報價、{len(all_lines)} 筆明細；未因到期刪除任何報價。明細/小計異常 {len(line_findings)} 筆；NetAmount 異常 {len(net_findings)} 筆；基準日 {BASE_DATE.isoformat()} 已逾期且未 Accepted {len(expired_unaccepted)} 份。'
qsheet, lsheet, ssheet = wb['Quotes'], wb['Lines'], wb['Summary']
for ws in (qsheet, lsheet, ssheet):
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)
quote_cols = [qsheet.cell(row=4, column=c).value for c in range(1, qsheet.max_column+1)]
line_cols = [lsheet.cell(row=4, column=c).value for c in range(1, lsheet.max_column+1)]
for r, q in enumerate(quotes, start=5):
    cur = q.get('Currency')
    net = money(q.get('NetAmount'))
    rate = fx_rates.get(cur)
    values = {'QuoteID': str(q.get('QuoteID','')) if q.get('QuoteID') is not None else None, 'CustomerID': str(q.get('CustomerID','')) if q.get('CustomerID') is not None else None, 'ProjectID': str(q.get('ProjectID')) if q.get('ProjectID') not in (None, '') else None, 'QuoteDate': q.get('QuoteDate'), 'ValidUntil': q.get('ValidUntil'), 'Status': q.get('Status'), 'Currency': cur, 'Subtotal': float(money(q.get('Subtotal'))), 'DiscountPercent': float(D(q.get('DiscountPercent'))), 'NetAmount': float(net), 'NetAmountTWD': float(money(net * rate)) if rate is not None else None}
    for c, col in enumerate(quote_cols, start=1):
        qsheet.cell(row=r, column=c, value=values.get(col))
    for c in [1,2,3,6,7]:
        qsheet.cell(row=r, column=c).number_format = '@'
    for c in [4,5]:
        qsheet.cell(row=r, column=c).number_format = 'yyyy-mm-dd'
    for c in [8,10,11]:
        qsheet.cell(row=r, column=c).number_format = '0.00'
    qsheet.cell(row=r, column=9).number_format = '0.00'
for r, line in enumerate(all_lines, start=5):
    values = {'QuoteID': str(line.get('QuoteID','')) if line.get('QuoteID') is not None else None, 'ProductID': str(line.get('ProductID','')) if line.get('ProductID') is not None else None, 'ProductName': line.get('ProductName'), 'Quantity': float(D(line.get('Quantity'))), 'UnitPrice': float(money(line.get('UnitPrice'))), 'LineAmount': float(money(line.get('LineAmount')))}
    for c, col in enumerate(line_cols, start=1):
        lsheet.cell(row=r, column=c, value=values.get(col))
    for c in [1,2,3]:
        lsheet.cell(row=r, column=c).number_format = '@'
    lsheet.cell(row=r, column=4).number_format = '0.00'
    for c in [5,6]:
        lsheet.cell(row=r, column=c).number_format = '0.00'
by_status_currency = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
for q in quotes:
    cur = q.get('Currency')
    net = money(q.get('NetAmount'))
    key = (q.get('Status',''), cur)
    by_status_currency[key]['amount'] += net
    by_status_currency[key]['twd'] += money(net * fx_rates.get(cur, Decimal('0')))
    by_status_currency[key]['count'] += 1
summary_rows = []
for (status, cur), agg in sorted(by_status_currency.items()):
    summary_rows.append(['Status', status, cur, float(money(agg['amount'])), float(money(agg['twd'])), agg['count']])
expired_twd = sum((money(D(e['NetAmount']) * fx_rates.get(e['Currency'], Decimal('0'))) for e in expired_unaccepted), Decimal('0'))
summary_rows.append(['Finding', f'ValidUntil < {BASE_DATE.isoformat()} and Status <> Accepted', 'TWD', None, float(money(expired_twd)), len(expired_unaccepted)])
for r, row in enumerate(summary_rows, start=5):
    for c, val in enumerate(row, start=1):
        ssheet.cell(row=r, column=c, value=val)
    for c in [1,2,3]: ssheet.cell(row=r, column=c).number_format = '@'
    for c in [4,5]: ssheet.cell(row=r, column=c).number_format = '0.00'
    ssheet.cell(row=r, column=6).number_format = '0'
last_quote_row = 4 + len(quotes) if quotes else 5
ssheet['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
qsheet.auto_filter.ref = f'A4:K{max(4+len(quotes),4)}'
lsheet.auto_filter.ref = f'A4:F{max(4+len(all_lines),4)}'
ssheet.auto_filter.ref = f'A4:F{max(4+len(summary_rows),4)}'
wb.save(REPORT)

def md_table(headers, rows):
    out = ['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |']
    for row in rows:
        out.append('| ' + ' | '.join(str(row.get(h, '')) for h in headers) + ' |')
    return out
md = []
md += ['# Findings - 2026 年 9 月報價檢查', '', f'- 查詢期間：{FROM_DATE} 至 {TO_DATE}（含起訖日）', f'- 逾期判定基準日：{BASE_DATE.isoformat()}', '- 資料來源：quote_cli_mock；匯率：input/fx_rates.csv', f'- 取得報價：{len(quotes)} 份；明細：{len(all_lines)} 筆', '- 處理原則：保留全部報價，不因到期而刪除。', '']
md += ['## 明細合計檢查']
if line_findings:
    md += [f'發現 {len(line_findings)} 筆明細金額或 Subtotal 不一致：'] + md_table(['QuoteID','ProductID','Quantity','UnitPrice','LineAmount','ExpectedLineAmount','Diff'], line_findings)
else:
    md += ['未發現 LineAmount = Quantity × UnitPrice 或 Subtotal = 明細 LineAmount 合計的不一致。']
md += ['', '## 折扣後 NetAmount 檢查']
if net_findings:
    md += [f'發現 {len(net_findings)} 筆 NetAmount 不一致：'] + md_table(['QuoteID','Subtotal','DiscountPercent','NetAmount','ExpectedNetAmount','Diff'], net_findings)
else:
    md += ['未發現 NetAmount = Subtotal × (1 − DiscountPercent / 100)（ROUND_HALF_UP 至兩位小數）的不一致。']
md += ['', '## ValidUntil 小於 2026-09-30 且尚未 Accepted 的報價']
if expired_unaccepted:
    md += [f'共 {len(expired_unaccepted)} 份：'] + md_table(['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','NetAmount'], expired_unaccepted)
else:
    md += ['未發現符合條件的報價。']
FINDINGS.write_text('\n'.join(md) + '\n', encoding='utf-8')
print('Wrote:', REPORT)
print('Wrote:', FINDINGS)
print('Line/subtotal findings:', len(line_findings))
print('Net findings:', len(net_findings))
print('Expired unaccepted:', len(expired_unaccepted))

# Reopen/check outputs
check_wb = load_workbook(REPORT, data_only=False)
assert check_wb.sheetnames == wb_probe.sheetnames, check_wb.sheetnames
cq, cl, cs = check_wb['Quotes'], check_wb['Lines'], check_wb['Summary']
assert [cq.cell(row=4, column=c).value for c in range(1, 12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [cl.cell(row=4, column=c).value for c in range(1, 7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
quote_rows = [r for r in range(5, cq.max_row+1) if cq.cell(row=r, column=1).value not in (None, '')]
line_rows = [r for r in range(5, cl.max_row+1) if cl.cell(row=r, column=1).value not in (None, '')]
assert len(quote_rows) == len(quotes), (len(quote_rows), len(quotes))
assert len(line_rows) == len(all_lines), (len(line_rows), len(all_lines))
assert cs['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})', cs['B2'].value
# text IDs when source provided; blank remains blank for missing ProjectID per instructions.
for r in quote_rows:
    for c in [1,2]:
        assert isinstance(cq.cell(row=r, column=c).value, str), (r,c,cq.cell(row=r,column=c).value)
    assert cq.cell(row=r, column=3).value is None or isinstance(cq.cell(row=r, column=3).value, str)
    for c in [8,9,10,11]:
        assert isinstance(cq.cell(row=r, column=c).value, (int, float)), (r,c,cq.cell(row=r,column=c).value,type(cq.cell(row=r,column=c).value))
for r in line_rows:
    for c in [1,2]:
        assert isinstance(cl.cell(row=r, column=c).value, str), (r,c,cl.cell(row=r,column=c).value)
    for c in [4,5,6]:
        assert isinstance(cl.cell(row=r, column=c).value, (int, float)), (r,c,cl.cell(row=r,column=c).value,type(cl.cell(row=r,column=c).value))
text = FINDINGS.read_text(encoding='utf-8')
assert '保留全部報價' in text
assert 'ValidUntil 小於 2026-09-30 且尚未 Accepted' in text
assert f'共 {len(expired_unaccepted)} 份' in text
print('CHECK OK: output workbook reopened; template sheet names kept; detail headers kept; Quotes rows =', len(quote_rows), '; Lines rows =', len(line_rows), '; text ID/numeric amount types checked; missing ProjectID remains blank; Summary B2 formula range =', cs['B2'].value, '; findings.md sections and counts checked.')
print('Output files:', [(p.name, p.stat().st_size) for p in sorted(OUTPUT.glob('*'))])
