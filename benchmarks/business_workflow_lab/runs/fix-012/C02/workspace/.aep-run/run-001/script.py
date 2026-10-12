from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('CWD:', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv preview:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i,row in zip(range(6), r):
            print(i, row)

manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\nquote_cli_manual.md first 120 lines:')
    text = manual.read_text(encoding='utf-8', errors='replace').splitlines()
    for i,line in enumerate(text[:120], 1):
        print(f'{i:03d}: {line}')

tpl = ws/'input'/'quote_template.xlsx'
if tpl.exists():
    print('\nWorkbook structure:')
    wb = load_workbook(tpl, data_only=False)
    for sh in wb.worksheets:
        print('Sheet:', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(row)
        print('---')
