import os, csv, json
from pathlib import Path
from openpyxl import load_workbook

ws = Path.cwd()
print('CWD', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# CSV headers/sample
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        rows = list(reader)
    print('\nfx_rates.csv rows:', len(rows))
    print('fx header:', rows[0] if rows else None)
    print('fx sample:', rows[1:6])

# manual first lines
manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    txt = manual.read_text(encoding='utf-8', errors='replace')
    print('\nquote_cli_manual.md first 120 lines:')
    for i, line in enumerate(txt.splitlines()[:120], 1):
        print(f'{i:03}: {line}')

# workbook structure
xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    wb = load_workbook(xlsx)
    print('\nWorkbook sheets:', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print(f'-- Sheet {s}: max_row={sh.max_row} max_col={sh.max_column}')
        for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(r)
        print('merged:', list(sh.merged_cells.ranges))
