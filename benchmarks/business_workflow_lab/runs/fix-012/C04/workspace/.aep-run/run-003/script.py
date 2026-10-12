import csv
import json
import subprocess
import sys
from copy import copy
from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)

TEMPLATE = INPUT / 'quote_template.xlsx'
FX_CSV = INPUT / 'fx_rates.csv'
CLI_PY = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
SOURCE_NAME = 'quote_cli synthetic quote system'
OUTFILE = OUTPUT / 'quote_report_2026-09_sent.xlsx'
FETCH_LOG = OUTPUT / 'quote_cli_fetch_log.json'
TW = timezone(timedelta(hours=8))
RUN_TS = datetime.now(TW).replace(microsecond=0).isoformat()

class CliFailure(Exception):
    pass

def read_fx_rates(path: Path):
    rates = {}
    with path.open('r', encoding='utf-8-sig', newline='') as f:
        reader = csv.reader(f)
        header = next(reader)
        if len(header) < 2 or header[0] != 'Currency' or header[1] != 'RateToTWD':
            raise ValueError(f'Unexpected FX CSV header: {header}')
        for row in reader:
            if not row or not row[0].strip():
                continue
            rates[row[0].strip()] = Decimal(row[1].strip())
    return rates

def run_cli_json(cli_args, label, max_retries_on_4=2):
    attempts = []
    retry_count = 0
    while True:
        cmd = [sys.executable, str(CLI_PY), *cli_args]
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True)
        rec = {'label': label, 'command': cmd, 'exit_code': proc.returncode, 'stderr': proc.stderr.strip(), 'stdout_length': len(proc.stdout or '')}
        attempts.append(rec)
        print(f"CLI {label} attempt {len(attempts)} exit={proc.returncode} stderr={proc.stderr.strip()!r} stdout_len={len(proc.stdout or '')}")
        if proc.returncode == 0:
            try:
                data = json.loads(proc.stdout)
            except Exception as e:
                rec['json_parse_error'] = str(e)
                raise CliFailure(f'{label}: exit code 0 but stdout was not parseable JSON; stdout not used. Parse error: {e}')
            rec['json_parsed'] = True
            return data, attempts
        if proc.returncode == 4:
            if retry_count < max_retries_on_4:
                retry_count += 1
                continue
            raise CliFailure(f'{label}: temporary service error after {len(attempts)} attempts. stderr: {proc.stderr.strip()}')
        if proc.returncode in (2, 3):
            raise CliFailure(f'{label}: CLI returned exit code {proc.returncode}; not retried. stderr: {proc.stderr.strip()}')
        raise CliFailure(f'{label}: CLI returned unexpected exit code {proc.returncode}. stderr: {proc.stderr.strip()}')

def clear_data_rows(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def set_id_cell(cell, value):
    if value is None or value == '':
        cell.value = None
    else:
        cell.value = str(value)
    cell.number_format = '@'

def set_date_cell(cell, value):
    cell.value = value
    cell.number_format = 'yyyy-mm-dd'

def set_money_cell(cell, value):
    cell.value = float(Decimal(value)) if value is not None else None
    cell.number_format = '0.00'

def set_number_cell(cell, value):
    if value is None:
        cell.value = None
    else:
        d = Decimal(str(value))
        cell.value = int(d) if d == d.to_integral_value() else float(d)

def style_data_row(ws, row_idx, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(4, col)
        dst = ws.cell(row_idx, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = Font(name=src.font.name or 'Calibri', size=src.font.sz or 11, bold=False)

def as_decimal(value):
    return Decimal(str(value))

def q2(value):
    return Decimal(value).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def nonempty_text_or_blank(cell):
    return cell.value is None or (cell.data_type == 's' and isinstance(cell.value, str))

def main():
    rates = read_fx_rates(FX_CSV)
    print('FX rates loaded:', {k: str(v) for k, v in rates.items()})

    all_attempts = []
    list_data, attempts = run_cli_json(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'], 'quotes list 2026-09 Sent')
    all_attempts.extend(attempts)
    if list_data.get('schema_version') != 1 or 'items' not in list_data:
        raise ValueError(f'Unexpected quotes list schema: keys={list(list_data.keys())}')
    quote_headers = list_data['items']
    print('List count field:', list_data.get('count'), 'items:', len(quote_headers))

    details = []
    lines = []
    for item in quote_headers:
        qid = str(item.get('QuoteID', ''))
        if not qid:
            raise ValueError(f'List item missing QuoteID: {item}')
        data, attempts = run_cli_json(['quotes', 'get', qid, '--format', 'json'], f'quotes get {qid}')
        all_attempts.extend(attempts)
        if data.get('schema_version') != 1 or 'header' not in data or 'line_items' not in data:
            raise ValueError(f'Unexpected quote detail schema for {qid}: keys={list(data.keys())}')
        header = data['header']
        if str(header.get('QuoteID')) != qid:
            raise ValueError(f'QuoteID mismatch: list {qid}, get {header.get("QuoteID")}')
        details.append(header)
        for li in data['line_items']:
            lines.append({'QuoteID': qid, **li})

    details.sort(key=lambda h: (h.get('QuoteDate', ''), h.get('QuoteID', '')))
    detail_order = {h['QuoteID']: i for i, h in enumerate(details)}
    lines.sort(key=lambda l: (detail_order.get(l.get('QuoteID'), 10**9), str(l.get('ProductID', '')), str(l.get('ProductName', ''))))

    wb = load_workbook(TEMPLATE)
    report = wb['Report']
    quotes_ws = wb['Quotes']
    lines_ws = wb['Lines']
    summary_ws = wb['Summary']
    clear_data_rows(quotes_ws, 5)
    clear_data_rows(lines_ws, 5)
    clear_data_rows(summary_ws, 5)

    fetched_ts = datetime.now(TW).replace(microsecond=0).isoformat()
    report['B3'] = FROM_DATE; report['B3'].number_format = 'yyyy-mm-dd'
    report['D3'] = TO_DATE; report['D3'].number_format = 'yyyy-mm-dd'
    report['B4'] = f'Status={STATUS}; QuoteDate between {FROM_DATE} and {TO_DATE} inclusive'
    report['B5'] = RUN_TS
    report['B7'] = f'{SOURCE_NAME}; fx_rates.csv'
    report['D7'] = fetched_ts
    report['B9'] = 'No exceptions. FX rates are fixed test rates from input/fx_rates.csv; stdout used only after exit code 0 and JSON parse success.'

    expected_q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
    if [quotes_ws.cell(4,c).value for c in range(1,12)] != expected_q_headers:
        raise ValueError('Quotes header mismatch')

    for idx, h in enumerate(details, start=5):
        style_data_row(quotes_ws, idx, len(expected_q_headers))
        set_id_cell(quotes_ws.cell(idx, 1), h.get('QuoteID'))
        set_id_cell(quotes_ws.cell(idx, 2), h.get('CustomerID'))
        set_id_cell(quotes_ws.cell(idx, 3), h.get('ProjectID'))
        set_date_cell(quotes_ws.cell(idx, 4), h.get('QuoteDate'))
        set_date_cell(quotes_ws.cell(idx, 5), h.get('ValidUntil'))
        quotes_ws.cell(idx, 6).value = h.get('Status')
        quotes_ws.cell(idx, 7).value = h.get('Currency')
        subtotal = q2(as_decimal(h.get('Subtotal')))
        discount = as_decimal(h.get('DiscountPercent'))
        net = q2(as_decimal(h.get('NetAmount')))
        expected_net = q2(subtotal * (Decimal('1') - discount / Decimal('100')))
        if net != expected_net:
            raise ValueError(f"NetAmount mismatch for {h.get('QuoteID')}: CLI={net} expected={expected_net}")
        set_money_cell(quotes_ws.cell(idx, 8), subtotal)
        set_number_cell(quotes_ws.cell(idx, 9), discount); quotes_ws.cell(idx, 9).number_format = '0.00'
        set_money_cell(quotes_ws.cell(idx, 10), net)
        cur = h.get('Currency')
        if cur not in rates:
            raise ValueError(f'Missing FX rate for {cur}')
        set_money_cell(quotes_ws.cell(idx, 11), q2(net * rates[cur]))

    expected_l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
    if [lines_ws.cell(4,c).value for c in range(1,7)] != expected_l_headers:
        raise ValueError('Lines header mismatch')
    for idx, l in enumerate(lines, start=5):
        style_data_row(lines_ws, idx, len(expected_l_headers))
        set_id_cell(lines_ws.cell(idx, 1), l.get('QuoteID'))
        set_id_cell(lines_ws.cell(idx, 2), l.get('ProductID'))
        lines_ws.cell(idx, 3).value = l.get('ProductName')
        set_number_cell(lines_ws.cell(idx, 4), l.get('Quantity'))
        set_money_cell(lines_ws.cell(idx, 5), q2(as_decimal(l.get('UnitPrice'))))
        amount = q2(as_decimal(l.get('LineAmount')))
        expected_amount = q2(as_decimal(l.get('Quantity')) * as_decimal(l.get('UnitPrice')))
        if amount != expected_amount:
            raise ValueError(f"LineAmount mismatch for {l.get('QuoteID')} {l.get('ProductID')}: CLI={amount} expected={expected_amount}")
        set_money_cell(lines_ws.cell(idx, 6), amount)

    by_currency = {}
    for h in details:
        cur = h['Currency']
        net = q2(as_decimal(h['NetAmount']))
        twd = q2(net * rates[cur])
        by_currency.setdefault(cur, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
        by_currency[cur]['amount'] += net
        by_currency[cur]['twd'] += twd
        by_currency[cur]['count'] += 1
    summary_rows = [('Status', STATUS, cur, q2(v['amount']), q2(v['twd']), v['count']) for cur, v in sorted(by_currency.items())]
    summary_ws['B2'] = f'=SUM(Quotes!K5:K{4+len(details)})' if details else 0
    summary_ws['B2'].number_format = '0.00'
    for idx, row in enumerate(summary_rows, start=5):
        style_data_row(summary_ws, idx, 6)
        for col, value in enumerate(row, start=1):
            cell = summary_ws.cell(idx, col)
            if col in (4,5):
                set_money_cell(cell, value)
            else:
                cell.value = value

    wb.save(OUTFILE)
    FETCH_LOG.write_text(json.dumps({
        'from': FROM_DATE, 'to': TO_DATE, 'status': STATUS,
        'run_timestamp': RUN_TS, 'fetched_timestamp': fetched_ts,
        'attempts': all_attempts,
        'list_count': list_data.get('count'), 'detail_count': len(details), 'line_count': len(lines),
    }, ensure_ascii=False, indent=2), encoding='utf-8')

    wb2 = load_workbook(OUTFILE, data_only=False)
    rq, rl, rs, rr = wb2['Quotes'], wb2['Lines'], wb2['Summary'], wb2['Report']
    checks = []
    checks.append(('output_file_nonempty', OUTFILE.exists() and OUTFILE.stat().st_size > 0))
    checks.append(('fetch_log_nonempty', FETCH_LOG.exists() and FETCH_LOG.stat().st_size > 0))
    checks.append(('report_period', rr['B3'].value == FROM_DATE and rr['D3'].value == TO_DATE))
    checks.append(('report_condition', STATUS in str(rr['B4'].value) and FROM_DATE in str(rr['B4'].value) and TO_DATE in str(rr['B4'].value)))
    checks.append(('quotes_headers_preserved', [rq.cell(4,c).value for c in range(1,12)] == expected_q_headers))
    checks.append(('lines_headers_preserved', [rl.cell(4,c).value for c in range(1,7)] == expected_l_headers))
    checks.append(('quote_row_count', rq.max_row == 4 + len(details)))
    checks.append(('line_row_count', rl.max_row == 4 + len(lines)))
    checks.append(('summary_row_count', rs.max_row == (4 + len(summary_rows) if summary_rows else 4)))
    quote_ids = []
    for r in range(5, 5 + len(details)):
        quote_ids.append(str(rq.cell(r,1).value))
        checks.append((f'quote_{r}_ids_text_or_blank', all(nonempty_text_or_blank(rq.cell(r,c)) for c in (1,2,3))))
        checks.append((f'quote_{r}_date_filter', FROM_DATE <= rq.cell(r,4).value <= TO_DATE))
        checks.append((f'quote_{r}_status_sent', rq.cell(r,6).value == STATUS))
        checks.append((f'quote_{r}_amounts_numeric', all(isinstance(rq.cell(r,c).value, (int,float)) for c in (8,9,10,11))))
    for r in range(5, 5 + len(lines)):
        checks.append((f'line_{r}_ids_text_or_blank', all(nonempty_text_or_blank(rl.cell(r,c)) for c in (1,2))))
        checks.append((f'line_{r}_numbers_numeric', isinstance(rl.cell(r,4).value, (int,float)) and isinstance(rl.cell(r,5).value, (int,float)) and isinstance(rl.cell(r,6).value, (int,float))))
    expected_formula = f'=SUM(Quotes!K5:K{4+len(details)})' if details else 0
    checks.append(('summary_total_formula', rs['B2'].value == expected_formula))
    for r in range(5, 5 + len(summary_rows)):
        checks.append((f'summary_{r}_amounts_numeric', isinstance(rs.cell(r,4).value, (int,float)) and isinstance(rs.cell(r,5).value, (int,float)) and isinstance(rs.cell(r,6).value, int)))
    failures = [name for name, ok in checks if not ok]
    if failures:
        raise AssertionError('Verification failed: ' + ', '.join(failures))

    saved_total_twd = q2(sum(Decimal(str(rq.cell(r,11).value)) for r in range(5, 5 + len(details)))) if details else Decimal('0.00')
    print('\nVERIFICATION PASSED')
    print('Output workbook:', OUTFILE)
    print('Fetch log:', FETCH_LOG)
    print('Quotes rows:', len(details))
    print('Line item rows:', len(lines))
    print('Summary rows:', len(summary_rows))
    print('Saved NetAmountTWD total:', saved_total_twd)
    print('QuoteIDs:', ', '.join(quote_ids) if quote_ids else '(none)')
    print('CLI attempts:', len(all_attempts))
    print('Exit code 4 attempts:', sum(1 for a in all_attempts if a['exit_code'] == 4))
    print('Exit code 2/3 attempts:', [a for a in all_attempts if a['exit_code'] in (2,3)])

if __name__ == '__main__':
    main()
