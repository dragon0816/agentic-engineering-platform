import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', sorted(str(p.relative_to(ws)) for p in (ws/'input').glob('*')))
# fx sample
fx=ws/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        rows=[]
        for i,row in enumerate(r):
            rows.append(row)
            if i>=5: break
    print('fx_rates sample:', rows)
# manual excerpt
man=ws/'input'/'sap_manual.md'
if man.exists():
    txt=man.read_text(encoding='utf-8')
    print('sap_manual first 4000 chars:\n', txt[:4000])
# workbook info
path=ws/'input'/'sales_template.xlsx'
if path.exists():
    wb=load_workbook(path)
    print('workbook sheets:', wb.sheetnames)
    for sh in wb.sheetnames:
        wsht=wb[sh]
        print('sheet', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
        for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
            print(row)
        print('---')
