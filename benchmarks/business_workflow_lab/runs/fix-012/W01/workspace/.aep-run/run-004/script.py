from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from collections import defaultdict
from openpyxl import load_workbook
from copy import copy

workspace = Path.cwd()
input_dir = workspace/'input'
output_dir = workspace/'output'
output_dir.mkdir(exist_ok=True)

csv_path = output_dir/'downloads'/'orders_export.csv'
template_path = input_dir/'sales_template.xlsx'
fx_path = input_dir/'fx_rates.csv'
out_path = output_dir/'september_open_orders.xlsx'

print('inputs:', csv_path, csv_path.exists(), template_path, template_path.exists(), fx_path, fx_path.exists())

# Read source CSV exported through browser from simulated SAP
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    orders = list(csv.DictReader(f))
print('csv columns:', orders[0].keys() if orders else [])
print('csv row count:', len(orders))
print('csv sample:', orders[:2])

# Detail fields observed through browser detail pages for these exported orders
details = {
    '0045000109': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10'},
    '0045000110': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000112': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15'},
    '0045000113': {'CustomerContact': '', 'PaymentTerms': 'NET15'},
    '0045000115': {'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET20'},
    '0045000116': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
    '0045000118': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET0'},
    '0045000119': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000120': {'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET10'},
    '0045000122': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET5'},
    '0045000124': {'CustomerContact': '', 'PaymentTerms': 'NET30'},
    '0045000126': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
    '0045000127': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET30'},
    '0045000128': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET5'},
}

# Read FX; CSV notes contain commas, so only use first two columns.
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0]:
            continue
        fx[row[0].strip()] = Decimal(row[1].strip())
print('fx:', fx)

q2 = Decimal('0.01')
def money(d):
    return Decimal(str(d)).quantize(q2, rounding=ROUND_HALF_UP)

orders.sort(key=lambda r: r['OrderID'])

wb = load_workbook(template_path)
wsr = wb['Report']
wsd = wb['Details']
wss = wb['Summary']

# Report fields
now = datetime.now().replace(microsecond=0).isoformat()
wsr['B3'] = '2026-09-01'
wsr['D3'] = '2026-09-30'
wsr['B4'] = 'Sales Orders: Status=Open; OrderDate 2026-09-01 to 2026-09-30; sorted by OrderID asc'
wsr['B5'] = now
wsr['B7'] = 'Lab ERP simulated SAP Sales Orders; input/fx_rates.csv'
wsr['D7'] = now
wsr['B9'] = 'Export CSV (all filtered rows) used for list fields; detail pages used for CustomerContact and PaymentTerms. Blank source values are kept blank.'

# Clear old Details data rows and Summary rows if present
for ws, start_row in [(wsd,5),(wss,5)]:
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row-start_row+1)

# Copy style from row 5 if available after insert? Template has none beyond row 4, use row 4 styles lightly.
headers = [wsd.cell(4,c).value for c in range(1, wsd.max_column+1)]
print('detail headers:', headers)

summary = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
last_detail_row = 4
for idx, r in enumerate(orders, start=5):
    oid = r['OrderID']
    curr = r['Currency']
    amt = money(r['NetAmount'])
    if curr not in fx:
        raise ValueError(f'Missing FX rate for {curr}')
    twd = money(amt * fx[curr])
    det = details.get(oid, {'CustomerContact':'', 'PaymentTerms':''})
    values = [
        oid,
        r['CustomerID'],
        r['CustomerName'],
        r.get('ProjectID',''),
        r['OrderDate'],
        r['Status'],
        curr,
        float(amt),
        float(twd),
        det.get('CustomerContact',''),
        det.get('PaymentTerms',''),
    ]
    for c, v in enumerate(values, start=1):
        cell = wsd.cell(idx, c, v)
        # Preserve basic row/header style where useful
        if c <= wsd.max_column:
            cell._style = copy(wsd.cell(4,c)._style)
        if c in [1,2,4]:
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in [8,9]:
            cell.number_format = '0.00'
    summary[curr]['amount'] += amt
    summary[curr]['twd'] += twd
    summary[curr]['count'] += 1
    last_detail_row = idx

# Summary formula
if orders:
    wss['B2'] = f'=SUM(Details!I5:I{last_detail_row})'
else:
    wss['B2'] = '=SUM(Details!I5:I5)'
wss['B2'].number_format = '0.00'

# Summary by currency and total TWD
row = 5
for curr in sorted(summary):
    s = summary[curr]
    vals = ['Currency', curr, curr, float(money(s['amount'])), float(money(s['twd'])), s['count']]
    for c,v in enumerate(vals, start=1):
        cell = wss.cell(row,c,v)
        cell._style = copy(wss.cell(4,c)._style)
        if c in [4,5]: cell.number_format = '0.00'
    row += 1

total_twd = sum((s['twd'] for s in summary.values()), Decimal('0'))
total_count = sum((s['count'] for s in summary.values()), 0)
vals = ['Total', 'All', 'TWD', None, float(money(total_twd)), total_count]
for c,v in enumerate(vals, start=1):
    cell = wss.cell(row,c,v)
    cell._style = copy(wss.cell(4,c)._style)
    if c in [4,5]: cell.number_format = '0.00'

# Re-apply text formats for ID columns
for r in range(5, last_detail_row+1):
    for c in [1,2,4]:
        wsd.cell(r,c).number_format='@'
    for c in [8,9]:
        wsd.cell(r,c).number_format='0.00'

wb.save(out_path)
print('wrote', out_path)

# Verification: reopen output and check requirements
wb2 = load_workbook(out_path, data_only=False)
wsd2 = wb2['Details']
wss2 = wb2['Summary']
rows = []
for r in range(5, wsd2.max_row+1):
    oid = wsd2.cell(r,1).value
    if oid is not None:
        rows.append([wsd2.cell(r,c).value for c in range(1,12)])
print('check rows:', len(rows))
print('first/last OrderID:', rows[0][0] if rows else None, rows[-1][0] if rows else None)
print('sorted:', [r[0] for r in rows] == sorted([r[0] for r in rows]))
print('all status Open:', all(r[5]=='Open' for r in rows))
print('all Sep 2026:', all('2026-09-01' <= str(r[4]) <= '2026-09-30' for r in rows))
print('OrderID text format first:', wsd2.cell(5,1).number_format, type(wsd2.cell(5,1).value).__name__, wsd2.cell(5,1).value)
print('summary formula B2:', wss2['B2'].value)
print('summary rows:')
for r in range(5, wss2.max_row+1):
    vals=[wss2.cell(r,c).value for c in range(1,7)]
    if any(v is not None for v in vals): print(vals)
# compute expected totals
print('computed total twd:', float(money(total_twd)))
print('output exists bytes:', out_path.stat().st_size)
