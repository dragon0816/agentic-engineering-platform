import os, json, csv
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
start_date = params['start_date']
end_date = params['end_date']
status = params['status']
output_filename = params['output_filename']
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
download_dir = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)

csv_path = download_dir / 'sales_orders_export.csv'
template_path = input_dir / 'sales_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
out_path = output_dir / output_filename

print('workflow params:', params)
print('using files:', {'csv': str(csv_path), 'template': str(template_path), 'fx': str(fx_path), 'output': str(out_path)})
for p in [csv_path, template_path, fx_path]:
    if not p.exists():
        raise FileNotFoundError(str(p))

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    orders = list(csv.DictReader(f))
print('downloaded row count:', len(orders))
print('downloaded columns:', list(orders[0].keys()) if orders else [])
print('downloaded sample:', orders[:2])

fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader, None)
    for row in reader:
        if row and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
print('fx currencies:', sorted(fx))

details = {
 '0045000101': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET15'},
 '0045000102': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET30'},
 '0045000103': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET30'},
 '0045000104': {'CustomerContact': '', 'PaymentTerms': 'NET0'},
 '0045000105': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15'},
 '0045000106': {'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET30'},
 '0045000107': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
 '0045000108': {'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET15'},
 '0045000109': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10'},
 '0045000110': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET30'},
 '0045000111': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET30'},
 '0045000112': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15'},
 '0045000113': {'CustomerContact': '', 'PaymentTerms': 'NET15'},
 '0045000114': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET5'},
 '0045000115': {'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET20'},
 '0045000116': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
 '0045000117': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET10'},
 '0045000118': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET0'},
 '0045000119': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
 '0045000120': {'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET10'},
 '0045000121': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET30'},
 '0045000122': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET5'},
 '0045000123': {'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET10'},
 '0045000124': {'CustomerContact': '', 'PaymentTerms': 'NET30'},
 '0045000125': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET0'},
 '0045000126': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
 '0045000127': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET30'},
 '0045000128': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET5'},
 '0045000129': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
 '0045000130': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET5'},
 '0045000131': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET15'},
 '0045000132': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': ''},
 '0045000133': {'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET30'},
 '0045000134': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15'},
 '0045000135': {'CustomerContact': '', 'PaymentTerms': 'NET30'},
 '0045000136': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET5'},
}

q2 = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(q2, rounding=ROUND_HALF_UP)

orders.sort(key=lambda r: r['OrderID'])

wb = load_workbook(template_path)
wsr = wb['Report']
wsd = wb['Details']
wss = wb['Summary']

now = datetime.now().replace(microsecond=0).isoformat()
wsr['B3'] = start_date
wsr['D3'] = end_date
wsr['B4'] = f'Sales Orders: Status={status}; OrderDate {start_date} to {end_date}; sorted by OrderID asc'
wsr['B5'] = now
wsr['B7'] = 'Lab ERP simulated SAP Sales Orders; input/fx_rates.csv'
wsr['D7'] = now

for ws, start_row in [(wsd,5),(wss,5)]:
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row-start_row+1)

summary = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
missing_detail_ids = []
last_detail_row = 4
for idx, r in enumerate(orders, start=5):
    oid = r['OrderID']
    curr = r['Currency']
    amt = money(r['NetAmount'])
    if curr not in fx:
        raise ValueError(f'Missing FX rate for {curr}')
    twd = money(amt * fx[curr])
    det = details.get(oid, {'CustomerContact':'', 'PaymentTerms':''})
    if oid not in details:
        missing_detail_ids.append(oid)
    values = [oid, r['CustomerID'], r['CustomerName'], r.get('ProjectID',''), r['OrderDate'], r['Status'], curr, float(amt), float(twd), det.get('CustomerContact',''), det.get('PaymentTerms','')]
    for c, v in enumerate(values, start=1):
        cell = wsd.cell(idx, c, v)
        cell._style = copy(wsd.cell(4,c)._style)
        if c in [1,2,4]: cell.number_format = '@'
        elif c == 5: cell.number_format = 'yyyy-mm-dd'
        elif c in [8,9]: cell.number_format = '0.00'
    summary[curr]['amount'] += amt
    summary[curr]['twd'] += twd
    summary[curr]['count'] += 1
    last_detail_row = idx

wss['B2'] = f'=SUM(Details!I5:I{last_detail_row})' if orders else '=SUM(Details!I5:I5)'
wss['B2'].number_format = '0.00'
row = 5
for curr in sorted(summary):
    s = summary[curr]
    vals = ['Currency', curr, curr, float(money(s['amount'])), float(money(s['twd'])), s['count']]
    for c, v in enumerate(vals, start=1):
        cell = wss.cell(row, c, v)
        cell._style = copy(wss.cell(4,c)._style)
        if c in [4,5]: cell.number_format = '0.00'
    row += 1

total_twd = money(sum((s['twd'] for s in summary.values()), Decimal('0')))
total_count = sum(s['count'] for s in summary.values())
for c, v in enumerate(['Total', 'All', 'TWD', None, float(total_twd), total_count], start=1):
    cell = wss.cell(row, c, v)
    cell._style = copy(wss.cell(4,c)._style)
    if c in [4,5]: cell.number_format = '0.00'

if missing_detail_ids:
    wsr['B9'] = 'Detail-only fields were not available for these OrderIDs and were left blank: ' + ', '.join(missing_detail_ids)
else:
    wsr['B9'] = 'Export CSV (all filtered rows) used for list fields; browser-read detail pages used for CustomerContact and PaymentTerms. Blank source values are kept blank.'

wb.save(out_path)
print('wrote:', out_path)

wb2 = load_workbook(out_path, data_only=False)
wsd2 = wb2['Details']
wss2 = wb2['Summary']
rows=[]
for rr in range(5, wsd2.max_row+1):
    if wsd2.cell(rr,1).value is not None:
        rows.append([wsd2.cell(rr,c).value for c in range(1,12)])
expected_total = sum(Decimal(str(r[8])).quantize(q2) for r in rows)
print('CHECK output_file_exists:', out_path.exists(), 'bytes', out_path.stat().st_size)
print('CHECK detail_row_count:', len(rows))
print('CHECK sorted_orderid:', [r[0] for r in rows] == sorted([r[0] for r in rows]))
print('CHECK status_matches:', all(r[5] == status for r in rows))
print('CHECK date_range:', all(start_date <= str(r[4]) <= end_date for r in rows))
print('CHECK orderid_text_type:', all(isinstance(r[0], str) for r in rows))
print('CHECK orderid_number_format:', wsd2.cell(5,1).number_format if rows else 'no rows')
print('CHECK summary_formula:', wss2['B2'].value)
print('CHECK summary_total_twd:', total_twd, 'expected_from_details:', expected_total)
print('CHECK missing_detail_ids:', missing_detail_ids)
print('OUTPUT_FILES:', [str(out_path.relative_to(workspace))])
