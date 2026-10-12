import os, json, csv
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
start_date = params['start_date']
end_date = params['end_date']
status = params['status']
output_filename = params['output_filename']
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
download_dir = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)

csv_path = download_dir / 'sales_orders_export.csv'
template_path = input_dir / 'sales_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
out_path = output_dir / output_filename

print('workflow params:', params)
print('using files:', {'csv': str(csv_path), 'template': str(template_path), 'fx': str(fx_path), 'output': str(out_path)})
for p in [csv_path, template_path, fx_path]:
    if not p.exists():
        raise FileNotFoundError(str(p))

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    orders = list(csv.DictReader(f))
print('downloaded row count:', len(orders))
print('downloaded columns:', list(orders[0].keys()) if orders else [])
print('downloaded sample:', orders[:2])

fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader, None)
    for row in reader:
        if row and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
print('fx currencies:', sorted(fx))

q2 = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(q2, rounding=ROUND_HALF_UP)

orders.sort(key=lambda r: r['OrderID'])

wb = load_workbook(template_path)
wsr = wb['Report']
wsd = wb['Details']
wss = wb['Summary']

now = datetime.now().replace(microsecond=0).isoformat()
wsr['B3'] = start_date
wsr['D3'] = end_date
wsr['B4'] = f'Sales Orders: Status={status}; OrderDate {start_date} to {end_date}; sorted by OrderID asc'
wsr['B5'] = now
wsr['B7'] = 'Lab ERP simulated SAP Sales Orders; input/fx_rates.csv'
wsr['D7'] = now

for ws, start_row in [(wsd,5),(wss,5)]:
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row-start_row+1)

summary = defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
last_detail_row = 4
for idx, r in enumerate(orders, start=5):
    oid = r['OrderID']
    curr = r['Currency']
    amt = money(r['NetAmount'])
    if curr not in fx:
        raise ValueError(f'Missing FX rate for {curr}')
    twd = money(amt * fx[curr])
    values = [oid, r['CustomerID'], r['CustomerName'], r.get('ProjectID',''), r['OrderDate'], r['Status'], curr, float(amt), float(twd), '', '']
    for c, v in enumerate(values, start=1):
        cell = wsd.cell(idx, c, v)
        cell._style = copy(wsd.cell(4,c)._style)
        if c in [1,2,4]:
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in [8,9]:
            cell.number_format = '0.00'
    summary[curr]['amount'] += amt
    summary[curr]['twd'] += twd
    summary[curr]['count'] += 1
    last_detail_row = idx

wss['B2'] = f'=SUM(Details!I5:I{last_detail_row})' if orders else '=SUM(Details!I5:I5)'
wss['B2'].number_format = '0.00'
row = 5
for curr in sorted(summary):
    s = summary[curr]
    vals = ['Currency', curr, curr, float(money(s['amount'])), float(money(s['twd'])), s['count']]
    for c, v in enumerate(vals, start=1):
        cell = wss.cell(row, c, v)
        cell._style = copy(wss.cell(4,c)._style)
        if c in [4,5]:
            cell.number_format = '0.00'
    row += 1

total_twd = money(sum((s['twd'] for s in summary.values()), Decimal('0')))
total_count = sum(s['count'] for s in summary.values())
for c, v in enumerate(['Total', 'All', 'TWD', None, float(total_twd), total_count], start=1):
    cell = wss.cell(row, c, v)
    cell._style = copy(wss.cell(4,c)._style)
    if c in [4,5]:
        cell.number_format = '0.00'

wsr['B9'] = 'Export CSV (all filtered rows) used for order list fields. CustomerContact and PaymentTerms are not in the export and are left blank.'

wb.save(out_path)
print('wrote:', out_path)

# Reopen and verify saved workbook requirements.
wb2 = load_workbook(out_path, data_only=False)
wsd2 = wb2['Details']
wss2 = wb2['Summary']
rows=[]
for rr in range(5, wsd2.max_row+1):
    if wsd2.cell(rr,1).value is not None:
        rows.append([wsd2.cell(rr,c).value for c in range(1,12)])
expected_total = sum(Decimal(str(r[8])).quantize(q2) for r in rows)
order_ids = [r[0] for r in rows]
checks = {
    'output_file_exists': out_path.exists() and out_path.stat().st_size > 0,
    'detail_row_count_matches_export': len(rows) == len(orders),
    'sorted_orderid': order_ids == sorted(order_ids),
    'status_matches': all(r[5] == status for r in rows),
    'date_range': all(start_date <= str(r[4]) <= end_date for r in rows),
    'orderid_text_type': all(isinstance(r[0], str) for r in rows),
    'summary_formula': wss2['B2'].value == (f'=SUM(Details!I5:I{last_detail_row})' if orders else '=SUM(Details!I5:I5)'),
    'summary_total_twd': money(sum((s['twd'] for s in summary.values()), Decimal('0'))) == expected_total,
}
for k, v in checks.items():
    print('CHECK', k, v)
if not all(checks.values()):
    raise AssertionError(checks)
print('OUTPUT_FILES:', [str(out_path.relative_to(workspace))])
