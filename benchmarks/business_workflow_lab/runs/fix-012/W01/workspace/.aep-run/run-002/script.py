from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
for sh in ['Instructions']:
    ws=wb[sh]
    for r in range(1, ws.max_row+1):
        print(r, [ws.cell(r,c).value for c in range(1, ws.max_column+1)])
