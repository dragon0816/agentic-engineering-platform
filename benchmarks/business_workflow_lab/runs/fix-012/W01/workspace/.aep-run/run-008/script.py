from pathlib import Path
from openpyxl import load_workbook
import csv, json, os

ws = Path.cwd()
out = ws/'output'
dl = out/'downloads'
print('workspace', ws)
print('output files', [str(p.relative_to(ws)) for p in sorted(out.rglob('*')) if p.is_file()])

csv_path = dl/'sales_orders_export.csv'
xlsx_path = out/'september_open_orders.xlsx'
print('csv_exists', csv_path.exists(), 'size', csv_path.stat().st_size if csv_path.exists() else None)
print('xlsx_exists', xlsx_path.exists(), 'size', xlsx_path.stat().st_size if xlsx_path.exists() else None)

if csv_path.exists():
    with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
        rows = list(csv.DictReader(f))
    print('csv_columns', rows[0].keys() if rows else [])
    print('csv_row_count', len(rows))
    print('csv_first_last_order', (rows[0]['OrderID'], rows[-1]['OrderID']) if rows else None)
    print('csv_statuses', sorted(set(r['Status'] for r in rows)))
    print('csv_dates_min_max', (min(r['OrderDate'] for r in rows), max(r['OrderDate'] for r in rows)) if rows else None)
    print('csv_sorted_by_orderid', [r['OrderID'] for r in rows] == sorted(r['OrderID'] for r in rows))

if xlsx_path.exists():
    wb = load_workbook(xlsx_path, data_only=False)
    print('sheets', wb.sheetnames)
    for s in wb.sheetnames:
        sh=wb[s]
        print('sheet', s, 'max_row', sh.max_row, 'max_column', sh.max_column)
    details = wb['Details'] if 'Details' in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row = None
    headers = []
    for r in range(1, min(details.max_row, 20)+1):
        vals=[details.cell(r,c).value for c in range(1, details.max_column+1)]
        if 'OrderID' in vals:
            header_row=r; headers=vals; break
    print('details_header_row', header_row)
    print('details_headers', headers)
    if header_row:
        data=[]
        for r in range(header_row+1, details.max_row+1):
            vals=[details.cell(r,c).value for c in range(1, details.max_column+1)]
            if any(v is not None for v in vals):
                data.append(dict(zip(headers, vals)))
        print('details_data_rows', len(data))
        print('details_orderids', [str(d.get('OrderID')) for d in data])
        print('details_sorted_by_orderid', [str(d.get('OrderID')) for d in data] == sorted(str(d.get('OrderID')) for d in data))
        print('details_statuses', sorted(set(str(d.get('Status')) for d in data)))
        print('details_dates_min_max', (min(str(d.get('OrderDate')) for d in data), max(str(d.get('OrderDate')) for d in data)) if data else None)
        print('details_currencies', sorted(set(str(d.get('Currency')) for d in data)))
    if 'Summary' in wb.sheetnames:
        sh=wb['Summary']
        rows=[]
        for r in range(1, sh.max_row+1):
            vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
            if any(v is not None for v in vals): rows.append(vals)
        print('summary_rows')
        for row in rows:
            print(row)
