import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:', sorted([p.name for p in (ws/'input').glob('*')]))
# Template details
 tpl = ws/'input'/'expense_template.xlsx'
print('template exists', tpl.exists(), 'size', tpl.stat().st_size if tpl.exists() else None)
if tpl.exists():
    wb = load_workbook(tpl, data_only=False)
    print('template sheets:', wb.sheetnames)
    for sh in wb.sheetnames:
        sheet = wb[sh]
        print('sheet', sh, 'max_row', sheet.max_row, 'max_col', sheet.max_column)
        for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 8), values_only=True):
            print(sh, row)
# fx rates headers/sample
fx = ws/'input'/'fx_rates.csv'
print('fx exists', fx.exists(), 'size', fx.stat().st_size if fx.exists() else None)
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        rdr=csv.reader(f)
        for i,row in zip(range(8),rdr): print('fx row', i+1, row)
# manual first lines
man = ws/'input'/'expense_manual.md'
print('manual exists', man.exists(), 'size', man.stat().st_size if man.exists() else None)
if man.exists():
    print(man.read_text(encoding='utf-8-sig')[:2000])