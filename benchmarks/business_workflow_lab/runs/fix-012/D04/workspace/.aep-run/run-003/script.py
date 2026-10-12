from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
print('Merged ranges:')
for sh in wb.sheetnames:
    print(sh, list(wb[sh].merged_cells.ranges))
print('Report dims/styles sample')
