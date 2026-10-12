from pathlib import Path
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
from openpyxl.styles import Alignment

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

employee = 'E006'
from_date = '2026-01-01'
to_date = '2026-01-31'
now_iso = '2026-10-06T00:00:00+08:00'
source_name = 'Expense Viewer'
# Proven by UI query: Total 0 rows · Page 0 / 0 and no-results dialog.
rows = []

# Inspect inputs again as first stage for this producing script.
print('Inputs found:', sorted(p.name for p in input_dir.glob('*')))
print('Query filter:', {'EmployeeID': employee, 'DateFrom': from_date, 'DateTo': to_date, 'Status': '(all)'})
print('Queried row count from Expense Viewer UI: 0')
print('Export file used: none (0 rows; export buttons disabled, so no nonexistent file is treated as success)')

# Load FX for completeness and validation of input availability.
fx_path = input_dir / 'fx_rates.csv'
fx_rates = {}
if fx_path.exists():
    import csv
    with fx_path.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        for r in reader:
            if r.get('Currency'):
                fx_rates[r['Currency']] = Decimal(str(r['RateToTWD']))
print('FX currencies:', sorted(fx_rates))

# Fill a copy of the template; do not rebuild workbook.
template = input_dir / 'expense_template.xlsx'
outfile = output_dir / f'expense_report_{employee}_{from_date}_to_{to_date}.xlsx'
wb = load_workbook(template)

# Report sheet
ws = wb['Report']
ws['B3'] = from_date
ws['D3'] = to_date
ws['B4'] = f'EmployeeID={employee}; ExpenseDate {from_date}..{to_date}; Status=ALL'
ws['B5'] = now_iso
ws['B7'] = source_name
ws['D7'] = now_iso
ws['B9'] = '查無符合條件的費用資料（0 筆）；Expense Viewer 於 0 筆結果時停用匯出，因此未產生匯出檔。'

# Clear old detail/summary/exception data rows if any while preserving headers and template structure.
for sheet_name in ['Details', 'Summary', 'Exceptions']:
    sh = wb[sheet_name]
    if sh.max_row >= 5:
        sh.delete_rows(5, sh.max_row - 4)

# For a normal empty report, leave data sheets with only headers (row 4) and no blank data rows.
# Preserve Summary B2 formula for zero rows per instruction.
summary = wb['Summary']
summary['B2'] = '=SUM(Details!I5:I5)'

# Save.
wb.save(outfile)
print('Wrote:', outfile)

# Verification by reopening output.
check_wb = load_workbook(outfile, data_only=False)
errors = []
# Sheets preserved
expected_sheets = ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
if check_wb.sheetnames != expected_sheets:
    errors.append(f'Sheet names/order mismatch: {check_wb.sheetnames}')
rep = check_wb['Report']
checks = {
    'B3': from_date,
    'D3': to_date,
    'B4': f'EmployeeID={employee}; ExpenseDate {from_date}..{to_date}; Status=ALL',
    'B5': now_iso,
    'B7': source_name,
    'D7': now_iso,
}
for cell, expected in checks.items():
    if rep[cell].value != expected:
        errors.append(f'Report {cell} expected {expected!r}, got {rep[cell].value!r}')
if '0 筆' not in str(rep['B9'].value):
    errors.append('Report B9 does not note zero rows')
# Header rows intact
expected_details_header = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
actual_details_header = [check_wb['Details'].cell(4,c).value for c in range(1,15)]
if actual_details_header != expected_details_header:
    errors.append(f'Details header mismatch: {actual_details_header}')
expected_summary_header = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
actual_summary_header = [check_wb['Summary'].cell(4,c).value for c in range(1,7)]
if actual_summary_header != expected_summary_header:
    errors.append(f'Summary header mismatch: {actual_summary_header}')
expected_exceptions_header = ['ExpenseID','EmployeeID','ExpenseDate','Reason']
actual_exceptions_header = [check_wb['Exceptions'].cell(4,c).value for c in range(1,5)]
if actual_exceptions_header != expected_exceptions_header:
    errors.append(f'Exceptions header mismatch: {actual_exceptions_header}')
# Row counts/no data rows
for sheet_name in ['Details','Summary','Exceptions']:
    sh = check_wb[sheet_name]
    if sh.max_row != 4:
        # openpyxl may keep dimension if formatting only; inspect actual values below row 4
        values = []
        for r in range(5, sh.max_row+1):
            row_values = [sh.cell(r,c).value for c in range(1, sh.max_column+1)]
            if any(v is not None for v in row_values):
                values.append((r,row_values))
        if values:
            errors.append(f'{sheet_name} has unexpected data rows: {values[:3]}')
if check_wb['Summary']['B2'].value != '=SUM(Details!I5:I5)':
    errors.append(f"Summary B2 formula mismatch: {check_wb['Summary']['B2'].value!r}")
# Merged cell preserved
if 'A1:F1' not in [str(r) for r in check_wb['Report'].merged_cells.ranges]:
    errors.append('Report A1:F1 merge not preserved')
# Output nonempty
if not outfile.exists() or outfile.stat().st_size <= 0:
    errors.append('Output file missing or empty')

print('Verification:')
print(' output_file_exists:', outfile.exists())
print(' output_size:', outfile.stat().st_size if outfile.exists() else None)
print(' sheets:', check_wb.sheetnames)
print(' report_B3_D3:', rep['B3'].value, rep['D3'].value)
print(' details_data_rows:', 0)
print(' summary_data_rows:', 0)
print(' exceptions_data_rows:', 0)
print(' summary_B2:', check_wb['Summary']['B2'].value)
if errors:
    print('ERRORS:')
    for e in errors:
        print('-', e)
    raise SystemExit(1)
print('All checks passed')