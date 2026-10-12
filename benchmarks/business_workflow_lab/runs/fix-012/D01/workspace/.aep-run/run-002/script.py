from pathlib import Path
import openpyxl, csv, os
ws=Path.cwd()
print('OUTPUT FILES')
for p in (ws/'output').glob('*'):
    print(p.name, p.stat().st_size)
exp=ws/'output'/'expense_export_E001_2026-09.xlsx'
wb=openpyxl.load_workbook(exp, data_only=True)
print('export sheets', wb.sheetnames)
for sh in wb.sheetnames:
    s=wb[sh]
    print('sheet', sh, s.max_row, s.max_column)
    for r in s.iter_rows(min_row=1, max_row=min(6,s.max_row), values_only=True): print(r)
print('Instructions full')
t=ws/'input'/'expense_template.xlsx'
wb2=openpyxl.load_workbook(t, data_only=True)
for row in wb2['Instructions'].iter_rows(values_only=True): print(row)