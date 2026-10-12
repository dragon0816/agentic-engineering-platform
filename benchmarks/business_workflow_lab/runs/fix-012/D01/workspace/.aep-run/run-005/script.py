from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
from datetime import datetime, timezone, timedelta
import csv, json, os, re
import openpyxl
from openpyxl.styles import Alignment

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
employee_id = params['employee_id']
from_date = params['from_date']
to_date = params['to_date']
output_filename = (params.get('output_filename') or '').strip()
if not output_filename:
    output_filename = f'{employee_id}_{from_date}_{to_date}_expenses.xlsx'
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'
# keep output inside workspace/output and make file name safe
output_filename = Path(output_filename).name

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'expense_template.xlsx'
FX = INPUT / 'fx_rates.csv'
EXPORT = OUTPUT / f'expense_export_{employee_id}_{from_date}_{to_date}.xlsx'
OUTFILE = OUTPUT / output_filename

print('inputs', {'template': str(TEMPLATE), 'fx': str(FX), 'export': str(EXPORT), 'output': str(OUTFILE)})
for p in [TEMPLATE, FX, EXPORT]:
    if not p.exists():
        raise FileNotFoundError(str(p))

rates = {}
with open(FX, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header', header)
    for row in reader:
        if row and row[0]:
            rates[row[0]] = Decimal(row[1])
print('rates', {k: str(v) for k, v in rates.items()})

exp_wb = openpyxl.load_workbook(EXPORT, data_only=True)
exp_ws = exp_wb[exp_wb.sheetnames[0]]
headers = [c.value for c in exp_ws[1]]
print('export headers', headers)
rows = []
for vals in exp_ws.iter_rows(min_row=2, values_only=True):
    if any(v is not None for v in vals):
        rows.append(dict(zip(headers, vals)))
print('export rows', len(rows))

required = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']
missing_headers = [h for h in required if h not in headers]
if missing_headers:
    raise ValueError('export missing headers: ' + ', '.join(missing_headers))

def q2(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def as_bool(v):
    if isinstance(v, bool):
        return v
    if v is None:
        return False
    return str(v).strip().lower() in ('true','1','yes','y')

def clone_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.protection = copy(src.protection)

def clear_data(ws, start_row, max_col):
    for r in range(start_row, ws.max_row + 1):
        for c in range(1, max_col + 1):
            ws.cell(r, c).value = None

wb = openpyxl.load_workbook(TEMPLATE)
now_tw = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report = wb['Report']
report['B3'] = from_date
report['D3'] = to_date
report['B4'] = f'EmployeeID={employee_id}; ExpenseDate={from_date}..{to_date}; Status=全部'
report['B5'] = now_tw
report['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D7'] = now_tw
report['B9'] = 'NeedsReview=true when MissingReceipt=true; Exceptions lists missing receipts.'

details = wb['Details']
clear_data(details, 5, 14)
for idx, rec in enumerate(rows, start=5):
    clone_row_style(details, 5, idx, 14)
    currency = str(rec.get('Currency') or '')
    if currency not in rates:
        raise ValueError(f'Missing FX rate for {currency}')
    amount = q2(rec.get('Amount') or '0')
    amount_twd = q2(amount * rates[currency])
    missing = as_bool(rec.get('MissingReceipt'))
    vals = [rec.get('ExpenseID'), rec.get('EmployeeID'), rec.get('EmployeeName'), rec.get('ProjectID'), rec.get('ExpenseDate'), rec.get('Category'), currency, float(amount), float(amount_twd), rec.get('Status'), rec.get('ReceiptID'), rec.get('Description'), 'true' if missing else 'false', 'true' if missing else 'false']
    for c, v in enumerate(vals, start=1):
        cell = details.cell(idx, c, v)
        if c in (1, 2, 4, 11):
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in (8, 9):
            cell.number_format = '0.00'
        elif c == 12:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
last_detail_row = 4 + len(rows) if rows else 5

summary = wb['Summary']
clear_data(summary, 5, 6)
summary['B2'] = f'=SUM(Details!I5:I{last_detail_row})'
groups = {}
for rec in rows:
    status = str(rec.get('Status') or '')
    currency = str(rec.get('Currency') or '')
    amount = q2(rec.get('Amount') or '0')
    amount_twd = q2(amount * rates[currency])
    key = (status, currency)
    groups.setdefault(key, {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0})
    groups[key]['amount'] += amount
    groups[key]['amount_twd'] += amount_twd
    groups[key]['count'] += 1
for r, ((status, currency), agg) in enumerate(sorted(groups.items()), start=5):
    clone_row_style(summary, 5, r, 6)
    vals = [status, '', currency, float(q2(agg['amount'])), float(q2(agg['amount_twd'])), agg['count']]
    for c, v in enumerate(vals, start=1):
        summary.cell(r, c).value = v
        if c in (4, 5):
            summary.cell(r, c).number_format = '0.00'

exceptions = wb['Exceptions']
clear_data(exceptions, 5, 4)
ex_rows = []
for rec in rows:
    if as_bool(rec.get('MissingReceipt')):
        ex_rows.append([rec.get('ExpenseID'), rec.get('EmployeeID'), rec.get('ExpenseDate'), 'MissingReceipt=true / receipt is missing'])
for r, vals in enumerate(ex_rows, start=5):
    clone_row_style(exceptions, 5, r, 4)
    for c, v in enumerate(vals, start=1):
        exceptions.cell(r, c).value = v
        if c in (1, 2):
            exceptions.cell(r, c).number_format = '@'
        elif c == 3:
            exceptions.cell(r, c).number_format = 'yyyy-mm-dd'

wb.save(OUTFILE)
print('saved', str(OUTFILE), OUTFILE.stat().st_size)

# Verification from saved output.
check = openpyxl.load_workbook(OUTFILE, data_only=False)
if check.sheetnames != ['Report','Details','Summary','Exceptions','Instructions']:
    raise AssertionError('sheet names changed: ' + repr(check.sheetnames))
cd = check['Details']
out_headers = [cd.cell(4, c).value for c in range(1, 15)]
expected = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
if out_headers != expected:
    raise AssertionError('detail headers mismatch')
detail_rows = []
for r in range(5, 5 + len(rows)):
    vals = [cd.cell(r, c).value for c in range(1, 15)]
    detail_rows.append(vals)
if len(detail_rows) != len(rows):
    raise AssertionError('detail row count mismatch')
if rows:
    if any(row[1] != employee_id for row in detail_rows):
        raise AssertionError('employee id mismatch in Details')
    if any(str(row[4]) < from_date or str(row[4]) > to_date for row in detail_rows):
        raise AssertionError('date outside range')
    if not all(isinstance(row[7], (int, float)) and isinstance(row[8], (int, float)) for row in detail_rows):
        raise AssertionError('amount columns must be numeric')
if check['Summary']['B2'].value != f'=SUM(Details!I5:I{last_detail_row})':
    raise AssertionError('Summary B2 formula mismatch')
computed_twd = sum(Decimal(str(row[8])) for row in detail_rows).quantize(Decimal('0.01'))
print('verified', {'detail_rows': len(detail_rows), 'summary_groups': len(groups), 'exceptions': len(ex_rows), 'twd_total': str(computed_twd), 'output_file': OUTFILE.name})
