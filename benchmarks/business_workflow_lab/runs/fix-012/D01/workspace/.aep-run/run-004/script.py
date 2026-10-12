from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
from datetime import datetime, timezone, timedelta
import csv, re, os, json
import openpyxl
from openpyxl.styles import Alignment

WORKSPACE = Path.cwd()
INPUT = WORKSPACE/'input'
OUTPUT = WORKSPACE/'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT/'expense_template.xlsx'
FX = INPUT/'fx_rates.csv'
EXPORT = OUTPUT/'expense_export_E001_2026-09.xlsx'
OUTFILE = OUTPUT/'E001_september_expenses.xlsx'

print('Inputs:')
for p in [TEMPLATE, FX, EXPORT]:
    print(p, 'exists=', p.exists(), 'size=', p.stat().st_size if p.exists() else None)

# Read fx rates (only first two columns are needed; notes may include unquoted commas in this lab file).
rates = {}
with open(FX, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0]] = Decimal(row[1])
print('rates:', rates)

# Read export workbook.
ewb = openpyxl.load_workbook(EXPORT, data_only=True)
es = ewb[ewb.sheetnames[0]]
headers = [c.value for c in es[1]]
print('export headers:', headers)
rows = []
for vals in es.iter_rows(min_row=2, values_only=True):
    if not any(v is not None for v in vals):
        continue
    rec = dict(zip(headers, vals))
    rows.append(rec)
print('export row count:', len(rows))
print('sample rows:', rows[:2])

# Helpers.
def q2(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def clone_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.protection:
            dst.protection = copy(src.protection)

def clear_data(ws, start_row, max_col):
    for r in range(start_row, ws.max_row + 1):
        for c in range(1, max_col+1):
            ws.cell(r,c).value = None

wb = openpyxl.load_workbook(TEMPLATE)
# Report
report = wb['Report']
generated_at = '2026-10-06T00:00:00+08:00'
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'EmployeeID=E001; ExpenseDate=2026-09-01..2026-09-30; Status=全部'
report['B5'] = generated_at
report['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D7'] = generated_at
report['B9'] = 'NeedsReview=true when MissingReceipt=true; Exceptions lists missing receipts.'

# Details
details = wb['Details']
clear_data(details, 5, 14)
# If previous data had many rows, do not delete rows to preserve template; write contiguous data from row 5.
for idx, rec in enumerate(rows, start=5):
    clone_row_style(details, 5, idx, 14)
    currency = str(rec.get('Currency') or '')
    amount = q2(rec.get('Amount') or '0')
    if currency not in rates:
        raise ValueError(f'Missing FX rate for {currency}')
    amount_twd = q2(amount * rates[currency])
    missing_raw = rec.get('MissingReceipt')
    missing = bool(missing_raw) if not isinstance(missing_raw, str) else missing_raw.strip().lower() == 'true'
    needs_review = missing
    values = [
        rec.get('ExpenseID'), rec.get('EmployeeID'), rec.get('EmployeeName'), rec.get('ProjectID'), rec.get('ExpenseDate'),
        rec.get('Category'), currency, float(amount), float(amount_twd), rec.get('Status'), rec.get('ReceiptID'),
        rec.get('Description'), 'true' if missing else 'false', 'true' if needs_review else 'false'
    ]
    for c, v in enumerate(values, start=1):
        cell = details.cell(idx, c, v)
        if c in [1,2,4,11]:
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in [8,9]:
            cell.number_format = '0.00'
        if c == 12:
            cell.alignment = copy(cell.alignment)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
last_detail = 4 + len(rows) if rows else 5

# Summary by Status and Currency.
summary = wb['Summary']
clear_data(summary, 5, 6)
summary['B2'] = f'=SUM(Details!I5:I{last_detail if rows else 5})'
summary_groups = {}
for rec in rows:
    status = str(rec.get('Status') or '')
    currency = str(rec.get('Currency') or '')
    amount = q2(rec.get('Amount') or '0')
    amount_twd = q2(amount * rates[currency])
    key = (status, currency)
    if key not in summary_groups:
        summary_groups[key] = {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0}
    summary_groups[key]['amount'] += amount
    summary_groups[key]['amount_twd'] += amount_twd
    summary_groups[key]['count'] += 1
for r, ((status, currency), agg) in enumerate(sorted(summary_groups.items()), start=5):
    clone_row_style(summary, 5, r, 6)
    vals = [status, '', currency, float(q2(agg['amount'])), float(q2(agg['amount_twd'])), agg['count']]
    for c, v in enumerate(vals, start=1):
        summary.cell(r,c).value = v
        if c in [4,5]:
            summary.cell(r,c).number_format = '0.00'

# Exceptions: missing receipts.
exceptions = wb['Exceptions']
clear_data(exceptions, 5, 4)
ex_rows = []
for rec in rows:
    missing_raw = rec.get('MissingReceipt')
    missing = bool(missing_raw) if not isinstance(missing_raw, str) else missing_raw.strip().lower() == 'true'
    if missing:
        ex_rows.append([rec.get('ExpenseID'), rec.get('EmployeeID'), rec.get('ExpenseDate'), 'MissingReceipt=true / receipt is missing'])
for r, vals in enumerate(ex_rows, start=5):
    clone_row_style(exceptions, 5, r, 4)
    for c, v in enumerate(vals, start=1):
        exceptions.cell(r,c).value = v
        if c in [1,2]:
            exceptions.cell(r,c).number_format = '@'
        if c == 3:
            exceptions.cell(r,c).number_format = 'yyyy-mm-dd'

wb.save(OUTFILE)
print('Saved', OUTFILE, OUTFILE.stat().st_size)

# Reopen and validate stated requirements and template-specific content.
check_wb = openpyxl.load_workbook(OUTFILE, data_only=False)
req_sheets = ['Report','Details','Summary','Exceptions','Instructions']
assert check_wb.sheetnames == req_sheets, check_wb.sheetnames
cd = check_wb['Details']
headers_out = [cd.cell(4,c).value for c in range(1,15)]
expected_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert headers_out == expected_headers, headers_out
actual_detail_rows = []
for r in range(5, 5+len(rows)):
    actual_detail_rows.append([cd.cell(r,c).value for c in range(1,15)])
assert len(actual_detail_rows) == 13
assert all(row[1] == 'E001' for row in actual_detail_rows)
assert min(row[4] for row in actual_detail_rows) >= '2026-09-01'
assert max(row[4] for row in actual_detail_rows) <= '2026-09-30'
assert all(isinstance(row[7], (int,float)) for row in actual_detail_rows), 'Amount not numeric'
assert all(isinstance(row[8], (int,float)) for row in actual_detail_rows), 'AmountTWD not numeric'
# Summary count and totals.
cs = check_wb['Summary']
assert cs['B2'].value == '=SUM(Details!I5:I17)', cs['B2'].value
summary_rows = []
r = 5
while any(cs.cell(r,c).value is not None for c in range(1,7)):
    summary_rows.append([cs.cell(r,c).value for c in range(1,7)])
    r += 1
assert len(summary_rows) == len(summary_groups)
# Reopen data_only after formula won't be calculated by openpyxl; compute total directly.
twd_total = sum(Decimal(str(row[8])) for row in actual_detail_rows).quantize(Decimal('0.01'))
expected_total = sum(q2(Decimal(str(rec['Amount'])) * rates[rec['Currency']]) for rec in rows).quantize(Decimal('0.01'))
assert twd_total == expected_total
ex = check_wb['Exceptions']
ex_count = 0
r = 5
while any(ex.cell(r,c).value is not None for c in range(1,5)):
    ex_count += 1
    r += 1
assert ex_count == 1
print('Validation OK')
print('detail_rows', len(actual_detail_rows))
print('summary_rows', len(summary_rows), summary_rows)
print('exceptions', ex_count)
print('TWD total', t企wd_total if False else twd_total)
print('output file', OUTFILE.name)
