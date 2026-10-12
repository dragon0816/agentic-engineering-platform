import os, csv, json
from pathlib import Path
import openpyxl
ws=Path.cwd()
print('CWD', ws)
for root, dirs, files in os.walk(ws/'input'):
    print('DIR', root)
    for f in files:
        p=Path(root)/f
        print('FILE', p.relative_to(ws), p.stat().st_size)
# template structure
p=ws/'input'/'expense_template.xlsx'
wb=openpyxl.load_workbook(p)
print('template sheets', wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('SHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
        print(' ', r)
# fx sample
p=ws/'input'/'fx_rates.csv'
print('fx_rates.csv')
with open(p, newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(10), rdr): print(row)
# manual excerpt
p=ws/'input'/'expense_manual.md'
print('manual excerpt')
print(p.read_text(encoding='utf-8')[:4000])