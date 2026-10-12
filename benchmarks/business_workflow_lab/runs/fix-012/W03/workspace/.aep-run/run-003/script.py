from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Border, Alignment
from openpyxl.cell.cell import MergedCell
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
import csv, shutil, json, os

workspace = Path.cwd()
input_dir = workspace / 'input'
out_dir = workspace / 'output'
out_dir.mkdir(exist_ok=True)

TEMPLATE = input_dir / 'project_template.xlsx'
FX = input_dir / 'fx_rates.csv'
CSV = out_dir / 'downloads' / 'invoices_2026-09.csv'
OUT = out_dir / 'project_report_2026-09_invoices.xlsx'

START = '2026-09-01'
END = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
FETCHED_AT = datetime.now().astimezone().replace(microsecond=0).isoformat()

# Helpers
TWOPLACES = Decimal('0.01')
def money(v):
    if v is None or v == '':
        return None
    return Decimal(str(v)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def to_float(d):
    if d is None:
        return None
    return float(d)

def col_letter(n):
    from openpyxl.utils import get_column_letter
    return get_column_letter(n)

def clear_data_rows(ws, start_row=5):
    # Delete existing data rows only; preserve title/header rows and sheet identity.
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def copy_row_style(ws, src_row, dst_row, max_col):
    for c in range(1, max_col + 1):
        src = ws.cell(src_row, c)
        dst = ws.cell(dst_row, c)
        if src.has_style:
            dst._style = src._style.copy()
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = src.alignment.copy()
        if src.fill:
            dst.fill = src.fill.copy()
        if src.font:
            dst.font = src.font.copy()
        if src.border:
            dst.border = src.border.copy()

# Read fx rates
rates = {}
with open(FX, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', rates)

# Read invoices
with open(CSV, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    invoices = list(reader)
print('Invoice rows read:', len(invoices))
print('Invoice headers:', reader.fieldnames)
print('First rows:', invoices[:3])

# Validate period all September 2026
for r in invoices:
    invd = date.fromisoformat(r['InvoiceDate'])
    assert date(2026,9,1) <= invd <= date(2026,9,30), (r['InvoiceID'], r['InvoiceDate'])
    assert r['Currency'] in rates, f"Missing FX rate {r['Currency']}"

# Load template
wb = load_workbook(TEMPLATE)

# Keep non-requested source sheets blank (only rows through 4)
for s in ['Orders', 'Expenses', 'Quotes']:
    clear_data_rows(wb[s], 5)

# Clear/fill requested sheets
clear_data_rows(wb['Invoices'], 5)
clear_data_rows(wb['Summary'], 5)
clear_data_rows(wb['Exceptions'], 5)

# Fill Report metadata without changing layout
rep = wb['Report']
rep['B3'] = START
rep['D3'] = END
rep['B4'] = 'Invoices: InvoiceDate 2026-09-01 to 2026-09-30; overdue calculated as Balance > 0 and DueDate < 2026-09-30 (due on base date not overdue).'
rep['B5'] = FETCHED_AT
rep['B7'] = 'Lab ERP Invoices; fx_rates.csv'
rep['D7'] = FETCHED_AT
rep['B9'] = 'Only Invoices and Summary populated per request; Orders, Expenses and Quotes source sheets left blank.'

# Fill Invoices
inv_ws = wb['Invoices']
inv_headers = [inv_ws.cell(4, c).value for c in range(1, inv_ws.max_column+1)]
expected_inv = ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
assert inv_headers == expected_inv, inv_headers

summary = {}
exceptions = []
for idx, r in enumerate(invoices, start=5):
    copy_row_style(inv_ws, 4, idx, len(expected_inv))
    amount = money(r['Amount'])
    paid = money(r['PaidAmount'])
    bal = money(r['Balance'])
    bal_twd = (bal * rates[r['Currency']]).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    row_vals = [
        r['InvoiceID'], r['OrderID'], r['ProjectID'] or None, r['CustomerID'],
        r['InvoiceDate'], r['DueDate'], r['Currency'],
        to_float(amount), to_float(paid), to_float(bal), to_float(bal_twd)
    ]
    for c, v in enumerate(row_vals, start=1):
        cell = inv_ws.cell(idx, c, v)
        if c in [1,2,3,4,7]:
            cell.number_format = '@'
        elif c in [5,6]:
            cell.number_format = 'yyyy-mm-dd'
        elif c in [8,9,10,11]:
            cell.number_format = '0.00'
    # summary by nonblank ProjectID
    project = r['ProjectID'].strip()
    if project:
        summary.setdefault(project, Decimal('0.00'))
        summary[project] += bal_twd
    due = date.fromisoformat(r['DueDate'])
    if bal > 0 and due < BASE_DATE:
        exceptions.append({'Source':'Invoices','DocumentID':r['InvoiceID'],'ProjectID': project or None,'Reason':'Overdue'})

# Fill Summary: one row per project found in invoice data, sorted by ProjectID
sum_ws = wb['Summary']
summary_headers = [sum_ws.cell(4, c).value for c in range(1, sum_ws.max_column+1)]
expected_sum = ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']
assert summary_headers == expected_sum, summary_headers
projects = sorted(summary.keys())
for i, project in enumerate(projects, start=5):
    copy_row_style(sum_ws, 4, i, len(expected_sum))
    vals = [project, 0.00, to_float(summary[project].quantize(TWOPLACES, rounding=ROUND_HALF_UP)), 0.00, 0.00]
    for c, v in enumerate(vals, start=1):
        cell = sum_ws.cell(i, c, v)
        if c == 1:
            cell.number_format = '@'
        else:
            cell.number_format = '0.00'
last = max(5, 4 + len(projects))
for c in range(2, 6):
    letter = col_letter(c)
    sum_ws.cell(2,c).value = f'=SUM({letter}5:{letter}{last})'

# Fill Exceptions
exc_ws = wb['Exceptions']
exc_headers = [exc_ws.cell(4, c).value for c in range(1, exc_ws.max_column+1)]
expected_exc = ['Source','DocumentID','ProjectID','Reason']
assert exc_headers == expected_exc, exc_headers
for i, e in enumerate(exceptions, start=5):
    copy_row_style(exc_ws, 4, i, len(expected_exc))
    vals = [e['Source'], e['DocumentID'], e['ProjectID'], e['Reason']]
    for c, v in enumerate(vals, start=1):
        cell = exc_ws.cell(i, c, v)
        if c in [1,2,3,4]:
            cell.number_format = '@'

# Ensure amount columns numeric; source sheets remain no rows after header
wb.save(OUT)
print('Wrote:', OUT)
print('Summary computed:', {k: str(v.quantize(TWOPLACES)) for k,v in summary.items()})
print('Exceptions:', exceptions)

# Verification: reopen saved output and confirm requirements
wb2 = load_workbook(OUT, data_only=False)
checks = []
# Sheets intact
checks.append(('sheets_intact', wb2.sheetnames == ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']))
# Other source worksheets blank below row 4
for s in ['Orders','Expenses','Quotes']:
    sh = wb2[s]
    nonempty = []
    for row in sh.iter_rows(min_row=5):
        if any(cell.value is not None for cell in row):
            nonempty.append(row[0].row)
    checks.append((f'{s}_blank_after_row4', not nonempty))
# Invoice rows count and headers
sh = wb2['Invoices']
rows = list(sh.iter_rows(min_row=5, max_row=4+len(invoices), values_only=False))
checks.append(('invoice_row_count_13', len(rows)==13 and sh.max_row==17))
checks.append(('invoice_headers_preserved', [sh.cell(4,c).value for c in range(1,12)] == expected_inv))
# ID as text-ish, leading zeros present; dates period; numeric amounts
for row in rows:
    vals = [cell.value for cell in row]
    assert isinstance(vals[0], str) and vals[0].startswith('009'), vals
    assert isinstance(vals[1], str) and vals[1].startswith('004'), vals
    assert vals[4] >= '2026-09-01' and vals[4] <= '2026-09-30', vals
    for idx in [7,8,9,10]:
        assert isinstance(vals[idx], (int,float)), (vals, idx, type(vals[idx]))
checks.append(('invoice_types_dates_values_ok', True))
# BalanceTWD check from source
for r, row in zip(invoices, rows):
    expected = float((money(r['Balance']) * rates[r['Currency']]).quantize(TWOPLACES, rounding=ROUND_HALF_UP))
    actual = row[10].value
    assert abs(actual - expected) < 0.0001, (r['InvoiceID'], actual, expected)
checks.append(('balance_twd_round_half_up_ok', True))
# Summary row checks
sum_sh = wb2['Summary']
sum_rows = list(sum_sh.iter_rows(min_row=5, max_row=4+len(projects), values_only=True))
expected_summary = {k: float(v.quantize(TWOPLACES, rounding=ROUND_HALF_UP)) for k,v in summary.items()}
actual_summary = {r[0]: r[2] for r in sum_rows}
checks.append(('summary_projects_match', set(actual_summary.keys()) == set(expected_summary.keys())))
for p,v in expected_summary.items():
    assert abs(actual_summary[p] - v) < 0.0001, (p, actual_summary[p], v)
for r in sum_rows:
    assert r[1] == 0 and r[3] == 0 and r[4] == 0, r
checks.append(('summary_invoice_balance_twd_ok_other_summary_cols_zero', True))
checks.append(('summary_formulas_extended', [sum_sh.cell(2,c).value for c in range(2,6)] == [f'=SUM(B5:B{last})', f'=SUM(C5:C{last})', f'=SUM(D5:D{last})', f'=SUM(E5:E{last})']))
# Exceptions exact Overdue: Balance>0 and DueDate < 2026-09-30, not equal
exc_sh = wb2['Exceptions']
exc_rows = list(exc_sh.iter_rows(min_row=5, max_row=4+len(exceptions), values_only=True))
actual_excs = [(r[0], r[1], r[2], r[3]) for r in exc_rows]
expected_excs = [(e['Source'], e['DocumentID'], e['ProjectID'], e['Reason']) for e in exceptions]
checks.append(('exceptions_match_overdue_rule', actual_excs == expected_excs))
# No exception for due date exactly 2026-09-30
bad_due30 = [r[0] for r in invoices if r['DueDate']=='2026-09-30' and Decimal(r['Balance'])>0 and r['InvoiceID'] in [e['DocumentID'] for e in exceptions]]
checks.append(('due_on_base_date_not_overdue', bad_due30 == []))

for name, ok in checks:
    print(f'CHECK {name}:', 'PASS' if ok else 'FAIL')
    assert ok, name
print('All verification checks passed.')
print('Output file:', OUT)
