import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('CWD', ws)
input_dir = ws/'input'
print('INPUT EXISTS', input_dir.exists())
print('FILES')
for p in sorted(input_dir.rglob('*')):
    print('-', p.relative_to(ws), 'size', p.stat().st_size)

# Inspect workbook structure
xlsx = input_dir/'project_template.xlsx'
if xlsx.exists():
    wb = load_workbook(xlsx)
    print('\nWORKBOOK SHEETS:', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print(f'\nSHEET {s}: max_row={sh.max_row} max_col={sh.max_column}')
        for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
            print(r)

# Inspect fx_rates.csv
csvp = input_dir/'fx_rates.csv'
if csvp.exists():
    print('\nFX RATES SAMPLE')
    with open(csvp, newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i,row in zip(range(10), reader):
            print(row)

# Inspect sap_manual.md
mdp = input_dir/'sap_manual.md'
if mdp.exists():
    print('\nSAP MANUAL HEAD')
    txt = mdp.read_text(encoding='utf-8-sig', errors='replace')
    print(txt[:4000])
