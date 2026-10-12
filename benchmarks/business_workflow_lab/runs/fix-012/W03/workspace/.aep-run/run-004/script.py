from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from copy import copy
import csv

workspace = Path.cwd()
input_dir = workspace / 'input'
out_dir = workspace / 'output'
out_dir.mkdir(exist_ok=True)

TEMPLATE = input_dir / 'project_template.xlsx'
FX = input_dir / 'fx_rates.csv'
CSV = out_dir / 'downloads' / 'invoices_2026-09.csv'
OUT = out_dir / 'project_report_2026-09_invoices.xlsx'

START = '2026-09-01'
END = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
FETCHED_AT = datetime.now().astimezone().replace(microsecond=0).isoformat()
TWOPLACES = Decimal('0.01')

def money(v):
    if v is None or v == '':
        return None
    return Decimal(str(v)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def to_float(d):
    if d is None:
        return None
    return float(d)

def clear_data_rows(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def copy_row_style(ws, src_row, dst_row, max_col):
    for c in range(1, max_col + 1):
        src = ws.cell(src_row, c)
        dst = ws.cell(dst_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.fill = copy(src.fill)
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.protection = copy(src.protection)

# Read/inspect sources
print('Input files:', sorted(str(p.relative_to(workspace)) for p in input_dir.iterdir()))
print('Downloaded invoice csv exists:', CSV.exists(), CSV.stat().st_size if CSV.exists() else None)

rates = {}
with open(FX, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header:', header)
    for row in reader:
        if not row or not row[0]:
            continue
        rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', {k: str(v) for k,v in rates.items()})

with open(CSV, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    invoices = list(reader)
    print('Invoice headers:', reader.fieldnames)
print('Invoice rows read:', len(invoices))
print('Sample rows:', invoices[:2])

for r in invoices:
    invd = date.fromisoformat(r['InvoiceDate'])
    assert date(2026,9,1) <= invd <= date(2026,9,30), (r['InvoiceID'], r['InvoiceDate'])
    assert r['Currency'] in rates, f"Missing FX rate {r['Currency']}"

wb = load_workbook(TEMPLATE)
print('Template sheets:', wb.sheetnames)

for s in ['Orders', 'Expenses', 'Quotes']:
    clear_data_rows(wb[s], 5)
clear_data_rows(wb['Invoices'], 5)
clear_data_rows(wb['Summary'], 5)
clear_data_rows(wb['Exceptions'], 5)

rep = wb['Report']
rep['B3'] = START
rep['D3'] = END
rep['B4'] = 'Invoices: InvoiceDate 2026-09-01 to 2026-09-30; overdue calculated as Balance > 0 and DueDate < 2026-09-30 (due on base date not overdue).'
rep['B5'] = FETCHED_AT
rep['B7'] = 'Lab ERP Invoices; fx_rates.csv'
rep['D7'] = FETCHED_AT
rep['B9'] = 'Only Invoices and Summary populated per request; Orders, Expenses and Quotes source sheets left blank.'

inv_ws = wb['Invoices']
expected_inv = ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
assert [inv_ws.cell(4, c).value for c in range(1,12)] == expected_inv

summary = {}
exceptions = []
for idx, r in enumerate(invoices, start=5):
    copy_row_style(inv_ws, 4, idx, len(expected_inv))
    amount = money(r['Amount'])
    paid = money(r['PaidAmount'])
    bal = money(r['Balance'])
    bal_twd = (bal * rates[r['Currency']]).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
    values = [r['InvoiceID'], r['OrderID'], r['ProjectID'] or None, r['CustomerID'], r['InvoiceDate'], r['DueDate'], r['Currency'], to_float(amount), to_float(paid), to_float(bal), to_float(bal_twd)]
    for c, v in enumerate(values, start=1):
        cell = inv_ws.cell(idx, c, v)
        if c in [1,2,3,4,7]: cell.number_format = '@'
        elif c in [5,6]: cell.number_format = 'yyyy-mm-dd'
        else: cell.number_format = '0.00'
    project = r['ProjectID'].strip()
    if project:
        summary[project] = summary.get(project, Decimal('0.00')) + bal_twd
    due = date.fromisoformat(r['DueDate'])
    if bal > 0 and due < BASE_DATE:
        exceptions.append({'Source':'Invoices','DocumentID':r['InvoiceID'],'ProjectID':project or None,'Reason':'Overdue'})

sum_ws = wb['Summary']
expected_sum = ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']
assert [sum_ws.cell(4,c).value for c in range(1,6)] == expected_sum
projects = sorted(summary.keys())
for i, project in enumerate(projects, start=5):
    copy_row_style(sum_ws, 4, i, len(expected_sum))
    vals = [project, 0.00, to_float(summary[project].quantize(TWOPLACES, rounding=ROUND_HALF_UP)), 0.00, 0.00]
    for c, v in enumerate(vals, start=1):
        cell = sum_ws.cell(i, c, v)
        cell.number_format = '@' if c == 1 else '0.00'
last = max(5, 4 + len(projects))
for c, letter in enumerate(['B','C','D','E'], start=2):
    sum_ws.cell(2,c).value = f'=SUM({letter}5:{letter}{last})'

exc_ws = wb['Exceptions']
expected_exc = ['Source','DocumentID','ProjectID','Reason']
assert [exc_ws.cell(4,c).value for c in range(1,5)] == expected_exc
for i, e in enumerate(exceptions, start=5):
    copy_row_style(exc_ws, 4, i, len(expected_exc))
    for c, v in enumerate([e['Source'], e['DocumentID'], e['ProjectID'], e['Reason']], start=1):
        cell = exc_ws.cell(i, c, v)
        cell.number_format = '@'

wb.save(OUT)
print('Wrote:', OUT)
print('Summary computed:', {k: str(v.quantize(TWOPLACES)) for k,v in summary.items()})
print('Exceptions computed:', exceptions)

# Verification by reopening the output
wb2 = load_workbook(OUT, data_only=False)
checks = []
checks.append(('output_file_exists_nonempty', OUT.exists() and OUT.stat().st_size > 0))
checks.append(('sheets_intact', wb2.sheetnames == ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']))

for s in ['Orders','Expenses','Quotes']:
    sh = wb2[s]
    nonempty = [cell.coordinate for row in sh.iter_rows(min_row=5) for cell in row if cell.value is not None]
    checks.append((f'{s}_blank_after_row4', nonempty == []))

sh = wb2['Invoices']
rows = list(sh.iter_rows(min_row=5, max_row=4+len(invoices), values_only=False))
checks.append(('invoice_row_count_13', len(rows)==13 and sh.max_row==17))
checks.append(('invoice_headers_preserved', [sh.cell(4,c).value for c in range(1,12)] == expected_inv))
for row in rows:
    vals = [cell.value for cell in row]
    assert isinstance(vals[0], str) and vals[0].startswith('009'), vals
    assert isinstance(vals[1], str) and vals[1].startswith('004'), vals
    if vals[2] is not None:
        assert isinstance(vals[2], str), vals
    assert isinstance(vals[3], str), vals
    assert '2026-09-01' <= vals[4] <= '2026-09-30', vals
    for idx in [7,8,9,10]:
        assert isinstance(vals[idx], (int,float)), (vals, idx, type(vals[idx]))
checks.append(('invoice_id_text_and_amount_numeric_ok', True))

for r, row in zip(invoices, rows):
    expected = float((money(r['Balance']) * rates[r['Currency']]).quantize(TWOPLACES, rounding=ROUND_HALF_UP))
    actual = row[10].value
    assert abs(actual - expected) < 0.0001, (r['InvoiceID'], actual, expected)
checks.append(('balance_twd_calculation_ok', True))

sum_sh = wb2['Summary']
sum_rows = list(sum_sh.iter_rows(min_row=5, max_row=4+len(projects), values_only=True))
expected_summary = {k: float(v.quantize(TWOPLACES, rounding=ROUND_HALF_UP)) for k,v in summary.items()}
actual_summary = {r[0]: r[2] for r in sum_rows}
checks.append(('summary_project_count', len(sum_rows)==len(projects)))
checks.append(('summary_projects_match', set(actual_summary.keys()) == set(expected_summary.keys())))
for p,v in expected_summary.items():
    assert abs(actual_summary[p] - v) < 0.0001, (p, actual_summary[p], v)
for r in sum_rows:
    assert r[1] == 0 and r[3] == 0 and r[4] == 0, r
checks.append(('summary_values_ok', True))
checks.append(('summary_formulas_extended', [sum_sh.cell(2,c).value for c in range(2,6)] == [f'=SUM(B5:B{last})', f'=SUM(C5:C{last})', f'=SUM(D5:D{last})', f'=SUM(E5:E{last})']))

exc_sh = wb2['Exceptions']
exc_rows = list(exc_sh.iter_rows(min_row=5, max_row=4+len(exceptions), values_only=True))
actual_excs = [(r[0], r[1], r[2], r[3]) for r in exc_rows]
expected_excs = [(e['Source'], e['DocumentID'], e['ProjectID'], e['Reason']) for e in exceptions]
checks.append(('exceptions_match_overdue_rule', actual_excs == expected_excs))
bad_due30 = [r['InvoiceID'] for r in invoices if r['DueDate']=='2026-09-30' and Decimal(r['Balance'])>0 and r['InvoiceID'] in [e['DocumentID'] for e in exceptions]]
checks.append(('due_on_base_date_not_overdue', bad_due30 == []))

for name, ok in checks:
    print(f'CHECK {name}:', 'PASS' if ok else 'FAIL')
    assert ok, name
print('All verification checks passed.')
print('Output file:', OUT)
