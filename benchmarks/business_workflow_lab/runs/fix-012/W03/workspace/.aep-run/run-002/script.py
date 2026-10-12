from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('Downloads:')
for p in (ws/'output'/'downloads').glob('*'):
    print(p, p.stat().st_size)
    with open(p, newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(20), r): print(row)
print('\nInstructions all:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True): print(row)
