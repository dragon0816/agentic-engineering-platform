from pathlib import Path
from openpyxl import load_workbook
import csv, json
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size)
# inspect workbook
path = ws/'input'/'sales_template.xlsx'
wb = load_workbook(path, data_only=False)
print('workbook sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('\nSHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
        print(r)
# inspect fx csv
fx = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample')
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(10), reader):
        print(row)
# inspect manual first chunk
man = ws/'input'/'sap_manual.md'
print('\nsap_manual.md first 4000 chars')
print(man.read_text(encoding='utf-8')[:4000])