from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True):
    print(r)