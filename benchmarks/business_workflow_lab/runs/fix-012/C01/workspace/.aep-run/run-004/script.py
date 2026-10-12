import os, json, csv
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import openpyxl

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
print('Workspace:', workspace)
print('Input files:', sorted(p.name for p in input_dir.glob('*')))
print('Output files:', sorted(p.name for p in output_dir.glob('*')))

out_file = output_dir / 'quote_report_C001_2026-09-01_2026-09-30.xlsx'
if not out_file.exists():
    raise FileNotFoundError(out_file)

wb = openpyxl.load_workbook(out_file, data_only=False)
print('Workbook sheets:', wb.sheetnames)
required_sheets = ['Report', 'Quotes', 'Lines', 'Summary']
missing = [s for s in required_sheets if s not in wb.sheetnames]
if missing:
    raise AssertionError(f'Missing sheets: {missing}')

for sheet in required_sheets:
    ws = wb[sheet]
    print(f'{sheet}: max_row={ws.max_row}, max_column={ws.max_column}')

# Read headers and rows. Workflow wrote headers on row 4 and data starting row 5.
quotes = wb['Quotes']
lines = wb['Lines']
summary = wb['Summary']
report = wb['Report']

quote_headers = [quotes.cell(4, c).value for c in range(1, quotes.max_column + 1)]
line_headers = [lines.cell(4, c).value for c in range(1, lines.max_column + 1)]
summary_headers = [summary.cell(4, c).value for c in range(1, summary.max_column + 1)]
print('Quotes headers:', quote_headers)
print('Lines headers:', line_headers)
print('Summary headers:', summary_headers)

quote_rows = []
for r in range(5, quotes.max_row + 1):
    vals = [quotes.cell(r, c).value for c in range(1, len(quote_headers) + 1)]
    if any(v is not None for v in vals):
        quote_rows.append(dict(zip(quote_headers, vals)))
line_rows = []
for r in range(5, lines.max_row + 1):
    vals = [lines.cell(r, c).value for c in range(1, len(line_headers) + 1)]
    if any(v is not None for v in vals):
        line_rows.append(dict(zip(line_headers, vals)))
summary_rows = []
for r in range(5, summary.max_row + 1):
    vals = [summary.cell(r, c).value for c in range(1, len(summary_headers) + 1)]
    if any(v is not None for v in vals):
        summary_rows.append(dict(zip(summary_headers, vals)))

print('Quote sample rows:', json.dumps(quote_rows[:2], ensure_ascii=False, default=str))
print('Line sample rows:', json.dumps(line_rows[:3], ensure_ascii=False, default=str))
print('Summary rows:', json.dumps(summary_rows, ensure_ascii=False, default=str))

# Requirement checks
expected_quote_ids = {'Q0010','Q0011','Q0014','Q0018'}
actual_quote_ids = {str(r.get('QuoteID')) for r in quote_rows}
if actual_quote_ids != expected_quote_ids:
    raise AssertionError(f'Quote IDs mismatch: {actual_quote_ids} != {expected_quote_ids}')
if len(quote_rows) != 4:
    raise AssertionError(f'Expected 4 quote rows, got {len(quote_rows)}')
if len(line_rows) != 11:
    raise AssertionError(f'Expected 11 line rows, got {len(line_rows)}')
if len(summary_rows) != 4:
    raise AssertionError(f'Expected 4 summary rows, got {len(summary_rows)}')

# Validate filters in Quote rows
for row in quote_rows:
    if str(row.get('CustomerID')) != 'C001':
        raise AssertionError(f'Wrong CustomerID in row: {row}')
    qdate = str(row.get('QuoteDate'))[:10]
    if not ('2026-09-01' <= qdate <= '2026-09-30'):
        raise AssertionError(f'QuoteDate out of range: {qdate}')
    # NetAmountTWD should be numeric
    if not isinstance(row.get('NetAmountTWD'), (int, float)):
        raise AssertionError(f'NetAmountTWD not numeric: {row}')

# Check line foreign keys
bad_lines = [r for r in line_rows if str(r.get('QuoteID')) not in expected_quote_ids]
if bad_lines:
    raise AssertionError(f'Lines for unexpected quote ids: {bad_lines}')

# Check summary grouping equals quote rows by Status + Currency
from collections import defaultdict
calc = defaultdict(lambda: {'Count':0, 'NetAmount':0.0, 'AmountTWD':0.0})
for row in quote_rows:
    key = (row.get('Status'), row.get('Currency'))
    calc[key]['Count'] += 1
    calc[key]['NetAmount'] += float(row.get('NetAmount') or 0)
    calc[key]['AmountTWD'] += float(row.get('NetAmountTWD') or 0)

summary_keys = set()
for row in summary_rows:
    key = (row.get('Status'), row.get('Currency'))
    summary_keys.add(key)
    if key not in calc:
        raise AssertionError(f'Unexpected summary key: {key}')
    if int(row.get('Count')) != calc[key]['Count']:
        raise AssertionError(f'Summary count mismatch for {key}: {row.get("Count")} vs {calc[key]["Count"]}')
    if abs(float(row.get('NetAmount')) - calc[key]['NetAmount']) > 0.01:
        raise AssertionError(f'Summary NetAmount mismatch for {key}')
    if abs(float(row.get('AmountTWD')) - calc[key]['AmountTWD']) > 0.05:
        raise AssertionError(f'Summary AmountTWD mismatch for {key}')
if summary_keys != set(calc.keys()):
    raise AssertionError(f'Summary keys mismatch: {summary_keys} vs {set(calc.keys())}')

# Report sheet metadata check: collect non-empty cells for visibility
report_nonempty = []
for row in report.iter_rows():
    vals = [cell.value for cell in row if cell.value not in (None, '')]
    if vals:
        report_nonempty.append(vals)
print('Report non-empty rows:', json.dumps(report_nonempty[:20], ensure_ascii=False, default=str))

print('CHECKS PASSED:', json.dumps({
    'file': str(out_file),
    'bytes': out_file.stat().st_size,
    'sheets_checked': required_sheets,
    'quote_rows': len(quote_rows),
    'line_rows': len(line_rows),
    'summary_rows': len(summary_rows),
    'quote_ids': sorted(actual_quote_ids),
}, ensure_ascii=False))
