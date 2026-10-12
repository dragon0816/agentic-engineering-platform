import csv
import json
import os
import re
import subprocess
import sys
import time
from copy import copy
from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)
DEFAULT_CLI = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
TWD = Decimal('0.01')

params_path = os.environ.get('AEP_WORKFLOW_PARAMS')
params = json.load(open(params_path, encoding='utf-8')) if params_path and Path(params_path).exists() else {}
customer_id = str(params.get('customer_id') or '').strip()
from_date = str(params.get('from_date') or '').strip()
to_date = str(params.get('to_date') or '').strip()
cli_path = str(params.get('cli_path') or DEFAULT_CLI).strip() or DEFAULT_CLI

if not customer_id:
    raise SystemExit('customer_id is required')
if not from_date or not to_date:
    raise SystemExit('from_date and to_date are required')
for label, value in [('from_date', from_date), ('to_date', to_date)]:
    try:
        datetime.strptime(value, '%Y-%m-%d')
    except ValueError as exc:
        raise SystemExit(f'{label} must be YYYY-MM-DD: {value}') from exc
if from_date > to_date:
    raise SystemExit(f'from_date must be <= to_date: {from_date} > {to_date}')

cli = Path(cli_path)
if not cli.exists():
    raise SystemExit(f'CLI path not found: {cli}')

def q2(x):
    return Decimal(str(x)).quantize(TWD, rounding=ROUND_HALF_UP)

def excel_number(x):
    if x is None or x == '':
        return None
    return float(q2(x))

def command_prefix():
    if cli.suffix.lower() == '.py':
        return [sys.executable, str(cli)]
    return [str(cli)]

def run_cli(args):
    cmd = command_prefix() + args
    last = None
    for attempt in range(3):
        cp = subprocess.run(cmd, capture_output=True, text=True)
        last = cp
        if cp.returncode == 0:
            try:
                return json.loads(cp.stdout)
            except json.JSONDecodeError as exc:
                raise RuntimeError(f'CLI returned non-JSON stdout for {args}: {cp.stdout[:500]}') from exc
        if cp.returncode == 4 and attempt < 2:
            time.sleep(1 + attempt)
            continue
        raise RuntimeError(f'CLI failed rc={cp.returncode} for {args}; stderr={cp.stderr.strip()}')
    raise RuntimeError(f'CLI failed after retries for {args}; rc={last.returncode}')

print('Workspace:', WORKSPACE)
print('Template:', INPUT / 'quote_template.xlsx')
print('FX rates:', INPUT / 'fx_rates.csv')
print('CLI:', cli)
print('Filters:', {'customer_id': customer_id, 'from_date': from_date, 'to_date': to_date})

fx_rates = {}
with (INPUT / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader, None)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            fx_rates[row[0].strip()] = Decimal(row[1].strip())
print('FX currencies:', sorted(fx_rates))

list_data = run_cli(['quotes', 'list', '--customer', customer_id, '--from', from_date, '--to', to_date, '--format', 'json'])
if list_data.get('schema_version') != 1:
    raise RuntimeError(f'Unsupported list schema_version: {list_data.get("schema_version")}')
quote_headers = list_data.get('items', [])
print('Quote count from list:', len(quote_headers), 'reported:', list_data.get('count'))

quotes = []
lines = []
for item in quote_headers:
    qid = item.get('QuoteID')
    detail = run_cli(['quotes', 'get', qid, '--format', 'json'])
    if detail.get('schema_version') != 1:
        raise RuntimeError(f'Unsupported get schema_version for {qid}: {detail.get("schema_version")}')
    h = detail.get('header') or {}
    if h.get('QuoteID') != qid:
        raise RuntimeError(f'QuoteID mismatch: list {qid}, get {h.get("QuoteID")}')
    currency = h.get('Currency')
    if currency not in fx_rates:
        raise RuntimeError(f'Missing FX rate for currency {currency} in quote {qid}')
    net = q2(h.get('NetAmount'))
    net_twd = q2(net * fx_rates[currency])
    h2 = dict(h)
    h2['NetAmountTWD'] = net_twd
    quotes.append(h2)
    for line in detail.get('line_items') or []:
        l2 = dict(line)
        l2['QuoteID'] = qid
        lines.append(l2)

template = INPUT / 'quote_template.xlsx'
wb = load_workbook(template)

def clear_rows(sh, min_row=5):
    for row in sh.iter_rows(min_row=min_row, max_row=sh.max_row):
        for cell in row:
            cell.value = None

for name in ['Quotes', 'Lines', 'Summary']:
    clear_rows(wb[name], 5)

tz = timezone(timedelta(hours=8))
now = datetime.now(tz).replace(microsecond=0).isoformat()
rep = wb['Report']
rep['B3'] = from_date
rep['D3'] = to_date
rep['B4'] = f'CustomerID={customer_id}; QuoteDate {from_date}..{to_date}; status=ALL; project=ALL'
rep['B5'] = now
rep['B7'] = 'quote_cli mock; fx_rates.csv'
rep['D7'] = now
rep['B9'] = '' if quotes else 'No quote records returned by CLI for this filter.'

def copy_row_style(sh, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        src = sh.cell(src_row, col)
        dst = sh.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

qsh = wb['Quotes']
quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for idx, q in enumerate(quotes, start=5):
    copy_row_style(qsh, 4, idx, len(quote_cols))
    values = [q.get('QuoteID'), q.get('CustomerID'), q.get('ProjectID'), q.get('QuoteDate'), q.get('ValidUntil'), q.get('Status'), q.get('Currency'), excel_number(q.get('Subtotal')), float(Decimal(str(q.get('DiscountPercent')))), excel_number(q.get('NetAmount')), excel_number(q.get('NetAmountTWD'))]
    for col, value in enumerate(values, start=1):
        cell = qsh.cell(idx, col, value)
        if col in [1,2,3,6,7]:
            cell.number_format = '@'
        elif col in [4,5]:
            cell.number_format = 'yyyy-mm-dd'
        elif col in [8,10,11]:
            cell.number_format = '0.00'

lsh = wb['Lines']
line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for idx, l in enumerate(lines, start=5):
    copy_row_style(lsh, 4, idx, len(line_cols))
    values = [l.get('QuoteID'), l.get('ProductID'), l.get('ProductName'), l.get('Quantity'), excel_number(l.get('UnitPrice')), excel_number(l.get('LineAmount'))]
    for col, value in enumerate(values, start=1):
        cell = lsh.cell(idx, col, value)
        if col in [1,2,3]:
            cell.number_format = '@'
        elif col in [5,6]:
            cell.number_format = '0.00'

summary = {}
for q in quotes:
    key = (q.get('Status') or '', q.get('Currency') or '')
    if key not in summary:
        summary[key] = {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0}
    summary[key]['amount'] += q2(q.get('NetAmount'))
    summary[key]['amount_twd'] += q2(q.get('NetAmountTWD'))
    summary[key]['count'] += 1

ssh = wb['Summary']
last_quote_row = max(5, 4 + len(quotes))
ssh['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
status_order = {'Draft': 0, 'Sent': 1, 'Accepted': 2, 'Expired': 3}
for idx, ((status, currency), agg) in enumerate(sorted(summary.items(), key=lambda kv: (status_order.get(kv[0][0], 99), kv[0][0], kv[0][1])), start=5):
    copy_row_style(ssh, 4, idx, 6)
    row_values = ['Status', status, currency, float(q2(agg['amount'])), float(q2(agg['amount_twd'])), int(agg['count'])]
    for col, value in enumerate(row_values, start=1):
        cell = ssh.cell(idx, col, value)
        if col in [4,5]:
            cell.number_format = '0.00'

safe = re.sub(r'[^A-Za-z0-9_.-]+', '_', f'{customer_id}_{from_date}_{to_date}')
outfile = OUTPUT / f'quote_report_{safe}.xlsx'
wb.save(outfile)
print('Wrote:', outfile)

check_wb = load_workbook(outfile, data_only=False)
required_sheets = ['Report','Quotes','Lines','Summary','Instructions']
assert check_wb.sheetnames == required_sheets, f'sheet names changed: {check_wb.sheetnames}'
q2sh = check_wb['Quotes']
l2sh = check_wb['Lines']
s2sh = check_wb['Summary']
q_rows = [[q2sh.cell(r,c).value for c in range(1,12)] for r in range(5, 5+len(quotes))]
l_rows = [[l2sh.cell(r,c).value for c in range(1,7)] for r in range(5, 5+len(lines))]
s_rows = [[s2sh.cell(r,c).value for c in range(1,7)] for r in range(5, 5+len(summary))]
assert len(q_rows) == len(quotes)
assert len(l_rows) == len(lines)
assert all(isinstance(row[0], str) and isinstance(row[1], str) for row in q_rows), 'Quote text IDs check failed'
assert all(isinstance(row[7], (int,float)) and isinstance(row[9], (int,float)) and isinstance(row[10], (int,float)) for row in q_rows), 'Quote numeric amount check failed'
assert all(isinstance(row[0], str) and isinstance(row[1], str) for row in l_rows), 'Line text IDs check failed'
assert all(isinstance(row[3], (int,float)) and isinstance(row[4], (int,float)) and isinstance(row[5], (int,float)) for row in l_rows), 'Line numeric check failed'
assert s2sh['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})', f'Summary B2 formula wrong: {s2sh["B2"].value}'
assert len(s_rows) == len(summary), 'Summary row count mismatch'
quote_twd_total = sum(Decimal(str(row[10])) for row in q_rows)
summary_twd_total = sum(Decimal(str(row[4])) for row in s_rows)
assert q2(quote_twd_total) == q2(summary_twd_total), f'TWD totals mismatch {quote_twd_total} vs {summary_twd_total}'
print('Checks passed:', json.dumps({'output_file': str(outfile), 'quote_rows': len(q_rows), 'line_rows': len(l_rows), 'summary_rows': len(s_rows), 'summary_b2_formula': s2sh['B2'].value, 'quote_twd_total': str(q2(quote_twd_total))}, ensure_ascii=False))
print('OUTPUT_FILES:', json.dumps([str(outfile)], ensure_ascii=False))
