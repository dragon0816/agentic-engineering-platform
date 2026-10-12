import json
from pathlib import Path
import openpyxl

workspace = Path.cwd()
out = workspace / 'output' / 'quote_report_C001_2026-09-01_2026-09-30.xlsx'
print('Workspace:', workspace)
print('Checking output:', out)
if not out.exists():
    raise FileNotFoundError(str(out))

wb = openpyxl.load_workbook(out, data_only=False)
required = ['Report','Quotes','Lines','Summary','Instructions']
print('Sheets:', wb.sheetnames)
if wb.sheetnames != required:
    raise AssertionError(f'Expected sheets {required}, got {wb.sheetnames}')

wsq = wb['Quotes']
wsl = wb['Lines']
wss = wb['Summary']
wsr = wb['Report']

def nonempty_rows(ws, header_row=4):
    rows=[]
    for r in range(header_row+1, ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        if any(v is not None for v in vals):
            rows.append(vals)
    return rows

quote_headers = [wsq.cell(4,c).value for c in range(1, wsq.max_column+1)]
line_headers = [wsl.cell(4,c).value for c in range(1, wsl.max_column+1)]
summary_headers = [wss.cell(4,c).value for c in range(1, wss.max_column+1)]
print('Quote headers:', quote_headers)
print('Line headers:', line_headers)
print('Summary headers:', summary_headers)

quotes = nonempty_rows(wsq)
lines = nonempty_rows(wsl)
summary = nonempty_rows(wss)
print('Quote rows:', len(quotes))
print('Line rows:', len(lines))
print('Summary rows:', len(summary))

if len(quotes) != 4:
    raise AssertionError(f'Expected 4 quote rows, got {len(quotes)}')
if len(lines) != 11:
    raise AssertionError(f'Expected 11 line rows, got {len(lines)}')
if len(summary) != 4:
    raise AssertionError(f'Expected 4 summary rows, got {len(summary)}')

expected_quote_ids = {'Q0010','Q0011','Q0014','Q0018'}
quote_id_col = quote_headers.index('QuoteID')
quote_ids = {str(r[quote_id_col]) for r in quotes}
if quote_ids != expected_quote_ids:
    raise AssertionError(f'Unexpected QuoteIDs: {quote_ids}')

line_quote_col = line_headers.index('QuoteID')
line_quote_ids = {str(r[line_quote_col]) for r in lines}
if not line_quote_ids <= expected_quote_ids:
    raise AssertionError(f'Line QuoteIDs not subset: {line_quote_ids}')

customer_col = quote_headers.index('CustomerID')
date_col = quote_headers.index('QuoteDate')
status_col = quote_headers.index('Status')
currency_col = quote_headers.index('Currency')
net_col = quote_headers.index('NetAmount')
twd_col = quote_headers.index('NetAmountTWD')
for r in quotes:
    if r[customer_col] != 'C001':
        raise AssertionError('Quote CustomerID mismatch')
    d = r[date_col]
    ds = d.isoformat() if hasattr(d, 'isoformat') else str(d)
    if not ('2026-09-01' <= ds[:10] <= '2026-09-30'):
        raise AssertionError(f'Quote date outside range: {ds}')
    if not isinstance(r[net_col], (int, float)) or not isinstance(r[twd_col], (int, float)):
        raise AssertionError('Amount columns must be numeric')

calc = {}
for r in quotes:
    key=(r[status_col], r[currency_col])
    calc.setdefault(key, [0.0,0.0,0])
    calc[key][0] += float(r[net_col])
    calc[key][1] += float(r[twd_col])
    calc[key][2] += 1

summary_group_col = summary_headers.index('Group')
summary_status_col = summary_headers.index('SubGroup')
summary_currency_col = summary_headers.index('Currency')
summary_amount_col = summary_headers.index('Amount')
summary_twd_col = summary_headers.index('AmountTWD')
summary_count_col = summary_headers.index('Count')
seen = {}
for r in summary:
    if r[summary_group_col] != 'Status':
        raise AssertionError(f'Summary Group should be Status, got {r[summary_group_col]}')
    key=(r[summary_status_col], r[summary_currency_col])
    seen[key] = [float(r[summary_amount_col]), float(r[summary_twd_col]), int(r[summary_count_col])]
if set(seen) != set(calc):
    raise AssertionError(f'Summary keys mismatch: seen={set(seen)} calc={set(calc)}')
for key, vals in calc.items():
    got = seen[key]
    if round(got[0],2) != round(vals[0],2) or round(got[1],2) != round(vals[1],2) or got[2] != vals[2]:
        raise AssertionError(f'Summary mismatch for {key}: got={got}, expected={vals}')

report_values = [wsr.cell(r,c).value for r in range(1, min(wsr.max_row,20)+1) for c in range(1, min(wsr.max_column,5)+1)]
report_text = '\n'.join('' if v is None else str(v) for v in report_values)
for token in ['C001','2026-09-01','2026-09-30']:
    if token not in report_text:
        raise AssertionError(f'Report missing {token}')

result = {
    'output_file': str(out),
    'bytes': out.stat().st_size,
    'quote_rows': len(quotes),
    'line_rows': len(lines),
    'summary_rows': len(summary),
    'quote_ids': sorted(quote_ids),
    'summary_keys': sorted([f'{k[0]} / {k[1]}' for k in seen]),
}
print('CHECKS_PASSED:', json.dumps(result, ensure_ascii=False, indent=2))