from pathlib import Path
from openpyxl import load_workbook
import subprocess, json, os
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx')
sh=wb['Instructions']
print('--- Instructions all rows ---')
for row in sh.iter_rows(values_only=True):
    print(row)
cli=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
print('\nCLI exists', cli, cli.exists())
if cli.exists():
    cmds=[ [str(cli),'quotes','list','--customer','C001','--from','2026-09-01','--to','2026-09-30','--format','json'] ]
    for cmd in cmds:
        print('\nRUN', cmd)
        cp=subprocess.run(cmd, capture_output=True, text=True, shell=False)
        print('returncode', cp.returncode)
        print('stdout', cp.stdout[:3000])
        print('stderr', cp.stderr[:1000])
        if cp.returncode==0:
            data=json.loads(cp.stdout)
            if data.get('items'):
                qid=data['items'][0]['QuoteID']
                cmd2=[str(cli),'quotes','get',qid,'--format','json']
                print('\nRUN', cmd2)
                cp2=subprocess.run(cmd2, capture_output=True, text=True, shell=False)
                print('returncode', cp2.returncode)
                print('stdout', cp2.stdout[:3000])
                print('stderr', cp2.stderr[:1000])
