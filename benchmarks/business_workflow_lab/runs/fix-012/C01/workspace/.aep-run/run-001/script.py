from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('CWD', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\n--- quote_cli_manual.md ---')
    print(manual.read_text(encoding='utf-8', errors='replace')[:4000])
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\n--- fx_rates.csv sample ---')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(10), r):
            print(row)
tpl = ws/'input'/'quote_template.xlsx'
if tpl.exists():
    print('\n--- workbook sheets/cells ---')
    wb=load_workbook(tpl)
    print(wb.sheetnames)
    for s in wb.sheetnames:
        sh=wb[s]
        print('\nSheet', s, 'max', sh.max_row, sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
            print(row)
