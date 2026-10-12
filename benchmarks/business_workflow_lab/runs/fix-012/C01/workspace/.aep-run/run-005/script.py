import os, json
from pathlib import Path
from collections import defaultdict
import openpyxl

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
print('Workspace:', workspace)
print('Input files:', sorted(p.name for p in input_dir.glob('*')))
print('Output files:', sorted(p.name for p in output_dir.glob('*')))

out_file = output_dir / 'quote_report_C001_2026-09-01_2026-09-30.xlsx'
if not out_file.exists():
    raise FileNotFoundError(out_file)
if out_file.stat().st_size <= 0:
    raise AssertionError('Output workbook is empty')

wb = openpyxl.load_workbook(out_file, data_only=False)
print('Workbook sheets:', wb.sheetnames)
required_sheets = ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
missing = [s for s in required_sheets if s not in wb.sheetnames]
if missing:
    raise AssertionError(f'Missing sheets: {missing}')

quotes = wb['Quotes']
lines = wb['Lines']
summary = wb['Summary']
report = wb['Report']

# Template headers are on row 4; data begins row 5.
def read_table(ws, header_row=4, data_start=5):
    headers = [ws.cell(header_row, c).value for c in range(1, ws.max_column + 1)]
    while headers and headers[-1] is None:
        headers.pop()
    rows = []
    for r in range(data_start, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, len(headers) + 1)]
        if any(v is not None for v in vals):
            rows.append(dict(zip(headers, vals)))
    return headers, rows

quote_headers, quote_rows = read_table(quotes)
line_headers, line_rows = read_table(lines)
summary_headers, summary_rows = read_table(summary)
print('Quotes headers:', quote_headers)
print('Lines headers:', line_headers)
print('Summary headers:', summary_headers)
print('Quote rows:', json.dumps(quote_rows, ensure_ascii=False, default=str))
print('Line row count:', len(line_rows))
print('Summary rows:', json.dumps(summary_rows, ensure_ascii=False, default=str))

expected_quote_headers = ['QuoteID', 'CustomerID', 'ProjectID', 'QuoteDate', 'ValidUntil', 'Status', 'Currency', 'Subtotal', 'DiscountPercent', 'NetAmount', 'NetAmountTWD']
expected_line_headers = ['QuoteID', 'ProductID', 'ProductName', 'Quantity', 'UnitPrice', 'LineAmount']
expected_summary_headers = ['Group', 'SubGroup', 'Currency', 'Amount', 'AmountTWD', 'Count']
if quote_headers != expected_quote_headers:
    raise AssertionError(f'Quotes headers changed: {quote_headers}')
if line_headers != expected_line_headers:
    raise AssertionError(f'Lines headers changed: {line_headers}')
if summary_headers != expected_summary_headers:
    raise AssertionError(f'Summary headers changed: {summary_headers}')

expected_quote_ids = {'Q0010','Q0011','Q0014','Q0018'}
actual_quote_ids = {str(r['QuoteID']) for r in quote_rows}
if actual_quote_ids != expected_quote_ids:
    raise AssertionError(f'Quote IDs mismatch: {actual_quote_ids} != {expected_quote_ids}')
if len(quote_rows) != 4:
    raise AssertionError(f'Expected 4 quote rows, got {len(quote_rows)}')
if len(line_rows) != 11:
    raise AssertionError(f'Expected 11 line rows, got {len(line_rows)}')
if len(summary_rows) != 4:
    raise AssertionError(f'Expected 4 summary rows, got {len(summary_rows)}')

for row in quote_rows:
    if row['CustomerID'] != 'C001':
        raise AssertionError(f'Wrong CustomerID: {row}')
    qdate = str(row['QuoteDate'])[:10]
    if not ('2026-09-01' <= qdate <= '2026-09-30'):
        raise AssertionError(f'QuoteDate out of range: {qdate}')
    for numeric_col in ['Subtotal', 'DiscountPercent', 'NetAmount', 'NetAmountTWD']:
        if not isinstance(row[numeric_col], (int, float)):
            raise AssertionError(f'{numeric_col} not numeric in row: {row}')

bad_lines = [r for r in line_rows if str(r['QuoteID']) not in expected_quote_ids]
if bad_lines:
    raise AssertionError(f'Lines for unexpected quote IDs: {bad_lines}')
for row in line_rows:
    for numeric_col in ['Quantity', 'UnitPrice', 'LineAmount']:
        if not isinstance(row[numeric_col], (int, float)):
            raise AssertionError(f'{numeric_col} not numeric in line: {row}')

# Summary rows represent Status+Currency as Group='Status', SubGroup=<Status>, Currency=<Currency>.
calc = defaultdict(lambda: {'Count':0, 'Amount':0.0, 'AmountTWD':0.0})
for row in quote_rows:
    key = (row['Status'], row['Currency'])
    calc[key]['Count'] += 1
    calc[key]['Amount'] += float(row['NetAmount'])
    calc[key]['AmountTWD'] += float(row['NetAmountTWD'])

seen = set()
for row in summary_rows:
    if row['Group'] != 'Status':
        raise AssertionError(f'Summary Group is not Status: {row}')
    key = (row['SubGroup'], row['Currency'])
    seen.add(key)
    if key not in calc:
        raise AssertionError(f'Unexpected summary key: {key}')
    if int(row['Count']) != calc[key]['Count']:
        raise AssertionError(f'Summary count mismatch for {key}: {row["Count"]} vs {calc[key]["Count"]}')
    if abs(float(row['Amount']) - calc[key]['Amount']) > 0.01:
        raise AssertionError(f'Summary Amount mismatch for {key}: {row["Amount"]} vs {calc[key]["Amount"]}')
    if abs(float(row['AmountTWD']) - calc[key]['AmountTWD']) > 0.05:
        raise AssertionError(f'Summary AmountTWD mismatch for {key}: {row["AmountTWD"]} vs {calc[key]["AmountTWD"]}')
if seen != set(calc.keys()):
    raise AssertionError(f'Summary keys mismatch: {seen} vs {set(calc.keys())}')

# Check Report has the requested filters visible somewhere.
report_values = []
for row in report.iter_rows():
    vals = [cell.value for cell in row if cell.value not in (None, '')]
    if vals:
        report_values.append(vals)
report_text = json.dumps(report_values, ensure_ascii=False, default=str)
print('Report non-empty rows:', report_text)
for required in ['C001', '2026-09-01', '2026-09-30']:
    if required not in report_text:
        raise AssertionError(f'Report sheet does not show {required}')

print('FINAL CHECKS PASSED:', json.dumps({
    'output_file': str(out_file),
    'bytes': out_file.stat().st_size,
    'sheets_checked': required_sheets,
    'quote_rows': len(quote_rows),
    'line_rows': len(line_rows),
    'summary_rows': len(summary_rows),
    'quote_ids': sorted(actual_quote_ids),
    'summary_keys': sorted([f'{k[0]}|{k[1]}' for k in seen]),
}, ensure_ascii=False))
