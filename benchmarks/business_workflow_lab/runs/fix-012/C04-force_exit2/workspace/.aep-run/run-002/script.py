from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
import subprocess, sys, json, csv, shutil, os, traceback
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'sent_quotes_2026-09.xlsx'
RUN_LOG = OUTPUT / 'quote_cli_run_log.json'
ERROR_TXT = OUTPUT / 'quote_cli_error.txt'

START_DATE = '2026-09-01'
END_DATE = '2026-09-30'
STATUS = 'Sent'
TZ = timezone(timedelta(hours=8))

# Clean outputs from previous reruns so validation reflects this run.
for p in [OUT_XLSX, RUN_LOG, ERROR_TXT]:
    if p.exists():
        p.unlink()

run_events = []

def run_cli(args, purpose):
    """Run quote CLI with retry contract. Return parsed JSON data or raise RuntimeError."""
    attempts = 0
    max_attempts = 3  # first try + up to two retries for exit code 4
    last = None
    while attempts < max_attempts:
        attempts += 1
        cmd = [sys.executable, str(CLI)] + args
        cp = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
        event = {
            'purpose': purpose,
            'attempt': attempts,
            'command': cmd,
            'exit_code': cp.returncode,
            'stdout_length': len(cp.stdout or ''),
            'stderr': cp.stderr,
        }
        run_events.append(event)
        last = cp
        if cp.returncode == 4:
            if attempts < max_attempts:
                continue
            raise RuntimeError(f'{purpose}: temporary service error after {attempts} attempts. stderr: {cp.stderr.strip()}')
        if cp.returncode in (2, 3):
            raise RuntimeError(f'{purpose}: quote CLI exited {cp.returncode}; no retry per instruction. stderr: {cp.stderr.strip()}')
        if cp.returncode != 0:
            raise RuntimeError(f'{purpose}: quote CLI exited {cp.returncode}. stderr: {cp.stderr.strip()}')
        # stdout is data only on success, but still must be parseable.
        try:
            return json.loads(cp.stdout)
        except Exception as e:
            raise RuntimeError(f'{purpose}: exit 0 but stdout was not parseable JSON; stdout ignored. parse error: {e}; stderr: {cp.stderr.strip()}')
    raise RuntimeError(f'{purpose}: unexpected retry loop end. last exit: {last.returncode if last else None}')

def load_fx_rates():
    rates = {}
    with FX_FILE.open('r', encoding='utf-8-sig', newline='') as f:
        reader = csv.reader(f)
        header = next(reader)
        for row in reader:
            if len(row) >= 2 and row[0]:
                rates[row[0]] = Decimal(row[1])
    return rates

def dec(v):
    return Decimal(str(v))

def money(v):
    return dec(v).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def write_failure(message):
    RUN_LOG.write_text(json.dumps({'status':'failed','message':message,'events':run_events}, ensure_ascii=False, indent=2), encoding='utf-8')
    ERROR_TXT.write_text(message + '\n\nCLI events:\n' + json.dumps(run_events, ensure_ascii=False, indent=2), encoding='utf-8')
    print('FAILED')
    print(message)
    print('WROTE', ERROR_TXT)
    print('WROTE', RUN_LOG)

def fill_template(list_data, detail_docs):
    rates = load_fx_rates()
    shutil.copy2(TEMPLATE, OUT_XLSX)
    wb = load_workbook(OUT_XLSX)
    now = datetime.now(TZ).isoformat(timespec='seconds')

    # Report sheet
    ws = wb['Report']
    ws['B3'] = START_DATE
    ws['D3'] = END_DATE
    ws['B4'] = f'Status={STATUS}; QuoteDate from {START_DATE} to {END_DATE}'
    ws['B5'] = now
    ws['B7'] = 'quote_cli; fx_rates.csv'
    ws['D7'] = now
    ws['B9'] = ''

    # Quotes sheet
    qws = wb['Quotes']
    headers = [qws.cell(4, c).value for c in range(1, qws.max_column + 1)]
    items = list_data.get('items', [])
    # deterministic order by QuoteID
    items = sorted(items, key=lambda x: str(x.get('QuoteID','')))
    for i, item in enumerate(items, start=5):
        currency = item.get('Currency')
        if currency not in rates:
            raise RuntimeError(f'Missing FX rate for currency {currency} in quote {item.get("QuoteID")}')
        net = money(item.get('NetAmount'))
        row_values = {
            'QuoteID': str(item.get('QuoteID','')),
            'CustomerID': str(item.get('CustomerID','')),
            'ProjectID': str(item.get('ProjectID','')),
            'QuoteDate': item.get('QuoteDate'),
            'ValidUntil': item.get('ValidUntil'),
            'Status': item.get('Status'),
            'Currency': currency,
            'Subtotal': money(item.get('Subtotal')),
            'DiscountPercent': dec(item.get('DiscountPercent')),
            'NetAmount': net,
            'NetAmountTWD': money(net * rates[currency]),
        }
        for c, h in enumerate(headers, start=1):
            v = row_values.get(h)
            cell = qws.cell(i, c, v)
            if h in ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency']:
                cell.number_format = '@'
            elif h in ['Subtotal','DiscountPercent','NetAmount','NetAmountTWD']:
                cell.number_format = '0.00'

    # Lines sheet
    lws = wb['Lines']
    line_headers = [lws.cell(4, c).value for c in range(1, lws.max_column + 1)]
    line_row = 5
    for qid in [item['QuoteID'] for item in items]:
        doc = detail_docs[qid]
        for line in doc.get('line_items', []):
            row_values = {
                'QuoteID': str(qid),
                'ProductID': str(line.get('ProductID','')),
                'ProductName': line.get('ProductName'),
                'Quantity': dec(line.get('Quantity')),
                'UnitPrice': money(line.get('UnitPrice')),
                'LineAmount': money(line.get('LineAmount')),
            }
            for c, h in enumerate(line_headers, start=1):
                cell = lws.cell(line_row, c, row_values.get(h))
                if h in ['QuoteID','ProductID','ProductName']:
                    cell.number_format = '@'
                elif h in ['Quantity']:
                    cell.number_format = '0'
                elif h in ['UnitPrice','LineAmount']:
                    cell.number_format = '0.00'
            line_row += 1

    # Summary sheet: by currency
    sws = wb['Summary']
    # update total formula range if rows present, otherwise to an empty-safe range
    last_q_row = 4 + len(items)
    sws['B2'] = f'=SUM(Quotes!K5:K{max(last_q_row,5)})'
    summary = {}
    for item in items:
        cur = item['Currency']
        net = money(item['NetAmount'])
        summary.setdefault(cur, {'amount': Decimal('0.00'), 'count': 0})
        summary[cur]['amount'] += net
        summary[cur]['count'] += 1
    r = 5
    for cur in sorted(summary):
        amount = money(summary[cur]['amount'])
        sws.cell(r,1,'Status')
        sws.cell(r,2,STATUS)
        sws.cell(r,3,cur)
        sws.cell(r,4,amount).number_format = '0.00'
        sws.cell(r,5,money(amount * rates[cur])).number_format = '0.00'
        sws.cell(r,6,summary[cur]['count'])
        r += 1

    wb.save(OUT_XLSX)
    RUN_LOG.write_text(json.dumps({'status':'succeeded','quote_count':len(items),'line_count':line_row-5,'events':run_events}, ensure_ascii=False, indent=2), encoding='utf-8')
    print('WROTE', OUT_XLSX)
    print('WROTE', RUN_LOG)

try:
    list_data = run_cli(['quotes','list','--from',START_DATE,'--to',END_DATE,'--status',STATUS,'--format','json'], 'list Sent quotes for 2026-09')
    if list_data.get('schema_version') != 1 or 'items' not in list_data or 'count' not in list_data:
        raise RuntimeError('list result JSON does not match expected schema_version/items/count; stdout ignored')
    details = {}
    for item in list_data['items']:
        qid = str(item.get('QuoteID',''))
        if not qid:
            raise RuntimeError('list result has item without QuoteID')
        doc = run_cli(['quotes','get',qid,'--format','json'], f'get quote {qid}')
        if doc.get('schema_version') != 1 or 'header' not in doc or 'line_items' not in doc:
            raise RuntimeError(f'quote {qid} detail JSON does not match expected schema')
        details[qid] = doc
    fill_template(list_data, details)
except Exception as e:
    write_failure(str(e))
