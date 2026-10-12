from pathlib import Path
import csv, json, zipfile
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
for p in sorted((ws/'input').glob('*')):
    print('INPUT_FILE', p.name, 'size', p.stat().st_size)

manual = ws/'input'/'quote_cli_manual.md'
print('\n--- quote_cli_manual.md ---')
print(manual.read_text(encoding='utf-8', errors='replace')[:4000])

fx = ws/'input'/'fx_rates.csv'
print('\n--- fx_rates.csv sample ---')
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i, row in zip(range(10), reader):
        print(i, row)

tpl = ws/'input'/'quote_template.xlsx'
print('\n--- workbook sheets/cells ---')
wb = load_workbook(tpl, data_only=False)
for sh in wb.sheetnames:
    s = wb[sh]
    print('SHEET', sh, 'max_row', s.max_row, 'max_col', s.max_column)
    for r in range(1, min(s.max_row, 12)+1):
        vals = [s.cell(r,c).value for c in range(1, min(s.max_column, 12)+1)]
        print(r, vals)
    print('---')
