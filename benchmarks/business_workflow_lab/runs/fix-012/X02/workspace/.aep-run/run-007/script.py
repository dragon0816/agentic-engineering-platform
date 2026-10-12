import csv, os, json, subprocess, sys, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from openpyxl import load_workbook

params=json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
customer_id=params['customer_id']
from_date=params['from_date']
to_date=params['to_date']
as_of_date=params.get('as_of_date') or to_date
WORKSPACE=Path.cwd(); INPUT=WORKSPACE/'input'; OUTPUT=WORKSPACE/'output'; DOWNLOADS=OUTPUT/'downloads'; OUTPUT.mkdir(exist_ok=True); DOWNLOADS.mkdir(exist_ok=True)
report_path=OUTPUT/f'project_report_{customer_id}_{from_date}_{to_date}.xlsx'
quotes_json_path=OUTPUT/f'quotes_{customer_id}_{from_date}_{to_date}.json'
orders_path=DOWNLOADS/'sap_orders_export.csv'
invoices_path=DOWNLOADS/'sap_invoices_export.csv'
expenses_path=OUTPUT/'expenses_all_approved_export.csv'

print('Workflow parameters:', params)
print('Input files:', sorted(p.name for p in INPUT.iterdir()))
for p in [orders_path, invoices_path, expenses_path]: print('Source', p, 'exists=', p.exists(), 'bytes=', p.stat().st_size if p.exists() else None)

def dec(v): return Decimal('0') if v is None or v=='' else Decimal(str(v))
def money(v): return dec(v).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
def d(s): return date.fromisoformat(str(s)) if s else None
def parse_bool(v): return str(v).strip().lower() in ('true','1','yes','y')
def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f: return list(csv.DictReader(f))
def amount_twd(amount, currency, fx):
    if currency not in fx: raise ValueError(f'Missing FX rate for {currency}')
    return (money(amount)*fx[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
def clear_data(ws):
    if ws.max_row>=5: ws.delete_rows(5, ws.max_row-4)
def write_text(ws,r,c,v):
    cell=ws.cell(r,c,value=v); cell.number_format='@'; cell.data_type='s'
def write_val(ws,r,c,v): ws.cell(r,c,value=v)
def set_money(ws,r,c,v):
    cell=ws.cell(r,c,value=float(v)); cell.number_format='0.00'
def set_date(ws,r,c,v):
    cell=ws.cell(r,c,value=d(v)); cell.number_format='yyyy-mm-dd'

fx={}
with open(INPUT/'fx_rates.csv', encoding='utf-8-sig', newline='') as f:
    reader=csv.reader(f); next(reader)
    for row in reader:
        if len(row)>=2 and row[0]: fx[row[0].strip()]=dec(row[1].strip())
print('FX rates loaded from attachment only:', fx)

orders_raw=read_csv(orders_path); invoices=read_csv(invoices_path); expenses_all=read_csv(expenses_path)
orders=[r for r in orders_raw if r.get('Status')!='Cancelled']
cli=r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py'
cmd=[sys.executable, cli, 'quotes', 'list', '--customer', customer_id, '--from', from_date, '--to', to_date, '--format', 'json']
proc=subprocess.run(cmd, text=True, capture_output=True)
print('quote cli returncode:', proc.returncode)
if proc.returncode!=0:
    print(proc.stderr); raise SystemExit(proc.returncode)
quotes_data=json.loads(proc.stdout); quotes=quotes_data['items']; quotes_json_path.write_text(json.dumps(quotes_data, ensure_ascii=False, indent=2), encoding='utf-8')
projects=sorted({r.get('ProjectID','') for r in orders+invoices+quotes if r.get('ProjectID','')})
expenses=[r for r in expenses_all if r.get('ProjectID') in projects and r.get('Status')=='Approved']
print('Customer related projects:', projects)
print('Counts:', {'orders_raw':len(orders_raw),'orders_after_excluding_cancelled':len(orders),'invoices':len(invoices),'expenses_all_approved':len(expenses_all),'expenses_after_project_filter':len(expenses),'quotes':len(quotes)})

shutil.copyfile(INPUT/'project_template.xlsx', report_path)
wb=load_workbook(report_path)
for sheet in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']: clear_data(wb[sheet])
rpt=wb['Report']; rpt['B3']=d(from_date); rpt['B3'].number_format='yyyy-mm-dd'; rpt['D3']=d(to_date); rpt['D3'].number_format='yyyy-mm-dd'
gen=datetime.now().astimezone().replace(microsecond=0).isoformat()
rpt['B4']=f'CustomerID={customer_id}; period={from_date}..{to_date}; orders exclude Cancelled; expenses Approved only for customer-related projects; quotes kept all; FX from input/fx_rates.csv'
rpt['B5']=gen; rpt['B7']='Lab ERP Sales Orders/Invoices; Expense Viewer; quote_cli; fx_rates.csv'; rpt['D7']=gen
rpt['B9']='Customer-related projects: '+', '.join(projects)+'; attention rules: invoice due <= as-of date with unpaid balance, approved expense missing receipt, quote ValidUntil < as-of date and Status != Accepted.'

ws=wb['Orders']
for i,row in enumerate(orders,5):
    write_text(ws,i,1,row['OrderID']); write_text(ws,i,2,row['ProjectID']); write_text(ws,i,3,row['CustomerID']); set_date(ws,i,4,row['OrderDate']); write_val(ws,i,5,row['Status']); write_val(ws,i,6,row['Currency']); set_money(ws,i,7,money(row['NetAmount'])); set_money(ws,i,8,amount_twd(row['NetAmount'],row['Currency'],fx))
ws=wb['Invoices']
for i,row in enumerate(invoices,5):
    write_text(ws,i,1,row['InvoiceID']); write_text(ws,i,2,row['OrderID']); write_text(ws,i,3,row['ProjectID']); write_text(ws,i,4,row['CustomerID']); set_date(ws,i,5,row['InvoiceDate']); set_date(ws,i,6,row['DueDate']); write_val(ws,i,7,row['Currency']); set_money(ws,i,8,money(row['Amount'])); set_money(ws,i,9,money(row['PaidAmount'])); set_money(ws,i,10,money(row['Balance'])); set_money(ws,i,11,amount_twd(row['Balance'],row['Currency'],fx))
ws=wb['Expenses']
for i,row in enumerate(expenses,5):
    write_text(ws,i,1,row['ExpenseID']); write_text(ws,i,2,row['ProjectID']); write_text(ws,i,3,row['EmployeeID']); set_date(ws,i,4,row['ExpenseDate']); write_val(ws,i,5,row['Status']); write_val(ws,i,6,row['Currency']); set_money(ws,i,7,money(row['Amount'])); set_money(ws,i,8,amount_twd(row['Amount'],row['Currency'],fx))
ws=wb['Quotes']
for i,row in enumerate(quotes,5):
    write_text(ws,i,1,row['QuoteID']); write_text(ws,i,2,row['ProjectID']); write_text(ws,i,3,row['CustomerID']); set_date(ws,i,4,row['QuoteDate']); set_date(ws,i,5,row['ValidUntil']); write_val(ws,i,6,row['Status']); write_val(ws,i,7,row['Currency']); set_money(ws,i,8,money(row['NetAmount'])); set_money(ws,i,9,amount_twd(row['NetAmount'],row['Currency'],fx))
summary={p:{'OrderTWD':Decimal('0.00'),'InvoiceBalanceTWD':Decimal('0.00'),'ApprovedExpenseTWD':Decimal('0.00'),'AcceptedQuoteTWD':Decimal('0.00')} for p in projects}
for row in orders: summary[row['ProjectID']]['OrderTWD']+=amount_twd(row['NetAmount'],row['Currency'],fx)
for row in invoices: summary[row['ProjectID']]['InvoiceBalanceTWD']+=amount_twd(row['Balance'],row['Currency'],fx)
for row in expenses: summary[row['ProjectID']]['ApprovedExpenseTWD']+=amount_twd(row['Amount'],row['Currency'],fx)
for row in quotes:
    if row['Status']=='Accepted': summary[row['ProjectID']]['AcceptedQuoteTWD']+=amount_twd(row['NetAmount'],row['Currency'],fx)
ws=wb['Summary']
for i,p in enumerate(sorted(summary),5):
    write_text(ws,i,1,p); set_money(ws,i,2,summary[p]['OrderTWD']); set_money(ws,i,3,summary[p]['InvoiceBalanceTWD']); set_money(ws,i,4,summary[p]['ApprovedExpenseTWD']); set_money(ws,i,5,summary[p]['AcceptedQuoteTWD'])
last=max(5,4+len(summary))
for c,l in enumerate(['B','C','D','E'],2): ws.cell(2,c,value=f'=SUM({l}5:{l}{last})').number_format='0.00'
exceptions=[]; asof=d(as_of_date)
for row in invoices:
    if d(row['DueDate']) <= asof and money(row['Balance'])>0: exceptions.append(('Invoices',row['InvoiceID'],row['ProjectID'],f'Overdue unpaid invoice as of {as_of_date}; Balance={row["Balance"]} {row["Currency"]}'))
for row in expenses:
    if parse_bool(row.get('MissingReceipt','')): exceptions.append(('Expenses',row['ExpenseID'],row['ProjectID'],'Approved expense missing receipt'))
for row in quotes:
    if d(row['ValidUntil']) < asof and row['Status']!='Accepted': exceptions.append(('Quotes',row['QuoteID'],row['ProjectID'],f'Quote expired before {as_of_date} and not Accepted; Status={row["Status"]}'))
ws=wb['Exceptions']
for i,(src,doc,proj,reason) in enumerate(exceptions,5): write_val(ws,i,1,src); write_text(ws,i,2,doc); write_text(ws,i,3,proj); write_val(ws,i,4,reason)
wb.save(report_path)

wb2=load_workbook(report_path, data_only=False)
def rows(sheet):
    sh=wb2[sheet]
    return list(sh.iter_rows(min_row=5, max_row=sh.max_row, values_only=True)) if sh.max_row>=5 else []
orders_o=rows('Orders'); invoices_o=rows('Invoices'); expenses_o=rows('Expenses'); quotes_o=rows('Quotes'); summary_o=rows('Summary'); exceptions_o=rows('Exceptions')
checks={
 'report_nonempty': report_path.exists() and report_path.stat().st_size>0,
 'orders_exclude_cancelled': all(r[4]!='Cancelled' for r in orders_o),
 'expenses_approved_only': all(r[4]=='Approved' for r in expenses_o),
 'expenses_limited_to_customer_projects': all(r[1] in projects for r in expenses_o),
 'quotes_all_kept': len(quotes_o)==len(quotes),
 'ids_text_preserved': (not orders_o or (wb2['Orders']['A5'].data_type=='s')),
 'amounts_numeric': all(isinstance(r[7],(int,float)) for r in orders_o) and all(isinstance(r[10],(int,float)) for r in invoices_o) and all(isinstance(r[7],(int,float)) for r in expenses_o) and all(isinstance(r[8],(int,float)) for r in quotes_o),
 'summary_formula_updated': wb2['Summary']['B2'].value==f'=SUM(B5:B{max(5,4+len(summary_o))})'
}
print('Validation checks:')
for k,v in checks.items(): print(k, v)
if not all(checks.values()): raise SystemExit('Validation failed')
print('OUTPUT_FILES')
print(report_path)
print(quotes_json_path)
print('Final row counts:', {'Orders':len(orders_o),'Invoices':len(invoices_o),'Expenses':len(expenses_o),'Quotes':len(quotes_o),'Summary':len(summary_o),'Exceptions':len(exceptions_o)})
print('Exceptions:', exceptions)
