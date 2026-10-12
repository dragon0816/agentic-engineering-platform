import json
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

ws = Path.cwd()
out = ws/'output'
print('workspace=', ws)
print('output files=')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(ws), 'bytes=', p.stat().st_size, 'mtime=', datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec='seconds'))

report = out/'project_report_C001_2026-09-01_2026-09-30.xlsx'
quotes_json = out/'quotes_C001_2026-09-01_2026-09-30.json'
orders_export = out/'downloads'/'sap_orders_export.csv'
invoices_export = out/'downloads'/'sap_invoices_export.csv'
expenses_export = out/'expenses_all_approved_export.csv'

errors=[]
def check(name, ok, detail):
    print(('PASS' if ok else 'FAIL'), name, detail)
    if not ok:
        errors.append(f'{name}: {detail}')

for label, path in [('report', report), ('quotes_json', quotes_json), ('orders_export', orders_export), ('invoices_export', invoices_export), ('expenses_export', expenses_export)]:
    check(f'{label}_exists_nonempty', path.exists() and path.stat().st_size > 0, f'{path} bytes={path.stat().st_size if path.exists() else "missing"}')

wb = load_workbook(report, data_only=False)
print('workbook_sheets=', wb.sheetnames)
check('template_sheets_present', all(s in wb.sheetnames for s in ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']), wb.sheetnames)

def rows(sheet_name, first_header):
    sh=wb[sheet_name]
    for r in range(1, min(sh.max_row, 25)+1):
        vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
        if vals and vals[0]==first_header:
            last=max(i for i,v in enumerate(vals,1) if v is not None)
            headers=vals[:last]
            data=[]
            for rr in range(r+1, sh.max_row+1):
                rowvals=[sh.cell(rr,c).value for c in range(1,last+1)]
                if any(v is not None for v in rowvals):
                    data.append(dict(zip(headers,rowvals)))
            return headers,data
    return [],[]

headers = {}
data = {}
for s,h in [('Orders','OrderID'),('Invoices','InvoiceID'),('Expenses','ExpenseID'),('Quotes','QuoteID'),('Summary','ProjectID'),('Exceptions','Source')]:
    headers[s], data[s] = rows(s,h)
    print(f'{s}: headers={headers[s]} rows={len(data[s])} sample={data[s][:2]}')

check('orders_count_3', len(data['Orders'])==3, len(data['Orders']))
check('invoices_count_3', len(data['Invoices'])==3, len(data['Invoices']))
check('expenses_count_9', len(data['Expenses'])==9, len(data['Expenses']))
check('quotes_count_4_all_retained', len(data['Quotes'])==4, len(data['Quotes']))
check('exceptions_count_4', len(data['Exceptions'])==4, len(data['Exceptions']))
check('orders_exclude_cancelled', all(r.get('Status')!='Cancelled' for r in data['Orders']), sorted(set(str(r.get('Status')) for r in data['Orders'])))
check('expenses_approved_only', all(r.get('Status')=='Approved' for r in data['Expenses']), sorted(set(str(r.get('Status')) for r in data['Expenses'])))
check('expenses_only_c001_projects', sorted(set(str(r.get('ProjectID')) for r in data['Expenses']))==['P001','P003'], sorted(set(str(r.get('ProjectID')) for r in data['Expenses'])))
check('customer_scope_orders_invoices_quotes', all(str(r.get('CustomerID'))=='C001' for s in ['Orders','Invoices','Quotes'] for r in data[s]), {s: sorted(set(str(r.get('CustomerID')) for r in data[s])) for s in ['Orders','Invoices','Quotes']})
check('summary_projects_p001_p003', sorted(str(r.get('ProjectID')) for r in data['Summary'])==['P001','P003'], data['Summary'])

id_ok=True; id_sample=[]
for s,col in [('Orders','OrderID'),('Invoices','InvoiceID'),('Expenses','ExpenseID'),('Quotes','QuoteID'),('Exceptions','DocumentID')]:
    v=data[s][0].get(col) if data[s] else None
    id_sample.append((s,col,v,type(v).__name__))
    id_ok = id_ok and isinstance(v,str)
check('id_fields_text_preserve_leading_zero', id_ok, id_sample)

twd_ok=True; twd_sample=[]
for s in ['Orders','Invoices','Expenses','Quotes','Summary']:
    for r in data[s]:
        for k,v in r.items():
            if k and 'TWD' in str(k).upper():
                twd_sample.append((s,k,v,type(v).__name__))
                twd_ok = twd_ok and (v is None or isinstance(v,(int,float)) or (isinstance(v,str) and v.startswith('=')))
check('twd_values_numeric', twd_ok, twd_sample[:12])

exc_text='\n'.join(json.dumps(r,ensure_ascii=False,default=str) for r in data['Exceptions'])
check('expected_attention_documents', all(x in exc_text for x in ['0090000219','EXP-00006','Q0010','Q0014']), exc_text)
check('attention_invoice_overdue_unpaid', '0090000219' in exc_text and 'Overdue unpaid invoice as of 2026-09-30' in exc_text and '386445.00 TWD' in exc_text, exc_text)
check('attention_missing_receipt', 'EXP-00006' in exc_text and 'missing receipt' in exc_text, exc_text)
check('attention_expired_quotes_nonaccepted', all(x in exc_text for x in ['Q0010','Q0014','not Accepted']), exc_text)

with open(quotes_json, encoding='utf-8') as f:
    qdoc=json.load(f)
qitems = qdoc.get('items') if isinstance(qdoc,dict) else qdoc
print('quotes_json schema=', qdoc.get('schema_version') if isinstance(qdoc,dict) else 'list', 'count_field=', qdoc.get('count') if isinstance(qdoc,dict) else None)
print('quotes_json_items_sample=', qitems[:2] if isinstance(qitems,list) else qitems)
check('quotes_json_has_four_items', isinstance(qitems,list) and len(qitems)==4, len(qitems) if isinstance(qitems,list) else type(qitems).__name__)
check('quotes_json_customer_c001', isinstance(qitems,list) and all(str(r.get('CustomerID'))=='C001' for r in qitems), sorted(set(str(r.get('CustomerID')) for r in qitems)) if isinstance(qitems,list) else 'not list')

if errors:
    print('ERRORS=', json.dumps(errors,ensure_ascii=False,indent=2))
    raise SystemExit(1)
print('All verification checks passed.')