from pathlib import Path
from openpyxl import load_workbook
wb = load_workbook('input/project_template.xlsx', data_only=False)
sh = wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print('\t'.join('' if v is None else str(v) for v in row))
