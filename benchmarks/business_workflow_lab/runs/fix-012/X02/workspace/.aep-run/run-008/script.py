import os, json, csv, glob
from pathlib import Path
from datetime import datetime, date
from openpyxl import load_workbook

ws = Path.cwd()
input_dir = ws/'input'
out = ws/'output'
print('workspace=', ws)
print('input files=', sorted(p.name for p in input_dir.glob('*')))
print('output files=')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(ws), 'bytes=', p.stat().st_size, 'mtime=', datetime.fromtimestamp(p.stat().st_mtime).isoformat(timespec='seconds'))

report = out/'project_report_C001_2026-09.xlsx'
quotes_json = out/'quotes_C001_2026-09.json'
checks = []
errors = []

def add_check(name, ok, detail):
    checks.append({'check': name, 'ok': bool(ok), 'detail': detail})
    if not ok:
        errors.append(f'{name}: {detail}')

add_check('report_exists_nonempty', report.exists() and report.stat().st_size > 0, str(report))
add_check('quotes_json_exists_nonempty', quotes_json.exists() and quotes_json.stat().st_size > 0, str(quotes_json))

if report.exists():
    wb = load_workbook(report, data_only=False)
    print('workbook_sheets=', wb.sheetnames)
    expected_sheets = ['Summary','Orders','Invoices','Expenses','Quotes','Exceptions']
    add_check('expected_sheets_present', all(s in wb.sheetnames for s in expected_sheets), wb.sheetnames)

    # locate headers and rows
    def rows_with_header(sheet_name):
        sh = wb[sheet_name]
        header_row = None
        headers = []
        for r in range(1, min(sh.max_row, 20)+1):
            vals = [sh.cell(r,c).value for c in range(1, sh.max_column+1)]
            if any(v is not None for v in vals):
                # use first row with multiple non-empty cells as header
                if sum(v is not None for v in vals) >= 2:
                    header_row = r
                    headers = vals
                    break
        if header_row is None:
            return [], []
        # Trim after last header
        last = max(i for i,v in enumerate(headers, start=1) if v is not None)
        headers = headers[:last]
        data = []
        for r in range(header_row+1, sh.max_row+1):
            vals = [sh.cell(r,c).value for c in range(1, last+1)]
            if any(v is not None for v in vals):
                data.append(dict(zip(headers, vals)))
        return headers, data

    counts = {}
    for s in expected_sheets:
        h, data = rows_with_header(s)
        counts[s] = len(data)
        print(f'{s}: headers={h} rows={len(data)} sample={data[:2]}')

    add_check('orders_count', counts.get('Orders') == 3, counts.get('Orders'))
    add_check('invoices_count', counts.get('Invoices') == 3, counts.get('Invoices'))
    add_check('expenses_count', counts.get('Expenses') == 9, counts.get('Expenses'))
    add_check('quotes_count', counts.get('Quotes') == 4, counts.get('Quotes'))
    add_check('exceptions_count', counts.get('Exceptions') == 4, counts.get('Exceptions'))

    h_orders, orders = rows_with_header('Orders')
    h_inv, invoices = rows_with_header('Invoices')
    h_exp, expenses = rows_with_header('Expenses')
    h_quotes, quotes = rows_with_header('Quotes')
    h_exc, exceptions = rows_with_header('Exceptions')

    # flexible key lookup
    def key(row, *names):
        lower_map = {str(k).strip().lower(): k for k in row.keys()}
        for n in names:
            k = lower_map.get(n.lower())
            if k is not None:
                return row.get(k)
        return None

    add_check('orders_exclude_cancelled', all(str(key(r,'Status') or '') != 'Cancelled' for r in orders), [key(r,'OrderID'), key(r,'Status')] if orders else 'no rows')
    add_check('expenses_approved_only', all(str(key(r,'Status') or '') == 'Approved' for r in expenses), sorted(set(str(key(r,'Status')) for r in expenses)))
    add_check('expenses_project_scope', all(str(key(r,'ProjectID','Project')) in {'P001','P003'} for r in expenses), sorted(set(str(key(r,'ProjectID','Project')) for r in expenses)))

    # Check IDs as text in workbook cells for key sheets: find first id col in header
    id_text_ok = True
    id_details = []
    for sheet_name, id_header in [('Orders','OrderID'),('Invoices','InvoiceID'),('Expenses','ExpenseID'),('Quotes','QuoteID')]:
        sh = wb[sheet_name]
        # find header cell
        found = None
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,20)):
            for cell in row:
                if cell.value == id_header:
                    found = cell
                    break
            if found: break
        if found:
            col = found.column
            for r in range(found.row+1, sh.max_row+1):
                v = sh.cell(r,col).value
                if v is not None:
                    if not isinstance(v, str):
                        id_text_ok = False
                    id_details.append((sheet_name, id_header, v, type(v).__name__))
                    break
    add_check('id_fields_are_text', id_text_ok, id_details)

    # Check exception IDs/reasons roughly
    exc_ids = sorted(str(key(r,'DocumentID','DocID','ID') or '') for r in exceptions)
    exc_text = '\n'.join(str(r) for r in exceptions)
    add_check('expected_exception_items', all(x in exc_text for x in ['0090000219','EXP-00006','Q0010','Q0014']), exc_text)

    # Amount TWD numeric checks: any header containing TWD
    twd_numeric_ok = True
    twd_values = []
    for sheet_name in ['Orders','Invoices','Expenses','Quotes','Summary']:
        sh = wb[sheet_name]
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,20)):
            headers = [c.value for c in row]
            if any(v is not None for v in headers):
                for idx, hv in enumerate(headers, start=1):
                    if hv and 'TWD' in str(hv).upper():
                        for rr in range(row[0].row+1, sh.max_row+1):
                            val = sh.cell(rr, idx).value
                            if val is not None:
                                twd_values.append((sheet_name, hv, val, type(val).__name__))
                                if not isinstance(val, (int,float)) and not (isinstance(val,str) and val.startswith('=')):
                                    twd_numeric_ok = False
                        break
                break
    add_check('twd_amounts_numeric_or_formula', twd_numeric_ok, twd_values[:20])

if quotes_json.exists():
    with open(quotes_json, encoding='utf-8') as f:
        q = json.load(f)
    print('quotes_json_type=', type(q).__name__, 'len=', len(q) if isinstance(q, list) else 'n/a', 'sample=', q[:2] if isinstance(q,list) else q)
    add_check('quotes_json_four_records', isinstance(q, list) and len(q)==4, len(q) if isinstance(q,list) else type(q).__name__)

print('CHECKS=', json.dumps(checks, ensure_ascii=False, indent=2, default=str))
if errors:
    print('ERRORS=', json.dumps(errors, ensure_ascii=False, indent=2))
    raise SystemExit(1)
print('All verification checks passed.')