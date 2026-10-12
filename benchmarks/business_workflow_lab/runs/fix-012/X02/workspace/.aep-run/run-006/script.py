from pathlib import Path
import csv, json, subprocess, sys, shutil, os
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE/'input'
OUTPUT = WORKSPACE/'output'
DOWNLOADS = OUTPUT/'downloads'
OUTPUT.mkdir(exist_ok=True)

customer_id = 'C001'
start_date = '2026-09-01'
end_date = '2026-09-30'
as_of = date.fromisoformat(end_date)
report_path = OUTPUT/f'project_report_{customer_id}_2026-09.xlsx'
quotes_json_path = OUTPUT/f'quotes_{customer_id}_2026-09.json'

print('Inputs:')
for p in sorted(INPUT.iterdir()): print(' -', p.name)
print('Source files expected:')
for p in [DOWNLOADS/'sap_orders_C001_2026-09.csv', DOWNLOADS/'sap_invoices_C001_2026-09.csv', OUTPUT/'expenses_all_approved_2026-09.csv']:
    print(' -', p, 'exists=', p.exists(), 'bytes=', p.stat().st_size if p.exists() else None)

# Helpers
def dec(v):
    if v is None or v == '': return Decimal('0')
    return Decimal(str(v))
def money(v):
    return dec(v).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
def parse_bool(v):
    return str(v).strip().lower() in ('true','1','yes','y')
def d(s):
    return date.fromisoformat(str(s)) if s else None
def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))
def write_cell(ws, row, col, value, num_fmt=None, as_text=False):
    c=ws.cell(row=row, column=col, value=value)
    if as_text:
        c.number_format='@'
        c.data_type='s'
    if num_fmt:
        c.number_format=num_fmt
    return c
def amount_twd(amount, currency, fx):
    if currency not in fx:
        raise ValueError(f'Missing FX rate for {currency}')
    return (money(amount) * fx[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def clear_data(ws):
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row-4)

def set_money_cell(ws, r, c, val):
    cell=ws.cell(r,c,value=float(val))
    cell.number_format='0.00'

def set_date_cell(ws, r, c, val):
    cell=ws.cell(r,c,value=d(val) if isinstance(val, str) else val)
    cell.number_format='yyyy-mm-dd'

# FX rates: use only attachment, first two columns; no live FX.
fx={}
with open(INPUT/'fx_rates.csv', encoding='utf-8-sig', newline='') as f:
    reader=csv.reader(f)
    header=next(reader)
    for row in reader:
        if len(row)>=2 and row[0]:
            fx[row[0].strip()] = dec(row[1].strip())
print('FX rates loaded from attachment:', fx)

orders_raw = read_csv(DOWNLOADS/'sap_orders_C001_2026-09.csv')
invoices = read_csv(DOWNLOADS/'sap_invoices_C001_2026-09.csv')
orders = [r for r in orders_raw if r.get('Status') != 'Cancelled']

# Quotes from CLI (no writes to quote system)
cli = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py'
cmd=[sys.executable, cli, 'quotes', 'list', '--customer', customer_id, '--from', start_date, '--to', end_date, '--format', 'json']
proc=subprocess.run(cmd, text=True, capture_output=True)
print('quote cli returncode:', proc.returncode)
if proc.returncode != 0:
    print(proc.stderr)
    raise SystemExit(proc.returncode)
quotes_data=json.loads(proc.stdout)
quotes=quotes_data['items']
quotes_json_path.write_text(json.dumps(quotes_data, ensure_ascii=False, indent=2), encoding='utf-8')

projects = sorted({r.get('ProjectID','') for r in orders+invoices+quotes if r.get('ProjectID','')})
print('C001 projects from orders/invoices/quotes:', projects)

expenses_all = read_csv(OUTPUT/'expenses_all_approved_2026-09.csv')
expenses = [r for r in expenses_all if r.get('ProjectID') in projects and r.get('Status') == 'Approved']

print('Counts:', {'orders_raw':len(orders_raw), 'orders_after_excluding_cancelled':len(orders), 'invoices':len(invoices), 'expenses_all_approved':len(expenses_all), 'expenses_after_project_filter':len(expenses), 'quotes':len(quotes)})
print('Order sample:', orders[:2])
print('Invoice sample:', invoices[:2])
print('Expense sample:', expenses[:2])
print('Quote sample:', quotes[:2])

# Load and fill a copy of the original template.
shutil.copyfile(INPUT/'project_template.xlsx', report_path)
wb=load_workbook(report_path)

for sheet in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    clear_data(wb[sheet])

# Report
rpt=wb['Report']
rpt['B3']=d(start_date); rpt['B3'].number_format='yyyy-mm-dd'
rpt['D3']=d(end_date); rpt['D3'].number_format='yyyy-mm-dd'
rpt['B4']=('CustomerID=C001; period=2026-09-01..2026-09-30; orders exclude Cancelled; '
           'expenses include Approved only for C001-related projects; quotes kept all; FX from input/fx_rates.csv')
gen=datetime.now().astimezone().replace(microsecond=0).isoformat()
rpt['B5']=gen
rpt['B7']='Lab ERP Sales Orders/Invoices; Expense Viewer; quote_cli; fx_rates.csv'
rpt['D7']=gen
rpt['B9']='C001 related projects: ' + ', '.join(projects) + '; attention rules: invoice due <= 2026-09-30 with unpaid balance, approved expense missing receipt, quote ValidUntil < 2026-09-30 and Status != Accepted.'

# Orders
ws=wb['Orders']
for idx,row in enumerate(orders, start=5):
    write_cell(ws,idx,1,row['OrderID'],as_text=True)
    write_cell(ws,idx,2,row['ProjectID'],as_text=True)
    write_cell(ws,idx,3,row['CustomerID'],as_text=True)
    set_date_cell(ws,idx,4,row['OrderDate'])
    write_cell(ws,idx,5,row['Status'])
    write_cell(ws,idx,6,row['Currency'])
    set_money_cell(ws,idx,7,money(row['NetAmount']))
    set_money_cell(ws,idx,8,amount_twd(row['NetAmount'], row['Currency'], fx))

# Invoices
ws=wb['Invoices']
for idx,row in enumerate(invoices, start=5):
    write_cell(ws,idx,1,row['InvoiceID'],as_text=True)
    write_cell(ws,idx,2,row['OrderID'],as_text=True)
    write_cell(ws,idx,3,row['ProjectID'],as_text=True)
    write_cell(ws,idx,4,row['CustomerID'],as_text=True)
    set_date_cell(ws,idx,5,row['InvoiceDate'])
    set_date_cell(ws,idx,6,row['DueDate'])
    write_cell(ws,idx,7,row['Currency'])
    set_money_cell(ws,idx,8,money(row['Amount']))
    set_money_cell(ws,idx,9,money(row['PaidAmount']))
    set_money_cell(ws,idx,10,money(row['Balance']))
    set_money_cell(ws,idx,11,amount_twd(row['Balance'], row['Currency'], fx))

# Expenses
ws=wb['Expenses']
for idx,row in enumerate(expenses, start=5):
    write_cell(ws,idx,1,row['ExpenseID'],as_text=True)
    write_cell(ws,idx,2,row['ProjectID'],as_text=True)
    write_cell(ws,idx,3,row['EmployeeID'],as_text=True)
    set_date_cell(ws,idx,4,row['ExpenseDate'])
    write_cell(ws,idx,5,row['Status'])
    write_cell(ws,idx,6,row['Currency'])
    set_money_cell(ws,idx,7,money(row['Amount']))
    set_money_cell(ws,idx,8,amount_twd(row['Amount'], row['Currency'], fx))

# Quotes
ws=wb['Quotes']
for idx,row in enumerate(quotes, start=5):
    write_cell(ws,idx,1,row['QuoteID'],as_text=True)
    write_cell(ws,idx,2,row['ProjectID'],as_text=True)
    write_cell(ws,idx,3,row['CustomerID'],as_text=True)
    set_date_cell(ws,idx,4,row['QuoteDate'])
    set_date_cell(ws,idx,5,row['ValidUntil'])
    write_cell(ws,idx,6,row['Status'])
    write_cell(ws,idx,7,row['Currency'])
    set_money_cell(ws,idx,8,money(row['NetAmount']))
    set_money_cell(ws,idx,9,amount_twd(row['NetAmount'], row['Currency'], fx))

# Summary
summary={p:{'OrderTWD':Decimal('0.00'),'InvoiceBalanceTWD':Decimal('0.00'),'ApprovedExpenseTWD':Decimal('0.00'),'AcceptedQuoteTWD':Decimal('0.00')} for p in projects}
for row in orders:
    summary[row['ProjectID']]['OrderTWD'] += amount_twd(row['NetAmount'], row['Currency'], fx)
for row in invoices:
    summary[row['ProjectID']]['InvoiceBalanceTWD'] += amount_twd(row['Balance'], row['Currency'], fx)
for row in expenses:
    summary[row['ProjectID']]['ApprovedExpenseTWD'] += amount_twd(row['Amount'], row['Currency'], fx)
for row in quotes:
    if row['Status'] == 'Accepted':
        summary.setdefault(row['ProjectID'], {'OrderTWD':Decimal('0.00'),'InvoiceBalanceTWD':Decimal('0.00'),'ApprovedExpenseTWD':Decimal('0.00'),'AcceptedQuoteTWD':Decimal('0.00')})
        summary[row['ProjectID']]['AcceptedQuoteTWD'] += amount_twd(row['NetAmount'], row['Currency'], fx)
ws=wb['Summary']
for idx,p in enumerate(sorted(summary), start=5):
    write_cell(ws,idx,1,p,as_text=True)
    set_money_cell(ws,idx,2,summary[p]['OrderTWD'])
    set_money_cell(ws,idx,3,summary[p]['InvoiceBalanceTWD'])
    set_money_cell(ws,idx,4,summary[p]['ApprovedExpenseTWD'])
    set_money_cell(ws,idx,5,summary[p]['AcceptedQuoteTWD'])
last=max(5, 4+len(summary))
for col,letter in enumerate(['B','C','D','E'], start=2):
    ws.cell(2,col).value=f'=SUM({letter}5:{letter}{last})'
    ws.cell(2,col).number_format='0.00'

# Exceptions
exceptions=[]
for row in invoices:
    if d(row['DueDate']) <= as_of and money(row['Balance']) > 0:
        exceptions.append(('Invoices', row['InvoiceID'], row['ProjectID'], f'Overdue unpaid invoice as of {end_date}; Balance={row["Balance"]} {row["Currency"]}'))
for row in expenses:
    if parse_bool(row.get('MissingReceipt','')):
        exceptions.append(('Expenses', row['ExpenseID'], row['ProjectID'], 'Approved expense missing receipt'))
for row in quotes:
    if d(row['ValidUntil']) < as_of and row['Status'] != 'Accepted':
        exceptions.append(('Quotes', row['QuoteID'], row['ProjectID'], f'Quote expired before {end_date} and not Accepted; Status={row["Status"]}'))
ws=wb['Exceptions']
for idx,(src,doc,proj,reason) in enumerate(exceptions, start=5):
    write_cell(ws,idx,1,src)
    write_cell(ws,idx,2,doc,as_text=True)
    write_cell(ws,idx,3,proj,as_text=True)
    write_cell(ws,idx,4,reason)

wb.save(report_path)

# Validation: reopen output and confirm requirements.
wb2=load_workbook(report_path, data_only=False)
checks=[]
def rows_of(sheet):
    sh=wb2[sheet]
    return list(sh.iter_rows(min_row=5, max_row=sh.max_row, values_only=True)) if sh.max_row>=5 else []
orders_out=rows_of('Orders'); invoices_out=rows_of('Invoices'); expenses_out=rows_of('Expenses'); quotes_out=rows_of('Quotes'); summary_out=rows_of('Summary'); exceptions_out=rows_of('Exceptions')
checks.append(('report_exists_nonempty', report_path.exists() and report_path.stat().st_size>0))
checks.append(('orders_exclude_cancelled', all(r[4] != 'Cancelled' for r in orders_out)))
checks.append(('expenses_approved_only', all(r[4] == 'Approved' for r in expenses_out)))
checks.append(('expenses_project_limited', all(r[1] in projects for r in expenses_out)))
checks.append(('quotes_all_kept_count', len(quotes_out)==len(quotes)))
checks.append(('id_text_preserved', wb2['Orders']['A5'].data_type == 's' and str(wb2['Orders']['A5'].value).startswith('0045')))
checks.append(('amounts_numeric', all(isinstance(r[7], (int,float)) for r in orders_out) and all(isinstance(r[10], (int,float)) for r in invoices_out) and all(isinstance(r[7], (int,float)) for r in expenses_out) and all(isinstance(r[8], (int,float)) for r in quotes_out)))
checks.append(('row_counts', (len(orders_out),len(invoices_out),len(expenses_out),len(quotes_out),len(summary_out),len(exceptions_out)) == (3,3,9,4,2,4)))
checks.append(('exception_ids', set(r[1] for r in exceptions_out) == {'0090000219','EXP-00006','Q0010','Q0014'}))
checks.append(('summary_formula_updated', wb2['Summary']['B2'].value == '=SUM(B5:B6)' and wb2['Summary']['E2'].value == '=SUM(E5:E6)'))

print('\nValidation checks:')
for name, ok in checks:
    print(f' - {name}: {ok}')
if not all(ok for _,ok in checks):
    raise SystemExit('Validation failed')
print('\nOutput workbook:', report_path)
print('Output quotes JSON:', quotes_json_path)
print('Final row counts:', {'Orders':len(orders_out),'Invoices':len(invoices_out),'Expenses':len(expenses_out),'Quotes':len(quotes_out),'Summary':len(summary_out),'Exceptions':len(exceptions_out)})
print('Exceptions:', exceptions)
