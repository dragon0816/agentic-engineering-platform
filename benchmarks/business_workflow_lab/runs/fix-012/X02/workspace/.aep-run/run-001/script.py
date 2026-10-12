from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
print('\nFX RATES SAMPLE:')
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i,row in zip(range(8), r):
            print(row)
print('\nTEMPLATE WORKBOOK:')
tpl = ws/'input'/'project_template.xlsx'
if tpl.exists():
    wb = load_workbook(tpl, data_only=False)
    print('sheets:', wb.sheetnames)
    for sh in wb.worksheets:
        print('SHEET', sh.title, 'max', sh.max_row, sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
            print(row)
        print('---')
print('\nMANUALS:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p = ws/'input'/name
    print('\n###', name)
    if p.exists():
        txt = p.read_text(encoding='utf-8', errors='replace')
        print(txt[:2500])
