import csv, os, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('output files:', sorted(p.name for p in (ws/'output').glob('*')))
for p in sorted((ws/'output').glob('expenses_*.csv')):
    print('\n---', p.name, '---')
    with open(p, encoding='utf-8-sig', newline='') as f:
        rows=list(csv.DictReader(f))
    print('count', len(rows), 'headers', rows[0].keys() if rows else [])
    for r in rows: print(r)
print('\n--- Instructions full ---')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=True)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True): print(row)
