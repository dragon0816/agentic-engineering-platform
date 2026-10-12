import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
input_dir=ws/'input'
print('workspace', ws)
print('input files:', sorted([p.name for p in input_dir.iterdir()]))
# print manuals
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=input_dir/name
    print('\n---', name, '---')
    print(p.read_text(encoding='utf-8'))
# fx sample
print('\n--- fx_rates.csv ---')
with open(input_dir/'fx_rates.csv', newline='', encoding='utf-8') as f:
    rows=list(csv.reader(f))
for r in rows[:20]: print(r)
# workbook structure
p=input_dir/'project_template.xlsx'
wb=load_workbook(p, data_only=False)
print('\n--- workbook sheets ---', wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('\nSHEET', sh, 'size', wsht.max_row, wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,10), values_only=True):
        print(row)
