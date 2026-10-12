from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# inspect fx
fx=ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    rows=list(r)
for row in rows[:10]: print(row)
# inspect manual first lines
manual=ws/'input'/'expense_manual.md'
print('\nexpense_manual.md first 80 lines:')
print('\n'.join(manual.read_text(encoding='utf-8').splitlines()[:80]))
# inspect workbook
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path)
print('\nWorkbook sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print(f'-- {sh}: max_row={wsht.max_row} max_col={wsht.max_column}')
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
        print(row)
    print('merged', list(wsht.merged_cells.ranges))
