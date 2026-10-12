from pathlib import Path
import csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from copy import copy
from datetime import date, datetime, timezone, timedelta
from openpyxl import load_workbook

ROOT = Path.cwd()
input_dir = ROOT / 'input'
output_dir = ROOT / 'output'
output_dir.mkdir(exist_ok=True)
template = input_dir / 'expense_template.xlsx'
export = output_dir / 'expense_export_2026-09_approved.csv'
fx_path = input_dir / 'fx_rates.csv'
outfile = output_dir / 'expense_report_2026-09_approved.xlsx'

FROM_DATE='2026-09-01'
TO_DATE='2026-09-30'
STATUS='Approved'
GENERATED='2026-10-06T00:00:00+08:00'

print('INPUTS:')
for p in sorted(input_dir.glob('*')):
    print('-', p.name, p.stat().st_size)
print('Export CSV:', export, 'exists=', export.exists(), 'bytes=', export.stat().st_size if export.exists() else None)

# Load FX rates (first two CSV columns are stable; Note text may contain commas)
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX headers:', header)
    for row in reader:
        if not row or not row[0].strip():
            continue
        fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates loaded:', {k: str(v) for k, v in fx.items()})

# Load exported data
with export.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    rows = list(reader)
headers = reader.fieldnames
print('Source headers:', headers)
print('Source row count:', len(rows))
print('First 3 rows:', rows[:3])

# Validate source filter
bad=[]
for r in rows:
    if not (FROM_DATE <= r['ExpenseDate'] <= TO_DATE and r['Status'] == STATUS):
        bad.append(r['ExpenseID'])
    if r['Currency'] not in fx:
        raise ValueError(f"Missing FX rate for {r['Currency']}")
if bad:
    raise ValueError(f'Rows outside requested filter: {bad}')

# Enrich rows
D = Decimal
for r in rows:
    amount = D(r['Amount'])
    rate = fx[r['Currency']]
    amount_twd = (amount * rate).quantize(D('0.01'), rounding=ROUND_HALF_UP)
    r['_AmountDecimal'] = amount
    r['_AmountTWDDecimal'] = amount_twd
    r['AmountTWD'] = f'{amount_twd:.2f}'
    r['NeedsReview'] = 'true' if amount_twd > D('5000.00') else 'false'
    mr = str(r.get('MissingReceipt','')).strip().lower()
    r['MissingReceipt'] = 'true' if mr == 'true' else 'false'

# Copy template, fill workbook
shutil.copyfile(template, outfile)
wb = load_workbook(outfile)

def clear_data(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row-min_row+1)

def copy_row_style(ws, source_row, target_row, max_col):
    # Preserve template style as much as possible; row 4 is the only formatted anchor row in the template.
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col+1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.protection:
            dst.protection = copy(src.protection)

# Report
rep = wb['Report']
rep['B3'] = date.fromisoformat(FROM_DATE)
rep['D3'] = date.fromisoformat(TO_DATE)
rep['B3'].number_format = 'yyyy-mm-dd'
rep['D3'].number_format = 'yyyy-mm-dd'
rep['B4'] = 'Status=Approved; Employee=All; Project=All; ExpenseDate includes start/end dates'
rep['B5'] = GENERATED
rep['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv fixed test rates'
rep['D7'] = GENERATED
rep['B9'] = 'AmountTWD = Amount × fixed RateToTWD, ROUND_HALF_UP to 2 decimals. NeedsReview=true when AmountTWD > 5000; false otherwise. Summary grouped by EmployeeID, Category, Currency to preserve original currency amounts.'

# Details
wsd = wb['Details']
clear_data(wsd, 5)
detail_cols = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
for i, r in enumerate(rows, start=5):
    copy_row_style(wsd, 4, i, len(detail_cols))
    values = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], date.fromisoformat(r['ExpenseDate']),
        r['Category'], r['Currency'], float(r['_AmountDecimal']), float(r['_AmountTWDDecimal']),
        r['Status'], r.get('ReceiptID','') or None, r.get('Description','') or None, r['MissingReceipt'], r['NeedsReview']
    ]
    for col, val in enumerate(values, start=1):
        c = wsd.cell(i, col, val)
        if col in [1,2,3,4,6,7,10,11,12,13,14]:
            c.number_format = '@'
        elif col == 5:
            c.number_format = 'yyyy-mm-dd'
        elif col in [8,9]:
            c.number_format = '0.00'
wsd.auto_filter.ref = f'A4:N{max(5, 4+len(rows))}'

# Summary grouped by employee + category + currency to preserve original-currency totals
wss = wb['Summary']
clear_data(wss, 5)
last_detail_row = 4 + len(rows) if rows else 5
wss['B2'] = f'=SUM(Details!I5:I{last_detail_row if rows else 5})'
summary = defaultdict(lambda: {'amount': D('0'), 'twd': D('0'), 'count': 0})
for r in rows:
    key = (r['EmployeeID'], r['Category'], r['Currency'])
    summary[key]['amount'] += r['_AmountDecimal']
    summary[key]['twd'] += r['_AmountTWDDecimal']
    summary[key]['count'] += 1
for idx, (key, vals) in enumerate(sorted(summary.items()), start=5):
    copy_row_style(wss, 4, idx, 6)
    emp, cat, cur = key
    rowvals = [emp, cat, cur, float(vals['amount'].quantize(D('0.01'), rounding=ROUND_HALF_UP)), float(vals['twd'].quantize(D('0.01'), rounding=ROUND_HALF_UP)), vals['count']]
    for col, val in enumerate(rowvals, start=1):
        cell = wss.cell(idx, col, val)
        if col in [1,2,3]: cell.number_format='@'
        elif col in [4,5]: cell.number_format='0.00'
        elif col == 6: cell.number_format='0'
wss.auto_filter.ref = f'A4:F{max(5, 4+len(summary))}'

# Exceptions: include rows with missing receipt, while still preserving them in Details.
wse = wb['Exceptions']
clear_data(wse, 5)
exceptions=[]
for r in rows:
    if r['MissingReceipt'] == 'true':
        exceptions.append((r['ExpenseID'], r['EmployeeID'], date.fromisoformat(r['ExpenseDate']), 'MissingReceipt=true; ReceiptID is blank'))
for idx, vals in enumerate(exceptions, start=5):
    copy_row_style(wse, 4, idx, 4)
    for col, val in enumerate(vals, start=1):
        c=wse.cell(idx, col, val)
        if col == 3: c.number_format='yyyy-mm-dd'
        else: c.number_format='@'
wse.auto_filter.ref = f'A4:D{max(5,4+len(exceptions))}'

wb.save(outfile)
print('WROTE', outfile)

# Reopen and verify requirements
wb2 = load_workbook(outfile, data_only=False)
d = wb2['Details']
s = wb2['Summary']
e = wb2['Exceptions']
read_headers=[d.cell(4,c).value for c in range(1,15)]
assert read_headers == detail_cols, read_headers
actual_detail=[]
for row in d.iter_rows(min_row=5, max_row=4+len(rows), max_col=14, values_only=True):
    actual_detail.append(row)
assert len(actual_detail)==len(rows)
assert all(row[0] for row in actual_detail)
id_to_src={r['ExpenseID']:r for r in rows}
true_review=[]
for row in actual_detail:
    expense_id=row[0]
    src=id_to_src[expense_id]
    date_cell = row[4].isoformat() if hasattr(row[4], 'isoformat') else str(row[4])
    assert date_cell == src['ExpenseDate']
    assert row[6] == src['Currency']
    assert Decimal(str(row[7])).quantize(D('0.01')) == src['_AmountDecimal'].quantize(D('0.01'))
    assert Decimal(str(row[8])).quantize(D('0.01')) == src['_AmountTWDDecimal']
    expected = 'true' if src['_AmountTWDDecimal'] > D('5000.00') else 'false'
    assert row[13] == expected
    if expected == 'true': true_review.append(expense_id)
summary_rows=[]
for row in s.iter_rows(min_row=5, max_row=4+len(summary), max_col=6, values_only=True):
    summary_rows.append(row)
assert len(summary_rows)==len(summary)
for group, sub, cur, amt, twd, cnt in summary_rows:
    vals=summary[(group, sub, cur)]
    assert Decimal(str(amt)).quantize(D('0.01')) == vals['amount'].quantize(D('0.01'))
    assert Decimal(str(twd)).quantize(D('0.01')) == vals['twd'].quantize(D('0.01'))
    assert cnt == vals['count']
assert s['B2'].value == f'=SUM(Details!I5:I{4+len(rows)})'
ex_rows=[]
if exceptions:
    for row in e.iter_rows(min_row=5, max_row=4+len(exceptions), max_col=4, values_only=True): ex_rows.append(row)
assert len(ex_rows)==len(exceptions)
# Confirm no extra non-empty Details row immediately after expected range.
extra = [d.cell(5+len(rows), c).value for c in range(1, 15)] if d.max_row >= 5+len(rows) else []
assert not extra or all(v is None for v in extra), extra

print('CHECKS PASSED')
print('detail_rows', len(actual_detail))
print('summary_rows', len(summary_rows))
print('exceptions_rows', len(ex_rows))
print('needs_review_true', true_review)
print('details_headers', read_headers)
print('summary_formula_B2', s['B2'].value)
print('output_file', str(outfile))
