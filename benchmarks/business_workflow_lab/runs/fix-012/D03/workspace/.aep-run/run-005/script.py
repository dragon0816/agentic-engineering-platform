from pathlib import Path
import csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from openpyxl import load_workbook

ROOT = Path.cwd()
input_dir = ROOT / 'input'
output_dir = ROOT / 'output'
output_dir.mkdir(exist_ok=True)
template = input_dir / 'expense_template.xlsx'
export = output_dir / 'expense_export_2026-09_approved.csv'
fx_path = input_dir / 'fx_rates.csv'
outfile = output_dir / 'expense_report_2026-09_approved.xlsx'

FROM_DATE='2026-09-01'
TO_DATE='2026-09-30'
STATUS='Approved'
GENERATED='2026-10-06T00:00:00+08:00'

print('INPUT FILES:')
for p in sorted(input_dir.glob('*')):
    print('-', p.name, p.stat().st_size)
print('USING EXPORT:', export.name, export.stat().st_size)

fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    print('FX HEADER/SAMPLE:', next(reader))
    for row in reader:
        if row and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', {k: str(v) for k, v in fx.items()})

with export.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    rows = list(reader)
headers = reader.fieldnames
print('Export headers:', headers)
print('Export rows:', len(rows))
print('Export sample:', rows[:2])

for r in rows:
    if not (FROM_DATE <= r['ExpenseDate'] <= TO_DATE):
        raise ValueError(f"Outside date range: {r}")
    if r['Status'] != STATUS:
        raise ValueError(f"Outside status: {r}")
    if r['Currency'] not in fx:
        raise ValueError(f"Missing FX rate for {r['Currency']}")
    amount = Decimal(r['Amount'])
    amount_twd = (amount * fx[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    r['_AmountDecimal'] = amount
    r['_AmountTWDDecimal'] = amount_twd
    r['MissingReceipt'] = 'true' if str(r.get('MissingReceipt','')).strip().lower() == 'true' else 'false'
    r['NeedsReview'] = 'true' if amount_twd > Decimal('5000.00') else 'false'

shutil.copyfile(template, outfile)
wb = load_workbook(outfile)

def clear_data(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row-min_row+1)

def copy_row_style(ws, source_row, target_row, max_col):
    ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height
    for col in range(1, max_col+1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)

rep=wb['Report']
rep['B3']=date.fromisoformat(FROM_DATE); rep['B3'].number_format='yyyy-mm-dd'
rep['D3']=date.fromisoformat(TO_DATE); rep['D3'].number_format='yyyy-mm-dd'
rep['B4']='Status=Approved; Employee=All; Project=All; ExpenseDate includes start/end dates'
rep['B5']=GENERATED
rep['B7']='Expense Viewer (synthetic lab); fx_rates.csv fixed test rates'
rep['D7']=GENERATED
rep['B9']='AmountTWD = Amount × fixed RateToTWD, ROUND_HALF_UP to 2 decimals. NeedsReview=true when AmountTWD > 5000; false otherwise. Summary grouped by EmployeeID, Category, Currency to preserve original currency amounts.'

# details
detail_cols=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
wsd=wb['Details']; clear_data(wsd,5)
for i,r in enumerate(rows, start=5):
    copy_row_style(wsd,4,i,14)
    vals=[r['ExpenseID'],r['EmployeeID'],r['EmployeeName'],r['ProjectID'],date.fromisoformat(r['ExpenseDate']),r['Category'],r['Currency'],float(r['_AmountDecimal']),float(r['_AmountTWDDecimal']),r['Status'],r.get('ReceiptID') or None,r.get('Description') or None,r['MissingReceipt'],r['NeedsReview']]
    for c,v in enumerate(vals, start=1):
        cell=wsd.cell(i,c,v)
        if c in (8,9): cell.number_format='0.00'
        elif c == 5: cell.number_format='yyyy-mm-dd'
        else: cell.number_format='@'
wsd.auto_filter.ref=f'A4:N{max(5,4+len(rows))}'

# summary
wss=wb['Summary']; clear_data(wss,5)
last=4+len(rows) if rows else 5
wss['B2']=f'=SUM(Details!I5:I{last})'
summary=defaultdict(lambda:{'amount':Decimal('0.00'),'twd':Decimal('0.00'),'count':0})
for r in rows:
    key=(r['EmployeeID'], r['Category'], r['Currency'])
    summary[key]['amount'] += r['_AmountDecimal']
    summary[key]['twd'] += r['_AmountTWDDecimal']
    summary[key]['count'] += 1
for i,(key,val) in enumerate(sorted(summary.items()), start=5):
    copy_row_style(wss,4,i,6)
    vals=[key[0],key[1],key[2],float(val['amount'].quantize(Decimal('0.01'))),float(val['twd'].quantize(Decimal('0.01'))),val['count']]
    for c,v in enumerate(vals, start=1):
        cell=wss.cell(i,c,v)
        if c in (4,5): cell.number_format='0.00'
        elif c == 6: cell.number_format='0'
        else: cell.number_format='@'
wss.auto_filter.ref=f'A4:F{max(5,4+len(summary))}'

# exceptions
wse=wb['Exceptions']; clear_data(wse,5)
exceptions=[]
for r in rows:
    if r['MissingReceipt']=='true':
        exceptions.append((r['ExpenseID'], r['EmployeeID'], date.fromisoformat(r['ExpenseDate']), 'MissingReceipt=true; ReceiptID is blank'))
for i, vals in enumerate(exceptions, start=5):
    copy_row_style(wse,4,i,4)
    for c,v in enumerate(vals, start=1):
        cell=wse.cell(i,c,v)
        cell.number_format='yyyy-mm-dd' if c==3 else '@'
wse.auto_filter.ref=f'A4:D{max(5,4+len(exceptions))}'

wb.save(outfile)
print('WROTE:', outfile)

# Verification: reopen all outputs and check requirements.
# CSV output check
with export.open(newline='', encoding='utf-8-sig') as f:
    csv_rows=list(csv.DictReader(f))
assert len(csv_rows)==15
# XLSX output check
wbv=load_workbook(outfile, data_only=False)
assert wbv.sheetnames == ['Report','Details','Summary','Exceptions','Instructions']
d=wbv['Details']; s=wbv['Summary']; e=wbv['Exceptions']
assert [d.cell(4,c).value for c in range(1,15)] == detail_cols
actual=[]
for row in d.iter_rows(min_row=5, max_row=4+len(rows), max_col=14, values_only=True):
    actual.append(row)
assert len(actual)==15
src={r['ExpenseID']:r for r in rows}
needs_true=[]; amount_twd_total=Decimal('0.00')
def cell_date_str(v):
    if isinstance(v, datetime): return v.date().isoformat()
    if isinstance(v, date): return v.isoformat()
    return str(v)
for row in actual:
    eid=row[0]
    r=src[eid]
    assert row[1] == r['EmployeeID']
    assert cell_date_str(row[4]) == r['ExpenseDate']
    assert row[6] == r['Currency']
    assert row[9] == 'Approved'
    assert Decimal(str(row[7])).quantize(Decimal('0.01')) == r['_AmountDecimal']
    assert Decimal(str(row[8])).quantize(Decimal('0.01')) == r['_AmountTWDDecimal']
    amount_twd_total += Decimal(str(row[8])).quantize(Decimal('0.01'))
    expected_review='true' if r['_AmountTWDDecimal'] > Decimal('5000.00') else 'false'
    assert row[13] == expected_review
    if row[13]=='true': needs_true.append(eid)
# must keep original currency and original amount in Details
assert all(d.cell(rn,7).value for rn in range(5,20))
assert all(isinstance(d.cell(rn,8).value, (int,float)) for rn in range(5,20))
# Summary check
summary_rows=list(s.iter_rows(min_row=5, max_row=4+len(summary), max_col=6, values_only=True))
assert len(summary_rows)==len(summary)
summary_twd=Decimal('0.00')
for group, subgroup, cur, amt, twd, cnt in summary_rows:
    val=summary[(group,subgroup,cur)]
    assert Decimal(str(amt)).quantize(Decimal('0.01')) == val['amount'].quantize(Decimal('0.01'))
    assert Decimal(str(twd)).quantize(Decimal('0.01')) == val['twd'].quantize(Decimal('0.01'))
    assert cnt == val['count']
    summary_twd += Decimal(str(twd)).quantize(Decimal('0.01'))
assert summary_twd == amount_twd_total
assert s['B2'].value == '=SUM(Details!I5:I19)'
# Exceptions check
exrows=list(e.iter_rows(min_row=5, max_row=4+len(exceptions), max_col=4, values_only=True)) if exceptions else []
assert len(exrows)==len(exceptions)
for er in exrows:
    assert src[er[0]]['MissingReceipt']=='true'
# no extra data row
if d.max_row >= 20:
    assert all(d.cell(20,c).value is None for c in range(1,15))

print('CHECKS PASSED')
print('Details rows:', len(actual))
print('Summary rows:', len(summary_rows))
print('Exceptions rows:', len(exrows))
print('NeedsReview=true:', needs_true)
print('AmountTWD total:', amount_twd_total)
print('Summary B2 formula:', s['B2'].value)
print('Output files checked:', export.name, outfile.name)
