from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
export=ws/'output'/'expense_export_2026-09_approved.csv'
print('EXPORTED EXISTS', export.exists(), export.stat().st_size if export.exists() else None)
with export.open(newline='', encoding='utf-8-sig') as f:
    reader=csv.DictReader(f)
    rows=list(reader)
print('Export headers:', reader.fieldnames)
print('Export row count:', len(rows))
for r in rows[:20]:
    print(r)
print('\nInstructions all rows:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
