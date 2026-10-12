from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
for name in ['submitted_2026-09.csv','approved_2026-09.csv']:
    p=ws/'output'/name
    print('\nCSV', name, 'exists', p.exists(), 'size', p.stat().st_size if p.exists() else None)
    if p.exists():
        with open(p, newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f)
            rows=list(r)
        print('rows incl header', len(rows))
        for row in rows[:6]: print(row)

wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
ins=wb['Instructions']
print('\nInstructions all:')
for row in ins.iter_rows(values_only=True):
    print(row)
print('\nReport merged:', list(wb['Report'].merged_cells.ranges))
print('Details merged:', list(wb['Details'].merged_cells.ranges))
print('Summary merged:', list(wb['Summary'].merged_cells.ranges))
print('Exceptions merged:', list(wb['Exceptions'].merged_cells.ranges))