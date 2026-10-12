from pathlib import Path
from openpyxl import load_workbook
import csv, json
ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size)

# workbook structure
xlsx = ws/'input'/'expense_template.xlsx'
wb = load_workbook(xlsx)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'-- sheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        print(r)

# csv headers/sample
csvp = ws/'input'/'fx_rates.csv'
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i,row in zip(range(8), reader):
        print('fx row', i+1, row)

md = ws/'input'/'expense_manual.md'
print('manual preview:')
print(md.read_text(encoding='utf-8')[:2000])