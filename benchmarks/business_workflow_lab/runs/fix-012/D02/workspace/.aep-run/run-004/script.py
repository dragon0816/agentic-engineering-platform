from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
import csv, sys
ws=Path.cwd()
out=ws/'output'/'expense_report_2026-09_submitted_approved.xlsx'
fx=ws/'input'/'fx_rates.csv'
print('checking', out, out.exists(), out.stat().st_size if out.exists() else None)
assert out.exists() and out.stat().st_size > 0

rates={}
with open(fx, newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f); next(r)
    for row in r:
        if len(row)>=2 and row[0]: rates[row[0]]=Decimal(row[1])

wb=load_workbook(out, data_only=False)
expected_sheets=['Report','Details','Summary','Exceptions','Instructions']
print('sheets', wb.sheetnames)
assert wb.sheetnames == expected_sheets

# Check headers preserved
expected_details=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
actual_details=[wb['Details'].cell(4,c).value for c in range(1,15)]
print('details headers', actual_details)
assert actual_details == expected_details
expected_ex=['ExpenseID','EmployeeID','ExpenseDate','Reason']
actual_ex=[wb['Exceptions'].cell(4,c).value for c in range(1,5)]
assert actual_ex == expected_ex

# Extract rows without relying on max_row extra blanks
sh=wb['Details']
detail_rows=[]
for row in sh.iter_rows(min_row=5, max_col=14, values_only=True):
    if row[0] is not None:
        detail_rows.append(row)
print('detail row count', len(detail_rows))
assert len(detail_rows)==22
ids=[r[0] for r in detail_rows]
assert len(ids)==len(set(ids))
for r in detail_rows:
    eid, emp, name, proj, date, cat, curr, amt, twd, status, receipt, desc, missing, needs = r
    assert isinstance(eid, str) and eid.startswith('EXP-')
    assert '2026-09-01' <= str(date) <= '2026-09-30', (eid,date)
    assert status in {'Submitted','Approved'}, (eid,status)
    assert curr in rates, curr
    expected_twd=(Decimal(str(amt))*rates[curr]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert Decimal(str(twd)).quantize(Decimal('0.01')) == expected_twd, (eid, amt, curr, twd, expected_twd)
    assert desc not in (None,''), ('missing desc', eid)
    assert str(missing).lower() in {'true','false'}
    assert str(needs).lower() in {'true','false'}

missing_ids=[r[0] for r in detail_rows if str(r[12]).lower()=='true']
print('missing in details', missing_ids)
assert set(missing_ids)=={'EXP-00015','EXP-00006'}
# Missing receipt rows should remain in Details and have blank receipt ID allowed
for r in detail_rows:
    if r[0] in missing_ids:
        assert r[10] in (None,''), (r[0], r[10])
        assert str(r[13]).lower()=='true'

# Exceptions exact missing ids and original IDs
ex=wb['Exceptions']
ex_rows=[]
for row in ex.iter_rows(min_row=5, max_col=4, values_only=True):
    if row[0] is not None:
        ex_rows.append(row)
print('exception rows', ex_rows)
assert len(ex_rows)==2
assert set(r[0] for r in ex_rows)==set(missing_ids)
assert all(r[3]=='Missing receipt' for r in ex_rows)

# Summary
summary=wb['Summary']
print('summary B2 formula', summary['B2'].value)
assert summary['B2'].value == '=SUM(Details!I5:I26)'
sum_rows=[]
for row in summary.iter_rows(min_row=5, max_col=6, values_only=True):
    if row[0] is not None:
        sum_rows.append(row)
print('summary rows', sum_rows)
assert len(sum_rows)==6
assert sum(r[5] for r in sum_rows)==22
# Totals match Details by status/currency
from collections import defaultdict
calc=defaultdict(lambda: [Decimal('0'),Decimal('0'),0])
for r in detail_rows:
    key=(r[9],r[6])
    calc[key][0]+=Decimal(str(r[7]))
    calc[key][1]+=Decimal(str(r[8]))
    calc[key][2]+=1
for group, sub, curr, amount, twd, count in sum_rows:
    key=(group,curr)
    assert key in calc, key
    assert Decimal(str(amount)).quantize(Decimal('0.01')) == calc[key][0].quantize(Decimal('0.01'))
    assert Decimal(str(twd)).quantize(Decimal('0.01')) == calc[key][1].quantize(Decimal('0.01'))
    assert count == calc[key][2]

# Report metadata
report=wb['Report']
assert report['B3'].value=='2026-09-01'
assert report['D3'].value=='2026-09-30'
assert 'Submitted' in report['B4'].value and 'Approved' in report['B4'].value
assert 'Expense Viewer' in report['B7'].value
print('CHECK PASSED: workbook requirements verified')