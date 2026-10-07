from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Workbook structure
xlsx = ws/'input'/'project_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'\nSheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in range(1, min(sh.max_row, 8)+1):
        vals = [sh.cell(r,c).value for c in range(1, min(sh.max_column, 12)+1)]
        print(' row', r, vals)

# FX rates
fx = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i,row in enumerate(reader):
        print(i+1, row)
        if i>=10: break

# Manual
md = ws/'input'/'sap_manual.md'
print('\nsap_manual.md first 200 lines:')
text = md.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:200], start=1):
    print(f'{i:03d}: {line}')
