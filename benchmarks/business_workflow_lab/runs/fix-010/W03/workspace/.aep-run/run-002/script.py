from pathlib import Path
from openpyxl import load_workbook
import csv

ws=Path.cwd()
print('Instructions rows 8-56:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in range(8, sh.max_row+1):
    vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
    print(r, vals)

csv_path=ws/'output'/'downloads'/'invoices_2026-09.csv'
print('\nDownloaded CSV exists:', csv_path.exists(), csv_path)
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in enumerate(reader):
        print(i+1,row)
        if i>=20: break
