from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Workbook details
xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    wb = load_workbook(xlsx, data_only=False)
    print('workbook sheets:', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print(f'sheet {s}: {sh.max_row} rows x {sh.max_column} cols')
        for r in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
            print('  ', r)

# FX csv details
csvp = ws/'input'/'fx_rates.csv'
if csvp.exists():
    with csvp.open(newline='', encoding='utf-8-sig') as f:
        rows = list(csv.reader(f))
    print('fx rows:', len(rows))
    for r in rows[:8]: print('  ', r)

# Manual excerpt
md = ws/'input'/'quote_cli_manual.md'
if md.exists():
    text = md.read_text(encoding='utf-8', errors='replace')
    print('manual first 4000 chars:')
    print(text[:4000])