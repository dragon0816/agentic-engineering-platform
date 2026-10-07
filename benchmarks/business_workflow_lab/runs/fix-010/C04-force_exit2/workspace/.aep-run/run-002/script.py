from pathlib import Path
import subprocess, json, csv, shutil, sys, time, os
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
CONDITION = f'Status={STATUS}; QuoteDate from {FROM_DATE} to {TO_DATE} inclusive'
SOURCE_NAME = 'quote_cli synthetic quote system'

CLI_CMD = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
CLI_PY = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

def base_cmd():
    if CLI_CMD.exists():
        return [str(CLI_CMD)]
    if CLI_PY.exists():
        return [sys.executable, str(CLI_PY)]
    raise FileNotFoundError(f'Neither CLI command nor Python CLI found: {CLI_CMD} / {CLI_PY}')

def run_cli(args, label, max_temp_retries=2):
    attempts = []
    retries = 0
    while True:
        cmd = base_cmd() + args
        cp = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
        rec = {
            'label': label,
            'args': args,
            'returncode': cp.returncode,
            'stdout_len': len(cp.stdout or ''),
            'stderr': (cp.stderr or '').strip(),
        }
        attempts.append(rec)
        if cp.returncode == 0:
            # stdout may be used only here, and only after it parses.
            try:
                data = json.loads(cp.stdout)
            except Exception as e:
                raise RuntimeError(f'{label}: CLI succeeded but stdout was not parseable JSON: {e}')
            return data, attempts
        if cp.returncode == 4 and retries < max_temp_retries:
            retries += 1
            time.sleep(0.2)
            continue
        if cp.returncode in (2, 3):
            raise RuntimeError(f'{label}: CLI returned exit code {cp.returncode}; stderr: {(cp.stderr or "").strip()}')
        if cp.returncode == 4:
            raise RuntimeError(f'{label}: CLI returned temporary service error exit code 4 after {len(attempts)} attempts; stderr: {(cp.stderr or "").strip()}')
        raise RuntimeError(f'{label}: CLI returned unexpected exit code {cp.returncode}; stderr: {(cp.stderr or "").strip()}')

run_log = {
    'query': {'from': FROM_DATE, 'to': TO_DATE, 'status': STATUS},
    'attempts': [],
    'result': None,
}
try:
    list_args = ['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
    listed, attempts = run_cli(list_args, 'quotes list')
    run_log['attempts'].extend(attempts)
    if listed.get('schema_version') != 1 or 'items' not in listed:
        raise RuntimeError('quotes list JSON did not match expected schema_version 1 with items')
    headers = listed['items']
    details = []
    for h in headers:
        qid = h.get('QuoteID')
        if not qid:
            raise RuntimeError(f'quotes list item missing QuoteID: {h!r}')
        got, gattempts = run_cli(['quotes', 'get', qid, '--format', 'json'], f'quotes get {qid}')
        run_log['attempts'].extend(gattempts)
        if got.get('schema_version') != 1 or 'header' not in got or 'line_items' not in got:
            raise RuntimeError(f'quotes get {qid} JSON did not match expected schema')
        details.append(got)

    # Load FX rates. CSV notes have unquoted commas in test fixture, so use first two columns only.
    rates = {}
    with (INPUT / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        hdr = next(reader)
        for row in reader:
            if len(row) >= 2 and row[0]:
                rates[row[0]] = Decimal(row[1])

    template = INPUT / 'quote_template.xlsx'
    out_xlsx = OUTPUT / 'quote_report_2026-09_Sent.xlsx'
    shutil.copyfile(template, out_xlsx)
    wb = load_workbook(out_xlsx)

    now_tpe = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()

    # Report sheet
    ws = wb['Report']
    ws['B3'] = FROM_DATE
    ws['D3'] = TO_DATE
    ws['B4'] = CONDITION
    ws['B5'] = now_tpe
    ws['B7'] = SOURCE_NAME
    ws['D7'] = now_tpe

    # Clear old rows from row 5 onward, if any
    def clear_data(sheet_name, start_row):
        sh = wb[sheet_name]
        if sh.max_row >= start_row:
            sh.delete_rows(start_row, sh.max_row - start_row + 1)

    clear_data('Quotes', 5)
    clear_data('Lines', 5)
    clear_data('Summary', 5)

    quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
    line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']

    qws = wb['Quotes']
    lws = wb['Lines']
    quote_row = 5
    line_row = 5
    summary = {}  # (Group, SubGroup, Currency) -> [Amount, AmountTWD, Count]
    total_twd = Decimal('0')

    for d in details:
        h = d['header']
        currency = h['Currency']
        if currency not in rates:
            raise RuntimeError(f'Missing FX rate for currency {currency}')
        rate = rates[currency]
        net = Decimal(str(h['NetAmount']))
        net_twd = (net * rate).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        total_twd += net_twd
        row_values = []
        for col in quote_cols:
            if col == 'NetAmountTWD':
                row_values.append(float(net_twd))
            elif col in ('Subtotal','DiscountPercent','NetAmount'):
                row_values.append(float(Decimal(str(h[col]))))
            else:
                row_values.append(h.get(col))
        for c, v in enumerate(row_values, 1):
            qws.cell(row=quote_row, column=c, value=v)
        quote_row += 1

        # summary by status and currency (plain grouped total required by Summary columns)
        key = ('Status', h.get('Status'), currency)
        if key not in summary:
            summary[key] = [Decimal('0'), Decimal('0'), 0]
        summary[key][0] += net
        summary[key][1] += net_twd
        summary[key][2] += 1

        for li in d['line_items']:
            for c, col in enumerate(line_cols, 1):
                val = li.get(col) if col != 'QuoteID' else h['QuoteID']
                if col in ('Quantity',):
                    val = int(val)
                elif col in ('UnitPrice','LineAmount'):
                    val = float(Decimal(str(val)))
                lws.cell(row=line_row, column=c, value=val)
            line_row += 1

    sws = wb['Summary']
    sws['B2'] = float(total_twd)
    sr = 5
    for (group, subgroup, currency), (amount, amount_twd, count) in sorted(summary.items()):
        vals = [group, subgroup, currency, float(amount.quantize(Decimal('0.01'))), float(amount_twd.quantize(Decimal('0.01'))), count]
        for c, v in enumerate(vals, 1):
            sws.cell(row=sr, column=c, value=v)
        sr += 1

    wb.save(out_xlsx)
    run_log['result'] = {'status': 'success', 'workbook': str(out_xlsx), 'quotes': len(details), 'lines': line_row - 5}
    (OUTPUT / 'quote_cli_run_log.json').write_text(json.dumps(run_log, ensure_ascii=False, indent=2), encoding='utf-8')
    print('SUCCESS workbook:', out_xlsx)
    print('quotes:', len(details), 'lines:', line_row - 5, 'cli attempts:', len(run_log['attempts']))

except Exception as e:
    run_log['result'] = {'status': 'failed', 'reason': str(e)}
    err_file = OUTPUT / 'quote_cli_error_2026-09_Sent.txt'
    log_file = OUTPUT / 'quote_cli_run_log.json'
    err_file.write_text(str(e) + '\n', encoding='utf-8')
    log_file.write_text(json.dumps(run_log, ensure_ascii=False, indent=2), encoding='utf-8')
    print('FAILED:', str(e))
    print('error_file:', err_file)
    print('log_file:', log_file)
    # Do not raise: keep generated diagnostic files for delivery.
