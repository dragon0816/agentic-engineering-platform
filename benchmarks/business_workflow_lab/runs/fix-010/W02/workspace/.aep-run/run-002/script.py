from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd(); wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
for s in wb.sheetnames:
    sh=wb[s]
    print('\nSHEET',s)
    print('merged:', list(sh.merged_cells.ranges))
    for r in sh.iter_rows(values_only=True):
        print(r)
