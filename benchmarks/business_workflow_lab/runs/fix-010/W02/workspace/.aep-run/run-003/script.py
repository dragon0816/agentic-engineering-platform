from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook
from openpyxl.styles import copy
import csv, json

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)
outfile = output_dir / 'sales_report_C001_2026-08-01_2026-10-31.xlsx'

# Data obtained from Lab ERP browser operations: exact Customer ID filter C001, 2026-08-01..2026-10-31, then each order detail page.
detail_records = [
    {"OrderID":"0045000101","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-08-03","Status":"Closed","Currency":"TWD","NetAmount":"95895.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET15"},
    {"OrderID":"0045000107","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P003","OrderDate":"2026-08-27","Status":"Closed","Currency":"USD","NetAmount":"6250.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET30"},
    {"OrderID":"0045000109","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P003","OrderDate":"2026-09-01","Status":"Open","Currency":"TWD","NetAmount":"150000.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET10"},
    {"OrderID":"0045000114","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-09-08","Status":"Cancelled","Currency":"USD","NetAmount":"11956.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET5"},
    {"OrderID":"0045000119","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-09-15","Status":"Open","Currency":"EUR","NetAmount":"8775.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET30"},
    {"OrderID":"0045000125","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-09-25","Status":"Closed","Currency":"TWD","NetAmount":"386445.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET0"},
    {"OrderID":"0045000129","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-10-01","Status":"Open","Currency":"USD","NetAmount":"1602.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET30"},
    {"OrderID":"0045000132","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P003","OrderDate":"2026-10-09","Status":"Open","Currency":"TWD","NetAmount":"50400.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":""},
    {"OrderID":"0045000136","CustomerID":"C001","CustomerName":"Acme Precision Taiwan Ltd.","ProjectID":"P001","OrderDate":"2026-10-31","Status":"Open","Currency":"EUR","NetAmount":"2250.00","CustomerContact":"Grace Lin <grace.lin@acme-precision.example>","PaymentTerms":"NET5"},
]

# Print true input facts for traceability.
print('Template:', input_dir/'sales_template.xlsx')
print('FX file:', input_dir/'fx_rates.csv')
print('Downloaded ERP CSV:', output_dir/'downloads'/'orders_C001_2026-08-01_2026-10-31.csv')
print('Detail record count:', len(detail_records))
print('Detail sample:', detail_records[0])

# Read rates, first two columns only because Note can contain commas in test fixture.
rates = {}
with (input_dir/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

# Cross-check downloaded filtered list before filling.
download = output_dir/'downloads'/'orders_C001_2026-08-01_2026-10-31.csv'
with download.open(newline='', encoding='utf-8-sig') as f:
    dl_rows = list(csv.DictReader(f))
print('Downloaded row count:', len(dl_rows))
assert len(dl_rows) == len(detail_records), 'downloaded row count differs from detail count'
assert all(r['CustomerID'] == 'C001' for r in dl_rows), 'download has non-C001 rows'
assert all('2026-08-01' <= r['OrderDate'] <= '2026-10-31' for r in dl_rows), 'download has out-of-range rows'
assert {r['OrderID'] for r in dl_rows} == {r['OrderID'] for r in detail_records}, 'detail order IDs differ from downloaded list'

q = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP)

def clone_row_style(sh, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = sh.cell(src_row, col)
        dst = sh.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

wb = load_workbook(input_dir/'sales_template.xlsx')
report = wb['Report']; details = wb['Details']; summary = wb['Summary']

# Report metadata
now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B3'] = '2026-08-01'
report['D3'] = '2026-10-31'
report['B4'] = 'CustomerID = C001 (exact match); OrderDate 2026-08-01 to 2026-10-31 inclusive; all statuses; sorted by OrderID ascending. Similar-name customers excluded.'
report['B5'] = now
report['B7'] = 'Lab ERP Sales Orders; fx_rates.csv'
report['D7'] = now

missing = []
for rec in detail_records:
    for field in ['CustomerContact', 'PaymentTerms']:
        if rec.get(field, '') == '':
            missing.append(f"{rec['OrderID']} missing {field}")
report['B9'] = '; '.join(missing) if missing else 'No missing CustomerContact or PaymentTerms fields. Similar-name customer C007 was not included.'

# Clear any old data rows (template has none but keep rerunnable).
for sh in (details, summary):
    if sh.max_row > 4:
        sh.delete_rows(5, sh.max_row-4)

headers = [details.cell(4, c).value for c in range(1, details.max_column+1)]
expected = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert headers == expected, headers

# Details rows
for idx, rec in enumerate(detail_records, start=5):
    clone_row_style(details, 4, idx, len(expected))
    amount = money(rec['NetAmount'])
    twd = money(amount * rates[rec['Currency']])
    vals = [
        rec['OrderID'], rec['CustomerID'], rec['CustomerName'], rec['ProjectID'], rec['OrderDate'], rec['Status'], rec['Currency'],
        float(amount), float(twd), rec['CustomerContact'], rec['PaymentTerms'] if rec['PaymentTerms'] != '' else None
    ]
    for col, val in enumerate(vals, start=1):
        cell = details.cell(idx, col, val)
        if col in (1,2,4):
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9):
            cell.number_format = '0.00'

last = 4 + len(detail_records)
summary['B2'] = f'=SUM(Details!I5:I{last})' if detail_records else '=SUM(Details!I5:I5)'

# Summary by Currency: avoids mixing original currencies while still providing required totals.
by_cur = {}
for rec in detail_records:
    cur = rec['Currency']
    amount = money(rec['NetAmount'])
    twd = money(amount * rates[cur])
    by_cur.setdefault(cur, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    by_cur[cur]['amount'] += amount
    by_cur[cur]['twd'] += twd
    by_cur[cur]['count'] += 1

for r, cur in enumerate(sorted(by_cur), start=5):
    clone_row_style(summary, 4, r, 6)
    data = by_cur[cur]
    vals = ['CustomerID', 'C001', cur, float(money(data['amount'])), float(money(data['twd'])), data['count']]
    for c, v in enumerate(vals, start=1):
        cell = summary.cell(r, c, v)
        if c in (4,5): cell.number_format = '0.00'

# Filters: extend existing / create if useful without changing layout
if details.auto_filter.ref:
    details.auto_filter.ref = f'A4:K{last}'
else:
    details.auto_filter.ref = f'A4:K{last}'
if summary.auto_filter.ref:
    summary.auto_filter.ref = f'A4:F{4+len(by_cur)}'
else:
    summary.auto_filter.ref = f'A4:F{4+len(by_cur)}'

wb.save(outfile)
print('Wrote', outfile)

# Reopen and verify stated/template requirements.
wb2 = load_workbook(outfile, data_only=False)
assert wb2.sheetnames == ['Report','Details','Summary','Instructions']
r = wb2['Report']; d = wb2['Details']; s = wb2['Summary']
assert str(r['B3'].value) == '2026-08-01' and str(r['D3'].value) == '2026-10-31'
assert 'C001' in r['B4'].value and 'Similar-name' in r['B4'].value
assert d.max_row == 13, d.max_row
read_headers = [d.cell(4,c).value for c in range(1,12)]
assert read_headers == expected
rows = []
for rr in range(5, d.max_row+1):
    row = {expected[c-1]: d.cell(rr,c).value for c in range(1,12)}
    rows.append(row)
assert len(rows) == 9
assert all(row['CustomerID'] == 'C001' for row in rows)
assert all('2026-08-01' <= row['OrderDate'] <= '2026-10-31' for row in rows)
assert all(isinstance(row['OrderID'], str) and row['OrderID'].startswith('0045') for row in rows)
assert all(isinstance(row['NetAmount'], (int,float)) for row in rows)
assert all(isinstance(row['NetAmountTWD'], (int,float)) for row in rows)
assert d.cell(12,11).value is None  # 0045000132 PaymentTerms blank
assert s['B2'].value == '=SUM(Details!I5:I13)'
calc_total = sum(Decimal(str(row['NetAmountTWD'])).quantize(q) for row in rows)
assert calc_total == Decimal('1702471.00'), calc_total
assert s.max_row == 7
assert [s.cell(rr,3).value for rr in range(5,8)] == ['EUR','TWD','USD']
print('Verification passed: output reopened; sheets retained; Details has 9 rows; all CustomerID exact C001; dates in range; OrderID text; money numeric; blank PaymentTerms preserved for 0045000132; Summary total formula extends to I13; TWD total =', calc_total)
print('Missing:', missing)
print('Output file:', outfile)