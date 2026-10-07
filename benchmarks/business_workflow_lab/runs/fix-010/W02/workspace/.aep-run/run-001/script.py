from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
# workbook structure
xlsx = ws/'input'/'sales_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
        print(' ', r)
# csv sample
csvp = ws/'input'/'fx_rates.csv'
if csvp.exists():
    print('fx_rates.csv sample:')
    with csvp.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i,row in zip(range(8), reader):
            print(' ', row)
# manual sample/full small
md = ws/'input'/'sap_manual.md'
if md.exists():
    txt = md.read_text(encoding='utf-8')
    print('sap_manual.md length', len(txt))
    print(txt[:4000])