from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from copy import copy
from openpyxl import load_workbook

workspace = Path.cwd()
in_dir = workspace / 'input'
out_dir = workspace / 'output'
out_dir.mkdir(exist_ok=True)

# Read source files exported from Expense Viewer (export includes detail fields)
records = []
for name in ['submitted_2026-09.csv', 'approved_2026-09.csv']:
    path = out_dir / name
    with path.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.DictReader(f)
        print(name, 'headers:', reader.fieldnames)
        rows = list(reader)
        print(name, 'rows:', len(rows), 'sample:', rows[:2])
        records.extend(rows)

# Read FX rates (first two columns only; note column intentionally has commas in source)
rates = {}
with (in_dir / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('rates:', rates)

# Validate filter: all dates Sep 2026 and status submitted/approved
allowed_status = {'Submitted','Approved'}
for r in records:
    assert '2026-09-01' <= r['ExpenseDate'] <= '2026-09-30', r
    assert r['Status'] in allowed_status, r
    assert r['Currency'] in rates, r

# Sort by date then ExpenseID for deterministic Details
records.sort(key=lambda r: (r['ExpenseDate'], r['ExpenseID']))

for r in records:
    amt = Decimal(r['Amount'])
    r['AmountTWD'] = (amt * rates[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    missing = str(r.get('MissingReceipt','')).strip().lower() == 'true'
    r['MissingReceiptOut'] = 'true' if missing else 'false'
    r['NeedsReview'] = 'true' if missing else 'false'

# Load template and fill a copy
wb = load_workbook(in_dir / 'expense_template.xlsx')
ws_report = wb['Report']
ws_details = wb['Details']
ws_summary = wb['Summary']
ws_ex = wb['Exceptions']

# helper clear rows from 5 down, preserving row 4/header
def clear_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

clear_data(ws_details)
clear_data(ws_summary)
clear_data(ws_ex)

# Report metadata
now_iso = datetime.now().astimezone().isoformat(timespec='seconds')
ws_report['B3'] = '2026-09-01'
ws_report['D3'] = '2026-09-30'
ws_report['B4'] = 'Status in Submitted, Approved; ExpenseDate 2026-09-01 through 2026-09-30'
ws_report['B5'] = now_iso
ws_report['B7'] = 'Expense Viewer; fx_rates.csv'
ws_report['D7'] = now_iso

# Details rows
cols = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceiptOut','NeedsReview']
for idx, r in enumerate(records, start=5):
    copy_row_style(ws_details, 4, idx, 14)
    values = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['ExpenseDate'],
        r['Category'], r['Currency'], float(Decimal(r['Amount'])), float(r['AmountTWD']), r['Status'],
        r.get('ReceiptID','') or None, r.get('Description','') or None, r['MissingReceiptOut'], r['NeedsReview']
    ]
    for col, val in enumerate(values, start=1):
        cell = ws_details.cell(idx, col, val)
        if col in (1,2,4,7,10,11,12,13,14):
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9):
            cell.number_format = '0.00'

last_detail = 4 + len(records) if records else 5
if records:
    ws_summary['B2'] = f'=SUM(Details!I5:I{last_detail})'
else:
    ws_summary['B2'] = '=SUM(Details!I5:I5)'
ws_summary['B2'].number_format = '0.00'

# Summary by Status + Currency; original Amount only summed within same currency
summary = {}
for r in records:
    key = (r['Status'], '', r['Currency'])
    if key not in summary:
        summary[key] = {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0}
    summary[key]['amount'] += Decimal(r['Amount'])
    summary[key]['twd'] += r['AmountTWD']
    summary[key]['count'] += 1

for idx, (key, vals) in enumerate(sorted(summary.items()), start=5):
    copy_row_style(ws_summary, 4, idx, 6)
    status, sub, cur = key
    rowvals = [status, sub or None, cur, float(vals['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), float(vals['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), vals['count']]
    for col, val in enumerate(rowvals, start=1):
        c = ws_summary.cell(idx, col, val)
        if col in (1,2,3):
            c.number_format = '@'
        elif col in (4,5):
            c.number_format = '0.00'
        elif col == 6:
            c.number_format = '0'

# Exceptions: missing receipts, keep original ExpenseID and do not remove Details rows
exceptions = [r for r in records if r['MissingReceiptOut'] == 'true']
for idx, r in enumerate(exceptions, start=5):
    copy_row_style(ws_ex, 4, idx, 4)
    vals = [r['ExpenseID'], r['EmployeeID'], r['ExpenseDate'], 'Missing receipt']
    for col, val in enumerate(vals, start=1):
        c = ws_ex.cell(idx, col, val)
        if col in (1,2,4):
            c.number_format = '@'
        elif col == 3:
            c.number_format = 'yyyy-mm-dd'

ws_report['B9'] = f'{len(records)} expense rows included. Missing receipt exceptions: {len(exceptions)}; all exception rows remain in Details.'

out_path = out_dir / 'expense_report_2026-09_submitted_approved.xlsx'
wb.save(out_path)
print('wrote', out_path)
print('detail_rows', len(records), 'exception_rows', len(exceptions), 'summary_rows', len(summary))
print('exceptions:', [r['ExpenseID'] for r in exceptions])

# Reopen and verify stated requirements
wb2 = load_workbook(out_path, data_only=False)
d = wb2['Details']; s = wb2['Summary']; e = wb2['Exceptions']; rep = wb2['Report']
headers = [d.cell(4,c).value for c in range(1,15)]
expected_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert headers == expected_headers, headers
# count nonempty detail rows
out_rows = []
for row in range(5, d.max_row+1):
    if d.cell(row,1).value:
        out_rows.append(row)
assert len(out_rows) == 22, len(out_rows)
statuses = {d.cell(r,10).value for r in out_rows}
assert statuses == {'Submitted','Approved'}, statuses
assert all('2026-09-01' <= str(d.cell(r,5).value) <= '2026-09-30' for r in out_rows)
# Detail fields present, missing receipt retained
expids = {d.cell(r,1).value for r in out_rows}
assert 'EXP-00015' in expids and 'EXP-00006' in expids
missing_details = [d.cell(r,1).value for r in out_rows if str(d.cell(r,13).value).lower() == 'true']
assert set(missing_details) == {'EXP-00015','EXP-00006'}, missing_details
assert all(d.cell(r,14).value in ('true','false') for r in out_rows)
# amount types numeric
assert all(isinstance(d.cell(r,8).value, (int,float)) for r in out_rows)
assert all(isinstance(d.cell(r,9).value, (int,float)) for r in out_rows)
# exceptions contain originals, not deleted
ex_rows = [r for r in range(5, e.max_row+1) if e.cell(r,1).value]
ex_ids = [e.cell(r,1).value for r in ex_rows]
assert ex_ids == ['EXP-00006','EXP-00015'], ex_ids
assert all(x in expids for x in ex_ids)
# report metadata and formula range
assert rep['B3'].value == '2026-09-01' and rep['D3'].value == '2026-09-30'
assert s['B2'].value == '=SUM(Details!I5:I26)', s['B2'].value
print('CHECK OK: reopened output; Details headers, 22 rows, Sep 2026 Submitted/Approved only, detail fields, numeric amounts/TWD, missing-receipt exceptions retained, formula range checked.')