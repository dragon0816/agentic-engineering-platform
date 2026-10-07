from pathlib import Path
from openpyxl import load_workbook
p=Path('input/expense_template.xlsx')
wb=load_workbook(p, data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
print('\nExport files:')
for p in Path('output').glob('*.csv'):
    print('---', p)
    print(p.read_text(encoding='utf-8-sig')[:1000])