from pathlib import Path
import csv, json, shutil, re
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from copy import copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

base=Path.cwd()
inp=base/'input'
out=base/'output'
dl=out/'downloads'
out.mkdir(exist_ok=True)

FROM='2026-09-01'
TO='2026-09-30'
CUSTOMER='C001'
GEN='2026-10-06T00:00:00+08:00'

def d(s): return datetime.strptime(s, '%Y-%m-%d').date()
def dec(s): return Decimal(str(s))
def money(x): return Decimal(x).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def rate_table():
    rates={}
    with (inp/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
        # robust parse because Note contains unquoted comma in lab file; first two cols are needed
        for i,row in enumerate(csv.reader(f)):
            if i==0: continue
            if len(row)>=2 and row[0]:
                rates[row[0]]=Decimal(row[1])
    return rates
rates=rate_table()
print('rates', rates)
def twd(amount, currency):
    if currency not in rates:
        raise ValueError(f'Missing FX rate for {currency}')
    return money(dec(amount)*rates[currency])

def read_csv(p):
    with Path(p).open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))
orders_raw=read_csv(dl/'erp_orders_C001_2026-09.csv')
invoices=read_csv(dl/'erp_invoices_C001_2026-09.csv')
orders=[r for r in orders_raw if r['Status']!='Cancelled']
quotes=json.loads((out/'quotes_C001_2026-09.json').read_text(encoding='utf-8'))['items']
# C001 corresponding projects: projects from non-cancelled orders, invoices, and quotes in period
projects=sorted({r['ProjectID'] for r in orders if r.get('ProjectID')} | {r['ProjectID'] for r in invoices if r.get('ProjectID')} | {r['ProjectID'] for r in quotes if r.get('ProjectID')})
print('projects', projects)
expenses=[]
for prj in projects:
    p=out/f'expenses_{prj}_2026-09_Approved.csv'
    if p.exists():
        expenses.extend(read_csv(p))
    else:
        print('missing expense file for', prj)
# enforce approved and project filter
expenses=[r for r in expenses if r.get('Status')=='Approved' and r.get('ProjectID') in projects]

# Output workbook copied from template, preserving sheets/layout
src=inp/'project_template.xlsx'
outfile=out/'project_report_C001_2026-09.xlsx'
shutil.copy2(src, outfile)
wb=load_workbook(outfile)

# helper clear data rows without removing sheets/formulas except data areas
def clear_rows(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row-min_row+1)

def append_row(ws, vals, text_cols=set(), date_cols=set(), money_cols=set()):
    r=ws.max_row+1
    if r < 5: r=5
    for c,val in enumerate(vals, start=1):
        cell=ws.cell(r,c,val)
        # copy style from row 4 header? better copy from row 5 if exists impossible after delete; use simple formats
        if c in text_cols:
            cell.number_format='@'
            if val is not None: cell.value=str(val)
        if c in date_cols:
            cell.number_format='yyyy-mm-dd'
            if isinstance(val, str) and val:
                cell.value=d(val)
        if c in money_cols:
            cell.number_format='0.00'
            if val not in (None,''):
                cell.value=float(val)
    return r

# report fields
ws=wb['Report']
ws['B3']=FROM; ws['D3']=TO
ws['B4']='CustomerID=C001; orders exclude Cancelled; expenses are Approved for C001-related projects; all quotes retained; FX from input/fx_rates.csv only.'
ws['B5']=GEN
ws['B7']='Lab ERP; Expense Viewer; quote_cli; fx_rates.csv'
ws['D7']=GEN
ws['B9']='關注條件：截至 2026-09-30 已逾期未收發票、缺收據費用、ValidUntil < 2026-09-30 且非 Accepted 的報價。'

# Details
for s in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    clear_rows(wb[s],5)

for r in orders:
    append_row(wb['Orders'], [r['OrderID'], r['ProjectID'], r['CustomerID'], r['OrderDate'], r['Status'], r['Currency'], money(r['NetAmount']), twd(r['NetAmount'], r['Currency'])], text_cols={1,2,3,5,6}, date_cols={4}, money_cols={7,8})
for r in invoices:
    bal=money(r['Balance'])
    append_row(wb['Invoices'], [r['InvoiceID'], r['OrderID'], r['ProjectID'], r['CustomerID'], r['InvoiceDate'], r['DueDate'], r['Currency'], money(r['Amount']), money(r['PaidAmount']), bal, twd(bal, r['Currency'])], text_cols={1,2,3,4,7}, date_cols={5,6}, money_cols={8,9,10,11})
for r in expenses:
    append_row(wb['Expenses'], [r['ExpenseID'], r['ProjectID'], r['EmployeeID'], r['ExpenseDate'], r['Status'], r['Currency'], money(r['Amount']), twd(r['Amount'], r['Currency'])], text_cols={1,2,3,5,6}, date_cols={4}, money_cols={7,8})
for r in quotes:
    append_row(wb['Quotes'], [r['QuoteID'], r['ProjectID'], r['CustomerID'], r['QuoteDate'], r['ValidUntil'], r['Status'], r['Currency'], money(r['NetAmount']), twd(r['NetAmount'], r['Currency'])], text_cols={1,2,3,6,7}, date_cols={4,5}, money_cols={8,9})

# Summary per project
from collections import defaultdict
summary={p:{'OrderTWD':Decimal('0'), 'InvoiceBalanceTWD':Decimal('0'), 'ApprovedExpenseTWD':Decimal('0'), 'AcceptedQuoteTWD':Decimal('0')} for p in projects}
for r in orders:
    summary[r['ProjectID']]['OrderTWD'] += twd(r['NetAmount'], r['Currency'])
for r in invoices:
    summary[r['ProjectID']]['InvoiceBalanceTWD'] += twd(r['Balance'], r['Currency'])
for r in expenses:
    summary[r['ProjectID']]['ApprovedExpenseTWD'] += twd(r['Amount'], r['Currency'])
for r in quotes:
    if r['Status']=='Accepted':
        summary[r['ProjectID']]['AcceptedQuoteTWD'] += twd(r['NetAmount'], r['Currency'])
for p in sorted(summary):
    vals=summary[p]
    append_row(wb['Summary'], [p, money(vals['OrderTWD']), money(vals['InvoiceBalanceTWD']), money(vals['ApprovedExpenseTWD']), money(vals['AcceptedQuoteTWD'])], text_cols={1}, money_cols={2,3,4,5})
last=max(5, wb['Summary'].max_row)
for col in range(2,6):
    letter=get_column_letter(col)
    wb['Summary'][f'{letter}2']=f'=SUM({letter}5:{letter}{last})'

# Exceptions
asof=d(TO)
exceptions=[]
for r in invoices:
    if d(r['DueDate']) <= asof and dec(r['Balance']) > 0:
        exceptions.append(['Invoices', r['InvoiceID'], r['ProjectID'], f"Overdue unpaid invoice: DueDate {r['DueDate']}, Balance {r['Balance']} {r['Currency']} as of {TO}"])
for r in expenses:
    mr=str(r.get('MissingReceipt','')).strip().lower()
    receipt=str(r.get('ReceiptID','')).strip()
    if mr in ('true','1','yes','y') or receipt=='':
        exceptions.append(['Expenses', r['ExpenseID'], r['ProjectID'], 'Missing receipt'])
for r in quotes:
    if d(r['ValidUntil']) < asof and r['Status']!='Accepted':
        exceptions.append(['Quotes', r['QuoteID'], r['ProjectID'], f"Quote expired/not accepted: ValidUntil {r['ValidUntil']}, Status {r['Status']}"])
for ex in exceptions:
    append_row(wb['Exceptions'], ex, text_cols={1,2,3,4})

wb.save(outfile)
print('wrote', outfile)
print('counts', {'orders_raw':len(orders_raw),'orders_after_excluding_cancelled':len(orders),'invoices':len(invoices),'expenses_approved':len(expenses),'quotes':len(quotes),'exceptions':len(exceptions)})
print('exceptions')
for ex in exceptions: print(ex)

# Validation reopen data_only False, inspect requirements
wb2=load_workbook(outfile, data_only=False)
checks=[]
def row_count(sheet):
    ws=wb2[sheet]
    return max(0, ws.max_row-4)
checks.append(('Orders rows', row_count('Orders'), len(orders)))
checks.append(('Invoices rows', row_count('Invoices'), len(invoices)))
checks.append(('Expenses rows', row_count('Expenses'), len(expenses)))
checks.append(('Quotes rows', row_count('Quotes'), len(quotes)))
checks.append(('Exceptions rows', row_count('Exceptions'), len(exceptions)))
# IDs are text style / values preserved
for sheet, col in [('Orders',1),('Invoices',1),('Expenses',1),('Quotes',1)]:
    ws=wb2[sheet]
    vals=[ws.cell(r,col).value for r in range(5, ws.max_row+1)]
    print(sheet, 'ID vals', vals)
# no cancelled orders
cancelled=[wb2['Orders'].cell(r,5).value for r in range(5, wb2['Orders'].max_row+1) if wb2['Orders'].cell(r,5).value=='Cancelled']
checks.append(('No cancelled orders', len(cancelled), 0))
# all expenses approved
bad_exp=[wb2['Expenses'].cell(r,5).value for r in range(5, wb2['Expenses'].max_row+1) if wb2['Expenses'].cell(r,5).value!='Approved']
checks.append(('All expenses Approved', len(bad_exp), 0))
# amount twd sample computed numeric
for sheet, cols in [('Orders',(7,8)),('Invoices',(10,11)),('Expenses',(7,8)),('Quotes',(8,9))]:
    ws=wb2[sheet]
    for r in range(5, ws.max_row+1):
        assert isinstance(ws.cell(r, cols[1]).value, (int,float)), f'{sheet} TWD not numeric row {r}'
print('VALIDATION')
for name, got, exp in checks:
    print(name, got, 'expected', exp, 'OK' if got==exp else 'FAIL')
    assert got==exp, name
print('Summary formulas', [wb2['Summary'][f'{c}2'].value for c in ['B','C','D','E']])
