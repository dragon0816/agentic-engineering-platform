from pathlib import Path
from decimal import Decimal
from datetime import datetime, date
import csv, json
from openpyxl import load_workbook

base=Path.cwd()
out=base/'output'
report=out/'project_report_C001_2026-09-01_2026-09-30.xlsx'
fx_path=base/'input'/'fx_rates.csv'
orders_csv=out/'downloads'/'erp_orders_C001_2026-09-01_2026-09-30.csv'
invoices_csv=out/'downloads'/'erp_invoices_C001_2026-09-01_2026-09-30.csv'
quotes_json=out/'quotes_C001_2026-09-01_2026-09-30.json'
expenses_csv=out/'expenses_all_approved_C001_2026-09-01_2026-09-30.csv'

for p in [report, fx_path, orders_csv, invoices_csv, quotes_json, expenses_csv]:
    assert p.exists() and p.stat().st_size>0, f'missing/empty {p}'

fx={}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader=csv.DictReader(f)
    print('fx headers', reader.fieldnames)
    for row in reader:
        cur=(row.get('Currency') or row.get('currency') or '').strip()
        rate=(row.get('TWD') or row.get('RateToTWD') or row.get('rate') or row.get('rate_to_twd') or '').strip()
        if cur and rate:
            fx[cur]=Decimal(rate)
assert fx=={'TWD':Decimal('1'), 'USD':Decimal('32'), 'EUR':Decimal('35')}, fx

wb=load_workbook(report, data_only=False)
assert wb.sheetnames==['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']

def norm_date(v):
    if isinstance(v, datetime): return v.date().isoformat()
    if isinstance(v, date): return v.isoformat()
    return str(v)

def sheet_rows(name):
    sh=wb[name]
    headers=[sh.cell(4,c).value for c in range(1, sh.max_column+1)]
    rows=[]
    for r in range(5, sh.max_row+1):
        vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
        if any(v is not None for v in vals):
            rows.append(dict(zip(headers, vals)))
    return headers, rows

rep=wb['Report']
assert rep['B3'].value=='2026-09-01'
assert rep['D3'].value=='2026-09-30'
assert 'CustomerID=C001' in rep['B4'].value
assert 'orders exclude Cancelled' in rep['B4'].value
assert 'Approved' in rep['B4'].value
assert 'FX from input/fx_rates.csv only' in rep['B4'].value

headers, orders=sheet_rows('Orders')
assert headers==['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD']
assert len(orders)==3
assert all(r['CustomerID']=='C001' for r in orders)
assert all(r['Status']!='Cancelled' for r in orders)
assert {r['OrderID'] for r in orders}=={'0045000109','0045000119','0045000125'}
for r in orders:
    assert isinstance(r['OrderID'], str)
    assert norm_date(r['OrderDate'])>='2026-09-01' and norm_date(r['OrderDate'])<='2026-09-30'
    assert isinstance(r['AmountTWD'], (int,float))
    assert abs(float(r['AmountTWD'])-float(Decimal(str(r['NetAmount']))*fx[r['Currency']]))<0.001

headers, invoices=sheet_rows('Invoices')
assert headers==['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
assert len(invoices)==3
assert all(r['CustomerID']=='C001' for r in invoices)
assert {r['InvoiceID'] for r in invoices}=={'0090000208','0090000214','0090000219'}
for r in invoices:
    assert isinstance(r['InvoiceID'], str) and isinstance(r['OrderID'], str)
    assert isinstance(r['BalanceTWD'], (int,float))
    assert abs(float(r['BalanceTWD'])-float(Decimal(str(r['Balance']))*fx[r['Currency']]))<0.001

headers, expenses=sheet_rows('Expenses')
assert headers==['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD']
assert len(expenses)==9
assert all(r['Status']=='Approved' for r in expenses)
assert {r['ProjectID'] for r in expenses}=={'P001','P003'}
for r in expenses:
    assert isinstance(r['AmountTWD'], (int,float))
    assert abs(float(r['AmountTWD'])-float(Decimal(str(r['Amount']))*fx[r['Currency']]))<0.001

headers, quotes=sheet_rows('Quotes')
assert headers==['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD']
assert len(quotes)==4
assert all(r['CustomerID']=='C001' for r in quotes)
assert {r['QuoteID'] for r in quotes}=={'Q0010','Q0011','Q0014','Q0018'}
for r in quotes:
    assert norm_date(r['QuoteDate'])>='2026-09-01' and norm_date(r['QuoteDate'])<='2026-09-30'
    assert isinstance(r['AmountTWD'], (int,float))
    assert abs(float(r['AmountTWD'])-float(Decimal(str(r['NetAmount']))*fx[r['Currency']]))<0.001

headers, exceptions=sheet_rows('Exceptions')
assert headers==['Source','DocumentID','ProjectID','Reason']
assert len(exceptions)==4
actual={(r['Source'], r['DocumentID'], r['ProjectID']) for r in exceptions}
expected={('Invoices','0090000219','P001'),('Expenses','EXP-00006','P001'),('Quotes','Q0010','P001'),('Quotes','Q0014','P001')}
assert actual==expected, actual

with orders_csv.open(newline='', encoding='utf-8-sig') as f: order_src=list(csv.DictReader(f))
with invoices_csv.open(newline='', encoding='utf-8-sig') as f: inv_src=list(csv.DictReader(f))
with expenses_csv.open(newline='', encoding='utf-8-sig') as f: exp_src=list(csv.DictReader(f))
qdata=json.load(open(quotes_json, encoding='utf-8'))
assert len(order_src)==4 and any(r['Status']=='Cancelled' for r in order_src)
assert len(inv_src)==3
assert len(exp_src)==15 and all(r['Status']=='Approved' for r in exp_src)
assert qdata.get('count')==4 and {r['QuoteID'] for r in qdata.get('items',[])}=={'Q0010','Q0011','Q0014','Q0018'}

files=[report, quotes_json, expenses_csv, orders_csv, invoices_csv]
print('VALIDATION PASSED')
print('Report rows:', {'Orders':len(orders),'Invoices':len(invoices),'Expenses':len(expenses),'Quotes':len(quotes),'Exceptions':len(exceptions)})
print('Exceptions:', exceptions)
print('Output files checked:')
for p in files:
    print('-', p.relative_to(base), p.stat().st_size, 'bytes')