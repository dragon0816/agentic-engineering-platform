from pathlib import Path
import csv, json, shutil, subprocess, sys, os, re
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
customer = params['customer_id']
from_date = params['from_date']
to_date = params['to_date']
as_of_date = params['as_of_date']

base = Path.cwd()
inp = base / 'input'
out = base / 'output'
dl = out / 'downloads'
out.mkdir(exist_ok=True)

print('run parameters', params)
print('input files', sorted(p.name for p in inp.glob('*')))
print('download files', sorted(p.name for p in dl.glob('*')) if dl.exists() else [])

def parse_date(s):
    return datetime.strptime(s, '%Y-%m-%d').date()

def dec(s):
    return Decimal(str(s))

def money(x):
    return Decimal(x).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def safe_part(s):
    return re.sub(r'[^A-Za-z0-9_.-]+', '_', str(s))

rates = {}
with (inp / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    for i, row in enumerate(csv.reader(f)):
        if i == 0:
            continue
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('fx rates from attachment', {k: str(v) for k, v in rates.items()})

def twd(amount, currency):
    if currency not in rates:
        raise ValueError(f'Missing FX rate for {currency}')
    return money(dec(amount) * rates[currency])

def read_csv(path):
    with Path(path).open(newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    print('read', path, 'rows', len(rows), 'headers', list(rows[0].keys()) if rows else [])
    return rows

suffix = f'{safe_part(customer)}_{safe_part(from_date)}_{safe_part(to_date)}'
orders_path = dl / 'erp_orders.csv'
invoices_path = dl / 'erp_invoices.csv'
expenses_path = out / 'expenses_all_approved.csv'
orders_raw = read_csv(orders_path)
invoices = read_csv(invoices_path)
all_approved_expenses = read_csv(expenses_path)

# Preserve run-specific copies so repeated runs do not collide in the final deliverables.
for src, dest in [(orders_path, dl / f'erp_orders_{suffix}.csv'), (invoices_path, dl / f'erp_invoices_{suffix}.csv'), (expenses_path, out / f'expenses_all_approved_{suffix}.csv')]:
    if src.exists():
        shutil.copy2(src, dest)
        print('copied run artifact', src, 'to', dest)

cli = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
cmd = [sys.executable, str(cli), 'quotes', 'list', '--customer', customer, '--from', from_date, '--to', to_date, '--format', 'json']
res = subprocess.run(cmd, text=True, capture_output=True)
print('quote_cli returncode', res.returncode)
if res.stderr:
    print('quote_cli stderr', res.stderr)
if res.returncode != 0:
    raise RuntimeError('quote_cli failed')
quote_data = json.loads(res.stdout)
quotes = quote_data.get('items', [])
quotes_json = out / f'quotes_{suffix}.json'
quotes_json.write_text(json.dumps(quote_data, ensure_ascii=False, indent=2), encoding='utf-8')
print('quotes rows', len(quotes), 'written', quotes_json)

orders = [r for r in orders_raw if r.get('Status') != 'Cancelled']
projects = sorted({r.get('ProjectID','') for r in orders if r.get('ProjectID')} | {r.get('ProjectID','') for r in invoices if r.get('ProjectID')} | {r.get('ProjectID','') for r in quotes if r.get('ProjectID')})
expenses = [r for r in all_approved_expenses if r.get('Status') == 'Approved' and r.get('ProjectID') in projects]
print('C001/customer-related projects', projects)
print('counts before workbook', {'orders_raw': len(orders_raw), 'orders_after_excluding_cancelled': len(orders), 'invoices': len(invoices), 'expenses_approved_for_projects': len(expenses), 'quotes': len(quotes)})

template = inp / 'project_template.xlsx'
outfile = out / f'project_report_{suffix}.xlsx'
shutil.copy2(template, outfile)
wb = load_workbook(outfile)

def clear_rows(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def append_row(ws, vals, text_cols=set(), date_cols=set(), money_cols=set()):
    r = max(ws.max_row + 1, 5)
    for c, val in enumerate(vals, start=1):
        cell = ws.cell(r, c, val)
        if c in text_cols:
            cell.number_format = '@'
            if val is not None:
                cell.value = str(val)
        if c in date_cols:
            cell.number_format = 'yyyy-mm-dd'
            if isinstance(val, str) and val:
                cell.value = parse_date(val)
        if c in money_cols:
            cell.number_format = '0.00'
            if val not in (None, ''):
                cell.value = float(val)
    return r

ws = wb['Report']
ws['B3'] = from_date
ws['D3'] = to_date
ws['B4'] = f'CustomerID={customer}; orders exclude Cancelled; expenses are Approved and filtered to customer-related projects; all quotes retained; FX from input/fx_rates.csv only.'
ws['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
ws['B7'] = 'Lab ERP; Expense Viewer; quote_cli; fx_rates.csv'
ws['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')
ws['B9'] = f'關注條件：截至 {as_of_date} 已逾期未收發票、缺收據費用、ValidUntil < {as_of_date} 且非 Accepted 的報價。'

for sheet in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    clear_rows(wb[sheet], 5)

for r in orders:
    append_row(wb['Orders'], [r['OrderID'], r['ProjectID'], r['CustomerID'], r['OrderDate'], r['Status'], r['Currency'], money(r['NetAmount']), twd(r['NetAmount'], r['Currency'])], text_cols={1,2,3,5,6}, date_cols={4}, money_cols={7,8})
for r in invoices:
    bal = money(r['Balance'])
    append_row(wb['Invoices'], [r['InvoiceID'], r['OrderID'], r['ProjectID'], r['CustomerID'], r['InvoiceDate'], r['DueDate'], r['Currency'], money(r['Amount']), money(r['PaidAmount']), bal, twd(bal, r['Currency'])], text_cols={1,2,3,4,7}, date_cols={5,6}, money_cols={8,9,10,11})
for r in expenses:
    append_row(wb['Expenses'], [r['ExpenseID'], r['ProjectID'], r['EmployeeID'], r['ExpenseDate'], r['Status'], r['Currency'], money(r['Amount']), twd(r['Amount'], r['Currency'])], text_cols={1,2,3,5,6}, date_cols={4}, money_cols={7,8})
for r in quotes:
    append_row(wb['Quotes'], [r['QuoteID'], r['ProjectID'], r['CustomerID'], r['QuoteDate'], r['ValidUntil'], r['Status'], r['Currency'], money(r['NetAmount']), twd(r['NetAmount'], r['Currency'])], text_cols={1,2,3,6,7}, date_cols={4,5}, money_cols={8,9})

summary = {p: {'OrderTWD': Decimal('0'), 'InvoiceBalanceTWD': Decimal('0'), 'ApprovedExpenseTWD': Decimal('0'), 'AcceptedQuoteTWD': Decimal('0')} for p in projects}
for r in orders:
    summary[r['ProjectID']]['OrderTWD'] += twd(r['NetAmount'], r['Currency'])
for r in invoices:
    summary[r['ProjectID']]['InvoiceBalanceTWD'] += twd(r['Balance'], r['Currency'])
for r in expenses:
    summary[r['ProjectID']]['ApprovedExpenseTWD'] += twd(r['Amount'], r['Currency'])
for r in quotes:
    if r.get('Status') == 'Accepted':
        summary[r['ProjectID']]['AcceptedQuoteTWD'] += twd(r['NetAmount'], r['Currency'])
for p in sorted(summary):
    s = summary[p]
    append_row(wb['Summary'], [p, money(s['OrderTWD']), money(s['InvoiceBalanceTWD']), money(s['ApprovedExpenseTWD']), money(s['AcceptedQuoteTWD'])], text_cols={1}, money_cols={2,3,4,5})
last = max(5, wb['Summary'].max_row)
for col in range(2, 6):
    letter = get_column_letter(col)
    wb['Summary'][f'{letter}2'] = f'=SUM({letter}5:{letter}{last})'

as_of = parse_date(as_of_date)
exceptions = []
for r in invoices:
    if parse_date(r['DueDate']) <= as_of and dec(r['Balance']) > 0:
        exceptions.append(['Invoices', r['InvoiceID'], r['ProjectID'], f"Overdue unpaid invoice: DueDate {r['DueDate']}, Balance {r['Balance']} {r['Currency']} as of {as_of_date}"])
for r in expenses:
    missing = str(r.get('MissingReceipt','')).strip().lower() in ('true','1','yes','y') or str(r.get('ReceiptID','')).strip() == ''
    if missing:
        exceptions.append(['Expenses', r['ExpenseID'], r['ProjectID'], 'Missing receipt'])
for r in quotes:
    if parse_date(r['ValidUntil']) < as_of and r.get('Status') != 'Accepted':
        exceptions.append(['Quotes', r['QuoteID'], r['ProjectID'], f"Quote expired/not accepted: ValidUntil {r['ValidUntil']}, Status {r['Status']}"])
for ex in exceptions:
    append_row(wb['Exceptions'], ex, text_cols={1,2,3,4})

wb.save(outfile)
print('wrote workbook', outfile)
print('exceptions', exceptions)

wb2 = load_workbook(outfile, data_only=False)
def row_count(sheet):
    return max(0, wb2[sheet].max_row - 4)
checks = [('Orders rows', row_count('Orders'), len(orders)), ('Invoices rows', row_count('Invoices'), len(invoices)), ('Expenses rows', row_count('Expenses'), len(expenses)), ('Quotes rows', row_count('Quotes'), len(quotes)), ('Exceptions rows', row_count('Exceptions'), len(exceptions))]
for name, got, exp in checks:
    print(name, got, 'expected', exp)
    if got != exp:
        raise AssertionError(name)
if any(wb2['Orders'].cell(r,5).value == 'Cancelled' for r in range(5, wb2['Orders'].max_row+1)):
    raise AssertionError('Cancelled order included')
if any(wb2['Expenses'].cell(r,5).value != 'Approved' for r in range(5, wb2['Expenses'].max_row+1)):
    raise AssertionError('Non-approved expense included')
for sheet, twd_col in [('Orders',8),('Invoices',11),('Expenses',8),('Quotes',9)]:
    for r in range(5, wb2[sheet].max_row+1):
        if not isinstance(wb2[sheet].cell(r,twd_col).value, (int,float)):
            raise AssertionError(f'{sheet} row {r} TWD value is not numeric')
print('validated workbook', outfile.name)
