from pathlib import Path
from decimal import Decimal
import csv, json
from openpyxl import load_workbook

base=Path.cwd()
out=base/'output'
report=out/'project_report_C001_2026-09-01_2026-09-30.xlsx'
fx_path=base/'input'/'fx_rates.csv'
orders_csv=out/'downloads'/'erp_orders_C001_2026-09-01_2026-09-30.csv'
invoices_csv=out/'downloads'/'erp_invoices_C001_2026-09-01_2026-09-30.csv'
quotes_json=out/'quotes_C001_2026-09-01_2026-09-30.json'
expenses_csv=out/'expenses_all_approved_C001_2026-09-01_2026-09-30.csv'

required=[report, fx_path, orders_csv, invoices_csv, quotes_json, expenses_csv]
for p in required:
    assert p.exists(), f'missing {p}'

# load fx rates, support common header names
fx={}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        cur=(row.get('Currency') or row.get('currency') or '').strip()
        rate=(row.get('TWD') or row.get('RateToTWD') or row.get('rate') or row.get('rate_to_twd') or '').strip()
        if cur and rate:
            fx[cur]=Decimal(rate)
assert fx=={'TWD':Decimal('1'), 'USD':Decimal('32'), 'EUR':Decimal('35')}, fx

wb=load_workbook(report, data_only=False)
expected_sheets=['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
assert wb.sheetnames==expected_sheets, wb.sheetnames

# helper read data from row 5 with headers at row 4
def sheet_rows(name):
    sh=wb[name]
    headers=[sh.cell(4,c).value for c in range(1, sh.max_column+1)]
    rows=[]
    for r in range(5, sh.max_row+1):
        vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
        if any(v is not None for v in vals):
            rows.append(dict(zip(headers, vals)))
    return headers, rows

checks=[]
# report period/condition
rep=wb['Report']
assert rep['B3'].value=='2026-09-01'
assert rep['D3'].value=='2026-09-30'
assert 'CustomerID=C001' in (rep['B4'].value or '')
assert 'FX from input/fx_rates.csv only' in (rep['B4'].value or '')
checks.append('Report sheet has customer C001, date range, and attachment FX note')

headers, orders=sheet_rows('Orders')
assert headers==['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD']
assert len(orders)==3, len(orders)
assert all(r['CustomerID']=='C001' for r in orders)
assert all(r['Status']!='Cancelled' for r in orders)
assert {r['OrderID'] for r in orders}=={'0045000109','0045000119','0045000125'}
for r in orders:
    assert isinstance(r['OrderID'], str) and r['OrderID'].startswith('00')
    expected=float(Decimal(str(r['NetAmount']))*fx[r['Currency']])
    assert abs(float(r['AmountTWD'])-expected)<0.001, (r, expected)
checks.append('Orders: 3 rows, C001 only, Cancelled excluded, IDs text, TWD amounts numeric and converted by input FX')

headers, invoices=sheet_rows('Invoices')
assert headers==['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
assert len(invoices)==3
assert all(r['CustomerID']=='C001' for r in invoices)
assert {r['InvoiceID'] for r in invoices}=={'0090000208','0090000214','0090000219'}
for r in invoices:
    assert isinstance(r['InvoiceID'], str) and r['InvoiceID'].startswith('00')
    expected=float(Decimal(str(r['Balance']))*fx[r['Currency']])
    assert abs(float(r['BalanceTWD'])-expected)<0.001, (r, expected)
checks.append('Invoices: 3 rows, C001 only, IDs text, BalanceTWD numeric and converted by input FX')

headers, expenses=sheet_rows('Expenses')
assert headers==['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD']
assert len(expenses)==9, len(expenses)
assert all(r['Status']=='Approved' for r in expenses)
assert {r['ProjectID'] for r in expenses}=={'P001','P003'}
for r in expenses:
    expected=float(Decimal(str(r['Amount']))*fx[r['Currency']])
    assert abs(float(r['AmountTWD'])-expected)<0.001, (r, expected)
checks.append('Expenses: 9 rows, Approved only, limited to C001-related projects P001/P003, AmountTWD numeric and converted by input FX')

headers, quotes=sheet_rows('Quotes')
assert headers==['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD']
assert len(quotes)==4, len(quotes)
assert all(r['CustomerID']=='C001' for r in quotes)
assert {r['QuoteID'] for r in quotes}=={'Q0010','Q0011','Q0014','Q0015'}
for r in quotes:
    expected=float(Decimal(str(r['NetAmount']))*fx[r['Currency']])
    assert abs(float(r['AmountTWD'])-expected)<0.001, (r, expected)
checks.append('Quotes: 4 rows, all C001 quotes retained, AmountTWD numeric and converted by input FX')

headers, exceptions=sheet_rows('Exceptions')
assert headers==['Source','DocumentID','ProjectID','Reason']
assert len(exceptions)==4, len(exceptions)
expected_ex={
    ('Invoices','0090000219','P001'),
    ('Expenses','EXP-00006','P001'),
    ('Quotes','Q0010','P001'),
    ('Quotes','Q0014','P001'),
}
actual_ex={(r['Source'], r['DocumentID'], r['ProjectID']) for r in exceptions}
assert actual_ex==expected_ex, actual_ex
assert any('overdue' in r['Reason'].lower() or '逾期' in r['Reason'] for r in exceptions if r['DocumentID']=='0090000219')
assert any('missing' in r['Reason'].lower() or '缺' in r['Reason'] for r in exceptions if r['DocumentID']=='EXP-00006')
assert all(('expired' in r['Reason'].lower() or 'ValidUntil' in r['Reason']) for r in exceptions if r['Source']=='Quotes')
checks.append('Exceptions: exactly overdue unpaid invoice, missing-receipt expense, and expired non-Accepted quotes as of 2026-09-30')

# Source exports counts
with orders_csv.open(newline='', encoding='utf-8-sig') as f: order_src=list(csv.DictReader(f))
with invoices_csv.open(newline='', encoding='utf-8-sig') as f: inv_src=list(csv.DictReader(f))
with expenses_csv.open(newline='', encoding='utf-8-sig') as f: exp_src=list(csv.DictReader(f))
qdata=json.load(open(quotes_json, encoding='utf-8'))
assert len(order_src)==4 # raw includes Cancelled, report excludes it
assert len(inv_src)==3
assert len(exp_src)==15 # all approved export; report filters to C001 projects
assert qdata.get('count')==4 and len(qdata.get('items',[]))==4
checks.append('Workflow source extracts present: orders raw 4, invoices 3, approved expenses raw 15, quotes 4')

# Output file list relevant to workflow run
files=[
    out/'project_report_C001_2026-09-01_2026-09-30.xlsx',
    out/'quotes_C001_2026-09-01_2026-09-30.json',
    out/'expenses_all_approved_C001_2026-09-01_2026-09-30.csv',
    out/'downloads'/'erp_orders_C001_2026-09-01_2026-09-30.csv',
    out/'downloads'/'erp_invoices_C001_2026-09-01_2026-09-30.csv',
]
for p in files:
    assert p.exists() and p.stat().st_size>0, p

print('VALIDATION PASSED')
for c in checks:
    print('-', c)
print('workflow output files:')
for p in files:
    print('-', p.relative_to(base), p.stat().st_size, 'bytes')
print('report row counts:', {'Orders':len(orders),'Invoices':len(invoices),'Expenses':len(expenses),'Quotes':len(quotes),'Exceptions':len(exceptions)})
print('exceptions:', exceptions)