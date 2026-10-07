from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r): print(row)
print('\nproject_template.xlsx structure:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('SHEET', sh, 'max', wsht.max_row, wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
        print(row)
    print('---')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    print('\n===== ', name, ' =====')
    text=(ws/'input'/name).read_text(encoding='utf-8')
    print(text[:4000])