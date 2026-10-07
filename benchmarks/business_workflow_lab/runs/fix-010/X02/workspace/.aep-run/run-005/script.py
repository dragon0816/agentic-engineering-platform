from pathlib import Path
from openpyxl import load_workbook
base=Path.cwd()
wb=load_workbook(base/'input'/'project_template.xlsx', data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    vals=[v for v in row if v is not None]
    if vals:
        print(vals)
