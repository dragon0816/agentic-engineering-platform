import os, csv, json
from pathlib import Path
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace', ws)
for sub in ['input','output','output/downloads']:
    p = ws / sub
    print('\n##', sub, 'exists=', p.exists())
    if p.exists():
        for f in sorted(p.iterdir()):
            if f.is_file():
                print(f.name, f.stat().st_size, 'bytes')

# print headers and sample rows for relevant csv/json/xlsx outputs
out = ws/'output'
print('\n## output csv/json samples')
for f in sorted(out.glob('*.csv')) + sorted((out/'downloads').glob('*.csv')) + sorted(out.glob('*.json')):
    print('\nFILE', f.relative_to(ws))
    try:
        if f.suffix.lower()=='.csv':
            with f.open(newline='', encoding='utf-8-sig') as fh:
                rows=list(csv.reader(fh))
            print('rows incl header', len(rows))
            print('header', rows[0] if rows else [])
            for r in rows[1:4]: print('sample', r)
        elif f.suffix.lower()=='.json':
            data=json.load(open(f, encoding='utf-8'))
            print('type', type(data).__name__, 'len', len(data) if hasattr(data,'__len__') else 'n/a')
            if isinstance(data, list):
                for item in data[:3]: print('sample', item)
            elif isinstance(data, dict):
                print('keys', list(data)[:10])
    except Exception as e:
        print('ERR', type(e).__name__, str(e))

print('\n## workbook summaries')
for f in sorted(out.glob('*.xlsx')):
    print('\nWORKBOOK', f.relative_to(ws), f.stat().st_size, 'bytes')
    wb=load_workbook(f, data_only=False)
    print('sheets', wb.sheetnames)
    for s in wb.sheetnames:
        sh=wb[s]
        print(s, 'max_row', sh.max_row, 'max_col', sh.max_column)
        header=[sh.cell(1,c).value for c in range(1, sh.max_column+1)]
        print('header', header)
        for r in range(2, min(sh.max_row,4)+1):
            print('row', r, [sh.cell(r,c).value for c in range(1, sh.max_column+1)])