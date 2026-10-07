from pathlib import Path
import csv
from openpyxl import load_workbook

base = Path.cwd()
print('WORKSPACE', base)
print('INPUT FILES')
for p in sorted((base/'input').glob('*')):
    print('-', p.name, p.stat().st_size)

# fx headers and sample
fx = base/'input'/'fx_rates.csv'
print('\nFX RATES')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i,row in zip(range(6), r):
        print(i, row)

# workbook sheets, dimensions, non-empty cells sample
wb_path = base/'input'/'project_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('\nWORKBOOK SHEETS')
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column)
    rows=[]
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=False):
        vals=[c.value for c in row[:min(ws.max_column,10)]]
        if any(v is not None for v in vals):
            rows.append(vals)
    for vals in rows[:8]:
        print(' ', vals)

# manuals first chunks
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p = base/'input'/name
    print(f'\n--- {name} ---')
    print(p.read_text(encoding='utf-8')[:4000])
