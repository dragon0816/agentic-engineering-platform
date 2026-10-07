from pathlib import Path
import json
from openpyxl import load_workbook
base=Path.cwd()
report=base/'output'/'project_report_C001_2026-09-01_2026-09-30.xlsx'
quotes=base/'output'/'quotes_C001_2026-09-01_2026-09-30.json'
print('json exists', quotes.exists(), quotes.stat().st_size if quotes.exists() else None)
if quotes.exists():
    data=json.load(open(quotes, encoding='utf-8'))
    print('json keys', data.keys())
    print('count', data.get('count'), 'items len', len(data.get('items',[])))
    for item in data.get('items',[]): print(item)
wb=load_workbook(report, data_only=False)
sh=wb['Quotes']
headers=[sh.cell(4,c).value for c in range(1, sh.max_column+1)]
print('sheet headers', headers)
for r in range(5, sh.max_row+1):
    vals=[sh.cell(r,c).value for c in range(1, sh.max_column+1)]
    if any(v is not None for v in vals): print(vals)

ex=wb['Exceptions']
headers=[ex.cell(4,c).value for c in range(1, ex.max_column+1)]
print('exceptions')
for r in range(5, ex.max_row+1):
    vals=[ex.cell(r,c).value for c in range(1, ex.max_column+1)]
    if any(v is not None for v in vals): print(dict(zip(headers, vals)))