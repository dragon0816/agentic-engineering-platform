from pathlib import Path
import csv, json, os
from openpyxl import load_workbook

ws = Path.cwd()
print('Workspace:', ws)
for sub in ['input','output','output/downloads']:
    p = ws / sub
    print('\n##', sub, 'exists=', p.exists())
    if p.exists():
        for f in sorted(p.iterdir()):
            if f.is_file():
                print('FILE', f.name, f.stat().st_size)

print('\n## input headers/samples')
for name in ['fx_rates.csv']:
    p = ws/'input'/name
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as fh:
            r = csv.reader(fh)
            rows = [row for _, row in zip(range(4), r)]
        print(name, rows)

template = ws/'input'/'project_template.xlsx'
if template.exists():
    wb = load_workbook(template, data_only=False)
    print('template sheets:', wb.sheetnames)
    for sh in wb.sheetnames:
        s=wb[sh]
        print('template sheet', sh, s.max_row, s.max_column, [s.cell(1,c).value for c in range(1, min(s.max_column,10)+1)])

print('\n## output CSV headers/samples')
for p in sorted((ws/'output').glob('*.csv')) + sorted((ws/'output'/'downloads').glob('*.csv')):
    with p.open(newline='', encoding='utf-8-sig') as fh:
        reader=csv.reader(fh)
        rows=[row for _, row in zip(range(5), reader)]
    print(p.relative_to(ws), rows)

print('\n## output workbook summaries')
for p in sorted((ws/'output').glob('*.xlsx')):
    print('WORKBOOK', p.name, p.stat().st_size)
    wb = load_workbook(p, data_only=False)
    print(' sheets:', wb.sheetnames)
    for sh in wb.sheetnames:
        s=wb[sh]
        headers=[s.cell(1,c).value for c in range(1, s.max_column+1)] if s.max_row else []
        print(' sheet', sh, 'rows', s.max_row, 'cols', s.max_column, 'headers', headers)
        for r in range(1, min(s.max_row,6)+1):
            print('  row', r, [s.cell(r,c).value for c in range(1, min(s.max_column,12)+1)])
