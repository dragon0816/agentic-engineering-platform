import os, json, csv, shutil, subprocess, sys
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from openpyxl import load_workbook

print('Inspecting workflow inputs and files...')
params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
date_from = params['date_from']
date_to = params['date_to']
period_label = params['period_label']
base = Path('.')
inp = base/'input'
out = base/'output'
downloads = out/'downloads'
out.mkdir(exist_ok=True)
print('Run values:', json.dumps(params, ensure_ascii=False))
print('Input files:', sorted(p.name for p in inp.iterdir() if p.is_file()))
print('Download files:', sorted(p.name for p in downloads.iterdir()) if downloads.exists() else [])

TWOPL = Decimal('0.01')
def money(x):
    if x is None or str(x).strip() == '':
        return Decimal('0')
    return Decimal(str(x).replace(',', '').strip())
def q2(d):
    return Decimal(d).quantize(TWOPL, rounding=ROUND_HALF_UP)
def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    print(f'{path}: headers={list(rows[0].keys()) if rows else []}, rows={len(rows)}, sample={rows[:2]}')
    return rows
def require_file(path):
    if not path.exists():
        raise FileNotFoundError(str(path))
    return path

fx = {}
with open(inp/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rdr = csv.DictReader(f)
    print('fx headers:', rdr.fieldnames)
    for row in rdr:
        keys = list(row.keys())
        cur = (row.get('Currency') or row.get(keys[0]) or '').strip()
        rate = (row.get('TWD') or row.get('Rate') or row.get(keys[1]) or '').strip() if len(keys) > 1 else ''
        if cur and rate:
            fx[cur] = Decimal(rate)
if not fx:
    raise RuntimeError('No FX rates loaded')
print('FX loaded:', {k: str(v) for k,v in fx.items()})

orders = []
for status in ['open', 'closed']:
    p = require_file(downloads/f'sap_orders_{status}.csv')
    for r in read_csv(p):
        if r.get('Status') != 'Cancelled':
            orders.append(r)
orders = [r for r in orders if date_from <= r.get('OrderDate','') <= date_to and r.get('Status') != 'Cancelled']
invoices = [r for r in read_csv(require_file(downloads/'sap_invoices.csv')) if date_from <= r.get('InvoiceDate','') <= date_to]
expenses = [r for r in read_csv(require_file(out/'expense_approved.csv')) if date_from <= r.get('ExpenseDate','') <= date_to and r.get('Status') == 'Approved']

cli_candidates = [Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py'), base/'quote_cli_mock'/'run_quote_cli.py', base.parent.parent/'quote_cli_mock'/'run_quote_cli.py']
cli = next((p for p in cli_candidates if p.exists()), None)
if cli is None:
    raise FileNotFoundError('quote CLI not found in candidates: ' + ', '.join(str(p) for p in cli_candidates))
cmd = [sys.executable, str(cli), 'quotes', 'list', '--from', date_from, '--to', date_to, '--status', 'Accepted', '--format', 'json']
print('Running quote CLI:', ' '.join(cmd))
proc = subprocess.run(cmd, text=True, capture_output=True)
print('quote CLI returncode:', proc.returncode)
if proc.stderr:
    print('quote CLI stderr:', proc.stderr)
if proc.returncode != 0:
    raise RuntimeError(f'quote_cli failed with exit {proc.returncode}: {proc.stderr}')
qdata = json.loads(proc.stdout)
quotes = [r for r in qdata.get('items', qdata if isinstance(qdata, list) else []) if date_from <= r.get('QuoteDate','') <= date_to and r.get('Status') == 'Accepted']
print('Quote rows:', len(quotes), 'sample:', quotes[:2])

exceptions = []
def twd(row, amount_field, docid, source):
    cur = (row.get('Currency') or '').strip()
    amt = money(row.get(amount_field))
    pid = (row.get('ProjectID') or '').strip()
    if not pid:
        exceptions.append({'Source': source, 'DocumentID': str(docid or ''), 'ProjectID': '', 'Reason': 'Missing ProjectID; not included in project summary'})
    if cur not in fx:
        exceptions.append({'Source': source, 'DocumentID': str(docid or ''), 'ProjectID': pid, 'Reason': f'Missing FX rate for {cur}'})
        return None
    return q2(amt * fx[cur])

outfile = out/f'project_integration_{period_label}.xlsx'
shutil.copyfile(inp/'project_template.xlsx', outfile)
wb = load_workbook(outfile)

def clear_from(ws, start=5):
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)

def write_rows(wsname, headers, rows):
    ws = wb[wsname]
    clear_from(ws, 5)
    for idx, row in enumerate(rows, start=5):
        for col, h in enumerate(headers, start=1):
            c = ws.cell(idx, col)
            val = row.get(h, '')
            if h.endswith('TWD') or h in ['NetAmount','Amount','PaidAmount','Balance']:
                c.value = float(val) if val not in ('', None) else None
                c.number_format = '0.00'
            else:
                c.value = str(val) if val is not None else ''
                if h.endswith('ID') or h in ['OrderID','InvoiceID','ExpenseID','QuoteID','CustomerID','EmployeeID','ProjectID']:
                    c.number_format = '@'
                if 'Date' in h or h == 'ValidUntil':
                    c.number_format = 'yyyy-mm-dd'

order_rows=[]
for r in orders:
    amount_twd = twd(r, 'NetAmount', r.get('OrderID',''), 'Sales Orders')
    order_rows.append({'OrderID':r.get('OrderID',''), 'ProjectID':r.get('ProjectID',''), 'CustomerID':r.get('CustomerID',''), 'OrderDate':r.get('OrderDate',''), 'Status':r.get('Status',''), 'Currency':r.get('Currency',''), 'NetAmount':q2(money(r.get('NetAmount'))), 'AmountTWD':amount_twd if amount_twd is not None else ''})
invoice_rows=[]
for r in invoices:
    balance_twd = twd(r, 'Balance', r.get('InvoiceID',''), 'Invoices')
    invoice_rows.append({'InvoiceID':r.get('InvoiceID',''), 'OrderID':r.get('OrderID',''), 'ProjectID':r.get('ProjectID',''), 'CustomerID':r.get('CustomerID',''), 'InvoiceDate':r.get('InvoiceDate',''), 'DueDate':r.get('DueDate',''), 'Currency':r.get('Currency',''), 'Amount':q2(money(r.get('Amount'))), 'PaidAmount':q2(money(r.get('PaidAmount'))), 'Balance':q2(money(r.get('Balance'))), 'BalanceTWD':balance_twd if balance_twd is not None else ''})
expense_rows=[]
for r in expenses:
    amount_twd = twd(r, 'Amount', r.get('ExpenseID',''), 'Expenses')
    expense_rows.append({'ExpenseID':r.get('ExpenseID',''), 'ProjectID':r.get('ProjectID',''), 'EmployeeID':r.get('EmployeeID',''), 'ExpenseDate':r.get('ExpenseDate',''), 'Status':r.get('Status',''), 'Currency':r.get('Currency',''), 'Amount':q2(money(r.get('Amount'))), 'AmountTWD':amount_twd if amount_twd is not None else ''})
quote_rows=[]
for r in quotes:
    amount_twd = twd(r, 'NetAmount', r.get('QuoteID',''), 'Quotes')
    quote_rows.append({'QuoteID':r.get('QuoteID',''), 'ProjectID':r.get('ProjectID',''), 'CustomerID':r.get('CustomerID',''), 'QuoteDate':r.get('QuoteDate',''), 'ValidUntil':r.get('ValidUntil',''), 'Status':r.get('Status',''), 'Currency':r.get('Currency',''), 'NetAmount':q2(money(r.get('NetAmount'))), 'AmountTWD':amount_twd if amount_twd is not None else ''})

write_rows('Orders', ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'], order_rows)
write_rows('Invoices', ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], invoice_rows)
write_rows('Expenses', ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'], expense_rows)
write_rows('Quotes', ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'], quote_rows)

summary = {}
def add(pid, key, val):
    if not pid or val in ('', None):
        return
    summary.setdefault(pid, {'ProjectID':pid, 'OrderTWD':Decimal('0'), 'InvoiceBalanceTWD':Decimal('0'), 'ApprovedExpenseTWD':Decimal('0'), 'AcceptedQuoteTWD':Decimal('0')})
    summary[pid][key] += Decimal(str(val))
for r in order_rows: add(r['ProjectID'], 'OrderTWD', r['AmountTWD'])
for r in invoice_rows: add(r['ProjectID'], 'InvoiceBalanceTWD', r['BalanceTWD'])
for r in expense_rows: add(r['ProjectID'], 'ApprovedExpenseTWD', r['AmountTWD'])
for r in quote_rows: add(r['ProjectID'], 'AcceptedQuoteTWD', r['AmountTWD'])
summary_rows = []
for pid in sorted(summary):
    d = summary[pid]
    summary_rows.append({'ProjectID':pid, 'OrderTWD':q2(d['OrderTWD']), 'InvoiceBalanceTWD':q2(d['InvoiceBalanceTWD']), 'ApprovedExpenseTWD':q2(d['ApprovedExpenseTWD']), 'AcceptedQuoteTWD':q2(d['AcceptedQuoteTWD'])})
write_rows('Summary', ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], summary_rows)
last = max(5, 4 + len(summary_rows))
for col in range(2,6):
    letter = chr(ord('A') + col - 1)
    wb['Summary'].cell(2, col).value = f'=SUM({letter}5:{letter}{last})'

write_rows('Exceptions', ['Source','DocumentID','ProjectID','Reason'], exceptions)

ws = wb['Report']
ws['B3'] = date_from
ws['D3'] = date_to
ws['B4'] = 'Sales Orders Status Open/Closed only (non-Cancelled); Invoices by InvoiceDate; Expenses Approved; Quotes Accepted; grouped by ProjectID.'
ws['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
ws['B7'] = 'Lab ERP SAP-like web; Expense Viewer; quote_cli; fx_rates.csv'
ws['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')
ws['B9'] = f'Orders and accepted quotes are reported separately and are not combined as revenue. Join/group key is ProjectID only; customer name is not used. Exceptions: {len(exceptions)}.'

wb.save(outfile)

check = load_workbook(outfile, data_only=False)
expected_headers = {
 'Orders':['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
 'Invoices':['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
 'Expenses':['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
 'Quotes':['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
 'Summary':['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
 'Exceptions':['Source','DocumentID','ProjectID','Reason'],
}
for s, heads in expected_headers.items():
    got = [check[s].cell(4,c).value for c in range(1,len(heads)+1)]
    assert got == heads, (s, got)
assert all(r['Status'] != 'Cancelled' for r in order_rows)
assert all(r['Status'] == 'Approved' for r in expense_rows)
assert all(r['Status'] == 'Accepted' for r in quote_rows)
for s in ['Orders','Invoices','Expenses','Quotes','Summary']:
    ws = check[s]
    for row in range(5, ws.max_row+1):
        if ws.cell(row,1).value is not None:
            assert isinstance(ws.cell(row,1).value, str), (s, row, ws.cell(row,1).value)
print(json.dumps({
    'output': str(outfile),
    'counts': {'orders_non_cancelled':len(order_rows), 'invoices':len(invoice_rows), 'approved_expenses':len(expense_rows), 'accepted_quotes':len(quote_rows), 'summary_projects':len(summary_rows), 'exceptions':len(exceptions)},
    'checks': ['template copy filled', 'headers preserved', 'ID cells text', 'orders exclude Cancelled', 'expenses Approved only', 'quotes Accepted only', 'separate summary columns by ProjectID', 'FX applied from fx_rates.csv']
}, ensure_ascii=False, indent=2))