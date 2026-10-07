from pathlib import Path
import csv, json
from openpyxl import load_workbook
base=Path('.')
print('Workspace:', base.resolve())
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print(' -', p, 'size=', p.stat().st_size)
print('\nMarkdown/manual previews:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=base/'input'/name
    print(f'\n--- {name} ---')
    print(p.read_text(encoding='utf-8')[:4000])
print('\nfx_rates.csv preview:')
with open(base/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(10),rdr):
        print(row)
print('\nproject_template.xlsx structure:')
wb=load_workbook(base/'input'/'project_template.xlsx', data_only=False)
print('sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('sheet', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,12), values_only=True):
        print(r)
    print('---')