from pathlib import Path
import csv
from openpyxl import load_workbook

base = Path('.')
print('CWD:', base.resolve())
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print(' -', p, 'size=', p.stat().st_size)

fx = base/'input'/'fx_rates.csv'
print('\nfx_rates.csv preview:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i, row in zip(range(8), r):
        print(i, row)

manual = base/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md first 120 lines:')
text = manual.read_text(encoding='utf-8')
for i, line in enumerate(text.splitlines()[:120], 1):
    print(f'{i:03d}: {line}')

wb_path = base/'input'/'quote_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print(f'\nSheet {ws.title}: max_row={ws.max_row}, max_col={ws.max_column}')
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 12), values_only=True):
        print(row)
