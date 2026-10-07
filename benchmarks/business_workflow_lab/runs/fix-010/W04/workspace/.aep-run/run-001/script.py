from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\n--- sap_manual.md head ---')
manual=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(manual[:4000])
print('\n--- fx_rates.csv ---')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(20), reader):
        print(row)
print('\n--- workbook structure ---')
wb=load_workbook(ws/'input'/'sales_template.xlsx')
print('sheets', wb.sheetnames)
for sh in wb.sheetnames:
    s=wb[sh]
    print('\nSHEET', sh, 'max', s.max_row, s.max_column)
    for r in s.iter_rows(min_row=1, max_row=min(s.max_row,12), values_only=True):
        print(r)