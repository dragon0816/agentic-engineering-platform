from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
import csv

p = Path('output/sales_orders_open_2026-09_filled.xlsx')
print('output exists', p.exists(), p, 'bytes', p.stat().st_size if p.exists() else None)
wb = load_workbook(p, data_only=False)
print('sheets', wb.sheetnames)
assert wb.sheetnames == ['Report', 'Details', 'Summary', 'Instructions']

# Check report fields
report = wb['Report']
checks = {
    'B3': report['B3'].value,
    'D3': report['D3'].value,
    'B4': report['B4'].value,
    'B7': report['B7'].value,
    'D7': report['D7'].value,
    'B9': report['B9'].value,
}
print('Report fields:', checks)
assert report['B3'].value == '2026-09-01'
assert report['D3'].value == '2026-09-30'
assert 'Open' in report['B4'].value

# Check details headers and rows
ws = wb['Details']
headers = [ws.cell(4,c).value for c in range(1,12)]
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
print('headers', headers)
assert headers == expected_headers
rows=[]
for r in range(5, ws.max_row+1):
    vals=[ws.cell(r,c).value for c in range(1,12)]
    if any(v is not None for v in vals):
        rows.append(vals)
print('detail row count', len(rows))
assert len(rows) == 14
ids=[r[0] for r in rows]
print('ids', ids)
assert len(ids)==len(set(ids))
assert all(isinstance(x,str) and x.startswith('0045') for x in ids)
assert all(r[5]=='Open' and '2026-09-01' <= r[4] <= '2026-09-30' for r in rows)
assert all(isinstance(r[7], (int,float)) and isinstance(r[8], (int,float)) for r in rows)
assert all(ws.cell(i,1).number_format == '@' for i in range(5, 19))
assert all(ws.cell(i,8).number_format == '0.00' and ws.cell(i,9).number_format == '0.00' for i in range(5, 19))
# blanks should remain empty, not N/A/0
for r in rows:
    if r[0] in ('0045000113','0045000124'):
        assert r[9] is None

# Cross-check against CSV order IDs
with open('output/downloads/open_orders_2026-09.csv', newline='', encoding='utf-8-sig') as f:
    csv_rows=list(csv.DictReader(f))
assert [r['OrderID'] for r in csv_rows] == ids

# Check AmountTWD values with fx
fx={'TWD':Decimal('1'), 'USD':Decimal('32'), 'EUR':Decimal('35')}
for excel_row in rows:
    cur=excel_row[6]
    net=Decimal(str(excel_row[7]))
    twd=Decimal(str(excel_row[8])).quantize(Decimal('0.01'))
    expected=(net*fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert twd == expected, (excel_row[0], twd, expected)

# Check summary formula and group totals
summary=wb['Summary']
print('Summary B2 formula', summary['B2'].value)
assert summary['B2'].value == '=SUM(Details!I5:I18)'
sum_rows=[]
for r in range(5, summary.max_row+1):
    vals=[summary.cell(r,c).value for c in range(1,7)]
    if any(v is not None for v in vals): sum_rows.append(vals)
print('summary rows', sum_rows)
assert len(sum_rows)==3
print('validated total TWD', sum(Decimal(str(r[8])).quantize(Decimal('0.01')) for r in rows))
print('VALIDATION OK')