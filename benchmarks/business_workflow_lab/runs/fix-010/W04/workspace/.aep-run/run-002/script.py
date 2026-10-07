from pathlib import Path
import csv
p=Path('output/downloads/open_orders_2026-09.csv')
print('exists', p.exists(), p)
with open(p, newline='', encoding='utf-8-sig') as f:
    rows=list(csv.DictReader(f))
print('headers', rows[0].keys() if rows else [])
print('count', len(rows))
print('OrderIDs', [r['OrderID'] for r in rows])
print('rows:')
for r in rows:
    print(r)

from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
ws=wb['Instructions']
print('\nInstructions all rows:')
for row in ws.iter_rows(values_only=True):
    print(row)