from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from copy import copy
import csv, json

base=Path.cwd()
input_dir=base/'input'
output_dir=base/'output'
output_dir.mkdir(exist_ok=True)
template=input_dir/'expense_template.xlsx'
export=output_dir/'expense_export_E001_2026-09.xlsx'
out=output_dir/'E001_september_expenses.xlsx'

# print actual inputs
print('Inputs:')
for p in sorted(input_dir.glob('*')): print('-', p.name, p.stat().st_size)
print('Export:', export, export.exists(), export.stat().st_size if export.exists() else None)

rates={}
with (input_dir/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader=csv.DictReader(f)
    print('FX headers:', reader.fieldnames)
    for row in reader:
        if row.get('Currency'):
            rates[row['Currency']]=Decimal(str(row['RateToTWD']))
print('FX rates:', {k:str(v) for k,v in rates.items()})

src_wb=load_workbook(export, data_only=True)
src_ws=src_wb.active
headers=[c.value for c in src_ws[1]]
print('Export headers:', headers)
rows=[]
for vals in src_ws.iter_rows(min_row=2, values_only=True):
    if not any(v is not None for v in vals): continue
    rows.append(dict(zip(headers, vals)))
print('Export row count:', len(rows))
print('Sample rows:', json.dumps(rows[:2], ensure_ascii=False, default=str))

def q2(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src=ws.cell(src_row, col); dst=ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.number_format = src.number_format

wb=load_workbook(template)
report=wb['Report']; details=wb['Details']; summary=wb['Summary']; exceptions=wb['Exceptions']

# preserve template row 5 style by not deleting row 5 before first use; clear from row5 onward values
for ws in (details, summary, exceptions):
    # ensure at least row 5 exists; clear all rows >=5
    if ws.max_row < 5:
        ws.append([None]*ws.max_column)
    maxr=max(ws.max_row,5)
    for r in range(5, maxr+1):
        for c in range(1, ws.max_column+1):
            ws.cell(r,c).value=None
    if maxr > 5:
        ws.delete_rows(6, maxr-5)

from_date='2026-09-01'; to_date='2026-09-30'; employee='E001'
taipei=timezone(timedelta(hours=8))
now=datetime.now(taipei).replace(microsecond=0).isoformat()
report['B3']=from_date
report['D3']=to_date
report['B4']=f'EmployeeID={employee}; ExpenseDate {from_date} to {to_date}; Status=ALL'
report['B5']=now
report['B7']='Expense Viewer (synthetic lab) Export Excel'
report['D7']=now
report['B9']='All exported rows included. NeedsReview/Exceptions mark MissingReceipt=true from source export. AmountTWD uses input/fx_rates.csv fixed test rates.'

summary_map={}
exceptions_out=[]
last_detail_row=4
for idx, rec in enumerate(rows, start=5):
    if idx > 5: copy_row_style(details, 5, idx, 14)
    cur=str(rec.get('Currency') or '')
    amt=q2(rec.get('Amount') or '0')
    if cur not in rates: raise ValueError(f'Missing FX rate for {cur}')
    amt_twd=q2(amt*rates[cur])
    missing=rec.get('MissingReceipt')
    missing_bool = bool(missing) if isinstance(missing,bool) else str(missing).strip().lower() in ('true','1','yes','y')
    vals=[rec.get('ExpenseID'),rec.get('EmployeeID'),rec.get('EmployeeName'),rec.get('ProjectID'),rec.get('ExpenseDate'),rec.get('Category'),cur,float(amt),float(amt_twd),rec.get('Status'),rec.get('ReceiptID'),rec.get('Description'),str(missing_bool).lower(),str(missing_bool).lower()]
    for col,val in enumerate(vals,1): details.cell(idx,col).value=val
    for col in [1,2,3,4,6,7,10,11,12,13,14]: details.cell(idx,col).number_format='@'
    details.cell(idx,5).number_format='yyyy-mm-dd'
    details.cell(idx,8).number_format='0.00'; details.cell(idx,9).number_format='0.00'
    key=(str(rec.get('Status') or ''),cur)
    summary_map.setdefault(key, {'amount':Decimal('0'), 'amount_twd':Decimal('0'), 'count':0})
    summary_map[key]['amount']+=amt; summary_map[key]['amount_twd']+=amt_twd; summary_map[key]['count']+=1
    if missing_bool:
        exceptions_out.append((rec.get('ExpenseID'),rec.get('EmployeeID'),rec.get('ExpenseDate'),'MissingReceipt=true'))
    last_detail_row=idx

summary['B2']=f'=SUM(Details!I5:I{last_detail_row})' if rows else '=SUM(Details!I5:I5)'
summary['B2'].number_format='0.00'
for r_idx, ((status, cur), vals) in enumerate(sorted(summary_map.items()), start=5):
    if r_idx > 5: copy_row_style(summary, 5, r_idx, 6)
    summary.cell(r_idx,1).value='Status'
    summary.cell(r_idx,2).value=status
    summary.cell(r_idx,3).value=cur
    summary.cell(r_idx,4).value=float(q2(vals['amount']))
    summary.cell(r_idx,5).value=float(q2(vals['amount_twd']))
    summary.cell(r_idx,6).value=vals['count']
    summary.cell(r_idx,4).number_format='0.00'; summary.cell(r_idx,5).number_format='0.00'
r_total=5+len(summary_map)
if r_total > 5: copy_row_style(summary, 5, r_total, 6)
summary.cell(r_total,1).value='GrandTotal'
summary.cell(r_total,2).value='TWD total'
summary.cell(r_total,3).value='TWD'
summary.cell(r_total,5).value='=B2'
summary.cell(r_total,6).value=len(rows)
summary.cell(r_total,5).number_format='0.00'

for r_idx, vals in enumerate(exceptions_out, start=5):
    if r_idx > 5: copy_row_style(exceptions, 5, r_idx, 4)
    for col,val in enumerate(vals,1): exceptions.cell(r_idx,col).value=val
    for col in [1,2,4]: exceptions.cell(r_idx,col).number_format='@'
    exceptions.cell(r_idx,3).number_format='yyyy-mm-dd'

wb.save(out)
print('Wrote', out, out.stat().st_size)

# Validation from saved file
check=load_workbook(out, data_only=False)
assert check.sheetnames == ['Report','Details','Summary','Exceptions','Instructions']
d=check['Details']; s=check['Summary']; e=check['Exceptions']; r=check['Report']
expected=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert [d.cell(4,c).value for c in range(1,15)]==expected
n=sum(1 for rr in range(5,d.max_row+1) if d.cell(rr,1).value)
assert n==13
assert r['B3'].value==from_date and r['D3'].value==to_date and employee in r['B4'].value
for rr in range(5,5+n):
    assert isinstance(d.cell(rr,1).value,str) and isinstance(d.cell(rr,2).value,str) and isinstance(d.cell(rr,4).value,str)
    assert isinstance(d.cell(rr,8).value,(int,float)) and isinstance(d.cell(rr,9).value,(int,float))
    assert d.cell(rr,13).value in ('true','false') and d.cell(rr,14).value in ('true','false')
# recompute summary
calc={}; total=Decimal('0'); exc_calc=0
for rr in range(5,5+n):
    status=d.cell(rr,10).value; cur=d.cell(rr,7).value
    amt=Decimal(str(d.cell(rr,8).value)).quantize(Decimal('0.01'))
    twd=Decimal(str(d.cell(rr,9).value)).quantize(Decimal('0.01'))
    calc.setdefault((status,cur), [Decimal('0'),Decimal('0'),0])
    calc[(status,cur)][0]+=amt; calc[(status,cur)][1]+=twd; calc[(status,cur)][2]+=1
    total+=twd
    if d.cell(rr,13).value=='true': exc_calc+=1
saved={}
for rr in range(5,s.max_row+1):
    if s.cell(rr,1).value=='Status':
        saved[(s.cell(rr,2).value,s.cell(rr,3).value)]=[Decimal(str(s.cell(rr,4).value)).quantize(Decimal('0.01')), Decimal(str(s.cell(rr,5).value)).quantize(Decimal('0.01')), s.cell(rr,6).value]
assert saved==calc, (saved, calc)
exc_rows=sum(1 for rr in range(5,e.max_row+1) if e.cell(rr,1).value)
assert exc_rows==exc_calc
assert s['B2'].value==f'=SUM(Details!I5:I{last_detail_row})'
print('Validated output file:', out.name)
print('Details rows:', n)
print('Exceptions rows:', exc_rows)
print('Summary B2 formula:', s['B2'].value)
print('TWD total from Details:', total)
print('Summary groups:')
for rr in range(5,s.max_row+1):
    if s.cell(rr,1).value:
        print([s.cell(rr,c).value for c in range(1,7)])
