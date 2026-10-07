from pathlib import Path
from openpyxl import load_workbook
import csv, json
base=Path.cwd()
exp=base/'output'/'expense_export_E001_2026-09.xlsx'
print('export exists', exp.exists(), exp.stat().st_size if exp.exists() else None)
if exp.exists():
    wb=load_workbook(exp, data_only=False)
    print('export sheets', wb.sheetnames)
    for ws in wb.worksheets:
        print('sheet', ws.title, ws.max_row, ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,8), values_only=True):
            print(row)
# full instructions
wb=load_workbook(base/'input'/'expense_template.xlsx', data_only=False)
ws=wb['Instructions']
print('Instructions:')
for r in range(1, ws.max_row+1):
    vals=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
    if any(v is not None for v in vals): print(r, vals)
