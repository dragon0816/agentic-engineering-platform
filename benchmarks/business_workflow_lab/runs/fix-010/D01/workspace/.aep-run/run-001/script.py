from pathlib import Path
import csv, json
from openpyxl import load_workbook
base=Path.cwd()
print('CWD', base)
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# fx headers/sample
fx=base/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        reader=csv.reader(f)
        rows=[]
        for i,row in zip(range(6), reader):
            rows.append(row)
    print('fx_rates sample:', json.dumps(rows, ensure_ascii=False))
# manual first lines
manual=base/'input'/'expense_manual.md'
if manual.exists():
    txt=manual.read_text(encoding='utf-8', errors='replace')
    print('manual head:\n', '\n'.join(txt.splitlines()[:80]))
# template sheets dimensions, non-empty cells
tpl=base/'input'/'expense_template.xlsx'
if tpl.exists():
    wb=load_workbook(tpl)
    print('template sheets:', wb.sheetnames)
    for ws in wb.worksheets:
        print('sheet', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,20), values_only=False):
            vals=[]
            anyv=False
            for c in row[:min(ws.max_column,12)]:
                if c.value is not None:
                    anyv=True
                vals.append(c.value)
            if anyv:
                print(ws.title, row[0].row, vals)
