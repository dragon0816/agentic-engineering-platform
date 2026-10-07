from pathlib import Path
import shutil
from openpyxl import load_workbook

workspace = Path.cwd()
out = workspace / 'output'
out.mkdir(exist_ok=True)
source = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-010\X01\workspace\output\project_integration_october_P003.xlsx')
target = out / 'october_P003.xlsx'
project = 'P003'
print('Workspace:', workspace)
print('Source workflow output:', source, 'exists=', source.exists())
if not source.exists():
    raise FileNotFoundError(source)

shutil.copyfile(source, target)
wb = load_workbook(target)

# Preserve workbook/template formatting; delete non-P003 data rows from copied workflow result.
def filter_sheet_by_project(ws_name, project_col_name='ProjectID'):
    ws = wb[ws_name]
    headers = [ws.cell(4, c).value for c in range(1, ws.max_column + 1)]
    if project_col_name not in headers:
        raise RuntimeError(f'{ws_name}: no {project_col_name} column in row 4: {headers}')
    pc = headers.index(project_col_name) + 1
    kept = 0
    deleted = 0
    for r in range(ws.max_row, 4, -1):
        val = ws.cell(r, pc).value
        if str(val) == project:
            kept += 1
        else:
            ws.delete_rows(r, 1)
            deleted += 1
    return kept, deleted

results = {}
for sheet in ['Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary']:
    results[sheet] = filter_sheet_by_project(sheet)

# Exceptions: keep only P003-related exceptions. Blank ProjectID exceptions are not P003 and are excluded.
results['Exceptions'] = filter_sheet_by_project('Exceptions')

# Recalculate Summary total formulas to the new P003-only range.
ws = wb['Summary']
last = max(5, ws.max_row)
for col in range(2, 6):
    letter = chr(ord('A') + col - 1)
    ws.cell(2, col).value = f'=SUM({letter}5:{letter}{last})'

# Update report metadata/note to make the project filter explicit, while preserving existing workbook layout.
if 'Report' in wb.sheetnames:
    rw = wb['Report']
    old_condition = rw['B4'].value or ''
    if 'ProjectID=P003' not in str(old_condition):
        rw['B4'] = f'{old_condition} Project filter: ProjectID=P003.'
    old_notes = rw['B9'].value or ''
    if 'Filtered output file contains P003 only' not in str(old_notes):
        rw['B9'] = f'{old_notes} Filtered output file contains P003 only; non-P003 rows removed after the installed workflow run.'

wb.save(target)
print('Wrote:', target, target.stat().st_size, 'bytes')
print('Filter results kept/deleted:', results)

# Verification from saved output.
check = load_workbook(target, data_only=False)
required_sheets = ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']
print('Sheets:', check.sheetnames)
missing = [s for s in required_sheets if s not in check.sheetnames]
if missing:
    raise AssertionError(f'Missing sheets: {missing}')

expected_headers = {
    'Orders':['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
    'Invoices':['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
    'Expenses':['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
    'Quotes':['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
    'Summary':['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
    'Exceptions':['Source','DocumentID','ProjectID','Reason'],
}
for s, heads in expected_headers.items():
    ws = check[s]
    got = [ws.cell(4, c).value for c in range(1, len(heads)+1)]
    print(s, 'headers:', got)
    if got != heads:
        raise AssertionError((s, got, heads))

row_counts = {}
project_values = {}
for s in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    ws = check[s]
    headers = [ws.cell(4,c).value for c in range(1, ws.max_column+1)]
    pc = headers.index('ProjectID') + 1
    vals = []
    data_rows = []
    for r in range(5, ws.max_row + 1):
        # Treat rows with any non-empty cell as data rows.
        row_vals = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        if any(v not in (None, '') for v in row_vals):
            data_rows.append(row_vals)
            vals.append(ws.cell(r,pc).value)
    row_counts[s] = len(data_rows)
    project_values[s] = vals
    bad = [v for v in vals if str(v) != project]
    if bad:
        raise AssertionError(f'{s} has non-P003 ProjectID values: {bad}')
    print(s, 'rows=', len(data_rows), 'ProjectIDs=', vals)

# Check workbook report dates and formulas.
report = check['Report']
assert str(report['B3'].value) == '2026-10-01', report['B3'].value
assert str(report['D3'].value) == '2026-10-31', report['D3'].value
summary = check['Summary']
assert summary['A5'].value == project, summary['A5'].value
assert summary['B2'].value == '=SUM(B5:B5)', summary['B2'].value
assert summary['C2'].value == '=SUM(C5:C5)', summary['C2'].value
assert summary['D2'].value == '=SUM(D5:D5)', summary['D2'].value
assert summary['E2'].value == '=SUM(E5:E5)', summary['E2'].value

# Check text ID formatting for ID columns and numeric formats for money columns.
for s, heads in expected_headers.items():
    ws = check[s]
    for idx, h in enumerate(heads, start=1):
        if h.endswith('ID') or h in ['OrderID','InvoiceID','ExpenseID','QuoteID','CustomerID','EmployeeID','ProjectID','DocumentID']:
            for r in range(5, ws.max_row+1):
                if ws.cell(r, idx).value not in (None, '') and ws.cell(r, idx).number_format != '@':
                    raise AssertionError(f'{s}!{ws.cell(r,idx).coordinate} ID not text formatted: {ws.cell(r,idx).number_format}')
        if h.endswith('TWD') or h in ['NetAmount','Amount','PaidAmount','Balance']:
            for r in range(5, ws.max_row+1):
                if ws.cell(r, idx).value not in (None, '') and not isinstance(ws.cell(r, idx).value, (int, float)):
                    raise AssertionError(f'{s}!{ws.cell(r,idx).coordinate} numeric column has non-numeric {ws.cell(r,idx).value!r}')

print('Verification OK')
print('Output file:', target)
print('Row counts:', row_counts)
