from pathlib import Path
import os
from openpyxl import load_workbook

workspace = Path.cwd()
print('Workspace:', workspace)
print('Input files:')
for p in sorted((workspace/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
print('Output files:')
out = workspace/'output'
out.mkdir(exist_ok=True)
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(out), p.stat().st_size, 'bytes')

target = out/'october_P003.xlsx'
print('Target exists:', target.exists())
if target.exists():
    wb = load_workbook(target, data_only=True)
    print('Workbook sheets:', wb.sheetnames)
    for ws in wb.worksheets:
        print('Sheet', ws.title, 'size', ws.max_row, 'x', ws.max_column)
        # print first 5 rows values
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,5), values_only=True):
            print(' ', row)
