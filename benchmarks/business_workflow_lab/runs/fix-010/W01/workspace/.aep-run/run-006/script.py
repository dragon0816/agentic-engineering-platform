from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
import csv, shutil, datetime, os, json

params=json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
ws=Path.cwd(); out=ws/'output'; downloads=out/'downloads'; out.mkdir(exist_ok=True); downloads.mkdir(exist_ok=True)
date_from=params['date_from']; date_to=params['date_to']; status=(params.get('status') or '').strip(); output_filename=params['output_filename']
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'
csv_path=downloads/'sales_orders_export.csv'
print('inputs:', sorted(p.name for p in (ws/'input').iterdir()))
print('download expected:', csv_path)
if not csv_path.exists():
    raise FileNotFoundError(csv_path)
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    rows=list(csv.DictReader(f))
print('download headers:', list(rows[0].keys()) if rows else [])
print('download row count:', len(rows))

rates={}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    header=next(reader, None)
    for row in reader:
        if row:
            rates[row[0]]=Decimal(row[1])
print('fx currencies:', sorted(rates))

filtered=[]
for r in rows:
    if not (date_from <= r['OrderDate'] <= date_to):
        raise ValueError(f"Row outside requested date range: {r}")
    if status and r['Status'] != status:
        raise ValueError(f"Row outside requested status: {r}")
    if r['Currency'] not in rates:
        raise ValueError(f"Missing FX rate for {r['Currency']}")
    filtered.append(r)
rows=sorted(filtered, key=lambda r: r['OrderID'])

outfile=out/output_filename
shutil.copyfile(ws/'input'/'sales_template.xlsx', outfile)
wb=load_workbook(outfile)
rep=wb['Report']; det=wb['Details']; summ=wb['Summary']
now=datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat(timespec='seconds')
rep['B3']=date_from; rep['D3']=date_to
rep['B4']=f"Sales Orders; Status={status or '(all)'}; Sort=OrderID asc"
rep['B5']=now; rep['B7']='Lab ERP (synthetic SAP-like)'; rep['D7']=now
rep['B9']='Export CSV (all filtered rows) used; FX rates from input/fx_rates.csv.'
for sh in (det, summ):
    if sh.max_row>=5:
        sh.delete_rows(5, sh.max_row-4)
expected=['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
headers=[det.cell(4,c).value for c in range(1,12)]
if headers!=expected:
    raise ValueError(f'Details headers mismatch: {headers}')
q=Decimal('0.01')
amount=defaultdict(Decimal); amount_twd=defaultdict(Decimal); count=defaultdict(int)
for idx,r in enumerate(rows, start=5):
    cur=r['Currency']; net=Decimal(r['NetAmount']); twd=(net*rates[cur]).quantize(q, rounding=ROUND_HALF_UP)
    vals=[r['OrderID'], r['CustomerID'], r['CustomerName'], r.get('ProjectID') or None, r['OrderDate'], r['Status'], cur, float(net), float(twd), None, None]
    for c,v in enumerate(vals,1):
        det.cell(idx,c,value=v)
    for c in [1,2,3,4,5,6,7,10,11]: det.cell(idx,c).number_format='@'
    for c in [8,9]: det.cell(idx,c).number_format='0.00'
    amount[cur]+=net; amount_twd[cur]+=twd; count[cur]+=1
last=4+len(rows)
summ['B2']=f'=SUM(Details!I5:I{last})' if rows else '=SUM(Details!I5:I5)'
srow=5
for cur in sorted(amount):
    vals=['Currency','',cur,float(amount[cur].quantize(q)),float(amount_twd[cur].quantize(q)),count[cur]]
    for c,v in enumerate(vals,1): summ.cell(srow,c,value=v)
    summ.cell(srow,4).number_format='0.00'; summ.cell(srow,5).number_format='0.00'; summ.cell(srow,6).number_format='0'
    srow+=1
total=sum(amount_twd.values(), Decimal('0')).quantize(q)
vals=['Total','All','TWD',None,float(total),len(rows)]
for c,v in enumerate(vals,1): summ.cell(srow,c,value=v)
summ.cell(srow,5).number_format='0.00'; summ.cell(srow,6).number_format='0'
if rows:
    det.auto_filter.ref=f'A4:K{last}'
summ.auto_filter.ref=f'A4:F{srow}'
wb.save(outfile)

# Reopen verification
wb2=load_workbook(outfile, data_only=False)
det2=wb2['Details']; summ2=wb2['Summary']; rep2=wb2['Report']
saved=[]
for r in range(5, det2.max_row+1):
    if det2.cell(r,1).value is not None:
        saved.append([det2.cell(r,c).value for c in range(1,12)])
assert rep2['B3'].value==date_from and rep2['D3'].value==date_to
assert [det2.cell(4,c).value for c in range(1,12)]==expected
assert len(saved)==len(rows)
assert [r[0] for r in saved]==sorted(r[0] for r in saved)
assert all(isinstance(r[0], str) for r in saved)
assert all(isinstance(r[7], (int,float)) and isinstance(r[8], (int,float)) for r in saved)
assert summ2['B2'].value == (f'=SUM(Details!I5:I{last})' if rows else '=SUM(Details!I5:I5)')
print('saved:', outfile)
print('checked row count:', len(saved))
print('checked sorted OrderIDs:', [r[0] for r in saved])
print('summary formula:', summ2['B2'].value)
print('total_twd:', total)