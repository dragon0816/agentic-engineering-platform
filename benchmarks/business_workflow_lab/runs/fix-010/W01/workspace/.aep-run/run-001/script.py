from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)

# workbook structure
wb_path=ws/'input'/'sales_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('\nworkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print(f'\nSheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in range(1, min(sh.max_row,10)+1):
        vals=[sh.cell(r,c).value for c in range(1,min(sh.max_column,12)+1)]
        print(r, vals)

# fx csv
fx_path=ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(10), rdr):
        print(i, row)

# manual
manual=ws/'input'/'sap_manual.md'
print('\nsap_manual.md first 120 lines:')
text=manual.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:120],1):
    print(f'{i}: {line}')