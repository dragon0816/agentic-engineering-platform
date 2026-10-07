from pathlib import Path
text=(Path.cwd()/'input'/'sales_template.xlsx').exists()
from openpyxl import load_workbook
wb=load_workbook(Path.cwd()/'input'/'sales_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in range(10, sh.max_row+1):
    print(r, [sh.cell(r,c).value for c in range(1,4)])