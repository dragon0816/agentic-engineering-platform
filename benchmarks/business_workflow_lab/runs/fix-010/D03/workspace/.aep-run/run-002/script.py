from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('output files:', [p.name for p in (ws/'output').iterdir()])
# exported csv sample
p=ws/'output'/'expense_2026-09_Approved.csv'
print('csv exists', p.exists(), 'size', p.stat().st_size if p.exists() else None)
with open(p, newline='', encoding='utf-8-sig') as f:
    r=csv.DictReader(f)
    print('csv fieldnames', r.fieldnames)
    rows=list(r)
print('csv row count', len(rows))
for row in rows[:5]: print(row)
# instructions full
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
