from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv, shutil
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook
from openpyxl.styles import Font

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)
export_csv = output_dir / 'expense_2026-09_Approved.csv'
template = input_dir / 'expense_template.xlsx'
fx_csv = input_dir / 'fx_rates.csv'
out_xlsx = output_dir / 'expense_report_2026-09_Approved_TWD.xlsx'

print('Inputs:')
print(' export_csv', export_csv, export_csv.exists(), export_csv.stat().st_size if export_csv.exists() else None)
print(' template', template, template.exists(), template.stat().st_size if template.exists() else None)
print(' fx_csv', fx_csv, fx_csv.exists(), fx_csv.stat().st_size if fx_csv.exists() else None)

# Load fixed FX rates. File notes contain commas; only first two columns are relevant.
fx = {}
with open(fx_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header:', header)
    for row in reader:
        if not row or not row[0].strip():
            continue
        fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', fx)

# Load exported expenses.
with open(export_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    rows = list(reader)
print('Export columns:', reader.fieldnames)
print('Export row count:', len(rows))
print('Export sample:', rows[:2])

# Compute TWD and NeedsReview.
q2 = Decimal('0.01')
for r in rows:
    cur = r['Currency'].strip()
    if cur not in fx:
        raise ValueError(f'Missing FX rate for {cur}')
    amount = Decimal(r['Amount'].strip())
    amount_twd = (amount * fx[cur]).quantize(q2, rounding=ROUND_HALF_UP)
    r['_AmountDecimal'] = amount
    r['_AmountTWDDecimal'] = amount_twd
    r['_NeedsReview'] = 'true' if amount_twd > Decimal('5000') else 'false'
    # normalize missing receipt as lower-case true/false text
    mr = str(r.get('MissingReceipt','')).strip().lower()
    r['_MissingReceipt'] = 'true' if mr == 'true' else 'false'

# Sort detail rows by ExpenseDate then ExpenseID for deterministic output (matches export order if same).
rows.sort(key=lambda r: (r['ExpenseDate'], r['ExpenseID']))

# Copy template, then fill sheets.
shutil.copyfile(template, out_xlsx)
wb = load_workbook(out_xlsx)

# Helper: clear rows 5 onward without altering headers/sheets.
def clear_data(sheet, start_row=5):
    if sheet.max_row >= start_row:
        sheet.delete_rows(start_row, sheet.max_row - start_row + 1)

# Report sheet.
report = wb['Report']
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Status=Approved; converted to TWD using input/fx_rates.csv; NeedsReview=true when AmountTWD > 5000'
report['B5'] = '2026-10-06T00:00:00+08:00'
report['B7'] = 'Expense Viewer desktop export; input/fx_rates.csv'
report['D7'] = '2026-10-06T00:00:00+08:00'
report['B9'] = f'Exported {len(rows)} Approved expense rows for September 2026. Original Currency and Amount retained. MissingReceipt values retained from source.'

# Details.
details = wb['Details']
clear_data(details, 5)
for idx, r in enumerate(rows, start=5):
    values = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['ExpenseDate'],
        r['Category'], r['Currency'], float(r['_AmountDecimal']), float(r['_AmountTWDDecimal']),
        r['Status'], r.get('ReceiptID',''), r.get('Description',''), r['_MissingReceipt'], r['_NeedsReview']
    ]
    for col, val in enumerate(values, start=1):
        cell = details.cell(row=idx, column=col, value=val)
        if col in (1,2,3,4,5,6,7,10,11,12,13,14):
            cell.number_format = '@'
        if col in (8,9):
            cell.number_format = '0.00'

# Summary by employee, category and currency so original-currency amounts are not mixed.
summary = wb['Summary']
clear_data(summary, 5)
summary['B2'] = f'=SUM(Details!I5:I{4+len(rows)})' if rows else '=SUM(Details!I5:I5)'
agg = {}
for r in rows:
    emp = f"{r['EmployeeID']} {r['EmployeeName']}"
    key = (emp, r['Category'], r['Currency'])
    if key not in agg:
        agg[key] = {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0}
    agg[key]['amount'] += r['_AmountDecimal']
    agg[key]['amount_twd'] += r['_AmountTWDDecimal']
    agg[key]['count'] += 1
for idx, (key, val) in enumerate(sorted(agg.items(), key=lambda kv: kv[0]), start=5):
    emp, cat, cur = key
    out = [emp, cat, cur, float(val['amount'].quantize(q2, rounding=ROUND_HALF_UP)), float(val['amount_twd'].quantize(q2, rounding=ROUND_HALF_UP)), val['count']]
    for col, cell_val in enumerate(out, start=1):
        cell = summary.cell(row=idx, column=col, value=cell_val)
        if col in (1,2,3):
            cell.number_format = '@'
        elif col in (4,5):
            cell.number_format = '0.00'
        elif col == 6:
            cell.number_format = '0'

# Exceptions: list missing receipts, while still keeping all in Details.
exceptions = wb['Exceptions']
clear_data(exceptions, 5)
ex_rows = [r for r in rows if r['_MissingReceipt'] == 'true']
for idx, r in enumerate(ex_rows, start=5):
    values = [r['ExpenseID'], r['EmployeeID'], r['ExpenseDate'], 'MissingReceipt=true']
    for col, val in enumerate(values, start=1):
        cell = exceptions.cell(row=idx, column=col, value=val)
        cell.number_format = '@'

# Keep workbook structure; save.
wb.save(out_xlsx)
print('Wrote', out_xlsx, out_xlsx.stat().st_size)
print('Summary rows:', len(agg), 'Exception rows:', len(ex_rows))

# Verification: reopen saved output and confirm requirements.
wb2 = load_workbook(out_xlsx, data_only=False)
print('VERIFY sheets:', wb2.sheetnames)
if wb2.sheetnames != ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']:
    raise AssertionError('Sheet names changed')
d = wb2['Details']; s = wb2['Summary']; e = wb2['Exceptions']
headers = [d.cell(4,c).value for c in range(1,15)]
expected_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
print('VERIFY details headers:', headers)
if headers != expected_headers:
    raise AssertionError('Details headers mismatch')
# Count non-empty detail rows.
detail_rows = []
for rr in range(5, d.max_row+1):
    if d.cell(rr,1).value is not None:
        detail_rows.append(rr)
print('VERIFY detail row count:', len(detail_rows))
if len(detail_rows) != 15:
    raise AssertionError('Expected 15 detail rows')
# Verify dates/status/amounts/twd/review flags and retained original currency+amount.
for rr in detail_rows:
    expid = d.cell(rr,1).value
    date = d.cell(rr,5).value
    cur = d.cell(rr,7).value
    amount = Decimal(str(d.cell(rr,8).value)).quantize(q2)
    twd = Decimal(str(d.cell(rr,9).value)).quantize(q2)
    status = d.cell(rr,10).value
    needs = str(d.cell(rr,14).value).lower()
    calc = (amount * fx[cur]).quantize(q2, rounding=ROUND_HALF_UP)
    if not ('2026-09-01' <= date <= '2026-09-30') or status != 'Approved':
        raise AssertionError(f'Filter failed for {expid}')
    if twd != calc:
        raise AssertionError(f'AmountTWD mismatch for {expid}: {twd} != {calc}')
    expected_need = 'true' if twd > Decimal('5000') else 'false'
    if needs != expected_need:
        raise AssertionError(f'NeedsReview mismatch for {expid}')
print('VERIFY all Details AmountTWD and NeedsReview flags OK')
# Verify summary formula and totals.
print('VERIFY Summary B2 formula:', s['B2'].value)
if s['B2'].value != '=SUM(Details!I5:I19)':
    raise AssertionError('Summary B2 formula not extended to last detail row')
sum_detail_twd = sum(Decimal(str(d.cell(rr,9).value)).quantize(q2) for rr in detail_rows)
sum_summary_twd = Decimal('0')
summary_count = 0
for rr in range(5, s.max_row+1):
    if s.cell(rr,1).value is not None:
        sum_summary_twd += Decimal(str(s.cell(rr,5).value)).quantize(q2)
        summary_count += int(s.cell(rr,6).value)
print('VERIFY total TWD details:', sum_detail_twd, 'summary:', sum_summary_twd, 'summary_count:', summary_count)
if sum_detail_twd != sum_summary_twd or summary_count != len(detail_rows):
    raise AssertionError('Summary totals/count do not match Details')
# Verify exceptions count equals missing receipt rows.
missing_detail = sum(1 for rr in detail_rows if str(d.cell(rr,13).value).lower() == 'true')
exception_count = sum(1 for rr in range(5, e.max_row+1) if e.cell(rr,1).value is not None)
print('VERIFY missing receipts:', missing_detail, 'exceptions:', exception_count)
if missing_detail != exception_count:
    raise AssertionError('Exceptions do not match MissingReceipt rows')
# Report values.
print('VERIFY report B3/D3/B4:', wb2['Report']['B3'].value, wb2['Report']['D3'].value, wb2['Report']['B4'].value)
print('Output file:', out_xlsx)
