import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
for p in sorted((ws/'input').iterdir()):
    print('INPUT', p.name, p.stat().st_size)
# inspect fx
fx=ws/'input'/'fx_rates.csv'
with open(fx, newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    rows=[]
    for i,row in zip(range(6), r): rows.append(row)
print('fx sample:', rows)
# inspect template
wb=load_workbook(ws/'input'/'expense_template.xlsx')
print('sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'max', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
        print(row)
# manual sample
m=ws/'input'/'expense_manual.md'
print('manual first lines')
print('\n'.join(m.read_text(encoding='utf-8', errors='replace').splitlines()[:40]))