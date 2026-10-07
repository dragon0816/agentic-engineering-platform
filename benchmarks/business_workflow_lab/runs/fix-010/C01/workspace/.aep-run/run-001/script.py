from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

# workbook structure and sample rows/cells
xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'[{s}] max_row={sh.max_row} max_col={sh.max_column}')
    for r in range(1, min(sh.max_row, 8)+1):
        vals = [sh.cell(r,c).value for c in range(1, min(sh.max_column, 12)+1)]
        print(' row', r, vals)

# fx headers/sample
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        rd = csv.reader(f)
        for i,row in zip(range(6), rd):
            print('fx row', i+1, row)

# manual first lines
manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    text = manual.read_text(encoding='utf-8')
    print('manual length', len(text))
    print('manual preview:\n' + text[:4000])