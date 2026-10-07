from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
for shname in ['Instructions','Report','Details','Summary','Exceptions']:
    sh=wb[shname]
    print('\nSHEET', shname)
    for row in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True):
        print(row)