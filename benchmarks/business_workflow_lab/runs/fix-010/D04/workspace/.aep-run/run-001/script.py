from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
# workbook structure and sample cells
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('template sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print('SHEET', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(r)
# fx rates headers/samples
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rows=list(csv.reader(f))
print('fx_rates rows sample:')
for r in rows[:10]: print(r)
# manual sample/full (not too long)
man=(ws/'input'/'expense_manual.md').read_text(encoding='utf-8')
print('manual length', len(man))
print(man[:4000])