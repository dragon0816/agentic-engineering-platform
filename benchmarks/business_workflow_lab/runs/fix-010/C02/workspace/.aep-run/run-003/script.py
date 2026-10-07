from pathlib import Path
import subprocess, sys, json, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = '2026-09-30'
SOURCE_NAME = 'quote_cli mock'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'quote_template.xlsx'
OUT_MD = OUTPUT / 'findings.md'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
MISSING = 'N/A'
TWOPLACES = Decimal('0.01')

print('Inputs resolved:')
for p in [TEMPLATE, FX, CLI]:
    print('-', p, 'exists=', p.exists())

def money(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def id_text(v):
    if v is None or str(v) == '':
        return MISSING
    return str(v)

def run_cli(args):
    cmd = [sys.executable, str(CLI)] + args
    cp = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True)
    if cp.returncode != 0:
        raise RuntimeError(f'CLI failed rc={cp.returncode}: {cmd}\nstderr={cp.stderr}\nstdout={cp.stdout}')
    return cp.stdout

# Read rates
fx_rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    fx_header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx_rates[row[0]] = Decimal(row[1])
print('FX header:', fx_header)
print('FX rates:', fx_rates)

# Fetch September 2026 quotes, no status filter.
list_payload = json.loads(run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json']))
headers = list_payload.get('items', [])
print('Listed quote count:', list_payload.get('count'), 'items length:', len(headers))
print('Header keys:', sorted(headers[0].keys()) if headers else [])
print('First 3 headers:', headers[:3])

details = []
for h in headers:
    details.append(json.loads(run_cli(['quotes', 'get', h['QuoteID'], '--format', 'json'])))
print('Fetched details:', len(details))
print('First detail line keys:', sorted(details[0].get('line_items',[{}])[0].keys()) if details else [])

details.sort(key=lambda d: (d['header'].get('QuoteDate',''), d['header'].get('QuoteID','')))

quote_rows, line_rows, findings, expired_unaccepted = [], [], [], []
status_currency, currency_totals = {}, {}
missing_fields = []
base = datetime.strptime(BASE_DATE, '%Y-%m-%d').date()

for doc in details:
    h = doc['header']
    qid = id_text(h.get('QuoteID'))
    cust = id_text(h.get('CustomerID'))
    proj = id_text(h.get('ProjectID'))
    for field, val in [('CustomerID', h.get('CustomerID')), ('ProjectID', h.get('ProjectID'))]:
        if val is None or str(val) == '':
            missing_fields.append((qid, field))
    currency = h['Currency']
    subtotal = money(h['Subtotal'])
    discount = Decimal(str(h['DiscountPercent']))
    net = money(h['NetAmount'])
    rate = fx_rates.get(currency)
    if rate is None:
        findings.append({'QuoteID': qid, 'Type': 'Missing FX rate', 'Detail': f'No RateToTWD for {currency}'})
        net_twd = None
    else:
        net_twd = money(net * rate)

    line_sum = Decimal('0.00')
    for line in doc.get('line_items', []):
        qty = Decimal(str(line['Quantity']))
        unit = Decimal(str(line['UnitPrice']))
        line_amount = money(line['LineAmount'])
        expected_line = money(qty * unit)
        if line_amount != expected_line:
            findings.append({'QuoteID': qid, 'Type': 'LineAmount mismatch', 'Detail': f"{line['ProductID']}: Quantity {qty} × UnitPrice {unit} = {expected_line}, reported {line_amount}"})
        line_sum += line_amount
        line_rows.append([qid, id_text(line.get('ProductID')), id_text(line.get('ProductName')), int(qty) if qty == qty.to_integral_value() else float(qty), float(unit), float(line_amount)])
    line_sum = money(line_sum)
    if line_sum != subtotal:
        findings.append({'QuoteID': qid, 'Type': 'Subtotal mismatch', 'Detail': f'LineAmount sum {line_sum} != Subtotal {subtotal}'})

    expected_net = money(subtotal * (Decimal('1') - discount / Decimal('100')))
    if expected_net != net:
        findings.append({'QuoteID': qid, 'Type': 'NetAmount mismatch', 'Detail': f'Subtotal {subtotal} × (1 - {discount}/100) = {expected_net}, reported {net}'})

    valid_until = datetime.strptime(h['ValidUntil'], '%Y-%m-%d').date()
    if valid_until < base and h['Status'] != 'Accepted':
        expired_unaccepted.append(qid)
        findings.append({'QuoteID': qid, 'Type': 'Expired and not Accepted', 'Detail': f"ValidUntil {h['ValidUntil']} < {BASE_DATE}; Status {h['Status']}"})

    quote_rows.append([qid, cust, proj, h['QuoteDate'], h['ValidUntil'], id_text(h['Status']), currency, float(subtotal), float(discount), float(net), None if net_twd is None else float(net_twd)])

    key = (h['Status'], currency)
    status_currency.setdefault(key, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    status_currency[key]['amount'] += net
    status_currency[key]['twd'] += net_twd or Decimal('0.00')
    status_currency[key]['count'] += 1
    currency_totals.setdefault(currency, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    currency_totals[currency]['amount'] += net
    currency_totals[currency]['twd'] += net_twd or Decimal('0.00')
    currency_totals[currency]['count'] += 1

# Fill copy of template, never alter input.
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

def clear_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def money_fmt(cell):
    cell.number_format = '0.00'

def text_fmt(cell):
    cell.number_format = '@'

report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = f'All quotes in September 2026; no status filter; expired flag basis ValidUntil < {BASE_DATE} and Status <> Accepted; all quotes retained.'
report['B5'] = GENERATED_AT
report['B7'] = SOURCE_NAME
report['D7'] = GENERATED_AT
note_parts = [f'{len(findings)} finding(s)', f'{len(expired_unaccepted)} expired and not Accepted', 'all quotes retained']
if missing_fields:
    note_parts.append(f'{len(missing_fields)} missing ID value(s) filled as {MISSING}')
report['B9'] = '; '.join(note_parts) + '. See findings.md.'

qws = wb['Quotes']
clear_data(qws)
for r, row in enumerate(quote_rows, 5):
    for c, val in enumerate(row, 1):
        cell = qws.cell(r, c, val)
        if c in (1,2,3,6,7):
            text_fmt(cell)
        elif c in (4,5):
            cell.number_format = 'yyyy-mm-dd'
        elif c in (8,9,10,11):
            money_fmt(cell)

lws = wb['Lines']
clear_data(lws)
for r, row in enumerate(line_rows, 5):
    for c, val in enumerate(row, 1):
        cell = lws.cell(r, c, val)
        if c in (1,2,3):
            text_fmt(cell)
        elif c in (5,6):
            money_fmt(cell)

sws = wb['Summary']
clear_data(sws)
last_quote_row = 4 + len(quote_rows)
sws['B2'] = 'Total NetAmountTWD'
sws['C2'] = f'=SUM(Quotes!K5:K{last_quote_row})' if quote_rows else 0
money_fmt(sws['C2'])
summary_rows = []
for currency in sorted(currency_totals):
    v = currency_totals[currency]
    summary_rows.append(['Currency total', '', currency, money(v['amount']), money(v['twd']), v['count']])
for (status, currency) in sorted(status_currency):
    v = status_currency[(status, currency)]
    summary_rows.append(['Status total', status, currency, money(v['amount']), money(v['twd']), v['count']])
summary_rows.append(['Exception', 'Expired and not Accepted', '', Decimal('0.00'), Decimal('0.00'), len(expired_unaccepted)])
summary_rows.append(['Exception', 'Calculation findings', '', Decimal('0.00'), Decimal('0.00'), len([f for f in findings if f['Type'] != 'Expired and not Accepted'])])
summary_rows.append(['Data quality', f'Missing ID filled as {MISSING}', '', Decimal('0.00'), Decimal('0.00'), len(missing_fields)])
for r, row in enumerate(summary_rows, 5):
    for c, val in enumerate(row, 1):
        cell = sws.cell(r, c, float(val) if isinstance(val, Decimal) else val)
        if c in (1,2,3):
            text_fmt(cell)
        elif c in (4,5):
            money_fmt(cell)

# findings.md
calc_findings = [f for f in findings if f['Type'] != 'Expired and not Accepted']
expired_findings = [f for f in findings if f['Type'] == 'Expired and not Accepted']
md = [
    '# Quotation Findings',
    '',
    f'- Query period: {FROM_DATE} through {TO_DATE} (inclusive)',
    f'- Source: {SOURCE_NAME}',
    f'- Generated at: {GENERATED_AT}',
    f'- Expiry basis: flag quotes where ValidUntil < {BASE_DATE} and Status is not Accepted',
    f'- Quotes retained in workbook: {len(quote_rows)}',
    f'- Line items retained in workbook: {len(line_rows)}',
    f'- Missing ID placeholder: {MISSING}',
    '',
    '## Calculation checks',
]
if calc_findings:
    md += ['| QuoteID | Finding | Detail |', '|---|---|---|']
    for f in calc_findings:
        md.append(f"| {f['QuoteID']} | {f['Type']} | {f['Detail']} |")
else:
    md.append('No line subtotal or NetAmount inconsistencies found.')
md += ['', '## Expired and not Accepted as of 2026-09-30']
if expired_findings:
    md += ['| QuoteID | Detail |', '|---|---|']
    for f in expired_findings:
        md.append(f"| {f['QuoteID']} | {f['Detail']} |")
else:
    md.append('No quotes matched the expired-and-not-Accepted condition.')
md += ['', '## Missing ID placeholders']
if missing_fields:
    md += ['| QuoteID | Field | Placeholder |', '|---|---|---|']
    for qid, field in missing_fields:
        md.append(f'| {qid} | {field} | {MISSING} |')
else:
    md.append('No missing ID values found.')
md += ['', '## Summary by status and currency', '| Status | Currency | Count | NetAmount | NetAmountTWD |', '|---|---|---:|---:|---:|']
for (status, currency) in sorted(status_currency):
    v = status_currency[(status, currency)]
    md.append(f"| {status} | {currency} | {v['count']} | {money(v['amount'])} | {money(v['twd'])} |")
OUT_MD.write_text('\n'.join(md) + '\n', encoding='utf-8')
wb.save(OUT_XLSX)
print('Wrote outputs:', OUT_XLSX, OUT_MD)

# Reopen every output and verify stated requirements.
wb2 = load_workbook(OUT_XLSX, data_only=False)
assert set(['Report','Quotes','Lines','Summary','Instructions']).issubset(set(wb2.sheetnames))
qws2, lws2, sws2, rws2 = wb2['Quotes'], wb2['Lines'], wb2['Summary'], wb2['Report']
assert [qws2.cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [lws2.cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert rws2['B3'].value == FROM_DATE and rws2['D3'].value == TO_DATE
assert qws2.max_row == 4 + len(quote_rows), (qws2.max_row, len(quote_rows))
assert lws2.max_row == 4 + len(line_rows), (lws2.max_row, len(line_rows))
assert [str(qws2.cell(r,1).value) for r in range(5, qws2.max_row+1)] == [row[0] for row in quote_rows]
for r in range(5, qws2.max_row+1):
    for c in (1,2,3,6,7):
        assert isinstance(qws2.cell(r,c).value, str) and qws2.cell(r,c).value != '', (r,c,qws2.cell(r,c).value)
    for c in (8,9,10,11):
        assert isinstance(qws2.cell(r,c).value, (int,float)), (r,c,qws2.cell(r,c).value,type(qws2.cell(r,c).value))
for r in range(5, lws2.max_row+1):
    for c in (1,2,3):
        assert isinstance(lws2.cell(r,c).value, str) and lws2.cell(r,c).value != '', (r,c,lws2.cell(r,c).value)
    for c in (4,5,6):
        assert isinstance(lws2.cell(r,c).value, (int,float)), (r,c,lws2.cell(r,c).value,type(lws2.cell(r,c).value))
md_text = OUT_MD.read_text(encoding='utf-8')
assert f'Quotes retained in workbook: {len(quote_rows)}' in md_text
assert 'No line subtotal or NetAmount inconsistencies found.' in md_text if not calc_findings else True
for qid in expired_unaccepted:
    assert qid in md_text
for qid, field in missing_fields:
    assert qid in md_text and field in md_text and MISSING in md_text

print('verification passed')
print('quotes retained:', len(quote_rows))
print('line rows retained:', len(line_rows))
print('calculation findings:', len(calc_findings))
print('expired not accepted:', len(expired_unaccepted), expired_unaccepted)
print('missing ID placeholders:', missing_fields)
print('output files:', OUT_XLSX.name, OUT_MD.name)
