from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)

# Inspect template workbook
p = ws/'input'/'quote_template.xlsx'
wb = load_workbook(p, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('\nSHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
        print(r)

# Inspect fx csv
p = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with p.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i,row in zip(range(10), reader):
        print(row)

# Inspect manual
p = ws/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md first 200 lines:')
text = p.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:200], 1):
    print(f'{i:03d}: {line}')
