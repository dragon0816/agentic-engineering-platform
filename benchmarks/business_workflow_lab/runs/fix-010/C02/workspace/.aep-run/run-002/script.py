from pathlib import Path
import subprocess, sys, json, csv, shutil, os
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = '2026-09-30'
SOURCE_NAME = 'quote_cli mock'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'quote_template.xlsx'
OUT_MD = OUTPUT / 'findings.md'

print('Inputs resolved:')
print(' template:', TEMPLATE, TEMPLATE.exists())
print(' fx:', FX, FX.exists())
print(' cli:', CLI, CLI.exists())
print(' output:', OUTPUT)

TWOPLACES = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [sys.executable, str(CLI)] + args
    cp = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True)
    if cp.returncode != 0:
        raise RuntimeError(f'CLI failed rc={cp.returncode}: {cmd}\nstderr={cp.stderr}\nstdout={cp.stdout}')
    return cp.stdout

# Read FX rates, keeping source's first two columns (Note may contain unquoted commas in this synthetic file).
fx_rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx_rates[row[0]] = Decimal(row[1])
print('fx currencies:', fx_rates)

# Fetch all September 2026 quote headers; no status filter.
list_payload = json.loads(run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json']))
headers = list_payload.get('items', [])
print('listed quotes count:', list_payload.get('count'), 'items:', len(headers))
print('sample headers:', headers[:3])

# Fetch all quote details.
details = []
for h in headers:
    qid = h['QuoteID']
    payload = json.loads(run_cli(['quotes', 'get', qid, '--format', 'json']))
    details.append(payload)
print('details fetched:', len(details))
if details:
    print('sample detail header:', details[0]['header'])
    print('sample detail lines:', details[0].get('line_items', [])[:3])

# Sort by QuoteDate then QuoteID for stable output.
details.sort(key=lambda d: (d['header'].get('QuoteDate',''), d['header'].get('QuoteID','')))

# Validate calculations and collect rows.
quote_rows = []
line_rows = []
findings = []
expired_unaccepted = []
status_counts = {}
currency_totals = {}
status_currency = {}

base = datetime.strptime(BASE_DATE, '%Y-%m-%d').date()

for doc in details:
    h = doc['header']
    qid = str(h['QuoteID'])
    currency = h['Currency']
    subtotal = money(h['Subtotal'])
    discount = Decimal(str(h['DiscountPercent']))
    net = money(h['NetAmount'])
    rate = fx_rates.get(currency)
    if rate is None:
        findings.append({'QuoteID': qid, 'Type': 'Missing FX rate', 'Detail': f'No RateToTWD for {currency}'})
        net_twd = None
    else:
        net_twd = money(net * rate)

    lines = doc.get('line_items', [])
    line_sum = Decimal('0.00')
    for line in lines:
        qty = Decimal(str(line['Quantity']))
        unit = Decimal(str(line['UnitPrice']))
        line_amount = money(line['LineAmount'])
        expected_line = money(qty * unit)
        if line_amount != expected_line:
            findings.append({'QuoteID': qid, 'Type': 'LineAmount mismatch', 'Detail': f"{line['ProductID']}: Quantity {qty} × UnitPrice {unit} = {expected_line}, reported {line_amount}"})
        line_sum += line_amount
        line_rows.append([
            qid,
            str(line['ProductID']),
            str(line['ProductName']),
            int(qty) if qty == qty.to_integral_value() else float(qty),
            float(unit),
            float(line_amount),
        ])
    line_sum = money(line_sum)
    if line_sum != subtotal:
        findings.append({'QuoteID': qid, 'Type': 'Subtotal mismatch', 'Detail': f'LineAmount sum {line_sum} != Subtotal {subtotal}'})

    expected_net = money(subtotal * (Decimal('1') - discount / Decimal('100')))
    if expected_net != net:
        findings.append({'QuoteID': qid, 'Type': 'NetAmount mismatch', 'Detail': f'Subtotal {subtotal} × (1 - {discount}/100) = {expected_net}, reported {net}'})

    valid_until = datetime.strptime(h['ValidUntil'], '%Y-%m-%d').date()
    if valid_until < base and h['Status'] != 'Accepted':
        expired_unaccepted.append(qid)
        findings.append({'QuoteID': qid, 'Type': 'Expired and not Accepted', 'Detail': f"ValidUntil {h['ValidUntil']} < {BASE_DATE}; Status {h['Status']}"})

    quote_rows.append([
        qid,
        str(h['CustomerID']),
        str(h['ProjectID']),
        h['QuoteDate'],
        h['ValidUntil'],
        h['Status'],
        currency,
        float(subtotal),
        float(discount),
        float(net),
        None if net_twd is None else float(net_twd),
    ])

    status = h['Status']
    status_counts[status] = status_counts.get(status, 0) + 1
    currency_totals.setdefault(currency, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    currency_totals[currency]['amount'] += net
    currency_totals[currency]['twd'] += net_twd or Decimal('0.00')
    currency_totals[currency]['count'] += 1
    key = (status, currency)
    status_currency.setdefault(key, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    status_currency[key]['amount'] += net
    status_currency[key]['twd'] += net_twd or Decimal('0.00')
    status_currency[key]['count'] += 1

# Fill a copy of the template workbook.
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

# Preserve Instructions, write specified report cells.
report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = f'All quotes in September 2026; no status filter; expired flag basis ValidUntil < {BASE_DATE} and Status <> Accepted; all quotes retained.'
# Use the provided environment date as the generated date context.
generated_at = '2026-10-06T00:00:00+08:00'
report['B5'] = generated_at
report['B7'] = SOURCE_NAME
report['D7'] = generated_at
if findings:
    report['B9'] = f'{len(findings)} finding(s); {len(expired_unaccepted)} expired and not Accepted. See findings.md.'
else:
    report['B9'] = 'No calculation mismatches and no expired unaccepted quotes found.'

# Clear detail sheets from row 5 down, preserving headers/formats as much as possible.
def clear_data(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def apply_money(cell):
    cell.number_format = '0.00'

def apply_date_text(cell):
    cell.number_format = 'yyyy-mm-dd'

quotes_ws = wb['Quotes']
clear_data(quotes_ws)
for r_idx, row in enumerate(quote_rows, 5):
    for c_idx, value in enumerate(row, 1):
        cell = quotes_ws.cell(r_idx, c_idx, value)
        if c_idx in (1,2,3,6,7):
            cell.number_format = '@'
        elif c_idx in (4,5):
            cell.number_format = 'yyyy-mm-dd'
        elif c_idx in (8,9,10,11):
            apply_money(cell)

lines_ws = wb['Lines']
clear_data(lines_ws)
for r_idx, row in enumerate(line_rows, 5):
    for c_idx, value in enumerate(row, 1):
        cell = lines_ws.cell(r_idx, c_idx, value)
        if c_idx in (1,2,3):
            cell.number_format = '@'
        elif c_idx in (4,5,6):
            apply_money(cell) if c_idx in (5,6) else None

summary_ws = wb['Summary']
clear_data(summary_ws)
summary_ws['B2'] = 'Total NetAmountTWD'
last_quote_row = 4 + len(quote_rows)
summary_ws['C2'] = f'=SUM(Quotes!K5:K{last_quote_row})' if quote_rows else 0
apply_money(summary_ws['C2'])
summary_rows = []
for currency in sorted(currency_totals):
    v = currency_totals[currency]
    summary_rows.append(['Currency total', '', currency, money(v['amount']), money(v['twd']), v['count']])
for (status, currency) in sorted(status_currency):
    v = status_currency[(status, currency)]
    summary_rows.append(['Status total', status, currency, money(v['amount']), money(v['twd']), v['count']])
summary_rows.append(['Exception', 'Expired and not Accepted', '', Decimal('0.00'), Decimal('0.00'), len(expired_unaccepted)])
summary_rows.append(['Exception', 'Calculation findings', '', Decimal('0.00'), Decimal('0.00'), len([f for f in findings if f['Type'] != 'Expired and not Accepted'])])
for r_idx, row in enumerate(summary_rows, 5):
    for c_idx, value in enumerate(row, 1):
        cell = summary_ws.cell(r_idx, c_idx, float(value) if isinstance(value, Decimal) else value)
        if c_idx in (1,2,3):
            cell.number_format = '@'
        elif c_idx in (4,5):
            apply_money(cell)

# Write findings.md.
md = []
md.append('# Quotation Findings')
md.append('')
md.append(f'- Query period: {FROM_DATE} through {TO_DATE} (inclusive)')
md.append(f'- Source: {SOURCE_NAME}')
md.append(f'- Generated at: {generated_at}')
md.append(f'- Expiry basis: flag quotes where ValidUntil < {BASE_DATE} and Status is not Accepted')
md.append(f'- Quotes retained in workbook: {len(quote_rows)}')
md.append(f'- Line items retained in workbook: {len(line_rows)}')
md.append('')
md.append('## Calculation checks')
calc_findings = [f for f in findings if f['Type'] != 'Expired and not Accepted']
if calc_findings:
    md.append('| QuoteID | Finding | Detail |')
    md.append('|---|---|---|')
    for f in calc_findings:
        md.append(f"| {f['QuoteID']} | {f['Type']} | {f['Detail']} |")
else:
    md.append('No line subtotal or NetAmount inconsistencies found.')
md.append('')
md.append('## Expired and not Accepted as of 2026-09-30')
expired_findings = [f for f in findings if f['Type'] == 'Expired and not Accepted']
if expired_findings:
    md.append('| QuoteID | Detail |')
    md.append('|---|---|')
    for f in expired_findings:
        md.append(f"| {f['QuoteID']} | {f['Detail']} |")
else:
    md.append('No quotes matched the expired-and-not-Accepted condition.')
md.append('')
md.append('## Summary by status and currency')
md.append('| Status | Currency | Count | NetAmount | NetAmountTWD |')
md.append('|---|---|---:|---:|---:|')
for (status, currency) in sorted(status_currency):
    v = status_currency[(status, currency)]
    md.append(f"| {status} | {currency} | {v['count']} | {money(v['amount'])} | {money(v['twd'])} |")
md.append('')
OUT_MD.write_text('\n'.join(md), encoding='utf-8')

wb.save(OUT_XLSX)
print('wrote:', OUT_XLSX)
print('wrote:', OUT_MD)

# Reopen and verify stated requirements from saved outputs.
wb2 = load_workbook(OUT_XLSX, data_only=False)
qws = wb2['Quotes']
lws = wb2['Lines']
sws = wb2['Summary']
rws = wb2['Report']
# Check row counts and headers.
expected_q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
expected_l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
actual_q_headers = [qws.cell(4,c).value for c in range(1,12)]
actual_l_headers = [lws.cell(4,c).value for c in range(1,7)]
assert actual_q_headers == expected_q_headers, actual_q_headers
assert actual_l_headers == expected_l_headers, actual_l_headers
assert rws['B3'].value == FROM_DATE and rws['D3'].value == TO_DATE
assert qws.max_row == 4 + len(quote_rows), (qws.max_row, len(quote_rows))
assert lws.max_row == 4 + len(line_rows), (lws.max_row, len(line_rows))
# Check every quote from list retained exactly once.
saved_qids = [str(qws.cell(r,1).value) for r in range(5, qws.max_row+1)]
source_qids = [row[0] for row in quote_rows]
assert saved_qids == source_qids, (saved_qids, source_qids)
# Check ID formats and numeric fields.
for r in range(5, qws.max_row+1):
    for c in (1,2,3,6,7):
        assert qws.cell(r,c).data_type == 's', (r,c,qws.cell(r,c).value,qws.cell(r,c).data_type)
    for c in (8,9,10,11):
        v = qws.cell(r,c).value
        assert isinstance(v, (int,float)), (r,c,v,type(v))
for r in range(5, lws.max_row+1):
    for c in (1,2,3):
        assert lws.cell(r,c).data_type == 's', (r,c,lws.cell(r,c).value,lws.cell(r,c).data_type)
    for c in (4,5,6):
        assert isinstance(lws.cell(r,c).value, (int,float)), (r,c,lws.cell(r,c).value,type(lws.cell(r,c).value))
# Verify findings.md mentions all expired unaccepted qids and the count.
md_text = OUT_MD.read_text(encoding='utf-8')
for qid in expired_unaccepted:
    assert qid in md_text, qid
assert f'Quotes retained in workbook: {len(quote_rows)}' in md_text

print('verification passed')
print('quotes retained:', len(quote_rows))
print('line rows retained:', len(line_rows))
print('findings total:', len(findings))
print('calc findings:', len(calc_findings))
print('expired not accepted:', len(expired_unaccepted), expired_unaccepted)
print('outputs:', [str(OUT_XLSX), str(OUT_MD)])
