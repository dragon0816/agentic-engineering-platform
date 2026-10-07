from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

manual = ws/'input'/'quote_cli_manual.md'
print('\n--- quote_cli_manual.md first 200 lines ---')
if manual.exists():
    text = manual.read_text(encoding='utf-8', errors='replace')
    for i, line in enumerate(text.splitlines()[:200], 1):
        print(f'{i:03d}: {line}')
else:
    print('manual missing')

fx = ws/'input'/'fx_rates.csv'
print('\n--- fx_rates.csv sample ---')
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            print(row)
            if i >= 10:
                break
else:
    print('fx missing')

tpl = ws/'input'/'quote_template.xlsx'
print('\n--- workbook structure ---')
if tpl.exists():
    wb = load_workbook(tpl, data_only=False)
    print('sheets:', wb.sheetnames)
    for sh in wb.worksheets:
        print(f'[{sh.title}] max_row={sh.max_row} max_column={sh.max_column}')
        for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(r)
        print('---')
else:
    print('template missing')
