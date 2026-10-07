from pathlib import Path
import subprocess, json, csv, os, sys, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'sent_quotes_2026-09.xlsx'
LOG_TXT = OUTPUT / 'sent_quotes_2026-09_cli_log.txt'

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
TZ = timezone(timedelta(hours=8))

print('Inputs confirmed:')
print(' CLI:', CLI, 'exists=', CLI.exists())
print(' TEMPLATE:', TEMPLATE, 'exists=', TEMPLATE.exists())
print(' FX:', FX, 'exists=', FX.exists())
print(' Output workbook:', OUT_XLSX)

if not CLI.exists():
    raise FileNotFoundError(f'CLI not found: {CLI}')
if not TEMPLATE.exists():
    raise FileNotFoundError(f'Template not found: {TEMPLATE}')
if not FX.exists():
    raise FileNotFoundError(f'FX rates not found: {FX}')

def q2(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def load_fx(path):
    rates = {}
    with path.open('r', encoding='utf-8-sig', newline='') as f:
        header = f.readline().rstrip('\n').split(',')
        for line in f:
            if not line.strip():
                continue
            parts = line.rstrip('\n').split(',', 2)
            if len(parts) >= 2:
                cur = parts[0].strip()
                rate = parts[1].strip()
                if cur:
                    rates[cur] = Decimal(rate)
    return rates

log_lines = []

def run_cli_json(args, description):
    """Run CLI. Exit code 4 retries at most twice; 2/3 never retry. Return parsed JSON only on code 0."""
    attempts = 0
    max_retries_for_4 = 2
    cmd = [str(CLI)] + args
    while True:
        attempts += 1
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True, shell=False)
        stderr = proc.stderr.strip()
        stdout = proc.stdout.strip()
        log_lines.append(f'[{datetime.now(TZ).isoformat(timespec="seconds")}] {description} attempt {attempts}: exit={proc.returncode}')
        if stderr:
            log_lines.append('stderr: ' + stderr)
        log_lines.append('stdout_bytes: ' + str(len(proc.stdout.encode('utf-8', errors='replace'))))
        if proc.returncode == 0:
            try:
                data = json.loads(stdout)
            except Exception as e:
                # stdout is success but not parseable; do not treat as data.
                raise RuntimeError(f'{description}: CLI succeeded but stdout was not parseable JSON: {e}') from e
            return data, attempts
        if proc.returncode in (2, 3):
            reason = stderr or '(stderr was empty)'
            raise RuntimeError(f'{description}: CLI exit code {proc.returncode}; not retried per rule. stderr: {reason}')
        if proc.returncode == 4:
            if attempts <= max_retries_for_4:
                continue
            reason = stderr or '(stderr was empty)'
            raise RuntimeError(f'{description}: CLI exit code 4 after {attempts} attempts (initial + {max_retries_for_4} retries). stderr: {reason}')
        reason = stderr or '(stderr was empty)'
        raise RuntimeError(f'{description}: unexpected CLI exit code {proc.returncode}. stderr: {reason}')

# Fetch data
list_args = ['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
quotes_list, list_attempts = run_cli_json(list_args, 'quotes list 2026-09 Sent')
if quotes_list.get('schema_version') != 1 or not isinstance(quotes_list.get('items'), list):
    raise RuntimeError('quotes list JSON schema invalid')
headers = quotes_list['items']
line_rows = []
get_attempts = {}
for h in headers:
    qid = h.get('QuoteID')
    if not qid:
        raise RuntimeError('quotes list item missing QuoteID')
    detail, attempts = run_cli_json(['quotes', 'get', qid, '--format', 'json'], f'quotes get {qid}')
    get_attempts[qid] = attempts
    if detail.get('schema_version') != 1 or not isinstance(detail.get('line_items'), list):
        raise RuntimeError(f'quotes get {qid} JSON schema invalid')
    # If header fields in get are richer, keep list header as requested list row source, but verify same ID.
    if detail.get('header', {}).get('QuoteID') not in (None, qid):
        raise RuntimeError(f'quotes get {qid} returned mismatched header QuoteID')
    for li in detail['line_items']:
        row = dict(li)
        row['QuoteID'] = qid
        line_rows.append(row)

rates = load_fx(FX)
missing_rates = sorted({h.get('Currency') for h in headers if h.get('Currency') not in rates})
if missing_rates:
    raise RuntimeError('Missing FX rate(s): ' + ', '.join(missing_rates))

# Fill template copy
if OUT_XLSX.exists():
    OUT_XLSX.unlink()
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

fetched_at = datetime.now(TZ).isoformat(timespec='seconds')
report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = f'Status={STATUS}; QuoteDate {FROM_DATE}..{TO_DATE} (inclusive)'
report['B5'] = fetched_at
report['B7'] = 'quote_cli synthetic quote system; fx_rates.csv'
report['D7'] = fetched_at
notes = []
if not headers:
    notes.append('No matching quotes returned.')
notes.append(f'CLI retry rule applied: exit code 4 max 2 retries; exit code 2/3 no retry; stdout used only after exit 0 JSON parse success. list attempts={list_attempts}; get attempts={get_attempts}.')
report['B9'] = ' '.join(notes)

# Helper for clearing old data rows without disturbing header rows. Template has only headers, but keep safe.
def clear_rows(ws, start_row, max_col):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

quotes_ws = wb['Quotes']
lines_ws = wb['Lines']
summary_ws = wb['Summary']
clear_rows(quotes_ws, 5, 11)
clear_rows(lines_ws, 5, 6)
clear_rows(summary_ws, 5, 6)

quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for idx, h in enumerate(headers, start=5):
    cur = h.get('Currency')
    net = q2(h.get('NetAmount', '0'))
    net_twd = q2(net * rates[cur])
    values = [
        h.get('QuoteID'), h.get('CustomerID'), h.get('ProjectID'), h.get('QuoteDate'), h.get('ValidUntil'),
        h.get('Status'), cur, q2(h.get('Subtotal', '0')), Decimal(str(h.get('DiscountPercent', '0'))), net, net_twd
    ]
    for col, val in enumerate(values, start=1):
        cell = quotes_ws.cell(idx, col, val)
        if col in (1,2,3,4,5,6,7):
            cell.number_format = '@' if col in (1,2,3) else 'yyyy-mm-dd' if col in (4,5) else 'General'
        if col in (8,10,11):
            cell.number_format = '0.00'
        if col == 9:
            cell.number_format = '0.00'

line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for idx, li in enumerate(line_rows, start=5):
    values = [li.get('QuoteID'), li.get('ProductID'), li.get('ProductName'), Decimal(str(li.get('Quantity', '0'))), q2(li.get('UnitPrice', '0')), q2(li.get('LineAmount', '0'))]
    for col, val in enumerate(values, start=1):
        cell = lines_ws.cell(idx, col, val)
        if col in (1,2):
            cell.number_format = '@'
        if col in (5,6):
            cell.number_format = '0.00'

# Summary: group by query status and currency, because TWD conversion must not mix original currencies.
summary = {}
for h in headers:
    key = (h.get('Status') or '', h.get('Currency') or '')
    net = q2(h.get('NetAmount', '0'))
    net_twd = q2(net * rates[h.get('Currency')])
    rec = summary.setdefault(key, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    rec['amount'] += net
    rec['twd'] += net_twd
    rec['count'] += 1
for idx, ((status, cur), rec) in enumerate(sorted(summary.items()), start=5):
    vals = ['Status', status, cur, q2(rec['amount']), q2(rec['twd']), rec['count']]
    for col, val in enumerate(vals, start=1):
        cell = summary_ws.cell(idx, col, val)
        if col in (4,5):
            cell.number_format = '0.00'

last_quote_row = max(5, 4 + len(headers))
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'

# Save
wb.save(OUT_XLSX)
LOG_TXT.write_text('\n'.join(log_lines) + '\n', encoding='utf-8')

# Verification by reopening output files
wbc = load_workbook(OUT_XLSX, data_only=False)
checks = []
checks.append(('output exists', OUT_XLSX.exists()))
checks.append(('log exists', LOG_TXT.exists()))
checks.append(('sheets preserved', wbc.sheetnames == ['Report','Quotes','Lines','Summary','Instructions']))
checks.append(('Report B3', wbc['Report']['B3'].value == FROM_DATE))
checks.append(('Report D3', wbc['Report']['D3'].value == TO_DATE))
checks.append(('Report B4 has Sent', 'Sent' in str(wbc['Report']['B4'].value)))
checks.append(('footer preserved', wbc['Report']['A20'].value == '合成測試資料 Synthetic lab data · template v1 · 請勿修改本行'))
checks.append(('Quotes headers preserved', [wbc['Quotes'].cell(4,c).value for c in range(1,12)] == quote_cols))
checks.append(('Lines headers preserved', [wbc['Lines'].cell(4,c).value for c in range(1,7)] == line_cols))
checks.append(('Summary headers preserved', [wbc['Summary'].cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']))
checks.append(('Quote row count', len([r for r in range(5, wbc['Quotes'].max_row+1) if wbc['Quotes'].cell(r,1).value]) == len(headers)))
checks.append(('Line row count', len([r for r in range(5, wbc['Lines'].max_row+1) if wbc['Lines'].cell(r,1).value]) == len(line_rows)))
checks.append(('Summary B2 formula', wbc['Summary']['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})'))
# IDs text format; numeric amounts have Decimal/int/float acceptable numeric cell values (openpyxl returns Decimal until saved? often Decimal preserved as numeric)
for r in range(5, 5+len(headers)):
    checks.append((f'Quotes row {r} IDs text format', all(wbc['Quotes'].cell(r,c).number_format == '@' for c in (1,2,3))))
    checks.append((f'Quotes row {r} amount formats', all(wbc['Quotes'].cell(r,c).number_format == '0.00' for c in (8,10,11))))
for r in range(5, 5+len(line_rows)):
    checks.append((f'Lines row {r} IDs text format', all(wbc['Lines'].cell(r,c).number_format == '@' for c in (1,2))))
    checks.append((f'Lines row {r} amount formats', all(wbc['Lines'].cell(r,c).number_format == '0.00' for c in (5,6))))
failed = [name for name, ok in checks if not ok]
print('\nVerification checks:')
for name, ok in checks:
    print(('OK ' if ok else 'FAIL ') + name)
if failed:
    raise RuntimeError('Verification failed: ' + '; '.join(failed))
print('\nFetched quote count:', len(headers))
print('Fetched line count:', len(line_rows))
print('Summary groups:', len(summary))
print('CLI attempts:', 'list=', list_attempts, 'get=', get_attempts)
print('Wrote:', OUT_XLSX)
print('Wrote:', LOG_TXT)
