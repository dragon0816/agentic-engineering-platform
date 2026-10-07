from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
wb = load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
sh = wb['Instructions']
for r in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True):
    print(r)
print('\nMerged ranges:')
for s in wb.sheetnames:
    print(s, list(wb[s].merged_cells.ranges))
print('\nReport cells around footer:')
for row in wb['Report'].iter_rows(min_row=1, max_row=25, values_only=False):
    vals=[c.value for c in row]
    if any(v is not None for v in vals):
        print(row[0].row, vals)
