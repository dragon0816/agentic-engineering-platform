import os, json, csv, sys, subprocess, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
FROM = params['from_date']
TO = params['to_date']
PERIOD = params['period_label']
ws = Path.cwd()
out = ws / 'output'
downloads = out / 'downloads'
out.mkdir(exist_ok=True)
downloads.mkdir(parents=True, exist_ok=True)
print('Inputs:', {'workspace': str(ws), 'from_date': FROM, 'to_date': TO, 'period_label': PERIOD})
print('Input files:', sorted(p.name for p in (ws/'input').glob('*')))
print('Download files before processing:', sorted(p.name for p in downloads.glob('*')))

TW = timezone(timedelta(hours=8))
now = datetime.now(TW).replace(microsecond=0).isoformat()
q2 = Decimal('0.01')
def dec(value):
    if value is None or str(value).strip() == '':
        return Decimal('0.00')
    return Decimal(str(value).replace(',', '').strip()).quantize(q2, rounding=ROUND_HALF_UP)
def twd(amount, currency, fx):
    return (dec(amount) * fx[currency]).quantize(q2, rounding=ROUND_HALF_UP)
def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))
def require_cols(rows, cols, source):
    if not rows:
        print(f'{source}: 0 rows')
        return
    missing = [c for c in cols if c not in rows[0]]
    if missing:
        raise RuntimeError(f'{source} missing columns: {missing}; headers={list(rows[0].keys())}')
def archive_download(src_name, dst_name):
    src = downloads / src_name
    if not src.exists():
        raise FileNotFoundError(src)
    dst = downloads / dst_name
    shutil.copyfile(src, dst)
    return dst

# Read FX rates from the attachment.
fx = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    for r in csv.DictReader(f):
        cur = (r.get('Currency') or '').strip()
        rate = (r.get('RateToTWD') or '').strip()
        if cur and rate:
            fx[cur] = Decimal(rate)
print('FX currencies:', sorted(fx.keys()))
if not fx:
    raise RuntimeError('No FX rates loaded from input/fx_rates.csv')

# Archive browser/desktop fixed download names with the requested period label.
open_path = archive_download('sap_orders_open.csv', f'sap_orders_open_{PERIOD}.csv')
closed_path = archive_download('sap_orders_closed.csv', f'sap_orders_closed_{PERIOD}.csv')
inv_path = archive_download('sap_invoices.csv', f'sap_invoices_{PERIOD}.csv')
exp_path = archive_download('expenses_approved.csv', f'expenses_approved_{PERIOD}.csv')

orders = read_csv(open_path) + read_csv(closed_path)
require_cols(orders, ['OrderID','CustomerID','ProjectID','OrderDate','Status','Currency','NetAmount'], 'SAP orders')
orders = [r for r in orders if FROM <= r['OrderDate'] <= TO and r['Status'] != 'Cancelled']

invoices = read_csv(inv_path)
require_cols(invoices, ['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance'], 'SAP invoices')
invoices = [r for r in invoices if FROM <= r['InvoiceDate'] <= TO]

expenses = read_csv(exp_path)
require_cols(expenses, ['ExpenseID','EmployeeID','ProjectID','ExpenseDate','Status','Currency','Amount'], 'Expenses')
expenses = [r for r in expenses if FROM <= r['ExpenseDate'] <= TO and r['Status'] == 'Approved']

# Quote CLI: Accepted quotes only, filtered by the same period.
cli = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py'
cmd = [sys.executable, cli, 'quotes', 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json']
print('Running quote CLI:', cmd)
cp = subprocess.run(cmd, capture_output=True, text=True)
print('quote_cli exit:', cp.returncode)
if cp.stderr:
    print('quote_cli stderr:', cp.stderr)
if cp.returncode != 0:
    raise RuntimeError(f'quote_cli failed: {cp.returncode}')
quote_data = json.loads(cp.stdout)
quotes = quote_data.get('items', quote_data if isinstance(quote_data, list) else [])
require_cols(quotes, ['QuoteID','CustomerID','ProjectID','QuoteDate','Status','Currency','NetAmount'], 'Quotes')
quotes = [r for r in quotes if FROM <= r['QuoteDate'] <= TO and r['Status'] == 'Accepted']
(downloads / f'quotes_accepted_{PERIOD}.json').write_text(json.dumps({'items': quotes}, ensure_ascii=False, indent=2), encoding='utf-8')
quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount']
with open(downloads / f'quotes_accepted_{PERIOD}.csv', 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.DictWriter(f, fieldnames=quote_cols)
    w.writeheader()
    for r in quotes:
        w.writerow({c: r.get(c, '') for c in quote_cols})

# Fill a copy of the real workbook template.
template = ws/'input'/'project_template.xlsx'
if not template.exists():
    raise FileNotFoundError(template)
output_xlsx = out / f'project_integration_report_{PERIOD}.xlsx'
shutil.copyfile(template, output_xlsx)
wb = load_workbook(output_xlsx)
required_sheets = ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions']
missing_sheets = [s for s in required_sheets if s not in wb.sheetnames]
if missing_sheets:
    raise RuntimeError(f'Template missing sheets: {missing_sheets}')
for sheet in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    sh = wb[sheet]
    if sh.max_row >= 5:
        sh.delete_rows(5, sh.max_row - 4)

exceptions = []
summary = {}
def project_bucket(pid):
    summary.setdefault(pid, {'OrderTWD': Decimal('0.00'), 'InvoiceBalanceTWD': Decimal('0.00'), 'ApprovedExpenseTWD': Decimal('0.00'), 'AcceptedQuoteTWD': Decimal('0.00')})
    return summary[pid]
def add_exception(source, doc, pid, reason):
    exceptions.append({'Source': source, 'DocumentID': doc or '', 'ProjectID': pid or '', 'Reason': reason})
def usable(source, doc, pid, currency):
    if not pid:
        add_exception(source, doc, pid, 'Missing ProjectID; excluded from project summary')
        return False
    if currency not in fx:
        add_exception(source, doc, pid, f'Missing FX rate for {currency}; excluded from project summary')
        return False
    project_bucket(pid)
    return True
def set_text(cell):
    cell.number_format = '@'

def write_row(sh, row, vals, text_cols):
    for c, v in enumerate(vals, 1):
        sh.cell(row, c, v)
    for c in text_cols:
        set_text(sh.cell(row, c))

# Detail sheets and separate project totals.
sh = wb['Orders']; row = 5
for r in sorted(orders, key=lambda x: (x.get('ProjectID',''), x.get('OrderID',''))):
    if not usable('Orders', r.get('OrderID'), r.get('ProjectID'), r.get('Currency')):
        continue
    amount_twd = twd(r['NetAmount'], r['Currency'], fx)
    write_row(sh, row, [r['OrderID'], r['ProjectID'], r['CustomerID'], r['OrderDate'], r['Status'], r['Currency'], float(dec(r['NetAmount'])), float(amount_twd)], [1,2,3,4,5,6])
    summary[r['ProjectID']]['OrderTWD'] += amount_twd
    row += 1

sh = wb['Invoices']; row = 5
for r in sorted(invoices, key=lambda x: (x.get('ProjectID',''), x.get('InvoiceID',''))):
    if not usable('Invoices', r.get('InvoiceID'), r.get('ProjectID'), r.get('Currency')):
        continue
    balance_twd = twd(r['Balance'], r['Currency'], fx)
    write_row(sh, row, [r['InvoiceID'], r['OrderID'], r['ProjectID'], r['CustomerID'], r['InvoiceDate'], r['DueDate'], r['Currency'], float(dec(r['Amount'])), float(dec(r['PaidAmount'])), float(dec(r['Balance'])), float(balance_twd)], [1,2,3,4,5,6,7])
    summary[r['ProjectID']]['InvoiceBalanceTWD'] += balance_twd
    row += 1

sh = wb['Expenses']; row = 5
for r in sorted(expenses, key=lambda x: (x.get('ProjectID',''), x.get('ExpenseID',''))):
    if not usable('Expenses', r.get('ExpenseID'), r.get('ProjectID'), r.get('Currency')):
        continue
    amount_twd = twd(r['Amount'], r['Currency'], fx)
    write_row(sh, row, [r['ExpenseID'], r['ProjectID'], r['EmployeeID'], r['ExpenseDate'], r['Status'], r['Currency'], float(dec(r['Amount'])), float(amount_twd)], [1,2,3,4,5,6])
    summary[r['ProjectID']]['ApprovedExpenseTWD'] += amount_twd
    row += 1

sh = wb['Quotes']; row = 5
for r in sorted(quotes, key=lambda x: (x.get('ProjectID',''), x.get('QuoteID',''))):
    if not usable('Quotes', r.get('QuoteID'), r.get('ProjectID'), r.get('Currency')):
        continue
    amount_twd = twd(r['NetAmount'], r['Currency'], fx)
    write_row(sh, row, [r['QuoteID'], r['ProjectID'], r['CustomerID'], r['QuoteDate'], r.get('ValidUntil',''), r['Status'], r['Currency'], float(dec(r['NetAmount'])), float(amount_twd)], [1,2,3,4,5,6,7])
    summary[r['ProjectID']]['AcceptedQuoteTWD'] += amount_twd
    row += 1

sh = wb['Summary']; row = 5
for pid in sorted(summary):
    s = summary[pid]
    write_row(sh, row, [pid, float(s['OrderTWD']), float(s['InvoiceBalanceTWD']), float(s['ApprovedExpenseTWD']), float(s['AcceptedQuoteTWD'])], [1])
    row += 1
last = max(5, row - 1)
for col in ['B','C','D','E']:
    sh[f'{col}2'] = f'=SUM({col}5:{col}{last})'

sh = wb['Exceptions']; row = 5
for e in exceptions:
    write_row(sh, row, [e['Source'], e['DocumentID'], e['ProjectID'], e['Reason']], [1,2,3,4])
    row += 1

rep = wb['Report']
rep['B3'] = FROM
rep['D3'] = TO
rep['B4'] = 'SAP Sales Orders: Open/Closed only (non-Cancelled); SAP Invoices by invoice date; Expense Viewer: Approved only; quote_cli: Accepted only; integrated by ProjectID; FX from input/fx_rates.csv; orders and quotes are reported separately, not added as revenue.'
rep['B5'] = now
rep['B7'] = 'SAP web, Expense desktop, quote CLI, fx_rates.csv'
rep['D7'] = now
wb.save(output_xlsx)

# Reopen and validate saved outputs.
wb2 = load_workbook(output_xlsx, data_only=False)
headers = {sheet: [wb2[sheet].cell(4, c).value for c in range(1, wb2[sheet].max_column + 1)] for sheet in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']}
counts = {}
for sheet in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions']:
    sh = wb2[sheet]
    counts[sheet] = sum(1 for rr in range(5, sh.max_row + 1) if any(sh.cell(rr, cc).value is not None for cc in range(1, sh.max_column + 1)))
summary_headers = [h for h in headers['Summary'] if h]
if 'CustomerName' in summary_headers or 'Customer' in summary_headers:
    raise RuntimeError('Summary contains customer-based key; expected ProjectID only')
if any(str(h).lower() == 'revenue' for h in summary_headers if h):
    raise RuntimeError('Summary contains Revenue field; orders and quotes must remain separate')
expected_summary = ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']
if summary_headers[:5] != expected_summary:
    raise RuntimeError(f'Summary headers mismatch: {summary_headers[:5]} != {expected_summary}')
# Verify filtered statuses in saved details.
for rr in range(5, wb2['Orders'].max_row + 1):
    status = wb2['Orders'].cell(rr, 5).value
    if status == 'Cancelled':
        raise RuntimeError('Cancelled order found in output workbook')
for rr in range(5, wb2['Expenses'].max_row + 1):
    status = wb2['Expenses'].cell(rr, 5).value
    if status and status != 'Approved':
        raise RuntimeError('Non-Approved expense found in output workbook')
for rr in range(5, wb2['Quotes'].max_row + 1):
    status = wb2['Quotes'].cell(rr, 6).value
    if status and status != 'Accepted':
        raise RuntimeError('Non-Accepted quote found in output workbook')
validation = {
    'period': {'from_date': FROM, 'to_date': TO, 'period_label': PERIOD},
    'generated_at': now,
    'files': {
        'workbook': str(output_xlsx.relative_to(ws)),
        'sap_orders_open': str((downloads / f'sap_orders_open_{PERIOD}.csv').relative_to(ws)),
        'sap_orders_closed': str((downloads / f'sap_orders_closed_{PERIOD}.csv').relative_to(ws)),
        'sap_invoices': str((downloads / f'sap_invoices_{PERIOD}.csv').relative_to(ws)),
        'expenses': str((downloads / f'expenses_approved_{PERIOD}.csv').relative_to(ws)),
        'quotes_csv': str((downloads / f'quotes_accepted_{PERIOD}.csv').relative_to(ws)),
        'quotes_json': str((downloads / f'quotes_accepted_{PERIOD}.json').relative_to(ws))
    },
    'row_counts': counts,
    'source_counts_after_filters': {'orders': len(orders), 'invoices': len(invoices), 'expenses': len(expenses), 'quotes': len(quotes)},
    'checks': [
        'filled a copy of input/project_template.xlsx',
        'joined and summarized by ProjectID, not customer name',
        'excluded Cancelled orders by using Open and Closed SAP exports and rechecking status',
        'included only Approved expenses',
        'included only Accepted quotes',
        'kept OrderTWD, InvoiceBalanceTWD, ApprovedExpenseTWD, and AcceptedQuoteTWD separate',
        'used input/fx_rates.csv for TWD conversion',
        'reopened workbook and validated headers, statuses, and row counts'
    ],
    'exceptions_count': len(exceptions)
}
validation_path = out / f'validation_{PERIOD}.json'
validation_path.write_text(json.dumps(validation, ensure_ascii=False, indent=2), encoding='utf-8')
print('Validation:', json.dumps(validation, ensure_ascii=False, indent=2))
print('Created files:', [validation['files'][k] for k in validation['files']] + [str(validation_path.relative_to(ws))])
