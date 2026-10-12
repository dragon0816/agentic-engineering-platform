from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size)
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=ws/'input'/name
    print('\n---', name, '---')
    print(p.read_text(encoding='utf-8'))
print('\n--- fx_rates.csv ---')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rows=list(csv.reader(f))
    for r in rows[:10]: print(r)
print('\n--- project_template.xlsx ---')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
print('sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)
