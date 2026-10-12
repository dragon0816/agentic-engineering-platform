import os, json, csv, math
from pathlib import Path
from openpyxl import load_workbook

ws = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-011\X01\workspace')
out = ws/'output'
print('Inputs actually present:', sorted(p.name for p in (ws/'input').iterdir() if p.is_file()))
print('Outputs actually present:')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(ws), p.stat().st_size)

workbook = out/'project_integration_report_2026-09.xlsx'
validation = out/'validation_2026-09.json'
required_files = [
    workbook,
    validation,
    out/'downloads'/'sap_orders_open_2026-09.csv',
    out/'downloads'/'sap_orders_closed_2026-09.csv',
    out/'downloads'/'sap_invoices_2026-09.csv',
    out/'downloads'/'expenses_approved_2026-09.csv',
    out/'downloads'/'quotes_accepted_2026-09.csv',
    out/'downloads'/'quotes_accepted_2026-09.json',
]
checks=[]
for p in required_files:
    checks.append((f'file exists: {p.relative_to(ws)}', p.exists() and p.stat().st_size > 0))

wb = load_workbook(workbook, data_only=True)
expected_sheets = ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']
checks.append(('expected workbook sheets exist', all(s in wb.sheetnames for s in expected_sheets)))

def header_map(sheet_name, header_row=4):
    sh=wb[sheet_name]
    headers=[sh.cell(header_row,c).value for c in range(1, sh.max_column+1)]
    return {h:i+1 for i,h in enumerate(headers) if h is not None}, sh

summary_headers, summary = header_map('Summary')
required_summary=['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']
checks.append(('Summary row 4 has required separate fields', all(c in summary_headers for c in required_summary)))
checks.append(('Summary has no CustomerID join key field', 'CustomerID' not in summary_headers))
combined_revenue = [h for h in summary_headers if str(h).strip().lower().replace(' ','') in {'revenue','totalrevenue','營收','totalincome'}]
checks.append(('Summary has no combined revenue column', not combined_revenue))
summary_projects=[]
summary_totals={k:0.0 for k in required_summary[1:]}
for r in range(5, summary.max_row+1):
    pid=summary.cell(r, summary_headers['ProjectID']).value
    if pid in (None,''):
        continue
    summary_projects.append(str(pid))
    for k in required_summary[1:]:
        v=summary.cell(r, summary_headers[k]).value
        checks.append((f'Summary {k} numeric at row {r}', isinstance(v,(int,float))))
        summary_totals[k]+=float(v or 0)
checks.append(('Summary has project IDs only, not customer names', all(str(p).startswith('P') for p in summary_projects)))
checks.append(('Summary contains 5 projects', len(summary_projects)==5))

# Detail sheet status/date/header validation
detail_specs = {
    'Orders': {'headers':['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'], 'status_col':'Status', 'allowed_status':{'Open','Closed'}, 'date_col':'OrderDate'},
    'Invoices': {'headers':['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], 'date_col':'InvoiceDate'},
    'Expenses': {'headers':['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'], 'status_col':'Status', 'allowed_status':{'Approved'}, 'date_col':'ExpenseDate'},
    'Quotes': {'headers':['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'], 'status_col':'Status', 'allowed_status':{'Accepted'}, 'date_col':'QuoteDate'},
}
computed = {'OrderTWD':{}, 'InvoiceBalanceTWD':{}, 'ApprovedExpenseTWD':{}, 'AcceptedQuoteTWD':{}}
row_counts = {}
for sheet_name, spec in detail_specs.items():
    hm, sh = header_map(sheet_name)
    checks.append((f'{sheet_name} row 4 headers match expected', all(h in hm for h in spec['headers'])))
    count=0
    for r in range(5, sh.max_row+1):
        first=sh.cell(r,1).value
        if first in (None,''):
            continue
        count += 1
        if 'status_col' in spec:
            status=sh.cell(r, hm[spec['status_col']]).value
            checks.append((f'{sheet_name} row {r} status allowed', status in spec['allowed_status']))
        date=str(sh.cell(r, hm[spec['date_col']]).value)
        checks.append((f'{sheet_name} row {r} date in Sep 2026', '2026-09-01' <= date <= '2026-09-30'))
        pid=sh.cell(r, hm['ProjectID']).value
        checks.append((f'{sheet_name} row {r} ProjectID present', pid not in (None,'')))
        # all included detail project ids are ProjectID, not customer name
        checks.append((f'{sheet_name} row {r} ProjectID format', str(pid).startswith('P')))
        if sheet_name=='Orders':
            amount=sh.cell(r, hm['AmountTWD']).value or 0
            computed['OrderTWD'].setdefault(pid,0.0); computed['OrderTWD'][pid]+=float(amount)
        elif sheet_name=='Invoices':
            amount=sh.cell(r, hm['BalanceTWD']).value or 0
            computed['InvoiceBalanceTWD'].setdefault(pid,0.0); computed['InvoiceBalanceTWD'][pid]+=float(amount)
        elif sheet_name=='Expenses':
            amount=sh.cell(r, hm['AmountTWD']).value or 0
            computed['ApprovedExpenseTWD'].setdefault(pid,0.0); computed['ApprovedExpenseTWD'][pid]+=float(amount)
        elif sheet_name=='Quotes':
            amount=sh.cell(r, hm['AmountTWD']).value or 0
            computed['AcceptedQuoteTWD'].setdefault(pid,0.0); computed['AcceptedQuoteTWD'][pid]+=float(amount)
    row_counts[sheet_name]=count

for metric in required_summary[1:]:
    for pid in summary_projects:
        sv=summary.cell(5+summary_projects.index(pid), summary_headers[metric]).value or 0
        cv=computed[metric].get(pid,0.0)
        checks.append((f'Summary {metric} equals detail sum for {pid}', abs(float(sv)-cv) < 0.01))

# Exceptions validation
exc_headers, exc = header_map('Exceptions')
checks.append(('Exceptions row 4 has Source/DocumentID/ProjectID/Reason', all(c in exc_headers for c in ['Source','DocumentID','ProjectID','Reason'])))
exc_count=0
valid_reasons=True
for r in range(5, exc.max_row+1):
    source=exc.cell(r, exc_headers['Source']).value if 'Source' in exc_headers else None
    doc=exc.cell(r, exc_headers['DocumentID']).value if 'DocumentID' in exc_headers else None
    reason=exc.cell(r, exc_headers['Reason']).value if 'Reason' in exc_headers else ''
    if source in (None,'') and doc in (None,''):
        continue
    exc_count += 1
    if not (('Missing ProjectID' in str(reason)) or ('Missing FX' in str(reason)) or ('excluded' in str(reason))):
        valid_reasons=False
checks.append(('Exceptions rows are missing ProjectID/FX exclusion notes', valid_reasons))

# Validation JSON content
vj=json.load(open(validation, encoding='utf-8'))
checks.append(('Validation period is 2026-09-01 to 2026-09-30', vj.get('period',{}).get('from_date')=='2026-09-01' and vj.get('period',{}).get('to_date')=='2026-09-30' and vj.get('period',{}).get('period_label')=='2026-09'))
for key, expected in row_counts.items():
    checks.append((f'Validation row count for {key}', vj.get('row_counts',{}).get(key)==expected))
checks.append(('Validation exception count matches workbook', vj.get('exceptions_count')==exc_count))
checks_text=' | '.join(vj.get('checks',[]))
checks.append(('Validation states ProjectID, not customer name', 'ProjectID' in checks_text and 'not customer name' in checks_text))
checks.append(('Validation states orders and quotes kept separate', 'separate' in checks_text and 'OrderTWD' in checks_text and 'AcceptedQuoteTWD' in checks_text))

print('\nWorkbook summary projects:', summary_projects)
print('Detail row counts:', row_counts)
print('Exception count:', exc_count)
print('Computed totals:', computed)
print('\nChecks:')
failed=[]
for name, ok in checks:
    print(('PASS' if ok else 'FAIL'), '-', name)
    if not ok:
        failed.append(name)
if failed:
    print('FAILED CHECKS:', failed)
    raise SystemExit(1)
print('ALL_CHECKS_PASSED')