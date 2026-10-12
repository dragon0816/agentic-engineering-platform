from pathlib import Path
from openpyxl import load_workbook
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
import json, csv

ws = Path.cwd()
out = ws/'output'
xlsx = out/'sent_quotes_2026-09.xlsx'
log_path = out/'quote_cli_run_log.json'
fx_path = ws/'input'/'fx_rates.csv'

print('fixing existing output without calling CLI again')
print('workbook:', xlsx, xlsx.exists(), xlsx.stat().st_size if xlsx.exists() else None)
print('log:', log_path, log_path.exists(), log_path.stat().st_size if log_path.exists() else None)

wb = load_workbook(xlsx)
qws = wb['Quotes']
# Convert QuoteDate / ValidUntil strings to real Excel dates and DiscountPercent strings to numbers.
for row in range(5, qws.max_row + 1):
    if not any(qws.cell(row, col).value is not None for col in range(1, 12)):
        continue
    for col in (4, 5):
        c = qws.cell(row, col)
        if isinstance(c.value, str) and len(c.value) == 10:
            y, m, d = map(int, c.value.split('-'))
            c.value = date(y, m, d)
        c.number_format = 'yyyy-mm-dd'
    c = qws.cell(row, 9)
    if isinstance(c.value, str):
        c.value = float(Decimal(c.value))
    c.number_format = '0.00'
    for col in (8, 10, 11):
        qws.cell(row, col).number_format = '0.00'
    for col in (1, 2, 3):
        qws.cell(row, col).number_format = '@'

lws = wb['Lines']
for row in range(5, lws.max_row + 1):
    if not any(lws.cell(row, col).value is not None for col in range(1, 7)):
        continue
    for col in (1, 2):
        lws.cell(row, col).number_format = '@'
    for col in (5, 6):
        lws.cell(row, col).number_format = '0.00'

wb.save(xlsx)
print('saved fixed workbook')

# Verification: reopen every output file.
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    next(r)
    for row in r:
        if len(row) >= 2 and row[0]:
            fx[row[0]] = Decimal(row[1])

def money(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

log = json.loads(log_path.read_text(encoding='utf-8'))
assert log['filter'] == {'from': '2026-09-01', 'to': '2026-09-30', 'status': 'Sent'}
assert log['quote_count'] == 4 and log['line_count'] == 8
list_attempts = [rec for rec in log['commands'] if rec['command'][1:3] == ['quotes','list']]
assert len(list_attempts) == 2, len(list_attempts)
assert list_attempts[0]['exit_code'] == 4 and list_attempts[0]['attempt'] == 1 and list_attempts[0]['stderr']
assert list_attempts[1]['exit_code'] == 0 and list_attempts[1]['attempt'] == 2
for rec in log['commands']:
    if rec['exit_code'] == 0:
        assert rec['stdout_length'] >= 0
    elif rec['exit_code'] == 4:
        assert rec['attempt'] <= 3 and rec['stderr']
    else:
        raise AssertionError(f'unexpected exit code; exit 2/3 should have aborted without retry: {rec}')

wb2 = load_workbook(xlsx, data_only=False)
assert wb2.sheetnames == ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
assert [str(rng) for rng in wb2['Report'].merged_cells.ranges] == ['A1:F1']
r = wb2['Report']
assert r['B3'].value == '2026-09-01'
assert r['D3'].value == '2026-09-30'
assert 'Sent' in r['B4'].value
assert r['B5'].value == '2026-10-06T00:00:00+08:00'
assert r['B7'].value == 'quote_cli; fx_rates.csv'
assert r['D7'].value == '2026-10-06T00:00:00+08:00'

qws = wb2['Quotes']
quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [qws.cell(4,c).value for c in range(1,12)] == quote_cols
qrows = [row for row in qws.iter_rows(min_row=5, max_col=11, values_only=True) if any(v is not None for v in row)]
assert len(qrows) == 4
quote_ids = set()
for row in qrows:
    qid, cust, proj, qdate, valid, status, cur, subtotal, disc, net, twd = row
    assert isinstance(qid, str) and qid.startswith('Q')
    assert isinstance(cust, str) and cust.startswith('C')
    assert proj is None or isinstance(proj, str)
    assert getattr(qdate, 'isoformat', lambda: '')().startswith('2026-09-')
    assert getattr(valid, 'isoformat', lambda: '')().count('-') == 2
    assert status == 'Sent'
    assert cur in fx
    for val in (subtotal, disc, net, twd):
        assert isinstance(val, (int, float)), (qid, val, type(val))
    expected_twd = float(money(Decimal(str(net)) * fx[cur]))
    assert abs(twd - expected_twd) < 0.001, (qid, twd, expected_twd)
    quote_ids.add(qid)

lws = wb2['Lines']
line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [lws.cell(4,c).value for c in range(1,7)] == line_cols
lrows = [row for row in lws.iter_rows(min_row=5, max_col=6, values_only=True) if any(v is not None for v in row)]
assert len(lrows) == 8
for row in lrows:
    qid, pid, pname, qty, unit, amount = row
    assert qid in quote_ids
    assert isinstance(qid, str) and isinstance(pid, str)
    assert isinstance(pname, str)
    assert isinstance(qty, (int, float))
    assert isinstance(unit, (int, float)) and isinstance(amount, (int, float))
    assert abs(float(qty) * float(unit) - float(amount)) < 0.001

sws = wb2['Summary']
assert sws['B2'].value == '=SUM(Quotes!K5:K8)'
srows = [row for row in sws.iter_rows(min_row=5, max_col=6, values_only=True) if any(v is not None for v in row)]
assert srows
assert sum(row[5] for row in srows) == len(qrows)
for row in srows:
    group, subgroup, cur, amount, twd, count = row
    assert group == 'Status'
    assert subgroup == 'Sent'
    assert cur in fx
    assert isinstance(amount, (int, float)) and isinstance(twd, (int, float)) and isinstance(count, int)

print('verification OK: reopened workbook and JSON log; checked files, template sheets/merge, report period/status/source/time, 4 Quotes rows, 8 Lines rows, blank missing values preserved, IDs text, dates as Excel dates, DiscountPercent and amounts numeric, FX TWD values, Summary formula/row counts, and CLI retry semantics logged (exit 4 retried once then success; no exit 2/3).')
print('outputs:')
for p in sorted(out.glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
