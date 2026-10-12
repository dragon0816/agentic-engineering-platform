from pathlib import Path
import subprocess, json, csv, shutil, sys, time, os
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
CLI = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
OUT_XLSX = OUTPUT / 'sent_quotes_2026-09.xlsx'
LOG_JSON = OUTPUT / 'quote_cli_run_log.json'

print('inputs really are:')
print(' workspace:', WORKSPACE)
print(' template:', TEMPLATE, TEMPLATE.exists())
print(' fx:', FX_FILE, FX_FILE.exists())
print(' cli:', CLI, Path(CLI).exists())
print(' filter:', FROM_DATE, TO_DATE, STATUS)

# Helpers
TWOPLACES = Decimal('0.01')
def money(x):
    if x is None or x == '':
        return None
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def excel_num(d):
    if d is None:
        return None
    return float(money(d))

def copy_row_format(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        s = ws.cell(src_row, col)
        d = ws.cell(dst_row, col)
        if s.has_style:
            d._style = copy(s._style)
        if s.number_format:
            d.number_format = s.number_format
        if s.alignment:
            d.alignment = copy(s.alignment)
        if s.font:
            d.font = copy(s.font)
        if s.fill:
            d.fill = copy(s.fill)
        if s.border:
            d.border = copy(s.border)

# CLI runner with requested retry semantics.
cli_log = []
def run_cli(args, expect_json=True):
    attempts = 0
    max_retries = 2
    while True:
        attempts += 1
        cmd = [CLI] + args
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True, shell=False)
        rec = {
            'command': cmd,
            'attempt': attempts,
            'exit_code': proc.returncode,
            'stderr': proc.stderr.strip(),
            'stdout_length': len(proc.stdout or ''),
        }
        cli_log.append(rec)
        print(f"CLI attempt {attempts}: exit={proc.returncode} args={' '.join(args)} stderr={proc.stderr.strip()!r}")
        if proc.returncode == 0:
            if not expect_json:
                return proc.stdout
            try:
                data = json.loads(proc.stdout)
            except Exception as e:
                raise RuntimeError(f"CLI stdout was not parseable JSON on success for args {args}: {e}") from e
            return data
        if proc.returncode == 4 and attempts <= 1 + max_retries:
            if attempts <= max_retries:
                time.sleep(0.2)
                continue
        if proc.returncode in (2, 3):
            raise RuntimeError(f"CLI returned exit code {proc.returncode}; not retried. stderr: {proc.stderr.strip()}")
        if proc.returncode == 4:
            raise RuntimeError(f"CLI returned exit code 4 after {attempts} attempts. stderr: {proc.stderr.strip()}")
        raise RuntimeError(f"CLI returned unexpected exit code {proc.returncode}. stderr: {proc.stderr.strip()}")

# Load FX rates; the Note column contains unquoted commas in the supplied fixture, so use first two columns only.
fx_rates = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx_rates[row[0]] = Decimal(row[1])
print('fx rates:', {k: str(v) for k, v in fx_rates.items()})

# Fetch matching quote headers and detail lines.
list_data = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'])
if list_data.get('schema_version') != 1 or 'items' not in list_data:
    raise RuntimeError(f"Unexpected quote list schema: keys={list(list_data.keys())}")
headers = list_data['items']
print('quote count:', len(headers), 'declared:', list_data.get('count'))
if list_data.get('count') != len(headers):
    raise RuntimeError(f"Quote list count mismatch: declared {list_data.get('count')} vs items {len(headers)}")

quote_details = {}
all_lines = []
for h in headers:
    qid = h['QuoteID']
    detail = run_cli(['quotes', 'get', qid, '--format', 'json'])
    if detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
        raise RuntimeError(f"Unexpected quote get schema for {qid}: keys={list(detail.keys())}")
    quote_details[qid] = detail
    for line in detail['line_items']:
        rec = dict(line)
        rec['QuoteID'] = qid
        all_lines.append(rec)
print('line count:', len(all_lines))

# Fill a copy of the template.
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

# Report sheet.
report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = f'Status={STATUS}; QuoteDate {FROM_DATE} to {TO_DATE}'
report['B5'] = GENERATED_AT
report['B7'] = 'quote_cli; fx_rates.csv'
report['D7'] = GENERATED_AT
notes = []
missing_fx = sorted({h.get('Currency') for h in headers if h.get('Currency') not in fx_rates})
if missing_fx:
    notes.append('Missing FX rate: ' + ', '.join(missing_fx))
if not headers:
    notes.append('No matching quotes returned by CLI.')
report['B9'] = '; '.join(notes) if notes else ''

# Clear previous data area if any, then write Quotes rows from row 5.
quotes_ws = wb['Quotes']
for row in quotes_ws.iter_rows(min_row=5, max_row=max(quotes_ws.max_row, 5), max_col=11):
    for cell in row:
        cell.value = None
quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
for idx, h in enumerate(headers, start=5):
    copy_row_format(quotes_ws, 5, idx, 11)
    cur = h.get('Currency')
    net = money(h.get('NetAmount'))
    rate = fx_rates.get(cur)
    twd = money(net * rate) if net is not None and rate is not None else None
    values = [
        h.get('QuoteID'), h.get('CustomerID'), h.get('ProjectID'), h.get('QuoteDate'), h.get('ValidUntil'),
        h.get('Status'), cur, excel_num(h.get('Subtotal')), h.get('DiscountPercent'), excel_num(h.get('NetAmount')), excel_num(twd)
    ]
    for col, val in enumerate(values, start=1):
        c = quotes_ws.cell(idx, col, val)
        if col in (1,2,3,4,5,6,7):
            c.number_format = '@' if col in (1,2,3) else 'yyyy-mm-dd' if col in (4,5) else 'General'
        if col in (8,10,11):
            c.number_format = '0.00'
        if col == 9:
            c.number_format = '0.00'

# Lines sheet.
lines_ws = wb['Lines']
for row in lines_ws.iter_rows(min_row=5, max_row=max(lines_ws.max_row, 5), max_col=6):
    for cell in row:
        cell.value = None
line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
for idx, line in enumerate(all_lines, start=5):
    copy_row_format(lines_ws, 5, idx, 6)
    values = [line.get('QuoteID'), line.get('ProductID'), line.get('ProductName'), line.get('Quantity'), excel_num(line.get('UnitPrice')), excel_num(line.get('LineAmount'))]
    for col, val in enumerate(values, start=1):
        c = lines_ws.cell(idx, col, val)
        if col in (1,2):
            c.number_format = '@'
        if col in (5,6):
            c.number_format = '0.00'

# Summary: group by Status and Currency, while keeping currencies separate.
summary_ws = wb['Summary']
for row in summary_ws.iter_rows(min_row=5, max_row=max(summary_ws.max_row, 5), max_col=6):
    for cell in row:
        cell.value = None
summary = {}
for h in headers:
    key = (h.get('Status') or '', h.get('Currency') or '')
    net = money(h.get('NetAmount')) or Decimal('0')
    rate = fx_rates.get(h.get('Currency'))
    twd = money(net * rate) if rate is not None else None
    if key not in summary:
        summary[key] = {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0, 'missing_twd': False}
    summary[key]['amount'] += net
    if twd is None:
        summary[key]['missing_twd'] = True
    else:
        summary[key]['twd'] += twd
    summary[key]['count'] += 1
for idx, ((status, cur), agg) in enumerate(sorted(summary.items()), start=5):
    copy_row_format(summary_ws, 5, idx, 6)
    values = ['Status', status, cur, excel_num(agg['amount']), None if agg['missing_twd'] else excel_num(agg['twd']), agg['count']]
    for col, val in enumerate(values, start=1):
        c = summary_ws.cell(idx, col, val)
        if col in (4,5):
            c.number_format = '0.00'
last_quote_row = 4 + len(headers) if headers else 5
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'

# Save workbook and log.
wb.save(OUT_XLSX)
LOG_JSON.write_text(json.dumps({'filter': {'from': FROM_DATE, 'to': TO_DATE, 'status': STATUS}, 'commands': cli_log, 'quote_count': len(headers), 'line_count': len(all_lines)}, ensure_ascii=False, indent=2), encoding='utf-8')
print('saved:', OUT_XLSX)
print('saved:', LOG_JSON)

# Verification: reopen every output and confirm stated requirements.
wb2 = load_workbook(OUT_XLSX, data_only=False)
assert wb2.sheetnames == ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions'], wb2.sheetnames
r = wb2['Report']
assert r['B3'].value == FROM_DATE and r['D3'].value == TO_DATE
assert STATUS in r['B4'].value
assert r['B5'].value == GENERATED_AT and r['D7'].value == GENERATED_AT
assert r['B7'].value == 'quote_cli; fx_rates.csv'
qws = wb2['Quotes']
assert [qws.cell(4,c).value for c in range(1,12)] == quote_cols
actual_qrows = []
for row in qws.iter_rows(min_row=5, max_col=11, values_only=True):
    if any(v is not None for v in row): actual_qrows.append(row)
assert len(actual_qrows) == len(headers), (len(actual_qrows), len(headers))
for row in actual_qrows:
    assert isinstance(row[0], str) and isinstance(row[1], str) and isinstance(row[2], str), row[:3]
    assert row[5] == STATUS, row
    assert isinstance(row[7], (int, float)) and isinstance(row[9], (int, float)) and isinstance(row[10], (int, float)), row
    cur = row[6]
    expected_twd = float(money(Decimal(str(row[9])) * fx_rates[cur]))
    assert abs(row[10] - expected_twd) < 0.001, (row[0], row[10], expected_twd)
lws = wb2['Lines']
assert [lws.cell(4,c).value for c in range(1,7)] == line_cols
actual_lrows = []
for row in lws.iter_rows(min_row=5, max_col=6, values_only=True):
    if any(v is not None for v in row): actual_lrows.append(row)
assert len(actual_lrows) == len(all_lines), (len(actual_lrows), len(all_lines))
quote_ids = {h['QuoteID'] for h in headers}
for row in actual_lrows:
    assert row[0] in quote_ids, row
    assert isinstance(row[0], str) and isinstance(row[1], str), row[:2]
    assert isinstance(row[3], (int, float)) and isinstance(row[4], (int, float)) and isinstance(row[5], (int, float)), row
sws = wb2['Summary']
assert sws['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})'
summary_rows = [row for row in sws.iter_rows(min_row=5, max_col=6, values_only=True) if any(v is not None for v in row)]
assert len(summary_rows) == len(summary)
log = json.loads(LOG_JSON.read_text(encoding='utf-8'))
assert log['quote_count'] == len(headers) and log['line_count'] == len(all_lines)
for rec in log['commands']:
    if rec['exit_code'] == 0:
        assert rec['stdout_length'] >= 0
print('verification OK: workbook reopened; sheet names, report cells, headers, row counts, text IDs, numeric amounts, FX-derived TWD values, Summary formula/rows, and JSON log checked.')
print('output files:')
for p in sorted(OUTPUT.glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
