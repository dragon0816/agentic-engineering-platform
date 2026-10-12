from pathlib import Path
from openpyxl import load_workbook
import csv

ws = Path.cwd()
print('workspace=', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

manual = ws/'input'/'quote_cli_manual.md'
print('\n--- quote_cli_manual.md first 200 lines ---')
if manual.exists():
    text = manual.read_text(encoding='utf-8', errors='replace')
    for i, line in enumerate(text.splitlines()[:200], 1):
        print(f'{i:03d}: {line}')
else:
    print('missing')

fx = ws/'input'/'fx_rates.csv'
print('\n--- fx_rates.csv headers/sample ---')
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i, row in zip(range(8), r):
            print(row)
else:
    print('missing')

tpl = ws/'input'/'quote_template.xlsx'
print('\n--- template workbook ---')
if tpl.exists():
    wb = load_workbook(tpl, data_only=False)
    print('sheets=', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print(f'[{s}] max_row={sh.max_row} max_col={sh.max_column}')
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(row)
else:
    print('missing')
