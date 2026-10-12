from pathlib import Path
from openpyxl import load_workbook
import json, csv
from decimal import Decimal, ROUND_HALF_UP

ws = Path.cwd()
out = ws/'output'
xlsx = out/'sent_quotes_2026-09.xlsx'
log_path = out/'quote_cli_run_log.json'
fx_path = ws/'input'/'fx_rates.csv'
print('checking output files:')
for p in [xlsx, log_path]:
    print(' -', p, 'exists=', p.exists(), 'bytes=', p.stat().st_size if p.exists() else None)

fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    next(r)
    for row in r:
        if len(row) >= 2 and row[0]:
            fx[row[0]] = Decimal(row[1])

def money(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

log = json.loads(log_path.read_text(encoding='utf-8'))
print('log filter:', log['filter'])
print('log quote_count:', log['quote_count'], 'line_count:', log['line_count'])
for rec in log['commands']:
    print(' command attempt:', ' '.join(rec['command'][1:]), 'attempt', rec['attempt'], 'exit', rec['exit_code'], 'stderr=', rec['stderr'])
    assert rec['exit_code'] in (0,4)
    if rec['exit_code'] == 4:
        assert 'temporarily unavailable' in rec['stderr'] or rec['stderr']

wb = load_workbook(xlsx, data_only=False)
print('sheets:', wb.sheetnames)
assert wb.sheetnames == ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
assert list(wb['Report'].merged_cells.ranges) == ['A1:F1']

r = wb['Report']
assert r['B3'].value == '2026-09-01'
assert r['D3'].value == '2026-09-30'
assert 'Sent' in r['B4'].value
assert r['B5'].value == '2026-10-06T00:00:00+08:00'
assert r['B7'].value == 'quote_cli; fx_rates.csv'
assert r['D7'].value == '2026-10-06T00:00:00+08:00'

qws = wb['Quotes']
quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [qws.cell(4,c).value for c in range(1,12)] == quote_cols
qrows = [row for row in qws.iter_rows(min_row=5, max_col=11, values_only=True) if any(v is not None for v in row)]
print('quote rows:', len(qrows))
assert len(qrows) == log['quote_count'] == 4
quote_ids = set()
for row in qrows:
    qid, cust, proj, qdate, valid, status, cur, subtotal, disc, net, twd = row
    assert isinstance(qid, str) and qid.startswith('Q')
    assert isinstance(cust, str) and cust.startswith('C')
    assert proj is None or isinstance(proj, str)  # missing source values remain blank
    assert isinstance(qdate, str) and qdate.startswith('2026-09-')
    assert isinstance(valid, str) and len(valid) == 10
    assert status == 'Sent'
    assert cur in fx
    for val in (subtotal, disc, net, twd):
        assert isinstance(val, (int, float)), (qid, val, type(val))
    expected_twd = float(money(Decimal(str(net)) * fx[cur]))
    assert abs(twd - expected_twd) < 0.001, (qid, twd, expected_twd)
    quote_ids.add(qid)

lws = wb['Lines']
line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [lws.cell(4,c).value for c in range(1,7)] == line_cols
lrows = [row for row in lws.iter_rows(min_row=5, max_col=6, values_only=True) if any(v is not None for v in row)]
print('line rows:', len(lrows))
assert len(lrows) == log['line_count'] == 8
for row in lrows:
    qid, pid, pname, qty, unit, amount = row
    assert qid in quote_ids
    assert isinstance(qid, str) and isinstance(pid, str)
    assert isinstance(pname, str)
    assert isinstance(qty, (int, float))
    assert isinstance(unit, (int, float)) and isinstance(amount, (int, float))
    assert abs(float(qty) * float(unit) - float(amount)) < 0.001

sws = wb['Summary']
assert sws['B2'].value == '=SUM(Quotes!K5:K8)'
srows = [row for row in sws.iter_rows(min_row=5, max_col=6, values_only=True) if any(v is not None for v in row)]
print('summary rows:', srows)
assert len(srows) >= 1
for row in srows:
    group, subgroup, cur, amount, twd, count = row
    assert group == 'Status'
    assert subgroup == 'Sent'
    assert cur in fx
    assert isinstance(amount, (int, float)) and isinstance(twd, (int, float)) and isinstance(count, int)

print('verification OK: reopened workbook and log; checked exact output files, template sheets/merge preserved, report period/status/source/generated time, Quotes headers and 4 Sent Sep-2026 rows, blank missing ProjectID allowed, text IDs, numeric money fields, FX-derived NetAmountTWD, Lines headers and 8 joined line rows, Summary formula and grouped rows.')
