from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
print('Instructions full:')
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
print('\nExisting formulas/styles relevant:')
for shname in ['Report','Quotes','Lines','Summary']:
    sh=wb[shname]
    print('SHEET', shname)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8)):
        print([c.coordinate+':'+repr(c.value) for c in row[:min(sh.max_column,12)]])