from pathlib import Path
import subprocess, json, csv, shutil, os
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date, timezone, timedelta
from copy import copy
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUT_FILE = OUTPUT / 'quote_report_P001-P005_2026-09_Accepted.xlsx'
CLI = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
PROJECTS = [f'P{i:03d}' for i in range(1, 6)]
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Accepted'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
FETCHED_AT = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()

TWOPLACES = Decimal('0.01')
def q2(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def dnum(x):
    return float(q2(x))

# Read FX rates, using first two columns because Note contains an unquoted comma in the supplied CSV.
rates = {}
with open(FX_FILE, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or len(row) < 2:
            continue
        cur = row[0].strip()
        if cur:
            rates[cur] = Decimal(row[1].strip())
print('FX rates:', {k: str(v) for k, v in rates.items()})

# Fetch quote headers and details.
headers = []
line_rows = []
cli_calls = []
for project in PROJECTS:
    cmd = [CLI, 'quotes', 'list', '--project', project, '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
    cp = subprocess.run(cmd, capture_output=True, text=True, shell=False)
    cli_calls.append(('list', project, cp.returncode, cp.stderr.strip()))
    if cp.returncode != 0:
        raise RuntimeError(f'quote_cli list failed for {project}: exit {cp.returncode}; stderr={cp.stderr}')
    data = json.loads(cp.stdout)
    if data.get('schema_version') != 1:
        raise RuntimeError(f'Unexpected schema_version for list {project}: {data.get("schema_version")}')
    for item in data.get('items', []):
        headers.append(item)

# Sort deterministically by ProjectID, QuoteDate, QuoteID.
headers.sort(key=lambda h: (h.get('ProjectID',''), h.get('QuoteDate',''), h.get('QuoteID','')))

for h in headers:
    qid = h['QuoteID']
    cmd = [CLI, 'quotes', 'get', qid, '--format', 'json']
    cp = subprocess.run(cmd, capture_output=True, text=True, shell=False)
    cli_calls.append(('get', qid, cp.returncode, cp.stderr.strip()))
    if cp.returncode != 0:
        raise RuntimeError(f'quote_cli get failed for {qid}: exit {cp.returncode}; stderr={cp.stderr}')
    doc = json.loads(cp.stdout)
    if doc.get('schema_version') != 1:
        raise RuntimeError(f'Unexpected schema_version for get {qid}: {doc.get("schema_version")}')
    # Use get header as authoritative but verify it matches list keys.
    gh = doc['header']
    for key in ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount']:
        if str(gh.get(key)) != str(h.get(key)):
            raise RuntimeError(f'Header mismatch for {qid} key {key}: list={h.get(key)} get={gh.get(key)}')
    for li in doc.get('line_items', []):
        line_rows.append({'QuoteID': qid, **li})

print('Fetched headers:', len(headers), headers)
print('Fetched line rows:', len(line_rows), line_rows)

# Prepare workbook from template copy.
shutil.copy2(TEMPLATE, OUT_FILE)
wb = load_workbook(OUT_FILE)

# Helpers.
def clear_data(sh, min_row, max_col):
    if sh.max_row >= min_row:
        for row in sh.iter_rows(min_row=min_row, max_row=sh.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

def copy_row_style(sh, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = sh.cell(src_row, col)
        dst = sh.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.protection:
            dst.protection = copy(src.protection)

# Report sheet.
report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = 'ProjectID P001-P005; QuoteDate 2026-09-01..2026-09-30; Status Accepted; NetAmountTWD by input/fx_rates.csv'
report['B5'] = GENERATED_AT
report['B7'] = 'quote_cli_mock; fx_rates.csv'
report['D7'] = FETCHED_AT
report['B9'] = 'P003 and P005 have no Accepted quotes in the requested period; Summary keeps one 0 row for each.'

# Quotes sheet.
quotes = wb['Quotes']
clear_data(quotes, 5, 11)
for idx, h in enumerate(headers, start=5):
    if idx > 5:
        copy_row_style(quotes, 5, idx, 11)
    currency = h['Currency']
    if currency not in rates:
        raise RuntimeError(f'Missing FX rate for {currency}')
    net = q2(h['NetAmount'])
    net_twd = q2(net * rates[currency])
    row_values = [
        h['QuoteID'], h['CustomerID'], h['ProjectID'], h['QuoteDate'], h['ValidUntil'], h['Status'], h['Currency'],
        dnum(h['Subtotal']), dnum(h['DiscountPercent']), dnum(net), dnum(net_twd)
    ]
    for col, val in enumerate(row_values, start=1):
        c = quotes.cell(idx, col, val)
        if col in [1,2,3,6,7]:
            c.number_format = '@'
        elif col in [4,5]:
            c.number_format = 'yyyy-mm-dd'
        elif col in [8,10,11]:
            c.number_format = '0.00'

# Lines sheet.
lines = wb['Lines']
clear_data(lines, 5, 6)
for idx, li in enumerate(line_rows, start=5):
    if idx > 5:
        copy_row_style(lines, 5, idx, 6)
    row_values = [li['QuoteID'], li['ProductID'], li.get('ProductName'), li.get('Quantity'), dnum(li.get('UnitPrice')), dnum(li.get('LineAmount'))]
    for col, val in enumerate(row_values, start=1):
        c = lines.cell(idx, col, val)
        if col in [1,2,3]:
            c.number_format = '@'
        elif col in [5,6]:
            c.number_format = '0.00'

# Summary by ProjectID, not mixing original currencies. Since each project has one row here, Currency shows quote currency;
# no-quote projects get blank currency, original amount 0, TWD amount 0, count 0.
summary = wb['Summary']
clear_data(summary, 5, 6)
by_project = {p: {'currency': None, 'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0} for p in PROJECTS}
# If a project had multiple currencies, keep original Amount blank to avoid mixed-currency arithmetic while still summing TWD.
project_currencies = {p: set() for p in PROJECTS}
for h in headers:
    p = h['ProjectID']
    cur = h['Currency']
    amt = q2(h['NetAmount'])
    amt_twd = q2(amt * rates[cur])
    by_project[p]['amount'] += amt
    by_project[p]['amount_twd'] += amt_twd
    by_project[p]['count'] += 1
    project_currencies[p].add(cur)

for idx, p in enumerate(PROJECTS, start=5):
    if idx > 5:
        copy_row_style(summary, 5, idx, 6)
    currencies = sorted(project_currencies[p])
    currency_cell = currencies[0] if len(currencies) == 1 else ('' if len(currencies) == 0 else 'MULTI')
    amount_val = dnum(by_project[p]['amount']) if len(currencies) <= 1 else None
    vals = [p, '', currency_cell, amount_val if amount_val is not None else None, dnum(by_project[p]['amount_twd']), by_project[p]['count']]
    for col, val in enumerate(vals, start=1):
        c = summary.cell(idx, col, val)
        if col in [1,2,3]:
            c.number_format = '@'
        elif col in [4,5]:
            c.number_format = '0.00'
# Extend formula range.
last_quote_row = 4 + max(len(headers), 1)
summary['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
summary['B2'].number_format = '0.00'

# Freeze panes / filters preserved; set autofilter refs if present? leave as template.
wb.save(OUT_FILE)
print('Saved:', OUT_FILE)

# Validation by reopening output.
wb2 = load_workbook(OUT_FILE, data_only=False)
errors = []
if wb2.sheetnames != ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']:
    errors.append(f'sheet names changed: {wb2.sheetnames}')
q = wb2['Quotes']; l = wb2['Lines']; s = wb2['Summary']; r = wb2['Report']
expected_q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
actual_q_headers = [q.cell(4,c).value for c in range(1,12)]
if actual_q_headers != expected_q_headers:
    errors.append(f'Quotes headers mismatch: {actual_q_headers}')
expected_l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if [l.cell(4,c).value for c in range(1,7)] != expected_l_headers:
    errors.append('Lines headers mismatch')
expected_s_headers = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
if [s.cell(4,c).value for c in range(1,7)] != expected_s_headers:
    errors.append('Summary headers mismatch')
# Row counts and filter correctness.
quote_rows = []
for row in range(5, 5+len(headers)):
    quote_rows.append([q.cell(row,c).value for c in range(1,12)])
if len(quote_rows) != 3:
    errors.append(f'Expected 3 quote rows, got {len(quote_rows)}')
for row in quote_rows:
    qid,cust,proj,qdate,valid,status,cur,subtotal,disc,net,nettwd = row
    if proj not in PROJECTS: errors.append(f'Unexpected project {proj}')
    if not (FROM_DATE <= qdate <= TO_DATE): errors.append(f'Date out of range {qdate}')
    if status != STATUS: errors.append(f'Unexpected status {status}')
    exp_twd = dnum(q2(Decimal(str(net)) * rates[cur]))
    if round(float(nettwd),2) != round(exp_twd,2): errors.append(f'Bad NetAmountTWD for {qid}: {nettwd} vs {exp_twd}')
    for text_col_idx, val in [(1,qid),(2,cust),(3,proj)]:
        if not isinstance(val, str): errors.append(f'ID not text in quote {qid} col{text_col_idx}: {val!r}')
# Lines refer to known quotes.
known_qids = {h['QuoteID'] for h in headers}
line_count = 0
for row in range(5, l.max_row+1):
    if l.cell(row,1).value is None:
        continue
    line_count += 1
    if l.cell(row,1).value not in known_qids:
        errors.append(f'Line row {row} unknown QuoteID {l.cell(row,1).value}')
    if not isinstance(l.cell(row,1).value, str) or not isinstance(l.cell(row,2).value, str):
        errors.append(f'Line ID not text row {row}')
if line_count != len(line_rows):
    errors.append(f'Line count mismatch: {line_count} vs {len(line_rows)}')
# Summary exactly five projects with zero rows for missing.
summary_rows = []
for row in range(5, 10):
    summary_rows.append([s.cell(row,c).value for c in range(1,7)])
if [row[0] for row in summary_rows] != PROJECTS:
    errors.append(f'Summary projects mismatch: {summary_rows}')
expected_counts = {'P001':1,'P002':1,'P003':0,'P004':1,'P005':0}
expected_twd = {'P001':331200.00,'P002':97450.85,'P003':0.00,'P004':282150.00,'P005':0.00}
for group,subgroup,cur,amount,amounttwd,count in summary_rows:
    if count != expected_counts[group]:
        errors.append(f'Summary count for {group} {count} != {expected_counts[group]}')
    if round(float(amounttwd),2) != round(expected_twd[group],2):
        errors.append(f'Summary TWD for {group} {amounttwd} != {expected_twd[group]}')
    if expected_counts[group] == 0 and (amount not in (0, 0.0) or amounttwd not in (0, 0.0) or cur not in (None, '')):
        errors.append(f'No-quote project not zero/blank currency: {group} row {group,subgroup,cur,amount,amounttwd,count}')
# Formula range.
if s['B2'].value != '=SUM(Quotes!K5:K7)':
    errors.append(f'Summary B2 formula not extended: {s["B2"].value}')
# Report cells.
if r['B3'].value != FROM_DATE or r['D3'].value != TO_DATE:
    errors.append('Report date cells wrong')
if errors:
    print('VALIDATION FAILED')
    for e in errors:
        print('ERROR:', e)
    raise SystemExit(1)
print('VALIDATION PASSED')
print('Checked: output workbook opens; sheets and headers preserved; Quotes has 3 Accepted rows for P001/P002/P004 in Sep 2026; Lines has', line_count, 'rows with matching QuoteIDs; NetAmountTWD uses Decimal ROUND_HALF_UP with fx_rates.csv; Summary has one row each for P001-P005 including zero rows for P003/P005; Summary B2 formula =SUM(Quotes!K5:K7).')
print('Summary rows:', summary_rows)
print('Quote rows:', quote_rows)