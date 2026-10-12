from pathlib import Path
import os, json, subprocess, sys, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from copy import copy
from collections import defaultdict
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

customer_id = str(params['customer_id']).strip()
from_date = str(params['from_date']).strip()
to_date = str(params['to_date']).strip()
cli_path = str(params['cli_path']).strip()

print('Run parameters:', {'customer_id': customer_id, 'from_date': from_date, 'to_date': to_date, 'cli_path': cli_path})
print('Input files:', sorted(p.name for p in input_dir.glob('*')))

if not customer_id:
    raise ValueError('customer_id is required')
if not from_date or not to_date:
    raise ValueError('from_date and to_date are required')
if from_date > to_date:
    raise ValueError('from_date must be <= to_date')
if not Path(cli_path).exists():
    raise FileNotFoundError(f'CLI path not found: {cli_path}')

template = input_dir / 'quote_template.xlsx'
fx_file = input_dir / 'fx_rates.csv'
if not template.exists():
    raise FileNotFoundError(f'Missing template: {template}')
if not fx_file.exists():
    raise FileNotFoundError(f'Missing FX rates: {fx_file}')

# Load and display template headers before doing work.
probe_wb = load_workbook(template, data_only=False)
print('Template sheets:', probe_wb.sheetnames)
for sheet_name, header_row, max_col in [('Quotes',4,11), ('Lines',4,6), ('Summary',4,6)]:
    ws = probe_wb[sheet_name]
    print(f'{sheet_name} headers:', [ws.cell(header_row, c).value for c in range(1, max_col+1)])

# Read FX rates; fx_rates.csv has an unquoted comma in Note in the test fixture, so use only first two columns.
rates = {}
with fx_file.open(newline='', encoding='utf-8-sig') as f:
    rdr = csv.reader(f)
    next(rdr, None)
    for row in rdr:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates loaded:', {k: str(v) for k, v in rates.items()})

def run_cli(args):
    p = Path(cli_path)
    if p.suffix.lower() == '.py':
        cmd = [sys.executable, str(p), *args]
        res = subprocess.run(cmd, capture_output=True, text=True)
    elif p.suffix.lower() in ('.cmd', '.bat'):
        cmdline = subprocess.list2cmdline([str(p), *args])
        res = subprocess.run(cmdline, capture_output=True, text=True, shell=True)
    else:
        res = subprocess.run([str(p), *args], capture_output=True, text=True)
    print('CLI:', ' '.join(args), 'exit', res.returncode)
    if res.stderr.strip():
        print('CLI stderr:', res.stderr.strip())
    if res.returncode != 0:
        raise RuntimeError(f'CLI failed with exit code {res.returncode}: {res.stderr.strip()}')
    try:
        return json.loads(res.stdout)
    except json.JSONDecodeError as e:
        raise RuntimeError(f'CLI returned non-JSON stdout for {args}: {res.stdout[:500]}') from e

def d2(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def to_excel_number(value):
    return float(d2(value))

listed = run_cli(['quotes', 'list', '--customer', customer_id, '--from', from_date, '--to', to_date, '--format', 'json'])
items = listed.get('items', [])
print('Listed quote count:', listed.get('count'), 'items:', len(items))
if listed.get('schema_version') != 1:
    raise RuntimeError(f"Unexpected list schema_version: {listed.get('schema_version')}")

quote_docs = []
for item in items:
    qid = item['QuoteID']
    doc = run_cli(['quotes', 'get', qid, '--format', 'json'])
    if doc.get('schema_version') != 1:
        raise RuntimeError(f"Unexpected get schema_version for {qid}: {doc.get('schema_version')}")
    quote_docs.append(doc)
print('Fetched quote documents:', len(quote_docs))

safe_stamp = datetime.now().astimezone().strftime('%Y%m%dT%H%M%S%z')
out_name = f'quote_report_{customer_id}_{from_date}_{to_date}_{safe_stamp}.xlsx'.replace(':','')
out_path = output_dir / out_name
shutil.copyfile(template, out_path)
wb = load_workbook(out_path)

report = wb['Report']
now = datetime.now().astimezone().isoformat(timespec='seconds')
report['B3'] = from_date
report['D3'] = to_date
report['B4'] = f'CustomerID={customer_id}; QuoteDate between {from_date} and {to_date} (inclusive)'
report['B5'] = now
report['B7'] = 'quote_cli'
report['D7'] = now
report['B9'] = ''

quotes_ws = wb['Quotes']
lines_ws = wb['Lines']
summary_ws = wb['Summary']

def clear_data(ws):
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

def copy_row_style(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)

clear_data(quotes_ws)
clear_data(lines_ws)
clear_data(summary_ws)

quote_rows = []
line_rows = []
for doc in quote_docs:
    h = doc['header']
    cur = h.get('Currency')
    if cur not in rates:
        raise RuntimeError(f'Missing FX rate for currency {cur}')
    net = d2(h.get('NetAmount'))
    net_twd = d2(net * rates[cur])
    quote_rows.append([
        h.get('QuoteID'), h.get('CustomerID'), h.get('ProjectID'), h.get('QuoteDate'), h.get('ValidUntil'),
        h.get('Status'), cur, to_excel_number(h.get('Subtotal')), float(Decimal(str(h.get('DiscountPercent')))),
        to_excel_number(net), to_excel_number(net_twd)
    ])
    for li in doc.get('line_items', []):
        line_rows.append([
            h.get('QuoteID'), li.get('ProductID'), li.get('ProductName'), int(li.get('Quantity')),
            to_excel_number(li.get('UnitPrice')), to_excel_number(li.get('LineAmount'))
        ])

for r_idx, row in enumerate(quote_rows, start=5):
    copy_row_style(quotes_ws, 4, r_idx, 11)
    for c_idx, val in enumerate(row, start=1):
        cell = quotes_ws.cell(r_idx, c_idx, value=val)
        if c_idx in (1,2,3,6,7):
            cell.number_format = '@'
        elif c_idx in (4,5):
            cell.number_format = 'yyyy-mm-dd'
        elif c_idx in (8,10,11):
            cell.number_format = '0.00'

for r_idx, row in enumerate(line_rows, start=5):
    copy_row_style(lines_ws, 4, r_idx, 6)
    for c_idx, val in enumerate(row, start=1):
        cell = lines_ws.cell(r_idx, c_idx, value=val)
        if c_idx in (1,2,3):
            cell.number_format = '@'
        elif c_idx in (5,6):
            cell.number_format = '0.00'

summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
for row in quote_rows:
    key = (row[5], row[6])  # status, currency
    summary[key]['amount'] += Decimal(str(row[9]))
    summary[key]['twd'] += Decimal(str(row[10]))
    summary[key]['count'] += 1

summary_rows = []
for (status, cur), vals in sorted(summary.items()):
    summary_rows.append(['Status', status, cur, to_excel_number(vals['amount']), to_excel_number(vals['twd']), vals['count']])

for r_idx, row in enumerate(summary_rows, start=5):
    copy_row_style(summary_ws, 4, r_idx, 6)
    for c_idx, val in enumerate(row, start=1):
        cell = summary_ws.cell(r_idx, c_idx, value=val)
        if c_idx in (4,5):
            cell.number_format = '0.00'

last_quote_row = max(5, 4 + len(quote_rows))
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
wb.save(out_path)
print('Wrote workbook:', out_path)

# Reopen and validate stated requirements.
check_wb = load_workbook(out_path, data_only=False)
q = check_wb['Quotes']
l = check_wb['Lines']
s = check_wb['Summary']
r = check_wb['Report']
expected_q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
expected_l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
expected_s_headers = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
actual_q_headers = [q.cell(4,c).value for c in range(1,12)]
actual_l_headers = [l.cell(4,c).value for c in range(1,7)]
actual_s_headers = [s.cell(4,c).value for c in range(1,7)]
if actual_q_headers != expected_q_headers:
    raise AssertionError(f'Quotes headers changed: {actual_q_headers}')
if actual_l_headers != expected_l_headers:
    raise AssertionError(f'Lines headers changed: {actual_l_headers}')
if actual_s_headers != expected_s_headers:
    raise AssertionError(f'Summary headers changed: {actual_s_headers}')
if r['B3'].value != from_date or r['D3'].value != to_date:
    raise AssertionError('Report date cells not set correctly')
if q.max_row != max(4, 4 + len(quote_rows)):
    raise AssertionError(f'Quotes row count mismatch: sheet max_row={q.max_row}, expected={4+len(quote_rows)}')
if l.max_row != max(4, 4 + len(line_rows)):
    raise AssertionError(f'Lines row count mismatch: sheet max_row={l.max_row}, expected={4+len(line_rows)}')
if s['B2'].value != f'=SUM(Quotes!K5:K{last_quote_row})':
    raise AssertionError(f'Summary B2 formula mismatch: {s["B2"].value}')

print('CHECK output exists:', out_path.exists(), out_path.stat().st_size)
print('CHECK quote rows:', len(quote_rows))
print('CHECK line rows:', len(line_rows))
print('CHECK summary rows:', len(summary_rows))
print('CHECK Summary B2 formula:', s['B2'].value)
print('CHECK first quote:', [q.cell(5,c).value for c in range(1,12)] if quote_rows else 'none')
print('CHECK first line:', [l.cell(5,c).value for c in range(1,7)] if line_rows else 'none')
print('CHECK summary data:', [[s.cell(rr,c).value for c in range(1,7)] for rr in range(5, 5+len(summary_rows))])
print(json.dumps({'output_file': str(out_path), 'quote_rows': len(quote_rows), 'line_rows': len(line_rows), 'summary_rows': len(summary_rows)}, ensure_ascii=False))
