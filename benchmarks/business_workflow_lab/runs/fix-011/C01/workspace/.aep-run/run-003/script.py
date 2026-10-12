from pathlib import Path
import subprocess, json, sys, os, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from copy import copy
from openpyxl import load_workbook
from openpyxl.styles import Font

workspace = Path.cwd()
output_dir = workspace/'output'
output_dir.mkdir(exist_ok=True)
cli_path = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
customer_id='C001'
from_date='2026-09-01'
to_date='2026-09-30'

def run_cli(args):
    p=Path(cli_path)
    if p.suffix.lower()=='.py':
        cmd=[sys.executable, str(p), *args]
        res=subprocess.run(cmd, capture_output=True, text=True)
    elif p.suffix.lower() in ('.cmd','.bat'):
        cmdline=subprocess.list2cmdline([str(p), *args])
        res=subprocess.run(cmdline, capture_output=True, text=True, shell=True)
    else:
        res=subprocess.run([str(p), *args], capture_output=True, text=True)
    print('CLI:', ' '.join(args), 'exit', res.returncode)
    if res.stderr.strip(): print('stderr:', res.stderr.strip())
    if res.returncode!=0:
        raise RuntimeError(f'CLI failed exit {res.returncode}: {res.stderr}')
    return json.loads(res.stdout)

print('Listing quotes...')
listed=run_cli(['quotes','list','--customer',customer_id,'--from',from_date,'--to',to_date,'--format','json'])
print('schema', listed.get('schema_version'), 'count', listed.get('count'), 'items', len(listed.get('items', [])))
for item in listed.get('items', [])[:5]: print('sample item', item)

quote_docs=[]
for item in listed.get('items', []):
    qid=item['QuoteID']
    doc=run_cli(['quotes','get',qid,'--format','json'])
    quote_docs.append(doc)
print('got docs', len(quote_docs))

# Load FX rates
rates={}
with (workspace/'input'/'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    header=next(rdr)
    for row in rdr:
        if len(row)>=2 and row[0]:
            rates[row[0]]=Decimal(row[1])
print('rates', rates)

def d2(x): return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
def to_num(x): return float(d2(x))

template=workspace/'input'/'quote_template.xlsx'
out=output_dir/f'quote_report_{customer_id}_{from_date}_{to_date}.xlsx'
shutil.copyfile(template, out)
wb=load_workbook(out)
# fill report
report=wb['Report']
now=datetime.now().astimezone().isoformat(timespec='seconds')
report['B3']=from_date
report['D3']=to_date
report['B4']=f'CustomerID={customer_id}; QuoteDate between {from_date} and {to_date} (inclusive)'
report['B5']=now
report['B7']='quote_cli'
report['D7']=now
report['B9']=''

# helper clear rows from row 5 down, preserving row 4 headers and template layout
def clear_data(ws):
    if ws.max_row>=5:
        ws.delete_rows(5, ws.max_row-4)

def copy_row_style(ws, source_row, target_row, max_col):
    for col in range(1, max_col+1):
        src=ws.cell(source_row, col)
        dst=ws.cell(target_row, col)
        if src.has_style:
            dst._style=copy(src._style)
        if src.number_format:
            dst.number_format=src.number_format
        if src.alignment:
            dst.alignment=copy(src.alignment)
        if src.font:
            dst.font=copy(src.font)
        if src.fill:
            dst.fill=copy(src.fill)
        if src.border:
            dst.border=copy(src.border)

quotes_ws=wb['Quotes']; lines_ws=wb['Lines']; summary_ws=wb['Summary']
clear_data(quotes_ws); clear_data(lines_ws); clear_data(summary_ws)
# ensure at least row 5 style exists by copying row4 style where needed later
quote_rows=[]; line_rows=[]
for doc in quote_docs:
    h=doc['header']
    cur=h.get('Currency')
    if cur not in rates:
        raise RuntimeError(f'Missing FX rate for {cur}')
    net=d2(h.get('NetAmount'))
    net_twd=d2(net*rates[cur])
    quote_rows.append([
        h.get('QuoteID'), h.get('CustomerID'), h.get('ProjectID'), h.get('QuoteDate'), h.get('ValidUntil'),
        h.get('Status'), cur, to_num(h.get('Subtotal')), float(Decimal(str(h.get('DiscountPercent')))), to_num(net), to_num(net_twd)
    ])
    for li in doc.get('line_items', []):
        line_rows.append([
            h.get('QuoteID'), li.get('ProductID'), li.get('ProductName'), int(li.get('Quantity')),
            to_num(li.get('UnitPrice')), to_num(li.get('LineAmount'))
        ])

for r_idx, row in enumerate(quote_rows, start=5):
    copy_row_style(quotes_ws, 4, r_idx, 11)
    for c_idx, val in enumerate(row, start=1):
        cell=quotes_ws.cell(r_idx,c_idx,value=val)
        if c_idx in (1,2,3,6,7): cell.number_format='@'
        if c_idx in (4,5): cell.number_format='yyyy-mm-dd'
        if c_idx in (8,10,11): cell.number_format='0.00'
for r_idx, row in enumerate(line_rows, start=5):
    copy_row_style(lines_ws, 4, r_idx, 6)
    for c_idx, val in enumerate(row, start=1):
        cell=lines_ws.cell(r_idx,c_idx,value=val)
        if c_idx in (1,2,3): cell.number_format='@'
        if c_idx in (5,6): cell.number_format='0.00'

# Summary by status and currency
from collections import defaultdict
summary=defaultdict(lambda: {'amount':Decimal('0.00'),'twd':Decimal('0.00'),'count':0})
for row in quote_rows:
    status=row[5]; cur=row[6]; net=Decimal(str(row[9])); twd=Decimal(str(row[10]))
    key=(status,cur)
    summary[key]['amount'] += net
    summary[key]['twd'] += twd
    summary[key]['count'] += 1
summary_rows=[]
for (status,cur), vals in sorted(summary.items()):
    summary_rows.append(['Status', status, cur, to_num(vals['amount']), to_num(vals['twd']), vals['count']])
for r_idx, row in enumerate(summary_rows, start=5):
    copy_row_style(summary_ws, 4, r_idx, 6)
    for c_idx, val in enumerate(row, start=1):
        cell=summary_ws.cell(r_idx,c_idx,value=val)
        if c_idx in (4,5): cell.number_format='0.00'
last_quote_row=max(5, 4+len(quote_rows))
summary_ws['B2']=f'=SUM(Quotes!K5:K{last_quote_row})'
wb.save(out)
print('wrote', out)

# Validate by reopening
wb2=load_workbook(out, data_only=False)
q=wb2['Quotes']; l=wb2['Lines']; s=wb2['Summary']; r=wb2['Report']
print('CHECK output exists', out.exists(), out.stat().st_size)
print('CHECK report dates', r['B3'].value, r['D3'].value, 'condition', r['B4'].value)
print('CHECK quotes header', [q.cell(4,c).value for c in range(1,12)])
print('CHECK lines header', [l.cell(4,c).value for c in range(1,7)])
print('CHECK summary header', [s.cell(4,c).value for c in range(1,7)])
print('CHECK quote rows', len(quote_rows), 'sheet max row', q.max_row)
print('CHECK line rows', len(line_rows), 'sheet max row', l.max_row)
print('CHECK summary rows', len(summary_rows), 'formula', s['B2'].value)
print('CHECK first quote', [q.cell(5,c).value for c in range(1,12)] if quote_rows else 'none')
print('CHECK first line', [l.cell(5,c).value for c in range(1,7)] if line_rows else 'none')
print('CHECK summary data', [[s.cell(rr,c).value for c in range(1,7)] for rr in range(5,5+len(summary_rows))])
