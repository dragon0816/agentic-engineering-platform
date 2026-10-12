from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for r in wb['Instructions'].iter_rows(values_only=True):
    print(r)
