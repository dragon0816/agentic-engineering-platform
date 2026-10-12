from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# workbook structure
wb_path=ws/'input'/'quote_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('workbook sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('\nSHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
        print(r)
# fx sample
fx_path=ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(8), rdr):
        print(row)
# manual sample
manual=ws/'input'/'quote_cli_manual.md'
print('\nmanual first 120 lines:')
text=manual.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:120],1):
    print(f'{i:03d}: {line}')
