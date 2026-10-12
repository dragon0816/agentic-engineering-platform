from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook(Path('input')/'expense_template.xlsx', data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)
print('\nMerged ranges:')
for ws in wb.worksheets:
    print(ws.title, list(ws.merged_cells.ranges))
print('\nReport row 20:', [wb['Report'].cell(20,c).value for c in range(1,7)])
