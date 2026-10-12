from pathlib import Path
import csv
from openpyxl import load_workbook

base = Path('.')
print('Workspace:', base.resolve())
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print(' -', p, 'size', p.stat().st_size)
print('Output files:')
for p in sorted((base/'output').glob('*')):
    print(' -', p, 'size', p.stat().st_size)

for csv_path in [base/'input'/'fx_rates.csv', base/'output'/'expenses_2026-09_approved.csv']:
    print('\nCSV:', csv_path)
    with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
        reader = csv.DictReader(f)
        print('Headers:', reader.fieldnames)
        for i, row in enumerate(reader):
            if i >= 5: break
            print('Row', i+1, row)

manual = base/'input'/'expense_manual.md'
print('\nManual first 1200 chars:')
print(manual.read_text(encoding='utf-8')[:1200])

template = base/'input'/'expense_template.xlsx'
wb = load_workbook(template)
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('\nSheet:', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
        print(r)
