from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
manual = ws/'input'/'expense_manual.md'
if manual.exists():
    text = manual.read_text(encoding='utf-8')
    print('\nmanual first 2000 chars:\n', text[:2000])
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv sample:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(8), r):
            print(row)
tmpl = ws/'input'/'expense_template.xlsx'
if tmpl.exists():
    wb = load_workbook(tmpl, data_only=False)
    print('\ntemplate sheets:', wb.sheetnames)
    for sh in wb.sheetnames:
        s=wb[sh]
        print('sheet', sh, 'size', s.max_row, s.max_column)
        for row in s.iter_rows(min_row=1, max_row=min(s.max_row,12), values_only=True):
            print(row)
        print('---')