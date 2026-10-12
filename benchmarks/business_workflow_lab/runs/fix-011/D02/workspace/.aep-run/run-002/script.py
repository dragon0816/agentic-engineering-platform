from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
for name in ['submitted_2026_09.csv','approved_2026_09.csv']:
    p=ws/'output'/name
    print('\nFILE', name, 'exists', p.exists(), 'size', p.stat().st_size if p.exists() else None)
    with open(p, newline='', encoding='utf-8-sig') as f:
        reader=csv.DictReader(f)
        print('headers:', reader.fieldnames)
        rows=[]
        for i,row in zip(range(5), reader): rows.append(row)
        for r in rows: print(r)
print('\nInstructions full:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    if any(v is not None for v in row): print(row)