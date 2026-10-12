from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
import csv, datetime, os

ws=Path.cwd(); outdir=ws/'output'; indir=ws/'input'
print('Reopening all output files for validation:')
for p in sorted(outdir.glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# read fx
fx={}
with open(indir/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f); next(reader)
    for row in reader:
        if row and row[0].strip(): fx[row[0].strip()] = Decimal(row[1].strip())

def read_export(name, expected_status, expected_count):
    p=outdir/name
    assert p.exists(), f'missing {name}'
    with open(p, newline='', encoding='utf-8-sig') as f:
        reader=csv.DictReader(f); rows=list(reader)
    expected_headers=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']
    assert reader.fieldnames == expected_headers, (name, reader.fieldnames)
    assert len(rows) == expected_count, (name, len(rows), expected_count)
    assert all(r['Status']==expected_status for r in rows)
    assert all('2026-09-01' <= r['ExpenseDate'] <= '2026-09-30' for r in rows)
    assert all(r['ExpenseID'].startswith('EXP-') for r in rows)
    print(f'{name}: {len(rows)} rows, headers/status/date range OK')
    return rows
submitted = read_export('submitted_2026_09.csv', 'Submitted', 7)
approved = read_export('approved_2026_09.csv', 'Approved', 15)
source_rows = sorted(submitted+approved, key=lambda r: (r['ExpenseDate'], r['ExpenseID']))
source_by_id = {r['ExpenseID']: r for r in source_rows}
missing_ids = [r['ExpenseID'] for r in source_rows if r['MissingReceipt'].lower()=='true']

xlsx=outdir/'expense_report_2026_09_submitted_approved.xlsx'
assert xlsx.exists(), 'missing workbook'
wb=load_workbook(xlsx, data_only=False)
assert wb.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions'], wb.sheetnames
print('Workbook sheets preserved:', wb.sheetnames)

# Report
rep=wb['Report']
assert rep['B3'].value == '2026-09-01'
assert rep['D3'].value == '2026-09-30'
assert 'Submitted' in rep['B4'].value and 'Approved' in rep['B4'].value
assert 'Missing receipts retained' in rep['B9'].value
print('Report metadata OK:', rep['B3'].value, rep['D3'].value, rep['B4'].value)

# Details
sh=wb['Details']
headers=[sh.cell(4,c).value for c in range(1,15)]
expected=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert headers==expected, headers
rows=[]
for r in range(5, sh.max_row+1):
    vals=[sh.cell(r,c).value for c in range(1,15)]
    if any(v is not None for v in vals): rows.append(vals)
assert len(rows)==22, len(rows)
ids=[r[0] for r in rows]
assert len(ids)==len(set(ids))==22
assert set(ids)==set(source_by_id)

def as_date(v):
    if isinstance(v, datetime.datetime): return v.date()
    if isinstance(v, datetime.date): return v
    return datetime.date.fromisoformat(v)
for vals in rows:
    eid, empid, empname, proj, dt, cat, curr, amt, twd, status, receipt, desc, missing, needs = vals
    src=source_by_id[eid]
    assert isinstance(eid, str) and isinstance(empid, str) and isinstance(proj, str)
    assert datetime.date(2026,9,1) <= as_date(dt) <= datetime.date(2026,9,30)
    assert status in ['Submitted','Approved']
    assert curr in fx
    assert isinstance(amt, (int,float)) and isinstance(twd, (int,float))
    expected_twd=(Decimal(src['Amount'])*fx[curr]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert Decimal(str(twd)).quantize(Decimal('0.01')) == expected_twd, (eid, twd, expected_twd)
    assert receipt == (src['ReceiptID'] or None)
    assert desc == (src['Description'] or None)
    assert missing in ['true','false'] and needs in ['true','false']
    if src['MissingReceipt'].lower()=='true':
        assert missing=='true' and needs=='true' and receipt is None
print('Details OK: 22 rows, ReceiptID/Description/MissingReceipt from export, missing rows retained:', missing_ids)

# Exceptions
ex=wb['Exceptions']
ex_headers=[ex.cell(4,c).value for c in range(1,5)]
assert ex_headers==['ExpenseID','EmployeeID','ExpenseDate','Reason']
ex_rows=[]
for r in range(5, ex.max_row+1):
    vals=[ex.cell(r,c).value for c in range(1,5)]
    if any(v is not None for v in vals): ex_rows.append(vals)
assert [r[0] for r in ex_rows] == missing_ids
for eid, empid, dt, reason in ex_rows:
    assert eid in ids
    assert empid == source_by_id[eid]['EmployeeID']
    assert reason == 'Missing receipt'
print('Exceptions OK: original ExpenseID listed for missing receipts:', [r[0] for r in ex_rows])

# Summary
sumsh=wb['Summary']
assert sumsh['B2'].value == '=SUM(Details!I5:I26)'
summary_rows=[]
for r in range(5, sumsh.max_row+1):
    vals=[sumsh.cell(r,c).value for c in range(1,7)]
    if any(v is not None for v in vals): summary_rows.append(vals)
# recompute by Status, Currency
expected_summary={}
for src in source_rows:
    key=(src['Status'], src['Currency'])
    expected_summary.setdefault(key, [Decimal('0'), Decimal('0'), 0])
    amt=Decimal(src['Amount']); twd=(amt*fx[src['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    expected_summary[key][0]+=amt; expected_summary[key][1]+=twd; expected_summary[key][2]+=1
assert len(summary_rows)==len(expected_summary)
for group, subgroup, curr, amount, twd, count in summary_rows:
    key=(group,curr); assert key in expected_summary
    exp_amt, exp_twd, exp_count = expected_summary[key]
    assert Decimal(str(amount)).quantize(Decimal('0.01')) == exp_amt.quantize(Decimal('0.01'))
    assert Decimal(str(twd)).quantize(Decimal('0.01')) == exp_twd.quantize(Decimal('0.01'))
    assert count == exp_count
print('Summary OK: formula and grouped totals by Status/Currency verified')
print('VALIDATION PASSED for all outputs.')