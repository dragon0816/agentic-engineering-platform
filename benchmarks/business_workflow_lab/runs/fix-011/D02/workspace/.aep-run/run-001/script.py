import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size)
# workbook sheets dimensions headers
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path)
print('workbook sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print('SHEET', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
        print(' ', r)
# fx sample
fx=ws/'input'/'fx_rates.csv'
print('fx_rates sample:')
with open(fx, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader): print(row)
# manual
man=ws/'input'/'expense_manual.md'
print('manual head:')
print(man.read_text(encoding='utf-8')[:2000])