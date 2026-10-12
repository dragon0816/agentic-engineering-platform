from pathlib import Path
from openpyxl import load_workbook
p=Path('input/sales_template.xlsx')
wb=load_workbook(p, data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)