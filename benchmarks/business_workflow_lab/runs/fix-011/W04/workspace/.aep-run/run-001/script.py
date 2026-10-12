from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)

# Workbook structure
xlsx = ws/'input'/'sales_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWORKBOOK SHEETS', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'-- Sheet {s}: max_row={sh.max_row} max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(r)
    print('merged', list(sh.merged_cells.ranges)[:10])

# fx rates sample
csvp = ws/'input'/'fx_rates.csv'
print('\nFX RATES')
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(10), reader): print(row)

# manual content
mdp = ws/'input'/'sap_manual.md'
print('\nSAP MANUAL')
text = mdp.read_text(encoding='utf-8')
print(text[:5000])