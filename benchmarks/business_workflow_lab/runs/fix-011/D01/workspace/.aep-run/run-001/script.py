from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
for p in sorted((ws/'input').glob('*')):
    print('INPUT', p.name, p.stat().st_size)
# workbook structure
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('template sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
        print(' ', row)
# fx csv preview
fx=ws/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open('r', encoding='utf-8-sig', newline='') as f:
        reader=csv.reader(f)
        for i,row in zip(range(8), reader):
            print('FXROW', row)
# manual preview
man=ws/'input'/'expense_manual.md'
if man.exists():
    print('MANUAL PREVIEW')
    print(man.read_text(encoding='utf-8')[:2000])