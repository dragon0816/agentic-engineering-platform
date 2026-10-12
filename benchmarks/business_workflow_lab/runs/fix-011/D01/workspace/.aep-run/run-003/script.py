from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv, re, os

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

employee_id = 'E001'
date_from = '2026-09-01'
date_to = '2026-09-30'
export_csv = output_dir / f'expense_export_{employee_id}_{date_from}_{date_to}.csv'
out_xlsx = output_dir / 'E001_september_expenses.xlsx'

def money(v):
    return Decimal(str(v)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def parse_bool(v):
    return str(v).strip().lower() == 'true'

# Read inputs
print('INPUT EXPORT', export_csv, export_csv.exists(), export_csv.stat().st_size if export_csv.exists() else 'missing')
with export_csv.open('r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.DictReader(f))
print('export headers', list(rows[0].keys()) if rows else [])
print('export row count', len(rows))
print('export sample', rows[:2])

rates = {}
with (input_dir / 'fx_rates.csv').open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    header = next(reader)
    for r in reader:
        if len(r) >= 2 and r[0].strip():
            rates[r[0].strip()] = Decimal(r[1].strip())
print('rates', rates)

# Validate exported result matches requested filter
for r in rows:
    assert r['EmployeeID'] == employee_id, f"Unexpected EmployeeID: {r}"
    assert date_from <= r['ExpenseDate'] <= date_to, f"Date outside range: {r}"
    assert r['Currency'] in rates, f"Missing FX rate: {r['Currency']}"

wb = load_workbook(input_dir / 'expense_template.xlsx')

# Helpers to clear data areas while preserving row 4 header/template workbook
def clear_from(sheet, start_row, max_col):
    sh = wb[sheet]
    if sh.max_row >= start_row:
        sh.delete_rows(start_row, sh.max_row - start_row + 1)

def copy_row_style(sh, source_row, target_row, max_col):
    for c in range(1, max_col + 1):
        src = sh.cell(source_row, c)
        dst = sh.cell(target_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

# Report sheet
now = datetime.now().astimezone().replace(microsecond=0).isoformat()
report = wb['Report']
report['B3'] = date_from
report['D3'] = date_to
report['B4'] = f'EmployeeID={employee_id}; Status=ALL; ExpenseDate {date_from}..{date_to}'
report['B5'] = now
report['B7'] = 'Expense Viewer export CSV; fx_rates.csv'
report['D7'] = now
missing_count = sum(1 for r in rows if parse_bool(r.get('MissingReceipt')))
report['B9'] = f'Exported all {len(rows)} rows via Expense Viewer CSV. NeedsReview/Exceptions marks MissingReceipt=true ({missing_count} row(s)).'

# Details
clear_from('Details', 5, 14)
details = wb['Details']
headers = [details.cell(4, c).value for c in range(1, 15)]
for idx, r in enumerate(rows, start=5):
    copy_row_style(details, 4, idx, 14)
    amount = money(r['Amount'])
    amount_twd = money(amount * rates[r['Currency']])
    missing = parse_bool(r.get('MissingReceipt'))
    needs_review = missing
    values = [
        r.get('ExpenseID') or None,
        r.get('EmployeeID') or None,
        r.get('EmployeeName') or None,
        r.get('ProjectID') or None,
        r.get('ExpenseDate') or None,
        r.get('Category') or None,
        r.get('Currency') or None,
        float(amount),
        float(amount_twd),
        r.get('Status') or None,
        r.get('ReceiptID') or None,
        r.get('Description') or None,
        'true' if missing else 'false',
        'true' if needs_review else 'false',
    ]
    for c, val in enumerate(values, start=1):
        cell = details.cell(idx, c, val)
        if c in (1,2,4,11):
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in (8,9):
            cell.number_format = '0.00'
        if c == 12:
            cell.alignment = copy(details.cell(4, c).alignment)
            cell.alignment = Alignment(wrap_text=True, vertical='top')
# Keep autofilter if desired
if rows:
    details.auto_filter.ref = f'A4:N{4+len(rows)}'
last_detail_row = 4 + max(len(rows), 1)

# Summary by Status and Currency
clear_from('Summary', 5, 6)
summary = wb['Summary']
if rows:
    summary['B2'] = f'=SUM(Details!I5:I{4+len(rows)})'
else:
    summary['B2'] = '=SUM(Details!I5:I5)'
summary_groups = {}
for r in rows:
    key = (r['Status'], r['Currency'])
    amount = money(r['Amount'])
    amount_twd = money(amount * rates[r['Currency']])
    g = summary_groups.setdefault(key, {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
    g['amount'] += amount
    g['amount_twd'] += amount_twd
    g['count'] += 1
for idx, ((status, currency), g) in enumerate(sorted(summary_groups.items()), start=5):
    copy_row_style(summary, 4, idx, 6)
    vals = ['Status', status, currency, float(money(g['amount'])), float(money(g['amount_twd'])), g['count']]
    for c, val in enumerate(vals, start=1):
        cell = summary.cell(idx, c, val)
        if c in (4,5):
            cell.number_format = '0.00'
if summary_groups:
    summary.auto_filter.ref = f'A4:F{4+len(summary_groups)}'

# Exceptions: missing receipt
clear_from('Exceptions', 5, 4)
exceptions = wb['Exceptions']
exc_rows = [r for r in rows if parse_bool(r.get('MissingReceipt'))]
for idx, r in enumerate(exc_rows, start=5):
    copy_row_style(exceptions, 4, idx, 4)
    vals = [r['ExpenseID'], r['EmployeeID'], r['ExpenseDate'], 'MissingReceipt=true / receipt is missing']
    for c, val in enumerate(vals, start=1):
        cell = exceptions.cell(idx, c, val)
        if c in (1,2): cell.number_format = '@'
        if c == 3: cell.number_format = 'yyyy-mm-dd'
if exc_rows:
    exceptions.auto_filter.ref = f'A4:D{4+len(exc_rows)}'

wb.save(out_xlsx)
print('wrote', out_xlsx, out_xlsx.stat().st_size)

# Reopen and verify
wb2 = load_workbook(out_xlsx, data_only=False)
sh = wb2['Details']
actual_detail_rows = sh.max_row - 4 if sh.max_row >= 5 else 0
# delete_rows can leave max row = 4; for data count use non-empty ExpenseID rows
actual_detail_rows = sum(1 for row in sh.iter_rows(min_row=5, max_col=1, values_only=True) if row[0])
assert actual_detail_rows == len(rows), (actual_detail_rows, len(rows))
assert [sh.cell(4,c).value for c in range(1,15)] == headers
assert wb2['Report']['B3'].value == date_from
assert wb2['Report']['D3'].value == date_to
assert wb2['Summary']['B2'].value == f'=SUM(Details!I5:I{4+len(rows)})'
# verify amounts and text IDs
calc_total = Decimal('0')
for i, r in enumerate(rows, start=5):
    assert sh.cell(i,1).value == r['ExpenseID']
    assert sh.cell(i,2).value == r['EmployeeID']
    assert sh.cell(i,8).data_type == 'n'
    assert sh.cell(i,9).data_type == 'n'
    expected_twd = money(Decimal(r['Amount']) * rates[r['Currency']])
    got = Decimal(str(sh.cell(i,9).value)).quantize(Decimal('0.01'))
    assert got == expected_twd, (i, got, expected_twd)
    calc_total += expected_twd
# Summary rows check
sum_sh = wb2['Summary']
summary_rows = []
for row in sum_sh.iter_rows(min_row=5, max_col=6, values_only=True):
    if row[0] is None: continue
    summary_rows.append(row)
assert len(summary_rows) == len(summary_groups), (summary_rows, summary_groups)
summary_twd = sum(Decimal(str(r[4])).quantize(Decimal('0.01')) for r in summary_rows)
assert summary_twd == calc_total, (summary_twd, calc_total)
# Exceptions count
exc_sh = wb2['Exceptions']
actual_exc = sum(1 for row in exc_sh.iter_rows(min_row=5, max_col=1, values_only=True) if row[0])
assert actual_exc == len(exc_rows), (actual_exc, len(exc_rows))
print('verified output workbook')
print('detail rows', actual_detail_rows)
print('summary groups', len(summary_rows), summary_rows)
print('exceptions', actual_exc)
print('TWD total', calc_total)
print('files:', out_xlsx.relative_to(workspace), export_csv.relative_to(workspace))