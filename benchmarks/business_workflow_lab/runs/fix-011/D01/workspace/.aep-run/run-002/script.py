from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('OUTPUTS')
for p in sorted((ws/'output').glob('*')):
    print(p.name, p.stat().st_size)
print('EXPORT PREVIEW')
exp=ws/'output'/'expense_export_E001_2026-09-01_2026-09-30.csv'
with exp.open('r', encoding='utf-8-sig', newline='') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader):
        print(row)
print('FULL INSTRUCTIONS SHEET')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
print('MANUAL FULL')
print((ws/'input'/'expense_manual.md').read_text(encoding='utf-8'))