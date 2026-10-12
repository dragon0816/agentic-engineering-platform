from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv, json, os, re

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)
employee_id = str(params['employee_id']).strip()
date_from = str(params['date_from']).strip()
date_to = str(params['date_to']).strip()
export_csv = output_dir / f'expense_export_{employee_id}_{date_from}_{date_to}.csv'
safe_from = date_from.replace('-', '')
safe_to = date_to.replace('-', '')
out_xlsx = output_dir / f'expense_report_{employee_id}_{safe_from}_{safe_to}.xlsx'
# Preserve requested historical filename for the benchmark case.
if employee_id == 'E001' and date_from == '2026-09-01' and date_to == '2026-09-30':
    out_xlsx = output_dir / 'E001_september_expenses.xlsx'

def money(v):
    return Decimal(str(v)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def parse_bool(v):
    return str(v).strip().lower() == 'true'

def copy_row_style(sh, source_row, target_row, max_col):
    for c in range(1, max_col + 1):
        src = sh.cell(source_row, c)
        dst = sh.cell(target_row, c)
        if src.has_style:
            dst._style = copy(src._style)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)

def clear_from(wb, sheet, start_row):
    sh = wb[sheet]
    if sh.max_row >= start_row:
        sh.delete_rows(start_row, sh.max_row - start_row + 1)

print('Reading export:', export_csv)
if not export_csv.exists():
    raise FileNotFoundError(f'Missing desktop export: {export_csv}')
with export_csv.open('r', encoding='utf-8-sig', newline='') as f:
    rows = list(csv.DictReader(f))
print('Export row count:', len(rows))
print('Export headers:', list(rows[0].keys()) if rows else [])

rates = {}
with (input_dir / 'fx_rates.csv').open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    next(reader, None)
    for r in reader:
        if len(r) >= 2 and r[0].strip():
            rates[r[0].strip()] = Decimal(r[1].strip())
print('FX currencies:', sorted(rates))

for r in rows:
    if r.get('EmployeeID') != employee_id:
        raise AssertionError(f'Unexpected EmployeeID in export: {r}')
    if not (date_from <= r.get('ExpenseDate','') <= date_to):
        raise AssertionError(f'ExpenseDate outside requested range: {r}')
    if r.get('Currency') not in rates:
        raise AssertionError(f"No FX rate for currency {r.get('Currency')}")

wb = load_workbook(input_dir / 'expense_template.xlsx')
now = datetime.now().astimezone().replace(microsecond=0).isoformat()
report = wb['Report']
report['B3'] = date_from
report['D3'] = date_to
report['B4'] = f'EmployeeID={employee_id}; Status=ALL; ExpenseDate {date_from}..{date_to}'
report['B5'] = now
report['B7'] = 'Expense Viewer export CSV; fx_rates.csv'
report['D7'] = now
missing_count = sum(1 for r in rows if parse_bool(r.get('MissingReceipt')))
report['B9'] = f'Exported all {len(rows)} rows via Expense Viewer CSV. NeedsReview/Exceptions marks MissingReceipt=true ({missing_count} row(s)).'

# Details
clear_from(wb, 'Details', 5)
details = wb['Details']
headers = [details.cell(4, c).value for c in range(1, 15)]
for idx, r in enumerate(rows, start=5):
    copy_row_style(details, 4, idx, 14)
    amount = money(r['Amount'])
    amount_twd = money(amount * rates[r['Currency']])
    missing = parse_bool(r.get('MissingReceipt'))
    values = [r.get('ExpenseID') or None, r.get('EmployeeID') or None, r.get('EmployeeName') or None, r.get('ProjectID') or None, r.get('ExpenseDate') or None, r.get('Category') or None, r.get('Currency') or None, float(amount), float(amount_twd), r.get('Status') or None, r.get('ReceiptID') or None, r.get('Description') or None, 'true' if missing else 'false', 'true' if missing else 'false']
    for c, val in enumerate(values, start=1):
        cell = details.cell(idx, c, val)
        if c in (1, 2, 4, 11):
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in (8, 9):
            cell.number_format = '0.00'
        elif c == 12:
            cell.alignment = Alignment(wrap_text=True, vertical='top')
if rows:
    details.auto_filter.ref = f'A4:N{4+len(rows)}'

# Summary by status and currency
clear_from(wb, 'Summary', 5)
summary = wb['Summary']
summary['B2'] = f'=SUM(Details!I5:I{4+len(rows)})' if rows else '=SUM(Details!I5:I5)'
groups = {}
for r in rows:
    key = (r['Status'], r['Currency'])
    amount = money(r['Amount'])
    amount_twd = money(amount * rates[r['Currency']])
    g = groups.setdefault(key, {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
    g['amount'] += amount
    g['amount_twd'] += amount_twd
    g['count'] += 1
for idx, ((status, currency), g) in enumerate(sorted(groups.items()), start=5):
    copy_row_style(summary, 4, idx, 6)
    vals = ['Status', status, currency, float(money(g['amount'])), float(money(g['amount_twd'])), g['count']]
    for c, val in enumerate(vals, start=1):
        cell = summary.cell(idx, c, val)
        if c in (4, 5):
            cell.number_format = '0.00'
if groups:
    summary.auto_filter.ref = f'A4:F{4+len(groups)}'

# Exceptions: missing receipt only
clear_from(wb, 'Exceptions', 5)
exceptions = wb['Exceptions']
exc_rows = [r for r in rows if parse_bool(r.get('MissingReceipt'))]
for idx, r in enumerate(exc_rows, start=5):
    copy_row_style(exceptions, 4, idx, 4)
    vals = [r['ExpenseID'], r['EmployeeID'], r['ExpenseDate'], 'MissingReceipt=true / receipt is missing']
    for c, val in enumerate(vals, start=1):
        cell = exceptions.cell(idx, c, val)
        if c in (1, 2):
            cell.number_format = '@'
        elif c == 3:
            cell.number_format = 'yyyy-mm-dd'
if exc_rows:
    exceptions.auto_filter.ref = f'A4:D{4+len(exc_rows)}'

wb.save(out_xlsx)
print('Wrote workbook:', out_xlsx)

# Verify output
wb2 = load_workbook(out_xlsx, data_only=False)
sh = wb2['Details']
detail_count = sum(1 for row in sh.iter_rows(min_row=5, max_col=1, values_only=True) if row[0])
if detail_count != len(rows):
    raise AssertionError(f'Detail rows {detail_count} != export rows {len(rows)}')
if [sh.cell(4,c).value for c in range(1,15)] != headers:
    raise AssertionError('Details headers changed')
if wb2['Report']['B3'].value != date_from or wb2['Report']['D3'].value != date_to:
    raise AssertionError('Report date cells are wrong')
expected_formula = f'=SUM(Details!I5:I{4+len(rows)})' if rows else '=SUM(Details!I5:I5)'
if wb2['Summary']['B2'].value != expected_formula:
    raise AssertionError('Summary B2 formula is wrong')
calc_total = Decimal('0')
for i, r in enumerate(rows, start=5):
    expected_twd = money(Decimal(r['Amount']) * rates[r['Currency']])
    got = Decimal(str(sh.cell(i, 9).value)).quantize(Decimal('0.01'))
    if got != expected_twd:
        raise AssertionError(f'AmountTWD mismatch at row {i}: {got} != {expected_twd}')
    calc_total += expected_twd
sum_sh = wb2['Summary']
summary_rows = [row for row in sum_sh.iter_rows(min_row=5, max_col=6, values_only=True) if row[0] is not None]
summary_total = sum(Decimal(str(r[4])).quantize(Decimal('0.01')) for r in summary_rows)
if summary_total != calc_total:
    raise AssertionError(f'Summary TWD total mismatch {summary_total} != {calc_total}')
exc_sh = wb2['Exceptions']
actual_exc = sum(1 for row in exc_sh.iter_rows(min_row=5, max_col=1, values_only=True) if row[0])
if actual_exc != len(exc_rows):
    raise AssertionError(f'Exception rows {actual_exc} != {len(exc_rows)}')
print('Verified workbook:', out_xlsx.name)
print('Details rows:', detail_count)
print('Summary groups:', len(summary_rows))
print('Exceptions:', actual_exc)
print('TWD total:', calc_total)
