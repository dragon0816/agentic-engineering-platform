from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# workbook structure
wb_path = ws/'input'/'project_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('\nSHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)
# csv sample
csv_path = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r): print(row)
# manual sample/full enough
md_path=ws/'input'/'sap_manual.md'
print('\nsap_manual.md:')
text=md_path.read_text(encoding='utf-8')
print(text[:4000])