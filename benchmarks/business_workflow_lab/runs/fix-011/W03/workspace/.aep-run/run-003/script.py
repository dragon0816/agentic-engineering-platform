from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
import csv, shutil, sys, json

BASE = Path.cwd()
IN = BASE / 'input'
OUT = BASE / 'output'
DL = OUT / 'downloads'
OUT.mkdir(exist_ok=True)

TEMPLATE = IN / 'project_template.xlsx'
INV_CSV = DL / 'invoices_2026-09.csv'
FX_CSV = IN / 'fx_rates.csv'
OUT_XLSX = OUT / 'project_invoice_report_2026-09.xlsx'
AS_OF = date(2026, 9, 30)
PERIOD_FROM = date(2026, 9, 1)
PERIOD_TO = date(2026, 9, 30)

# Helpers
def q2(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def parse_date(s):
    if not s:
        return None
    return datetime.strptime(s, '%Y-%m-%d').date()

def clear_data_rows(ws, start=5):
    # Clear existing data rows only; keep row 1-4 and formats/layout as much as possible.
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)

def set_text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.number_format = '@'

def set_date(cell, value):
    cell.value = value
    cell.number_format = 'yyyy-mm-dd'

def set_money(cell, value):
    cell.value = float(q2(value)) if value is not None else None
    cell.number_format = '0.00'

# Read and print true inputs
print('Using template:', TEMPLATE)
print('Using invoices CSV:', INV_CSV)
print('Using fx CSV:', FX_CSV)

# Read FX; tolerate malformed Note commas by using first two columns only
rates = {}
with open(FX_CSV, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('rates:', rates)

# Read invoices
with open(INV_CSV, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('invoice headers:', reader.fieldnames)
    invoices = list(reader)
print('invoice count from CSV:', len(invoices))
print('invoice sample:', invoices[:3])

# Validate filter data and calculate
calc_rows = []
summary = {}
exceptions = []
for r in invoices:
    inv_date = parse_date(r['InvoiceDate'])
    due_date = parse_date(r['DueDate'])
    if not (PERIOD_FROM <= inv_date <= PERIOD_TO):
        raise ValueError(f"Invoice outside requested period: {r}")
    currency = r['Currency']
    if currency not in rates:
        raise ValueError(f"Missing FX rate for {currency}")
    amount = q2(r['Amount'])
    paid = q2(r['PaidAmount'])
    balance = q2(amount - paid)
    # Verify source Balance, then use computed amount per instruction
    source_balance = q2(r['Balance'])
    if source_balance != balance:
        raise ValueError(f"Balance mismatch for {r['InvoiceID']}: source {source_balance}, calc {balance}")
    balance_twd = q2(balance * rates[currency])
    row = {
        'InvoiceID': r['InvoiceID'],
        'OrderID': r['OrderID'],
        'ProjectID': r['ProjectID'],
        'CustomerID': r['CustomerID'],
        'InvoiceDate': inv_date,
        'DueDate': due_date,
        'Currency': currency,
        'Amount': amount,
        'PaidAmount': paid,
        'Balance': balance,
        'BalanceTWD': balance_twd,
    }
    calc_rows.append(row)
    # Summary by actual project only; blank ProjectID cannot be assigned to a project.
    if row['ProjectID']:
        summary.setdefault(row['ProjectID'], Decimal('0.00'))
        summary[row['ProjectID']] += balance_twd
    if balance > 0 and due_date < AS_OF:
        exceptions.append({'Source':'Invoices','DocumentID':r['InvoiceID'],'ProjectID':r['ProjectID'],'Reason':'Overdue'})

print('calculated invoice rows:', len(calc_rows))
print('summary projects:', {k: str(q2(v)) for k,v in sorted(summary.items())})
print('exceptions:', exceptions)

# Load template and fill a copy
wb = load_workbook(TEMPLATE)
# Preserve template sheets; clear any data rows in detail/summary/exceptions
for sheet in ['Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions']:
    clear_data_rows(wb[sheet], 5)

# Report metadata
rep = wb['Report']
rep['B3'] = PERIOD_FROM
rep['B3'].number_format = 'yyyy-mm-dd'
rep['D3'] = PERIOD_TO
rep['D3'].number_format = 'yyyy-mm-dd'
rep['B4'] = 'Invoices InvoiceDate 2026-09-01 to 2026-09-30; as-of 2026-09-30; overdue when Balance > 0 and DueDate < 2026-09-30.'
rep['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B7'] = 'Lab ERP Invoices; fx_rates.csv'
rep['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B9'] = 'Orders, Expenses, Quotes left blank per request. Summary excludes invoices with blank ProjectID because they are not attributable to a project.'

# Invoices sheet
inv_ws = wb['Invoices']
for i, r in enumerate(calc_rows, start=5):
    vals = [r['InvoiceID'], r['OrderID'], r['ProjectID'], r['CustomerID'], r['InvoiceDate'], r['DueDate'], r['Currency'], r['Amount'], r['PaidAmount'], r['Balance'], r['BalanceTWD']]
    # A-D text IDs
    set_text(inv_ws.cell(i,1), vals[0])
    set_text(inv_ws.cell(i,2), vals[1])
    set_text(inv_ws.cell(i,3), vals[2])
    set_text(inv_ws.cell(i,4), vals[3])
    set_date(inv_ws.cell(i,5), vals[4])
    set_date(inv_ws.cell(i,6), vals[5])
    inv_ws.cell(i,7).value = vals[6]
    for c, val in zip(range(8,12), vals[7:]):
        set_money(inv_ws.cell(i,c), val)

# Summary sheet: one row per project with invoice balance TWD; other sources blank
sum_ws = wb['Summary']
for i, (project, bal_twd) in enumerate(sorted(summary.items()), start=5):
    set_text(sum_ws.cell(i,1), project)
    # Other source columns intentionally blank per request
    sum_ws.cell(i,2).value = None
    set_money(sum_ws.cell(i,3), q2(bal_twd))
    sum_ws.cell(i,4).value = None
    sum_ws.cell(i,5).value = None
last_summary_row = max(5, 4 + len(summary))
# Extend totals formulas to last summary row; if no data, B5:B5 per instruction.
for col in range(2,6):
    col_letter = get_column_letter(col)
    sum_ws.cell(2,col).value = f'=SUM({col_letter}5:{col_letter}{last_summary_row})'

# Exceptions sheet
ex_ws = wb['Exceptions']
for i, r in enumerate(exceptions, start=5):
    ex_ws.cell(i,1).value = r['Source']
    set_text(ex_ws.cell(i,2), r['DocumentID'])
    set_text(ex_ws.cell(i,3), r['ProjectID'])
    ex_ws.cell(i,4).value = r['Reason']

# Ensure other source sheets have only header rows (no data rows)
for sheet in ['Orders','Expenses','Quotes']:
    if wb[sheet].max_row > 4:
        raise AssertionError(f'{sheet} unexpectedly has data rows')

wb.save(OUT_XLSX)
print('saved', OUT_XLSX)

# Reopen and validate every stated requirement
wb2 = load_workbook(OUT_XLSX, data_only=False)
checks = []
# Sheet presence/name check
expected_sheets = ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']
assert wb2.sheetnames == expected_sheets, wb2.sheetnames
checks.append('all original worksheet names preserved')
# Other source worksheets blank
for sheet in ['Orders','Expenses','Quotes']:
    sh = wb2[sheet]
    assert sh.max_row == 4, f'{sheet} max_row={sh.max_row}'
checks.append('Orders, Expenses, Quotes have no rows below headers')
# Invoices check
sh = wb2['Invoices']
headers = [sh.cell(4,c).value for c in range(1,12)]
assert headers == ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], headers
actual_inv_count = sh.max_row - 4
assert actual_inv_count == 13, actual_inv_count
for idx, source in enumerate(calc_rows, start=5):
    assert isinstance(sh.cell(idx,1).value, str) and sh.cell(idx,1).value.startswith('009'), sh.cell(idx,1).value
    assert isinstance(sh.cell(idx,2).value, str) and sh.cell(idx,2).value.startswith('004'), sh.cell(idx,2).value
    assert sh.cell(idx,5).value == source['InvoiceDate'], (idx, sh.cell(idx,5).value, source['InvoiceDate'])
    assert sh.cell(idx,6).value == source['DueDate'], (idx, sh.cell(idx,6).value, source['DueDate'])
    assert round(float(sh.cell(idx,10).value),2) == float(source['Balance']), (idx, sh.cell(idx,10).value, source['Balance'])
    assert round(float(sh.cell(idx,11).value),2) == float(source['BalanceTWD']), (idx, sh.cell(idx,11).value, source['BalanceTWD'])
checks.append('Invoices sheet has 13 September rows, text IDs, dates, numeric balances and BalanceTWD')
# Exceptions: exactly overdue with due date strictly before as-of
ex = wb2['Exceptions']
ex_rows = []
for row in ex.iter_rows(min_row=5, values_only=True):
    if any(v is not None and v != '' for v in row):
        ex_rows.append(row)
expected_ex = [(e['Source'], e['DocumentID'], e['ProjectID'], e['Reason']) for e in exceptions]
assert ex_rows == expected_ex, (ex_rows, expected_ex)
assert all(r[3] == 'Overdue' for r in ex_rows)
assert '0090000211' not in [r[1] for r in ex_rows] and '0090000213' not in [r[1] for r in ex_rows] and '0090000219' not in [r[1] for r in ex_rows]
checks.append('Exceptions marks only Balance > 0 and DueDate < 2026-09-30 as Overdue; due on 2026-09-30 excluded')
# Summary check
sumsh = wb2['Summary']
summary_rows = []
for r in range(5, sumsh.max_row+1):
    vals = [sumsh.cell(r,c).value for c in range(1,6)]
    if any(v is not None and v != '' for v in vals):
        summary_rows.append(vals)
expected_projects = sorted(summary.items())
assert len(summary_rows) == len(expected_projects), (summary_rows, expected_projects)
for vals, (project, amt) in zip(summary_rows, expected_projects):
    assert vals[0] == project, vals
    assert vals[1] in (None, ''), vals
    assert round(float(vals[2]),2) == float(q2(amt)), (vals, amt)
    assert vals[3] in (None, '') and vals[4] in (None, ''), vals
# formulas
last = max(5, 4+len(summary))
assert sumsh['B2'].value == f'=SUM(B5:B{last})'
assert sumsh['C2'].value == f'=SUM(C5:C{last})'
assert sumsh['D2'].value == f'=SUM(D5:D{last})'
assert sumsh['E2'].value == f'=SUM(E5:E{last})'
checks.append('Summary has one row per nonblank ProjectID and sums unpaid BalanceTWD; other source summary columns blank; totals formulas extended')
# Print summary rows for final
print('CHECKS PASSED:')
for c in checks:
    print('-', c)
print('output:', OUT_XLSX)
print('final summary rows:')
for vals in summary_rows:
    print(vals)
print('final exception rows:')
for vals in ex_rows:
    print(vals)
