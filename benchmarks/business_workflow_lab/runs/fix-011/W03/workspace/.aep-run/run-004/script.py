from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
import csv

BASE = Path.cwd()
IN = BASE / 'input'
OUT = BASE / 'output'
DL = OUT / 'downloads'
OUT.mkdir(exist_ok=True)
TEMPLATE = IN / 'project_template.xlsx'
INV_CSV = DL / 'invoices_2026-09.csv'
FX_CSV = IN / 'fx_rates.csv'
OUT_XLSX = OUT / 'project_invoice_report_2026-09.xlsx'
AS_OF = date(2026, 9, 30)
PERIOD_FROM = date(2026, 9, 1)
PERIOD_TO = date(2026, 9, 30)

def q2(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def parse_date(s):
    return datetime.strptime(s, '%Y-%m-%d').date() if s else None

def as_date(v):
    if isinstance(v, datetime): return v.date()
    return v

def clear_data_rows(ws, start=5):
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)

def set_text(cell, value):
    cell.value = '' if value is None else str(value)
    cell.number_format = '@'

def set_date(cell, value):
    cell.value = value
    cell.number_format = 'yyyy-mm-dd'

def set_money(cell, value):
    cell.value = float(q2(value)) if value is not None else None
    cell.number_format = '0.00'

print('Inputs in workspace:')
for p in sorted((BASE/'input').glob('*')):
    print('-', p.name)
print('Downloaded invoice file:', INV_CSV, 'exists=', INV_CSV.exists())

rates = {}
with open(FX_CSV, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('fx rates:', {k: str(v) for k,v in rates.items()})

with open(INV_CSV, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('invoice headers:', reader.fieldnames)
    invoices = list(reader)
print('invoice row count:', len(invoices))
print('first rows:', invoices[:2])

calc_rows = []
summary = {}
exceptions = []
for r in invoices:
    inv_date = parse_date(r['InvoiceDate'])
    due_date = parse_date(r['DueDate'])
    if not (PERIOD_FROM <= inv_date <= PERIOD_TO):
        raise ValueError(f"Invoice outside requested period: {r['InvoiceID']} {inv_date}")
    currency = r['Currency']
    if currency not in rates:
        raise ValueError(f"Missing FX rate for {currency}")
    amount = q2(r['Amount'])
    paid = q2(r['PaidAmount'])
    balance = q2(amount - paid)
    if q2(r['Balance']) != balance:
        raise ValueError(f"Balance mismatch for {r['InvoiceID']}")
    balance_twd = q2(balance * rates[currency])
    row = dict(r)
    row.update({'InvoiceDate': inv_date, 'DueDate': due_date, 'Amount': amount, 'PaidAmount': paid, 'Balance': balance, 'BalanceTWD': balance_twd})
    calc_rows.append(row)
    if r['ProjectID']:
        summary.setdefault(r['ProjectID'], Decimal('0.00'))
        summary[r['ProjectID']] += balance_twd
    if balance > 0 and due_date < AS_OF:
        exceptions.append({'Source':'Invoices','DocumentID':r['InvoiceID'],'ProjectID':r['ProjectID'],'Reason':'Overdue'})

wb = load_workbook(TEMPLATE)
for sheet in ['Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions']:
    clear_data_rows(wb[sheet], 5)

rep = wb['Report']
rep['B3'] = PERIOD_FROM; rep['B3'].number_format = 'yyyy-mm-dd'
rep['D3'] = PERIOD_TO; rep['D3'].number_format = 'yyyy-mm-dd'
rep['B4'] = 'Invoices InvoiceDate 2026-09-01 to 2026-09-30; as-of 2026-09-30; overdue when Balance > 0 and DueDate < 2026-09-30.'
rep['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B7'] = 'Lab ERP Invoices; fx_rates.csv'
rep['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')
rep['B9'] = 'Orders, Expenses, Quotes left blank per request. Summary excludes invoices with blank ProjectID because they are not attributable to a project.'

inv_ws = wb['Invoices']
for i, r in enumerate(calc_rows, start=5):
    set_text(inv_ws.cell(i,1), r['InvoiceID'])
    set_text(inv_ws.cell(i,2), r['OrderID'])
    set_text(inv_ws.cell(i,3), r['ProjectID'])
    set_text(inv_ws.cell(i,4), r['CustomerID'])
    set_date(inv_ws.cell(i,5), r['InvoiceDate'])
    set_date(inv_ws.cell(i,6), r['DueDate'])
    inv_ws.cell(i,7).value = r['Currency']
    set_money(inv_ws.cell(i,8), r['Amount'])
    set_money(inv_ws.cell(i,9), r['PaidAmount'])
    set_money(inv_ws.cell(i,10), r['Balance'])
    set_money(inv_ws.cell(i,11), r['BalanceTWD'])

sum_ws = wb['Summary']
for i, (project, bal_twd) in enumerate(sorted(summary.items()), start=5):
    set_text(sum_ws.cell(i,1), project)
    sum_ws.cell(i,2).value = None
    set_money(sum_ws.cell(i,3), q2(bal_twd))
    sum_ws.cell(i,4).value = None
    sum_ws.cell(i,5).value = None
last_summary_row = max(5, 4 + len(summary))
for col, letter in [(2,'B'),(3,'C'),(4,'D'),(5,'E')]:
    sum_ws.cell(2,col).value = f'=SUM({letter}5:{letter}{last_summary_row})'

ex_ws = wb['Exceptions']
for i, r in enumerate(exceptions, start=5):
    ex_ws.cell(i,1).value = r['Source']
    set_text(ex_ws.cell(i,2), r['DocumentID'])
    set_text(ex_ws.cell(i,3), r['ProjectID'])
    ex_ws.cell(i,4).value = r['Reason']

wb.save(OUT_XLSX)
print('saved:', OUT_XLSX)

# Reopen checks
wb2 = load_workbook(OUT_XLSX, data_only=False)
checks=[]
assert wb2.sheetnames == ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
checks.append('all original worksheet names preserved')
for sheet in ['Orders','Expenses','Quotes']:
    assert wb2[sheet].max_row == 4, f'{sheet} not blank below row 4'
checks.append('Orders, Expenses, Quotes remain blank below headers')
sh=wb2['Invoices']
assert sh.max_row - 4 == len(calc_rows) == 13
for idx, source in enumerate(calc_rows, start=5):
    assert isinstance(sh.cell(idx,1).value, str)
    assert isinstance(sh.cell(idx,2).value, str)
    assert isinstance(sh.cell(idx,3).value, str)
    assert isinstance(sh.cell(idx,4).value, str)
    assert as_date(sh.cell(idx,5).value) == source['InvoiceDate']
    assert as_date(sh.cell(idx,6).value) == source['DueDate']
    assert round(float(sh.cell(idx,8).value),2) == float(source['Amount'])
    assert round(float(sh.cell(idx,9).value),2) == float(source['PaidAmount'])
    assert round(float(sh.cell(idx,10).value),2) == float(source['Balance'])
    assert round(float(sh.cell(idx,11).value),2) == float(source['BalanceTWD'])
checks.append('Invoices has 13 September invoices with text IDs, date cells, numeric Amount/Paid/Balance/BalanceTWD')
ex=wb2['Exceptions']
ex_rows=[]
for row in ex.iter_rows(min_row=5, values_only=True):
    if any(v is not None and v != '' for v in row):
        ex_rows.append(row)
expected_ex=[(e['Source'],e['DocumentID'],e['ProjectID'],e['Reason']) for e in exceptions]
assert ex_rows == expected_ex, (ex_rows, expected_ex)
for same_day in ['0090000211','0090000213','0090000219']:
    assert same_day not in [r[1] for r in ex_rows]
checks.append('Exceptions contains only overdue rows where Balance > 0 and DueDate is before 2026-09-30; same-day due invoices are not marked')
sumsh=wb2['Summary']
summary_rows=[]
for r in range(5, sumsh.max_row+1):
    vals=[sumsh.cell(r,c).value for c in range(1,6)]
    if any(v is not None and v != '' for v in vals):
        summary_rows.append(vals)
expected=sorted(summary.items())
assert len(summary_rows)==len(expected)
for vals,(project,amt) in zip(summary_rows, expected):
    assert vals[0]==project
    assert vals[1] in (None,'')
    assert round(float(vals[2]),2)==float(q2(amt))
    assert vals[3] in (None,'')
    assert vals[4] in (None,'')
last=max(5,4+len(expected))
assert [sumsh.cell(2,c).value for c in range(2,6)] == [f'=SUM(B5:B{last})', f'=SUM(C5:C{last})', f'=SUM(D5:D{last})', f'=SUM(E5:E{last})']
checks.append('Summary has project-level unpaid TWD in InvoiceBalanceTWD only, with total formulas extended')
assert as_date(wb2['Report']['B3'].value)==PERIOD_FROM and as_date(wb2['Report']['D3'].value)==PERIOD_TO
checks.append('Report period metadata set to 2026-09-01 through 2026-09-30')
print('CHECKS PASSED:')
for c in checks: print('-', c)
print('SUMMARY_ROWS')
for vals in summary_rows: print(vals)
print('EXCEPTION_ROWS')
for vals in ex_rows: print(vals)
print('OUTPUT_FILE', OUT_XLSX)
