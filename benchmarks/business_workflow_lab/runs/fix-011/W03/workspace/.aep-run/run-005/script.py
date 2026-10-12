from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
import csv

BASE=Path.cwd(); OUT=BASE/'output'; IN=BASE/'input'; DL=OUT/'downloads'
OUT_XLSX=OUT/'project_invoice_report_2026-09.xlsx'
INV_CSV=DL/'invoices_2026-09.csv'
FX_CSV=IN/'fx_rates.csv'
AS_OF=date(2026,9,30); PERIOD_FROM=date(2026,9,1); PERIOD_TO=date(2026,9,30)

def q2(x): return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
def parse_date(s): return datetime.strptime(s,'%Y-%m-%d').date()
def as_date(v): return v.date() if isinstance(v, datetime) else v

# Reconstruct expected data from saved source files
rates={}
with open(FX_CSV,newline='',encoding='utf-8-sig') as f:
    r=csv.reader(f); next(r)
    for row in r:
        if len(row)>=2 and row[0]: rates[row[0]]=Decimal(row[1])
with open(INV_CSV,newline='',encoding='utf-8-sig') as f:
    invoices=list(csv.DictReader(f))
calc=[]; summary={}; exceptions=[]
for r in invoices:
    inv_date=parse_date(r['InvoiceDate']); due_date=parse_date(r['DueDate'])
    assert PERIOD_FROM <= inv_date <= PERIOD_TO
    amount=q2(r['Amount']); paid=q2(r['PaidAmount']); bal=q2(amount-paid); assert bal==q2(r['Balance'])
    bal_twd=q2(bal*rates[r['Currency']])
    row={**r,'InvoiceDate':inv_date,'DueDate':due_date,'Amount':amount,'PaidAmount':paid,'Balance':bal,'BalanceTWD':bal_twd}
    calc.append(row)
    if r['ProjectID']:
        summary.setdefault(r['ProjectID'],Decimal('0.00')); summary[r['ProjectID']]+=bal_twd
    if bal>0 and due_date<AS_OF:
        exceptions.append(('Invoices', r['InvoiceID'], r['ProjectID'] or None, 'Overdue'))

wb=load_workbook(OUT_XLSX, data_only=False)
checks=[]
assert wb.sheetnames==['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
checks.append('all worksheet names preserved')
# Report cells
assert as_date(wb['Report']['B3'].value)==PERIOD_FROM
assert as_date(wb['Report']['D3'].value)==PERIOD_TO
assert '2026-09-30' in str(wb['Report']['B4'].value)
checks.append('Report period/condition metadata set')
# Source sheets blank
for s in ['Orders','Expenses','Quotes']:
    sh=wb[s]
    assert sh.max_row==4, (s, sh.max_row)
checks.append('Orders, Expenses and Quotes contain only template header rows')
# Invoice detail
sh=wb['Invoices']
assert [sh.cell(4,c).value for c in range(1,12)]==['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
assert sh.max_row-4==13
for idx, r in enumerate(calc, start=5):
    assert sh.cell(idx,1).value==r['InvoiceID'] and isinstance(sh.cell(idx,1).value,str)
    assert sh.cell(idx,2).value==r['OrderID'] and isinstance(sh.cell(idx,2).value,str)
    # Missing ProjectID is saved as a blank cell (None); nonmissing is text
    cell_project=sh.cell(idx,3).value
    assert (cell_project==r['ProjectID']) or (r['ProjectID']=='' and cell_project is None), (idx, cell_project, r['ProjectID'])
    if r['ProjectID']:
        assert isinstance(cell_project,str)
    assert sh.cell(idx,4).value==r['CustomerID'] and isinstance(sh.cell(idx,4).value,str)
    assert as_date(sh.cell(idx,5).value)==r['InvoiceDate']
    assert as_date(sh.cell(idx,6).value)==r['DueDate']
    assert sh.cell(idx,7).value==r['Currency']
    assert round(float(sh.cell(idx,8).value),2)==float(r['Amount'])
    assert round(float(sh.cell(idx,9).value),2)==float(r['PaidAmount'])
    assert round(float(sh.cell(idx,10).value),2)==float(r['Balance'])
    assert round(float(sh.cell(idx,11).value),2)==float(r['BalanceTWD'])
checks.append('Invoices sheet has all 13 downloaded September invoices, text IDs, date fields, and numeric TWD balances')
# Summary
sumsh=wb['Summary']
summary_rows=[]
for rr in range(5,sumsh.max_row+1):
    vals=[sumsh.cell(rr,c).value for c in range(1,6)]
    if any(v is not None and v!='' for v in vals): summary_rows.append(vals)
expected=sorted(summary.items())
assert len(summary_rows)==len(expected)
for vals,(project,amt) in zip(summary_rows,expected):
    assert vals[0]==project
    assert vals[1] is None
    assert round(float(vals[2]),2)==float(q2(amt))
    assert vals[3] is None and vals[4] is None
last=max(5,4+len(expected))
assert [sumsh.cell(2,c).value for c in range(2,6)]==[f'=SUM(B5:B{last})',f'=SUM(C5:C{last})',f'=SUM(D5:D{last})',f'=SUM(E5:E{last})']
checks.append('Summary groups unpaid BalanceTWD by nonblank ProjectID; other source amount columns remain blank; total formulas extended')
# Exceptions
ex=wb['Exceptions']; ex_rows=[]
for row in ex.iter_rows(min_row=5,values_only=True):
    if any(v is not None and v!='' for v in row): ex_rows.append(row)
# Blank ProjectID is None in workbook
assert ex_rows==exceptions, (ex_rows, exceptions)
assert all(row[3]=='Overdue' for row in ex_rows)
assert all(row[1] not in ['0090000211','0090000213','0090000219'] for row in ex_rows)
checks.append('Exceptions marks only Balance > 0 and DueDate before 2026-09-30; invoices due on 2026-09-30 are not overdue')

print('CHECKS PASSED')
for c in checks: print('-', c)
print('OUTPUT_FILE', OUT_XLSX)
print('INVOICE_COUNT', len(calc))
print('SUMMARY_ROWS')
for vals in summary_rows: print(vals)
print('EXCEPTION_ROWS')
for vals in ex_rows: print(vals)
