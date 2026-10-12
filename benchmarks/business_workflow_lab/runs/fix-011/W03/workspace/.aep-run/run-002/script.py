from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('Downloaded files:')
for p in sorted((ws/'output'/'downloads').glob('*')):
    print(p.name, p.stat().st_size)
    print(p.read_text(encoding='utf-8-sig')[:2000])
print('\nInstructions full:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)