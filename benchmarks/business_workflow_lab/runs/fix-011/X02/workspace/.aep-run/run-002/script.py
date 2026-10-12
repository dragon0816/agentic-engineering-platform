from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True):
    print(r)
