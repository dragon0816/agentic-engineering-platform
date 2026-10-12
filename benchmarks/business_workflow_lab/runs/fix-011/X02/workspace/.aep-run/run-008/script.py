from pathlib import Path
import csv, json, os
from openpyxl import load_workbook

ws = Path.cwd()
out = ws / 'output'
dl = out / 'downloads'
print('workspace=', ws)
print('output exists=', out.exists())
print('output files:')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(ws), p.stat().st_size)

report = out / 'project_report_C001_2026-09-01_to_2026-09-30.xlsx'
assert report.exists(), f'missing report {report}'
wb = load_workbook(report, data_only=False)
print('sheets=', wb.sheetnames)

# Print dimensions, headers, and sample rows for every sheet.
for sh in wb.sheetnames:
    wsht = wb[sh]
    print('\nSHEET', sh, 'rows', wsht.max_row, 'cols', wsht.max_column)
    headers = [wsht.cell(1, c).value for c in range(1, wsht.max_column+1)]
    print('headers=', headers)
    for r in range(2, min(wsht.max_row, 5)+1):
        print('row', r, [wsht.cell(r, c).value for c in range(1, wsht.max_column+1)])

expected = {
    'Orders': 3,
    'Invoices': 3,
    'Expenses': 9,
    'Quotes': 4,
    'Summary': 2,
    'Exceptions': 3,
}
actual = {}
for sh, n in expected.items():
    assert sh in wb.sheetnames, f'missing sheet {sh}'
    actual[sh] = max(0, wb[sh].max_row - 1)
print('\nrow_counts=', actual)
for sh, n in expected.items():
    assert actual[sh] == n, f'{sh} row count {actual[sh]} != {n}'

# Requirement checks.
orders = wb['Orders']
headers = [orders.cell(1,c).value for c in range(1, orders.max_column+1)]
status_col = headers.index('Status') + 1
order_id_col = headers.index('OrderID') + 1
for r in range(2, orders.max_row+1):
    assert orders.cell(r, status_col).value != 'Cancelled', f'Cancelled order in row {r}'
    assert isinstance(orders.cell(r, order_id_col).value, str), f'OrderID not text row {r}'

expenses = wb['Expenses']
eh = [expenses.cell(1,c).value for c in range(1, expenses.max_column+1)]
proj_col = eh.index('ProjectID') + 1
status_col = eh.index('Status') + 1
exp_id_col = eh.index('ExpenseID') + 1
allowed_projects = {'P001','P003'}
for r in range(2, expenses.max_row+1):
    assert expenses.cell(r, proj_col).value in allowed_projects, f'expense non-C001 project row {r}'
    assert expenses.cell(r, status_col).value == 'Approved', f'expense not Approved row {r}'
    assert isinstance(expenses.cell(r, exp_id_col).value, str), f'ExpenseID not text row {r}'

quotes = wb['Quotes']
qh = [quotes.cell(1,c).value for c in range(1, quotes.max_column+1)]
quote_statuses = [quotes.cell(r, qh.index('Status')+1).value for r in range(2, quotes.max_row+1)]
print('quote_statuses=', quote_statuses)
assert len(quote_statuses) == 4, 'quotes not all retained count mismatch'

# Amount numeric / TWD conversion columns if present.
for sh in ['Orders','Invoices','Expenses','Quotes']:
    sht = wb[sh]
    hdr = [sht.cell(1,c).value for c in range(1, sht.max_column+1)]
    for cname in hdr:
        if cname and ('TWD' in str(cname) or str(cname) in {'Amount','PaidAmount','Balance'}):
            cidx = hdr.index(cname)+1
            vals = [sht.cell(r,cidx).value for r in range(2, sht.max_row+1) if sht.cell(r,cidx).value is not None]
            if vals:
                assert all(isinstance(v, (int,float)) or (isinstance(v,str) and v.startswith('=')) for v in vals), f'{sh}.{cname} has non-numeric values {vals}'

exceptions = wb['Exceptions']
xh = [exceptions.cell(1,c).value for c in range(1, exceptions.max_column+1)]
print('exception_headers=', xh)
exception_rows = []
for r in range(2, exceptions.max_row+1):
    row = {xh[c-1]: exceptions.cell(r,c).value for c in range(1, exceptions.max_column+1)}
    exception_rows.append(row)
print('exception_rows=', json.dumps(exception_rows, ensure_ascii=False, default=str, indent=2))
ids = json.dumps(exception_rows, ensure_ascii=False)
for required in ['EXP-00006','Q0010','Q0014']:
    assert required in ids, f'missing exception {required}'
assert '0090000219' not in ids, 'invoice due on cutoff should not be overdue before/through 2026-09-30'

# Workflow downloads should exist from approved run.
for name in ['sap_orders_workflow.csv', 'sap_invoices_workflow.csv']:
    path = dl / name
    assert path.exists(), f'missing workflow download {name}'
    print(name, 'bytes=', path.stat().st_size)

print('\nCHECK PASSED: report and workflow output files satisfy stated filters and attention rules.')