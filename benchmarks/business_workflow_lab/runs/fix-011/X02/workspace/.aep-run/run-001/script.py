import os, csv
from pathlib import Path
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace=', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(8), rdr):
        print(row)
print('\nmanuals preview:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    print(f'--- {name} ---')
    txt=(ws/'input'/name).read_text(encoding='utf-8')
    print(txt[:3000])
print('\ntemplate workbook:')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
print('sheets=', wb.sheetnames)
for sh in wb.worksheets:
    print('\nSHEET', sh.title, 'max', sh.max_row, sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(r)
