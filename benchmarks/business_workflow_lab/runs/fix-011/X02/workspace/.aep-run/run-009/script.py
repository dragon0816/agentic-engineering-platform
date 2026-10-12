from pathlib import Path
import json
from openpyxl import load_workbook

base = Path.cwd()
out = base / 'output'
report = out / 'project_report_C001_2026-09-01_to_2026-09-30.xlsx'
assert report.exists(), f'missing report: {report}'
wb = load_workbook(report, data_only=False)
print('report=', report)
print('sheets=', wb.sheetnames)

def table(sheet_name):
    ws = wb[sheet_name]
    headers = [ws.cell(4, c).value for c in range(1, ws.max_column+1)]
    rows = []
    for r in range(5, ws.max_row+1):
        if any(ws.cell(r, c).value is not None for c in range(1, ws.max_column+1)):
            rows.append({headers[c-1]: ws.cell(r,c).value for c in range(1, ws.max_column+1)})
    return headers, rows

expected_counts = {
    'Orders': 3,
    'Invoices': 3,
    'Expenses': 9,
    'Quotes': 4,
    'Summary': 2,
    'Exceptions': 3,
}
for s, expected in expected_counts.items():
    hdr, rows = table(s)
    print(f'{s}: headers={hdr}; data_rows={len(rows)}')
    print(' sample=', rows[:3])
    assert len(rows) == expected, f'{s} data rows {len(rows)} != {expected}'

# Orders: C001, September, excluding Cancelled; IDs kept as text.
_, orders = table('Orders')
for row in orders:
    assert row['CustomerID'] == 'C001', f"wrong order customer {row}"
    assert row['Status'] != 'Cancelled', f"cancelled order included {row}"
    assert '2026-09-01' <= row['OrderDate'] <= '2026-09-30', f"order outside month {row}"
    assert isinstance(row['OrderID'], str), f"OrderID not text {row}"
    assert isinstance(row['AmountTWD'], (int, float)), f"Order AmountTWD not numeric {row}"

# Invoices: C001, September, IDs text and TWD balance numeric.
_, invoices = table('Invoices')
for row in invoices:
    assert row['CustomerID'] == 'C001', f"wrong invoice customer {row}"
    assert '2026-09-01' <= row['InvoiceDate'] <= '2026-09-30', f"invoice outside month {row}"
    assert isinstance(row['InvoiceID'], str), f"InvoiceID not text {row}"
    assert isinstance(row['BalanceTWD'], (int, float)), f"Invoice BalanceTWD not numeric {row}"

# Expenses: approved only, only C001 projects P001/P003, September, IDs text and TWD amount numeric.
_, expenses = table('Expenses')
for row in expenses:
    assert row['ProjectID'] in {'P001', 'P003'}, f"wrong expense project {row}"
    assert row['Status'] == 'Approved', f"non-approved expense {row}"
    assert '2026-09-01' <= row['ExpenseDate'] <= '2026-09-30', f"expense outside month {row}"
    assert isinstance(row['ExpenseID'], str), f"ExpenseID not text {row}"
    assert isinstance(row['AmountTWD'], (int, float)), f"Expense AmountTWD not numeric {row}"

# Quotes: all statuses retained, C001, IDs text and AmountTWD numeric.
_, quotes = table('Quotes')
statuses = sorted({row['Status'] for row in quotes})
print('quote statuses retained=', statuses)
for row in quotes:
    assert row['CustomerID'] == 'C001', f"wrong quote customer {row}"
    assert isinstance(row['QuoteID'], str), f"QuoteID not text {row}"
    assert isinstance(row['AmountTWD'], (int, float)), f"Quote AmountTWD not numeric {row}"
assert {'Accepted','Expired','Sent'}.issubset(set(statuses)), 'not all expected quote statuses retained'

# Exceptions: exactly stated attention rules.
_, exceptions = table('Exceptions')
print('exceptions=', json.dumps(exceptions, ensure_ascii=False, default=str, indent=2))
ids = {str(r['DocumentID']) for r in exceptions}
assert ids == {'EXP-00006', 'Q0010', 'Q0014'}, f'unexpected exception IDs {ids}'
reasons = {str(r['Reason']) for r in exceptions}
assert any('Missing receipt' in r for r in reasons), 'missing receipt exception absent'
assert any('ValidUntil 2026-09-20' in r and 'not Accepted' in r for r in reasons), 'Q0010 quote exception rule absent'
assert any('ValidUntil 2026-09-28' in r and 'not Accepted' in r for r in reasons), 'Q0014 quote exception rule absent'

# No invoice overdue unpaid as of cutoff: due date must be < 2026-09-30 and balance > 0; verify none.
for row in invoices:
    if row['DueDate'] < '2026-09-30' and float(row['Balance']) > 0:
        raise AssertionError(f"overdue unpaid invoice should be exception: {row}")
assert '0090000219' not in ids, 'invoice due on cutoff is not before cutoff and should not be flagged as overdue'

# Template was filled, not rebuilt: Instructions sheet preserved and report sheet populated.
assert 'Instructions' in wb.sheetnames, 'Instructions sheet not preserved'
rep = wb['Report']
assert rep['B3'].value == '2026-09-01' and rep['D3'].value == '2026-09-30', 'Report dates wrong'
assert 'fx_rates.csv only' in str(rep['B4'].value), 'Report condition does not record attachment FX use'

# Workflow step outputs present.
for rel in [
    'output/downloads/sap_orders_workflow.csv',
    'output/downloads/sap_invoices_workflow.csv',
    'output/expense_all_approved_2026-09-01_to_2026-09-30.csv',
    'output/project_report_C001_2026-09-01_to_2026-09-30.xlsx',
]:
    p = base / rel
    assert p.exists() and p.stat().st_size > 0, f'missing/empty {rel}'
    print(rel, p.stat().st_size)

print('\nCHECK PASSED: workflow output matches the requested period, customer, filters, template, TWD conversion, and attention rules.')