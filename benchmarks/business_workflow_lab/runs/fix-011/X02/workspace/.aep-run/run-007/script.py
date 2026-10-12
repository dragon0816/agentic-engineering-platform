import csv, json, subprocess, sys, os, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

params=json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
customer_id=params['customer_id']
start=params['from_date']
end=params['to_date']
cutoff=params['attention_cutoff']
ws=Path.cwd(); input_dir=ws/'input'; out_dir=ws/'output'; dl_dir=out_dir/'downloads'; out_dir.mkdir(exist_ok=True); dl_dir.mkdir(exist_ok=True)

def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def money(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

fx={}
with open(input_dir/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f); next(rdr)
    for row in rdr:
        if len(row)>=2 and row[0]: fx[row[0]]=Decimal(row[1])

def twd(amount, currency):
    if currency not in fx: raise ValueError(f'Missing FX rate for {currency}')
    return (Decimal(str(amount))*fx[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def parse_date(s):
    return datetime.strptime(s, '%Y-%m-%d').date()

orders_path=dl_dir/'sap_orders_workflow.csv'
invoices_path=dl_dir/'sap_invoices_workflow.csv'
expenses_path=out_dir/f'expense_all_approved_{start}_to_{end}.csv'
print('input files:', orders_path, orders_path.exists(), invoices_path, invoices_path.exists(), expenses_path, expenses_path.exists())
orders_raw=read_csv(orders_path); invoices_raw=read_csv(invoices_path); expenses_all=read_csv(expenses_path)
orders=[r for r in orders_raw if r['CustomerID']==customer_id and start <= r['OrderDate'] <= end and r['Status']!='Cancelled']
projects=sorted({r['ProjectID'] for r in orders if r.get('ProjectID')})
invoices=[r for r in invoices_raw if r['CustomerID']==customer_id and start <= r['InvoiceDate'] <= end]
expenses=[r for r in expenses_all if r.get('ProjectID') in projects and r.get('Status')=='Approved' and start <= r.get('ExpenseDate','') <= end]
cmd=[sys.executable, r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py','quotes','list','--customer',customer_id,'--from',start,'--to',end,'--format','json']
cp=subprocess.run(cmd,capture_output=True,text=True)
print('quote_cli return', cp.returncode)
if cp.returncode!=0:
    print(cp.stderr); raise SystemExit(cp.returncode)
quotes=json.loads(cp.stdout)['items']

outfile=out_dir/f'project_report_{customer_id}_{start}_to_{end}.xlsx'
shutil.copy2(input_dir/'project_template.xlsx', outfile)
wb=load_workbook(outfile)
now=datetime.now().astimezone().isoformat(timespec='seconds')
rep=wb['Report']; rep['B3']=start; rep['D3']=end
rep['B4']=f'CustomerID={customer_id}; orders exclude Cancelled; expenses Approved only for customer projects {",".join(projects)}; quotes all statuses; FX from fx_rates.csv only; attention cutoff {cutoff}.'
rep['B5']=now; rep['B7']='Lab ERP; Expense Viewer; quote_cli; fx_rates.csv'; rep['D7']=now
rep['B9']='關注條件：截至 cutoff 已逾期未收發票（DueDate < cutoff 且 Balance > 0）、缺收據費用、ValidUntil < cutoff 且非 Accepted 的報價。'
text_cols={'Orders':['A','B','C','E','F'],'Invoices':['A','B','C','D','G'],'Expenses':['A','B','C','E','F'],'Quotes':['A','B','C','F','G'],'Summary':['A'],'Exceptions':['A','B','C','D']}
def clear(sheet):
    sh=wb[sheet]
    if sh.max_row>=5: sh.delete_rows(5, sh.max_row-4)
def write(sheet, rows):
    sh=wb[sheet]; clear(sheet)
    for i,row in enumerate(rows, start=5):
        for j,val in enumerate(row, start=1):
            cell=sh.cell(i,j,val); col=cell.column_letter
            if col in text_cols.get(sheet,[]): cell.number_format='@'; cell.value='' if val is None else str(val)
            elif isinstance(val, Decimal): cell.value=float(val); cell.number_format='0.00'
            if sheet in ['Orders','Invoices','Expenses','Quotes'] and col in ['G','H','I','J','K'] and isinstance(cell.value,(int,float)): cell.number_format='0.00'
orders_rows=[[r['OrderID'],r['ProjectID'],r['CustomerID'],r['OrderDate'],r['Status'],r['Currency'],money(r['NetAmount']),twd(r['NetAmount'],r['Currency'])] for r in orders]
invoice_rows=[[r['InvoiceID'],r['OrderID'],r['ProjectID'],r['CustomerID'],r['InvoiceDate'],r['DueDate'],r['Currency'],money(r['Amount']),money(r['PaidAmount']),money(r['Balance']),twd(r['Balance'],r['Currency'])] for r in invoices]
expense_rows=[[r['ExpenseID'],r['ProjectID'],r['EmployeeID'],r['ExpenseDate'],r['Status'],r['Currency'],money(r['Amount']),twd(r['Amount'],r['Currency'])] for r in expenses]
quote_rows=[[r['QuoteID'],r['ProjectID'],r['CustomerID'],r['QuoteDate'],r['ValidUntil'],r['Status'],r['Currency'],money(r['NetAmount']),twd(r['NetAmount'],r['Currency'])] for r in quotes]
for s,rows in [('Orders',orders_rows),('Invoices',invoice_rows),('Expenses',expense_rows),('Quotes',quote_rows)]: write(s,rows)
summary_projects=sorted(set([r[1] for r in orders_rows]+[r[2] for r in invoice_rows]+[r[1] for r in expense_rows]+[r[1] for r in quote_rows if r[5]=='Accepted']))
summary=[]
for p in summary_projects:
    summary.append([p, sum((r[7] for r in orders_rows if r[1]==p),Decimal('0.00')), sum((r[10] for r in invoice_rows if r[2]==p),Decimal('0.00')), sum((r[7] for r in expense_rows if r[1]==p),Decimal('0.00')), sum((r[8] for r in quote_rows if r[1]==p and r[5]=='Accepted'),Decimal('0.00'))])
write('Summary', summary); last=max(5,4+len(summary))
for c in ['B','C','D','E']: wb['Summary'][f'{c}2']=f'=SUM({c}5:{c}{last})'
cut=parse_date(cutoff); exceptions=[]
for r in invoices:
    if parse_date(r['DueDate']) < cut and Decimal(r['Balance']) > 0: exceptions.append(['Invoices',r['InvoiceID'],r['ProjectID'],f"Overdue unpaid invoice: DueDate {r['DueDate']} < {cutoff}, Balance {r['Balance']} {r['Currency']}"])
for r in expenses:
    if str(r.get('MissingReceipt','')).strip().lower()=='true': exceptions.append(['Expenses',r['ExpenseID'],r['ProjectID'],'Missing receipt'])
for r in quotes:
    if parse_date(r['ValidUntil']) < cut and r['Status']!='Accepted': exceptions.append(['Quotes',r['QuoteID'],r['ProjectID'],f"Expired/non-accepted quote: ValidUntil {r['ValidUntil']} < {cutoff}, Status {r['Status']}"])
write('Exceptions', exceptions)
fill=PatternFill(fill_type='solid', fgColor='FFF2CC'); ids={e[1] for e in exceptions}
for sheet in ['Invoices','Expenses','Quotes']:
    sh=wb[sheet]
    for row in range(5, sh.max_row+1):
        if sh.cell(row,1).value in ids:
            for col in range(1, sh.max_column+1): sh.cell(row,col).fill=fill
wb.save(outfile)
# verification
wb2=load_workbook(outfile, data_only=False)
def n(sheet): return max(0, wb2[sheet].max_row-4)
print('saved', outfile)
print('verified counts', {'Orders':n('Orders'),'Invoices':n('Invoices'),'Expenses':n('Expenses'),'Quotes':n('Quotes'),'Summary':n('Summary'),'Exceptions':n('Exceptions')})
print('projects', projects)
print('exceptions', exceptions)
