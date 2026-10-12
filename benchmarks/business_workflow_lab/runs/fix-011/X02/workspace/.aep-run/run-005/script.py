import csv, json, subprocess, sys, os, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from openpyxl import load_workbook
from openpyxl.styles import PatternFill

ws=Path.cwd()
input_dir=ws/'input'
out_dir=ws/'output'
dl_dir=out_dir/'downloads'
out_dir.mkdir(exist_ok=True)

customer_id='C001'
start='2026-09-01'
end='2026-09-30'
cutoff='2026-09-30'

print('inputs really used:')
for p in [input_dir/'project_template.xlsx', input_dir/'fx_rates.csv', dl_dir/'sap_orders_C001_2026-09.csv', dl_dir/'sap_invoices_C001_2026-09.csv', out_dir/'expense_all_approved_2026-09.csv']:
    print(p, 'exists=', p.exists(), 'bytes=', p.stat().st_size if p.exists() else None)

def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

# FX rates: use only attachment rates, first two columns are authoritative.
fx={}
with open(input_dir/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    header=next(rdr)
    print('fx header/sample=', header)
    for row in rdr:
        if len(row)>=2 and row[0]:
            fx[row[0]]=Decimal(row[1])
print('fx=', fx)

def money(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def twd(amount, currency):
    if currency not in fx:
        raise ValueError(f'Missing FX rate for {currency}')
    return (Decimal(str(amount))*fx[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def parse_date(s):
    return datetime.strptime(s, '%Y-%m-%d').date()

orders_raw=read_csv(dl_dir/'sap_orders_C001_2026-09.csv')
invoices_raw=read_csv(dl_dir/'sap_invoices_C001_2026-09.csv')
expenses_all=read_csv(out_dir/'expense_all_approved_2026-09.csv')
print('orders raw rows=', len(orders_raw), 'headers=', list(orders_raw[0].keys()) if orders_raw else [])
print('invoices raw rows=', len(invoices_raw), 'headers=', list(invoices_raw[0].keys()) if invoices_raw else [])
print('expenses all approved rows=', len(expenses_all), 'headers=', list(expenses_all[0].keys()) if expenses_all else [])

orders=[r for r in orders_raw if r['CustomerID']==customer_id and start <= r['OrderDate'] <= end and r['Status']!='Cancelled']
projects=sorted({r['ProjectID'] for r in orders if r.get('ProjectID')})
print('C001 projects from non-cancelled orders=', projects)
invoices=[r for r in invoices_raw if r['CustomerID']==customer_id and start <= r['InvoiceDate'] <= end]
expenses=[r for r in expenses_all if r.get('ProjectID') in projects and r.get('Status')=='Approved' and start <= r.get('ExpenseDate','') <= end]

quote_cmd=[sys.executable, r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py', 'quotes','list','--customer',customer_id,'--from',start,'--to',end,'--format','json']
cp=subprocess.run(quote_cmd, capture_output=True, text=True)
print('quote cli return=', cp.returncode)
if cp.returncode!=0:
    print('stderr=', cp.stderr)
    raise SystemExit(cp.returncode)
quotes_obj=json.loads(cp.stdout)
quotes=quotes_obj['items']
print('quotes rows=', len(quotes), 'headers=', list(quotes[0].keys()) if quotes else [])

# Copy and fill the exact template.
outfile=out_dir/f'project_report_{customer_id}_{start}_to_{end}.xlsx'
shutil.copy2(input_dir/'project_template.xlsx', outfile)
wb=load_workbook(outfile)

# Report metadata
rep=wb['Report']
now=datetime.now().astimezone().isoformat(timespec='seconds')
rep['B3']=start
rep['D3']=end
rep['B4']=f'CustomerID={customer_id}; orders exclude Cancelled; expenses Approved only for C001 projects {",".join(projects)}; quotes all statuses; FX from fx_rates.csv only; attention cutoff {cutoff}.'
rep['B5']=now
rep['B7']='Lab ERP; Expense Viewer; quote_cli; fx_rates.csv'
rep['D7']=now
rep['B9']='關注條件：截至 2026-09-30 已逾期未收發票（DueDate < cutoff 且 Balance > 0）、缺收據費用、ValidUntil < 2026-09-30 且非 Accepted 的報價。'

# Helpers
text_cols_by_sheet={
    'Orders':['A','B','C','E','F'],
    'Invoices':['A','B','C','D','G'],
    'Expenses':['A','B','C','E','F'],
    'Quotes':['A','B','C','F','G'],
    'Summary':['A'],
    'Exceptions':['A','B','C','D'],
}
def clear_data(sheet):
    sh=wb[sheet]
    if sh.max_row>=5:
        sh.delete_rows(5, sh.max_row-4)
def write_rows(sheet, rows):
    sh=wb[sheet]
    clear_data(sheet)
    for i,row in enumerate(rows, start=5):
        for j,val in enumerate(row, start=1):
            cell=sh.cell(i,j,val)
            col=cell.column_letter
            if col in text_cols_by_sheet.get(sheet,[]):
                cell.number_format='@'
                if val is not None:
                    cell.value=str(val)
            elif isinstance(val, Decimal):
                cell.value=float(val)
                cell.number_format='0.00'
            elif col in ['D','E','F'] and sheet in ['Orders','Invoices','Expenses','Quotes']:
                cell.number_format='yyyy-mm-dd'
            if sheet in ['Orders','Invoices','Expenses','Quotes'] and col in ['G','H','I','J','K']:
                if isinstance(cell.value, (int,float)):
                    cell.number_format='0.00'
    return sh.max_row

orders_rows=[]
for r in orders:
    orders_rows.append([r['OrderID'], r['ProjectID'], r['CustomerID'], r['OrderDate'], r['Status'], r['Currency'], money(r['NetAmount']), twd(r['NetAmount'], r['Currency'])])

invoice_rows=[]
for r in invoices:
    invoice_rows.append([r['InvoiceID'], r['OrderID'], r['ProjectID'], r['CustomerID'], r['InvoiceDate'], r['DueDate'], r['Currency'], money(r['Amount']), money(r['PaidAmount']), money(r['Balance']), twd(r['Balance'], r['Currency'])])

expense_rows=[]
for r in expenses:
    expense_rows.append([r['ExpenseID'], r['ProjectID'], r['EmployeeID'], r['ExpenseDate'], r['Status'], r['Currency'], money(r['Amount']), twd(r['Amount'], r['Currency'])])

quote_rows=[]
for r in quotes:
    quote_rows.append([r['QuoteID'], r['ProjectID'], r['CustomerID'], r['QuoteDate'], r['ValidUntil'], r['Status'], r['Currency'], money(r['NetAmount']), twd(r['NetAmount'], r['Currency'])])

write_rows('Orders', orders_rows)
write_rows('Invoices', invoice_rows)
write_rows('Expenses', expense_rows)
write_rows('Quotes', quote_rows)

# Summary per project across present project sources.
summary_projects=sorted(set([r[1] for r in orders_rows]+[r[2] for r in invoice_rows]+[r[1] for r in expense_rows]+[r[1] for r in quote_rows if r[5]=='Accepted']))
summary_rows=[]
for p in summary_projects:
    order_twd=sum((r[7] for r in orders_rows if r[1]==p), Decimal('0.00'))
    inv_twd=sum((r[10] for r in invoice_rows if r[2]==p), Decimal('0.00'))
    exp_twd=sum((r[7] for r in expense_rows if r[1]==p), Decimal('0.00'))
    acc_quote_twd=sum((r[8] for r in quote_rows if r[1]==p and r[5]=='Accepted'), Decimal('0.00'))
    summary_rows.append([p, order_twd, inv_twd, exp_twd, acc_quote_twd])
write_rows('Summary', summary_rows)
sum_sh=wb['Summary']
last=max(5, 4+len(summary_rows))
for col in ['B','C','D','E']:
    sum_sh[f'{col}2']=f'=SUM({col}5:{col}{last})'
    for row in range(5,last+1):
        sum_sh[f'{col}{row}'].number_format='0.00'

# Exceptions
cut=parse_date(cutoff)
exceptions=[]
for r in invoices:
    if parse_date(r['DueDate']) < cut and Decimal(r['Balance']) > 0:
        exceptions.append(['Invoices', r['InvoiceID'], r['ProjectID'], f"Overdue unpaid invoice: DueDate {r['DueDate']} < {cutoff}, Balance {r['Balance']} {r['Currency']}"])
for r in expenses:
    if str(r.get('MissingReceipt','')).strip().lower()=='true':
        exceptions.append(['Expenses', r['ExpenseID'], r['ProjectID'], 'Missing receipt'])
for r in quotes:
    if parse_date(r['ValidUntil']) < cut and r['Status']!='Accepted':
        exceptions.append(['Quotes', r['QuoteID'], r['ProjectID'], f"Expired/non-accepted quote: ValidUntil {r['ValidUntil']} < {cutoff}, Status {r['Status']}"])
write_rows('Exceptions', exceptions)

# Highlight exception document rows on source sheets
fill=PatternFill(fill_type='solid', fgColor='FFF2CC')
exception_ids={e[1] for e in exceptions}
for sheet in ['Invoices','Expenses','Quotes']:
    sh=wb[sheet]
    for row in range(5, sh.max_row+1):
        if sh.cell(row,1).value in exception_ids:
            for col in range(1, sh.max_column+1):
                sh.cell(row,col).fill=fill

wb.save(outfile)
print('saved', outfile)

# Verification: reopen output and confirm requirements.
wb2=load_workbook(outfile, data_only=False)
checks=[]
def values(sheet):
    sh=wb2[sheet]
    return [[c.value for c in row] for row in sh.iter_rows(min_row=5, max_row=sh.max_row)] if sh.max_row>=5 else []
checks.append(('orders_rows', len(values('Orders'))))
checks.append(('orders_no_cancelled', all(r[4] != 'Cancelled' for r in values('Orders'))))
checks.append(('orders_customer', all(r[2] == customer_id for r in values('Orders'))))
checks.append(('invoices_rows', len(values('Invoices'))))
checks.append(('expenses_rows', len(values('Expenses'))))
checks.append(('expenses_approved_projects', all(r[4]=='Approved' and r[1] in projects for r in values('Expenses'))))
checks.append(('quotes_rows', len(values('Quotes'))))
checks.append(('exceptions_rows', len(values('Exceptions'))))
checks.append(('summary_rows', len(values('Summary'))))
checks.append(('summary_formula_B2', wb2['Summary']['B2'].value))
# ID cells text format
checks.append(('id_text_formats', all(wb2[s].cell(5,1).number_format=='@' for s in ['Orders','Invoices','Expenses','Quotes'] if wb2[s].max_row>=5)))
checks.append(('numeric_amounts', all(isinstance(r[-1], (int,float)) for s in ['Orders','Invoices','Expenses','Quotes'] for r in values(s))))
for k,v in checks:
    print(k, '=', v)
print('row detail counts:', {'Orders':len(orders_rows),'Invoices':len(invoice_rows),'Expenses':len(expense_rows),'Quotes':len(quote_rows),'Summary':len(summary_rows),'Exceptions':len(exceptions)})
print('exceptions=', exceptions)
