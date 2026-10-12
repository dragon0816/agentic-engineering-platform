from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
ins=wb['Instructions']
print('Instructions full:')
for row in ins.iter_rows(values_only=True):
    print(row)
print('\nReport merged ranges:', list(wb['Report'].merged_cells.ranges))
print('Summary existing rows:')
for r in wb['Summary'].iter_rows(min_row=1, max_row=20, values_only=True):
    print(r)