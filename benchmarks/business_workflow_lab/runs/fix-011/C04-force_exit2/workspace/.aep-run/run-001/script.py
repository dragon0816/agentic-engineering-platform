from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size, 'bytes')

manual = ws/'input'/'quote_cli_manual.md'
print('\nMANUAL START:')
print(manual.read_text(encoding='utf-8')[:4000])
print('MANUAL END SAMPLE')

fx = ws/'input'/'fx_rates.csv'
print('\nFX CSV SAMPLE:')
with fx.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    for i, row in zip(range(8), reader):
        print(i, row)

tpl = ws/'input'/'quote_template.xlsx'
wb = load_workbook(tpl, data_only=False)
print('\nWORKBOOK SHEETS:', wb.sheetnames)
for sh in wb.sheetnames:
    wsht = wb[sh]
    print('SHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 12), values_only=True):
        print(r)
    print('---')