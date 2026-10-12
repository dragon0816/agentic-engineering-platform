from pathlib import Path
from openpyxl import load_workbook
p=Path('input/sales_template.xlsx')
wb=load_workbook(p)
for sh in wb.worksheets:
    print('---', sh.title)
    for row in sh.iter_rows(values_only=True):
        print(row)
