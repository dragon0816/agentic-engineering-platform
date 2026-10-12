from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta, date
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
downloads = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)

start_date = '2026-09-01'
end_date = '2026-09-30'
status = 'Open'
csv_path = downloads / 'sales_orders_2026-09-01_2026-09-30_Open.csv'
template_path = input_dir / 'sales_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
out_path = output_dir / 'september_open_orders.xlsx'

details = {
  '0045000109': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10'},
  '0045000110': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET30'},
  '0045000112': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15'},
  '0045000113': {'CustomerContact': '', 'PaymentTerms': 'NET15'},
  '0045000115': {'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET20'},
  '0045000116': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
  '0045000118': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET0'},
  '0045000119': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
  '0045000120': {'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET10'},
  '0045000122': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET5'},
  '0045000124': {'CustomerContact': '', 'PaymentTerms': 'NET30'},
  '0045000126': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
  '0045000127': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET30'},
  '0045000128': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET5'},
}

print('inputs:', template_path.name, template_path.exists(), fx_path.name, fx_path.exists(), csv_path, csv_path.exists())

fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    print('fx header:', next(reader))
    for row in reader:
        if row and row[0]:
            fx[row[0]] = Decimal(row[1])
print('fx rates:', fx)

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
print('csv headers:', list(rows[0].keys()) if rows else [])
print('csv row count before filter:', len(rows))
rows = [r for r in rows if start_date <= r['OrderDate'] <= end_date and r['Status'] == status]
rows.sort(key=lambda r: r['OrderID'])
print('filtered rows:', len(rows))
print('first/last:', rows[0]['OrderID'] if rows else None, rows[-1]['OrderID'] if rows else None)

for r in rows:
    r.update(details.get(r['OrderID'], {'CustomerContact':'', 'PaymentTerms':''}))
    amt = Decimal(r['NetAmount'])
    r['NetAmountTWD'] = (amt * fx[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

wb = load_workbook(template_path)
report = wb['Report']; dws = wb['Details']; sws = wb['Summary']
if dws.max_row >= 5:
    dws.delete_rows(5, dws.max_row - 4)
if sws.max_row >= 5:
    sws.delete_rows(5, sws.max_row - 4)

now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = f'Sales Orders; Status={status}; Sort by OrderID asc; all filtered rows exported and detail pages read'
report['B5'] = now
report['B7'] = 'Lab ERP synthetic SAP-like web; fx_rates.csv'
report['D7'] = now
report['B9'] = 'CustomerContact is blank where the source detail page had a blank value.'

for ridx, r in enumerate(rows, start=5):
    vals = [r['OrderID'], r['CustomerID'], r['CustomerName'], r.get('ProjectID') or None,
            date.fromisoformat(r['OrderDate']), r['Status'], r['Currency'],
            float(Decimal(r['NetAmount'])), float(r['NetAmountTWD']),
            r.get('CustomerContact') or None, r.get('PaymentTerms') or None]
    for c, v in enumerate(vals, start=1):
        cell = dws.cell(ridx, c, v)
        if c in (1,2,4): cell.number_format = '@'
        if c == 5: cell.number_format = 'yyyy-mm-dd'
        if c in (8,9): cell.number_format = '0.00'

last_row = 4 + len(rows)
sws['B2'] = f'=SUM(Details!I5:I{last_row if rows else 5})'
by_cur = {}
for r in rows:
    cur = r['Currency']
    by_cur.setdefault(cur, {'amount':Decimal('0'), 'twd':Decimal('0'), 'count':0})
    by_cur[cur]['amount'] += Decimal(r['NetAmount'])
    by_cur[cur]['twd'] += r['NetAmountTWD']
    by_cur[cur]['count'] += 1
srow = 5
for cur in sorted(by_cur):
    data = by_cur[cur]
    sws.cell(srow,1,'Currency Total')
    sws.cell(srow,2,None)
    sws.cell(srow,3,cur)
    sws.cell(srow,4,float(data['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)))
    sws.cell(srow,5,float(data['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)))
    sws.cell(srow,6,data['count'])
    sws.cell(srow,4).number_format = '0.00'; sws.cell(srow,5).number_format = '0.00'
    srow += 1
grand_twd = sum((d['twd'] for d in by_cur.values()), Decimal('0')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
sws.cell(srow,1,'Grand Total')
sws.cell(srow,2,'All currencies')
sws.cell(srow,3,'TWD')
sws.cell(srow,4,None)
sws.cell(srow,5,float(grand_twd))
sws.cell(srow,6,len(rows))
sws.cell(srow,5).number_format = '0.00'

wb.save(out_path)
print('saved:', out_path)

# Reopen and verify requirements.
wb2 = load_workbook(out_path, data_only=False)
cd = wb2['Details']; cs = wb2['Summary']; cr = wb2['Report']
row_count = cd.max_row - 4
order_ids = [str(cd.cell(r,1).value) for r in range(5, cd.max_row+1)]
statuses = [cd.cell(r,6).value for r in range(5, cd.max_row+1)]
def cell_date_iso(v):
    if hasattr(v, 'date'):
        return v.date().isoformat()
    if hasattr(v, 'isoformat'):
        return v.isoformat()
    return str(v)
date_isos = [cell_date_iso(cd.cell(r,5).value) for r in range(5, cd.max_row+1)]
net_amount_types = [type(cd.cell(r,8).value).__name__ for r in range(5, cd.max_row+1)]
twd_amount_types = [type(cd.cell(r,9).value).__name__ for r in range(5, cd.max_row+1)]
summary_rows = [[cs.cell(r,c).value for c in range(1,7)] for r in range(5, cs.max_row+1)]
contacts = [cd.cell(r,10).value for r in range(5, cd.max_row+1)]

print('CHECK output exists:', out_path.exists())
print('CHECK file name:', out_path.name)
print('CHECK Details headers:', [cd.cell(4,c).value for c in range(1,12)])
print('CHECK row count:', row_count)
print('CHECK OrderID sorted:', order_ids == sorted(order_ids), order_ids)
print('CHECK all status Open:', all(s == status for s in statuses))
print('CHECK dates in range:', all(start_date <= d <= end_date for d in date_isos), date_isos)
print('CHECK ID formats text:', [cd.cell(5,c).number_format for c in (1,2,4)])
print('CHECK amount formats:', [cd.cell(5,c).number_format for c in (8,9)])
print('CHECK amount value types:', sorted(set(net_amount_types)), sorted(set(twd_amount_types)))
print('CHECK blank contacts preserved count:', sum(1 for x in contacts if x is None))
print('CHECK Summary B2 formula:', cs['B2'].value)
print('CHECK Summary rows:', summary_rows)
print('CHECK grand total TWD:', grand_twd)
print('CHECK Report fields:', cr['B3'].value, cr['D3'].value, cr['B4'].value)

assert out_path.exists()
assert out_path.name == 'september_open_orders.xlsx'
assert [cd.cell(4,c).value for c in range(1,12)] == ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert row_count == 14
assert order_ids == sorted(order_ids)
assert all(s == status for s in statuses)
assert all(start_date <= d <= end_date for d in date_isos)
assert all(isinstance(cd.cell(r,8).value, (int,float)) for r in range(5, cd.max_row+1))
assert all(isinstance(cd.cell(r,9).value, (int,float)) for r in range(5, cd.max_row+1))
assert cs['B2'].value == '=SUM(Details!I5:I18)'
assert summary_rows[-1] == ['Grand Total', 'All currencies', 'TWD', None, float(grand_twd), 14]
