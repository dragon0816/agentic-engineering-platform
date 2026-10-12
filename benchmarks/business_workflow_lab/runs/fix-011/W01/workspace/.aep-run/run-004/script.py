from pathlib import Path
import csv, shutil, json
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta, date
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
downloads = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)

start_date = '2026-09-01'
end_date = '2026-09-30'
status = 'Open'
source_url = 'http://127.0.0.1:8810/'

csv_path = downloads / 'sales_orders_2026-09-01_2026-09-30_Open.csv'
template_path = input_dir / 'sales_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
out_path = output_dir / 'september_open_orders.xlsx'

# Browser-read detail records. Empty strings are retained as true source blanks.
details = {
  '0045000109': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10'},
  '0045000110': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET30'},
  '0045000112': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15'},
  '0045000113': {'CustomerContact': '', 'PaymentTerms': 'NET15'},
  '0045000115': {'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET20'},
  '0045000116': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
  '0045000118': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET0'},
  '0045000119': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
  '0045000120': {'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET10'},
  '0045000122': {'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET5'},
  '0045000124': {'CustomerContact': '', 'PaymentTerms': 'NET30'},
  '0045000126': {'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30'},
  '0045000127': {'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET30'},
  '0045000128': {'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET5'},
}

print('inputs:', template_path, template_path.exists(), fx_path, fx_path.exists(), csv_path, csv_path.exists())
print('output:', out_path)

# Parse FX rates; csv rows have Note with commas, but first two columns are stable.
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if not row or not row[0]:
            continue
        fx[row[0]] = Decimal(row[1])
print('fx:', fx)

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
print('csv headers:', list(rows[0].keys()) if rows else [])
print('csv rows:', len(rows))
print('csv sample:', rows[:3])

# Keep only the requested criteria and sort by OrderID text.
rows = [r for r in rows if start_date <= r['OrderDate'] <= end_date and r['Status'] == status]
rows.sort(key=lambda r: r['OrderID'])

# Enrich and compute amounts.
for r in rows:
    oid = r['OrderID']
    if oid in details:
        r.update(details[oid])
    else:
        r['CustomerContact'] = ''
        r['PaymentTerms'] = ''
    amt = Decimal(r['NetAmount'])
    rate = fx[r['Currency']]
    r['NetAmountTWD'] = (amt * rate).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

wb = load_workbook(template_path)
report = wb['Report']
detail_ws = wb['Details']
summary = wb['Summary']

# Clear any existing data rows in Details and Summary while preserving headers/template structure.
if detail_ws.max_row >= 5:
    detail_ws.delete_rows(5, detail_ws.max_row - 4)
if summary.max_row >= 5:
    summary.delete_rows(5, summary.max_row - 4)

now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
fetched_at = now
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = f'Sales Orders; Status={status}; Sort by OrderID asc; all filtered rows exported and detail pages read'
report['B5'] = now
report['B7'] = 'Lab ERP synthetic SAP-like web; fx_rates.csv'
report['D7'] = fetched_at
report['B9'] = 'CustomerContact is blank where the source detail page had a blank value.'

headers = [detail_ws.cell(4, c).value for c in range(1, 12)]
print('detail template headers:', headers)

for idx, r in enumerate(rows, start=5):
    vals = [
        r['OrderID'], r['CustomerID'], r['CustomerName'], r.get('ProjectID') or None,
        date.fromisoformat(r['OrderDate']), r['Status'], r['Currency'],
        float(Decimal(r['NetAmount'])), float(r['NetAmountTWD']),
        r.get('CustomerContact') or None, r.get('PaymentTerms') or None,
    ]
    for c, v in enumerate(vals, start=1):
        cell = detail_ws.cell(idx, c, v)
        if c in (1, 2, 4):
            cell.number_format = '@'
        elif c == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif c in (8, 9):
            cell.number_format = '0.00'

last_row = 4 + len(rows)
summary['B2'] = f'=SUM(Details!I5:I{last_row if rows else 5})'

# Currency summaries
by_cur = {}
for r in rows:
    cur = r['Currency']
    by_cur.setdefault(cur, {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
    by_cur[cur]['amount'] += Decimal(r['NetAmount'])
    by_cur[cur]['twd'] += r['NetAmountTWD']
    by_cur[cur]['count'] += 1

srow = 5
for cur in sorted(by_cur):
    data = by_cur[cur]
    summary.cell(srow, 1, 'Currency Total')
    summary.cell(srow, 2, None)
    summary.cell(srow, 3, cur)
    summary.cell(srow, 4, float(data['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)))
    summary.cell(srow, 5, float(data['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)))
    summary.cell(srow, 6, data['count'])
    summary.cell(srow, 4).number_format = '0.00'
    summary.cell(srow, 5).number_format = '0.00'
    srow += 1

grand_twd = sum((data['twd'] for data in by_cur.values()), Decimal('0')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
summary.cell(srow, 1, 'Grand Total')
summary.cell(srow, 2, 'All currencies')
summary.cell(srow, 3, 'TWD')
summary.cell(srow, 4, None)
summary.cell(srow, 5, float(grand_twd))
summary.cell(srow, 6, len(rows))
summary.cell(srow, 5).number_format = '0.00'

# Set numeric formats in Summary.
for r in range(5, srow + 1):
    summary.cell(r, 4).number_format = '0.00'
    summary.cell(r, 5).number_format = '0.00'

wb.save(out_path)
print('saved', out_path)

# Verification by reopening output.
check = load_workbook(out_path, data_only=False)
cd = check['Details']
cs = check['Summary']
cr = check['Report']
actual_rows = cd.max_row - 4
order_ids = [str(cd.cell(r, 1).value) for r in range(5, cd.max_row + 1)]
statuses = [cd.cell(r, 6).value for r in range(5, cd.max_row + 1)]
dates = [cd.cell(r, 5).value for r in range(5, cd.max_row + 1)]
summary_rows = []
for r in range(5, cs.max_row + 1):
    summary_rows.append([cs.cell(r,c).value for c in range(1,7)])
print('CHECK row_count', actual_rows)
print('CHECK sorted', order_ids == sorted(order_ids), order_ids)
print('CHECK all_status_open', all(s == status for s in statuses))
print('CHECK dates_in_range', all(start_date <= (d.isoformat() if hasattr(d, 'isoformat') else str(d)) <= end_date for d in dates))
print('CHECK B2 formula', cs['B2'].value)
print('CHECK summary_rows', summary_rows)
print('CHECK grand_twd', grand_twd)
print('CHECK report', cr['B3'].value, cr['D3'].value, cr['B4'].value)
print('CHECK id_number_format', [cd.cell(5,c).number_format for c in (1,2,4)])
print('CHECK amount_number_format', [cd.cell(5,c).number_format for c in (8,9)])
assert actual_rows == 14
assert order_ids == sorted(order_ids)
assert all(s == status for s in statuses)
assert all(start_date <= (d.isoformat() if hasattr(d, 'isoformat') else str(d)) <= end_date for d in dates)
assert cs['B2'].value == '=SUM(Details!I5:I18)'
assert out_path.exists()
