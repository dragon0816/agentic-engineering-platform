from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
inp = ws / 'input'
out = ws / 'output'
print('workspace=', ws)
print('input files=', [p.name for p in inp.iterdir()] if inp.exists() else 'NO INPUT')

manual = inp / 'sap_manual.md'
if manual.exists():
    print('\n--- sap_manual.md ---')
    print(manual.read_text(encoding='utf-8', errors='replace')[:5000])
else:
    print('no sap_manual.md')

fx = inp / 'fx_rates.csv'
if fx.exists():
    print('\n--- fx_rates.csv ---')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        rows = list(csv.reader(f))
    print('rows=', len(rows))
    for r in rows[:20]: print(r)
else:
    print('no fx_rates.csv')

tpl = inp / 'sales_template.xlsx'
if tpl.exists():
    print('\n--- sales_template.xlsx ---')
    wb = load_workbook(tpl)
    print('sheets=', wb.sheetnames)
    for sh in wb.worksheets:
        print('sheet', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(row)
        print('---')
else:
    print('no sales_template.xlsx')
