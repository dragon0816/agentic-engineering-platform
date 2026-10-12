from pathlib import Path
import csv, json, os, re
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta, date
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
downloads = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
print('workflow parameter type:', type(params).__name__)
print('workflow parameter keys:', list(params.keys()) if isinstance(params, dict) else 'n/a')

def extract_detail_map(obj):
    out = {}
    if not isinstance(obj, dict):
        return out
    candidates = obj.get('items') if isinstance(obj.get('items'), dict) else obj
    for oid, payload in candidates.items():
        detail = None
        if isinstance(payload, dict):
            detail = payload.get('detail')
            if detail is None and isinstance(payload.get('data'), dict):
                detail = payload['data'].get('detail')
        if isinstance(detail, dict):
            out[str(oid)] = detail
    return out

detail_map = extract_detail_map(params)
print('detail records returned:', len(detail_map))

csv_files = sorted(downloads.glob('sales_orders*.csv'), key=lambda p: p.stat().st_mtime, reverse=True)
if not csv_files:
    raise FileNotFoundError('No downloaded sales_orders*.csv found under output/downloads')
csv_path = csv_files[0]
print('using csv:', csv_path)

fx = {}
with (input_dir / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('fx header:', header)
    for row in reader:
        if row and row[0]:
            fx[row[0]] = Decimal(row[1])
print('fx rates:', fx)

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rows = list(csv.DictReader(f))
headers = list(rows[0].keys()) if rows else []
print('csv headers:', headers)
print('csv rows:', len(rows))
rows.sort(key=lambda r: r['OrderID'])

# Derive the run filter from the exported data and/or filename, so the workflow can rerun with changed values
# even though the browser download step uses a fixed safe filename.
file_m = re.match(r'sales_orders_(\d{4}-\d{2}-\d{2})_(\d{4}-\d{2}-\d{2})_(.+)\.csv$', csv_path.name)
if file_m:
    start_date, end_date, status = file_m.group(1), file_m.group(2), file_m.group(3)
elif rows:
    order_dates = sorted(str(r['OrderDate']) for r in rows)
    statuses = sorted({str(r['Status']) for r in rows})
    start_date, end_date = order_dates[0], order_dates[-1]
    status = statuses[0] if len(statuses) == 1 else 'Mixed'
else:
    start_date = end_date = status = ''
print('derived filter:', start_date, end_date, status)

for r in rows:
    detail = detail_map.get(r['OrderID'], {})
    r['CustomerContact'] = detail.get('CustomerContact', '') if isinstance(detail, dict) else ''
    r['PaymentTerms'] = detail.get('PaymentTerms', '') if isinstance(detail, dict) else ''
    cur = r['Currency']
    if cur not in fx:
        raise KeyError(f'Missing FX rate for {cur}')
    amt = Decimal(r['NetAmount'])
    r['NetAmountTWD'] = (amt * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

if start_date == '2026-09-01' and end_date == '2026-09-30' and status == 'Open':
    out_path = output_dir / 'september_open_orders.xlsx'
else:
    safe_status = re.sub(r'[^A-Za-z0-9_-]+', '_', status or 'all')
    out_path = output_dir / f'sales_orders_{start_date}_{end_date}_{safe_status}.xlsx'

template_path = input_dir / 'sales_template.xlsx'
wb = load_workbook(template_path)
report = wb['Report']; dws = wb['Details']; sws = wb['Summary']
if dws.max_row >= 5:
    dws.delete_rows(5, dws.max_row - 4)
if sws.max_row >= 5:
    sws.delete_rows(5, sws.max_row - 4)

now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = f'Sales Orders; Status={status}; Sort by OrderID asc; all filtered rows exported and detail pages read'
report['B5'] = now
report['B7'] = 'Lab ERP synthetic SAP-like web; fx_rates.csv'
report['D7'] = now
missing_contacts = sum(1 for r in rows if not r.get('CustomerContact'))
report['B9'] = f'CustomerContact blank count: {missing_contacts}; blanks reflect source detail pages or missing detail records.'

for ridx, r in enumerate(rows, start=5):
    vals = [r['OrderID'], r['CustomerID'], r['CustomerName'], r.get('ProjectID') or None,
            date.fromisoformat(r['OrderDate']), r['Status'], r['Currency'],
            float(Decimal(r['NetAmount'])), float(r['NetAmountTWD']),
            r.get('CustomerContact') or None, r.get('PaymentTerms') or None]
    for c, v in enumerate(vals, start=1):
        cell = dws.cell(ridx, c, v)
        if c in (1,2,4): cell.number_format = '@'
        elif c == 5: cell.number_format = 'yyyy-mm-dd'
        elif c in (8,9): cell.number_format = '0.00'

last_row = 4 + len(rows)
sws['B2'] = f'=SUM(Details!I5:I{last_row if rows else 5})'
by_cur = {}
for r in rows:
    cur = r['Currency']
    by_cur.setdefault(cur, {'amount':Decimal('0'), 'twd':Decimal('0'), 'count':0})
    by_cur[cur]['amount'] += Decimal(r['NetAmount'])
    by_cur[cur]['twd'] += r['NetAmountTWD']
    by_cur[cur]['count'] += 1
srow = 5
for cur in sorted(by_cur):
    data = by_cur[cur]
    sws.cell(srow,1,'Currency Total')
    sws.cell(srow,2,None)
    sws.cell(srow,3,cur)
    sws.cell(srow,4,float(data['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)))
    sws.cell(srow,5,float(data['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)))
    sws.cell(srow,6,data['count'])
    sws.cell(srow,4).number_format = '0.00'; sws.cell(srow,5).number_format = '0.00'
    srow += 1
grand_twd = sum((d['twd'] for d in by_cur.values()), Decimal('0')).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
sws.cell(srow,1,'Grand Total')
sws.cell(srow,2,'All currencies')
sws.cell(srow,3,'TWD')
sws.cell(srow,4,None)
sws.cell(srow,5,float(grand_twd))
sws.cell(srow,6,len(rows))
sws.cell(srow,5).number_format = '0.00'

wb.save(out_path)
print('saved:', out_path)

# Verification from saved workbook.
wb2 = load_workbook(out_path, data_only=False)
cd = wb2['Details']; cs = wb2['Summary']
row_count = max(0, cd.max_row - 4)
order_ids = [str(cd.cell(r,1).value) for r in range(5, cd.max_row+1)] if row_count else []
statuses = [cd.cell(r,6).value for r in range(5, cd.max_row+1)] if row_count else []
def cell_date_iso(v):
    if hasattr(v, 'date'):
        return v.date().isoformat()
    if hasattr(v, 'isoformat'):
        return v.isoformat()
    return str(v)
date_isos = [cell_date_iso(cd.cell(r,5).value) for r in range(5, cd.max_row+1)] if row_count else []
summary_rows = [[cs.cell(r,c).value for c in range(1,7)] for r in range(5, cs.max_row+1)]
print('CHECK output exists:', out_path.exists())
print('CHECK output name:', out_path.name)
print('CHECK Details headers:', [cd.cell(4,c).value for c in range(1,12)])
print('CHECK row count:', row_count)
print('CHECK OrderID sorted:', order_ids == sorted(order_ids))
print('CHECK all statuses in workbook:', sorted(set(statuses)))
print('CHECK dates min/max:', (min(date_isos), max(date_isos)) if date_isos else None)
print('CHECK Summary B2 formula:', cs['B2'].value)
print('CHECK Summary rows:', summary_rows)
print('CHECK grand TWD:', grand_twd)
assert out_path.exists()
assert [cd.cell(4,c).value for c in range(1,12)] == ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert order_ids == sorted(order_ids)
assert row_count == len(rows)
assert all(isinstance(cd.cell(r,8).value, (int,float)) for r in range(5, cd.max_row+1))
assert all(isinstance(cd.cell(r,9).value, (int,float)) for r in range(5, cd.max_row+1))
if status and status != 'Mixed':
    assert all(s == status for s in statuses)
if start_date and end_date:
    assert all(start_date <= d <= end_date for d in date_isos)
