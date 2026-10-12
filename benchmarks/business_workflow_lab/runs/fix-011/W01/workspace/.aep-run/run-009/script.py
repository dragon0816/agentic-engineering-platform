import csv, json, os
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
import openpyxl

ws = Path.cwd()
print('WORKSPACE', ws)
input_dir = ws / 'input'
output_dir = ws / 'output'

csv_path = output_dir / 'downloads' / 'sales_orders_export.csv'
xlsx_path = output_dir / 'september_open_orders.xlsx'
fx_path = input_dir / 'fx_rates.csv'
template_path = input_dir / 'sales_template.xlsx'

print('FILES')
for p in [template_path, fx_path, csv_path, xlsx_path]:
    print(f' - {p.relative_to(ws)} exists={p.exists()} size={p.stat().st_size if p.exists() else None}')
for p in [template_path, fx_path, csv_path, xlsx_path]:
    if not p.exists():
        raise SystemExit(f'Missing required file: {p}')

# Read source CSV from workflow browser export.
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    csv_headers = reader.fieldnames
    csv_rows = list(reader)
print('CSV_HEADERS', csv_headers)
print('CSV_ROW_COUNT', len(csv_rows))
print('CSV_FIRST_LAST', csv_rows[0] if csv_rows else None, csv_rows[-1] if csv_rows else None)

# Read FX rates using actual attached schema.
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    fx_headers = reader.fieldnames
    for row in reader:
        cur = (row.get('Currency') or '').strip()
        rate_raw = (row.get('RateToTWD') or '').strip()
        if cur and rate_raw:
            fx[cur] = Decimal(rate_raw)
print('FX_HEADERS', fx_headers)
print('FX_RATES', {k: str(v) for k, v in sorted(fx.items())})

# Read template/output workbook structures.
twb = openpyxl.load_workbook(template_path, data_only=False)
wb_f = openpyxl.load_workbook(xlsx_path, data_only=False)
wb_v = openpyxl.load_workbook(xlsx_path, data_only=True)
print('TEMPLATE_SHEETS', twb.sheetnames)
print('OUTPUT_SHEETS', wb_f.sheetnames)
assert wb_f.sheetnames == twb.sheetnames, 'Output workbook did not preserve template sheet list/order'

# Validate report sheet metadata.
report = wb_v['Report']
print('REPORT_B3_D3_B4', report['B3'].value, report['D3'].value, report['B4'].value)
assert str(report['B3'].value)[:10] == '2026-09-01', 'Report start date mismatch'
assert str(report['D3'].value)[:10] == '2026-09-30', 'Report end date mismatch'
assert 'Open' in str(report['B4'].value), 'Report condition does not mention Open'

# Read Details.
details_f = wb_f['Details']
details_v = wb_v['Details']
headers = [details_f.cell(4, c).value for c in range(1, details_f.max_column + 1)]
print('DETAIL_HEADERS', headers)
required = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
hmap = {h: i+1 for i, h in enumerate(headers) if h}
missing = [h for h in required if h not in hmap]
assert not missing, f'Missing detail columns: {missing}'

rows = []
for r in range(5, details_v.max_row + 1):
    oid = details_v.cell(r, hmap['OrderID']).value
    if oid is None or str(oid).strip() == '':
        continue
    rows.append({h: details_v.cell(r, hmap[h]).value for h in required})
print('DETAIL_ROW_COUNT', len(rows))
print('DETAIL_FIRST', rows[0])
print('DETAIL_LAST', rows[-1])

# Validate all exported rows and workbook rows.
csv_ids = [row['OrderID'] for row in csv_rows]
ids = [str(row['OrderID']) for row in rows]
assert len(csv_rows) == 14, f'Expected 14 CSV rows, got {len(csv_rows)}'
assert len(rows) == 14, f'Expected 14 workbook detail rows, got {len(rows)}'
assert csv_ids == sorted(csv_ids), 'CSV rows are not sorted by OrderID ascending'
assert ids == sorted(ids), 'Workbook rows are not sorted by OrderID ascending'
assert ids == csv_ids, 'Workbook detail rows do not match exported all-filtered CSV rows'
assert all(row['Status'] == 'Open' for row in csv_rows), 'CSV contains non-Open rows'
assert all(row['Status'] == 'Open' for row in rows), 'Workbook contains non-Open rows'
assert all('2026-09-01' <= row['OrderDate'] <= '2026-09-30' for row in csv_rows), 'CSV contains rows outside date range'
for row in rows:
    d = row['OrderDate']
    dstr = d.strftime('%Y-%m-%d') if hasattr(d, 'strftime') else str(d)[:10]
    assert '2026-09-01' <= dstr <= '2026-09-30', f'Workbook date outside range: {row}'
    assert isinstance(row['OrderID'], str) and row['OrderID'] == row['OrderID'].zfill(10), f'OrderID not preserved as text: {row["OrderID"]}'
    assert isinstance(row['NetAmount'], (int, float)), f'NetAmount is not numeric: {row}'
    assert isinstance(row['NetAmountTWD'], (int, float)), f'NetAmountTWD is not numeric: {row}'

# Check conversion and aggregate totals.
by_cur_amount = defaultdict(Decimal)
by_cur_twd = defaultdict(Decimal)
by_cur_count = defaultdict(int)
for row in rows:
    cur = row['Currency']
    assert cur in fx, f'Missing FX rate for {cur}'
    amt = Decimal(str(row['NetAmount'])).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    actual_twd = Decimal(str(row['NetAmountTWD'])).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    expected_twd = (amt * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert actual_twd == expected_twd, f'FX mismatch for {row["OrderID"]}: {actual_twd} != {expected_twd}'
    by_cur_amount[cur] += amt
    by_cur_twd[cur] += actual_twd
    by_cur_count[cur] += 1
expected_summary = {cur: {'amount': by_cur_amount[cur].quantize(Decimal('0.01')), 'twd': by_cur_twd[cur].quantize(Decimal('0.01')), 'count': by_cur_count[cur]} for cur in sorted(by_cur_count)}
expected_total = sum(by_cur_twd.values()).quantize(Decimal('0.01'))
print('EXPECTED_SUMMARY', {k: {'amount': str(v['amount']), 'twd': str(v['twd']), 'count': v['count']} for k, v in expected_summary.items()})
print('EXPECTED_TOTAL_TWD', str(expected_total))

# Validate Summary sheet.
summary_f = wb_f['Summary']
summary_v = wb_v['Summary']
print('SUMMARY_B2_FORMULA', summary_f['B2'].value)
assert summary_f['B2'].value == '=SUM(Details!I5:I18)', f'Unexpected Summary B2 formula: {summary_f["B2"].value}'
summary_by_cur = {}
grand = None
for r in range(5, summary_v.max_row + 1):
    group = summary_v.cell(r, 1).value
    cur = summary_v.cell(r, 3).value
    amount = summary_v.cell(r, 4).value
    twd = summary_v.cell(r, 5).value
    count = summary_v.cell(r, 6).value
    if group == 'Currency Total' and cur:
        summary_by_cur[cur] = {'amount': Decimal(str(amount)).quantize(Decimal('0.01')), 'twd': Decimal(str(twd)).quantize(Decimal('0.01')), 'count': int(count)}
    if group == 'Grand Total':
        grand = {'currency': cur, 'twd': Decimal(str(twd)).quantize(Decimal('0.01')), 'count': int(count)}
print('SUMMARY_BY_CUR', {k: {'amount': str(v['amount']), 'twd': str(v['twd']), 'count': v['count']} for k, v in summary_by_cur.items()})
print('GRAND', {k: str(v) if isinstance(v, Decimal) else v for k, v in grand.items()})
assert set(summary_by_cur) == set(expected_summary), 'Summary currencies mismatch'
for cur, exp in expected_summary.items():
    got = summary_by_cur[cur]
    assert got == exp, f'Summary mismatch for {cur}: got {got}, expected {exp}'
assert grand == {'currency': 'TWD', 'twd': expected_total, 'count': len(rows)}, f'Grand total mismatch: {grand}'

# Blank placeholders / blanks preserved: compare ProjectID blank count and CustomerContact detail fields populated/blanks as available.
blank_project_csv = sum(1 for row in csv_rows if row.get('ProjectID','') == '')
blank_project_wb = sum(1 for row in rows if row.get('ProjectID') in (None, ''))
blank_contact_wb = sum(1 for row in rows if row.get('CustomerContact') in (None, ''))
print('BLANK_COUNTS', {'ProjectID_csv': blank_project_csv, 'ProjectID_workbook': blank_project_wb, 'CustomerContact_workbook': blank_contact_wb})
assert blank_project_wb == blank_project_csv, 'Blank ProjectID count was not preserved'
assert blank_contact_wb == 2, f'Expected 2 blank CustomerContact cells, got {blank_contact_wb}'

result = {
    'checked_files': [str(csv_path), str(xlsx_path)],
    'detail_rows': len(rows),
    'sorted_by_order_id': True,
    'status': 'Open',
    'date_range': '2026-09-01..2026-09-30',
    'summary': {k: {'amount': str(v['amount']), 'twd': str(v['twd']), 'count': v['count']} for k, v in expected_summary.items()},
    'grand_total_twd': str(expected_total),
    'workflow_output_workbook': str(xlsx_path),
    'workflow_download_csv': str(csv_path)
}
print('CHECKS PASSED')
print(json.dumps(result, ensure_ascii=False, indent=2))