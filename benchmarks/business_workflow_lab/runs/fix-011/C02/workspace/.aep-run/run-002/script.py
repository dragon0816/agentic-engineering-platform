from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from datetime import datetime, date
import csv, json, subprocess, sys, os, shutil
from openpyxl import load_workbook
from openpyxl.styles import numbers

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(parents=True, exist_ok=True)

START_DATE = '2026-09-01'
END_DATE = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
CONDITION = 'QuoteDate 2026-09-01..2026-09-30; all statuses; expired flag: ValidUntil < 2026-09-30 and Status != Accepted; keep all quotes.'
SOURCE = 'quote_cli synthetic quote system; fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'quote_template_filled.xlsx'
OUT_MD = OUTPUT / 'findings.md'

Q2 = Decimal('0.01')
def D(x):
    if isinstance(x, Decimal): return x
    return Decimal(str(x))

def money(x):
    return D(x).quantize(Q2, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [sys.executable, str(CLI)] + args
    cp = subprocess.run(cmd, cwd=str(WORKSPACE), capture_output=True, text=True, encoding='utf-8')
    print('CLI:', ' '.join(args), 'exit', cp.returncode)
    if cp.stderr.strip():
        print('STDERR:', cp.stderr.strip())
    if cp.returncode != 0:
        raise RuntimeError(f'CLI failed ({cp.returncode}) for {args}')
    return cp.stdout

# Print real inputs and template instruction facts for audit.
print('WORKSPACE:', WORKSPACE)
print('Input files:', sorted(p.name for p in INPUT.glob('*')))
wb0 = load_workbook(TEMPLATE, data_only=False)
print('Template sheets:', wb0.sheetnames)
for s in ['Report','Quotes','Lines','Summary','Instructions']:
    sh = wb0[s]
    print(f'Sheet {s}: rows={sh.max_row}, cols={sh.max_column}')
    if s in ['Quotes','Lines','Summary']:
        print('Header row 4:', [sh.cell(4,c).value for c in range(1, sh.max_column+1)])
print('Instructions rows:')
for r in range(1, wb0['Instructions'].max_row+1):
    vals = [wb0['Instructions'].cell(r,c).value for c in range(1,4)]
    if any(v is not None for v in vals):
        print(r, vals)

# Load FX rates, handling note text with comma by taking first two columns only.
fx_rates = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx_rates[row[0]] = D(row[1])
print('FX rates:', {k: str(v) for k, v in fx_rates.items()})

# Fetch all September 2026 quote headers, no status filter.
list_stdout = run_cli(['quotes','list','--from',START_DATE,'--to',END_DATE,'--format','json'])
listed = json.loads(list_stdout)
headers = listed.get('items', [])
print('Quote count from list:', listed.get('count'), 'items length:', len(headers))
print('Quote IDs:', [h.get('QuoteID') for h in headers])
if listed.get('count') != len(headers):
    raise AssertionError('CLI count does not match item length')

# Fetch full details for every listed quote.
quotes = []
all_lines = []
findings = []
for h in headers:
    qid = h['QuoteID']
    detail = json.loads(run_cli(['quotes','get',qid,'--format','json']))
    header = detail['header']
    lines = detail.get('line_items', [])
    # Preserve CLI get header as source of truth; validate matches list key fields.
    for key in ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount']:
        if key in h and str(h[key]) != str(header.get(key)):
            findings.append({
                'QuoteID': qid, 'Type': 'HeaderMismatch',
                'Message': f'list {key}={h[key]} differs from get {key}={header.get(key)}'
            })
    line_sum = sum((money(D(li['LineAmount'])) for li in lines), Decimal('0.00'))
    subtotal = money(header['Subtotal'])
    discount_percent = D(header['DiscountPercent'])
    expected_net = money(subtotal * (Decimal('1') - discount_percent / Decimal('100')))
    net = money(header['NetAmount'])
    recomputed_line_amount_issues = []
    for idx, li in enumerate(lines, 1):
        calc_line = money(D(li['Quantity']) * D(li['UnitPrice']))
        given_line = money(li['LineAmount'])
        if calc_line != given_line:
            recomputed_line_amount_issues.append((idx, li.get('ProductID'), given_line, calc_line))
            findings.append({
                'QuoteID': qid, 'Type': 'LineAmountMismatch',
                'Message': f'line {idx} {li.get("ProductID")}: LineAmount {given_line} != Quantity×UnitPrice {calc_line}'
            })
    if line_sum != subtotal:
        findings.append({
            'QuoteID': qid, 'Type': 'SubtotalMismatch',
            'Message': f'line detail sum {line_sum} != header Subtotal {subtotal}'
        })
    if expected_net != net:
        findings.append({
            'QuoteID': qid, 'Type': 'NetAmountMismatch',
            'Message': f'expected NetAmount {expected_net} from Subtotal {subtotal} and DiscountPercent {discount_percent}% != header NetAmount {net}'
        })
    valid_until = date.fromisoformat(header['ValidUntil'])
    expired_unaccepted = (valid_until < BASE_DATE and header['Status'] != 'Accepted')
    if expired_unaccepted:
        findings.append({
            'QuoteID': qid, 'Type': 'ExpiredNotAccepted',
            'Message': f'ValidUntil {header["ValidUntil"]} < {BASE_DATE.isoformat()} and Status {header["Status"]} is not Accepted'
        })
    # store computed audit fields outside workbook details
    header['_LineSum'] = line_sum
    header['_ExpectedNet'] = expected_net
    header['_ExpiredNotAccepted'] = expired_unaccepted
    quotes.append({'header': header, 'lines': lines})
    for li in lines:
        row = dict(li)
        row['QuoteID'] = qid
        all_lines.append(row)

# Sort stable by QuoteID for deterministic workbook, preserving all fetched quotes.
quotes.sort(key=lambda q: q['header']['QuoteID'])
all_lines.sort(key=lambda li: (li['QuoteID'], li['ProductID'], li.get('ProductName','')))

# Fill a copy of the actual template.
shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

# Report sheet
report = wb['Report']
now_iso = datetime.now().astimezone().isoformat(timespec='seconds')
report['B3'] = START_DATE
report['D3'] = END_DATE
report['B4'] = CONDITION
report['B5'] = now_iso
report['B7'] = SOURCE
report['D7'] = now_iso
note_parts = [
    f'Fetched {len(quotes)} quotes and {len(all_lines)} line items.',
    f'Validation base date: {BASE_DATE.isoformat()}.',
    'All fetched quotes retained, including expired/unaccepted quotes.',
]
if findings:
    note_parts.append(f'{len(findings)} finding(s); see findings.md.')
else:
    note_parts.append('No validation findings.')
report['B9'] = ' '.join(note_parts)

# Clear old detail data rows from row 5 down if any.
def clear_from_row(sheet, start_row=5):
    for row in sheet.iter_rows(min_row=start_row, max_row=max(sheet.max_row, start_row), max_col=sheet.max_column):
        for cell in row:
            cell.value = None

quotes_sh = wb['Quotes']
lines_sh = wb['Lines']
summary_sh = wb['Summary']
clear_from_row(quotes_sh)
clear_from_row(lines_sh)
clear_from_row(summary_sh)

quote_headers = [quotes_sh.cell(4, c).value for c in range(1, quotes_sh.max_column+1)]
line_headers = [lines_sh.cell(4, c).value for c in range(1, lines_sh.max_column+1)]
summary_headers = [summary_sh.cell(4, c).value for c in range(1, summary_sh.max_column+1)]

# Write Quotes rows.
for ridx, q in enumerate(quotes, start=5):
    h = q['header']
    rate = fx_rates.get(h['Currency'])
    if rate is None:
        findings.append({'QuoteID': h['QuoteID'], 'Type': 'MissingFXRate', 'Message': f'No FX rate for {h["Currency"]}'})
        net_twd = None
    else:
        net_twd = money(D(h['NetAmount']) * rate)
    values = {
        'QuoteID': h['QuoteID'],
        'CustomerID': h['CustomerID'],
        'ProjectID': h['ProjectID'],
        'QuoteDate': h['QuoteDate'],
        'ValidUntil': h['ValidUntil'],
        'Status': h['Status'],
        'Currency': h['Currency'],
        'Subtotal': float(money(h['Subtotal'])),
        'DiscountPercent': float(D(h['DiscountPercent'])),
        'NetAmount': float(money(h['NetAmount'])),
        'NetAmountTWD': float(net_twd) if net_twd is not None else None,
    }
    for c, name in enumerate(quote_headers, start=1):
        cell = quotes_sh.cell(ridx, c)
        cell.value = values.get(name)
        if name in ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency']:
            cell.number_format = '@'
        elif name in ['Subtotal','NetAmount','NetAmountTWD']:
            cell.number_format = '0.00'
        elif name == 'DiscountPercent':
            cell.number_format = '0.00'

# Write line rows.
for ridx, li in enumerate(all_lines, start=5):
    values = {
        'QuoteID': li['QuoteID'],
        'ProductID': li['ProductID'],
        'ProductName': li['ProductName'],
        'Quantity': int(D(li['Quantity'])),
        'UnitPrice': float(money(li['UnitPrice'])),
        'LineAmount': float(money(li['LineAmount'])),
    }
    for c, name in enumerate(line_headers, start=1):
        cell = lines_sh.cell(ridx, c)
        cell.value = values.get(name)
        if name in ['QuoteID','ProductID','ProductName']:
            cell.number_format = '@'
        elif name in ['UnitPrice','LineAmount']:
            cell.number_format = '0.00'
        elif name == 'Quantity':
            cell.number_format = '0'

# Summary rows: NetAmount by Status/Currency plus Findings summary.
summary_rows = []
from collections import defaultdict
by_status_cur = defaultdict(lambda: {'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0})
for q in quotes:
    h = q['header']
    cur = h['Currency']
    net = money(h['NetAmount'])
    twd = money(net * fx_rates[cur]) if cur in fx_rates else Decimal('0.00')
    key = ('Status', h['Status'], cur)
    by_status_cur[key]['Amount'] += net
    by_status_cur[key]['AmountTWD'] += twd
    by_status_cur[key]['Count'] += 1
for (group, subgroup, cur), agg in sorted(by_status_cur.items()):
    summary_rows.append({
        'Group': group, 'SubGroup': subgroup, 'Currency': cur,
        'Amount': money(agg['Amount']), 'AmountTWD': money(agg['AmountTWD']), 'Count': agg['Count']
    })
by_type = defaultdict(int)
for f in findings:
    by_type[f['Type']] += 1
for typ, count in sorted(by_type.items()):
    summary_rows.append({'Group': 'Finding', 'SubGroup': typ, 'Currency': 'TWD', 'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': count})
for ridx, row in enumerate(summary_rows, start=5):
    values = {
        'Group': row['Group'], 'SubGroup': row['SubGroup'], 'Currency': row['Currency'],
        'Amount': float(row['Amount']), 'AmountTWD': float(row['AmountTWD']), 'Count': row['Count'],
    }
    for c, name in enumerate(summary_headers, start=1):
        cell = summary_sh.cell(ridx, c)
        cell.value = values.get(name)
        if name in ['Group','SubGroup','Currency']:
            cell.number_format = '@'
        elif name in ['Amount','AmountTWD']:
            cell.number_format = '0.00'
        elif name == 'Count':
            cell.number_format = '0'
# Update total formula on summary sheet to include all quote rows.
last_quote_row = 4 + len(quotes)
summary_sh['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})' if len(quotes) else '=SUM(Quotes!K5:K5)'

# Save workbook.
wb.save(OUT_XLSX)

# Write findings.md.
expired = [f for f in findings if f['Type'] == 'ExpiredNotAccepted']
validation = [f for f in findings if f['Type'] != 'ExpiredNotAccepted']
lines = []
lines.append('# Quote validation findings')
lines.append('')
lines.append(f'- Query period: {START_DATE} to {END_DATE} inclusive')
lines.append(f'- Base date for expiration flag: {BASE_DATE.isoformat()}')
lines.append(f'- Source: {SOURCE}')
lines.append(f'- Quotes retained in workbook: {len(quotes)}')
lines.append(f'- Line items retained in workbook: {len(all_lines)}')
lines.append('')
lines.append('## Validation rules')
lines.append('- Detail subtotal check: sum(LineAmount) per quote must equal header Subtotal.')
lines.append('- NetAmount check: Subtotal × (1 − DiscountPercent / 100), rounded ROUND_HALF_UP to 2 decimals, must equal header NetAmount.')
lines.append('- LineAmount check: Quantity × UnitPrice, rounded to 2 decimals, must equal LineAmount.')
lines.append(f'- Expired/unaccepted flag: ValidUntil < {BASE_DATE.isoformat()} and Status != Accepted.')
lines.append('')
lines.append('## Summary')
lines.append(f'- Total findings: {len(findings)}')
lines.append(f'- Expired and not Accepted quotes: {len(expired)}')
for typ in sorted(by_type):
    lines.append(f'- {typ}: {by_type[typ]}')
lines.append('')
lines.append('## Expired and not Accepted as of 2026-09-30')
if expired:
    lines.append('| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |')
    lines.append('|---|---|---|---|---|---|---|---:|')
    for f in expired:
        h = next(q['header'] for q in quotes if q['header']['QuoteID'] == f['QuoteID'])
        lines.append(f"| {h['QuoteID']} | {h['CustomerID']} | {h['ProjectID']} | {h['QuoteDate']} | {h['ValidUntil']} | {h['Status']} | {h['Currency']} | {money(h['NetAmount'])} |")
else:
    lines.append('None.')
lines.append('')
lines.append('## Detail findings')
if findings:
    lines.append('| QuoteID | Type | Finding |')
    lines.append('|---|---|---|')
    for f in sorted(findings, key=lambda x: (x['QuoteID'], x['Type'], x['Message'])):
        msg = f['Message'].replace('|', '\\|')
        lines.append(f"| {f['QuoteID']} | {f['Type']} | {msg} |")
else:
    lines.append('No validation findings.')
lines.append('')
lines.append('## Completeness')
lines.append('- No quotes were removed because of expiration status.')
lines.append('- Workbook Quotes sheet contains every QuoteID returned by the September 2026 list query.')
OUT_MD.write_text('\n'.join(lines) + '\n', encoding='utf-8')

# Verification: reopen outputs and check requirements.
print('\nVERIFY OUTPUTS')
if not OUT_XLSX.exists():
    raise AssertionError('Missing output workbook')
if not OUT_MD.exists():
    raise AssertionError('Missing findings.md')
wbv = load_workbook(OUT_XLSX, data_only=False)
qv = wbv['Quotes']
lv = wbv['Lines']
sv = wbv['Summary']
rv = wbv['Report']
q_header_v = [qv.cell(4,c).value for c in range(1, qv.max_column+1)]
l_header_v = [lv.cell(4,c).value for c in range(1, lv.max_column+1)]
s_header_v = [sv.cell(4,c).value for c in range(1, sv.max_column+1)]
print('Quotes header:', q_header_v)
print('Lines header:', l_header_v)
print('Summary header:', s_header_v)
assert q_header_v == quote_headers, 'Quotes headers changed'
assert l_header_v == line_headers, 'Lines headers changed'
assert s_header_v == summary_headers, 'Summary headers changed'
# Count non-empty data rows.
quote_rows = []
for r in range(5, qv.max_row+1):
    if qv.cell(r,1).value is not None:
        quote_rows.append(r)
line_rows = []
for r in range(5, lv.max_row+1):
    if lv.cell(r,1).value is not None:
        line_rows.append(r)
print('Quote data rows:', len(quote_rows), 'Expected:', len(quotes))
print('Line data rows:', len(line_rows), 'Expected:', len(all_lines))
assert len(quote_rows) == len(quotes), 'Quote row count mismatch'
assert len(line_rows) == len(all_lines), 'Line row count mismatch'
# Verify every QuoteID retained and text formatting for IDs.
out_qids = [str(qv.cell(r,1).value) for r in quote_rows]
expected_qids = [q['header']['QuoteID'] for q in quotes]
print('Output QuoteIDs:', out_qids)
assert out_qids == expected_qids, 'Quote IDs not all retained or order mismatch'
for r in quote_rows:
    for c in [1,2,3]:
        assert qv.cell(r,c).number_format == '@', f'ID/date-ish text format missing at Quotes {r},{c}'
    for c in [8,9,10,11]:
        assert isinstance(qv.cell(r,c).value, (int,float)), f'Numeric value expected at Quotes {r},{c}'
for r in line_rows:
    for c in [1,2,3]:
        assert lv.cell(r,c).number_format == '@', f'Text format missing at Lines {r},{c}'
    for c in [4,5,6]:
        assert isinstance(lv.cell(r,c).value, (int,float)), f'Numeric value expected at Lines {r},{c}'
assert rv['B3'].value == START_DATE and rv['D3'].value == END_DATE, 'Report dates not filled'
md = OUT_MD.read_text(encoding='utf-8')
assert 'Expired and not Accepted' in md, 'findings.md missing expiration section'
assert 'No quotes were removed' in md, 'findings.md missing completeness statement'
for f in findings:
    assert f['QuoteID'] in md, f'finding quote {f["QuoteID"]} missing in md'
print('findings.md bytes:', OUT_MD.stat().st_size)
print('Output files:')
print('-', OUT_XLSX)
print('-', OUT_MD)
print('\nFINDINGS SUMMARY:', dict(sorted(by_type.items())))
