from pathlib import Path
import csv, json, os
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

# Inspect CSV headers/sample
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i, row in zip(range(6), r):
            print(i, row)

# Inspect manual first part
manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\nquote_cli_manual.md first 120 lines:')
    txt = manual.read_text(encoding='utf-8')
    for i, line in enumerate(txt.splitlines()[:120], 1):
        print(f'{i:03d}: {line}')

# Inspect workbook sheets dimensions, headers, sample cells
xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    print('\nquote_template.xlsx:')
    wb = load_workbook(xlsx, data_only=False)
    for sh in wb.worksheets:
        print('SHEET', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(row)
        print('---')
