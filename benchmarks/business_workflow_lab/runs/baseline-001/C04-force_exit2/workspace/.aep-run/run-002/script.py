from pathlib import Path
import subprocess, sys, json, os, csv
from openpyxl import load_workbook
ws=Path.cwd()
cli=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
cmd=[sys.executable, str(cli), 'quotes','list','--from','2026-09-01','--to','2026-09-30','--status','Sent','--format','json']
print('Running:', ' '.join(cmd))
p=subprocess.run(cmd, cwd=ws, text=True, capture_output=True)
print('exit', p.returncode)
print('stderr:', p.stderr)
print('stdout:', p.stdout[:1000])
print('\nFX rates:')
print((ws/'input'/'fx_rates.csv').read_text())
print('\nWorkbook sheets/dimensions/cells:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for sh in wb.worksheets:
    print('SHEET', sh.title, sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,15), values_only=False):
        vals=[]
        for c in row[:min(sh.max_column,8)]:
            vals.append(f'{c.coordinate}={c.value!r}')
        print(' | '.join(vals))
    print('---')
