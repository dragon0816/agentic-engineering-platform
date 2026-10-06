from pathlib import Path
import sys, os, csv
from openpyxl import load_workbook
ws=Path.cwd()
cli=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
print('CLI exists', cli.exists(), cli)
print(cli.read_text(encoding='utf-8', errors='replace')[:4000])
print('\nFX rates:')
print((ws/'input'/'fx_rates.csv').read_text())
print('\nWorkbook sheets/dimensions/cells:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for sh in wb.worksheets:
    print('SHEET', sh.title, sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,20), values_only=False):
        print(' | '.join(f'{c.coordinate}={c.value!r}' for c in row[:min(sh.max_column,10)]))
    print('---')
