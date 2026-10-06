from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/quote_template.xlsx', data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)
print('\nReport footer-ish rows:')
r=wb['Report']
for i in range(1, r.max_row+1):
    print(i, [r.cell(i,j).value for j in range(1,r.max_column+1)])
print('\nSummary rows')
s=wb['Summary']
for i in range(1, s.max_row+1): print(i, [s.cell(i,j).value for j in range(1,s.max_column+1)])
