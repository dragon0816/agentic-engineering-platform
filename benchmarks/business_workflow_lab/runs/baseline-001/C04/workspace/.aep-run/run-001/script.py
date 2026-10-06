from pathlib import Path
from openpyxl import load_workbook
import csv, json, os
ws = Path('.')
print('cwd', ws.resolve())
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p, p.stat().st_size)
manual = ws/'input'/'quote_cli_manual.md'
print('\n--- quote_cli_manual.md ---')
print(manual.read_text(encoding='utf-8', errors='replace'))
print('\n--- fx_rates.csv ---')
print((ws/'input'/'fx_rates.csv').read_text(encoding='utf-8', errors='replace'))
print('\n--- workbook structure ---')
wb = load_workbook(ws/'input'/'quote_template.xlsx')
print('sheets', wb.sheetnames)
for sh in wb.worksheets:
    print('sheet', sh.title, 'max', sh.max_row, sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(r)
    print('merged', list(sh.merged_cells.ranges))
