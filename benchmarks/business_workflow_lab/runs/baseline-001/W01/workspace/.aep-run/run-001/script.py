from pathlib import Path
import os, json, csv
from openpyxl import load_workbook
base=Path('.')
print('cwd', Path.cwd())
for p in sorted(Path('input').glob('*')):
    print('INPUT', p, p.stat().st_size)
print('\n--- sap_manual.md ---')
print(Path('input/sap_manual.md').read_text(encoding='utf-8'))
print('\n--- fx_rates.csv ---')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8'))
wb=load_workbook('input/sales_template.xlsx')
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('Sheet', ws.title, 'max', ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=True):
        print(row)
    print('---')