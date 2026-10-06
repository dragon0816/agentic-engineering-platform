from pathlib import Path
from openpyxl import load_workbook
import csv, json, os
base=Path('.')
print('cwd', base.resolve())
print('input files:', [p.name for p in (base/'input').iterdir()])
manual=(base/'input'/'quote_cli_manual.md').read_text(encoding='utf-8')
print('--- manual ---')
print(manual[:4000])
wb=load_workbook(base/'input'/'quote_template.xlsx')
print('sheets', wb.sheetnames)
for ws in wb.worksheets:
    print('SHEET', ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=True):
        print(row)
    print('---')
print('fx_rates')
print((base/'input'/'fx_rates.csv').read_text(encoding='utf-8')[:1000])