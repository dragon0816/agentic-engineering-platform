from pathlib import Path
import os
from openpyxl import load_workbook
ws=Path('.')
print('cwd', ws.resolve())
print('input files:')
for p in Path('input').glob('*'):
    print(p, p.stat().st_size)
print('output exists', Path('output').exists())
# inspect workbook sheets and dimensions/headers
xlsx=Path('input/sales_template.xlsx')
wb=load_workbook(xlsx)
print('sheets', wb.sheetnames)
for sh in wb.worksheets:
    print('sheet', sh.title, sh.max_row, sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(r)
    print('---')
print('fx_rates.csv head:')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8-sig')[:1000])
print('sap_manual.md head:')
print(Path('input/sap_manual.md').read_text(encoding='utf-8-sig')[:2000])
