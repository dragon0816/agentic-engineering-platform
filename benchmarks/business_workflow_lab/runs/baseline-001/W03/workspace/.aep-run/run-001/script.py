from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path('.')
print('cwd', ws.resolve())
print('input files:', [str(p) for p in Path('input').glob('*')])
print('\n--- sap_manual.md ---')
print(Path('input/sap_manual.md').read_text(encoding='utf-8', errors='replace')[:5000])
print('\n--- fx_rates.csv ---')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8', errors='replace'))
wb=load_workbook('input/project_template.xlsx')
print('sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('\nSHEET', s, 'max', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,10), values_only=True):
        print(row)
