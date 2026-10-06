from pathlib import Path
import openpyxl, csv, json
base=Path('.')
print('cwd', base.resolve())
print('input files:', [str(p) for p in Path('input').glob('*')])
print('\n--- quote_cli_manual.md ---')
print(Path('input/quote_cli_manual.md').read_text(encoding='utf-8'))
print('\n--- fx_rates.csv ---')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8'))
wb=openpyxl.load_workbook('input/quote_template.xlsx')
print('sheets', wb.sheetnames)
for ws in wb.worksheets:
    print('\nSHEET', ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=True):
        print(row)
