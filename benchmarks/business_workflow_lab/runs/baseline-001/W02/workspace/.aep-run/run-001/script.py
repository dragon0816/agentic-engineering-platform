from pathlib import Path
from openpyxl import load_workbook
import csv, json, os
base=Path('.')
print('cwd', base.resolve())
print('tree:')
for p in sorted(base.rglob('*')):
    print(p)
print('\nmanual preview')
mp=base/'input'/'sap_manual.md'
print(mp.read_text(encoding='utf-8')[:4000])
print('\nfx')
print((base/'input'/'fx_rates.csv').read_text(encoding='utf-8')[:2000])
print('\nworkbook sheets/cells')
wb=load_workbook(base/'input'/'sales_template.xlsx')
for ws in wb.worksheets:
    print('SHEET', ws.title, ws.max_row, ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,15), values_only=True):
        print(r)
    print('merged', list(ws.merged_cells.ranges))