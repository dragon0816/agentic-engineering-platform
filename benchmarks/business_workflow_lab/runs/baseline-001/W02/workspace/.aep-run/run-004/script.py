from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx')
for row in wb['Instructions'].iter_rows(values_only=True): print(row)
# fetch details pages print field text
import urllib.request, urllib.parse, http.cookiejar, re, html
base='http://127.0.0.1:8810'; cj=http.cookiejar.CookieJar(); opener=urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
def open_path(path,data=None):
    if data is not None:
        req=urllib.request.Request(base+path, data=urllib.parse.urlencode(data).encode(), method='POST')
    else: req=urllib.request.Request(base+path)
    with opener.open(req,timeout=10) as r: return r.read().decode('utf-8',errors='replace')
open_path('/login'); open_path('/login', {'username':'lab.analyst','password':'Lab#2026'})
ids=['0045000101','0045000107','0045000109','0045000114','0045000119','0045000125','0045000129','0045000132','0045000136']
Path('output/workflow/details').mkdir(parents=True, exist_ok=True)
for oid in ids:
    s=open_path('/orders/'+oid)
    Path(f'output/workflow/details/{oid}.html').write_text(s,encoding='utf-8')
    print('\n---',oid,'len',len(s))
    # strip tags rudimentary around main
    main=s[s.find('<main'):s.find('</main>')]
    text=re.sub('<[^>]+>','\n',main)
    text=html.unescape(re.sub('\n+','\n',text))
    print(text[:2000])