from pathlib import Path
import os, json
from openpyxl import load_workbook
base=Path('.')
print('cwd', Path.cwd())
print('input files:')
for p in sorted(Path('input').glob('*')):
    print(p, p.stat().st_size)
    if p.suffix.lower() in ['.md','.csv','.txt']:
        print('---', p, '---')
        print(p.read_text(encoding='utf-8', errors='replace')[:4000])
        print('--- end preview ---')
# inspect workbook sheets/cells
wb=load_workbook('input/project_template.xlsx', data_only=False)
print('workbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('\nSHEET', ws.title, 'max', ws.max_row, ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=True):
        print(r)
