from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/quote_template.xlsx')
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)
