from pathlib import Path
from openpyxl import load_workbook
ws=Path('.')
print('cwd', ws.resolve())
print('input files:')
for p in Path('input').glob('*'):
    print(p, p.stat().st_size)
manual=Path('input/quote_cli_manual.md').read_text(encoding='utf-8')
print('\n--- manual ---')
print(manual[:4000])
print('\n--- fx ---')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8'))
wb=load_workbook('input/quote_template.xlsx')
print('sheets', wb.sheetnames)
for sh in wb.sheetnames:
    ws=wb[sh]
    print('\nSHEET', sh, 'max', ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=True):
        print(row)
