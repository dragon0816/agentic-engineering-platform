from pathlib import Path
import os, json
base=Path('.')
print('CWD', Path.cwd())
for p in sorted(base.rglob('*')):
    print(p, 'DIR' if p.is_dir() else f'FILE {p.stat().st_size}')
print('\nManual previews:')
for name in ['input/sap_manual.md','input/expense_manual.md','input/quote_cli_manual.md','input/fx_rates.csv']:
    p=Path(name)
    print('\n---', name, '---')
    print(p.read_text(encoding='utf-8')[:4000])

# inspect xlsx sheets dimensions and cells
from openpyxl import load_workbook
wb=load_workbook('input/project_template.xlsx')
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('Sheet', ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,20), values_only=True):
        print(row)
    print('---')
