from pathlib import Path
from openpyxl import load_workbook
print(Path('input/expense_manual.md').read_text(encoding='utf-8'))
print('\nFX:')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8'))
wb=load_workbook('input/expense_template.xlsx')
print('sheets', wb.sheetnames)
for ws in wb.worksheets:
    print('\nSHEET', ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=True):
        print(row)
