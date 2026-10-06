from pathlib import Path
import json, csv
from openpyxl import load_workbook
ws=Path('.')
print('cwd', ws.resolve())
for p in ws.rglob('*'):
    print(p, 'file' if p.is_file() else 'dir')
print('\n--- manual ---')
mp=Path('input/expense_manual.md')
print(mp.read_text(encoding='utf-8'))
print('\n--- fx sample ---')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8')[:2000])
print('\n--- template ---')
wb=load_workbook('input/expense_template.xlsx')
print(wb.sheetnames)
for sh in wb.worksheets:
    print('sheet', sh.title, 'max', sh.max_row, sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,15), values_only=True):
        print(r)
    print('---')