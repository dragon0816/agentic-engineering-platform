from pathlib import Path
from openpyxl import load_workbook
import json, csv
# print instructions all rows
wb=load_workbook('input/expense_template.xlsx', data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True): print(row)
# print app data metadata
p=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\operator\data\source.json')
print('\nsource exists', p.exists(), p)
if p.exists():
    data=json.loads(p.read_text(encoding='utf-8'))
    print(type(data), data.keys() if isinstance(data,dict) else 'list')
    if isinstance(data, dict):
        for k,v in data.items():
            print('key',k,'type',type(v),'len', len(v) if hasattr(v,'__len__') else '')
            print(str(v)[:1000])
