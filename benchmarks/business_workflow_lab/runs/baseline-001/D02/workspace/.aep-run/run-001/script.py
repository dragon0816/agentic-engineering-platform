from pathlib import Path
from openpyxl import load_workbook
import csv, json, os
base=Path('.')
print('cwd', base.resolve())
print('input files:', [str(p) for p in Path('input').glob('*')])
print('workspace tree sample:')
for p in list(Path('.').rglob('*'))[:80]: print(p)
manual=Path('input/expense_manual.md')
print('\nMANUAL:\n', manual.read_text(encoding='utf-8'))
print('\nFX:')
print(Path('input/fx_rates.csv').read_text(encoding='utf-8'))
wb=load_workbook('input/expense_template.xlsx')
print('sheets', wb.sheetnames)
for ws in wb.worksheets:
    print('sheet', ws.title, 'max', ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,10), values_only=True): print(row)
