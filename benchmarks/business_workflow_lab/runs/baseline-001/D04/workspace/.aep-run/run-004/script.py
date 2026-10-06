from openpyxl import load_workbook
from pathlib import Path
wb=load_workbook('input/expense_template.xlsx')
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)
# inspect formulas/styles/dims maybe merged ranges
for s in wb.sheetnames:
    ws=wb[s]
    print('SHEET',s,'merged', list(ws.merged_cells.ranges))
    print('freeze', ws.freeze_panes)
    for cell in ['B3','D3','B4','B5','B7','D7','B9','B2','A20']:
        if cell in ws:
            c=ws[cell]
            print(cell, c.value, c.number_format)
