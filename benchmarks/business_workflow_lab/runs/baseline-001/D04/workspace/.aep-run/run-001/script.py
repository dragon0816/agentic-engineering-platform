from pathlib import Path
from openpyxl import load_workbook
import csv, json
base=Path('.')
print('cwd', base.resolve())
for p in sorted((base/'input').glob('**/*')):
    print('INPUT', p, p.stat().st_size)
manual=base/'input'/'expense_manual.md'
if manual.exists():
    print('\n--- manual ---')
    print(manual.read_text(encoding='utf-8'))
fx=base/'input'/'fx_rates.csv'
if fx.exists():
    print('\n--- fx ---')
    print(fx.read_text(encoding='utf-8'))
tmpl=base/'input'/'expense_template.xlsx'
if tmpl.exists():
    wb=load_workbook(tmpl)
    print('\n--- workbook ---')
    print('sheets', wb.sheetnames)
    for ws in wb.worksheets:
        print('sheet', ws.title, 'max', ws.max_row, ws.max_column)
        for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,15), values_only=True):
            print(r)
        print('---')
