from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/expense_template.xlsx')
for sh in ['Instructions','Report']:
    s=wb[sh]
    print('\nSHEET', sh)
    for row in s.iter_rows(values_only=True):
        print(row)
# Read app data summary, not full descriptions maybe
p=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\operator\data\source.json')
print('\ndata exists', p.exists(), p)
if p.exists():
    import json
    data=json.loads(p.read_text(encoding='utf-8'))
    print(data.keys())
    for k,v in data.items():
        print(k, type(v), len(v) if isinstance(v,list) else None)
    for e in data['expenses'][:3]: print(e)
