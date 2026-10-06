from pathlib import Path
import os, json, csv
from openpyxl import load_workbook
ws=Path('.')
print('cwd', ws.resolve())
for p in sorted(ws.rglob('*')):
    print(p, 'dir' if p.is_dir() else p.stat().st_size)
# inspect manual and rates and workbook sheets/headers
inp=ws/'input'
for name in ['expense_manual.md','fx_rates.csv']:
    p=inp/name
    print('\n---', name, '---')
    print(p.read_text(encoding='utf-8', errors='replace')[:4000])
wb=load_workbook(inp/'expense_template.xlsx')
print('workbook sheets', wb.sheetnames)
for sh in wb.sheetnames:
    s=wb[sh]
    print('sheet', sh, 'max', s.max_row, s.max_column)
    for row in s.iter_rows(min_row=1, max_row=min(s.max_row,5), values_only=True):
        print(row)
