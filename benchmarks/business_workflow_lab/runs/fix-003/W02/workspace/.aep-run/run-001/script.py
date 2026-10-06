import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
for p in sorted((ws/'input').iterdir()):
    print('INPUT', p.name, p.stat().st_size)
print('\nmanual head:')
mp=ws/'input'/'sap_manual.md'
print(mp.read_text(encoding='utf-8')[:2000])
print('\nfx sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(5), r): print(row)
print('\ntemplate sheets/headers/sample:')
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    s=wb[sh]
    print('SHEET', sh, 'max', s.max_row, s.max_column)
    for row in s.iter_rows(min_row=1, max_row=min(5,s.max_row), values_only=True): print(row)
