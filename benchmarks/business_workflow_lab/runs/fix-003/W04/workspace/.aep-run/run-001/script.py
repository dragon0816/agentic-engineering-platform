from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in (ws/'input').iterdir():
    print('-', p.name, p.stat().st_size)
# workbook sheets and headers/sample dimensions
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
print('workbook sheets:', wb.sheetnames)
for sh in wb.worksheets:
    print('sheet', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(5, sh.max_row), values_only=True):
        print(' ', r)
# fx headers
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rows=list(csv.reader(f))
print('fx rows', len(rows))
for r in rows[:5]: print('fx', r)
# manual excerpt
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print('manual chars', len(text))
print(text[:4000])