from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
for sh in ['Instructions']:
    s=wb[sh]
    for row in s.iter_rows(values_only=True): print(row)