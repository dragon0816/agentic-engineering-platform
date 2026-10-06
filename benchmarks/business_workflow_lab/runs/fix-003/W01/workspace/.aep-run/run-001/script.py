import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv preview:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(6), r): print(row)
print('\nsap_manual.md preview:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[:2000])
print('\nsales_template.xlsx structure:')
wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
print(wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('sheet', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(8, wsht.max_row), values_only=True):
        print(row)
    print('---')