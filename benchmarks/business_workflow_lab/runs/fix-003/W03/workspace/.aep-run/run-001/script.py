from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
for p in sorted((ws/'input').iterdir()):
    print('INPUT', p.name, p.stat().st_size)

# inspect workbook
wb_path = ws/'input'/'project_template.xlsx'
wb = load_workbook(wb_path, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('\nSHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
        print(r)

# inspect fx csv
csv_path = ws/'input'/'fx_rates.csv'
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader):
        print('CSV', i, row)

# inspect sap_manual beginning
md_path=ws/'input'/'sap_manual.md'
text=md_path.read_text(encoding='utf-8', errors='replace')
print('\nMD first 2000 chars:\n', text[:2000])