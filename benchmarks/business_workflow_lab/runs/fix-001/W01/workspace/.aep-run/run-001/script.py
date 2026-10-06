from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
input_dir=ws/'input'
print('input files:', [p.name for p in input_dir.iterdir()])
# workbook details
wb_path=input_dir/'sales_template.xlsx'
wb=load_workbook(wb_path)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)
# fx csv
fx=input_dir/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in enumerate(r):
        print(row)
        if i>=10: break
# manual
manual=input_dir/'sap_manual.md'
print('\nsap_manual.md first 120 lines:')
text=manual.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:120],1): print(f'{i}: {line}')