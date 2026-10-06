from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
for p in sorted((ws/'input').iterdir()):
    print('INPUT', p.name, p.stat().st_size)
print('\nmanual head:')
print((ws/'input'/'expense_manual.md').read_text(encoding='utf-8')[:4000])
print('\nfx_rates head:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    for i,row in zip(range(10), csv.reader(f)):
        print(row)
print('\ntemplate sheets/cells:')
wb=load_workbook(ws/'input'/'expense_template.xlsx')
print(wb.sheetnames)
for sh in wb.worksheets:
    print('SHEET', sh.title, sh.max_row, sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,15), values_only=True):
        print(r)
    print('---')