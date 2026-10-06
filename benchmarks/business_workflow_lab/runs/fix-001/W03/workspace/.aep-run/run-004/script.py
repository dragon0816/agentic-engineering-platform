from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
download_dir = output_dir / 'downloads'
src_template = input_dir / 'project_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
invoice_csv = download_dir / 'invoices_2026-09.csv'
out_path = output_dir / 'project_report_2026-09_invoices.xlsx'

print('Final rebuild and verification')
print('Input files and samples:')
for p in [src_template, fx_path, invoice_csv]:
    print('-', p, 'exists=', p.exists(), 'size=', p.stat().st_size if p.exists() else None)

rates = {}
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    fx_header = next(reader)
    for row in reader:
        if row and row[0].strip():
            rates[row[0].strip()] = Decimal(row[1].strip())
print('FX header:', fx_header)
print('FX rates used:', {k: str(v) for k, v in rates.items()})

with open(invoice_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    invoices = list(reader)
    invoice_headers = reader.fieldnames
print('Invoice CSV headers:', invoice_headers)
print('Invoice CSV row count:', len(invoices))
print('Invoice CSV sample rows:', invoices[:3])

required_csv_headers = ['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance']
if invoice_headers != required_csv_headers:
    raise RuntimeError(f'Unexpected invoice CSV headers: {invoice_headers}')

wb = load_workbook(src_template)

def clear_data(sheet_name):
    ws = wb[sheet_name]
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)
    return ws

# Keep other source worksheets blank.
for sheet in ['Orders', 'Expenses', 'Quotes']:
    clear_data(sheet)

# Report metadata.
ws_report = wb['Report']
ws_report['B3'] = '2026-09-01'
ws_report['D3'] = '2026-09-30'
ws_report['B4'] = 'Invoices only; outstanding balance as of 2026-09-30; Orders/Expenses/Quotes left blank'
ws_report['B5'] = '2026-10-06 Asia/Taipei'

# Fill invoice detail.
ws_inv = clear_data('Invoices')
expected_invoice_sheet_headers = ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
if [ws_inv.cell(4,c).value for c in range(1,12)] != expected_invoice_sheet_headers:
    raise RuntimeError('Template Invoices headers do not match expected structure')

summary = {}
exceptions = []
base_date = date(2026, 9, 30)
for row in sorted(invoices, key=lambda r: r['InvoiceID']):
    inv_date = datetime.strptime(row['InvoiceDate'], '%Y-%m-%d').date()
    if not (date(2026,9,1) <= inv_date <= date(2026,9,30)):
        raise RuntimeError(f"Out-of-period invoice: {row['InvoiceID']} {row['InvoiceDate']}")
    if row['Currency'] not in rates:
        raise RuntimeError(f"Missing rate for {row['Currency']}")

for excel_row, row in enumerate(sorted(invoices, key=lambda r: r['InvoiceID']), start=5):
    amount = Decimal(row['Amount'])
    paid = Decimal(row['PaidAmount'])
    balance = Decimal(row['Balance'])
    balance_twd = (balance * rates[row['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    vals = [row['InvoiceID'], row['OrderID'], row['ProjectID'], row['CustomerID'], row['InvoiceDate'], row['DueDate'], row['Currency'], float(amount), float(paid), float(balance), float(balance_twd)]
    for col, val in enumerate(vals, 1):
        cell = ws_inv.cell(row=excel_row, column=col, value=val)
        cell.number_format = '@' if col <= 7 else '#,##0.00'
    if row['ProjectID']:
        summary[row['ProjectID']] = summary.get(row['ProjectID'], Decimal('0.00')) + balance_twd
    due = datetime.strptime(row['DueDate'], '%Y-%m-%d').date()
    if balance > 0 and due < base_date:
        exceptions.append(['Invoices', row['InvoiceID'], row['ProjectID'], 'Overdue'])

# Fill summary by project.
ws_sum = clear_data('Summary')
expected_summary_headers = ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']
if [ws_sum.cell(4,c).value for c in range(1,6)] != expected_summary_headers:
    raise RuntimeError('Template Summary headers do not match expected structure')
for excel_row, pid in enumerate(sorted(summary), start=5):
    vals = [pid, 0.0, float(summary[pid].quantize(Decimal('0.01'))), 0.0, 0.0]
    for col, val in enumerate(vals, 1):
        cell = ws_sum.cell(row=excel_row, column=col, value=val)
        cell.number_format = '@' if col == 1 else '#,##0.00'
last_summary_row = max(5, 4 + len(summary))
for col in 'BCDE':
    ws_sum[f'{col}2'] = f'=SUM({col}5:{col}{last_summary_row})'

# Fill exceptions.
ws_exc = clear_data('Exceptions')
expected_exception_headers = ['Source','DocumentID','ProjectID','Reason']
if [ws_exc.cell(4,c).value for c in range(1,5)] != expected_exception_headers:
    raise RuntimeError('Template Exceptions headers do not match expected structure')
for excel_row, vals in enumerate(exceptions, start=5):
    for col, val in enumerate(vals, 1):
        cell = ws_exc.cell(row=excel_row, column=col, value=val if val != '' else None)
        cell.number_format = '@'

wb.save(out_path)
print('Saved output workbook:', out_path)

# Reopen saved workbook and verify requirements.
wb2 = load_workbook(out_path, data_only=False)
errors = []
def norm_blank(x):
    return '' if x is None else x
def expect(name, actual, expected):
    ok = actual == expected
    print(f'- {name}: {"PASS" if ok else "FAIL"} actual={actual!r} expected={expected!r}')
    if not ok:
        errors.append(name)

print('\nVerification checks from saved output:')
expect('Output file exists', out_path.exists(), True)
expect('Invoices row count', wb2['Invoices'].max_row - 4, len(invoices))
for s in ['Orders','Expenses','Quotes']:
    expect(f'{s} remains blank below header', wb2[s].max_row - 4, 0)
expect('Invoices columns/header names', [wb2['Invoices'].cell(4,c).value for c in range(1,12)], expected_invoice_sheet_headers)
expect('Summary columns/header names', [wb2['Summary'].cell(4,c).value for c in range(1,6)], expected_summary_headers)
expect('Exceptions columns/header names', [wb2['Exceptions'].cell(4,c).value for c in range(1,5)], expected_exception_headers)

saved_inv = []
for r in range(5, wb2['Invoices'].max_row + 1):
    saved_inv.append([wb2['Invoices'].cell(r,c).value for c in range(1,12)])
expect('All invoices are in September 2026', all('2026-09-01' <= v[4] <= '2026-09-30' for v in saved_inv), True)
expect('Invoice IDs and Order IDs saved as text', all(isinstance(v[0], str) and isinstance(v[1], str) for v in saved_inv), True)
expect('Amount, PaidAmount, Balance, BalanceTWD are numeric', all(isinstance(v[c], (int,float)) for v in saved_inv for c in [7,8,9,10]), True)

# Exact BalanceTWD conversion per row from saved data.
row_expected_baltwd = {}
for row in invoices:
    row_expected_baltwd[row['InvoiceID']] = (Decimal(row['Balance']) * rates[row['Currency']]).quantize(Decimal('0.01'))
row_actual_baltwd = {v[0]: Decimal(str(v[10])).quantize(Decimal('0.01')) for v in saved_inv}
expect('Per-invoice BalanceTWD conversion', row_actual_baltwd, row_expected_baltwd)

saved_ex = []
for r in range(5, wb2['Exceptions'].max_row + 1):
    saved_ex.append([norm_blank(wb2['Exceptions'].cell(r,c).value) for c in range(1,5)])
expect('Exceptions are Balance > 0 and DueDate < 2026-09-30 only', saved_ex, exceptions)
due_on_base = [v[0] for v in saved_inv if v[5] == '2026-09-30' and Decimal(str(v[9])) > 0]
expect('Invoices due exactly 2026-09-30 not overdue', all(doc not in [e[1] for e in saved_ex] for doc in due_on_base), True)

saved_summary = {wb2['Summary'].cell(r,1).value: Decimal(str(wb2['Summary'].cell(r,3).value)).quantize(Decimal('0.01')) for r in range(5, wb2['Summary'].max_row + 1)}
expected_summary = {k: v.quantize(Decimal('0.01')) for k,v in summary.items()}
expect('Summary project IDs', sorted(saved_summary), sorted(expected_summary))
expect('Summary InvoiceBalanceTWD by project', saved_summary, expected_summary)
expect('Summary other source TWD columns zero', all(wb2['Summary'].cell(r,c).value == 0 for r in range(5, wb2['Summary'].max_row + 1) for c in [2,4,5]), True)
expect('Summary total formulas set', [wb2['Summary'][f'{c}2'].value for c in 'BCDE'], [f'=SUM({c}5:{c}{last_summary_row})' for c in 'BCDE'])

print('\nComputed summary:', {k: str(v.quantize(Decimal('0.01'))) for k,v in sorted(summary.items())})
print('Computed exceptions:', exceptions)
if errors:
    raise RuntimeError('Verification failed: ' + ', '.join(errors))
print('All verification checks passed.')