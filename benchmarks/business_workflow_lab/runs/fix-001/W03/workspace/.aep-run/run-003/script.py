from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
download_dir = output_dir / 'downloads'
src_template = input_dir / 'project_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
invoice_csv = download_dir / 'invoices_2026-09.csv'
out_path = output_dir / 'project_report_2026-09_invoices.xlsx'

print('Rebuilding and verifying workbook')
print('Files:')
for p in [src_template, fx_path, invoice_csv]:
    print(p.name, 'exists=', p.exists(), 'size=', p.stat().st_size if p.exists() else None)

rates = {}
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    print('FX header/sample:', next(reader))
    for row in reader:
        if row and row[0].strip():
            rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates used:', {k: str(v) for k, v in rates.items()})

with open(invoice_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    invoices = list(reader)
    csv_headers = reader.fieldnames
print('Invoice CSV headers:', csv_headers)
print('Invoice count:', len(invoices))
print('Invoice sample:', invoices[:2])

wb = load_workbook(src_template)

def clear_data(sheet_name):
    ws = wb[sheet_name]
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)
    return ws

for s in ['Orders', 'Expenses', 'Quotes']:
    clear_data(s)

ws_report = wb['Report']
ws_report['B3'] = '2026-09-01'
ws_report['D3'] = '2026-09-30'
ws_report['B4'] = 'Invoices only; outstanding balance as of 2026-09-30; Orders/Expenses/Quotes left blank'
ws_report['B5'] = '2026-10-06 Asia/Taipei'

ws_inv = clear_data('Invoices')
summary = {}
exceptions = []
base_date = date(2026, 9, 30)

for row in invoices:
    inv_date = datetime.strptime(row['InvoiceDate'], '%Y-%m-%d').date()
    if inv_date < date(2026, 9, 1) or inv_date > date(2026, 9, 30):
        raise RuntimeError(f'Out of range invoice {row}')
    cur = row['Currency']
    if cur not in rates:
        raise RuntimeError(f'Missing FX {cur}')

for excel_row, row in enumerate(sorted(invoices, key=lambda r: r['InvoiceID']), start=5):
    amount = Decimal(row['Amount'])
    paid = Decimal(row['PaidAmount'])
    balance = Decimal(row['Balance'])
    bal_twd = (balance * rates[row['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    vals = [
        row['InvoiceID'], row['OrderID'], row['ProjectID'], row['CustomerID'],
        row['InvoiceDate'], row['DueDate'], row['Currency'],
        float(amount), float(paid), float(balance), float(bal_twd),
    ]
    for c, v in enumerate(vals, 1):
        cell = ws_inv.cell(excel_row, c, v)
        cell.number_format = '@' if c <= 7 else '#,##0.00'
    if row['ProjectID']:
        summary[row['ProjectID']] = summary.get(row['ProjectID'], Decimal('0.00')) + bal_twd
    due = datetime.strptime(row['DueDate'], '%Y-%m-%d').date()
    if balance > 0 and due < base_date:
        exceptions.append((row['InvoiceID'], row['ProjectID']))

ws_sum = clear_data('Summary')
for excel_row, pid in enumerate(sorted(summary), start=5):
    vals = [pid, 0.0, float(summary[pid]), 0.0, 0.0]
    for c, v in enumerate(vals, 1):
        cell = ws_sum.cell(excel_row, c, v)
        cell.number_format = '@' if c == 1 else '#,##0.00'
last = max(5, 4 + len(summary))
for col in 'BCDE':
    ws_sum[f'{col}2'] = f'=SUM({col}5:{col}{last})'

ws_exc = clear_data('Exceptions')
for excel_row, (docid, pid) in enumerate(exceptions, start=5):
    for c, v in enumerate(['Invoices', docid, pid, 'Overdue'], 1):
        cell = ws_exc.cell(excel_row, c, v)
        cell.number_format = '@'

wb.save(out_path)
print('Saved:', out_path)

# Reopen and verify saved output.
wb2 = load_workbook(out_path, data_only=False)
errors = []
def expect(name, actual, expected):
    ok = actual == expected
    print(f'{name}: {"PASS" if ok else "FAIL"} actual={actual!r} expected={expected!r}')
    if not ok:
        errors.append(name)

expect('Invoices rows', wb2['Invoices'].max_row - 4, 13)
expect('Orders blank rows', wb2['Orders'].max_row - 4, 0)
expect('Expenses blank rows', wb2['Expenses'].max_row - 4, 0)
expect('Quotes blank rows', wb2['Quotes'].max_row - 4, 0)
expect('Invoices headers', [wb2['Invoices'].cell(4,c).value for c in range(1,12)], ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'])
expect('Summary headers', [wb2['Summary'].cell(4,c).value for c in range(1,6)], ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'])
expect('Exceptions headers', [wb2['Exceptions'].cell(4,c).value for c in range(1,5)], ['Source','DocumentID','ProjectID','Reason'])

# Verify value types, dates, formulas and amounts.
rows = []
for r in range(5, wb2['Invoices'].max_row+1):
    rows.append([wb2['Invoices'].cell(r,c).value for c in range(1,12)])
expect('Invoice IDs remain text', all(isinstance(v[0], str) and v[0].startswith('009') for v in rows), True)
expect('Order IDs remain text', all(isinstance(v[1], str) and v[1].startswith('004') for v in rows), True)
expect('All invoice dates in 2026-09', all('2026-09-01' <= v[4] <= '2026-09-30' for v in rows), True)
expect('Amount/Paid/Balance/BalanceTWD numeric', all(isinstance(v[c], (int,float)) for v in rows for c in [7,8,9,10]), True)

saved_ex = []
for r in range(5, wb2['Exceptions'].max_row+1):
    saved_ex.append([wb2['Exceptions'].cell(r,c).value for c in range(1,5)])
expected_ex = [['Invoices', docid, pid, 'Overdue'] for docid, pid in exceptions]
expect('Overdue rows only DueDate < 2026-09-30 and Balance > 0', saved_ex, expected_ex)
due_today_positive = [v[0] for v in rows if v[5] == '2026-09-30' and Decimal(str(v[9])) > 0]
expect('Due 2026-09-30 not marked overdue', all(doc not in [x[1] for x in saved_ex] for doc in due_today_positive), True)

saved_summary = {}
for r in range(5, wb2['Summary'].max_row+1):
    saved_summary[wb2['Summary'].cell(r,1).value] = Decimal(str(wb2['Summary'].cell(r,3).value)).quantize(Decimal('0.01'))
expected_summary = {k: v.quantize(Decimal('0.01')) for k,v in summary.items()}
expect('Summary project IDs', sorted(saved_summary), sorted(expected_summary))
expect('Summary InvoiceBalanceTWD amounts', saved_summary, expected_summary)
expect('Summary other source columns zero', all(wb2['Summary'].cell(r,c).value == 0 for r in range(5, wb2['Summary'].max_row+1) for c in [2,4,5]), True)
expect('Summary total formulas', [wb2['Summary'][f'{c}2'].value for c in 'BCDE'], [f'=SUM({c}5:{c}{last})' for c in 'BCDE'])

if errors:
    raise RuntimeError('Verification failed: ' + ', '.join(errors))
print('All verification checks passed.')