from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv preview:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r):
        print(row)
print('\nsap_manual.md preview:')
text=(ws/'input'/'sap_manual.md').read_text(encoding='utf-8-sig')
print(text[:2000])
print('\nworkbook sheets and headers:')
wb=load_workbook(ws/'input'/'project_template.xlsx')
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('SHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,5), values_only=True):
        print(r)
    print('---')