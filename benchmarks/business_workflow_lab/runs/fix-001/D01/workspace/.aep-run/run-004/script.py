from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv, shutil

workspace=Path('.')
input_dir=workspace/'input'
output_dir=workspace/'output'
output_dir.mkdir(exist_ok=True)
export_path=output_dir/'E001_september_expenses_export.xlsx'
template_path=input_dir/'expense_template.xlsx'
fx_path=input_dir/'fx_rates.csv'
out_path=output_dir/'E001_september_expenses.xlsx'

print('Inputs:')
for p in [template_path, fx_path, export_path]:
    print(p, p.exists(), p.stat().st_size if p.exists() else None)

exp_wb=load_workbook(export_path, data_only=True)
exp_ws=exp_wb.active
headers=[c.value for c in exp_ws[1]]
print('Export headers:', headers)
rows=[]
for r in exp_ws.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in r):
        continue
    rows.append({headers[i]: r[i] for i in range(len(headers))})
print('Export row count:', len(rows))
print('First rows:', rows[:3])

rates={}
with fx_path.open(newline='', encoding='utf-8-sig') as fh:
    reader=csv.reader(fh)
    next(reader)
    for row in reader:
        if len(row)>=2 and row[0]:
            rates[row[0]]=Decimal(str(row[1]))
print('FX rates:', rates)

shutil.copyfile(template_path, out_path)
wb=load_workbook(out_path)
report=wb['Report']; details=wb['Details']; summary=wb['Summary']; exceptions=wb['Exceptions']
for ws in [details, summary, exceptions]:
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row-4)

employee_id='E001'; from_date='2026-09-01'; to_date='2026-09-30'
now=datetime.now().astimezone().isoformat(timespec='seconds')
report['B3']=from_date; report['D3']=to_date
report['B4']=f'EmployeeID={employee_id}; ExpenseDate between {from_date} and {to_date}; Status=ALL; Project=ALL'
report['B5']=now
report['B7']='Expense Viewer (synthetic lab) export'; report['D7']=now
report['B9']='All filtered results were obtained using the application Export Excel function. NeedsReview=true when MissingReceipt is true, Status is Draft, or Status is Rejected.'

def q2(d):
    if not isinstance(d, Decimal): d=Decimal(str(d))
    return d.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def bool_text(v):
    if isinstance(v, bool): return 'true' if v else 'false'
    if v is None: return 'false'
    return 'true' if str(v).strip().lower() in ('true','1','yes','y') else 'false'

summary_map={}; exceptions_rows=[]
for idx,d in enumerate(rows, start=5):
    currency=str(d['Currency']); amount=q2(d['Amount']); rate=rates[currency]; amount_twd=q2(amount*rate)
    missing=bool_text(d.get('MissingReceipt')); status=str(d['Status'])
    needs_review='true' if (missing=='true' or status in ('Draft','Rejected')) else 'false'
    values=[str(d['ExpenseID']), str(d['EmployeeID']), str(d['EmployeeName']), str(d['ProjectID']), str(d['ExpenseDate']), str(d['Category']), currency, float(amount), float(amount_twd), status, '' if d.get('ReceiptID') is None else str(d.get('ReceiptID')), '' if d.get('Description') is None else str(d.get('Description')), missing, needs_review]
    for col,val in enumerate(values, start=1):
        cell=details.cell(row=idx, column=col, value=val)
        if col in (1,2,4,11): cell.number_format='@'
        elif col==5: cell.number_format='yyyy-mm-dd'
        elif col in (8,9): cell.number_format='0.00'
    key=(status,currency)
    summary_map.setdefault(key, {'amount':Decimal('0'), 'amount_twd':Decimal('0'), 'count':0})
    summary_map[key]['amount']+=amount; summary_map[key]['amount_twd']+=amount_twd; summary_map[key]['count']+=1
    reasons=[]
    if missing=='true': reasons.append('MissingReceipt=true')
    if status=='Draft': reasons.append('Status=Draft')
    if status=='Rejected': reasons.append('Status=Rejected')
    if reasons: exceptions_rows.append([str(d['ExpenseID']), str(d['EmployeeID']), str(d['ExpenseDate']), '; '.join(reasons)])

for r_idx, ((status,currency), agg) in enumerate(sorted(summary_map.items()), start=5):
    vals=['Status', status, currency, float(q2(agg['amount'])), float(q2(agg['amount_twd'])), agg['count']]
    for col,val in enumerate(vals, start=1):
        c=summary.cell(row=r_idx, column=col, value=val)
        if col in (4,5): c.number_format='0.00'
last_detail_row=4+len(rows) if rows else 5
summary['B2']=f'=SUM(Details!I5:I{last_detail_row})'
summary['D2']='TWD Total'
summary['E2']=float(q2(sum(Decimal(str(details.cell(row=i,column=9).value)) for i in range(5,5+len(rows))))) if rows else 0.0
summary['E2'].number_format='0.00'

for r_idx, vals in enumerate(exceptions_rows, start=5):
    for col,val in enumerate(vals, start=1):
        c=exceptions.cell(row=r_idx, column=col, value=val)
        if col in (1,2): c.number_format='@'
        elif col==3: c.number_format='yyyy-mm-dd'

for ws in [details, summary, exceptions, report]:
    for col_idx in range(1, ws.max_column+1):
        max_len=0
        for row_idx in range(1, ws.max_row+1):
            v=ws.cell(row=row_idx, column=col_idx).value
            max_len=max(max_len, len(str(v)) if v is not None else 0)
        ws.column_dimensions[get_column_letter(col_idx)].width=min(max(max_len+2, 10), 45)

wb.save(out_path)
print('Wrote', out_path)

# Verification re-open from saved outputs
check_wb=load_workbook(out_path, data_only=False)
assert set(['Report','Details','Summary','Exceptions','Instructions']).issubset(check_wb.sheetnames)
cd=check_wb['Details']; cs=check_wb['Summary']; ce=check_wb['Exceptions']; cr=check_wb['Report']
expected=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
actual=[cd.cell(4,i).value for i in range(1,15)]
assert actual==expected, actual
assert cd.max_row==4+len(rows), (cd.max_row,len(rows))
for i in range(5, cd.max_row+1):
    assert cd.cell(i,2).value==employee_id
    assert from_date <= str(cd.cell(i,5).value) <= to_date
    assert isinstance(cd.cell(i,1).value, str)
    assert isinstance(cd.cell(i,2).value, str)
    assert isinstance(cd.cell(i,4).value, str)
    assert isinstance(cd.cell(i,8).value, (int,float))
    assert isinstance(cd.cell(i,9).value, (int,float))
    assert str(cd.cell(i,13).value) in ('true','false')
    assert str(cd.cell(i,14).value) in ('true','false')
assert [cs.cell(4,i).value for i in range(1,7)]==['Group','SubGroup','Currency','Amount','AmountTWD','Count']
assert cs.max_row-4==len(summary_map), (cs.max_row, len(summary_map))
calc_total=q2(sum(Decimal(str(cd.cell(i,9).value)) for i in range(5, cd.max_row+1)))
assert cs['B2'].value==f'=SUM(Details!I5:I{cd.max_row})'
assert q2(cs['E2'].value)==calc_total
assert ce.max_row-4==len(exceptions_rows)
print('CHECK PASSED')
print('Output file:', out_path)
print('Details rows:', len(rows))
print('Summary groups:')
for key, agg in sorted(summary_map.items()):
    print(key, 'amount', q2(agg['amount']), 'amount_twd', q2(agg['amount_twd']), 'count', agg['count'])
print('Exceptions rows:', len(exceptions_rows))
print('TWD total:', calc_total)
