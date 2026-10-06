from pathlib import Path
from openpyxl import load_workbook
import csv
for p in ['output/E001_september_expenses_export.xlsx']:
    print(p, Path(p).exists(), Path(p).stat().st_size if Path(p).exists() else None)
    if Path(p).exists():
        wb=load_workbook(p, data_only=False)
        print('sheets', wb.sheetnames)
        for sh in wb.sheetnames:
            ws=wb[sh]
            print('SHEET',sh,ws.max_row,ws.max_column)
            for row in ws.iter_rows(min_row=1,max_row=min(ws.max_row,8), values_only=True):
                print(row)
print('Instructions:')
wb=load_workbook('input/expense_template.xlsx', data_only=False)
ws=wb['Instructions']
for row in ws.iter_rows(values_only=True):
    print(row)
