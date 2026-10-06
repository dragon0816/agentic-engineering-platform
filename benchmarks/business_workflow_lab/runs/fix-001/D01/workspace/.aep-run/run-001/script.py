import os, json, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path('.')
print('CWD', ws.resolve())
for base in ['input','output']:
    p=Path(base)
    print(base, 'exists', p.exists())
    if p.exists():
        for f in sorted(p.rglob('*')):
            print(' -', f, 'size', f.stat().st_size if f.is_file() else 'DIR')
manual=Path('input/expense_manual.md')
if manual.exists():
    print('\nMANUAL HEAD:')
    print(manual.read_text(encoding='utf-8')[:3000])
fx=Path('input/fx_rates.csv')
if fx.exists():
    print('\nFX SAMPLE:')
    with fx.open(newline='', encoding='utf-8-sig') as fh:
        r=csv.reader(fh)
        for i,row in zip(range(10), r): print(row)
tpl=Path('input/expense_template.xlsx')
if tpl.exists():
    wb=load_workbook(tpl, data_only=False)
    print('\nTEMPLATE SHEETS', wb.sheetnames)
    for sh in wb.sheetnames:
        ws=wb[sh]
        print('SHEET', sh, 'max', ws.max_row, ws.max_column)
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,8), values_only=True):
            print(row)
