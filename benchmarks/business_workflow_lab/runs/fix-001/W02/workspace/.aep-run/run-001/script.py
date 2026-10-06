from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:')
for p in (ws/'input').iterdir():
    print('-', p.name, p.stat().st_size)
# inspect xlsx
xlsx=ws/'input'/'sales_template.xlsx'
wb=load_workbook(xlsx, data_only=False)
print('workbook sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(r)
# inspect csv
csvp=ws/'input'/'fx_rates.csv'
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(10), reader): print('fx', i, row)
# inspect manual head
md=ws/'input'/'sap_manual.md'
text=md.read_text(encoding='utf-8')
print('manual first 3000 chars:\n', text[:3000])