import os, csv
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('CWD', ws)
print('Input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size)
# Inspect CSV
csvp=ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with open(csvp, newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(10), rdr): print(row)
# Inspect md first lines
mdp=ws/'input'/'expense_manual.md'
print('\nmanual first 80 lines:')
with open(mdp, encoding='utf-8') as f:
    for i,line in zip(range(80), f): print(f'{i+1}: {line.rstrip()}')
# Inspect workbook sheets/headers/dimensions and sample values
wbp=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wbp, data_only=False)
print('\nWorkbook sheets:')
for sh in wb.worksheets:
    print(sh.title, sh.max_row, sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print('  ', r)
    print(' merged:', list(sh.merged_cells.ranges))
