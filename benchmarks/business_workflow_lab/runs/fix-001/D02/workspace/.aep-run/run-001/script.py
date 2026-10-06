from pathlib import Path
import os
base=Path.cwd()
print('cwd', base)
for p in sorted((base/'input').iterdir()):
    print('INPUT', p.name, p.stat().st_size)
    if p.suffix.lower() in ['.csv','.md','.txt']:
        print('---', p.name, 'head ---')
        txt=p.read_text(encoding='utf-8-sig', errors='replace')
        print('\n'.join(txt.splitlines()[:30]))
    elif p.suffix.lower() in ['.xlsx']:
        import openpyxl
        wb=openpyxl.load_workbook(p, data_only=False)
        print('sheets', wb.sheetnames)
        for ws in wb.worksheets:
            print('sheet', ws.title, 'max', ws.max_row, ws.max_column)
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,5), values_only=True):
                print(row)
            print('---')