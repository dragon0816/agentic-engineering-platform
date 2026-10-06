from pathlib import Path
from openpyxl import load_workbook
import csv
ws = Path.cwd()
print('CWD', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Inspect workbook
xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx)
print('Workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('Sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        print(' ', r)

# Inspect CSV
csvp = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i,row in zip(range(10), reader):
        print(i, row)

# Inspect manual
md = ws/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md head:')
print(md.read_text(encoding='utf-8')[:2000])