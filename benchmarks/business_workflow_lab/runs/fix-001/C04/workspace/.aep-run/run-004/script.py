from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
out=ws/'output'/'sent_quotes_2026-09.xlsx'
wb=load_workbook(out)
q=wb['Quotes']
print('Quotes rows:')
for r in range(4, q.max_row+1):
    print(r, [q.cell(r,c).value for c in range(1,12)], [q.cell(r,c).number_format for c in range(1,4)])
l=wb['Lines']
print('Lines rows:')
for r in range(4, l.max_row+1):
    print(r, [l.cell(r,c).value for c in range(1,7)], [l.cell(r,c).number_format for c in range(1,3)])
s=wb['Summary']
print('Summary B2', s['B2'].value)
for r in range(4, s.max_row+1):
    print(r, [s.cell(r,c).value for c in range(1,7)])
