import csv, json, subprocess, sys, shutil, os, time
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from datetime import datetime
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUTFILE = OUTPUT / 'sent_quotes_2026-09.xlsx'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
NOW = datetime.now().astimezone().replace(microsecond=0).isoformat()

Q2 = Decimal('0.01')
def money(d):
    if not isinstance(d, Decimal):
        d = Decimal(str(d))
    return d.quantize(Q2, rounding=ROUND_HALF_UP)

def parse_json_stdout(stdout, what):
    try:
        return json.loads(stdout)
    except json.JSONDecodeError as e:
        raise RuntimeError(f'{what}: stdout was not valid JSON on successful exit: {e}; stdout={stdout[:500]!r}')

def run_quote_cli(args, what):
    cmd = [str(CLI)] + args
    attempts = 0
    max_retries = 2
    while True:
        attempts += 1
        proc = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace')
        print(f'CLI {what} attempt {attempts}: exit={proc.returncode}')
        if proc.stderr.strip():
            print(f'CLI {what} stderr: {proc.stderr.strip()}')
        if proc.returncode == 0:
            # Only now treat stdout as data.
            return parse_json_stdout(proc.stdout, what)
        if proc.returncode == 4 and attempts <= max_retries + 1:
            if attempts <= max_retries:
                print(f'CLI {what}: temporary service error; retrying ({attempts}/{max_retries})')
                time.sleep(0.2)
                continue
        if proc.returncode in (2, 3):
            raise RuntimeError(f'CLI {what} failed with exit code {proc.returncode}; not retrying. stderr: {proc.stderr.strip()}')
        if proc.returncode == 4:
            raise RuntimeError(f'CLI {what} failed with exit code 4 after {attempts-1} retries. stderr: {proc.stderr.strip()}')
        raise RuntimeError(f'CLI {what} failed with unexpected exit code {proc.returncode}. stderr: {proc.stderr.strip()}')

# Load FX rates; notes contain unquoted commas in the synthetic file, so use first two columns only.
fx = {}
with open(FX_FILE, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header:', header)
    for row in reader:
        if len(row) >= 2 and row[0].strip():
            fx[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', {k: str(v) for k,v in fx.items()})

# Query quote headers and details.
listing = run_quote_cli(['quotes','list','--from',FROM_DATE,'--to',TO_DATE,'--status',STATUS,'--format','json'], 'quotes list September 2026 Sent')
if listing.get('schema_version') != 1 or 'items' not in listing:
    raise RuntimeError(f'Unexpected list schema: keys={list(listing.keys())}')
headers = listing['items']
print(f'List returned {len(headers)} headers; reported count={listing.get("count")}')

# stable sort for deterministic output
headers = sorted(headers, key=lambda h: h.get('QuoteID',''))
details_by_id = {}
lines = []
for h in headers:
    qid = str(h['QuoteID'])
    detail = run_quote_cli(['quotes','get',qid,'--format','json'], f'quotes get {qid}')
    if detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
        raise RuntimeError(f'Unexpected get schema for {qid}: keys={list(detail.keys())}')
    details_by_id[qid] = detail
    for li in detail['line_items']:
        row = dict(li)
        row['QuoteID'] = qid
        lines.append(row)
print(f'Fetched {len(details_by_id)} quote detail docs and {len(lines)} line items')

# Use detail headers as source of truth, preserving list order.
quote_rows = []
for h in headers:
    qid = str(h['QuoteID'])
    dh = details_by_id[qid]['header']
    # Validate filter.
    if not (FROM_DATE <= str(dh['QuoteDate']) <= TO_DATE and dh['Status'] == STATUS):
        raise RuntimeError(f'Quote {qid} does not match requested filters: {dh}')
    currency = dh['Currency']
    if currency not in fx:
        raise RuntimeError(f'Missing FX rate for {currency}')
    net = money(dh['NetAmount'])
    net_twd = money(net * fx[currency])
    quote_rows.append({
        'QuoteID': str(dh['QuoteID']),
        'CustomerID': str(dh['CustomerID']),
        'ProjectID': str(dh['ProjectID']),
        'QuoteDate': str(dh['QuoteDate']),
        'ValidUntil': str(dh['ValidUntil']),
        'Status': str(dh['Status']),
        'Currency': str(currency),
        'Subtotal': money(dh['Subtotal']),
        'DiscountPercent': Decimal(str(dh['DiscountPercent'])),
        'NetAmount': net,
        'NetAmountTWD': net_twd,
    })

line_rows = []
for li in sorted(lines, key=lambda r: (str(r['QuoteID']), str(r.get('ProductID','')), str(r.get('ProductName','')))):
    qty = Decimal(str(li['Quantity']))
    unit = money(li['UnitPrice'])
    amount = money(li['LineAmount'])
    expected = money(qty * unit)
    if amount != expected:
        raise RuntimeError(f'LineAmount mismatch for {li}: expected {expected}, got {amount}')
    line_rows.append({
        'QuoteID': str(li['QuoteID']),
        'ProductID': str(li['ProductID']),
        'ProductName': str(li['ProductName']),
        'Quantity': int(qty) if qty == qty.to_integral_value() else float(qty),
        'UnitPrice': unit,
        'LineAmount': amount,
    })

# Copy template to output and fill.
shutil.copyfile(TEMPLATE, OUTFILE)
wb = load_workbook(OUTFILE)

# Helpers for keeping formatting from template row 5 or row 4.
def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

def clear_data(ws, start_row, max_col):
    if ws.max_row >= start_row:
        for row in range(start_row, ws.max_row+1):
            for col in range(1, max_col+1):
                ws.cell(row, col).value = None

# Report
rpt = wb['Report']
rpt['B3'] = FROM_DATE
rpt['D3'] = TO_DATE
rpt['B4'] = f'Status={STATUS}; QuoteDate {FROM_DATE} to {TO_DATE} inclusive'
rpt['B5'] = NOW
rpt['B7'] = 'quote_cli_mock; fx_rates.csv'
rpt['D7'] = NOW
rpt['B9'] = f'CLI stdout used only after exit code 0 and valid JSON parse; exit code 4 retried at most twice. Quotes={len(quote_rows)}, Lines={len(line_rows)}.'

# Quotes sheet
qws = wb['Quotes']
q_headers = [qws.cell(4,c).value for c in range(1,12)]
expected_q_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
if q_headers != expected_q_headers:
    raise RuntimeError(f'Unexpected Quotes headers: {q_headers}')
clear_data(qws, 5, 11)
for i, row in enumerate(quote_rows, start=5):
    copy_row_style(qws, 4, i, 11)
    values = [row[h] for h in expected_q_headers]
    for c, v in enumerate(values, start=1):
        cell = qws.cell(i,c)
        if c in (1,2,3):
            cell.value = str(v)
            cell.number_format = '@'
        elif c in (8,10,11):
            cell.value = float(v)
            cell.number_format = '0.00'
        elif c == 9:
            cell.value = float(v)
        else:
            cell.value = v

# Lines sheet
lws = wb['Lines']
l_headers = [lws.cell(4,c).value for c in range(1,7)]
expected_l_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if l_headers != expected_l_headers:
    raise RuntimeError(f'Unexpected Lines headers: {l_headers}')
clear_data(lws, 5, 6)
for i, row in enumerate(line_rows, start=5):
    copy_row_style(lws, 4, i, 6)
    values = [row[h] for h in expected_l_headers]
    for c, v in enumerate(values, start=1):
        cell = lws.cell(i,c)
        if c in (1,2):
            cell.value = str(v)
            cell.number_format = '@'
        elif c in (5,6):
            cell.value = float(v)
            cell.number_format = '0.00'
        else:
            cell.value = v

# Summary: group by Status and Currency, so no cross-currency mixing in Amount.
sws = wb['Summary']
s_headers = [sws.cell(4,c).value for c in range(1,7)]
expected_s_headers = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
if s_headers != expected_s_headers:
    raise RuntimeError(f'Unexpected Summary headers: {s_headers}')
clear_data(sws, 5, 6)
summary = {}
for row in quote_rows:
    key = (row['Status'], row['Currency'])
    if key not in summary:
        summary[key] = {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0}
    summary[key]['amount'] += row['NetAmount']
    summary[key]['twd'] += row['NetAmountTWD']
    summary[key]['count'] += 1
summary_rows = []
for (status, currency), agg in sorted(summary.items()):
    summary_rows.append({
        'Group': 'Status',
        'SubGroup': status,
        'Currency': currency,
        'Amount': money(agg['amount']),
        'AmountTWD': money(agg['twd']),
        'Count': agg['count'],
    })
for i, row in enumerate(summary_rows, start=5):
    copy_row_style(sws, 4, i, 6)
    vals = [row[h] for h in expected_s_headers]
    for c, v in enumerate(vals, start=1):
        cell = sws.cell(i,c)
        if c in (4,5):
            cell.value = float(v)
            cell.number_format = '0.00'
        else:
            cell.value = v
last_q_row = 4 + max(1, len(quote_rows))
sws['B2'] = f'=SUM(Quotes!K5:K{last_q_row})'

wb.save(OUTFILE)
print('Wrote', OUTFILE)

# Reopen and verify saved output.
wb2 = load_workbook(OUTFILE, data_only=False)
errors = []
# sheets preserved
if wb2.sheetnames != ['Report','Quotes','Lines','Summary','Instructions']:
    errors.append(f'sheet names changed: {wb2.sheetnames}')
# report cells
r = wb2['Report']
for cell, expected in [('B3', FROM_DATE), ('D3', TO_DATE)]:
    if r[cell].value != expected:
        errors.append(f'Report {cell} expected {expected}, got {r[cell].value}')
if STATUS not in str(r['B4'].value):
    errors.append('Report B4 does not include status condition')
# quotes
q = wb2['Quotes']
if [q.cell(4,c).value for c in range(1,12)] != expected_q_headers:
    errors.append('Quotes headers not preserved')
q_count = len(quote_rows)
for idx, source in enumerate(quote_rows, start=5):
    vals = [q.cell(idx,c).value for c in range(1,12)]
    if vals[0] != source['QuoteID'] or q.cell(idx,1).number_format != '@': errors.append(f'QuoteID text check failed row {idx}')
    if vals[1] != source['CustomerID'] or q.cell(idx,2).number_format != '@': errors.append(f'CustomerID text check failed row {idx}')
    if vals[2] != source['ProjectID'] or q.cell(idx,3).number_format != '@': errors.append(f'ProjectID text check failed row {idx}')
    if vals[3] != source['QuoteDate'] or vals[4] != source['ValidUntil']: errors.append(f'Date check failed row {idx}')
    if vals[5] != STATUS: errors.append(f'Status check failed row {idx}')
    for col in (8,9,10,11):
        if not isinstance(q.cell(idx,col).value, (int,float)):
            errors.append(f'Quotes numeric check failed row {idx} col {col}: {q.cell(idx,col).value!r}')
    # Recalculate TWD from saved values using Decimal(str(float)) may carry .0 only; compare rounded.
    cur = vals[6]
    saved_twd = money(str(vals[10]))
    expected_twd = money(source['NetAmount'] * fx[cur])
    if saved_twd != expected_twd:
        errors.append(f'NetAmountTWD mismatch row {idx}: {saved_twd} vs {expected_twd}')
# no extra quote rows
if q.cell(5+q_count,1).value is not None:
    errors.append('Quotes has extra non-empty row after expected data')
# lines
l = wb2['Lines']
if [l.cell(4,c).value for c in range(1,7)] != expected_l_headers:
    errors.append('Lines headers not preserved')
for idx, source in enumerate(line_rows, start=5):
    if l.cell(idx,1).value != source['QuoteID'] or l.cell(idx,1).number_format != '@': errors.append(f'Line QuoteID text check failed row {idx}')
    if l.cell(idx,2).value != source['ProductID'] or l.cell(idx,2).number_format != '@': errors.append(f'ProductID text check failed row {idx}')
    for col in (4,5,6):
        if not isinstance(l.cell(idx,col).value, (int,float)):
            errors.append(f'Lines numeric check failed row {idx} col {col}')
if l.cell(5+len(line_rows),1).value is not None:
    errors.append('Lines has extra non-empty row after expected data')
# summary
s = wb2['Summary']
if s['B2'].value != f'=SUM(Quotes!K5:K{last_q_row})':
    errors.append(f'Summary B2 formula wrong: {s["B2"].value}')
if [s.cell(4,c).value for c in range(1,7)] != expected_s_headers:
    errors.append('Summary headers not preserved')
for idx in range(5, 5+len(summary_rows)):
    if not isinstance(s.cell(idx,4).value, (int,float)) or not isinstance(s.cell(idx,5).value, (int,float)) or not isinstance(s.cell(idx,6).value, int):
        errors.append(f'Summary numeric/count check failed row {idx}')
# total integrity
saved_total = sum(Decimal(str(q.cell(i,11).value)) for i in range(5,5+q_count))
summary_total = sum(Decimal(str(s.cell(i,5).value)) for i in range(5,5+len(summary_rows)))
if money(saved_total) != money(summary_total):
    errors.append(f'Summary total {summary_total} does not match quote TWD total {saved_total}')
# lines join to quotes
quote_ids = {row['QuoteID'] for row in quote_rows}
for idx in range(5,5+len(line_rows)):
    if l.cell(idx,1).value not in quote_ids:
        errors.append(f'Line row {idx} QuoteID has no Quotes match')

if errors:
    print('VERIFY FAILED')
    for e in errors:
        print('-', e)
    raise SystemExit(1)
print('VERIFY OK')
print(f'Output file: {OUTFILE}')
print(f'Checked: sheets/header preservation, Report cells, quote row count={q_count}, line row count={len(line_rows)}, ID text formats, numeric amount types/formats, FX TWD calculations, Summary formula/ranges/totals, line QuoteID joins.')
