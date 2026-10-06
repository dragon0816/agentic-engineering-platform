from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('cwd=', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
manual = ws/'input'/'quote_cli_manual.md'
print('\n--- quote_cli_manual.md head ---')
print(manual.read_text(encoding='utf-8')[:4000])
print('\n--- fx_rates.csv sample ---')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(10), rdr):
        print(i, row)
print('\n--- workbook sheets and used ranges ---')
wb = load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('sheet', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for r in wsht.iter_rows(min_row=1, max_row=min(8, wsht.max_row), values_only=True):
        print(' ', r)
