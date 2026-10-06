from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv
from openpyxl import load_workbook

ws = Path.cwd()
out = ws/'output'/'sent_quotes_2026-09.xlsx'
fx_file = ws/'input'/'fx_rates.csv'
FROM_DATE='2026-09-01'
TO_DATE='2026-09-30'
STATUS='Sent'
Q2=Decimal('0.01')
def money(x): return Decimal(str(x)).quantize(Q2, rounding=ROUND_HALF_UP)

# Load fx rates (only first two columns are needed; notes contain unquoted commas).
fx={}
with open(fx_file, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    next(reader)
    for row in reader:
        if len(row)>=2 and row[0].strip():
            fx[row[0].strip()]=Decimal(row[1].strip())

wb = load_workbook(out)
q = wb['Quotes']
missing_project=[]
quote_rows=[]
row=5
while q.cell(row,1).value is not None:
    rec = {
        'QuoteID': q.cell(row,1).value,
        'CustomerID': q.cell(row,2).value,
        'ProjectID': q.cell(row,3).value,
        'QuoteDate': q.cell(row,4).value,
        'ValidUntil': q.cell(row,5).value,
        'Status': q.cell(row,6).value,
        'Currency': q.cell(row,7).value,
        'Subtotal': q.cell(row,8).value,
        'DiscountPercent': q.cell(row,9).value,
        'NetAmount': q.cell(row,10).value,
        'NetAmountTWD': q.cell(row,11).value,
    }
    quote_rows.append(rec)
    if rec['ProjectID'] in (None, ''):
        missing_project.append(rec['QuoteID'])
    row += 1
l = wb['Lines']
line_count=0
row=5
while l.cell(row,1).value is not None:
    line_count += 1
    row += 1
# Append missing-data note if needed.
if missing_project:
    note = wb['Report']['B9'].value or ''
    add = 'Missing ProjectID for ' + ', '.join(missing_project) + '; cells left blank.'
    if add not in note:
        wb['Report']['B9'] = (note + ' ' + add).strip()
wb.save(out)
print('Updated missing-data note if needed:', missing_project)

# Reopen and verify saved output from disk.
wb = load_workbook(out, data_only=False)
errors=[]
expected_sheets=['Report','Quotes','Lines','Summary','Instructions']
if wb.sheetnames != expected_sheets:
    errors.append(f'sheet names expected {expected_sheets}, got {wb.sheetnames}')
r=wb['Report']
if r['B3'].value != FROM_DATE: errors.append(f'Report B3 wrong: {r["B3"].value}')
if r['D3'].value != TO_DATE: errors.append(f'Report D3 wrong: {r["D3"].value}')
if STATUS not in str(r['B4'].value): errors.append('Report B4 missing Sent condition')
if 'quote_cli' not in str(r['B7'].value): errors.append('Report B7 missing quote_cli source')
if missing_project and not all(qid in str(r['B9'].value) for qid in missing_project):
    errors.append('Report B9 missing ProjectID exception note')

q=wb['Quotes']
qh=['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
if [q.cell(4,c).value for c in range(1,12)] != qh:
    errors.append('Quotes column names changed')
quote_count=0
seen_quotes=set()
row=5
while q.cell(row,1).value is not None:
    quote_count += 1
    qid=q.cell(row,1).value
    seen_quotes.add(qid)
    if not isinstance(qid, str) or q.cell(row,1).number_format != '@': errors.append(f'QuoteID not text row {row}')
    if not isinstance(q.cell(row,2).value, str) or q.cell(row,2).number_format != '@': errors.append(f'CustomerID not text row {row}')
    proj=q.cell(row,3).value
    if proj not in (None,'') and not isinstance(proj, str): errors.append(f'ProjectID not text/blank row {row}')
    if q.cell(row,3).number_format != '@': errors.append(f'ProjectID number_format not text row {row}')
    if not (FROM_DATE <= q.cell(row,4).value <= TO_DATE): errors.append(f'QuoteDate outside range row {row}')
    if q.cell(row,6).value != STATUS: errors.append(f'Status not Sent row {row}')
    cur=q.cell(row,7).value
    if cur not in fx: errors.append(f'Missing FX for saved currency {cur}')
    for col in (8,9,10,11):
        if not isinstance(q.cell(row,col).value, (int,float)):
            errors.append(f'Quotes numeric value expected row {row} col {col}')
    if cur in fx:
        expected_twd=money(money(q.cell(row,10).value)*fx[cur])
        if money(q.cell(row,11).value) != expected_twd:
            errors.append(f'NetAmountTWD mismatch row {row}: {q.cell(row,11).value} vs {expected_twd}')
    row += 1
if quote_count != 4:
    errors.append(f'Expected 4 quote rows from CLI result, found {quote_count}')
if q.cell(5+quote_count,1).value is not None:
    errors.append('Quotes has extra non-empty row after expected data')

l=wb['Lines']
lh=['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if [l.cell(4,c).value for c in range(1,7)] != lh:
    errors.append('Lines column names changed')
line_count=0
row=5
while l.cell(row,1).value is not None:
    line_count += 1
    qid=l.cell(row,1).value
    if qid not in seen_quotes: errors.append(f'Line row {row} QuoteID not in Quotes: {qid}')
    if not isinstance(qid, str) or l.cell(row,1).number_format != '@': errors.append(f'Line QuoteID not text row {row}')
    if not isinstance(l.cell(row,2).value, str) or l.cell(row,2).number_format != '@': errors.append(f'ProductID not text row {row}')
    for col in (4,5,6):
        if not isinstance(l.cell(row,col).value, (int,float)):
            errors.append(f'Lines numeric value expected row {row} col {col}')
    if money(Decimal(str(l.cell(row,4).value))*Decimal(str(l.cell(row,5).value))) != money(l.cell(row,6).value):
        errors.append(f'Line amount mismatch row {row}')
    row += 1
if line_count != 8:
    errors.append(f'Expected 8 line rows from CLI detail results, found {line_count}')
if l.cell(5+line_count,1).value is not None:
    errors.append('Lines has extra non-empty row after expected data')

s=wb['Summary']
sh=['Group','SubGroup','Currency','Amount','AmountTWD','Count']
if [s.cell(4,c).value for c in range(1,7)] != sh:
    errors.append('Summary column names changed')
expected_formula=f'=SUM(Quotes!K5:K{4+max(1,quote_count)})'
if s['B2'].value != expected_formula:
    errors.append(f'Summary B2 formula expected {expected_formula}, got {s["B2"].value}')
summary_total=Decimal('0')
summary_count=0
row=5
while s.cell(row,1).value is not None:
    if s.cell(row,1).value != 'Status' or s.cell(row,2).value != STATUS:
        errors.append(f'Summary grouping unexpected row {row}')
    if s.cell(row,3).value not in fx:
        errors.append(f'Summary currency missing fx row {row}')
    for col in (4,5):
        if not isinstance(s.cell(row,col).value, (int,float)):
            errors.append(f'Summary amount numeric expected row {row} col {col}')
    if not isinstance(s.cell(row,6).value, int): errors.append(f'Summary count integer expected row {row}')
    summary_total += Decimal(str(s.cell(row,5).value))
    summary_count += s.cell(row,6).value
    row += 1
quote_twd_total=sum(Decimal(str(q.cell(i,11).value)) for i in range(5,5+quote_count))
if money(summary_total) != money(quote_twd_total):
    errors.append(f'Summary total {summary_total} != Quotes TWD total {quote_twd_total}')
if summary_count != quote_count:
    errors.append(f'Summary count {summary_count} != quote count {quote_count}')

if errors:
    print('VERIFY FAILED')
    for e in errors: print('-', e)
    raise SystemExit(1)
print('VERIFY OK')
print(f'Output file: {out}')
print('Checked from saved workbook: file exists; sheet names unchanged; column names preserved; Report dates/condition/source/notes; 4 quote rows and 8 line rows; ID columns stored as text where present and ProjectID blanks noted; dates in September 2026; all statuses Sent; amount/count columns numeric; NetAmountTWD equals NetAmount × fx_rates.csv with ROUND_HALF_UP; Summary B2 extends to K8; Summary totals/counts match Quotes; every Lines QuoteID joins to Quotes.')
