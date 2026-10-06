from pathlib import Path
import csv, zipfile
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
print('\nquote_cli_manual.md preview:')
manual = ws/'input'/'quote_cli_manual.md'
print(manual.read_text(encoding='utf-8', errors='replace')[:4000])
print('\nfx_rates.csv headers/sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i,row in zip(range(6), r):
        print(i, row)
print('\nquote_template.xlsx workbook:')
wb = load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
print('sheets', wb.sheetnames)
for sh in wb.worksheets:
    print('sheet', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(10, sh.max_row), values_only=True):
        print(row)
    print('---')
