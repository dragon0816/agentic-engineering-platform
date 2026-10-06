from pathlib import Path
import subprocess, sys, json, csv, shutil, os
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
OUT_XLSX = OUTPUT / 'sent_quotes_2026-09.xlsx'
ERR_TXT = OUTPUT / 'quote_cli_error.txt'
RUN_LOG = OUTPUT / 'quote_cli_run_log.json'

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
TZ = timezone(timedelta(hours=8))

run_events = []

def now_iso():
    return datetime.now(TZ).isoformat(timespec='seconds')

def run_cli(args, parse_json=True, max_temp_retries=2):
    """Run quote CLI. Retry only exit code 4. For exit 2/3 raise without retry. Parse stdout only for rc=0."""
    cmd = [sys.executable, str(CLI)] + args
    attempts = 0
    while True:
        attempts += 1
        proc = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
        event = {
            'cmd': cmd,
            'attempt': attempts,
            'returncode': proc.returncode,
            'stderr': proc.stderr,
            'stdout_len': len(proc.stdout or ''),
        }
        run_events.append(event)
        if proc.returncode == 0:
            if not parse_json:
                return proc.stdout
            try:
                return json.loads(proc.stdout)
            except Exception as e:
                raise RuntimeError(f"CLI stdout was not parseable JSON for command {' '.join(args)}: {e}")
        if proc.returncode == 4:
            if attempts <= max_temp_retries:
                continue
            raise RuntimeError(f"CLI temporary service error after {attempts} attempts for command {' '.join(args)}. stderr: {proc.stderr.strip()}")
        if proc.returncode in (2, 3):
            raise RuntimeError(f"CLI returned exit code {proc.returncode}; no retry. stderr: {proc.stderr.strip()}")
        raise RuntimeError(f"CLI returned unexpected exit code {proc.returncode}; stderr: {proc.stderr.strip()}")

def dec(v):
    return Decimal(str(v))

def money(v):
    return dec(v).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def write_error(reason):
    ERR_TXT.write_text(reason + '\n\nRun log:\n' + json.dumps(run_events, ensure_ascii=False, indent=2), encoding='utf-8')
    RUN_LOG.write_text(json.dumps(run_events, ensure_ascii=False, indent=2), encoding='utf-8')
    print('ERROR:', reason)
    print('WROTE', ERR_TXT)
    print('WROTE', RUN_LOG)

print('Inputs:', [p.name for p in sorted(INPUT.iterdir())])
print('CLI exists:', CLI.exists(), CLI)
print('Query:', FROM_DATE, TO_DATE, STATUS)

try:
    list_data = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'])
    if not isinstance(list_data, dict) or list_data.get('schema_version') != 1 or not isinstance(list_data.get('items'), list):
        raise RuntimeError('quotes list returned rc=0 but JSON schema was not recognized')
    headers = list_data['items']
    quotes = []
    for h in headers:
        qid = h.get('QuoteID')
        if not qid:
            raise RuntimeError('quotes list item missing QuoteID')
        detail = run_cli(['quotes', 'get', qid, '--format', 'json'])
        if not isinstance(detail, dict) or detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
            raise RuntimeError(f'quotes get {qid} returned rc=0 but JSON schema was not recognized')
        quotes.append(detail)
except Exception as e:
    write_error(str(e))
    raise SystemExit(0)

# Load FX rates. Note field may contain commas in malformed sample, but first two columns are sufficient.
fx = {}
with open(FX, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    hdr = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx[row[0]] = dec(row[1])

shutil.copy2(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

# Metadata
report = wb['Report']
fetched_at = now_iso()
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = f'Status={STATUS}; QuoteDate {FROM_DATE}..{TO_DATE}'
report['B5'] = fetched_at
report['B7'] = 'quote_cli'
report['D7'] = fetched_at
report['B9'] = 'stdout was used only for exit code 0 with parseable JSON; exit code 4 retried up to 2 times; exit code 2/3 not retried.'

# Clear existing data rows if any
for sheet_name in ['Quotes', 'Lines', 'Summary']:
    ws = wb[sheet_name]
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

quotes_ws = wb['Quotes']
lines_ws = wb['Lines']
summary_ws = wb['Summary']

quote_rows = []
line_rows = []
summary = {}

for detail in quotes:
    h = detail['header']
    cur = h['Currency']
    if cur not in fx:
        raise RuntimeError(f'Missing FX rate for currency {cur}')
    net = money(h['NetAmount'])
    net_twd = money(net * fx[cur])
    row = [
        str(h['QuoteID']), str(h['CustomerID']), str(h['ProjectID']), str(h['QuoteDate']), str(h['ValidUntil']),
        str(h['Status']), str(cur), float(money(h['Subtotal'])), float(dec(h['DiscountPercent'])), float(net), float(net_twd)
    ]
    quote_rows.append(row)
    key = (str(h['Status']), str(cur))
    if key not in summary:
        summary[key] = {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0}
    summary[key]['amount'] += net
    summary[key]['twd'] += net_twd
    summary[key]['count'] += 1
    for li in detail['line_items']:
        line_rows.append([
            str(h['QuoteID']), str(li['ProductID']), str(li['ProductName']),
            int(li['Quantity']), float(money(li['UnitPrice'])), float(money(li['LineAmount']))
        ])

# Sort deterministic by QuoteID then ProductID
quote_rows.sort(key=lambda r: r[0])
line_rows.sort(key=lambda r: (r[0], r[1], r[2]))

for r_idx, row in enumerate(quote_rows, start=5):
    for c_idx, val in enumerate(row, start=1):
        cell = quotes_ws.cell(r_idx, c_idx, val)
        if c_idx in [1,2,3,4,5,6,7]:
            cell.number_format = '@'
        if c_idx in [8,10,11]:
            cell.number_format = '0.00'

for r_idx, row in enumerate(line_rows, start=5):
    for c_idx, val in enumerate(row, start=1):
        cell = lines_ws.cell(r_idx, c_idx, val)
        if c_idx in [1,2,3]:
            cell.number_format = '@'
        if c_idx in [5,6]:
            cell.number_format = '0.00'

for r_idx, ((status, cur), vals) in enumerate(sorted(summary.items()), start=5):
    row = ['Status', status, cur, float(money(vals['amount'])), float(money(vals['twd'])), vals['count']]
    for c_idx, val in enumerate(row, start=1):
        cell = summary_ws.cell(r_idx, c_idx, val)
        if c_idx in [4,5]:
            cell.number_format = '0.00'

last_quote_row = 5 + len(quote_rows) - 1 if quote_rows else 5
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'

wb.save(OUT_XLSX)
RUN_LOG.write_text(json.dumps(run_events, ensure_ascii=False, indent=2), encoding='utf-8')

# Reopen and verify requirements from saved output
wb2 = load_workbook(OUT_XLSX, data_only=False)
qws = wb2['Quotes']; lws = wb2['Lines']; sws = wb2['Summary']; rws = wb2['Report']
assert [qws.cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [lws.cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [sws.cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
assert rws['B3'].value == FROM_DATE and rws['D3'].value == TO_DATE
assert len(quote_rows) == (list_data.get('count') if list_data.get('count') is not None else len(headers))
for row in range(5, 5+len(quote_rows)):
    assert isinstance(qws.cell(row,1).value, str)
    assert qws.cell(row,6).value == STATUS
    assert FROM_DATE <= qws.cell(row,4).value <= TO_DATE
    assert isinstance(qws.cell(row,8).value, (int,float))
    assert isinstance(qws.cell(row,10).value, (int,float))
    assert isinstance(qws.cell(row,11).value, (int,float))
for row in range(5, 5+len(line_rows)):
    assert isinstance(lws.cell(row,1).value, str)
    assert isinstance(lws.cell(row,4).value, int)
    assert isinstance(lws.cell(row,5).value, (int,float))
    assert isinstance(lws.cell(row,6).value, (int,float))
assert sws['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row})'
print('SUCCESS wrote', OUT_XLSX)
print('WROTE', RUN_LOG)
print('quote_count', len(quote_rows), 'line_count', len(line_rows), 'summary_rows', len(summary))
print('checked: headers, dates/status filter, text IDs, numeric money fields, summary formula, saved workbook reopened')
