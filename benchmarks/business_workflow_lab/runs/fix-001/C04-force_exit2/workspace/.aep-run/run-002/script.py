from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
print('Full Instructions:')
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
