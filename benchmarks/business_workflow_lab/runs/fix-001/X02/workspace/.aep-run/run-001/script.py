import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path('.')
print('cwd', ws.resolve())
for p in sorted(Path('input').glob('*')):
    print('FILE', p, 'size', p.stat().st_size)
    if p.suffix.lower()=='.csv':
        with p.open(newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f)
            for i,row in zip(range(5), r): print('CSVROW', row)
    elif p.suffix.lower()=='.md':
        txt=p.read_text(encoding='utf-8')
        print('MD_HEAD', txt[:1200].replace('\n','\\n'))
    elif p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print('XLSX sheets', wb.sheetnames)
        for sh in wb.sheetnames:
            sheet=wb[sh]
            print('SHEET', sh, 'max', sheet.max_row, sheet.max_column)
            for row in sheet.iter_rows(min_row=1, max_row=min(10,sheet.max_row), values_only=True):
                print('ROW', row)
