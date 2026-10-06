from pathlib import Path
from openpyxl import load_workbook
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    print('\n====', name, '====')
    print(Path('input',name).read_text(encoding='utf-8'))
wb=load_workbook('input/project_template.xlsx', data_only=False)
sh=wb['Instructions']
print('\n==== Instructions all ====')
for row in sh.iter_rows(values_only=True): print(row)
