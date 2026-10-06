from pathlib import Path
from openpyxl import load_workbook
import csv, os
ws=Path.cwd()
print('cwd', ws)
print('input files:')
for p in (ws/'input').iterdir():
    print('-', p.name, p.stat().st_size)
print('\nmanual head:')
print((ws/'input'/'quote_cli_manual.md').read_text(encoding='utf-8')[:2000])
print('\nfx_rates sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r): print(row)
print('\nworkbook sheets/headers/sample:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(5, sh.max_row), values_only=True):
        print(row)
