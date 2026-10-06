from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
print('\nReport cells values/formulas:')
r=wb['Report']
for row in r.iter_rows(min_row=1,max_row=20,min_col=1,max_col=6,values_only=False):
    print([c.value for c in row])
