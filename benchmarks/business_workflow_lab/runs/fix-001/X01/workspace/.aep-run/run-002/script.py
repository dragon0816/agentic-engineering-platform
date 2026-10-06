from pathlib import Path
import csv, subprocess, json, os, sys
from openpyxl import load_workbook
ws=Path.cwd()
print('downloads:')
for p in sorted((ws/'output'/'downloads').glob('*')):
    print('-', p.name, p.stat().st_size)
    with open(p, newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(6), r): print(' ', row)
print('\ntemplate sheets/cells:')
wb=load_workbook(ws/'input'/'project_template.xlsx')
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('SHEET', sh, 'max', wsht.max_row, wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True): print(row)
print('\nquote cli test:')
cmd=[sys.executable, r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py','quotes','list','--from','2026-09-01','--to','2026-09-30','--status','Accepted','--format','json']
res=subprocess.run(cmd,capture_output=True,text=True)
print('return',res.returncode)
print('stdout',res.stdout[:2000])
print('stderr',res.stderr[:1000])
