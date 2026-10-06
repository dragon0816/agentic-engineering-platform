from pathlib import Path
import csv, json
from openpyxl import load_workbook
base=Path('.')
print('Workspace', base.resolve())
print('Input files:')
for p in sorted((base/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
print('\nfx_rates.csv sample:')
with open(base/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(5), r): print(row)
print('\nsap_manual.md head:')
text=(base/'input'/'sap_manual.md').read_text(encoding='utf-8')
print(text[:3000])
print('\nsales_template.xlsx sheets/headers:')
wb=load_workbook(base/'input'/'sales_template.xlsx')
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,5), values_only=True):
        print(row)
    print('---')