from pathlib import Path
from openpyxl import load_workbook
import subprocess, json, sys
ws=Path('.')
# print full Instructions rows
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
sh=wb['Instructions']
print('Instructions:')
for row in sh.iter_rows(values_only=True):
    if any(v is not None for v in row): print(row)
cli = r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd'
cmd=[cli,'quotes','list','--from','2026-09-01','--to','2026-09-30','--format','json']
print('\nRunning:', cmd)
r=subprocess.run(cmd,capture_output=True,text=True,encoding='utf-8')
print('exit', r.returncode, 'stderr', r.stderr[:500])
print('stdout head', r.stdout[:2000])
if r.returncode==0:
    data=json.loads(r.stdout)
    print('count',data.get('count'))
    for item in data.get('items',[])[:3]: print(item)
    if data.get('items'):
        qid=data['items'][0]['QuoteID']
        r2=subprocess.run([cli,'quotes','get',qid,'--format','json'],capture_output=True,text=True,encoding='utf-8')
        print('\nGET',qid,'exit',r2.returncode,'stderr',r2.stderr[:200])
        print(r2.stdout[:2000])