from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path('.')
print('Workspace:', ws.resolve())
for p in sorted((ws/'input').iterdir()):
    print('INPUT', p.name, p.stat().st_size)
print('\nquote_cli_manual.md head:')
print((ws/'input'/'quote_cli_manual.md').read_text(encoding='utf-8')[:2000])
print('\nfx_rates.csv sample:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(10), r): print(row)
print('\nWorkbook sheets/cells:')
wb=load_workbook(ws/'input'/'quote_template.xlsx')
print(wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('Sheet', s, 'max', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True): print(row)
    print('---')