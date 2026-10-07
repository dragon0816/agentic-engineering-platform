from pathlib import Path
from openpyxl import load_workbook
import csv, json

base = Path('.')
print('Workspace:', base.resolve())
for sub in ['input','output','output/downloads']:
    p = base / sub
    print(f'\n{sub}:')
    if p.exists():
        for f in sorted(p.iterdir()):
            print(' -', f.name, 'size=', f.stat().st_size)
    else:
        print(' MISSING')

xlsx = base / 'output' / 'project_C001_2026-09_report.xlsx'
if xlsx.exists():
    wb = load_workbook(xlsx, data_only=False)
    print('\nWorkbook sheets:', wb.sheetnames)
    for ws in wb.worksheets:
        print(ws.title, 'max_row=', ws.max_row, 'max_col=', ws.max_column)
        headers = [ws.cell(4,c).value for c in range(1, ws.max_column+1)] if ws.max_row >= 4 else []
        print(' row4 headers:', headers)
        if ws.max_row >= 5:
            sample = [ws.cell(5,c).value for c in range(1, min(ws.max_column,8)+1)]
            print(' row5 sample:', sample)
else:
    print('\nReport workbook MISSING:', xlsx)