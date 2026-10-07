from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size, 'bytes')
print('\nfx_rates.csv preview:')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i,row in zip(range(8), r):
        print(i, row)
print('\nWorkbook sheets and dimensions:')
wb = load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    s = wb[sh]
    print('SHEET', sh, s.max_row, s.max_column)
    for row in s.iter_rows(min_row=1, max_row=min(s.max_row,8), values_only=True):
        print(' ', row)
print('\nManual previews:')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p = ws/'input'/name
    print('\n---', name, '---')
    txt = p.read_text(encoding='utf-8')
    print(txt[:3000])