from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', sorted(p.name for p in (ws/'input').iterdir()))
# workbook structure
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_column', sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(' ', row)
# csv headers/sample
csv_path=ws/'input'/'fx_rates.csv'
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    rows=list(csv.reader(f))
print('fx rows sample:')
for r in rows[:8]: print(r)
# manual sample
md_path=ws/'input'/'expense_manual.md'
text=md_path.read_text(encoding='utf-8')
print('manual first 2000 chars:\n', text[:2000])