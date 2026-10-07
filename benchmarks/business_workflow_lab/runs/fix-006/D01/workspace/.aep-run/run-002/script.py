from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('output files currently:', sorted(p.name for p in (ws/'output').iterdir()))
export=ws/'output'/'E001_september_expenses_export.csv'
with open(export, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    rows=list(reader)
print('export row count incl header:', len(rows))
print('export header:', rows[0])
for r in rows[1:6]: print('sample', r)
# full instructions rows
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
print('Instructions:')
for row in sh.iter_rows(values_only=True):
    print(row)