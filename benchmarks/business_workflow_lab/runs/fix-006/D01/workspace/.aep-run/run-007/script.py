from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
import openpyxl

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
print('Workspace:', ws)
print('Input files:', sorted(p.name for p in input_dir.glob('*')))
print('Output files:', sorted(p.name for p in output_dir.glob('*')))

workbook_path = output_dir / 'E001_september_expenses.xlsx'
export_path = output_dir / 'expense_export_E001_2026-09-01_2026-09-30.csv'
legacy_export_path = output_dir / 'E001_september_expenses_export.csv'

for p in [workbook_path, export_path, legacy_export_path]:
    print(f'{p.name}: exists={p.exists()} size={p.stat().st_size if p.exists() else None}')
assert workbook_path.exists()
assert export_path.exists()
assert legacy_export_path.exists()

with export_path.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    rows = list(reader)
    headers = reader.fieldnames
print('Export headers:', headers)
print('Export row count:', len(rows))
print('Export sample row:', rows[0] if rows else None)

bad_emp = [r for r in rows if r.get('EmployeeID') != 'E001']
bad_date = [r for r in rows if not ('2026-09-01' <= r.get('ExpenseDate','') <= '2026-09-30')]
print('Rows not E001:', len(bad_emp))
print('Rows outside September 2026:', len(bad_date))
assert len(rows) == 13
assert not bad_emp
assert not bad_date

# Read FX rates using actual headers.
fx_path = input_dir / 'fx_rates.csv'
with fx_path.open('r', encoding='utf-8-sig', newline='') as f:
    fx_reader = csv.DictReader(f)
    fx_rows = list(fx_reader)
    fx_headers = fx_reader.fieldnames
print('FX headers:', fx_headers)
print('FX rows:', fx_rows)

currency_col = next((c for c in fx_headers if c.lower() in {'currency', 'ccy'}), None)
rate_col = next((c for c in fx_headers if 'twd' in c.lower() or 'rate' in c.lower()), None)
if not currency_col or not rate_col:
    raise AssertionError(f'Cannot identify FX columns from {fx_headers}')
fx = {r[currency_col]: Decimal(str(r[rate_col])) for r in fx_rows}
print('FX mapping:', fx)

summary = defaultdict(lambda: {'count': 0, 'amount': Decimal('0'), 'twd': Decimal('0')})
exceptions = []
for r in rows:
    amount = Decimal(str(r['Amount']))
    cur = r['Currency']
    status = r['Status']
    twd = (amount * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    key = (status, cur)
    summary[key]['count'] += 1
    summary[key]['amount'] += amount
    summary[key]['twd'] += twd
    if str(r.get('MissingReceipt','')).strip().lower() == 'true':
        exceptions.append(r)
expected_twd_total = sum((v['twd'] for v in summary.values()), Decimal('0'))
print('Expected summary rows:', len(summary))
print('Expected exceptions:', len(exceptions))
print('Expected TWD total:', expected_twd_total)

wb = openpyxl.load_workbook(workbook_path, data_only=False)
print('Workbook sheets:', wb.sheetnames)
for s in ['Details', 'Summary', 'Exceptions']:
    assert s in wb.sheetnames

# Helper to get rows by header row 4.
def get_table(sheet_name):
    sh = wb[sheet_name]
    headers = [sh.cell(4, c).value for c in range(1, sh.max_column+1)]
    # trim trailing None headers
    while headers and headers[-1] is None:
        headers.pop()
    data = []
    for rr in range(5, sh.max_row+1):
        vals = [sh.cell(rr, c).value for c in range(1, len(headers)+1)]
        if any(v is not None for v in vals):
            data.append(vals)
    return headers, data

d_headers, d_data = get_table('Details')
s_headers, s_data = get_table('Summary')
e_headers, e_data = get_table('Exceptions')
print('Details headers:', d_headers)
print('Details row count:', len(d_data))
print('Details first row:', d_data[0] if d_data else None)
print('Summary B2:', wb['Summary']['B2'].value)
print('Summary headers:', s_headers)
print('Summary row count:', len(s_data))
print('Summary rows:', s_data)
print('Exceptions headers:', e_headers)
print('Exceptions row count:', len(e_data))
print('Exceptions rows:', e_data)

assert len(d_data) == 13
assert len(s_data) == 7
assert len(e_data) == 1
assert wb['Summary']['B2'].value == '=SUM(Details!I5:I17)'

# Validate Details columns and types.
expected_detail_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','TWD Amount','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert d_headers[:len(expected_detail_headers)] == expected_detail_headers, f'Details headers mismatch: {d_headers}'
emp_idx = d_headers.index('EmployeeID')
expense_id_idx = d_headers.index('ExpenseID')
amount_idx = d_headers.index('Amount')
twd_idx = d_headers.index('TWD Amount')
missing_idx = d_headers.index('MissingReceipt')
review_idx = d_headers.index('NeedsReview')
for row in d_data:
    assert isinstance(row[expense_id_idx], str), f'ExpenseID not text: {row}'
    assert isinstance(row[emp_idx], str), f'EmployeeID not text: {row}'
    assert isinstance(row[amount_idx], (int, float)), f'Amount not numeric: {row[amount_idx]}'
    assert isinstance(row[twd_idx], (int, float)), f'TWD Amount not numeric: {row[twd_idx]}'
    assert str(row[missing_idx]).lower() in {'true','false'}, f'MissingReceipt not boolean text: {row[missing_idx]}'
    assert str(row[review_idx]).lower() in {'true','false'}, f'NeedsReview not boolean text: {row[review_idx]}'

# Validate Summary contents against computed summary.
for row in s_data:
    record = dict(zip(s_headers, row))
    key = (record['Status'], record['Currency'])
    assert key in summary, f'Unexpected summary key {key}'
    assert int(record['Count']) == summary[key]['count'], f'Count mismatch {key}'
    assert Decimal(str(record['Amount'])) == summary[key]['amount'], f'Amount mismatch {key}'
    assert Decimal(str(record['TWD Amount'])) == summary[key]['twd'], f'TWD mismatch {key}'
workbook_twd_total = sum(Decimal(str(dict(zip(s_headers, r))['TWD Amount'])) for r in s_data)
print('Workbook Summary TWD total:', workbook_twd_total)
assert workbook_twd_total == expected_twd_total == Decimal('31877.90')

# Validate Exceptions are missing receipt and need review.
for row in e_data:
    record = dict(zip(e_headers, row))
    assert str(record['MissingReceipt']).lower() == 'true'
    assert str(record['NeedsReview']).lower() == 'true'

print('Checks passed: workflow output files exist; CSV has 13 E001 rows in Sep 2026; workbook has Details 13, Summary 7, Exceptions 1; Summary formula and TWD total 31877.90 verified; ID/text and numeric amount fields verified.')
