from pathlib import Path
from openpyxl import load_workbook
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from collections import defaultdict
import csv, json, os, re

params=json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
EMPLOYEE=str(params['employee_id']).strip()
DATE_FROM=str(params['date_from']).strip()
DATE_TO=str(params['date_to']).strip()
OUTNAME=str(params.get('output_file') or '').strip()
if not OUTNAME:
    OUTNAME=f"{EMPLOYEE}_{DATE_FROM}_{DATE_TO}_expenses.xlsx"
if not OUTNAME.lower().endswith('.xlsx'):
    OUTNAME += '.xlsx'
# keep output inside output folder and avoid path separators from parameter
OUTNAME=Path(OUTNAME).name

WORKSPACE=Path.cwd()
INPUT=WORKSPACE/'input'
OUTPUT=WORKSPACE/'output'
TEMPLATE=INPUT/'expense_template.xlsx'
EXPORT=OUTPUT/f'expense_export_{EMPLOYEE}_{DATE_FROM}_{DATE_TO}.csv'
OUTFILE=OUTPUT/OUTNAME
TWOPLACES=Decimal('0.01')

def q2(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def truthy(v):
    return str(v).strip().lower() in ('true','1','yes','y')

def copy_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src=ws.cell(src_row,col); dst=ws.cell(dst_row,col)
        if src.has_style:
            dst._style=copy(src._style)
        dst.number_format=src.number_format
        dst.alignment=copy(src.alignment)
        dst.protection=copy(src.protection)

if not EXPORT.exists():
    raise FileNotFoundError(f'Expected exported CSV not found: {EXPORT}')

rates={}
with open(INPUT/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    next(reader)
    for row in reader:
        if len(row)>=2 and row[0].strip():
            rates[row[0].strip()]=Decimal(row[1].strip())

with open(EXPORT, newline='', encoding='utf-8-sig') as f:
    reader=csv.DictReader(f)
    rows=list(reader)
headers=reader.fieldnames or []
expected=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']
if headers != expected:
    raise AssertionError(f'Unexpected export headers: {headers}')

bad=[]
for r in rows:
    if r['EmployeeID'] != EMPLOYEE:
        bad.append(f"wrong employee {r['ExpenseID']} {r['EmployeeID']}")
    if not (DATE_FROM <= r['ExpenseDate'] <= DATE_TO):
        bad.append(f"out of range {r['ExpenseID']} {r['ExpenseDate']}")
    if r['Currency'] not in rates:
        bad.append(f"missing fx rate {r['Currency']} for {r['ExpenseID']}")
if bad:
    raise AssertionError('\n'.join(bad))
rows.sort(key=lambda r:(r['ExpenseDate'], r['ExpenseID']))

wb=load_workbook(TEMPLATE)
report=wb['Report']; details=wb['Details']; summary=wb['Summary']; exc=wb['Exceptions']
now=datetime.now().astimezone().isoformat(timespec='seconds')
report['B3']=DATE_FROM
report['D3']=DATE_TO
report['B4']=f'EmployeeID={EMPLOYEE}; ExpenseDate {DATE_FROM} to {DATE_TO}; all statuses'
report['B5']=now
report['B7']='Expense Viewer (synthetic lab); fx_rates.csv'
report['D7']=now
report['B9']='NeedsReview=true and Exceptions are set for MissingReceipt=true. AmountTWD uses fx_rates.csv with ROUND_HALF_UP to 2 decimals.'

for ws in (details, summary, exc):
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row-4)

last_detail_row=4 if not rows else 4+len(rows)
for idx,r in enumerate(rows, start=5):
    copy_row_style(details, 4, idx, 14)
    amount=q2(r['Amount']); amount_twd=q2(amount*rates[r['Currency']])
    missing=truthy(r.get('MissingReceipt',''))
    vals=[r['ExpenseID'],r['EmployeeID'],r['EmployeeName'],r['ProjectID'],r['ExpenseDate'],r['Category'],r['Currency'],float(amount),float(amount_twd),r['Status'],r.get('ReceiptID',''),r.get('Description',''),'true' if missing else 'false','true' if missing else 'false']
    for col,val in enumerate(vals, start=1):
        details.cell(idx,col,val)
    for col in [1,2,3,4,6,7,10,11,12,13,14]:
        details.cell(idx,col).number_format='@'
    details.cell(idx,5).number_format='yyyy-mm-dd'
    details.cell(idx,8).number_format='0.00'
    details.cell(idx,9).number_format='0.00'

agg=defaultdict(lambda:{'amount':Decimal('0'),'amount_twd':Decimal('0'),'count':0})
for r in rows:
    amount=q2(r['Amount']); amount_twd=q2(amount*rates[r['Currency']])
    key=(r['Status'], r['Currency'])
    agg[key]['amount'] += amount
    agg[key]['amount_twd'] += amount_twd
    agg[key]['count'] += 1
status_order={'Draft':0,'Submitted':1,'Approved':2,'Rejected':3}
summary_rows=[]
for (status,currency),vals in sorted(agg.items(), key=lambda kv:(status_order.get(kv[0][0],99), kv[0][0], kv[0][1])):
    summary_rows.append((status,'',currency,q2(vals['amount']),q2(vals['amount_twd']),vals['count']))
for idx,vals in enumerate(summary_rows, start=5):
    copy_row_style(summary, 4, idx, 6)
    for col,val in enumerate(vals, start=1):
        summary.cell(idx,col, float(val) if isinstance(val, Decimal) else val)
    summary.cell(idx,4).number_format='0.00'
    summary.cell(idx,5).number_format='0.00'
    summary.cell(idx,6).number_format='0'
summary['B2']=f'=SUM(Details!I5:I{last_detail_row if rows else 5})'

exceptions=[]
for r in rows:
    if truthy(r.get('MissingReceipt','')):
        exceptions.append((r['ExpenseID'],r['EmployeeID'],r['ExpenseDate'],'Missing receipt'))
for idx,vals in enumerate(exceptions, start=5):
    copy_row_style(exc, 4, idx, 4)
    for col,val in enumerate(vals, start=1):
        exc.cell(idx,col,val)
    for col in [1,2,4]:
        exc.cell(idx,col).number_format='@'
    exc.cell(idx,3).number_format='yyyy-mm-dd'

wb.save(OUTFILE)

# Reopen and verify
wb2=load_workbook(OUTFILE, data_only=False)
assert wb2.sheetnames == ['Report','Details','Summary','Exceptions','Instructions']
rep=wb2['Report']; det=wb2['Details']; summ=wb2['Summary']; ex=wb2['Exceptions']
assert rep['B3'].value == DATE_FROM and rep['D3'].value == DATE_TO and EMPLOYEE in rep['B4'].value
expected_detail=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
assert [det.cell(4,c).value for c in range(1,15)] == expected_detail
actual_detail_count=sum(1 for row in det.iter_rows(min_row=5, max_col=1, values_only=True) if row[0] is not None)
assert actual_detail_count == len(rows)
for rnum in range(5,5+len(rows)):
    assert isinstance(det.cell(rnum,1).value, str)
    assert isinstance(det.cell(rnum,2).value, str)
    assert isinstance(det.cell(rnum,4).value, str)
    assert isinstance(det.cell(rnum,8).value, (int,float))
    assert isinstance(det.cell(rnum,9).value, (int,float))
    assert det.cell(rnum,13).value in ('true','false')
    assert det.cell(rnum,14).value in ('true','false')
assert summ['B2'].value == f'=SUM(Details!I5:I{last_detail_row if rows else 5})'
summary_count=sum(1 for row in summ.iter_rows(min_row=5, max_col=1, values_only=True) if row[0] is not None)
assert summary_count == len(summary_rows)
summary_twd=sum(Decimal(str(summ.cell(r,5).value)).quantize(TWOPLACES) for r in range(5,5+summary_count))
detail_twd=sum(Decimal(str(det.cell(r,9).value)).quantize(TWOPLACES) for r in range(5,5+len(rows)))
assert summary_twd == detail_twd
exception_count=sum(1 for row in ex.iter_rows(min_row=5, max_col=1, values_only=True) if row[0] is not None)
assert exception_count == len(exceptions)
print(f'Saved {OUTFILE}')
print(f'Details rows: {actual_detail_count}')
print(f'Summary rows: {summary_count}')
print(f'Exceptions rows: {exception_count}')
print(f'TWD total: {detail_twd}')
print(f'Export source: {EXPORT}')