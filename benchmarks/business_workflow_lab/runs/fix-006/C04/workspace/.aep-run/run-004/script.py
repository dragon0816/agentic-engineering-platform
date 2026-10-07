import csv
import json
import os
import shutil
import subprocess
import sys
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from copy import copy

from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
OUTFILE = OUTPUT / 'quote_report_2026-09_sent.xlsx'
METADATA = OUTPUT / 'quote_report_2026-09_sent_runlog.json'

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
TZ = timezone(timedelta(hours=8))
NOW = datetime.now(TZ).replace(microsecond=0).isoformat()

class CliFailure(RuntimeError):
    def __init__(self, command, code, stderr):
        super().__init__(f'CLI failed with exit code {code}: {stderr}')
        self.command = command
        self.code = code
        self.stderr = stderr

class CliParseFailure(RuntimeError):
    pass

def run_cli(args, parse_json=True):
    cmd = [sys.executable, str(CLI)] + args
    attempts = []
    max_attempts = 3
    for attempt_no in range(1, max_attempts + 1):
        proc = subprocess.run(cmd, capture_output=True, text=True)
        attempts.append({
            'attempt': attempt_no,
            'returncode': proc.returncode,
            'stderr': proc.stderr.strip(),
            'stdout_length': len(proc.stdout or ''),
        })
        if proc.returncode == 0:
            if not parse_json:
                return proc.stdout, attempts
            try:
                return json.loads(proc.stdout), attempts
            except json.JSONDecodeError as e:
                raise CliParseFailure(f'CLI stdout was not parseable JSON for command {args}: {e}')
        if proc.returncode in (2, 3):
            raise CliFailure(args, proc.returncode, proc.stderr.strip())
        if proc.returncode == 4 and attempt_no < max_attempts:
            continue
        raise CliFailure(args, proc.returncode, proc.stderr.strip())

fx_rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx_rates[row[0]] = Decimal(row[1])
print('FX rates:', {k: str(v) for k, v in fx_rates.items()})

runlog = {'commands': []}
list_args = ['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
list_data, attempts = run_cli(list_args)
runlog['commands'].append({'args': list_args, 'attempts': attempts})
if list_data.get('schema_version') != 1 or 'items' not in list_data:
    raise RuntimeError(f'Unexpected quotes list schema: {list_data}')
items = list_data['items']
print(f'Fetched list count={list_data.get("count")} items={len(items)}')

headers = sorted(items, key=lambda x: x.get('QuoteID', ''))
details_by_qid = {}
all_lines = []
for h in headers:
    qid = h['QuoteID']
    get_args = ['quotes', 'get', qid, '--format', 'json']
    detail, attempts = run_cli(get_args)
    runlog['commands'].append({'args': get_args, 'attempts': attempts})
    if detail.get('schema_version') != 1 or 'header' not in detail or 'line_items' not in detail:
        raise RuntimeError(f'Unexpected quotes get schema for {qid}: {detail}')
    details_by_qid[qid] = detail
    for line in detail['line_items']:
        row = dict(line)
        row['QuoteID'] = qid
        all_lines.append(row)
print(f'Fetched details={len(details_by_qid)} line_items={len(all_lines)}')

def money(value):
    if value is None or value == '':
        return None
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def dec(value):
    if value is None or value == '':
        return None
    return Decimal(str(value))

def to_float(d):
    if d is None:
        return None
    return float(d)

def amount_twd(amount, currency):
    if amount is None:
        return None
    if currency not in fx_rates:
        raise RuntimeError(f'Missing FX rate for currency {currency}')
    return (amount * fx_rates[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def clear_data_rows(ws, start_row=5):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def copy_row_style(ws, source_row, target_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(source_row, col)
        dst = ws.cell(target_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.protection = copy(src.protection)
        dst.number_format = src.number_format

shutil.copy2(TEMPLATE, OUTFILE)
wb = load_workbook(OUTFILE)
report = wb['Report']
quotes_ws = wb['Quotes']
lines_ws = wb['Lines']
summary_ws = wb['Summary']

report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = f'Status={STATUS}; QuoteDate between {FROM_DATE} and {TO_DATE} inclusive'
report['B5'] = NOW
report['B7'] = 'quote_cli; fx_rates.csv'
report['D7'] = NOW
report['B9'] = (f'Fetched {len(headers)} Sent quote header(s) and {len(all_lines)} line item(s). '
                'CLI exit-code handling applied: code 4 retried up to two times; '
                'code 2/3 would stop without retry and report stderr.')

clear_data_rows(quotes_ws, 5)
clear_data_rows(lines_ws, 5)
clear_data_rows(summary_ws, 5)

quote_columns = [quotes_ws.cell(4, c).value for c in range(1, quotes_ws.max_column + 1)]
quote_values_for_summary = []
for idx, h in enumerate(headers, start=5):
    subtotal = money(h.get('Subtotal'))
    discount = dec(h.get('DiscountPercent'))
    net = money(h.get('NetAmount'))
    currency = h.get('Currency')
    net_twd = amount_twd(net, currency)
    values = {
        'QuoteID': h.get('QuoteID'),
        'CustomerID': h.get('CustomerID'),
        'ProjectID': h.get('ProjectID'),
        'QuoteDate': h.get('QuoteDate'),
        'ValidUntil': h.get('ValidUntil'),
        'Status': h.get('Status'),
        'Currency': currency,
        'Subtotal': to_float(subtotal),
        'DiscountPercent': to_float(discount),
        'NetAmount': to_float(net),
        'NetAmountTWD': to_float(net_twd),
    }
    copy_row_style(quotes_ws, 4, idx, len(quote_columns))
    for c, colname in enumerate(quote_columns, start=1):
        cell = quotes_ws.cell(idx, c)
        cell.value = values.get(colname)
        if colname in ('QuoteID', 'CustomerID', 'ProjectID', 'Status', 'Currency'):
            cell.number_format = '@'
        elif colname in ('QuoteDate', 'ValidUntil'):
            cell.number_format = 'yyyy-mm-dd'
        elif colname in ('Subtotal', 'NetAmount', 'NetAmountTWD'):
            cell.number_format = '0.00'
    quote_values_for_summary.append({'Status': h.get('Status'), 'Currency': currency, 'NetAmount': net, 'NetAmountTWD': net_twd})

line_columns = [lines_ws.cell(4, c).value for c in range(1, lines_ws.max_column + 1)]
for idx, line in enumerate(sorted(all_lines, key=lambda x: (x.get('QuoteID',''), x.get('ProductID',''))), start=5):
    qty = dec(line.get('Quantity'))
    unit = money(line.get('UnitPrice'))
    amt = money(line.get('LineAmount'))
    values = {
        'QuoteID': line.get('QuoteID'),
        'ProductID': line.get('ProductID'),
        'ProductName': line.get('ProductName'),
        'Quantity': to_float(qty),
        'UnitPrice': to_float(unit),
        'LineAmount': to_float(amt),
    }
    copy_row_style(lines_ws, 4, idx, len(line_columns))
    for c, colname in enumerate(line_columns, start=1):
        cell = lines_ws.cell(idx, c)
        cell.value = values.get(colname)
        if colname in ('QuoteID', 'ProductID', 'ProductName'):
            cell.number_format = '@'
        elif colname in ('UnitPrice', 'LineAmount'):
            cell.number_format = '0.00'

summary_columns = [summary_ws.cell(4, c).value for c in range(1, summary_ws.max_column + 1)]
summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
for r in quote_values_for_summary:
    key = (r['Status'], r['Currency'])
    summary[key]['amount'] += r['NetAmount'] or Decimal('0.00')
    summary[key]['twd'] += r['NetAmountTWD'] or Decimal('0.00')
    summary[key]['count'] += 1
for idx, ((status, currency), agg) in enumerate(sorted(summary.items()), start=5):
    copy_row_style(summary_ws, 4, idx, len(summary_columns))
    values = {
        'Group': 'Status',
        'SubGroup': status,
        'Currency': currency,
        'Amount': to_float(agg['amount'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)),
        'AmountTWD': to_float(agg['twd'].quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)),
        'Count': agg['count'],
    }
    for c, colname in enumerate(summary_columns, start=1):
        cell = summary_ws.cell(idx, c)
        cell.value = values.get(colname)
        if colname in ('Amount', 'AmountTWD'):
            cell.number_format = '0.00'

last_quote_row = 4 + len(headers)
summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})' if headers else '=SUM(Quotes!K5:K5)'
summary_ws['B2'].number_format = '0.00'

wb.save(OUTFILE)
runlog.update({
    'from': FROM_DATE,
    'to': TO_DATE,
    'status': STATUS,
    'generated_at': NOW,
    'quote_count': len(headers),
    'line_count': len(all_lines),
    'output': str(OUTFILE),
})
METADATA.write_text(json.dumps(runlog, ensure_ascii=False, indent=2), encoding='utf-8')
print('Wrote:', OUTFILE)
print('Wrote:', METADATA)

# Verification
wb2 = load_workbook(OUTFILE, data_only=False)
tmpl = load_workbook(TEMPLATE, read_only=True)
assert wb2.sheetnames == tmpl.sheetnames, f'Sheet names/order changed: {wb2.sheetnames} vs {tmpl.sheetnames}'
qws = wb2['Quotes']; lws = wb2['Lines']; sws = wb2['Summary']; rws = wb2['Report']
assert [qws.cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [lws.cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert [sws.cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
actual_quote_rows = max(0, qws.max_row - 4)
actual_line_rows = max(0, lws.max_row - 4)
actual_summary_rows = max(0, sws.max_row - 4)
assert actual_quote_rows == len(headers), (actual_quote_rows, len(headers))
assert actual_line_rows == len(all_lines), (actual_line_rows, len(all_lines))
assert actual_summary_rows == len(summary), (actual_summary_rows, len(summary))
assert rws['B3'].value == FROM_DATE and rws['D3'].value == TO_DATE
assert STATUS in rws['B4'].value
assert sws['B2'].value == (f'=SUM(Quotes!K5:K{last_quote_row})' if headers else '=SUM(Quotes!K5:K5)')

for row in range(5, 5 + len(headers)):
    qid = qws.cell(row,1).value
    currency = qws.cell(row,7).value
    net = Decimal(str(qws.cell(row,10).value)).quantize(Decimal('0.01'))
    twd = Decimal(str(qws.cell(row,11).value)).quantize(Decimal('0.01'))
    assert qws.cell(row,1).data_type == 's', f'QuoteID not text at row {row}'
    assert qws.cell(row,2).data_type == 's', f'CustomerID not text at row {row}'
    assert qws.cell(row,3).data_type == 's', f'ProjectID not text at row {row}'
    assert qws.cell(row,6).value == STATUS, f'Status mismatch at row {row}'
    assert qws.cell(row,8).number_format == '0.00'
    assert qws.cell(row,10).number_format == '0.00'
    assert qws.cell(row,11).number_format == '0.00'
    expected_twd = amount_twd(net, currency)
    assert twd == expected_twd, f'TWD mismatch {qid}: {twd} != {expected_twd}'

quote_ids = {qws.cell(row,1).value for row in range(5, 5 + len(headers))}
for row in range(5, 5 + len(all_lines)):
    assert lws.cell(row,1).value in quote_ids, f'Line row {row} quote id missing in Quotes'
    assert lws.cell(row,1).data_type == 's'
    assert lws.cell(row,2).data_type == 's'
    assert isinstance(lws.cell(row,4).value, (int, float)), f'Quantity not numeric row {row}'
    assert isinstance(lws.cell(row,5).value, (int, float)), f'UnitPrice not numeric row {row}'
    assert isinstance(lws.cell(row,6).value, (int, float)), f'LineAmount not numeric row {row}'
    assert lws.cell(row,5).number_format == '0.00'
    assert lws.cell(row,6).number_format == '0.00'

sum_quote_twd = sum(Decimal(str(qws.cell(row,11).value)).quantize(Decimal('0.01')) for row in range(5, 5 + len(headers)))
sum_summary_twd = sum(Decimal(str(sws.cell(row,5).value)).quantize(Decimal('0.01')) for row in range(5, sws.max_row + 1) if sws.cell(row,5).value is not None)
assert sum_quote_twd == sum_summary_twd, f'Summary TWD total mismatch {sum_quote_twd} != {sum_summary_twd}'

# Validate no disallowed retry behavior happened: exit 4 attempts may be >1, exit 2/3 would be first and failed; all successful here.
for cmd in runlog['commands']:
    codes = [a['returncode'] for a in cmd['attempts']]
    assert codes[-1] == 0, f'Unsuccessful command in runlog: {cmd}'
    assert len(codes) <= 3, f'Too many attempts: {cmd}'

print('CHECK PASSED')
print(json.dumps({
    'output_file': str(OUTFILE),
    'runlog_file': str(METADATA),
    'quotes': len(headers),
    'lines': len(all_lines),
    'summary_rows': actual_summary_rows,
    'summary_b2_formula': sws['B2'].value,
    'quote_twd_total': str(sum_quote_twd),
    'cli_commands': len(runlog['commands']),
    'attempts_per_command': [len(c['attempts']) for c in runlog['commands']],
}, ensure_ascii=False, indent=2))
