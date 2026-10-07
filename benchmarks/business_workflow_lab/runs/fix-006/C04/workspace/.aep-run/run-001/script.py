import os, csv, json
from pathlib import Path
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace:', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# CSV headers/sample
csv_path = ws/'input'/'fx_rates.csv'
if csv_path.exists():
    print('\nfx_rates.csv preview:')
    with csv_path.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i, row in enumerate(reader):
            print(i, row)
            if i >= 5:
                break

# Workbook structure
xlsx_path = ws/'input'/'quote_template.xlsx'
if xlsx_path.exists():
    print('\nquote_template.xlsx sheets/cells preview:')
    wb = load_workbook(xlsx_path, data_only=False)
    for sh in wb.sheetnames:
        wsht = wb[sh]
        print('sheet:', sh, 'max_row:', wsht.max_row, 'max_col:', wsht.max_column)
        for r in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 10), values_only=False):
            vals = [c.value for c in r[:min(wsht.max_column, 12)]]
            print(vals)
        print('---')

# Manual preview
manual_path = ws/'input'/'quote_cli_manual.md'
if manual_path.exists():
    print('\nquote_cli_manual.md preview:')
    text = manual_path.read_text(encoding='utf-8', errors='replace')
    print(text[:4000])
