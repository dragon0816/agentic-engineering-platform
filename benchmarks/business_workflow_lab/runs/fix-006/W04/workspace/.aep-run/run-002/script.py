from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('downloads:')
for p in sorted((ws/'output'/'downloads').glob('*')):
    print(p, p.stat().st_size)
    print(p.read_text(encoding='utf-8-sig')[:1000])

wb=load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
sh=wb['Instructions']
print('Instructions all rows:')
for row in sh.iter_rows(values_only=True):
    print(row)