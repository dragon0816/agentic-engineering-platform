from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from datetime import date
import csv, shutil, json

ws = Path.cwd()
input_xlsx = ws/'input'/'sales_template.xlsx'
output_dir = ws/'output'
out_xlsx = output_dir/'sales_template_filled.xlsx'
output_dir.mkdir(exist_ok=True)

# Read downloaded filtered list for authoritative OrderID order and list fields
csv_path = output_dir/'downloads'/'open_orders_2026_09.csv'
orders = []
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        orders.append(row)
print('Downloaded filtered rows:', len(orders))
print('Downloaded headers:', reader.fieldnames)
print('OrderIDs:', [r['OrderID'] for r in orders])

# Details read through browser from ERP detail pages
raw_details = {
"0045000109": {"OrderID": "0045000109", "CustomerID": "C001", "CustomerName": "Acme Precision Taiwan Ltd.", "ProjectID": "P003", "OrderDate": "2026-09-01", "Status": "Open", "Currency": "TWD", "NetAmount": "150000.00", "CustomerContact": "Grace Lin <grace.lin@acme-precision.example>", "PaymentTerms": "NET10"},
"0045000110": {"OrderID": "0045000110", "CustomerID": "C002", "CustomerName": "台灣精密股份有限公司", "ProjectID": "P002", "OrderDate": "2026-09-02", "Status": "Open", "Currency": "TWD", "NetAmount": "293572.50", "CustomerContact": "張美玲 <mei@tw-precision.example>", "PaymentTerms": "NET30"},
"0045000112": {"OrderID": "0045000112", "CustomerID": "C007", "CustomerName": "Acme Precision Taiwan Limited", "ProjectID": "P005", "OrderDate": "2026-09-04", "Status": "Open", "Currency": "USD", "NetAmount": "12311.00", "CustomerContact": "Peter Ho <peter.ho@acme-limited.example>", "PaymentTerms": "NET15"},
"0045000113": {"OrderID": "0045000113", "CustomerID": "C004", "CustomerName": "藍海電子有限公司", "ProjectID": "", "OrderDate": "2026-09-05", "Status": "Open", "Currency": "TWD", "NetAmount": "15390.00", "CustomerContact": "", "PaymentTerms": "NET15"},
"0045000115": {"OrderID": "0045000115", "CustomerID": "C005", "CustomerName": "Pacific Test Systems Inc.", "ProjectID": "", "OrderDate": "2026-09-09", "Status": "Open", "Currency": "USD", "NetAmount": "6391.00", "CustomerContact": "Sam Patel <sam@pacifictest.example>", "PaymentTerms": "NET20"},
"0045000116": {"OrderID": "0045000116", "CustomerID": "C008", "CustomerName": "新竹量測科技", "ProjectID": "", "OrderDate": "2026-09-10", "Status": "Open", "Currency": "TWD", "NetAmount": "29400.00", "CustomerContact": "李建宏 <jh.lee@hsinchu-metrology.example>", "PaymentTerms": "NET30"},
"0045000118": {"OrderID": "0045000118", "CustomerID": "C003", "CustomerName": "Northwind Components GmbH", "ProjectID": "P004", "OrderDate": "2026-09-14", "Status": "Open", "Currency": "EUR", "NetAmount": "17484.75", "CustomerContact": "Jonas Krüger <jk@northwind.example>", "PaymentTerms": "NET0"},
"0045000119": {"OrderID": "0045000119", "CustomerID": "C001", "CustomerName": "Acme Precision Taiwan Ltd.", "ProjectID": "P001", "OrderDate": "2026-09-15", "Status": "Open", "Currency": "EUR", "NetAmount": "8775.00", "CustomerContact": "Grace Lin <grace.lin@acme-precision.example>", "PaymentTerms": "NET30"},
"0045000120": {"OrderID": "0045000120", "CustomerID": "C006", "CustomerName": "Helios Optics B.V.", "ProjectID": "", "OrderDate": "2026-09-16", "Status": "Open", "Currency": "EUR", "NetAmount": "461.70", "CustomerContact": "Eva de Vries <eva@helios.example>", "PaymentTerms": "NET10"},
"0045000122": {"OrderID": "0045000122", "CustomerID": "C002", "CustomerName": "台灣精密股份有限公司", "ProjectID": "P002", "OrderDate": "2026-09-21", "Status": "Open", "Currency": "USD", "NetAmount": "4141.00", "CustomerContact": "張美玲 <mei@tw-precision.example>", "PaymentTerms": "NET5"},
"0045000124": {"OrderID": "0045000124", "CustomerID": "C004", "CustomerName": "藍海電子有限公司", "ProjectID": "", "OrderDate": "2026-09-24", "Status": "Open", "Currency": "TWD", "NetAmount": "45000.00", "CustomerContact": "", "PaymentTerms": "NET30"},
"0045000126": {"OrderID": "0045000126", "CustomerID": "C008", "CustomerName": "新竹量測科技", "ProjectID": "", "OrderDate": "2026-09-28", "Status": "Open", "Currency": "USD", "NetAmount": "9171.00", "CustomerContact": "李建宏 <jh.lee@hsinchu-metrology.example>", "PaymentTerms": "NET30"},
"0045000127": {"OrderID": "0045000127", "CustomerID": "C007", "CustomerName": "Acme Precision Taiwan Limited", "ProjectID": "P005", "OrderDate": "2026-09-30", "Status": "Open", "Currency": "EUR", "NetAmount": "2214.92", "CustomerContact": "Peter Ho <peter.ho@acme-limited.example>", "PaymentTerms": "NET30"},
"0045000128": {"OrderID": "0045000128", "CustomerID": "C003", "CustomerName": "Northwind Components GmbH", "ProjectID": "P004", "OrderDate": "2026-09-30", "Status": "Open", "Currency": "TWD", "NetAmount": "104160.00", "CustomerContact": "Jonas Krüger <jk@northwind.example>", "PaymentTerms": "NET5"},
}

# Read fx rates: file note has extra comma, only first two columns matter
rates = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        rates[row['Currency']] = Decimal(row['RateToTWD'])
print('FX rates:', rates)

# Check completeness vs downloaded list
ids = [r['OrderID'] for r in orders]
missing = [oid for oid in ids if oid not in raw_details]
extra = [oid for oid in raw_details if oid not in ids]
assert not missing, f'Missing details: {missing}'
assert not extra, f'Unexpected details not in filtered list: {extra}'
assert len(ids) == len(set(ids)), 'Duplicate OrderIDs in source list'

# Cross-check detail header fields match downloaded list exactly (after decimal normalization)
for src in orders:
    d = raw_details[src['OrderID']]
    for k in ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency']:
        assert (src.get(k) or '') == (d.get(k) or ''), f'Mismatch {src["OrderID"]} {k}: csv={src.get(k)} detail={d.get(k)}'
    assert Decimal(src['NetAmount']) == Decimal(d['NetAmount']), f'NetAmount mismatch {src["OrderID"]}'

# Copy and fill original template
shutil.copy2(input_xlsx, out_xlsx)
wb = load_workbook(out_xlsx)

# Report sheet
rep = wb['Report']
rep['B3'] = '2026-09-01'
rep['D3'] = '2026-09-30'
rep['B4'] = 'Sales Orders: OrderDate 2026-09-01..2026-09-30 inclusive; Status=Open; sorted by OrderID asc'
rep['B5'] = '2026-10-06T00:00:00+08:00'
rep['B7'] = 'Lab ERP (synthetic) Sales Orders; fx_rates.csv'
rep['D7'] = '2026-10-06T00:00:00+08:00'
rep['B9'] = 'ERP session expired once and was restored by SSO login; all 14 filtered orders have matching detail records. Blank source values are left blank.'

# Details sheet: preserve row 4 headers, data from row 5
headers = [wb['Details'].cell(4,c).value for c in range(1,12)]
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
assert headers == expected_headers, headers
det = wb['Details']
# clear pre-existing rows 5+ within relevant columns if any
for row in range(5, max(det.max_row,5)+1):
    for col in range(1, 12):
        det.cell(row, col).value = None

def money(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

rows_for_summary = []
for idx, oid in enumerate(ids, start=5):
    d = raw_details[oid]
    amt = Decimal(d['NetAmount'])
    cur = d['Currency']
    twd = (amt * rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    values = [
        d['OrderID'], d['CustomerID'], d['CustomerName'], d['ProjectID'] or None,
        date.fromisoformat(d['OrderDate']), d['Status'], cur, float(amt), float(twd),
        d['CustomerContact'] or None, d['PaymentTerms'] or None
    ]
    for col, val in enumerate(values, start=1):
        cell = det.cell(idx, col)
        cell.value = val
        if col in [1,2,4]:
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in [8,9]:
            cell.number_format = '0.00'
    rows_for_summary.append((d['Status'], cur, amt, twd))

# Summary: total formula and grouped by Status/Currency
summary = wb['Summary']
last_row = 4 + len(ids)
summary['B2'] = f'=SUM(Details!I5:I{last_row})' if ids else '=SUM(Details!I5:I5)'
# clear existing summary data rows 5+
for row in range(5, max(summary.max_row,5)+1):
    for col in range(1,7):
        summary.cell(row,col).value = None

# group by (Status, Currency); SubGroup unused left blank
from collections import OrderedDict
groups = OrderedDict()
for status, cur, amt, twd in rows_for_summary:
    key = (status, cur)
    if key not in groups:
        groups[key] = [Decimal('0.00'), Decimal('0.00'), 0]
    groups[key][0] += amt
    groups[key][1] += twd
    groups[key][2] += 1

r = 5
for (status, cur), (amt_sum, twd_sum, count) in groups.items():
    vals = [status, None, cur, float(amt_sum.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), float(twd_sum.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)), count]
    for c, val in enumerate(vals, start=1):
        summary.cell(r,c).value = val
        if c in [4,5]:
            summary.cell(r,c).number_format = '0.00'
    r += 1

# Save
wb.save(out_xlsx)
print('Wrote', out_xlsx)

# Reopen and verify every stated/template requirement we can check
wb2 = load_workbook(out_xlsx, data_only=False)
assert wb2.sheetnames == ['Report','Details','Summary','Instructions']
d2 = wb2['Details']
assert [d2.cell(4,c).value for c in range(1,12)] == expected_headers
# row count and no blank rows among data rows
read_ids = [d2.cell(r,1).value for r in range(5, 5+len(ids))]
assert read_ids == ids, f'OrderID order mismatch: {read_ids} vs {ids}'
assert all(isinstance(x, str) and x.startswith('00') for x in read_ids)
assert d2.cell(5,5).number_format.lower() in ('yyyy-mm-dd','yyyy-mm-dd')
for r in range(5, 5+len(ids)):
    assert d2.cell(r,6).value == 'Open'
    assert str(d2.cell(r,5).value)[:10] >= '2026-09-01' and str(d2.cell(r,5).value)[:10] <= '2026-09-30'
    assert isinstance(d2.cell(r,8).value, (int,float)), f'NetAmount not numeric row {r}'
    assert isinstance(d2.cell(r,9).value, (int,float)), f'NetAmountTWD not numeric row {r}'
    assert d2.cell(r,8).number_format == '0.00'
    assert d2.cell(r,9).number_format == '0.00'
# ensure next row blank in key cols
assert all(d2.cell(5+len(ids), c).value is None for c in range(1,12)), 'Unexpected data after expected rows'
# Validate TWD conversions from saved file
saved_total = Decimal('0.00')
for r in range(5, 5+len(ids)):
    cur = d2.cell(r,7).value
    amt = Decimal(str(d2.cell(r,8).value))
    expected = (amt * rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    actual = Decimal(str(d2.cell(r,9).value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    assert actual == expected, f'TWD conversion mismatch row {r}: {actual} != {expected}'
    saved_total += actual
s2 = wb2['Summary']
assert s2['B2'].value == '=SUM(Details!I5:I18)'
summary_rows = []
for r in range(5, 8):
    summary_rows.append(tuple(s2.cell(r,c).value for c in range(1,7)))
assert len([x for x in summary_rows if x[0]]) == 3
summary_total = sum(Decimal(str(x[4])).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP) for x in summary_rows)
assert summary_total == saved_total.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
# Report fields
assert wb2['Report']['B3'].value == '2026-09-01'
assert wb2['Report']['D3'].value == '2026-09-30'
assert 'Status=Open' in wb2['Report']['B4'].value
print('Verification passed:')
print('- output workbook exists and sheets preserved')
print('- Details headers preserved; 14 rows written from row 5 with no duplicate/missing OrderID')
print('- OrderID/CustomerID/ProjectID stored as text where applicable; blank source values left blank')
print('- dates within 2026-09-01..2026-09-30 and status Open')
print('- NetAmount and NetAmountTWD numeric with 0.00 format; TWD values match fx_rates.csv ROUND_HALF_UP')
print('- Summary B2 formula extends to Details!I5:I18 and grouped totals equal Details total')
print('Summary rows:', summary_rows)
print('Total NetAmountTWD:', saved_total)
