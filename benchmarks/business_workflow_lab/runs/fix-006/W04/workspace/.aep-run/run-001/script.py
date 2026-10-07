from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

# workbook structure
xlsx = ws/'input'/'sales_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(' ', r)

# fx rates
csvp = ws/'input'/'fx_rates.csv'
with open(csvp, newline='', encoding='utf-8-sig') as f:
    rdr = csv.reader(f)
    rows=[]
    for i,row in enumerate(rdr):
        rows.append(row)
        if i>=8: break
print('fx_rates sample:')
for r in rows: print(' ', r)

# manual sample
md = ws/'input'/'sap_manual.md'
text = md.read_text(encoding='utf-8')
print('manual first 2000 chars:')
print(text[:2000])