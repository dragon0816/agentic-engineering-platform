from pathlib import Path
import csv, json
from openpyxl import load_workbook
base=Path('.')
print('Workspace:', base.resolve())
print('Input files:')
for p in sorted((base/'input').iterdir()):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Workbook structure
wb_path=base/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print(f'\nSheet {ws.title}: {ws.max_row} rows x {ws.max_column} cols')
    for r in range(1, min(ws.max_row,5)+1):
        print([ws.cell(r,c).value for c in range(1, min(ws.max_column,12)+1)])

# CSV sample
csv_path=base/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(8), rdr):
        print(row)

# Manual excerpt
md_path=base/'input'/'expense_manual.md'
print('\nexpense_manual.md first lines:')
with open(md_path, encoding='utf-8') as f:
    for i,line in zip(range(80), f):
        print(line.rstrip())