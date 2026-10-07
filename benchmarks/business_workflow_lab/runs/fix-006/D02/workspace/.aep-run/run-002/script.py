from pathlib import Path
import csv
from openpyxl import load_workbook
base=Path('.')
for name in ['submitted_sep2026.csv','approved_sep2026.csv']:
    p=base/'output'/name
    print('\nCSV', name, 'exists', p.exists(), 'size', p.stat().st_size if p.exists() else None)
    if p.exists():
        with open(p, newline='', encoding='utf-8-sig') as f:
            rdr=csv.reader(f)
            for i,row in zip(range(5), rdr):
                print(row)

wb=load_workbook(base/'input'/'expense_template.xlsx', data_only=False)
ws=wb['Instructions']
print('\nInstructions all rows:')
for row in ws.iter_rows(values_only=True):
    print(row)