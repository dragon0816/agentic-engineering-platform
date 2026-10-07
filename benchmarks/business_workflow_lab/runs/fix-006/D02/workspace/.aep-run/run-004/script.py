from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv, shutil, copy
from collections import defaultdict
from openpyxl import load_workbook

base=Path('.')
in_dir=base/'input'
out_dir=base/'output'
out_dir.mkdir(exist_ok=True)

submitted=out_dir/'submitted_sep2026.csv'
approved=out_dir/'approved_sep2026.csv'
template=in_dir/'expense_template.xlsx'
fx_path=in_dir/'fx_rates.csv'
out_path=out_dir/'expense_report_2026-09_submitted_approved.xlsx'

rates={}
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    next(rdr)
    for row in rdr:
        if len(row)>=2 and row[0]:
            rates[row[0]]=Decimal(row[1])
print('FX rates:', rates)

rows=[]
for p in [submitted, approved]:
    if not p.exists():
        raise FileNotFoundError(p)
    with open(p, newline='', encoding='utf-8-sig') as f:
        rdr=csv.DictReader(f)
        print(f'{p.name} headers:', rdr.fieldnames)
        for r in rdr:
            rows.append(r)

required_cols=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']
for r in rows:
    missing=[c for c in required_cols if c not in r]
    if missing:
        raise ValueError(f'Missing columns {missing}')
    if r['Status'] not in {'Submitted','Approved'}:
        raise ValueError(f'Unexpected status {r}')
    if not ('2026-09-01' <= r['ExpenseDate'] <= '2026-09-30'):
        raise ValueError(f'Unexpected date {r}')
    if r['Currency'] not in rates:
        raise ValueError(f'Missing FX rate for {r["Currency"]}')

rows.sort(key=lambda r: (r['ExpenseDate'], r['ExpenseID']))
for r in rows:
    amt=Decimal(r['Amount'])
    r['_AmountDec']=amt
    r['_AmountTWDDec']=(amt * rates[r['Currency']]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    r['_Missing']=str(r['MissingReceipt']).strip().lower() == 'true'
    r['_NeedsReview']=r['_Missing']

shutil.copy2(template, out_path)
wb=load_workbook(out_path)

def clear_data(ws, start_row):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)

def copy_row_style(ws, src_row, dst_row, max_col):
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height
    for c in range(1, max_col+1):
        src=ws.cell(src_row, c)
        dst=ws.cell(dst_row, c)
        if src.has_style:
            dst._style = copy.copy(src._style)
        dst.number_format = src.number_format
        dst.font = copy.copy(src.font)
        dst.fill = copy.copy(src.fill)
        dst.border = copy.copy(src.border)
        dst.alignment = copy.copy(src.alignment)
        dst.protection = copy.copy(src.protection)

# Report sheet
ws=wb['Report']
ws['B3']='2026-09-01'
ws['D3']='2026-09-30'
ws['B4']='ExpenseDate 2026-09-01 through 2026-09-30; Status in Submitted, Approved; all employees/projects'
ws['B5']='2026-10-06T00:00:00+08:00'
ws['B7']='Expense Viewer (synthetic lab); fx_rates.csv'
ws['D7']='2026-10-06T00:00:00+08:00'
missing_count=sum(1 for r in rows if r['_Missing'])
ws['B9']=f'{len(rows)} expenses included; {missing_count} missing receipt expense(s) listed in Exceptions and retained in Details.'

# Details
ws=wb['Details']
clear_data(ws, 5)
for idx,r in enumerate(rows, start=5):
    copy_row_style(ws, 4, idx, 14)
    values=[r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], datetime.strptime(r['ExpenseDate'], '%Y-%m-%d').date(), r['Category'], r['Currency'], float(r['_AmountDec']), float(r['_AmountTWDDec']), r['Status'], r['ReceiptID'] if r['ReceiptID'] else None, r['Description'] if r['Description'] else None, 'true' if r['_Missing'] else 'false', 'true' if r['_NeedsReview'] else 'false']
    for c,val in enumerate(values, start=1):
        cell=ws.cell(idx,c,val)
        if c in [1,2,4,11,13,14]:
            cell.number_format='@'
        elif c==5:
            cell.number_format='yyyy-mm-dd'
        elif c in [8,9]:
            cell.number_format='0.00'
last_detail_row=4+len(rows) if rows else 5
ws.auto_filter.ref=f'A4:N{last_detail_row}'

# Summary
ws=wb['Summary']
clear_data(ws, 5)
ws['B2']=f'=SUM(Details!I5:I{last_detail_row})'
summary=defaultdict(lambda: {'amount': Decimal('0'), 'twd': Decimal('0'), 'count':0})
for r in rows:
    key=(r['Status'], '', r['Currency'])
    summary[key]['amount'] += r['_AmountDec']
    summary[key]['twd'] += r['_AmountTWDDec']
    summary[key]['count'] += 1
for idx,(key, agg) in enumerate(sorted(summary.items()), start=5):
    copy_row_style(ws, 4, idx, 6)
    vals=[key[0], None, key[2], float(agg['amount'].quantize(Decimal('0.01'))), float(agg['twd'].quantize(Decimal('0.01'))), agg['count']]
    for c,val in enumerate(vals, start=1):
        cell=ws.cell(idx,c,val)
        if c in [4,5]: cell.number_format='0.00'
        if c==6: cell.number_format='0'
last_summary=4+len(summary) if summary else 5
ws.auto_filter.ref=f'A4:F{last_summary}'

# Exceptions
ws=wb['Exceptions']
clear_data(ws, 5)
exceptions=[r for r in rows if r['_Missing']]
for idx,r in enumerate(exceptions, start=5):
    copy_row_style(ws, 4, idx, 4)
    vals=[r['ExpenseID'], r['EmployeeID'], datetime.strptime(r['ExpenseDate'], '%Y-%m-%d').date(), 'MissingReceipt=true']
    for c,val in enumerate(vals, start=1):
        cell=ws.cell(idx,c,val)
        if c in [1,2]: cell.number_format='@'
        elif c==3: cell.number_format='yyyy-mm-dd'
last_ex=4+len(exceptions) if exceptions else 5
ws.auto_filter.ref=f'A4:D{last_ex}'

wb.save(out_path)
print('Wrote', out_path)

# Reopen and validate every stated requirement.
wb2=load_workbook(out_path, data_only=False)
expected_sheets=['Report','Details','Summary','Exceptions','Instructions']
assert wb2.sheetnames == expected_sheets, wb2.sheetnames

ws=wb2['Details']
expected_headers=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
out_headers=[ws.cell(4,c).value for c in range(1,15)]
assert out_headers == expected_headers, out_headers
detail_rows=[r for r in range(5, ws.max_row+1) if ws.cell(r,1).value]
assert len(detail_rows)==22, len(detail_rows)
ids=[ws.cell(r,1).value for r in detail_rows]
assert len(ids)==len(set(ids))==22
assert set(ids)==set(r['ExpenseID'] for r in rows)
assert all(ws.cell(r,10).value in ['Submitted','Approved'] for r in detail_rows)
assert all(ws.cell(r,5).value.strftime('%Y-%m-%d') >= '2026-09-01' and ws.cell(r,5).value.strftime('%Y-%m-%d') <= '2026-09-30' for r in detail_rows)
assert all(isinstance(ws.cell(r,1).value, str) for r in detail_rows)
assert all(isinstance(ws.cell(r,8).value, (int,float)) for r in detail_rows)
assert all(isinstance(ws.cell(r,9).value, (int,float)) for r in detail_rows)
assert all(ws.cell(r,13).value in ['true','false'] for r in detail_rows)
assert all(ws.cell(r,14).value in ['true','false'] for r in detail_rows)
# Detail fields populated from source: Description nonblank, ReceiptID either source ID or blank when missing.
assert all(ws.cell(r,12).value not in [None, ''] for r in detail_rows)
for r in detail_rows:
    if ws.cell(r,13).value == 'true':
        assert ws.cell(r,11).value in [None, '']

ws_ex=wb2['Exceptions']
ex_rows=[r for r in range(5, ws_ex.max_row+1) if ws_ex.cell(r,1).value]
ex_ids=[ws_ex.cell(r,1).value for r in ex_rows]
expected_ex_ids=[r['ExpenseID'] for r in exceptions]
assert ex_ids==expected_ex_ids, (ex_ids, expected_ex_ids)
assert set(ex_ids).issubset(set(ids))
assert all(ws_ex.cell(r,4).value=='MissingReceipt=true' for r in ex_rows)

ws_sum=wb2['Summary']
assert ws_sum['B2'].value == f'=SUM(Details!I5:I{last_detail_row})'
summary_rows=[r for r in range(5, ws_sum.max_row+1) if ws_sum.cell(r,1).value]
assert len(summary_rows)==len(summary)

submitted_count=sum(1 for r in rows if r['Status']=='Submitted')
approved_count=sum(1 for r in rows if r['Status']=='Approved')
total_twd=sum(r['_AmountTWDDec'] for r in rows)
print('Validated output workbook:')
print(' - Details rows:', len(detail_rows))
print(' - Submitted:', submitted_count, 'Approved:', approved_count)
print(' - Exception rows:', len(ex_rows), ex_ids)
print(' - Exception IDs retained in Details:', set(ex_ids).issubset(set(ids)))
print(' - Summary rows:', len(summary_rows), 'Summary B2:', ws_sum['B2'].value)
print(' - Total AmountTWD:', total_twd)
print(' - Output file:', out_path.resolve())