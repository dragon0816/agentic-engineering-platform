from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook(Path('input')/'quote_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
print('\nSummary existing formulas/styles:')
for sheet in ['Report','Quotes','Lines','Summary']:
    sh=wb[sheet]
    print(sheet)
    for r in range(1, min(sh.max_row,8)+1):
        vals=[]
        for c in range(1, sh.max_column+1):
            cell=sh.cell(r,c)
            vals.append((cell.coordinate, cell.value, cell.number_format))
        print(vals)
