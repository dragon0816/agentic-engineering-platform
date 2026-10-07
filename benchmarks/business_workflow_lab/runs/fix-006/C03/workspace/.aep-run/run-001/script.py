from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('Workspace:', ws)
print('Input files:')
for p in sorted((ws/'input').iterdir()):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Inspect template workbook
xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'\nSheet {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        print(r)

# Inspect fx csv
csvp = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv preview:')
with csvp.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i,row in zip(range(8), reader):
        print(row)

# Inspect manual
manual = ws/'input'/'quote_cli_manual.md'
print('\nquote_cli_manual.md first 120 lines:')
text = manual.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:120], 1):
    print(f'{i:03}: {line}')