from pathlib import Path
import subprocess, sys, json, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUTFILE = OUTPUT / 'quote_template_filled.xlsx'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')

PROJECTS = ['P001','P002','P003','P004','P005']
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Accepted'
GENERATED_AT = '2026-10-06T00:00:00+08:00'
FETCHED_AT = datetime.now().astimezone().isoformat(timespec='seconds')
TWOPLACES = Decimal('0.01')

def money(x):
    return Decimal(str(x)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def run_cli(args):
    cmd = [sys.executable, str(CLI)] + args
    res = subprocess.run(cmd, text=True, capture_output=True)
    if res.returncode != 0:
        raise RuntimeError(f'CLI failed ({res.returncode}) for {args}: {res.stderr}')
    return json.loads(res.stdout)

def read_fx(path):
    rates = {}
    with path.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        header = next(reader)
        if header[:2] != ['Currency','RateToTWD']:
            raise ValueError(f'Unexpected FX header: {header}')
        for row in reader:
            if not row or not row[0]:
                continue
            rates[row[0]] = Decimal(row[1])
    return rates

def clear_data(sh, start_row, max_col):
    if sh.max_row >= start_row:
        for row in sh.iter_rows(min_row=start_row, max_row=sh.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

def copy_row_style(sh, src_row, dst_row, max_col):
    # If the template has no row 5 style, row 4 header is still not copied; set formats explicitly later.
    for col in range(1, max_col + 1):
        src = sh.cell(src_row, col)
        dst = sh.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

# Fetch quote headers and line items.
rates = read_fx(FX_FILE)
headers = []
line_items = []
fetch_counts = {}
for project in PROJECTS:
    listed = run_cli(['quotes','list','--project',project,'--from',FROM_DATE,'--to',TO_DATE,'--status',STATUS,'--format','json'])
    if listed.get('schema_version') != 1:
        raise ValueError(f'Unexpected schema_version for {project}: {listed.get("schema_version")}')
    items = listed.get('items', [])
    fetch_counts[project] = len(items)
    for item in items:
        if item['ProjectID'] != project or item['Status'] != STATUS or not (FROM_DATE <= item['QuoteDate'] <= TO_DATE):
            raise ValueError(f'CLI returned out-of-filter item: {item}')
        curr = item['Currency']
        if curr not in rates:
            raise ValueError(f'Missing FX rate for currency {curr}')
        net_twd = money(Decimal(item['NetAmount']) * rates[curr])
        row = dict(item)
        row['NetAmountTWD'] = net_twd
        headers.append(row)
        detail = run_cli(['quotes','get',item['QuoteID'],'--format','json'])
        if detail.get('schema_version') != 1:
            raise ValueError(f'Unexpected get schema_version for {item["QuoteID"]}')
        if detail['header']['QuoteID'] != item['QuoteID']:
            raise ValueError(f'Header mismatch for {item["QuoteID"]}')
        for li in detail.get('line_items', []):
            lirow = dict(li)
            lirow['QuoteID'] = item['QuoteID']
            line_items.append(lirow)

# Sort for stable output.
headers.sort(key=lambda r: (r['ProjectID'], r['QuoteDate'], r['QuoteID']))
line_items.sort(key=lambda r: (r['QuoteID'], r['ProductID']))

# Copy template and fill.
shutil.copy2(TEMPLATE, OUTFILE)
wb = load_workbook(OUTFILE)

# Report metadata.
rep = wb['Report']
rep['B3'] = FROM_DATE
rep['D3'] = TO_DATE
rep['B4'] = 'ProjectID P001-P005; QuoteDate 2026-09-01..2026-09-30 inclusive; Status Accepted; NetAmountTWD = NetAmount × fx_rates.csv RateToTWD'
rep['B5'] = GENERATED_AT
rep['B7'] = 'quote_cli_mock; fx_rates.csv'
rep['D7'] = FETCHED_AT
missing_projects = [p for p in PROJECTS if fetch_counts.get(p, 0) == 0]
rep['B9'] = 'No Accepted quotes retained as zero in Summary: ' + ', '.join(missing_projects) if missing_projects else ''

# Quotes sheet.
qsh = wb['Quotes']
clear_data(qsh, 5, 11)
quote_headers = [qsh.cell(4, c).value for c in range(1, 12)]
expected_quote_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
if quote_headers != expected_quote_headers:
    raise ValueError(f'Unexpected Quotes headers: {quote_headers}')
for idx, item in enumerate(headers, start=5):
    qsh.cell(idx, 1, item['QuoteID']); qsh.cell(idx, 1).number_format='@'
    qsh.cell(idx, 2, item['CustomerID']); qsh.cell(idx, 2).number_format='@'
    qsh.cell(idx, 3, item['ProjectID']); qsh.cell(idx, 3).number_format='@'
    qsh.cell(idx, 4, item['QuoteDate']); qsh.cell(idx, 4).number_format='yyyy-mm-dd'
    qsh.cell(idx, 5, item['ValidUntil']); qsh.cell(idx, 5).number_format='yyyy-mm-dd'
    qsh.cell(idx, 6, item['Status'])
    qsh.cell(idx, 7, item['Currency'])
    qsh.cell(idx, 8, float(money(item['Subtotal']))); qsh.cell(idx, 8).number_format='0.00'
    qsh.cell(idx, 9, float(Decimal(item['DiscountPercent']))); qsh.cell(idx, 9).number_format='0'
    qsh.cell(idx, 10, float(money(item['NetAmount']))); qsh.cell(idx, 10).number_format='0.00'
    qsh.cell(idx, 11, float(item['NetAmountTWD'])); qsh.cell(idx, 11).number_format='0.00'

# Lines sheet.
lsh = wb['Lines']
clear_data(lsh, 5, 6)
line_headers = [lsh.cell(4, c).value for c in range(1, 7)]
expected_line_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if line_headers != expected_line_headers:
    raise ValueError(f'Unexpected Lines headers: {line_headers}')
for idx, li in enumerate(line_items, start=5):
    lsh.cell(idx,1,li['QuoteID']); lsh.cell(idx,1).number_format='@'
    lsh.cell(idx,2,li['ProductID']); lsh.cell(idx,2).number_format='@'
    lsh.cell(idx,3,li['ProductName'])
    lsh.cell(idx,4,int(li['Quantity']) if Decimal(li['Quantity']) == Decimal(li['Quantity']).to_integral_value() else float(Decimal(li['Quantity'])))
    lsh.cell(idx,5,float(money(li['UnitPrice']))); lsh.cell(idx,5).number_format='0.00'
    lsh.cell(idx,6,float(money(li['LineAmount']))); lsh.cell(idx,6).number_format='0.00'

# Summary by ProjectID. Amount is summed only within each ProjectID/currency bucket; current data has <=1 currency per project.
ssh = wb['Summary']
clear_data(ssh, 5, 6)
summary_rows = []
for p in PROJECTS:
    proj_quotes = [h for h in headers if h['ProjectID'] == p]
    if not proj_quotes:
        summary_rows.append({'Group': p, 'SubGroup':'', 'Currency':'', 'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0})
    else:
        by_currency = {}
        for h in proj_quotes:
            curr = h['Currency']
            by_currency.setdefault(curr, {'Amount':Decimal('0.00'), 'AmountTWD':Decimal('0.00'), 'Count':0})
            by_currency[curr]['Amount'] += money(h['NetAmount'])
            by_currency[curr]['AmountTWD'] += h['NetAmountTWD']
            by_currency[curr]['Count'] += 1
        for curr in sorted(by_currency):
            vals = by_currency[curr]
            summary_rows.append({'Group': p, 'SubGroup':'', 'Currency':curr, 'Amount': money(vals['Amount']), 'AmountTWD': money(vals['AmountTWD']), 'Count': vals['Count']})
for idx, r in enumerate(summary_rows, start=5):
    ssh.cell(idx,1,r['Group']); ssh.cell(idx,1).number_format='@'
    ssh.cell(idx,2,r['SubGroup'])
    ssh.cell(idx,3,r['Currency'])
    ssh.cell(idx,4,float(r['Amount'])); ssh.cell(idx,4).number_format='0.00'
    ssh.cell(idx,5,float(r['AmountTWD'])); ssh.cell(idx,5).number_format='0.00'
    ssh.cell(idx,6,int(r['Count'])); ssh.cell(idx,6).number_format='0'

last_quote_row = 4 + max(1, len(headers))
ssh['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
ssh['B2'].number_format = '#,##0.00'

# Preserve filters if absent no-op; set data ranges for sheets with rows.
q_last = 4 + max(1, len(headers))
l_last = 4 + max(1, len(line_items))
s_last = 4 + max(1, len(summary_rows))
qsh.auto_filter.ref = f'A4:K{q_last}'
lsh.auto_filter.ref = f'A4:F{l_last}'
ssh.auto_filter.ref = f'A4:F{s_last}'

wb.save(OUTFILE)

# Reopen and verify output requirements.
wbv = load_workbook(OUTFILE, data_only=False)
errors = []
if wbv.sheetnames != ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']:
    errors.append(f'sheet names changed: {wbv.sheetnames}')
qv, lv, sv, rv = wbv['Quotes'], wbv['Lines'], wbv['Summary'], wbv['Report']
if [qv.cell(4,c).value for c in range(1,12)] != expected_quote_headers:
    errors.append('Quotes column names mismatch')
if [lv.cell(4,c).value for c in range(1,7)] != expected_line_headers:
    errors.append('Lines column names mismatch')
if [sv.cell(4,c).value for c in range(1,7)] != ['Group','SubGroup','Currency','Amount','AmountTWD','Count']:
    errors.append('Summary column names mismatch')
if rv['B3'].value != FROM_DATE or rv['D3'].value != TO_DATE:
    errors.append('Report date range mismatch')
if sv['B2'].value != f'=SUM(Quotes!K5:K{last_quote_row})':
    errors.append(f'Summary B2 formula mismatch: {sv["B2"].value}')
# Collect rows.
qrows = []
for r in range(5, 5+len(headers)):
    qrows.append({expected_quote_headers[c-1]: qv.cell(r,c).value for c in range(1,12)})
if len(qrows) != len(headers):
    errors.append('Quote row count mismatch')
for r in qrows:
    if r['ProjectID'] not in PROJECTS or r['Status'] != STATUS or not (FROM_DATE <= r['QuoteDate'] <= TO_DATE):
        errors.append(f'Quote filter mismatch row {r}')
    if not isinstance(r['QuoteID'], str) or not isinstance(r['CustomerID'], str) or not isinstance(r['ProjectID'], str):
        errors.append(f'ID type mismatch row {r}')
    expected_twd = money(Decimal(str(r['NetAmount'])) * rates[r['Currency']])
    actual_twd = money(Decimal(str(r['NetAmountTWD'])))
    if actual_twd != expected_twd:
        errors.append(f'NetAmountTWD mismatch for {r["QuoteID"]}: {actual_twd} vs {expected_twd}')
# Summary rows exactly P001-P005 once each for current data.
srows = []
for r in range(5, 5+len(summary_rows)):
    srows.append({['Group','SubGroup','Currency','Amount','AmountTWD','Count'][c-1]: sv.cell(r,c).value for c in range(1,7)})
if [r['Group'] for r in srows] != PROJECTS:
    errors.append(f'Summary project rows mismatch: {[r["Group"] for r in srows]}')
for p in PROJECTS:
    sr = next(r for r in srows if r['Group'] == p)
    proj = [q for q in qrows if q['ProjectID'] == p]
    if not proj:
        if money(sr['Amount']) != Decimal('0.00') or money(sr['AmountTWD']) != Decimal('0.00') or sr['Count'] != 0:
            errors.append(f'Zero project summary mismatch for {p}: {sr}')
    else:
        currencies = sorted(set(q['Currency'] for q in proj))
        if len(currencies) != 1:
            errors.append(f'Unexpected multi-currency project in one-row check {p}: {currencies}')
        curr = currencies[0]
        exp_amount = money(sum((money(q['NetAmount']) for q in proj), Decimal('0.00')))
        exp_twd = money(sum((money(q['NetAmountTWD']) for q in proj), Decimal('0.00')))
        if sr['Currency'] != curr or money(sr['Amount']) != exp_amount or money(sr['AmountTWD']) != exp_twd or sr['Count'] != len(proj):
            errors.append(f'Summary mismatch for {p}: {sr}; expected {curr} {exp_amount} {exp_twd} {len(proj)}')
# Line rows verify all QuoteID in Quotes and numeric values.
quote_ids = {q['QuoteID'] for q in qrows}
line_count = len(line_items)
for r in range(5, 5+line_count):
    qid = lv.cell(r,1).value
    if qid not in quote_ids:
        errors.append(f'Line row {r} QuoteID {qid} missing from Quotes')
    if not isinstance(lv.cell(r,2).value, str):
        errors.append(f'Line row {r} ProductID is not text')
    for c in [4,5,6]:
        if not isinstance(lv.cell(r,c).value, (int,float)):
            errors.append(f'Line row {r} col {c} not numeric: {lv.cell(r,c).value!r}')
for r in range(5, 5+len(headers)):
    for c in [8,9,10,11]:
        if not isinstance(qv.cell(r,c).value, (int,float)):
            errors.append(f'Quote row {r} col {c} not numeric: {qv.cell(r,c).value!r}')
for r in range(5, 5+len(summary_rows)):
    for c in [4,5,6]:
        if not isinstance(sv.cell(r,c).value, (int,float)):
            errors.append(f'Summary row {r} col {c} not numeric: {sv.cell(r,c).value!r}')

if errors:
    print('VALIDATION FAILED')
    for e in errors:
        print(' -', e)
    raise SystemExit(1)

print('Created:', OUTFILE)
print('FX rates used:', {k: str(v) for k,v in rates.items()})
print('Fetched Accepted quotes by project:', fetch_counts)
print('Quote rows:', len(headers))
print('Line rows:', len(line_items))
print('Summary rows:', len(summary_rows))
print('Summary B2 formula:', sv['B2'].value)
print('Quotes:')
for r in qrows:
    print(' ', r)
print('Summary:')
for r in srows:
    print(' ', r)
print('Validation checked: workbook sheet names, template headers, report dates, B2 formula range, quote filters, text IDs, numeric amount columns, NetAmountTWD by FX, Summary P001-P005 rows including zero projects, and Lines QuoteID joins.')