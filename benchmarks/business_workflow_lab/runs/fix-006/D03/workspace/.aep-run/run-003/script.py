from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv, shutil
from copy import copy
from openpyxl import load_workbook

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)
export_path = output_dir / 'expense_export.csv'
fx_path = input_dir / 'fx_rates.csv'
template_path = input_dir / 'expense_template.xlsx'
out_path = output_dir / 'expense_report.xlsx'

print('Inputs:')
for p in [template_path, fx_path, export_path]:
    print(f' - {p}: exists={p.exists()} size={p.stat().st_size if p.exists() else None}')

# Read fixed FX rates (only Currency and RateToTWD columns are relevant; notes may contain commas)
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('FX headers:', reader.fieldnames)
    for row in reader:
        cur = (row.get('Currency') or '').strip()
        rate = (row.get('RateToTWD') or '').strip()
        if cur:
            fx[cur] = Decimal(rate)
print('FX rates:', {k: str(v) for k, v in fx.items()})

# Read exported Expense data
with export_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('Expense export headers:', reader.fieldnames)
    raw_rows = list(reader)
print('Export row count:', len(raw_rows))
print('Expense sample:', raw_rows[:2])

from_date = '2026-09-01'
to_date = '2026-09-30'
condition = 'ExpenseDate 2026-09-01 to 2026-09-30; Status=Approved'
now = datetime.now().astimezone().replace(microsecond=0).isoformat()

def q2(x: Decimal) -> Decimal:
    return x.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

processed = []
for r in raw_rows:
    if r['Status'] != 'Approved':
        raise ValueError(f"Unexpected status in export: {r}")
    if not (from_date <= r['ExpenseDate'] <= to_date):
        raise ValueError(f"Unexpected date outside period: {r}")
    cur = r['Currency']
    if cur not in fx:
        raise ValueError(f"Missing FX rate for {cur}")
    amount = q2(Decimal(r['Amount']))
    amount_twd = q2(amount * fx[cur])
    row = dict(r)
    row['AmountDecimal'] = amount
    row['AmountTWDDecimal'] = amount_twd
    row['NeedsReview'] = 'true' if amount_twd > Decimal('5000') else 'false'
    row['MissingReceipt'] = 'true' if str(r.get('MissingReceipt','')).strip().lower() == 'true' else 'false'
    processed.append(row)

# Copy template and fill it, preserving workbook sheets/layout.
shutil.copy2(template_path, out_path)
wb = load_workbook(out_path)

# Helper: clear a rectangular data area from row 5 down, preserving template row 5 style if present.
def clear_data(sheet, min_row=5, max_col=None):
    if max_col is None:
        max_col = sheet.max_column
    if sheet.max_row >= min_row:
        for row in sheet.iter_rows(min_row=min_row, max_row=sheet.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

# Helper copy style from a template row cell-by-cell to newly used row.
def copy_row_style(sheet, src_row, dst_row, max_col):
    for col in range(1, max_col+1):
        src = sheet.cell(src_row, col)
        dst = sheet.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

# Report sheet
report = wb['Report']
report['B3'] = from_date
report['D3'] = to_date
report['B4'] = condition
report['B5'] = now
report['B7'] = 'Expense Viewer (synthetic lab); fx_rates.csv'
report['D7'] = now
report['B9'] = 'AmountTWD calculated using fixed rates in fx_rates.csv. NeedsReview=true when AmountTWD > 5000.'

# Details sheet
details = wb['Details']
headers = [details.cell(4, c).value for c in range(1, details.max_column+1)]
print('Details template headers:', headers)
clear_data(details, 5, 14)
for i, r in enumerate(processed, start=5):
    copy_row_style(details, 5, i, 14)
    values = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['ExpenseDate'],
        r['Category'], r['Currency'], float(r['AmountDecimal']), float(r['AmountTWDDecimal']), r['Status'],
        r['ReceiptID'], r['Description'], r['MissingReceipt'], r['NeedsReview']
    ]
    for c, v in enumerate(values, start=1):
        cell = details.cell(i, c)
        cell.value = v
    # Required formats
    for c in [1,2,3,4,6,7,10,11,12,13,14]:
        details.cell(i, c).number_format = '@'
    details.cell(i, 5).number_format = 'yyyy-mm-dd'
    details.cell(i, 8).number_format = '0.00'
    details.cell(i, 9).number_format = '0.00'

last_detail = 4 + len(processed)

# Summary: group by EmployeeID and Category, with Currency included so original currency amounts are not mixed.
summary = wb['Summary']
clear_data(summary, 5, 6)
summary['B2'] = f'=SUM(Details!I5:I{last_detail if processed else 5})'
summary_groups = {}
for r in processed:
    key = (r['EmployeeID'], r['Category'], r['Currency'])
    if key not in summary_groups:
        summary_groups[key] = {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0}
    summary_groups[key]['amount'] += r['AmountDecimal']
    summary_groups[key]['amount_twd'] += r['AmountTWDDecimal']
    summary_groups[key]['count'] += 1

for i, (key, agg) in enumerate(sorted(summary_groups.items()), start=5):
    copy_row_style(summary, 5, i, 6)
    emp, cat, cur = key
    values = [emp, cat, cur, float(q2(agg['amount'])), float(q2(agg['amount_twd'])), agg['count']]
    for c, v in enumerate(values, start=1):
        summary.cell(i, c).value = v
    for c in [1,2,3]:
        summary.cell(i, c).number_format = '@'
    summary.cell(i, 4).number_format = '0.00'
    summary.cell(i, 5).number_format = '0.00'
    summary.cell(i, 6).number_format = '0'

# Exceptions: list rows that need review or are missing receipt, while retaining them in Details.
exceptions = wb['Exceptions']
clear_data(exceptions, 5, 4)
exception_rows = []
for r in processed:
    reasons = []
    if r['NeedsReview'] == 'true':
        reasons.append('AmountTWD > 5000')
    if r['MissingReceipt'] == 'true':
        reasons.append('MissingReceipt=true')
    if reasons:
        exception_rows.append((r['ExpenseID'], r['EmployeeID'], r['ExpenseDate'], '; '.join(reasons)))
for i, vals in enumerate(exception_rows, start=5):
    copy_row_style(exceptions, 5, i, 4)
    for c, v in enumerate(vals, start=1):
        exceptions.cell(i, c).value = v
    exceptions.cell(i, 1).number_format = '@'
    exceptions.cell(i, 2).number_format = '@'
    exceptions.cell(i, 3).number_format = 'yyyy-mm-dd'
    exceptions.cell(i, 4).number_format = '@'

wb.save(out_path)
print('Wrote', out_path)
print('Processed rows:', len(processed))
print('Summary rows:', len(summary_groups))
print('Exception rows:', len(exception_rows))
print('NeedsReview true:', sum(1 for r in processed if r['NeedsReview']=='true'))
print('NeedsReview false:', sum(1 for r in processed if r['NeedsReview']=='false'))
