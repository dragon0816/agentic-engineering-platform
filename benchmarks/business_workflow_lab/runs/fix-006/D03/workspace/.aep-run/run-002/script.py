from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
exp=ws/'output'/'expense_export.csv'
print('export exists', exp.exists(), exp)
if exp.exists():
    for enc in ['utf-8-sig','utf-8','cp950']:
        try:
            with exp.open(newline='', encoding=enc) as f:
                rows=list(csv.reader(f))
            print('encoding', enc, 'rows', len(rows), 'headers', rows[0] if rows else None, 'sample', rows[1:4])
            break
        except Exception as e:
            print('encoding failed', enc, repr(e))
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
inst=wb['Instructions']
print('Instructions rows:')
for row in inst.iter_rows(values_only=True):
    if any(v is not None for v in row): print(row)
