from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', sorted([p.name for p in (ws/'input').iterdir()]))
# CSV info
csv_path=ws/'input'/'fx_rates.csv'
with csv_path.open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    rows=list(r)
print('fx_rates.csv rows', len(rows))
print('fx_rates headers', rows[0] if rows else None)
print('fx_rates sample', rows[1:6])
# manual
manual=ws/'input'/'expense_manual.md'
print('manual exists', manual.exists())
print('manual preview:')
print(manual.read_text(encoding='utf-8')[:3000])
# workbook
wb_path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path, data_only=False)
print('workbook sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    vals=[]
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        vals.append(row)
    print('preview', vals)