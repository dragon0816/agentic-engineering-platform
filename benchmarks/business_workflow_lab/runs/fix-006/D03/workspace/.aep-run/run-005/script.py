from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
import csv
from openpyxl import load_workbook

ws=Path.cwd()
out=ws/'output'/'expense_report.xlsx'
export=ws/'output'/'expense_export.csv'
fx_path=ws/'input'/'fx_rates.csv'
print('Reopening output:', out, 'exists=', out.exists(), 'size=', out.stat().st_size if out.exists() else None)

fx={}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    for row in csv.DictReader(f):
        fx[row['Currency'].strip()] = Decimal(row['RateToTWD'].strip())
with export.open(newline='', encoding='utf-8-sig') as f:
    exported=list(csv.DictReader(f))

wb=load_workbook(out, data_only=False)
assert wb.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions'], wb.sheetnames
print('Sheet names OK:', wb.sheetnames)

def blank_norm(v):
    return '' if v is None else str(v)

q2=lambda d: d.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# Details checks
det=wb['Details']
expected_headers=['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
headers=[det.cell(4,c).value for c in range(1,15)]
assert headers==expected_headers, headers
rows=[]
r=5
while det.cell(r,1).value is not None:
    rows.append([det.cell(r,c).value for c in range(1,15)])
    r+=1
assert len(rows)==len(exported)==15, (len(rows), len(exported))
print('Details headers and row count OK:', len(rows))

needs_true=needs_false=0
for i,(outrow,src) in enumerate(zip(rows, exported), start=5):
    for idx, name in enumerate(['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency'], start=0):
        assert blank_norm(outrow[idx]) == src[name], (i, name, outrow[idx], src[name])
    assert outrow[9] == 'Approved', (i, outrow[9])
    assert blank_norm(outrow[10]) == src['ReceiptID'], (i, outrow[10], src['ReceiptID'])
    assert blank_norm(outrow[11]) == src['Description'], (i, outrow[11], src['Description'])
    assert str(outrow[12]) in ['true','false'], (i, outrow[12])
    assert str(outrow[13]) in ['true','false'], (i, outrow[13])
    amount=Decimal(str(outrow[7]))
    assert amount == q2(Decimal(src['Amount'])), (i, amount, src['Amount'])
    amount_twd=Decimal(str(outrow[8])).quantize(Decimal('0.01'))
    expected=q2(Decimal(src['Amount']) * fx[src['Currency']])
    assert amount_twd == expected, (i, amount_twd, expected)
    expected_review = 'true' if expected > Decimal('5000') else 'false'
    assert outrow[13] == expected_review, (i, outrow[13], expected_review)
    if expected_review == 'true': needs_true += 1
    else: needs_false += 1
print('Details source preservation, blank missing values, numeric Amount/AmountTWD, fixed-FX conversion, and NeedsReview rule OK')
print('NeedsReview true/false:', needs_true, needs_false)

# Summary checks
sumsh=wb['Summary']
assert [sumsh.cell(4,c).value for c in range(1,7)] == ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
assert sumsh['B2'].value == '=SUM(Details!I5:I19)', sumsh['B2'].value
expected={}
for src in exported:
    key=(src['EmployeeID'], src['Category'], src['Currency'])
    expected.setdefault(key, [Decimal('0'), Decimal('0'), 0])
    amt=q2(Decimal(src['Amount']))
    expected[key][0]+=amt
    expected[key][1]+=q2(amt*fx[src['Currency']])
    expected[key][2]+=1
summary_rows=[]
r=5
while sumsh.cell(r,1).value is not None:
    key=(sumsh.cell(r,1).value, sumsh.cell(r,2).value, sumsh.cell(r,3).value)
    vals=(Decimal(str(sumsh.cell(r,4).value)).quantize(Decimal('0.01')), Decimal(str(sumsh.cell(r,5).value)).quantize(Decimal('0.01')), int(sumsh.cell(r,6).value))
    summary_rows.append((key, vals))
    assert key in expected, (r, key)
    exp=(expected[key][0].quantize(Decimal('0.01')), expected[key][1].quantize(Decimal('0.01')), expected[key][2])
    assert vals==exp, (r, key, vals, exp)
    r+=1
assert len(summary_rows)==len(expected)==14, (len(summary_rows), len(expected))
print('Summary by EmployeeID/Category/Currency OK:', len(summary_rows), 'formula', sumsh['B2'].value)

# Exceptions checks
exc=wb['Exceptions']
assert [exc.cell(4,c).value for c in range(1,5)] == ['ExpenseID','EmployeeID','ExpenseDate','Reason']
expected_exceptions=[]
for src in exported:
    amt_twd=q2(Decimal(src['Amount'])*fx[src['Currency']])
    reasons=[]
    if amt_twd > Decimal('5000'): reasons.append('AmountTWD > 5000')
    if str(src['MissingReceipt']).strip().lower()=='true': reasons.append('MissingReceipt=true')
    if reasons:
        expected_exceptions.append((src['ExpenseID'],src['EmployeeID'],src['ExpenseDate'],'; '.join(reasons)))
actual=[]
r=5
while exc.cell(r,1).value is not None:
    actual.append(tuple(exc.cell(r,c).value for c in range(1,5)))
    r+=1
assert actual==expected_exceptions, (actual, expected_exceptions)
print('Exceptions OK:', len(actual))

# Report checks
rep=wb['Report']
assert rep['B3'].value=='2026-09-01'
assert rep['D3'].value=='2026-09-30'
assert 'Approved' in rep['B4'].value
assert 'Expense Viewer' in rep['B7'].value and 'fx_rates.csv' in rep['B7'].value
print('Report metadata OK')
print('All stated requirements verified from saved output file.')
