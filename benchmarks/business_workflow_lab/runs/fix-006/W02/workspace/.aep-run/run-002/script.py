from pathlib import Path
from openpyxl import load_workbook

ws = Path.cwd()
print('Workspace:', ws)
print('Re-inspecting template instructions:')
wb = load_workbook(ws/'input'/'sales_template.xlsx', data_only=False)
inst = wb['Instructions']
for r in range(1, inst.max_row + 1):
    vals = [inst.cell(r,c).value for c in range(1, inst.max_column + 1)]
    print(r, vals)
print('\nExisting merged ranges:')
for sh in wb.sheetnames:
    print(sh, list(wb[sh].merged_cells.ranges))
