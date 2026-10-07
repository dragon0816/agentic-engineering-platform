from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta, date
from copy import copy
import csv
from openpyxl import load_workbook

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
output_dir.mkdir(exist_ok=True)

print('Workspace:', ws)
print('Inputs present:', sorted(p.name for p in input_dir.glob('*')))

# Inputs captured from Lab ERP via browser detail pages. Only exact CustomerID C001 from list filter.
orders = [
    {'OrderID':'0045000101','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-08-03','Status':'Closed','Currency':'TWD','NetAmount':'95895.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET15'},
    {'OrderID':'0045000107','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P003','OrderDate':'2026-08-27','Status':'Closed','Currency':'USD','NetAmount':'6250.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET30'},
    {'OrderID':'0045000109','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P003','OrderDate':'2026-09-01','Status':'Open','Currency':'TWD','NetAmount':'150000.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET10'},
    {'OrderID':'0045000114','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-09-08','Status':'Cancelled','Currency':'USD','NetAmount':'11956.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET5'},
    {'OrderID':'0045000119','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-09-15','Status':'Open','Currency':'EUR','NetAmount':'8775.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET30'},
    {'OrderID':'0045000125','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-09-25','Status':'Closed','Currency':'TWD','NetAmount':'386445.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET0'},
    {'OrderID':'0045000129','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-10-01','Status':'Open','Currency':'USD','NetAmount':'1602.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET30'},
    {'OrderID':'0045000132','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P003','OrderDate':'2026-10-09','Status':'Open','Currency':'TWD','NetAmount':'50400.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':''},
    {'OrderID':'0045000136','CustomerID':'C001','CustomerName':'Acme Precision Taiwan Ltd.','ProjectID':'P001','OrderDate':'2026-10-31','Status':'Open','Currency':'EUR','NetAmount':'2250.00','CustomerContact':'Grace Lin <grace.lin@acme-precision.example>','PaymentTerms':'NET5'},
]
print('Order sample:', orders[0])
print('Order count:', len(orders))

# Read FX rates; first two columns are the reliable fields in this lab CSV.
rates = {}
with (input_dir / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header:', header)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

for o in orders:
    amount = Decimal(o['NetAmount'])
    rate = rates[o['Currency']]
    o['NetAmountDecimal'] = amount
    o['NetAmountTWD'] = (amount * rate).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# Fill a copy of the provided template.
wb = load_workbook(input_dir / 'sales_template.xlsx')
report = wb['Report']
details = wb['Details']
summary = wb['Summary']

# Report metadata
start = '2026-08-01'
end = '2026-10-31'
tz = timezone(timedelta(hours=8))
now_iso = datetime.now(tz).replace(microsecond=0).isoformat()
fetched_iso = now_iso
report['B3'] = start
report['D3'] = end
report['B4'] = 'CustomerID = C001 (exact match); OrderDate 2026-08-01 through 2026-10-31; all statuses'
report['B5'] = now_iso
report['B7'] = 'Lab ERP Sales Orders; fx_rates.csv'
report['D7'] = fetched_iso
missing_payment = [o['OrderID'] for o in orders if not o['PaymentTerms']]
missing_contact = [o['OrderID'] for o in orders if not o['CustomerContact']]
missing_parts = []
if missing_payment:
    missing_parts.append('PaymentTerms blank: ' + ', '.join(missing_payment))
if missing_contact:
    missing_parts.append('CustomerContact blank: ' + ', '.join(missing_contact))
if not missing_parts:
    missing_parts.append('No missing CustomerContact or PaymentTerms')
report['B9'] = '; '.join(missing_parts) + '. Exact CustomerID filter C001 used; similarly named C007 Acme Precision Taiwan Limited was excluded.'

# Clear possible old data rows below headers without touching row 4 headers.
for sheet in (details, summary):
    if sheet.max_row >= 5:
        sheet.delete_rows(5, sheet.max_row - 4)

# Details rows. Copy row 4 style to inserted data cells to preserve template appearance as much as possible.
headers = [details.cell(4, c).value for c in range(1, details.max_column + 1)]
print('Details headers:', headers)
for idx, o in enumerate(orders, start=5):
    vals = [
        o['OrderID'], o['CustomerID'], o['CustomerName'], o['ProjectID'],
        date.fromisoformat(o['OrderDate']), o['Status'], o['Currency'],
        float(o['NetAmountDecimal']), float(o['NetAmountTWD']),
        o['CustomerContact'] or None, o['PaymentTerms'] or None,
    ]
    for col, val in enumerate(vals, start=1):
        cell = details.cell(idx, col, val)
        src = details.cell(4, col)
        if src.has_style:
            cell._style = copy(src._style)
        cell.font = copy(src.font)
        cell.fill = copy(src.fill)
        cell.border = copy(src.border)
        cell.alignment = copy(src.alignment)
        cell.protection = copy(src.protection)
        if col in (1,2,4):
            cell.number_format = '@'
        elif col == 5:
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,9):
            cell.number_format = '0.00'

# Summary grouped by Currency to avoid mixing original currencies.
summary_groups = []
for cur in sorted(set(o['Currency'] for o in orders)):
    group_orders = [o for o in orders if o['Currency'] == cur]
    amount = sum((o['NetAmountDecimal'] for o in group_orders), Decimal('0.00'))
    amount_twd = sum((o['NetAmountTWD'] for o in group_orders), Decimal('0.00'))
    summary_groups.append(('Currency', '', cur, amount, amount_twd, len(group_orders)))

for idx, (group, subgroup, cur, amount, amount_twd, count) in enumerate(summary_groups, start=5):
    vals = [group, subgroup or None, cur, float(amount), float(amount_twd), count]
    for col, val in enumerate(vals, start=1):
        cell = summary.cell(idx, col, val)
        src = summary.cell(4, col)
        if src.has_style:
            cell._style = copy(src._style)
        cell.font = copy(src.font)
        cell.fill = copy(src.fill)
        cell.border = copy(src.border)
        cell.alignment = copy(src.alignment)
        cell.protection = copy(src.protection)
        if col in (4,5):
            cell.number_format = '0.00'
        if col == 6:
            cell.number_format = '0'

last_row = 4 + len(orders)
summary['B2'] = f'=SUM(Details!I5:I{last_row})'

# Preserve sheets; save output.
outfile = output_dir / 'sales_report_C001_2026-08-01_2026-10-31.xlsx'
wb.save(outfile)
print('Saved:', outfile)

# Reopen and validate the stated requirements.
wb2 = load_workbook(outfile, data_only=False)
assert wb2.sheetnames == ['Report', 'Details', 'Summary', 'Instructions'], wb2.sheetnames
r2 = wb2['Report']; d2 = wb2['Details']; s2 = wb2['Summary']
assert r2['B3'].value == start and r2['D3'].value == end
assert 'CustomerID = C001' in r2['B4'].value
assert '0045000132' in r2['B9'].value and 'PaymentTerms blank' in r2['B9'].value
assert 'C007' in r2['B9'].value and 'excluded' in r2['B9'].value
assert d2.max_row == 13, d2.max_row
assert [d2.cell(4,c).value for c in range(1,12)] == ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
ids = []
for row in range(5, 14):
    ids.append(d2.cell(row, 1).value)
    assert d2.cell(row, 1).data_type == 's', f'OrderID not text at row {row}'
    assert d2.cell(row, 1).value.startswith('00')
    assert d2.cell(row, 2).value == 'C001', f'Non-C001 row {row}: {d2.cell(row,2).value}'
    assert d2.cell(row, 2).data_type == 's'
    assert start <= d2.cell(row, 5).value.isoformat() <= end
    assert isinstance(d2.cell(row, 8).value, (int, float))
    assert isinstance(d2.cell(row, 9).value, (int, float))
assert ids == [o['OrderID'] for o in orders], ids
assert '0045000105' not in ids, 'similar-name C007 order included unexpectedly'
# Missing PaymentTerms left as blank cell, not N/A/0.
row_0132 = ids.index('0045000132') + 5
assert d2.cell(row_0132, 11).value is None, d2.cell(row_0132, 11).value
assert all(d2.cell(row, 10).value for row in range(5,14)), 'A CustomerContact is missing unexpectedly'
assert s2['B2'].value == '=SUM(Details!I5:I13)', s2['B2'].value
assert s2.max_row == 7, s2.max_row
summary_rows = [[s2.cell(r,c).value for c in range(1,7)] for r in range(5,8)]
print('Summary rows:', summary_rows)
# Formula not evaluated by openpyxl, so validate detail sum directly.
total_twd = sum(Decimal(str(d2.cell(r,9).value)) for r in range(5,14)).quantize(Decimal('0.01'))
assert total_twd == Decimal('1702471.00'), total_twd
print('Validation passed: output reopened; sheets preserved; 9 C001 rows only; OrderID/CustomerID stored as text; dates in range; numeric amounts; missing PaymentTerms blank; Summary formula range and TWD total checked.')
