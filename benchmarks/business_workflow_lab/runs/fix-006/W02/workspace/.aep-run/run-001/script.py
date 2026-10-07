from pathlib import Path
import csv
from openpyxl import load_workbook

ws = Path.cwd()
print('Workspace:', ws)
print('Input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Inspect CSV
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nfx_rates.csv sample:')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i, row in zip(range(6), reader):
            print(i, row)

# Inspect manual beginning
manual = ws/'input'/'sap_manual.md'
if manual.exists():
    print('\nsap_manual.md first 120 lines:')
    text = manual.read_text(encoding='utf-8', errors='replace')
    for i, line in enumerate(text.splitlines()[:120], 1):
        print(f'{i:03d}: {line}')

# Inspect workbook sheets/headers/sample cells
xlsx = ws/'input'/'sales_template.xlsx'
if xlsx.exists():
    print('\nsales_template.xlsx structure:')
    wb = load_workbook(xlsx, data_only=False)
    for sh in wb.sheetnames:
        s = wb[sh]
        print('Sheet:', sh, 'max_row=', s.max_row, 'max_col=', s.max_column)
        for r in range(1, min(s.max_row, 8)+1):
            vals = [s.cell(r,c).value for c in range(1, min(s.max_column, 12)+1)]
            print(' ', r, vals)
