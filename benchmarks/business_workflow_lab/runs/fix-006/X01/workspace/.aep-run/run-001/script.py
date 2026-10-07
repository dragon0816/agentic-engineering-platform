from pathlib import Path
import csv, json, os
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size, 'bytes')

# CSV preview
fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\nFX CSV PREVIEW')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i,row in zip(range(6), r):
            print(i, row)

# manuals preview
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p = ws/'input'/name
    if p.exists():
        print(f'\n{name} PREVIEW')
        txt = p.read_text(encoding='utf-8', errors='replace')
        print(txt[:3000])

# workbook structure
wbp = ws/'input'/'project_template.xlsx'
if wbp.exists():
    print('\nWORKBOOK STRUCTURE')
    wb = load_workbook(wbp, data_only=False)
    for sh in wb.sheetnames:
        wsht = wb[sh]
        print('SHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
        for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row,8), values_only=True):
            print(row)
        print('---')
