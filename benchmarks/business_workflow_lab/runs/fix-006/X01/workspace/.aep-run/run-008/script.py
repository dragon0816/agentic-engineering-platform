from pathlib import Path
from openpyxl import load_workbook
ws=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-006\X01\workspace')
for path in [ws/'input'/'project_template.xlsx', ws/'output'/'project_integration_report.xlsx']:
    print('\nWORKBOOK', path.relative_to(ws))
    wb=load_workbook(path, data_only=False)
    for s in ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']:
        sh=wb[s]
        print('SHEET', s)
        maxr=min(sh.max_row,10); maxc=min(sh.max_column,12)
        for r in range(1,maxr+1):
            vals=[sh.cell(r,c).value for c in range(1,maxc+1)]
            print(r, vals)
        print('---')