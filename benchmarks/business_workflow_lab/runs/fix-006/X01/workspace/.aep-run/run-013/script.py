from pathlib import Path
import csv, json, os
from openpyxl import load_workbook

root = Path('.')
print('Workspace:', root.resolve())
print('Input files:')
for p in sorted((root/'input').glob('*')):
    print(' -', p, p.stat().st_size)
print('Output files:')
for p in sorted((root/'output').rglob('*')):
    if p.is_file():
        print(' -', p, p.stat().st_size)

# Print headers/sample rows for current generic outputs
for rel in ['output/downloads/sap_orders.csv','output/downloads/sap_invoices.csv','output/expense_approved.csv','output/project_summary_sep2026.csv']:
    p = root/rel
    if p.exists():
        print('\nCSV sample:', rel)
        with p.open(newline='', encoding='utf-8-sig') as f:
            rows = list(csv.reader(f))
        print(' rows including header:', len(rows))
        print(' header:', rows[0] if rows else None)
        for r in rows[1:4]: print(' sample:', r)

for rel in ['output/quotes_accepted.json','output/quotes_accepted_sep2026.json']:
    p = root/rel
    if p.exists():
        print('\nJSON sample:', rel)
        data = json.loads(p.read_text(encoding='utf-8'))
        print(' type:', type(data).__name__, 'len:', len(data) if hasattr(data,'__len__') else 'n/a')
        if isinstance(data, list):
            print(' first:', data[0] if data else None)
        elif isinstance(data, dict):
            print(' keys:', list(data)[:10])

p = root/'output/project_integration_report.xlsx'
if p.exists():
    print('\nWorkbook sheets:', p)
    wb = load_workbook(p, data_only=True)
    for ws in wb.worksheets:
        print(' sheet:', ws.title, 'max_row:', ws.max_row, 'max_col:', ws.max_column)
        vals = [[ws.cell(r,c).value for c in range(1, min(ws.max_column,8)+1)] for r in range(1, min(ws.max_row,4)+1)]
        for row in vals: print('  ', row)
