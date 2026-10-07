from pathlib import Path
import csv, json
from openpyxl import load_workbook

root = Path('.')
out = root/'output'
dl = out/'downloads'

def read_csv(path):
    with path.open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def ensure(cond, msg):
    if not cond:
        raise AssertionError(msg)

def load_quotes(path):
    data = json.loads(path.read_text(encoding='utf-8'))
    if isinstance(data, dict):
        return data.get('items', [])
    return data

print('Rechecking explicit August and September outputs')

aug_files = [
    dl/'sap_orders_aug2026.csv',
    dl/'sap_invoices_aug2026.csv',
    out/'expense_approved_aug2026.csv',
    out/'quotes_accepted_aug2026.json',
    out/'project_summary_aug2026.csv',
    out/'project_integration_report_aug2026.xlsx',
]
sep_files = [
    dl/'sap_orders_sep2026.csv',
    dl/'sap_invoices_sep2026.csv',
    out/'expense_approved_sep2026.csv',
    out/'quotes_accepted_sep2026.json',
    out/'project_summary_sep2026.csv',
    out/'project_integration_report_sep2026.xlsx',
]
for p in aug_files + sep_files:
    ensure(p.exists(), f'Missing file: {p}')
    print(' -', p, p.stat().st_size)

# August validation
orders = read_csv(dl/'sap_orders_aug2026.csv')
invoices = read_csv(dl/'sap_invoices_aug2026.csv')
expenses = read_csv(out/'expense_approved_aug2026.csv')
summary = read_csv(out/'project_summary_aug2026.csv')
quotes = load_quotes(out/'quotes_accepted_aug2026.json')

print('\nAugust counts:')
print(' orders', len(orders))
print(' invoices', len(invoices))
print(' expenses', len(expenses))
print(' quotes', len(quotes))
print(' summary rows', len(summary))
print(' quote sample:', quotes[:2])

ensure(len(orders) == 7, f'August orders count expected 7 got {len(orders)}')
ensure(all('2026-08-01' <= r['OrderDate'] <= '2026-08-31' for r in orders), 'August orders outside date range')
ensure(all(r['Status'] != 'Cancelled' for r in orders), 'August orders include Cancelled')
ensure(len(invoices) == 6, f'August invoices count expected 6 got {len(invoices)}')
ensure(all('2026-08-01' <= r['InvoiceDate'] <= '2026-08-31' for r in invoices), 'August invoices outside date range')
ensure(len(expenses) == 4, f'August approved expenses count expected 4 got {len(expenses)}')
ensure(all('2026-08-01' <= r['ExpenseDate'] <= '2026-08-31' for r in expenses), 'August expenses outside date range')
ensure(all(r['Status'] == 'Approved' for r in expenses), 'August expenses include non-Approved')
ensure(len(quotes) == 4, f'August accepted quotes count expected 4 got {len(quotes)}')
for q in quotes:
    status = q.get('Status') or q.get('status')
    qd = q.get('QuoteDate') or q.get('quote_date') or q.get('date')
    ensure(status in (None, 'Accepted'), f'August quote not Accepted: {q}')
    if qd:
        ensure('2026-08-01' <= qd <= '2026-08-31', f'August quote outside date range: {q}')
ensure(len(summary) == 5, f'August summary project rows expected 5 got {len(summary)}')
ensure(list(summary[0].keys()) == ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], 'August summary columns changed')

wb = load_workbook(out/'project_integration_report_aug2026.xlsx', data_only=True)
ensure(str(wb['Report']['B3'].value) == '2026-08-01', f'August report From mismatch: {wb["Report"]["B3"].value}')
ensure(str(wb['Report']['D3'].value) == '2026-08-31', f'August report To mismatch: {wb["Report"]["D3"].value}')
expected_rows = {'Orders': 4+len(orders), 'Invoices': 4+len(invoices), 'Expenses': 4+len(expenses), 'Quotes': 4+len(quotes), 'Summary': 4+len(summary)}
for sheet, expected in expected_rows.items():
    ensure(wb[sheet].max_row == expected, f'August {sheet} row count expected {expected} got {wb[sheet].max_row}')

# September validation restored
sep_orders = read_csv(dl/'sap_orders_sep2026.csv')
sep_invoices = read_csv(dl/'sap_invoices_sep2026.csv')
sep_expenses = read_csv(out/'expense_approved_sep2026.csv')
sep_summary = read_csv(out/'project_summary_sep2026.csv')
sep_quotes = load_quotes(out/'quotes_accepted_sep2026.json')
print('\nSeptember restored counts:')
print(' orders', len(sep_orders))
print(' invoices', len(sep_invoices))
print(' expenses', len(sep_expenses))
print(' quotes', len(sep_quotes))
print(' summary rows', len(sep_summary))

ensure(len(sep_orders) == 18, f'September orders expected 18 got {len(sep_orders)}')
ensure(all('2026-09-01' <= r['OrderDate'] <= '2026-09-30' for r in sep_orders), 'September orders outside date range')
ensure(all(r['Status'] != 'Cancelled' for r in sep_orders), 'September orders include Cancelled')
ensure(len(sep_invoices) == 13, f'September invoices expected 13 got {len(sep_invoices)}')
ensure(all('2026-09-01' <= r['InvoiceDate'] <= '2026-09-30' for r in sep_invoices), 'September invoices outside date range')
ensure(len(sep_expenses) == 15, f'September expenses expected 15 got {len(sep_expenses)}')
ensure(all(r['Status'] == 'Approved' for r in sep_expenses), 'September expenses include non-Approved')
ensure(len(sep_quotes) == 4, f'September accepted quotes expected 4 got {len(sep_quotes)}')
ensure(len(sep_summary) == 5, f'September summary rows expected 5 got {len(sep_summary)}')
sep_wb = load_workbook(out/'project_integration_report_sep2026.xlsx', data_only=True)
ensure(str(sep_wb['Report']['B3'].value) == '2026-09-01', f'September report From mismatch: {sep_wb["Report"]["B3"].value}')
ensure(str(sep_wb['Report']['D3'].value) == '2026-09-30', f'September report To mismatch: {sep_wb["Report"]["D3"].value}')

print('\nValidation passed: explicit August outputs exist, September _sep2026 outputs restored, and date/status/count checks passed.')
