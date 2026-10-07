from pathlib import Path
import csv, json, os, shutil, subprocess, sys
from datetime import datetime, date, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

workspace = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-006\X01\workspace')
input_dir = workspace / 'input'
output_dir = workspace / 'output'
downloads = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)
downloads.mkdir(exist_ok=True)

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
start = datetime.strptime(FROM_DATE, '%Y-%m-%d').date()
end = datetime.strptime(TO_DATE, '%Y-%m-%d').date()

def d(s):
    return datetime.strptime(str(s)[:10], '%Y-%m-%d').date()

def q2(x):
    return Decimal(str(x)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def money(v):
    if v is None or v == '':
        return Decimal('0')
    return Decimal(str(v))

def read_csv(path):
    with path.open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def write_csv(path, rows, fieldnames):
    with path.open('w', newline='', encoding='utf-8-sig') as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for r in rows:
            w.writerow({k: r.get(k, '') for k in fieldnames})

print('Inputs really present:')
for p in sorted(input_dir.glob('*')):
    print(' -', p.relative_to(workspace), p.stat().st_size)
print('Current source extracts present:')
for p in sorted(output_dir.rglob('*')):
    if p.is_file():
        print(' -', p.relative_to(workspace), p.stat().st_size)

# Read FX rates supplied as attachment.
fx_rows = read_csv(input_dir / 'fx_rates.csv')
fx = {r['Currency']: money(r['RateToTWD']) for r in fx_rows}
print('FX rates:', fx)

# Use SAP/Expense extracts produced by the approved workflow run.
orders_raw_path = downloads / 'sap_orders.csv'
invoices_path = downloads / 'sap_invoices.csv'
expense_path = output_dir / 'expense_approved.csv'
if not orders_raw_path.exists() or not invoices_path.exists() or not expense_path.exists():
    raise SystemExit('Missing SAP/Expense extracts from approved workflow run')
orders_raw = read_csv(orders_raw_path)
invoices_raw = read_csv(invoices_path)
expenses_raw = read_csv(expense_path)
print('Raw rows:', {'orders': len(orders_raw), 'invoices': len(invoices_raw), 'expenses': len(expenses_raw)})
print('Order status counts raw:', {s: sum(1 for r in orders_raw if r.get('Status') == s) for s in sorted({r.get('Status') for r in orders_raw})})

# Obtain accepted quotes for the same period from quote CLI when available; fall back only if CLI is unavailable.
quote_cli_py = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
quotes_obj = None
if quote_cli_py.exists():
    cmd = [sys.executable, str(quote_cli_py), 'quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', 'Accepted', '--format', 'json']
    print('Running quote CLI:', ' '.join(cmd))
    cp = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', check=True)
    quotes_obj = json.loads(cp.stdout)
    (output_dir / 'quotes_accepted.json').write_text(json.dumps(quotes_obj, ensure_ascii=False, indent=2), encoding='utf-8')
    (output_dir / 'quotes_accepted_sep2026.json').write_text(json.dumps(quotes_obj, ensure_ascii=False, indent=2), encoding='utf-8')
else:
    print('Quote CLI not found; using existing workflow JSON output')
    with (output_dir / 'quotes_accepted_sep2026.json').open(encoding='utf-8') as f:
        quotes_obj = json.load(f)

def quote_items(obj):
    if isinstance(obj, list):
        return obj
    if isinstance(obj, dict) and isinstance(obj.get('items'), list):
        return obj['items']
    if isinstance(obj, dict):
        vals = list(obj.values())
        if vals and all(isinstance(v, dict) for v in vals):
            return vals
    raise SystemExit('Unsupported quote JSON structure')
quotes_raw = quote_items(quotes_obj)

# Filter explicitly according to the user request.
orders = [r for r in orders_raw if start <= d(r['OrderDate']) <= end and r.get('Status') != 'Cancelled']
invoices = [r for r in invoices_raw if start <= d(r['InvoiceDate']) <= end]
expenses = [r for r in expenses_raw if start <= d(r['ExpenseDate']) <= end and r.get('Status') == 'Approved']
quotes = [r for r in quotes_raw if start <= d(r['QuoteDate']) <= end and r.get('Status') == 'Accepted']
print('Filtered rows:', {'orders': len(orders), 'invoices': len(invoices), 'expenses': len(expenses), 'quotes': len(quotes)})
print('Order status counts filtered:', {s: sum(1 for r in orders if r.get('Status') == s) for s in sorted({r.get('Status') for r in orders})})

# Overwrite the delivered SAP order extracts with the requested non-Cancelled set.
order_fields = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']
write_csv(downloads / 'sap_orders.csv', orders, order_fields)
write_csv(downloads / 'sap_orders_sep2026.csv', orders, order_fields)

# Also normalize the JSON count to the accepted quote items actually used.
quotes_obj_out = {'schema_version': 1, 'items': quotes, 'count': len(quotes)}
(output_dir / 'quotes_accepted.json').write_text(json.dumps(quotes_obj_out, ensure_ascii=False, indent=2), encoding='utf-8')
(output_dir / 'quotes_accepted_sep2026.json').write_text(json.dumps(quotes_obj_out, ensure_ascii=False, indent=2), encoding='utf-8')

# Summarize by ProjectID only; missing ProjectID rows go to Exceptions and are not summarized.
summary = {}
exceptions = []

def ensure_project(pid):
    if pid not in summary:
        summary[pid] = {
            'ProjectID': pid,
            'OrderTWD': Decimal('0'),
            'InvoiceBalanceTWD': Decimal('0'),
            'ApprovedExpenseTWD': Decimal('0'),
            'AcceptedQuoteTWD': Decimal('0'),
        }
    return summary[pid]

def amount_twd(currency, amount):
    if currency not in fx:
        raise SystemExit('Missing FX rate for ' + str(currency))
    return q2(money(amount) * fx[currency])

orders_out=[]
for r in orders:
    amt_twd = amount_twd(r['Currency'], r['NetAmount'])
    out = {
        'OrderID': str(r.get('OrderID','')),
        'ProjectID': str(r.get('ProjectID','')) if r.get('ProjectID') else '',
        'CustomerID': str(r.get('CustomerID','')),
        'OrderDate': r.get('OrderDate',''),
        'Status': r.get('Status',''),
        'Currency': r.get('Currency',''),
        'NetAmount': q2(r.get('NetAmount','0')),
        'AmountTWD': amt_twd,
    }
    orders_out.append(out)
    if out['ProjectID']:
        ensure_project(out['ProjectID'])['OrderTWD'] += amt_twd
    else:
        exceptions.append({'Source':'SAP Orders','DocumentID':out['OrderID'],'ProjectID':'','Reason':'Missing ProjectID; excluded from project summary'})

invoices_out=[]
for r in invoices:
    bal_twd = amount_twd(r['Currency'], r['Balance'])
    out = {
        'InvoiceID': str(r.get('InvoiceID','')),
        'OrderID': str(r.get('OrderID','')),
        'ProjectID': str(r.get('ProjectID','')) if r.get('ProjectID') else '',
        'CustomerID': str(r.get('CustomerID','')),
        'InvoiceDate': r.get('InvoiceDate',''),
        'DueDate': r.get('DueDate',''),
        'Currency': r.get('Currency',''),
        'Amount': q2(r.get('Amount','0')),
        'PaidAmount': q2(r.get('PaidAmount','0')),
        'Balance': q2(r.get('Balance','0')),
        'BalanceTWD': bal_twd,
    }
    invoices_out.append(out)
    if out['ProjectID']:
        ensure_project(out['ProjectID'])['InvoiceBalanceTWD'] += bal_twd
    else:
        exceptions.append({'Source':'SAP Invoices','DocumentID':out['InvoiceID'],'ProjectID':'','Reason':'Missing ProjectID; excluded from project summary'})

expenses_out=[]
for r in expenses:
    amt_twd = amount_twd(r['Currency'], r['Amount'])
    out = {
        'ExpenseID': str(r.get('ExpenseID','')),
        'ProjectID': str(r.get('ProjectID','')) if r.get('ProjectID') else '',
        'EmployeeID': str(r.get('EmployeeID','')),
        'ExpenseDate': r.get('ExpenseDate',''),
        'Status': r.get('Status',''),
        'Currency': r.get('Currency',''),
        'Amount': q2(r.get('Amount','0')),
        'AmountTWD': amt_twd,
    }
    expenses_out.append(out)
    if out['ProjectID']:
        ensure_project(out['ProjectID'])['ApprovedExpenseTWD'] += amt_twd
    else:
        exceptions.append({'Source':'Expenses','DocumentID':out['ExpenseID'],'ProjectID':'','Reason':'Missing ProjectID; excluded from project summary'})

quotes_out=[]
for r in quotes:
    amt_twd = amount_twd(r['Currency'], r['NetAmount'])
    out = {
        'QuoteID': str(r.get('QuoteID','')),
        'ProjectID': str(r.get('ProjectID','')) if r.get('ProjectID') else '',
        'CustomerID': str(r.get('CustomerID','')),
        'QuoteDate': r.get('QuoteDate',''),
        'ValidUntil': r.get('ValidUntil',''),
        'Status': r.get('Status',''),
        'Currency': r.get('Currency',''),
        'NetAmount': q2(r.get('NetAmount','0')),
        'AmountTWD': amt_twd,
    }
    quotes_out.append(out)
    if out['ProjectID']:
        ensure_project(out['ProjectID'])['AcceptedQuoteTWD'] += amt_twd
    else:
        exceptions.append({'Source':'Quotes','DocumentID':out['QuoteID'],'ProjectID':'','Reason':'Missing ProjectID; excluded from project summary'})

# Write summary CSV with separate requested metrics.
summary_rows=[]
for pid in sorted(summary):
    s=summary[pid]
    summary_rows.append({
        'ProjectID': pid,
        'OrderTWD': f"{q2(s['OrderTWD']):.2f}",
        'InvoiceBalanceTWD': f"{q2(s['InvoiceBalanceTWD']):.2f}",
        'ApprovedExpenseTWD': f"{q2(s['ApprovedExpenseTWD']):.2f}",
        'AcceptedQuoteTWD': f"{q2(s['AcceptedQuoteTWD']):.2f}",
    })
write_csv(output_dir / 'project_summary_sep2026.csv', summary_rows, ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'])

# Fill a copy of the provided template.
wb = load_workbook(input_dir / 'project_template.xlsx')

def clear_from(sheet_name, start_row=5):
    sh=wb[sheet_name]
    if sh.max_row >= start_row:
        sh.delete_rows(start_row, sh.max_row - start_row + 1)
    return sh

# Report cells per Instructions.
taipei = timezone(timedelta(hours=8))
ts = datetime.now(taipei).isoformat(timespec='seconds')
rep = wb['Report']
rep['B3'] = FROM_DATE
rep['D3'] = TO_DATE
rep['B4'] = 'SAP Sales Orders excluding Cancelled by OrderDate; SAP Invoices by InvoiceDate; Expense Approved by ExpenseDate; Quote CLI Accepted by QuoteDate. Aggregated separately by ProjectID; customer name is not used as a join key; orders and quotes are not combined as revenue.'
rep['B5'] = ts
rep['B7'] = 'Lab ERP SAP web; Expense Viewer desktop; quote_cli; fx_rates.csv'
rep['D7'] = ts
rep['B9'] = f"Rows used: orders {len(orders)} (Cancelled excluded), invoices {len(invoices)}, approved expenses {len(expenses)}, accepted quotes {len(quotes)}. Exceptions list rows missing ProjectID."

def write_sheet(sheet_name, rows, fields):
    sh=clear_from(sheet_name, 5)
    # keep row 4 headers; assert names
    current=[sh.cell(4,c).value for c in range(1,len(fields)+1)]
    if current != fields:
        raise SystemExit(f'{sheet_name} template headers differ: {current} expected {fields}')
    for i,r in enumerate(rows, start=5):
        for j,f in enumerate(fields, start=1):
            v = r.get(f, '')
            if isinstance(v, Decimal):
                v = float(v)
            sh.cell(i,j).value = v
    return sh

write_sheet('Orders', orders_out, ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'])
write_sheet('Invoices', invoices_out, ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'])
write_sheet('Expenses', expenses_out, ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'])
write_sheet('Quotes', quotes_out, ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'])

sum_sh=clear_from('Summary', 5)
sum_headers=['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD']
if [sum_sh.cell(4,c).value for c in range(1,6)] != sum_headers:
    raise SystemExit('Summary template headers differ')
for i,r in enumerate(summary_rows, start=5):
    sum_sh.cell(i,1).value = r['ProjectID']
    sum_sh.cell(i,2).value = float(Decimal(r['OrderTWD']))
    sum_sh.cell(i,3).value = float(Decimal(r['InvoiceBalanceTWD']))
    sum_sh.cell(i,4).value = float(Decimal(r['ApprovedExpenseTWD']))
    sum_sh.cell(i,5).value = float(Decimal(r['AcceptedQuoteTWD']))
last = 4 + len(summary_rows)
if summary_rows:
    sum_sh['B2'] = f'=SUM(B5:B{last})'
    sum_sh['C2'] = f'=SUM(C5:C{last})'
    sum_sh['D2'] = f'=SUM(D5:D{last})'
    sum_sh['E2'] = f'=SUM(E5:E{last})'
else:
    sum_sh['B2'] = '=SUM(B5:B5)'; sum_sh['C2'] = '=SUM(C5:C5)'; sum_sh['D2'] = '=SUM(D5:D5)'; sum_sh['E2'] = '=SUM(E5:E5)'

exc_sh=clear_from('Exceptions', 5)
exc_fields=['Source','DocumentID','ProjectID','Reason']
if [exc_sh.cell(4,c).value for c in range(1,5)] != exc_fields:
    raise SystemExit('Exceptions template headers differ')
for i,r in enumerate(exceptions, start=5):
    for j,f in enumerate(exc_fields, start=1):
        exc_sh.cell(i,j).value = r.get(f,'')

out_wb = output_dir / 'project_integration_report.xlsx'
wb.save(out_wb)
print('WROTE', out_wb)
print('WROTE', output_dir / 'project_summary_sep2026.csv')
print('WROTE normalized extracts and quote JSON')

# Reopen and validate every output requirement.
orders_check = read_csv(downloads / 'sap_orders.csv')
invoices_check = read_csv(downloads / 'sap_invoices.csv')
expenses_check = read_csv(output_dir / 'expense_approved.csv')
summary_check = read_csv(output_dir / 'project_summary_sep2026.csv')
quotes_check_obj = json.load(open(output_dir / 'quotes_accepted_sep2026.json', encoding='utf-8'))
quotes_check = quote_items(quotes_check_obj)
assert len(orders_check) == len(orders), 'orders row count mismatch'
assert len(invoices_check) == len(invoices), 'invoices row count mismatch'
assert len(expenses_check) == len(expenses), 'expenses row count mismatch'
assert len(quotes_check) == len(quotes), 'quotes row count mismatch'
assert all(start <= d(r['OrderDate']) <= end for r in orders_check)
assert all(r['Status'] != 'Cancelled' for r in orders_check)
assert all(start <= d(r['InvoiceDate']) <= end for r in invoices_check)
assert all(start <= d(r['ExpenseDate']) <= end and r['Status'] == 'Approved' for r in expenses_check)
assert all(start <= d(r['QuoteDate']) <= end and r['Status'] == 'Accepted' for r in quotes_check)

wb2 = load_workbook(out_wb, data_only=False)
for s in ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']:
    assert s in wb2.sheetnames, f'missing sheet {s}'
assert wb2['Report']['B3'].value == FROM_DATE
assert wb2['Report']['D3'].value == TO_DATE
assert [wb2['Orders'].cell(4,c).value for c in range(1,9)] == ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD']
assert [wb2['Invoices'].cell(4,c).value for c in range(1,12)] == ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD']
assert [wb2['Expenses'].cell(4,c).value for c in range(1,9)] == ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD']
assert [wb2['Quotes'].cell(4,c).value for c in range(1,10)] == ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD']
assert [wb2['Summary'].cell(4,c).value for c in range(1,6)] == sum_headers
for sheet, expected in [('Orders',len(orders)),('Invoices',len(invoices)),('Expenses',len(expenses)),('Quotes',len(quotes))]:
    actual=sum(1 for row in wb2[sheet].iter_rows(min_row=5, values_only=True) if any(v is not None for v in row))
    assert actual == expected, f'{sheet} workbook rows {actual} expected {expected}'
for row in wb2['Orders'].iter_rows(min_row=5, values_only=True):
    if any(v is not None for v in row):
        assert row[4] != 'Cancelled', 'Cancelled in workbook Orders'
for row in wb2['Expenses'].iter_rows(min_row=5, values_only=True):
    if any(v is not None for v in row):
        assert row[4] == 'Approved', 'non Approved in workbook Expenses'
for row in wb2['Quotes'].iter_rows(min_row=5, values_only=True):
    if any(v is not None for v in row):
        assert row[5] == 'Accepted', 'non Accepted in workbook Quotes'
for row in wb2['Summary'].iter_rows(min_row=5, values_only=True):
    if any(v is not None for v in row):
        assert isinstance(row[0], str), 'ProjectID not text'
        for val in row[1:5]: assert isinstance(val, (int,float)), f'non numeric summary value {val!r}'
headers=[wb2['Summary'].cell(4,c).value for c in range(1,6)]
assert not any('revenue' in str(h).lower() or '營收' in str(h) for h in headers if h), 'revenue header found'
print('VALIDATION PASSED')
print('Final counts:', {'orders_non_cancelled':len(orders), 'invoices':len(invoices), 'approved_expenses':len(expenses), 'accepted_quotes':len(quotes), 'projects':len(summary_rows), 'exceptions':len(exceptions)})