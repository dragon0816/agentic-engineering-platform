from pathlib import Path
import csv, json, os
from datetime import datetime
from openpyxl import load_workbook

ws = Path.cwd()
out = ws/'output'
print('WORKSPACE', ws)
print('OUTPUT EXISTS', out.exists())
for p in sorted(out.rglob('*')):
    if p.is_file():
        rel = p.relative_to(ws)
        st = p.stat()
        print(f'FILE {rel} size={st.st_size} mtime={datetime.fromtimestamp(st.st_mtime).isoformat(timespec="seconds")}')

# Inspect key CSV date ranges and status/project blanks without changing anything.
def inspect_csv(rel):
    p = ws/rel
    if not p.exists():
        print('MISSING', rel)
        return
    with p.open(newline='', encoding='utf-8-sig') as f:
        rows = list(csv.DictReader(f))
    print('CSV', rel, 'rows', len(rows), 'headers', rows[0].keys() if rows else [])
    if rows:
        print(' first', rows[0])
        print(' last', rows[-1])
        for date_col in ['OrderDate','InvoiceDate','ExpenseDate','QuoteDate']:
            if date_col in rows[0]:
                vals = sorted({r.get(date_col,'') for r in rows})
                print(' ', date_col, vals[:3], '...', vals[-3:])
        if 'Status' in rows[0]:
            print(' Status counts', {s: sum(1 for r in rows if r.get('Status')==s) for s in sorted({r.get('Status') for r in rows})})
        if 'ProjectID' in rows[0]:
            print(' blank ProjectID', sum(1 for r in rows if not (r.get('ProjectID') or '').strip()))

for rel in [
    'output/downloads/sap_orders.csv',
    'output/downloads/sap_invoices.csv',
    'output/downloads/sap_orders_sep2026.csv',
    'output/downloads/sap_invoices_sep2026.csv',
    'output/expense_approved.csv',
    'output/expense_approved_sep2026.csv',
    'output/project_summary_sep2026.csv',
]:
    inspect_csv(rel)

# Inspect workbook if available.
wbp = out/'project_integration_report.xlsx'
if wbp.exists():
    wb = load_workbook(wbp, data_only=False)
    print('WORKBOOK', wbp.relative_to(ws), 'sheets', wb.sheetnames)
    for s in wb.sheetnames:
        sh = wb[s]
        print(' SHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
        vals = [sh.cell(1,c).value for c in range(1, min(sh.max_column, 12)+1)]
        print('  row1', vals)
