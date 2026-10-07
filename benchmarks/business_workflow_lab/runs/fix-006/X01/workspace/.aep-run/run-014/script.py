from pathlib import Path
import csv, json, shutil, sys
from datetime import date
from decimal import Decimal
from openpyxl import load_workbook

root = Path('.')
out = root/'output'
dl = out/'downloads'
preserved = out/'preserved_2026_09_before_august_rerun'

print('Preparing August copies and restoring September outputs')
print('Output files before:')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(' -', p, p.stat().st_size)

required_current_aug = [
    dl/'sap_orders.csv',
    dl/'sap_invoices.csv',
    out/'expense_approved.csv',
    out/'quotes_accepted.json',
    out/'project_integration_report.xlsx',
    out/'project_summary_sep2026.csv',  # contains August from workflow fixed naming
]
missing = [str(p) for p in required_current_aug if not p.exists()]
if missing:
    raise SystemExit('Missing current August workflow outputs: ' + ', '.join(missing))

# Copy August run outputs to non-overwriting names.
aug_copies = {
    dl/'sap_orders.csv': dl/'sap_orders_aug2026.csv',
    dl/'sap_invoices.csv': dl/'sap_invoices_aug2026.csv',
    out/'expense_approved.csv': out/'expense_approved_aug2026.csv',
    out/'quotes_accepted.json': out/'quotes_accepted_aug2026.json',
    out/'project_integration_report.xlsx': out/'project_integration_report_aug2026.xlsx',
    out/'project_summary_sep2026.csv': out/'project_summary_aug2026.csv',
}
for src, dst in aug_copies.items():
    shutil.copy2(src, dst)
    print('August copy:', src, '->', dst)

# Restore September outputs from backup to explicit _sep2026 names.
sep_restore = {
    preserved/'downloads'/'sap_orders_sep2026.csv': dl/'sap_orders_sep2026.csv',
    preserved/'downloads'/'sap_invoices_sep2026.csv': dl/'sap_invoices_sep2026.csv',
    preserved/'expense_approved_sep2026.csv': out/'expense_approved_sep2026.csv',
    preserved/'quotes_accepted_sep2026.json': out/'quotes_accepted_sep2026.json',
    preserved/'project_summary_sep2026.csv': out/'project_summary_sep2026.csv',
    preserved/'project_integration_report.xlsx': out/'project_integration_report_sep2026.xlsx',
}
missing_sep = [str(p) for p in sep_restore if not p.exists()]
if missing_sep:
    raise SystemExit('Missing preserved September files: ' + ', '.join(missing_sep))
for src, dst in sep_restore.items():
    shutil.copy2(src, dst)
    print('September restore:', src, '->', dst)

# Also leave generic workbook as the August run artifact, but explicit files are what should be used.

def read_csv(path):
    with path.open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def money(x):
    return Decimal(str(x or '0'))

def ensure(cond, msg):
    if not cond:
        raise AssertionError(msg)

# Validate August source details.
orders = read_csv(dl/'sap_orders_aug2026.csv')
invoices = read_csv(dl/'sap_invoices_aug2026.csv')
expenses = read_csv(out/'expense_approved_aug2026.csv')
summary = read_csv(out/'project_summary_aug2026.csv')
quotes_data = json.loads((out/'quotes_accepted_aug2026.json').read_text(encoding='utf-8'))
quotes = quotes_data.get('items', quotes_data if isinstance(quotes_data, list) else [])

ensure(len(orders) == 7, f'August orders count expected 7 got {len(orders)}')
ensure(all('2026-08-01' <= r['OrderDate'] <= '2026-08-31' for r in orders), 'August orders outside date range')
ensure(all(r['Status'] != 'Cancelled' for r in orders), 'August orders include Cancelled')
ensure(len(invoices) == 6, f'August invoices count expected 6 got {len(invoices)}')
ensure(all('2026-08-01' <= r['InvoiceDate'] <= '2026-08-31' for r in invoices), 'August invoices outside date range')
ensure(len(expenses) == 4, f'August approved expenses count expected 4 got {len(expenses)}')
ensure(all('2026-08-01' <= r['ExpenseDate'] <= '2026-08-31' for r in expenses), 'August expenses outside date range')
ensure(all(r['Status'] == 'Approved' for r in expenses), 'August expenses include non-Approved')
ensure(len(quotes) == 3, f'August accepted quotes count expected 3 got {len(quotes)}')
# quote field names vary; check all have Accepted and August date if present
for q in quotes:
    if 'Status' in q:
        ensure(q['Status'] == 'Accepted', 'Quote not Accepted')
    qd = q.get('QuoteDate') or q.get('quote_date') or q.get('date')
    if qd:
        ensure('2026-08-01' <= qd <= '2026-08-31', 'Quote outside August')
ensure(len(summary) == 5, f'August summary project rows expected 5 got {len(summary)}')
ensure(set(summary[0].keys()) == {'ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'}, 'Summary columns changed')

# Validate August workbook sheets and core date cells.
wb = load_workbook(out/'project_integration_report_aug2026.xlsx', data_only=True)
ensure('Report' in wb.sheetnames, 'Missing Report sheet')
ws = wb['Report']
ensure(str(ws['B3'].value) == '2026-08-01', f'Report From expected 2026-08-01 got {ws["B3"].value}')
ensure(str(ws['D3'].value) == '2026-08-31', f'Report To expected 2026-08-31 got {ws["D3"].value}')
# Check sheet row counts: max_row includes 4 header/instruction rows.
expected_rows = {'Orders': 4+len(orders), 'Invoices': 4+len(invoices), 'Expenses': 4+len(expenses), 'Quotes': 4+len(quotes), 'Summary': 4+len(summary)}
for sheet, expected in expected_rows.items():
    ensure(sheet in wb.sheetnames, f'Missing sheet {sheet}')
    ensure(wb[sheet].max_row == expected, f'{sheet} max_row expected {expected} got {wb[sheet].max_row}')

# Validate September explicit files are restored and not August. Use known date windows/counts from earlier validation.
sep_orders = read_csv(dl/'sap_orders_sep2026.csv')
sep_invoices = read_csv(dl/'sap_invoices_sep2026.csv')
sep_expenses = read_csv(out/'expense_approved_sep2026.csv')
sep_summary = read_csv(out/'project_summary_sep2026.csv')
ensure(len(sep_orders) == 18, f'September orders expected 18 got {len(sep_orders)}')
ensure(all('2026-09-01' <= r['OrderDate'] <= '2026-09-30' for r in sep_orders), 'September orders outside date range')
ensure(all(r['Status'] != 'Cancelled' for r in sep_orders), 'September orders include Cancelled')
ensure(len(sep_invoices) == 13, f'September invoices expected 13 got {len(sep_invoices)}')
ensure(all('2026-09-01' <= r['InvoiceDate'] <= '2026-09-30' for r in sep_invoices), 'September invoices outside date range')
ensure(len(sep_expenses) == 15, f'September expenses expected 15 got {len(sep_expenses)}')
ensure(all(r['Status'] == 'Approved' for r in sep_expenses), 'September expenses include non-Approved')
ensure(len(sep_summary) == 5, f'September summary rows expected 5 got {len(sep_summary)}')
sep_wb = load_workbook(out/'project_integration_report_sep2026.xlsx', data_only=True)
ensure(str(sep_wb['Report']['B3'].value) == '2026-09-01', 'September report from date not restored')
ensure(str(sep_wb['Report']['D3'].value) == '2026-09-30', 'September report to date not restored')

print('\nValidation passed.')
print('August outputs:')
for p in [dl/'sap_orders_aug2026.csv', dl/'sap_invoices_aug2026.csv', out/'expense_approved_aug2026.csv', out/'quotes_accepted_aug2026.json', out/'project_summary_aug2026.csv', out/'project_integration_report_aug2026.xlsx']:
    print(' -', p, p.stat().st_size)
print('Restored September outputs:')
for p in [dl/'sap_orders_sep2026.csv', dl/'sap_invoices_sep2026.csv', out/'expense_approved_sep2026.csv', out/'quotes_accepted_sep2026.json', out/'project_summary_sep2026.csv', out/'project_integration_report_sep2026.xlsx']:
    print(' -', p, p.stat().st_size)
