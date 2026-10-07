from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-006\X01\workspace')
print('INPUTS')
for p in sorted((ws/'input').glob('*')):
    print(p.name, p.stat().st_size)
    if p.suffix.lower() in ['.csv','.md']:
        print(p.read_text(encoding='utf-8-sig')[:1000].replace('\r',''))
        print('---')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
print('TEMPLATE sheets', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'rows', sh.max_row, 'cols', sh.max_column, 'headers', [c.value for c in sh[1]])
print('CURRENT ORDERS STATUS COUNTS')
with (ws/'output'/'downloads'/'sap_orders.csv').open(newline='', encoding='utf-8-sig') as f:
    rows=list(csv.DictReader(f))
counts={}
for r in rows: counts[r['Status']]=counts.get(r['Status'],0)+1
print(counts)
print('Cancelled rows')
for r in rows:
    if r['Status']=='Cancelled': print(r)
print('QUOTE JSON top')
obj=json.load(open(ws/'output'/'quotes_accepted_sep2026.json',encoding='utf-8'))
print(type(obj).__name__, obj.keys() if isinstance(obj,dict) else 'list', str(obj)[:1000])