from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace=', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# fx sample
fx=ws/'input'/'fx_rates.csv'
if fx.exists():
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r=csv.reader(f)
        for i,row in zip(range(6), r):
            print('fx row', i, row)
# manual excerpt
man=ws/'input'/'expense_manual.md'
if man.exists():
    text=man.read_text(encoding='utf-8')
    print('manual length', len(text))
    print('manual excerpt:\n', text[:4000])
# template workbook structure
path=ws/'input'/'expense_template.xlsx'
wb=load_workbook(path, data_only=False)
print('template sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('sheet', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 12), values_only=False):
        vals=[]
        for c in row[:min(wsht.max_column, 12)]:
            vals.append(c.value)
        print(vals)
    print('---')