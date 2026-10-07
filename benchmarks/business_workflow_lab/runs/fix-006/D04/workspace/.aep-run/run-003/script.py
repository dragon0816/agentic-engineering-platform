from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime, timezone, timedelta

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

template = input_dir / 'expense_template.xlsx'
output_file = output_dir / 'expense_report_E006_2026-01-01_to_2026-01-31.xlsx'

employee = 'E006'
start_date = '2026-01-01'
end_date = '2026-01-31'
status_desc = '全部狀態（Status 空白）'
source_name = 'Expense Viewer (synthetic lab)'
# User supplied current date/context is 2026-10-06 Asia/Taipei; use deterministic ISO 8601 within that date.
generated_at = '2026-10-06T00:00:00+08:00'
fetched_at = generated_at

# No CSV is expected for the zero-row query; do not treat a missing export as success.
export_csv = output_dir / 'expense_E006_2026-01.csv'
print('expected export csv exists?', export_csv.exists())

wb = load_workbook(template)

# Fill report metadata in the template copy; keep sheets, merged cells, styles, footer intact.
report = wb['Report']
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = f'EmployeeID={employee}; ExpenseDate {start_date} 至 {end_date}（含起訖兩日）; Status={status_desc}'
report['B5'] = generated_at
report['B7'] = source_name
report['D7'] = fetched_at
report['B9'] = 'Expense Viewer 查詢結果為 0 筆；未產生匯出檔，依要求產出正常空報表。'

# Ensure detail sheets have no data rows after row 4.
for sheet_name in ['Details', 'Summary', 'Exceptions']:
    ws = wb[sheet_name]
    if ws.max_row > 4:
        ws.delete_rows(5, ws.max_row - 4)

# Summary formula for zero rows per template instructions.
summary = wb['Summary']
summary['B2'] = '=SUM(Details!I5:I5)'

# Apply basic formats to potential first data row/ranges without creating data values.
details = wb['Details']
for col in ['A','B','D','K','M','N']:
    details[f'{col}5'].number_format = '@'
details['E5'].number_format = 'yyyy-mm-dd'
for col in ['H','I']:
    details[f'{col}5'].number_format = '0.00'

# Set autofilter over header row if template did not already have it.
if not details.auto_filter.ref:
    details.auto_filter.ref = 'A4:N4'
if not summary.auto_filter.ref:
    summary.auto_filter.ref = 'A4:F4'
exceptions = wb['Exceptions']
if not exceptions.auto_filter.ref:
    exceptions.auto_filter.ref = 'A4:D4'

wb.save(output_file)
print('wrote', output_file)

# Verification: reopen saved output and confirm stated/template requirements.
check = load_workbook(output_file, data_only=False)
errors = []
if check.sheetnames != ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']:
    errors.append(f'sheet names changed: {check.sheetnames}')
if str(list(check['Report'].merged_cells.ranges)) != '[<MergedCellRange A1:F1>]':
    errors.append(f'Report merged ranges changed: {list(check["Report"].merged_cells.ranges)}')
rep = check['Report']
checks = {
    'Report!B3': (rep['B3'].value, start_date),
    'Report!D3': (rep['D3'].value, end_date),
    'Report!B7': (rep['B7'].value, source_name),
    'Summary!B2': (check['Summary']['B2'].value, '=SUM(Details!I5:I5)'),
}
for name,(got,expected) in checks.items():
    if got != expected:
        errors.append(f'{name} expected {expected!r} got {got!r}')
# headers unchanged
expected_details_headers = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
actual_details_headers = [check['Details'].cell(4,c).value for c in range(1,15)]
if actual_details_headers != expected_details_headers:
    errors.append(f'Details headers mismatch: {actual_details_headers}')
expected_summary_headers = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
actual_summary_headers = [check['Summary'].cell(4,c).value for c in range(1,7)]
if actual_summary_headers != expected_summary_headers:
    errors.append(f'Summary headers mismatch: {actual_summary_headers}')
expected_ex_headers = ['ExpenseID','EmployeeID','ExpenseDate','Reason']
actual_ex_headers = [check['Exceptions'].cell(4,c).value for c in range(1,5)]
if actual_ex_headers != expected_ex_headers:
    errors.append(f'Exceptions headers mismatch: {actual_ex_headers}')
# count non-empty data rows from row 5 downward.
def nonempty_rows(ws, min_col, max_col):
    count=0
    for r in range(5, ws.max_row+1):
        if any(ws.cell(r,c).value is not None for c in range(min_col, max_col+1)):
            count += 1
    return count
counts = {
    'Details': nonempty_rows(check['Details'],1,14),
    'Summary': nonempty_rows(check['Summary'],1,6),
    'Exceptions': nonempty_rows(check['Exceptions'],1,4),
}
print('data row counts:', counts)
for sh,cnt in counts.items():
    if cnt != 0:
        errors.append(f'{sh} expected 0 data rows got {cnt}')
# Verify output file is a valid xlsx and no CSV is being claimed as success.
if not output_file.exists() or output_file.stat().st_size <= 0:
    errors.append('output xlsx missing or empty')
print('output xlsx size:', output_file.stat().st_size)
print('export csv exists after processing?', export_csv.exists())

if errors:
    print('VERIFICATION FAILED')
    for e in errors:
        print('ERROR:', e)
    raise SystemExit(1)
print('VERIFICATION PASSED: template copy saved, headers/sheets/merged range preserved, zero data rows, zero-row formula retained, metadata filled.')