from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('output files before:')
out=ws/'output'
out.mkdir(exist_ok=True)
for p in sorted(out.glob('*')):
    print('-', p.name, p.stat().st_size)
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
print('Full Instructions:')
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)
print('Merged ranges:')
for sh in wb.sheetnames:
    print(sh, list(wb[sh].merged_cells.ranges))