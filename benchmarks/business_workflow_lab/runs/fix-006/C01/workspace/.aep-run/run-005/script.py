from pathlib import Path
from openpyxl import load_workbook
import json, os, textwrap
ws = Path.cwd()
print('Workspace:', ws)
print('Input files:', sorted(p.name for p in (ws/'input').glob('*')))
print('Output files:', [(p.name, p.stat().st_size) for p in sorted((ws/'output').glob('*'))])
run_dir = ws/'.aep-run'
if run_dir.exists():
    runs = sorted([p for p in run_dir.iterdir() if p.is_dir()])
    print('Run dirs:', [p.name for p in runs])
    if runs:
        latest = runs[-1]
        print('Latest run dir:', latest)
        sp = latest/'script.py'
        if sp.exists():
            lines = sp.read_text(encoding='utf-8', errors='replace').splitlines()
            print('Latest script line count:', len(lines))
            for i, line in enumerate(lines[-80:], start=max(1,len(lines)-79)):
                if 'assert' in line or 'Summary' in line or 'B3' in line or 'D3' in line:
                    print(f'{i}: {line}')
for out in sorted((ws/'output').glob('*.xlsx')):
    print('\nInspect workbook:', out.name)
    wb = load_workbook(out, data_only=False)
    print('Sheets:', wb.sheetnames)
    for s in ['Report','Quotes','Lines','Summary']:
        sh = wb[s]
        print(s, 'max_row', sh.max_row, 'max_col', sh.max_column)
        sample = []
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
            sample.append(row[:min(sh.max_column,12)])
        print(s, 'sample:', sample)
    cr = wb['Report']
    print('Report B3/D3 raw:', repr(cr['B3'].value), type(cr['B3'].value), repr(cr['D3'].value), type(cr['D3'].value))
    print('Summary B2:', wb['Summary']['B2'].value)
