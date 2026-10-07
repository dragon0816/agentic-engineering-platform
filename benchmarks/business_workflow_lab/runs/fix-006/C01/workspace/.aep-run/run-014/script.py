from pathlib import Path
from openpyxl import load_workbook

ws = Path.cwd()
print('Workspace:', ws)
print('Input files:', sorted([p.name for p in (ws/'input').iterdir()]) if (ws/'input').exists() else 'NO input folder')
out = ws/'output'
print('Output files:', sorted([p.name for p in out.iterdir()]) if out.exists() else 'NO output folder')
for p in sorted(out.glob('*.xlsx')):
    wb = load_workbook(p, read_only=True, data_only=False)
    print('Workbook:', p.name)
    print('  sheets:', wb.sheetnames)
    for s in ['Report','Quotes','Lines','Summary']:
        if s in wb.sheetnames:
            sh = wb[s]
            rows = []
            for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,5), values_only=True):
                rows.append(row)
            print('  sample', s, rows)
    wb.close()
