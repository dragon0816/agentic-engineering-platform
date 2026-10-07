import csv
import json
import os
import shutil
import subprocess
import sys
import time
from collections import defaultdict
from copy import copy
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / "input"
OUTPUT = WORKSPACE / "output"
OUTPUT.mkdir(parents=True, exist_ok=True)

params_path = os.environ.get("AEP_WORKFLOW_PARAMS")
if not params_path:
    raise RuntimeError("This script must be run by the workflow runtime with AEP_WORKFLOW_PARAMS.")
with open(params_path, "r", encoding="utf-8") as f:
    PARAMS = json.load(f)

CUSTOMER_ID = str(PARAMS.get("customer_id", "")).strip()
START_DATE = str(PARAMS.get("start_date", "")).strip()
END_DATE = str(PARAMS.get("end_date", "")).strip()
CLI_PATH = str(PARAMS.get("cli_path", "")).strip()

if not CUSTOMER_ID:
    raise ValueError("customer_id is required")
if not START_DATE or not END_DATE:
    raise ValueError("start_date and end_date are required")
try:
    start_dt = date.fromisoformat(START_DATE)
    end_dt = date.fromisoformat(END_DATE)
except Exception as exc:
    raise ValueError("start_date and end_date must be YYYY-MM-DD") from exc
if start_dt > end_dt:
    raise ValueError("start_date must be on or before end_date")
if not CLI_PATH:
    raise ValueError("cli_path is required")
cli = Path(CLI_PATH)
if not cli.exists():
    raise FileNotFoundError(f"CLI path does not exist: {CLI_PATH}")

template_path = INPUT / "quote_template.xlsx"
fx_path = INPUT / "fx_rates.csv"
if not template_path.exists():
    raise FileNotFoundError(template_path)
if not fx_path.exists():
    raise FileNotFoundError(fx_path)

print("Workspace:", WORKSPACE)
print("Input files:", sorted(p.name for p in INPUT.glob("*")))
print("Parameters:", {"customer_id": CUSTOMER_ID, "start_date": START_DATE, "end_date": END_DATE, "cli_path": CLI_PATH})

# Read FX rates.  The note field in the supplied CSV can contain unquoted commas;
# only the first two columns are part of the contract used here.
fx = {}
with fx_path.open("r", newline="", encoding="utf-8-sig") as f:
    reader = csv.reader(f)
    header = next(reader, None)
    if not header or header[0] != "Currency" or header[1] != "RateToTWD":
        raise ValueError("fx_rates.csv must start with Currency,RateToTWD")
    for row in reader:
        if not row or not row[0].strip():
            continue
        fx[row[0].strip()] = Decimal(row[1].strip())
print("FX currencies:", sorted(fx.keys()))

TWOPLACES = Decimal("0.01")

def q2(value):
    return Decimal(str(value)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def decimal_to_number(value):
    d = q2(value)
    # openpyxl stores native numeric types; float is acceptable after Decimal rounding.
    return float(d)

def parse_decimal(value):
    if value is None or value == "":
        return None
    return Decimal(str(value))

def parse_date_cell(value):
    if value in (None, ""):
        return None
    return date.fromisoformat(str(value))

def build_command(args):
    suffix = cli.suffix.lower()
    if suffix == ".py":
        return [sys.executable, str(cli)] + args
    if suffix in (".cmd", ".bat"):
        line = subprocess.list2cmdline([str(cli)] + args)
        return [os.environ.get("ComSpec", "cmd.exe"), "/d", "/s", "/c", line]
    return [str(cli)] + args

def run_cli_json(args):
    cmd = build_command(args)
    last = None
    for attempt in range(1, 4):
        completed = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True, encoding="utf-8", errors="replace")
        last = completed
        if completed.returncode == 4 and attempt < 3:
            time.sleep(attempt)
            continue
        break
    if last.returncode != 0:
        raise RuntimeError(f"quote_cli failed rc={last.returncode} args={args}; stderr={last.stderr.strip()}")
    out = last.stdout.strip()
    try:
        return json.loads(out)
    except json.JSONDecodeError as exc:
        raise RuntimeError(f"quote_cli returned non-JSON stdout for args={args}: {out[:500]}") from exc

list_payload = run_cli_json([
    "quotes", "list",
    "--customer", CUSTOMER_ID,
    "--from", START_DATE,
    "--to", END_DATE,
    "--format", "json",
])
if list_payload.get("schema_version") != 1:
    raise ValueError(f"Unsupported list schema_version: {list_payload.get('schema_version')}")
quote_refs = list_payload.get("items", []) or []
print("Quotes listed:", len(quote_refs), "count field:", list_payload.get("count"))

quote_docs = []
for item in quote_refs:
    qid = str(item.get("QuoteID", "")).strip()
    if not qid:
        raise ValueError(f"Quote list item missing QuoteID: {item}")
    doc = run_cli_json(["quotes", "get", qid, "--format", "json"])
    if doc.get("schema_version") != 1:
        raise ValueError(f"Unsupported get schema_version for {qid}: {doc.get('schema_version')}")
    header = doc.get("header") or {}
    lines = doc.get("line_items") or []
    if str(header.get("QuoteID", "")) != qid:
        raise ValueError(f"QuoteID mismatch in get payload for {qid}: {header.get('QuoteID')}")
    quote_docs.append({"header": header, "line_items": lines})
print("Quote documents fetched:", len(quote_docs))
print("Line items fetched:", sum(len(d["line_items"]) for d in quote_docs))

out_name = f"quote_report_{CUSTOMER_ID}_{START_DATE}_to_{END_DATE}.xlsx"
out_path = OUTPUT / out_name
shutil.copy2(template_path, out_path)
wb = load_workbook(out_path)
required_sheets = ["Report", "Quotes", "Lines", "Summary", "Instructions"]
missing = [s for s in required_sheets if s not in wb.sheetnames]
if missing:
    raise ValueError(f"Template missing sheets: {missing}")

report = wb["Report"]
quotes_ws = wb["Quotes"]
lines_ws = wb["Lines"]
summary_ws = wb["Summary"]

quote_headers = ["QuoteID", "CustomerID", "ProjectID", "QuoteDate", "ValidUntil", "Status", "Currency", "Subtotal", "DiscountPercent", "NetAmount", "NetAmountTWD"]
line_headers = ["QuoteID", "ProductID", "ProductName", "Quantity", "UnitPrice", "LineAmount"]
summary_headers = ["Group", "SubGroup", "Currency", "Amount", "AmountTWD", "Count"]
if [quotes_ws.cell(4, c).value for c in range(1, 12)] != quote_headers:
    raise ValueError("Quotes row 4 headers do not match template contract")
if [lines_ws.cell(4, c).value for c in range(1, 7)] != line_headers:
    raise ValueError("Lines row 4 headers do not match template contract")
if [summary_ws.cell(4, c).value for c in range(1, 7)] != summary_headers:
    raise ValueError("Summary row 4 headers do not match template contract")

# Clear existing data rows without changing workbook structure/sheets.
def clear_rows(ws, min_row, max_col):
    if ws.max_row >= min_row:
        for row in ws.iter_rows(min_row=min_row, max_row=ws.max_row, max_col=max_col):
            for cell in row:
                cell.value = None

clear_rows(quotes_ws, 5, len(quote_headers))
clear_rows(lines_ws, 5, len(line_headers))
clear_rows(summary_ws, 5, len(summary_headers))

# Preserve likely detail row styling by copying row 4 styles down to data cells.
def apply_style_from_header(ws, row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(4, col)
        dst = ws.cell(row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)
        if src.alignment:
            dst.alignment = copy(src.alignment)

now = datetime.now(ZoneInfo("Asia/Taipei")).replace(microsecond=0).isoformat()
report["B3"] = start_dt
report["B3"].number_format = "yyyy-mm-dd"
report["D3"] = end_dt
report["D3"].number_format = "yyyy-mm-dd"
report["B4"] = f"CustomerID={CUSTOMER_ID}; QuoteDate {START_DATE} to {END_DATE} inclusive"
report["B5"] = now
report["B7"] = "quote_cli; fx_rates.csv"
report["D7"] = now
report["B9"] = "" if quote_docs else "No quotes returned for the requested inclusive range."

summary = defaultdict(lambda: {"amount": Decimal("0"), "amount_twd": Decimal("0"), "count": 0})
quote_ids = set()
line_count = 0

for idx, doc in enumerate(quote_docs, start=5):
    h = doc["header"]
    qid = str(h.get("QuoteID", ""))
    currency = str(h.get("Currency", ""))
    if currency not in fx:
        raise ValueError(f"No FX rate for currency {currency} in quote {qid}")
    net = q2(parse_decimal(h.get("NetAmount")))
    net_twd = q2(net * fx[currency])
    quote_ids.add(qid)
    values = [
        qid,
        str(h.get("CustomerID", "")),
        str(h.get("ProjectID", "")),
        parse_date_cell(h.get("QuoteDate")),
        parse_date_cell(h.get("ValidUntil")),
        str(h.get("Status", "")),
        currency,
        decimal_to_number(h.get("Subtotal")),
        float(Decimal(str(h.get("DiscountPercent", "0")))),
        decimal_to_number(net),
        decimal_to_number(net_twd),
    ]
    apply_style_from_header(quotes_ws, idx, len(quote_headers))
    for col, val in enumerate(values, start=1):
        cell = quotes_ws.cell(idx, col, val)
        if col in (1, 2, 3, 6, 7):
            cell.number_format = "@"
        elif col in (4, 5):
            cell.number_format = "yyyy-mm-dd"
        elif col in (8, 10, 11):
            cell.number_format = "0.00"
        elif col == 9:
            cell.number_format = "0.00"
    key = (str(h.get("Status", "")), currency)
    summary[key]["amount"] += net
    summary[key]["amount_twd"] += net_twd
    summary[key]["count"] += 1

line_row = 5
for doc in quote_docs:
    qid = str(doc["header"].get("QuoteID", ""))
    for line in doc["line_items"]:
        apply_style_from_header(lines_ws, line_row, len(line_headers))
        values = [
            qid,
            str(line.get("ProductID", "")),
            line.get("ProductName", ""),
            int(Decimal(str(line.get("Quantity", "0")))),
            decimal_to_number(line.get("UnitPrice")),
            decimal_to_number(line.get("LineAmount")),
        ]
        for col, val in enumerate(values, start=1):
            cell = lines_ws.cell(line_row, col, val)
            if col in (1, 2):
                cell.number_format = "@"
            elif col in (5, 6):
                cell.number_format = "0.00"
            elif col == 4:
                cell.number_format = "0"
        line_row += 1
        line_count += 1

summary_row = 5
for (status, currency), agg in sorted(summary.items()):
    apply_style_from_header(summary_ws, summary_row, len(summary_headers))
    vals = ["Status", status, currency, decimal_to_number(agg["amount"]), decimal_to_number(agg["amount_twd"]), int(agg["count"])]
    for col, val in enumerate(vals, start=1):
        cell = summary_ws.cell(summary_row, col, val)
        if col in (1, 2, 3):
            cell.number_format = "@"
        elif col in (4, 5):
            cell.number_format = "0.00"
        elif col == 6:
            cell.number_format = "0"
    summary_row += 1

last_quote_row = 4 + len(quote_docs)
if len(quote_docs) == 0:
    summary_ws["B2"] = "=SUM(Quotes!K5:K5)"
else:
    summary_ws["B2"] = f"=SUM(Quotes!K5:K{last_quote_row})"

wb.save(out_path)

# Reopen and verify saved output against stated requirements.
check_wb = load_workbook(out_path, data_only=False)
assert set(required_sheets).issubset(check_wb.sheetnames)
cq = check_wb["Quotes"]
cl = check_wb["Lines"]
cs = check_wb["Summary"]
cr = check_wb["Report"]
assert [cq.cell(4, c).value for c in range(1, 12)] == quote_headers
assert [cl.cell(4, c).value for c in range(1, 7)] == line_headers
assert [cs.cell(4, c).value for c in range(1, 7)] == summary_headers
assert cr["B3"].value == start_dt and cr["D3"].value == end_dt
assert cq.max_row == (4 if not quote_docs else 4 + len(quote_docs))
assert cl.max_row == (4 if not quote_docs else 4 + line_count)
assert cq["A5"].data_type == "s" if quote_docs else True
for r in range(5, 5 + len(quote_docs)):
    for c in (1, 2, 3):
        assert isinstance(cq.cell(r, c).value, str), f"Quotes ID column {c} row {r} is not text"
    for c in (8, 9, 10, 11):
        assert isinstance(cq.cell(r, c).value, (int, float)), f"Quotes numeric column {c} row {r} is not numeric"
for r in range(5, 5 + line_count):
    assert isinstance(cl.cell(r, 1).value, str), f"Lines QuoteID row {r} is not text"
    assert isinstance(cl.cell(r, 2).value, str), f"Lines ProductID row {r} is not text"
    for c in (4, 5, 6):
        assert isinstance(cl.cell(r, c).value, (int, float)), f"Lines numeric column {c} row {r} is not numeric"
expected_formula = "=SUM(Quotes!K5:K5)" if not quote_docs else f"=SUM(Quotes!K5:K{last_quote_row})"
assert cs["B2"].value == expected_formula

saved_summary = {}
for r in range(5, cs.max_row + 1):
    if cs.cell(r, 1).value is None:
        continue
    key = (cs.cell(r, 2).value, cs.cell(r, 3).value)
    saved_summary[key] = {
        "amount": q2(cs.cell(r, 4).value),
        "amount_twd": q2(cs.cell(r, 5).value),
        "count": int(cs.cell(r, 6).value),
    }
expected_summary = {k: {"amount": q2(v["amount"]), "amount_twd": q2(v["amount_twd"]), "count": v["count"]} for k, v in summary.items()}
assert saved_summary == expected_summary, f"Summary mismatch: saved={saved_summary} expected={expected_summary}"

print("Output workbook:", out_path)
print("Checked: template sheets preserved; Report dates/condition set; Quotes headers and", len(quote_docs), "rows; Lines headers and", line_count, "rows; Summary by status/currency", len(expected_summary), "rows; B2 formula", expected_formula, "; ID columns text; amount columns numeric.")
