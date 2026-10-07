from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date

ws = Path.cwd()
out = ws / 'output'
first = out / 'quote_report_C001_2026-09-01_to_2026-09-30.xlsx'
second = out / 'quote_report_C002_2026-08-01_to_2026-10-31.xlsx'
print('Workspace:', ws)
print('Output files:', sorted(p.name for p in out.iterdir()))

assert first.exists(), f'Missing first output: {first}'
assert second.exists(), f'Missing second output: {second}'
assert first.name != second.name, 'Second output overwrote first output name'
assert first.stat().st_size > 0, 'First output is empty'
assert second.stat().st_size > 0, 'Second output is empty'

wb = load_workbook(second, data_only=False)
try:
    expected_sheets = ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
    assert wb.sheetnames == expected_sheets, f'Sheet order mismatch: {wb.sheetnames}'
    report = wb['Report']
    quotes = wb['Quotes']
    lines = wb['Lines']
    summary = wb['Summary']

    def as_date(v):
        if isinstance(v, datetime):
            return v.date()
        if isinstance(v, date):
            return v
        return v

    assert as_date(report['B3'].value) == date(2026,8,1), f"Report B3 mismatch: {report['B3'].value!r}"
    assert as_date(report['D3'].value) == date(2026,10,31), f"Report D3 mismatch: {report['D3'].value!r}"
    cond = str(report['B4'].value)
    assert 'CustomerID=C002' in cond, f'Condition missing C002: {cond}'
    assert '2026-08-01' in cond and '2026-10-31' in cond, f'Condition dates mismatch: {cond}'

    q_headers = [quotes.cell(4,c).value for c in range(1,12)]
    assert q_headers == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD'], q_headers
    l_headers = [lines.cell(4,c).value for c in range(1,7)]
    assert l_headers == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount'], l_headers
    s_headers = [summary.cell(4,c).value for c in range(1,7)]
    assert s_headers == ['Group','SubGroup','Currency','Amount','AmountTWD','Count'], s_headers

    quote_rows = []
    for r in range(5, quotes.max_row+1):
        qid = quotes.cell(r,1).value
        if qid is not None:
            quote_rows.append(r)
            assert isinstance(qid, str), f'QuoteID not text at row {r}: {qid!r}'
            assert quotes.cell(r,2).value == 'C002', f'CustomerID mismatch row {r}: {quotes.cell(r,2).value!r}'
            assert isinstance(quotes.cell(r,3).value, str), f'ProjectID not text row {r}'
            assert date(2026,8,1) <= as_date(quotes.cell(r,4).value) <= date(2026,10,31), f'QuoteDate out of range row {r}: {quotes.cell(r,4).value!r}'
            for c in [8,9,10,11]:
                assert isinstance(quotes.cell(r,c).value, (int,float)), f'Quote numeric col {c} row {r}: {quotes.cell(r,c).value!r}'
    assert len(quote_rows) == 6, f'Expected 6 quote rows, got {len(quote_rows)}'

    line_rows = []
    quote_ids = {quotes.cell(r,1).value for r in quote_rows}
    for r in range(5, lines.max_row+1):
        qid = lines.cell(r,1).value
        if qid is not None:
            line_rows.append(r)
            assert isinstance(qid, str), f'Line QuoteID not text row {r}'
            assert qid in quote_ids, f'Line row {r} references quote outside selected set: {qid}'
            assert isinstance(lines.cell(r,2).value, str), f'ProductID not text row {r}'
            assert isinstance(lines.cell(r,3).value, str), f'ProductName not text row {r}'
            for c in [4,5,6]:
                assert isinstance(lines.cell(r,c).value, (int,float)), f'Line numeric col {c} row {r}: {lines.cell(r,c).value!r}'
    assert len(line_rows) == 11, f'Expected 11 line rows, got {len(line_rows)}'

    assert summary['B2'].value == '=SUM(Quotes!K5:K10)', f"Summary B2 formula mismatch: {summary['B2'].value!r}"
    summary_rows = [r for r in range(5, summary.max_row+1) if summary.cell(r,1).value is not None]
    assert len(summary_rows) == 5, f'Expected 5 summary rows, got {len(summary_rows)}'
    for r in summary_rows:
        assert isinstance(summary.cell(r,1).value, str), f'Summary group not text row {r}'
        assert isinstance(summary.cell(r,2).value, str), f'Summary subgroup not text row {r}'
        assert isinstance(summary.cell(r,3).value, str), f'Summary currency not text row {r}'
        for c in [4,5,6]:
            assert isinstance(summary.cell(r,c).value, (int,float)), f'Summary numeric col {c} row {r}: {summary.cell(r,c).value!r}'
finally:
    wb.close()

print('VALIDATION PASSED')
print('First output preserved:', first.name, first.stat().st_size, 'bytes')
print('Second output:', second.name, second.stat().st_size, 'bytes')
print('Second output checks: sheets ok; C002 date range ok; 6 quote rows; 11 line rows; 5 summary rows; formula =SUM(Quotes!K5:K10); IDs text; numeric amount/count fields numeric.')
