from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime, date
from decimal import Decimal

workspace = Path.cwd()
print('Workspace:', workspace)
print('Input files:', sorted(p.name for p in (workspace/'input').iterdir()) if (workspace/'input').exists() else 'NO INPUT DIR')
print('Output files:', sorted(p.name for p in (workspace/'output').iterdir()) if (workspace/'output').exists() else 'NO OUTPUT DIR')

out = workspace / 'output' / 'quote_report_C001_2026-09-01_to_2026-09-30.xlsx'
assert out.exists(), f'Missing output file: {out}'

wb = load_workbook(out, data_only=False)
print('Workbook sheets:', wb.sheetnames)
expected_sheets = ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']
assert wb.sheetnames == expected_sheets, f'Sheet names/order mismatch: {wb.sheetnames}'

# Print quick samples from saved output
for s in expected_sheets[:4]:
    ws = wb[s]
    sample = []
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 8), values_only=True):
        sample.append(row)
    print(f'Sample {s}:', sample)

report = wb['Report']
quotes = wb['Quotes']
lines = wb['Lines']
summary = wb['Summary']

# Normalize date-like cells for comparison.
def norm_date(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, str):
        return v.strip()[:10]
    raise AssertionError(f'Unexpected date value/type: {v!r} ({type(v)})')

assert norm_date(report['B3'].value) == '2026-09-01', f"Report B3 mismatch: {report['B3'].value!r}"
assert norm_date(report['D3'].value) == '2026-09-30', f"Report D3 mismatch: {report['D3'].value!r}"
assert 'C001' in str(report['B4'].value), f"Report B4 should mention C001: {report['B4'].value!r}"

quote_headers = [quotes.cell(4, c).value for c in range(1, 12)]
expected_quote_headers = ['QuoteID', 'CustomerID', 'ProjectID', 'QuoteDate', 'ValidUntil', 'Status', 'Currency', 'Subtotal', 'DiscountPercent', 'NetAmount', 'NetAmountTWD']
assert quote_headers == expected_quote_headers, f'Quote headers mismatch: {quote_headers}'

line_headers = [lines.cell(4, c).value for c in range(1, 7)]
expected_line_headers = ['QuoteID', 'ProductID', 'ProductName', 'Quantity', 'UnitPrice', 'LineAmount']
assert line_headers == expected_line_headers, f'Line headers mismatch: {line_headers}'

quote_rows = []
for r in range(5, quotes.max_row + 1):
    vals = [quotes.cell(r, c).value for c in range(1, 12)]
    if any(v is not None for v in vals):
        quote_rows.append((r, vals))

line_rows = []
for r in range(5, lines.max_row + 1):
    vals = [lines.cell(r, c).value for c in range(1, 7)]
    if any(v is not None for v in vals):
        line_rows.append((r, vals))

print('Quote data row count:', len(quote_rows))
print('Line data row count:', len(line_rows))
assert len(quote_rows) == 4, f'Expected 4 quote rows, got {len(quote_rows)}'
assert len(line_rows) > 0, 'Expected at least one line row'

for r, vals in quote_rows:
    quote_id, customer_id, project_id, quote_date, valid_until, status, currency, subtotal, discount_percent, net_amount, net_amount_twd = vals
    assert isinstance(quote_id, str), f'QuoteID row {r} not text: {quote_id!r}'
    assert isinstance(customer_id, str), f'CustomerID row {r} not text: {customer_id!r}'
    assert customer_id == 'C001', f'CustomerID row {r} mismatch: {customer_id!r}'
    assert isinstance(project_id, str), f'ProjectID row {r} not text: {project_id!r}'
    assert norm_date(quote_date) >= '2026-09-01' and norm_date(quote_date) <= '2026-09-30', f'QuoteDate row {r} out of range: {quote_date!r}'
    assert currency in {'TWD','USD','EUR'}, f'Currency row {r} unexpected: {currency!r}'
    for name, value in [('Subtotal', subtotal), ('DiscountPercent', discount_percent), ('NetAmount', net_amount), ('NetAmountTWD', net_amount_twd)]:
        assert isinstance(value, (int, float)), f'{name} row {r} not numeric: {value!r} ({type(value)})'

quote_ids = {vals[0] for _, vals in quote_rows}
for r, vals in line_rows:
    quote_id, product_id, product_name, quantity, unit_price, line_amount = vals
    assert isinstance(quote_id, str), f'Line QuoteID row {r} not text: {quote_id!r}'
    assert quote_id in quote_ids, f'Line row {r} QuoteID not in Quotes: {quote_id!r}'
    assert isinstance(product_id, str), f'ProductID row {r} not text: {product_id!r}'
    assert isinstance(product_name, str) and product_name, f'ProductName row {r} blank/non-text: {product_name!r}'
    for name, value in [('Quantity', quantity), ('UnitPrice', unit_price), ('LineAmount', line_amount)]:
        assert isinstance(value, (int, float)), f'{name} row {r} not numeric: {value!r} ({type(value)})'

last_quote_row = 4 + len(quote_rows)
expected_formula = f'=SUM(Quotes!K5:K{last_quote_row})'
assert summary['B2'].value == expected_formula, f"Summary B2 formula mismatch: {summary['B2'].value!r}, expected {expected_formula!r}"

summary_headers = [summary.cell(4, c).value for c in range(1, 7)]
expected_summary_headers = ['Group', 'SubGroup', 'Currency', 'Amount', 'AmountTWD', 'Count']
assert summary_headers == expected_summary_headers, f'Summary headers mismatch: {summary_headers}'
summary_rows = []
for r in range(5, summary.max_row + 1):
    vals = [summary.cell(r, c).value for c in range(1, 7)]
    if any(v is not None for v in vals):
        summary_rows.append((r, vals))
print('Summary data row count:', len(summary_rows))
assert summary_rows, 'Expected summary rows'
for r, vals in summary_rows:
    group, subgroup, currency, amount, amount_twd, count = vals
    assert isinstance(group, str) and group, f'Summary Group row {r} invalid: {group!r}'
    assert isinstance(currency, str) and currency in {'TWD','USD','EUR'}, f'Summary Currency row {r} invalid: {currency!r}'
    for name, value in [('Amount', amount), ('AmountTWD', amount_twd), ('Count', count)]:
        assert isinstance(value, (int, float)), f'Summary {name} row {r} not numeric: {value!r}'

print('CHECK PASSED')
print('Verified output file:', str(out))
print('Final output files:', sorted(p.name for p in (workspace/'output').iterdir()))