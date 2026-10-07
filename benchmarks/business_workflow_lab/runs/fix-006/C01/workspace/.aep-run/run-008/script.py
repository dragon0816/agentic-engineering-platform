from pathlib import Path
import json, os, zipfile
from openpyxl import load_workbook
ws = Path.cwd()
print('Workspace:', ws)
print('Input files:', sorted(p.name for p in (ws/'input').glob('*')) if (ws/'input').exists() else 'NO input')
print('Output files:', [(p.name, p.stat().st_size) for p in sorted((ws/'output').glob('*'))] if (ws/'output').exists() else 'NO output')
run_root = ws/'.aep-run'
print('Run dirs:', [p.name for p in sorted(run_root.glob('*'))] if run_root.exists() else 'NO .aep-run')
for run in sorted(run_root.glob('run-*'))[-5:]:
    print('\n==', run.name, '==')
    for p in sorted(run.iterdir()):
        print(' ', p.name, p.stat().st_size)
    sp = run/'script.py'
    if sp.exists():
        text = sp.read_text(encoding='utf-8', errors='replace')
        print(' script first line:', text.splitlines()[0] if text.splitlines() else '')
        # Print likely assertion/check lines near bottom for current failure diagnosis
        lines = text.splitlines()
        print(' script tail:')
        for i, line in list(enumerate(lines, start=1))[-35:]:
            print(f'{i}: {line}')

out = ws/'output'/'quote_report_C001_2026-09-01_to_2026-09-30.xlsx'
if out.exists():
    wb = load_workbook(out, data_only=False)
    print('\nWorkbook sheets:', wb.sheetnames)
    for s in ['Report','Quotes','Lines','Summary']:
        sh = wb[s]
        print('\nSheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 10), values_only=True):
            print(row)
    wb.close()
