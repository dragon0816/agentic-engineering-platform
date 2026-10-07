import csv, json, os, shutil, subprocess, sys
from collections import defaultdict
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from pathlib import Path
from openpyxl import load_workbook

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)
params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
customer_id = str(params.get('customer_id', '')).strip()
start_s = str(params.get('start_date', '')).strip()[:10]
end_s = str(params.get('end_date', '')).strip()[:10]
cli_path = Path(str(params.get('cli_path', '')).strip())
if not customer_id or not start_s or not end_s or not str(cli_path):
    raise SystemExit('Missing required parameter: customer_id, start_date, end_date, cli_path')
start_dt = datetime.strptime(start_s, '%Y-%m-%d').date()
end_dt = datetime.strptime(end_s, '%Y-%m-%d').date()
if end_dt < start_dt:
    raise SystemExit('end_date must be on or after start_date')
if not cli_path.exists():
    raise SystemExit(f'CLI path does not exist: {cli_path}')

template_path = input_dir / 'quote_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
manual_path = input_dir / 'quote_cli_manual.md'
for p in (template_path, fx_path, manual_path):
    if not p.exists():
        raise SystemExit(f'Missing input file: {p}')
print('Workspace:', workspace)
print('Input files:', sorted(p.name for p in input_dir.iterdir()))
print('Parameters:', {'customer_id': customer_id, 'start_date': start_s, 'end_date': end_s, 'cli_path': str(cli_path)})
probe_wb = load_workbook(template_path, data_only=False)
print('Template sheets:', probe_wb.sheetnames)
for sheet_name in ['Report','Quotes','Lines','Summary','Instructions']:
    ws = probe_wb[sheet_name]
    sample = []
    for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), values_only=True):
        sample.append(row[:min(ws.max_column, 12)])
    print(f'Sheet sample {sheet_name}:', sample)
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    rdr = csv.DictReader(f)
    fx_rows = list(rdr)
print('FX headers:', rdr.fieldnames)
print('FX currencies:', [r.get('Currency') for r in fx_rows])
manual_preview = manual_path.read_text(encoding='utf-8', errors='replace')[:800]
print('Manual preview:', manual_preview.replace('\n',' ')[:300])

def run_cli(args):
    cmd = [str(cli_path)] + args
    p = subprocess.run(cmd, cwd=str(workspace), text=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, shell=False)
    if p.returncode != 0:
        raise RuntimeError(f'CLI failed rc={p.returncode} cmd={cmd}\nSTDOUT:{p.stdout}\nSTDERR:{p.stderr}')
    txt = (p.stdout or '').strip()
    if not txt:
        raise RuntimeError(f'CLI returned empty output for {cmd}')
    try:
        return json.loads(txt)
    except Exception as e:
        raise RuntimeError(f'CLI output was not JSON for {cmd}: {txt[:500]}') from e

def pick_list(obj):
    if isinstance(obj, list):
        return obj
    if isinstance(obj, dict):
        for k in ['quotes','items','data','results','records']:
            if isinstance(obj.get(k), list):
                return obj[k]
    raise ValueError('Cannot find list in CLI output: ' + repr(obj)[:500])

def norm_key(s):
    return ''.join(ch.lower() for ch in str(s) if ch.isalnum())

def get_any(d, names, default=None):
    if not isinstance(d, dict):
        return default
    direct = {str(k): v for k, v in d.items()}
    for n in names:
        if n in direct:
            return direct[n]
    nmap = {norm_key(k): v for k, v in d.items()}
    for n in names:
        nk = norm_key(n)
        if nk in nmap:
            return nmap[nk]
    return default

def to_dec(v, default='0'):
    if v is None or v == '':
        return Decimal(default)
    try:
        return Decimal(str(v))
    except InvalidOperation:
        return Decimal(default)

def q2(v):
    return Decimal(str(v)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def dec_num(v):
    d = q2(v)
    i = int(d)
    return i if d == Decimal(i) else float(d)

def parse_date_value(v):
    if v is None or v == '':
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return datetime.strptime(str(v)[:10], '%Y-%m-%d').date()

# Fetch list. The mock manual supports these common options; fallback variants make reruns robust.
list_variants = [
    ['list-quotes','--customer',customer_id,'--from',start_s,'--to',end_s],
    ['list','--customer',customer_id,'--from',start_s,'--to',end_s],
    ['quotes','--customer',customer_id,'--from',start_s,'--to',end_s],
    ['list-quotes','--customer-id',customer_id,'--start-date',start_s,'--end-date',end_s],
]
errors = []
quote_rows = None
for args in list_variants:
    try:
        quote_rows = pick_list(run_cli(args))
        print('CLI list command:', args)
        break
    except Exception as e:
        errors.append(str(e))
if quote_rows is None:
    raise RuntimeError('Could not list quotes:\n' + '\n'.join(errors))

def row_quote_id(row):
    return str(get_any(row, ['QuoteID','quote_id','id','quoteId'], '')).strip()

def fetch_detail(qid):
    variants = [
        ['get-quote','--quote-id',qid],
        ['get', '--quote-id', qid],
        ['quote','--quote-id',qid],
        ['get-quote', qid],
    ]
    detail_errors = []
    for args in variants:
        try:
            obj = run_cli(args)
            print('CLI detail command:', args)
            return obj
        except Exception as e:
            detail_errors.append(str(e))
    raise RuntimeError('Could not fetch quote ' + qid + ':\n' + '\n'.join(detail_errors))

quote_docs = []
for row in quote_rows:
    qid = row_quote_id(row)
    if not qid:
        raise ValueError('Quote row has no ID: ' + repr(row))
    doc = fetch_detail(qid)
    if isinstance(doc, dict) and isinstance(doc.get('quote'), dict):
        doc = doc['quote']
    if not isinstance(doc, dict):
        raise ValueError('Quote detail is not an object: ' + repr(doc))
    # Merge list row values as fallback.
    merged = dict(row)
    merged.update(doc)
    qdate = parse_date_value(get_any(merged, ['QuoteDate','quote_date','date']))
    cust = str(get_any(merged, ['CustomerID','customer_id','customerId'], customer_id)).strip()
    if cust == customer_id and start_dt <= qdate <= end_dt:
        quote_docs.append(merged)
quote_docs.sort(key=lambda d: (parse_date_value(get_any(d, ['QuoteDate','quote_date','date'])), str(get_any(d, ['QuoteID','quote_id','id','quoteId']))))
print('Quotes listed:', len(quote_rows), 'Quote documents in range:', len(quote_docs))

fx = {}
for r in fx_rows:
    cur = str(r.get('Currency','')).strip()
    if cur:
        fx[cur] = to_dec(r.get('RateToTWD','1'), '1')
print('FX loaded:', fx)

wb = load_workbook(template_path)
required_sheets = ['Report','Quotes','Lines','Summary','Instructions']
for s in required_sheets:
    if s not in wb.sheetnames:
        raise SystemExit(f'Missing required template sheet: {s}')
report_ws = wb['Report']; quotes_ws = wb['Quotes']; lines_ws = wb['Lines']; summary_ws = wb['Summary']
quote_header_row = 4
line_header_row = 4
summary_header_row = 4
quote_headers = [quotes_ws.cell(quote_header_row,c).value for c in range(1, quotes_ws.max_column+1) if quotes_ws.cell(quote_header_row,c).value is not None]
line_headers = [lines_ws.cell(line_header_row,c).value for c in range(1, lines_ws.max_column+1) if lines_ws.cell(line_header_row,c).value is not None]
summary_headers = [summary_ws.cell(summary_header_row,c).value for c in range(1, summary_ws.max_column+1) if summary_ws.cell(summary_header_row,c).value is not None]
print('Quote headers:', quote_headers)
print('Line headers:', line_headers)
print('Summary headers:', summary_headers)

def hmap(ws, row):
    return {str(ws.cell(row,c).value): c for c in range(1, ws.max_column+1) if ws.cell(row,c).value is not None}
qh = hmap(quotes_ws, quote_header_row); lh = hmap(lines_ws, line_header_row); sh = hmap(summary_ws, summary_header_row)

def clear_from(ws, start_row):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)
clear_from(quotes_ws, 5)
clear_from(lines_ws, 5)
clear_from(summary_ws, 5)

def style_from_header(ws, row, max_col):
    for c in range(1, max_col+1):
        src = ws.cell(row-1, c)
        dst = ws.cell(row, c)
        if src.has_style:
            dst._style = src._style.copy()
        if src.number_format:
            dst.number_format = src.number_format

def set_cell(ws, row, header_map, header, value):
    if header in header_map:
        ws.cell(row, header_map[header], value)

def get_lines(doc):
    lines = get_any(doc, ['Lines','lines','LineItems','line_items','items'], [])
    if isinstance(lines, dict):
        lines = pick_list(lines)
    if not isinstance(lines, list):
        lines = []
    return lines

summary = defaultdict(lambda: {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
line_count = 0
for idx, doc in enumerate(quote_docs, start=5):
    style_from_header(quotes_ws, idx, len(quote_headers))
    qid = str(get_any(doc, ['QuoteID','quote_id','id','quoteId'], '')).strip()
    cust = str(get_any(doc, ['CustomerID','customer_id','customerId'], customer_id)).strip()
    project = str(get_any(doc, ['ProjectID','project_id','projectId'], '') or '')
    qdate = parse_date_value(get_any(doc, ['QuoteDate','quote_date','date']))
    valid = parse_date_value(get_any(doc, ['ValidUntil','valid_until','validUntil']))
    status = str(get_any(doc, ['Status','status'], '') or '')
    currency = str(get_any(doc, ['Currency','currency'], 'TWD') or 'TWD')
    subtotal = q2(get_any(doc, ['Subtotal','subtotal'], '0'))
    disc = q2(get_any(doc, ['DiscountPercent','discount_percent','discountPercent'], '0'))
    net = get_any(doc, ['NetAmount','net_amount','netAmount','TotalAmount','total_amount','totalAmount'], None)
    if net is None:
        net = subtotal * (Decimal('1') - disc / Decimal('100'))
    net = q2(net)
    rate = fx.get(currency, Decimal('1'))
    net_twd = q2(net * rate)
    values = {
        'QuoteID': qid, 'CustomerID': cust, 'ProjectID': project, 'QuoteDate': qdate,
        'ValidUntil': valid, 'Status': status, 'Currency': currency, 'Subtotal': dec_num(subtotal),
        'DiscountPercent': dec_num(disc), 'NetAmount': dec_num(net), 'NetAmountTWD': dec_num(net_twd),
    }
    for h, v in values.items():
        set_cell(quotes_ws, idx, qh, h, v)
    for h in ['QuoteID','CustomerID','ProjectID','Status','Currency']:
        if h in qh: quotes_ws.cell(idx, qh[h]).number_format = '@'
    for h in ['QuoteDate','ValidUntil']:
        if h in qh: quotes_ws.cell(idx, qh[h]).number_format = 'yyyy-mm-dd'
    for h in ['Subtotal','DiscountPercent','NetAmount','NetAmountTWD']:
        if h in qh: quotes_ws.cell(idx, qh[h]).number_format = '0.00'
    summary[(status, currency)]['amount'] += net
    summary[(status, currency)]['amount_twd'] += net_twd
    summary[(status, currency)]['count'] += 1
    for line in get_lines(doc):
        row = 5 + line_count
        style_from_header(lines_ws, row, len(line_headers))
        product_id = str(get_any(line, ['ProductID','product_id','productId','LineID','line_id'], '') or '')
        product_name = str(get_any(line, ['ProductName','product_name','productName','Product','product'], '') or '')
        qty = q2(get_any(line, ['Quantity','quantity','qty'], '0'))
        unit = q2(get_any(line, ['UnitPrice','unit_price','unitPrice','price'], '0'))
        amount = get_any(line, ['LineAmount','line_amount','lineAmount','Amount','amount'], None)
        if amount is None:
            amount = qty * unit
        amount = q2(amount)
        line_values = {'QuoteID': qid, 'ProductID': product_id, 'ProductName': product_name, 'Quantity': dec_num(qty), 'UnitPrice': dec_num(unit), 'LineAmount': dec_num(amount)}
        for h, v in line_values.items():
            set_cell(lines_ws, row, lh, h, v)
        for h in ['QuoteID','ProductID','ProductName']:
            if h in lh: lines_ws.cell(row, lh[h]).number_format = '@'
        for h in ['Quantity','UnitPrice','LineAmount']:
            if h in lh: lines_ws.cell(row, lh[h]).number_format = '0.00' if h != 'Quantity' else '0'
        line_count += 1

for i, ((status, currency), agg) in enumerate(sorted(summary.items()), start=5):
    style_from_header(summary_ws, i, len(summary_headers))
    vals = {'Group': 'Status', 'SubGroup': status, 'Currency': currency, 'Amount': dec_num(agg['amount']), 'AmountTWD': dec_num(agg['amount_twd']), 'Count': int(agg['count'])}
    for h, v in vals.items():
        set_cell(summary_ws, i, sh, h, v)
    for h in ['Group','SubGroup','Currency']:
        if h in sh: summary_ws.cell(i, sh[h]).number_format = '@'
    for h in ['Amount','AmountTWD']:
        if h in sh: summary_ws.cell(i, sh[h]).number_format = '0.00'
    if 'Count' in sh: summary_ws.cell(i, sh['Count']).number_format = '0'
last_quote_row = 4 + len(quote_docs)
summary_ws['B2'] = '=SUM(Quotes!K5:K5)' if not quote_docs else f'=SUM(Quotes!K5:K{last_quote_row})'
report_ws['B3'] = start_dt
report_ws['D3'] = end_dt
report_ws['B3'].number_format = 'yyyy-mm-dd'
report_ws['D3'].number_format = 'yyyy-mm-dd'
report_ws['B4'] = f'CustomerID={customer_id}; QuoteDate {start_s} to {end_s} inclusive'
now = datetime.now().astimezone().replace(microsecond=0).isoformat()
report_ws['B5'] = now
report_ws['B7'] = 'quote_cli; fx_rates.csv'
report_ws['D7'] = now

out_name = f'quote_report_{customer_id}_{start_s}_to_{end_s}.xlsx'
out_path = output_dir / out_name
if out_path.exists():
    out_path.unlink()
wb.save(out_path)

# Verify saved file.
check = load_workbook(out_path, data_only=False)
cq = check['Quotes']; cl = check['Lines']; cs = check['Summary']; cr = check['Report']
def as_date(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if isinstance(v, str): return datetime.strptime(v[:10], '%Y-%m-%d').date()
    return v
assert [cq.cell(4,c).value for c in range(1, len(quote_headers)+1)] == quote_headers, 'Quotes headers changed'
assert [cl.cell(4,c).value for c in range(1, len(line_headers)+1)] == line_headers, 'Lines headers changed'
assert [cs.cell(4,c).value for c in range(1, len(summary_headers)+1)] == summary_headers, 'Summary headers changed'
assert as_date(cr['B3'].value) == start_dt, f'Report!B3 {cr["B3"].value!r} != {start_dt}'
assert as_date(cr['D3'].value) == end_dt, f'Report!D3 {cr["D3"].value!r} != {end_dt}'
assert cq.max_row == (4 if not quote_docs else 4 + len(quote_docs)), f'Quotes rows {cq.max_row}'
assert cl.max_row == (4 if not quote_docs else 4 + line_count), f'Lines rows {cl.max_row}'
for r in range(5, 5 + len(quote_docs)):
    for c in [qh.get('QuoteID'), qh.get('CustomerID'), qh.get('ProjectID')]:
        if c: assert isinstance(cq.cell(r,c).value, str), f'Quotes text column {c} row {r}'
    for c in [qh.get('Subtotal'), qh.get('DiscountPercent'), qh.get('NetAmount'), qh.get('NetAmountTWD')]:
        if c: assert isinstance(cq.cell(r,c).value, (int,float)), f'Quotes numeric column {c} row {r}'
for r in range(5, 5 + line_count):
    for c in [lh.get('QuoteID'), lh.get('ProductID')]:
        if c: assert isinstance(cl.cell(r,c).value, str), f'Lines text column {c} row {r}'
    for c in [lh.get('Quantity'), lh.get('UnitPrice'), lh.get('LineAmount')]:
        if c: assert isinstance(cl.cell(r,c).value, (int,float)), f'Lines numeric column {c} row {r}'
expected_formula = '=SUM(Quotes!K5:K5)' if not quote_docs else f'=SUM(Quotes!K5:K{last_quote_row})'
assert cs['B2'].value == expected_formula, f'Summary!B2 {cs["B2"].value!r} != {expected_formula!r}'
saved_summary = {}
for r in range(5, cs.max_row+1):
    if cs.cell(r,1).value is None:
        continue
    key = (cs.cell(r, sh['SubGroup']).value, cs.cell(r, sh['Currency']).value)
    saved_summary[key] = {'amount': q2(cs.cell(r, sh['Amount']).value), 'amount_twd': q2(cs.cell(r, sh['AmountTWD']).value), 'count': int(cs.cell(r, sh['Count']).value)}
expected_summary = {k: {'amount': q2(v['amount']), 'amount_twd': q2(v['amount_twd']), 'count': int(v['count'])} for k, v in summary.items()}
assert saved_summary == expected_summary, f'Summary mismatch saved={saved_summary} expected={expected_summary}'
print('Output workbook:', out_path)
print('Checked: template sheets preserved; Report dates and condition set; Quotes rows:', len(quote_docs), '; Lines rows:', line_count, '; Summary rows:', len(expected_summary), '; Summary!B2:', expected_formula, '; ID columns text; numeric amount columns numeric.')
print('Output files:', sorted(p.name for p in output_dir.iterdir()))
