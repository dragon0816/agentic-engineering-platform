from pathlib import Path
from openpyxl import load_workbook
import csv

ws = Path.cwd()
print('workspace=', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\n--- quote_cli_manual.md first 4000 chars ---')
    print(manual.read_text(encoding='utf-8', errors='replace')[:4000])

fx = ws/'input'/'fx_rates.csv'
if fx.exists():
    print('\n--- fx_rates.csv sample ---')
    with fx.open(newline='', encoding='utf-8-sig') as f:
        r = csv.reader(f)
        for i,row in zip(range(8), r):
            print(row)

tpl = ws/'input'/'quote_template.xlsx'
if tpl.exists():
    print('\n--- workbook sheets/dimensions/headers/sample rows ---')
    wb = load_workbook(tpl, data_only=False)
    for sh in wb.sheetnames:
        wsht = wb[sh]
        print(f'SHEET {sh}: max_row={wsht.max_row} max_col={wsht.max_column}')
        for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 8), values_only=True):
            print(row)
        print('---')
