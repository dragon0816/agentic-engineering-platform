from pathlib import Path
import subprocess, json, csv, os, sys
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from datetime import datetime, date
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TEMPLATE = INPUT / 'quote_template.xlsx'
FX_FILE = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'quote_template.xlsx'
FINDINGS = OUTPUT / 'findings.md'
FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
GENERATED_AT = datetime.now().astimezone().isoformat(timespec='seconds')

TWOPL = Decimal('0.01')
def D(x):
    if x is None or x == '':
        return None
    return Decimal(str(x))

def q2(x: Decimal) -> Decimal:
    return x.quantize(TWOPL, rounding=ROUND_HALF_UP)

def run_cli(args, retries=2):
    cmd = [sys.executable, str(CLI)] + args
    last = None
    for attempt in range(retries + 1):
        cp = subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
        last = cp
        if cp.returncode == 0:
            if cp.stdout.strip():
                return json.loads(cp.stdout)
            return None
        if cp.returncode == 4 and attempt < retries:
            continue
        raise RuntimeError(f'CLI failed exit={cp.returncode} args={args}\nSTDERR={cp.stderr}\nSTDOUT={cp.stdout}')
    raise RuntimeError('unreachable')

print('Input files:')
for p in sorted(INPUT.iterdir()):
    print(f'- {p.name} ({p.stat().st_size} bytes)')
print('Using CLI:', CLI)

# Load FX rates robustly because Note contains commas in this synthetic CSV.
fx = {}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    rdr = csv.reader(f)
    header = next(rdr)
    print('FX header:', header)
    for row in rdr:
        if not row or not row[0]:
            continue
        fx[row[0]] = D(row[1])
print('FX rates:', {k: str(v) for k, v in fx.items()})

listing = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json'])
items = listing.get('items', [])
print(f'CLI list count: reported={listing.get("count")} actual_items={len(items)}')
print('List fields:', list(items[0].keys()) if items else [])
print('List sample:', items[:2])

# Get each quote detail.
quotes = []
all_lines = []
for it in items:
    qid = it['QuoteID']
    detail = run_cli(['quotes', 'get', qid, '--format', 'json'])
    h = detail['header']
    lines = detail.get('line_items', [])
    quotes.append(h)
    for ln in lines:
        row = dict(ln)
        row['QuoteID'] = qid
        all_lines.append(row)
print(f'Detail fetched: quotes={len(quotes)} lines={len(all_lines)}')
print('Header fields:', list(quotes[0].keys()) if quotes else [])
print('Line fields:', list(all_lines[0].keys()) if all_lines else [])
print('First quote detail:', quotes[0] if quotes else None)
print('First line:', all_lines[0] if all_lines else None)

# Sort to make output deterministic.
quotes.sort(key=lambda h: (h.get('QuoteDate',''), h.get('QuoteID','')))
quote_order = {h['QuoteID']: i for i,h in enumerate(quotes)}
all_lines.sort(key=lambda l: (quote_order.get(l['QuoteID'], 10**9), l.get('ProductID',''), l.get('ProductName','')))

# Build consistency checks.
lines_by_q = defaultdict(list)
for ln in all_lines:
    lines_by_q[ln['QuoteID']].append(ln)

check_rows = []
expired_unaccepted = []
missing_fx = set()
for h in quotes:
    qid = h['QuoteID']
    currency = h.get('Currency')
    subtotal = D(h.get('Subtotal'))
    discount = D(h.get('DiscountPercent')) or Decimal('0')
    net = D(h.get('NetAmount'))
    line_sum = Decimal('0.00')
    bad_lines = []
    for idx, ln in enumerate(lines_by_q[qid], start=1):
        qty = D(ln.get('Quantity')) or Decimal('0')
        unit = D(ln.get('UnitPrice')) or Decimal('0')
        line_amount = D(ln.get('LineAmount')) or Decimal('0')
        expected_line = q2(qty * unit)
        line_sum += line_amount
        if q2(line_amount) != expected_line:
            bad_lines.append(f"line {idx} {ln.get('ProductID')}: LineAmount {line_amount} != Quantity*UnitPrice {expected_line}")
    line_sum = q2(line_sum)
    expected_net = q2(subtotal * (Decimal('1') - discount / Decimal('100'))) if subtotal is not None else None
    subtotal_ok = (subtotal is not None and q2(subtotal) == line_sum)
    net_ok = (net is not None and expected_net is not None and q2(net) == expected_net)
    line_formula_ok = not bad_lines
    status = h.get('Status') or ''
    vu = date.fromisoformat(h['ValidUntil']) if h.get('ValidUntil') else None
    is_expired_unaccepted = bool(vu and vu < BASE_DATE and status != 'Accepted')
    if is_expired_unaccepted:
        expired_unaccepted.append(h)
    if currency not in fx:
        missing_fx.add(currency)
    check_rows.append({
        'QuoteID': qid,
        'Currency': currency,
        'Status': status,
        'ValidUntil': h.get('ValidUntil'),
        'LineCount': len(lines_by_q[qid]),
        'LineSum': line_sum,
        'Subtotal': subtotal,
        'SubtotalOK': subtotal_ok,
        'DiscountPercent': discount,
        'ExpectedNetAmount': expected_net,
        'NetAmount': net,
        'NetAmountOK': net_ok,
        'LineFormulaOK': line_formula_ok,
        'LineFormulaIssues': '; '.join(bad_lines),
        'ExpiredUnaccepted': is_expired_unaccepted,
    })

if missing_fx:
    raise RuntimeError(f'Missing FX rate(s): {sorted(missing_fx)}')

# Fill workbook from template copy.
wb = load_workbook(TEMPLATE)

# Helper clearing existing detail rows from row 5 down while preserving sheets.
def clear_rows(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def style_like(ws, row, src_row=5):
    # If no template data row exists, copy from header row lightly for newly inserted data row.
    source = src_row if ws.max_row >= src_row else 4
    for col in range(1, ws.max_column + 1):
        src = ws.cell(source, col)
        dst = ws.cell(row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

# Report sheet
rpt = wb['Report']
rpt['B3'] = FROM_DATE
rpt['D3'] = TO_DATE
rpt['B4'] = 'QuoteDate 2026-09-01..2026-09-30 inclusive; all quotes retained; expired flag uses ValidUntil < 2026-09-30 and Status <> Accepted.'
rpt['B5'] = GENERATED_AT
rpt['B7'] = 'quote_cli mock; fx_rates.csv'
rpt['D7'] = GENERATED_AT
bad_count = sum(1 for c in check_rows if not (c['SubtotalOK'] and c['NetAmountOK'] and c['LineFormulaOK']))
rpt['B9'] = f'Quotes={len(quotes)}, lines={len(all_lines)}, consistency issues={bad_count}, expired-unaccepted={len(expired_unaccepted)}. See findings.md.'

# Quotes sheet
qws = wb['Quotes']
clear_rows(qws, 5)
quote_headers = [qws.cell(4, c).value for c in range(1, qws.max_column + 1)]
for idx, h in enumerate(quotes, start=5):
    style_like(qws, idx)
    vals = [
        h.get('QuoteID'), h.get('CustomerID'), h.get('ProjectID'),
        date.fromisoformat(h['QuoteDate']) if h.get('QuoteDate') else None,
        date.fromisoformat(h['ValidUntil']) if h.get('ValidUntil') else None,
        h.get('Status'), h.get('Currency'),
        float(q2(D(h.get('Subtotal')))),
        float(D(h.get('DiscountPercent'))),
        float(q2(D(h.get('NetAmount')))),
        float(q2(D(h.get('NetAmount')) * fx[h.get('Currency')]))
    ]
    for col, val in enumerate(vals, start=1):
        cell = qws.cell(idx, col, val)
        if col in (1,2,3,6,7):
            cell.number_format = '@'
        elif col in (4,5):
            cell.number_format = 'yyyy-mm-dd'
        elif col in (8,10,11):
            cell.number_format = '0.00'
        elif col == 9:
            cell.number_format = '0.00'

# Lines sheet
lws = wb['Lines']
clear_rows(lws, 5)
for idx, ln in enumerate(all_lines, start=5):
    style_like(lws, idx)
    vals = [
        ln.get('QuoteID'), ln.get('ProductID'), ln.get('ProductName'),
        int(D(ln.get('Quantity'))) if D(ln.get('Quantity')) == D(ln.get('Quantity')).to_integral_value() else float(D(ln.get('Quantity'))),
        float(q2(D(ln.get('UnitPrice')))),
        float(q2(D(ln.get('LineAmount')))),
    ]
    for col, val in enumerate(vals, start=1):
        cell = lws.cell(idx, col, val)
        if col in (1,2,3):
            cell.number_format = '@'
        elif col in (5,6):
            cell.number_format = '0.00'

# Summary sheet: group by Status/Currency and all/currency total.
sws = wb['Summary']
clear_rows(sws, 5)
last_quote_row = 4 + len(quotes) if quotes else 5
sws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row if quotes else 5})'
summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0})
for h in quotes:
    key = (h.get('Status'), h.get('Currency'))
    net = q2(D(h.get('NetAmount')))
    amount_twd = q2(net * fx[h.get('Currency')])
    summary[key]['amount'] += net
    summary[key]['amount_twd'] += amount_twd
    summary[key]['count'] += 1
row = 5
for (status, curr), val in sorted(summary.items(), key=lambda kv: (kv[0][0] or '', kv[0][1] or '')):
    style_like(sws, row)
    values = ['Status', status, curr, float(q2(val['amount'])), float(q2(val['amount_twd'])), val['count']]
    for col, v in enumerate(values, start=1):
        cell = sws.cell(row, col, v)
        if col in (4,5):
            cell.number_format = '0.00'
    row += 1
# also summarize by Currency as a separate Group for useful review
for curr in sorted({h.get('Currency') for h in quotes}):
    val = {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0}
    for h in quotes:
        if h.get('Currency') == curr:
            net = q2(D(h.get('NetAmount')))
            val['amount'] += net
            val['amount_twd'] += q2(net * fx[curr])
            val['count'] += 1
    style_like(sws, row)
    values = ['Currency', curr, curr, float(q2(val['amount'])), float(q2(val['amount_twd'])), val['count']]
    for col, v in enumerate(values, start=1):
        cell = sws.cell(row, col, v)
        if col in (4,5):
            cell.number_format = '0.00'
    row += 1

# Findings markdown.
lines = []
lines.append('# Quote findings')
lines.append('')
lines.append(f'- Query period: {FROM_DATE} to {TO_DATE} inclusive')
lines.append(f'- Baseline date for expiry check: {BASE_DATE.isoformat()}')
lines.append(f'- Source: quote_cli mock; FX: fx_rates.csv')
lines.append(f'- Generated at: {GENERATED_AT}')
lines.append(f'- Quotes retained in workbook: {len(quotes)}')
lines.append(f'- Line items retained in workbook: {len(all_lines)}')
lines.append('')
lines.append('## Consistency checks')
lines.append('')
lines.append('Rules: `Subtotal` must equal the sum of line `LineAmount`; each line `LineAmount` should equal `Quantity × UnitPrice`; `NetAmount` must equal `Subtotal × (1 - DiscountPercent / 100)`, rounded ROUND_HALF_UP to 2 decimals.')
lines.append('')
issues = [c for c in check_rows if not (c['SubtotalOK'] and c['NetAmountOK'] and c['LineFormulaOK'])]
if not issues:
    lines.append('No consistency differences found.')
else:
    lines.append('| QuoteID | Subtotal | LineSum | SubtotalOK | DiscountPercent | ExpectedNetAmount | NetAmount | NetAmountOK | LineAmount issues |')
    lines.append('|---|---:|---:|:---:|---:|---:|---:|:---:|---|')
    for c in issues:
        lines.append(f"| {c['QuoteID']} | {c['Subtotal']} | {c['LineSum']} | {c['SubtotalOK']} | {c['DiscountPercent']} | {c['ExpectedNetAmount']} | {c['NetAmount']} | {c['NetAmountOK']} | {c['LineFormulaIssues'] or '-'} |")
lines.append('')
lines.append('### Per-quote check table')
lines.append('')
lines.append('| QuoteID | Status | Currency | LineCount | LineSum | Subtotal | SubtotalOK | DiscountPercent | ExpectedNetAmount | NetAmount | NetAmountOK |')
lines.append('|---|---|---|---:|---:|---:|:---:|---:|---:|---:|:---:|')
for c in check_rows:
    lines.append(f"| {c['QuoteID']} | {c['Status']} | {c['Currency']} | {c['LineCount']} | {c['LineSum']} | {c['Subtotal']} | {c['SubtotalOK']} | {c['DiscountPercent']} | {c['ExpectedNetAmount']} | {c['NetAmount']} | {c['NetAmountOK']} |")
lines.append('')
lines.append('## ValidUntil < 2026-09-30 and not Accepted')
lines.append('')
if not expired_unaccepted:
    lines.append('No quotes match this condition.')
else:
    lines.append('| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |')
    lines.append('|---|---|---|---|---|---|---|---:|')
    for h in sorted(expired_unaccepted, key=lambda x: (x.get('ValidUntil',''), x.get('QuoteID',''))):
        lines.append(f"| {h.get('QuoteID')} | {h.get('CustomerID')} | {h.get('ProjectID')} | {h.get('QuoteDate')} | {h.get('ValidUntil')} | {h.get('Status')} | {h.get('Currency')} | {q2(D(h.get('NetAmount')))} |")
lines.append('')
lines.append('## Workbook notes')
lines.append('')
lines.append('- All September quote headers and line items returned by the CLI were written; expired quotes were flagged in this findings file only and were not removed from the workbook.')
lines.append('- NetAmountTWD uses fixed test rates from fx_rates.csv and ROUND_HALF_UP to 2 decimals.')
FINDINGS.write_text('\n'.join(lines) + '\n', encoding='utf-8')

wb.save(OUT_XLSX)
print('Wrote:', OUT_XLSX)
print('Wrote:', FINDINGS)

# Validation: reopen outputs and verify stated requirements.
print('\nVALIDATION')
wb2 = load_workbook(OUT_XLSX, data_only=False)
assert wb2.sheetnames == ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions'], wb2.sheetnames
q2ws = wb2['Quotes']
l2ws = wb2['Lines']
s2ws = wb2['Summary']
actual_quote_rows = max(0, q2ws.max_row - 4)
actual_line_rows = max(0, l2ws.max_row - 4)
assert actual_quote_rows == len(quotes), (actual_quote_rows, len(quotes))
assert actual_line_rows == len(all_lines), (actual_line_rows, len(all_lines))
assert [q2ws.cell(4,c).value for c in range(1,12)] == ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [l2ws.cell(4,c).value for c in range(1,7)] == ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert s2ws['B2'].value == f'=SUM(Quotes!K5:K{last_quote_row if quotes else 5})'
# Ensure IDs are strings and numeric fields are numeric cells.
for r in range(5, 5 + len(quotes)):
    assert isinstance(q2ws.cell(r,1).value, str)
    assert isinstance(q2ws.cell(r,2).value, str)
    assert isinstance(q2ws.cell(r,3).value, str)
    assert isinstance(q2ws.cell(r,8).value, (int,float))
    assert isinstance(q2ws.cell(r,10).value, (int,float))
    assert isinstance(q2ws.cell(r,11).value, (int,float))
for r in range(5, 5 + len(all_lines)):
    assert isinstance(l2ws.cell(r,1).value, str)
    assert isinstance(l2ws.cell(r,2).value, str)
    assert isinstance(l2ws.cell(r,4).value, (int,float))
    assert isinstance(l2ws.cell(r,5).value, (int,float))
    assert isinstance(l2ws.cell(r,6).value, (int,float))
md = FINDINGS.read_text(encoding='utf-8')
assert 'ValidUntil < 2026-09-30 and not Accepted' in md
for h in expired_unaccepted:
    assert h['QuoteID'] in md
# Check every workbook QuoteID from original list appears exactly once in Quotes.
out_ids = [q2ws.cell(r,1).value for r in range(5, 5 + len(quotes))]
assert sorted(out_ids) == sorted([h['QuoteID'] for h in quotes])
print(f'Checked workbook sheets, headers, row counts quotes={actual_quote_rows}, lines={actual_line_rows}, Summary!B2 range, text IDs, numeric amount cells, and findings content.')
print(f'Consistency issues={len(issues)}; expired-unaccepted={len(expired_unaccepted)}')
print('Expired-unaccepted IDs:', ', '.join(h['QuoteID'] for h in sorted(expired_unaccepted, key=lambda x: (x.get('ValidUntil',''), x.get('QuoteID','')))) or '(none)')
print('Consistency issue IDs:', ', '.join(c['QuoteID'] for c in issues) or '(none)')
