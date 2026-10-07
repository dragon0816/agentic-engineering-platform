from pathlib import Path
import subprocess, json, csv, sys
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from collections import defaultdict
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd(); INPUT = WORKSPACE/'input'; OUTPUT = WORKSPACE/'output'; OUTPUT.mkdir(exist_ok=True)
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TEMPLATE = INPUT/'quote_template.xlsx'; FX_FILE = INPUT/'fx_rates.csv'
OUT_XLSX = OUTPUT/'quote_template.xlsx'; FINDINGS = OUTPUT/'findings.md'
FROM_DATE='2026-09-01'; TO_DATE='2026-09-30'; BASE_DATE=date.fromisoformat('2026-09-30')
GENERATED_AT=datetime.now().astimezone().isoformat(timespec='seconds')
TWOPL=Decimal('0.01')
def D(x): return None if x is None or x=='' else Decimal(str(x))
def q2(x): return x.quantize(TWOPL, rounding=ROUND_HALF_UP)
def run_cli(args, retries=2):
    cmd=[sys.executable, str(CLI)] + args
    for attempt in range(retries+1):
        cp=subprocess.run(cmd, cwd=str(WORKSPACE), text=True, capture_output=True)
        if cp.returncode==0:
            return json.loads(cp.stdout) if cp.stdout.strip() else None
        if cp.returncode==4 and attempt<retries: continue
        raise RuntimeError(f'CLI failed exit={cp.returncode} args={args}\nSTDERR={cp.stderr}\nSTDOUT={cp.stdout}')

print('Input files:')
for p in sorted(INPUT.iterdir()): print(f'- {p.name} ({p.stat().st_size} bytes)')
print('Using CLI:', CLI)

fx={}
with FX_FILE.open(newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f); header=next(rdr); print('FX header:', header)
    for row in rdr:
        if row and row[0]: fx[row[0]]=D(row[1])
print('FX rates:', {k:str(v) for k,v in fx.items()})

listing=run_cli(['quotes','list','--from',FROM_DATE,'--to',TO_DATE,'--format','json'])
items=listing.get('items', [])
print(f'CLI list count: reported={listing.get("count")} actual_items={len(items)}')
print('List fields:', list(items[0].keys()) if items else [])
print('List sample:', items[:2])

quotes=[]; all_lines=[]
for it in items:
    detail=run_cli(['quotes','get',it['QuoteID'],'--format','json'])
    h=detail['header']; quotes.append(h)
    for ln in detail.get('line_items', []):
        row=dict(ln); row['QuoteID']=h['QuoteID']; all_lines.append(row)
print(f'Detail fetched: quotes={len(quotes)} lines={len(all_lines)}')
print('Header fields:', list(quotes[0].keys()) if quotes else [])
print('Line fields:', list(all_lines[0].keys()) if all_lines else [])
print('First quote detail:', quotes[0] if quotes else None)
print('First line:', all_lines[0] if all_lines else None)

quotes.sort(key=lambda h:(h.get('QuoteDate',''), h.get('QuoteID','')))
order={h['QuoteID']:i for i,h in enumerate(quotes)}
all_lines.sort(key=lambda l:(order.get(l['QuoteID'], 10**9), l.get('ProductID',''), l.get('ProductName','')))

lines_by_q=defaultdict(list)
for ln in all_lines: lines_by_q[ln['QuoteID']].append(ln)
check_rows=[]; expired_unaccepted=[]; missing_fx=set()
for h in quotes:
    qid=h['QuoteID']; subtotal=D(h.get('Subtotal')); discount=D(h.get('DiscountPercent')) or Decimal('0'); net=D(h.get('NetAmount'))
    line_sum=Decimal('0.00'); bad_lines=[]
    for idx, ln in enumerate(lines_by_q[qid], start=1):
        qty=D(ln.get('Quantity')) or Decimal('0'); unit=D(ln.get('UnitPrice')) or Decimal('0'); line_amt=D(ln.get('LineAmount')) or Decimal('0')
        exp_line=q2(qty*unit); line_sum += line_amt
        if q2(line_amt) != exp_line: bad_lines.append(f"line {idx} {ln.get('ProductID')}: LineAmount {line_amt} != Quantity*UnitPrice {exp_line}")
    line_sum=q2(line_sum); expected_net=q2(subtotal*(Decimal('1')-discount/Decimal('100'))) if subtotal is not None else None
    subtotal_ok=(subtotal is not None and q2(subtotal)==line_sum); net_ok=(net is not None and expected_net is not None and q2(net)==expected_net)
    vu=date.fromisoformat(h['ValidUntil']) if h.get('ValidUntil') else None
    is_exp=bool(vu and vu<BASE_DATE and h.get('Status')!='Accepted')
    if is_exp: expired_unaccepted.append(h)
    if h.get('Currency') not in fx: missing_fx.add(h.get('Currency'))
    check_rows.append({'QuoteID':qid,'Currency':h.get('Currency'),'Status':h.get('Status') or '', 'ValidUntil':h.get('ValidUntil'), 'LineCount':len(lines_by_q[qid]), 'LineSum':line_sum, 'Subtotal':subtotal, 'SubtotalOK':subtotal_ok, 'DiscountPercent':discount, 'ExpectedNetAmount':expected_net, 'NetAmount':net, 'NetAmountOK':net_ok, 'LineFormulaOK':not bad_lines, 'LineFormulaIssues':'; '.join(bad_lines), 'ExpiredUnaccepted':is_exp})
if missing_fx: raise RuntimeError(f'Missing FX rate(s): {sorted(missing_fx)}')

wb=load_workbook(TEMPLATE)
def clear_rows(ws, min_row=5):
    if ws.max_row>=min_row: ws.delete_rows(min_row, ws.max_row-min_row+1)
def style_like(ws, row):
    source=5 if ws.max_row>=5 else 4
    for col in range(1, ws.max_column+1):
        src=ws.cell(source,col); dst=ws.cell(row,col)
        if src.has_style: dst._style=copy(src._style)
        dst.number_format=src.number_format
        dst.alignment=copy(src.alignment); dst.font=copy(src.font); dst.fill=copy(src.fill); dst.border=copy(src.border)

rpt=wb['Report']
rpt['B3']=FROM_DATE; rpt['D3']=TO_DATE
rpt['B4']='QuoteDate 2026-09-01..2026-09-30 inclusive; all quotes retained; expired flag uses ValidUntil < 2026-09-30 and Status <> Accepted.'
rpt['B5']=GENERATED_AT; rpt['B7']='quote_cli mock; fx_rates.csv'; rpt['D7']=GENERATED_AT
bad_count=sum(1 for c in check_rows if not (c['SubtotalOK'] and c['NetAmountOK'] and c['LineFormulaOK']))
rpt['B9']=f'Quotes={len(quotes)}, lines={len(all_lines)}, consistency issues={bad_count}, expired-unaccepted={len(expired_unaccepted)}. See findings.md.'

qws=wb['Quotes']; clear_rows(qws,5)
for r,h in enumerate(quotes, start=5):
    style_like(qws,r)
    vals=[h.get('QuoteID'), h.get('CustomerID'), h.get('ProjectID') or None, date.fromisoformat(h['QuoteDate']) if h.get('QuoteDate') else None, date.fromisoformat(h['ValidUntil']) if h.get('ValidUntil') else None, h.get('Status'), h.get('Currency'), float(q2(D(h.get('Subtotal')))), float(D(h.get('DiscountPercent'))), float(q2(D(h.get('NetAmount')))), float(q2(D(h.get('NetAmount'))*fx[h.get('Currency')]))]
    for c,v in enumerate(vals, start=1):
        cell=qws.cell(r,c,v)
        if c in (1,2,3,6,7): cell.number_format='@'
        elif c in (4,5): cell.number_format='yyyy-mm-dd'
        elif c in (8,9,10,11): cell.number_format='0.00'

lws=wb['Lines']; clear_rows(lws,5)
for r,ln in enumerate(all_lines, start=5):
    style_like(lws,r)
    qty=D(ln.get('Quantity'))
    vals=[ln.get('QuoteID'), ln.get('ProductID'), ln.get('ProductName'), int(qty) if qty==qty.to_integral_value() else float(qty), float(q2(D(ln.get('UnitPrice')))), float(q2(D(ln.get('LineAmount'))))]
    for c,v in enumerate(vals, start=1):
        cell=lws.cell(r,c,v)
        if c in (1,2,3): cell.number_format='@'
        elif c in (5,6): cell.number_format='0.00'

sws=wb['Summary']; clear_rows(sws,5)
last_quote_row=4+len(quotes) if quotes else 5
sws['B2']=f'=SUM(Quotes!K5:K{last_quote_row if quotes else 5})'
summary=defaultdict(lambda:{'amount':Decimal('0.00'),'amount_twd':Decimal('0.00'),'count':0})
for h in quotes:
    key=(h.get('Status'), h.get('Currency')); net=q2(D(h.get('NetAmount'))); twd=q2(net*fx[h.get('Currency')])
    summary[key]['amount']+=net; summary[key]['amount_twd']+=twd; summary[key]['count']+=1
r=5
for (status,curr),v in sorted(summary.items(), key=lambda kv:(kv[0][0] or '', kv[0][1] or '')):
    style_like(sws,r)
    for c,val in enumerate(['Status',status,curr,float(q2(v['amount'])),float(q2(v['amount_twd'])),v['count']], start=1):
        cell=sws.cell(r,c,val)
        if c in (4,5): cell.number_format='0.00'
    r+=1
for curr in sorted({h.get('Currency') for h in quotes}):
    amt=sum((q2(D(h.get('NetAmount'))) for h in quotes if h.get('Currency')==curr), Decimal('0.00'))
    twd=sum((q2(q2(D(h.get('NetAmount')))*fx[curr]) for h in quotes if h.get('Currency')==curr), Decimal('0.00'))
    cnt=sum(1 for h in quotes if h.get('Currency')==curr)
    style_like(sws,r)
    for c,val in enumerate(['Currency',curr,curr,float(q2(amt)),float(q2(twd)),cnt], start=1):
        cell=sws.cell(r,c,val)
        if c in (4,5): cell.number_format='0.00'
    r+=1

md=[]
md += ['# Quote findings','',f'- Query period: {FROM_DATE} to {TO_DATE} inclusive',f'- Baseline date for expiry check: {BASE_DATE.isoformat()}','- Source: quote_cli mock; FX: fx_rates.csv',f'- Generated at: {GENERATED_AT}',f'- Quotes retained in workbook: {len(quotes)}',f'- Line items retained in workbook: {len(all_lines)}','']
md += ['## Consistency checks','','Rules: `Subtotal` must equal the sum of line `LineAmount`; each line `LineAmount` should equal `Quantity × UnitPrice`; `NetAmount` must equal `Subtotal × (1 - DiscountPercent / 100)`, rounded ROUND_HALF_UP to 2 decimals.','']
issues=[c for c in check_rows if not (c['SubtotalOK'] and c['NetAmountOK'] and c['LineFormulaOK'])]
if not issues: md.append('No consistency differences found.')
else:
    md += ['| QuoteID | Subtotal | LineSum | SubtotalOK | DiscountPercent | ExpectedNetAmount | NetAmount | NetAmountOK | LineAmount issues |','|---|---:|---:|:---:|---:|---:|---:|:---:|---|']
    for c in issues: md.append(f"| {c['QuoteID']} | {c['Subtotal']} | {c['LineSum']} | {c['SubtotalOK']} | {c['DiscountPercent']} | {c['ExpectedNetAmount']} | {c['NetAmount']} | {c['NetAmountOK']} | {c['LineFormulaIssues'] or '-'} |")
md += ['','### Per-quote check table','','| QuoteID | Status | Currency | LineCount | LineSum | Subtotal | SubtotalOK | DiscountPercent | ExpectedNetAmount | NetAmount | NetAmountOK |','|---|---|---|---:|---:|---:|:---:|---:|---:|---:|:---:|']
for c in check_rows: md.append(f"| {c['QuoteID']} | {c['Status']} | {c['Currency']} | {c['LineCount']} | {c['LineSum']} | {c['Subtotal']} | {c['SubtotalOK']} | {c['DiscountPercent']} | {c['ExpectedNetAmount']} | {c['NetAmount']} | {c['NetAmountOK']} |")
md += ['','## ValidUntil < 2026-09-30 and not Accepted','']
if not expired_unaccepted: md.append('No quotes match this condition.')
else:
    md += ['| QuoteID | CustomerID | ProjectID | QuoteDate | ValidUntil | Status | Currency | NetAmount |','|---|---|---|---|---|---|---|---:|']
    for h in sorted(expired_unaccepted, key=lambda x:(x.get('ValidUntil',''), x.get('QuoteID',''))): md.append(f"| {h.get('QuoteID')} | {h.get('CustomerID')} | {h.get('ProjectID') or ''} | {h.get('QuoteDate')} | {h.get('ValidUntil')} | {h.get('Status')} | {h.get('Currency')} | {q2(D(h.get('NetAmount')))} |")
md += ['','## Workbook notes','','- All September quote headers and line items returned by the CLI were written; expired quotes were flagged in this findings file only and were not removed from the workbook.','- Missing source values are left blank.','- NetAmountTWD uses fixed test rates from fx_rates.csv and ROUND_HALF_UP to 2 decimals.']
FINDINGS.write_text('\n'.join(md)+'\n', encoding='utf-8')
wb.save(OUT_XLSX)
print('Wrote:', OUT_XLSX); print('Wrote:', FINDINGS)

print('\nVALIDATION')
wb2=load_workbook(OUT_XLSX, data_only=False)
assert wb2.sheetnames==['Report','Quotes','Lines','Summary','Instructions']
q=l= None
qws=wb2['Quotes']; lws=wb2['Lines']; sws=wb2['Summary']
assert max(0,qws.max_row-4)==len(quotes), (qws.max_row, len(quotes))
assert max(0,lws.max_row-4)==len(all_lines), (lws.max_row, len(all_lines))
assert [qws.cell(4,c).value for c in range(1,12)]==['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
assert [lws.cell(4,c).value for c in range(1,7)]==['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
assert sws['B2'].value==f'=SUM(Quotes!K5:K{last_quote_row if quotes else 5})'
for rr in range(5,5+len(quotes)):
    for cc in (1,2,3,6,7):
        val=qws.cell(rr,cc).value
        assert val is None or isinstance(val,str), (rr,cc,val,type(val))
    for cc in (8,9,10,11): assert isinstance(qws.cell(rr,cc).value,(int,float)), (rr,cc,qws.cell(rr,cc).value)
for rr in range(5,5+len(all_lines)):
    for cc in (1,2,3): assert lws.cell(rr,cc).value is None or isinstance(lws.cell(rr,cc).value,str), (rr,cc)
    for cc in (4,5,6): assert isinstance(lws.cell(rr,cc).value,(int,float)), (rr,cc,lws.cell(rr,cc).value)
text=FINDINGS.read_text(encoding='utf-8')
assert 'ValidUntil < 2026-09-30 and not Accepted' in text
for h in expired_unaccepted: assert h['QuoteID'] in text
out_ids=[qws.cell(rr,1).value for rr in range(5,5+len(quotes))]
assert sorted(out_ids)==sorted([h['QuoteID'] for h in quotes])
print(f'Checked output workbook sheet names, unchanged headers, quote row count={len(quotes)}, line row count={len(all_lines)}, Summary!B2 range, ID/text fields including blanks, numeric amount fields, all QuoteIDs retained, and findings.md expired list.')
print(f'Consistency issue quote count={len(issues)}; IDs={", ".join(c["QuoteID"] for c in issues) or "(none)"}')
print(f'Expired-unaccepted quote count={len(expired_unaccepted)}; IDs={", ".join(h["QuoteID"] for h in sorted(expired_unaccepted, key=lambda x:(x.get("ValidUntil",""), x.get("QuoteID","")))) or "(none)"}')
