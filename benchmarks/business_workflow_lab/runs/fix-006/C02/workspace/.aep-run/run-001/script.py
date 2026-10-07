from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size, 'bytes')

# workbook structure
xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWORKBOOK SHEETS:', wb.sheetnames)
for sh in wb.worksheets:
    print(f'-- sheet {sh.title}: {sh.max_row}x{sh.max_column}')
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(row)

# fx rates sample
csvp = ws/'input'/'fx_rates.csv'
print('\nFX_RATES')
with csvp.open(newline='', encoding='utf-8-sig') as f:
    rdr = csv.reader(f)
    for i,row in zip(range(8), rdr):
        print(row)

# manual first lines
md = ws/'input'/'quote_cli_manual.md'
print('\nMANUAL PREVIEW')
text = md.read_text(encoding='utf-8')
print(text[:4000])