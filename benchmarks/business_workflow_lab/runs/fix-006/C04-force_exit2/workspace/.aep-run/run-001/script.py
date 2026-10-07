from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
print('OUTPUT EXISTS', (ws/'output').exists())

# Inspect workbook
xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('WORKBOOK SHEETS', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'SHEET {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
        print('  ', r)

# Inspect FX CSV
csvp = ws/'input'/'fx_rates.csv'
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    rows = []
    for i, row in enumerate(reader):
        rows.append(row)
        if i >= 5: break
print('FX SAMPLE')
for row in rows:
    print(' ', row)

# Inspect manual
md = ws/'input'/'quote_cli_manual.md'
text = md.read_text(encoding='utf-8')
print('MANUAL FIRST 2000 CHARS')
print(text[:2000])