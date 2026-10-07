from pathlib import Path
import subprocess, sys, json, csv, os, shutil, datetime, decimal, traceback
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook
from openpyxl.styles import numbers

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

CLI_PY = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
TEMPLATE = input_dir / 'quote_template.xlsx'
FX_CSV = input_dir / 'fx_rates.csv'
OUT_XLSX = output_dir / 'sent_quotes_2026-09.xlsx'
LOG = output_dir / 'quote_cli_run_log.txt'

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'

print('TASK INPUTS')
print('workspace=', workspace)
print('template=', TEMPLATE, TEMPLATE.exists())
print('fx_rates=', FX_CSV, FX_CSV.exists())
print('cli_py=', CLI_PY, CLI_PY.exists())
print('query=', FROM_DATE, TO_DATE, STATUS)

# Clean prior outputs from this script to avoid stale delivery on failure.
for p in [OUT_XLSX, LOG]:
    if p.exists():
        p.unlink()

log_lines = []
def log(msg):
    print(msg)
    log_lines.append(str(msg))

class CliFailure(Exception):
    def __init__(self, exit_code, stderr, command):
        super().__init__(f'CLI failed with exit code {exit_code}: {stderr}')
        self.exit_code = exit_code
        self.stderr = stderr
        self.command = command


def run_cli(args, parse_json=True):
    # exit code 4: max two retries after first attempt; 2/3: no retry.
    max_attempts = 3
    last = None
    for attempt in range(1, max_attempts + 1):
        cmd = [sys.executable, str(CLI_PY)] + args
        log(f'CLI attempt {attempt}: ' + ' '.join(cmd))
        cp = subprocess.run(cmd, cwd=str(workspace), text=True, capture_output=True)
        log(f'  exit_code={cp.returncode}')
        if cp.stderr:
            log('  stderr=' + cp.stderr.strip())
        # stdout is considered data only if success and parseable.
        if cp.returncode == 0:
            if parse_json:
                try:
                    data = json.loads(cp.stdout)
                except Exception as e:
                    raise CliFailure(0, f'stdout was not parseable JSON: {e}', cmd)
                return data
            return cp.stdout
        if cp.returncode in (2, 3):
            raise CliFailure(cp.returncode, cp.stderr.strip(), cmd)
        if cp.returncode == 4:
            last = cp
            if attempt < max_attempts:
                continue
            raise CliFailure(cp.returncode, cp.stderr.strip(), cmd)
        # Unspecified non-zero code: do not treat stdout as data and fail.
        raise CliFailure(cp.returncode, cp.stderr.strip(), cmd)

try:
    # Load FX rates preserving decimal precision. CSV notes have commas but rates are in first two columns.
    fx = {}
    with open(FX_CSV, newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        header = next(reader)
        for row in reader:
            if not row or not row[0]:
                continue
            fx[row[0]] = Decimal(row[1])
    log('FX rates: ' + repr(fx))

    list_data = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json'])
    if list_data.get('schema_version') != 1 or 'items' not in list_data:
        raise RuntimeError('Unexpected quotes list schema')
    headers = list_data['items']
    log(f'Quote list count={list_data.get("count")} items={len(headers)}')

    details = []
    for h in headers:
        qid = h['QuoteID']
        d = run_cli(['quotes', 'get', str(qid), '--format', 'json'])
        if d.get('schema_version') != 1 or 'header' not in d or 'line_items' not in d:
            raise RuntimeError(f'Unexpected quote get schema for {qid}')
        details.append(d)
    fetched_at = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).isoformat(timespec='seconds')

    # Fill a copy of the actual template, never rebuild it.
    shutil.copyfile(TEMPLATE, OUT_XLSX)
    wb = load_workbook(OUT_XLSX)
    report = wb['Report']
    quotes_ws = wb['Quotes']
    lines_ws = wb['Lines']
    summary_ws = wb['Summary']

    # Clear existing detail data rows if any.
    for ws in [quotes_ws, lines_ws, summary_ws]:
        if ws.max_row >= 5:
            ws.delete_rows(5, ws.max_row - 4)

    report['B3'] = FROM_DATE
    report['D3'] = TO_DATE
    report['B4'] = f'QuoteDate {FROM_DATE} to {TO_DATE}, Status={STATUS}'
    report['B5'] = fetched_at
    report['B7'] = 'quote_cli'
    report['D7'] = fetched_at
    report['B9'] = ''

    money_fmt = '0.00'
    date_fmt = 'yyyy-mm-dd'
    text_fmt = '@'

    quote_rows = []
    line_rows = []
    currency_summary = {}
    customer_summary = {}

    for d in details:
        h = d['header']
        currency = h['Currency']
        if currency not in fx:
            raise RuntimeError(f'Missing FX rate for currency {currency}')
        subtotal = Decimal(str(h['Subtotal'])).quantize(Decimal('0.01'))
        disc = Decimal(str(h['DiscountPercent'])).quantize(Decimal('0.01'))
        net = Decimal(str(h['NetAmount'])).quantize(Decimal('0.01'))
        net_twd = (net * fx[currency]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
        qr = [str(h['QuoteID']), str(h['CustomerID']), str(h['ProjectID']), h['QuoteDate'], h['ValidUntil'], h['Status'], currency, float(subtotal), float(disc), float(net), float(net_twd)]
        quote_rows.append(qr)
        currency_summary.setdefault(currency, {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
        currency_summary[currency]['amount'] += net
        currency_summary[currency]['twd'] += net_twd
        currency_summary[currency]['count'] += 1
        cust = str(h['CustomerID'])
        customer_summary.setdefault(cust, {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0, 'currency': currency})
        customer_summary[cust]['amount'] += net
        customer_summary[cust]['twd'] += net_twd
        customer_summary[cust]['count'] += 1
        for li in d['line_items']:
            qty = Decimal(str(li['Quantity']))
            unit = Decimal(str(li['UnitPrice'])).quantize(Decimal('0.01'))
            amount = Decimal(str(li['LineAmount'])).quantize(Decimal('0.01'))
            line_rows.append([str(h['QuoteID']), str(li['ProductID']), str(li['ProductName']), int(qty), float(unit), float(amount)])

    for r_idx, row in enumerate(quote_rows, start=5):
        for c_idx, val in enumerate(row, start=1):
            cell = quotes_ws.cell(r_idx, c_idx, val)
            if c_idx in [1,2,3,6,7]:
                cell.number_format = text_fmt
            elif c_idx in [4,5]:
                cell.number_format = date_fmt
            elif c_idx in [8,9,10,11]:
                cell.number_format = money_fmt

    for r_idx, row in enumerate(line_rows, start=5):
        for c_idx, val in enumerate(row, start=1):
            cell = lines_ws.cell(r_idx, c_idx, val)
            if c_idx in [1,2,3]:
                cell.number_format = text_fmt
            elif c_idx in [5,6]:
                cell.number_format = money_fmt

    # Summary rows: by currency then by customer.
    sum_row = 5
    for cur in sorted(currency_summary):
        s = currency_summary[cur]
        row = ['Currency', cur, cur, float(s['amount'].quantize(Decimal('0.01'))), float(s['twd'].quantize(Decimal('0.01'))), int(s['count'])]
        for c_idx, val in enumerate(row, start=1):
            cell = summary_ws.cell(sum_row, c_idx, val)
            if c_idx in [4,5]: cell.number_format = money_fmt
            if c_idx in [1,2,3]: cell.number_format = text_fmt
        sum_row += 1
    for cust in sorted(customer_summary):
        s = customer_summary[cust]
        row = ['Customer', cust, s['currency'], float(s['amount'].quantize(Decimal('0.01'))), float(s['twd'].quantize(Decimal('0.01'))), int(s['count'])]
        for c_idx, val in enumerate(row, start=1):
            cell = summary_ws.cell(sum_row, c_idx, val)
            if c_idx in [4,5]: cell.number_format = money_fmt
            if c_idx in [1,2,3]: cell.number_format = text_fmt
        sum_row += 1

    # Update summary total formula range if there are rows.
    if quote_rows:
        summary_ws['B2'] = f'=SUM(Quotes!K5:K{4+len(quote_rows)})'
    else:
        summary_ws['B2'] = '=SUM(Quotes!K5:K5)'

    wb.save(OUT_XLSX)
    log(f'Saved workbook: {OUT_XLSX}')

    # Required saved-output checks.
    check_wb = load_workbook(OUT_XLSX, data_only=False)
    qws = check_wb['Quotes']; lws = check_wb['Lines']; sws = check_wb['Summary']; rws = check_wb['Report']
    expected_q_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
    expected_l_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
    expected_s_cols = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
    assert [qws.cell(4,c).value for c in range(1,12)] == expected_q_cols
    assert [lws.cell(4,c).value for c in range(1,7)] == expected_l_cols
    assert [sws.cell(4,c).value for c in range(1,7)] == expected_s_cols
    assert rws['B3'].value == FROM_DATE and rws['D3'].value == TO_DATE
    assert rws['B4'].value == f'QuoteDate {FROM_DATE} to {TO_DATE}, Status={STATUS}'
    actual_quote_rows = max(0, qws.max_row - 4)
    actual_line_rows = max(0, lws.max_row - 4)
    assert actual_quote_rows == len(quote_rows), (actual_quote_rows, len(quote_rows))
    assert actual_line_rows == len(line_rows), (actual_line_rows, len(line_rows))
    for r in range(5, 5 + len(quote_rows)):
        assert isinstance(qws.cell(r,1).value, str)
        assert qws.cell(r,6).value == STATUS
        assert FROM_DATE <= qws.cell(r,4).value <= TO_DATE
        assert isinstance(qws.cell(r,10).value, (int,float))
        assert isinstance(qws.cell(r,11).value, (int,float))
    for r in range(5, 5 + len(line_rows)):
        assert isinstance(lws.cell(r,1).value, str)
        assert isinstance(lws.cell(r,4).value, int)
        assert isinstance(lws.cell(r,5).value, (int,float))
        assert isinstance(lws.cell(r,6).value, (int,float))
    log('CHECKS PASSED: reopened output workbook; verified headers, report cells, row counts, text IDs, Sent/date filter, numeric money/quantity fields.')

except CliFailure as e:
    log(f'CLI_FAILURE exit_code={e.exit_code}')
    log(f'STOPPED_WITHOUT_RETRY_FOR_2_OR_3={e.exit_code in (2,3)}')
    log('STDERR_REASON=' + (e.stderr or '(empty stderr)'))
    # Do not create workbook from non-success stdout.
    LOG.write_text('\n'.join(log_lines) + '\n', encoding='utf-8')
    print('DELIVERY: no workbook was generated because CLI did not return successful parseable data.')
    print('ERROR_LOG:', LOG)
    # Exit 0 so the agent can report the controlled CLI failure, not a Python crash.
    sys.exit(0)
except Exception:
    log('UNEXPECTED_FAILURE')
    log(traceback.format_exc())
    LOG.write_text('\n'.join(log_lines) + '\n', encoding='utf-8')
    raise

LOG.write_text('\n'.join(log_lines) + '\n', encoding='utf-8')
print('OUTPUT FILES')
for p in sorted(output_dir.glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
