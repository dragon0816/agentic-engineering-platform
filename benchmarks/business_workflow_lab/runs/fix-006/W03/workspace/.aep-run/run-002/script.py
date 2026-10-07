from pathlib import Path
import csv
from openpyxl import load_workbook

base=Path('.')
print('Downloaded files:')
for p in sorted((base/'output'/'downloads').glob('*')):
    print(' -', p, p.stat().st_size)
    with p.open('r', encoding='utf-8-sig', newline='') as f:
        reader=csv.reader(f)
        for i,row in zip(range(20), reader):
            print(' row', i+1, row)

print('\nInstructions full:')
wb=load_workbook(base/'input'/'project_template.xlsx', data_only=False)
ws=wb['Instructions']
for r in range(1, ws.max_row+1):
    print(r, [ws.cell(r,c).value for c in range(1, ws.max_column+1)])
