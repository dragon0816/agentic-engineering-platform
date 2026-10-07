from pathlib import Path
import csv
from openpyxl import load_workbook

base = Path('.')
print('Workspace:', base.resolve())
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Inspect workbook
xlsx = base/'input'/'project_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print(f'\nSheet: {ws.title} max_row={ws.max_row} max_col={ws.max_column}')
    for r in range(1, min(ws.max_row, 8)+1):
        vals = [ws.cell(r,c).value for c in range(1, min(ws.max_column, 12)+1)]
        print(' row', r, vals)

# Inspect CSVs
for csv_path in sorted((base/'input').glob('*.csv')):
    print(f'\nCSV: {csv_path.name}')
    with csv_path.open('r', encoding='utf-8-sig', newline='') as f:
        reader = csv.reader(f)
        for i, row in zip(range(8), reader):
            print(' row', i+1, row)

# Inspect markdown start
md = base/'input'/'sap_manual.md'
if md.exists():
    print('\nMarkdown sap_manual.md first 120 lines:')
    text = md.read_text(encoding='utf-8', errors='replace').splitlines()
    for i,line in enumerate(text[:120], 1):
        print(f'{i:03d}: {line}')
