from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
import csv
from collections import defaultdict
from openpyxl import load_workbook

workspace = Path('.').resolve()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
download_dir = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)

TEMPLATE = input_dir / 'project_template.xlsx'
INVOICE_CSV = download_dir / 'invoices_2026-09.csv'
FX_CSV = input_dir / 'fx_rates.csv'
OUT = output_dir / 'project_invoices_2026-09.xlsx'

PERIOD_START = date(2026, 9, 1)
PERIOD_END = date(2026, 9, 30)
AS_OF = date(2026, 9, 30)
GENERATED_AT = datetime.now().astimezone().isoformat(timespec='seconds')
FETCHED_AT = GENERATED_AT

TWOPLACES = Decimal('0.01')
def money(d):
    return Decimal(str(d)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def parse_date(s):
    return datetime.strptime(s, '%Y-%m-%d').date()

def clear_data(ws, min_row=5):
    # Delete any existing data rows below header while preserving first four template rows.
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def write_cell(ws, row, col, value, number_format=None):
    cell = ws.cell(row=row, column=col, value=value)
    if number_format:
        cell.number_format = number_format
    return cell

# Load FX rates. CSV notes contain commas, but first two columns are stable.
rates = {}
with FX_CSV.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0].strip():
            continue
        rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates loaded:', rates)

# Load downloaded invoice rows.
invoices = []
with INVOICE_CSV.open('r', encoding='utf-8-sig', newline='') as f:
    reader = csv.DictReader(f)
    print('Invoice CSV headers:', reader.fieldnames)
    for raw in reader:
        inv_date = parse_date(raw['InvoiceDate'])
        if not (PERIOD_START <= inv_date <= PERIOD_END):
            raise AssertionError(f"Invoice outside requested September 2026 range: {raw}")
        currency = raw['Currency']
        if currency not in rates:
            raise AssertionError(f'Missing FX rate for {currency}')
        amount = money(raw['Amount'])
        paid = money(raw['PaidAmount'])
        balance = money(raw['Balance'])
        expected_balance = money(amount - paid)
        if balance != expected_balance:
            raise AssertionError(f"Balance mismatch for {raw['InvoiceID']}: {balance} != {expected_balance}")
        balance_twd = money(balance * rates[currency])
        row = {
            'InvoiceID': raw['InvoiceID'],
            'OrderID': raw['OrderID'],
            'ProjectID': raw['ProjectID'],
            'CustomerID': raw['CustomerID'],
            'InvoiceDate': inv_date,
            'DueDate': parse_date(raw['DueDate']),
            'Currency': currency,
            'Amount': amount,
            'PaidAmount': paid,
            'Balance': balance,
            'BalanceTWD': balance_twd,
        }
        invoices.append(row)

print('Invoice rows loaded:', len(invoices))
for r in invoices[:5]:
    print(' sample invoice:', r)

wb = load_workbook(TEMPLATE)

# Clear detail data rows in all source/detail sheets and summary/exceptions output rows.
for sheet in ['Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions']:
    clear_data(wb[sheet], 5)

# Fill Report metadata while keeping title/footer/layout.
ws_report = wb['Report']
ws_report['B3'] = PERIOD_START
ws_report['B3'].number_format = 'yyyy-mm-dd'
ws_report['D3'] = PERIOD_END
ws_report['D3'].number_format = 'yyyy-mm-dd'
ws_report['B4'] = 'Invoices: InvoiceDate 2026-09-01 through 2026-09-30; as-of 2026-09-30; overdue when Balance > 0 and DueDate < 2026-09-30.'
ws_report['B5'] = GENERATED_AT
ws_report['B7'] = 'Lab ERP Invoices; fx_rates.csv'
ws_report['D7'] = FETCHED_AT
ws_report['B9'] = 'Orders, Expenses, and Quotes detail sheets intentionally left blank per request. DueDate on 2026-09-30 is not marked overdue.'

# Fill Invoices row 5+ in template column order.
ws_inv = wb['Invoices']
invoice_headers = [ws_inv.cell(4, c).value for c in range(1, ws_inv.max_column + 1)]
expected_inv_headers = ['InvoiceID', 'OrderID', 'ProjectID', 'CustomerID', 'InvoiceDate', 'DueDate', 'Currency', 'Amount', 'PaidAmount', 'Balance', 'BalanceTWD']
if invoice_headers != expected_inv_headers:
    raise AssertionError(f'Unexpected Invoices headers: {invoice_headers}')

for idx, inv in enumerate(invoices, start=5):
    values = [inv[h] for h in expected_inv_headers]
    for col, val in enumerate(values, start=1):
        cell = ws_inv.cell(idx, col, val)
        if col in [1,2,3,4,7]:
            cell.number_format = '@'
        elif col in [5,6]:
            cell.number_format = 'yyyy-mm-dd'
        elif col in [8,9,10,11]:
            cell.value = float(val)
            cell.number_format = '0.00'

# Summary: one row per ProjectID present in invoice data, including blank ProjectID if invoices have no project.
# Only InvoiceBalanceTWD is populated; other source columns remain blank because those source sheets are intentionally blank.
summary = defaultdict(Decimal)
for inv in invoices:
    summary[inv['ProjectID']] += inv['BalanceTWD']

# Sort blank ProjectID last, named projects ascending.
project_ids = sorted([p for p in summary.keys() if p]) + ([p for p in summary.keys() if not p])
ws_sum = wb['Summary']
summary_headers = [ws_sum.cell(4, c).value for c in range(1, ws_sum.max_column + 1)]
expected_sum_headers = ['ProjectID', 'OrderTWD', 'InvoiceBalanceTWD', 'ApprovedExpenseTWD', 'AcceptedQuoteTWD']
if summary_headers != expected_sum_headers:
    raise AssertionError(f'Unexpected Summary headers: {summary_headers}')

for idx, project_id in enumerate(project_ids, start=5):
    c = ws_sum.cell(idx, 1, project_id if project_id else None)
    c.number_format = '@'
    # B blank, C invoice balance, D/E blank
    amt_cell = ws_sum.cell(idx, 3, float(money(summary[project_id])))
    amt_cell.number_format = '0.00'

last_summary_row = 4 + len(project_ids) if project_ids else 5
if project_ids:
    for col in range(2, 6):
        col_letter = chr(ord('A') + col - 1)
        ws_sum.cell(2, col, f'=SUM({col_letter}5:{col_letter}{last_summary_row})')
else:
    for col in range(2, 6):
        col_letter = chr(ord('A') + col - 1)
        ws_sum.cell(2, col, f'=SUM({col_letter}5:{col_letter}5)')

# Exceptions: Balance > 0 and DueDate < 2026-09-30. Due on 2026-09-30 is not overdue.
ws_ex = wb['Exceptions']
exception_rows = []
for inv in invoices:
    if inv['Balance'] > 0 and inv['DueDate'] < AS_OF:
        exception_rows.append(inv)
for idx, inv in enumerate(exception_rows, start=5):
    vals = ['Invoices', inv['InvoiceID'], inv['ProjectID'], 'Overdue']
    for col, val in enumerate(vals, start=1):
        cell = ws_ex.cell(idx, col, val if val != '' else None)
        if col in [1,2,3,4]:
            cell.number_format = '@'

# Format all monetary columns in blank detail sheets from row 5 if any future rows; they remain empty.

wb.save(OUT)
print('Saved:', OUT)

# Verification by reopening output workbook.
wb2 = load_workbook(OUT, data_only=False)
checks = []

# Sheets preserved.
checks.append(('sheet_names_preserved', wb2.sheetnames == ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']))

# Invoices row count, headers, ID text, date objects, numeric amounts.
ws = wb2['Invoices']
headers = [ws.cell(4, c).value for c in range(1, 12)]
checks.append(('invoice_headers_exact', headers == expected_inv_headers))
actual_invoice_rows = []
for r in range(5, ws.max_row + 1):
    if any(ws.cell(r,c).value is not None for c in range(1,12)):
        actual_invoice_rows.append(r)
checks.append(('invoice_row_count_13', len(actual_invoice_rows) == 13))
checks.append(('invoice_ids_text_leading_zero', all(isinstance(ws.cell(r,1).value, str) and ws.cell(r,1).value.startswith('00') and ws.cell(r,1).number_format == '@' for r in actual_invoice_rows)))
checks.append(('order_ids_text_leading_zero', all(isinstance(ws.cell(r,2).value, str) and ws.cell(r,2).value.startswith('00') and ws.cell(r,2).number_format == '@' for r in actual_invoice_rows)))
checks.append(('invoice_dates_are_dates', all(isinstance(ws.cell(r,5).value, datetime) or isinstance(ws.cell(r,5).value, date) for r in actual_invoice_rows)))
checks.append(('due_dates_are_dates', all(isinstance(ws.cell(r,6).value, datetime) or isinstance(ws.cell(r,6).value, date) for r in actual_invoice_rows)))
checks.append(('invoice_amounts_numeric', all(isinstance(ws.cell(r,c).value, (int,float)) for r in actual_invoice_rows for c in [8,9,10,11])))

# Check BalanceTWD values.
for r, inv in zip(actual_invoice_rows, invoices):
    got = money(ws.cell(r,11).value)
    exp = inv['BalanceTWD']
    if got != exp:
        raise AssertionError(f"BalanceTWD mismatch row {r}: {got} != {exp}")
checks.append(('balance_twd_recomputed_all_rows', True))

# Other source worksheets blank from row 5 onward.
for sheet in ['Orders','Expenses','Quotes']:
    wsx = wb2[sheet]
    nonblank = []
    for r in range(5, wsx.max_row + 1):
        if any(wsx.cell(r,c).value is not None for c in range(1, wsx.max_column+1)):
            nonblank.append(r)
    checks.append((f'{sheet}_data_rows_blank', len(nonblank) == 0))

# Summary check.
ws = wb2['Summary']
summary_rows = []
for r in range(5, ws.max_row + 1):
    if any(ws.cell(r,c).value is not None for c in range(1,6)):
        summary_rows.append(r)
checks.append(('summary_row_count_projects_including_blank', len(summary_rows) == len(project_ids)))
actual_summary = {}
for r in summary_rows:
    pid = ws.cell(r,1).value or ''
    actual_summary[pid] = money(ws.cell(r,3).value or 0)
checks.append(('summary_invoice_balance_twd_exact', actual_summary == {p: money(v) for p,v in summary.items()}))
checks.append(('summary_other_source_columns_blank', all(ws.cell(r,c).value is None for r in summary_rows for c in [2,4,5])))
checks.append(('summary_formulas_extended', [ws.cell(2,c).value for c in range(2,6)] == [f'=SUM(B5:B{last_summary_row})', f'=SUM(C5:C{last_summary_row})', f'=SUM(D5:D{last_summary_row})', f'=SUM(E5:E{last_summary_row})']))

# Exceptions check.
ws = wb2['Exceptions']
actual_ex = []
for r in range(5, ws.max_row + 1):
    if any(ws.cell(r,c).value is not None for c in range(1,5)):
        actual_ex.append([ws.cell(r,c).value for c in range(1,5)])
expected_ex = [['Invoices', inv['InvoiceID'], inv['ProjectID'] or None, 'Overdue'] for inv in exception_rows]
checks.append(('exceptions_exact_overdue_only', actual_ex == expected_ex))
checks.append(('due_2026_09_30_not_overdue', all(row[1] not in ['0090000211','0090000213','0090000219'] for row in actual_ex)))

print('\nVerification checks:')
all_ok = True
for name, ok in checks:
    print(f' - {name}:', 'OK' if ok else 'FAIL')
    all_ok = all_ok and ok
if not all_ok:
    raise AssertionError('One or more verification checks failed')

print('\nSummary computed:')
for p in project_ids:
    print(f" - ProjectID {p if p else '(blank)'}: InvoiceBalanceTWD {money(summary[p])}")
print('\nExceptions:')
for row in actual_ex:
    print(' -', row)
print('\nOutput file:', OUT)
