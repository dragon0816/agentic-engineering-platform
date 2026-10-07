import os, csv, json
from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
print('cwd', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
# workbook structure
xlsx=ws/'input'/'sales_template.xlsx'
wb=load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row,8), values_only=True):
        print(' ', r)
# csv
csvp=ws/'input'/'fx_rates.csv'
print('fx_rates sample:')
with open(csvp, newline='', encoding='utf-8-sig') as f:
    rdr=csv.reader(f)
    for i,row in zip(range(8), rdr): print(row)
# manual
md=ws/'input'/'sap_manual.md'
print('sap_manual first 80 lines:')
with open(md, encoding='utf-8') as f:
    for i,line in zip(range(80), f): print(f'{i+1}: {line.rstrip()}')
