from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

ws = Path.cwd()
out_xlsx = ws/'output'/'september_open_orders.xlsx'
csv_path = ws/'output'/'downloads'/'orders_export.csv'
fx_path = ws/'input'/'fx_rates.csv'
template_path = ws/'input'/'sales_template.xlsx'

print('Workspace:', ws)
print('Inputs:', sorted(p.name for p in (ws/'input').glob('*')))
print('Outputs:', sorted(str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()))
for p in [template_path, fx_path, csv_path, out_xlsx]:
    if not p.exists():
        raise FileNotFoundError(str(p))
    print(f'Found {p.relative_to(ws)} bytes={p.stat().st_size}')

# Read actual source structures
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    fx_rows = list(csv.DictReader(f))
print('fx_rates.csv headers:', list(fx_rows[0].keys()))
rate_col = 'RateToTWD' if 'RateToTWD' in fx_rows[0] else 'TWD_per_unit'
fx = {r['Currency']: Decimal(str(r[rate_col])) for r in fx_rows}
print('FX rates:', {k: str(v) for k, v in fx.items()})

with open(csv_path, newline='', encoding='utf-8-sig') as f:
    src_rows = list(csv.DictReader(f))
print('orders_export.csv headers:', list(src_rows[0].keys()) if src_rows else [])
print('orders_export.csv rows:', len(src_rows))
print('orders_export.csv sample:', src_rows[:2])

wb = load_workbook(out_xlsx, data_only=False)
print('Workbook sheets:', wb.sheetnames)
assert wb.sheetnames == ['Report', 'Details', 'Summary', 'Instructions'], wb.sheetnames

d = wb['Details']
headers = [d.cell(4, c).value for c in range(1, d.max_column+1)]
print('Details headers:', headers)
required = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD']
for h in required:
    assert h in headers, f'Missing Details column {h}'
col = {h: headers.index(h)+1 for h in headers if h}
rows=[]
for r in range(5, d.max_row+1):
    if d.cell(r, col['OrderID']).value in (None, ''):
        continue
    rows.append({h: d.cell(r, col[h]).value for h in required})
print('Details rows:', len(rows))
print('Details first/last:', rows[0], rows[-1])

assert len(src_rows) == 14, len(src_rows)
assert len(rows) == 14, len(rows)
ids = [str(r['OrderID']) for r in rows]
assert ids == sorted(ids), ids
assert ids == [str(r['OrderID']) for r in src_rows], 'Details rows differ from exported full filtered CSV'
assert all(len(x) == 10 and x.startswith('0') for x in ids), ids

by_cur = {}
total_twd = Decimal('0.00')
for r in rows:
    od = r['OrderDate']
    od_s = od.strftime('%Y-%m-%d') if hasattr(od, 'strftime') else str(od)[:10]
    assert '2026-09-01' <= od_s <= '2026-09-30', od_s
    assert str(r['Status']) == 'Open', r
    cur = str(r['Currency'])
    amt = Decimal(str(r['NetAmount'])).quantize(Decimal('0.01'))
    expected_twd = (amt * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    actual_twd = Decimal(str(r['NetAmountTWD'])).quantize(Decimal('0.01'))
    assert actual_twd == expected_twd, (r['OrderID'], actual_twd, expected_twd)
    by_cur.setdefault(cur, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    by_cur[cur]['amount'] += amt
    by_cur[cur]['twd'] += actual_twd
    by_cur[cur]['count'] += 1
    total_twd += actual_twd
print('Computed currency summary:', {k:{kk:str(vv) for kk,vv in v.items()} for k,v in by_cur.items()})
print('Computed total TWD:', total_twd)

# Validate Summary sheet contains per-currency original amount and TWD totals, plus all-currency TWD total.
s = wb['Summary']
print('Summary!B2 formula:', s['B2'].value)
assert s['B2'].value == '=SUM(Details!I5:I18)', s['B2'].value
summary_records = []
for r in range(1, s.max_row+1):
    vals = [s.cell(r,c).value for c in range(1, s.max_column+1)]
    if any(v is not None for v in vals):
        summary_records.append(vals)
print('Summary records:', summary_records)

# Header row in this template: Group, SubGroup, Currency, Amount, AmountTWD
summary_by_currency = {}
all_total_twd = None
for vals in summary_records:
    group = vals[0] if len(vals) > 0 else None
    subgroup = vals[1] if len(vals) > 1 else None
    currency = vals[2] if len(vals) > 2 else None
    amount = vals[3] if len(vals) > 3 else None
    amount_twd = vals[4] if len(vals) > 4 else None
    if group == 'Currency' and currency:
        summary_by_currency[str(currency)] = (Decimal(str(amount)).quantize(Decimal('0.01')), Decimal(str(amount_twd)).quantize(Decimal('0.01')))
    if group == 'Total' and subgroup == 'All currencies' and currency == 'TWD':
        all_total_twd = Decimal(str(amount_twd)).quantize(Decimal('0.01'))

expected = {
    'EUR': {'amount': Decimal('28936.37'), 'twd': Decimal('1012772.95'), 'count': 4},
    'TWD': {'amount': Decimal('637522.50'), 'twd': Decimal('637522.50'), 'count': 6},
    'USD': {'amount': Decimal('32014.00'), 'twd': Decimal('1024448.00'), 'count': 4},
}
for cur, exp in expected.items():
    assert by_cur[cur] == exp, (cur, by_cur[cur], exp)
    assert summary_by_currency[cur] == (exp['amount'], exp['twd']), (cur, summary_by_currency.get(cur), exp)
assert total_twd == Decimal('2674743.45'), total_twd
assert all_total_twd == Decimal('2674743.45'), all_total_twd

print('FINAL VALIDATION OK')
print('Verified: all pages export has 14 rows; Details has 14 Open Sep-2026 rows sorted by text OrderID; FX conversion uses fx_rates.csv; Summary has per-currency original amounts and all-currency TWD total.')