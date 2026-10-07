from pathlib import Path
import csv
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

ws = Path.cwd()
print('Workspace:', ws)
print('Input files:', sorted([p.name for p in (ws/'input').glob('*')]))
print('Output files:', sorted([str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()]))

out_xlsx = ws/'output'/'september_open_orders.xlsx'
csv_path = ws/'output'/'downloads'/'orders_export.csv'
fx_path = ws/'input'/'fx_rates.csv'
template_path = ws/'input'/'sales_template.xlsx'

for p in [template_path, fx_path, csv_path, out_xlsx]:
    if not p.exists():
        raise FileNotFoundError(p)
    print('Exists:', p.relative_to(ws), 'bytes=', p.stat().st_size)

with open(fx_path, newline='', encoding='utf-8-sig') as f:
    fx_rows = list(csv.DictReader(f))
fx_headers = list(fx_rows[0].keys()) if fx_rows else []
print('FX headers:', fx_headers)
print('FX sample:', fx_rows[:3])
rate_col = 'RateToTWD' if 'RateToTWD' in fx_headers else 'TWD_per_unit'
fx = {r['Currency']: Decimal(str(r[rate_col])) for r in fx_rows}
print('FX rates used:', fx)

with open(csv_path, newline='', encoding='utf-8-sig') as f:
    src_rows = list(csv.DictReader(f))
print('CSV headers:', list(src_rows[0].keys()) if src_rows else [])
print('CSV row count:', len(src_rows))
print('CSV first rows:', src_rows[:3])

wb = load_workbook(out_xlsx, data_only=False)
print('Workbook sheets:', wb.sheetnames)
for s in ['Report','Details','Summary','Instructions']:
    if s not in wb.sheetnames:
        raise AssertionError(f'Missing sheet {s}')

d = wb['Details']
headers = [d.cell(4,c).value for c in range(1, d.max_column+1)]
print('Details headers row 4:', headers)
required = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD']
for h in required:
    if h not in headers:
        raise AssertionError(f'Missing required Details column {h}')
col = {h:i+1 for i,h in enumerate(headers) if h}
rows=[]
for r in range(5, d.max_row+1):
    if d.cell(r, col['OrderID']).value in (None, ''):
        continue
    item={h:d.cell(r, col[h]).value for h in required}
    for h in ['CustomerContact','PaymentTerms']:
        if h in col:
            item[h]=d.cell(r, col[h]).value
    rows.append(item)
print('Details data row count:', len(rows))
print('Details first/last OrderID:', rows[0]['OrderID'] if rows else None, rows[-1]['OrderID'] if rows else None)
print('Details first rows:', rows[:3])

if len(src_rows) != 14:
    raise AssertionError(f'CSV rows expected 14, got {len(src_rows)}')
if len(rows) != 14:
    raise AssertionError(f'Details rows expected 14, got {len(rows)}')

ids = [str(r['OrderID']) for r in rows]
if ids != sorted(ids):
    raise AssertionError('OrderID not sorted ascending')
if not all(len(x)==10 and x.startswith('0') for x in ids):
    raise AssertionError('OrderID not preserved as text with leading zeroes')

csv_ids = [str(r['OrderID']) for r in src_rows]
if csv_ids != ids:
    raise AssertionError('Workbook Details OrderIDs do not match exported all-filtered CSV order')

for r in rows:
    if str(r['Status']) != 'Open':
        raise AssertionError('Non-Open status found')
    od = r['OrderDate']
    od_s = od.strftime('%Y-%m-%d') if hasattr(od, 'strftime') else str(od)[:10]
    if not ('2026-09-01' <= od_s <= '2026-09-30'):
        raise AssertionError(f'Date outside September 2026: {od_s}')
    cur = str(r['Currency'])
    expected_twd = (Decimal(str(r['NetAmount'])) * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    actual_twd = Decimal(str(r['NetAmountTWD'])).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    if actual_twd != expected_twd:
        raise AssertionError(f'FX mismatch {r["OrderID"]}: {actual_twd} vs {expected_twd}')

# Summary check
s = wb['Summary']
summary_rows = []
for r in range(1, s.max_row+1):
    vals = [s.cell(r,c).value for c in range(1, min(s.max_column, 5)+1)]
    if any(v is not None for v in vals):
        summary_rows.append((r, vals))
print('Summary non-empty rows:', summary_rows)
print('Summary!B2:', s['B2'].value)
if s['B2'].value != '=SUM(Details!I5:I18)':
    raise AssertionError(f'Unexpected Summary!B2 formula: {s["B2"].value}')

by_cur = {}
total_twd = Decimal('0.00')
for r in rows:
    cur = str(r['Currency'])
    amt = Decimal(str(r['NetAmount'])).quantize(Decimal('0.01'))
    twd = Decimal(str(r['NetAmountTWD'])).quantize(Decimal('0.01'))
    by_cur.setdefault(cur, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count':0})
    by_cur[cur]['amount'] += amt
    by_cur[cur]['twd'] += twd
    by_cur[cur]['count'] += 1
    total_twd += twd
print('Computed summary:', {k:{kk:str(vv) for kk,vv in val.items()} for k,val in by_cur.items()}, 'TotalTWD=', total_twd)

expected = {
    'EUR': (Decimal('28936.37'), Decimal('1012772.95'), 4),
    'TWD': (Decimal('637522.50'), Decimal('637522.50'), 6),
    'USD': (Decimal('32014.00'), Decimal('1024448.00'), 4),
}
for cur, (amt, twd, cnt) in expected.items():
    got = by_cur.get(cur)
    if got is None or got['amount'] != amt or got['twd'] != twd or got['count'] != cnt:
        raise AssertionError(f'Summary mismatch for {cur}: {got}')
if total_twd != Decimal('2674743.45'):
    raise AssertionError(f'Total TWD mismatch: {total_twd}')

# Validate Summary sheet contains the expected aggregate rows by looking for currency labels and total label.
summary_text = [[s.cell(r,c).value for c in range(1, s.max_column+1)] for r in range(1, s.max_row+1)]
flat = ' | '.join(str(v) for row in summary_text for v in row if v is not None)
for cur in expected:
    if cur not in flat:
        raise AssertionError(f'Summary sheet missing currency label {cur}')
if 'Total TWD' not in flat:
    raise AssertionError('Summary sheet missing Total TWD label')

print('FINAL VALIDATION OK: workflow output matches requested September 2026 Open Sales Orders report.')