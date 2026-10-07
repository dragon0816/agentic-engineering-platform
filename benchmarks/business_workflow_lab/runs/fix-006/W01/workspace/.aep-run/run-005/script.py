import os, csv, re, datetime, json
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
start_date = params['start_date']
end_date = params['end_date']
status = params['status']
output_filename = params.get('output_filename') or 'sales_orders_report.xlsx'
if Path(output_filename).name != output_filename or not output_filename.lower().endswith('.xlsx'):
    raise ValueError('output_filename must be a simple .xlsx file name')

ws = Path.cwd()
in_dir = ws / 'input'
out_dir = ws / 'output'
download = out_dir / 'downloads' / 'orders_export.csv'
template = in_dir / 'sales_template.xlsx'
fx_path = in_dir / 'fx_rates.csv'
out_path = out_dir / output_filename
for p in [template, fx_path, download]:
    if not p.exists():
        raise FileNotFoundError(str(p))
print('Inputs OK:', template.name, fx_path.name, download)

# Preserve contact/payment terms from an existing output workbook if present.
preserved = {}
if out_path.exists():
    old = load_workbook(out_path, data_only=False)
    if 'Details' in old.sheetnames:
        sh = old['Details']
        for i in range(5, sh.max_row + 1):
            oid = sh.cell(i, 1).value
            if oid:
                preserved[str(oid)] = {
                    'CustomerContact': sh.cell(i, 10).value or '',
                    'PaymentTerms': sh.cell(i, 11).value or '',
                }
print('Preserved detail-only fields for', len(preserved), 'orders')

rates = {}
with open(fx_path, encoding='utf-8-sig', newline='') as f:
    rdr = csv.reader(f)
    next(rdr)
    for row in rdr:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('FX rates:', rates)

rows = []
with open(download, encoding='utf-8-sig', newline='') as f:
    rdr = csv.DictReader(f)
    expected = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']
    if rdr.fieldnames != expected:
        raise ValueError(f'Unexpected export headers: {rdr.fieldnames}')
    rows = list(rdr)
print('Downloaded rows:', len(rows))
filtered = [r for r in rows if start_date <= r['OrderDate'] <= end_date and r['Status'] == status]
filtered.sort(key=lambda r: r['OrderID'])
print('Filtered rows:', len(filtered), 'first/last:', (filtered[0]['OrderID'], filtered[-1]['OrderID']) if filtered else None)

def q2(x):
    return Decimal(x).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

wb = load_workbook(template)
report = wb['Report']; details = wb['Details']; summary = wb['Summary']
if details.max_row > 4:
    details.delete_rows(5, details.max_row - 4)
if summary.max_row > 4:
    summary.delete_rows(5, summary.max_row - 4)

now = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = f'Sales Orders; Status={status}; Sort by OrderID asc'
report['B5'] = now
report['B7'] = 'Lab ERP synthetic SAP-like web system'
report['D7'] = now
report['B9'] = 'Export CSV (all filtered rows) used; CustomerContact/PaymentTerms preserved from existing output when available, otherwise blank. fx_rates.csv fixed test rates used.'

for idx, r in enumerate(filtered, start=5):
    cur = r['Currency']
    if cur not in rates:
        raise ValueError(f'Missing FX rate for {cur}')
    amt = Decimal(r['NetAmount'])
    twd = q2(amt * rates[cur])
    extra = preserved.get(r['OrderID'], {'CustomerContact': '', 'PaymentTerms': ''})
    vals = [r['OrderID'], r['CustomerID'], r['CustomerName'], r.get('ProjectID',''), r['OrderDate'], r['Status'], cur, float(amt), float(twd), extra['CustomerContact'], extra['PaymentTerms']]
    for col, val in enumerate(vals, start=1):
        c = details.cell(idx, col)
        c.value = val
        src = details.cell(4, col)
        c.font = copy(src.font); c.fill = copy(src.fill); c.border = copy(src.border); c.alignment = copy(src.alignment)
        if col in (1,2,4): c.number_format = '@'
        elif col == 5: c.number_format = 'yyyy-mm-dd'
        elif col in (8,9): c.number_format = '0.00'

last_row = 4 + len(filtered) if filtered else 5
summary['B2'] = f'=SUM(Details!I5:I{last_row if filtered else 5})'
by_cur = {}
for r in filtered:
    cur = r['Currency']; amt = Decimal(r['NetAmount']); twd = q2(amt * rates[cur])
    rec = by_cur.setdefault(cur, {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0})
    rec['amount'] += amt; rec['twd'] += twd; rec['count'] += 1
sum_rows = []
for cur in sorted(by_cur):
    rec = by_cur[cur]
    sum_rows.append(['Currency', cur, cur, float(q2(rec['amount'])), float(q2(rec['twd'])), rec['count']])
total_twd = sum((rec['twd'] for rec in by_cur.values()), Decimal('0.00'))
total_count = sum((rec['count'] for rec in by_cur.values()), 0)
sum_rows.append(['Total', 'All currencies', 'TWD', None, float(q2(total_twd)), total_count])
for idx, vals in enumerate(sum_rows, start=5):
    for col, val in enumerate(vals, start=1):
        c = summary.cell(idx, col)
        c.value = val
        src = summary.cell(4, col)
        c.font = copy(src.font); c.fill = copy(src.fill); c.border = copy(src.border); c.alignment = copy(src.alignment)
        if col in (4,5): c.number_format = '0.00'
        if col == 6: c.number_format = '0'

details.auto_filter.ref = f'A4:K{last_row}'
summary.auto_filter.ref = f'A4:F{4+len(sum_rows)}'
out_dir.mkdir(parents=True, exist_ok=True)
wb.save(out_path)
print('Saved:', out_path)

wb2 = load_workbook(out_path, data_only=False)
rd = wb2['Details']; sm = wb2['Summary']
headers = [rd.cell(4,c).value for c in range(1,12)]
assert headers == ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
out_rows = [[rd.cell(i,c).value for c in range(1,12)] for i in range(5, 5+len(filtered))]
assert len(out_rows) == len(filtered)
assert [r[0] for r in out_rows] == sorted(r[0] for r in out_rows)
assert all(isinstance(r[0], str) and re.match(r'^\d{10}$', r[0]) for r in out_rows)
assert all(r[5] == status for r in out_rows)
assert all(start_date <= r[4] <= end_date for r in out_rows)
for row in out_rows:
    assert q2(Decimal(str(row[8]))) == q2(Decimal(str(row[7])) * rates[row[6]])
print('Validation OK:', output_filename, len(out_rows), 'rows; Summary B2:', sm['B2'].value)
