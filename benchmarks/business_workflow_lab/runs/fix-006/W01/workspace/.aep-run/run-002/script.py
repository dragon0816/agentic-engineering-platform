from pathlib import Path
import csv, os
p=Path('output/downloads/orders_export.csv')
print('download exists', p.exists(), p.stat().st_size if p.exists() else None)
if p.exists():
    with open(p, encoding='utf-8-sig', newline='') as f:
        rdr=csv.reader(f)
        for i,row in zip(range(20), rdr): print(i,row)
# print rest of instructions
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
sh=wb['Instructions']
for r in sh.iter_rows(min_row=1, max_row=sh.max_row, values_only=True): print(r)
