from pathlib import Path
from openpyxl import load_workbook
import csv, os, json

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')
print('OUTPUT FILES')
for p in sorted((ws/'output').glob('**/*')):
    if p.is_file():
        print(' -', p.relative_to(ws), p.stat().st_size, 'bytes')

out = ws/'output'/'september_open_orders.xlsx'
print('TARGET_EXISTS', out.exists())
if not out.exists():
    raise SystemExit(0)

wb = load_workbook(out, data_only=False)
print('SHEETS', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('SHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    sample = []
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        sample.append(row)
    print('SAMPLE', repr(sample))

# find details header row
sh = wb['Details'] if 'Details' in wb.sheetnames else wb[wb.sheetnames[0]]
header_row = None
headers = None
for r in range(1, min(sh.max_row, 20)+1):
    vals = [sh.cell(r,c).value for c in range(1, sh.max_column+1)]
    if 'OrderID' in vals:
        header_row = r
        headers = vals
        break
print('DETAIL_HEADER_ROW', header_row)
print('DETAIL_HEADERS', headers)

if header_row:
    idx = {h:i+1 for i,h in enumerate(headers) if h is not None}
    rows=[]
    for r in range(header_row+1, sh.max_row+1):
        if any(sh.cell(r,c).value is not None for c in range(1, sh.max_column+1)):
            rows.append(r)
    print('DETAIL_DATA_ROWS', len(rows), rows[:3], rows[-3:])
    order_ids = [str(sh.cell(r, idx['OrderID']).value) for r in rows if 'OrderID' in idx]
    print('ORDER_IDS', order_ids)
    print('ORDER_IDS_SORTED', order_ids == sorted(order_ids))
    if 'Status' in idx:
        print('STATUSES', sorted(set(sh.cell(r, idx['Status']).value for r in rows)))
    if 'OrderDate' in idx:
        print('DATE_RANGE', min(sh.cell(r, idx['OrderDate']).value for r in rows), max(sh.cell(r, idx['OrderDate']).value for r in rows))
    if 'Currency' in idx and 'NetAmount' in idx and 'NetAmountTWD' in idx:
        by_cur = {}
        total_twd = 0
        for r in rows:
            cur = sh.cell(r, idx['Currency']).value
            amt = sh.cell(r, idx['NetAmount']).value
            twd = sh.cell(r, idx['NetAmountTWD']).value
            by_cur.setdefault(cur, [0,0,0])
            by_cur[cur][0] += float(amt)
            by_cur[cur][1] += float(twd)
            by_cur[cur][2] += 1
            total_twd += float(twd)
        print('BY_CURRENCY', json.dumps(by_cur, ensure_ascii=False, sort_keys=True))
        print('TOTAL_TWD', round(total_twd, 2))

if 'Summary' in wb.sheetnames:
    ss = wb['Summary']
    print('SUMMARY_B2_FORMULA_OR_VALUE', ss['B2'].value)
    summary_rows=[]
    for row in ss.iter_rows(min_row=1, max_row=ss.max_row, values_only=True):
        summary_rows.append(row)
    print('SUMMARY_ROWS', repr(summary_rows))
