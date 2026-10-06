from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size, 'bytes')

# Inspect FX rates
fx = ws/'input'/'fx_rates.csv'
print('\nFX RATES PREVIEW')
with fx.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i, row in zip(range(10), reader):
        print(i, row)

# Inspect workbook
path = ws/'input'/'quote_template.xlsx'
wb = load_workbook(path)
print('\nWORKBOOK SHEETS', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'\nSHEET {s}: max_row={sh.max_row}, max_col={sh.max_column}')
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        print(r)

# Inspect manual
manual = ws/'input'/'quote_cli_manual.md'
print('\nMANUAL PREVIEW')
text = manual.read_text(encoding='utf-8', errors='replace')
print(text[:4000])