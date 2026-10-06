from pathlib import Path
from openpyxl import load_workbook
import subprocess, json, sys
ws = Path.cwd()
print('FULL INSTRUCTIONS')
wb = load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)

cli = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
print('\nCLI EXISTS', cli.exists(), cli)
cmd = [str(cli), 'quotes', 'list', '--project', 'P001', '--from', '2026-09-01', '--to', '2026-09-30', '--status', 'Accepted', '--format', 'json']
print('RUN SAMPLE', cmd)
p = subprocess.run(cmd, capture_output=True, text=True, encoding='utf-8', errors='replace')
print('exit', p.returncode)
print('stdout', p.stdout[:2000])
print('stderr', p.stderr[:2000])
if p.returncode == 0 and p.stdout.strip():
    data=json.loads(p.stdout)
    print('count', data.get('count'))
    if data.get('items'):
        qid=data['items'][0]['QuoteID']
        p2 = subprocess.run([str(cli),'quotes','get',qid,'--format','json'], capture_output=True, text=True, encoding='utf-8', errors='replace')
        print('GET', qid, 'exit', p2.returncode)
        print(p2.stdout[:2000])
        print(p2.stderr[:1000])