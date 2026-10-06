from pathlib import Path
import json
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

ws=Path.cwd()
out=ws/'output'/'sent_quotes_2026-09.xlsx'
logf=ws/'output'/'quote_cli_attempts.json'
print('Checking existing outputs only; no CLI call.')
print('Workbook exists:', out.exists(), out)
print('Log exists:', logf.exists(), logf)
wb=load_workbook(out, data_only=False)
qws=wb['Quotes']; lws=wb['Lines']; sws=wb['Summary']; rws=wb['Report']
expected_quote_headers=['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
expected_line_headers=['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
expected_sum_headers=['Group','SubGroup','Currency','Amount','AmountTWD','Count']
assert [qws.cell(4,c).value for c in range(1,12)] == expected_quote_headers
assert [lws.cell(4,c).value for c in range(1,7)] == expected_line_headers
assert [sws.cell(4,c).value for c in range(1,7)] == expected_sum_headers
qrows=[]
for row in qws.iter_rows(min_row=5, max_row=qws.max_row, max_col=11, values_only=True):
    if any(v is not None for v in row): qrows.append(row)
lrows=[]
for row in lws.iter_rows(min_row=5, max_row=lws.max_row, max_col=6, values_only=True):
    if any(v is not None for v in row): lrows.append(row)
srows=[]
for row in sws.iter_rows(min_row=5, max_row=sws.max_row, max_col=6, values_only=True):
    if any(v is not None for v in row): srows.append(row)
log=json.load(logf.open(encoding='utf-8'))
assert len(qrows)==log['quote_count']==4, (len(qrows), log.get('quote_count'))
assert len(lrows)==log['line_count']==8, (len(lrows), log.get('line_count'))
assert len(srows)==3, len(srows)
assert rws['B3'].value=='2026-09-01' and rws['D3'].value=='2026-09-30'
assert 'Sent' in rws['B4'].value
assert sws['B2'].value=='=SUM(Quotes!K5:K8)', sws['B2'].value
# Type and value checks, allowing blanks for missing source values.
for i,row in enumerate(qrows, start=5):
    assert isinstance(row[0], str) and row[0], (i,row)
    assert row[1] is None or isinstance(row[1], str), (i,row)
    assert row[2] is None or isinstance(row[2], str), (i,row)
    assert '2026-09-01' <= row[3] <= '2026-09-30', (i,row)
    assert row[5]=='Sent', (i,row)
    assert row[6] in ('TWD','USD','EUR'), (i,row)
    for v in (row[7], row[8], row[9], row[10]):
        assert isinstance(v, (int,float)), (i,row,type(v))
for i,row in enumerate(lrows, start=5):
    assert isinstance(row[0], str) and row[0], (i,row)
    assert isinstance(row[1], str) and row[1], (i,row)
    assert row[2] is None or isinstance(row[2], str), (i,row)
    assert isinstance(row[3], (int,float)), (i,row)
    assert isinstance(row[4], (int,float)), (i,row)
    assert isinstance(row[5], (int,float)), (i,row)
# Summary totals reconcile to quote data by status/currency.
from collections import defaultdict
agg=defaultdict(lambda: [0.0,0])
for row in qrows:
    agg[(row[5], row[6])][0]+=float(row[10])
    agg[(row[5], row[6])][1]+=1
for row in srows:
    key=(row[0], row[2])
    assert key in agg, row
    assert abs(float(row[4])-agg[key][0]) < 0.005, (row, agg[key])
    assert row[5]==agg[key][1], (row, agg[key])
# Attempt log checks: exit 4 attempts are not counted as data; exit 2/3 would have aborted.
assert log['query']=={'from':'2026-09-01','to':'2026-09-30','status':'Sent'}
attempts=log['attempts']
assert attempts, 'missing attempts'
# For each label, if code 4 occurred, attempts for that label do not exceed 3 total.
by_label=defaultdict(list)
for a in attempts:
    by_label[a['label']].append(a['returncode'])
    assert a['returncode'] not in (2,3), f'unexpected non-retry error in log: {a}'
for label,codes in by_label.items():
    if 4 in codes:
        assert len(codes) <= 3, (label,codes)
        assert codes[-1] == 0, (label,codes)
print('CHECKED OK')
print('Quote rows:', len(qrows))
print('Line rows:', len(lrows))
print('Summary rows:', len(srows))
print('Summary B2:', sws['B2'].value)
print('CLI attempts by label:', dict(by_label))
print('Quote IDs:', [r[0] for r in qrows])
print('Outputs:', out, logf)
