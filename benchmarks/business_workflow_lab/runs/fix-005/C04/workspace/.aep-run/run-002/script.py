from pathlib import Path
from openpyxl import load_workbook
import subprocess, json, sys, os

ws=Path.cwd()
wb=load_workbook(ws/'input'/'quote_template.xlsx', data_only=False)
print('FULL INSTRUCTIONS:')
for row in wb['Instructions'].iter_rows(values_only=True):
    print(row)

cli = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
print('\nCLI exists:', cli.exists(), cli)
cmd=[str(cli),'quotes','list','--from','2026-09-01','--to','2026-09-30','--status','Sent','--format','json']
print('RUN sample list:', cmd)
p = subprocess.run(cmd, capture_output=True, text=True, shell=False)
print('returncode=', p.returncode)
print('stderr=', p.stderr[:1000])
print('stdout head=', p.stdout[:2000])
if p.returncode==0:
    try:
        data=json.loads(p.stdout)
        print('count', data.get('count'), 'first ids', [it.get('QuoteID') for it in data.get('items',[])[:10]])
    except Exception as e:
        print('parse error', repr(e))
