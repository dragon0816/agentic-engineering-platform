from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

manual = ws/'input'/'quote_cli_manual.md'
print('\nMANUAL HEAD:')
print(manual.read_text(encoding='utf-8', errors='replace')[:4000])

fx = ws/'input'/'fx_rates.csv'
print('\nFX CSV SAMPLE:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i,row in zip(range(8), r):
        print(i, row)

xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('\nWORKBOOK SHEETS:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print(f'\nSHEET {s}: max_row={sh.max_row} max_col={sh.max_column}')
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=False):
        vals=[]
        for c in row[:min(sh.max_column, 12)]:
            vals.append(c.value)
        print(vals)
