from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/sales_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    ws=wb[sh]
    print('\nSHEET', sh)
    for r in ws.iter_rows(values_only=True):
        print(r)
