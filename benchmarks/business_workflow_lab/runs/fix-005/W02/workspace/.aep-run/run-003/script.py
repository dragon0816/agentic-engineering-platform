from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from copy import copy
import csv, json
from openpyxl import load_workbook
from openpyxl.styles import Alignment

workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(exist_ok=True)

source_csv = output_dir / 'downloads' / 'orders_C001_2026-08-01_2026-10-31.csv'
template = input_dir / 'sales_template.xlsx'
fx_csv = input_dir / 'fx_rates.csv'
out_xlsx = output_dir / 'sales_report_C001_20260801_20261031.xlsx'
out_summary = output_dir / 'sales_report_C001_20260801_20261031_summary.txt'

# Details obtained by opening each order detail page through browser__operate.
detail_lookup = {
    '0045000101': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET15'},
    '0045000107': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000109': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10'},
    '0045000114': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET5'},
    '0045000119': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000125': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET0'},
    '0045000129': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30'},
    '0045000132': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': ''},
    '0045000136': {'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET5'},
}

# Load CSV exported from the ERP after filtering exact Customer ID C001 and inclusive date window.
with open(source_csv, newline='', encoding='utf-8-sig') as f:
    orders = list(csv.DictReader(f))
print('exported orders:', len(orders))
print('export headers:', orders[0].keys() if orders else [])
print('export sample:', orders[:2])

# Load fixed exchange rates; CSV notes contain commas, so use first two columns only.
fx = {}
with open(fx_csv, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx[row[0]] = Decimal(row[1])
print('fx rates:', fx)

# Validate source before producing report.
expected_customer = 'C001'
start_date = '2026-08-01'
end_date = '2026-10-31'
for r in orders:
    assert r['CustomerID'] == expected_customer, f"unexpected customer {r['CustomerID']} in {r['OrderID']}"
    assert start_date <= r['OrderDate'] <= end_date, f"date outside range {r['OrderID']} {r['OrderDate']}"
    assert r['OrderID'] in detail_lookup, f"missing detail lookup for {r['OrderID']}"
assert len(orders) == 9, f"expected 9 filtered rows, got {len(orders)}"

# Keep OrderID text and sort by export order (OrderID asc).
orders.sort(key=lambda r: r['OrderID'])

TWOPLACES = Decimal('0.01')
def money(value):
    return Decimal(str(value)).quantize(TWOPLACES, rounding=ROUND_HALF_UP)

def copy_row_style(ws, src_row, dst_row, min_col, max_col):
    for col in range(min_col, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)

wb = load_workbook(template)
report = wb['Report']
details = wb['Details']
summary = wb['Summary']

# Report header cells per template instructions.
report['B3'] = start_date
report['D3'] = end_date
report['B4'] = 'Sales Orders; Customer ID = C001 (exact match); OrderDate inclusive 2026-08-01 to 2026-10-31; all statuses'
report['B5'] = '2026-10-06T00:00:00+08:00'
report['B7'] = 'Lab ERP Sales Orders export + Sales Order detail pages; fx_rates.csv'
report['D7'] = '2026-10-06T00:00:00+08:00'

# Clear any old data rows without touching header/footer/sheet names.
for ws in (details, summary):
    if ws.max_row > 4:
        ws.delete_rows(5, ws.max_row - 4)

# Fill Details sheet.
headers = [details.cell(4, c).value for c in range(1, details.max_column + 1)]
missing = []
for idx, r in enumerate(orders, start=5):
    copy_row_style(details, 4, idx, 1, 11)
    oid = r['OrderID']
    net = money(r['NetAmount'])
    rate = fx[r['Currency']]
    net_twd = money(net * rate)
    det = detail_lookup[oid]
    row_values = {
        'OrderID': oid,
        'CustomerID': r['CustomerID'],
        'CustomerName': r['CustomerName'],
        'ProjectID': r.get('ProjectID') or '',
        'OrderDate': r['OrderDate'],
        'Status': r['Status'],
        'Currency': r['Currency'],
        'NetAmount': float(net),
        'NetAmountTWD': float(net_twd),
        'CustomerContact': det.get('CustomerContact') or '',
        'PaymentTerms': det.get('PaymentTerms') or '',
    }
    for col, h in enumerate(headers, start=1):
        cell = details.cell(idx, col)
        cell.value = row_values[h]
        if h in {'OrderID', 'CustomerID', 'ProjectID', 'CustomerContact', 'PaymentTerms', 'CustomerName', 'Status', 'Currency'}:
            cell.number_format = '@'
        if h in {'NetAmount', 'NetAmountTWD'}:
            cell.number_format = '0.00'
        if h == 'OrderDate':
            cell.number_format = 'yyyy-mm-dd'
    for field in ['CustomerContact', 'PaymentTerms']:
        if not row_values[field]:
            missing.append((oid, field))

last_row = 4 + len(orders)
# Set autofilter if present/useful
if len(orders):
    details.auto_filter.ref = f'A4:K{last_row}'

# Fill Summary sheet grouped by Status + Currency.
summary_rows = {}
for r in orders:
    det = detail_lookup[r['OrderID']]
    net = money(r['NetAmount'])
    net_twd = money(net * fx[r['Currency']])
    key = (r['Status'], r['Currency'])
    if key not in summary_rows:
        summary_rows[key] = {'Amount': Decimal('0'), 'AmountTWD': Decimal('0'), 'Count': 0}
    summary_rows[key]['Amount'] += net
    summary_rows[key]['AmountTWD'] += net_twd
    summary_rows[key]['Count'] += 1

# Stable, readable ordering.
status_order = {'Open': 1, 'Closed': 2, 'Cancelled': 3}
for row_idx, ((status, curr), vals) in enumerate(sorted(summary_rows.items(), key=lambda kv: (status_order.get(kv[0][0], 99), kv[0][1])), start=5):
    copy_row_style(summary, 4, row_idx, 1, 6)
    summary.cell(row_idx, 1).value = 'Status'
    summary.cell(row_idx, 2).value = status
    summary.cell(row_idx, 3).value = curr
    summary.cell(row_idx, 4).value = float(vals['Amount'].quantize(TWOPLACES, rounding=ROUND_HALF_UP))
    summary.cell(row_idx, 5).value = float(vals['AmountTWD'].quantize(TWOPLACES, rounding=ROUND_HALF_UP))
    summary.cell(row_idx, 6).value = vals['Count']
    summary.cell(row_idx, 4).number_format = '0.00'
    summary.cell(row_idx, 5).number_format = '0.00'

summary_last = 4 + len(summary_rows)
summary['B2'] = f'=SUM(Details!I5:I{last_row})' if orders else '=SUM(Details!I5:I5)'
if summary_rows:
    summary.auto_filter.ref = f'A4:F{summary_last}'

missing_text = '缺少資料：' + ('；'.join(f'{oid} {field}' for oid, field in missing) if missing else '無')
report['B9'] = f'ERP 篩選結果共 {len(orders)} 筆；僅納入 CustomerID 精確等於 C001，未納入名稱相似客戶（例如 C007）。{missing_text}。'
report['B9'].alignment = Alignment(wrap_text=True)

# Save report.
wb.save(out_xlsx)

# Reopen and verify from saved output.
check = load_workbook(out_xlsx, data_only=False)
cd = check['Details']
cs = check['Summary']
cr = check['Report']
errors = []
if cr['B3'].value != start_date or cr['D3'].value != end_date:
    errors.append('Report date cells incorrect')
if cd.max_row != 13:
    errors.append(f'Details row count wrong: max_row={cd.max_row}')
saved_headers = [cd.cell(4, c).value for c in range(1, 12)]
if saved_headers != headers:
    errors.append(f'Details headers changed: {saved_headers}')
ids = [cd.cell(r, 1).value for r in range(5, cd.max_row + 1)]
if ids != [r['OrderID'] for r in orders]:
    errors.append(f'Order IDs mismatch: {ids}')
if any(cd.cell(r, 2).value != 'C001' for r in range(5, cd.max_row + 1)):
    errors.append('Non-C001 row present')
if any(not str(cd.cell(r, 1).value).startswith('0045') for r in range(5, cd.max_row + 1)):
    errors.append('OrderID not preserved as text-like leading zero value')
if any(cd.cell(r, 8).data_type not in ('n',) or cd.cell(r, 9).data_type not in ('n',) for r in range(5, cd.max_row + 1)):
    errors.append('Amount columns are not numeric')
# Missing data check: only order 0045000132 PaymentTerms is blank.
blank_contacts = [cd.cell(r, 1).value for r in range(5, cd.max_row + 1) if cd.cell(r, 10).value in (None, '')]
blank_terms = [cd.cell(r, 1).value for r in range(5, cd.max_row + 1) if cd.cell(r, 11).value in (None, '')]
if blank_contacts:
    errors.append(f'Unexpected blank CustomerContact: {blank_contacts}')
if blank_terms != ['0045000132']:
    errors.append(f'PaymentTerms blanks mismatch: {blank_terms}')
if cs['B2'].value != '=SUM(Details!I5:I13)':
    errors.append(f'Summary B2 formula incorrect: {cs["B2"].value}')
# Validate computed TWD total from saved values.
total_twd = sum(Decimal(str(cd.cell(r, 9).value)) for r in range(5, cd.max_row + 1)).quantize(TWOPLACES)
if total_twd != Decimal('1702471.00'):
    errors.append(f'TWD total mismatch: {total_twd}')

summary_lines = [
    'Sales order report for CustomerID C001, 2026-08-01 to 2026-10-31 (inclusive).',
    f'Filtered ERP export rows: {len(orders)}; all rows verified CustomerID=C001. Similar customer names were not included.',
    f'Total NetAmountTWD: {total_twd}',
    missing_text + '。',
    'Output workbook: ' + str(out_xlsx),
]
out_summary.write_text('\n'.join(summary_lines), encoding='utf-8')

print('wrote:', out_xlsx)
print('wrote:', out_summary)
print('detail rows:', len(orders), 'summary rows:', len(summary_rows), 'total_twd:', total_twd)
print('missing:', missing)
print('checks:', 'PASS' if not errors else 'FAIL')
if errors:
    for e in errors:
        print('ERROR:', e)
    raise SystemExit(1)
