from pathlib import Path
import csv, json
from openpyxl import load_workbook

ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size, 'bytes')

# Inspect workbook sheets/dimensions/headers/sample rows
xlsx = ws/'input'/'sales_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for sh in wb.sheetnames:
    sheet = wb[sh]
    print(f'[{sh}] max_row={sheet.max_row} max_col={sheet.max_column}')
    for r in range(1, min(sheet.max_row, 5)+1):
        vals = [sheet.cell(r,c).value for c in range(1, sheet.max_column+1)]
        print(' row', r, vals)

# CSV headers/sample
csvp = ws/'input'/'fx_rates.csv'
if csvp.exists():
    with open(csvp, newline='', encoding='utf-8-sig') as f:
        rows = list(csv.reader(f))
    print('fx_rates rows:', len(rows))
    for row in rows[:5]: print(' fx', row)

# manual preview
md = ws/'input'/'sap_manual.md'
if md.exists():
    text = md.read_text(encoding='utf-8')
    print('sap_manual first 2000 chars:\n', text[:2000])