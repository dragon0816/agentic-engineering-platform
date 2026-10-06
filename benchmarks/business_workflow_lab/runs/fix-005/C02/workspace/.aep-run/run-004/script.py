from pathlib import Path
import subprocess, json, sys, csv, shutil
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
from copy import copy
from openpyxl import load_workbook

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
OUT_XLSX = OUTPUT / 'quote_template.xlsx'
OUT_MD = OUTPUT / 'findings.md'

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
BASE_DATE = date.fromisoformat('2026-09-30')
GENERATED_AT = datetime.now().astimezone().isoformat(timespec='seconds')

TWOPLACES = Decimal('0.01')
def D(v):
    if v is None or v == '':
        return None
    return Decimal(str(v))
def money(v):
    return D(v).quantize(TWOPLACES, rounding=ROUND_HALF_UP)
def iso_date(s):
    if not s:
        return None
    return date.fromisoformat(s)
def run_cli(args):
    cmd = [sys.executable, str(CLI)] + args
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        raise RuntimeError(f'CLI failed rc={res.returncode}: {cmd}\nstderr={res.stderr}\nstdout={res.stdout}')
    return json.loads(res.stdout)

# Read fx rates; note text may contain an unquoted comma, so only first two columns are used.
fx_rates = {}
with open(FX, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            fx_rates[row[0]] = D(row[1])
print('fx_rates', fx_rates)

# Fetch all September 2026 quotes (inclusive) and all line-item details.
listed = run_cli(['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--format', 'json'])
if listed.get('schema_version') != 1:
    raise RuntimeError(f'Unexpected list schema_version: {listed.get("schema_version")}')
headers = listed.get('items', [])
print('quote_count_from_list', listed.get('count'), 'items_len', len(headers))
all_docs = []
for h in headers:
    qid = h['QuoteID']
    doc = run_cli(['quotes', 'get', qid, '--format', 'json'])
    if doc.get('schema_version') != 1:
        raise RuntimeError(f'Unexpected get schema_version for {qid}: {doc.get("schema_version")}')
    all_docs.append(doc)
print('fetched_details', len(all_docs))
print('quote_ids', [d['header']['QuoteID'] for d in all_docs])

# Validate and compute findings.
quote_rows = []
line_rows = []
checks = []
expired_unaccepted = []
for doc in all_docs:
    h = doc['header']
    qid = h['QuoteID']
    currency = h.get('Currency')
    if currency not in fx_rates:
        raise RuntimeError(f'Missing FX rate for currency {currency} quote {qid}')
    lines = doc.get('line_items', [])
    line_sum = Decimal('0.00')
    line_issues = []
    for li in lines:
        qty = D(li.get('Quantity'))
        unit = D(li.get('UnitPrice'))
        source_line_amt = D(li.get('LineAmount'))
        expected_line_amt = money(qty * unit)
        if source_line_amt != expected_line_amt:
            line_issues.append(f"line {li.get('ProductID')} amount {source_line_amt} != Quantity*UnitPrice {expected_line_amt}")
        line_sum += source_line_amt
        line_rows.append([
            qid,
            li.get('ProductID') or None,
            li.get('ProductName') or None,
            int(qty) if qty == qty.to_integral_value() else float(qty),
            float(money(unit)),
            float(money(source_line_amt)),
        ])
    line_sum = money(line_sum)
    subtotal = money(h.get('Subtotal'))
    discount = D(h.get('DiscountPercent'))
    net_amount = money(h.get('NetAmount'))
    expected_net = money(subtotal * (Decimal('1') - (discount / Decimal('100'))))
    subtotal_ok = (line_sum == subtotal)
    net_ok = (expected_net == net_amount)
    valid_until = iso_date(h.get('ValidUntil'))
    status = h.get('Status') or ''
    is_expired_unaccepted = (valid_until is not None and valid_until < BASE_DATE and status != 'Accepted')
    if is_expired_unaccepted:
        expired_unaccepted.append({
            'QuoteID': qid,
            'CustomerID': h.get('CustomerID') or '',
            'ProjectID': h.get('ProjectID') or '',
            'ValidUntil': h.get('ValidUntil') or '',
            'Status': status,
            'NetAmount': net_amount,
            'Currency': currency,
        })
    checks.append({
        'QuoteID': qid,
        'LineCount': len(lines),
        'LineSum': line_sum,
        'Subtotal': subtotal,
        'SubtotalOK': subtotal_ok,
        'ExpectedNetAmount': expected_net,
        'NetAmount': net_amount,
        'NetOK': net_ok,
        'LineIssues': line_issues,
    })
    net_twd = money(net_amount * fx_rates[currency])
    quote_rows.append([
        qid,
        h.get('CustomerID') or None,
        h.get('ProjectID') or None,
        iso_date(h.get('QuoteDate')),
        valid_until,
        status or None,
        currency,
        float(subtotal),
        float(discount),
        float(net_amount),
        float(net_twd),
    ])

# Fill workbook.
shutil.copyfile(TEMPLATE, OUT_XLSX)
wb = load_workbook(OUT_XLSX)

# Report sheet
report = wb['Report']
report['B3'] = FROM_DATE
report['D3'] = TO_DATE
report['B4'] = '2026 年 9 月全部報價；檢查明細合計、折扣後 NetAmount；標示 ValidUntil < 2026-09-30 且 Status != Accepted。'
report['B5'] = GENERATED_AT
report['B7'] = 'quote_cli_mock; fx_rates.csv'
report['D7'] = GENERATED_AT
report['B9'] = f'共取得 {len(quote_rows)} 份報價、{len(line_rows)} 筆明細；保留全部報價，不因到期刪除。'

# Helper to clear and write detail data.
def clear_data_rows(ws, min_row=5):
    if ws.max_row >= min_row:
        ws.delete_rows(min_row, ws.max_row - min_row + 1)

def clone_row_style(ws, src_row, dst_row, max_col):
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        dst = ws.cell(dst_row, col)
        if src.has_style:
            dst._style = copy(src._style)
        if src.number_format:
            dst.number_format = src.number_format
        if src.alignment:
            dst.alignment = copy(src.alignment)
        if src.font:
            dst.font = copy(src.font)
        if src.fill:
            dst.fill = copy(src.fill)
        if src.border:
            dst.border = copy(src.border)

quotes_ws = wb['Quotes']
clear_data_rows(quotes_ws)
for r_idx, row in enumerate(quote_rows, start=5):
    clone_row_style(quotes_ws, 4, r_idx, 11)
    for c_idx, val in enumerate(row, start=1):
        cell = quotes_ws.cell(r_idx, c_idx, val)
        if c_idx in [1,2,3,6,7]:
            cell.number_format = '@'
        elif c_idx in [4,5]:
            cell.number_format = 'yyyy-mm-dd'
        elif c_idx in [8,10,11]:
            cell.number_format = '0.00'
        elif c_idx == 9:
            cell.number_format = '0.##'
quotes_ws.auto_filter.ref = f'A4:K{max(4, 4 + len(quote_rows))}'

lines_ws = wb['Lines']
clear_data_rows(lines_ws)
for r_idx, row in enumerate(line_rows, start=5):
    clone_row_style(lines_ws, 4, r_idx, 6)
    for c_idx, val in enumerate(row, start=1):
        cell = lines_ws.cell(r_idx, c_idx, val)
        if c_idx in [1,2,3]:
            cell.number_format = '@'
        elif c_idx == 4:
            cell.number_format = '0.##'
        elif c_idx in [5,6]:
            cell.number_format = '0.00'
lines_ws.auto_filter.ref = f'A4:F{max(4, 4 + len(line_rows))}'

# Summary by Status and Currency, preserving columns and formula requirement.
summary = wb['Summary']
clear_data_rows(summary)
last_quote_row = 4 + len(quote_rows)
summary['B2'] = f'=SUM(Quotes!K5:K{last_quote_row if quote_rows else 5})'
summary['B2'].number_format = '0.00'
# Group/SubGroup: Status/Currency.
groups = {}
for row in quote_rows:
    status = row[5] or ''
    currency = row[6] or ''
    amount = Decimal(str(row[9])).quantize(TWOPLACES)
    twd = Decimal(str(row[10])).quantize(TWOPLACES)
    key = (status, currency)
    if key not in groups:
        groups[key] = {'amount': Decimal('0.00'), 'twd': Decimal('0.00'), 'count': 0}
    groups[key]['amount'] += amount
    groups[key]['twd'] += twd
    groups[key]['count'] += 1
for r_idx, ((status, currency), vals) in enumerate(sorted(groups.items()), start=5):
    clone_row_style(summary, 4, r_idx, 6)
    row = ['Status', status, currency, float(money(vals['amount'])), float(money(vals['twd'])), vals['count']]
    for c_idx, val in enumerate(row, start=1):
        cell = summary.cell(r_idx, c_idx, val)
        if c_idx in [1,2,3]:
            cell.number_format = '@'
        elif c_idx in [4,5]:
            cell.number_format = '0.00'
        elif c_idx == 6:
            cell.number_format = '0'
summary.auto_filter.ref = f'A4:F{max(4, 4 + len(groups))}'

wb.save(OUT_XLSX)

# findings.md
mismatched_subtotals = [c for c in checks if not c['SubtotalOK']]
mismatched_net = [c for c in checks if not c['NetOK']]
line_amount_issues = [c for c in checks if c['LineIssues']]
all_ok = not mismatched_subtotals and not mismatched_net and not line_amount_issues

lines = []
lines.append('# Findings - 2026-09 Quotes')
lines.append('')
lines.append(f'- 查詢期間：{FROM_DATE} 至 {TO_DATE}（含起訖日）')
lines.append(f'- 到期判斷基準日：{BASE_DATE.isoformat()}')
lines.append(f'- 資料來源：quote_cli_mock；匯率：input/fx_rates.csv')
lines.append(f'- 取得報價數：{len(quote_rows)}')
lines.append(f'- 取得明細列數：{len(line_rows)}')
lines.append('- 保留原則：全部報價皆保留於 Excel，不因到期或狀態刪除。')
lines.append('')
lines.append('## 金額一致性檢查')
if all_ok:
    lines.append('- 全部報價的明細 LineAmount 合計均等於表頭 Subtotal。')
    lines.append('- 全部報價的折扣後 NetAmount 均符合 `Subtotal × (1 − DiscountPercent / 100)` 並 ROUND_HALF_UP 至兩位小數。')
    lines.append('- 全部明細列的 LineAmount 均符合 `Quantity × UnitPrice`。')
else:
    if mismatched_subtotals:
        lines.append('### 明細合計與 Subtotal 不一致')
        lines.append('| QuoteID | Line Sum | Header Subtotal |')
        lines.append('|---|---:|---:|')
        for c in mismatched_subtotals:
            lines.append(f"| {c['QuoteID']} | {c['LineSum']} | {c['Subtotal']} |")
    if mismatched_net:
        lines.append('### 折扣後 NetAmount 不一致')
        lines.append('| QuoteID | Expected NetAmount | Header NetAmount |')
        lines.append('|---|---:|---:|')
        for c in mismatched_net:
            lines.append(f"| {c['QuoteID']} | {c['ExpectedNetAmount']} | {c['NetAmount']} |")
    if line_amount_issues:
        lines.append('### 明細列 LineAmount 不一致')
        for c in line_amount_issues:
            lines.append(f"- {c['QuoteID']}: " + '; '.join(c['LineIssues']))
lines.append('')
lines.append('## ValidUntil < 2026-09-30 且尚未 Accepted 的報價')
if expired_unaccepted:
    lines.append('| QuoteID | CustomerID | ProjectID | ValidUntil | Status | Currency | NetAmount |')
    lines.append('|---|---|---|---|---|---|---:|')
    for e in expired_unaccepted:
        lines.append(f"| {e['QuoteID']} | {e['CustomerID']} | {e['ProjectID']} | {e['ValidUntil']} | {e['Status']} | {e['Currency']} | {e['NetAmount']} |")
else:
    lines.append('- 無。')
lines.append('')
lines.append('## 每份報價檢查明細')
lines.append('| QuoteID | Lines | Line Sum | Subtotal OK | Expected NetAmount | NetAmount OK |')
lines.append('|---|---:|---:|---|---:|---|')
for c in checks:
    lines.append(f"| {c['QuoteID']} | {c['LineCount']} | {c['LineSum']} | {'OK' if c['SubtotalOK'] else 'NG'} | {c['ExpectedNetAmount']} | {'OK' if c['NetOK'] else 'NG'} |")
OUT_MD.write_text('\n'.join(lines) + '\n', encoding='utf-8')

# Reopen and verify output requirements from saved files.
wb2 = load_workbook(OUT_XLSX, data_only=False)
errors = []
if wb2.sheetnames != ['Report', 'Quotes', 'Lines', 'Summary', 'Instructions']:
    errors.append(f'sheet names changed: {wb2.sheetnames}')
qws = wb2['Quotes']
lws = wb2['Lines']
sws = wb2['Summary']
expected_quote_headers = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
expected_line_headers = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if [qws.cell(4,c).value for c in range(1,12)] != expected_quote_headers:
    errors.append('Quotes headers mismatch')
if [lws.cell(4,c).value for c in range(1,7)] != expected_line_headers:
    errors.append('Lines headers mismatch')
quote_data_rows = [r for r in qws.iter_rows(min_row=5, max_row=qws.max_row, values_only=True) if any(v is not None for v in r)]
line_data_rows = [r for r in lws.iter_rows(min_row=5, max_row=lws.max_row, values_only=True) if any(v is not None for v in r)]
if len(quote_data_rows) != len(quote_rows):
    errors.append(f'quote row count {len(quote_data_rows)} != {len(quote_rows)}')
if len(line_data_rows) != len(line_rows):
    errors.append(f'line row count {len(line_data_rows)} != {len(line_rows)}')
# IDs stored as text and dates/numbers are proper types.
for idx, row in enumerate(quote_data_rows, start=5):
    if not isinstance(qws.cell(idx,1).value, str):
        errors.append(f'QuoteID at row {idx} is not text')
    if qws.cell(idx,4).value is not None and not isinstance(qws.cell(idx,4).value, (datetime, date)):
        errors.append(f'QuoteDate at row {idx} is not date')
    for c in [8,9,10,11]:
        if not isinstance(qws.cell(idx,c).value, (int,float)):
            errors.append(f'Quote numeric col {c} row {idx} not numeric: {qws.cell(idx,c).value!r}')
for idx, row in enumerate(line_data_rows, start=5):
    if not isinstance(lws.cell(idx,1).value, str):
        errors.append(f'Line QuoteID at row {idx} is not text')
    for c in [4,5,6]:
        if not isinstance(lws.cell(idx,c).value, (int,float)):
            errors.append(f'Line numeric col {c} row {idx} not numeric: {lws.cell(idx,c).value!r}')
expected_formula = f'=SUM(Quotes!K5:K{4 + len(quote_rows) if quote_rows else 5})'
if sws['B2'].value != expected_formula:
    errors.append(f'Summary B2 formula {sws["B2"].value!r} != {expected_formula!r}')
md_text = OUT_MD.read_text(encoding='utf-8')
for e in expired_unaccepted:
    if e['QuoteID'] not in md_text:
        errors.append(f'expired unaccepted quote {e["QuoteID"]} missing from findings')
# ensure no September quote removed by expiry: all qids in output workbook.
out_qids = [r[0] for r in quote_data_rows]
expected_qids = [r[0] for r in quote_rows]
if out_qids != expected_qids:
    errors.append(f'output qids differ: {out_qids} != {expected_qids}')

if errors:
    print('VALIDATION FAILED')
    for e in errors:
        print('-', e)
    raise SystemExit(1)

print('WROTE', OUT_XLSX)
print('WROTE', OUT_MD)
print('validation_ok')
print('checked_quote_rows', len(quote_data_rows))
print('checked_line_rows', len(line_data_rows))
print('expired_unaccepted_qids', [e['QuoteID'] for e in expired_unaccepted])
print('amount_check_all_ok', all_ok)
print('summary_formula', expected_formula)
print('files_in_output', sorted(p.name for p in OUTPUT.iterdir()))