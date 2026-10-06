from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('cwd', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)

# quote template structure
xlsx = ws/'input'/'quote_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('\nSHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        print(r)

# fx rates headers/sample
csvp = ws/'input'/'fx_rates.csv'
print('\nfx_rates sample:')
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for i,row in zip(range(8), reader):
        print(row)

# manual
md = ws/'input'/'quote_cli_manual.md'
print('\nmanual:')
print(md.read_text(encoding='utf-8')[:4000])