from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
for p in (ws/'input').iterdir():
    print('\nFILE', p.name, 'size', p.stat().st_size)
    if p.suffix.lower()=='.csv':
        with p.open(newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f)
            for i,row in zip(range(6), r): print(' row', i, row)
    elif p.suffix.lower()=='.md':
        print(p.read_text(encoding='utf-8')[:1200])
    elif p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print(' sheets', wb.sheetnames)
        for s in wb.sheetnames:
            sh=wb[s]
            print(' sheet', s, 'max', sh.max_row, sh.max_column)
            for row in sh.iter_rows(min_row=1, max_row=min(8, sh.max_row), values_only=True):
                print('  ', row)
