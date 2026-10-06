from pathlib import Path
import csv, shutil, datetime, json
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import Font

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)
EXPORT = OUTPUT / 'expense_export.csv'
TEMPLATE = INPUT / 'expense_template.xlsx'
FX = INPUT / 'fx_rates.csv'
OUTFILE = OUTPUT / 'expense_report_2026-09_approved.xlsx'

print('Inputs:')
for p in [EXPORT, TEMPLATE, FX]:
    print(f' - {p.relative_to(WORKSPACE)} exists={p.exists()} size={p.stat().st_size if p.exists() else None}')

# Robust CSV read for fx_rates.csv because Note contains an unquoted comma in this synthetic fixture.
rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    print('FX header:', header)
    for row in reader:
        if len(row) < 2:
            continue
        currency = row[0].strip()
        rate = Decimal(row[1].strip())
        rates[currency] = rate
print('Rates loaded:', {k: str(v) for k, v in rates.items()})

with EXPORT.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    source_headers = reader.fieldnames
    rows = list(reader)
print('Export headers:', source_headers)
print('Export row count:', len(rows))
print('Export sample:', rows[:2])

# Validate and transform rows
start = datetime.date(2026, 9, 1)
end = datetime.date(2026, 9, 30)
required = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt']
missing_cols = [c for c in required if c not in source_headers]
if missing_cols:
    raise RuntimeError(f'Missing export columns: {missing_cols}')

transformed = []
exceptions = []
for r in rows:
    expense_date = datetime.date.fromisoformat(r['ExpenseDate'])
    if not (start <= expense_date <= end):
        raise RuntimeError(f'Row outside target date range: {r}')
    if r['Status'] != 'Approved':
        raise RuntimeError(f'Row not Approved: {r}')
    cur = r['Currency'].strip()
    if cur not in rates:
        raise RuntimeError(f'Missing FX rate for {cur}')
    amount = Decimal(r['Amount']).quantize(Decimal('0.01'))
    amount_twd = (amount * rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    needs_review = amount_twd > Decimal('5000')
    missing_receipt = str(r.get('MissingReceipt','')).strip().lower() == 'true'
    out = dict(r)
    out['ExpenseDate'] = expense_date
    out['Amount'] = amount
    out['AmountTWD'] = amount_twd
    out['NeedsReview'] = needs_review
    out['MissingReceipt'] = missing_receipt
    transformed.append(out)
    if missing_receipt:
        exceptions.append({
            'ExpenseID': r['ExpenseID'],
            'EmployeeID': r['EmployeeID'],
            'ExpenseDate': expense_date,
            'Reason': 'MissingReceipt is true'
        })

# Sort deterministically by date then ExpenseID as exported is already sorted but this makes output stable.
transformed.sort(key=lambda x: (x['ExpenseDate'], x['ExpenseID']))

# Copy template to output
shutil.copyfile(TEMPLATE, OUTFILE)
wb = load_workbook(OUTFILE)

# Report sheet
report = wb['Report']
report['B3'] = start.isoformat()
report['D3'] = end.isoformat()
report['B4'] = 'Status=Approved; ExpenseDate 2026-09-01..2026-09-30; AmountTWD uses fixed rates from fx_rates.csv; NeedsReview=true when AmountTWD>5000'
now = datetime.datetime.now(datetime.timezone(datetime.timedelta(hours=8))).replace(microsecond=0).isoformat()
report['B5'] = now
report['B7'] = 'Expense Viewer; fx_rates.csv'
report['D7'] = now
# Fill any obvious count/total placeholders if present in lower report area, without changing labels.
for row in range(1, report.max_row + 1):
    for col in range(1, report.max_column + 1):
        val = report.cell(row, col).value
        if isinstance(val, str):
            low = val.lower()
            right = report.cell(row, col + 1) if col < report.max_column else None
            if right is not None and right.value is None:
                if 'count' in low or '筆數' in val:
                    right.value = len(transformed)
                elif 'total' in low or '總計' in val:
                    right.value = float(sum((x['AmountTWD'] for x in transformed), Decimal('0.00')))

# Details sheet
headers_details = ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
details = wb['Details']
# Clear existing data rows from row 5 onward
if details.max_row >= 5:
    details.delete_rows(5, details.max_row - 4)
# Ensure headers match template
for c, h in enumerate(headers_details, 1):
    if details.cell(4, c).value != h:
        raise RuntimeError(f'Details header mismatch col {c}: expected {h}, got {details.cell(4,c).value}')
for r_idx, r in enumerate(transformed, 5):
    values = [
        r['ExpenseID'], r['EmployeeID'], r['EmployeeName'], r['ProjectID'], r['ExpenseDate'], r['Category'], r['Currency'],
        float(r['Amount']), float(r['AmountTWD']), r['Status'], r['ReceiptID'], r['Description'], r['MissingReceipt'], r['NeedsReview']
    ]
    for c_idx, v in enumerate(values, 1):
        details.cell(r_idx, c_idx).value = v
    details.cell(r_idx, 5).number_format = 'yyyy-mm-dd'
    details.cell(r_idx, 8).number_format = '#,##0.00'
    details.cell(r_idx, 9).number_format = '#,##0.00'

# Summary by employee and category, with currency included to preserve source currency amount if an employee/category has multiple currencies.
summary = wb['Summary']
if summary.max_row >= 5:
    summary.delete_rows(5, summary.max_row - 4)
summary_headers = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
for c, h in enumerate(summary_headers, 1):
    if summary.cell(4, c).value != h:
        raise RuntimeError(f'Summary header mismatch col {c}: expected {h}, got {summary.cell(4,c).value}')
agg = defaultdict(lambda: {'Amount': Decimal('0.00'), 'AmountTWD': Decimal('0.00'), 'Count': 0})
for r in transformed:
    # Group shows employee id and name, SubGroup is Category. Currency remains separate so original currency totals are meaningful.
    key = (f"{r['EmployeeID']} {r['EmployeeName']}", r['Category'], r['Currency'])
    agg[key]['Amount'] += r['Amount']
    agg[key]['AmountTWD'] += r['AmountTWD']
    agg[key]['Count'] += 1
for idx, (key, vals) in enumerate(sorted(agg.items()), 5):
    group, cat, cur = key
    summary.cell(idx, 1).value = group
    summary.cell(idx, 2).value = cat
    summary.cell(idx, 3).value = cur
    summary.cell(idx, 4).value = float(vals['Amount'].quantize(Decimal('0.01')))
    summary.cell(idx, 5).value = float(vals['AmountTWD'].quantize(Decimal('0.01')))
    summary.cell(idx, 6).value = vals['Count']
    summary.cell(idx, 4).number_format = '#,##0.00'
    summary.cell(idx, 5).number_format = '#,##0.00'
summary['B2'] = f'=SUM(Details!I5:I{4+len(transformed)})'

# Exceptions sheet: template exists; populate missing receipt exceptions only, because the instruction says Details include MissingReceipt.
exceptions_sh = wb['Exceptions']
if exceptions_sh.max_row >= 5:
    exceptions_sh.delete_rows(5, exceptions_sh.max_row - 4)
for c, h in enumerate(['ExpenseID','EmployeeID','ExpenseDate','Reason'], 1):
    if exceptions_sh.cell(4, c).value != h:
        raise RuntimeError(f'Exceptions header mismatch col {c}: expected {h}, got {exceptions_sh.cell(4,c).value}')
for idx, ex in enumerate(exceptions, 5):
    exceptions_sh.cell(idx, 1).value = ex['ExpenseID']
    exceptions_sh.cell(idx, 2).value = ex['EmployeeID']
    exceptions_sh.cell(idx, 3).value = ex['ExpenseDate']
    exceptions_sh.cell(idx, 3).number_format = 'yyyy-mm-dd'
    exceptions_sh.cell(idx, 4).value = ex['Reason']

# Freeze panes / autofilter (non-destructive convenience)
details.auto_filter.ref = f'A4:N{4+len(transformed)}'
details.freeze_panes = 'A5'
summary.auto_filter.ref = f'A4:F{4+len(agg)}'
summary.freeze_panes = 'A5'
exceptions_sh.auto_filter.ref = f'A4:D{4+len(exceptions)}'
exceptions_sh.freeze_panes = 'A5'

wb.save(OUTFILE)
print('Wrote:', OUTFILE.relative_to(WORKSPACE))

# Verification from saved workbook
wb2 = load_workbook(OUTFILE, data_only=False)
d = wb2['Details']
s = wb2['Summary']
e = wb2['Exceptions']
# Row counts
actual_details = d.max_row - 4
if actual_details != len(transformed):
    raise RuntimeError(f'Details row count mismatch {actual_details} != {len(transformed)}')
# Check headers
for c, h in enumerate(headers_details, 1):
    assert d.cell(4, c).value == h
for c, h in enumerate(summary_headers, 1):
    assert s.cell(4, c).value == h
# Check all details values, dates, statuses, original currency/amount and NeedsReview rule
bad = []
for row in range(5, d.max_row + 1):
    cur = d.cell(row, 7).value
    amt = Decimal(str(d.cell(row, 8).value)).quantize(Decimal('0.01'))
    twd = Decimal(str(d.cell(row, 9).value)).quantize(Decimal('0.01'))
    status = d.cell(row, 10).value
    nr = d.cell(row, 14).value
    expected_twd = (amt * rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    expected_nr = expected_twd > Decimal('5000')
    if status != 'Approved' or twd != expected_twd or nr != expected_nr or cur not in rates:
        bad.append((row, cur, amt, twd, expected_twd, status, nr, expected_nr))
if bad:
    raise RuntimeError('Details verification failed: ' + repr(bad[:5]))
# Check date range and type
for row in range(5, d.max_row + 1):
    dt = d.cell(row, 5).value
    if not isinstance(dt, datetime.datetime) and not isinstance(dt, datetime.date):
        raise RuntimeError(f'ExpenseDate not date type at row {row}: {dt!r}')
    dd = dt.date() if isinstance(dt, datetime.datetime) else dt
    if not (start <= dd <= end):
        raise RuntimeError(f'Date outside range at row {row}: {dd}')
# Check summary reconciles to details
summary_count = 0
summary_twd = Decimal('0.00')
summary_orig = defaultdict(Decimal)
for row in range(5, s.max_row + 1):
    cnt = s.cell(row, 6).value
    if cnt is None:
        continue
    summary_count += int(cnt)
    cur = s.cell(row, 3).value
    orig = Decimal(str(s.cell(row, 4).value)).quantize(Decimal('0.01'))
    twd = Decimal(str(s.cell(row, 5).value)).quantize(Decimal('0.01'))
    summary_orig[cur] += orig
    summary_twd += twd
if summary_count != len(transformed):
    raise RuntimeError(f'Summary count mismatch {summary_count} != {len(transformed)}')
detail_twd = sum((Decimal(str(d.cell(row, 9).value)).quantize(Decimal('0.01')) for row in range(5, d.max_row + 1)), Decimal('0.00'))
if summary_twd != detail_twd:
    raise RuntimeError(f'Summary TWD total mismatch {summary_twd} != {detail_twd}')
# Check exception count equals MissingReceipt true
missing_count = sum(1 for row in range(5, d.max_row + 1) if d.cell(row, 13).value is True)
exception_rows = max(0, e.max_row - 4)
if exception_rows != missing_count:
    raise RuntimeError(f'Exceptions rows {exception_rows} != MissingReceipt true count {missing_count}')
review_count = sum(1 for row in range(5, d.max_row + 1) if d.cell(row, 14).value is True)
print('Verification OK')
print(json.dumps({
    'details_rows': len(transformed),
    'summary_rows': len(agg),
    'exceptions_rows': exception_rows,
    'needs_review_true': review_count,
    'total_twd': str(detail_twd),
    'output_file': str(OUTFILE.relative_to(WORKSPACE)),
}, ensure_ascii=False, indent=2))
