from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
print('manual rest:')
text=(ws/'input'/'expense_manual.md').read_text(encoding='utf-8')
print(text)