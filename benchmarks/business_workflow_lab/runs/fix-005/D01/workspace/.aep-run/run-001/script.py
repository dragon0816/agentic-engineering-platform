from pathlib import Path
import csv
from openpyxl import load_workbook
base=Path('.')
print('Workspace:', base.resolve())
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print(' -', p, p.stat().st_size, 'bytes')
# workbook sheets and headers/sample
wb_path=base/'input'/'expense_template.xlsx'
wb=load_workbook(wb_path)
print('Workbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('SHEET', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
    for r in ws.iter_rows(min_row=1, max_row=min(ws.max_row,5), values_only=True):
        print(' ', r)
# csv headers/sample
csv_path=base/'input'/'fx_rates.csv'
print('CSV sample:')
with open(csv_path, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(6), reader):
        print(row)
# manual first part
md_path=base/'input'/'expense_manual.md'
print('Manual preview:')
print(md_path.read_text(encoding='utf-8')[:2000])