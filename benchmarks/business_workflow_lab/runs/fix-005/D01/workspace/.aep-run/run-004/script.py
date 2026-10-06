from pathlib import Path
from openpyxl import load_workbook
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv, os, json

workspace = Path('.').resolve()
input_dir = workspace/'input'
output_dir = workspace/'output'
output_dir.mkdir(exist_ok=True)

template = input_dir/'expense_template.xlsx'
export_path = output_dir/'expense_export_E001_2026-09-01_2026-09-30.xlsx'
fx_path = input_dir/'fx_rates.csv'
out_path = output_dir/'E001_september_expenses.xlsx'

# Read rates; csv has commas in note field but first two columns are stable.
rates = {}
with open(fx_path, newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0].strip():
            continue
        rates[row[0].strip()] = Decimal(row[1].strip())
print('FX rates:', rates)

# Read exported workbook.
exp_wb = load_workbook(export_path, data_only=True)
exp_ws = exp_wb[exp_wb.sheetnames[0]]
headers = [c.value for c in exp_ws[1]]
print('Export headers:', headers)
records=[]
for row in exp_ws.iter_rows(min_row=2, values_only=True):
    if all(v is None for v in row):
        continue
    rec=dict(zip(headers,row))
    records.append(rec)
print('Export record count:', len(records))

# Fill template.
wb = load_workbook(template)
ws_report = wb['Report']
ws_details = wb['Details']
ws_summary = wb['Summary']
ws_ex = wb['Exceptions']

# Clear old data rows while preserving row 4 headers / sheet layout as much as possible.
def clear_from(ws, start_row, max_col):
    for row in ws.iter_rows(min_row=start_row, max_row=ws.max_row, max_col=max_col):
        for cell in row:
            cell.value = None
clear_from(ws_details, 5, 14)
clear_from(ws_summary, 5, 6)
clear_from(ws_ex, 5, 4)

employee = 'E001'
start_date = '2026-09-01'
end_date = '2026-09-30'
now = datetime.now().astimezone().replace(microsecond=0).isoformat()
ws_report['B3'] = start_date
ws_report['D3'] = end_date
ws_report['B4'] = f'EmployeeID={employee}; ExpenseDate {start_date} to {end_date}; Status=ALL; Project=ALL'
ws_report['B5'] = now
ws_report['B7'] = 'Expense Viewer (synthetic lab) export'
ws_report['D7'] = now

summary = {}
exceptions = []
quant = Decimal('0.01')
for idx, rec in enumerate(records, start=5):
    currency = str(rec.get('Currency') or '').strip()
    if currency not in rates:
        raise ValueError(f'Missing FX rate for {currency}')
    amount = Decimal(str(rec.get('Amount'))).quantize(quant, rounding=ROUND_HALF_UP)
    amount_twd = (amount * rates[currency]).quantize(quant, rounding=ROUND_HALF_UP)
    missing = rec.get('MissingReceipt')
    missing_bool = str(missing).strip().lower() in ('true','1','yes','y') if not isinstance(missing, bool) else missing
    needs_review = missing_bool
    vals = [
        rec.get('ExpenseID'), rec.get('EmployeeID'), rec.get('EmployeeName'), rec.get('ProjectID'),
        rec.get('ExpenseDate'), rec.get('Category'), currency, float(amount), float(amount_twd),
        rec.get('Status'), rec.get('ReceiptID'), rec.get('Description'), str(missing_bool).lower(), str(needs_review).lower()
    ]
    for col, val in enumerate(vals, start=1):
        cell = ws_details.cell(idx, col, val)
        if col in (1,2,4,11,13,14):
            cell.number_format = '@'
        if col == 5:
            cell.number_format = 'yyyy-mm-dd'
        if col in (8,9):
            cell.number_format = '0.00'
    key=(str(rec.get('Status') or ''), currency)
    if key not in summary:
        summary[key]={'amount':Decimal('0.00'),'twd':Decimal('0.00'),'count':0}
    summary[key]['amount'] += amount
    summary[key]['twd'] += amount_twd
    summary[key]['count'] += 1
    if missing_bool:
        exceptions.append((rec.get('ExpenseID'), rec.get('EmployeeID'), rec.get('ExpenseDate'), 'MissingReceipt=true'))

last_detail_row = 4 + len(records) if records else 5
if records:
    ws_summary['B2'] = f'=SUM(Details!I5:I{last_detail_row})'
else:
    ws_summary['B2'] = '=SUM(Details!I5:I5)'

for r, ((status,currency), vals) in enumerate(sorted(summary.items()), start=5):
    row_vals=['Status', status, currency, float(vals['amount'].quantize(quant)), float(vals['twd'].quantize(quant)), vals['count']]
    for c, val in enumerate(row_vals, start=1):
        cell=ws_summary.cell(r,c,val)
        if c in (4,5): cell.number_format='0.00'

for r, vals in enumerate(exceptions, start=5):
    for c, val in enumerate(vals, start=1):
        cell=ws_ex.cell(r,c,val)
        if c in (1,2): cell.number_format='@'
        if c == 3: cell.number_format='yyyy-mm-dd'

note = f'匯出 {len(records)} 筆；FX rates from fx_rates.csv；NeedsReview/Exceptions based on MissingReceipt=true.'
ws_report['B9'] = note

wb.save(out_path)
print('Wrote', out_path)
print('Summary rows:', len(summary), 'Exceptions:', len(exceptions))
print('Total TWD:', sum(v['twd'] for v in summary.values()).quantize(quant))

# Verification from saved file.
check = load_workbook(out_path, data_only=False)
assert check.sheetnames == ['Report', 'Details', 'Summary', 'Exceptions', 'Instructions']
d = check['Details']; s=check['Summary']; e=check['Exceptions']; r=check['Report']
# headers unchanged
assert [d.cell(4,c).value for c in range(1,15)] == ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','AmountTWD','Status','ReceiptID','Description','MissingReceipt','NeedsReview']
# row count
actual_details = sum(1 for row in d.iter_rows(min_row=5, max_col=14, values_only=True) if any(v is not None for v in row))
assert actual_details == len(records), (actual_details, len(records))
# dates/employee filter
for row in d.iter_rows(min_row=5, max_row=4+len(records), values_only=True):
    assert row[1] == 'E001'
    assert '2026-09-01' <= str(row[4]) <= '2026-09-30'
    assert isinstance(row[7], (int,float))
    assert isinstance(row[8], (int,float))
    assert row[12] in ('true','false')
    assert row[13] in ('true','false')
# Summary formula and counts
assert s['B2'].value == f'=SUM(Details!I5:I{4+len(records)})'
actual_summary_rows = sum(1 for row in s.iter_rows(min_row=5, max_col=6, values_only=True) if any(v is not None for v in row))
assert actual_summary_rows == len(summary)
summary_count = sum(s.cell(row,6).value or 0 for row in range(5,5+len(summary)))
assert summary_count == len(records)
# TWD total equivalence
sum_details = Decimal('0.00')
for row in range(5,5+len(records)):
    sum_details += Decimal(str(d.cell(row,9).value)).quantize(quant)
sum_summary = Decimal('0.00')
for row in range(5,5+len(summary)):
    sum_summary += Decimal(str(s.cell(row,5).value)).quantize(quant)
assert sum_details == sum_summary
actual_ex = sum(1 for row in e.iter_rows(min_row=5, max_col=4, values_only=True) if any(v is not None for v in row))
assert actual_ex == len(exceptions)
print('CHECK PASSED')
print('Output:', out_path)
print('Details rows checked:', actual_details)
print('Summary groups checked:', actual_summary_rows)
print('Exceptions checked:', actual_ex)
print('TWD total checked:', sum_details)