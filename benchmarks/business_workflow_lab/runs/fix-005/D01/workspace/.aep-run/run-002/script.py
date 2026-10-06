from pathlib import Path
from openpyxl import load_workbook
base=Path('.')
print('Output files currently:')
for p in sorted((base/'output').glob('*')):
    print(' -', p, p.stat().st_size, 'bytes')
exp=base/'output'/'expense_export_E001_2026-09-01_2026-09-30.xlsx'
wb=load_workbook(exp, data_only=False)
print('Export sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('Sheet', ws.title, 'rows', ws.max_row, 'cols', ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,6), values_only=True):
        print(row)