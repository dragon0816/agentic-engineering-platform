from pathlib import Path
from openpyxl import load_workbook
import csv
ws=Path.cwd()
print('workspace', ws)
print('input files:', [p.name for p in (ws/'input').iterdir()])
print('output files:', [p.name for p in (ws/'output').iterdir()] if (ws/'output').exists() else [])
for p in (ws/'input').iterdir():
    print('\nFILE', p.name, 'size', p.stat().st_size)
    if p.suffix.lower()=='.xlsx':
        wb=load_workbook(p, data_only=False)
        print('sheets', wb.sheetnames)
        for s in wb.sheetnames:
            sh=wb[s]
            print('sheet', s, 'max', sh.max_row, sh.max_column)
            rows=[]
            for row in sh.iter_rows(min_row=1, max_row=min(5, sh.max_row), values_only=True):
                rows.append(row)
            for r in rows: print(r)
    elif p.suffix.lower()=='.csv':
        with p.open(newline='', encoding='utf-8-sig') as f:
            rdr=csv.reader(f)
            for i,row in zip(range(5), rdr): print(row)
    elif p.suffix.lower()=='.md':
        print(p.read_text(encoding='utf-8')[:1000])