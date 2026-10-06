from pathlib import Path
from openpyxl import load_workbook
import csv
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)

xlsx = ws/'input'/'project_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('\nSHEET', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=True):
        print(r)

csvp = ws/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with csvp.open(newline='', encoding='utf-8-sig') as f:
    rdr = csv.reader(f)
    for i,row in zip(range(10), rdr):
        print(row)

md = ws/'input'/'sap_manual.md'
print('\nsap_manual.md first 120 lines:')
text = md.read_text(encoding='utf-8')
for i,line in enumerate(text.splitlines()[:120], start=1):
    print(f'{i:03d}: {line}')