from pathlib import Path
from openpyxl import load_workbook
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, date
import csv

ws = Path.cwd()
input_dir = ws / 'input'
output_dir = ws / 'output'
download_dir = output_dir / 'downloads'
output_dir.mkdir(exist_ok=True)

print('Workspace:', ws)
print('Input files:', [p.name for p in sorted(input_dir.glob('*'))])
print('Downloaded files:', [p.name for p in sorted(download_dir.glob('*'))] if download_dir.exists() else [])

csv_path = download_dir / 'invoices_2026_09.csv'
template_path = input_dir / 'project_template.xlsx'
fx_path = input_dir / 'fx_rates.csv'
out_path = output_dir / 'project_report_2026-09_invoices.xlsx'

if not csv_path.exists():
    raise FileNotFoundError(f'Missing downloaded invoice CSV: {csv_path}')

rates = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    rdr = csv.reader(f)
    header = next(rdr)
    print('fx header:', header)
    for row in rdr:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('fx rates:', {k: str(v) for k,v in rates.items()})

with csv_path.open(newline='', encoding='utf-8-sig') as f:
    rdr = csv.DictReader(f)
    rows = list(rdr)
    csv_headers = rdr.fieldnames
print('invoice csv headers:', csv_headers)
print('invoice row count:', len(rows))
print('invoice sample rows:', rows[:3])

required = ['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance']
if csv_headers != required:
    raise AssertionError(f'Unexpected invoice headers: {csv_headers}')

from_date = date(2026,9,1)
to_date = date(2026,9,30)
asof = date(2026,9,30)
invoice_rows = []
for r in rows:
    inv_date = datetime.strptime(r['InvoiceDate'], '%Y-%m-%d').date()
    if not (from_date <= inv_date <= to_date):
        continue
    due_date = datetime.strptime(r['DueDate'], '%Y-%m-%d').date()
    amount = Decimal(r['Amount'])
    paid = Decimal(r['PaidAmount'])
    balance = Decimal(r['Balance'])
    if balance != amount - paid:
        raise AssertionError(f"Balance mismatch for {r['InvoiceID']}: {balance} != {amount-paid}")
    cur = r['Currency']
    if cur not in rates:
        raise AssertionError(f'Missing FX rate for {cur}')
    balance_twd = (balance * rates[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    rr = dict(r)
    rr['_InvoiceDate'] = inv_date
    rr['_DueDate'] = due_date
    rr['_Amount'] = amount
    rr['_PaidAmount'] = paid
    rr['_Balance'] = balance
    rr['_BalanceTWD'] = balance_twd
    invoice_rows.append(rr)
invoice_rows.sort(key=lambda r: r['InvoiceID'])
if len(invoice_rows) != 13:
    raise AssertionError(f'Expected 13 September invoice rows from query, got {len(invoice_rows)}')

wb = load_workbook(template_path)
START_ROW = 5

def clear_data(wsht, min_row=5):
    if wsht.max_row >= min_row:
        wsht.delete_rows(min_row, wsht.max_row - min_row + 1)

def write_row(wsht, row_idx, values, template_row=5):
    for col_idx, value in enumerate(values, start=1):
        cell = wsht.cell(row=row_idx, column=col_idx, value=value)
        src = wsht.cell(row=template_row, column=col_idx)
        if src.has_style:
            cell._style = copy(src._style)
        if src.number_format:
            cell.number_format = src.number_format

# Report metadata
report = wb['Report']
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Invoices InvoiceDate between 2026-09-01 and 2026-09-30; Balance as of 2026-09-30; overdue when Balance > 0 and DueDate < 2026-09-30.'
report['B5'] = datetime.now().astimezone().isoformat(timespec='seconds')
report['B7'] = 'Lab ERP Invoices; fx_rates.csv'
report['D7'] = datetime.now().astimezone().isoformat(timespec='seconds')

for sheet_name in ['Orders', 'Expenses', 'Quotes']:
    clear_data(wb[sheet_name], START_ROW)

inv_sheet = wb['Invoices']
clear_data(inv_sheet, START_ROW)
for idx, r in enumerate(invoice_rows, start=START_ROW):
    write_row(inv_sheet, idx, [
        str(r['InvoiceID']), str(r['OrderID']), str(r['ProjectID']), str(r['CustomerID']),
        r['_InvoiceDate'], r['_DueDate'], str(r['Currency']),
        float(r['_Amount']), float(r['_PaidAmount']), float(r['_Balance']), float(r['_BalanceTWD'])
    ])
    for c in [1,2,3,4,7]:
        inv_sheet.cell(idx,c).number_format = '@'
    for c in [5,6]:
        inv_sheet.cell(idx,c).number_format = 'yyyy-mm-dd'
    for c in [8,9,10,11]:
        inv_sheet.cell(idx,c).number_format = '#,##0.00'

exc_sheet = wb['Exceptions']
clear_data(exc_sheet, START_ROW)
overdues = [r for r in invoice_rows if r['_Balance'] > 0 and r['_DueDate'] < asof]
for idx, r in enumerate(overdues, start=START_ROW):
    write_row(exc_sheet, idx, ['Invoices', str(r['InvoiceID']), str(r['ProjectID']) if r['ProjectID'] else None, 'Overdue'])
    for c in [1,2,3,4]:
        exc_sheet.cell(idx,c).number_format = '@'

summary = wb['Summary']
clear_data(summary, START_ROW)
by_project = {}
for r in invoice_rows:
    pid = str(r['ProjectID']).strip()
    if not pid:
        continue
    by_project[pid] = by_project.get(pid, Decimal('0.00')) + r['_BalanceTWD']
for idx, pid in enumerate(sorted(by_project), start=START_ROW):
    write_row(summary, idx, [pid, 0.0, float(by_project[pid]), 0.0, 0.0])
    summary.cell(idx,1).number_format = '@'
    for c in [2,3,4,5]:
        summary.cell(idx,c).number_format = '#,##0.00'
last_summary_row = max(START_ROW, START_ROW + len(by_project) - 1)
summary['B2'] = f'=SUM(B5:B{last_summary_row})'
summary['C2'] = f'=SUM(C5:C{last_summary_row})'
summary['D2'] = f'=SUM(D5:D{last_summary_row})'
summary['E2'] = f'=SUM(E5:E{last_summary_row})'

expected_headers = {
    'Invoices': ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
    'Summary': ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
    'Exceptions': ['Source','DocumentID','ProjectID','Reason'],
    'Orders': ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
    'Expenses': ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
    'Quotes': ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
}
for s, hdr in expected_headers.items():
    actual = [wb[s].cell(4,c).value for c in range(1, len(hdr)+1)]
    if actual != hdr:
        raise AssertionError(f'Header changed on {s}: {actual}')

wb.save(out_path)
print('Saved:', out_path)

# Reopen and validate saved output.
check_wb = load_workbook(out_path, data_only=False)
cinv = check_wb['Invoices']
csum = check_wb['Summary']
cexc = check_wb['Exceptions']

for s in ['Orders','Expenses','Quotes']:
    sh = check_wb[s]
    nonempty = []
    for row in sh.iter_rows(min_row=5):
        vals = [cell.value for cell in row]
        if any(v is not None for v in vals):
            nonempty.append(vals)
    if nonempty:
        raise AssertionError(f'{s} is not blank below header: {nonempty[:3]}')

for s, hdr in expected_headers.items():
    actual = [check_wb[s].cell(4,c).value for c in range(1, len(hdr)+1)]
    if actual != hdr:
        raise AssertionError(f'Saved header changed on {s}: {actual}')

invoice_saved_count = sum(1 for row in cinv.iter_rows(min_row=5, max_col=11) if any(c.value is not None for c in row))
if invoice_saved_count != len(invoice_rows):
    raise AssertionError(f'Invoice saved count {invoice_saved_count} != expected {len(invoice_rows)}')

for i, r in enumerate(invoice_rows, start=START_ROW):
    checks = [
        (1, r['InvoiceID'], str), (2, r['OrderID'], str), (3, r['ProjectID'] or None, (str, type(None))),
        (4, r['CustomerID'], str), (7, r['Currency'], str),
    ]
    for col, expected, typ in checks:
        got = cinv.cell(i,col).value
        if got != expected or not isinstance(got, typ):
            raise AssertionError(f'Cell mismatch at Invoices row {i} col {col}: {got!r} != {expected!r}')
    if cinv.cell(i,5).value.date() != r['_InvoiceDate'] if hasattr(cinv.cell(i,5).value, 'date') else cinv.cell(i,5).value != r['_InvoiceDate']:
        raise AssertionError(f'InvoiceDate mismatch row {i}')
    if cinv.cell(i,6).value.date() != r['_DueDate'] if hasattr(cinv.cell(i,6).value, 'date') else cinv.cell(i,6).value != r['_DueDate']:
        raise AssertionError(f'DueDate mismatch row {i}')
    for col, dec in [(8,r['_Amount']),(9,r['_PaidAmount']),(10,r['_Balance']),(11,r['_BalanceTWD'])]:
        got = Decimal(str(cinv.cell(i,col).value)).quantize(Decimal('0.01'))
        if got != dec:
            raise AssertionError(f'Numeric mismatch row {i} col {col}: {got} != {dec}')

saved_excs = []
for row in cexc.iter_rows(min_row=5, max_col=4, values_only=True):
    if any(v is not None for v in row):
        saved_excs.append(row)
expected_excs = [('Invoices', r['InvoiceID'], r['ProjectID'] or None, 'Overdue') for r in overdues]
if saved_excs != expected_excs:
    raise AssertionError(f'Exceptions mismatch. saved={saved_excs}, expected={expected_excs}')
# Check no due-on-asof is marked overdue.
due_on_asof_ids = {r['InvoiceID'] for r in invoice_rows if r['_DueDate'] == asof and r['_Balance'] > 0}
marked_ids = {row[1] for row in saved_excs}
if due_on_asof_ids & marked_ids:
    raise AssertionError(f'Due-on-asof invoices incorrectly marked overdue: {due_on_asof_ids & marked_ids}')

saved_summary = {}
for row in csum.iter_rows(min_row=5, max_col=5, values_only=True):
    if any(v is not None for v in row):
        pid, order_twd, inv_bal_twd, exp_twd, quote_twd = row
        if order_twd != 0 or exp_twd != 0 or quote_twd != 0:
            raise AssertionError(f'Non-invoice summary value not zero for {pid}: {row}')
        saved_summary[pid] = Decimal(str(inv_bal_twd)).quantize(Decimal('0.01'))
if saved_summary != by_project:
    raise AssertionError(f'Summary mismatch. saved={saved_summary}, expected={by_project}')

print('Validation OK')
print('Output file:', out_path)
print('Invoices rows saved:', invoice_saved_count)
print('Invoice IDs:', [r['InvoiceID'] for r in invoice_rows])
print('Overdue exceptions:', len(saved_excs), saved_excs)
print('Due on 2026-09-30 with balance, not overdue:', sorted(due_on_asof_ids))
print('Summary rows:', len(saved_summary), {k: str(v) for k,v in saved_summary.items()})
print('Other source sheets blank below row 4: Orders, Expenses, Quotes')
print('Summary total InvoiceBalanceTWD:', str(sum(by_project.values(), Decimal('0.00'))))
