from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES:')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')
print('\nOUTPUT FILES:')
out = ws/'output'
out.mkdir(exist_ok=True)
for p in sorted(out.glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

manual = ws/'input'/'quote_cli_manual.md'
print('\nMANUAL HEAD:')
print(manual.read_text(encoding='utf-8')[:4000])

fx = ws/'input'/'fx_rates.csv'
print('\nFX SAMPLE:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r = csv.reader(f)
    for i,row in zip(range(6), r):
        print(row)

tmpl = ws/'input'/'quote_template.xlsx'
wb = load_workbook(tmpl, data_only=False)
print('\nWORKBOOK SHEETS:', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print(f'-- {s}: max_row={sh.max_row} max_col={sh.max_column}')
    for row in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 8), values_only=False):
        print([cell.value for cell in row[:min(sh.max_column, 12)]])
    print('tables:', list(sh.tables.keys()))

print('\nPOSSIBLE CLI FILES IN WORKSPACE TREE:')
for p in ws.rglob('*'):
    if p.is_file() and any(x in p.name.lower() for x in ['quote','cli']):
        print(p.relative_to(ws), p.stat().st_size)
