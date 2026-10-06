import csv
import json
import os
import shutil
import subprocess
import sys
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook

params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
workspace = Path.cwd()
input_dir = workspace / 'input'
output_dir = workspace / 'output'
output_dir.mkdir(parents=True, exist_ok=True)

customer_id = str(params['customer_id']).strip()
from_date = str(params['from_date']).strip()
to_date = str(params['to_date']).strip()
cli_path = str(params['cli_path']).strip().strip('"')
if not customer_id:
    raise SystemExit('customer_id is required')
# Validate date strings and inclusive range.
from_dt = datetime.strptime(from_date, '%Y-%m-%d').date()
to_dt = datetime.strptime(to_date, '%Y-%m-%d').date()
if from_dt > to_dt:
    raise SystemExit('from_date must be on or before to_date')

cli_file = Path(cli_path)
if not cli_file.exists():
    found = shutil.which(cli_path)
    if found:
        cli_file = Path(found)
    else:
        raise SystemExit(f'CLI path not found: {cli_path}')

def q2(value):
    return Decimal(str(value)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

def money_to_excel(dec):
    return float(q2(dec))

def run_cli(args):
    if cli_file.suffix.lower() == '.py':
        cmd = [sys.executable, str(cli_file)] + args
        completed = subprocess.run(cmd, cwd=str(workspace), text=True, capture_output=True)
    elif cli_file.suffix.lower() in {'.cmd', '.bat'}:
        # Windows command scripts need shell=True. Keep arguments quoted by subprocess.list2cmdline.
        cmdline = subprocess.list2cmdline([str(cli_file)] + args)
        completed = subprocess.run(cmdline, cwd=str(workspace), text=True, capture_output=True, shell=True)
    else:
        cmd = [str(cli_file)] + args
        completed = subprocess.run(cmd, cwd=str(workspace), text=True, capture_output=True)
    if completed.returncode != 0:
        raise RuntimeError(f'CLI failed rc={completed.returncode}\nSTDERR:\n{completed.stderr}\nSTDOUT:\n{completed.stdout}')
    if completed.stderr.strip():
        print('CLI stderr:', completed.stderr.strip())
    return completed.stdout

# Load FX rates. CSV notes intentionally contain commas, so only first two columns are meaningful.
fx_rates = {}
with (input_dir / 'fx_rates.csv').open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if not row or not row[0].strip():
            continue
        fx_rates[row[0].strip()] = Decimal(row[1].strip())
if not fx_rates:
    raise SystemExit('No FX rates loaded')
print('Loaded FX currencies:', ', '.join(sorted(fx_rates)))

# Fetch list, then detail for each QuoteID.
list_json = run_cli(['quotes', 'list', '--customer', customer_id, '--from', from_date, '--to', to_date, '--format', 'json'])
list_data = json.loads(list_json)
if list_data.get('schema_version') != 1:
    raise SystemExit(f'Unsupported quote list schema_version: {list_data.get("schema_version")}')
items = list_data.get('items', [])
quote_ids = [item['QuoteID'] for item in items]
print(f'Fetched quote list: {len(quote_ids)} quote(s) for {customer_id} from {from_date} to {to_date}')

headers = []
lines = []
for qid in quote_ids:
    detail_json = run_cli(['quotes', 'get', qid, '--format', 'json'])
    detail = json.loads(detail_json)
    if detail.get('schema_version') != 1:
        raise SystemExit(f'Unsupported quote detail schema_version for {qid}: {detail.get("schema_version")}')
    header = detail['header']
    # Defensive re-check of filter contract.
    if header.get('CustomerID') != customer_id:
        raise SystemExit(f'Quote {qid} has unexpected CustomerID {header.get("CustomerID")}')
    qdate = datetime.strptime(header['QuoteDate'], '%Y-%m-%d').date()
    if qdate < from_dt or qdate > to_dt:
        raise SystemExit(f'Quote {qid} QuoteDate {qdate} outside requested range')
    currency = header['Currency']
    if currency not in fx_rates:
        raise SystemExit(f'Missing FX rate for currency {currency}')
    net = q2(header['NetAmount'])
    rate = fx_rates[currency]
    header['_NetAmountTWD'] = q2(net * rate)
    headers.append(header)
    for line in detail.get('line_items', []):
        enriched = dict(line)
        enriched['QuoteID'] = qid
        lines.append(enriched)

# Stable sort for reproducible output.
headers.sort(key=lambda h: (h['QuoteDate'], h['QuoteID']))
line_order = {h['QuoteID']: i for i, h in enumerate(headers)}
lines.sort(key=lambda l: (line_order.get(l['QuoteID'], 999999), l['QuoteID'], l.get('ProductID', '')))

# Prepare workbook from template.
template_path = input_dir / 'quote_template.xlsx'
out_path = output_dir / f'quotation_report_{customer_id}_{from_date}_{to_date}.xlsx'
wb = load_workbook(template_path)

# Helpers.
def clear_data_rows(ws, start_row, max_col):
    if ws.max_row >= start_row:
        ws.delete_rows(start_row, ws.max_row - start_row + 1)
    # Ensure header row still defines intended columns; no action needed for rows below deleted.

# Report sheet.
taipei_now = datetime.now(timezone(timedelta(hours=8))).replace(microsecond=0).isoformat()
report = wb['Report']
report['B3'] = from_date
report['D3'] = to_date
report['B4'] = f'CustomerID={customer_id}; QuoteDate {from_date} to {to_date} inclusive'
report['B5'] = taipei_now
report['B7'] = f'quote_cli; fx_rates.csv'
report['D7'] = taipei_now

# Quotes sheet.
quotes_ws = wb['Quotes']
quote_cols = [quotes_ws.cell(4, c).value for c in range(1, quotes_ws.max_column + 1)]
expected_quote_cols = ['QuoteID','CustomerID','ProjectID','QuoteDate','ValidUntil','Status','Currency','Subtotal','DiscountPercent','NetAmount','NetAmountTWD']
if quote_cols[:len(expected_quote_cols)] != expected_quote_cols:
    raise SystemExit(f'Unexpected Quotes header: {quote_cols}')
clear_data_rows(quotes_ws, 5, len(expected_quote_cols))
for r, h in enumerate(headers, start=5):
    values = [
        h['QuoteID'], h['CustomerID'], h.get('ProjectID'), h['QuoteDate'], h.get('ValidUntil'),
        h['Status'], h['Currency'], money_to_excel(h['Subtotal']), money_to_excel(h['DiscountPercent']),
        money_to_excel(h['NetAmount']), money_to_excel(h['_NetAmountTWD'])
    ]
    for c, value in enumerate(values, start=1):
        quotes_ws.cell(r, c).value = value

# Lines sheet.
lines_ws = wb['Lines']
line_cols = [lines_ws.cell(4, c).value for c in range(1, lines_ws.max_column + 1)]
expected_line_cols = ['QuoteID','ProductID','ProductName','Quantity','UnitPrice','LineAmount']
if line_cols[:len(expected_line_cols)] != expected_line_cols:
    raise SystemExit(f'Unexpected Lines header: {line_cols}')
clear_data_rows(lines_ws, 5, len(expected_line_cols))
for r, line in enumerate(lines, start=5):
    values = [
        line['QuoteID'], line.get('ProductID'), line.get('ProductName'),
        int(line.get('Quantity', 0)), money_to_excel(line.get('UnitPrice', '0')),
        money_to_excel(line.get('LineAmount', '0'))
    ]
    for c, value in enumerate(values, start=1):
        lines_ws.cell(r, c).value = value

# Summary by status and currency.
summary_ws = wb['Summary']
summary_cols = [summary_ws.cell(4, c).value for c in range(1, summary_ws.max_column + 1)]
expected_summary_cols = ['Group','SubGroup','Currency','Amount','AmountTWD','Count']
if summary_cols[:len(expected_summary_cols)] != expected_summary_cols:
    raise SystemExit(f'Unexpected Summary header: {summary_cols}')
clear_data_rows(summary_ws, 5, len(expected_summary_cols))
summary = defaultdict(lambda: {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
for h in headers:
    key = (h['Status'], h['Currency'])
    summary[key]['amount'] += q2(h['NetAmount'])
    summary[key]['amount_twd'] += q2(h['_NetAmountTWD'])
    summary[key]['count'] += 1
for r, ((status, currency), agg) in enumerate(sorted(summary.items()), start=5):
    values = ['Status', status, currency, money_to_excel(agg['amount']), money_to_excel(agg['amount_twd']), int(agg['count'])]
    for c, value in enumerate(values, start=1):
        summary_ws.cell(r, c).value = value
last_quote_row = 4 + len(headers)
if headers:
    summary_ws['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
else:
    summary_ws['B2'] = 0

# Save output only; never modify input template.
wb.save(out_path)
print('Wrote', out_path)

# Verification by reopening saved file.
check_wb = load_workbook(out_path, data_only=False)
checks = []
qws = check_wb['Quotes']
lws = check_wb['Lines']
sws = check_wb['Summary']
rws = check_wb['Report']
checks.append(('Report B3 from_date', rws['B3'].value == from_date))
checks.append(('Report D3 to_date', rws['D3'].value == to_date))
checks.append(('Quotes header names', [qws.cell(4,c).value for c in range(1,12)] == expected_quote_cols))
checks.append(('Lines header names', [lws.cell(4,c).value for c in range(1,7)] == expected_line_cols))
checks.append(('Summary header names', [sws.cell(4,c).value for c in range(1,7)] == expected_summary_cols))
checks.append(('Quotes row count', max(qws.max_row - 4, 0) == len(headers)))
checks.append(('Lines row count', max(lws.max_row - 4, 0) == len(lines)))
checks.append(('Summary row count', max(sws.max_row - 4, 0) == len(summary)))
# Text IDs remain strings.
for row in range(5, 5 + len(headers)):
    checks.append((f'QuoteID text row {row}', isinstance(qws.cell(row,1).value, str)))
    checks.append((f'CustomerID text row {row}', isinstance(qws.cell(row,2).value, str)))
# Numeric money/count fields are numeric in workbook.
for row in range(5, 5 + len(headers)):
    for col in [8,9,10,11]:
        checks.append((f'Quotes numeric r{row}c{col}', isinstance(qws.cell(row,col).value, (int, float))))
for row in range(5, 5 + len(lines)):
    checks.append((f'Lines quantity numeric row {row}', isinstance(lws.cell(row,4).value, int)))
    for col in [5,6]:
        checks.append((f'Lines numeric r{row}c{col}', isinstance(lws.cell(row,col).value, (int, float))))
for row in range(5, 5 + len(summary)):
    for col in [4,5,6]:
        checks.append((f'Summary numeric r{row}c{col}', isinstance(sws.cell(row,col).value, (int, float))))
# Verify summary values match workbook Quotes data.
book_summary = defaultdict(lambda: {'amount': Decimal('0'), 'amount_twd': Decimal('0'), 'count': 0})
for row in range(5, 5 + len(headers)):
    st = qws.cell(row,6).value
    cur = qws.cell(row,7).value
    book_summary[(st,cur)]['amount'] += q2(qws.cell(row,10).value)
    book_summary[(st,cur)]['amount_twd'] += q2(qws.cell(row,11).value)
    book_summary[(st,cur)]['count'] += 1
for row in range(5, 5 + len(summary)):
    key = (sws.cell(row,2).value, sws.cell(row,3).value)
    checks.append((f'Summary amount {key}', q2(sws.cell(row,4).value) == q2(book_summary[key]['amount'])))
    checks.append((f'Summary amount TWD {key}', q2(sws.cell(row,5).value) == q2(book_summary[key]['amount_twd'])))
    checks.append((f'Summary count {key}', sws.cell(row,6).value == book_summary[key]['count']))

failed = [name for name, ok in checks if not ok]
if failed:
    raise SystemExit('Verification failed: ' + '; '.join(failed))
print('Verified output workbook:')
print(f'- file: {out_path.name}')
print(f'- Quotes rows: {len(headers)}')
print(f'- Lines rows: {len(lines)}')
print(f'- Summary groups: {len(summary)}')
print('- Checked sheet names, column names, row counts, text IDs, numeric amounts/counts, and Summary totals by status/currency')
