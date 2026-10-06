from pathlib import Path
from openpyxl import load_workbook
import csv
ws = Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').glob('*')):
    print(' -', p.name, p.stat().st_size)
# workbook structure
xlsx = ws/'input'/'expense_template.xlsx'
wb = load_workbook(xlsx, data_only=False)
print('template sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in range(1, min(sh.max_row,8)+1):
        vals=[sh.cell(r,c).value for c in range(1,min(sh.max_column,10)+1)]
        print(' row', r, vals)
# csv
csvp=ws/'input'/'fx_rates.csv'
with open(csvp, newline='', encoding='utf-8-sig') as f:
    rows=list(csv.reader(f))
print('fx rows sample:')
for row in rows[:8]: print(row)
# manual
md=ws/'input'/'expense_manual.md'
text=md.read_text(encoding='utf-8')
print('manual first 2000 chars:')
print(text[:2000])