from pathlib import Path
from openpyxl import load_workbook
wb=load_workbook('input/expense_template.xlsx', data_only=False)
for wsname in ['Instructions','Report','Details','Summary','Exceptions']:
    ws=wb[wsname]
    print('\n---', wsname, '---')
    for r in range(1, ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        print(r, vals)
print('\nMerged ranges:')
for s in wb.sheetnames:
    print(s, list(wb[s].merged_cells.ranges))