from pathlib import Path
import csv, json
from openpyxl import load_workbook
base=Path('.')
print('CWD', Path.cwd())
print('input files:')
for p in sorted((base/'input').iterdir()):
    print('-', p.name, p.stat().st_size)

fx=base/'input'/'fx_rates.csv'
print('\nfx_rates.csv sample:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r):
        print(i, row)

xlsx=base/'input'/'project_template.xlsx'
wb=load_workbook(xlsx, data_only=False)
print('\nWorkbook sheets:')
for ws in wb.worksheets:
    print(ws.title, ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,8), values_only=True):
        print(' ', row)

for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=base/'input'/name
    print(f'\n===== {name} =====')
    text=p.read_text(encoding='utf-8')
    print(text[:5000])
    if len(text)>5000:
        print('...<truncated>...')
