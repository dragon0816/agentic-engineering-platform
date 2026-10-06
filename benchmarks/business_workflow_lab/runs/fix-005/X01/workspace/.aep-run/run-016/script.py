from pathlib import Path
import csv
from decimal import Decimal
from openpyxl import load_workbook

out=Path('output')
required={
 'project_integration_report.xlsx': None,
 'project_summary.csv': ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
 'orders_normalized.csv': ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],
 'invoices_normalized.csv': ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],
 'expenses_normalized.csv': ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],
 'quotes_normalized.csv': ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],
 'exceptions.csv': ['Source','DocumentID','ProjectID','Reason'],
 'expense_approved.csv': ['ExpenseID','EmployeeID','EmployeeName','ProjectID','ExpenseDate','Category','Currency','Amount','Status','ReceiptID','Description','MissingReceipt'],
 'downloads/sap_orders.csv': ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount'],
 'downloads/sap_invoices.csv': ['InvoiceID','OrderID','CustomerID','ProjectID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance'],
}
checks=[]
errors=[]
rows_by_name={}
for name, header in required.items():
    p=out/name
    if not p.exists() or p.stat().st_size==0:
        errors.append(f'missing/empty {name}')
        continue
    checks.append(f'{name} exists {p.stat().st_size} bytes')
    if header:
        with p.open(newline='', encoding='utf-8-sig') as f:
            data=list(csv.DictReader(f))
            rows_by_name[name]=data
            actual=list(data[0].keys()) if data else []
        # handle empty exceptions possible, but here not empty
        with p.open(newline='', encoding='utf-8-sig') as f:
            reader=csv.reader(f); actual_header=next(reader)
        if actual_header != header:
            errors.append(f'header mismatch {name}: {actual_header}')
        checks.append(f'{name} rows {len(data)}')

# filters/status rules
if any(r.get('Status')=='Cancelled' for r in rows_by_name.get('orders_normalized.csv', [])):
    errors.append('orders_normalized contains Cancelled')
if any(r.get('Status')!='Approved' for r in rows_by_name.get('expenses_normalized.csv', [])):
    errors.append('expenses_normalized contains non-Approved')
if any(r.get('Status')!='Accepted' for r in rows_by_name.get('quotes_normalized.csv', [])):
    errors.append('quotes_normalized contains non-Accepted')
if any(not r.get('ProjectID') for r in rows_by_name.get('project_summary.csv', [])):
    errors.append('project_summary contains blank ProjectID')

# numeric fields parse
numeric_fields={
 'project_summary.csv':['OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],
 'orders_normalized.csv':['NetAmount','AmountTWD'],
 'invoices_normalized.csv':['Amount','PaidAmount','Balance','BalanceTWD'],
 'expenses_normalized.csv':['Amount','AmountTWD'],
 'quotes_normalized.csv':['NetAmount','AmountTWD'],
}
for name, fields in numeric_fields.items():
    for i,r in enumerate(rows_by_name.get(name, []), start=2):
        for field in fields:
            try:
                Decimal(str(r.get(field,'')))
            except Exception:
                errors.append(f'{name} row {i} field {field} not numeric: {r.get(field)!r}')
                break

# Excel checks
xlsx=out/'project_integration_report.xlsx'
wb=load_workbook(xlsx, data_only=False)
expected_sheets=['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
if wb.sheetnames != expected_sheets:
    errors.append(f'xlsx sheets mismatch: {wb.sheetnames}')
else:
    checks.append('xlsx sheets match expected')
header_rows={
 'Orders': required['orders_normalized.csv'],
 'Invoices': required['invoices_normalized.csv'],
 'Expenses': required['expenses_normalized.csv'],
 'Quotes': required['quotes_normalized.csv'],
 'Summary': required['project_summary.csv'],
 'Exceptions': required['exceptions.csv'],
}
row_source={
 'Orders':'orders_normalized.csv', 'Invoices':'invoices_normalized.csv', 'Expenses':'expenses_normalized.csv',
 'Quotes':'quotes_normalized.csv', 'Summary':'project_summary.csv', 'Exceptions':'exceptions.csv'
}
for sheet, header in header_rows.items():
    ws=wb[sheet]
    actual=[ws.cell(4,c).value for c in range(1,len(header)+1)]
    if actual != header:
        errors.append(f'xlsx {sheet} header mismatch: {actual}')
    expected_rows=len(rows_by_name[row_source[sheet]])
    actual_data_rows=max(0, ws.max_row-4)
    if actual_data_rows != expected_rows:
        errors.append(f'xlsx {sheet} row count {actual_data_rows} != {expected_rows}')
    checks.append(f'xlsx {sheet} data rows {actual_data_rows}')

print('CHECKS PASSED:' if not errors else 'CHECKS FAILED:')
for c in checks:
    print('-', c)
if errors:
    for e in errors:
        print('ERROR', e)
    raise SystemExit(1)
