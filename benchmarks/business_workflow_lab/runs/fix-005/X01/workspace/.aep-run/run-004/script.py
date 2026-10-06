from pathlib import Path
import csv, json, subprocess, sys
from openpyxl import load_workbook
base=Path('.')
print('output files:', sorted(str(p) for p in (base/'output').rglob('*') if p.is_file()))
for rel in ['output/downloads/sap_orders.csv','output/downloads/sap_invoices.csv','output/expense_approved.csv','input/fx_rates.csv']:
    p=base/rel
    print('\n---', rel, 'exists=', p.exists(), '---')
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as f:
            r=csv.DictReader(f)
            print('fieldnames:', r.fieldnames)
            for i,row in zip(range(5), r): print(row)
wb=load_workbook(base/'input'/'project_template.xlsx')
print('\nWorkbook sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('Sheet', ws.title, 'size', ws.max_row, ws.max_column)
    for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row,8), values_only=True):
        print(row)
cli=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
cmd=[sys.executable, str(cli), 'quotes', 'list', '--from','2026-09-01','--to','2026-09-30','--status','Accepted','--format','json']
print('\nCLI cmd:', cmd)
cp=subprocess.run(cmd, capture_output=True, text=True)
print('returncode:', cp.returncode)
print('stderr:', cp.stderr[:500])
print('stdout:', cp.stdout[:2000])