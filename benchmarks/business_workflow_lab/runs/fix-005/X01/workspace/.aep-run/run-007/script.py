from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd(); wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for s in ['Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Report']:
    sh=wb[s]
    print('\n',s)
    for r in range(1, min(sh.max_row,8)+1):
        print(r, [sh.cell(r,c).value for c in range(1, min(sh.max_column,12)+1)])
