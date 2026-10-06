from pathlib import Path
import csv
from openpyxl import load_workbook

base=Path('.')
out=base/'output'
print('workspace', base.resolve())
print('output exists', out.exists())
print('output files:')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(p.as_posix(), p.stat().st_size)

for name in ['project_summary.csv','orders_normalized.csv','invoices_normalized.csv','expenses_normalized.csv','quotes_normalized.csv','exceptions.csv','expense_approved.csv','downloads/sap_orders.csv','downloads/sap_invoices.csv']:
    p=out/name
    print('\nCSV', name, 'exists', p.exists(), 'size', p.stat().st_size if p.exists() else None)
    if p.exists():
        with p.open(newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f)
            rows=[]
            for i,row in enumerate(r):
                rows.append(row)
                if i>=3: break
        for row in rows:
            print(row)

xlsx=out/'project_integration_report.xlsx'
print('\nXLSX exists', xlsx.exists(), 'size', xlsx.stat().st_size if xlsx.exists() else None)
if xlsx.exists():
    wb=load_workbook(xlsx, data_only=False)
    print('sheets', wb.sheetnames)
    for ws in wb.worksheets:
        print(ws.title, ws.max_row, ws.max_column)
        vals=[]
        for row in ws.iter_rows(min_row=1, max_row=min(5, ws.max_row), values_only=True):
            vals.append(row)
        print(vals)
