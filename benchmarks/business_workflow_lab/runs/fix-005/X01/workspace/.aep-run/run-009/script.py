import csv, json, os, subprocess, sys
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

params=json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
FROM=params['from_date']; TO=params['to_date']
ws=Path.cwd(); input_dir=ws/'input'; output_dir=ws/'output'; output_dir.mkdir(exist_ok=True)

def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def dec(x):
    if x is None or str(x).strip()=='': return Decimal('0')
    return Decimal(str(x).strip())

def money(x):
    return dec(x).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

# FX rates from the attachment. Some notes contain commas, so only the first two columns are contractual here.
fx={}
with open(input_dir/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    header=next(r)
    for row in r:
        if len(row)>=2 and row[0].strip():
            fx[row[0].strip()]=dec(row[1])
if not fx:
    raise SystemExit('No FX rates found')
print('FX', fx)

def to_twd(amount, currency):
    c=(currency or '').strip()
    if c not in fx:
        raise SystemExit(f'Missing FX rate for {c}')
    return money(dec(amount)*fx[c])

orders_raw=read_csv(output_dir/'downloads'/'sap_orders.csv')
invoices_raw=read_csv(output_dir/'downloads'/'sap_invoices.csv')
expenses_raw=read_csv(output_dir/'expense_approved.csv')
print('raw counts', len(orders_raw), len(invoices_raw), len(expenses_raw))

# Quote CLI: public CLI, read-only.
cli=Path('C:/Users/Chi_L/Documents/ChatGPT/agentic-engineering-platform/benchmarks/business_workflow_lab/quote_cli_mock/run_quote_cli.py')
cmd=[sys.executable, str(cli), 'quotes', 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json']
cp=subprocess.run(cmd, capture_output=True, text=True)
if cp.returncode!=0:
    raise SystemExit(f'quote CLI failed rc={cp.returncode}: {cp.stderr}')
quote_doc=json.loads(cp.stdout)
quotes_raw=quote_doc.get('items', [])
print('quotes count', len(quotes_raw))

orders=[]; invoices=[]; expenses=[]; quotes=[]; exceptions=[]
summary={}

def bucket(pid):
    if pid not in summary:
        summary[pid]={'OrderTWD':Decimal('0.00'),'InvoiceBalanceTWD':Decimal('0.00'),'ApprovedExpenseTWD':Decimal('0.00'),'AcceptedQuoteTWD':Decimal('0.00')}
    return summary[pid]

def add_exception(source, docid, pid, reason):
    exceptions.append({'Source':source,'DocumentID':docid,'ProjectID':pid or '','Reason':reason})

# Non-cancelled orders. Use ProjectID only for grouping; do not use CustomerName as a key.
for r in orders_raw:
    if r.get('Status')=='Cancelled':
        continue
    pid=(r.get('ProjectID') or '').strip()
    if not pid:
        add_exception('Orders', r.get('OrderID',''), pid, 'Missing ProjectID; excluded from project summary')
        continue
    amt_twd=to_twd(r.get('NetAmount'), r.get('Currency'))
    out={'OrderID':r.get('OrderID',''), 'ProjectID':pid, 'CustomerID':r.get('CustomerID',''), 'OrderDate':r.get('OrderDate',''), 'Status':r.get('Status',''), 'Currency':r.get('Currency',''), 'NetAmount':money(r.get('NetAmount')), 'AmountTWD':amt_twd}
    orders.append(out); bucket(pid)['OrderTWD']+=amt_twd

# Invoices: monthly invoices, summarize unpaid Balance only.
for r in invoices_raw:
    pid=(r.get('ProjectID') or '').strip()
    if not pid:
        add_exception('Invoices', r.get('InvoiceID',''), pid, 'Missing ProjectID; excluded from project summary')
        continue
    bal_twd=to_twd(r.get('Balance'), r.get('Currency'))
    out={'InvoiceID':r.get('InvoiceID',''), 'OrderID':r.get('OrderID',''), 'ProjectID':pid, 'CustomerID':r.get('CustomerID',''), 'InvoiceDate':r.get('InvoiceDate',''), 'DueDate':r.get('DueDate',''), 'Currency':r.get('Currency',''), 'Amount':money(r.get('Amount')), 'PaidAmount':money(r.get('PaidAmount')), 'Balance':money(r.get('Balance')), 'BalanceTWD':bal_twd}
    invoices.append(out); bucket(pid)['InvoiceBalanceTWD']+=bal_twd

# Approved expenses; desktop query already filters status, verify defensively.
for r in expenses_raw:
    if r.get('Status')!='Approved':
        continue
    pid=(r.get('ProjectID') or '').strip()
    if not pid:
        add_exception('Expenses', r.get('ExpenseID',''), pid, 'Missing ProjectID; excluded from project summary')
        continue
    amt_twd=to_twd(r.get('Amount'), r.get('Currency'))
    out={'ExpenseID':r.get('ExpenseID',''), 'ProjectID':pid, 'EmployeeID':r.get('EmployeeID',''), 'ExpenseDate':r.get('ExpenseDate',''), 'Status':r.get('Status',''), 'Currency':r.get('Currency',''), 'Amount':money(r.get('Amount')), 'AmountTWD':amt_twd}
    expenses.append(out); bucket(pid)['ApprovedExpenseTWD']+=amt_twd

# Accepted quotes. Keep separate from orders; do not treat orders+quotes as revenue.
for r in quotes_raw:
    if r.get('Status')!='Accepted':
        continue
    pid=(r.get('ProjectID') or '').strip()
    if not pid:
        add_exception('Quotes', r.get('QuoteID',''), pid, 'Missing ProjectID; excluded from project summary')
        continue
    amt_twd=to_twd(r.get('NetAmount'), r.get('Currency'))
    out={'QuoteID':r.get('QuoteID',''), 'ProjectID':pid, 'CustomerID':r.get('CustomerID',''), 'QuoteDate':r.get('QuoteDate',''), 'ValidUntil':r.get('ValidUntil',''), 'Status':r.get('Status',''), 'Currency':r.get('Currency',''), 'NetAmount':money(r.get('NetAmount')), 'AmountTWD':amt_twd}
    quotes.append(out); bucket(pid)['AcceptedQuoteTWD']+=amt_twd

# Write normalized audit CSV files.
def write_csv(rel, fieldnames, rows):
    p=output_dir/rel
    with open(p, 'w', newline='', encoding='utf-8-sig') as f:
        w=csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        for row in rows:
            w.writerow({k:(str(row.get(k,'')) if isinstance(row.get(k,''), Decimal) else row.get(k,'')) for k in fieldnames})
    print('wrote', rel, len(rows))

write_csv('orders_normalized.csv', ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'], orders)
write_csv('invoices_normalized.csv', ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], invoices)
write_csv('expenses_normalized.csv', ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'], expenses)
write_csv('quotes_normalized.csv', ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'], quotes)
write_csv('exceptions.csv', ['Source','DocumentID','ProjectID','Reason'], exceptions)
summary_rows=[]
for pid in sorted(summary):
    vals=summary[pid]
    summary_rows.append({'ProjectID':pid, 'OrderTWD':money(vals['OrderTWD']), 'InvoiceBalanceTWD':money(vals['InvoiceBalanceTWD']), 'ApprovedExpenseTWD':money(vals['ApprovedExpenseTWD']), 'AcceptedQuoteTWD':money(vals['AcceptedQuoteTWD'])})
write_csv('project_summary.csv', ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], summary_rows)

# Fill workbook template.
wb=load_workbook(input_dir/'project_template.xlsx')

def clear_from(sh, start_row=5):
    if sh.max_row >= start_row:
        sh.delete_rows(start_row, sh.max_row-start_row+1)

def fill_sheet(name, headers, rows):
    sh=wb[name]
    clear_from(sh, 5)
    # preserve source header row names; verify expected headers exactly.
    actual=[sh.cell(4,c).value for c in range(1,len(headers)+1)]
    if actual!=headers:
        raise SystemExit(f'{name} header mismatch: {actual} != {headers}')
    for i,row in enumerate(rows, start=5):
        for c,h in enumerate(headers, start=1):
            v=row.get(h,'')
            if isinstance(v, Decimal): v=float(v)
            sh.cell(i,c).value=v

fill_sheet('Orders', ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'], orders)
fill_sheet('Invoices', ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], invoices)
fill_sheet('Expenses', ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'], expenses)
fill_sheet('Quotes', ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'], quotes)
fill_sheet('Summary', ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], summary_rows)
fill_sheet('Exceptions', ['Source','DocumentID','ProjectID','Reason'], exceptions)

sh=wb['Summary']
last=max(5, 4+len(summary_rows))
sh['B2']=f'=SUM(B5:B{last})'; sh['C2']=f'=SUM(C5:C{last})'; sh['D2']=f'=SUM(D5:D{last})'; sh['E2']=f'=SUM(E5:E{last})'
report=wb['Report']
report['B3']=FROM; report['D3']=TO
report['B4']='SAP orders: non-Cancelled; SAP invoices: Balance only; Expense: Approved; Quotes: Accepted; grouped only by ProjectID; orders and quotes kept separate.'
report['B5']=datetime.now(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S Asia/Taipei')
report['A8']='SAP Sales Orders'; report['C8']=report['B5']; report['A9']='SAP Invoices'; report['C9']=report['B5']; report['A10']='Expense Viewer'; report['C10']=report['B5']; report['A11']='quote_cli'; report['C11']=report['B5']
out_xlsx=output_dir/'project_integration_report.xlsx'
wb.save(out_xlsx)
print('wrote', out_xlsx.relative_to(ws))

# Reopen and verify outputs from saved files.
checks=[]
wb2=load_workbook(out_xlsx, data_only=False)
checks.append(('sheets', set(['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions']).issubset(wb2.sheetnames)))
for name, headers, expected_count in [
    ('Orders',['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],len(orders)),
    ('Invoices',['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],len(invoices)),
    ('Expenses',['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],len(expenses)),
    ('Quotes',['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],len(quotes)),
    ('Summary',['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],len(summary_rows)),
    ('Exceptions',['Source','DocumentID','ProjectID','Reason'],len(exceptions)),
]:
    sh=wb2[name]
    checks.append((name+'_headers', [sh.cell(4,c).value for c in range(1,len(headers)+1)]==headers))
    actual=max(0, sh.max_row-4) if sh.max_row>=5 else 0
    checks.append((name+'_rows', actual==expected_count))
# Validate filters and separation.
checks.append(('orders_non_cancelled', all(r['Status']!='Cancelled' for r in orders)))
checks.append(('expenses_approved', all(r['Status']=='Approved' for r in expenses)))
checks.append(('quotes_accepted', all(r['Status']=='Accepted' for r in quotes)))
checks.append(('no_blank_project_in_summary', all(r['ProjectID'] for r in summary_rows)))
checks.append(('summary_separate_columns', wb2['Summary']['B4'].value=='OrderTWD' and wb2['Summary']['E4'].value=='AcceptedQuoteTWD'))
failed=[k for k,v in checks if not v]
print('CHECKS', checks)
if failed:
    raise SystemExit('Verification failed: '+', '.join(failed))
print('VERIFIED all outputs')
print('COUNTS', json.dumps({'orders':len(orders),'invoices':len(invoices),'expenses':len(expenses),'quotes':len(quotes),'summary_projects':len(summary_rows),'exceptions':len(exceptions)}, ensure_ascii=False))
