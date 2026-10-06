from pathlib import Path
import csv
from openpyxl import load_workbook

base = Path('output')
print('workspace output exists:', base.exists())
for p in sorted(base.rglob('*')):
    if p.is_file():
        print(p.as_posix(), p.stat().st_size)
        if p.suffix.lower() == '.csv':
            try:
                with p.open('r', encoding='utf-8-sig', newline='') as f:
                    rows = list(csv.reader(f))
                print('  csv rows incl header:', len(rows), 'headers:', rows[0] if rows else [])
                if len(rows) > 1:
                    print('  sample:', rows[1])
            except Exception as e:
                print('  csv read error:', repr(e))
        elif p.suffix.lower() == '.xlsx':
            try:
                wb = load_workbook(p, read_only=True, data_only=True)
                print('  xlsx sheets:', wb.sheetnames)
                for ws in wb.worksheets:
                    print('   ', ws.title, 'rows:', ws.max_row, 'cols:', ws.max_column)
                wb.close()
            except Exception as e:
                print('  xlsx read error:', repr(e))