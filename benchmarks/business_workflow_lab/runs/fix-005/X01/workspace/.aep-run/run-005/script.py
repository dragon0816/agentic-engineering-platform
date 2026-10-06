from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws=Path.cwd()
print('WORKSPACE', ws)
for p in sorted((ws/'input').glob('*')):
    print('INPUT', p.name, p.stat().st_size)
print('OUTPUT FILES')
for p in sorted((ws/'output').rglob('*')):
    if p.is_file(): print(str(p.relative_to(ws)), p.stat().st_size)
for fn in ['output/downloads/sap_orders.csv','output/downloads/sap_invoices.csv','output/expense_approved.csv','input/fx_rates.csv']:
    p=ws/fn
    print('\nFILE', fn, 'exists=', p.exists())
    if p.exists() and p.suffix.lower()=='.csv':
        with p.open(newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f); rows=[]
            for i,row in zip(range(4),r): rows.append(row)
        print('sample', rows)
wb=load_workbook(ws/'input'/'project_template.xlsx')
print('\nTEMPLATE SHEETS', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('SHEET', s, 'max', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(3,sh.max_row), values_only=True): print(' ', row)
