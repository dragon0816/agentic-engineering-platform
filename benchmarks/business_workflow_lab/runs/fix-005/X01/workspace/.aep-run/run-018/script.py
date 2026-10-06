import csv, json, os, subprocess, sys
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime, timezone, timedelta
from openpyxl import load_workbook

params=json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
FROM=params['from_date']
TO=params['to_date']
ws=Path.cwd()
input_dir=ws/'input'
output_dir=ws/'output'
downloads=output_dir/'downloads'
output_dir.mkdir(exist_ok=True)
D2=Decimal('0.01')
TW=timezone(timedelta(hours=8))

def dec(x):
    if x is None or str(x).strip()=='':
        return Decimal('0')
    return Decimal(str(x).replace(',', '').strip())

def money(x):
    return dec(x).quantize(D2, rounding=ROUND_HALF_UP)

def fmt(x):
    return format(money(x), 'f')

def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def write_csv(name, fields, rows):
    p=output_dir/name
    with p.open('w', newline='', encoding='utf-8-sig') as f:
        w=csv.DictWriter(f, fieldnames=fields)
        w.writeheader()
        for row in rows:
            w.writerow({k: row.get(k, '') for k in fields})
    print('wrote', p.relative_to(ws), len(rows))

# FX rates from input/fx_rates.csv. csv.DictReader preserves the named fields even if the note column contains commas.
fx={}
for r in read_csv(input_dir/'fx_rates.csv'):
    cur=(r.get('Currency') or '').strip()
    rate=(r.get('RateToTWD') or '').strip()
    if cur and rate:
        fx[cur]=dec(rate)
print('FX', fx)

def twd(amount, currency):
    cur=(currency or '').strip()
    if cur not in fx:
        raise ValueError(f'Missing FX rate for {cur}')
    return money(dec(amount)*fx[cur])

orders_raw=read_csv(downloads/'sap_orders.csv')
invoices_raw=read_csv(downloads/'sap_invoices.csv')
expenses_raw=read_csv(output_dir/'expense_approved.csv')
print('raw counts', len(orders_raw), len(invoices_raw), len(expenses_raw))

quote_script=Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\run_quote_cli.py')
cmd=[sys.executable, str(quote_script), 'quotes', 'list', '--from', FROM, '--to', TO, '--status', 'Accepted', '--format', 'json']
cp=subprocess.run(cmd, text=True, capture_output=True)
if cp.returncode!=0:
    raise RuntimeError('quote_cli failed: '+cp.stderr)
quote_payload=json.loads(cp.stdout)
quotes_raw=quote_payload.get('items', [])
print('quotes count', len(quotes_raw))

exceptions=[]
orders=[]
for r in orders_raw:
    status=(r.get('Status') or '').strip()
    oid=(r.get('OrderID') or '').strip()
    pid=(r.get('ProjectID') or '').strip()
    if status=='Cancelled':
        exceptions.append({'Source':'Orders','DocumentID':oid,'ProjectID':pid,'Reason':'Cancelled order; excluded from project summary'})
        continue
    if not pid:
        exceptions.append({'Source':'Orders','DocumentID':oid,'ProjectID':pid,'Reason':'Missing ProjectID; excluded from project summary'})
        continue
    orders.append({'OrderID':oid,'ProjectID':pid,'CustomerID':(r.get('CustomerID') or '').strip(),'OrderDate':(r.get('OrderDate') or '').strip(),'Status':status,'Currency':(r.get('Currency') or '').strip(),'NetAmount':fmt(r.get('NetAmount')),'AmountTWD':fmt(twd(r.get('NetAmount'), r.get('Currency')))})

invoices=[]
for r in invoices_raw:
    iid=(r.get('InvoiceID') or '').strip()
    pid=(r.get('ProjectID') or '').strip()
    if not pid:
        exceptions.append({'Source':'Invoices','DocumentID':iid,'ProjectID':pid,'Reason':'Missing ProjectID; excluded from project summary'})
        continue
    invoices.append({'InvoiceID':iid,'OrderID':(r.get('OrderID') or '').strip(),'ProjectID':pid,'CustomerID':(r.get('CustomerID') or '').strip(),'InvoiceDate':(r.get('InvoiceDate') or '').strip(),'DueDate':(r.get('DueDate') or '').strip(),'Currency':(r.get('Currency') or '').strip(),'Amount':fmt(r.get('Amount')),'PaidAmount':fmt(r.get('PaidAmount')),'Balance':fmt(r.get('Balance')),'BalanceTWD':fmt(twd(r.get('Balance'), r.get('Currency')))})

expenses=[]
for r in expenses_raw:
    eid=(r.get('ExpenseID') or '').strip()
    pid=(r.get('ProjectID') or '').strip()
    status=(r.get('Status') or '').strip()
    if status!='Approved':
        exceptions.append({'Source':'Expenses','DocumentID':eid,'ProjectID':pid,'Reason':'Not Approved; excluded from project summary'})
        continue
    if not pid:
        exceptions.append({'Source':'Expenses','DocumentID':eid,'ProjectID':pid,'Reason':'Missing ProjectID; excluded from project summary'})
        continue
    expenses.append({'ExpenseID':eid,'ProjectID':pid,'EmployeeID':(r.get('EmployeeID') or '').strip(),'ExpenseDate':(r.get('ExpenseDate') or '').strip(),'Status':status,'Currency':(r.get('Currency') or '').strip(),'Amount':fmt(r.get('Amount')),'AmountTWD':fmt(twd(r.get('Amount'), r.get('Currency')))})

quotes=[]
for r in quotes_raw:
    qid=(r.get('QuoteID') or '').strip()
    pid=(r.get('ProjectID') or '').strip()
    status=(r.get('Status') or '').strip()
    if status!='Accepted':
        exceptions.append({'Source':'Quotes','DocumentID':qid,'ProjectID':pid,'Reason':'Not Accepted; excluded from project summary'})
        continue
    if not pid:
        exceptions.append({'Source':'Quotes','DocumentID':qid,'ProjectID':pid,'Reason':'Missing ProjectID; excluded from project summary'})
        continue
    quotes.append({'QuoteID':qid,'ProjectID':pid,'CustomerID':(r.get('CustomerID') or '').strip(),'QuoteDate':(r.get('QuoteDate') or '').strip(),'ValidUntil':(r.get('ValidUntil') or '').strip(),'Status':status,'Currency':(r.get('Currency') or '').strip(),'NetAmount':fmt(r.get('NetAmount')),'AmountTWD':fmt(twd(r.get('NetAmount'), r.get('Currency')))})

project_ids=sorted({r['ProjectID'] for r in orders+invoices+expenses+quotes if r.get('ProjectID')})

def sum_field(rows, field, pid):
    total=Decimal('0')
    for r in rows:
        if r.get('ProjectID')==pid:
            total += dec(r.get(field))
    return money(total)

summary_rows=[]
for pid in project_ids:
    summary_rows.append({'ProjectID':pid,'OrderTWD':fmt(sum_field(orders,'AmountTWD',pid)),'InvoiceBalanceTWD':fmt(sum_field(invoices,'BalanceTWD',pid)),'ApprovedExpenseTWD':fmt(sum_field(expenses,'AmountTWD',pid)),'AcceptedQuoteTWD':fmt(sum_field(quotes,'AmountTWD',pid))})

write_csv('orders_normalized.csv', ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'], orders)
write_csv('invoices_normalized.csv', ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], invoices)
write_csv('expenses_normalized.csv', ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'], expenses)
write_csv('quotes_normalized.csv', ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'], quotes)
write_csv('exceptions.csv', ['Source','DocumentID','ProjectID','Reason'], exceptions)
write_csv('project_summary.csv', ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], summary_rows)

wb=load_workbook(input_dir/'project_template.xlsx')

def fill_sheet(sheet_name, headers, rows):
    sh=wb[sheet_name]
    if sh.max_row>=5:
        sh.delete_rows(5, sh.max_row-4)
    for c,h in enumerate(headers,1):
        sh.cell(4,c).value=h
    for r_idx,row in enumerate(rows,5):
        for c_idx,h in enumerate(headers,1):
            val=row.get(h, '')
            if h.endswith('TWD') or h in ['NetAmount','Amount','PaidAmount','Balance']:
                try:
                    sh.cell(r_idx,c_idx).value=float(dec(val))
                    continue
                except Exception:
                    pass
            sh.cell(r_idx,c_idx).value=val

fill_sheet('Orders', ['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'], orders)
fill_sheet('Invoices', ['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'], invoices)
fill_sheet('Expenses', ['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'], expenses)
fill_sheet('Quotes', ['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'], quotes)
fill_sheet('Summary', ['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], summary_rows)
fill_sheet('Exceptions', ['Source','DocumentID','ProjectID','Reason'], exceptions)

sh=wb['Summary']
last=max(5, 4+len(summary_rows))
sh['B2']=f'=SUM(B5:B{last})'
sh['C2']=f'=SUM(C5:C{last})'
sh['D2']=f'=SUM(D5:D{last})'
sh['E2']=f'=SUM(E5:E{last})'
report=wb['Report']
run_time=datetime.now(TW).strftime('%Y-%m-%d %H:%M:%S Asia/Taipei')
report['B3']=FROM
report['D3']=TO
report['B4']='SAP orders: non-Cancelled; SAP invoices: Balance only; Expense: Approved; Quotes: Accepted; grouped only by ProjectID; orders and quotes kept separate.'
report['B5']=run_time
report['A8']='SAP Sales Orders'; report['C8']=run_time
report['A9']='SAP Invoices'; report['C9']=run_time
report['A10']='Expense Viewer'; report['C10']=run_time
report['A11']='quote_cli'; report['C11']=run_time
out_xlsx=output_dir/'project_integration_report.xlsx'
wb.save(out_xlsx)
print('wrote', out_xlsx.relative_to(ws))

# Reopen every output and verify the stated requirements.
for rel in ['project_integration_report.xlsx','project_summary.csv','orders_normalized.csv','invoices_normalized.csv','expenses_normalized.csv','quotes_normalized.csv','exceptions.csv','expense_approved.csv','downloads/sap_orders.csv','downloads/sap_invoices.csv']:
    p=output_dir/rel
    if not p.exists() or p.stat().st_size<=0:
        raise SystemExit(f'Missing or empty output/{rel}')

wb2=load_workbook(out_xlsx, data_only=False)
checks=[]
checks.append(('sheets', set(['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions']).issubset(wb2.sheetnames)))
for name, headers, expected_count in [
    ('Orders',['OrderID','ProjectID','CustomerID','OrderDate','Status','Currency','NetAmount','AmountTWD'],len(orders)),
    ('Invoices',['InvoiceID','OrderID','ProjectID','CustomerID','InvoiceDate','DueDate','Currency','Amount','PaidAmount','Balance','BalanceTWD'],len(invoices)),
    ('Expenses',['ExpenseID','ProjectID','EmployeeID','ExpenseDate','Status','Currency','Amount','AmountTWD'],len(expenses)),
    ('Quotes',['QuoteID','ProjectID','CustomerID','QuoteDate','ValidUntil','Status','Currency','NetAmount','AmountTWD'],len(quotes)),
    ('Summary',['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'],len(summary_rows)),
    ('Exceptions',['Source','DocumentID','ProjectID','Reason'],len(exceptions)),
]:
    sheet=wb2[name]
    checks.append((name+'_headers', [sheet.cell(4,c).value for c in range(1,len(headers)+1)]==headers))
    actual=max(0, sheet.max_row-4) if sheet.max_row>=5 else 0
    checks.append((name+'_rows', actual==expected_count))

def reread(name):
    return read_csv(output_dir/name)
orders_csv=reread('orders_normalized.csv')
invoices_csv=reread('invoices_normalized.csv')
expenses_csv=reread('expenses_normalized.csv')
quotes_csv=reread('quotes_normalized.csv')
summary_csv=reread('project_summary.csv')
checks.append(('orders_non_cancelled', all(r['Status']!='Cancelled' for r in orders_csv)))
checks.append(('expenses_approved', all(r['Status']=='Approved' for r in expenses_csv)))
checks.append(('quotes_accepted', all(r['Status']=='Accepted' for r in quotes_csv)))
checks.append(('no_blank_project_in_summary', all(r['ProjectID'] for r in summary_csv)))
checks.append(('summary_separate_columns_csv', list(summary_csv[0].keys())==['ProjectID','OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'] if summary_csv else True))
checks.append(('summary_separate_columns_xlsx', wb2['Summary']['B4'].value=='OrderTWD' and wb2['Summary']['E4'].value=='AcceptedQuoteTWD'))
# Check money columns are parseable as numbers in CSV.
for filename, money_cols in [('orders_normalized.csv',['NetAmount','AmountTWD']),('invoices_normalized.csv',['Amount','PaidAmount','Balance','BalanceTWD']),('expenses_normalized.csv',['Amount','AmountTWD']),('quotes_normalized.csv',['NetAmount','AmountTWD']),('project_summary.csv',['OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'])]:
    for row in reread(filename):
        for col in money_cols:
            dec(row[col])
failed=[name for name, ok in checks if not ok]
print('CHECKS', checks)
if failed:
    raise SystemExit('Verification failed: '+', '.join(failed))
print('VERIFIED all outputs')
print('COUNTS', json.dumps({'orders':len(orders),'invoices':len(invoices),'expenses':len(expenses),'quotes':len(quotes),'summary_projects':len(summary_rows),'exceptions':len(exceptions)}, ensure_ascii=False))
