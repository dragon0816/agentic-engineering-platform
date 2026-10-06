from pathlib import Path
import csv
from datetime import date
from decimal import Decimal
from openpyxl import load_workbook

out = Path('output')
required = [
    'project_integration_report.xlsx',
    'project_summary.csv',
    'orders_normalized.csv',
    'invoices_normalized.csv',
    'expenses_normalized.csv',
    'quotes_normalized.csv',
    'exceptions.csv',
    'expense_approved.csv',
    'downloads/sap_orders.csv',
    'downloads/sap_invoices.csv',
]
start = date.fromisoformat('2026-08-01')
end = date.fromisoformat('2026-08-31')

def read_csv(rel):
    p = out / rel
    with p.open('r', encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def check(cond, msg):
    if not cond:
        raise AssertionError(msg)
    print('OK:', msg)

for rel in required:
    p = out / rel
    check(p.exists(), f'{rel} exists')
    check(p.stat().st_size > 0, f'{rel} is non-empty')

orders = read_csv('orders_normalized.csv')
invoices = read_csv('invoices_normalized.csv')
expenses = read_csv('expenses_normalized.csv')
quotes = read_csv('quotes_normalized.csv')
summary = read_csv('project_summary.csv')
exceptions = read_csv('exceptions.csv')
raw_orders = read_csv('downloads/sap_orders.csv')
raw_invoices = read_csv('downloads/sap_invoices.csv')
expense_raw = read_csv('expense_approved.csv')

check(len(raw_orders) == 8, 'raw SAP orders has 8 data rows for Aug')
check(len(raw_invoices) == 6, 'raw SAP invoices has 6 data rows for Aug')
check(len(expense_raw) == 4, 'exported approved expenses has 4 data rows for Aug')
check(len(orders) == 5, 'normalized orders has 5 data rows')
check(len(invoices) == 5, 'normalized invoices has 5 data rows')
check(len(expenses) == 4, 'normalized expenses has 4 data rows')
check(len(quotes) == 4, 'normalized accepted quotes has 4 data rows')
check(len(summary) == 5, 'project summary has 5 data rows')
check(len(exceptions) == 4, 'exceptions has 4 data rows')

for r in raw_orders:
    d = date.fromisoformat(r['OrderDate'])
    check(start <= d <= end, f"raw order {r['OrderID']} date in Aug")
for r in raw_invoices:
    d = date.fromisoformat(r['InvoiceDate'])
    check(start <= d <= end, f"raw invoice {r['InvoiceID']} date in Aug")
for r in expense_raw:
    d = date.fromisoformat(r['ExpenseDate'])
    check(start <= d <= end and r['Status'] == 'Approved', f"expense {r['ExpenseID']} approved and date in Aug")

check(all(r['Status'] != 'Cancelled' for r in orders), 'normalized orders exclude Cancelled')
check(all(r['Status'] == 'Approved' for r in expenses), 'normalized expenses are Approved')
check(all(r['Status'] == 'Accepted' for r in quotes), 'normalized quotes are Accepted')
check(all(r['ProjectID'] for r in summary), 'summary has no blank ProjectID')

for rows, cols, name in [
    (orders, ['NetAmount','AmountTWD'], 'orders'),
    (invoices, ['Amount','PaidAmount','Balance','BalanceTWD'], 'invoices'),
    (expenses, ['Amount','AmountTWD'], 'expenses'),
    (quotes, ['NetAmount','AmountTWD'], 'quotes'),
    (summary, ['OrderTWD','InvoiceBalanceTWD','ApprovedExpenseTWD','AcceptedQuoteTWD'], 'summary'),
]:
    for r in rows:
        for c in cols:
            Decimal(r[c])
    check(True, f'{name} numeric amount columns parse as Decimal')

xlsx = out / 'project_integration_report.xlsx'
wb = load_workbook(xlsx, read_only=True, data_only=True)
expected_sheets = ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']
check(wb.sheetnames == expected_sheets, 'Excel sheets match expected order')
for s in expected_sheets:
    ws = wb[s]
    check(ws.max_row >= 1 and ws.max_column >= 1, f'Excel sheet {s} is readable')
wb.close()

print('SUMMARY_COUNTS', {
    'sap_orders': len(raw_orders),
    'sap_invoices': len(raw_invoices),
    'expense_approved': len(expense_raw),
    'orders_normalized': len(orders),
    'invoices_normalized': len(invoices),
    'expenses_normalized': len(expenses),
    'quotes_normalized': len(quotes),
    'project_summary': len(summary),
    'exceptions': len(exceptions),
})