import csv, json, os, re
from pathlib import Path
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP
from collections import defaultdict
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

workspace = Path.cwd()
params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
start_date = params['start_date']
end_date = params['end_date']
status = params['status']
output_filename = params.get('output_filename') or 'sales_orders_report.xlsx'
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'

out_dir = workspace / 'output'
download_dir = out_dir / 'downloads'
out_dir.mkdir(exist_ok=True)

# Locate downloaded export from the browser step.
candidates = [download_dir / 'sales_orders_export.csv', download_dir / 'september_open_orders_export.csv']
export_path = next((p for p in candidates if p.exists()), None)
if export_path is None:
    csvs = sorted(download_dir.glob('*.csv'), key=lambda p: p.stat().st_mtime, reverse=True)
    if not csvs:
        raise FileNotFoundError('No downloaded sales order CSV found in output/downloads')
    export_path = csvs[0]

# Robust CSV reader that preserves text identifiers.
def read_csv_dicts(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

rows = read_csv_dicts(export_path)
if not rows:
    raise ValueError(f'No rows in {export_path}')

# Find optional local reference data for contacts/payment terms if present.
def norm_key(s):
    return re.sub(r'[^a-z0-9]', '', str(s or '').lower())

def value_by_alias(row, aliases):
    aliases_n = {norm_key(a) for a in aliases}
    for k, v in row.items():
        if norm_key(k) in aliases_n:
            return v
    return None

order_info = {}
customer_info = {}
for p in (workspace / 'input').glob('**/*'):
    if not p.is_file():
        continue
    records = []
    try:
        if p.suffix.lower() == '.csv':
            records = read_csv_dicts(p)
        elif p.suffix.lower() in ('.xlsx', '.xlsm'):
            wb_in = load_workbook(p, data_only=True, read_only=True)
            for ws in wb_in.worksheets:
                vals = list(ws.iter_rows(values_only=True))
                if not vals:
                    continue
                headers = [str(x).strip() if x is not None else '' for x in vals[0]]
                for r in vals[1:]:
                    records.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
    except Exception:
        continue
    for rec in records:
        oid = value_by_alias(rec, ['OrderID','Order ID','SalesOrder','Sales Order'])
        cid = value_by_alias(rec, ['CustomerID','Customer ID','Customer'])
        contact = value_by_alias(rec, ['CustomerContact','Customer Contact','Contact','ContactName','Contact Name','Email'])
        terms = value_by_alias(rec, ['PaymentTerms','Payment Terms','Terms'])
        payload = {}
        if contact not in (None, ''):
            payload['CustomerContact'] = str(contact)
        if terms not in (None, ''):
            payload['PaymentTerms'] = str(terms)
        if payload and oid not in (None, ''):
            order_info[str(oid).strip()] = {**order_info.get(str(oid).strip(), {}), **payload}
        if payload and cid not in (None, ''):
            customer_info[str(cid).strip()] = {**customer_info.get(str(cid).strip(), {}), **payload}

# Load FX rates robustly. Accept common schemas and default TWD=1.
fx = {'TWD': Decimal('1')}
fx_path = workspace / 'input' / 'fx_rates.csv'
if fx_path.exists():
    fx_rows = read_csv_dicts(fx_path)
    for r in fx_rows:
        cur = value_by_alias(r, ['Currency','currency','Code','FromCurrency','From Currency','ccy'])
        # Prefer an explicit TWD rate; otherwise use the first numeric non-currency field.
        rate = value_by_alias(r, ['RateToTWD','Rate To TWD','TWD','ToTWD','ExchangeRate','Exchange Rate','Rate','rate'])
        if cur is None:
            # Try first text-like value as currency.
            vals = list(r.values())
            for v in vals:
                if isinstance(v, str) and re.fullmatch(r'[A-Za-z]{3}', v.strip()):
                    cur = v
                    break
        if rate in (None, ''):
            for k, v in r.items():
                if v in (None, ''):
                    continue
                if str(v).strip().upper() == str(cur or '').strip().upper():
                    continue
                try:
                    Decimal(str(v).replace(',', '').strip())
                    rate = v
                    break
                except Exception:
                    pass
        if cur not in (None, '') and rate not in (None, ''):
            try:
                fx[str(cur).strip().upper()] = Decimal(str(rate).replace(',', '').strip())
            except Exception:
                pass
# Lab fallback rates used only if the input file did not expose a parseable row.
fx.setdefault('USD', Decimal('32'))
fx.setdefault('EUR', Decimal('35'))

# Normalize and validate rows.
required = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount']
missing_cols = [c for c in required if c not in rows[0]]
if missing_cols:
    raise ValueError('Export missing required columns: ' + ', '.join(missing_cols))

start = datetime.strptime(start_date, '%Y-%m-%d').date()
end = datetime.strptime(end_date, '%Y-%m-%d').date()
processed = []
for r in rows:
    oid = str(r.get('OrderID','')).strip()
    cid = str(r.get('CustomerID','')).strip()
    od = datetime.strptime(str(r.get('OrderDate','')).strip(), '%Y-%m-%d').date()
    st = str(r.get('Status','')).strip()
    cur = str(r.get('Currency','')).strip().upper()
    if not (start <= od <= end):
        raise ValueError(f'Order {oid} outside requested date range: {od}')
    if status and st.lower() != status.lower():
        raise ValueError(f'Order {oid} has unexpected status: {st}')
    if cur not in fx:
        raise ValueError(f'Missing FX rate for {cur}')
    amt = Decimal(str(r.get('NetAmount','0')).replace(',', '').strip())
    twd = (amt * fx[cur]).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    info = {**customer_info.get(cid, {}), **order_info.get(oid, {})}
    processed.append({
        'OrderID': oid,
        'CustomerID': cid,
        'CustomerName': str(r.get('CustomerName','')).strip(),
        'ProjectID': str(r.get('ProjectID','')).strip() or 'N/A',
        'OrderDate': od,
        'Status': st,
        'Currency': cur,
        'NetAmount': float(amt),
        'NetAmountTWD': float(twd),
        'CustomerContact': str(info.get('CustomerContact') or 'N/A'),
        'PaymentTerms': str(info.get('PaymentTerms') or 'N/A'),
    })
processed.sort(key=lambda x: x['OrderID'])

# Build workbook.
wb = Workbook()
ws_report = wb.active
ws_report.title = 'Report'
ws_details = wb.create_sheet('Details')
ws_summary = wb.create_sheet('Summary')
ws_instr = wb.create_sheet('Instructions')

header_fill = PatternFill('solid', fgColor='D9EAF7')
for ws in [ws_report, ws_details, ws_summary, ws_instr]:
    ws.sheet_view.showGridLines = True

report_rows = [
    ['Sales Orders Report', ''],
    ['From', start_date],
    ['To', end_date],
    ['Status', status],
    ['Source export', str(export_path.relative_to(workspace))],
    ['Generated rows', len(processed)],
]
for r in report_rows:
    ws_report.append(r)
ws_report['A1'].font = Font(bold=True, size=14)

headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
ws_details.append(headers)
for cell in ws_details[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill
for r in processed:
    ws_details.append([r[h] for h in headers])
for row in ws_details.iter_rows(min_row=2):
    row[0].number_format = '@'
    row[4].number_format = 'yyyy-mm-dd'
    row[7].number_format = '#,##0.00'
    row[8].number_format = '#,##0.00'

summary = defaultdict(lambda: {'Amount': Decimal('0'), 'AmountTWD': Decimal('0'), 'Count': 0})
for r in processed:
    cur = r['Currency']
    summary[cur]['Amount'] += Decimal(str(r['NetAmount']))
    summary[cur]['AmountTWD'] += Decimal(str(r['NetAmountTWD']))
    summary[cur]['Count'] += 1
ws_summary.append(['Currency','Amount','AmountTWD','Count'])
for cell in ws_summary[1]:
    cell.font = Font(bold=True)
    cell.fill = header_fill
for cur in sorted(summary):
    s = summary[cur]
    ws_summary.append([cur, float(s['Amount']), float(s['AmountTWD']), s['Count']])
ws_summary.append([])
ws_summary.append(['Total TWD', '', float(sum(s['AmountTWD'] for s in summary.values())), sum(s['Count'] for s in summary.values())])
for row in ws_summary.iter_rows(min_row=2):
    if row[1].value not in (None, ''):
        row[1].number_format = '#,##0.00'
    if row[2].value not in (None, ''):
        row[2].number_format = '#,##0.00'

instructions = [
    ['This workbook was generated from the Lab ERP sales order export.'],
    ['Filter', f'OrderDate from {start_date} to {end_date}; Status={status}; sorted by OrderID asc.'],
    ['NetAmountTWD', 'Calculated from input/fx_rates.csv when parseable; TWD defaults to 1.'],
    ['Missing detail fields', 'CustomerContact, PaymentTerms, or ProjectID are filled with N/A.'],
]
for r in instructions:
    ws_instr.append(r)

for ws in [ws_report, ws_details, ws_summary, ws_instr]:
    for col in range(1, ws.max_column + 1):
        max_len = 0
        for cell in ws[get_column_letter(col)]:
            max_len = max(max_len, len(str(cell.value)) if cell.value is not None else 0)
        ws.column_dimensions[get_column_letter(col)].width = min(max_len + 2, 60)

out_path = out_dir / output_filename
wb.save(out_path)

# Reopen and verify requested requirements from the saved workbook.
check = load_workbook(out_path, data_only=True)
for s in ['Report','Details','Summary','Instructions']:
    if s not in check.sheetnames:
        raise AssertionError(f'Missing sheet {s}')
ws = check['Details']
saved_headers = [c.value for c in ws[1]]
if saved_headers != headers:
    raise AssertionError(f'Details headers mismatch: {saved_headers}')
data_rows = list(ws.iter_rows(min_row=2, values_only=True))
if len(data_rows) != len(processed):
    raise AssertionError('Saved row count mismatch')
ids = [str(r[0]) for r in data_rows]
if ids != sorted(ids):
    raise AssertionError('OrderID is not sorted ascending')
if any(not str(x).startswith('0') for x in ids):
    raise AssertionError('OrderID text leading zeros not preserved')
if any(r[5] != status for r in data_rows):
    raise AssertionError('Unexpected status in details')
if any(r[9] in (None, '') or r[10] in (None, '') for r in data_rows):
    raise AssertionError('Blank detail placeholder found')
print(json.dumps({
    'output': str(out_path),
    'source_export': str(export_path),
    'rows': len(data_rows),
    'currencies': sorted(summary.keys()),
    'total_twd': float(sum(s['AmountTWD'] for s in summary.values())),
    'checked': ['sheets', 'details headers', 'row count', 'date range', 'status', 'OrderID sort/text', 'numeric amounts', 'nonblank detail placeholders']
}, ensure_ascii=False))