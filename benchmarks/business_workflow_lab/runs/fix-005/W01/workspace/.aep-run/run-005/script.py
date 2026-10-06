from pathlib import Path
from openpyxl import load_workbook
import csv, math

ws = Path.cwd()
print('workspace=', ws)
print('input files=', sorted(str(p.relative_to(ws)) for p in (ws/'input').glob('*')))
print('output files=', sorted(str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()))

out = ws/'output'/'september_open_orders.xlsx'
if not out.exists():
    raise SystemExit('MISSING output/september_open_orders.xlsx')

wb = load_workbook(out, data_only=True)
print('sheets=', wb.sheetnames)
required_sheets = ['Report', 'Details', 'Summary']
for sh in required_sheets:
    if sh not in wb.sheetnames:
        raise SystemExit(f'MISSING sheet {sh}')

# Report sheet checks
report = wb['Report']
print('Report B3=', report['B3'].value, 'D3=', report['D3'].value)
print('Report B4=', report['B4'].value)
assert str(report['B3'].value) == '2026-09-01', 'Report from date mismatch'
assert str(report['D3'].value) == '2026-09-30', 'Report to date mismatch'
assert 'Open' in str(report['B4'].value), 'Report condition missing Open'

# Details uses template header row 4 and data from row 5
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
d = wb['Details']
headers = [d.cell(row=4, column=i).value for i in range(1, len(expected_headers)+1)]
print('details headers row 4=', headers)
if headers != expected_headers:
    raise SystemExit('Details headers mismatch')
rows = []
for row_idx in range(5, d.max_row + 1):
    vals = [d.cell(row=row_idx, column=i).value for i in range(1, len(expected_headers)+1)]
    if any(v is not None for v in vals):
        rows.append(dict(zip(expected_headers, vals)))
print('details row count=', len(rows))
print('order ids=', [r['OrderID'] for r in rows])
print('first 3 rows=', rows[:3])
print('last 3 rows=', rows[-3:])

if len(rows) != 14:
    raise SystemExit(f'Expected 14 rows, got {len(rows)}')
order_ids = [str(r['OrderID']) for r in rows]
if order_ids != sorted(order_ids):
    raise SystemExit('OrderID not sorted ascending')
if not all(len(str(r['OrderID'])) == 10 and str(r['OrderID']).startswith('0') for r in rows):
    raise SystemExit('OrderID text/leading zero check failed')
if not all(r['Status'] == 'Open' for r in rows):
    raise SystemExit('Non-Open status found')
if not all(str(r['OrderDate'])[:7] == '2026-09' for r in rows):
    raise SystemExit('OrderDate outside 2026-09 found')
if not all(isinstance(r['NetAmount'], (int, float)) for r in rows):
    raise SystemExit('NetAmount not numeric')
if not all(isinstance(r['NetAmountTWD'], (int, float)) for r in rows):
    raise SystemExit('NetAmountTWD not numeric')
if not all(r['CustomerContact'] not in (None, '') for r in rows):
    raise SystemExit('CustomerContact missing')
if not all(r['PaymentTerms'] not in (None, '') for r in rows):
    raise SystemExit('PaymentTerms missing')

# Read FX rates and validate conversions
rates = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    print('fx headers=', reader.fieldnames)
    for rec in reader:
        cur = rec.get('Currency') or rec.get('currency')
        rate = rec.get('TWD') or rec.get('RateToTWD') or rec.get('Rate') or rec.get('rate')
        if cur and rate:
            rates[cur] = float(rate)
print('fx rates=', rates)

currency_amount = {}
currency_twd = {}
currency_count = {}
for r in rows:
    cur = r['Currency']
    if cur not in rates:
        raise SystemExit(f'Missing FX rate for {cur}')
    amt = float(r['NetAmount'])
    expected_twd = round(amt * rates[cur], 2)
    got_twd = round(float(r['NetAmountTWD']), 2)
    if got_twd != expected_twd:
        raise SystemExit(f'TWD mismatch {r["OrderID"]}: got {got_twd}, expected {expected_twd}')
    currency_amount[cur] = currency_amount.get(cur, 0.0) + amt
    currency_twd[cur] = currency_twd.get(cur, 0.0) + got_twd
    currency_count[cur] = currency_count.get(cur, 0) + 1
currency_amount = {k: round(v, 2) for k, v in sorted(currency_amount.items())}
currency_twd = {k: round(v, 2) for k, v in sorted(currency_twd.items())}
print('recomputed currency amount=', currency_amount)
print('recomputed currency twd=', currency_twd)
print('recomputed currency count=', dict(sorted(currency_count.items())))
recomputed_total_twd = round(sum(currency_twd.values()), 2)
print('recomputed total twd=', recomputed_total_twd)

# Summary header row 4 and values row 5+
s = wb['Summary']
summary_total = round(float(s['B2'].value), 2)
if summary_total != recomputed_total_twd:
    raise SystemExit(f'Summary B2 mismatch: {summary_total} vs {recomputed_total_twd}')
summary_headers = [s.cell(row=4, column=i).value for i in range(1, 7)]
print('summary headers row 4=', summary_headers)
expected_summary_headers = ['Group', 'SubGroup', 'Currency', 'Amount', 'AmountTWD', 'Count']
if summary_headers != expected_summary_headers:
    raise SystemExit('Summary headers mismatch')
summary_rows = []
for row_idx in range(5, s.max_row + 1):
    vals = [s.cell(row=row_idx, column=i).value for i in range(1, 7)]
    if any(v is not None for v in vals):
        summary_rows.append(dict(zip(expected_summary_headers, vals)))
print('summary rows=', summary_rows)
summary_by_cur = {r['Currency']: r for r in summary_rows if r.get('Group') == 'Status' and r.get('SubGroup') == 'Open'}
for cur in sorted(currency_amount):
    r = summary_by_cur.get(cur)
    if not r:
        raise SystemExit(f'Missing summary row for {cur}')
    if round(float(r['Amount']), 2) != currency_amount[cur]:
        raise SystemExit(f'Summary Amount mismatch for {cur}')
    if round(float(r['AmountTWD']), 2) != currency_twd[cur]:
        raise SystemExit(f'Summary AmountTWD mismatch for {cur}')
    if int(r['Count']) != currency_count[cur]:
        raise SystemExit(f'Summary Count mismatch for {cur}')

# Check export exists and row count if present
export = ws/'output'/'downloads'/'september_open_orders_export.csv'
if export.exists():
    with open(export, newline='', encoding='utf-8-sig') as f:
        export_rows = list(csv.DictReader(f))
    print('download export rows=', len(export_rows))
    print('download export headers=', list(export_rows[0].keys()) if export_rows else [])
    if len(export_rows) != 14:
        raise SystemExit(f'Expected 14 export rows, got {len(export_rows)}')

print('CHECKS PASSED')