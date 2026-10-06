from pathlib import Path
import csv, json
from openpyxl import load_workbook
base=Path('.')
print('Workspace:', base.resolve())
print('Input files:')
for p in sorted((base/'input').glob('*')):
    print('-', p, 'size=', p.stat().st_size)

# Inspect fx_rates.csv
fx=base/'input'/'fx_rates.csv'
print('\nfx_rates.csv:')
with fx.open(newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    rows=[]
    for i,row in zip(range(10), r):
        rows.append(row)
    for row in rows:
        print(row)

# Inspect workbook sheets/headers/sample dimensions
wb_path=base/'input'/'sales_template.xlsx'
wb=load_workbook(wb_path)
print('\nsales_template.xlsx sheets:', wb.sheetnames)
for ws in wb.worksheets:
    print('Sheet', ws.title, 'max_row', ws.max_row, 'max_col', ws.max_column)
    for rr in range(1, min(ws.max_row, 8)+1):
        vals=[ws.cell(rr,c).value for c in range(1, min(ws.max_column, 12)+1)]
        print(rr, vals)

# Print manual beginning and lines with url/pagination/export/date/status keywords
man=base/'input'/'sap_manual.md'
text=man.read_text(encoding='utf-8')
print('\nsap_manual.md first 4000 chars:\n', text[:4000])
print('\nKeyword lines:')
for line in text.splitlines():
    low=line.lower()
    if any(k in low for k in ['http', 'url', 'page', 'next', 'status', 'date', 'order', 'export', 'sales']):
        print(line)
