import os, json, csv, re, glob, shutil
from pathlib import Path
from decimal import Decimal, ROUND_HALF_UP
from openpyxl import load_workbook

workspace = Path.cwd()
params = json.load(open(os.environ['AEP_WORKFLOW_PARAMS'], encoding='utf-8'))
start_date = params['start_date']
end_date = params['end_date']
status = params['status']
output_filename = params['output_filename']
if not output_filename.lower().endswith('.xlsx'):
    output_filename += '.xlsx'
input_dir = workspace / 'input'
out_dir = workspace / 'output'
download_dir = out_dir / 'downloads'
out_dir.mkdir(exist_ok=True)

candidates = list(download_dir.glob('sales_orders_export*.csv')) + list(download_dir.glob('*orders*export*.csv')) + list(download_dir.glob('*.csv'))
if not candidates:
    raise FileNotFoundError(f'No downloaded CSV found in {download_dir}')
export_path = max(candidates, key=lambda p: p.stat().st_mtime)

fx_path = input_dir / 'fx_rates.csv'
fx = {}
with fx_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    for row in reader:
        lowered = {k.strip().lower(): v for k, v in row.items() if k is not None}
        cur = (lowered.get('currency') or lowered.get('ccy') or lowered.get('from') or '').strip()
        rate_s = (lowered.get('rate') or lowered.get('twd') or lowered.get('to_twd') or lowered.get('rate_to_twd') or lowered.get('twd_rate') or '').strip()
        if cur and rate_s:
            fx[cur] = Decimal(rate_s)
fx.setdefault('TWD', Decimal('1'))

def norm_key(k):
    return re.sub(r'[^a-z0-9]', '', (k or '').lower())

def pick(row, names):
    by_norm = {norm_key(k): v for k, v in row.items()}
    for name in names:
        if norm_key(name) in by_norm:
            return by_norm[norm_key(name)]
    return ''

def dec_value(v):
    s = str(v or '').replace(',', '').strip()
    if s == '':
        return Decimal('0')
    return Decimal(s)

def money(v):
    return Decimal(v).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)

with export_path.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.DictReader(f)
    rows = list(reader)

records = []
for row in rows:
    order_id = str(pick(row, ['OrderID','Order ID','Order'])).strip()
    order_date = str(pick(row, ['OrderDate','Order Date','Date'])).strip()
    row_status = str(pick(row, ['Status'])).strip()
    if not order_id:
        continue
    if order_date and not (start_date <= order_date <= end_date):
        continue
    if status and row_status.lower() != status.lower():
        continue
    currency = str(pick(row, ['Currency','Curr'])).strip()
    net = money(dec_value(pick(row, ['NetAmount','Net Amount','Amount'])))
    if currency not in fx:
        raise ValueError(f'Missing FX rate for {currency}')
    net_twd = money(net * fx[currency])
    contact = str(pick(row, ['CustomerContact','Customer Contact','Contact'])).strip()
    payment = str(pick(row, ['PaymentTerms','Payment Terms','Terms'])).strip()
    if not contact:
        contact = 'N/A'
    if not payment:
        payment = 'N/A'
    records.append({
        'OrderID': order_id,
        'CustomerID': str(pick(row, ['CustomerID','Customer ID'])).strip(),
        'CustomerName': str(pick(row, ['CustomerName','Customer Name'])).strip(),
        'ProjectID': str(pick(row, ['ProjectID','Project ID'])).strip(),
        'OrderDate': order_date,
        'Status': row_status,
        'Currency': currency,
        'NetAmount': net,
        'NetAmountTWD': net_twd,
        'CustomerContact': contact,
        'PaymentTerms': payment,
    })

records.sort(key=lambda r: r['OrderID'])

template_path = input_dir / 'sales_template.xlsx'
out_path = out_dir / output_filename
shutil.copyfile(template_path, out_path)
wb = load_workbook(out_path)
for s in ['Report','Details','Summary']:
    if s not in wb.sheetnames:
        wb.create_sheet(s)

wsr = wb['Report']
wsr['A1'] = 'Sales Orders Report'
wsr['A2'] = 'From'
wsr['B2'] = start_date
wsr['A3'] = 'To'
wsr['B3'] = end_date
wsr['A4'] = 'Status'
wsr['B4'] = status
wsr['A5'] = 'Generated Rows'
wsr['B5'] = len(records)

headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
ws = wb['Details']
ws.delete_rows(1, ws.max_row)
for c, h in enumerate(headers, 1):
    ws.cell(1, c).value = h
for r_idx, rec in enumerate(records, 2):
    for c_idx, h in enumerate(headers, 1):
        val = rec[h]
        if h in ['OrderID','CustomerID','ProjectID']:
            val = str(val)
            ws.cell(r_idx, c_idx).number_format = '@'
        elif h in ['NetAmount','NetAmountTWD']:
            val = float(val)
            ws.cell(r_idx, c_idx).number_format = '#,##0.00'
        ws.cell(r_idx, c_idx).value = val

summary = {}
for rec in records:
    cur = rec['Currency']
    summary.setdefault(cur, {'Amount': Decimal('0'), 'AmountTWD': Decimal('0'), 'Count': 0})
    summary[cur]['Amount'] += rec['NetAmount']
    summary[cur]['AmountTWD'] += rec['NetAmountTWD']
    summary[cur]['Count'] += 1

wss = wb['Summary']
wss.delete_rows(1, wss.max_row)
for c, h in enumerate(['Currency','Amount','AmountTWD','Count'], 1):
    wss.cell(1, c).value = h
row_i = 2
grand = Decimal('0')
for cur in sorted(summary):
    vals = summary[cur]
    wss.cell(row_i, 1).value = cur
    wss.cell(row_i, 2).value = float(money(vals['Amount']))
    wss.cell(row_i, 3).value = float(money(vals['AmountTWD']))
    wss.cell(row_i, 4).value = vals['Count']
    wss.cell(row_i, 2).number_format = '#,##0.00'
    wss.cell(row_i, 3).number_format = '#,##0.00'
    grand += vals['AmountTWD']
    row_i += 1
wss.cell(row_i, 1).value = 'Total TWD'
wss.cell(row_i, 3).value = float(money(grand))
wss.cell(row_i, 3).number_format = '#,##0.00'
wss.cell(row_i, 4).value = len(records)

wb.save(out_path)

wb2 = load_workbook(out_path, data_only=True)
ws2 = wb2['Details']
headers2 = [ws2.cell(1, c).value for c in range(1, len(headers)+1)]
if headers2 != headers:
    raise AssertionError(f'Details headers mismatch: {headers2}')
loaded = []
for r in range(2, ws2.max_row+1):
    if ws2.cell(r,1).value is None:
        continue
    loaded.append({h: ws2.cell(r, i+1).value for i,h in enumerate(headers)})
if len(loaded) != len(records):
    raise AssertionError(f'Row count mismatch {len(loaded)} != {len(records)}')
ids = [str(r['OrderID']) for r in loaded]
if ids != sorted(ids):
    raise AssertionError('OrderID is not sorted')
if any(not (start_date <= str(r['OrderDate']) <= end_date) for r in loaded):
    raise AssertionError('Date outside requested window')
if any(str(r['Status']).lower() != status.lower() for r in loaded):
    raise AssertionError('Status outside requested filter')
if any(not isinstance(r['NetAmount'], (int, float)) or not isinstance(r['NetAmountTWD'], (int, float)) for r in loaded):
    raise AssertionError('Amount columns are not numeric')
if any(not str(r['OrderID']).startswith('0') for r in loaded):
    raise AssertionError('OrderID leading zero not preserved')
if any(r['CustomerContact'] in (None, '') or r['PaymentTerms'] in (None, '') for r in loaded):
    raise AssertionError('Missing detail placeholders')
print('export_path=', export_path)
print('output_path=', out_path)
print('rows=', len(loaded))
print('ids=', ','.join(ids))
print('summary=', {k: {'Amount': str(money(v['Amount'])), 'AmountTWD': str(money(v['AmountTWD'])), 'Count': v['Count']} for k,v in summary.items()})
print('grand_twd=', str(money(grand)))