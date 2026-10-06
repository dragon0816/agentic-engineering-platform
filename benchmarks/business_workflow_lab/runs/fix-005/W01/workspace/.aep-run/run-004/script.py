from pathlib import Path
from openpyxl import load_workbook
import csv, json, os

ws = Path.cwd()
print('workspace=', ws)
print('input files=', sorted(str(p.relative_to(ws)) for p in (ws/'input').glob('*')))
print('output files=', sorted(str(p.relative_to(ws)) for p in (ws/'output').rglob('*') if p.is_file()))

out = ws/'output'/'september_open_orders.xlsx'
if not out.exists():
    raise SystemExit('MISSING output/september_open_orders.xlsx')

wb = load_workbook(out, data_only=True)
print('sheets=', wb.sheetnames)
for s in wb.sheetnames:
    sh = wb[s]
    print('sheet', s, 'size', sh.max_row, sh.max_column)
    for row in sh.iter_rows(min_row=1, max_row=min(5, sh.max_row), values_only=True):
        print(' ', row)

# Detect Details/Summary sheets and headers
if 'Details' not in wb.sheetnames:
    raise SystemExit('MISSING Details sheet')
if 'Summary' not in wb.sheetnames:
    raise SystemExit('MISSING Summary sheet')

d = wb['Details']
headers = [c.value for c in d[1]]
print('details headers=', headers)
rows = []
for r in d.iter_rows(min_row=2, values_only=True):
    if any(v is not None for v in r):
        rows.append(dict(zip(headers, r)))
print('details row count=', len(rows))
print('first order=', rows[0].get('OrderID') if rows else None)
print('last order=', rows[-1].get('OrderID') if rows else None)
print('order ids=', [r.get('OrderID') for r in rows])

# Validate stated requirements
expected_headers = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
missing = [h for h in expected_headers if h not in headers]
if missing:
    raise SystemExit('Missing Details headers: ' + repr(missing))

order_ids = [str(r['OrderID']) for r in rows]
assert order_ids == sorted(order_ids), 'OrderID not sorted ascending'
assert all(r['Status'] == 'Open' for r in rows), 'Non-Open status found'
assert all(str(r['OrderDate'])[:7] == '2026-09' for r in rows), 'OrderDate outside 2026-09 found'
assert all(isinstance(r['NetAmount'], (int, float)) for r in rows), 'NetAmount not numeric'
assert all(isinstance(r['NetAmountTWD'], (int, float)) for r in rows), 'NetAmountTWD not numeric'
assert all(len(str(r['OrderID'])) == 10 and str(r['OrderID']).startswith('0') for r in rows), 'OrderID text/leading zero check failed'

# Recompute from fx rates
rates = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    for rec in csv.DictReader(f):
        cur = rec.get('Currency') or rec.get('currency')
        rate = rec.get('TWD') or rec.get('RateToTWD') or rec.get('rate') or rec.get('Rate')
        if cur and rate:
            rates[cur] = float(rate)
print('fx rates=', rates)

currency_amount = {}
currency_twd = {}
for r in rows:
    cur = r['Currency']
    amt = float(r['NetAmount'])
    twd_expected = round(amt * rates[cur], 2)
    if round(float(r['NetAmountTWD']), 2) != twd_expected:
        raise SystemExit(f'TWD mismatch {r["OrderID"]}: got {r["NetAmountTWD"]}, expected {twd_expected}')
    currency_amount[cur] = currency_amount.get(cur, 0.0) + amt
    currency_twd[cur] = currency_twd.get(cur, 0.0) + float(r['NetAmountTWD'])
print('recomputed currency amount=', {k: round(v,2) for k,v in sorted(currency_amount.items())})
print('recomputed currency twd=', {k: round(v,2) for k,v in sorted(currency_twd.items())})
print('recomputed total twd=', round(sum(currency_twd.values()),2))

s = wb['Summary']
summary_values = [[c for c in row] for row in s.iter_rows(values_only=True) if any(v is not None for v in row)]
print('summary non-empty rows=')
for row in summary_values:
    print(' ', row)

print('CHECKS PASSED')