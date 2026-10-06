from pathlib import Path
import csv, json
from openpyxl import load_workbook
ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size)
print('\nFX RATES HEAD')
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    r=csv.reader(f)
    for i,row in zip(range(8), r):
        print(row)
print('\nTEMPLATE SHEETS/CELLS')
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
for sh in wb.sheetnames:
    wsht=wb[sh]
    print('SHEET', sh, 'max_row', wsht.max_row, 'max_col', wsht.max_column)
    for row in wsht.iter_rows(min_row=1, max_row=min(wsht.max_row, 12), values_only=True):
        print(row)
    print('---')
print('\nMANUAL PREVIEWS')
for name in ['sap_manual.md','expense_manual.md','quote_cli_manual.md']:
    p=ws/'input'/name
    print('\n###', name)
    text=p.read_text(encoding='utf-8')
    print(text[:4000])