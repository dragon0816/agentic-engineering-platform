from pathlib import Path
import csv, json, re
from openpyxl import load_workbook

ws = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-005\X02\workspace')
out = ws/'output'
report = out/'project_C001_2026-09_report.xlsx'
print('WORKSPACE', ws)
print('REPORT', report, 'exists=', report.exists(), 'bytes=', report.stat().st_size if report.exists() else None)

required_files = [
    out/'downloads'/'orders_C001_2026-09.csv',
    out/'downloads'/'invoices_C001_2026-09.csv',
    out/'expenses_P001_2026-09.csv',
    out/'expenses_P003_2026-09.csv',
    out/'quotes_C001_2026-09.csv',
    out/'quotes_C001_2026-09.json',
    report,
]
for p in required_files:
    print('FILE', p.relative_to(ws), 'exists=', p.exists(), 'bytes=', p.stat().st_size if p.exists() else None)
    if not p.exists():
        raise SystemExit(f'Missing required file: {p}')

# Source counts and fields.
def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))
orders = read_csv(out/'downloads'/'orders_C001_2026-09.csv')
invoices = read_csv(out/'downloads'/'invoices_C001_2026-09.csv')
expenses = read_csv(out/'expenses_P001_2026-09.csv') + read_csv(out/'expenses_P003_2026-09.csv')
quotes = read_csv(out/'quotes_C001_2026-09.csv')
with open(out/'quotes_C001_2026-09.json', encoding='utf-8') as f:
    quotes_json = json.load(f)
print('SOURCE_COUNTS', {'orders':len(orders),'invoices':len(invoices),'expenses':len(expenses),'quotes_csv':len(quotes),'quotes_json':len(quotes_json) if isinstance(quotes_json, list) else 'non-list'})
print('SOURCE_HEADERS', {
    'orders': list(orders[0].keys()) if orders else [],
    'invoices': list(invoices[0].keys()) if invoices else [],
    'expenses': list(expenses[0].keys()) if expenses else [],
    'quotes': list(quotes[0].keys()) if quotes else [],
})

wb = load_workbook(report, data_only=False)
print('SHEETS', wb.sheetnames)
expected_sheets = ['Report','Summary','Orders','Invoices','Expenses','Quotes','Exceptions','Instructions']
missing = [s for s in expected_sheets if s not in wb.sheetnames]
if missing:
    raise AssertionError(f'Missing sheets: {missing}')

# Find header row by expected first column labels.
expected_first = {
    'Summary': 'ProjectID',
    'Orders': 'OrderID',
    'Invoices': 'InvoiceID',
    'Expenses': 'ExpenseID',
    'Quotes': 'QuoteID',
    'Exceptions': 'Source',
}

def find_header(sheet_name, first_label):
    sh = wb[sheet_name]
    for r in range(1, sh.max_row+1):
        if sh.cell(r,1).value == first_label:
            headers = [sh.cell(r,c).value for c in range(1, sh.max_column+1)]
            return r, headers
    raise AssertionError(f'Could not find header {first_label} in {sheet_name}')

def data_rows(sheet_name, first_label):
    sh = wb[sheet_name]
    hr, headers = find_header(sheet_name, first_label)
    rows = []
    for r in range(hr+1, sh.max_row+1):
        vals = [sh.cell(r,c).value for c in range(1, len(headers)+1)]
        if any(v is not None for v in vals):
            rows.append((r, vals))
    return hr, headers, rows

expected_counts = {'Orders':len(orders),'Invoices':len(invoices),'Expenses':len(expenses),'Quotes':len(quotes)}
for sheet, count in expected_counts.items():
    hr, headers, rows = data_rows(sheet, expected_first[sheet])
    print('SHEET_DATA', sheet, 'header_row=', hr, 'rows=', len(rows), 'headers=', headers)
    if len(rows) != count:
        raise AssertionError(f'{sheet} row count {len(rows)} != source count {count}')

# Summary should contain project rows P001/P003 plus grand total (if present). Validate ProjectIDs and numeric TWD totals.
hr, headers, rows = data_rows('Summary', 'ProjectID')
summary_projects = [row[0][0] for row in []]
project_ids = [vals[0] for _, vals in rows if vals[0] not in (None, 'Total', 'Grand Total')]
print('SUMMARY_PROJECT_IDS', project_ids)
if not {'P001','P003'}.issubset(set(project_ids)):
    raise AssertionError('Summary missing P001/P003')
for r, vals in rows:
    for h, v in zip(headers, vals):
        if h and 'TWD' in str(h) and v is not None:
            # formulas or numbers are acceptable in template summary
            if not (isinstance(v, (int,float)) or (isinstance(v, str) and v.startswith('='))):
                raise AssertionError(f'Summary {h} row {r} not numeric/formula: {v!r}')

# ID columns text, amounts numeric/formula, TWD columns present.
for sheet in ['Orders','Invoices','Expenses','Quotes']:
    hr, headers, rows = data_rows(sheet, expected_first[sheet])
    id_cols = [i for i,h in enumerate(headers) if h and str(h).lower().endswith('id')]
    twd_cols = [i for i,h in enumerate(headers) if h and 'TWD' in str(h)]
    if not twd_cols:
        raise AssertionError(f'{sheet} missing TWD column')
    for r, vals in rows:
        for i in id_cols:
            if vals[i] is not None and not isinstance(vals[i], str):
                raise AssertionError(f'{sheet} ID column {headers[i]} row {r} is not text: {vals[i]!r}')
        for i in twd_cols:
            if vals[i] is not None and not isinstance(vals[i], (int,float)):
                raise AssertionError(f'{sheet} TWD column {headers[i]} row {r} is not numeric: {vals[i]!r}')
print('TYPE_CHECKS passed')

# Recompute attention rule count from source using as-of 2026-10-06.
from datetime import date
as_of = date(2026,10,6)
exceptions = []
for rec in orders:
    if rec.get('Status') == 'Cancelled':
        exceptions.append(('Orders', rec.get('OrderID')))
for rec in invoices:
    bal = float(rec.get('Balance') or rec.get('BalanceAmount') or 0)
    due = date.fromisoformat(rec.get('DueDate')) if rec.get('DueDate') else None
    if bal > 0 and due and due < as_of:
        exceptions.append(('Invoices', rec.get('InvoiceID')))
for rec in expenses:
    status = rec.get('Status')
    receipt = str(rec.get('ReceiptAttached') or rec.get('Receipt') or '').strip().lower()
    exp_id = rec.get('ExpenseID')
    if status != 'Approved':
        exceptions.append(('Expenses', exp_id))
    if receipt in ('false','no','0','n',''):
        exceptions.append(('Expenses', exp_id))
for rec in quotes:
    status = rec.get('Status')
    valid_until = rec.get('ValidUntil') or rec.get('ValidTo') or rec.get('ExpiryDate')
    expired_by_date = False
    if valid_until:
        try:
            expired_by_date = date.fromisoformat(valid_until) < as_of
        except Exception:
            expired_by_date = False
    if status in ('Draft','Expired') or (expired_by_date and status != 'Accepted'):
        exceptions.append(('Quotes', rec.get('QuoteID')))

hr, headers, rows = data_rows('Exceptions', 'Source')
print('EXCEPTION_EXPECTED_COUNT', len(exceptions), 'WORKBOOK_EXCEPTION_ROWS', len(rows), 'headers=', headers)
if len(rows) != len(exceptions):
    raise AssertionError(f'Exceptions row count {len(rows)} != expected {len(exceptions)}')

# Check some expected summary numbers in data_only mode.
wbv = load_workbook(report, data_only=True)
sh = wbv['Summary']
# find header row again in data-only workbook
hdr_row = None
for r in range(1, sh.max_row+1):
    if sh.cell(r,1).value == 'ProjectID':
        hdr_row = r; break
headers_v = [sh.cell(hdr_row,c).value for c in range(1, sh.max_column+1)]
summary = {}
for r in range(hdr_row+1, sh.max_row+1):
    pid = sh.cell(r,1).value
    if pid in ('P001','P003'):
        summary[pid] = {headers_v[c-1]: sh.cell(r,c).value for c in range(1, sh.max_column+1)}
print('SUMMARY_VALUES', summary)
expected_summary = {
    'P001': {'OrderTWD': 1076162.00, 'InvoiceBalanceTWD': 386445.00, 'ApprovedExpenseTWD': 18305.00, 'AcceptedQuoteTWD': 331200.00},
    'P003': {'OrderTWD': 150000.00, 'InvoiceBalanceTWD': 0.00, 'ApprovedExpenseTWD': 13037.50, 'AcceptedQuoteTWD': 0.00},
}
for pid, checks in expected_summary.items():
    for col, exp in checks.items():
        got = summary.get(pid, {}).get(col)
        if got is None or abs(float(got)-exp) > 0.01:
            raise AssertionError(f'Summary {pid} {col}: got {got}, expected {exp}')
print('SUMMARY_CHECK passed')

# Highlight check: at least exception-related detail rows have a yellow-ish fill.
for sheet in ['Orders','Invoices','Expenses','Quotes']:
    hr, headers, rows = data_rows(sheet, expected_first[sheet])
    highlighted = 0
    for r, vals in rows:
        fill = wb[sheet].cell(r,1).fill
        color = fill.fgColor.rgb or fill.fgColor.indexed or fill.fgColor.theme
        if fill.fill_type and color not in (None, '00000000'):
            highlighted += 1
    print('HIGHLIGHTED_ROWS', sheet, highlighted)
print('CHECK_COMPLETE all required output files readable; workbook sheets, rows, text IDs, numeric TWD amounts, exceptions, and summary values verified')