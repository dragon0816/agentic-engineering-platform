from pathlib import Path
import csv, json
from datetime import date
from openpyxl import load_workbook

ws = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-005\X02\workspace')
out = ws/'output'
report = out/'project_C001_2026-09_report.xlsx'

print('FINAL RESULT CHECK')
print('workspace:', ws)

required = [
    out/'downloads'/'orders_C001_2026-09.csv',
    out/'downloads'/'invoices_C001_2026-09.csv',
    out/'expenses_P001_2026-09.csv',
    out/'expenses_P003_2026-09.csv',
    out/'quotes_C001_2026-09.csv',
    out/'quotes_C001_2026-09.json',
    report,
]
for p in required:
    print('file:', p.relative_to(ws), 'exists:', p.exists(), 'bytes:', p.stat().st_size if p.exists() else None)
    assert p.exists() and p.stat().st_size > 0

def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))
orders = read_csv(out/'downloads'/'orders_C001_2026-09.csv')
invoices = read_csv(out/'downloads'/'invoices_C001_2026-09.csv')
expenses = read_csv(out/'expenses_P001_2026-09.csv') + read_csv(out/'expenses_P003_2026-09.csv')
quotes = read_csv(out/'quotes_C001_2026-09.csv')
with open(out/'quotes_C001_2026-09.json', encoding='utf-8') as f:
    quotes_json = json.load(f)

print('source counts:', {'orders': len(orders), 'invoices': len(invoices), 'expenses': len(expenses), 'quotes': len(quotes)})
print('quotes_json top-level:', type(quotes_json).__name__, 'keys:', list(quotes_json) if isinstance(quotes_json, dict) else None)

wb = load_workbook(report, data_only=False)
assert wb.sheetnames == ['Report', 'Orders', 'Invoices', 'Expenses', 'Quotes', 'Summary', 'Exceptions', 'Instructions']
print('workbook sheets:', wb.sheetnames)

first_labels = {'Summary':'ProjectID','Orders':'OrderID','Invoices':'InvoiceID','Expenses':'ExpenseID','Quotes':'QuoteID','Exceptions':'Source'}
def table(sheet):
    sh = wb[sheet]
    first = first_labels[sheet]
    hr = None
    for r in range(1, sh.max_row+1):
        if sh.cell(r,1).value == first:
            hr = r; break
    assert hr, f'no header in {sheet}'
    headers = [sh.cell(hr,c).value for c in range(1, sh.max_column+1)]
    rows=[]
    for r in range(hr+1, sh.max_row+1):
        vals=[sh.cell(r,c).value for c in range(1, len(headers)+1)]
        if any(v is not None for v in vals):
            d=dict(zip(headers, vals)); d['__row__']=r; rows.append(d)
    return headers, rows

# Row counts and field checks.
expected_counts = {'Orders':len(orders),'Invoices':len(invoices),'Expenses':len(expenses),'Quotes':len(quotes)}
for sheet, count in expected_counts.items():
    headers, rows = table(sheet)
    print(sheet, 'headers:', headers, 'rows:', len(rows))
    assert len(rows) == count
    assert any('TWD' in str(h) for h in headers)
    for row in rows:
        for h in headers:
            if h and str(h).endswith('ID') and row.get(h) is not None:
                assert isinstance(row[h], str), f'{sheet} {h} not text at row {row["__row__"]}'
            if h and 'TWD' in str(h) and row.get(h) is not None:
                assert isinstance(row[h], (int,float)), f'{sheet} {h} not numeric at row {row["__row__"]}'
print('detail sheet checks passed')

# Summary values.
headers, rows = table('Summary')
summary = {r['ProjectID']: r for r in rows if r.get('ProjectID') in ('P001','P003')}
expected_summary = {
    'P001': {'OrderTWD': 1076162.00, 'InvoiceBalanceTWD': 386445.00, 'ApprovedExpenseTWD': 18305.00, 'AcceptedQuoteTWD': 331200.00},
    'P003': {'OrderTWD': 150000.00, 'InvoiceBalanceTWD': 0.00, 'ApprovedExpenseTWD': 13037.50, 'AcceptedQuoteTWD': 0.00},
}
for pid, cols in expected_summary.items():
    assert pid in summary
    for col, val in cols.items():
        got = summary[pid].get(col)
        assert got is not None and abs(float(got)-val) < 0.01, (pid, col, got, val)
print('summary values passed:', expected_summary)

# Exceptions: verify every exception row references an existing source document and source doc has at least one attention condition.
as_of = date(2026,10,6)
source_maps = {
    'Orders': {r['OrderID']: r for r in orders},
    'Invoices': {r['InvoiceID']: r for r in invoices},
    'Expenses': {r['ExpenseID']: r for r in expenses},
    'Quotes': {r['QuoteID']: r for r in quotes},
}
def has_attention(src, rec):
    if src == 'Orders':
        return rec.get('Status') == 'Cancelled'
    if src == 'Invoices':
        return float(rec.get('Balance') or 0) > 0 and rec.get('DueDate') and date.fromisoformat(rec['DueDate']) < as_of
    if src == 'Expenses':
        missing = str(rec.get('MissingReceipt','')).strip().lower() in ('true','yes','1','y')
        return rec.get('Status') != 'Approved' or missing
    if src == 'Quotes':
        valid = date.fromisoformat(rec['ValidUntil']) if rec.get('ValidUntil') else None
        return rec.get('Status') in ('Draft','Expired') or (valid and valid < as_of and rec.get('Status') != 'Accepted')
    return False

headers, exc_rows = table('Exceptions')
print('Exceptions headers:', headers, 'rows:', len(exc_rows))
for r in exc_rows:
    src = str(r.get('Source'))
    doc = str(r.get('DocumentID'))
    reason = str(r.get('Reason'))
    print('exception:', src, doc, r.get('ProjectID'), reason)
    assert src in source_maps, f'bad source {src}'
    assert doc in source_maps[src], f'{src} document not in source: {doc}'
    assert has_attention(src, source_maps[src][doc]), f'{src} {doc} has no attention condition'
    assert reason and reason != 'None'
print('exception traceability passed')

# Ensure all source attention documents appear at least once in Exceptions.
expected_attention_docs=set()
for src, records in source_maps.items():
    id_field = {'Orders':'OrderID','Invoices':'InvoiceID','Expenses':'ExpenseID','Quotes':'QuoteID'}[src]
    for rec in records.values():
        if has_attention(src, rec):
            expected_attention_docs.add((src, rec[id_field]))
actual_docs={(str(r['Source']), str(r['DocumentID'])) for r in exc_rows}
print('expected attention documents:', len(expected_attention_docs), sorted(expected_attention_docs))
print('actual exception documents:', len(actual_docs), sorted(actual_docs))
assert expected_attention_docs <= actual_docs
print('all attention documents represented; duplicate rows are separate reasons for the same document if present')

# Highlight checks.
for sheet in ['Orders','Invoices','Expenses','Quotes']:
    headers, rows = table(sheet)
    id_field = first_labels[sheet]
    expected_doc_ids = {doc for src, doc in expected_attention_docs if src == sheet}
    highlighted=[]
    for row in rows:
        if row.get(id_field) in expected_doc_ids:
            fill = wb[sheet].cell(row['__row__'],1).fill
            has_fill = bool(fill.fill_type) and (fill.fgColor.rgb not in (None,'00000000') or fill.fgColor.indexed is not None or fill.fgColor.theme is not None)
            if has_fill:
                highlighted.append(row.get(id_field))
    print('highlight:', sheet, len(highlighted), 'of', len(expected_doc_ids), sorted(highlighted))
    assert set(highlighted) == expected_doc_ids
print('highlight checks passed')

print('FINAL_CHECK_COMPLETE: workbook and exported source files are readable, counts match, TWD numeric fields and text IDs are preserved, summary totals match, and all attention documents are represented/highlighted.')