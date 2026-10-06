from pathlib import Path
from openpyxl import load_workbook
ws=Path.cwd()
wb=load_workbook(ws/'input'/'project_template.xlsx', data_only=False)
ins=wb['Instructions']
for row in ins.iter_rows(min_row=1, max_row=ins.max_row, values_only=True):
    print(row)
print('\nDownloaded files:')
for p in sorted((ws/'output'/'downloads').glob('*')):
    print(p.name, p.stat().st_size)
    print(p.read_text(encoding='utf-8-sig')[:1000])