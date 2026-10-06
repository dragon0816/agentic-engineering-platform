from pathlib import Path
import csv, json
from datetime import date
from openpyxl import load_workbook

ws = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\runs\fix-005\X02\workspace')
out = ws/'output'
report = out/'project_C001_2026-09_report.xlsx'

print('=== FILE INVENTORY ===')
for p in sorted(out.rglob('*')):
    if p.is_file():
        print(p.relative_to(ws), p.stat().st_size)

required = {
    'workbook': report,
    'orders': out/'downloads'/'orders_C001_2026-09.csv',
    'invoices': out/'downloads'/'invoices_C001_2026-09.csv',
    'expenses_P001': out/'expenses_P001_2026-09.csv',
    'expenses_P003': out/'expenses_P003_2026-09.csv',
    'quotes_csv': out/'quotes_C001_2026-09.csv',
    'quotes_json': out/'quotes_C001_2026-09.json',
}
for name, path in required.items():
    if not path.exists() or path.stat().st_size == 0:
        raise AssertionError(f'{name} missing or empty: {path}')
print('REQUIRED_FILES_CHECK passed')

def read_csv(path):
    with open(path, newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

orders = read_csv(required['orders'])
invoices = read_csv(required['invoices'])
expenses = read_csv(required['expenses_P001']) + read_csv(required['expenses_P003'])
quotes = read_csv(required['quotes_csv'])
with open(required['quotes_json'], encoding='utf-8') as f:
    qj = json.load(f)
print('SOURCE_COUNTS', {'orders':len(orders), 'invoices':len(invoices), 'expenses':len(expenses), 'quotes_csv':len(quotes), 'quotes_json_type':type(qj).__name__})

wb = load_workbook(report, data_only=False)
print('WORKBOOK_SHEETS', wb.sheetnames)
expected_sheets = ['Report','Orders','Invoices','Expenses','Quotes','Summary','Exceptions','Instructions']
assert wb.sheetnames == expected_sheets, f'sheet order/name mismatch: {wb.sheetnames}'

expected_first = {'Summary':'ProjectID','Orders':'OrderID','Invoices':'InvoiceID','Expenses':'ExpenseID','Quotes':'QuoteID','Exceptions':'Source'}

def table(sheet, first_label, wb_obj=wb):
    sh = wb_obj[sheet]
    hr = None
    for r in range(1, sh.max_row+1):
        if sh.cell(r,1).value == first_label:
            hr = r; break
    if not hr:
        raise AssertionError(f'Header not found for {sheet}')
    headers = [sh.cell(hr,c).value for c in range(1, sh.max_column+1)]
    rows = []
    for r in range(hr+1, sh.max_row+1):
        vals = [sh.cell(r,c).value for c in range(1, len(headers)+1)]
        if any(v is not None for v in vals):
            rows.append(dict(zip(headers, vals)) | {'__row__': r})
    return hr, headers, rows

for sheet, src_count in [('Orders',len(orders)),('Invoices',len(invoices)),('Expenses',len(expenses)),('Quotes',len(quotes))]:
    hr, headers, rows = table(sheet, expected_first[sheet])
    print('TABLE', sheet, 'header_row', hr, 'rows', len(rows), 'headers', headers)
    assert len(rows) == src_count, f'{sheet}: {len(rows)} != {src_count}'
    assert any('TWD' in str(h) for h in headers), f'{sheet} lacks TWD column'
    id_headers = [h for h in headers if h and str(h).endswith('ID')]
    for row in rows:
        for h in id_headers:
            if row.get(h) is not None:
                assert isinstance(row[h], str), f'{sheet} {h} row {row["__row__"]} is not text'
        for h in headers:
            if h and 'TWD' in str(h) and row.get(h) is not None:
                assert isinstance(row[h], (int,float)), f'{sheet} {h} row {row["__row__"]} not numeric'
print('ROW_COUNT_AND_TYPE_CHECK passed')

# Summary values in the workbook.
_, s_headers, s_rows = table('Summary','ProjectID')
summary = {r['ProjectID']: r for r in s_rows if r.get('ProjectID') in ('P001','P003')}
print('SUMMARY_ROWS', summary)
expected_summary = {
    'P001': {'OrderTWD': 1076162.00, 'InvoiceBalanceTWD': 386445.00, 'ApprovedExpenseTWD': 18305.00, 'AcceptedQuoteTWD': 331200.00},
    'P003': {'OrderTWD': 150000.00, 'InvoiceBalanceTWD': 0.00, 'ApprovedExpenseTWD': 13037.50, 'AcceptedQuoteTWD': 0.00},
}
for pid, vals in expected_summary.items():
    assert pid in summary, f'Summary missing {pid}'
    for col, expected in vals.items():
        got = summary[pid].get(col)
        assert got is not None, f'Summary missing {pid} {col}'
        assert abs(float(got) - expected) <= 0.01, f'{pid} {col}: {got} != {expected}'
print('SUMMARY_VALUE_CHECK passed')

# Expected exceptions: one row per document with any attention condition.
as_of = date(2026,10,6)
expected_exc = {}
def add(src, doc, project, reason):
    key = (src, doc)
    if key not in expected_exc:
        expected_exc[key] = {'Source':src, 'DocumentID':doc, 'ProjectID':project, 'Reasons':[]}
    expected_exc[key]['Reasons'].append(reason)

for r in orders:
    if r.get('Status') == 'Cancelled':
        add('Orders', r.get('OrderID'), r.get('ProjectID'), 'Cancelled')
for r in invoices:
    bal = float(r.get('Balance') or 0)
    due = date.fromisoformat(r['DueDate']) if r.get('DueDate') else None
    if bal > 0 and due and due < as_of:
        add('Invoices', r.get('InvoiceID'), r.get('ProjectID'), 'Overdue unpaid balance')
for r in expenses:
    reasons = []
    if r.get('Status') != 'Approved': reasons.append('Not approved')
    mr = str(r.get('MissingReceipt','')).strip().lower()
    if mr in ('true','yes','1','y'): reasons.append('Missing receipt')
    for reason in reasons:
        add('Expenses', r.get('ExpenseID'), r.get('ProjectID'), reason)
for r in quotes:
    status = r.get('Status')
    valid = date.fromisoformat(r['ValidUntil']) if r.get('ValidUntil') else None
    if status in ('Draft','Expired') or (valid and valid < as_of and status != 'Accepted'):
        add('Quotes', r.get('QuoteID'), r.get('ProjectID'), 'Draft/expired/not accepted')

_, e_headers, e_rows = table('Exceptions','Source')
actual_keys = {(str(r.get('Source')), str(r.get('DocumentID'))) for r in e_rows}
expected_keys = {(src, doc) for (src, doc) in expected_exc.keys()}
print('EXPECTED_EXCEPTION_COUNT_UNIQUE_DOCS', len(expected_keys))
print('ACTUAL_EXCEPTION_ROWS', len(e_rows))
print('ACTUAL_EXCEPTION_KEYS', sorted(actual_keys))
assert len(e_rows) == len(expected_keys), f'Exceptions: {len(e_rows)} != {len(expected_keys)}'
missing = expected_keys - actual_keys
extra = actual_keys - expected_keys
assert not missing and not extra, f'Exception keys mismatch missing={missing} extra={extra}'
print('EXCEPTIONS_CHECK passed')

# Check yellow-ish highlight exists for attention rows.
for sheet in ['Orders','Invoices','Expenses','Quotes']:
    _, headers, rows = table(sheet, expected_first[sheet])
    id_col = expected_first[sheet]
    expected_doc_ids = {doc for src, doc in expected_keys if src == sheet}
    highlighted_expected = 0
    for row in rows:
        if row.get(id_col) in expected_doc_ids:
            cell = wb[sheet].cell(row['__row__'], 1)
            fill = cell.fill
            if fill.fill_type and (fill.fgColor.rgb not in (None, '00000000') or fill.fgColor.indexed not in (None, '00000000') or fill.fgColor.theme is not None):
                highlighted_expected += 1
    print('HIGHLIGHT_CHECK', sheet, 'expected_attention_docs', len(expected_doc_ids), 'highlighted', highlighted_expected)
    assert highlighted_expected == len(expected_doc_ids), f'{sheet} not all attention rows highlighted'
print('HIGHLIGHT_CHECK passed')

# Report sheet metadata present.
report_text = []
for row in wb['Report'].iter_rows(values_only=True):
    for v in row:
        if v is not None:
            report_text.append(str(v))
joined = '\n'.join(report_text)
for token in ['C001', '2026-09-01', '2026-09-30', 'TWD']:
    assert token in joined, f'Report sheet missing {token}'
print('REPORT_METADATA_CHECK passed')
print('FINAL_CHECK_COMPLETE')