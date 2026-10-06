from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('workspace', ws)
print('input files:')
for p in sorted((ws/'input').iterdir()):
    print('-', p.name, p.stat().st_size)
# workbook structure
xlsx=ws/'input'/'sales_template.xlsx'
wb=load_workbook(xlsx, data_only=False)
print('workbook sheets:', wb.sheetnames)
for s in wb.sheetnames:
    sh=wb[s]
    print('sheet', s, 'max_row', sh.max_row, 'max_col', sh.max_column)
    for r in sh.iter_rows(min_row=1, max_row=min(5, sh.max_row), values_only=True):
        print(' ', r)
# fx csv
csvp=ws/'input'/'fx_rates.csv'
with open(csvp, newline='', encoding='utf-8-sig') as f:
    reader=csv.reader(f)
    for i,row in zip(range(8), reader):
        print('fx row', i+1, row)
# manual excerpt
md=ws/'input'/'sap_manual.md'
text=md.read_text(encoding='utf-8')
print('manual first 3000 chars:')
print(text[:3000])