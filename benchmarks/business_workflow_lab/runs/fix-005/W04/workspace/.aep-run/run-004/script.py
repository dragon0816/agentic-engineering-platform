from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Alignment, Protection
from copy import copy
from decimal import Decimal, ROUND_HALF_UP
from datetime import datetime
import csv, json

ws = Path.cwd()
infile = ws/'input'/'sales_template.xlsx'
outdir = ws/'output'
outdir.mkdir(exist_ok=True)
outfile = outdir/'sales_template.xlsx'

# Data collected via browser detail pages
headers = {
'0045000109': {'OrderID': '0045000109', 'CustomerID': 'C001', 'CustomerName': 'Acme Precision Taiwan Ltd.', 'ProjectID': 'P003', 'OrderDate': '2026-09-01', 'Status': 'Open', 'Currency': 'TWD', 'NetAmount': '150000.00', 'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET10', 'SalesRep': 'E001'},
'0045000110': {'OrderID': '0045000110', 'CustomerID': 'C002', 'CustomerName': '台灣精密股份有限公司', 'ProjectID': 'P002', 'OrderDate': '2026-09-02', 'Status': 'Open', 'Currency': 'TWD', 'NetAmount': '293572.50', 'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET30', 'SalesRep': 'E004'},
'0045000112': {'OrderID': '0045000112', 'CustomerID': 'C007', 'CustomerName': 'Acme Precision Taiwan Limited', 'ProjectID': 'P005', 'OrderDate': '2026-09-04', 'Status': 'Open', 'Currency': 'USD', 'NetAmount': '12311.00', 'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET15', 'SalesRep': 'E001'},
'0045000113': {'OrderID': '0045000113', 'CustomerID': 'C004', 'CustomerName': '藍海電子有限公司', 'ProjectID': '', 'OrderDate': '2026-09-05', 'Status': 'Open', 'Currency': 'TWD', 'NetAmount': '15390.00', 'CustomerContact': '', 'PaymentTerms': 'NET15', 'SalesRep': 'E005'},
'0045000115': {'OrderID': '0045000115', 'CustomerID': 'C005', 'CustomerName': 'Pacific Test Systems Inc.', 'ProjectID': '', 'OrderDate': '2026-09-09', 'Status': 'Open', 'Currency': 'USD', 'NetAmount': '6391.00', 'CustomerContact': 'Sam Patel <sam@pacifictest.example>', 'PaymentTerms': 'NET20', 'SalesRep': 'E004'},
'0045000116': {'OrderID': '0045000116', 'CustomerID': 'C008', 'CustomerName': '新竹量測科技', 'ProjectID': '', 'OrderDate': '2026-09-10', 'Status': 'Open', 'Currency': 'TWD', 'NetAmount': '29400.00', 'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30', 'SalesRep': 'E005'},
'0045000118': {'OrderID': '0045000118', 'CustomerID': 'C003', 'CustomerName': 'Northwind Components GmbH', 'ProjectID': 'P004', 'OrderDate': '2026-09-14', 'Status': 'Open', 'Currency': 'EUR', 'NetAmount': '17484.75', 'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET0', 'SalesRep': 'E006'},
'0045000119': {'OrderID': '0045000119', 'CustomerID': 'C001', 'CustomerName': 'Acme Precision Taiwan Ltd.', 'ProjectID': 'P001', 'OrderDate': '2026-09-15', 'Status': 'Open', 'Currency': 'EUR', 'NetAmount': '8775.00', 'CustomerContact': 'Grace Lin <grace.lin@acme-precision.example>', 'PaymentTerms': 'NET30', 'SalesRep': 'E004'},
'0045000120': {'OrderID': '0045000120', 'CustomerID': 'C006', 'CustomerName': 'Helios Optics B.V.', 'ProjectID': '', 'OrderDate': '2026-09-16', 'Status': 'Open', 'Currency': 'EUR', 'NetAmount': '461.70', 'CustomerContact': 'Eva de Vries <eva@helios.example>', 'PaymentTerms': 'NET10', 'SalesRep': 'E003'},
'0045000122': {'OrderID': '0045000122', 'CustomerID': 'C002', 'CustomerName': '台灣精密股份有限公司', 'ProjectID': 'P002', 'OrderDate': '2026-09-21', 'Status': 'Open', 'Currency': 'USD', 'NetAmount': '4141.00', 'CustomerContact': '張美玲 <mei@tw-precision.example>', 'PaymentTerms': 'NET5', 'SalesRep': 'E006'},
'0045000124': {'OrderID': '0045000124', 'CustomerID': 'C004', 'CustomerName': '藍海電子有限公司', 'ProjectID': '', 'OrderDate': '2026-09-24', 'Status': 'Open', 'Currency': 'TWD', 'NetAmount': '45000.00', 'CustomerContact': '', 'PaymentTerms': 'NET30', 'SalesRep': 'E003'},
'0045000126': {'OrderID': '0045000126', 'CustomerID': 'C008', 'CustomerName': '新竹量測科技', 'ProjectID': '', 'OrderDate': '2026-09-28', 'Status': 'Open', 'Currency': 'USD', 'NetAmount': '9171.00', 'CustomerContact': '李建宏 <jh.lee@hsinchu-metrology.example>', 'PaymentTerms': 'NET30', 'SalesRep': 'E003'},
'0045000127': {'OrderID': '0045000127', 'CustomerID': 'C007', 'CustomerName': 'Acme Precision Taiwan Limited', 'ProjectID': 'P005', 'OrderDate': '2026-09-30', 'Status': 'Open', 'Currency': 'EUR', 'NetAmount': '2214.92', 'CustomerContact': 'Peter Ho <peter.ho@acme-limited.example>', 'PaymentTerms': 'NET30', 'SalesRep': 'E006'},
'0045000128': {'OrderID': '0045000128', 'CustomerID': 'C003', 'CustomerName': 'Northwind Components GmbH', 'ProjectID': 'P004', 'OrderDate': '2026-09-30', 'Status': 'Open', 'Currency': 'TWD', 'NetAmount': '104160.00', 'CustomerContact': 'Jonas Krüger <jk@northwind.example>', 'PaymentTerms': 'NET5', 'SalesRep': 'E006'},
}

# Read exported order list and reconcile order IDs
csvp = ws/'output'/'downloads'/'open_orders_2026_09.csv'
with open(csvp, newline='', encoding='utf-8-sig') as f:
    order_rows = list(csv.DictReader(f))
order_ids = [r['OrderID'] for r in order_rows]
assert len(order_ids) == len(set(order_ids)) == 14, order_ids
assert set(order_ids) == set(headers), (order_ids, sorted(headers))

# Read FX rates (only first two columns needed)
rates = {}
with open(ws/'input'/'fx_rates.csv', newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])
print('rates', rates)

q = Decimal('0.01')
def money(x):
    return Decimal(str(x)).quantize(q, rounding=ROUND_HALF_UP)

def twd(amount, currency):
    return (Decimal(str(amount)) * rates[currency]).quantize(q, rounding=ROUND_HALF_UP)

wb = load_workbook(infile)
report = wb['Report']
details = wb['Details']
summary = wb['Summary']

# Fill report metadata
fetched_at = '2026-10-06T00:00:00+08:00'
report['B3'] = '2026-09-01'
report['D3'] = '2026-09-30'
report['B4'] = 'Sales Orders: Status=Open; OrderDate 2026-09-01..2026-09-30; sorted by OrderID asc'
report['B5'] = fetched_at
report['B7'] = 'Lab ERP (synthetic)'
report['D7'] = fetched_at
report['B9'] = 'Recovered once from session timeout; all 14 exported Open orders reconciled with detail headers. Blank ProjectID/CustomerContact cells reflect source blanks.'

# Clear possible old data rows while preserving sheets/headers
for sh, maxc in [(details, 11), (summary, 6)]:
    for row in range(5, sh.max_row + 1):
        for col in range(1, maxc+1):
            sh.cell(row, col).value = None

# Details rows
columns = ['OrderID','CustomerID','CustomerName','ProjectID','OrderDate','Status','Currency','NetAmount','NetAmountTWD','CustomerContact','PaymentTerms']
for idx, oid in enumerate(order_ids, start=5):
    h = headers[oid]
    vals = [
        h['OrderID'], h['CustomerID'], h['CustomerName'], h.get('ProjectID') or None,
        h['OrderDate'], h['Status'], h['Currency'], float(money(h['NetAmount'])),
        float(twd(h['NetAmount'], h['Currency'])), h.get('CustomerContact') or None, h.get('PaymentTerms') or None
    ]
    for col, val in enumerate(vals, start=1):
        cell = details.cell(idx, col)
        cell.value = val
        # copy style from row 4 as conservative template-like styling
        src = details.cell(4, col)
        if src.has_style:
            cell.font = copy(src.font); cell.fill = copy(src.fill); cell.border = copy(src.border)
            cell.alignment = copy(src.alignment); cell.protection = copy(src.protection); cell.number_format = src.number_format
    for col in [1,2,4]:
        details.cell(idx, col).number_format = '@'
    details.cell(idx, 5).number_format = 'yyyy-mm-dd'
    for col in [8,9]:
        details.cell(idx, col).number_format = '0.00'

last_row = 4 + len(order_ids)
summary['B2'] = f'=SUM(Details!I5:I{last_row})'

# Summary by Status + Currency (no cross-currency original amount mixing)
groups = {}
for oid in order_ids:
    h = headers[oid]
    key = (h['Status'], h['Currency'])
    groups.setdefault(key, {'amount': Decimal('0'), 'twd': Decimal('0'), 'count': 0})
    groups[key]['amount'] += Decimal(h['NetAmount'])
    groups[key]['twd'] += twd(h['NetAmount'], h['Currency'])
    groups[key]['count'] += 1
for ridx, ((status, currency), g) in enumerate(sorted(groups.items(), key=lambda kv: kv[0][1]), start=5):
    vals = [status, '', currency, float(g['amount'].quantize(q, rounding=ROUND_HALF_UP)), float(g['twd'].quantize(q, rounding=ROUND_HALF_UP)), g['count']]
    for col, val in enumerate(vals, start=1):
        cell = summary.cell(ridx, col)
        cell.value = val
        src = summary.cell(4, col)
        if src.has_style:
            cell.font = copy(src.font); cell.fill = copy(src.fill); cell.border = copy(src.border)
            cell.alignment = copy(src.alignment); cell.protection = copy(src.protection); cell.number_format = src.number_format
    summary.cell(ridx, 4).number_format = '0.00'
    summary.cell(ridx, 5).number_format = '0.00'

wb.save(outfile)
print('wrote', outfile)

# Verification by reopening output
wb2 = load_workbook(outfile, data_only=False)
rep2, det2, sum2 = wb2['Report'], wb2['Details'], wb2['Summary']
headers_row = [det2.cell(4,c).value for c in range(1,12)]
print('details headers', headers_row)
assert headers_row == columns
read_rows = []
for r in range(5, 5+len(order_ids)):
    row = [det2.cell(r,c).value for c in range(1,12)]
    read_rows.append(row)
# no extra data row immediately after
assert all(det2.cell(5+len(order_ids), c).value is None for c in range(1,12))
assert len(read_rows) == 14
assert [r[0] for r in read_rows] == order_ids
assert all(isinstance(r[0], str) and len(r[0]) == 10 and r[0].startswith('00') for r in read_rows)
assert all(r[5] == 'Open' for r in read_rows)
assert all('2026-09-01' <= r[4] <= '2026-09-30' for r in read_rows)
assert all(isinstance(r[7], (int,float)) and isinstance(r[8], (int,float)) for r in read_rows)
# Check exact amount TWD values
for row in read_rows:
    expected = float(twd(Decimal(str(row[7])), row[6]))
    assert round(row[8],2) == round(expected,2), (row[0], row[8], expected)
assert sum2['B2'].value == '=SUM(Details!I5:I18)'
summary_rows = []
for r in range(5, 5+len(groups)):
    summary_rows.append([sum2.cell(r,c).value for c in range(1,7)])
print('summary rows', summary_rows)
assert sum(x[5] for x in summary_rows) == 14
# Formula won't calculate in openpyxl; verify range and expected grand total manually
expected_total = sum(Decimal(str(r[8])) for r in read_rows).quantize(q)
print('expected NetAmountTWD total', expected_total)
print('report range', rep2['B3'].value, rep2['D3'].value, 'condition', rep2['B4'].value)
print('row_count_checked', len(read_rows))
print('output_file', outfile)
