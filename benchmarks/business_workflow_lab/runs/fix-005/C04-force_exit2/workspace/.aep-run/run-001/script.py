from pathlib import Path
import csv
from openpyxl import load_workbook
ws = Path.cwd()
print('WORKSPACE', ws)
print('INPUT FILES')
for p in sorted((ws/'input').glob('*')):
    print('-', p.name, p.stat().st_size, 'bytes')

# CSV headers/sample
csv_path = ws/'input'/'fx_rates.csv'
if csv_path.exists():
    print('\nFX CSV')
    with csv_path.open(newline='', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for i, row in zip(range(5), reader):
            print(i, row)

# Workbook sheets and dimensions, sample cells
xlsx = ws/'input'/'quote_template.xlsx'
if xlsx.exists():
    print('\nWORKBOOK')
    wb = load_workbook(xlsx, data_only=False)
    print('sheets:', wb.sheetnames)
    for sh in wb.worksheets:
        print('sheet', sh.title, 'max_row', sh.max_row, 'max_col', sh.max_column)
        for r in sh.iter_rows(min_row=1, max_row=min(sh.max_row, 12), values_only=True):
            print(r)
        print('---')

# Manual excerpt
manual = ws/'input'/'quote_cli_manual.md'
if manual.exists():
    print('\nMANUAL FIRST 200 LINES')
    text = manual.read_text(encoding='utf-8', errors='replace').splitlines()
    for i, line in enumerate(text[:200], 1):
        print(f'{i:03d}: {line}')
