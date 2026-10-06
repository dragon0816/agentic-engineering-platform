from pathlib import Path
import subprocess, sys, json, csv, os
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from datetime import datetime, timezone, timedelta
from collections import defaultdict
from openpyxl import load_workbook
from openpyxl.styles import numbers

WORKSPACE = Path.cwd()
INPUT = WORKSPACE / 'input'
OUTPUT = WORKSPACE / 'output'
OUTPUT.mkdir(exist_ok=True)

FROM_DATE = '2026-09-01'
TO_DATE = '2026-09-30'
STATUS = 'Sent'
CLI = Path(r'C:\Users\Chi_L\Documents\ChatGPT\agentic-engineering-platform\benchmarks\business_workflow_lab\quote_cli_mock\quote_cli.cmd')
TEMPLATE = INPUT / 'quote_template.xlsx'
FX = INPUT / 'fx_rates.csv'
OUT_XLSX = OUTPUT / 'quote_report_2026-09_sent.xlsx'
RUN_LOG = OUTPUT / 'quote_cli_run_log.json'
ERROR_TXT = OUTPUT / 'quote_cli_error.txt'

TAIPEI = timezone(timedelta(hours=8))
now_iso = datetime.now(TAIPEI).replace(microsecond=0).isoformat()

def write_error(message, log=None):
    ERROR_TXT.write_text(message + '\n', encoding='utf-8')
    if log is not None:
        RUN_LOG.write_text(json.dumps(log, ensure_ascii=False, indent=2), encoding='utf-8')
    print('ERROR_WRITTEN', ERROR_TXT)
    print(message)


def run_cli(args, label):
    attempts = []
    max_attempts = 3  # initial + up to two retries for exit code 4
    for attempt in range(1, max_attempts + 1):
        cmd = [str(CLI)] + args
        cp = subprocess.run(cmd, capture_output=True, text=True, shell=False)
        rec = {
            'label': label,
            'attempt': attempt,
            'command': cmd,
            'exit_code': cp.returncode,
            'stdout_preview': cp.stdout[:500],
            'stderr': cp.stderr,
        }
        attempts.append(rec)
        if cp.returncode == 0:
            try:
                data = json.loads(cp.stdout)
            except Exception as e:
                return None, attempts, f'{label}: exit code 0 but stdout is not parseable JSON: {e}'
            return data, attempts, None
        if cp.returncode in (2, 3):
            return None, attempts, f'{label}: quote_cli returned exit code {cp.returncode}; stderr: {cp.stderr.strip()}'
        if cp.returncode == 4:
            if attempt < max_attempts:
                continue
            return None, attempts, f'{label}: quote_cli returned temporary service error exit code 4 after {max_attempts} attempts; stderr: {cp.stderr.strip()}'
        return None, attempts, f'{label}: quote_cli returned unexpected exit code {cp.returncode}; stderr: {cp.stderr.strip()}'
    return None, attempts, f'{label}: unknown retry loop failure'

# Clean previous generated files from this task so final state is unambiguous.
for p in [OUT_XLSX, RUN_LOG, ERROR_TXT]:
    if p.exists():
        p.unlink()

print('WORKSPACE', WORKSPACE)
print('TEMPLATE_EXISTS', TEMPLATE.exists(), TEMPLATE)
print('FX_EXISTS', FX.exists(), FX)
print('CLI_EXISTS', CLI.exists(), CLI)

all_attempts = []
list_args = ['quotes', 'list', '--from', FROM_DATE, '--to', TO_DATE, '--status', STATUS, '--format', 'json']
list_data, attempts, err = run_cli(list_args, 'quotes list Sent Sep 2026')
all_attempts.extend(attempts)
if err:
    log = {'generated_at': now_iso, 'status': 'failed', 'attempts': all_attempts, 'error': err}
    write_error(err, log)
    raise SystemExit(0)

if not isinstance(list_data, dict) or list_data.get('schema_version') != 1 or not isinstance(list_data.get('items'), list):
    err = 'quotes list: parsed stdout does not match expected schema_version=1 with items list'
    log = {'generated_at': now_iso, 'status': 'failed', 'attempts': all_attempts, 'error': err, 'list_data_type': str(type(list_data))}
    write_error(err, log)
    raise SystemExit(0)

headers = list_data['items']
# Fetch line items for each quote.
details = []
for item in headers:
    qid = str(item.get('QuoteID', ''))
    if not qid:
        err = 'quotes list item missing QuoteID; cannot fetch detail safely'
        log = {'generated_at': now_iso, 'status': 'failed', 'attempts': all_attempts, 'error': err, 'bad_item': item}
        write_error(err, log)
        raise SystemExit(0)
    data, attempts, err = run_cli(['quotes', 'get', qid, '--format', 'json'], f'quotes get {qid}')
    all_attempts.extend(attempts)
    if err:
        log = {'generated_at': now_iso, 'status': 'failed', 'attempts': all_attempts, 'error': err}
        write_error(err, log)
        raise SystemExit(0)
    if not isinstance(data, dict) or data.get('schema_version') != 1 or not isinstance(data.get('header'), dict) or not isinstance(data.get('line_items'), list):
        err = f'quotes get {qid}: parsed stdout does not match expected schema_version=1 with header and line_items'
        log = {'generated_at': now_iso, 'status': 'failed', 'attempts': all_attempts, 'error': err}
        write_error(err, log)
        raise SystemExit(0)
    details.append(data)

# Load rates; keep first two columns because Note may contain unquoted commas.
rates = {}
with FX.open(newline='', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    header = next(reader)
    for row in reader:
        if len(row) >= 2 and row[0]:
            rates[row[0]] = Decimal(row[1])

# Prepare workbook.
wb = load_workbook(TEMPLATE)
ws_report = wb['Report']
ws_quotes = wb['Quotes']
ws_lines = wb['Lines']
ws_summary = wb['Summary']

# Clear detail rows from row 5 down.
for ws in [ws_quotes, ws_lines, ws_summary]:
    if ws.max_row >= 5:
        ws.delete_rows(5, ws.max_row - 4)

# Fill report metadata.
ws_report['B3'] = FROM_DATE
ws_report['D3'] = TO_DATE
ws_report['B4'] = f'Status = {STATUS}; QuoteDate from {FROM_DATE} to {TO_DATE} inclusive'
ws_report['B5'] = now_iso
ws_report['B7'] = 'quote_cli; fx_rates.csv'
ws_report['D7'] = now_iso
ws_report['B9'] = 'CLI stdout used only after exit code 0 and JSON parse success; exit code 4 retried up to two times.'

money_fmt = '0.00'
quote_headers = [c.value for c in ws_quotes[4]]
line_headers = [c.value for c in ws_lines[4]]
summary_headers = [c.value for c in ws_summary[4]]

quote_row = 5
line_row = 5
summary = defaultdict(lambda: {'amount': Decimal('0.00'), 'amount_twd': Decimal('0.00'), 'count': 0})

D2 = Decimal('0.01')
for detail in details:
    h = detail['header']
    currency = str(h.get('Currency', ''))
    if currency not in rates:
        err = f'Missing FX rate for currency {currency}'
        log = {'generated_at': now_iso, 'status': 'failed', 'attempts': all_attempts, 'error': err}
        write_error(err, log)
        raise SystemExit(0)
    net = Decimal(str(h.get('NetAmount'))).quantize(D2, rounding=ROUND_HALF_UP)
    net_twd = (net * rates[currency]).quantize(D2, rounding=ROUND_HALF_UP)
    row_values = {
        'QuoteID': str(h.get('QuoteID', '')),
        'CustomerID': str(h.get('CustomerID', '')),
        'ProjectID': str(h.get('ProjectID', '')),
        'QuoteDate': str(h.get('QuoteDate', '')),
        'ValidUntil': str(h.get('ValidUntil', '')),
        'Status': str(h.get('Status', '')),
        'Currency': currency,
        'Subtotal': Decimal(str(h.get('Subtotal'))).quantize(D2, rounding=ROUND_HALF_UP),
        'DiscountPercent': Decimal(str(h.get('DiscountPercent'))),
        'NetAmount': net,
        'NetAmountTWD': net_twd,
    }
    for col_idx, col_name in enumerate(quote_headers, 1):
        val = row_values.get(col_name)
        cell = ws_quotes.cell(quote_row, col_idx, val)
        if col_name in ('QuoteID', 'CustomerID', 'ProjectID'):
            cell.number_format = '@'
        if col_name in ('Subtotal', 'DiscountPercent', 'NetAmount', 'NetAmountTWD'):
            cell.value = float(val)
            cell.number_format = money_fmt
    summary[('Status', row_values['Status'], currency)]['amount'] += net
    summary[('Status', row_values['Status'], currency)]['amount_twd'] += net_twd
    summary[('Status', row_values['Status'], currency)]['count'] += 1
    summary[('Currency', currency, currency)]['amount'] += net
    summary[('Currency', currency, currency)]['amount_twd'] += net_twd
    summary[('Currency', currency, currency)]['count'] += 1
    quote_row += 1

    for li in detail['line_items']:
        line_values = {
            'QuoteID': str(h.get('QuoteID', '')),
            'ProductID': str(li.get('ProductID', '')),
            'ProductName': str(li.get('ProductName', '')),
            'Quantity': Decimal(str(li.get('Quantity'))),
            'UnitPrice': Decimal(str(li.get('UnitPrice'))).quantize(D2, rounding=ROUND_HALF_UP),
            'LineAmount': Decimal(str(li.get('LineAmount'))).quantize(D2, rounding=ROUND_HALF_UP),
        }
        for col_idx, col_name in enumerate(line_headers, 1):
            val = line_values.get(col_name)
            cell = ws_lines.cell(line_row, col_idx, val)
            if col_name in ('QuoteID', 'ProductID'):
                cell.number_format = '@'
            if col_name in ('Quantity', 'UnitPrice', 'LineAmount'):
                cell.value = float(val)
                cell.number_format = money_fmt if col_name != 'Quantity' else '0.########'
        line_row += 1

summary_row = 5
for key in sorted(summary.keys()):
    group, subgroup, currency = key
    vals = {
        'Group': group,
        'SubGroup': subgroup,
        'Currency': currency,
        'Amount': summary[key]['amount'].quantize(D2, rounding=ROUND_HALF_UP),
        'AmountTWD': summary[key]['amount_twd'].quantize(D2, rounding=ROUND_HALF_UP),
        'Count': summary[key]['count'],
    }
    for col_idx, col_name in enumerate(summary_headers, 1):
        val = vals.get(col_name)
        cell = ws_summary.cell(summary_row, col_idx, val)
        if col_name in ('Amount', 'AmountTWD'):
            cell.value = float(val)
            cell.number_format = money_fmt
        elif col_name == 'Count':
            cell.value = int(val)
    summary_row += 1

# Fix total formula range to actual rows; if zero rows use 0.
last_quote_row = quote_row - 1
if last_quote_row >= 5:
    ws_summary['B2'] = f'=SUM(Quotes!K5:K{last_quote_row})'
else:
    ws_summary['B2'] = 0
ws_summary['B2'].number_format = money_fmt

wb.save(OUT_XLSX)
RUN_LOG.write_text(json.dumps({
    'generated_at': now_iso,
    'status': 'success',
    'query': {'from': FROM_DATE, 'to': TO_DATE, 'status': STATUS},
    'quote_count': len(headers),
    'detail_count': len(details),
    'line_count': line_row - 5,
    'attempts': all_attempts,
}, ensure_ascii=False, indent=2), encoding='utf-8')
print('OUTPUT_XLSX', OUT_XLSX)
print('RUN_LOG', RUN_LOG)
print('QUOTE_COUNT', len(headers))
print('LINE_COUNT', line_row - 5)
