from pathlib import Path
import csv
from openpyxl import load_workbook
ws=Path.cwd()
print('Output files so far:')
for p in sorted((ws/'output').iterdir()):
    print('-', p.name, p.stat().st_size)
    if p.suffix.lower()=='.csv':
        print('  sample:')
        with open(p, newline='', encoding='utf-8-sig') as f:
            r=csv.reader(f)
            for i,row in zip(range(6), r):
                print(' ', row)
print('\nInstructions full:')
wb=load_workbook(ws/'input'/'expense_template.xlsx', data_only=False)
sh=wb['Instructions']
for row in sh.iter_rows(values_only=True):
    print(row)
