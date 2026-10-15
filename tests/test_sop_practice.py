"""The SOP to Workflow practice target (SOP_TO_WORKFLOW W1)."""

import json
import subprocess
import sys
import urllib.parse
import urllib.request
from pathlib import Path

import pytest

from evaluation.sop_practice import (
    HEADERS,
    SHEET_NAME,
    PracticeSite,
    write_sop_pdf,
    write_workbook,
)


def fetch(url: str, data: dict[str, str] | None = None) -> str:
    body = urllib.parse.urlencode(data).encode() if data is not None else None
    with urllib.request.urlopen(url, data=body, timeout=10) as reply:
        page: bytes = reply.read()
        return page.decode("utf-8")


def test_the_site_searches_lists_and_records_a_notification() -> None:
    site = PracticeSite()
    site.start()
    try:
        assert "客戶代碼" in fetch(site.url + "/")
        listing = fetch(site.url + "/orders?customer=c001")
        assert listing.count("<tr>") == 5 and "SO-1004" in listing and "SO-2001" not in listing
        assert listing.count("標記已通知") == 2
        after = fetch(site.url + "/notify", {"order": "SO-1002"})
        assert after.count("標記已通知") == 1 and "已通知" in after
        assert json.loads(fetch(site.url + "/state.json")) == {"notified": ["SO-1002"]}
        assert "查無訂單" in fetch(site.url + "/orders?customer=C999")
    finally:
        site.close()


def test_the_workbook_holds_a_stale_and_an_unrelated_order(tmp_path: Path) -> None:
    from openpyxl import load_workbook

    sheet = load_workbook(write_workbook(tmp_path / "book.xlsx"))[SHEET_NAME]
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == (*HEADERS, "備註")
    assert [row[0] for row in rows[1:]] == ["SO-1001", "SO-0999"]


@pytest.mark.skipif(sys.platform != "win32", reason="prints with the installed Edge")
def test_the_sop_pdf_carries_the_routine_as_text(tmp_path: Path) -> None:
    from pypdf import PdfReader

    from integrations.browser import browser_path

    if browser_path() is None:
        pytest.skip("no Chromium browser installed")
    from capabilities.runtime import InstalledCapabilities
    from capabilities.workflow_author.handlers import DraftWorkflowHandler

    pdf = write_sop_pdf(tmp_path / "sop.pdf")
    raw = "".join(page.extract_text() or "" for page in PdfReader(pdf).pages)
    # Found here: Edge's PDF makes pypdf return 客⼾ (a Kangxi radical) for 客戶.
    assert "⼾" in raw and "客戶代碼" not in raw
    author = DraftWorkflowHandler(
        InstalledCapabilities(), model=None, model_alias="none", workspace_root=tmp_path
    )
    text = author._read_pdf(str(pdf))
    assert text is not None
    for phrase in ("客戶代碼", "出貨追蹤.xlsx", "標記已通知", "訂單號"):
        assert phrase in text, phrase


#: Stands in for headless Edge as measured on 2026-10-07: it writes the PDF
#: (here in two halves, so a half-written file is seen) and then never exits.
_PRINTS_THEN_HANGS = r"""
import sys, time
target = next(a.split("=", 1)[1] for a in sys.argv if a.startswith("--print-to-pdf="))
body = b"%PDF-1.4\n" + b"x" * 4096 + b"\n%%EOF\n"
with open(target, "wb") as out:
    out.write(body[:100]); out.flush(); time.sleep(1.0); out.write(body[100:])
time.sleep(600)
"""
_FAKE_PDF = b"%PDF-1.4\n" + b"x" * 4096 + b"\n%%EOF\n"


def _fake_browser(monkeypatch: pytest.MonkeyPatch, script: str) -> list["subprocess.Popen[bytes]"]:
    """Run ``script`` where the browser would start; anything else (taskkill)
    starts for real."""
    started: list[subprocess.Popen[bytes]] = []
    real_popen = subprocess.Popen

    def popen(arguments: list[str], **options: object) -> subprocess.Popen[bytes]:
        if arguments[0] == "fake-edge":
            arguments = [sys.executable, "-c", script, *arguments[1:]]
        process: subprocess.Popen[bytes] = real_popen(arguments, **options)  # type: ignore[call-overload]
        started.append(process)
        return process

    monkeypatch.setattr(subprocess, "Popen", popen)  # the module sop_practice uses
    return started


def test_a_browser_that_never_exits_after_printing_is_stopped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import time

    started = _fake_browser(monkeypatch, _PRINTS_THEN_HANGS)
    began = time.monotonic()

    pdf = write_sop_pdf(tmp_path / "out" / "sop.pdf", browser="fake-edge")

    assert time.monotonic() - began < 30, "waited on the process instead of the PDF"
    assert pdf.read_bytes() == _FAKE_PDF
    assert started[0].poll() is not None, "the hung browser was left running"


def test_a_browser_that_exits_without_a_pdf_fails_plainly(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    import evaluation.sop_practice as practice

    _fake_browser(monkeypatch, "pass")
    monkeypatch.setattr(practice, "PRINT_AFTER_EXIT_SECONDS", 0.5)

    with pytest.raises(RuntimeError, match="did not write the SOP PDF"):
        write_sop_pdf(tmp_path / "sop.pdf", browser="fake-edge")
    assert not (tmp_path / "sop.pdf").exists()
