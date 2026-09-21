"""Office extraction: a generated PDF, PPTX and DOCX corpus round-trips to Raw
with its page/slide/image relationships. Needs the `office` extra."""

import io
import struct
import zlib
from datetime import date
from pathlib import Path

import pytest
from test_raw import drop, make_vault

from knowledge.office import DocxExtractor, PdfExtractor, PptxExtractor, table_markdown
from knowledge.raw import DropIntake, ExtractionError, StagedAssets, parse

pytest.importorskip("pypdf")
pytest.importorskip("pptx")
pytest.importorskip("docx")

TODAY = date(2026, 9, 21)


def make_png() -> bytes:
    """A valid 1x1 red PNG built from the spec, so the corpus needs no fixture file."""

    def chunk(kind: bytes, body: bytes) -> bytes:
        return (
            struct.pack(">I", len(body))
            + kind
            + body
            + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)
        )

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    pixels = zlib.compress(b"\x00\xff\x00\x00")
    return (
        b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", pixels) + chunk(b"IEND", b"")
    )


PNG = make_png()


def make_pdf(pages: list[str], *, image: bool = False, broken_image: bool = False) -> bytes:
    """A minimal PDF with one line of Helvetica text per page and, optionally,
    a 1x1 RGB image on the first page; offsets are computed so the xref is real."""
    objects: list[bytes] = []
    kids = []
    # Objects 3.. are page/content pairs; the font and the image follow them.
    font = 3 + 2 * len(pages)
    image_object = font + 1
    for index, text in enumerate(pages):
        page_number = 3 + 2 * index
        content_number = page_number + 1
        kids.append(f"{page_number} 0 R")
        resources = f"/Font << /F1 {font} 0 R >>"
        stream = f"BT /F1 12 Tf 20 100 Td ({text}) Tj ET"
        if image and index == 0:
            resources += f" /XObject << /Im1 {image_object} 0 R >>"
            stream += "\nq 10 0 0 10 20 20 cm /Im1 Do Q"
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 200 200] /Contents {content_number} 0 R"
            f" /Resources << {resources} >> >>".encode()
        )
        objects.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    if image:
        objects.append(
            b"<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceRGB"
            b" /BitsPerComponent 8"
            + (b" /Filter /JBIG2Decode" if broken_image else b"")
            + b" /Length 3 >>\nstream\n\xff\x00\x00\nendstream"
        )
    header = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(pages)} >>".encode(),
    ]
    body = b"%PDF-1.4\n"
    offsets = []
    for number, obj in enumerate(header + objects, 1):
        offsets.append(len(body))
        body += f"{number} 0 obj\n".encode() + obj + b"\nendobj\n"
    xref = len(body)
    body += f"xref\n0 {len(offsets) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        body += f"{offset:010d} 00000 n \n".encode()
    body += (
        f"trailer\n<< /Size {len(offsets) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return body


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    blank = deck.slide_layouts[6]
    first = deck.slides.add_slide(blank)
    box = first.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Slide one title"
    first.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(2))
    second = deck.slides.add_slide(blank)
    second.shapes.add_textbox(
        Inches(1), Inches(1), Inches(4), Inches(1)
    ).text_frame.text = "Before the table"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    table.cell(0, 0).text, table.cell(0, 1).text = "Channel", "Rate | unit"
    table.cell(1, 0).text, table.cell(1, 1).text = "A", "1.5"
    second.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(4))
    out = io.BytesIO()
    deck.save(out)
    return out.getvalue()


def make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Report title", level=1)
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Name", "Value"
    table.cell(1, 0).text, table.cell(1, 1).text = "x", "42"
    document.add_paragraph("After the table.")
    document.add_picture(io.BytesIO(PNG))
    document.add_heading("Closing", level=2)
    document.add_paragraph("Done.")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_table_markdown_flattens_cells_and_escapes_pipes() -> None:
    assert table_markdown([["a", "b | c"], ["1"]]) == "| a | b \\| c |\n| --- | --- |\n| 1 |  |"
    assert (
        table_markdown([["multi\nline", " spaced   out "]])
        == "| multi line | spaced out |\n| --- | --- |"
    )
    assert table_markdown([]) == ""


def test_pdf_keeps_text_and_images_by_page() -> None:
    staged = StagedAssets("raw/report/assets")
    sections = PdfExtractor().extract(
        make_pdf(["Hello page one", "Second page"], image=True), staged
    )
    assert [(s.kind, s.page) for s in sections] == [("text", 1), ("image", 1), ("text", 2)]
    assert sections[0].text == "Hello page one" and sections[2].text == "Second page"
    assert sections[1].image_ref is not None and sections[1].image_ref.startswith(
        "raw/report/assets/"
    )
    assert staged.items[sections[1].image_ref].startswith(b"\x89PNG")
    assert PdfExtractor().extract(make_pdf([" "]), StagedAssets("raw/x/assets")) == ()
    with pytest.raises(ExtractionError) as raised:
        PdfExtractor().extract(b"%PDF-1.4 garbage", StagedAssets("raw/x/assets"))
    assert raised.value.code == "undecodable"


def test_pptx_keeps_text_tables_and_pictures_by_slide() -> None:
    staged = StagedAssets("raw/deck/assets")
    sections = PptxExtractor().extract(make_pptx(), staged)
    assert [(s.kind, s.slide) for s in sections] == [
        ("text", 1),
        ("image", 1),
        ("text", 2),
        ("table", 2),
        ("image", 2),
    ]
    assert sections[0].text == "Slide one title" and sections[2].text == "Before the table"
    assert sections[3].text == "| Channel | Rate \\| unit |\n| --- | --- |\n| A | 1.5 |"
    # The same picture on two slides is one content-named asset, referenced twice.
    assert sections[1].image_ref == sections[4].image_ref and len(staged.items) == 1
    with pytest.raises(ExtractionError, match="undecodable"):
        PptxExtractor().extract(b"PK\x03\x04 not a deck", staged)


def test_docx_keeps_order_headings_tables_and_pictures() -> None:
    staged = StagedAssets("raw/doc/assets")
    sections = DocxExtractor().extract(make_docx(), staged)
    assert [s.kind for s in sections] == ["text", "table", "text", "image", "text"]
    assert sections[0].text == "# Report title\n\nFirst paragraph.\n\nSecond paragraph."
    assert sections[1].text == "| Name | Value |\n| --- | --- |\n| x | 42 |"
    assert sections[2].text == "After the table."
    assert sections[4].text == "## Closing\n\nDone."
    assert all(s.page is None and s.slide is None for s in sections)
    with pytest.raises(ExtractionError, match="undecodable"):
        DocxExtractor().extract(b"not a document", staged)


def test_the_corpus_round_trips_to_raw_with_assets(tmp_path: Path) -> None:
    vault = make_vault(tmp_path)
    intake = DropIntake(vault, [PdfExtractor(), PptxExtractor(), DocxExtractor()])
    drop(vault, "drop/q3/report.pdf", make_pdf(["Quarterly numbers"], image=True))
    dropped = make_pptx()
    drop(vault, "drop/q3/deck.pptx", dropped)
    drop(vault, "drop/q3/memo.docx", make_docx())
    preview = intake.intake_all(
        ["drop/q3/report.pdf", "drop/q3/deck.pptx", "drop/q3/memo.docx"], today=TODAY
    )
    assert [o.status for o in preview] == ["written"] * 3
    assert preview[0].written[0] == "raw/q3/report.md"
    assert (
        preview[0].written[1].startswith("raw/q3/report/assets/") and len(preview[0].written) == 2
    )
    assert not (tmp_path / "raw" / "q3").exists()

    applied = intake.intake_all(
        ["drop/q3/report.pdf", "drop/q3/deck.pptx", "drop/q3/memo.docx"], mode="apply", today=TODAY
    )
    assert [o.written for o in applied] == [o.written for o in preview]
    deck = parse(vault.read("raw/q3/deck.md"))
    assert [(s.kind, s.slide) for s in deck.sections][:2] == [("text", 1), ("image", 1)]
    image = deck.sections[1].image_ref
    assert image is not None and (tmp_path / image).read_bytes() == PNG
    assert vault.raw_files() == ("raw/q3/deck.md", "raw/q3/memo.md", "raw/q3/report.md")
    # Assets are write-once too: the same bytes again are a no-op, other bytes refused.
    assert vault.write_raw_bytes(image, PNG) is False
    with pytest.raises(Exception, match="raw_exists"):
        vault.write_raw_bytes(image, b"other")
    assert (tmp_path / "drop" / "q3" / "deck.pptx").read_bytes() == dropped


def test_an_unreadable_image_does_not_lose_the_pages_text() -> None:
    staged = StagedAssets("raw/report/assets")
    sections = PdfExtractor().extract(
        make_pdf(["Text survives", "Second page"], image=True, broken_image=True), staged
    )
    assert [(s.kind, s.page, s.text) for s in sections] == [
        ("text", 1, "Text survives"),
        ("text", 2, "Second page"),
    ]
    assert staged.items == {}


def test_pictures_in_placeholders_and_groups_are_kept() -> None:
    from pptx import Presentation
    from pptx.enum.shapes import PP_PLACEHOLDER
    from pptx.util import Inches

    deck = Presentation()
    captioned = deck.slides.add_slide(deck.slide_layouts[8])  # Picture with Caption
    assert captioned.shapes.title is not None
    captioned.shapes.title.text = "Captioned"
    for placeholder in captioned.placeholders:
        if placeholder.placeholder_format.type == PP_PLACEHOLDER.PICTURE:
            placeholder.insert_picture(io.BytesIO(PNG))
    grouped = deck.slides.add_slide(deck.slide_layouts[6])
    group = grouped.shapes.add_group_shape()
    group.shapes.add_textbox(
        Inches(1), Inches(1), Inches(3), Inches(1)
    ).text_frame.text = "grouped text"
    group.shapes.add_picture(io.BytesIO(PNG), Inches(1), Inches(2))
    out = io.BytesIO()
    deck.save(out)

    sections = PptxExtractor().extract(out.getvalue(), StagedAssets("raw/deck/assets"))
    assert [(s.kind, s.slide) for s in sections] == [
        ("text", 1),
        ("image", 1),
        ("text", 2),
        ("image", 2),
    ]
    assert sections[0].text == "Captioned" and sections[2].text == "grouped text"


def test_docx_content_controls_and_cell_pictures_are_kept() -> None:
    from docx import Document
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

    document = Document()
    document.add_paragraph("before")
    control = parse_xml(
        f"<w:sdt {nsdecls('w')}><w:sdtContent><w:p><w:r><w:t>inside content control"
        "</w:t></w:r></w:p></w:sdtContent></w:sdt>"
    )
    body = document.element.body
    body.insert(len(body) - 1, control)  # before the final sectPr
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "cell text"
    table.cell(0, 1).paragraphs[0].add_run().add_picture(io.BytesIO(PNG))
    document.add_paragraph("after")
    out = io.BytesIO()
    document.save(out)

    sections = DocxExtractor().extract(out.getvalue(), StagedAssets("raw/doc/assets"))
    assert [s.kind for s in sections] == ["text", "table", "image", "text"]
    assert sections[0].text == "before\n\ninside content control"
    assert sections[1].text.startswith("| cell text |")
    assert sections[3].text == "after"
