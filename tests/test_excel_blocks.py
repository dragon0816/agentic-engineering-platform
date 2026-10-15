"""`excel/read-sheet` and `excel/upsert-rows` on the practice workbook (W3)."""

from __future__ import annotations

import asyncio
import shutil
import zipfile
from pathlib import Path

import pytest
from openpyxl import load_workbook
from test_gateway import request

from capabilities.excel_blocks.handlers import (
    READ_SHEET_SPEC,
    UPSERT_SPEC,
    ReadSheetHandler,
    ReadSheetRequest,
    UpsertHandler,
    UpsertRequest,
)
from capabilities.runtime import CapabilityRefused
from evaluation.sop_practice import ORDERS, SHEET_NAME, write_workbook

KEY = "訂單號"


def web_rows(customer: str = "C001") -> list[dict[str, str]]:
    """The rows `web/read` returns for the practice site's table."""
    return [
        {
            "訂單號": o.order,
            "品項": o.item,
            "數量": str(o.quantity),
            "預計出貨日": o.ship_date,
            "狀態": o.status,
            "通知": "標記已通知" if o.status == "延遲" else "",
        }
        for o in ORDERS
        if o.customer == customer
    ]


def ask(book: Path, **changes: object) -> UpsertRequest:
    return UpsertRequest.model_validate(
        {
            "workbook": str(book),
            "sheet": SHEET_NAME,
            "key": KEY,
            "rows": web_rows(),
            "columns": ["訂單號", "品項", "數量", "預計出貨日", "狀態"],
            **changes,
        }
    )


def sheet_rows(book: Path) -> list[tuple[object, ...]]:
    return list(load_workbook(book)[SHEET_NAME].iter_rows(values_only=True))


def test_preview_says_what_would_change_and_writes_nothing(tmp_path: Path) -> None:
    book = write_workbook(tmp_path / "book.xlsx")
    before = book.read_bytes()
    plan = UpsertHandler(tmp_path / "backups").preview(ask(book))
    assert plan.added == ("SO-1002", "SO-1003", "SO-1004") and not plan.written
    (change,) = plan.updated
    assert change.key == "SO-1001" and change.fields == {"狀態": ("備料中", "準時")}
    assert book.read_bytes() == before and not (tmp_path / "backups").exists()


def test_upsert_updates_appends_keeps_notes_and_other_rows_and_backs_up(tmp_path: Path) -> None:
    book = write_workbook(tmp_path / "book.xlsx")
    result = asyncio.run(UpsertHandler(tmp_path / "backups")(request("run"), ask(book)))
    assert result.written and result.backup is not None  # type: ignore[attr-defined]
    assert Path(result.backup).is_file()  # type: ignore[attr-defined]
    rows = sheet_rows(book)
    by_key = {row[0]: row for row in rows[1:]}
    assert by_key["SO-1001"] == (
        "SO-1001",
        "CMW500 校正套件",
        2,
        "2026-10-08",
        "準時",
        "客戶催過一次",
    )
    assert by_key["SO-0999"][4] == "已出貨"  # another customer's row stays
    assert by_key["SO-1004"] == ("SO-1004", "電源模組", 4, "2026-10-01", "延遲", None)
    assert len(rows) == 6
    again = asyncio.run(UpsertHandler(tmp_path / "backups")(request("run"), ask(book)))
    assert again.unchanged == 4 and not again.written  # type: ignore[attr-defined]


def test_read_sheet_returns_rows_by_header(tmp_path: Path) -> None:
    book = write_workbook(tmp_path / "book.xlsx")
    rows = asyncio.run(
        ReadSheetHandler()(request("run"), ReadSheetRequest(workbook=str(book), sheet=SHEET_NAME))
    )
    assert rows.rows[0]["訂單號"] == "SO-1001" and rows.rows[0]["備註"] == "客戶催過一次"  # type: ignore[attr-defined]


def test_unsafe_workbooks_are_refused(tmp_path: Path) -> None:
    book = write_workbook(tmp_path / "book.xlsx")
    handler = UpsertHandler(tmp_path / "backups")
    cases = []
    (tmp_path / "~$book.xlsx").write_text("lock")
    cases.append((ask(book), "excel_workbook_open"))
    drawn = tmp_path / "drawn.xlsx"
    shutil.copy(book, drawn)
    with zipfile.ZipFile(drawn, "a") as package:
        package.writestr("xl/drawings/drawing1.xml", "<xdr:wsDr/>")
    cases.append((ask(drawn), "excel_workbook_has_drawings"))
    cases.append(
        (
            ask(book, sheet="Nope", workbook=str(drawn).replace("drawn", "missing")),
            "excel_workbook_missing",
        )
    )
    cases.append(
        (ask(write_workbook(tmp_path / "b2.xlsx"), key="客戶"), "excel_key_column_missing")
    )
    for item, code in cases:
        with pytest.raises(CapabilityRefused) as refused:
            asyncio.run(handler(request("run"), item))
        assert refused.value.code == code


def test_the_specs_classify_reading_and_writing() -> None:
    assert READ_SHEET_SPEC.side_effect == "read"
    assert UPSERT_SPEC.side_effect == "write" and UPSERT_SPEC.policy.approval_required
