"""Office extraction: PDF, PPTX and DOCX originals into Raw sections.

Each extractor keeps the relationship the Roadmap asks for — text, tables and
images with the page or slide they came from — and hands image bytes to the
intake's `AssetSink` rather than writing anything. The libraries are the
optional extra `office`; nothing else in the platform imports them, and every
import here is lazy so a host without the extra still loads this module.

Two rules shape the error handling. A corrupt or unreadable *original* is the
closed status `undecodable`, mapped from the parsers' documented error types
with the cause chained, never from a bare `Exception`. An unreadable *image*
inside a readable original is skipped, not fatal: losing one picture is better
than losing every page's text behind a status that blames the file.
"""

import io
import zipfile
from collections.abc import Iterable, Iterator, Sequence
from pathlib import PurePosixPath
from typing import Any

from knowledge.raw import AssetSink, ExtractionError, RawSection

# What the parsers raise on corrupt input, beyond their own base classes:
# broken streams surface as these from the standard library and lxml.
_CORRUPT = (ValueError, TypeError, KeyError, IndexError, OSError, zipfile.BadZipFile)


def table_markdown(rows: Sequence[Sequence[str]]) -> str:
    """A GitHub-style table; the first row is the header. Cells are flattened
    to one line and pipes escaped, so the table stays a table in Markdown."""
    if not rows:
        return ""
    width = max(len(row) for row in rows)

    def cell(text: str) -> str:
        return " ".join(text.split()).replace("|", "\\|")

    padded = [[cell(text) for text in row] + [""] * (width - len(row)) for row in rows]
    lines = ["| " + " | ".join(padded[0]) + " |", "|" + " --- |" * width]
    lines += ["| " + " | ".join(row) + " |" for row in padded[1:]]
    return "\n".join(lines)


class _Sections:
    """Gathers consecutive text into one section and flushes it before a
    table or an image, so reading order is kept."""

    def __init__(self) -> None:
        self.items: list[RawSection] = []
        self._texts: list[str] = []

    def text(self, value: str) -> None:
        if value.strip():
            self._texts.append(value)

    def flush(self, **place: int | None) -> None:
        body = "\n\n".join(self._texts)
        self._texts = []
        if body.strip():
            self.items.append(RawSection(kind="text", text=body, **place))

    def table(self, rows: Sequence[Sequence[str]], **place: int | None) -> None:
        self.flush(**place)
        markdown = table_markdown(rows)
        if markdown:
            self.items.append(RawSection(kind="table", text=markdown, **place))

    def image(self, ref: str, **place: int | None) -> None:
        self.flush(**place)
        self.items.append(RawSection(kind="image", image_ref=ref, **place))

    def done(self, **place: int | None) -> tuple[RawSection, ...]:
        self.flush(**place)
        return tuple(self.items)


def _undecodable(error: BaseException) -> ExtractionError:
    failure = ExtractionError("undecodable")
    failure.__cause__ = error
    return failure


class PdfExtractor:
    """Text per page and embedded images per page. PDF has no table model, so
    tables arrive as text in reading order; that is a known limitation."""

    name = "pdf.v1"
    suffixes = (".pdf",)

    def extract(self, data: bytes, assets: AssetSink) -> tuple[RawSection, ...]:
        from pypdf import PdfReader
        from pypdf.errors import DependencyError, PyPdfError

        out = _Sections()
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted and not reader.decrypt(""):
                raise ExtractionError("undecodable")
            for number, page in enumerate(reader.pages, 1):
                out.text(page.extract_text() or "")
                out.flush(page=number)
                for key in list(page.images.keys()):
                    try:
                        # Decoding happens on access; one image with a filter
                        # pypdf cannot decode (or no Pillow) must not cost the
                        # page's text, which is still evidence.
                        image = page.images[key]
                        blob = image.data
                    except (PyPdfError, DependencyError, ImportError, *_CORRUPT):
                        # DependencyError (a missing external decoder) is not
                        # a PyPdfError in pypdf's hierarchy.
                        continue
                    suffix = PurePosixPath(image.name).suffix or ".png"
                    out.image(assets.put(blob, suffix), page=number)
        except (PyPdfError, *_CORRUPT) as error:
            raise _undecodable(error) from error
        return out.done()


class PptxExtractor:
    """Per slide, in shape order and into groups: text frames gathered into
    one text section, tables as Markdown, pictures (placeholders included) as
    image sections. Speaker notes are not extracted."""

    name = "pptx.v1"
    suffixes = (".pptx",)

    def extract(self, data: bytes, assets: AssetSink) -> tuple[RawSection, ...]:
        from lxml.etree import LxmlError
        from pptx import Presentation
        from pptx.exc import PythonPptxError
        from pptx.shapes.group import GroupShape
        from pptx.shapes.picture import Picture

        out = _Sections()

        def walk(shapes: Iterable[Any], number: int) -> None:
            for shape in shapes:
                # Type checks, not `shape_type`: an element python-pptx does not
                # model raises from that property, and a picture in a placeholder
                # is still a Picture.
                if isinstance(shape, GroupShape):
                    walk(shape.shapes, number)
                elif isinstance(shape, Picture):
                    try:
                        blob, ext = shape.image.blob, shape.image.ext
                    except (PythonPptxError, *_CORRUPT):
                        continue  # a linked, non-embedded picture has no bytes
                    out.image(assets.put(blob, ext), slide=number)
                elif shape.has_table:
                    rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                    out.table(rows, slide=number)
                elif shape.has_text_frame:
                    out.text(shape.text_frame.text)

        try:
            deck = Presentation(io.BytesIO(data))
            for number, slide in enumerate(deck.slides, 1):
                walk(slide.shapes, number)
                out.flush(slide=number)
        except (PythonPptxError, LxmlError, *_CORRUPT) as error:
            raise _undecodable(error) from error
        return out.done()


class DocxExtractor:
    """Body blocks in document order, content controls (`w:sdt`) included:
    paragraphs gathered into text sections (headings keep their level as
    Markdown `#`), tables as Markdown with the pictures inside their cells
    following them, inline pictures as image sections. Word has no fixed pages,
    so sections carry no page number."""

    name = "docx.v1"
    suffixes = (".docx",)

    def extract(self, data: bytes, assets: AssetSink) -> tuple[RawSection, ...]:
        from docx import Document
        from docx.opc.exceptions import OpcError
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        from lxml.etree import LxmlError

        out = _Sections()
        paragraph_tag, table_tag, control_tag = qn("w:p"), qn("w:tbl"), qn("w:sdt")

        def blocks(element: Any) -> Iterator[Any]:
            for child in element.iterchildren():
                if child.tag in (paragraph_tag, table_tag):
                    yield child
                elif child.tag == control_tag:
                    for content in child.iterchildren(qn("w:sdtContent")):
                        yield from blocks(content)

        def pictures(element: Any) -> Iterator[str]:
            for blip in element.iter(qn("a:blip")):
                relationship = blip.get(qn("r:embed"))
                if not relationship:
                    continue
                try:
                    part = document.part.related_parts[relationship]
                except (OpcError, *_CORRUPT):
                    continue
                suffix = PurePosixPath(str(part.partname)).suffix or ".png"
                yield assets.put(part.blob, suffix)

        def paragraph(element: Any) -> None:
            item = Paragraph(element, document)
            style = item.style.name if item.style is not None and item.style.name else ""
            out.text(_heading(style, item.text))
            for ref in pictures(element):
                out.image(ref)

        def table(element: Any) -> None:
            item = Table(element, document)
            out.table([[cell.text for cell in row.cells] for row in item.rows])
            for ref in pictures(element):
                out.image(ref)

        try:
            document = Document(io.BytesIO(data))
            for block in blocks(document.element.body):
                if block.tag == table_tag:
                    table(block)
                else:
                    paragraph(block)
        except (OpcError, LxmlError, *_CORRUPT) as error:
            raise _undecodable(error) from error
        return out.done()


def _heading(style: str, text: str) -> str:
    if style.startswith("Heading ") and style[8:].isdigit() and text.strip():
        return "#" * min(int(style[8:]), 6) + " " + text.strip()
    return text


def office_extractors() -> tuple[PdfExtractor, PptxExtractor, DocxExtractor]:
    return (PdfExtractor(), PptxExtractor(), DocxExtractor())
