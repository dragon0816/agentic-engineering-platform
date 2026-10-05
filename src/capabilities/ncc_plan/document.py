"""Edit the example plan into this station's plan: Workflow 14's `render` step.

Ported from rs_workflow_system@896046e `host-bridge/jobs/ncc_improvement_plan.py`
(the "docx editing" helpers and the `render` part of `run`). python-docx comes
from AEP's `office` extra and is imported lazily.
"""

from __future__ import annotations

import copy
from collections.abc import Sequence
from pathlib import Path
from typing import Any

from capabilities.ncc_plan.contracts import NccDocument

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


class DocumentError(Exception):
    """A stable code; the template's text never travels in it."""


def set_text(paragraph: Any, text: str) -> None:
    """Replace a paragraph's words, keeping the first run's formatting."""
    runs = paragraph.runs
    if not runs:
        paragraph.add_run(text)
        return
    runs[0].text = text
    for run in runs[1:]:
        run._element.getparent().remove(run._element)


def replace_prefixed(doc: Any, prefix: str, text: str, warnings: list[str]) -> bool:
    """Rewrite the first paragraph starting with `prefix`. Reports if absent."""
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            set_text(paragraph, text)
            return True
    warnings.append(f"template has no paragraph starting {prefix!r}; left as it was")
    return False


def replace_numbered_block(
    doc: Any, after_prefix: str, lines: Sequence[str], warnings: list[str]
) -> bool:
    """Replace the （二）（三）… paragraphs following `after_prefix` with `lines`."""
    from docx.text.paragraph import Paragraph

    paragraphs = doc.paragraphs
    anchor_i = next((i for i, para in enumerate(paragraphs)
                     if para.text.strip().startswith(after_prefix)), None)
    if anchor_i is None:
        warnings.append(f"template has no paragraph starting {after_prefix!r}; section left")
        return False
    victims = []
    for para in paragraphs[anchor_i + 1:]:
        text = para.text.strip()
        if not text:
            continue
        if text.startswith("（") and "）" in text[:4]:
            victims.append(para)
            continue
        break   # the next heading ends the section
    if not victims:
        warnings.append(f"nothing numbered after {after_prefix!r}; section left alone")
        return False
    model = victims[0]._p
    previous = paragraphs[anchor_i]._p
    for line in lines:
        clone = copy.deepcopy(model)
        previous.addnext(clone)
        previous = clone
        set_text(Paragraph(clone, paragraphs[anchor_i]._parent), line)
    for para in victims:
        para._p.getparent().remove(para._p)
    return True


def resize_table(table: Any, wanted_rows: int) -> None:
    """Grow or shrink a table's body to `wanted_rows`, header excluded."""
    body = len(table.rows) - 1
    while body < wanted_rows:
        table._tbl.append(copy.deepcopy(table.rows[-1]._tr))
        body += 1
    while body > wanted_rows:
        table._tbl.remove(table.rows[-1]._tr)
        body -= 1


def fill_table(table: Any, rows: Sequence[Sequence[str]]) -> None:
    """Write `rows` into the table below its header row."""
    resize_table(table, len(rows))
    for row_cells, target in zip(rows, table.rows[1:], strict=False):
        for value, cell in zip(row_cells, target.cells, strict=False):
            paragraphs = cell.paragraphs
            set_text(paragraphs[0], str(value))
            for extra in paragraphs[1:]:
                extra._element.getparent().remove(extra._element)


def set_fault_captions(doc: Any, captions: Sequence[str], warnings: list[str]) -> None:
    """One single-column block per channel, captioned, with the photo removed."""
    blocks = [t for t in doc.tables if len(t.columns) == 1]
    if not blocks:
        warnings.append("template has no single-column photo blocks")
        return
    model = blocks[0]._tbl
    while len(blocks) < len(captions):
        clone = copy.deepcopy(model)
        blocks[-1]._tbl.addnext(clone)
        blocks = [t for t in doc.tables if len(t.columns) == 1]
    for extra in blocks[len(captions):]:
        extra._tbl.getparent().remove(extra._tbl)
    blocks = [t for t in doc.tables if len(t.columns) == 1]
    for table, caption in zip(blocks, captions, strict=False):
        # A photograph of the template's station would be a lie about this one.
        for cell in (c for row in table.rows for c in row.cells):
            for tag in ("drawing", "pict"):
                for element in list(cell._tc.iter(_W + tag)):
                    parent = element.getparent()
                    if parent is not None:
                        parent.remove(element)
        target = table.rows[-1].cells[0]
        set_text(target.paragraphs[0], caption)
        for extra in target.paragraphs[1:]:
            extra._element.getparent().remove(extra._element)


def render(template: Path, station: str, plan: NccDocument, output: Path) -> list[str]:
    """Write the plan; returns warnings about template sections left alone."""
    from docx import Document

    warnings: list[str] = []
    doc = Document(str(template))
    replace_prefixed(doc, "標的區域所在機關", plan.authority_line, warnings)
    # The title carries the station too, and it is a tab-separated line.
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith("標的區域") and station not in paragraph.text:
            set_text(paragraph, plan.title_line)
            break
    if plan.contractor_line:
        replace_prefixed(doc, "履約廠商", plan.contractor_line, warnings)
    replace_prefixed(doc, "（一）數位無線電視為民眾提供", plan.channel_sentence, warnings)
    replace_prefixed(doc, "1、數位無線電視服務提供起始年度", plan.start_year_line, warnings)
    replace_prefixed(doc, "2、服務區域", plan.area_line, warnings)
    if plan.confirm_line:
        replace_prefixed(doc, "（一）標的區域仍有機關", plan.confirm_line, warnings)

    wide = [t for t in doc.tables if len(t.columns) >= 3]
    if len(wide) < 2:
        raise DocumentError("template_tables_missing")
    fill_table(wide[0], plan.system_rows)
    fill_table(wide[-1], plan.service_rows)
    if plan.plan_lines:
        replace_numbered_block(doc, "（一）標的區域仍有機關", plan.plan_lines, warnings)
    else:
        warnings.append("no planLines configured; 三、 left as the template had it")
    if plan.fault_intro:
        replace_prefixed(doc, "（一）5W", plan.fault_intro, warnings)
        replace_prefixed(doc, "（一）25W", plan.fault_intro, warnings)
    if plan.fault_captions:
        set_fault_captions(doc, plan.fault_captions, warnings)
    else:
        warnings.append("no faultLines configured; the photo blocks were left alone")

    try:
        doc.save(str(output))
    except PermissionError:
        # Word holds an exclusive lock on an open document; reviewing the last
        # run's output and running again is the normal way to use this.
        raise DocumentError("document_open_in_word") from None
    return warnings
