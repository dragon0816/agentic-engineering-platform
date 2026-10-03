"""Read-only Workflow 10 adapter around the ported source transformation."""

import asyncio
import datetime as dt
import hashlib
import json
import re
from collections.abc import Callable, Mapping, Sequence
from pathlib import Path
from typing import Any

from capabilities.chipset_report.contracts import (
    ChipsetPreview,
    ChipsetPreviewInput,
    ChipsetPreviewStepInput,
    ChipsetReportSettings,
    ChipsetSourceChange,
    ChipsetSourceSummary,
    DroppedChipsetEvidence,
)
from capabilities.chipset_report.rules import (
    TRACE_HEADERS,
    SourceRow,
    row_colour,
    transform,
    watched_but_missing,
)
from capabilities.chipset_report.ruleset import RulesetError, compile_rules, load_ruleset
from capabilities.chipset_report.source_snapshot import prepare_diff
from capabilities.runtime import CapabilityRefused
from common.base import Contract
from common.execution import RequestContext
from integrations import excel
from integrations.excel import WorkbookError

DEFAULT_SOURCE_SHEETS = ("Project List_Sales update", "Project List")


class ChipsetPreviewError(Exception):
    """A stable refusal code that never includes company paths or cell data."""


def _read_table(
    path: str,
    sheet: str | None,
    *,
    header_row: int,
    max_rows: int,
    source: bool,
) -> tuple[str, tuple[str, ...], tuple[dict[str, Any], ...]]:
    try:
        import openpyxl
    except ImportError:
        raise ChipsetPreviewError("library_missing") from None
    if not Path(path).is_file():
        raise ChipsetPreviewError("workbook_missing")
    try:
        book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - corrupt and unsupported workbook types
        raise ChipsetPreviewError("workbook_unreadable") from None
    try:
        chosen = sheet
        if chosen is None and source:
            chosen = next((name for name in DEFAULT_SOURCE_SHEETS if name in book.sheetnames), None)
        if chosen is None or chosen not in book.sheetnames:
            raise ChipsetPreviewError("sheet_missing")
        rows = book[chosen].iter_rows(values_only=True)
        for _ in range(header_row - 1):
            next(rows, ())
        raw_headers = next(rows, ())
        headers = tuple("" if cell is None else str(cell) for cell in raw_headers)
        records: list[dict[str, Any]] = []
        for index, row in enumerate(rows):
            if index >= max_rows:
                break
            if all(value is None or isinstance(value, str) and not value.strip() for value in row):
                records.append({})  # Retain source row numbers exactly.
            else:
                records.append(
                    {name: value for name, value in zip(headers, row, strict=False) if name}
                )
        return chosen, tuple(name for name in headers if name), tuple(records)
    finally:
        book.close()


def _fiscal_year(path: str) -> str:
    match = re.search(r"FY\s*(\d{4})", path, re.IGNORECASE)
    return f"FY{match.group(1)}" if match else Path(path).stem


def _digest(headers: Sequence[str], records: Sequence[Mapping[str, Any]]) -> str:
    payload = json.dumps([headers, records], ensure_ascii=False, default=str)
    return hashlib.sha256(payload.encode("utf-8")).hexdigest()


def build_preview(
    settings: ChipsetReportSettings,
    request: ChipsetPreviewInput | None = None,
    *,
    today: dt.date,
    state_dir: Path | None = None,
) -> ChipsetPreview:
    """Mirror W1's read/filter/transform order without touching the target."""
    settings = ChipsetReportSettings.model_validate(settings)
    request = ChipsetPreviewInput.model_validate(request or ChipsetPreviewInput())
    max_rows = min(settings.max_rows, request.max_rows or settings.max_rows)
    compiled = compile_rules(load_ruleset(settings.rules_file))
    source_rows: list[SourceRow] = []
    summaries: list[ChipsetSourceSummary] = []
    source_changes: list[ChipsetSourceChange] = []
    for index, source in enumerate(settings.source_files):
        if not source.enabled:
            summaries.append(ChipsetSourceSummary(source_index=index, skipped=True))
            continue
        sheet, _, records = _read_table(
            source.path, source.sheet, header_row=1, max_rows=max_rows, source=True
        )
        fiscal_year = source.fiscal_year or _fiscal_year(source.path)
        used = 0
        rows_here: list[SourceRow] = []
        for row_number, record in enumerate(records, start=2):
            if not record:
                continue
            row = SourceRow.from_record(record, fiscal_year, sheet, row_number=row_number)
            if not any((row.brand, row.mfg, row.project_name, row.chipset_models)):
                continue
            source_rows.append(row)
            rows_here.append(row)
            used += 1
        summaries.append(
            ChipsetSourceSummary(
                source_index=index, sheet=sheet, fiscal_year=fiscal_year, rows=used
            )
        )
        if state_dir is not None:
            difference, candidate = prepare_diff(state_dir, source.path, rows_here)
            source_changes.append(
                ChipsetSourceChange(
                    source_index=index,
                    summary=difference.summary(),
                    critical_count=len(difference.critical_changes),
                    diff=difference.as_dict(),
                    candidate=candidate,
                )
            )
    target_sheet, headers, existing = _read_table(
        settings.target_file,
        settings.target_sheet,
        header_row=settings.target_header_row,
        max_rows=settings.max_rows,
        source=False,
    )
    if not headers:
        # The source has fallback headers for explicit unmatched previews. A
        # configured real target with no headers is instead a configuration error.
        raise ChipsetPreviewError("target_headers_missing")
    temp_present = settings.temp_sheet in excel.sheet_names(Path(settings.target_file))
    temp_digest: str | None = None
    if temp_present:
        _, temp_headers, temp_rows = _read_table(
            settings.target_file,
            settings.temp_sheet,
            header_row=1,
            max_rows=settings.max_rows,
            source=False,
        )
        temp_digest = _digest(temp_headers, temp_rows)
    result = transform(
        source_rows,
        existing,
        headers,
        compiled,
        today=today,
        include_statuses=settings.include_statuses,
        exclude_statuses=settings.exclude_statuses,
        target_header_row=settings.target_header_row,
    )
    out_headers = (*headers, *(TRACE_HEADERS if settings.write_trace_columns else ()))
    out_records = tuple(
        {header: record.get(header, "") for header in out_headers} for record in result.records
    )
    return ChipsetPreview(
        target_sheet=target_sheet,
        temp_sheet=settings.temp_sheet,
        target_digest=_digest(headers, existing),
        temp_digest=temp_digest,
        temp_present=temp_present,
        managed_headers=headers,
        headers=out_headers,
        records=out_records,
        format_records=result.records,
        row_fills=tuple(row_colour(record, compiled) for record in result.records),
        dropped=tuple(
            DroppedChipsetEvidence(
                project=entry.project,
                fiscal_year=entry.fiscal_year,
                row_number=entry.row_number,
                raw=entry.raw,
                parsed=entry.parsed,
                reason=entry.reason,
            )
            for entry in result.dropped
        ),
        sources=tuple(summaries),
        source_changes=tuple(source_changes),
        projects_read=result.projects_read,
        projects_included=result.projects_included,
        rows_generated=result.rows_generated,
        chipsets_extracted=result.chipsets_extracted,
        unresolved_technology=result.unresolved_technology,
        match_counts=dict(result.match_counts),
        dropped_count=len(result.dropped),
        watched_missing=tuple(watched_but_missing(headers, compiled)),
    )


class ChipsetPreviewHandler:
    def __init__(
        self,
        settings: ChipsetReportSettings,
        *,
        today: Callable[[], dt.date] = dt.date.today,
        state_dir: Path | None = None,
    ) -> None:
        self.settings = ChipsetReportSettings.model_validate(settings)
        self.today = today
        self.state_dir = state_dir

    async def __call__(self, context: RequestContext, inputs: Contract) -> Contract:
        request = ChipsetPreviewStepInput.model_validate(inputs).request
        try:
            return await asyncio.to_thread(
                build_preview,
                self.settings,
                request,
                today=self.today(),
                state_dir=self.state_dir,
            )
        except (ChipsetPreviewError, RulesetError, WorkbookError) as error:
            raise CapabilityRefused(str(error)) from None
