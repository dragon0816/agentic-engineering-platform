"""`excel/read-sheet` and `excel/upsert-rows`: Excel building blocks (SOP_TO_WORKFLOW W3).

Read one sheet as rows keyed by the header row, and upsert rows by a key
column: existing rows change only in the named columns, new keys are
appended, nothing is ever deleted and other columns (a person's notes) are
left alone. `preview` computes exactly what an upsert would add and change
without writing, which is what a Workflow preview run shows.

openpyxl rewrites the whole file and drops what it does not model (charts,
pictures, shapes), so a workbook holding drawings is refused rather than
damaged. A workbook open in Excel is refused too (its `~$` lock file is
there). Before writing, the workbook is copied to the host's backup folder.
"""

from __future__ import annotations

import asyncio
import re
import shutil
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any

from pydantic import Field, JsonValue

from capabilities.contracts import CapabilitySpec
from capabilities.runtime import CapabilityRefused
from common.base import Contract, Text
from common.execution import RequestContext
from integrations import excel

READ_SHEET_SPEC = CapabilitySpec.model_validate({
    "identity": {"namespace": "excel", "name": "read-sheet", "version": "1.0.0"},
    "name": "excel.read_sheet",
    "description": "Read one workbook sheet as rows keyed by its header row",
    "input_contract": "excel.read-sheet.input.v1",
    "output_contract": "excel.rows.v1",
    "side_effect": "read",
    "policy": {"required_permissions": ["excel.read"], "policy_refs": ["excel-policy"]},
})

UPSERT_SPEC = CapabilitySpec.model_validate({
    "identity": {"namespace": "excel", "name": "upsert-rows", "version": "1.0.0"},
    "name": "excel.upsert_rows",
    "description": (
        "Update rows by a key column and append new keys; never deletes rows or "
        "touches other columns"
    ),
    "input_contract": "excel.upsert-rows.input.v1",
    "output_contract": "excel.upsert-rows.output.v1",
    "side_effect": "write",
    "policy": {
        "risk": "low",
        "approval_required": True,
        "required_permissions": ["excel.write"],
        "policy_refs": ["excel-policy"],
    },
})

_INTEGER = re.compile(r"^-?\d{1,15}$")
_DECIMAL = re.compile(r"^-?\d{1,15}\.\d{1,10}$")


class ReadSheetRequest(Contract):
    workbook: Text
    sheet: Text


class SheetRows(Contract):
    headers: tuple[Text, ...]
    rows: tuple[dict[Text, JsonValue], ...]


class UpsertRequest(Contract):
    workbook: Text
    sheet: Text
    key: Text
    rows: tuple[dict[Text, JsonValue], ...] = Field(max_length=5000)
    # Columns to write; empty means every sheet column the rows carry.
    columns: tuple[Text, ...] = ()


class RowChange(Contract):
    key: Text
    fields: dict[Text, tuple[Text, Text]] = Field(default_factory=dict)


class UpsertResult(Contract):
    added: tuple[Text, ...]
    updated: tuple[RowChange, ...]
    unchanged: int
    written: bool
    backup: Text | None = None


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _typed(text: str) -> int | float | str:
    if _INTEGER.match(text):
        return int(text)
    if _DECIMAL.match(text):
        return float(text)
    return text


def _refuse(error: excel.WorkbookError) -> CapabilityRefused:
    return CapabilityRefused(f"excel_{error.code}")


class ReadSheetHandler:
    async def __call__(self, context: RequestContext, inputs: Contract) -> Contract:
        request = ReadSheetRequest.model_validate(inputs)
        try:
            headers, rows = await asyncio.to_thread(
                excel.read_rows, Path(request.workbook), request.sheet)
        except excel.WorkbookError as error:
            raise _refuse(error) from None
        return SheetRows(
            headers=headers,
            rows=tuple({h: cell for h, cell in zip(headers, row, strict=False) if h}
                       for row in rows),
        )


class UpsertHandler:
    def __init__(self, backups: Path) -> None:
        self.backups = backups

    async def __call__(self, context: RequestContext, inputs: Contract) -> Contract:
        request = UpsertRequest.model_validate(inputs)
        return await asyncio.to_thread(self.upsert, request, write=True)

    def preview(self, inputs: Contract) -> UpsertResult:
        """What `upsert` would do, computed from the workbook; nothing written."""
        return self.upsert(UpsertRequest.model_validate(inputs), write=False)

    def upsert(self, request: UpsertRequest, *, write: bool) -> UpsertResult:
        path = Path(request.workbook)
        if path.suffix.lower() != ".xlsx" or not path.is_file():
            raise CapabilityRefused("excel_workbook_missing")
        if write and (path.parent / f"~${path.name}").exists():
            raise CapabilityRefused("excel_workbook_open")
        with zipfile.ZipFile(path) as package:
            if any(name.startswith("xl/drawings/") for name in package.namelist()):
                raise CapabilityRefused("excel_workbook_has_drawings")
        import openpyxl

        book = openpyxl.load_workbook(path)
        if request.sheet not in book.sheetnames:
            raise CapabilityRefused("excel_sheet_missing")
        sheet = book[request.sheet]
        headers = [_cell_text(cell.value) for cell in sheet[1]]
        if request.key not in headers:
            raise CapabilityRefused("excel_key_column_missing")
        columns = request.columns or tuple(
            name for name in headers if name and any(name in row for row in request.rows))
        missing = [name for name in columns if name not in headers]
        if missing:
            raise CapabilityRefused("excel_column_missing")
        index = {name: position + 1 for position, name in enumerate(headers) if name}
        key_column = index[request.key]
        existing = {
            _cell_text(sheet.cell(row=number, column=key_column).value): number
            for number in range(2, sheet.max_row + 1)
            if _cell_text(sheet.cell(row=number, column=key_column).value)
        }
        added: list[str] = []
        updated: list[RowChange] = []
        unchanged = 0
        next_row = sheet.max_row + 1
        for row in request.rows:
            key = _cell_text(row.get(request.key))
            if not key:
                raise CapabilityRefused("excel_row_without_key")
            if key in existing:
                number = existing[key]
                changes: dict[str, tuple[str, str]] = {}
                for name in columns:
                    if name == request.key or name not in row:
                        continue
                    cell = sheet.cell(row=number, column=index[name])
                    before, after = _cell_text(cell.value), _cell_text(row[name])
                    if before != after:
                        changes[name] = (before, after)
                        cell.value = _typed(after)
                if changes:
                    updated.append(RowChange(key=key, fields=changes))
                else:
                    unchanged += 1
            else:
                for name in columns:
                    if name in row:
                        sheet.cell(row=next_row, column=index[name]).value = _typed(
                            _cell_text(row[name]))
                existing[key] = next_row
                next_row += 1
                added.append(key)
        backup: str | None = None
        changed = bool(added or updated)
        if write and changed:
            self.backups.mkdir(parents=True, exist_ok=True)
            target = self.backups / f"{path.stem}.{datetime.now():%Y%m%d-%H%M%S}{path.suffix}"
            shutil.copy2(path, target)
            backup = str(target)
            book.save(path)
        book.close()
        return UpsertResult(added=tuple(added), updated=tuple(updated), unchanged=unchanged,
                            written=write and changed, backup=backup)
